from __future__ import annotations

import re
from collections import defaultdict
from io import BytesIO

from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from sklad.models import (
    WarehouseContainer,
    WarehouseEvent,
    WarehouseOperation,
    WarehouseReserve,
    WarehouseStockSnapshot,
)


_OS_LINE_DISPLAY_LABELS = {
    1: "0",
    2: "A",
    3: "B",
    4: "C",
    5: "D",
    6: "E",
    7: "F",
    8: "G",
    9: "I",
}

_EVENT_LABELS = {
    "receiving_arrived": "Поступил на приемку",
    "placement_started": "Начали размещение",
    "placement_completed": "Положили в короб",
    "putaway_requested": "Запросили размещение на склад",
    "putaway_completed": "Разместили на склад",
    "processing_requested": "Запросили обработку",
    "processing_reserved": "Зарезервировали под обработку",
    "processing_reserve_released": "Сняли резерв обработки",
    "processing_zone_arrived": "Приехал в обработку",
    "processing_started": "Начали обработку",
    "processing_completed": "Завершили обработку",
    "processing_consumed": "Списали в обработку",
    "shipping_reserved": "Зарезервировали под отгрузку",
    "shipping_reserve_released": "Сняли резерв отгрузки",
    "otg_requested": "Запросили перемещение в OTG",
    "otg_arrived": "Приехал в OTG",
    "palletization_started": "Начали паллетизацию",
    "palletization_completed": "Завершили паллетизацию",
    "ready_for_loading": "Готов к погрузке",
    "assigned_to_trip": "Назначили в рейс",
    "loading_started": "Начали погрузку",
    "loaded_to_vehicle": "Погрузили в машину",
    "shipped": "Отгрузили",
    "movement_requested": "Запросили перемещение",
    "movement_task_created": "Создали задачу перемещения",
    "movement_started": "Начали перемещение",
    "movement_completed": "Завершили перемещение",
    "movement_canceled": "Отменили перемещение",
    "stock_returned_to_storage": "Вернули на склад",
    "manual_stock_restore": "Ручное восстановление остатка",
    "manual_stock_released": "Ручное восстановление после списания",
    "warehouse_context_canceled": "Отменили складской контекст",
}

_STATE_LABELS = {
    "unknown": "Неизвестно",
    "received_unplaced": "Принят, не размещен",
    "placed_in_receiving": "В приемке",
    "stored": "На складе",
    "reserved_for_processing": "Резерв обработки",
    "moving_to_processing": "Едет в обработку",
    "in_processing_zone": "В обработке",
    "processing_in_progress": "Обрабатывается",
    "processing_consumed": "Списан в обработку",
    "placed_after_processing": "После обработки",
    "reserved_for_shipping": "Резерв отгрузки",
    "moving_to_otg": "Едет в OTG",
    "in_otg": "В OTG",
    "palletizing": "Паллетизация",
    "ready_for_loading": "Готов к погрузке",
    "assigned_to_trip": "В рейсе",
    "loading_in_progress": "Погрузка",
    "loaded_to_vehicle": "В машине",
    "shipped": "Отгружен",
    "partially_shipped": "Частично отгружен",
    "canceled": "Отменен",
}

_EVENT_FILL_COLORS = {
    "receiving_arrived": "E2F0D9",
    "placement_started": "C6E0B4",
    "placement_completed": "A9D18E",
    "putaway_requested": "D9EAF7",
    "putaway_completed": "BDD7EE",
    "processing_requested": "EADCF8",
    "processing_reserved": "FCE4D6",
    "processing_reserve_released": "F8CBAD",
    "processing_zone_arrived": "D9EAD3",
    "processing_started": "B6D7A8",
    "processing_completed": "93C47D",
    "processing_consumed": "D5A6BD",
    "shipping_reserved": "FFF2CC",
    "shipping_reserve_released": "FCE5CD",
    "otg_requested": "DDEBF7",
    "otg_arrived": "9FC5E8",
    "palletization_started": "D9D2E9",
    "palletization_completed": "B4A7D6",
    "ready_for_loading": "C9DAF8",
    "assigned_to_trip": "A4C2F4",
    "loading_started": "F4CCCC",
    "loaded_to_vehicle": "EA9999",
    "shipped": "E06666",
    "movement_requested": "E4DFEC",
    "movement_task_created": "CCC0DA",
    "movement_started": "B4A7D6",
    "movement_completed": "8E7CC3",
    "movement_canceled": "E7E6E6",
    "stock_returned_to_storage": "D9D9D9",
    "manual_stock_restore": "C6E0B4",
    "manual_stock_released": "C6E0B4",
    "warehouse_context_canceled": "BFBFBF",
}

_MOVEMENT_KIND_BY_EVENT_TYPE = {
    "receiving_arrived": "Приход",
    "placement_started": "На месте",
    "placement_completed": "На месте",
    "putaway_requested": "На месте",
    "putaway_completed": "На месте",
    "processing_requested": "На месте",
    "processing_reserved": "На месте",
    "processing_reserve_released": "На месте",
    "processing_zone_arrived": "На месте",
    "processing_started": "На месте",
    "processing_completed": "Приход",
    "processing_consumed": "Расход",
    "shipping_reserved": "На месте",
    "shipping_reserve_released": "На месте",
    "otg_requested": "На месте",
    "otg_arrived": "На месте",
    "palletization_started": "На месте",
    "palletization_completed": "На месте",
    "ready_for_loading": "На месте",
    "assigned_to_trip": "На месте",
    "loading_started": "На месте",
    "loaded_to_vehicle": "На месте",
    "shipped": "Расход",
    "movement_requested": "На месте",
    "movement_task_created": "На месте",
    "movement_started": "На месте",
    "movement_completed": "На месте",
    "movement_canceled": "На месте",
    "stock_returned_to_storage": "На месте",
    "manual_stock_restore": "Приход",
    "manual_stock_released": "Приход",
    "warehouse_context_canceled": "На месте",
}

_STOCK_RESTORE_CONTEXT_TYPES = {
    "manual_restore",
    "inventory_reconciliation",
}

_STOCK_RESTORE_SOURCE_DOCUMENT_TYPES = {
    "manual_restore",
    "processing_order_correction",
}

_HIDDEN_JOURNAL_EVENT_TYPES = {
    "assigned_to_trip",
}

_UNSTABLE_PARENT_PALLET_EVENT_TYPES = {
    "placement_completed",
}

_LOCATION_INSENSITIVE_DUPLICATE_EVENT_TYPES = {
    "otg_arrived",
    "processing_zone_arrived",
}

_ROLE_LABELS = {
    "driver": "Водитель",
    "head_manager": "Главный менеджер",
    "logistician": "Логист",
    "manager": "Менеджер",
    "processor": "Обработка",
    "reachtruck": "Ричтрак",
    "reachtruck_driver": "Ричтрак",
    "storekeeper": "Кладовщик",
}


def build_product_check_workbook(article: str) -> BytesIO:
    article = _clean_text(article)
    snapshots = _product_snapshots(article)
    reserves = _product_reserves(article)
    operations = _product_operations(snapshots, reserves)
    events = _product_events(article, snapshots, reserves, operations)

    workbook = Workbook()
    journal_sheet = workbook.active
    journal_sheet.title = "Журнал"
    _write_journal_sheet(journal_sheet, article, snapshots, events)
    stock_sheet = workbook.create_sheet("Текущие остатки")
    _write_stock_sheet(stock_sheet, article, snapshots)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def product_check_filename(article: str) -> str:
    safe_article = _safe_filename_part(article) or "article"
    return f"{safe_article}_{timezone.localdate().isoformat()}.xlsx"


def _product_snapshots(article: str) -> list[WarehouseStockSnapshot]:
    if not article:
        return []
    return list(
        WarehouseStockSnapshot.objects.filter(
            Q(sku_code__iexact=article) | Q(sku_ref__sku_code__iexact=article)
        )
        .select_related(
            "agency",
            "sku_ref",
            "location",
            "active_operation",
            "last_event",
        )
        .order_by("created_at", "id")
    )


def _product_reserves(article: str) -> list[WarehouseReserve]:
    if not article:
        return []
    return list(
        WarehouseReserve.objects.filter(
            Q(sku_code__iexact=article) | Q(sku_ref__sku_code__iexact=article)
        )
        .select_related("agency", "sku_ref")
        .order_by("created_at", "id")
    )


def _product_operations(
    snapshots: list[WarehouseStockSnapshot],
    reserves: list[WarehouseReserve],
) -> list[WarehouseOperation]:
    reserve_ids = {int(reserve.id or 0) for reserve in reserves if int(reserve.id or 0)}
    operation_ids = {
        int(snapshot.active_operation_id or 0)
        for snapshot in snapshots
        if int(snapshot.active_operation_id or 0)
    }
    if not reserve_ids and not operation_ids:
        return []
    query = Q()
    if reserve_ids:
        query |= Q(reserve_id__in=reserve_ids)
    if operation_ids:
        query |= Q(id__in=operation_ids)
    return list(
        WarehouseOperation.objects.filter(query)
        .select_related("agency", "reserve", "source_location", "destination_location")
        .order_by("created_at", "id")
    )


def _product_events(
    article: str,
    snapshots: list[WarehouseStockSnapshot],
    reserves: list[WarehouseReserve],
    operations: list[WarehouseOperation],
) -> list[WarehouseEvent]:
    container_ids = set()
    snapshot_ids = set()
    last_event_ids = set()
    source_context_pairs = set()
    for snapshot in snapshots:
        snapshot_ids.add(int(snapshot.id or 0))
        if int(snapshot.container_id or 0):
            container_ids.add(int(snapshot.container_id))
        if int(snapshot.parent_container_id or 0):
            container_ids.add(int(snapshot.parent_container_id))
        if int(snapshot.last_event_id or 0):
            last_event_ids.add(int(snapshot.last_event_id))
        context_type = _clean_text(snapshot.source_context_type)
        context_id = _clean_text(snapshot.source_context_id)
        if context_type and context_id:
            source_context_pairs.add((context_type, context_id))

    reserve_ids = {int(reserve.id or 0) for reserve in reserves if int(reserve.id or 0)}
    operation_ids = {int(operation.id or 0) for operation in operations if int(operation.id or 0)}

    query = Q(payload__sku_code=article) | Q(payload__sku=article)
    if snapshot_ids:
        query |= Q(payload__snapshot_id__in=list(snapshot_ids))
    if container_ids:
        query |= Q(container_id__in=container_ids)
    if reserve_ids:
        query |= Q(reserve_id__in=reserve_ids)
    if operation_ids:
        query |= Q(operation_id__in=operation_ids)
    if last_event_ids:
        query |= Q(id__in=last_event_ids)
    for context_type, context_id in source_context_pairs:
        query |= Q(
            event_type="receiving_arrived",
            stock_context_type=context_type,
            stock_context_id=context_id,
        )

    return list(
        WarehouseEvent.objects.filter(query)
        .select_related(
            "agency",
            "operation",
            "operation__reserve",
            "operation_task",
            "reserve",
            "reserve__sku_ref",
            "from_location",
            "to_location",
            "performed_by",
        )
        .order_by("occurred_at", "id")
    )


def _container_map_for_report(
    snapshots: list[WarehouseStockSnapshot],
    events: list[WarehouseEvent],
) -> dict[int, dict]:
    container_ids = set()
    for snapshot in snapshots:
        if int(snapshot.container_id or 0):
            container_ids.add(int(snapshot.container_id or 0))
        if int(snapshot.parent_container_id or 0):
            container_ids.add(int(snapshot.parent_container_id or 0))
    for event in events:
        if int(event.container_id or 0):
            container_ids.add(int(event.container_id or 0))
        task = event.operation_task
        if task is not None and int(task.container_id or 0):
            container_ids.add(int(task.container_id or 0))
    if not container_ids:
        return {}

    rows = list(
        WarehouseContainer.objects.filter(id__in=container_ids).values(
            "id",
            "container_type",
            "container_code",
            "parent_container_id",
        )
    )
    parent_ids = {
        int(row.get("parent_container_id") or 0)
        for row in rows
        if int(row.get("parent_container_id") or 0)
    }
    missing_parent_ids = parent_ids - {int(row.get("id") or 0) for row in rows}
    if missing_parent_ids:
        rows.extend(
            WarehouseContainer.objects.filter(id__in=missing_parent_ids).values(
                "id",
                "container_type",
                "container_code",
                "parent_container_id",
            )
        )
    return {int(row.get("id") or 0): row for row in rows if int(row.get("id") or 0)}


def _snapshot_row(snapshot: WarehouseStockSnapshot | None, container_map: dict[int, dict]) -> dict:
    if snapshot is None:
        return {}
    sku_ref = snapshot.sku_ref
    box_code, pallet_code = _container_codes(int(snapshot.container_id or 0), container_map)
    if not pallet_code and int(snapshot.parent_container_id or 0):
        _, pallet_code = _container_codes(int(snapshot.parent_container_id or 0), container_map)
        parent = container_map.get(int(snapshot.parent_container_id or 0))
        if parent and parent.get("container_type") in {WarehouseContainer.TYPE_PALLET, WarehouseContainer.TYPE_MIXED_PALLET}:
            pallet_code = _clean_text(parent.get("container_code"))
    if not box_code and not pallet_code and _clean_text(snapshot.container_code):
        pallet_code = _clean_text(snapshot.container_code)
    return {
        "sku": _clean_text(getattr(sku_ref, "sku_code", "") or snapshot.sku_code),
        "name": _clean_text(snapshot.name) or _clean_text(getattr(sku_ref, "name", "")),
        "warehouse_state_code": _clean_text(snapshot.warehouse_state_code),
        "qty": int(snapshot.qty or 0),
        "available_qty": int(snapshot.available_qty or 0),
        "processing_reserved_qty": int(snapshot.processing_reserved_qty or 0),
        "shipping_reserved_qty": int(snapshot.shipping_reserved_qty or 0),
        "box_code": box_code,
        "pallet_code": pallet_code,
        "location": _location_label(snapshot.location),
    }


def _product_header_info(
    article: str,
    snapshots: list[WarehouseStockSnapshot],
    events: list[WarehouseEvent],
) -> dict[str, str]:
    clients = []
    names = []
    for snapshot in snapshots:
        client = _agency_label(snapshot.agency)
        if client:
            clients.append(client)
        name = _clean_text(snapshot.name) or _clean_text(getattr(snapshot.sku_ref, "name", ""))
        if name:
            names.append(name)
    for event in events:
        client = _agency_label(event.agency)
        if client:
            clients.append(client)
        name = _event_product_name(event, {})
        if name:
            names.append(name)
    return {
        "article": article,
        "client": _unique_join(clients),
        "name": _unique_join(names),
        "generated_at": _datetime_text(timezone.now()),
    }


def _write_report_header(sheet, title: str, info: dict[str, str], last_column: str) -> int:
    sheet.append([title])
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Клиент", info.get("client") or ""])
    sheet.append(["Наименование", info.get("name") or ""])
    sheet.append(["Артикул", info.get("article") or ""])
    sheet.append(["Дата формирования", info.get("generated_at") or ""])
    sheet.append([])
    for row_number in range(2, 6):
        sheet.cell(row=row_number, column=1).font = Font(bold=True)
    return 7


def _write_journal_report_header(sheet, info: dict[str, str], totals: dict[str, int], last_column: str) -> int:
    sheet.append(["Проверка товара"])
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Клиент", info.get("client") or "", "", "Приход", int(totals.get("Приход") or 0)])
    sheet.append(["Наименование", info.get("name") or "", "", "Расход", int(totals.get("Расход") or 0)])
    sheet.append(["Артикул", info.get("article") or "", "", "Остаток", int(totals.get("Остаток") or 0)])
    sheet.append(["Дата формирования", info.get("generated_at") or ""])
    sheet.append([])
    for row_number in range(2, 6):
        sheet.cell(row=row_number, column=1).font = Font(bold=True)
        sheet.cell(row=row_number, column=4).font = Font(bold=True)
    return 7


def _write_journal_sheet(sheet, article: str, snapshots: list[WarehouseStockSnapshot], events: list[WarehouseEvent]) -> None:
    snapshots_by_container = _snapshots_by_container(snapshots)
    snapshots_by_context = _snapshots_by_source_context(snapshots)
    snapshots_by_last_event = _snapshots_by_last_event(snapshots)
    snapshots_by_id = _snapshots_by_id(snapshots)
    container_map = _container_map_for_report(snapshots, events)

    rows = []
    rows_by_key = {}
    totals = {"Приход": 0, "Расход": 0}
    for event in events:
        event_type = _clean_text(event.event_type)
        if event_type in _HIDDEN_JOURNAL_EVENT_TYPES:
            continue
        if _event_payload_article_mismatch(event, article):
            continue
        related_snapshots = _snapshots_for_event(
            event,
            snapshots_by_id=snapshots_by_id,
            snapshots_by_container=snapshots_by_container,
            snapshots_by_context=snapshots_by_context,
            snapshots_by_last_event=snapshots_by_last_event,
        )
        snapshot = related_snapshots[0] if related_snapshots else None
        row = _snapshot_row(snapshot, container_map) if snapshot else {}
        box_code, pallet_code = _journal_container_codes(event, row, container_map)
        qty = _event_qty_for_article(event, related_snapshots, article)
        movement_kind = _movement_kind_label(event)
        qty_value = int(qty or 0)
        action = _event_action_label(event)
        from_label = _location_label(event.from_location)
        to_label = _location_label(event.to_location)
        document = _document_label(event)
        note = _event_note(event)
        row_data = {
            "event_type": event_type,
            "movement_kind": movement_kind,
            "qty_value": qty_value,
            "cells": [
                "",
                _date_text(event.occurred_at),
                _time_text(event.occurred_at),
                action,
                qty_value if movement_kind == "Приход" and qty is not None else "",
                qty_value if movement_kind == "Расход" and qty is not None else "",
                qty_value if movement_kind == "На месте" and qty is not None else "",
                box_code,
                pallet_code,
                from_label,
                to_label,
                document,
                _state_label(row.get("warehouse_state_code") or ""),
                _user_label(event.performed_by, event.performed_by_role),
                note,
            ],
        }
        group_key = _journal_group_key(
            event,
            action=action,
            movement_kind=movement_kind,
            box_code=box_code,
            pallet_code=pallet_code,
            from_label=from_label,
            to_label=to_label,
            document=document,
        )
        existing = rows_by_key.get(group_key)
        if existing is not None:
            _merge_journal_row(existing, row_data)
            continue
        rows_by_key[group_key] = row_data
        rows.append(row_data)

    for row_data in rows:
        movement_kind = row_data["movement_kind"]
        if movement_kind in totals:
            totals[movement_kind] += int(row_data["qty_value"] or 0)

    totals["Остаток"] = int(totals.get("Приход") or 0) - int(totals.get("Расход") or 0)
    info = _product_header_info(article, snapshots, events)
    header_row = _write_journal_report_header(sheet, info, totals, "O")
    _write_event_legend(sheet, row_number=6, events=events)
    headers = [
        "№",
        "Дата",
        "Время",
        "Действие",
        "Приход",
        "Расход",
        "На месте",
        "Короб",
        "Палета",
        "Откуда",
        "Куда",
        "Документ",
        "Текущий статус",
        "Исполнитель",
        "Примечание",
    ]
    sheet.append(headers)
    _style_header(sheet, row_number=header_row)
    sheet.freeze_panes = "A8"

    if not rows:
        sheet.append(["", "", "", "По артикулу не найдены складские события"])
    for index, row_data in enumerate(rows, start=1):
        row_data["cells"][0] = index
        sheet.append(row_data["cells"])
        _apply_event_fill(sheet, sheet.max_row, row_data["event_type"], max_column=15)

    sheet.auto_filter.ref = f"A{header_row}:O{sheet.max_row}"
    _set_widths(
        sheet,
        {
            "A": 6,
            "B": 14,
            "C": 12,
            "D": 30,
            "E": 10,
            "F": 10,
            "G": 10,
            "H": 24,
            "I": 24,
            "J": 18,
            "K": 18,
            "L": 24,
            "M": 26,
            "N": 24,
            "O": 36,
        },
    )


def _write_stock_sheet(sheet, article: str, snapshots: list[WarehouseStockSnapshot]) -> None:
    info = _product_header_info(article, snapshots, [])
    header_row = _write_report_header(sheet, "Текущие остатки", info, "M")
    headers = [
        "№",
        "Всего",
        "Доступно",
        "Резерв обработки",
        "Резерв отгрузки",
        "Состояние",
        "Короб",
        "Палета",
        "Место",
        "Процесс",
        "Последнее событие",
        "Обновлен",
        "Архив",
    ]
    sheet.append(headers)
    _style_header(sheet, row_number=header_row)
    sheet.freeze_panes = "A8"
    container_map = _container_map_for_report(snapshots, [])
    if not snapshots:
        sheet.append(["", "", "", "", "", "Остатков не найдено"])
    for index, snapshot in enumerate(snapshots, start=1):
        row = _snapshot_row(snapshot, container_map)
        sheet.append(
            [
                index,
                int(snapshot.qty or 0),
                int(snapshot.available_qty or 0),
                int(snapshot.processing_reserved_qty or 0),
                int(snapshot.shipping_reserved_qty or 0),
                _state_label(snapshot.warehouse_state_code),
                row.get("box_code") or "",
                row.get("pallet_code") or "",
                _location_label(snapshot.location) or row.get("location") or "",
                _document_label_from_parts(snapshot.source_context_type, snapshot.source_context_id),
                _EVENT_LABELS.get(
                    _clean_text(getattr(snapshot.last_event, "event_type", "")),
                    _clean_text(getattr(snapshot.last_event, "event_type", "")),
                ),
                _datetime_text(snapshot.updated_at),
                "Да" if snapshot.is_archived else "",
            ]
        )
    sheet.auto_filter.ref = f"A{header_row}:M{sheet.max_row}"
    _set_widths(
        sheet,
        {
            "A": 6,
            "B": 10,
            "C": 10,
            "D": 18,
            "E": 16,
            "F": 24,
            "G": 22,
            "H": 22,
            "I": 18,
            "J": 26,
            "K": 26,
            "L": 20,
            "M": 10,
        },
    )


def _snapshots_by_container(snapshots: list[WarehouseStockSnapshot]) -> dict[int, list[WarehouseStockSnapshot]]:
    grouped: dict[int, list[WarehouseStockSnapshot]] = defaultdict(list)
    for snapshot in snapshots:
        for container_id in {
            int(snapshot.container_id or 0),
            int(snapshot.parent_container_id or 0),
        }:
            if container_id:
                grouped[container_id].append(snapshot)
    return grouped


def _snapshots_by_source_context(snapshots: list[WarehouseStockSnapshot]) -> dict[tuple[str, str], list[WarehouseStockSnapshot]]:
    grouped: dict[tuple[str, str], list[WarehouseStockSnapshot]] = defaultdict(list)
    for snapshot in snapshots:
        context_type = _clean_text(snapshot.source_context_type)
        context_id = _clean_text(snapshot.source_context_id)
        if context_type and context_id:
            grouped[(context_type, context_id)].append(snapshot)
    return grouped


def _snapshots_by_last_event(snapshots: list[WarehouseStockSnapshot]) -> dict[int, list[WarehouseStockSnapshot]]:
    grouped: dict[int, list[WarehouseStockSnapshot]] = defaultdict(list)
    for snapshot in snapshots:
        event_id = int(snapshot.last_event_id or 0)
        if event_id:
            grouped[event_id].append(snapshot)
    return grouped


def _snapshots_by_id(snapshots: list[WarehouseStockSnapshot]) -> dict[int, WarehouseStockSnapshot]:
    return {int(snapshot.id): snapshot for snapshot in snapshots if int(snapshot.id or 0)}


def _snapshots_for_event(
    event: WarehouseEvent,
    *,
    snapshots_by_id: dict[int, WarehouseStockSnapshot],
    snapshots_by_container: dict[int, list[WarehouseStockSnapshot]],
    snapshots_by_context: dict[tuple[str, str], list[WarehouseStockSnapshot]],
    snapshots_by_last_event: dict[int, list[WarehouseStockSnapshot]],
) -> list[WarehouseStockSnapshot]:
    result = []
    payload_snapshot_id = _payload_int(event.payload, "snapshot_id")
    if payload_snapshot_id and payload_snapshot_id in snapshots_by_id:
        result.append(snapshots_by_id[payload_snapshot_id])
    result.extend(snapshots_by_last_event.get(int(event.id or 0), []))
    container_id = _event_container_id(event)
    if container_id:
        result.extend(snapshots_by_container.get(container_id, []))
    context_key = (_clean_text(event.stock_context_type), _clean_text(event.stock_context_id))
    if event.event_type == "receiving_arrived" and all(context_key):
        result.extend(snapshots_by_context.get(context_key, []))
    seen = set()
    unique = []
    for snapshot in result:
        snapshot_id = int(snapshot.id or 0)
        if snapshot_id in seen:
            continue
        seen.add(snapshot_id)
        unique.append(snapshot)
    return unique


def _event_container_id(event: WarehouseEvent) -> int:
    if int(event.container_id or 0):
        return int(event.container_id or 0)
    task = event.operation_task
    if task is not None and task.container_id:
        return int(task.container_id or 0)
    return 0


def _container_codes(container_id: int, container_map: dict[int, dict]) -> tuple[str, str]:
    container = container_map.get(int(container_id or 0))
    if not container:
        return "", ""
    container_code = _clean_text(container.get("container_code"))
    parent = container_map.get(int(container.get("parent_container_id") or 0))
    parent_code = _clean_text(parent.get("container_code")) if parent else ""
    if container.get("container_type") == WarehouseContainer.TYPE_BOX:
        return container_code, parent_code
    if container.get("container_type") in {WarehouseContainer.TYPE_PALLET, WarehouseContainer.TYPE_MIXED_PALLET}:
        return "", container_code
    return container_code, parent_code


def _journal_container_codes(
    event: WarehouseEvent,
    snapshot_row: dict,
    container_map: dict[int, dict],
) -> tuple[str, str]:
    event_type = _clean_text(event.event_type)
    if event_type == "receiving_arrived":
        return "", ""
    payload_box = _payload_value(event.payload, "box_code", "source_box_code", "target_box_code")
    payload_pallet = _payload_value(
        event.payload,
        "target_pallet_code",
        "destination_pallet_code",
        "pallet_code",
        "source_pallet_code",
    )
    box_code = payload_box
    pallet_code = payload_pallet
    event_container_is_pallet = False

    container = container_map.get(_event_container_id(event))
    if container:
        container_code = _clean_text(container.get("container_code"))
        container_type = container.get("container_type")
        if container_type == WarehouseContainer.TYPE_BOX and not box_code:
            box_code = container_code
        elif container_type in {WarehouseContainer.TYPE_PALLET, WarehouseContainer.TYPE_MIXED_PALLET}:
            event_container_is_pallet = True
            if not pallet_code:
                pallet_code = container_code

    if not box_code and not event_container_is_pallet:
        box_code = snapshot_row.get("box_code") or ""
    if not pallet_code and _clean_text(event.event_type) not in _UNSTABLE_PARENT_PALLET_EVENT_TYPES:
        pallet_code = snapshot_row.get("pallet_code") or ""
    return box_code, pallet_code


def _same_article(value, article: str) -> bool:
    return _clean_text(value).casefold() == _clean_text(article).casefold()


def _event_payload_article_values(event: WarehouseEvent) -> list[str]:
    payload = event.payload
    if not isinstance(payload, dict):
        return []
    values = []
    for key in ("sku_code", "sku", "article"):
        value = _clean_text(payload.get(key))
        if value:
            values.append(value)
    previous_snapshot = payload.get("previous_snapshot")
    if isinstance(previous_snapshot, dict):
        value = _clean_text(previous_snapshot.get("sku_code") or previous_snapshot.get("sku"))
        if value:
            values.append(value)
    return values


def _event_payload_matches_article(event: WarehouseEvent, article: str) -> bool:
    return any(_same_article(value, article) for value in _event_payload_article_values(event))


def _event_payload_article_mismatch(event: WarehouseEvent, article: str) -> bool:
    if event.reserve_id and event.reserve is not None and _clean_text(event.reserve.sku_code):
        return not _same_article(event.reserve.sku_code, article)
    values = _event_payload_article_values(event)
    return bool(values) and not any(_same_article(value, article) for value in values)


def _payload_int(payload, key: str) -> int:
    if not isinstance(payload, dict):
        return 0
    return _int_value(payload.get(key))


def _event_qty_for_article(event: WarehouseEvent, snapshots: list[WarehouseStockSnapshot], article: str) -> int | None:
    if event.reserve_id and event.reserve is not None and _same_article(event.reserve.sku_code, article):
        if int(event.qty or 0):
            return int(event.qty or 0)
        return int(event.reserve.qty_reserved or 0)
    if _event_payload_matches_article(event, article) and int(event.qty or 0):
        return int(event.qty or 0)
    if snapshots:
        return sum(int(snapshot.qty or 0) for snapshot in snapshots)
    if int(event.qty or 0):
        return int(event.qty or 0)
    return None


def _event_product_name(event: WarehouseEvent, snapshot_row: dict) -> str:
    if snapshot_row.get("name"):
        return snapshot_row["name"]
    reserve = event.reserve
    if reserve is not None and reserve.sku_ref is not None:
        return _clean_text(reserve.sku_ref.name)
    payload_name = _payload_value(event.payload, "name", "product_name")
    return payload_name


def _event_note(event: WarehouseEvent) -> str:
    notes = []
    if event.reserve_id and event.reserve is not None:
        notes.append(
            "Резерв: "
            + ", ".join(
                part
                for part in (
                    _clean_text(event.reserve.reserve_type),
                    _clean_text(event.reserve.status),
                    f"{int(event.reserve.qty_reserved or 0)} шт",
                )
                if part
            )
        )
    if isinstance(event.payload, dict):
        reason = _clean_text(event.payload.get("reason"))
        if reason:
            notes.append(reason)
        for label, key in (
            ("из короба", "source_box_code"),
            ("в короб", "target_box_code"),
            ("из палеты", "source_pallet_code"),
            ("в палету", "target_pallet_code"),
        ):
            value = _clean_text(event.payload.get(key))
            if value:
                notes.append(f"{label}: {value}")
    return " | ".join(notes)


def _operation_label(operation) -> str:
    if operation is None:
        return ""
    operation_types = dict(WarehouseOperation.TYPE_CHOICES)
    label = operation_types.get(operation.operation_type, operation.operation_type)
    context = _context_label(operation.context_type, operation.context_id)
    status = _clean_text(operation.status)
    parts = [label]
    if context:
        parts.append(context)
    if status:
        parts.append(status)
    return " · ".join(parts)


def _location_label(location) -> str:
    if location is None:
        return ""
    zone = _clean_text(location.zone_code).upper()
    row = _int_value(location.row_no)
    section = _int_value(location.section_no)
    tier = _int_value(location.tier_no)
    cell = _int_value(location.cell_no)
    if zone == "OS":
        line_label = _OS_LINE_DISPLAY_LABELS.get(section, str(section or "").strip())
        if line_label and row and tier and cell:
            return f"OS: {line_label}-{row}/{tier}-{cell}"
        if line_label and row:
            return f"OS: {line_label}-{row}"
        return "OS"
    if zone == "MR":
        return f"MR-{row}" if row else "MR"
    if zone in {"PR", "OBR", "OTG"}:
        return zone
    return _clean_text(location.display_name) or _clean_text(location.location_code) or zone


def _state_label(value) -> str:
    code = _clean_text(value)
    if not code:
        return ""
    label = _STATE_LABELS.get(code, code)
    return f"{label} ({code})" if label != code else code


def _movement_kind_label(event_or_type) -> str:
    event_type = _clean_text(getattr(event_or_type, "event_type", event_or_type))
    if event_type == "stock_returned_to_storage" and _stock_return_is_restore(event_or_type):
        return "Приход"
    return _MOVEMENT_KIND_BY_EVENT_TYPE.get(event_type, "На месте")


def _stock_return_is_restore(event) -> bool:
    if not isinstance(event, WarehouseEvent):
        return False
    context_type = _clean_text(event.stock_context_type)
    source_document_type = _clean_text(event.source_document_type)
    payload = event.payload if isinstance(event.payload, dict) else {}
    return (
        context_type in _STOCK_RESTORE_CONTEXT_TYPES
        or source_document_type in _STOCK_RESTORE_SOURCE_DOCUMENT_TYPES
        or bool(payload.get("manual_restore"))
    )


def _event_action_label(event: WarehouseEvent) -> str:
    event_type = _clean_text(event.event_type)
    operation_type = _clean_text(getattr(event.operation, "operation_type", ""))
    if event_type == "movement_started":
        if operation_type == WarehouseOperation.TYPE_PUTAWAY:
            return "Начали размещение на склад"
        if operation_type == WarehouseOperation.TYPE_MOVE_TO_OTG:
            return "Начали перемещение в OTG"
        if operation_type == WarehouseOperation.TYPE_MOVE_TO_PROCESSING:
            return "Начали перемещение в обработку"
        if operation_type == WarehouseOperation.TYPE_INTERNAL_RELOCATION:
            return "Начали внутреннее перемещение"
    if event_type == "movement_completed":
        if operation_type == WarehouseOperation.TYPE_PUTAWAY:
            return "Разместили на склад"
        if operation_type == WarehouseOperation.TYPE_MOVE_TO_OTG:
            return "Приехал в OTG"
        if operation_type == WarehouseOperation.TYPE_MOVE_TO_PROCESSING:
            return "Приехал в обработку"
        if operation_type == WarehouseOperation.TYPE_INTERNAL_RELOCATION:
            return "Завершили внутреннее перемещение"
    if event_type == "stock_returned_to_storage" and _stock_return_is_restore(event):
        return "Восстановили остаток"
    return _EVENT_LABELS.get(event_type, event_type)


def _journal_group_key(
    event: WarehouseEvent,
    *,
    action: str,
    movement_kind: str,
    box_code: str,
    pallet_code: str,
    from_label: str,
    to_label: str,
    document: str,
) -> tuple:
    event_type = _clean_text(event.event_type)
    payload_marker = _payload_value(event.payload, "trip_id", "processing_order_id", "source_document_id")
    if event_type in _LOCATION_INSENSITIVE_DUPLICATE_EVENT_TYPES:
        from_label = ""
        to_label = ""
    return (
        event_type,
        action,
        movement_kind,
        box_code,
        pallet_code,
        from_label,
        to_label,
        document,
        payload_marker,
        timezone.localtime(event.occurred_at).strftime("%Y-%m-%d %H:%M") if event.occurred_at else "",
    )


def _merge_journal_row(existing: dict, incoming: dict) -> None:
    if int(incoming.get("qty_value") or 0) > int(existing.get("qty_value") or 0):
        existing["qty_value"] = incoming["qty_value"]
        _set_movement_qty(existing["cells"], existing["movement_kind"], existing["qty_value"])
    if not existing["cells"][13] and incoming["cells"][13]:
        existing["cells"][13] = incoming["cells"][13]
    if incoming["cells"][14] and incoming["cells"][14] not in existing["cells"][14]:
        existing["cells"][14] = " | ".join(part for part in (existing["cells"][14], incoming["cells"][14]) if part)


def _set_movement_qty(cells: list, movement_kind: str, qty_value: int) -> None:
    cells[4] = ""
    cells[5] = ""
    cells[6] = ""
    if movement_kind == "Приход":
        cells[4] = qty_value
    elif movement_kind == "Расход":
        cells[5] = qty_value
    elif movement_kind == "На месте":
        cells[6] = qty_value


def _context_label(context_type, context_id) -> str:
    context_type = _clean_text(context_type)
    context_id = _clean_text(context_id)
    if context_type and context_id:
        return f"{context_type}:{context_id}"
    return context_type or context_id


def _document_label(event: WarehouseEvent) -> str:
    source = _document_label_from_parts(event.source_document_type, event.source_document_id)
    if source:
        return source
    return _document_label_from_parts(event.stock_context_type, event.stock_context_id)


def _document_label_from_parts(context_type, context_id) -> str:
    context_type = _clean_text(context_type)
    context_id = _clean_text(context_id)
    if not context_type and not context_id:
        return ""
    if context_type == "placement_act":
        return f"{context_id}_PN" if context_id else "PN"
    if context_type == "processing_placement":
        return f"Обработка №{context_id}" if context_id else "Обработка"
    if context_type == "stock_move":
        return f"Перемещение №{context_id}" if context_id else "Перемещение"
    labels = {
        "processing": "Обработка",
        "processing_order": "Обработка",
        "shipping": "Отгрузка",
        "shipping_order": "Отгрузка",
        "trip": "Рейс",
    }
    if context_type in {"receiving", "receiving_order"}:
        return f"{context_id}_PN" if context_id else "PN"
    label = labels.get(context_type, context_type)
    if label and context_id:
        return f"{label} №{context_id}"
    return label or context_id


def _agency_label(agency) -> str:
    if agency is None:
        return ""
    return _clean_text(getattr(agency, "short_name", "")) or _clean_text(getattr(agency, "agn_name", "")) or str(agency)


def _user_label(user, role: str = "") -> str:
    role = _clean_text(role)
    role_label = _ROLE_LABELS.get(role, role)
    if user is None:
        return role_label
    name = _clean_text(getattr(user, "get_full_name", lambda: "")()) or _clean_text(getattr(user, "username", ""))
    if name and role_label and name.casefold() == role_label.casefold():
        return name
    if role_label and name:
        return f"{name} ({role_label})"
    return name or role_label


def _payload_value(payload, *keys: str) -> str:
    if not isinstance(payload, dict):
        return ""
    for key in keys:
        value = _clean_text(payload.get(key))
        if value:
            return value
    return ""


def _unique_join(values: list[str], *, limit: int = 3) -> str:
    result = []
    seen = set()
    for value in values:
        text = _clean_text(value)
        if not text:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(text)
        if len(result) >= limit:
            break
    if len(seen) > limit:
        result.append("...")
    return ", ".join(result)


def _datetime_text(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M:%S")


def _date_text(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y")


def _time_text(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%H:%M:%S")


def _clean_text(value) -> str:
    return str(value or "").strip()


def _int_value(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _safe_filename_part(value: str) -> str:
    text = _clean_text(value)
    text = re.sub(r'[\\/:*?"<>|]+', "_", text)
    text = re.sub(r"\s+", "_", text)
    return text.strip("._")


def _style_header(sheet, row_number: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    font = Font(bold=True)
    for cell in sheet[row_number]:
        cell.fill = fill
        cell.font = font


def _write_event_legend(sheet, row_number: int, events: list[WarehouseEvent]) -> None:
    sheet.cell(row=row_number, column=1, value="Цвета")
    sheet.cell(row=row_number, column=1).font = Font(bold=True)
    event_types = []
    seen = set()
    for event in events:
        event_type = _clean_text(event.event_type)
        if event_type in _HIDDEN_JOURNAL_EVENT_TYPES:
            continue
        if not event_type or event_type in seen:
            continue
        seen.add(event_type)
        event_types.append(event_type)
    for index, event_type in enumerate(event_types, start=2):
        cell = sheet.cell(row=row_number, column=index, value=_EVENT_LABELS.get(event_type, event_type))
        color = _EVENT_FILL_COLORS.get(event_type)
        if color:
            cell.fill = PatternFill(fill_type="solid", fgColor=color)
        cell.font = Font(bold=True)


def _apply_event_fill(sheet, row_number: int, event_type: str, *, max_column: int) -> None:
    color = _EVENT_FILL_COLORS.get(_clean_text(event_type))
    if not color:
        return
    fill = PatternFill(fill_type="solid", fgColor=color)
    for column_number in range(1, max_column + 1):
        sheet.cell(row=row_number, column=column_number).fill = fill


def _set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
