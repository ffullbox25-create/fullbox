from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from employees.models import Employee


class RevisorProductCheckExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="head", password="x")
        Employee.objects.create(
            full_name="Главный менеджер",
            role="head_manager",
            user=self.user,
            is_active=True,
        )
        self.client.force_login(self.user)

    def test_product_check_requires_article(self):
        response = self.client.get("/revisor/product-check/")

        self.assertEqual(response.status_code, 400)

    def test_product_check_downloads_xlsx_for_article(self):
        article = "полка017"
        response = self.client.get("/revisor/product-check/", {"sku": article})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        disposition = response["Content-Disposition"]
        self.assertIn("attachment", disposition)
        self.assertIn(timezone.localdate().isoformat(), disposition)
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("Журнал", workbook.sheetnames)
        self.assertIn("Текущие остатки", workbook.sheetnames)

    def test_revisor_page_has_product_check_download_button(self):
        response = self.client.get("/revisor/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-url="/revisor/product-check/"')
