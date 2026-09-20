from __future__ import annotations

import io
import re
from collections import defaultdict

from django.http import Http404, HttpResponse
from django.shortcuts import redirect
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from employees.access import get_request_role
from audit.models import OrderAuditEntry
from reachtruck.models import MoveTask
from sklad.models import WarehouseEvent, WarehouseStockSnapshot
from sku.models import SKU

from .web_ui import _client_agency_from_request, _processing_params_from_payload, _repair_mojibake_text


def processing_report_xlsx(request, order_id):
    order_key = str(order_id or "").strip()
    if not getattr(request.user, "is_authenticated", False):
        return redirect(f"/accounts/login/?next=/orders/processing/{order_key}/report.xlsx")

    entries = list(
        OrderAuditEntry.objects.filter(order_type="processing", order_id=order_key)
        .select_related("agency")
        .order_by("created_at", "id")
    )
    if not entries:
        raise Http404("Заявка на обработку не найдена")

    latest = entries[-1]
    agency = latest.agency
    client_agency = _client_agency_from_request(request)
    if client_agency:
        if int(client_agency.id) != int(latest.agency_id or 0):
            raise Http404("Заявка на обработку не найдена")
    else:
        role = str(get_request_role(request) or "").strip().lower()
        allowed_roles = {"manager", "storekeeper", "head_manager", "director", "admin", "processing_head"}
        if role not in allowed_roles:
            return HttpResponse("Недостаточно прав", status=403)

    payload = {}
    for entry in entries:
        if isinstance(entry.payload, dict):
            payload.update(entry.payload)

    task_rows = _processing_task_rows(agency, order_key)
    received_rows = _received_rows(agency, order_key, task_rows)
    result_rows = payload.get("processing_results") or []
    if isinstance(result_rows, dict):
        result_rows = list(result_rows.values())
    result_rows = [row for row in result_rows if isinstance(row, dict)]
    final_snapshots = list(
        WarehouseStockSnapshot.objects.filter(
            agency=agency,
            source_context_type="processing",
            source_context_id=order_key,
            is_archived=False,
            qty__gt=0,
        )
        .exclude(warehouse_state_code="processing_consumed")
        .select_related("container", "parent_container", "location")
        .order_by("parent_container__container_code", "container_code", "sku_code", "size", "id")
    )
    article_codes = {
        str(row.get("article") or "").strip()
        for row in received_rows
        if str(row.get("article") or "").strip()
    }
    article_codes.update(
        _article(row).strip()
        for row in result_rows
        if _article(row).strip()
    )
    article_codes.update(
        str(snapshot.sku_code or "").strip()
        for snapshot in final_snapshots
        if str(snapshot.sku_code or "").strip()
    )
    sku_names = {
        str(sku.sku_code or "").strip().casefold(): str(sku.name or "").strip()
        for sku in SKU.objects.filter(
            agency=agency,
            deleted=False,
            sku_code__in=article_codes,
        ).only("sku_code", "name")
    }

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    received_sheet = workbook.create_sheet("Получено в OBR")
    work_sheet = workbook.create_sheet("Выполненные работы")
    placement_sheet = workbook.create_sheet("Итоговое размещение")
    styles = _ReportStyles()

    _fill_summary(
        summary_sheet,
        styles,
        entries=entries,
        payload=payload,
        agency=agency,
        order_key=order_key,
        received_rows=received_rows,
        result_rows=result_rows,
        final_snapshots=final_snapshots,
    )
    _fill_received(received_sheet, styles, received_rows, sku_names)
    _fill_work(work_sheet, styles, result_rows, _processing_params_from_payload(payload), sku_names)
    _fill_placement(placement_sheet, styles, final_snapshots, sku_names)

    stream = io.BytesIO()
    workbook.save(stream)
    safe_order_key = re.sub(r"[^A-Za-z0-9_-]+", "_", order_key) or "processing"
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="processing_{safe_order_key}_report.xlsx"'
    return response


def _processing_task_rows(agency, order_key: str) -> list[dict]:
    rows = []
    seen = set()
    tasks = (
        MoveTask.objects.filter(
            request__agency=agency,
            request__context_type="processing",
            request__context_id=order_key,
            request__destination_zone="OBR",
        )
        .select_related("request")
        .order_by("id")
    )
    for task in tasks:
        payload = task.payload if isinstance(task.payload, dict) else {}
        for raw_row in payload.get("requested_rows") or []:
            if not isinstance(raw_row, dict):
                continue
            row = dict(raw_row)
            row.setdefault("pallet_code", task.pallet_code or payload.get("pallet_code") or "")
            row.setdefault("from_label", payload.get("from_label") or payload.get("from_code") or "")
            identity = (
                str(row.get("box_code") or "").strip(),
                str(row.get("pallet_code") or "").strip(),
                _article(row),
                _barcode(row),
                _number(row.get("qty")),
            )
            if identity in seen:
                continue
            seen.add(identity)
            rows.append(row)
    return rows


def _received_rows(agency, order_key: str, task_rows: list[dict]) -> list[dict]:
    rows_by_box = defaultdict(list)
    for row in task_rows:
        box_key = str(row.get("box_code") or "").strip().lower()
        if box_key:
            rows_by_box[box_key].append(row)

    events = list(
        WarehouseEvent.objects.filter(
            agency=agency,
            stock_context_type="processing",
            stock_context_id=order_key,
            event_type="processing_zone_arrived",
        )
        .select_related("container", "from_location", "to_location")
        .order_by("occurred_at", "id")
    )
    result = []
    used_rows = set()
    for event in events:
        event_payload = event.payload if isinstance(event.payload, dict) else {}
        box_code = str(
            event_payload.get("source_box_code")
            or getattr(event.container, "container_code", "")
            or ""
        ).strip()
        candidates = rows_by_box.get(box_code.lower(), []) if box_code else []
        task_row = next((row for row in candidates if id(row) not in used_rows), {})
        if task_row:
            used_rows.add(id(task_row))
        source_place = _pick(task_row, "from_label", "from_code")
        if not source_place and event.from_location:
            source_place = (
                getattr(event.from_location, "display_name", "")
                or getattr(event.from_location, "location_code", "")
                or str(event.from_location)
            )
        result.append(
            {
                "box": box_code,
                "pallet": str(event_payload.get("source_pallet_code") or _pick(task_row, "pallet_code") or ""),
                "article": _article(task_row),
                "name": str(_pick(task_row, "requested_name", "sku_name", "goods_name", "name") or ""),
                "size": str(_pick(task_row, "requested_size", "size") or ""),
                "barcode": _barcode(task_row),
                "goods_type": str(_pick(task_row, "requested_goods_type", "goods_type") or ""),
                "qty": _number(event.qty) or _number(task_row.get("qty")),
                "source_place": str(source_place or ""),
                "arrived_at": event.occurred_at,
            }
        )

    rows = result or [
        {
            "box": str(row.get("box_code") or ""),
            "pallet": str(row.get("pallet_code") or ""),
            "article": _article(row),
            "name": str(_pick(row, "requested_name", "sku_name", "goods_name", "name") or ""),
            "size": str(_pick(row, "requested_size", "size") or ""),
            "barcode": _barcode(row),
            "goods_type": str(_pick(row, "requested_goods_type", "goods_type") or ""),
            "qty": _number(row.get("qty")),
            "source_place": str(_pick(row, "from_label", "from_code") or ""),
            "arrived_at": None,
        }
        for row in task_rows
    ]
    product_by_box = _product_metadata_by_box(agency, [row.get("box") for row in rows])
    for row in rows:
        product = product_by_box.get(str(row.get("box") or "").strip().casefold(), {})
        row["article"] = row.get("article") or product.get("article") or ""
        row["name"] = row.get("name") or product.get("name") or ""
        row["size"] = row.get("size") or product.get("size") or ""
        row["barcode"] = row.get("barcode") or product.get("barcode") or ""
        row["goods_type"] = row.get("goods_type") or product.get("goods_type") or ""
    return rows


def _product_metadata_by_box(agency, box_codes) -> dict[str, dict]:
    prepared_codes = {
        str(code or "").strip()
        for code in box_codes
        if str(code or "").strip()
    }
    if not prepared_codes:
        return {}
    result = {}
    events = (
        WarehouseEvent.objects.filter(
            agency=agency,
            container__container_code__in=prepared_codes,
        )
        .select_related("container")
        .order_by("occurred_at", "id")
    )
    for event in events:
        payload = event.payload if isinstance(event.payload, dict) else {}
        box_code = str(getattr(event.container, "container_code", "") or "").strip().casefold()
        if not box_code:
            continue
        product = result.setdefault(box_code, {})
        for target, keys in {
            "article": ("sku_code", "sku", "article"),
            "name": ("name", "sku_name", "goods_name"),
            "size": ("size",),
            "barcode": ("barcode", "shk", "ean"),
            "goods_type": ("goods_type",),
        }.items():
            value = _pick(payload, *keys)
            if value not in (None, "", [], {}):
                product[target] = str(value)
    return result


class _ReportStyles:
    def __init__(self):
        self.title_fill = PatternFill("solid", fgColor="D9A617")
        self.header_fill = PatternFill("solid", fgColor="F5E7B0")
        self.section_fill = PatternFill("solid", fgColor="E8F3EC")
        thin = Side(style="thin", color="D9D9D9")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def title(self, sheet, text: str, columns: int):
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        cell = sheet.cell(1, 1, text)
        cell.fill = self.title_fill
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 25

    def headers(self, sheet, row_no: int, values: list[str]):
        for column_no, value in enumerate(values, 1):
            cell = sheet.cell(row_no, column_no, value)
            cell.fill = self.header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.border

    def row(self, sheet, row_no: int, values: list, fill=None):
        for column_no, value in enumerate(values, 1):
            cell = sheet.cell(row_no, column_no, value)
            cell.border = self.border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill


def _fill_summary(sheet, styles, *, entries, payload, agency, order_key, received_rows, result_rows, final_snapshots):
    styles.title(sheet, "Отчёт об обработке", 2)
    styles.headers(sheet, 3, ["Показатель", "Значение"])
    client_name = (
        getattr(agency, "name", "")
        or getattr(agency, "short_name", "")
        or getattr(agency, "legal_name", "")
        or str(agency or "")
    )
    total_processed = sum(
        _number(_pick(row, "processed", "processed_qty", "qty_processed", "fact_qty"))
        for row in result_rows
    )
    values = [
        ("Заявка", f"№ {order_key}_OBR"),
        ("Клиент", client_name),
        ("Статус", _text(payload.get("status_label") or payload.get("status") or getattr(entries[-1], "status", ""))),
        ("Создана", _date_text(entries[0].created_at)),
        ("Завершена / последнее изменение", _date_text(entries[-1].created_at)),
        ("Получено в OBR, шт.", sum(_number(row.get("qty")) for row in received_rows)),
        ("Обработано, шт.", total_processed),
        (
            "Разница: получено - обработано, шт.",
            sum(_number(row.get("qty")) for row in received_rows) - total_processed,
        ),
        ("Размещено после обработки, шт.", sum(_number(row.qty) for row in final_snapshots)),
        ("Исходных коробов", len({row["box"] for row in received_rows if row["box"]})),
        ("Итоговых коробов", len({row.container_code for row in final_snapshots if row.container_code})),
        ("Итоговых паллет", len({getattr(row.parent_container, "container_code", "") for row in final_snapshots if row.parent_container})),
        ("Отчёт сформирован", _date_text(timezone.now())),
    ]
    for row_no, row in enumerate(values, 4):
        styles.row(sheet, row_no, list(row))
    sheet.freeze_panes = "A4"
    _set_widths(sheet, [36, 70])


def _fill_received(sheet, styles, rows, sku_names):
    headers = ["Короб", "Исходная паллета", "Артикул", "Наименование", "Размер", "ШК", "Тип товара", "Количество, шт.", "Откуда забрали", "Прибыл в OBR"]
    styles.title(sheet, "Фактически получено в OBR", len(headers))
    styles.headers(sheet, 3, headers)
    for row_no, row in enumerate(rows, 4):
        name = row["name"] or sku_names.get(str(row["article"] or "").strip().casefold(), "")
        styles.row(sheet, row_no, [row["box"] or "-", row["pallet"] or "-", row["article"] or "-", name or "-", row["size"] or "-", row["barcode"] or "-", row["goods_type"] or "-", row["qty"], row["source_place"] or "-", _date_text(row["arrived_at"])])
    sheet.freeze_panes = "A4"
    if rows:
        sheet.auto_filter.ref = f"A3:J{len(rows) + 3}"
    _set_widths(sheet, [26, 26, 20, 42, 12, 20, 14, 16, 28, 20])


def _fill_work(sheet, styles, result_rows, processing_params, sku_names):
    styles.title(sheet, "Выполненные работы", 8)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=8)
    sheet.cell(3, 1, "Результаты по товару").fill = styles.section_fill
    sheet.cell(3, 1).font = Font(bold=True)
    styles.headers(sheet, 4, ["Артикул", "Новый артикул", "Наименование", "Размер", "ШК", "Получено", "Обработано", "Разница"])
    row_no = 5
    for row in result_rows:
        received = _number(_pick(row, "received", "received_qty", "qty", "quantity"))
        processed = _number(_pick(row, "processed", "processed_qty", "qty_processed", "fact_qty"))
        article = _article(row)
        name = _pick(row, "name", "sku_name", "product_name") or sku_names.get(article.strip().casefold(), "")
        styles.row(sheet, row_no, [article or "-", _text(_pick(row, "new_article", "target_article", "result_article")), _text(name), _text(_pick(row, "size")), _text(_pick(row, "barcode", "shk")), received, processed, processed - received], styles.header_fill if received != processed else None)
        row_no += 1
    row_no += 1
    sheet.merge_cells(start_row=row_no, start_column=1, end_row=row_no, end_column=8)
    sheet.cell(row_no, 1, "Параметры обработки").fill = styles.section_fill
    sheet.cell(row_no, 1).font = Font(bold=True)
    row_no += 1
    styles.headers(sheet, row_no, ["Параметр", "Значение"])
    for param in processing_params:
        if not isinstance(param, dict):
            continue
        row_no += 1
        styles.row(sheet, row_no, [_text(_pick(param, "label", "name", "title")), _text(_pick(param, "display_value", "value"))])
    sheet.freeze_panes = "A5"
    _set_widths(sheet, [26, 26, 42, 12, 22, 14, 14, 14])


def _fill_placement(sheet, styles, snapshots, sku_names):
    headers = ["Короб", "Паллета", "Артикул", "Наименование", "Размер", "ШК", "Тип товара", "Количество, шт.", "Зона", "Место хранения", "Состояние"]
    styles.title(sheet, "Итоговое размещение после обработки", len(headers))
    styles.headers(sheet, 3, headers)
    for row_no, snapshot in enumerate(snapshots, 4):
        location = snapshot.location
        location_label = ""
        if location:
            location_label = getattr(location, "display_name", "") or getattr(location, "location_code", "") or str(location)
        pallet = getattr(snapshot.parent_container, "container_code", "") if snapshot.parent_container else ""
        location_label = _repair_mojibake_text(location_label)
        name = sku_names.get(str(snapshot.sku_code or "").strip().casefold(), "")
        styles.row(sheet, row_no, [snapshot.container_code or getattr(snapshot.container, "container_code", "") or "-", pallet or "-", snapshot.sku_code or "-", name or "-", snapshot.size or "-", snapshot.barcode or "-", snapshot.goods_type or "-", _number(snapshot.qty), snapshot.zone_code or "-", location_label or snapshot.zone_code or "-", snapshot.warehouse_state_code or "-"])
    sheet.freeze_panes = "A4"
    if snapshots:
        sheet.auto_filter.ref = f"A3:K{len(snapshots) + 3}"
    _set_widths(sheet, [26, 26, 20, 42, 12, 20, 14, 16, 14, 30, 24])


def _pick(row, *keys):
    if not isinstance(row, dict):
        return ""
    for key in keys:
        value = row.get(key)
        if value not in (None, "", [], {}):
            return value
    return ""


def _number(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _article(row) -> str:
    return str(_pick(row, "requested_article", "requested_sku", "requested_sku_code", "article", "sku", "sku_code") or "")


def _barcode(row) -> str:
    value = _pick(row, "requested_barcode", "barcode", "shk", "ean")
    if value:
        return str(value)
    values = _pick(row, "requested_barcodes", "barcodes")
    return str(values[0]) if isinstance(values, (list, tuple)) and values else ""


def _text(value) -> str:
    if value in (None, "", [], {}):
        return "-"
    if isinstance(value, bool):
        return "Да" if value else "Нет"
    if isinstance(value, (list, tuple, set)):
        return ", ".join(_text(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_text(item)}" for key, item in value.items())
    return str(value)


def _date_text(value) -> str:
    if not value:
        return "-"
    try:
        value = timezone.localtime(value)
    except (TypeError, ValueError):
        pass
    return value.strftime("%d.%m.%Y %H:%M") if hasattr(value, "strftime") else _text(value)


def _set_widths(sheet, widths):
    for column_no, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column_no)].width = width
