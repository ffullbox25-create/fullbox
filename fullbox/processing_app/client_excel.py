"""Excel import and physical-box planning for client processing requests."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook


MAX_EXCEL_SIZE = 10 * 1024 * 1024
MAX_EXCEL_ROWS = 5000
MAX_EXCEL_COLUMNS = 50


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value: Any) -> str:
    text = _text(value).casefold().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]+", " ", text).strip()


def _positive_int(value: Any) -> int:
    text = _text(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not text:
        return 0
    try:
        number = float(text)
    except (TypeError, ValueError):
        return 0
    if number <= 0 or not number.is_integer():
        return 0
    return int(number)


def _normalize_barcode(value: Any) -> str:
    text = _text(value)
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


@dataclass(frozen=True)
class ProcessingExcelLine:
    row_no: int
    source_article: str
    source_barcode: str
    target_article: str
    target_barcode: str
    qty: int
    declared_boxes: int = 0
    declared_box_qty: int = 0
    sheet_name: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "row_no": self.row_no,
            "sheet_name": self.sheet_name,
            "source_article": self.source_article,
            "source_barcode": self.source_barcode,
            "target_article": self.target_article,
            "target_barcode": self.target_barcode,
            "qty": self.qty,
            "declared_boxes": self.declared_boxes,
            "declared_box_qty": self.declared_box_qty,
        }


_VERTICAL_LABELS = {
    "article": {
        "артикул продавца",
        "артикул товара",
        "артикул",
        "sku",
    },
    "barcode": {"баркод", "штрихкод", "штрих код", "шк"},
    "qty": {"шт", "штук", "количество", "кол во товаров", "количество товаров"},
    "boxes": {"кол во коробов", "количество коробов", "коробов"},
    "box_qty": {"квантовка", "кратность", "в коробе", "шт в коробе"},
}


def _vertical_kind(value: Any) -> str:
    normalized = _key(value)
    for kind, aliases in _VERTICAL_LABELS.items():
        if normalized in aliases:
            return kind
    return ""


def _vertical_line(rows: list[tuple], sheet_name: str) -> tuple[ProcessingExcelLine | None, list[str]]:
    found: dict[str, tuple[int, tuple]] = {}
    for row_no, row in enumerate(rows, start=1):
        if not row:
            continue
        kind = _vertical_kind(row[0])
        if kind and kind not in found:
            found[kind] = (row_no, row)
    if "article" not in found or "qty" not in found:
        return None, []

    article_row_no, article_row = found["article"]
    qty_row_no, qty_row = found["qty"]
    source_article = _text(article_row[1] if len(article_row) > 1 else "")
    target_article = _text(article_row[2] if len(article_row) > 2 else "")
    source_qty = _positive_int(qty_row[1] if len(qty_row) > 1 else "")
    target_qty = _positive_int(qty_row[2] if len(qty_row) > 2 else "")
    errors: list[str] = []
    if source_qty and target_qty and source_qty != target_qty:
        errors.append(
            f"Лист «{sheet_name}»: количество до и после перемаркировки различается "
            f"({source_qty} и {target_qty} шт.)."
        )
    qty = source_qty or target_qty
    barcode_row = found.get("barcode", (0, ()))
    boxes_row = found.get("boxes", (0, ()))
    box_qty_row = found.get("box_qty", (0, ()))
    source_barcode = _normalize_barcode(barcode_row[1][1] if len(barcode_row[1]) > 1 else "")
    target_barcode = _normalize_barcode(barcode_row[1][2] if len(barcode_row[1]) > 2 else "")
    declared_boxes = _positive_int(boxes_row[1][1] if len(boxes_row[1]) > 1 else "")
    declared_box_qty = _positive_int(box_qty_row[1][1] if len(box_qty_row[1]) > 1 else "")
    if not source_article and not source_barcode:
        errors.append(f"Лист «{sheet_name}»: не указан исходный артикул или баркод.")
    if qty <= 0:
        errors.append(f"Лист «{sheet_name}»: в строке {qty_row_no} укажите количество в штуках.")
    if errors:
        return None, errors
    return ProcessingExcelLine(
        row_no=article_row_no,
        source_article=source_article,
        source_barcode=source_barcode,
        target_article=target_article,
        target_barcode=target_barcode,
        qty=qty,
        declared_boxes=declared_boxes,
        declared_box_qty=declared_box_qty,
        sheet_name=sheet_name,
    ), []


def _table_header_kind(value: Any) -> str:
    normalized = _key(value)
    if not normalized:
        return ""
    target = any(token in normalized for token in ("нов", "после", "результ", "на что", "артикул б"))
    if "баркод" in normalized or "штрихкод" in normalized or normalized == "шк":
        return "target_barcode" if target else "source_barcode"
    if "артикул" in normalized or normalized in {"sku", "offer id", "offerid"}:
        return "target_article" if target else "source_article"
    if normalized in _VERTICAL_LABELS["qty"] or "количество шт" in normalized:
        return "qty"
    if normalized in _VERTICAL_LABELS["boxes"]:
        return "declared_boxes"
    if normalized in _VERTICAL_LABELS["box_qty"] or "шт короб" in normalized:
        return "declared_box_qty"
    return ""


def _tabular_lines(rows: list[tuple], sheet_name: str) -> tuple[list[ProcessingExcelLine], list[str]]:
    header_index = -1
    columns: dict[str, int] = {}
    for index, row in enumerate(rows[:30]):
        candidate: dict[str, int] = {}
        for column_index, value in enumerate(row[:MAX_EXCEL_COLUMNS]):
            kind = _table_header_kind(value)
            if kind and kind not in candidate:
                candidate[kind] = column_index
        if "qty" in candidate and ({"source_article", "source_barcode"} & set(candidate)):
            header_index = index
            columns = candidate
            break
    if header_index < 0:
        return [], []

    lines: list[ProcessingExcelLine] = []
    errors: list[str] = []
    for index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        value = lambda name: row[columns[name]] if name in columns and columns[name] < len(row) else ""
        source_article = _text(value("source_article"))
        source_barcode = _normalize_barcode(value("source_barcode"))
        qty_text = _text(value("qty"))
        if not any((source_article, source_barcode, qty_text)):
            continue
        qty = _positive_int(qty_text)
        if not source_article and not source_barcode:
            errors.append(f"Лист «{sheet_name}», строка {index}: нет исходного артикула или баркода.")
            continue
        if qty <= 0:
            errors.append(f"Лист «{sheet_name}», строка {index}: укажите количество в штуках.")
            continue
        lines.append(
            ProcessingExcelLine(
                row_no=index,
                source_article=source_article,
                source_barcode=source_barcode,
                target_article=_text(value("target_article")),
                target_barcode=_normalize_barcode(value("target_barcode")),
                qty=qty,
                declared_boxes=_positive_int(value("declared_boxes")),
                declared_box_qty=_positive_int(value("declared_box_qty")),
                sheet_name=sheet_name,
            )
        )
    return lines, errors


def parse_processing_excel(file_obj) -> tuple[list[ProcessingExcelLine], list[str]]:
    filename = _text(getattr(file_obj, "name", ""))
    if filename and not filename.casefold().endswith(".xlsx"):
        return [], ["Для автоматического заполнения загрузите файл Excel в формате .xlsx."]
    file_size = int(getattr(file_obj, "size", 0) or 0)
    if file_size > MAX_EXCEL_SIZE:
        return [], ["Файл Excel больше 10 МБ. Уменьшите файл и повторите загрузку."]
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        return [], ["Не удалось прочитать Excel. Проверьте, что файл сохранён в формате .xlsx."]

    lines: list[ProcessingExcelLine] = []
    errors: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = list(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(int(worksheet.max_row or 1), MAX_EXCEL_ROWS),
                    min_col=1,
                    max_col=min(int(worksheet.max_column or 1), MAX_EXCEL_COLUMNS),
                    values_only=True,
                )
            )
            vertical, vertical_errors = _vertical_line(rows, worksheet.title)
            if vertical is not None or vertical_errors:
                if vertical is not None:
                    lines.append(vertical)
                errors.extend(vertical_errors)
                continue
            table_lines, table_errors = _tabular_lines(rows, worksheet.title)
            lines.extend(table_lines)
            errors.extend(table_errors)
    finally:
        workbook.close()
    if not lines and not errors:
        errors.append(
            "В Excel не найден состав обработки. Укажите исходный артикул или баркод и количество в штуках."
        )
    return lines, errors


def _box_codes(row: dict) -> list[str]:
    raw = row.get("box_codes") or row.get("box_code") or []
    if isinstance(raw, str):
        raw = raw.split(",")
    if not isinstance(raw, (list, tuple)):
        return []
    return list(dict.fromkeys(_text(code) for code in raw if _text(code)))


def _row_matches_line(row: dict, line: ProcessingExcelLine) -> bool:
    row_article = _key(row.get("sku_code") or row.get("sku") or row.get("article"))
    row_barcode = _normalize_barcode(row.get("barcode")).casefold()
    line_article = _key(line.source_article)
    line_barcode = _normalize_barcode(line.source_barcode).casefold()
    if line_barcode and row_barcode != line_barcode:
        return False
    if line_article and row_article != line_article:
        return False
    return bool(line_article or line_barcode)


def _exact_full_box_selection(candidates: list[dict], target_qty: int) -> list[dict] | None:
    if target_qty <= 0:
        return []
    states: dict[int, tuple[int, ...]] = {0: ()}
    for index, row in enumerate(candidates):
        box_qty = max(int(row.get("box_qty") or 0), 0)
        if box_qty <= 0 or box_qty > target_qty:
            continue
        for subtotal, selected in list(states.items())[::-1]:
            candidate = subtotal + box_qty
            if candidate > target_qty or candidate in states:
                continue
            states[candidate] = (*selected, index)
        if target_qty in states:
            return [candidates[item] for item in states[target_qty]]
    return None


def build_processing_excel_plan(
    lines: list[ProcessingExcelLine],
    stock_rows: list[dict],
) -> dict[str, Any]:
    used_box_codes: set[str] = set()
    planned_lines: list[dict[str, Any]] = []
    errors: list[str] = []
    warnings: list[str] = []

    for line in lines:
        pool = []
        for row in stock_rows or []:
            if not isinstance(row, dict):
                continue
            codes = _box_codes(row)
            if (
                not _row_matches_line(row, line)
                or len(codes) != 1
                or codes[0].casefold() in used_box_codes
                or bool(row.get("is_mixed_box"))
                or int(row.get("available_qty") or 0) <= 0
            ):
                continue
            pool.append({**row, "_box_code": codes[0]})
        identity = line.source_article or line.source_barcode
        if not pool:
            errors.append(
                f"Лист «{line.sheet_name}», строка {line.row_no}: {identity} не найден "
                "в доступных немиксованных коробах."
            )
            continue

        full_candidates = [
            row
            for row in pool
            if int(row.get("box_qty") or 0) > 0
            and int(row.get("available_qty") or 0) >= int(row.get("box_qty") or 0)
        ]
        full_candidates.sort(
            key=lambda row: (
                0 if line.declared_box_qty and int(row.get("box_qty") or 0) == line.declared_box_qty else 1,
                -int(row.get("box_qty") or 0),
                _key(row.get("source_state")),
                _key(row.get("_box_code")),
            )
        )
        exact = _exact_full_box_selection(full_candidates, line.qty)
        selected_full: list[dict] = exact or []
        remaining = 0 if exact is not None else line.qty
        if exact is None:
            for row in full_candidates:
                box_qty = int(row.get("box_qty") or 0)
                if box_qty <= remaining:
                    selected_full.append(row)
                    remaining -= box_qty

        selected_codes = {str(row["_box_code"]).casefold() for row in selected_full}
        partial_row = None
        if remaining > 0:
            partial_candidates = [
                row
                for row in pool
                if str(row["_box_code"]).casefold() not in selected_codes
                and int(row.get("available_qty") or 0) >= remaining
            ]
            partial_candidates.sort(
                key=lambda row: (
                    int(row.get("available_qty") or 0),
                    0 if line.declared_box_qty and int(row.get("box_qty") or 0) == line.declared_box_qty else 1,
                    _key(row.get("_box_code")),
                )
            )
            partial_row = partial_candidates[0] if partial_candidates else None
        if remaining > 0 and partial_row is None:
            available_total = sum(int(row.get("available_qty") or 0) for row in pool)
            errors.append(
                f"Лист «{line.sheet_name}», строка {line.row_no}: для {identity} запрошено "
                f"{line.qty} шт., но безопасно подобрать нужное количество из доступных коробов "
                f"не удалось (доступно {available_total} шт.)."
            )
            continue

        allocations: list[dict[str, Any]] = []
        for row in selected_full:
            box_qty = int(row.get("box_qty") or 0)
            allocations.append(
                {
                    "stock_key": _text(row.get("key")),
                    "box_code": row["_box_code"],
                    "qty": box_qty,
                    "box_qty": box_qty,
                    "mode": "whole",
                    "sku_code": _text(row.get("sku_code") or row.get("sku")),
                    "name": _text(row.get("name")),
                    "size": _text(row.get("size")),
                    "barcode": _normalize_barcode(row.get("barcode")),
                    "goods_type": _text(row.get("goods_type")),
                    "source_state": _text(row.get("source_state")),
                    "source_zone": _text(row.get("source_zone")),
                }
            )
            used_box_codes.add(str(row["_box_code"]).casefold())
        if partial_row is not None and remaining > 0:
            allocations.append(
                {
                    "stock_key": _text(partial_row.get("key")),
                    "box_code": partial_row["_box_code"],
                    "qty": remaining,
                    "box_qty": int(partial_row.get("box_qty") or partial_row.get("available_qty") or 0),
                    "mode": "piece",
                    "sku_code": _text(partial_row.get("sku_code") or partial_row.get("sku")),
                    "name": _text(partial_row.get("name")),
                    "size": _text(partial_row.get("size")),
                    "barcode": _normalize_barcode(partial_row.get("barcode")),
                    "goods_type": _text(partial_row.get("goods_type")),
                    "source_state": _text(partial_row.get("source_state")),
                    "source_zone": _text(partial_row.get("source_zone")),
                }
            )
            used_box_codes.add(str(partial_row["_box_code"]).casefold())

        whole_qty = sum(row["qty"] for row in allocations if row["mode"] == "whole")
        piece_qty = sum(row["qty"] for row in allocations if row["mode"] == "piece")
        if line.declared_boxes and line.declared_boxes != len(allocations):
            warnings.append(
                f"{identity}: в Excel указано {line.declared_boxes} кор., по фактическому складу "
                f"используется {len(allocations)} кор."
            )
        planned_lines.append(
            {
                **line.as_dict(),
                "whole_box_count": len(selected_full),
                "whole_box_qty": whole_qty,
                "piece_pick_qty": piece_qty,
                "physical_box_count": len(allocations),
                "allocations": allocations,
            }
        )

    whole_box_count = sum(int(row.get("whole_box_count") or 0) for row in planned_lines)
    whole_box_qty = sum(int(row.get("whole_box_qty") or 0) for row in planned_lines)
    piece_pick_qty = sum(int(row.get("piece_pick_qty") or 0) for row in planned_lines)
    requested_qty = sum(int(row.get("qty") or 0) for row in planned_lines)
    return {
        "ok": bool(planned_lines) and not errors,
        "lines": planned_lines,
        "errors": errors,
        "warnings": warnings,
        "totals": {
            "positions": len(planned_lines),
            "requested_qty": requested_qty,
            "whole_box_count": whole_box_count,
            "whole_box_qty": whole_box_qty,
            "piece_pick_qty": piece_pick_qty,
            "piece_pick_box_count": sum(1 for row in planned_lines for item in row["allocations"] if item["mode"] == "piece"),
            "requires_piece_pick_confirmation": piece_pick_qty > 0,
        },
    }


def validate_processing_box_rows(
    submitted_rows: list[dict],
    available_rows: list[dict],
) -> tuple[list[dict], dict[str, int | bool], str]:
    available_by_code: dict[str, list[dict]] = {}
    for row in available_rows or []:
        codes = _box_codes(row)
        if len(codes) != 1:
            continue
        available_by_code.setdefault(codes[0].casefold(), []).append(row)

    normalized: list[dict] = []
    used_row_keys: set[str] = set()
    for submitted in submitted_rows or []:
        if not isinstance(submitted, dict):
            continue
        codes = _box_codes(submitted)
        if len(codes) != 1:
            return [], {}, "Для каждой строки обработки нужно выбрать один конкретный короб."
        box_code = codes[0]
        candidates = available_by_code.get(box_code.casefold(), [])
        if not candidates:
            return [], {}, (
                "Короба уже заняты другой заявкой или доставлены в обработку: "
                f"{box_code}. Обновите состав заявки."
            )
        source_article = _key(submitted.get("article") or submitted.get("sku"))
        source_barcode = _normalize_barcode(submitted.get("barcode")).casefold()
        source_size = _key(submitted.get("size"))
        source_type = _key(submitted.get("goods_type"))

        def matches(row: dict) -> bool:
            if source_barcode and _normalize_barcode(row.get("barcode")).casefold() != source_barcode:
                return False
            if source_article and _key(row.get("sku_code") or row.get("sku")) != source_article:
                return False
            if source_size and _key(row.get("size")) != source_size:
                return False
            if source_type and _key(row.get("goods_type")) != source_type:
                return False
            return True

        matches_rows = [row for row in candidates if matches(row)]
        if len(matches_rows) != 1:
            return [], {}, (
                f"Короб {box_code} больше не соответствует выбранному товару. "
                "Обновите состав заявки."
            )
        available = matches_rows[0]
        row_key = _text(available.get("key")) or f"{box_code.casefold()}:{source_article}:{source_barcode}:{source_size}:{source_type}"
        if row_key in used_row_keys:
            return [], {}, f"Короб {box_code} выбран повторно для одной позиции. Обновите состав заявки."
        used_row_keys.add(row_key)
        qty = _positive_int(submitted.get("qty"))
        available_qty = max(int(available.get("available_qty") or 0), 0)
        if qty <= 0 or qty > available_qty:
            return [], {}, (
                f"В коробе {box_code} доступно {available_qty} шт., запрошено {qty} шт. "
                "Обновите состав заявки."
            )
        box_qty = max(int(available.get("box_qty") or available_qty), 0)
        is_whole = not bool(available.get("is_mixed_box")) and box_qty > 0 and qty == box_qty
        row = dict(submitted)
        row.update(
            {
                "article": _text(available.get("sku_code") or available.get("sku")),
                "size": _text(available.get("size")),
                "barcode": _normalize_barcode(available.get("barcode")),
                "goods_type": _text(available.get("goods_type")),
                "qty": qty,
                "box_qty": box_qty,
                "boxes": 1 if is_whole else 0,
                "partial_boxes": 0 if is_whole else 1,
                "box_codes": [box_code],
                "source_state": _text(available.get("source_state")),
                "source_zone": _text(available.get("source_zone")),
                "strict_selected_box": True,
                "is_mixed_box": bool(available.get("is_mixed_box")),
                "mixed_group": _text(available.get("mixed_group")),
            }
        )
        if is_whole:
            row.pop("allow_partial_box_reserve", None)
            row.pop("partial_box_split", None)
        else:
            row["allow_partial_box_reserve"] = True
            row["partial_box_split"] = {
                "qty": qty,
                "box_qty": box_qty,
                "box_code": box_code,
                "mode": "units",
            }
        normalized.append(row)

    whole_codes = {
        _box_codes(row)[0].casefold()
        for row in normalized
        if int(row.get("boxes") or 0) > 0 and _box_codes(row)
    }
    piece_rows = [row for row in normalized if int(row.get("partial_boxes") or 0) > 0]
    piece_codes = {_box_codes(row)[0].casefold() for row in piece_rows if _box_codes(row)}
    summary: dict[str, int | bool] = {
        "requested_qty": sum(int(row.get("qty") or 0) for row in normalized),
        "whole_box_count": len(whole_codes),
        "whole_box_qty": sum(int(row.get("qty") or 0) for row in normalized if int(row.get("boxes") or 0) > 0),
        "piece_pick_qty": sum(int(row.get("qty") or 0) for row in piece_rows),
        "piece_pick_box_count": len(piece_codes),
        "requires_piece_pick_confirmation": bool(piece_rows),
    }
    return normalized, summary, ""
