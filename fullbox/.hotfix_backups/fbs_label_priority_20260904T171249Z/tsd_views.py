from collections import defaultdict
from copy import copy
from datetime import date, timedelta
from decimal import Decimal
from functools import wraps
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlencode, urlsplit
import re
import unicodedata
import uuid

from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Exists, F, Max, Min, OuterRef, Prefetch, Q, Sum
from django.db.models.functions import Coalesce
from django.http import (
    FileResponse,
    Http404,
    HttpResponse,
    HttpResponseForbidden,
    JsonResponse,
    QueryDict,
)
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from openpyxl import Workbook

from agent.models import AgentEvent
from employees.access import (
    get_request_employee,
    get_request_role,
    get_request_roles,
    resolve_cabinet_url,
    role_required,
)
from sku.models import Agency, SKU, SKUBarcode
from processing_app.models import ProcessingPrintJob
from sklad.models import WarehouseLocation
from sklad.topology import os_location_code, os_location_label

from .exceptions import FbsError, FbsMarkingAlreadyUsedError, FbsScanMismatchError
from .flags import feature_enabled, module_enabled
from .controller_session import get_controller_workstation
from .services.printing import (
    fbs_desktop_print_lease_seconds,
    recover_stale_fbs_order_label_print_job,
)
from .services.pick_restock import (
    confirm_marketplace_rejected_order_return,
    confirm_ozon_canceled_order_return,
    order_is_client_canceled_by_marketplace,
)
from .services.physical_locations import fbs_box_physical_location_label
from .workspace import workspace_shell_enabled
from .models import (
    FbsControllerCheckTote,
    FbsControllerPickTote,
    FbsControllerSession,
    FbsControllerToteOrder,
    FbsOrder,
    FbsHandoverBatch,
    FbsHandoverBox,
    FbsHandoverOrder,
    FbsHandoverOrderAssignment,
    FbsIntegrationProfile,
    FbsInventorySession,
    FbsMarketplaceMetadataTransfer,
    FbsMarketplaceCommand,
    FbsOrderLabel,
    FbsOrderItem,
    FbsOrderStockAllocation,
    FbsInternalMovement,
    FbsPickBatch,
    FbsPickException,
    FbsPickRestockRequest,
    FbsPickRestockScan,
    FbsPickScanEvent,
    FbsPickTask,
    FbsPickVerificationProgress,
    FbsPickingCart,
    FbsReplenishmentAllocation,
    FbsReplenishmentLine,
    FbsReplenishmentPlan,
    FbsReplenishmentPreparedBox,
    FbsStorageCell,
    FbsPallet,
    FbsBox,
    FbsStockBalance,
    FbsToteBinding,
    FbsWorkstation,
)
from .services import (
    ACTIVE_PICK_RESTOCK_STATUSES,
    HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES,
    activate_drained_inventory,
    add_handover_box,
    add_wb_handover_boxes,
    add_order_to_handover_box,
    approve_inventory,
    analyze_replenishment_demand,
    cancel_replenishment_plan,
    claim_pick_batch,
    claim_pick_batch_verification,
    claim_next_pick_batch,
    claim_internal_movement,
    claim_pick_task,
    claim_pick_restock_request,
    claim_replenishment_plan,
    close_handover_box,
    complete_replenishment_allocation,
    complete_pick_allocation,
    confirm_order_label_scan,
    confirm_replenishment_plan,
    create_handover_batch,
    create_inventory_session,
    create_order_pick_restock_request,
    dispatch_handover_batch,
    handover_pick_batch_for_verification,
    finish_inventory_count,
    generate_replenishment_plans,
    pack_staged_item_plan,
    prepare_pick_queue,
    pick_missing_group_quantity,
    configured_fbs_print_workstations,
    queue_fbs_handover_box_label_print,
    queue_fbs_handover_supply_label_print,
    queue_fbs_order_label_print,
    queue_verification_problem_restock,
    request_invalid_kiz_reroute,
    record_inventory_scan,
    refresh_handover_acceptance,
    request_wb_handover_delivery,
    report_pick_missing_quantity,
    release_order_pick_restock_to_queue,
    retry_order_pick_restock_marketplace,
    scan_internal_movement,
    scan_handover_box,
    scan_pick_restock,
    stage_replenishment_allocation,
    validate_replenishment_source_scan,
    validate_pick_box_scan,
    validate_pick_cell_scan,
    verify_handover_order_label,
    verify_pick_allocation_unit,
    wb_handover_uses_marketplace_boxes,
    pick_restock_state,
)
from .services.picking import (
    ACTIVE_TASK_STATUSES,
    assign_pick_handover_workstation,
    find_pick_verification_allocation,
    format_queue_rejections,
    resolve_verification_item_scan,
    wb_optional_marking_available,
)
from .services.floor_replenishment import (
    analyze_floor_replenishment_needs,
    create_floor_replenishment_movement,
)
from .services.handover import (
    approve_handover_verification_override,
    archive_handover_batch,
    get_ready_handover_box_for_order,
    handover_archive_readiness,
    handover_composition_readiness,
    handover_verification_override_status,
    has_handover_verification_override,
)
from .services.marketplace import (
    is_final_wb_marking_rejection,
    retry_wb_marking_code,
)
from .services.scanning import record_order_label_scan_events
from .services.labels import is_preloaded_ozon_order_label
from .services.printing import queue_fbs_preloaded_ozon_order_label_print
from .services.sync import marketplace_terminal_order_q
from .services.totes import (
    auto_release_pick_tote_if_complete,
    confirm_order_label_to_check_tote,
    controller_pick_tote_for_batch,
    mark_pick_tote_awaiting_empty,
    record_extra_problem_tote_item,
)
from .staging import movement_staging_container_code


FULLBOX_DESKTOP_DIRECT_PRINT_MIN_VERSION = (1, 0, 12)


def _fullbox_desktop_version(request) -> tuple[int, int, int] | None:
    user_agent = str(request.META.get("HTTP_USER_AGENT") or "")
    match = re.search(r"(?:^|\s)FullboxDesktop/(\d+)\.(\d+)\.(\d+)(?:\s|$)", user_agent)
    if match is None:
        return None
    return tuple(int(value) for value in match.groups())


def _fullbox_desktop_supports_direct_print(request) -> bool:
    version = _fullbox_desktop_version(request)
    return version is not None and version >= FULLBOX_DESKTOP_DIRECT_PRINT_MIN_VERSION


def _desktop_label_print_target(request, *, shipping: bool = False) -> dict:
    if str(request.POST.get("fullbox_desktop") or "").strip() != "1":
        return {}
    if not _fullbox_desktop_supports_direct_print(request):
        return {}
    workstation_id = str(request.POST.get("desktop_workstation_id") or "").strip()
    printer_field = "desktop_shipping_printer" if shipping else "desktop_label_printer"
    fallback_field = "desktop_label_printer" if shipping else "desktop_shipping_printer"
    printer_name = str(request.POST.get(printer_field) or "").strip()
    if not printer_name:
        # Наклейки отгрузки и заказа одного размера, поэтому если нужный принтер
        # в Fullbox Desktop не выбран, берем второй выбранный.
        printer_name = str(request.POST.get(fallback_field) or "").strip()
    if not workstation_id or len(workstation_id) > 80 or not re.fullmatch(r"[A-Za-z0-9._-]+", workstation_id):
        return {}
    if not printer_name:
        # Ни один принтер не выбран. Отдаем печать рабочему месту вместо отказа:
        # у него принтер и агент настроены, иначе оператор видел бы только
        # ошибку и не мог напечатать вообще ничего.
        return {}
    return {
        "desktop_agent_id": f"desktop:{workstation_id}",
        "desktop_printer_name": printer_name,
    }


STOREKEEPER_ROLES = ("storekeeper", "head_manager", "director", "admin")
CONTROLLER_ROLES = ("fbs_controller", *STOREKEEPER_ROLES)
INVENTORY_MANAGER_ROLES = ("head_manager", "director", "admin")
REACHTRUCK_ROLES = ("reachtruck_driver", "director", "admin")
DEMAND_DISPATCH_ORDER_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
)
DEMAND_MAPPING_ORDER_STATUSES = (
    *DEMAND_DISPATCH_ORDER_STATUSES,
    FbsOrder.STATUS_VALIDATION_FAILED,
)
DEMAND_MAPPING_STATUS_LABELS = {
    "barcode_missing": "Маркетплейс не передал ШК.",
    "barcode_ambiguous": "Маркетплейс передал несколько ШК; нужен один точный ШК.",
    "barcode_not_found": "ШК не найден в номенклатуре клиента.",
}
DEMAND_STATE_CHOICES = (
    ("action", "Требуют действия"),
    ("mapping", "Ошибка сопоставления"),
    ("ready", "Готовы к плану"),
    ("in_plan", "Уже в планах"),
    ("blocked", "Есть блокировка"),
    ("covered", "Потребность закрыта"),
)
PICKER_ROLES = ("picker", "processing_worker", "storekeeper", "head_manager", "director", "admin")
TSD_ROLES = tuple(
    dict.fromkeys((*CONTROLLER_ROLES, *REACHTRUCK_ROLES, *PICKER_ROLES))
)
OPEN_PLAN_STATUSES = (
    FbsReplenishmentPlan.STATUS_CONFIRMED,
    FbsReplenishmentPlan.STATUS_IN_PROGRESS,
)
STOREKEEPER_OPEN_PLAN_STATUSES = (
    *OPEN_PLAN_STATUSES,
    FbsReplenishmentPlan.STATUS_AWAITING_PACK,
)
OPEN_ALLOCATION_STATUSES = (
    FbsReplenishmentAllocation.STATUS_RESERVED,
    FbsReplenishmentAllocation.STATUS_IN_PROGRESS,
)
ACTIVE_ORDER_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_VALIDATION_FAILED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
    FbsOrder.STATUS_PICKING,
    FbsOrder.STATUS_PICKED,
    FbsOrder.STATUS_READY_FOR_HANDOVER,
    FbsOrder.STATUS_EXCEPTION,
)
STOREKEEPER_WORK_ORDER_STATUSES = tuple(
    status for status in ACTIVE_ORDER_STATUSES if status != FbsOrder.STATUS_EXCEPTION
)
ORDER_ATTENTION_STALE_HOURS = 2
ORDER_ATTENTION_NOT_STARTED_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_VALIDATION_FAILED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_EXCEPTION,
)
ORDER_ATTENTION_NOT_PICKED_STATUSES = (
    *ORDER_ATTENTION_NOT_STARTED_STATUSES,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
    FbsOrder.STATUS_PICKING,
)
ACTIVE_PICK_BATCH_STATUSES = (
    FbsPickBatch.STATUS_QUEUED,
    FbsPickBatch.STATUS_IN_PROGRESS,
    FbsPickBatch.STATUS_VERIFICATION,
)
SEPARATED_PICK_RESTOCK_STATUSES = (
    FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
    FbsPickRestockRequest.STATUS_QUEUED,
    FbsPickRestockRequest.STATUS_IN_PROGRESS,
    FbsPickRestockRequest.STATUS_COMPLETED,
    FbsPickRestockRequest.STATUS_FAILED,
)
PICK_EQUIPMENT_SESSION_KEY = "fbs_tsd_pick_equipment"
FBS_STOCK_SORT_FIELDS = {
    "client": "agency__agn_name",
    "sku": "sku_code",
    "name": "name",
    "barcode": "barcode",
    "qty": "qty",
    "available": "available_qty",
    "reserved": "reserved_qty",
    "box": "box__box_code",
    "pallet": "box__pallet__pallet_code",
    "cell": "box__pallet__cell__cell_code",
    "expiry": "expiry_date",
    "updated": "updated_at",
}
FBS_STOCK_DESCENDING_DEFAULTS = {"qty", "available", "reserved", "updated"}
FBS_STOCK_BOX_SORT_FIELDS = {
    "client": "agency__agn_name",
    "sku": "group_sku",
    "name": "group_name",
    "barcode": "group_barcode",
    "qty": "group_qty",
    "available": "group_available_qty",
    "reserved": "group_reserved_qty",
    "box": "box__box_code",
    "pallet": "box__pallet__pallet_code",
    "cell": "box__pallet__cell__cell_code",
    "expiry": "group_expiry",
    "updated": "group_updated",
}


def fbs_module_required(view_func):
    @wraps(view_func)
    def wrapped(request, *args, **kwargs):
        if not module_enabled():
            raise Http404
        return view_func(request, *args, **kwargs)

    return wrapped


def _base_context(request, **extra):
    route_name = str(getattr(getattr(request, "resolver_match", None), "url_name", "") or "")
    role = get_request_role(request)
    if "controller" in route_name:
        operator_section = "controller"
    elif "storekeeper_stock" in route_name:
        operator_section = "stock"
    elif route_name == "tsd_storekeeper":
        operator_section = "overview"
    elif "inventory" in route_name:
        operator_section = "inventory"
    elif "handover_reports" in route_name:
        operator_section = "handover_reports"
    elif "handover" in route_name:
        operator_section = "handover"
    elif "demand" in route_name or "plan" in route_name:
        operator_section = "demand"
    elif "metadata" in route_name:
        operator_section = "metadata"
    elif "label" in route_name:
        operator_section = "labels"
    elif "picking" in route_name or "pick_" in route_name:
        operator_section = "waves"
    else:
        operator_section = ""
    context = {
        "employee": get_request_employee(request),
        "request_role": role,
        "workspace_shell": workspace_shell_enabled(role),
        "cabinet_url": resolve_cabinet_url(role),
        "warehouse_writes_enabled": feature_enabled("warehouse_writes"),
        "operator_section": operator_section,
    }
    context.update(extra)
    workspace_section = str(context.get("operator_section") or "")
    workspace_presentation = {
        "overview": (
            "Пульт смены FBS",
            "Всё, что требует решения кладовщика прямо сейчас",
            "Пульт смены",
        ),
        "stock": (
            "Остатки FBS",
            "Физический, доступный, зарезервированный и заблокированный",
            "Остатки FBS",
        ),
        "inventory": ("Инвентаризация", "", "Инвентаризация"),
        "handover": (
            "Отгрузки и короба",
            "Набор коробов, сканирование, передача водителю",
            "Отгрузки",
        ),
        "handover_reports": (
            "Отчёты по отгрузкам",
            "Отгрузки, заказы и товары отдельного контура FBS",
            "Отгрузки / Отчёты",
        ),
        "demand": ("Подсорт", "Что просит сборка", "Подсорт"),
        "labels": ("Этикетки", "", "Этикетки"),
    }
    workspace_title, workspace_subtitle, workspace_crumb = workspace_presentation.get(
        workspace_section,
        (str(context.get("page_title") or "FBS"), "", "FBS"),
    )
    context.setdefault("workspace_section", workspace_section)
    context.setdefault("workspace_title", workspace_title)
    context.setdefault("workspace_subtitle", workspace_subtitle)
    context.setdefault("workspace_crumb", workspace_crumb)
    return context


def _plan_queryset():
    lines = FbsReplenishmentLine.objects.select_related(
        "sku_ref",
        "source_container",
        "target_box",
    ).order_by("id")
    return (
        FbsReplenishmentPlan.objects.select_related(
            "agency",
            "target_cell__location",
            "target_pallet",
            "target_box",
            "staging_location",
            "assigned_to",
            "confirmed_by",
        )
        .prefetch_related(Prefetch("lines", queryset=lines))
        .order_by("-created_at", "-id")
    )


def _order_attention_started_before_q(moment) -> Q:
    return Q(ordered_at__lte=moment) | Q(
        ordered_at__isnull=True,
        imported_at__lte=moment,
    )


def _storekeeper_list_context(request, *, error=""):
    plan_summaries = FbsReplenishmentPlan.objects.select_related("agency").order_by(
        "-created_at",
        "-id",
    )
    now = timezone.now()
    today = timezone.localdate()
    active_orders_queryset = FbsOrder.objects.filter(
        internal_status__in=STOREKEEPER_WORK_ORDER_STATUSES
    ).exclude(marketplace_terminal_order_q())
    sla_counts = active_orders_queryset.aggregate(
        tracked=Count("id", filter=Q(cutoff_at__isnull=False)),
        on_time=Count("id", filter=Q(cutoff_at__gte=now)),
        urgent=Count(
            "id",
            filter=Q(cutoff_at__gte=now, cutoff_at__lte=now + timedelta(hours=2)),
        ),
        overdue=Count("id", filter=Q(cutoff_at__lt=now)),
        without_cutoff=Count("id", filter=Q(cutoff_at__isnull=True)),
    )
    tracked_count = int(sla_counts["tracked"] or 0)
    overdue_count = int(sla_counts["overdue"] or 0)
    sla_counts["risk_percent"] = (
        round(overdue_count * 100 / tracked_count) if tracked_count else 0
    )
    sla_counts["next_cutoff"] = (
        active_orders_queryset.filter(cutoff_at__gte=now)
        .order_by("cutoff_at", "id")
        .values_list("cutoff_at", flat=True)
        .first()
    )
    attention_orders_queryset = (
        FbsOrder.objects.filter(
            internal_status__in=ORDER_ATTENTION_NOT_PICKED_STATUSES,
        )
        .exclude(marketplace_terminal_order_q())
        .select_related("profile__agency")
        .annotate(
            has_active_pick_task=Exists(
                FbsPickTask.objects.filter(
                    order_id=OuterRef("pk"),
                    status__in=ACTIVE_TASK_STATUSES,
                )
            ),
        )
    )
    stale_attention_orders = attention_orders_queryset.filter(
        internal_status__in=ORDER_ATTENTION_NOT_STARTED_STATUSES,
        has_active_pick_task=False,
    ).filter(
        _order_attention_started_before_q(
            now - timedelta(hours=ORDER_ATTENTION_STALE_HOURS)
        )
    )
    local_midnight = timezone.localtime(now).replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )
    previous_day_attention_orders = attention_orders_queryset.filter(
        _order_attention_started_before_q(local_midnight)
    )
    stale_attention_counts = stale_attention_orders.aggregate(
        total=Count("id"),
        awaiting_stock=Count(
            "id",
            filter=Q(internal_status=FbsOrder.STATUS_AWAITING_STOCK),
        ),
        without_cutoff=Count("id", filter=Q(cutoff_at__isnull=True)),
    )
    order_attention_counts = {
        "stale_unstarted": int(stale_attention_counts["total"] or 0),
        "previous_day_unpicked": previous_day_attention_orders.count(),
        "awaiting_stock": int(stale_attention_counts["awaiting_stock"] or 0),
        "without_cutoff": int(stale_attention_counts["without_cutoff"] or 0),
    }
    daily_flow = FbsOrder.objects.aggregate(
        received=Count("id", filter=Q(imported_at__date=today)),
        dispatched=Count(
            "id",
            filter=Q(
                handover_order__status=FbsHandoverOrder.STATUS_ACTIVE,
                handover_order__box__batch__dispatched_at__date=today,
            ),
            distinct=True,
        ),
        accepted=Count(
            "id",
            filter=Q(
                handover_order__status=FbsHandoverOrder.STATUS_ACTIVE,
                handover_order__box__batch__accepted_at__date=today,
            ),
            distinct=True,
        ),
        canceled=Count(
            "id",
            filter=Q(
                internal_status=FbsOrder.STATUS_CANCELLED,
                updated_at__date=today,
            ),
        ),
    )
    active_batches_queryset = (
        FbsPickBatch.objects.select_related(
            "agency",
            "assigned_to",
            "workstation",
            "cart",
        )
        .filter(status__in=ACTIVE_PICK_BATCH_STATUSES)
        .annotate(order_count=Count("tasks"))
        .order_by("created_at", "id")
    )
    active_batch_count = active_batches_queryset.count()
    active_batches = list(active_batches_queryset[:8])

    open_exceptions = FbsPickException.objects.filter(
        status=FbsPickException.STATUS_OPEN
    )
    recent_exceptions = list(
        open_exceptions.select_related(
            "task__order__profile__agency",
            "allocation__order_item",
            "created_by",
        ).order_by("-created_at", "-id")[:6]
    )

    label_counts = FbsOrderLabel.objects.aggregate(
        requested=Count("id", filter=Q(status=FbsOrderLabel.STATUS_REQUESTED)),
        ready=Count("id", filter=Q(status=FbsOrderLabel.STATUS_READY)),
        errors=Count("id", filter=Q(status=FbsOrderLabel.STATUS_ERROR)),
    )
    handover_counts = FbsHandoverBox.objects.aggregate(
        waiting_scan=Count("id", filter=Q(status=FbsHandoverBox.STATUS_CLOSED)),
        ready=Count("id", filter=Q(status=FbsHandoverBox.STATUS_SCANNED)),
        problems=Count("id", filter=Q(status=FbsHandoverBox.STATUS_PROBLEM)),
    )
    metadata_counts = FbsMarketplaceMetadataTransfer.objects.aggregate(
        waiting=Count(
            "id",
            filter=Q(
                status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
                    FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
                    FbsMarketplaceMetadataTransfer.STATUS_SENT,
                    FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                )
            ),
        ),
        conflicts=Count(
            "id",
            filter=Q(
                status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                    FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                )
            ),
        ),
    )
    order_counts = FbsOrder.objects.exclude(marketplace_terminal_order_q()).aggregate(
        received=Count("id", filter=Q(internal_status=FbsOrder.STATUS_RECEIVED)),
        awaiting_stock=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_AWAITING_STOCK)
        ),
        reserved=Count("id", filter=Q(internal_status=FbsOrder.STATUS_RESERVED)),
        queued=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_QUEUED_FOR_PICK)
        ),
        picking=Count("id", filter=Q(internal_status=FbsOrder.STATUS_PICKING)),
        picked=Count("id", filter=Q(internal_status=FbsOrder.STATUS_PICKED)),
        ready_for_handover=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_READY_FOR_HANDOVER)
        ),
        validation_failed=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_VALIDATION_FAILED)
        ),
        exception=Count("id", filter=Q(internal_status=FbsOrder.STATUS_EXCEPTION)),
        active=Count(
            "id", filter=Q(internal_status__in=STOREKEEPER_WORK_ORDER_STATUSES)
        ),
        problem=Count(
            "id",
            filter=Q(
                internal_status__in=(
                    FbsOrder.STATUS_VALIDATION_FAILED,
                    FbsOrder.STATUS_EXCEPTION,
                )
            ),
        ),
    )

    workstations = list(
        FbsWorkstation.objects.filter(is_active=True)
        .select_related("device_agent")
        .order_by("name", "id")
    )
    online_threshold = now - timedelta(seconds=60)
    for workstation in workstations:
        agent = workstation.device_agent
        workstation.agent_online = bool(
            agent and agent.last_seen and agent.last_seen >= online_threshold
        )
    active_cart_count = FbsPickingCart.objects.filter(is_active=True).count()
    online_workstation_count = sum(
        1 for workstation in workstations if workstation.agent_online
    )

    open_exception_count = open_exceptions.count()
    problem_count = (
        int(order_counts["problem"] or 0)
        + open_exception_count
        + int(label_counts["errors"] or 0)
        + int(handover_counts["problems"] or 0)
        + int(metadata_counts["conflicts"] or 0)
    )
    profile_counts = FbsIntegrationProfile.objects.aggregate(
        total=Count("id"),
        active=Count("id", filter=Q(is_active=True)),
    )
    proposed_plans = plan_summaries.filter(
        status=FbsReplenishmentPlan.STATUS_PROPOSED
    )[:30]
    active_plans = plan_summaries.filter(
        status__in=STOREKEEPER_OPEN_PLAN_STATUSES
    )[:30]
    return _base_context(
        request,
        page_title="FBS · Смена",
        operator_section="overview",
        back_url="/sklad/",
        proposed_plans=proposed_plans,
        active_plans=active_plans,
        active_batches=active_batches,
        recent_exceptions=recent_exceptions,
        order_counts=order_counts,
        label_counts=label_counts,
        handover_counts=handover_counts,
        metadata_counts=metadata_counts,
        profile_counts=profile_counts,
        sla_counts=sla_counts,
        order_attention_counts=order_attention_counts,
        daily_flow=daily_flow,
        active_inventory_count=FbsInventorySession.objects.exclude(
            status__in=(
                FbsInventorySession.STATUS_DONE,
                FbsInventorySession.STATUS_CANCELED,
            )
        ).count(),
        equipment_summary={
            "workstations": len(workstations),
            "workstations_online": online_workstation_count,
            "carts": active_cart_count,
        },
        workstations=workstations[:4],
        dashboard_counts={
            "new": int(order_counts["received"] or 0)
            + int(order_counts["awaiting_stock"] or 0),
            "to_pick": int(order_counts["reserved"] or 0)
            + int(order_counts["queued"] or 0),
            "waves": active_batch_count,
            "problems": problem_count,
        },
        problem_breakdown={
            "orders": int(order_counts["problem"] or 0),
            "pick": open_exception_count,
            "labels": int(label_counts["errors"] or 0),
            "handover": int(handover_counts["problems"] or 0),
            "metadata": int(metadata_counts["conflicts"] or 0),
        },
        error=error,
    )


def _storekeeper_plan_context(request, plan, *, error="", ok_message=""):
    staged_qty = int(
        FbsReplenishmentAllocation.objects.filter(line__plan=plan).aggregate(
            total=Sum("qty_staged")
        )["total"]
        or 0
    )
    boxes = FbsBox.objects.filter(
        agency=plan.agency,
        status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        pallet__status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        pallet__cell__is_active=True,
    ).select_related("pallet__cell").order_by(
        "pallet__cell__cell_code", "pallet__pallet_code", "box_code", "id"
    )
    pallets = FbsPallet.objects.filter(
        agency=plan.agency,
        status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        cell__is_active=True,
    ).select_related("cell").annotate(
        active_box_count=Count(
            "boxes",
            filter=Q(boxes__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE)),
        )
    ).order_by("cell__cell_code", "pallet_code", "id")
    prepared_boxes = list(
        FbsReplenishmentPreparedBox.objects.filter(plan=plan)
        .select_related("fbs_box__pallet__cell")
        .order_by("sequence_no", "id")
    )
    return _base_context(
        request,
        page_title=f"План FBS #{plan.id}",
        back_url=reverse("fbs:tsd_storekeeper"),
        plan=plan,
        staging_container_code=(
            movement_staging_container_code(plan.client_movement_request_id)
            if plan.client_movement_request_id
            and plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and not prepared_boxes
            else ""
        ),
        staged_qty=staged_qty,
        can_pack=(
            not prepared_boxes
            and plan.status == FbsReplenishmentPlan.STATUS_AWAITING_PACK
        ),
        prepared_boxes=prepared_boxes,
        prepared_scanned_qty=sum(int(box.scanned_qty or 0) for box in prepared_boxes),
        prepared_placed_qty=sum(int(box.placed_qty or 0) for box in prepared_boxes),
        dynamic_destination=(
            plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and plan.status != FbsReplenishmentPlan.STATUS_DONE
        ),
        awaiting_placement=(
            not prepared_boxes
            and
            plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and plan.status == FbsReplenishmentPlan.STATUS_IN_PROGRESS
            and plan.target_box_id is not None
            and int(plan.moved_qty or 0) < int(plan.planned_qty or 0)
        ),
        boxes=boxes,
        pallets=pallets,
        error=error,
        ok_message=ok_message,
    )


def _demand_row_state(row) -> tuple[str, str, str]:
    if row.general_shortage_qty and row.proposed_qty:
        return "partial", "Частично готово", "conflict"
    if row.general_shortage_qty:
        return "blocked_stock", "Нет общего остатка", "failed"
    if row.destination_blocked_qty:
        return "blocked_destination", "Нет FBS-короба", "failed"
    if row.proposed_qty:
        return "ready", "Готово к плану", "confirmed"
    if row.open_plan_qty:
        return "in_plan", "Уже в планах", "in_progress"
    return "covered", "Потребность закрыта", "done"


def _demand_mapping_issue(item, valid_barcode_pairs) -> tuple[str, str]:
    requirements = item.requirements if isinstance(item.requirements, dict) else {}
    mapping_status = str(requirements.get("sku_mapping") or "").strip()
    imported_reason = DEMAND_MAPPING_STATUS_LABELS.get(mapping_status, "")
    barcode = str(item.barcode or "").strip()
    if item.sku_id is None:
        return mapping_status or "sku_missing", imported_reason or "SKU не привязан."
    if not barcode:
        return mapping_status or "barcode_missing", imported_reason or "ШК не указан."
    if item.sku.agency_id != item.order.profile.agency_id:
        return "foreign_sku", "SKU принадлежит другому клиенту."
    if item.sku.deleted:
        return "deleted_sku", "SKU удален из номенклатуры клиента."
    if (item.sku_id, barcode) not in valid_barcode_pairs:
        return "barcode_not_bound", "ШК не привязан к SKU клиента."
    return "", ""


def _demand_context(request, *, error="", ok_message=""):
    analysis = analyze_replenishment_demand()
    floor_analysis = analyze_floor_replenishment_needs()
    params = request.GET if request.method == "GET" else request.POST
    query = str(params.get("q") or "").strip()
    agency_filter = str(params.get("agency") or "").strip()
    state_filter = str(params.get("state") or "").strip()
    deadline_filter = str(params.get("deadline") or "").strip()
    mapping_items = list(
        FbsOrderItem.objects.filter(
            order__internal_status__in=DEMAND_MAPPING_ORDER_STATUSES,
        )
        .select_related("order__profile__agency", "sku")
        .order_by("order__cutoff_at", "order_id", "id")
    )
    agency_ids = {row.agency_id for row in analysis.rows} | {
        item.order.profile.agency_id for item in mapping_items
    }
    sku_ids = {row.sku_id for row in analysis.rows}
    analysis_barcodes = {
        str(getattr(row, "barcode", "") or "").strip()
        for row in analysis.rows
        if str(getattr(row, "barcode", "") or "").strip()
    }
    if not agency_filter.isdigit() or int(agency_filter) not in agency_ids:
        agency_filter = ""
    if state_filter not in {value for value, _ in DEMAND_STATE_CHOICES}:
        state_filter = ""
    if deadline_filter not in {"overdue", "urgent", "today", "none"}:
        deadline_filter = ""

    order_items = (
        FbsOrderItem.objects.filter(
            order__internal_status__in=DEMAND_DISPATCH_ORDER_STATUSES,
            order__profile__agency_id__in=agency_ids,
        )
        .filter(Q(sku_id__in=sku_ids) | Q(barcode__in=analysis_barcodes))
        .select_related("order__profile", "order__profile__agency")
        .order_by("order__cutoff_at", "order_id", "id")
    )
    orders_by_sku = defaultdict(dict)
    orders_by_barcode = defaultdict(dict)
    barcodes_by_sku = defaultdict(set)
    now = timezone.now()
    today = timezone.localdate()

    def add_order(target, key, item):
        entry = target[key].get(item.order_id)
        if entry is None:
            cutoff = item.order.cutoff_at
            entry = SimpleNamespace(
                id=item.order_id,
                external_order_id=item.order.external_order_id,
                marketplace=item.order.profile.get_marketplace_display(),
                quantity=0,
                cutoff_at=cutoff,
                is_overdue=bool(cutoff and cutoff < now),
                is_urgent=bool(cutoff and now <= cutoff <= now + timedelta(hours=2)),
                is_today=bool(cutoff and timezone.localtime(cutoff).date() == today),
            )
            target[key][item.order_id] = entry
        entry.quantity += int(item.quantity or 0)

    for item in order_items:
        agency_id = item.order.profile.agency_id
        barcode = str(item.barcode or "").strip()
        if item.sku_id:
            sku_key = (agency_id, item.sku_id)
            add_order(orders_by_sku, sku_key, item)
            if barcode:
                barcodes_by_sku[sku_key].add(barcode)
        if barcode:
            add_order(orders_by_barcode, (agency_id, barcode), item)

    mapping_sku_ids = {item.sku_id for item in mapping_items if item.sku_id}
    mapping_barcodes = {
        str(item.barcode or "").strip()
        for item in mapping_items
        if str(item.barcode or "").strip()
    }
    valid_mapping_pairs = set(
        SKUBarcode.objects.filter(
            sku_id__in=mapping_sku_ids,
            value__in=mapping_barcodes,
        ).values_list("sku_id", "value")
    )
    all_mapping_rows = []
    for item in mapping_items:
        reason_code, reason = _demand_mapping_issue(item, valid_mapping_pairs)
        if not reason_code:
            continue
        requirements = item.requirements if isinstance(item.requirements, dict) else {}
        marketplace_barcodes = [
            str(value or "").strip()
            for value in requirements.get("marketplace_barcodes", ())
            if str(value or "").strip()
        ]
        barcode = str(item.barcode or "").strip()
        display_barcode = barcode or ", ".join(marketplace_barcodes[:3]) or "Не передан"
        if len(marketplace_barcodes) > 3 and not barcode:
            display_barcode += f" +{len(marketplace_barcodes) - 3}"
        cutoff = item.order.cutoff_at
        all_mapping_rows.append(
            SimpleNamespace(
                id=item.id,
                order_id=item.order_id,
                external_order_id=item.order.external_order_id,
                agency_id=item.order.profile.agency_id,
                agency_name=str(item.order.profile.agency),
                marketplace=item.order.profile.get_marketplace_display(),
                external_sku=str(item.external_sku or "").strip() or "Без артикула",
                product_name=(
                    str(item.product_name or "").strip()
                    or str(getattr(item.sku, "name", "") or "").strip()
                    or "Товар не определен"
                ),
                display_barcode=display_barcode,
                quantity=int(item.quantity or 0),
                reason_code=reason_code,
                reason=reason,
                cutoff_at=cutoff,
                is_overdue=bool(cutoff and cutoff < now),
                is_urgent=bool(cutoff and now <= cutoff <= now + timedelta(hours=2)),
                is_today=bool(cutoff and timezone.localtime(cutoff).date() == today),
            )
        )

    mapping_rows = []
    for row in all_mapping_rows:
        if agency_filter and row.agency_id != int(agency_filter):
            continue
        query_folded = query.casefold()
        if query_folded and not any(
            query_folded in value.casefold()
            for value in (
                row.external_order_id,
                row.agency_name,
                row.external_sku,
                row.product_name,
                row.display_barcode,
                row.reason,
            )
        ):
            continue
        if state_filter not in {"", "action", "blocked", "mapping"}:
            continue
        if deadline_filter == "overdue" and not row.is_overdue:
            continue
        if deadline_filter == "urgent" and not row.is_urgent:
            continue
        if deadline_filter == "today" and not row.is_today:
            continue
        if deadline_filter == "none" and row.cutoff_at is not None:
            continue
        mapping_rows.append(row)
    mapping_rows.sort(
        key=lambda row: (
            0 if row.is_overdue else 1 if row.is_urgent else 2,
            row.cutoff_at is None,
            row.cutoff_at or now + timedelta(days=36500),
            row.agency_name.casefold(),
            row.external_order_id,
            row.id,
        )
    )

    demand_rows = []
    for row in analysis.rows:
        row_barcode = str(getattr(row, "barcode", "") or "").strip()
        sku_key = (row.agency_id, row.sku_id)
        if row_barcode:
            order_rows = list(
                orders_by_barcode[(row.agency_id, row_barcode)].values()
            )
            barcode_label = row_barcode
            order_search = row_barcode
        else:
            order_rows = list(orders_by_sku[sku_key].values())
            row_barcodes = sorted(barcodes_by_sku[sku_key])
            barcode_label = ", ".join(row_barcodes[:3]) or "не указан"
            if len(row_barcodes) > 3:
                barcode_label += f" +{len(row_barcodes) - 3}"
            order_search = row.sku_code
        order_rows.sort(
            key=lambda order: (
                order.cutoff_at is None,
                order.cutoff_at or now + timedelta(days=36500),
                order.id,
            )
        )
        if deadline_filter == "overdue":
            visible_order_rows = [order for order in order_rows if order.is_overdue]
        elif deadline_filter == "urgent":
            visible_order_rows = [order for order in order_rows if order.is_urgent]
        elif deadline_filter == "today":
            visible_order_rows = [order for order in order_rows if order.is_today]
        elif deadline_filter == "none":
            visible_order_rows = [
                order for order in order_rows if order.cutoff_at is None
            ]
        else:
            visible_order_rows = order_rows
        state, state_label, state_css = _demand_row_state(row)
        nearest_cutoff = next(
            (
                order.cutoff_at
                for order in visible_order_rows
                if order.cutoff_at is not None
            ),
            None,
        )
        decorated = SimpleNamespace(
            **row.__dict__,
            barcode_label=barcode_label,
            orders=visible_order_rows[:8],
            order_ids={order.id for order in visible_order_rows},
            order_count=len(visible_order_rows),
            hidden_order_count=max(len(visible_order_rows) - 8, 0),
            overdue_order_ids={
                order.id for order in visible_order_rows if order.is_overdue
            },
            urgent_order_ids={
                order.id for order in visible_order_rows if order.is_urgent
            },
            today_order_ids={
                order.id for order in visible_order_rows if order.is_today
            },
            no_cutoff_order_ids={
                order.id for order in visible_order_rows if order.cutoff_at is None
            },
            nearest_cutoff=nearest_cutoff,
            state=state,
            state_label=state_label,
            state_css=state_css,
            orders_url=(
                reverse("fbs:operator_orders")
                + "?"
                + urlencode(
                    {"agency": row.agency_id, "q": order_search, "page_size": 100}
                )
            ),
        )
        if agency_filter and row.agency_id != int(agency_filter):
            continue
        query_folded = query.casefold()
        if query_folded and not any(
            query_folded in value.casefold()
            for value in (
                row.agency_name,
                row.sku_code,
                row.sku_name,
                barcode_label,
                " ".join(order.external_order_id for order in order_rows),
            )
        ):
            continue
        if state_filter == "action" and not (
            row.proposed_qty or row.general_shortage_qty or row.destination_blocked_qty
        ):
            continue
        if state_filter == "mapping":
            continue
        if state_filter == "ready" and not (
            row.proposed_qty
            and not row.general_shortage_qty
            and not row.destination_blocked_qty
        ):
            continue
        if state_filter == "in_plan" and not row.open_plan_qty:
            continue
        if state_filter == "blocked" and not (
            row.general_shortage_qty or row.destination_blocked_qty
        ):
            continue
        if state_filter == "covered" and row.uncovered_qty:
            continue
        if deadline_filter and not decorated.order_ids:
            continue
        demand_rows.append(decorated)

    state_priority = {
        "blocked_stock": 0,
        "blocked_destination": 0,
        "partial": 1,
        "ready": 2,
        "in_plan": 3,
        "covered": 4,
    }
    demand_rows.sort(
        key=lambda row: (
            0 if row.overdue_order_ids else 1 if row.urgent_order_ids else 2,
            state_priority[row.state],
            row.nearest_cutoff is None,
            row.nearest_cutoff or now + timedelta(days=36500),
            row.agency_name.casefold(),
            row.sku_code.casefold(),
            row.barcode_label,
        )
    )
    affected_order_ids = (
        set().union(*(row.order_ids for row in demand_rows)) if demand_rows else set()
    )
    overdue_order_ids = (
        set().union(*(row.overdue_order_ids for row in demand_rows)) if demand_rows else set()
    )
    urgent_order_ids = (
        set().union(*(row.urgent_order_ids for row in demand_rows)) if demand_rows else set()
    )
    return _base_context(
        request,
        page_title="FBS · Потребность",
        back_url=reverse("fbs:tsd_storekeeper"),
        analysis=analysis,
        floor_analysis=floor_analysis,
        demand_rows=demand_rows,
        demand_agencies=Agency.objects.filter(id__in=agency_ids).order_by("agn_name", "id"),
        demand_state_choices=DEMAND_STATE_CHOICES,
        demand_filters={
            "q": query,
            "agency": agency_filter,
            "state": state_filter,
            "deadline": deadline_filter,
        },
        dispatch_summary={
            "rows": len(demand_rows),
            "orders": len(affected_order_ids),
            "overdue": len(overdue_order_ids),
            "urgent": len(urgent_order_ids),
        },
        mapping_rows=mapping_rows[:100],
        mapping_summary={
            "total_lines": len(all_mapping_rows),
            "total_orders": len({row.order_id for row in all_mapping_rows}),
            "total_qty": sum(row.quantity for row in all_mapping_rows),
            "filtered_lines": len(mapping_rows),
            "filtered_orders": len({row.order_id for row in mapping_rows}),
            "filtered_qty": sum(row.quantity for row in mapping_rows),
            "hidden_lines": max(len(mapping_rows) - 100, 0),
        },
        error=error,
        ok_message=ok_message,
    )


def _reachtruck_list_context(request, *, error="", ok_message=""):
    plans = (
        _plan_queryset()
        .filter(status__in=OPEN_PLAN_STATUSES)
        .filter(Q(assigned_to__isnull=True) | Q(assigned_to=request.user))
    )
    completed = _plan_queryset().filter(
        status=FbsReplenishmentPlan.STATUS_DONE,
        assigned_to=request.user,
    )[:10]
    return _base_context(
        request,
        page_title="FBS · Задания",
        back_url="/reachtruck/",
        available_plans=plans[:30],
        available_movements=FbsInternalMovement.objects.select_related(
            "agency",
            "source_box",
            "target_pallet__cell",
            "target_box",
            "assigned_to",
        ).filter(
            status__in=(
                FbsInternalMovement.STATUS_PROPOSED,
                FbsInternalMovement.STATUS_IN_PROGRESS,
            )
        ).filter(
            Q(assigned_to__isnull=True) | Q(assigned_to=request.user)
        ).order_by("created_at", "id")[:30],
        completed_plans=completed,
        error=error,
        ok_message=ok_message,
    )


def _allocation_queryset():
    return FbsReplenishmentAllocation.objects.select_related(
        "line__plan__agency",
        "line__plan__assigned_to",
        "line__plan__target_cell__location",
        "line__plan__target_pallet",
        "line__source_container",
        "source_snapshot__container",
        "source_snapshot__location",
        "source_snapshot__sku_ref",
        "target_box",
        "warehouse_task",
    )


def _pick_task_queryset():
    return FbsPickTask.objects.select_related(
        "batch__agency",
        "order__profile__agency",
        "assigned_to",
    ).order_by("batch_id", "sort_order", "id")


def _pick_allocation_queryset():
    return FbsOrderStockAllocation.objects.select_related(
        "pick_task__batch__agency",
        "pick_task__order__profile__agency",
        "pick_task__assigned_to",
        "order_item__sku",
        "order_item__order__profile",
        "balance__box__pallet__cell__location",
        "balance__box__source_container__current_location",
        "balance__sku_ref",
        "traceability",
        "verification_progress",
    )


def _picker_cart(request):
    stored = request.session.get(PICK_EQUIPMENT_SESSION_KEY)
    if not isinstance(stored, dict):
        return None
    try:
        cart_id = int(stored.get("cart_id"))
    except (TypeError, ValueError):
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        return None
    cart = FbsPickingCart.objects.filter(pk=cart_id, is_active=True).first()
    if cart is None:
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        return None
    return cart


def _store_picker_cart(request, cart) -> None:
    request.session[PICK_EQUIPMENT_SESSION_KEY] = {
        "cart_id": cart.id,
    }


def _open_pick_allocations(batch):
    return (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            status__in=(
                FbsOrderStockAllocation.STATUS_RESERVED,
                FbsOrderStockAllocation.STATUS_PICKING,
            ),
        )
        .order_by(
            "balance__box__pallet__cell__location__row_no",
            "balance__box__pallet__cell__location__section_no",
            "balance__box__pallet__cell__location__tier_no",
            "balance__box__pallet__cell__location__cell_no",
            "balance__box__box_code",
            "balance__sku_code",
            "pick_task__sort_order",
            "id",
        )
    )


def _next_pick_allocation_id(batch) -> int | None:
    return _open_pick_allocations(batch).values_list("id", flat=True).first()


def _picking_context(request, *, error="", ok_message=""):
    now = timezone.now()
    overdue_boundary = now - timedelta(hours=15)
    batches = (
        FbsPickBatch.objects.select_related(
            "agency",
            "assigned_to",
            "workstation",
            "cart",
        )
        .filter(
            status__in=(
                FbsPickBatch.STATUS_QUEUED,
                FbsPickBatch.STATUS_IN_PROGRESS,
                FbsPickBatch.STATUS_VERIFICATION,
            )
        )
        .filter(Q(assigned_to__isnull=True) | Q(assigned_to=request.user))
        .exclude(pick_restock_request__status__in=ACTIVE_PICK_RESTOCK_STATUSES)
        .annotate(
            oldest_order_at=Min(
                Coalesce("tasks__order__ordered_at", "tasks__order__imported_at")
            )
        )
        .order_by("created_at", "id")
    )
    my_active_batches = list(
        batches.filter(assigned_to=request.user).order_by("created_at", "id")
    )
    free_batches = list(
        batches.filter(assigned_to__isnull=True).order_by("created_at", "id")[:50]
    )
    available_batches = [*my_active_batches, *free_batches]
    for batch in available_batches:
        batch.order_count = batch.tasks.count()
        batch.next_allocation_id = _next_pick_allocation_id(batch)
        batch.is_overdue = bool(
            batch.oldest_order_at
            and batch.oldest_order_at <= overdue_boundary
        )
        batch.overdue_age_hours = (
            int((now - batch.oldest_order_at).total_seconds() // 3600)
            if batch.is_overdue
            else 0
        )
    cart = _picker_cart(request)
    free_batch_count = FbsPickBatch.objects.filter(
        status=FbsPickBatch.STATUS_QUEUED,
        assigned_to__isnull=True,
    ).count()
    active_wave_count = FbsPickBatch.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsPickBatch.STATUS_IN_PROGRESS,
            FbsPickBatch.STATUS_VERIFICATION,
        ),
        picking_completed_at__isnull=True,
    ).count()
    cart_busy = bool(
        cart
        and FbsPickBatch.objects.filter(
            cart=cart,
            status__in=(
                FbsPickBatch.STATUS_IN_PROGRESS,
                FbsPickBatch.STATUS_VERIFICATION,
            ),
            cart_released_at__isnull=True,
        ).exists()
    )
    request_role = get_request_role(request)
    restock_requests = list(
        FbsPickRestockRequest.objects.select_related(
            "batch__workstation",
            "assigned_to",
            "source_tote",
            "quarantine_box__pallet__cell__location",
        )
        .filter(
            Q(
                status=FbsPickRestockRequest.STATUS_QUEUED,
                assigned_to__isnull=True,
            )
            | Q(
                status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
                assigned_to=request.user,
            )
        )
        .filter(Q(order__isnull=True) | Q(source_tote__isnull=False))
        .annotate(source_box_count=Count("lines__source_box", distinct=True))
        .order_by("created_at", "id")[:20]
    )
    overdue_unassembled = (
        FbsOrder.objects.select_related("profile__agency")
        .filter(
            profile__is_active=True,
            internal_status__in=(
                FbsOrder.STATUS_RECEIVED,
                FbsOrder.STATUS_VALIDATION_FAILED,
                FbsOrder.STATUS_AWAITING_STOCK,
                FbsOrder.STATUS_RESERVED,
                FbsOrder.STATUS_QUEUED_FOR_PICK,
                FbsOrder.STATUS_PICKING,
            ),
        )
        .annotate(age_started_at=Coalesce("ordered_at", "imported_at"))
        .filter(age_started_at__lte=overdue_boundary)
        .order_by("age_started_at", "id")
    )
    overdue_oldest = overdue_unassembled.first()
    overdue_client_rows = list(
        overdue_unassembled.annotate(
            client_name=F("profile__agency__agn_name")
        )
        .values("client_name")
        .annotate(order_count=Count("id"))
        .order_by("-order_count", "client_name")
    )
    return _base_context(
        request,
        page_title="Сборка FBS" if request_role == "picker" else "FBS · Отбор",
        back_url=(
            reverse("fbs:tsd_home")
            if request_role == "picker"
            else "/sklad/" if request_role in STOREKEEPER_ROLES else ""
        ),
        available_batches=available_batches,
        my_active_batches=my_active_batches,
        free_batch_preview=free_batches[:3],
        completed_batches=FbsPickBatch.objects.select_related("agency").filter(
            status=FbsPickBatch.STATUS_DONE,
            assigned_to=request.user,
        ).order_by("-completed_at", "-id")[:10],
        received_count=FbsOrder.objects.filter(
            internal_status__in=(
                FbsOrder.STATUS_RECEIVED,
                FbsOrder.STATUS_AWAITING_STOCK,
            )
        ).count(),
        reserved_count=FbsOrder.objects.filter(
            internal_status=FbsOrder.STATUS_RESERVED
        ).count(),
        can_prepare_queue=(
            request_role in STOREKEEPER_ROLES
            or request_role == "developer"
        ),
        picker_cart=cart,
        free_batch_count=free_batch_count,
        active_wave_count=active_wave_count,
        picker_cart_busy=cart_busy,
        restock_requests=restock_requests,
        restock_request_count=len(restock_requests),
        overdue_unassembled_count=overdue_unassembled.count(),
        overdue_unassembled_oldest_hours=(
            int((now - overdue_oldest.age_started_at).total_seconds() // 3600)
            if overdue_oldest is not None
            else 0
        ),
        overdue_unassembled_clients=overdue_client_rows,
        can_claim_next=bool(
            cart
            and free_batch_count
            and not cart_busy
        ),
        error=error,
        ok_message=ok_message,
    )


def _picker_home_context(request):
    active_wave_count = FbsPickBatch.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsPickBatch.STATUS_IN_PROGRESS,
            FbsPickBatch.STATUS_VERIFICATION,
        ),
        picking_completed_at__isnull=True,
    ).count()
    free_batch_count = FbsPickBatch.objects.filter(
        status=FbsPickBatch.STATUS_QUEUED,
        assigned_to__isnull=True,
    ).count()
    restock_request_count = FbsPickRestockRequest.objects.filter(
        Q(
            status=FbsPickRestockRequest.STATUS_QUEUED,
            assigned_to__isnull=True,
        )
        | Q(
            status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
            assigned_to=request.user,
        )
    ).filter(Q(order__isnull=True) | Q(source_tote__isnull=False)).count()
    movement_task_count = FbsInternalMovement.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsInternalMovement.STATUS_PROPOSED,
            FbsInternalMovement.STATUS_IN_PROGRESS,
        ),
    ).count()
    return _base_context(
        request,
        page_title="Главное меню",
        back_url="",
        active_wave_count=active_wave_count,
        free_batch_count=free_batch_count,
        task_menu_count=active_wave_count + free_batch_count,
        restock_request_count=restock_request_count,
        movement_task_count=movement_task_count,
    )


def _normalized_picker_location_scan(value):
    normalized = unicodedata.normalize("NFKC", str(value or "")).strip().upper()
    normalized = normalized.translate(
        str.maketrans(
            {
                "А": "A",
                "В": "B",
                "С": "C",
                "Е": "E",
                "Н": "H",
                "К": "K",
                "М": "M",
                "О": "O",
                "Р": "P",
                "Т": "T",
                "Х": "X",
                "У": "Y",
            }
        )
    )
    normalized = re.sub(r"^FBS\s*[@:/-]?\s*", "", normalized)
    return "-".join(re.findall(r"[A-Z]+|\d+", normalized))


def _picker_location_candidates(cell):
    location = cell.location
    candidates = {
        str(cell.cell_code or ""),
        str(location.location_code or ""),
    }
    if all(
        int(getattr(location, field, 0) or 0) > 0
        for field in ("row_no", "section_no", "tier_no", "cell_no")
    ):
        candidates.add(
            os_location_code(
                row=location.row_no,
                section=location.section_no,
                tier=location.tier_no,
                cell=location.cell_no,
            )
        )
    return {
        _normalized_picker_location_scan(candidate)
        for candidate in candidates
        if candidate
    }


def _picker_location_label(cell):
    location = cell.location
    if all(
        int(getattr(location, field, 0) or 0) > 0
        for field in ("row_no", "section_no", "tier_no", "cell_no")
    ):
        return os_location_label(
            row=location.row_no,
            section=location.section_no,
            tier=location.tier_no,
            cell=location.cell_no,
        )
    return str(location.display_name or location.location_code or cell.cell_code)


def _pick_box_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_pick_box_{allocation_id}"


def _pick_cell_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_pick_cell_{allocation_id}"


def _pick_location_label(cell) -> str:
    location = getattr(cell, "location", None)
    if location is not None and str(location.zone_code or "").strip().upper() == "OS":
        coordinates = (
            int(location.row_no or 0),
            int(location.section_no or 0),
            int(location.tier_no or 0),
            int(location.cell_no or 0),
        )
        if all(value > 0 for value in coordinates):
            return os_location_label(
                row=coordinates[0],
                section=coordinates[1],
                tier=coordinates[2],
                cell=coordinates[3],
            )
    return str(
        (location.display_name if location else "")
        or (location.location_code if location else "")
        or cell.cell_code
    )


def _pick_box_number_suffix(box_code: str, *, length: int = 4) -> str:
    digits = "".join(re.findall(r"\d", str(box_code or "")))
    return digits[-length:] if digits else ""


def _pick_allocation_context(
    request,
    allocation,
    *,
    error="",
    ok_message="",
    scan_error="",
):
    task = allocation.pick_task
    open_count = FbsOrderStockAllocation.objects.filter(
        pick_task=task,
        status__in=(
            FbsOrderStockAllocation.STATUS_RESERVED,
            FbsOrderStockAllocation.STATUS_PICKING,
        ),
    ).count()
    balance = allocation.balance
    expected_item_code = balance.barcode or allocation.order_item.barcode
    photo = balance.sku_ref.photos.order_by("sort_order", "id").first() if balance.sku_ref_id else None
    batch = task.batch
    pick_location_label = fbs_box_physical_location_label(balance.box)
    box_progress = FbsOrderStockAllocation.objects.filter(
        pick_task__batch=batch,
        balance__box_id=balance.box_id,
    ).aggregate(
        planned_qty=Sum("qty_reserved"),
        picked_qty=Sum("qty_picked"),
        order_count=Count("pick_task__order_id", distinct=True),
    )
    box_planned_qty = int(box_progress["planned_qty"] or 0)
    box_picked_qty = int(box_progress["picked_qty"] or 0)
    return _base_context(
        request,
        page_title=f"FBS · Заказ {task.order.external_order_id}",
        back_url=reverse("fbs:tsd_picking"),
        allocation=allocation,
        task=task,
        order=task.order,
        balance=balance,
        open_count=open_count,
        cell_confirmed=bool(request.session.get(_pick_cell_session_key(allocation.id))),
        box_confirmed=bool(request.session.get(_pick_box_session_key(allocation.id))),
        allocation_remaining=max(
            int(allocation.qty_reserved or 0) - int(allocation.qty_picked or 0), 0
        ),
        missing_group_qty=pick_missing_group_quantity(
            allocation_id=allocation.id,
            assigned_to=request.user,
        ),
        position_picked_qty=int(allocation.qty_picked or 0),
        source_box_code=balance.box.box_code,
        box_planned_qty=box_planned_qty,
        box_picked_qty=box_picked_qty,
        box_remaining_qty=max(box_planned_qty - box_picked_qty, 0),
        box_order_count=int(box_progress["order_count"] or 0),
        wave_picked_qty=int(batch.picked_qty or 0),
        wave_planned_qty=int(batch.planned_qty or 0),
        recent_scan_events=FbsPickScanEvent.objects.filter(batch=batch)
        .select_related("created_by")
        .order_by("-created_at", "-id")[:12],
        expected_item_code=expected_item_code,
        pick_location_label=pick_location_label,
        item_scan_label="Штрихкод товара",
        product_photo_url=photo.url if photo else "",
        scan_error=scan_error,
        error=error,
        ok_message=ok_message,
    )


def _next_verification_allocation(batch):
    return (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            pick_task__status=FbsPickTask.STATUS_PICKED,
            status=FbsOrderStockAllocation.STATUS_PICKED,
        )
        .exclude(
            pick_task__order__pick_restock_requests__status__in=(
                SEPARATED_PICK_RESTOCK_STATUSES
            )
        )
        .filter(
            Q(verification_progress__isnull=True)
            | Q(verification_progress__qty_verified__lt=F("qty_picked"))
        )
        .order_by("pick_task__sort_order", "pick_task_id", "id")
        .first()
    )


def _next_verification_label(batch):
    tasks = list(
        batch.tasks.filter(status=FbsPickTask.STATUS_PICKED)
        .exclude(
            order__pick_restock_requests__status__in=SEPARATED_PICK_RESTOCK_STATUSES
        )
        .annotate(
            ui_picked_qty=Sum("allocations__qty_picked"),
            ui_verified_qty=Sum("allocations__verification_progress__qty_verified"),
        )
        .order_by("sort_order", "id")
    )
    if not tasks:
        return None
    active_pick_context = (
        FbsControllerPickTote.objects.filter(
            pick_batch=batch,
            status=FbsControllerPickTote.STATUS_PROCESSING,
        )
        .order_by("id")
        .first()
    )
    current_tote_order_ids = set()
    if active_pick_context is not None:
        current_tote_order_ids = set(
            active_pick_context.orders.exclude(
                status=FbsControllerToteOrder.STATUS_REMOVED,
            ).values_list("order_id", flat=True)
        )
    scanned_label_order_ids = FbsPickScanEvent.objects.filter(
        batch=batch,
        task__isnull=False,
        stage=FbsPickScanEvent.STAGE_ORDER_LABEL,
        result=FbsPickScanEvent.RESULT_SUCCESS,
    )
    if batch.verification_started_at is not None:
        scanned_label_order_ids = scanned_label_order_ids.filter(
            created_at__gte=batch.verification_started_at,
        )
    scanned_label_order_ids = set(
        scanned_label_order_ids.values_list("task__order_id", flat=True)
    )
    latest_labels = {}
    for label in (
        FbsOrderLabel.objects.filter(
            order_id__in=[task.order_id for task in tasks]
        )
        .exclude(status=FbsOrderLabel.STATUS_CANCELED)
        .order_by("order_id", "-requested_at", "-id")
    ):
        latest_labels.setdefault(label.order_id, label)
    for task in tasks:
        if task.order_id in scanned_label_order_ids:
            continue
        picked_qty = int(task.ui_picked_qty or 0)
        verified_qty = int(task.ui_verified_qty or 0)
        label = latest_labels.get(task.order_id)
        if (
            picked_qty
            and verified_qty >= picked_qty
            and label is not None
            and (
                active_pick_context is None
                or task.order_id not in current_tote_order_ids
            )
        ):
            return label
    return None


def _next_incomplete_verification_allocation(batch):
    return (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            pick_task__status=FbsPickTask.STATUS_PICKED,
            status=FbsOrderStockAllocation.STATUS_PICKED,
            qty_picked__gt=0,
            pick_task__exceptions__exception_type=FbsPickException.TYPE_NOT_FOUND,
            pick_task__exceptions__status=FbsPickException.STATUS_OPEN,
        )
        .exclude(
            pick_task__order__pick_restock_requests__status__in=(
                SEPARATED_PICK_RESTOCK_STATUSES
            )
        )
        .order_by("pick_task__sort_order", "pick_task_id", "id")
        .first()
    )


def _open_pick_shortage_issue(task):
    return (
        FbsPickException.objects.filter(
            task=task,
            exception_type=FbsPickException.TYPE_NOT_FOUND,
            status=FbsPickException.STATUS_OPEN,
        )
        .order_by("id")
        .first()
    )


def _verification_order_allocations(task, *, current_allocation_id=None):
    rows = list(
        _pick_allocation_queryset()
        .filter(pick_task=task)
        .order_by("id")
    )
    for row in rows:
        row_progress = getattr(row, "verification_progress", None)
        row.verified_qty = int(getattr(row_progress, "qty_verified", 0) or 0)
        row.is_verified = row.verified_qty >= int(row.qty_picked or 0)
        row.is_current = row.id == current_allocation_id
    return rows


def _ozon_multi_item_hold_prompt(allocation):
    """Describe the remaining composition after one accepted Ozon unit scan."""
    task = allocation.pick_task
    if (
        task is None
        or task.order.profile.marketplace
        != FbsIntegrationProfile.MARKETPLACE_OZON
    ):
        return None
    rows = _verification_order_allocations(
        task,
        current_allocation_id=allocation.id,
    )
    planned_qty = sum(int(row.qty_picked or 0) for row in rows)
    verified_qty = sum(int(row.verified_qty or 0) for row in rows)
    remaining_qty = max(planned_qty - verified_qty, 0)
    if planned_qty <= 1 or remaining_qty <= 0:
        return None

    remaining_by_item = {}
    for row in rows:
        remaining = max(int(row.qty_picked or 0) - int(row.verified_qty or 0), 0)
        if not remaining:
            continue
        item = row.order_item
        item_summary = remaining_by_item.setdefault(
            item.id,
            {
                "article": str(item.external_sku or row.balance.sku_code or "").strip(),
                "name": str(item.product_name or row.balance.name or "Товар").strip(),
                "barcode": str(item.barcode or row.balance.barcode or "").strip(),
                "quantity": 0,
            },
        )
        item_summary["quantity"] += remaining

    scanned_item = allocation.order_item
    return SimpleNamespace(
        order_number=task.order.external_order_id,
        planned_qty=planned_qty,
        verified_qty=verified_qty,
        remaining_qty=remaining_qty,
        scanned_article=str(
            scanned_item.external_sku or allocation.balance.sku_code or ""
        ).strip(),
        scanned_name=str(
            scanned_item.product_name or allocation.balance.name or "Товар"
        ).strip(),
        scanned_barcode=str(
            scanned_item.barcode or allocation.balance.barcode or ""
        ).strip(),
        remaining_items=[SimpleNamespace(**item) for item in remaining_by_item.values()],
    )


def _verification_wave_rows(batch, *, current_task_id=None):
    tasks = list(
        batch.tasks.filter(status=FbsPickTask.STATUS_PICKED)
        .exclude(
            order__pick_restock_requests__status__in=SEPARATED_PICK_RESTOCK_STATUSES
        )
        .select_related("order")
        .annotate(
            ui_verified_qty=Sum("allocations__verification_progress__qty_verified")
        )
        .order_by("sort_order", "id")
    )
    order_ids = [task.order_id for task in tasks]
    active_pick_context = (
        FbsControllerPickTote.objects.filter(
            pick_batch=batch,
            status=FbsControllerPickTote.STATUS_PROCESSING,
        )
        .order_by("id")
        .first()
    )
    current_tote_order_ids = set()
    if active_pick_context is not None:
        current_tote_order_ids = set(
            active_pick_context.orders.exclude(
                status=FbsControllerToteOrder.STATUS_REMOVED,
            ).values_list("order_id", flat=True)
        )
    shortage_issues = {}
    for issue in FbsPickException.objects.filter(
        task_id__in=[task.id for task in tasks],
        exception_type=FbsPickException.TYPE_NOT_FOUND,
        status=FbsPickException.STATUS_OPEN,
    ).order_by("task_id", "id"):
        shortage_issues.setdefault(issue.task_id, issue)
    task_barcodes = defaultdict(list)
    for task_id, balance_barcode, item_barcode in (
        FbsOrderStockAllocation.objects.filter(
            pick_task_id__in=[task.id for task in tasks],
            qty_picked__gt=0,
        )
        .values_list("pick_task_id", "balance__barcode", "order_item__barcode")
        .order_by("pick_task_id", "id")
    ):
        barcode = str(balance_barcode or item_barcode or "").strip()
        if barcode and barcode not in task_barcodes[task_id]:
            task_barcodes[task_id].append(barcode)
    latest_labels = {}
    for label in (
        FbsOrderLabel.objects.filter(order_id__in=order_ids)
        .exclude(status=FbsOrderLabel.STATUS_CANCELED)
        .order_by("order_id", "-requested_at", "-id")
    ):
        latest_labels.setdefault(label.order_id, label)
    handover_links = {
        link.order_id: link
        for link in FbsHandoverOrder.objects.filter(
            order_id__in=order_ids,
            status=FbsHandoverOrder.STATUS_ACTIVE,
        )
        .select_related("box")
        .order_by("id")
    }
    handover_assignments = {
        assignment.order_id: assignment
        for assignment in FbsHandoverOrderAssignment.objects.filter(
            order_id__in=order_ids,
        )
        .exclude(status=FbsHandoverOrderAssignment.STATUS_CANCELED)
        .select_related("batch")
        .order_by("id")
    }
    target_handover_batch = (
        FbsHandoverBatch.objects.filter(
            controller_check_tote__pick_totes__pick_batch=batch,
        )
        .order_by("id")
        .first()
    )

    rows = []
    for task in tasks:
        shortage_issue = shortage_issues.get(task.id)
        is_incomplete = shortage_issue is not None
        verified_qty = int(task.ui_verified_qty or 0)
        planned_qty = int(task.planned_qty or 0)
        label = latest_labels.get(task.order_id)
        handover_link = handover_links.get(task.order_id)
        handover_assignment = handover_assignments.get(task.order_id)
        row_handover_batch = (
            handover_assignment.batch
            if handover_assignment is not None
            else target_handover_batch
        )
        product_complete = bool(planned_qty and verified_qty >= planned_qty)
        label_complete = bool(
            label is not None and label.status == FbsOrderLabel.STATUS_APPLIED
            and (
                active_pick_context is None
                or task.order_id in current_tote_order_ids
            )
        )
        handover_added = handover_assignment is not None
        if is_incomplete:
            state = "problem"
            state_label = "Недокомплект · передать в проблемные"
        elif handover_link is not None:
            state = "done"
            state_label = f"В коробе {handover_link.box.qr_code}"
        elif (
            task.order_id in current_tote_order_ids
            and label is not None
            and label.status == FbsOrderLabel.STATUS_REQUESTED
        ):
            state = "in_progress"
            state_label = "Заказ закреплён · Ozon проверяет"
        elif (
            task.order_id in current_tote_order_ids
            and label is not None
            and label.status == FbsOrderLabel.STATUS_READY
        ):
            state = "in_progress"
            state_label = "Этикетка Ozon готова · ожидает упаковки"
        elif label_complete and handover_assignment is not None:
            state = "done"
            state_label = f"В отгрузке #{handover_assignment.batch_id} · короб ожидается"
        elif label_complete:
            state = "in_progress"
            state_label = "QR отсканирован · отгрузка создается"
        elif label is not None and label.status in {
            FbsOrderLabel.STATUS_READY,
            FbsOrderLabel.STATUS_APPLIED,
        }:
            state = "in_progress"
            state_label = "QR готов, требуется скан"
        elif label is not None and label.status == FbsOrderLabel.STATUS_REQUESTED:
            state = "in_progress"
            state_label = "Готовится QR заказа"
        elif label is not None and label.status == FbsOrderLabel.STATUS_ERROR:
            state = "problem"
            state_label = "Ошибка этикетки"
        elif product_complete:
            state = "in_progress"
            state_label = "QR заказа не создан"
        elif verified_qty:
            state = "in_progress"
            state_label = f"Осталось проверить {max(planned_qty - verified_qty, 0)} шт."
        else:
            state = "pending"
            state_label = f"Остался товар {planned_qty} шт."
        rows.append(
            SimpleNamespace(
                task=task,
                order=task.order,
                verified_qty=verified_qty,
                planned_qty=planned_qty,
                label=label,
                handover_link=handover_link,
                handover_assignment=handover_assignment,
                target_handover_batch=row_handover_batch,
                is_product_complete=product_complete,
                is_label_complete=label_complete,
                is_handover_added=handover_added,
                is_verification_complete=bool(
                    product_complete
                    and label_complete
                    and handover_added
                    and not is_incomplete
                ),
                is_incomplete=is_incomplete,
                problem_reason=(
                    str(task.order.problem_reason or "").strip()
                    if is_incomplete
                    else ""
                ),
                state=state,
                state_label=state_label,
                is_current=task.id == current_task_id,
                item_barcodes=tuple(task_barcodes.get(task.id, ())),
            )
        )
    return rows


def _verification_wave_context(batch, *, current_task_id=None):
    rows = _verification_wave_rows(batch, current_task_id=current_task_id)
    pending_rows = [row for row in rows if not row.is_verification_complete]
    handover_batches = []
    seen_handover_ids = set()
    for row in rows:
        handover_batch = row.target_handover_batch
        if handover_batch is None or handover_batch.id in seen_handover_ids:
            continue
        seen_handover_ids.add(handover_batch.id)
        handover_batches.append(handover_batch)
    return {
        "verification_wave_rows": rows,
        "verification_pending_rows": pending_rows,
        "verification_pending_count": len(pending_rows),
        "verification_handover_batches": handover_batches,
        "verification_primary_handover": handover_batches[0] if handover_batches else None,
    }


def _verification_success_events(batch):
    return (
        FbsPickScanEvent.objects.filter(
            batch=batch,
            stage=FbsPickScanEvent.STAGE_VERIFY_ITEM,
            result=FbsPickScanEvent.RESULT_SUCCESS,
        )
        .select_related("task__order", "created_by")
        .order_by("-created_at", "-id")[:30]
    )


def _marketplace_order_is_canceled(order) -> bool:
    from .services.pick_restock import WB_CLIENT_CANCEL_STATUSES

    marketplace_statuses = {
        str(order.marketplace_status or "").strip().casefold(),
        str(order.marketplace_substatus or "").strip().casefold(),
    }
    return bool(marketplace_statuses & WB_CLIENT_CANCEL_STATUSES)


def _label_requires_problem_tote(label) -> bool:
    if label.status == FbsOrderLabel.STATUS_ERROR:
        return True
    if label.status != FbsOrderLabel.STATUS_REQUESTED:
        return False
    profile = label.order.profile
    return not profile.is_active or not profile.outbox_enabled


def _verification_service_tote_prompt(
    *,
    batch,
    allocation,
    controller,
    problem_kind: str,
    reason: str,
    restock_request=None,
):
    pick_tote_context = (
        FbsControllerPickTote.objects.select_related(
            "session__problem_tote",
            "session__canceled_tote",
        )
        .filter(
            pick_batch=batch,
            session__controller=controller,
            session__status="active",
            status__in=(
                FbsControllerPickTote.STATUS_PROCESSING,
                FbsControllerPickTote.STATUS_AWAITING_EMPTY,
            ),
        )
        .order_by("-id")
        .first()
    )
    session = pick_tote_context.session if pick_tote_context else None
    is_canceled = problem_kind == "canceled"
    is_marketplace_rejected = problem_kind == "marketplace_rejected"
    uses_canceled_tote = is_canceled or is_marketplace_rejected
    service_tote = (
        session.canceled_tote if uses_canceled_tote and session else None
    ) or (session.problem_tote if not uses_canceled_tote and session else None)
    if is_canceled:
        title = "Заказ отменен"
    elif is_marketplace_rejected:
        title = "Маркетплейс отказал в приеме заказа"
    elif problem_kind == "label":
        title = "Этикетка заказа не создана"
    else:
        title = "Честный знак не принят"
    if is_canceled:
        exception_type = FbsPickException.TYPE_NOT_FOUND
    elif is_marketplace_rejected:
        exception_type = FbsPickException.TYPE_OTHER
    elif problem_kind == "label":
        exception_type = FbsPickException.TYPE_OTHER
    else:
        exception_type = FbsPickException.TYPE_BARCODE
    marketplace_check_pending = bool(
        is_marketplace_rejected
        and restock_request is not None
        and restock_request.status
        in {
            FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
            FbsPickRestockRequest.STATUS_FAILED,
        }
    )
    return SimpleNamespace(
        kind=problem_kind,
        title=title,
        reason=str(reason or "").strip(),
        restock_request=restock_request,
        ready_to_scan=bool(
            restock_request is None
            or restock_request.status == FbsPickRestockRequest.STATUS_QUEUED
            or marketplace_check_pending
        ),
        marketplace_check_pending=marketplace_check_pending,
        marketplace_check_failed=bool(
            restock_request is not None
            and restock_request.status == FbsPickRestockRequest.STATUS_FAILED
        ),
        requires_order_scan=is_canceled,
        uses_canceled_tote=uses_canceled_tote,
        order_scan_code=str(allocation.pick_task.order.external_order_id or "").strip(),
        service_tote=service_tote,
        service_tote_label=(
            "тара отмененных заказов" if uses_canceled_tote else "проблемная тара"
        ),
        exception_type=exception_type,
        action_url=reverse(
            "fbs:tsd_pick_verification",
            kwargs={"batch_id": batch.id},
        ) if is_marketplace_rejected else reverse(
            "fbs:tsd_pick_verification_problem",
            kwargs={"allocation_id": allocation.id},
        ),
    )


def _pick_verification_base_template(request) -> str:
    if (
        get_request_role(request) == "fbs_controller"
        and request.headers.get("X-Requested-With") == "fetch"
    ):
        return "fbs/tsd_verification_fragment_base.html"
    return "fbs/tsd_base.html"


def _pick_verification_context(
    request,
    batch,
    allocation,
    *,
    error="",
    ok_message="",
    submitted_item_scan="",
    submitted_marking_scan="",
    retry_marking=False,
    item_selected=False,
    used_marking_info=None,
    unknown_scan_info=None,
    service_tote_prompt=None,
    multi_item_hold_prompt=None,
):
    task = allocation.pick_task
    progress = getattr(allocation, "verification_progress", None)
    verified_qty = int(getattr(progress, "qty_verified", 0) or 0)
    wave_verified_qty = int(
        FbsPickVerificationProgress.objects.filter(
            allocation__pick_task__batch=batch
        ).aggregate(total=Sum("qty_verified"))["total"]
        or 0
    )
    balance = allocation.balance
    from .services.traceability import (
        controller_legacy_marking_exception_allowed,
        controller_marking_scan_required,
        metadata_requirements,
    )

    metadata = metadata_requirements(allocation.order_item)
    trace_expiry = allocation.traceability.expiry_date
    raw_min_expiry = (
        allocation.order_item.requirements.get("min_expiry_date")
        if isinstance(allocation.order_item.requirements, dict)
        else None
    )
    try:
        minimum_expiry = date.fromisoformat(str(raw_min_expiry)) if raw_min_expiry else None
    except ValueError:
        minimum_expiry = None
    minimum_expiry = max(
        value for value in (timezone.localdate(), minimum_expiry) if value is not None
    )
    photo = balance.sku_ref.photos.order_by("sort_order", "id").first() if balance.sku_ref_id else None
    order_allocations = _verification_order_allocations(
        task,
        current_allocation_id=allocation.id,
    )
    order_verified_qty = sum(row.verified_qty for row in order_allocations)
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if (
        get_request_role(request) == "fbs_controller"
        and batch.workstation_id
        and batch.workstation.device_agent_id
    ):
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    verification_wave_context = _verification_wave_context(
        batch,
        current_task_id=task.id,
    )
    current_wave_row = next(
        (
            row
            for row in verification_wave_context["verification_wave_rows"]
            if row.task.id == task.id
        ),
        None,
    )
    order_incomplete = (
        bool(current_wave_row.is_incomplete)
        if current_wave_row is not None
        else _open_pick_shortage_issue(task) is not None
    )
    marking_required = controller_marking_scan_required(allocation)
    legacy_marking_exception_allowed = (
        controller_legacy_marking_exception_allowed(allocation)
    )
    marking_optional = bool(
        not marking_required and wb_optional_marking_available(allocation.order_item)
    )
    return _base_context(
        request,
        verification_base_template=_pick_verification_base_template(request),
        page_title=f"FBS · Проверка волны #{batch.id}",
        back_url=(
            reverse("fbs:controller_home")
            if get_request_role(request) == "fbs_controller"
            else reverse("fbs:operator_wave_detail", kwargs={"batch_id": batch.id})
        ),
        batch=batch,
        task=task,
        order=task.order,
        order_incomplete=order_incomplete,
        order_problem_reason=(
            str(task.order.problem_reason or "").strip()
            if order_incomplete
            else ""
        ),
        allocation=allocation,
        balance=balance,
        verified_qty=verified_qty,
        allocation_remaining=max(int(allocation.qty_picked or 0) - verified_qty, 0),
        wave_verified_qty=wave_verified_qty,
        wave_planned_qty=int(batch.planned_qty or 0),
        order_verified_qty=order_verified_qty,
        order_planned_qty=int(task.planned_qty or 0),
        order_allocations=order_allocations,
        **verification_wave_context,
        verified_product_events=_verification_success_events(batch),
        expected_item_code=balance.barcode or allocation.order_item.barcode,
        expected_marking_code=balance.marking_code,
        marking_required=marking_required,
        legacy_marking_exception_allowed=legacy_marking_exception_allowed,
        marking_optional=marking_optional,
        submitted_item_scan=submitted_item_scan,
        submitted_marking_scan=submitted_marking_scan,
        retry_marking=retry_marking,
        item_selected=item_selected,
        used_marking_info=used_marking_info,
        unknown_scan_info=unknown_scan_info,
        service_tote_prompt=service_tote_prompt,
        multi_item_hold_prompt=multi_item_hold_prompt,
        expiry_required=metadata.expiry_required,
        expiry_entry_required=metadata.expiry_required and trace_expiry is None,
        trace_expiry_date=trace_expiry,
        minimum_expiry_date=minimum_expiry,
        item_scan_label="Штрихкод товара",
        agent_scan_poll_url=agent_scan_poll_url,
        agent_scan_event_id=agent_scan_event_id,
        product_photo_url=photo.url if photo else "",
        recent_scan_events=FbsPickScanEvent.objects.filter(batch=batch)
        .select_related("created_by", "task__order")
        .order_by("-created_at", "-id")[:12],
        error=error,
        ok_message=ok_message,
    )


def _pick_verification_label_context(
    request,
    batch,
    label,
    *,
    error="",
    ok_message="",
):
    context = _label_context(request, label, error=error, ok_message=ok_message)
    context["verification_base_template"] = _pick_verification_base_template(
        request
    )
    task = (
        batch.tasks.filter(
            order_id=label.order_id,
            status=FbsPickTask.STATUS_PICKED,
        )
        .order_by("sort_order", "id")
        .first()
    )
    if task is None:
        return context
    shortage_issue = _open_pick_shortage_issue(task)
    order_allocations = _verification_order_allocations(task)
    workstation = (
        FbsWorkstation.objects.select_related(
            "active_handover_box__batch__profile__agency"
        )
        .filter(pk=batch.workstation_id, is_active=True)
        .first()
        if batch.workstation_id
        else None
    )
    assignment = (
        FbsHandoverOrderAssignment.objects.select_related("batch__profile__agency")
        .filter(order_id=label.order_id)
        .first()
    )
    active_box = workstation.active_handover_box if workstation is not None else None
    active_box_compatible = bool(
        active_box is not None
        and assignment is not None
        and active_box.batch_id == assignment.batch_id
        and active_box.status == FbsHandoverBox.STATUS_OPEN
        and active_box.batch.status == FbsHandoverBatch.STATUS_OPEN
    )
    packing_box = (
        active_box
        if active_box_compatible
        else get_ready_handover_box_for_order(order_id=label.order_id)
    )
    controller_pick_tote = controller_pick_tote_for_batch(batch_id=batch.id)
    service_tote_prompt = None
    problem_allocation = next(
        (row for row in order_allocations if int(row.qty_picked or 0) > 0),
        None,
    )
    if (
        get_request_role(request) == "fbs_controller"
        and problem_allocation is not None
        and _label_requires_problem_tote(label)
    ):
        service_tote_prompt = _verification_service_tote_prompt(
            batch=batch,
            allocation=problem_allocation,
            controller=request.user,
            problem_kind="label",
            reason=(
                str(label.error or "").strip()
                or "Этикетку заказа не удалось получить или напечатать."
            ),
        )
    verification_wave_context = _verification_wave_context(
        batch,
        current_task_id=task.id,
    )
    context.update(
        {
            "page_title": f"FBS · Проверка волны #{batch.id}",
            "back_url": reverse("fbs:controller_home"),
            "verification_label_phase": True,
            "preloaded_ozon_order_label": is_preloaded_ozon_order_label(label),
            "task": task,
            "order_incomplete": shortage_issue is not None,
            "order_problem_reason": (
                str(task.order.problem_reason or "").strip()
                if shortage_issue is not None
                else ""
            ),
            "order_allocations": order_allocations,
            **verification_wave_context,
            "verified_product_events": _verification_success_events(batch),
            "order_verified_qty": sum(row.verified_qty for row in order_allocations),
            "order_planned_qty": int(task.planned_qty or 0),
            "wave_verified_qty": int(
                FbsPickVerificationProgress.objects.filter(
                    allocation__pick_task__batch=batch,
                    allocation__pick_task__status=FbsPickTask.STATUS_PICKED,
                ).aggregate(total=Sum("qty_verified"))["total"]
                or 0
            ),
            "wave_planned_qty": int(
                batch.tasks.filter(status=FbsPickTask.STATUS_PICKED).aggregate(
                    total=Sum("planned_qty")
                )["total"]
                or 0
            ),
            "verification_workstation": workstation,
            "handover_assignment": assignment,
            "handover_assignment_confirmed": bool(
                assignment is not None
                and assignment.status
                == FbsHandoverOrderAssignment.STATUS_CONFIRMED
            ),
            "active_handover_box": active_box,
            "active_box_compatible": active_box_compatible,
            "packing_handover_box": packing_box,
            "automatic_box_selection": bool(
                packing_box is not None and not active_box_compatible
            ),
            "controller_pick_tote": controller_pick_tote,
            "controller_check_tote": (
                controller_pick_tote.check_tote
                if controller_pick_tote is not None
                else None
            ),
            "service_tote_prompt": service_tote_prompt,
        }
    )
    return context


def _label_queryset():
    return FbsOrderLabel.objects.select_related(
        "order__profile__agency",
        "requested_by",
        "applied_by",
    ).order_by("requested_at", "id")


def _accessible_label_queryset(request):
    labels = _label_queryset()
    if get_request_role(request) == "fbs_controller":
        labels = labels.filter(
            order__pick_tasks__batch__verification_assigned_to=request.user
        ).distinct()
    return labels


def _labels_context(request, *, error="", ok_message=""):
    labels = _accessible_label_queryset(request)
    return _base_context(
        request,
        page_title="FBS · Этикетки заказов",
        back_url=(
            reverse("fbs:controller_home")
            if get_request_role(request) == "fbs_controller"
            else reverse("fbs:tsd_picking")
        ),
        requested_labels=labels.filter(status=FbsOrderLabel.STATUS_REQUESTED)[:50],
        ready_labels=labels.filter(status=FbsOrderLabel.STATUS_READY)[:50],
        applied_labels=labels.filter(status=FbsOrderLabel.STATUS_APPLIED).order_by(
            "-applied_at", "-id"
        )[:10],
        error=error,
        ok_message=ok_message,
    )


def _label_context(request, label, *, error="", ok_message=""):
    batch = (
        FbsPickBatch.objects.filter(
            tasks__order=label.order,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .select_related("workstation__device_agent")
        .order_by("created_at", "id")
        .first()
    )
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if (
        get_request_role(request) == "fbs_controller"
        and batch is not None
        and batch.workstation_id
        and batch.workstation.device_agent_id
    ):
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    preloaded_ozon_order_label = is_preloaded_ozon_order_label(label)
    print_prefix = (
        f"fbs:ozon-order-qr:{label.id}:"
        if preloaded_ozon_order_label
        else f"fbs:order-label:{label.id}:"
    )
    print_jobs = ProcessingPrintJob.objects.filter(card_id__startswith=print_prefix)
    if batch is not None and batch.verification_started_at is not None:
        print_jobs = print_jobs.filter(
            created_at__gte=batch.verification_started_at,
        )
    if get_request_role(request) == "fbs_controller":
        recover_stale_fbs_order_label_print_job(
            label_id=label.id,
            not_before=(
                batch.verification_started_at
                if batch is not None
                else None
            ),
        )
    print_job = print_jobs.order_by("-created_at", "-id").first()
    print_status_labels = {
        ProcessingPrintJob.STATUS_PENDING: "В очереди",
        ProcessingPrintJob.STATUS_PRINTING: "Печатается",
        ProcessingPrintJob.STATUS_PRINTED: "Напечатана",
        ProcessingPrintJob.STATUS_FAILED: "Ошибка печати",
    }
    print_job_delay_warning = ""
    print_job_retry_ready = False
    if print_job is not None and print_job.status in {
        ProcessingPrintJob.STATUS_PENDING,
        ProcessingPrintJob.STATUS_PRINTING,
    }:
        job_age_seconds = max(
            int((timezone.now() - print_job.updated_at).total_seconds()),
            0,
        )
        lease_seconds = fbs_desktop_print_lease_seconds()
        if (
            print_job.status == ProcessingPrintJob.STATUS_PRINTING
            and job_age_seconds >= lease_seconds
        ):
            print_job_retry_ready = True
            print_job_delay_warning = (
                "Fullbox Desktop не подтвердил печать за установленный срок. "
                "Задание доступно для безопасного повторного запуска."
            )
        elif job_age_seconds >= 20:
            if print_job.status == ProcessingPrintJob.STATUS_PENDING:
                print_job_delay_warning = (
                    "Задание больше 20 секунд ожидает Fullbox Desktop. "
                    "Проверьте, что Desktop запущен и выбран правильный принтер."
                )
            else:
                print_job_delay_warning = (
                    "Fullbox Desktop забрал задание, но ещё не подтвердил печать. "
                    "Проверьте принтер и не запускайте вторую копию до окончания ожидания."
                )
    return _base_context(
        request,
        page_title=f"FBS · Этикетка {label.order.external_order_id}",
        back_url=(
            reverse("fbs:tsd_pick_verification", kwargs={"batch_id": batch.id})
            if batch
            else reverse("fbs:tsd_labels")
        ),
        label=label,
        order=label.order,
        batch=batch,
        print_job=print_job,
        print_status_label=(print_status_labels.get(print_job.status, "") if print_job else ""),
        print_job_delay_warning=print_job_delay_warning,
        print_job_retry_ready=print_job_retry_ready,
        preloaded_ozon_order_label=preloaded_ozon_order_label,
        agent_scan_poll_url=agent_scan_poll_url,
        agent_scan_event_id=agent_scan_event_id,
        error=error,
        ok_message=ok_message,
    )


def _source_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_source_scan_{allocation_id}"


def _query_int(request, key: str) -> int:
    try:
        raw_value = request.POST.get(key) if key in request.POST else request.GET.get(key)
        return max(int(raw_value or 0), 0)
    except (TypeError, ValueError):
        return 0


def _allocation_context(request, allocation, *, error="", ok_message=""):
    source_snapshot = allocation.source_snapshot
    source_code = (
        str(source_snapshot.container_code or "").strip()
        or str(getattr(source_snapshot.container, "container_code", "") or "").strip()
        or str(getattr(source_snapshot.location, "location_code", "") or "").strip()
    )
    source_confirmed = bool(request.session.get(_source_session_key(allocation.id)))
    open_count = FbsReplenishmentAllocation.objects.filter(
        line__plan=allocation.line.plan,
        status__in=OPEN_ALLOCATION_STATUSES,
    ).count()
    return _base_context(
        request,
        page_title=f"FBS · Задание #{allocation.id}",
        back_url=reverse("fbs:tsd_reachtruck"),
        allocation=allocation,
        plan=allocation.line.plan,
        source_code=source_code,
        source_confirmed=source_confirmed,
        open_count=open_count,
        is_staging_transfer=bool(
            allocation.line.plan.client_movement_request_id
            and allocation.line.plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and allocation.line.plan.target_box_id is None
            and allocation.line.plan.staging_location_id
        ),
        staging_container_code=(
            movement_staging_container_code(
                allocation.line.plan.client_movement_request_id
            )
            if allocation.line.plan.client_movement_request_id
            and allocation.line.plan.mode == FbsReplenishmentPlan.MODE_ITEM
            else ""
        ),
        error=error,
        ok_message=ok_message,
    )


@fbs_module_required
@role_required(*TSD_ROLES)
@require_GET
def tsd_home(request):
    role = get_request_role(request)
    if role == "reachtruck_driver":
        return redirect("/reachtruck/")
    if role == "picker":
        return render(request, "fbs/tsd_picker_home.html", _picker_home_context(request))
    if role == "processing_worker":
        return redirect("fbs:tsd_picking")
    if role == "fbs_controller":
        return redirect("fbs:controller_home")
    if role in STOREKEEPER_ROLES or role == "developer":
        return redirect("fbs:tsd_storekeeper")
    if role in REACHTRUCK_ROLES:
        return redirect("/reachtruck/")
    if role in PICKER_ROLES:
        return redirect("fbs:tsd_picking")
    return HttpResponseForbidden("Доступ запрещен")


@fbs_module_required
@role_required("picker")
@require_GET
def tsd_picker_movements(request):
    movements = list(
        FbsInternalMovement.objects.select_related(
            "source_box",
            "source_box__pallet__cell__location",
            "target_pallet__cell__location",
            "target_box",
            "source_balance",
        )
        .filter(
            assigned_to=request.user,
            status__in=(
                FbsInternalMovement.STATUS_PROPOSED,
                FbsInternalMovement.STATUS_IN_PROGRESS,
            ),
        )
        .order_by("created_at", "id")[:50]
    )
    for movement in movements:
        movement.source_location_label = _picker_location_label(
            movement.source_box.pallet.cell
        )
        movement.target_location_label = _picker_location_label(
            movement.target_pallet.cell
        )
    return render(
        request,
        "fbs/tsd_picker_movements.html",
        _base_context(
            request,
            page_title="Перемещение товаров",
            back_url=reverse("fbs:tsd_home"),
            movements=movements,
        ),
    )


@fbs_module_required
@role_required("picker")
@require_GET
def tsd_picker_returns(request):
    restock_requests = list(
        FbsPickRestockRequest.objects.select_related(
            "batch__workstation",
            "assigned_to",
            "source_tote",
            "quarantine_box__pallet__cell__location",
        )
        .filter(
            Q(
                status=FbsPickRestockRequest.STATUS_QUEUED,
                assigned_to__isnull=True,
            )
            | Q(
                status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
                assigned_to=request.user,
            )
        )
        .filter(Q(order__isnull=True) | Q(source_tote__isnull=False))
        .annotate(source_box_count=Count("lines__source_box", distinct=True))
        .order_by("created_at", "id")[:50]
    )
    return render(
        request,
        "fbs/tsd_picker_returns.html",
        _base_context(
            request,
            page_title="Возврат отборов",
            back_url=reverse("fbs:tsd_home"),
            restock_requests=restock_requests,
        ),
    )


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_location(request):
    scan = str(request.POST.get("scan") or "").strip() if request.method == "POST" else ""
    error = ""
    cell = None
    balances = []
    totals = {"qty": 0, "available": 0, "reserved": 0, "boxes": 0}
    if request.method == "POST":
        normalized_scan = _normalized_picker_location_scan(scan)
        if not normalized_scan:
            error = "Скан не получен. Повторите сканирование QR места."
        else:
            cells = FbsStorageCell.objects.select_related("location").filter(
                is_active=True,
                location__is_active=True,
            )
            cell = next(
                (
                    candidate
                    for candidate in cells
                    if normalized_scan in _picker_location_candidates(candidate)
                ),
                None,
            )
            if cell is None:
                error = "Место не найдено в FBS. Проверьте QR и отсканируйте еще раз."
            else:
                balance_queryset = FbsStockBalance.objects.select_related("box").filter(
                    box__pallet__cell=cell,
                    box__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
                    qty__gt=0,
                )
                summary = balance_queryset.aggregate(
                    qty=Sum("qty"),
                    available=Sum("available_qty"),
                    reserved=Sum("reserved_qty"),
                    boxes=Count("box_id", distinct=True),
                )
                balances = list(
                    balance_queryset.order_by("box__box_code", "sku_code", "id")[:200]
                )
                totals = {
                    "qty": int(summary["qty"] or 0),
                    "available": int(summary["available"] or 0),
                    "reserved": int(summary["reserved"] or 0),
                    "boxes": int(summary["boxes"] or 0),
                }
    return render(
        request,
        "fbs/tsd_picker_location.html",
        _base_context(
            request,
            page_title="Информация о месте",
            back_url=reverse("fbs:tsd_home"),
            scan=scan,
            cell=cell,
            location_label=_picker_location_label(cell) if cell else "",
            balances=balances,
            totals=totals,
            error=error,
        ),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper(request):
    return render(request, "fbs/tsd_storekeeper_list.html", _storekeeper_list_context(request))


def _fbs_stock_base_queryset():
    return FbsStockBalance.objects.filter(
        qty__gt=0,
        box__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        box__pallet__status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        box__pallet__cell__is_active=True,
    ).select_related(
        "agency",
        "sku_ref",
        "box__pallet__cell__location",
    )


def _fbs_stock_apply_filters(request, queryset, *, include_sku=True):
    q = str(request.GET.get("q") or "").strip()
    sku = str(request.GET.get("sku") or "").strip()
    client_id = str(request.GET.get("client") or "").strip()
    stock_status = str(request.GET.get("stock_status") or "").strip()
    cell = str(request.GET.get("cell") or "").strip()
    box = str(request.GET.get("box") or "").strip()
    pallet = str(request.GET.get("pallet") or "").strip()

    if q:
        exact_code_queryset = queryset.filter(
            Q(sku_code=q)
            | Q(barcode=q)
            | Q(marking_code=q)
            | Q(lot_code=q)
            | Q(box__box_code=q)
            | Q(box__pallet__pallet_code=q)
            | Q(box__pallet__cell__cell_code=q)
        )
        exact_code_found = (
            len(q) >= 4
            and not any(character.isspace() for character in q)
            and exact_code_queryset.exists()
        )
        if exact_code_found:
            queryset = exact_code_queryset
        else:
            queryset = queryset.filter(
                Q(agency__agn_name__icontains=q)
                | Q(agency__short_name__icontains=q)
                | Q(sku_code__icontains=q)
                | Q(name__icontains=q)
                | Q(size__icontains=q)
                | Q(barcode__icontains=q)
                | Q(marking_code__icontains=q)
                | Q(lot_code__icontains=q)
                | Q(box__box_code__icontains=q)
                | Q(box__pallet__pallet_code__icontains=q)
                | Q(box__pallet__cell__cell_code__icontains=q)
            )
    if include_sku and sku:
        queryset = queryset.filter(
            Q(sku_code__icontains=sku) | Q(name__icontains=sku)
        )
    if client_id.isdigit():
        queryset = queryset.filter(agency_id=int(client_id))
    if stock_status == "available":
        queryset = queryset.filter(available_qty__gt=0)
    elif stock_status == "reserved":
        queryset = queryset.filter(reserved_qty__gt=0)
    elif stock_status == "unavailable":
        queryset = queryset.filter(available_qty=0)
    else:
        stock_status = ""
    if cell:
        queryset = queryset.filter(box__pallet__cell__cell_code__icontains=cell)
    if box:
        queryset = queryset.filter(box__box_code__icontains=box)
    if pallet:
        queryset = queryset.filter(box__pallet__pallet_code__icontains=pallet)

    return queryset, {
        "q": q,
        "sku": sku,
        "client_id": client_id if client_id.isdigit() else "",
        "stock_status": stock_status,
        "cell": cell,
        "box": box,
        "pallet": pallet,
    }


def _fbs_stock_sort(request, queryset):
    sort_key = str(request.GET.get("sort") or "client").strip()
    if sort_key not in FBS_STOCK_SORT_FIELDS:
        sort_key = "client"
    default_dir = "desc" if sort_key in FBS_STOCK_DESCENDING_DEFAULTS else "asc"
    sort_dir = str(request.GET.get("dir") or default_dir).strip().lower()
    if sort_dir not in {"asc", "desc"}:
        sort_dir = default_dir
    field = FBS_STOCK_SORT_FIELDS[sort_key]
    ordering = f"-{field}" if sort_dir == "desc" else field
    return queryset.order_by(ordering, "agency_id", "sku_code", "id"), sort_key, sort_dir


def _fbs_stock_grouped_sort(request, queryset):
    sort_key = str(request.GET.get("sort") or "client").strip()
    if sort_key not in FBS_STOCK_BOX_SORT_FIELDS:
        sort_key = "client"
    default_dir = "desc" if sort_key in FBS_STOCK_DESCENDING_DEFAULTS else "asc"
    sort_dir = str(request.GET.get("dir") or default_dir).strip().lower()
    if sort_dir not in {"asc", "desc"}:
        sort_dir = default_dir

    grouped_queryset = queryset.values(
        "box_id",
        "agency_id",
        "agency__agn_name",
        "agency__short_name",
        "box__box_code",
        "box__pallet_id",
        "box__pallet__pallet_code",
        "box__pallet__cell_id",
        "box__pallet__cell__cell_code",
    ).annotate(
        group_sku=Min("sku_code"),
        group_name=Min("name"),
        group_barcode=Min("barcode"),
        group_expiry=Min("expiry_date"),
        group_updated=Max("updated_at"),
        group_qty=Sum("qty"),
        group_available_qty=Sum("available_qty"),
        group_reserved_qty=Sum("reserved_qty"),
    )
    field = FBS_STOCK_BOX_SORT_FIELDS[sort_key]
    ordering = f"-{field}" if sort_dir == "desc" else field
    return (
        grouped_queryset.order_by(ordering, "box__box_code", "box_id"),
        sort_key,
        sort_dir,
    )


def _fbs_stock_box_groups(page_rows, balances) -> list[dict]:
    balances_by_box = defaultdict(list)
    for balance in balances:
        balances_by_box[balance.box_id].append(balance)

    groups = []
    for page_row in page_rows:
        box_balances = balances_by_box.get(page_row["box_id"], [])
        if not box_balances:
            continue

        items_by_identity = {}
        for balance in box_balances:
            identity = (
                balance.sku_ref_id,
                balance.sku_code,
                balance.name,
                balance.size,
                balance.barcode,
                balance.goods_type,
                balance.lot_code,
                balance.expiry_date,
            )
            item = items_by_identity.setdefault(
                identity,
                {
                    "balance": balance,
                    "qty": 0,
                    "available_qty": 0,
                    "reserved_qty": 0,
                    "marking_codes": [],
                    "balance_count": 0,
                },
            )
            item["qty"] += int(balance.qty or 0)
            item["available_qty"] += int(balance.available_qty or 0)
            item["reserved_qty"] += int(balance.reserved_qty or 0)
            item["balance_count"] += 1
            marking_code = str(balance.marking_code or "").strip()
            if marking_code:
                item["marking_codes"].append(marking_code)

        items = sorted(
            items_by_identity.values(),
            key=lambda item: (
                str(item["balance"].sku_code or "").casefold(),
                str(item["balance"].name or "").casefold(),
                str(item["balance"].size or "").casefold(),
                str(item["balance"].barcode or "").casefold(),
            ),
        )
        for item in items:
            item["marking_codes"].sort()
            item["marking_count"] = len(item["marking_codes"])

        first_balance = box_balances[0]
        groups.append(
            {
                "box_id": first_balance.box_id,
                "box": first_balance.box,
                "agency": first_balance.agency,
                "items": items,
                "item_count": len(items),
                "balance_count": len(box_balances),
                "marking_count": sum(item["marking_count"] for item in items),
                "qty": sum(item["qty"] for item in items),
                "available_qty": sum(item["available_qty"] for item in items),
                "reserved_qty": sum(item["reserved_qty"] for item in items),
                "updated_at": max(
                    (
                        balance.updated_at
                        for balance in box_balances
                        if balance.updated_at is not None
                    ),
                    default=None,
                ),
            }
        )
    return groups


def _fbs_stock_sort_urls(request, current_sort: str, current_dir: str) -> dict[str, str]:
    urls = {}
    for key in FBS_STOCK_SORT_FIELDS:
        params = request.GET.copy()
        params.pop("page", None)
        params["sort"] = key
        params["dir"] = (
            "desc"
            if key == current_sort and current_dir == "asc"
            else "asc"
        )
        urls[key] = f"{request.path}?{params.urlencode()}"
    return urls


def _fbs_stock_pagination_links(page_obj, *, radius: int = 2) -> list[dict]:
    total_pages = int(page_obj.paginator.num_pages or 0)
    current = int(page_obj.number or 1)
    pages = {1, total_pages}
    pages.update(range(max(1, current - radius), min(total_pages, current + radius) + 1))
    links = []
    previous = 0
    for number in sorted(page for page in pages if page > 0):
        if previous and number - previous > 1:
            links.append({"ellipsis": True})
        links.append({"number": number, "current": number == current})
        previous = number
    return links


def _fbs_stock_client_options():
    return Agency.objects.filter(
        fbs_stock_balances__qty__gt=0,
        fbs_stock_balances__box__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        fbs_stock_balances__box__pallet__status__in=(
            FbsPallet.STATUS_PLANNED,
            FbsPallet.STATUS_ACTIVE,
        ),
        fbs_stock_balances__box__pallet__cell__is_active=True,
    ).distinct().order_by("short_name", "agn_name", "id")


def _fbs_stock_sku_options(queryset, *, limit: int = 1000) -> list[dict[str, str]]:
    options = []
    seen_values = set()
    rows = (
        queryset.values_list("sku_code", "name")
        .order_by("sku_code", "name")
        .distinct()[:limit]
    )
    for sku_code, name in rows:
        sku_value = str(sku_code or "").strip()
        name_value = str(name or "").strip()
        value = sku_value or name_value
        normalized_value = value.casefold()
        if not value or normalized_value in seen_values:
            continue
        seen_values.add(normalized_value)
        options.append(
            {
                "value": value,
                "label": " · ".join(part for part in (sku_value, name_value) if part),
            }
        )
    return options


def _fbs_stock_summary(queryset) -> dict:
    totals = queryset.aggregate(
        row_count=Count("id"),
        client_count=Count("agency_id", distinct=True),
        box_count=Count("box_id", distinct=True),
        pallet_count=Count("box__pallet_id", distinct=True),
        qty=Sum("qty"),
        available_qty=Sum("available_qty"),
        reserved_qty=Sum("reserved_qty"),
    )
    totals["sku_count"] = queryset.values(
        "agency_id", "sku_code", "size", "barcode", "goods_type"
    ).distinct().count()
    return {key: int(value or 0) for key, value in totals.items()}


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_stock(request):
    base_queryset = _fbs_stock_base_queryset()
    option_queryset, _ = _fbs_stock_apply_filters(
        request,
        base_queryset,
        include_sku=False,
    )
    queryset, filters = _fbs_stock_apply_filters(request, base_queryset)
    summary = _fbs_stock_summary(queryset)
    grouped_queryset, sort_key, sort_dir = _fbs_stock_grouped_sort(request, queryset)
    paginator = Paginator(grouped_queryset, 100)
    page_obj = paginator.get_page(request.GET.get("page"))
    page_rows = list(page_obj.object_list)
    page_box_ids = [row["box_id"] for row in page_rows]
    page_balances = list(
        queryset.filter(box_id__in=page_box_ids).order_by(
            "box_id",
            "sku_code",
            "name",
            "size",
            "barcode",
            "goods_type",
            "lot_code",
            "expiry_date",
            "marking_code",
            "id",
        )
    )

    page_params = request.GET.copy()
    page_params.pop("page", None)
    export_params = request.GET.copy()
    export_params.pop("page", None)
    export_url = reverse("fbs:tsd_storekeeper_stock_export")
    if export_params:
        export_url = f"{export_url}?{export_params.urlencode()}"

    return render(
        request,
        "fbs/tsd_stock_list.html",
        _base_context(
            request,
            page_title="FBS · Остатки",
            operator_section="stock",
            back_url=reverse("fbs:tsd_storekeeper"),
            box_groups=_fbs_stock_box_groups(page_rows, page_balances),
            page_obj=page_obj,
            page_links=_fbs_stock_pagination_links(page_obj),
            page_query=page_params.urlencode(),
            client_options=_fbs_stock_client_options(),
            sku_options=_fbs_stock_sku_options(option_queryset),
            summary=summary,
            stock_sort_key=sort_key,
            stock_sort_dir=sort_dir,
            stock_sort_urls=_fbs_stock_sort_urls(request, sort_key, sort_dir),
            stock_groups_expanded=bool(filters["q"] or filters["sku"] or filters["box"]),
            export_url=export_url,
            reset_url=reverse("fbs:tsd_storekeeper_stock"),
            **filters,
        ),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_stock_export(request):
    queryset, _ = _fbs_stock_apply_filters(request, _fbs_stock_base_queryset())
    queryset, _, _ = _fbs_stock_sort(request, queryset)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title="Остатки FBS")
    sheet.append(
        [
            "Клиент",
            "SKU",
            "Наименование",
            "Размер",
            "Штрихкод",
            "Тип товара",
            "Честный знак",
            "Партия",
            "Срок годности",
            "Всего",
            "Доступно",
            "В резерве",
            "Короб FBS",
            "Паллета FBS",
            "Ячейка FBS",
            "Обновлено",
        ]
    )
    for balance in queryset.iterator(chunk_size=1000):
        updated_at = (
            timezone.localtime(balance.updated_at).strftime("%d.%m.%Y %H:%M")
            if balance.updated_at
            else ""
        )
        sheet.append(
            [
                balance.agency.short_name or balance.agency.agn_name or str(balance.agency),
                balance.sku_code,
                balance.name,
                balance.size,
                balance.barcode,
                balance.goods_type,
                balance.marking_code,
                balance.lot_code,
                balance.expiry_date.strftime("%d.%m.%Y") if balance.expiry_date else "",
                int(balance.qty or 0),
                int(balance.available_qty or 0),
                int(balance.reserved_qty or 0),
                balance.box.box_code,
                balance.box.pallet.pallet_code,
                balance.box.pallet.cell.cell_code,
                updated_at,
            ]
        )
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="fbs_stock_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


def _inventory_queryset():
    return FbsInventorySession.objects.select_related(
        "agency",
        "cell",
        "pallet",
        "box",
        "sku",
        "first_counter",
        "second_counter",
        "approved_by",
    ).prefetch_related("lines__balance__box__pallet__cell")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_inventory(request):
    error = ""
    if request.method == "POST":
        role = get_request_role(request)
        if role not in INVENTORY_MANAGER_ROLES and role != "developer":
            return HttpResponseForbidden("Создать инвентаризацию может только начальник склада")
        target_value = str(request.POST.get("target") or "").strip()
        mode = str(request.POST.get("mode") or "").strip()
        scan_mode = str(
            request.POST.get("scan_mode") or FbsInventorySession.SCAN_MODE_BARCODE
        ).strip()
        try:
            scope_type, raw_object_id = target_value.split(":", 1)
            object_id = int(raw_object_id or 0)
        except (TypeError, ValueError):
            scope_type = ""
            object_id = 0
        kwargs = {}
        model_by_scope = {
            FbsInventorySession.SCOPE_AGENCY: (Agency, "agency"),
            FbsInventorySession.SCOPE_CELL: (FbsStorageCell, "cell"),
            FbsInventorySession.SCOPE_PALLET: (FbsPallet, "pallet"),
            FbsInventorySession.SCOPE_BOX: (FbsBox, "box"),
            FbsInventorySession.SCOPE_SKU: (SKU, "sku"),
        }
        try:
            if scope_type == FbsInventorySession.SCOPE_ALL:
                pass
            elif scope_type in model_by_scope:
                model, key = model_by_scope[scope_type]
                kwargs[key] = model.objects.get(pk=object_id)
                if scope_type == FbsInventorySession.SCOPE_SKU:
                    kwargs["agency"] = kwargs["sku"].agency
            else:
                raise ValueError("Выберите поддерживаемую область инвентаризации.")
            session = create_inventory_session(
                scope_type=scope_type,
                mode=mode,
                scan_mode=scan_mode,
                created_by=request.user,
                **kwargs,
            )
            return redirect("fbs:tsd_inventory_detail", session_id=session.id)
        except (
            FbsError,
            ValueError,
            Agency.DoesNotExist,
            FbsStorageCell.DoesNotExist,
            FbsPallet.DoesNotExist,
            FbsBox.DoesNotExist,
            SKU.DoesNotExist,
        ) as exc:
            error = str(exc)
    context = _base_context(
        request,
        page_title="FBS · Инвентаризация",
        back_url=reverse("fbs:tsd_storekeeper"),
        sessions=_inventory_queryset().order_by("-created_at")[:50],
        agencies=Agency.objects.filter(fbs_integration_profiles__isnull=False).distinct().order_by("agn_name"),
        cells=FbsStorageCell.objects.filter(is_active=True).order_by("cell_code"),
        pallets=FbsPallet.objects.exclude(status=FbsPallet.STATUS_ARCHIVED).order_by("pallet_code"),
        boxes=FbsBox.objects.exclude(status=FbsBox.STATUS_ARCHIVED).order_by("box_code"),
        skus=SKU.objects.filter(
            agency__fbs_integration_profiles__isnull=False,
            deleted=False,
        ).select_related("agency").distinct().order_by("agency__agn_name", "sku_code")[:1000],
        mode_choices=FbsInventorySession.MODE_CHOICES,
        scan_mode_choices=FbsInventorySession.SCAN_MODE_CHOICES,
        can_manage=get_request_role(request) in INVENTORY_MANAGER_ROLES or get_request_role(request) == "developer",
        error=error,
    )
    return render(request, "fbs/tsd_inventory_list.html", context, status=400 if error else 200)


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_inventory_detail(request, session_id: int):
    session = get_object_or_404(_inventory_queryset(), pk=session_id)
    error = ""
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        try:
            if action == "activate":
                activate_drained_inventory(session_id=session.id)
            elif action == "scan":
                record_inventory_scan(
                    session_id=session.id,
                    scan_code=request.POST.get("scan_code", ""),
                    counted_by=request.user,
                )
            elif action == "finish":
                finish_inventory_count(session_id=session.id, counted_by=request.user)
            elif action == "approve":
                if get_request_role(request) not in INVENTORY_MANAGER_ROLES and get_request_role(request) != "developer":
                    return HttpResponseForbidden("Утверждение доступно только начальнику склада")
                final_counts = {
                    line.id: int(request.POST[f"final_{line.id}"])
                    for line in session.lines.all()
                    if str(request.POST.get(f"final_{line.id}") or "").strip()
                }
                approve_inventory(
                    session_id=session.id,
                    approved_by=request.user,
                    final_counts=final_counts,
                )
            else:
                raise ValueError("Неизвестная команда инвентаризации.")
            return redirect("fbs:tsd_inventory_detail", session_id=session.id)
        except (FbsError, ValueError) as exc:
            error = str(exc)
        session = get_object_or_404(_inventory_queryset(), pk=session_id)
    return render(
        request,
        "fbs/tsd_inventory_detail.html",
        _base_context(
            request,
            page_title=f"FBS · Инвентаризация #{session.id}",
            back_url=reverse("fbs:tsd_inventory"),
            inventory=session,
            can_manage=get_request_role(request) in INVENTORY_MANAGER_ROLES or get_request_role(request) == "developer",
            error=error,
        ),
        status=400 if error else 200,
    )


def _handover_queryset():
    list_items = FbsOrderItem.objects.select_related(
        "sku", "order__profile"
    ).prefetch_related(
        Prefetch(
            "metadata_transfers",
            queryset=FbsMarketplaceMetadataTransfer.objects.select_related(
                "order_item__order__profile"
            ).only(
                "id",
                "order_item_id",
                "metadata_type",
                "status",
                "is_required",
                "external_status",
                "order_item__order__profile__marketplace",
            ),
        )
    )
    list_assignments = (
        FbsHandoverOrderAssignment.objects.exclude(
            status=FbsHandoverOrderAssignment.STATUS_CANCELED
        )
        .select_related("order")
        .prefetch_related(
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.only("id", "order_id", "status"),
            ),
            Prefetch("order__items", queryset=list_items),
        )
        .order_by("id")
    )
    list_box_orders = (
        FbsHandoverOrder.objects.filter(status=FbsHandoverOrder.STATUS_ACTIVE)
        .select_related("order")
        .prefetch_related(
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.only("id", "order_id", "status"),
            ),
            Prefetch("order__items", queryset=list_items),
        )
        .order_by("id")
    )
    list_boxes = FbsHandoverBox.objects.prefetch_related(
        Prefetch("orders", queryset=list_box_orders)
    ).order_by("id")
    list_commands = FbsMarketplaceCommand.objects.only(
        "id", "handover_batch_id", "order_id", "status"
    ).order_by("id")
    return FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by", "dispatch_location"
    ).prefetch_related(
        Prefetch("boxes", queryset=list_boxes, to_attr="list_boxes"),
        Prefetch(
            "order_assignments",
            queryset=list_assignments,
            to_attr="list_assignments",
        ),
        Prefetch(
            "marketplace_commands",
            queryset=list_commands,
            to_attr="list_commands",
        ),
    )


def _handover_relevant_commands(commands, *, order_ids):
    current_order_ids = set(order_ids)
    return [
        command
        for command in commands
        if command.order_id is None or command.order_id in current_order_ids
    ]


def _prepare_handover_list_rows(batches):
    from .services.traceability import (
        marketplace_metadata_transfer_resolved,
        metadata_requirements,
        ozon_marking_transfer_not_required,
    )

    for batch in batches:
        boxes = list(batch.list_boxes)
        assignments = list(batch.list_assignments)
        commands = list(batch.list_commands)
        boxed_links = [link for box in boxes for link in box.orders.all()]
        orders_by_id = {
            assignment.order_id: assignment.order for assignment in assignments
        }
        orders_by_id.update({link.order_id: link.order for link in boxed_links})
        batch.list_unit_count = 0
        batch.list_weight_kg = Decimal("0")
        batch.list_weight_has_value = False
        batch.list_weight_complete = bool(orders_by_id)
        for order in orders_by_id.values():
            for item in order.items.all():
                quantity = int(item.quantity or 0)
                batch.list_unit_count += quantity
                unit_weight = None
                if item.sku is not None:
                    unit_weight = item.sku.weight_gross_kg or item.sku.weight_kg
                if unit_weight is None:
                    batch.list_weight_complete = False
                    continue
                batch.list_weight_has_value = True
                batch.list_weight_kg += Decimal(unit_weight) * quantity
        assigned_order_ids = {assignment.order_id for assignment in assignments}
        boxed_order_ids = {link.order_id for link in boxed_links}
        transfers = [
            transfer
            for order in orders_by_id.values()
            for item in order.items.all()
            for transfer in item.metadata_transfers.all()
        ]
        metadata_missing_count = 0
        metadata_unconfirmed_count = 0
        metadata_blocked_order_numbers = []
        for order in orders_by_id.values():
            order_has_unconfirmed_metadata = False
            for item in order.items.all():
                transfer_by_type = {
                    transfer.metadata_type: transfer
                    for transfer in item.metadata_transfers.all()
                }
                requirements = metadata_requirements(item)
                required_types = set()
                marking_transfer = transfer_by_type.get(
                    FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                )
                if (
                    requirements.marketplace_marking_required
                    or (
                        marking_transfer is not None
                        and marking_transfer.is_required
                        and not ozon_marking_transfer_not_required(
                            marking_transfer
                        )
                    )
                ):
                    required_types.add(
                        FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                    )
                if requirements.expiry_required:
                    required_types.add(
                        FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION
                    )
                for metadata_type in required_types:
                    transfer = transfer_by_type.get(metadata_type)
                    if marketplace_metadata_transfer_resolved(transfer):
                        continue
                    metadata_unconfirmed_count += 1
                    metadata_missing_count += int(transfer is None)
                    order_has_unconfirmed_metadata = True
            if order_has_unconfirmed_metadata:
                metadata_blocked_order_numbers.append(order.external_order_id)
        batch.box_count = len(boxes)
        batch.scanned_box_count = sum(
            box.status
            in {
                FbsHandoverBox.STATUS_SCANNED,
                FbsHandoverBox.STATUS_DISPATCHED,
                FbsHandoverBox.STATUS_ACCEPTED,
            }
            for box in boxes
        )
        batch.assigned_order_count = len(assigned_order_ids)
        batch.boxed_order_count = len(boxed_order_ids)
        batch.label_ready_count = sum(
            any(
                label.status
                in {FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED}
                for label in order.marketplace_labels.all()
            )
            for order in orders_by_id.values()
        )
        batch.metadata_waiting_count = sum(
            transfer.status
            in {
                FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
                FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
                FbsMarketplaceMetadataTransfer.STATUS_SENT,
                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
            }
            for transfer in transfers
        )
        batch.metadata_problem_count = sum(
            transfer.status
            in {
                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED,
                FbsMarketplaceMetadataTransfer.STATUS_CANCELED,
            }
            and not marketplace_metadata_transfer_resolved(transfer)
            for transfer in transfers
        )
        batch.metadata_missing_count = metadata_missing_count
        batch.metadata_unconfirmed_count = metadata_unconfirmed_count
        batch.metadata_blocked_order_numbers = tuple(
            dict.fromkeys(metadata_blocked_order_numbers)
        )
        batch.supply_print_blocked = bool(
            batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and batch.metadata_unconfirmed_count
        )
        batch.supply_print_block_reason = (
            "КИЗ/срок не подтверждены: "
            f"{batch.metadata_unconfirmed_count}"
            + (
                " · заказы "
                + ", ".join(batch.metadata_blocked_order_numbers[:3])
                if batch.metadata_blocked_order_numbers
                else ""
            )
        )
        relevant_commands = _handover_relevant_commands(
            commands,
            order_ids=orders_by_id,
        )
        batch.command_active_count = sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_PENDING,
                FbsMarketplaceCommand.STATUS_SENT,
                FbsMarketplaceCommand.STATUS_RETRY,
            }
            for command in relevant_commands
        )
        batch.command_problem_count = sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_FAILED,
                FbsMarketplaceCommand.STATUS_CONFLICT,
            }
            for command in relevant_commands
        )
        batch.list_order_count = max(
            int(batch.assigned_order_count or 0),
            int(batch.boxed_order_count or 0),
        )
        stages = _handover_stage_rows(batch)
        batch.current_stage = next(
            (stage for stage in stages if stage.state in {"in_progress", "problem"}),
            stages[-1],
        )
        if batch.metadata_missing_count:
            batch.metadata_state = "problem"
            batch.metadata_label = (
                f"Не подтверждены: {batch.metadata_unconfirmed_count}"
            )
        elif batch.metadata_problem_count:
            batch.metadata_state = "problem"
            batch.metadata_label = f"Ошибок: {batch.metadata_problem_count}"
        elif batch.metadata_waiting_count:
            batch.metadata_state = "in_progress"
            batch.metadata_label = f"Ожидают: {batch.metadata_waiting_count}"
        else:
            batch.metadata_state = "done" if batch.list_order_count else "pending"
            batch.metadata_label = "Без ошибок" if batch.list_order_count else "Нет заказов"
        batch.has_problem = bool(
            batch.status != FbsHandoverBatch.STATUS_ARCHIVED
            and (
                batch.status == FbsHandoverBatch.STATUS_PROBLEM
                or batch.marketplace_state == FbsHandoverBatch.MARKETPLACE_ERROR
                or batch.metadata_missing_count
                or batch.metadata_problem_count
                or batch.command_problem_count
            )
        )
        if batch.status == FbsHandoverBatch.STATUS_ARCHIVED:
            batch.ui_status_label = "Архив"
            batch.ui_status_state = "pending"
        elif batch.status == FbsHandoverBatch.STATUS_ACCEPTED:
            batch.ui_status_label = "Принята маркетплейсом"
            batch.ui_status_state = "done"
        elif batch.status == FbsHandoverBatch.STATUS_PROBLEM:
            batch.ui_status_label = "Отклонена / проблема"
            batch.ui_status_state = "problem"
        elif batch.status == FbsHandoverBatch.STATUS_DISPATCHED:
            batch.ui_status_label = "В пути"
            batch.ui_status_state = "in_progress"
        elif (
            batch.status == FbsHandoverBatch.STATUS_READY
            and batch.metadata_unconfirmed_count
        ):
            batch.ui_status_label = "КИЗ/срок не подтверждены"
            batch.ui_status_state = (
                "problem"
                if batch.metadata_missing_count or batch.metadata_problem_count
                else "in_progress"
            )
        elif batch.status == FbsHandoverBatch.STATUS_READY:
            batch.ui_status_label = "Проверена"
            batch.ui_status_state = "done"
        elif batch.list_order_count:
            batch.ui_status_label = "Проверяется"
            batch.ui_status_state = "in_progress"
        else:
            batch.ui_status_label = "Новая"
            batch.ui_status_state = "pending"
    return batches


HANDOVER_REPORT_CHOICES = (
    ("status", "Статус сборки"),
    ("count", "Количество заказов"),
    ("orders", "Отгруженные заказы"),
    ("goods", "Отгруженные товары"),
)
HANDOVER_REPORT_ROW_LIMIT = 50_000


def _handover_report_filter_context(request):
    today = timezone.localdate()
    report_type = str(request.GET.get("report") or "status").strip()
    if report_type not in dict(HANDOVER_REPORT_CHOICES):
        report_type = "status"
    agency_filter = str(request.GET.get("agency_id") or "").strip()
    marketplace_filter = str(request.GET.get("marketplace") or "").strip()
    date_from_filter = str(request.GET.get("date_from") or today.replace(day=1).isoformat()).strip()
    date_to_filter = str(request.GET.get("date_to") or today.isoformat()).strip()
    filter_error = ""
    try:
        date_from = date.fromisoformat(date_from_filter)
        date_to = date.fromisoformat(date_to_filter)
    except ValueError:
        date_from = today.replace(day=1)
        date_to = today
        filter_error = "Укажите даты в корректном формате."
    if date_from > date_to:
        filter_error = "Дата начала не может быть позже даты окончания."
    if not agency_filter.isdigit():
        agency_filter = ""
    if marketplace_filter not in dict(FbsIntegrationProfile.MARKETPLACE_CHOICES):
        marketplace_filter = ""
    return {
        "report_type": report_type,
        "agency_filter": agency_filter,
        "marketplace_filter": marketplace_filter,
        "date_from_filter": date_from_filter,
        "date_to_filter": date_to_filter,
        "date_from": date_from,
        "date_to": date_to,
        "filter_error": filter_error,
    }


def _handover_report_xlsx(*, title, headers, rows, warning=""):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=title[:31])
    if warning:
        sheet.append((warning,))
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        'attachment; filename="fullbox-fbs-shipments-report.xlsx"'
    )
    return response


def _handover_detail_queryset():
    order_items = FbsOrderItem.objects.select_related("sku").prefetch_related(
        Prefetch(
            "metadata_transfers",
            queryset=FbsMarketplaceMetadataTransfer.objects.select_related(
                "traceability__allocation__pick_task__batch__workstation"
            ).order_by("metadata_type", "id"),
        )
    )
    handover_orders = (
        FbsHandoverOrder.objects.select_related(
            "order", "added_by", "excluded_by", "verified_by", "verified_label"
        )
        .prefetch_related(
            Prefetch("order__items", queryset=order_items),
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.order_by("-requested_at", "-id"),
            ),
        )
        .order_by("id")
    )
    boxes = (
        FbsHandoverBox.objects.select_related("scanned_by")
        .prefetch_related(Prefetch("orders", queryset=handover_orders))
        .order_by("id")
    )
    return FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by"
    ).prefetch_related(
        Prefetch("boxes", queryset=boxes),
        Prefetch(
            "order_assignments",
            queryset=FbsHandoverOrderAssignment.objects.select_related(
                "order", "assigned_by"
            )
            .prefetch_related(
                Prefetch("order__items", queryset=order_items),
                Prefetch(
                    "order__marketplace_labels",
                    queryset=FbsOrderLabel.objects.order_by("-requested_at", "-id"),
                ),
            )
            .order_by("id"),
        ),
        Prefetch(
            "marketplace_commands",
            queryset=FbsMarketplaceCommand.objects.select_related(
                "order", "handover_box", "requested_by"
            ).order_by("-created_at", "-id"),
        ),
    )


def _handover_stage_rows(batch: FbsHandoverBatch) -> tuple[SimpleNamespace, ...]:
    statuses = {
        FbsHandoverBatch.STATUS_OPEN: 2,
        FbsHandoverBatch.STATUS_READY: 3,
        FbsHandoverBatch.STATUS_DISPATCHED: 4,
        FbsHandoverBatch.STATUS_ACCEPTED: 5,
        FbsHandoverBatch.STATUS_PROBLEM: 4,
    }
    current = statuses.get(batch.status, 1)
    labels = (
        "Создана",
        "Проверяется",
        "Проверена",
        "Передана",
        "Принята",
    )
    rows = []
    for position, label in enumerate(labels, start=1):
        if batch.status == FbsHandoverBatch.STATUS_ARCHIVED:
            state = "done" if position == 1 else "pending"
            state_label = "Архив" if position == 1 else "Не выполняется"
        elif batch.status == FbsHandoverBatch.STATUS_PROBLEM and position == 5:
            state, state_label = "problem", "Требует решения"
        elif position < current or (
            batch.status == FbsHandoverBatch.STATUS_ACCEPTED and position == current
        ):
            state, state_label = "done", "Готово"
        elif position == current:
            state, state_label = "in_progress", "Текущий этап"
        else:
            state, state_label = "pending", "Ожидает"
        rows.append(
            SimpleNamespace(
                number=position,
                label=label,
                state=state,
                state_label=state_label,
            )
        )
    return tuple(rows)


def _handover_actor_name(user) -> str:
    if user is None:
        return "Система"
    full_name = str(user.get_full_name() or "").strip()
    return full_name or str(user.get_username() or "Система")


def _handover_history_rows(
    batch: FbsHandoverBatch,
    *,
    boxes,
    assignments,
    commands,
) -> tuple[tuple[SimpleNamespace, ...], int]:
    rows = []
    background_reads = {}

    def add(occurred_at, title, detail="", *, user=None, state="pending"):
        if occurred_at is None:
            return
        rows.append(
            SimpleNamespace(
                occurred_at=occurred_at,
                title=title,
                detail=detail,
                actor=_handover_actor_name(user),
                state=state,
            )
        )

    add(
        batch.created_at,
        "Создана отгрузка",
        f"{batch.profile.get_marketplace_display()} · склад "
        f"{batch.profile.external_warehouse_id or 'не указан'}",
        user=batch.created_by,
        state="done",
    )
    if batch.archived_at:
        add(
            batch.archived_at,
            "Отгрузка перенесена в архив",
            batch.archive_reason,
            user=batch.archived_by,
            state="done",
        )
    for assignment in assignments:
        order_number = assignment.order.external_order_id
        add(
            assignment.created_at,
            f"Заказ {order_number} добавлен в состав",
            assignment.get_status_display(),
            user=assignment.assigned_by,
            state="done" if assignment.status == assignment.STATUS_CONFIRMED else "in_progress",
        )
        if assignment.confirmed_at:
            add(
                assignment.confirmed_at,
                f"Marketplace подтвердил заказ {order_number}",
                "Заказ включен в поставку",
                state="done",
            )
        elif assignment.status in {
            assignment.STATUS_ERROR,
            assignment.STATUS_CANCELED,
        }:
            add(
                assignment.updated_at,
                f"Заказ {order_number}: {assignment.get_status_display()}",
                assignment.error,
                user=assignment.assigned_by,
                state="problem",
            )

    for box in boxes:
        add(
            box.created_at,
            f"Создан короб {box.qr_code}",
            box.external_box_id,
            state="done",
        )
        if box.label_ready_at:
            add(
                box.label_ready_at,
                f"Готова этикетка короба {box.qr_code}",
                box.label_format.upper() if box.label_format else "",
                state="done",
            )
        if box.scanned_at:
            add(
                box.scanned_at,
                f"Проверен короб {box.qr_code}",
                "Состав подтвержден сканированием",
                user=box.scanned_by,
                state="done",
            )
        if box.accepted_at:
            add(
                box.accepted_at,
                f"Marketplace принял короб {box.qr_code}",
                state="done",
            )
        for link in box.orders.all():
            order_number = link.order.external_order_id
            add(
                link.added_at,
                f"Заказ {order_number} помещен в короб",
                box.qr_code,
                user=link.added_by,
                state="done",
            )
            if link.verified_at:
                add(
                    link.verified_at,
                    f"Проверен заказ {order_number}",
                    "Этикетка заказа подтверждена сканированием",
                    user=link.verified_by,
                    state="done",
                )
            if link.excluded_at:
                add(
                    link.excluded_at,
                    f"Заказ {order_number} исключен",
                    link.exclusion_reason,
                    user=link.excluded_by,
                    state="problem",
                )

    for command in commands:
        if command.status not in {
            command.STATUS_CONFIRMED,
            command.STATUS_FAILED,
            command.STATUS_CONFLICT,
            command.STATUS_CANCELLED,
        }:
            continue
        order_suffix = (
            f" · заказ {command.order.external_order_id}" if command.order_id else ""
        )
        occurred_at = command.confirmed_at or command.updated_at
        if (
            command.status == command.STATUS_CONFIRMED
            and str(command.command_type or "").startswith(("wb_read_", "ozon_read_"))
        ):
            key = (command.command_type, command.order_id)
            group = background_reads.setdefault(
                key,
                {
                    "count": 0,
                    "occurred_at": occurred_at,
                    "command": command,
                    "order_suffix": order_suffix,
                },
            )
            group["count"] += 1
            if occurred_at and (
                group["occurred_at"] is None or occurred_at > group["occurred_at"]
            ):
                group["occurred_at"] = occurred_at
                group["command"] = command
                group["order_suffix"] = order_suffix
            continue
        add(
            occurred_at,
            f"Marketplace: {command.get_status_display()}",
            f"{command.command_type}{order_suffix}"
            + (f" · {command.error}" if command.error else ""),
            user=command.requested_by,
            state=(
                "done"
                if command.status == command.STATUS_CONFIRMED
                else "problem"
            ),
        )

    for group in background_reads.values():
        command = group["command"]
        count = int(group["count"] or 0)
        detail = f"{command.command_type}{group['order_suffix']}"
        if count > 1:
            detail += f" · {count} фоновых проверок"
        add(
            group["occurred_at"],
            "Marketplace: фоновая проверка подтверждена",
            detail,
            user=command.requested_by,
            state="done",
        )

    if batch.dispatched_at:
        add(
            batch.dispatched_at,
            "Отгрузка передана водителю",
            batch.external_supply_id,
            user=batch.dispatched_by,
            state="done",
        )
    if batch.accepted_at:
        add(
            batch.accepted_at,
            "Отгрузка принята маркетплейсом",
            batch.external_supply_id,
            state="done",
        )
    elif batch.status == FbsHandoverBatch.STATUS_PROBLEM:
        add(
            batch.updated_at,
            "Отгрузка требует решения",
            batch.get_marketplace_state_display(),
            state="problem",
        )

    rows.sort(key=lambda row: row.occurred_at, reverse=True)
    total = len(rows)
    return tuple(rows[:300]), total


def _handover_next_action(
    batch: FbsHandoverBatch,
    *,
    boxes,
    assignment_count: int,
    missing_order_count: int,
    physical_missing_order_count: int = 0,
    unverified_order_count: int = 0,
    all_assignments_confirmed: bool,
    all_orders_ready: bool,
    composition_ready: bool,
    blocking_return_count: int = 0,
    blocking_return_order_numbers=(),
    blocking_return_statuses=(),
    background_return_count: int = 0,
    background_return_order_numbers=(),
) -> SimpleNamespace:
    open_boxes = [box for box in boxes if box.status == FbsHandoverBox.STATUS_OPEN]
    filled_open_boxes = [box for box in open_boxes if box.order_count]
    empty_open_boxes = [box for box in open_boxes if not box.order_count]
    closed_boxes = [box for box in boxes if box.status == FbsHandoverBox.STATUS_CLOSED]
    blocking_return_order_numbers = tuple(
        number
        for number in dict.fromkeys(
            str(number).strip() for number in blocking_return_order_numbers
        )
        if number
    )
    blocking_return_statuses = {
        str(status).strip() for status in blocking_return_statuses if str(status).strip()
    }
    background_return_order_numbers = tuple(
        number
        for number in dict.fromkeys(
            str(number).strip() for number in background_return_order_numbers
        )
        if number
    )

    def background_return_copy():
        numbers = ", ".join(background_return_order_numbers[:5])
        suffix = f" Заказы: {numbers}." if numbers else ""
        return (
            "Проверка продолжена — ждём сверку WB",
            (
                "Контролёр может перейти к следующей работе. Финальная отправка "
                "поставки автоматически останется закрыта до подтверждения WB."
                f"{suffix}"
            ),
        )

    def blocking_return_copy(*, ready: bool):
        if FbsPickRestockRequest.STATUS_FAILED in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Положите товар в назначенную тару отменённых заказов, "
                    "отсканируйте её и продолжайте проверку. Финальная отправка "
                    "останется заблокирована до успешной сверки WB."
                ),
            )
        if FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Положите товар в назначенную тару отменённых заказов и "
                    "отсканируйте её. Сверка WB продолжится в фоне."
                ),
            )
        if FbsPickRestockRequest.STATUS_QUEUED in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Контролёру: положите весь заказ в тару отменённых заказов "
                    "и подтвердите передачу подборщику."
                ),
            )
        if len(blocking_return_order_numbers) == 1:
            order_number = blocking_return_order_numbers[0]
            title = f"Завершить возврат заказа {order_number}"
            if ready:
                description = (
                    f"После физического возврата заказа {order_number} "
                    "он перестанет блокировать состав."
                )
            else:
                description = (
                    f"Подборщик должен вернуть товар заказа {order_number} "
                    "в исходный FBS-короб."
                )
            return title, description
        if blocking_return_order_numbers:
            visible_numbers = ", ".join(blocking_return_order_numbers[:5])
            remaining_count = len(blocking_return_order_numbers) - 5
            suffix = f" и ещё {remaining_count}" if remaining_count > 0 else ""
            return (
                f"Завершить возврат проблемных заказов: {len(blocking_return_order_numbers)}",
                f"Заказы: {visible_numbers}{suffix}. Верните товар в исходные FBS-короба.",
            )
        return (
            "Завершить возврат проблемного заказа",
            (
                "После физического возврата заказ перестанет блокировать состав."
                if ready
                else "Подборщик должен вернуть товар в исходный FBS-короб."
            ),
        )

    def action(code, step, title, description, *, box=None, blocked=False):
        return SimpleNamespace(
            code=code,
            step=step,
            title=title,
            description=description,
            box=box,
            blocked=blocked,
        )

    if batch.status == FbsHandoverBatch.STATUS_ARCHIVED:
        return action(
            "archived",
            1,
            "Отгрузка в архиве",
            "Активных действий нет. История marketplace сохранена.",
        )

    if batch.status == FbsHandoverBatch.STATUS_ACCEPTED:
        return action(
            "complete",
            5,
            "Отгрузка завершена",
            "Marketplace подтвердил приемку. Дополнительных действий не требуется.",
        )
    if batch.status in {
        FbsHandoverBatch.STATUS_DISPATCHED,
        FbsHandoverBatch.STATUS_PROBLEM,
    }:
        return action(
            "refresh",
            5,
            "Проверить приемку marketplace",
            "Обновите статусы после передачи. Ошибки будут показаны отдельно.",
        )
    if batch.status == FbsHandoverBatch.STATUS_READY:
        if blocking_return_count:
            title, description = blocking_return_copy(ready=True)
            return action(
                "await_return",
                2,
                title,
                description,
                blocked=True,
            )
        if background_return_count:
            title, description = background_return_copy()
            return action(
                "wait_marketplace_exclusion",
                3,
                title,
                description,
                blocked=True,
            )
        if not composition_ready:
            return action(
                "resolve_composition",
                2,
                "Устранить ошибки состава",
                "Проверка откроется после подтверждения всех заказов, этикеток, КИЗов и сроков.",
                blocked=True,
            )
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and unverified_order_count
        ):
            return action(
                "verify_orders",
                2,
                "Проверить состав сканированием",
                f"Отсканируйте WB-этикетки оставшихся заказов: {unverified_order_count}.",
            )
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and batch.marketplace_state != FbsHandoverBatch.MARKETPLACE_COMPLETE
        ):
            return action(
                "wait_marketplace_delivery",
                4,
                "Автоматическая отправка в WB",
                "Поставка проверена. Система автоматически передает ее в доставку Wildberries.",
            )
        return action(
            "dispatch",
            4,
            "Передать короб водителю",
            "Все короба подтверждены. Зафиксируйте фактическую передачу водителю.",
        )

    if batch.status != FbsHandoverBatch.STATUS_OPEN:
        return action(
            "refresh_page",
            2,
            "Обновить состояние отгрузки",
            "Статус изменился. Обновите экран перед следующим действием.",
        )

    if blocking_return_count:
        title, description = blocking_return_copy(ready=False)
        return action(
            "await_return",
            2,
            title,
            description,
            blocked=True,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
        and assignment_count
        and not physical_missing_order_count
        and not composition_ready
    ):
        return action(
            "resolve_composition",
            2,
            "Ozon проверяет заказы в фоне",
            (
                "Заказы уже добавлены в транспортный короб по первичному скану. "
                "Повторно сканировать QR не нужно; дождитесь официальных этикеток "
                "или устраните показанную ошибку Ozon."
            ),
            blocked=True,
        )

    if assignment_count and missing_order_count:
        target_box = filled_open_boxes[0] if filled_open_boxes else (
            open_boxes[0] if open_boxes else None
        )
        return action(
            "pack_orders",
            2,
            f"Контрольный скан заказов: {missing_order_count}",
            (
                "Откройте тару проверки. Сканируйте WB-этикетку каждого "
                "действующего заказа и сразу перекладывайте товар в "
                "транспортный короб. После последнего скана закройте тару."
                if batch.profile.marketplace
                == FbsIntegrationProfile.MARKETPLACE_WB
                else
                "Для заказов, подтверждённых до включения одного скана, "
                "отсканируйте QR заказа Ozon при фактическом помещении в "
                "транспортный короб. Новые заказы повторного скана не требуют."
            ),
            box=target_box,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and unverified_order_count
    ):
        if not composition_ready and not background_return_count:
            return action(
                "resolve_composition",
                2,
                "Устранить ошибки состава",
                "Проверка откроется после подтверждения всех заказов, этикеток, КИЗов и сроков.",
                blocked=True,
            )
        return action(
            "verify_orders",
            2,
            "Проверить состав сканированием",
            f"Отсканируйте WB-этикетки оставшихся заказов: {unverified_order_count}.",
        )

    if empty_open_boxes and assignment_count and not missing_order_count:
        box = empty_open_boxes[0]
        return action(
            "blocked_empty_box",
            2,
            "Пустой короб блокирует отгрузку",
            f"Короб {box.qr_code} создан, но в нем нет заказа. Уберите лишний короб до отправки.",
            box=box,
            blocked=True,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and boxes
        and all_assignments_confirmed
        and all_orders_ready
        and not background_return_count
    ):
        return action(
            "wait_marketplace_delivery",
            3,
            "Автоматическая отправка в WB",
            "Состав проверен. Система автоматически передает поставку в доставку Wildberries.",
        )

    if filled_open_boxes:
        box = filled_open_boxes[0]
        return action(
            "close_box",
            2,
            "Закрыть заполненный короб",
            f"Заказы уже находятся в коробе {box.qr_code}. Закройте его, чтобы открыть подтверждение к отгрузке.",
            box=box,
        )

    if background_return_count:
        title, description = background_return_copy()
        return action(
            "wait_marketplace_exclusion",
            3,
            title,
            description,
            blocked=True,
        )

    if closed_boxes:
        box = closed_boxes[0]
        return action(
            "scan_box",
            3,
            "Подтвердить закрытый короб",
            f"Короб {box.qr_code} закрыт. Отсканируйте его QR для подтверждения к отгрузке.",
            box=box,
        )

    if empty_open_boxes:
        box = empty_open_boxes[0]
        return action(
            "pack_orders",
            1,
            "Добавить заказ в открытый короб",
            f"Короб {box.qr_code} открыт. Отсканируйте этикетку следующего заказа.",
            box=box,
        )

    if not boxes:
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and wb_handover_uses_marketplace_boxes(batch)
        ):
            if all_assignments_confirmed and all_orders_ready:
                return action(
                    "create_wb_boxes",
                    1,
                    "Получить QR коробов WB",
                    "Состав поставки подтвержден. Укажите нужное количество транспортных коробов.",
                )
            return action(
                "wait_marketplace",
                1,
                "Дождаться подтверждения состава",
                "Короба можно создать после подтверждения всех заказов marketplace.",
            )
        if batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
            return action(
                "wait_marketplace",
                1,
                "Дождаться транспортного короба",
                "Для поставки на склад или СЦ короб Fullbox появится после подготовки заказа.",
            )
        return action(
            "add_box",
            1,
            "Добавить транспортный короб",
            "Сканируйте QR пустого короба, в который будут складываться проверенные заказы.",
        )

    return action(
        "refresh_page",
        3,
        "Обновить состояние отгрузки",
        "Все доступные короба обработаны. Обновите экран для следующего шага.",
    )


def _handover_restock_blocks_handover(request) -> bool:
    if request.source_tote_id is not None:
        return False
    return bool(
        request.status in HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES
        or (
            request.status == FbsPickRestockRequest.STATUS_QUEUED
            and request.source_tote_id is None
        )
    )


def _handover_assignment_display_priority(assignment, exclusion_by_order_id):
    exclusion_request = exclusion_by_order_id.get(assignment.order_id)
    if exclusion_request is not None and _handover_restock_blocks_handover(
        exclusion_request
    ):
        return (0, exclusion_request.id)
    return (1, assignment.id)


def _handover_exclusion_check_copy(exclusion_request):
    if exclusion_request.status == FbsPickRestockRequest.STATUS_FAILED:
        return (
            "Снимите заказ и освободите проверку",
            (
                "Контролёру: положить весь заказ в назначенную тару отменённых "
                "заказов, отсканировать её и продолжить проверку. Сверка WB "
                "останется отдельным финальным барьером."
            ),
        )
    if exclusion_request.status == FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE:
        return (
            "Снимите заказ и продолжите проверку",
            (
                "Контролёру: положить весь заказ в назначенную тару отменённых "
                "заказов и отсканировать её. Сверка WB продолжится в фоне."
            ),
        )
    if (
        exclusion_request.status == FbsPickRestockRequest.STATUS_QUEUED
        and exclusion_request.source_tote_id is None
    ):
        return (
            "Передать в возврат отбора",
            (
                "Контролёру: положить весь заказ в тару отменённых заказов "
                "и подтвердить передачу подборщику."
            ),
        )
    return (
        exclusion_request.get_status_display(),
        exclusion_request.reason or "Заказ исключается из поставки.",
    )


def _handover_detail_summary(
    batch: FbsHandoverBatch,
    *,
    compact_controller: bool = False,
) -> dict:
    from billing.models import WarehouseServiceFact

    from .services.handover import (
        MARKETPLACE_PROBLEM_STATUSES,
        handover_order_has_accepted_marketplace_status,
    )
    from .services.traceability import (
        marketplace_metadata_transfer_resolved,
        metadata_requirements,
        ozon_marking_transfer_not_required,
    )

    boxes = list(batch.boxes.all())
    all_assignments = list(batch.order_assignments.all())
    assignments = [
        assignment
        for assignment in all_assignments
        if assignment.status != FbsHandoverOrderAssignment.STATUS_CANCELED
    ]
    excluded_assignments = [
        assignment
        for assignment in all_assignments
        if assignment.status == FbsHandoverOrderAssignment.STATUS_CANCELED
    ]
    exclusion_requests = list(
        FbsPickRestockRequest.objects.filter(
            handover_assignment__batch=batch,
            order__isnull=False,
        )
        .select_related("order", "assigned_to", "created_by", "source_tote")
        .order_by("created_at", "id")
    )
    exclusion_by_order_id = {
        request.order_id: request for request in exclusion_requests
    }
    assignments.sort(
        key=lambda assignment: _handover_assignment_display_priority(
            assignment,
            exclusion_by_order_id,
        )
    )
    commands = list(batch.marketplace_commands.all())
    relevant_commands = _handover_relevant_commands(
        commands,
        order_ids=(assignment.order_id for assignment in assignments),
    )
    order_ids = []
    counted_order_ids = set()
    total_units = 0
    total_lines = 0
    known_weight = Decimal("0")
    weight_complete = True
    label_document_count = 0
    metadata_problem_count = 0
    metadata_waiting_count = 0
    accepted_order_count = 0
    rejected_order_count = 0
    handover_links_by_order_id = {}
    acceptance_is_active = batch.status in {
        FbsHandoverBatch.STATUS_DISPATCHED,
        FbsHandoverBatch.STATUS_PROBLEM,
        FbsHandoverBatch.STATUS_ACCEPTED,
    }

    problem_transfer_statuses = {
        FbsMarketplaceMetadataTransfer.STATUS_FAILED,
        FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
        FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED,
        FbsMarketplaceMetadataTransfer.STATUS_CANCELED,
    }
    waiting_transfer_statuses = {
        FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
        FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
        FbsMarketplaceMetadataTransfer.STATUS_SENT,
        FbsMarketplaceMetadataTransfer.STATUS_RETRY,
    }

    def prepare_order(order):
        nonlocal total_units, total_lines, known_weight, weight_complete
        nonlocal label_document_count, metadata_problem_count, metadata_waiting_count
        nonlocal accepted_order_count, rejected_order_count

        items = list(order.items.all())
        labels = list(order.marketplace_labels.all())
        order_units = 0
        order_weight = Decimal("0")
        order_weight_complete = True
        required_transfer_count = 0
        confirmed_required_count = 0
        order_problem = False
        order_problem_label = ""
        order_has_invalid_kiz = False
        invalid_kiz_workstation_id = None
        order_waiting = False

        for item in items:
            item.photo_url = str(getattr(item.sku, "img", "") or "")
            item.transfers = list(item.metadata_transfers.all())
            requirements = metadata_requirements(item)
            transfer_by_type = {
                transfer.metadata_type: transfer for transfer in item.transfers
            }
            item.metadata_checks = []
            for metadata_type, label, required in (
                (
                    FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE,
                    "КИЗ",
                    requirements.marketplace_marking_required,
                ),
                (
                    FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION,
                    "Срок годности",
                    requirements.expiry_required,
                ),
            ):
                transfer = transfer_by_type.get(metadata_type)
                effective_required = bool(
                    required
                    or (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and transfer is not None
                        and transfer.is_required
                        and not ozon_marking_transfer_not_required(transfer)
                    )
                )
                if not effective_required:
                    state, state_label = "pending", "Не требуется"
                elif marketplace_metadata_transfer_resolved(transfer):
                    state, state_label = "done", "Подтверждено"
                elif transfer is None or transfer.status in waiting_transfer_statuses:
                    state, state_label = "in_progress", "Ожидает подтверждения"
                elif transfer.status in problem_transfer_statuses:
                    state, state_label = "problem", transfer.get_status_display()
                else:
                    state, state_label = "in_progress", transfer.get_status_display()
                item.metadata_checks.append(
                    SimpleNamespace(
                        label=label,
                        required=effective_required,
                        state=state,
                        state_label=state_label,
                        transfer=transfer,
                        can_retry=(
                            metadata_type
                            == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                            and transfer is not None
                            and transfer.status
                            in {
                                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                            }
                            and batch.status
                            in {
                                FbsHandoverBatch.STATUS_OPEN,
                                FbsHandoverBatch.STATUS_READY,
                            }
                            and batch.profile.marketplace
                            == FbsIntegrationProfile.MARKETPLACE_WB
                            and str(order.marketplace_status or "").strip().casefold()
                            == "confirm"
                        ),
                        retry_block_reason=(
                            ""
                            if metadata_type
                            != FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                            or transfer is None
                            or transfer.status
                            not in {
                                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                            }
                            else (
                                "Повтор доступен только в открытой отгрузке."
                                if batch.status
                                not in {
                                    FbsHandoverBatch.STATUS_OPEN,
                                    FbsHandoverBatch.STATUS_READY,
                                }
                                else (
                                    "Повтор доступен только для WB."
                                    if batch.profile.marketplace
                                    != FbsIntegrationProfile.MARKETPLACE_WB
                                    else (
                                        "WB разрешает повтор КИЗа только в статусе confirm; "
                                        f"сейчас {order.marketplace_status or 'статус не получен'}."
                                        if str(order.marketplace_status or "")
                                        .strip()
                                        .casefold()
                                        != "confirm"
                                        else ""
                                    )
                                )
                            )
                        ),
                    )
                )
            required_types = set()
            marking_transfer = transfer_by_type.get(
                FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
            )
            if requirements.marketplace_marking_required or (
                marking_transfer is not None
                and marking_transfer.is_required
                and not ozon_marking_transfer_not_required(marking_transfer)
            ):
                required_types.add(FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE)
            if requirements.expiry_required:
                required_types.add(FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION)
            required_transfer_count += len(required_types)
            for metadata_type in required_types:
                transfer = transfer_by_type.get(metadata_type)
                if marketplace_metadata_transfer_resolved(transfer):
                    confirmed_required_count += 1
                elif transfer is None or transfer.status in waiting_transfer_statuses:
                    order_waiting = True
                elif transfer.status in problem_transfer_statuses:
                    order_problem = True
                    if (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and is_final_wb_marking_rejection(transfer)
                    ):
                        allocation = (
                            transfer.traceability.allocation
                            if transfer.traceability_id
                            else None
                        )
                        if (
                            allocation is not None
                            and allocation.status
                            == FbsOrderStockAllocation.STATUS_PICKED
                            and allocation.pick_task_id
                        ):
                            order_has_invalid_kiz = True
                            invalid_kiz_workstation_id = (
                                allocation.pick_task.batch.workstation_id
                            )
                    if (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and is_final_wb_marking_rejection(transfer)
                    ):
                        order_problem_label = "Невалиден"
                    elif transfer.status in {
                        FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                        FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                    }:
                        order_problem_label = "Ошибка проверки WB"
                    elif (
                        not order_problem_label
                        and transfer.status
                        == FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED
                    ):
                        order_problem_label = "Проверка недоступна"
                    elif not order_problem_label:
                        order_problem_label = transfer.get_status_display()

            quantity = int(item.quantity or 0)
            order_units += quantity
            unit_weight = None
            if item.sku is not None:
                unit_weight = item.sku.weight_gross_kg or item.sku.weight_kg
            if unit_weight is None:
                order_weight_complete = False
            else:
                order_weight += Decimal(unit_weight) * quantity

        if order_problem:
            order.metadata_state = "problem"
            order.metadata_state_label = order_problem_label or "Есть ошибка"
        elif order_waiting:
            order.metadata_state = "in_progress"
            order.metadata_state_label = "Ожидает"
        elif required_transfer_count:
            order.metadata_state = "done"
            order.metadata_state_label = (
                f"Подтверждено {confirmed_required_count} из {required_transfer_count}"
            )
        else:
            order.metadata_state = "pending"
            order.metadata_state_label = "Не требуется"

        order.ui_items = items
        order.has_invalid_kiz = order_has_invalid_kiz
        order.invalid_kiz_workstation_id = invalid_kiz_workstation_id
        order.ui_labels = labels
        order.latest_label = labels[0] if labels else None
        kiz_retry_label_number = ""
        if order.latest_label is not None:
            kiz_retry_label_number = (
                str(order.latest_label.external_label_id or "").strip()
                or str(order.latest_label.barcode or "").strip()
            )
        for item in items:
            item.kiz_retry_label_number = kiz_retry_label_number
            item.kiz_retry_product_barcode = str(item.barcode or "").strip()
        normalized_marketplace_status = str(order.marketplace_status or "").strip().casefold()
        if acceptance_is_active and normalized_marketplace_status in MARKETPLACE_PROBLEM_STATUSES:
            order.acceptance_state = "problem"
            order.acceptance_label = "Отклонен"
        elif acceptance_is_active and handover_order_has_accepted_marketplace_status(
            order,
            marketplace=batch.profile.marketplace,
        ):
            order.acceptance_state = "done"
            order.acceptance_label = "Принят"
        elif acceptance_is_active:
            order.acceptance_state = "in_progress"
            order.acceptance_label = "Не принят"
        else:
            order.acceptance_state = "pending"
            order.acceptance_label = "Не передан"
        order.unit_count = order_units
        order.line_count = len(items)
        order.weight_kg = order_weight
        order.weight_complete = order_weight_complete

        if order.id not in counted_order_ids:
            counted_order_ids.add(order.id)
            order_ids.append(order.id)
            total_units += order_units
            total_lines += len(items)
            known_weight += order_weight
            weight_complete = weight_complete and order_weight_complete
            label_document_count += sum(bool(label.document_url) for label in labels)
            metadata_problem_count += int(order_problem)
            metadata_waiting_count += int(order_waiting and not order_problem)
            accepted_order_count += int(order.acceptance_state == "done")
            rejected_order_count += int(order.acceptance_state == "problem")
        return order_units, order_weight, order_weight_complete

    for assignment in assignments:
        prepare_order(assignment.order)
        assignment.latest_label = assignment.order.latest_label
        label_payload = (
            assignment.latest_label.payload
            if assignment.latest_label is not None
            and isinstance(assignment.latest_label.payload, dict)
            else {}
        )
        assignment.uses_preconfirmed_ozon_qr = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and assignment.latest_label is not None
            and label_payload.get("_preloaded_ozon_order_barcode") is True
            and label_payload.get("_preconfirmed_order_barcode_at")
        )

    retryable_kiz_items = []
    retryable_kiz_item_ids = set()
    for assignment in assignments:
        for item in assignment.order.ui_items:
            if item.id in retryable_kiz_item_ids:
                continue
            if not any(
                check.label == "КИЗ" and check.can_retry
                for check in item.metadata_checks
            ):
                continue
            retryable_kiz_item_ids.add(item.id)
            retryable_kiz_items.append(
                SimpleNamespace(
                    order_item_id=item.id,
                    order_number=assignment.order.external_order_id,
                    product_name=item.product_name or item.external_sku,
                    label_number=item.kiz_retry_label_number,
                    product_barcode=item.kiz_retry_product_barcode,
                )
            )

    invalid_kiz_workstation_ids = {
        assignment.order.invalid_kiz_workstation_id
        for assignment in assignments
        if assignment.order.has_invalid_kiz
        and assignment.order.invalid_kiz_workstation_id
    }
    invalid_kiz_sessions_by_workstation = {
        session.workstation_id: session
        for session in FbsControllerSession.objects.select_related(
            "problem_tote__binding", "workstation"
        ).filter(
            workstation_id__in=invalid_kiz_workstation_ids,
            status=FbsControllerSession.STATUS_ACTIVE,
            problem_tote__isnull=False,
        )
    }

    for box in boxes:
        box_units = 0
        box_weight = Decimal("0")
        box_weight_complete = True
        links = [
            link
            for link in box.orders.all()
            if link.status == FbsHandoverOrder.STATUS_ACTIVE
        ]
        for link in links:
            handover_links_by_order_id[link.order_id] = link
            order_units, order_weight, order_weight_complete = prepare_order(link.order)
            box_units += order_units
            box_weight += order_weight
            box_weight_complete = box_weight_complete and order_weight_complete

        box.ui_orders = links
        box.order_count = len(links)
        box.is_empty_history = not links and box.orders.count() > 0
        box.unit_count = box_units
        box.weight_kg = box_weight
        box.weight_complete = box_weight_complete

    verification_rows = []
    requires_composition_scan = (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
    )
    verification_override_active = has_handover_verification_override(batch)
    for assignment in assignments:
        order = assignment.order
        order.display_created_at = order.ordered_at or order.imported_at
        exclusion_request = exclusion_by_order_id.get(order.id)
        link = handover_links_by_order_id.get(order.id)
        assignment.exclusion_request = exclusion_request
        assignment.can_controller_stage_exclusion = bool(
            exclusion_request is not None
            and exclusion_request.source_tote_id is None
            and exclusion_request.status
            in {
                FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
                FbsPickRestockRequest.STATUS_QUEUED,
                FbsPickRestockRequest.STATUS_FAILED,
            }
        )
        assignment.ozon_client_canceled = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and order_is_client_canceled_by_marketplace(order)
        )
        assignment.can_return_ozon_canceled = bool(
            assignment.ozon_client_canceled
            and batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and assignment.status
            in {
                FbsHandoverOrderAssignment.STATUS_PENDING,
                FbsHandoverOrderAssignment.STATUS_CONFIRMED,
                FbsHandoverOrderAssignment.STATUS_ERROR,
            }
            and (
                exclusion_request is None
                or (
                    exclusion_request.status == FbsPickRestockRequest.STATUS_QUEUED
                    and exclusion_request.source_tote_id is None
                )
            )
        )
        assignment.can_exclude = (
            batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and exclusion_request is None
            and link is not None
            and assignment.status
            in {
                FbsHandoverOrderAssignment.STATUS_CONFIRMED,
                FbsHandoverOrderAssignment.STATUS_ERROR,
            }
        )
        assignment.exclusion_requires_seller_cancel = (
            assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        )
        assignment.invalid_kiz_route_available = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and batch.status
            in {FbsHandoverBatch.STATUS_OPEN, FbsHandoverBatch.STATUS_READY}
            and assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
            and order.has_invalid_kiz
            and str(order.marketplace_status or "").strip().casefold() == "confirm"
        )
        invalid_kiz_session = invalid_kiz_sessions_by_workstation.get(
            order.invalid_kiz_workstation_id
        )
        invalid_kiz_problem_tote = (
            invalid_kiz_session.problem_tote if invalid_kiz_session is not None else None
        )
        invalid_kiz_binding = (
            getattr(invalid_kiz_problem_tote, "binding", None)
            if invalid_kiz_problem_tote is not None
            else None
        )
        assignment.invalid_kiz_problem_tote = invalid_kiz_problem_tote
        assignment.can_reroute_invalid_kiz = bool(
            assignment.invalid_kiz_route_available
            and invalid_kiz_binding is not None
            and invalid_kiz_binding.state == FbsToteBinding.STATE_AT_CONTROL
            and invalid_kiz_binding.workstation_id
            == order.invalid_kiz_workstation_id
            and invalid_kiz_binding.controller_session_id
            == invalid_kiz_session.id
        )
        label = (
            link.verified_label
            if link is not None and link.verified_label_id
            else assignment.latest_label
        )
        label_scanned = bool(
            label is not None and label.status == FbsOrderLabel.STATUS_APPLIED
        )
        composition_verified = bool(
            link is not None
            and (
                link.verified_at is not None or verification_override_active
                if requires_composition_scan
                else label_scanned
            )
        )
        if (
            exclusion_request is not None
            and exclusion_request.status != FbsPickRestockRequest.STATUS_COMPLETED
        ):
            state = "problem"
            state_label, state_detail = _handover_exclusion_check_copy(
                exclusion_request
            )
        elif assignment.status in {
            FbsHandoverOrderAssignment.STATUS_ERROR,
            FbsHandoverOrderAssignment.STATUS_CANCELED,
        }:
            state = "problem"
            state_label = "Не ОК"
            state_detail = assignment.error or assignment.get_status_display()
        elif assignment.status == FbsHandoverOrderAssignment.STATUS_PENDING:
            state = "in_progress"
            state_label = "Сверяется с WB"
            state_detail = assignment.error or "Ожидается подтверждение состава поставки WB."
        elif composition_verified:
            state = "done"
            if link.verified_at is not None:
                state_label = "ОК · Проверен"
                state_detail = f"Проверен в коробе {link.box.qr_code}."
            else:
                state_label = "ОК · Разрешено"
                state_detail = (
                    "Первичный скан подтвержден; повторную проверку разрешил "
                    f"начальник склада · короб {link.box.qr_code}."
                )
        elif label_scanned and link is not None:
            state = "in_progress"
            state_label = "Ожидает проверки"
            state_detail = f"Нужен контрольный скан · короб {link.box.qr_code}."
        elif label_scanned:
            state = "in_progress"
            state_label = "Ожидает упаковки"
            state_detail = "Этикетка отсканирована, короб не назначен."
        else:
            state = "problem"
            state_label = "Не проверен"
            state_detail = "Заказ не отсканирован."
        assignment.check_state = state
        assignment.check_label = state_label
        assignment.check_detail = state_detail
        assignment.handover_link = link
        verification_rows.append(
            SimpleNamespace(
                assignment=assignment,
                order=order,
                label=label,
                handover_link=link,
                state=state,
                state_label=state_label,
                state_detail=state_detail,
                is_scanned=composition_verified,
                exclusion_request=exclusion_request,
                can_exclude=assignment.can_exclude,
            )
        )

    verified_order_count = sum(row.is_scanned for row in verification_rows)
    missing_order_count = len(verification_rows) - verified_order_count
    unboxed_order_count = sum(
        row.handover_link is None or row.label is None
        or row.label.status != FbsOrderLabel.STATUS_APPLIED
        for row in verification_rows
    )
    physical_missing_order_count = sum(
        row.handover_link is None for row in verification_rows
    )
    assignment_confirmed_count = sum(
        assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        for assignment in assignments
    )
    all_assignments_confirmed = bool(assignments) and all(
        assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        for assignment in assignments
    )
    all_orders_ready = bool(assignments) and all(
        assignment.order.internal_status == FbsOrder.STATUS_READY_FOR_HANDOVER
        for assignment in assignments
    )
    active_boxes = [box for box in boxes if box.order_count]
    active_check_tote = (
        FbsControllerCheckTote.objects.select_related(
            "tote", "session__workstation"
        )
        .filter(
            handover_batch=batch,
            status__in=(
                FbsControllerCheckTote.STATUS_OPEN,
                FbsControllerCheckTote.STATUS_WAITING_KIZ,
                FbsControllerCheckTote.STATUS_READY,
                FbsControllerCheckTote.STATUS_COMPOSITION,
            ),
        )
        .first()
    )
    if compact_controller and batch.status in {
        FbsHandoverBatch.STATUS_ACCEPTED,
        FbsHandoverBatch.STATUS_ARCHIVED,
    }:
        # The authoritative composition gate is only actionable before dispatch.
        # Re-running it for a finished controller screen performs one query per
        # order and cannot change the next action or any available operation.
        composition_readiness = SimpleNamespace(ready=True, reasons=())
    else:
        composition_readiness = handover_composition_readiness(batch)
    archive_readiness = (
        SimpleNamespace(ready=False, reasons=())
        if compact_controller
        else handover_archive_readiness(batch)
    )
    blocking_return_requests = [
        request
        for request in exclusion_requests
        if _handover_restock_blocks_handover(request)
    ]
    blocking_return_count = len(blocking_return_requests)
    background_return_requests = [
        request
        for request in exclusion_requests
        if request.source_tote_id is not None
        and request.status in HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES
    ]
    background_return_count = len(background_return_requests)
    blocking_return_order_numbers = tuple(
        str(request.order.external_order_id or request.order_id).strip()
        for request in blocking_return_requests
    )
    next_action = _handover_next_action(
        batch,
        boxes=active_boxes,
        assignment_count=len(assignments),
        missing_order_count=unboxed_order_count,
        physical_missing_order_count=physical_missing_order_count,
        unverified_order_count=missing_order_count,
        all_assignments_confirmed=all_assignments_confirmed,
        all_orders_ready=all_orders_ready,
        composition_ready=composition_readiness.ready,
        blocking_return_count=blocking_return_count,
        blocking_return_order_numbers=blocking_return_order_numbers,
        blocking_return_statuses=tuple(
            request.status for request in blocking_return_requests
        ),
        background_return_count=background_return_count,
        background_return_order_numbers=tuple(
            str(request.order.external_order_id or request.order_id).strip()
            for request in background_return_requests
        ),
    )

    box_count = len(active_boxes)
    accepted_boxes = sum(
        box.status == FbsHandoverBox.STATUS_ACCEPTED for box in active_boxes
    )
    problem_boxes = sum(
        box.status == FbsHandoverBox.STATUS_PROBLEM for box in active_boxes
    )
    unique_order_count = len(counted_order_ids)
    if unique_order_count and accepted_order_count == unique_order_count:
        acceptance_label, acceptance_class = "Принята полностью", "done"
    elif accepted_order_count:
        acceptance_label, acceptance_class = "Принята частично", "in_progress"
    elif rejected_order_count or problem_boxes or batch.status == FbsHandoverBatch.STATUS_PROBLEM:
        acceptance_label, acceptance_class = "Отклонена", "problem"
    elif batch.status == FbsHandoverBatch.STATUS_ACCEPTED and box_count:
        acceptance_label, acceptance_class = "Принята полностью", "done"
    elif accepted_boxes:
        acceptance_label, acceptance_class = "Принята частично", "in_progress"
    elif batch.status == FbsHandoverBatch.STATUS_DISPATCHED:
        acceptance_label, acceptance_class = "Ожидает приемки", "in_progress"
    else:
        acceptance_label, acceptance_class = "Не передана", "pending"

    billing_facts = []
    handover_history = ()
    handover_history_total = 0
    if not compact_controller:
        billing_order_ids = [f"FBS-ORDER-{order_id}" for order_id in set(order_ids)]
        billing_order_ids.append(f"FBS-DELIVERY-{batch.id}")
        billing_facts = list(
            WarehouseServiceFact.objects.filter(
                client=batch.profile.agency,
                order_type=WarehouseServiceFact.ORDER_FBS,
                order_id__in=billing_order_ids,
            )
            .select_related("service", "charge")
            .order_by("performed_at", "id")
        )
        for fact in billing_facts:
            fact.display_name = fact.service_name_snapshot or fact.service.name

        handover_history, handover_history_total = _handover_history_rows(
            batch,
            boxes=boxes,
            assignments=all_assignments,
            commands=commands,
        )
    pending_verification_rows = []
    related_pick_batches = (
        FbsPickBatch.objects.filter(
            controller_pick_tote__check_tote__handover_batch=batch,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .order_by("id")
        .distinct()
    )
    assigned_order_ids = {assignment.order_id for assignment in assignments}
    for pick_batch in related_pick_batches:
        for row in _verification_wave_rows(pick_batch):
            if row.order.id in assigned_order_ids:
                continue
            row.pick_batch = pick_batch
            if not row.is_product_complete:
                row.missing_reason = (
                    f"Товар проверен {row.verified_qty} из {row.planned_qty} шт."
                )
            elif not row.is_label_complete:
                row.missing_reason = "QR заказа не создан или не отсканирован."
            else:
                row.missing_reason = "QR отсканирован, заказ еще не добавлен в отгрузку."
            pending_verification_rows.append(row)

    reconciliation_available = (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and batch.status
        in {
            FbsHandoverBatch.STATUS_OPEN,
            FbsHandoverBatch.STATUS_READY,
        }
        and bool(verification_rows or pending_verification_rows)
    )
    reconciliation_missing_count = missing_order_count + len(pending_verification_rows)

    return {
        "handover_boxes": active_boxes,
        "handover_check_tote": active_check_tote,
        "handover_excluded_boxes": [box for box in boxes if not box.order_count],
        "handover_assignments": assignments,
        "handover_excluded_assignments": excluded_assignments,
        "handover_exclusion_requests": exclusion_requests,
        "handover_kiz_retry_items": retryable_kiz_items,
        "handover_kiz_retry_count": len(retryable_kiz_items),
        "handover_blocking_return_count": blocking_return_count,
        "handover_verification_rows": verification_rows,
        "handover_verified_order_count": verified_order_count,
        "handover_missing_order_count": missing_order_count,
        "handover_unboxed_order_count": unboxed_order_count,
        "handover_physical_missing_order_count": physical_missing_order_count,
        "handover_reconciliation_available": reconciliation_available,
        "handover_reconciliation_total_count": (
            len(verification_rows) + len(pending_verification_rows)
        ),
        "handover_reconciliation_missing_count": reconciliation_missing_count,
        "handover_scan_complete": bool(verification_rows)
        and not reconciliation_missing_count,
        "handover_composition_ready": composition_readiness.ready,
        "handover_composition_blockers": composition_readiness.reasons,
        "handover_archive_ready": archive_readiness.ready,
        "handover_archive_blockers": archive_readiness.reasons,
        "handover_exclusion_reason_choices": FbsPickRestockRequest.REASON_CHOICES,
        "handover_assignment_confirmed_count": assignment_confirmed_count,
        "handover_printable_order_label_count": sum(
            bool(
                assignment.latest_label
                and assignment.latest_label.status
                in {FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED}
                and assignment.latest_label.file
                and assignment.latest_label.label_format
                in {FbsOrderLabel.FORMAT_PNG, FbsOrderLabel.FORMAT_PDF}
            )
            for assignment in assignments
        ),
        "handover_all_assignments_confirmed": all_assignments_confirmed,
        "handover_all_orders_ready": all_orders_ready,
        "handover_next_action": next_action,
        "handover_commands": commands,
        "handover_commands_active": any(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_PENDING,
                FbsMarketplaceCommand.STATUS_SENT,
                FbsMarketplaceCommand.STATUS_RETRY,
            }
            for command in relevant_commands
        ),
        "handover_command_problem_count": sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_FAILED,
                FbsMarketplaceCommand.STATUS_CONFLICT,
            }
            for command in relevant_commands
        ),
        "handover_stages": _handover_stage_rows(batch),
        "handover_order_count": unique_order_count,
        "handover_unit_count": total_units,
        "handover_line_count": total_lines,
        "handover_weight_kg": known_weight,
        "handover_weight_complete": weight_complete and bool(order_ids),
        "handover_accepted_boxes": accepted_boxes,
        "handover_problem_boxes": problem_boxes,
        "handover_accepted_orders": accepted_order_count,
        "handover_rejected_orders": rejected_order_count,
        "handover_acceptance_label": acceptance_label,
        "handover_acceptance_class": acceptance_class,
        "handover_label_document_count": label_document_count,
        "handover_metadata_problem_count": metadata_problem_count,
        "handover_metadata_waiting_count": metadata_waiting_count,
        "handover_driver_manifest_ready": bool(boxes)
        and batch.status != FbsHandoverBatch.STATUS_OPEN,
        "handover_billing_facts": billing_facts,
        "handover_billing_charged_count": sum(
            fact.status == WarehouseServiceFact.STATUS_CHARGED for fact in billing_facts
        ),
        "handover_print_workstations": configured_fbs_print_workstations(),
        "handover_uses_marketplace_boxes": wb_handover_uses_marketplace_boxes(batch),
        "handover_history": handover_history,
        "handover_history_total": handover_history_total,
        "handover_pending_verification_rows": pending_verification_rows,
        "handover_pending_verification_count": len(pending_verification_rows),
        "handover_fbs_otg_locations_configured": WarehouseLocation.objects.filter(
            warehouse_code="MSK",
            zone_code="OTG",
            is_active=True,
            is_topology_visible=False,
            is_fbs_visible=True,
        ).exists(),
    }


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_handover(request):
    error = ""
    filter_error = ""
    if request.method == "POST":
        try:
            action = str(request.POST.get("action") or "").strip()
            if action == "bulk_print_supply_labels":
                raw_batch_ids = request.POST.getlist("batch_ids")
                try:
                    batch_ids = list(dict.fromkeys(int(value) for value in raw_batch_ids))
                except (TypeError, ValueError):
                    raise ValueError("Некорректный список отгрузок для печати.")
                if not batch_ids:
                    raise ValueError("Выберите хотя бы одну отгрузку.")
                if len(batch_ids) > 100:
                    raise ValueError("За один раз можно напечатать не более 100 ШК поставок.")
                selected_batches = list(
                    FbsHandoverBatch.objects.filter(pk__in=batch_ids)
                    .only("id", "supply_label_file")
                    .order_by("id")
                )
                if len(selected_batches) != len(batch_ids):
                    raise ValueError("Одна из выбранных отгрузок не найдена.")
                missing_labels = [batch.id for batch in selected_batches if not batch.supply_label_file]
                if missing_labels:
                    missing_text = ", ".join(f"#{batch_id}" for batch_id in missing_labels[:10])
                    raise ValueError(f"ШК поставки еще не готов: {missing_text}.")
                workstation_id = int(request.POST.get("workstation_id") or 0)
                desktop_target = _desktop_label_print_target(request, shipping=True)
                with transaction.atomic():
                    for selected_batch in selected_batches:
                        queue_fbs_handover_supply_label_print(
                            batch_id=selected_batch.id,
                            workstation_id=workstation_id,
                            requested_by=request.user,
                            **desktop_target,
                        )
                redirect_query = request.GET.copy()
                redirect_query["printed_supplies"] = str(len(selected_batches))
                redirect_url = reverse("fbs:tsd_handover")
                if redirect_query:
                    redirect_url = f"{redirect_url}?{redirect_query.urlencode()}"
                return redirect(redirect_url)
            profile = FbsIntegrationProfile.objects.get(pk=int(request.POST.get("profile_id") or 0))
            batch = create_handover_batch(
                profile=profile,
                external_supply_id=request.POST.get("external_supply_id", ""),
                created_by=request.user,
            )
            return redirect("fbs:tsd_handover_detail", batch_id=batch.id)
        except (FbsError, ValueError, FbsIntegrationProfile.DoesNotExist) as exc:
            error = str(exc)
    queryset = _handover_queryset()
    search = str(request.GET.get("q") or "").strip()
    status_filter = str(request.GET.get("status") or "").strip()
    marketplace_filter = str(request.GET.get("marketplace") or "").strip()
    agency_filter = str(request.GET.get("agency_id") or "").strip()
    date_field_filter = str(request.GET.get("date_field") or "created").strip()
    date_from_filter = str(request.GET.get("date_from") or "").strip()
    date_to_filter = str(request.GET.get("date_to") or "").strip()
    date_field_choices = (
        ("created", "Дата создания"),
        ("verified", "Дата проверки"),
        ("dispatched", "Дата отправки"),
    )
    if date_field_filter not in dict(date_field_choices):
        date_field_filter = "created"
    if search:
        search_filter = (
            Q(external_supply_id__icontains=search)
            | Q(external_name__icontains=search)
            | Q(profile__agency__agn_name__icontains=search)
            | Q(profile__external_warehouse_id__icontains=search)
        )
        if search.isdigit():
            search_filter |= Q(id=int(search))
        queryset = queryset.filter(search_filter)
    if status_filter in dict(FbsHandoverBatch.STATUS_CHOICES):
        queryset = queryset.filter(status=status_filter)
    else:
        queryset = queryset.exclude(status=FbsHandoverBatch.STATUS_ARCHIVED)
    if marketplace_filter in dict(FbsIntegrationProfile.MARKETPLACE_CHOICES):
        queryset = queryset.filter(profile__marketplace=marketplace_filter)
    if agency_filter.isdigit():
        queryset = queryset.filter(profile__agency_id=int(agency_filter))

    parsed_date_from = None
    parsed_date_to = None
    try:
        if date_from_filter:
            parsed_date_from = date.fromisoformat(date_from_filter)
        if date_to_filter:
            parsed_date_to = date.fromisoformat(date_to_filter)
    except ValueError:
        filter_error = "Укажите даты в формате ДД.ММ.ГГГГ."
    if parsed_date_from and parsed_date_to and parsed_date_from > parsed_date_to:
        filter_error = "Дата начала не может быть позже даты окончания."
    if not filter_error and (parsed_date_from or parsed_date_to):
        date_lookup = {
            "created": "created_at__date",
            "verified": "boxes__orders__verified_at__date",
            "dispatched": "dispatched_at__date",
        }[date_field_filter]
        if parsed_date_from:
            queryset = queryset.filter(**{f"{date_lookup}__gte": parsed_date_from})
        if parsed_date_to:
            queryset = queryset.filter(**{f"{date_lookup}__lte": parsed_date_to})
        if date_field_filter == "verified":
            queryset = queryset.distinct()
    page_size_choices = (25, 50, 100, 200, 500)
    try:
        page_size = int(request.GET.get("page_size") or 50)
    except (TypeError, ValueError):
        page_size = 50
    if page_size not in page_size_choices:
        page_size = 50
    paginator = Paginator(queryset.order_by("-created_at", "-id"), page_size)
    page_obj = paginator.get_page(request.GET.get("page"))
    batches = _prepare_handover_list_rows(list(page_obj.object_list))
    handover_totals = queryset.aggregate(
        active=Count(
            "id",
            filter=Q(
                status__in=(
                    FbsHandoverBatch.STATUS_OPEN,
                    FbsHandoverBatch.STATUS_READY,
                    FbsHandoverBatch.STATUS_DISPATCHED,
                )
            ),
        ),
        ready=Count("id", filter=Q(status=FbsHandoverBatch.STATUS_READY)),
        dispatched=Count("id", filter=Q(status=FbsHandoverBatch.STATUS_DISPATCHED)),
        problem=Count("id", filter=Q(status=FbsHandoverBatch.STATUS_PROBLEM)),
    )
    handover_totals["archived"] = _handover_queryset().filter(
        status=FbsHandoverBatch.STATUS_ARCHIVED
    ).count()
    client_profiles = (
        FbsIntegrationProfile.objects.filter(handover_batches__isnull=False)
        .select_related("agency")
        .order_by("agency__agn_name", "agency_id", "id")
        .distinct()
    )
    handover_clients = []
    seen_agency_ids = set()
    for profile in client_profiles:
        if profile.agency_id in seen_agency_ids:
            continue
        seen_agency_ids.add(profile.agency_id)
        profile.agency.handover_selected = agency_filter == str(profile.agency_id)
        handover_clients.append(profile.agency)
    stale_batches = list(
        FbsHandoverBatch.objects.select_related("profile__agency")
        .filter(
            status__in=(
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            ),
            created_at__lt=timezone.now() - timedelta(days=4),
        )
        .order_by("created_at", "id")[:20]
    )
    query_without_page = request.GET.copy()
    query_without_page.pop("page", None)
    query_without_page_encoded = query_without_page.urlencode()
    return render(
        request,
        "fbs/tsd_handover_list.html",
        _base_context(
            request,
            page_title="FBS · Отгрузки",
            back_url=(
                reverse("fbs:controller_home")
                if get_request_role(request) == "fbs_controller"
                else reverse("fbs:tsd_storekeeper")
            ),
            batches=batches,
            page_obj=page_obj,
            handover_page_size=page_size,
            handover_page_size_choices=page_size_choices,
            handover_query_without_page=query_without_page_encoded,
            handover_totals=handover_totals,
            handover_search=search,
            handover_status_filter=status_filter,
            handover_marketplace_filter=marketplace_filter,
            handover_agency_filter=agency_filter,
            handover_date_field_filter=date_field_filter,
            handover_date_from_filter=date_from_filter,
            handover_date_to_filter=date_to_filter,
            handover_status_choices=FbsHandoverBatch.STATUS_CHOICES,
            handover_marketplace_choices=FbsIntegrationProfile.MARKETPLACE_CHOICES,
            handover_date_field_choices=date_field_choices,
            handover_clients=handover_clients,
            handover_stale_batches=stale_batches,
            handover_print_workstations=configured_fbs_print_workstations(),
            ok_message=(
                f"В очередь печати добавлено ШК поставок: {request.GET.get('printed_supplies')}."
                if str(request.GET.get("printed_supplies") or "").isdigit()
                else ""
            ),
            profiles=FbsIntegrationProfile.objects.filter(
                is_active=True,
                marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
            ).select_related("agency"),
            error=error,
            filter_error=filter_error,
        ),
        status=400 if error else 200,
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_reports(request):
    filters = _handover_report_filter_context(request)
    report_type = filters["report_type"]
    report_rows = []
    report_totals = {"shipments": 0, "orders": 0, "units": 0, "boxes": 0}
    report_truncation_warning = ""
    marketplace_labels = dict(FbsIntegrationProfile.MARKETPLACE_CHOICES)

    batch_queryset = FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by"
    ).filter(
        created_at__date__gte=filters["date_from"],
        created_at__date__lte=filters["date_to"],
    )
    if filters["agency_filter"]:
        batch_queryset = batch_queryset.filter(
            profile__agency_id=int(filters["agency_filter"])
        )
    if filters["marketplace_filter"]:
        batch_queryset = batch_queryset.filter(
            profile__marketplace=filters["marketplace_filter"]
        )

    if filters["filter_error"]:
        pass
    elif report_type == "status":
        aggregates = {
            row["status"]: row
            for row in batch_queryset.values("status").annotate(
                shipment_count=Count("id", distinct=True),
                order_count=Count("order_assignments__order_id", distinct=True),
                box_count=Count("boxes__id", distinct=True),
            )
        }
        for status, label in FbsHandoverBatch.STATUS_CHOICES:
            values = aggregates.get(status, {})
            report_rows.append(
                SimpleNamespace(
                    status=status,
                    label=label,
                    shipment_count=int(values.get("shipment_count") or 0),
                    order_count=int(values.get("order_count") or 0),
                    box_count=int(values.get("box_count") or 0),
                )
            )
        report_totals["shipments"] = sum(row.shipment_count for row in report_rows)
        report_totals["orders"] = sum(row.order_count for row in report_rows)
        report_totals["boxes"] = sum(row.box_count for row in report_rows)
    else:
        item_queryset = FbsOrderItem.objects.select_related("sku").order_by("id")
        link_queryset = (
            FbsHandoverOrder.objects.filter(
                status=FbsHandoverOrder.STATUS_ACTIVE,
                box__batch__dispatched_at__isnull=False,
                box__batch__dispatched_at__date__gte=filters["date_from"],
                box__batch__dispatched_at__date__lte=filters["date_to"],
            )
            .select_related(
                "order",
                "box",
                "box__batch",
                "box__batch__profile",
                "box__batch__profile__agency",
            )
            .prefetch_related(Prefetch("order__items", queryset=item_queryset))
            .order_by("-box__batch__dispatched_at", "-id")
        )
        if filters["agency_filter"]:
            link_queryset = link_queryset.filter(
                box__batch__profile__agency_id=int(filters["agency_filter"])
            )
        if filters["marketplace_filter"]:
            link_queryset = link_queryset.filter(
                box__batch__profile__marketplace=filters["marketplace_filter"]
            )
        link_total_count = link_queryset.count()
        links = list(link_queryset[:HANDOVER_REPORT_ROW_LIMIT])
        if link_total_count > HANDOVER_REPORT_ROW_LIMIT:
            report_truncation_warning = (
                f"Показаны первые {HANDOVER_REPORT_ROW_LIMIT} строк из "
                f"{link_total_count}. Сузьте период или фильтры."
            )
        shipment_ids = set()
        box_ids = set()
        total_units = 0
        for link in links:
            link.report_items = list(link.order.items.all())
            link.report_unit_count = sum(int(item.quantity or 0) for item in link.report_items)
            shipment_ids.add(link.box.batch_id)
            box_ids.add(link.box_id)
            total_units += link.report_unit_count
        report_totals = {
            "shipments": len(shipment_ids),
            "orders": len(links),
            "units": total_units,
            "boxes": len(box_ids),
        }

        if report_type == "count":
            grouped = {}
            for link in links:
                profile = link.box.batch.profile
                key = (profile.agency_id, profile.marketplace)
                row = grouped.setdefault(
                    key,
                    SimpleNamespace(
                        agency=profile.agency,
                        marketplace=marketplace_labels.get(
                            profile.marketplace, profile.marketplace
                        ),
                        shipment_ids=set(),
                        order_count=0,
                        unit_count=0,
                    ),
                )
                row.shipment_ids.add(link.box.batch_id)
                row.order_count += 1
                row.unit_count += link.report_unit_count
            report_rows = sorted(
                grouped.values(),
                key=lambda row: (str(row.agency), row.marketplace),
            )
            for row in report_rows:
                row.shipment_count = len(row.shipment_ids)
        elif report_type == "orders":
            report_rows = links
        else:
            for link in links:
                for item in link.report_items:
                    report_rows.append(
                        SimpleNamespace(
                            link=link,
                            item=item,
                            dispatched_at=link.box.batch.dispatched_at,
                        )
                    )

    if (
        not filters["filter_error"]
        and str(request.GET.get("export") or "").strip() == "xlsx"
    ):
        if report_type == "status":
            headers = ("Статус", "Отгрузки", "Заказы", "Короба")
            rows = (
                (row.label, row.shipment_count, row.order_count, row.box_count)
                for row in report_rows
            )
        elif report_type == "count":
            headers = ("Клиент", "Площадка", "Отгрузки", "Заказы", "Товары")
            rows = (
                (
                    str(row.agency),
                    row.marketplace,
                    row.shipment_count,
                    row.order_count,
                    row.unit_count,
                )
                for row in report_rows
            )
        elif report_type == "orders":
            headers = (
                "Дата отправки",
                "Отгрузка",
                "Клиент",
                "Площадка",
                "Заказ",
                "Короб",
                "Товаров",
                "Статус Fullbox",
                "Статус marketplace",
            )
            rows = (
                (
                    timezone.localtime(row.box.batch.dispatched_at).strftime("%d.%m.%Y %H:%M"),
                    row.box.batch_id,
                    str(row.box.batch.profile.agency),
                    row.box.batch.profile.get_marketplace_display(),
                    row.order.external_order_id,
                    row.box.qr_code,
                    row.report_unit_count,
                    row.order.get_internal_status_display(),
                    row.order.marketplace_status,
                )
                for row in report_rows
            )
        else:
            headers = (
                "Дата отправки",
                "Отгрузка",
                "Клиент",
                "Площадка",
                "Заказ",
                "Артикул",
                "Штрихкод",
                "Товар",
                "Количество",
            )
            rows = (
                (
                    timezone.localtime(row.dispatched_at).strftime("%d.%m.%Y %H:%M"),
                    row.link.box.batch_id,
                    str(row.link.box.batch.profile.agency),
                    row.link.box.batch.profile.get_marketplace_display(),
                    row.link.order.external_order_id,
                    row.item.external_sku,
                    row.item.barcode,
                    row.item.product_name,
                    int(row.item.quantity or 0),
                )
                for row in report_rows
            )
        return _handover_report_xlsx(
            title=dict(HANDOVER_REPORT_CHOICES)[report_type],
            headers=headers,
            rows=rows,
            warning=report_truncation_warning,
        )

    page_size_choices = (25, 50, 100, 200, 500)
    try:
        page_size = int(request.GET.get("page_size") or 100)
    except (TypeError, ValueError):
        page_size = 100
    if page_size not in page_size_choices:
        page_size = 100
    page_obj = Paginator(report_rows, page_size).get_page(request.GET.get("page"))
    query_without_page = request.GET.copy()
    query_without_page.pop("page", None)
    export_query = request.GET.copy()
    export_query["export"] = "xlsx"
    clients = list(
        Agency.objects.filter(fbs_integration_profiles__handover_batches__isnull=False)
        .order_by("agn_name", "id")
        .distinct()
    )
    return render(
        request,
        "fbs/tsd_handover_reports.html",
        _base_context(
            request,
            page_title="FBS · Отчеты по отгрузкам",
            back_url=reverse("fbs:tsd_handover"),
            handover_report_choices=HANDOVER_REPORT_CHOICES,
            handover_report_type=report_type,
            handover_report_title=dict(HANDOVER_REPORT_CHOICES)[report_type],
            handover_report_period_label=(
                "Дата создания отгрузки"
                if report_type == "status"
                else "Дата отправки"
            ),
            handover_report_rows=page_obj.object_list,
            handover_report_totals=report_totals,
            handover_report_truncation_warning=report_truncation_warning,
            handover_report_clients=clients,
            handover_report_agency_filter=filters["agency_filter"],
            handover_report_marketplace_filter=filters["marketplace_filter"],
            handover_report_date_from=filters["date_from_filter"],
            handover_report_date_to=filters["date_to_filter"],
            handover_report_marketplace_choices=FbsIntegrationProfile.MARKETPLACE_CHOICES,
            handover_report_page_size=page_size,
            handover_report_page_size_choices=page_size_choices,
            handover_report_query_without_page=query_without_page.urlencode(),
            handover_report_export_query=export_query.urlencode(),
            page_obj=page_obj,
            filter_error=filters["filter_error"],
        ),
    )


def _handover_agent_scan_context(request) -> dict:
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if get_request_role(request) == "fbs_controller":
        workstation = get_controller_workstation(request)
        if workstation is not None and workstation.device_agent_id:
            agent_scan_poll_url = reverse("fbs:controller_scan_events")
            agent_scan_event_id = int(
                AgentEvent.objects.filter(
                    agent_id=workstation.device_agent.agent_id,
                    event_type=AgentEvent.EVENT_SCAN,
                )
                .order_by("-id")
                .values_list("id", flat=True)
                .first()
                or 0
            )
    return {
        "agent_scan_poll_url": agent_scan_poll_url,
        "agent_scan_event_id": agent_scan_event_id,
    }


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_handover_detail(request, batch_id: int):
    batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    request_role = get_request_role(request)
    error = ""
    action = ""
    kiz_rescan_all_requested = (
        str(request.GET.get("kiz_rescan_all") or "").strip() == "1"
    )
    verified_order = str(request.GET.get("verified_order") or "").strip()
    verified_box = str(request.GET.get("verified_box") or "").strip()
    handover_check_message = (
        f"Заказ {verified_order} подтвержден в коробе {verified_box}."
        if verified_order
        else ""
    )
    ok_message = {
        "box": "QR короба отправлен в очередь печати.",
        "supply": "QR поставки отправлен в очередь печати.",
        "order": "Этикетка заказа отправлена в очередь печати.",
        "ozon_qr": "Рабочий QR Ozon отправлен на повторную печать.",
    }.get(str(request.GET.get("printed") or "").strip(), "")
    printed_orders = str(request.GET.get("printed_orders") or "").strip()
    skipped_orders = str(request.GET.get("skipped_orders") or "").strip()
    if printed_orders.isdigit():
        ok_message = f"В очередь печати добавлено этикеток заказов: {printed_orders}."
        if skipped_orders.isdigit() and int(skipped_orders):
            ok_message += f" Пропущено без готовой этикетки: {skipped_orders}."
    if str(request.GET.get("ozon_return") or "").strip() == "queued":
        ok_message = (
            "Отмененный клиентом заказ убран из отгрузки и помещен в тару "
            "возврата. Остаток восстановится после физического возврата товара."
        )
    retried_order = str(request.GET.get("kiz_requeued") or "").strip()
    if retried_order:
        ok_message = f"Новый КИЗ заказа {retried_order} поставлен в очередь WB."
    kiz_route = str(request.GET.get("kiz_route") or "").strip()
    target_handover = str(request.GET.get("target_handover") or "").strip()
    problem_tote = str(request.GET.get("problem_tote") or "").strip()
    if kiz_route in {"rewave", "quarantine"} and target_handover:
        ok_message = (
            f"Товар зарегистрирован в проблемной таре {problem_tote}. "
            "Заказ не отменен в WB. Запущен безопасный перенос в отгрузку "
            f"№{target_handover}; после подтверждения WB "
            + (
                "система создаст новую волну."
                if kiz_route == "rewave"
                else "задача останется в карантине кладовщика."
            )
        )
    if str(request.GET.get("exclusion") or "").strip() == "queued":
        ok_message = (
            "Комментарий сохранен. Система сверяет состав с WB; после подтверждения "
            "товар появится в разделе «Возврат отбора»."
        )
    if str(request.GET.get("exclusion") or "").strip() == "return_queued":
        ok_message = (
            "Заказ исключен из отгрузки и передан в «Возврат отбора». "
            "FBS-остаток восстановится только после физического возврата товара."
        )
    if str(request.GET.get("exclusion") or "").strip() == "controller_removed":
        ok_message = (
            "Заказ снят с текущей проверки и помещён в назначенную тару "
            "отменённых заказов. Проверяйте следующий заказ; сверка WB "
            "продолжится отдельно, без изменения остатков и резервов."
        )
    if str(request.GET.get("verification_override") or "").strip() == "approved":
        ok_message = (
            "Начальник склада разрешил передачу без повторной проверки. "
            "Состав, короба, этикетки, КИЗ и сроки остаются обязательными."
        )
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        kiz_rescan_all_requested = (
            str(request.POST.get("kiz_rescan_all") or "").strip() == "1"
        )
        try:
            if (
                batch.status == FbsHandoverBatch.STATUS_ARCHIVED
                and action != "archive"
            ):
                raise ValueError("Отгрузка находится в архиве.")
            if action == "archive":
                archive_handover_batch(
                    batch_id=batch.id,
                    archived_by=request.user,
                    reason=request.POST.get("archive_reason", ""),
                )
            elif action == "approve_verification_override":
                approve_handover_verification_override(
                    batch_id=batch.id,
                    reason=request.POST.get("override_reason", ""),
                    approved_by=request.user,
                )
            elif action == "add_box":
                if batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
                    raise ValueError("Для WB создайте транспортные короба через API.")
                add_handover_box(
                    batch_id=batch.id,
                    qr_code=request.POST.get("box_qr", ""),
                    external_box_id=request.POST.get("external_box_id", ""),
                )
            elif action == "add_order":
                add_order_to_handover_box(
                    box_id=int(request.POST.get("box_id") or 0),
                    order_label_scan=request.POST.get("order_label_scan", ""),
                    added_by=request.user,
                )
            elif action == "verify_order":
                link = verify_handover_order_label(
                    batch_id=batch.id,
                    order_label_scan=request.POST.get("order_label_scan", ""),
                    verified_by=request.user,
                )
            elif action == "exclude_order":
                order_id = int(request.POST.get("order_id") or 0)
                pick_batch_id = (
                    FbsPickTask.objects.filter(
                        order_id=order_id,
                        batch__status=FbsPickBatch.STATUS_VERIFICATION,
                        batch__picking_completed_at__isnull=False,
                    )
                    .order_by("-batch_id")
                    .values_list("batch_id", flat=True)
                    .first()
                )
                if not pick_batch_id:
                    raise ValueError("Завершенная волна заказа не найдена.")
                create_order_pick_restock_request(
                    batch_id=pick_batch_id,
                    handover_batch_id=batch.id,
                    order_id=order_id,
                    reason_code=FbsPickRestockRequest.REASON_OTHER,
                    reason=request.POST.get("reason", ""),
                    confirm_seller_cancel=True,
                    created_by=request.user,
                )
            elif action == "retry_exclusion":
                retry_order_pick_restock_marketplace(
                    request_id=int(request.POST.get("request_id") or 0),
                    requested_by=request.user,
                )
            elif action == "release_exclusion":
                if request_role != "fbs_controller":
                    raise ValueError("Снять заказ может только FBS-контролёр этой смены.")
                release_order_pick_restock_to_queue(
                    request_id=int(request.POST.get("request_id") or 0),
                    comment=request.POST.get("reason", ""),
                    confirm_physical=True,
                    canceled_tote_scan=request.POST.get("canceled_tote_scan", ""),
                    handover_batch_id=batch.id,
                    performed_by=request.user,
                )
            elif action == "confirm_ozon_canceled_return":
                if request_role != "fbs_controller":
                    raise ValueError("Возврат отмененного заказа выполняет FBS-контролер.")
                confirm_ozon_canceled_order_return(
                    handover_batch_id=batch.id,
                    order_id=int(request.POST.get("order_id") or 0),
                    order_scan=request.POST.get("order_scan", ""),
                    canceled_tote_scan=request.POST.get("canceled_tote_scan", ""),
                    comment=request.POST.get("reason", ""),
                    performed_by=request.user,
                )
            elif action == "close_box":
                close_handover_box(box_id=int(request.POST.get("box_id") or 0))
            elif action == "scan_box":
                scan_handover_box(
                    batch_id=batch.id,
                    box_qr_scan=request.POST.get("box_qr_scan", ""),
                    scanned_by=request.user,
                )
            elif action == "create_wb_boxes":
                add_wb_handover_boxes(
                    batch_id=batch.id,
                    amount=int(request.POST.get("box_count") or 0),
                    requested_by=request.user,
                )
            elif action == "deliver_marketplace":
                request_wb_handover_delivery(
                    batch_id=batch.id,
                    requested_by=request.user,
                )
            elif action == "print_box_label":
                queue_fbs_handover_box_label_print(
                    batch_id=batch.id,
                    box_id=int(request.POST.get("box_id") or 0),
                    workstation_id=int(request.POST.get("workstation_id") or 0),
                    requested_by=request.user,
                    **_desktop_label_print_target(request, shipping=True),
                )
            elif action == "print_supply_label":
                queue_fbs_handover_supply_label_print(
                    batch_id=batch.id,
                    workstation_id=int(request.POST.get("workstation_id") or 0),
                    requested_by=request.user,
                    **_desktop_label_print_target(request, shipping=True),
                )
            elif action == "print_order_label":
                label_id = int(request.POST.get("label_id") or 0)
                if not batch.order_assignments.exclude(
                    status=FbsHandoverOrderAssignment.STATUS_CANCELED
                ).filter(order__marketplace_labels__id=label_id).exists():
                    raise ValueError("Этикетка не относится к этой отгрузке.")
                queue_fbs_order_label_print(
                    label_id=label_id,
                    requested_by=request.user,
                    force=True,
                    **_desktop_label_print_target(request),
                )
            elif action == "reprint_ozon_order_qr":
                if batch.profile.marketplace != FbsIntegrationProfile.MARKETPLACE_OZON:
                    raise ValueError("Рабочий QR Ozon доступен только для отгрузки Ozon.")
                label_id = int(request.POST.get("label_id") or 0)
                if not batch.order_assignments.exclude(
                    status=FbsHandoverOrderAssignment.STATUS_CANCELED
                ).filter(order__marketplace_labels__id=label_id).exists():
                    raise ValueError("Этикетка не относится к этой отгрузке.")
                queue_fbs_preloaded_ozon_order_label_print(
                    label_id=label_id,
                    requested_by=request.user,
                    force=True,
                    **_desktop_label_print_target(request),
                )
            elif action == "print_all_order_labels":
                raise ValueError(
                    "Массовая повторная печать этикеток Ozon отключена. Рабочий QR "
                    "печатается и сканируется один раз во время контроля заказа."
                )
            elif action == "dispatch":
                dispatch_handover_batch(
                    batch_id=batch.id,
                    dispatched_by=request.user,
                    dispatch_location_scan=request.POST.get("dispatch_location_scan", ""),
                )
            elif action == "refresh":
                refresh_handover_acceptance(batch_id=batch.id)
            elif action == "retry_kiz":
                command = retry_wb_marking_code(
                    batch_id=batch.id,
                    order_item_id=int(request.POST.get("order_item_id") or 0),
                    marking_scan=request.POST.get("marking_scan", ""),
                    requested_by=request.user,
                )
            elif action == "reroute_invalid_kiz":
                reroute = request_invalid_kiz_reroute(
                    handover_batch_id=batch.id,
                    order_id=int(request.POST.get("order_id") or 0),
                    route=request.POST.get("route", ""),
                    problem_tote_scan=request.POST.get("problem_tote_scan", ""),
                    requested_by=request.user,
                )
            else:
                raise ValueError("Неизвестная команда передачи.")
            detail_url = reverse("fbs:tsd_handover_detail", kwargs={"batch_id": batch.id})
            if action == "archive":
                archive_url = reverse("fbs:tsd_handover")
                return redirect(f"{archive_url}?status={FbsHandoverBatch.STATUS_ARCHIVED}")
            if action == "approve_verification_override":
                return redirect(f"{detail_url}?verification_override=approved")
            if action == "verify_order":
                query = urlencode(
                    {
                        "check": "1",
                        "verified_order": link.order.external_order_id,
                        "verified_box": link.box.qr_code,
                    }
                )
                return redirect(f"{detail_url}?{query}")
            if action in {"exclude_order", "retry_exclusion"}:
                return redirect(f"{detail_url}?exclusion=queued")
            if action == "release_exclusion":
                return redirect(f"{detail_url}?exclusion=controller_removed")
            if action == "confirm_ozon_canceled_return":
                return redirect(f"{detail_url}?ozon_return=queued")
            if action == "print_box_label":
                return redirect(f"{detail_url}?printed=box")
            if action == "print_supply_label":
                return redirect(f"{detail_url}?printed=supply")
            if action == "print_order_label":
                return redirect(f"{detail_url}?printed=order")
            if action == "reprint_ozon_order_qr":
                return redirect(f"{detail_url}?printed=ozon_qr")
            if action == "retry_kiz":
                query = {"kiz_requeued": command.order.external_order_id}
                if kiz_rescan_all_requested:
                    query["kiz_rescan_all"] = "1"
                return redirect(
                    f"{detail_url}?" + urlencode(query)
                )
            if action == "reroute_invalid_kiz":
                return redirect(
                    f"{detail_url}?"
                    + urlencode(
                        {
                            "kiz_route": reroute.route,
                            "target_handover": reroute.target_batch_id,
                            "problem_tote": request.POST.get(
                                "problem_tote_scan", ""
                            ),
                        }
                    )
                )
            return redirect(detail_url)
        except (FbsError, ValueError) as exc:
            error = str(exc)
        batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    detail_context = _handover_detail_summary(
        batch,
        compact_controller=request_role == "fbs_controller",
    )
    retryable_kiz_items = detail_context["handover_kiz_retry_items"]
    selected_kiz_retry = None
    if action == "retry_kiz":
        submitted_item_id = str(request.POST.get("order_item_id") or "").strip()
        selected_kiz_retry = next(
            (
                item
                for item in retryable_kiz_items
                if str(item.order_item_id) == submitted_item_id
            ),
            None,
        )
    elif kiz_rescan_all_requested and retryable_kiz_items:
        selected_kiz_retry = retryable_kiz_items[0]
    if kiz_rescan_all_requested and not retryable_kiz_items and not error:
        ok_message = "Все проблемные КИЗы отгрузки отсканированы и поставлены в очередь WB."
    can_manage_verification_override = bool(
        get_request_roles(request).intersection(
            {"head_manager", "director", "admin"}
        )
    )
    active_verification_override = None
    verification_override_blocker = ""
    if (
        can_manage_verification_override
        and batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and batch.status
        in (FbsHandoverBatch.STATUS_OPEN, FbsHandoverBatch.STATUS_READY)
    ):
        (
            active_verification_override,
            verification_override_blocker,
        ) = handover_verification_override_status(batch)
    controller_canceled_tote = None
    if request_role == "fbs_controller":
        controller_session = (
            FbsControllerSession.objects.select_related("canceled_tote")
            .filter(
                controller=request.user,
                status=FbsControllerSession.STATUS_ACTIVE,
            )
            .order_by("-id")
            .first()
        )
        controller_canceled_tote = (
            controller_session.canceled_tote if controller_session else None
        )
    return render(
        request,
        "fbs/tsd_handover_detail.html",
        _base_context(
            request,
            page_title=f"FBS · Передача #{batch.id}",
            back_url=reverse("fbs:tsd_handover"),
            handover=batch,
            error=error,
            ok_message=ok_message,
            handover_check_open=(
                detail_context["handover_reconciliation_available"]
                and (
                    str(request.GET.get("check") or "") == "1"
                    or action == "verify_order"
                )
            ),
            handover_check_message=handover_check_message,
            handover_check_error=error if action == "verify_order" else "",
            kiz_retry_open=bool(error and action == "retry_kiz")
            or bool(kiz_rescan_all_requested and selected_kiz_retry),
            kiz_retry_item_id=(
                str(selected_kiz_retry.order_item_id)
                if selected_kiz_retry is not None
                else str(request.POST.get("order_item_id") or "")
                if action == "retry_kiz"
                else ""
            ),
            kiz_retry_caption=(
                f"Заказ {selected_kiz_retry.order_number} · "
                f"{selected_kiz_retry.product_name}"
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_retry_label_number=(
                selected_kiz_retry.label_number
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_retry_product_barcode=(
                selected_kiz_retry.product_barcode
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_rescan_all_requested=kiz_rescan_all_requested,
            kiz_retry_remaining_count=len(retryable_kiz_items),
            kiz_retry_error=error if action == "retry_kiz" else "",
            can_manage_verification_override=can_manage_verification_override,
            active_verification_override=active_verification_override,
            verification_override_blocker=verification_override_blocker,
            handover_controller_canceled_tote=controller_canceled_tote,
            **_handover_agent_scan_context(request),
            **detail_context,
        ),
        status=400 if error else 200,
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_box_label(request, batch_id: int, box_id: int):
    box = get_object_or_404(
        FbsHandoverBox.objects.select_related("batch__profile__agency"),
        pk=box_id,
        batch_id=batch_id,
    )
    if not box.label_file:
        if box.batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
            raise Http404
        response = render(
            request,
            "fbs/tsd_handover_internal_box_label.html",
            {"box": box, "handover": box.batch},
        )
        response["Cache-Control"] = "private, no-store"
        response["X-Robots-Tag"] = "noindex, nofollow"
        return response
    response = FileResponse(
        box.label_file.open("rb"),
        as_attachment=False,
        filename=Path(box.label_file.name).name,
        content_type="image/png",
    )
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_supply_label(request, batch_id: int):
    batch = get_object_or_404(FbsHandoverBatch, pk=batch_id)
    if not batch.supply_label_file:
        raise Http404
    if request.GET.get("raw") == "1":
        response = FileResponse(
            batch.supply_label_file.open("rb"),
            as_attachment=False,
            filename=Path(batch.supply_label_file.name).name,
            content_type="image/png",
        )
        response["X-Content-Type-Options"] = "nosniff"
        response["Cache-Control"] = "private, no-store"
        return response
    response = render(
        request,
        "fbs/tsd_handover_supply_label_print.html",
        {
            "handover": batch,
            "label_image_url": f"{request.path}?raw=1",
        },
    )
    response["Cache-Control"] = "private, no-store"
    response["X-Robots-Tag"] = "noindex, nofollow"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_driver_manifest(request, batch_id: int):
    batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    detail_context = _handover_detail_summary(batch)
    return render(
        request,
        "fbs/tsd_handover_driver_manifest.html",
        {
            "handover": batch,
            "generated_at": timezone.now(),
            **detail_context,
        },
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_demand(request):
    created = _query_int(request, "created")
    skipped = _query_int(request, "skipped")
    ok_message = ""
    if created:
        ok_message = f"Создано планов: {created}."
        if skipped:
            ok_message += f" Клиентов с блокировками: {skipped}."
    elif request.GET.get("generated") == "1":
        ok_message = "Новая потребность для планирования отсутствует."
        if skipped:
            ok_message = f"Планы не созданы. Клиентов с блокировками: {skipped}."
    floor_movement_id = _query_int(request, "floor_movement")
    if floor_movement_id:
        ok_message = (
            f"Заявка #{floor_movement_id} на перемещение товара на первый ярус "
            "передана водителю ричтрака."
        )
    return render(
        request,
        "fbs/tsd_demand.html",
        _demand_context(request, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_create_floor_movement(request):
    try:
        source_box_id = int(request.POST.get("source_box_id") or 0)
        target_pallet_id = int(request.POST.get("target_pallet_id") or 0)
        needed_qty = int(request.POST.get("needed_qty") or 0)
    except (TypeError, ValueError):
        source_box_id = target_pallet_id = needed_qty = 0
    try:
        result = create_floor_replenishment_movement(
            source_box_id=source_box_id,
            target_pallet_id=target_pallet_id,
            barcode=str(request.POST.get("barcode") or ""),
            needed_qty=needed_qty,
            requested_by=request.user,
            comment=str(request.POST.get("comment") or ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_demand.html",
            _demand_context(request, error=str(exc)),
            status=409,
        )
    return redirect(
        f"{reverse('fbs:tsd_demand')}?floor_movement={result.movement.id}"
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_generate_plans(request):
    try:
        result = generate_replenishment_plans(
            requested_by=request.user,
            requested_by_role=get_request_role(request),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_demand.html",
            _demand_context(request, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_demand")
    return redirect(
        f"{url}?generated=1&created={len(result.plans)}"
        f"&skipped={len(result.skipped_agencies)}"
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_GET
def tsd_picking(request):
    ok_message = ""
    error = ""
    missing_qty = _query_int(request, "missing_qty")
    missing_orders = _query_int(request, "missing_orders")
    if missing_qty:
        ok_message = (
            f"Снято с текущей волны: {missing_qty} шт. "
            f"Затронуто заказов: {missing_orders}. "
            "Заказы возвращены в очередь повторного резервирования."
        )
        if request.GET.get("source_quarantined") == "1":
            ok_message += " Остаток в указанном коробе закрыт до пересчета."
    elif request.GET.get("equipment") == "1":
        ok_message = "ТСД привязан к рабочему месту и таре."
    elif request.GET.get("done") == "1":
        ok_message = "Заказ отобран."
    elif request.GET.get("released") == "1":
        ok_message = (
            "Заказ исключен из текущей волны и возвращен в ожидание товара."
        )
    elif request.GET.get("shortage") == "1":
        ok_message = (
            "Позиция отмечена как отсутствующая. "
            "Заказ поедет на проверку неполным."
        )
    elif request.GET.get("restock_done") == "1":
        ok_message = "Ошибочный отбор полностью возвращен в исходные FBS-короба."
    elif request.GET.get("prepared") == "1":
        tasks = _query_int(request, "tasks")
        waiting = _query_int(request, "waiting")
        failed = _query_int(request, "failed")
        ok_message = f"Создано заданий: {tasks}."
        if waiting or failed:
            ok_message += f" Ожидают товар: {waiting}; ошибки SKU: {failed}."
    if request.GET.get("wave_stopped") == "1":
        error = (
            "Задание этой волны остановлено. Список обновлен; "
            "выберите доступную волну."
        )
    return render(
        request,
        "fbs/tsd_picking_list.html",
        _picking_context(request, error=error, ok_message=ok_message),
    )


def _pick_restock_context(request, restock, *, error="", ok_message=""):
    state = pick_restock_state(restock.id)
    line = state["line"]
    recent_scans = restock.scans.select_related("line").order_by("-created_at", "-id")[:8]
    stage = state["stage"]
    stage_copy = {
        FbsPickRestockScan.STAGE_WORKSTATION: (
            "Забрать из тары" if state["source_tote"] else "Забрать со стола",
            (
                "Подойдите к указанной служебной таре и подтвердите ее QR."
                if state["source_tote"]
                else "Подойдите к указанному рабочему столу и подтвердите его QR."
            ),
            "QR служебной тары" if state["source_tote"] else "QR рабочего стола",
        ),
        FbsPickRestockScan.STAGE_PICKUP_ITEM: (
            "Забрать товар из тары" if state["source_tote"] else "Снять товар со стола",
            (
                "Сканируйте каждую единицу, которую забираете из тары. "
                "Перейти к размещению можно только после полного сканирования."
                if state["source_tote"]
                else "Сканируйте каждую единицу, которую забираете со стола. "
                "Перейти к размещению можно только после полного сканирования."
            ),
            "Штрихкод товара из тары" if state["source_tote"] else "Штрихкод товара со стола",
        ),
        FbsPickRestockScan.STAGE_CELL: (
            "Карантинная ячейка" if state["is_quarantine"] else "Исходная ячейка",
            (
                "Отсканируйте QR карантинного места, указанного контролером."
                if state["is_quarantine"]
                else "Отсканируйте QR места, из которого товар был отобран."
            ),
            "QR ячейки",
        ),
        FbsPickRestockScan.STAGE_BOX: (
            "Карантинный короб" if state["is_quarantine"] else "Исходный короб",
            (
                "Отсканируйте QR карантинного FBS-короба."
                if state["is_quarantine"]
                else "Отсканируйте QR исходного короба. Все товары этого короба вернутся одним проходом."
            ),
            "QR короба",
        ),
        FbsPickRestockScan.STAGE_ITEM: (
            "Товар",
            "Сканируйте штрихкод каждой единицы. КИЗ для этой позиции не требуется. После последней единицы система сама даст следующий короб.",
            "Штрихкод товара",
        ),
    }
    title, instruction, input_label = stage_copy.get(
        stage,
        ("Завершено", "Все единицы возвращены.", "Скан"),
    )
    destination_cell = state.get("destination_cell")
    destination_box = state.get("destination_box")
    location_label = _pick_location_label(destination_cell) if destination_cell else ""
    if stage == FbsPickRestockScan.STAGE_BOX and destination_box is not None:
        box_number_suffix = _pick_box_number_suffix(destination_box.box_code)
        if box_number_suffix:
            title = f"{title} · {box_number_suffix}"
    return _base_context(
        request,
        page_title=f"Возврат отбора #{restock.id}",
        back_url=reverse("fbs:tsd_picking"),
        restock=state["request"],
        restock_state=state,
        source_workstation=state["source_workstation"],
        source_tote=state["source_tote"],
        destination_cell=destination_cell,
        destination_box=destination_box,
        is_quarantine=state["is_quarantine"],
        line=line,
        stage=stage,
        stage_title=title,
        stage_instruction=instruction,
        scan_input_label=input_label,
        source_location_label=location_label,
        request_token=str(uuid.uuid4()),
        recent_scans=recent_scans,
        error=error,
        ok_message=ok_message,
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_restock(request, request_id: int):
    restock = get_object_or_404(
        FbsPickRestockRequest.objects.select_related("batch__workstation", "assigned_to"),
        pk=request_id,
    )
    if restock.assigned_to_id not in {None, request.user.id}:
        return HttpResponseForbidden("Задание возврата уже выполняет другой подборщик.")
    if request.method == "POST":
        action = str(request.POST.get("action") or "scan").strip()
        try:
            if action == "claim":
                claim_pick_restock_request(
                    request_id=restock.id,
                    assigned_to=request.user,
                )
                return redirect("fbs:tsd_pick_restock", request_id=restock.id)
            if action != "scan":
                raise FbsError("Неизвестная команда возврата отбора.")
            result = scan_pick_restock(
                request_id=restock.id,
                stage=request.POST.get("stage", ""),
                scan_value=request.POST.get("scan_value", ""),
                request_token=request.POST.get("request_token", ""),
                performed_by=request.user,
            )
        except FbsError as exc:
            restock.refresh_from_db()
            return render(
                request,
                "fbs/tsd_pick_restock_detail.html",
                _pick_restock_context(request, restock, error=str(exc)),
                status=409,
            )
        if result.request.status == FbsPickRestockRequest.STATUS_COMPLETED:
            return redirect(f"{reverse('fbs:tsd_picking')}?restock_done=1")
        restock.refresh_from_db()
        return render(
            request,
            "fbs/tsd_pick_restock_detail.html",
            _pick_restock_context(
                request,
                restock,
                ok_message=result.event.message,
            ),
        )
    return render(
        request,
        "fbs/tsd_pick_restock_detail.html",
        _pick_restock_context(request, restock),
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_bind_pick_equipment(request):
    try:
        batch = claim_next_pick_batch(
            assigned_to=request.user,
            workstation_scan="",
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    _store_picker_cart(request, batch.cart)
    if batch.status == FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:tsd_pick_handover", batch_id=batch.id)
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_clear_pick_equipment(request):
    request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
    return redirect("fbs:tsd_picking")


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_next_pick_batch(request):
    cart = _picker_cart(request)
    if cart is None:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(
                request,
                error="Сначала отсканируйте свободную тару.",
            ),
            status=409,
        )
    try:
        batch = claim_next_pick_batch(
            assigned_to=request.user,
            workstation_scan="",
            cart_scan=cart.barcode,
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    if batch.status == FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:tsd_pick_handover", batch_id=batch.id)
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_labels(request):
    ok_message = ""
    if request.GET.get("requested") == "1":
        ok_message = "Заказ отобран. Этикетка маркетплейса запрошена."
    elif request.GET.get("done") == "1":
        ok_message = "Этикетка подтверждена. Заказ готов к передаче."
    return render(
        request,
        "fbs/tsd_label_list.html",
        _labels_context(request, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_label_file(request, label_id: int):
    label = get_object_or_404(
        _accessible_label_queryset(request),
        pk=label_id,
        status__in=(FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED),
    )
    if not label.file:
        raise Http404
    content_type = {
        FbsOrderLabel.FORMAT_PDF: "application/pdf",
        FbsOrderLabel.FORMAT_PNG: "image/png",
        FbsOrderLabel.FORMAT_ZPL: "application/octet-stream",
    }.get(label.label_format, "application/octet-stream")
    response = FileResponse(
        label.file.open("rb"),
        as_attachment=False,
        filename=Path(label.file.name).name,
        content_type=content_type,
    )
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_label_detail(request, label_id: int):
    label = get_object_or_404(_accessible_label_queryset(request), pk=label_id)
    if get_request_role(request) == "fbs_controller":
        controller_batch = (
            FbsPickBatch.objects.filter(
                tasks__order_id=label.order_id,
                verification_assigned_to=request.user,
                status=FbsPickBatch.STATUS_VERIFICATION,
            )
            .order_by("verification_started_at", "id")
            .first()
        )
        if controller_batch is not None:
            return redirect(
                "fbs:tsd_pick_verification",
                batch_id=controller_batch.id,
            )
        if request.method == "POST":
            return render(
                request,
                "fbs/tsd_label_detail.html",
                _label_context(
                    request,
                    label,
                    error=(
                        "Контролер подтверждает этикетку только внутри активной "
                        "проверки волны. Вернитесь на рабочий стол и отсканируйте тару."
                    ),
                ),
                status=409,
            )
    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(
                request,
                label,
                ok_message=(
                    "Этикетка отправлена в очередь печати."
                    if request.GET.get("printed") == "1"
                    else ""
                ),
            ),
        )
    action = str(request.POST.get("action") or "confirm").strip()
    if action == "print":
        try:
            queue_fbs_order_label_print(
                label_id=label.id,
                requested_by=request.user,
                force=True,
                **_desktop_label_print_target(request),
            )
        except FbsError as exc:
            label.refresh_from_db()
            return render(
                request,
                "fbs/tsd_label_detail.html",
                _label_context(request, label, error=str(exc)),
                status=400,
            )
        return redirect(f"{reverse('fbs:tsd_label_detail', kwargs={'label_id': label.id})}?printed=1")
    if action != "confirm":
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(request, label, error="Неизвестная команда этикетки."),
            status=400,
        )
    label_scan = str(request.POST.get("label_scan") or "").strip()
    try:
        confirm_order_label_scan(
            label_id=label.id,
            label_scan=label_scan,
            performed_by=request.user,
        )
    except FbsError as exc:
        record_order_label_scan_events(
            order_id=label.order_id,
            scan_value=label_scan,
            expected_value=label.barcode,
            result=FbsPickScanEvent.RESULT_ERROR,
            message=str(exc),
            created_by=request.user,
        )
        label.refresh_from_db()
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(request, label, error=str(exc)),
            status=400,
        )
    next_batch = (
        FbsPickBatch.objects.filter(
            tasks__order_id=label.order_id,
            verification_assigned_to=request.user,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .order_by("created_at", "id")
        .first()
    )
    if next_batch is not None:
        return redirect("fbs:tsd_pick_verification", batch_id=next_batch.id)
    completed_batch = (
        FbsPickBatch.objects.filter(
            tasks__order_id=label.order_id,
            verification_assigned_to=request.user,
            status=FbsPickBatch.STATUS_DONE,
        )
        .order_by("-completed_at", "-id")
        .first()
    )
    if completed_batch is not None:
        if get_request_role(request) == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect(
            f"{reverse('fbs:operator_wave_detail', kwargs={'batch_id': completed_batch.id})}?done=1"
        )
    return redirect(f"{reverse('fbs:tsd_labels')}?done=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_prepare_picking(request):
    try:
        max_orders_per_batch = int(request.POST.get("wave_size") or 50)
    except (TypeError, ValueError):
        max_orders_per_batch = 0
    try:
        result = prepare_pick_queue(
            max_orders_per_batch=max_orders_per_batch,
            max_units_per_batch=100,
            performed_by=request.user,
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    rejection_message = format_queue_rejections(
        getattr(result, "rejected_orders", ())
    )
    if rejection_message and not result.tasks_created:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(
                request,
                error=f"Волна не создана. {rejection_message}",
            ),
            status=409,
        )
    if rejection_message:
        messages.warning(request, rejection_message)
    url = reverse("fbs:tsd_picking")
    return redirect(
        f"{url}?prepared=1&tasks={result.tasks_created}"
        f"&waiting={result.awaiting_stock_orders}"
        f"&failed={result.validation_failed_orders}"
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_pick_task(request, task_id: int):
    get_object_or_404(FbsPickTask, pk=task_id)
    try:
        task = claim_pick_task(
            task_id=task_id,
            assigned_to=request.user,
            workstation_scan=request.POST.get("workstation_scan", ""),
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    allocation = (
        _pick_allocation_queryset()
        .filter(
            pick_task=task,
            status__in=(
                FbsOrderStockAllocation.STATUS_RESERVED,
                FbsOrderStockAllocation.STATUS_PICKING,
            ),
        )
        .order_by("balance__box__pallet__cell__cell_code", "id")
        .first()
    )
    if allocation is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation.id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_pick_batch(request, batch_id: int):
    get_object_or_404(FbsPickBatch, pk=batch_id)
    # Очередь строгая: волны выдаются по возрасту, сборщик их не выбирает.
    # В интерфейсе кнопки взятия конкретной волны нет, но адрес открыт
    # напрямую, поэтому порядок проверяем здесь. Свою уже начатую волну
    # продолжить можно всегда.
    own_active = FbsPickBatch.objects.filter(
        pk=batch_id,
        assigned_to=request.user,
        picking_completed_at__isnull=True,
    ).exists()
    if not own_active:
        oldest_free_id = (
            FbsPickBatch.objects.filter(
                status=FbsPickBatch.STATUS_QUEUED,
                assigned_to__isnull=True,
            )
            .order_by("created_at", "id")
            .values_list("id", flat=True)
            .first()
        )
        if oldest_free_id != batch_id:
            return render(
                request,
                "fbs/tsd_picking_list.html",
                _picking_context(
                    request,
                    error=(
                        "Волны выдаются по очереди, с самой старой. "
                        "Нажмите «Получить следующую волну»."
                    ),
                ),
                status=409,
            )
    try:
        batch = claim_pick_batch(
            batch_id=batch_id,
            assigned_to=request.user,
            workstation_scan=request.POST.get("workstation_scan", ""),
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_allocation(request, allocation_id: int):
    allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation_id)
    task = allocation.pick_task
    if task is None:
        return HttpResponseForbidden("Позиция не включена в задание")
    if task.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Задание назначено другому сборщику")
    if allocation.status == FbsOrderStockAllocation.STATUS_PICKED:
        return redirect("fbs:tsd_picking")
    if (
        task.status != FbsPickTask.STATUS_IN_PROGRESS
        or task.batch.status != FbsPickBatch.STATUS_IN_PROGRESS
        or allocation.status != FbsOrderStockAllocation.STATUS_PICKING
    ):
        request.session.pop(_pick_cell_session_key(allocation.id), None)
        request.session.pop(_pick_box_session_key(allocation.id), None)
        return redirect(f"{reverse('fbs:tsd_picking')}?wave_stopped=1")
    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(request, allocation),
        )

    action = str(request.POST.get("action") or "").strip()
    cell_session_key = _pick_cell_session_key(allocation.id)
    session_key = _pick_box_session_key(allocation.id)
    if action == "not_found":
        if request.POST.get("confirm_not_found") != "1":
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request,
                    allocation,
                    error="Подтвердите, что указанной позиции нет на месте.",
                ),
                status=400,
            )
        try:
            missing_qty = int(request.POST.get("missing_qty") or 0)
        except (TypeError, ValueError):
            missing_qty = 0
        try:
            result = report_pick_missing_quantity(
                allocation_id=allocation.id,
                missing_qty=missing_qty,
                reason=request.POST.get("reason", ""),
                reported_by=request.user,
            )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session.pop(cell_session_key, None)
        request.session.pop(session_key, None)
        query = urlencode(
            {
                "missing_qty": result.released_qty,
                "missing_orders": result.affected_order_count,
                "source_quarantined": "1" if result.source_quarantined else "0",
            }
        )
        return redirect(f"{reverse('fbs:tsd_picking')}?{query}")
    if action == "scan_cell":
        cell_scan = str(request.POST.get("cell_scan") or "").strip()
        try:
            validate_pick_cell_scan(
                allocation_id=allocation.id,
                cell_scan=cell_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(cell_session_key, None)
            request.session.pop(session_key, None)
            original_error = str(exc)
            scan_error = (
                "Неверно отсканирована ячейка. "
                "Отсканируйте указанную ячейку еще раз."
                if "Скан ячейки FBS не совпадает с маршрутом" in original_error
                else original_error
            )
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request,
                    allocation,
                    error=scan_error,
                    scan_error=scan_error,
                ),
                status=400,
            )
        request.session[cell_session_key] = cell_scan
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(request, allocation, ok_message="Ячейка подтверждена."),
        )
    if action == "scan_box":
        if not str(request.session.get(cell_session_key) or "").strip():
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error="Сначала подтвердите ячейку."),
                status=400,
            )
        box_scan = str(request.POST.get("box_scan") or "").strip()
        try:
            validate_pick_box_scan(
                allocation_id=allocation.id,
                box_scan=box_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(session_key, None)
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session[session_key] = box_scan
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(request, allocation, ok_message="Короб подтвержден."),
        )

    if action == "complete":
        cell_scan = str(request.session.get(cell_session_key) or "").strip()
        box_scan = str(request.session.get(session_key) or "").strip()
        item_scan = str(request.POST.get("item_scan") or "").strip()
        if not cell_scan or not box_scan:
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request, allocation, error="Сначала подтвердите ячейку и короб."
                ),
                status=400,
            )
        try:
            complete_pick_allocation(
                allocation_id=allocation.id,
                cell_scan=cell_scan,
                box_scan=box_scan,
                item_scan=item_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        allocation.refresh_from_db()
        if allocation.status == FbsOrderStockAllocation.STATUS_PICKED:
            next_allocation = _open_pick_allocations(task.batch).first()
            request.session.pop(cell_session_key, None)
            request.session.pop(session_key, None)
            if next_allocation is not None:
                current_cell_id = allocation.balance.box.pallet.cell_id
                next_cell_id = next_allocation.balance.box.pallet.cell_id
                if current_cell_id == next_cell_id:
                    request.session[
                        _pick_cell_session_key(next_allocation.id)
                    ] = cell_scan
                if (
                    allocation.balance.box_id == next_allocation.balance.box_id
                    and current_cell_id == next_cell_id
                ):
                    request.session[
                        _pick_box_session_key(next_allocation.id)
                    ] = box_scan
                return redirect(
                    "fbs:tsd_pick_allocation",
                    allocation_id=next_allocation.id,
                )
        else:
            return redirect("fbs:tsd_pick_allocation", allocation_id=allocation.id)
        task.batch.refresh_from_db()
        if task.batch.status == FbsPickBatch.STATUS_VERIFICATION:
            return redirect("fbs:tsd_pick_handover", batch_id=task.batch_id)
        return redirect("fbs:tsd_picking")

    return render(
        request,
        "fbs/tsd_pick_allocation.html",
        _pick_allocation_context(request, allocation, error="Неизвестная команда."),
        status=400,
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_handover(request, batch_id: int):
    batch = get_object_or_404(
        FbsPickBatch.objects.select_related("agency", "assigned_to", "workstation", "cart"),
        pk=batch_id,
    )
    if batch.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Волна назначена другому сборщику")
    handover_redirect_url = f"{reverse('fbs:tsd_picking')}?handed_over=1"
    if batch.picking_completed_at is not None:
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        if request.GET.get("handover_status") == "1":
            response = JsonResponse(
                {
                    "handed_over": True,
                    "redirect_url": handover_redirect_url,
                }
            )
            response["Cache-Control"] = "no-store"
            return response
        return redirect(handover_redirect_url)
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return HttpResponseForbidden("Волна еще не собрана полностью")
    if request.GET.get("handover_status") == "1":
        response = JsonResponse({"handed_over": False})
        response["Cache-Control"] = "no-store"
        return response
    error = ""
    if request.method == "POST":
        try:
            handover_pick_batch_for_verification(
                batch_id=batch.id,
                workstation_scan=request.POST.get("workstation_scan", ""),
                performed_by=request.user,
            )
            request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
            return redirect(handover_redirect_url)
        except FbsError as exc:
            error = str(exc)
    workstation_recommendation = None
    try:
        workstation_recommendation = assign_pick_handover_workstation(
            batch_id=batch.id
        )
        batch.refresh_from_db(fields=["workstation", "updated_at"])
    except FbsError as exc:
        if not error:
            error = str(exc)
    return render(
        request,
        "fbs/tsd_pick_handover.html",
        _base_context(
            request,
            page_title=f"FBS · Сдать волну #{batch.id}",
            back_url=reverse("fbs:tsd_picking"),
            batch=batch,
            workstation_recommendation=workstation_recommendation,
            error=error,
        ),
        status=400 if error else 200,
    )


def _tsd_pick_verification_response(request, batch_id: int):
    batch = get_object_or_404(
        FbsPickBatch.objects.select_related("agency", "assigned_to", "workstation", "cart"),
        pk=batch_id,
    )
    request_role = get_request_role(request)
    if batch.verification_assigned_to_id is None:
        if request_role == "fbs_controller":
            workstation = get_controller_workstation(request)
            if workstation is None:
                return redirect("fbs:controller_home")
            if batch.workstation_id != workstation.id:
                return HttpResponseForbidden(
                    "Волна доставлена на другое рабочее место контролера"
                )
            return redirect("fbs:controller_home")
        if request.method == "POST" and request.POST.get("action") == "claim":
            try:
                claim_pick_batch_verification(batch_id=batch.id, assigned_to=request.user)
            except FbsError as exc:
                return render(
                    request,
                    "fbs/tsd_pick_verification_claim.html",
                    _base_context(
                        request,
                        batch=batch,
                        back_url=(
                            reverse("fbs:controller_home")
                            if request_role == "fbs_controller"
                            else reverse("fbs:operator_waves")
                        ),
                        error=str(exc),
                    ),
                    status=409,
                )
            return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
        return render(
            request,
            "fbs/tsd_pick_verification_claim.html",
            _base_context(
                request,
                batch=batch,
                back_url=(
                    reverse("fbs:controller_home")
                    if request_role == "fbs_controller"
                    else reverse("fbs:operator_waves")
                ),
            ),
        )
    if batch.verification_assigned_to_id != request.user.id:
        return HttpResponseForbidden("Проверку волны выполняет другой оператор")
    marketplace_rejection_return = None
    marketplace_rejection_allocation = None
    if request_role == "fbs_controller":
        marketplace_rejection_return = (
            FbsPickRestockRequest.objects.select_related(
                "order",
                "handover_assignment",
            )
            .filter(
                batch=batch,
                reason_code=FbsPickRestockRequest.REASON_MARKETPLACE,
                status__in=(
                    FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
                    FbsPickRestockRequest.STATUS_QUEUED,
                    FbsPickRestockRequest.STATUS_FAILED,
                ),
                source_tote__isnull=True,
            )
            .order_by("created_at", "id")
            .first()
        )
        if marketplace_rejection_return is not None:
            marketplace_rejection_allocation = (
                _pick_allocation_queryset()
                .filter(
                    pick_task__batch=batch,
                    pick_task__order_id=marketplace_rejection_return.order_id,
                    status=FbsOrderStockAllocation.STATUS_PICKED,
                    qty_picked__gt=0,
                )
                .order_by("id")
                .first()
            )
    if (
        marketplace_rejection_return is not None
        and marketplace_rejection_allocation is not None
    ):
        error = ""
        if (
            request.method == "POST"
            and str(request.POST.get("action") or "").strip()
            == "confirm_marketplace_return_tote"
        ):
            try:
                confirm_marketplace_rejected_order_return(
                    request_id=marketplace_rejection_return.id,
                    canceled_tote_scan=str(
                        request.POST.get("service_tote_scan") or ""
                    ).strip(),
                    performed_by=request.user,
                )
            except FbsError as exc:
                error = str(exc)
                marketplace_rejection_return.refresh_from_db()
            else:
                return redirect(
                    f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}"
                    "?problem=1"
                )
        prompt = _verification_service_tote_prompt(
            batch=batch,
            allocation=marketplace_rejection_allocation,
            controller=request.user,
            problem_kind="marketplace_rejected",
            reason=marketplace_rejection_return.reason,
            restock_request=marketplace_rejection_return,
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                marketplace_rejection_allocation,
                error=error,
                item_selected=True,
                service_tote_prompt=prompt,
            ),
            status=409 if error else 200,
        )
    if batch.status == FbsPickBatch.STATUS_DONE:
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return HttpResponseForbidden("Волна еще не готова к проверке")
    label = _next_verification_label(batch)
    allocation = None if label is not None else _next_verification_allocation(batch)
    if allocation is None and label is None:
        allocation = _next_incomplete_verification_allocation(batch)
    if allocation is None:
        if label is not None:
            if request_role == "fbs_controller":
                error = ""
                ok_message = (
                    "Этикетка отправлена в очередь печати."
                    if request.GET.get("printed") == "1"
                    else ""
                )
                if request.method == "GET" and not _fullbox_desktop_supports_direct_print(request):
                    # Экран контролера дошел до шага этикетки — ставим печать
                    # сразу, не дожидаясь фонового обработчика очереди.
                    # Дедупликация по card_id делает вызов идемпотентным:
                    # если задание уже создано, вернется существующее.
                    try:
                        if is_preloaded_ozon_order_label(label):
                            queue_fbs_preloaded_ozon_order_label_print(label_id=label.id)
                        else:
                            queue_fbs_order_label_print(label_id=label.id)
                    except FbsError:
                        pass
                if request.method == "POST":
                    action = str(request.POST.get("action") or "confirm_label").strip()
                    if action == "print_label":
                        try:
                            print_mode = str(
                                request.POST.get("print_mode") or "automatic"
                            ).strip()
                            print_scope = ""
                            if print_mode != "manual":
                                verification_started = (
                                    batch.verification_started_at.isoformat()
                                    if batch.verification_started_at is not None
                                    else "not-started"
                                )
                                print_scope = (
                                    f"verification:{batch.id}:{verification_started}"
                                )
                            if is_preloaded_ozon_order_label(label):
                                queue_fbs_preloaded_ozon_order_label_print(
                                    label_id=label.id,
                                    requested_by=request.user,
                                    force=True,
                                    **_desktop_label_print_target(request),
                                )
                            else:
                                queue_fbs_order_label_print(
                                    label_id=label.id,
                                    requested_by=request.user,
                                    force=True,
                                    print_scope=print_scope,
                                    **_desktop_label_print_target(request),
                                )
                        except FbsError as exc:
                            error = str(exc)
                        else:
                            return redirect(
                                f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}?printed=1"
                            )
                    elif action == "confirm_label":
                        label_scan = str(request.POST.get("label_scan") or "").strip()
                        try:
                            if batch.workstation_id is None:
                                raise ValueError("У волны не указано рабочее место контролера.")
                            pick_tote_context = controller_pick_tote_for_batch(
                                batch_id=batch.id
                            )
                            if pick_tote_context is None:
                                raise ValueError(
                                    "Сначала примите тару подбора и выберите тару проверки."
                                )
                            confirm_order_label_to_check_tote(
                                label_id=label.id,
                                label_scan=label_scan,
                                pick_batch_id=batch.id,
                                performed_by=request.user,
                            )
                        except (FbsError, ValueError) as exc:
                            record_order_label_scan_events(
                                order_id=label.order_id,
                                scan_value=label_scan,
                                expected_value=label.barcode,
                                result=FbsPickScanEvent.RESULT_ERROR,
                                message=str(exc),
                                created_by=request.user,
                            )
                            error = str(exc)
                        else:
                            # Вся волна отсканирована — тара освобождается сама,
                            # нажимать ничего не нужно. Если освободить нельзя
                            # (выключен флаг, чужой контролер), остается прежний
                            # путь с подтверждением пустоты на рабочем столе.
                            released_tote = auto_release_pick_tote_if_complete(
                                pick_batch_id=batch.id
                            )
                            if released_tote is not None:
                                released_query = urlencode(
                                    {
                                        "tote_released": "1",
                                        "tote": released_tote.tote.barcode,
                                    }
                                )
                                return redirect(
                                    f"{reverse('fbs:controller_home')}?{released_query}"
                                )
                            if mark_pick_tote_awaiting_empty(
                                pick_batch_id=batch.id
                            ):
                                return redirect(
                                    f"{reverse('fbs:controller_home')}?empty={batch.id}"
                                )
                            return redirect(
                                "fbs:tsd_pick_verification",
                                batch_id=batch.id,
                            )
                    else:
                        error = "Неизвестная команда этикетки."
                    label.refresh_from_db()
                return render(
                    request,
                    "fbs/tsd_pick_verification.html",
                    _pick_verification_label_context(
                        request,
                        batch,
                        label,
                        error=error,
                        ok_message=ok_message,
                    ),
                    status=400 if error else 200,
                )
            return redirect("fbs:tsd_label_detail", label_id=label.id)
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if request.method == "GET":
        ok_message = ""
        if request.GET.get("problem") == "1":
            ok_message = (
                "Проблемный заказ исключен из проверки и передан комплектовщику "
                "для размещения в карантин."
            )
        elif request.GET.get("problem_returned") == "1":
            ok_message = (
                "Товар возвращен из проблемной тары в тару проверки. "
                "Отсканируйте товар для продолжения проверки."
            )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                allocation,
                ok_message=ok_message,
            ),
        )
    action = str(request.POST.get("action") or "select_item").strip()
    if action not in {"select_item", "verify"}:
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                allocation,
                error="Неизвестная команда проверки.",
            ),
            status=400,
        )
    item_scan = str(request.POST.get("item_scan") or "").strip()
    marking_scan = str(request.POST.get("marking_scan") or "").strip()
    legacy_marking_absent = (
        str(request.POST.get("legacy_marking_absent") or "").strip() == "1"
    )
    expiry_date_value = str(request.POST.get("expiry_date") or "").strip()
    selected_allocation = allocation
    inferred_marking_scan = ""
    if action == "verify":
        try:
            selected_allocation_id = int(request.POST.get("allocation_id") or 0)
        except (TypeError, ValueError):
            selected_allocation_id = 0
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(),
            pk=selected_allocation_id,
            pick_task__batch=batch,
        )
        if request_role == "fbs_controller" and _marketplace_order_is_canceled(
            selected_allocation.pick_task.order
        ):
            canceled_reason = "Маркетплейс отменил этот заказ."
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    error=canceled_reason,
                    submitted_item_scan=item_scan,
                    item_selected=True,
                    service_tote_prompt=_verification_service_tote_prompt(
                        batch=batch,
                        allocation=selected_allocation,
                        controller=request.user,
                        problem_kind="canceled",
                        reason=canceled_reason,
                    ),
                ),
                status=409,
            )
    else:
        try:
            selected_allocation = find_pick_verification_allocation(
                batch_id=batch.id,
                item_scan=item_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            error_text = str(exc)
            unknown_scan_info = None
            if (
                request_role == "fbs_controller"
                and item_scan
                and "отсутствует в текущей таре" in error_text.casefold()
            ):
                try:
                    problem_item = record_extra_problem_tote_item(
                        pick_batch_id=batch.id,
                        scanned_value=item_scan,
                        performed_by=request.user,
                    )
                except FbsError as problem_exc:
                    error_text = str(problem_exc)
                else:
                    unknown_scan_info = SimpleNamespace(
                        scanned_value=item_scan,
                        tote_name=problem_item.problem_tote.name,
                        tote_barcode=problem_item.problem_tote.barcode,
                    )
                    error_text = (
                        f"Товара со штрихкодом {item_scan} в этой таре нет. "
                        "Скан не засчитан; положите лишний товар в проблемную тару "
                        f"{problem_item.problem_tote.name} "
                        f"({problem_item.problem_tote.barcode})."
                    )
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    allocation,
                    error=error_text,
                    unknown_scan_info=unknown_scan_info,
                ),
                status=400,
            )

        if request_role == "fbs_controller" and _marketplace_order_is_canceled(
            selected_allocation.pick_task.order
        ):
            canceled_reason = "Маркетплейс отменил этот заказ."
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    error=canceled_reason,
                    submitted_item_scan=item_scan,
                    item_selected=True,
                    service_tote_prompt=_verification_service_tote_prompt(
                        batch=batch,
                        allocation=selected_allocation,
                        controller=request.user,
                        problem_kind="canceled",
                        reason=canceled_reason,
                    ),
                ),
                status=409,
            )

        matched_item_scan, inferred_marking_scan = resolve_verification_item_scan(
            selected_allocation,
            item_scan,
        )
        if inferred_marking_scan:
            item_scan = matched_item_scan
            marking_scan = inferred_marking_scan

        from .services.traceability import (
            controller_marking_scan_required,
            metadata_requirements,
        )

        requirements = metadata_requirements(selected_allocation.order_item)
        optional_wb_marking = wb_optional_marking_available(
            selected_allocation.order_item
        )
        needs_follow_up = (
            controller_marking_scan_required(selected_allocation)
            or requirements.expiry_required
            or optional_wb_marking
        )
        if needs_follow_up and (not inferred_marking_scan or requirements.expiry_required):
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    submitted_item_scan=item_scan,
                    submitted_marking_scan=inferred_marking_scan,
                    item_selected=True,
                ),
            )
    try:
        verify_pick_allocation_unit(
            allocation_id=selected_allocation.id,
            item_scan=item_scan,
            marking_scan=marking_scan,
            expiry_date_value=expiry_date_value,
            legacy_marking_absent=legacy_marking_absent,
            performed_by=request.user,
        )
    except FbsScanMismatchError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=str(exc),
                submitted_item_scan=item_scan,
                retry_marking=True,
                item_selected=True,
            ),
            status=400,
        )
    except FbsMarkingAlreadyUsedError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        error_text = str(exc)
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=error_text,
                submitted_item_scan=item_scan,
                retry_marking=True,
                item_selected=True,
            ),
            status=409,
        )
    except FbsError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        expected_item_code = str(
            selected_allocation.balance.barcode
            or selected_allocation.order_item.barcode
            or ""
        ).strip()
        error_text = str(exc)
        marking_error = any(
            marker in error_text.casefold()
            for marker in ("киз", "чз", "data matrix", "маркиров")
        )
        retry_marking = bool(
            marking_error
            and expected_item_code
            and item_scan.casefold() == expected_item_code.casefold()
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=error_text,
                submitted_item_scan=item_scan,
                retry_marking=retry_marking,
                item_selected=True,
            ),
            status=400,
        )
    selected_allocation.refresh_from_db()
    next_allocation = _next_verification_allocation(batch)
    multi_item_hold_prompt = None
    if request_role == "fbs_controller":
        multi_item_hold_prompt = _ozon_multi_item_hold_prompt(selected_allocation)
    if multi_item_hold_prompt is not None:
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                next_allocation or selected_allocation,
                ok_message=(
                    "Товар принят. Заказ состоит из нескольких товаров — "
                    "отложите эту единицу отдельно до завершения заказа."
                ),
                multi_item_hold_prompt=multi_item_hold_prompt,
            ),
        )
    if next_allocation is None:
        label = _next_verification_label(batch)
        if label is not None:
            if request_role == "fbs_controller":
                return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
            return redirect("fbs:tsd_label_detail", label_id=label.id)
        if _next_incomplete_verification_allocation(batch) is not None:
            return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if (
        request_role == "fbs_controller"
        and request.headers.get("X-Requested-With") == "fetch"
    ):
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                next_allocation,
                ok_message="Товар принят. Сканируйте следующую единицу.",
            ),
        )
    return redirect("fbs:tsd_pick_verification", batch_id=batch.id)


def _pick_verification_follow_up_request(request, location):
    """Build a GET copy for a same-screen post/redirect/get response."""
    target = urlsplit(location)
    follow_up = copy(request)
    follow_up.method = "GET"
    follow_up.path = target.path
    follow_up.path_info = target.path
    follow_up.META = request.META.copy()
    follow_up.META["REQUEST_METHOD"] = "GET"
    follow_up.META["QUERY_STRING"] = target.query
    follow_up.GET = QueryDict(target.query)
    follow_up.POST = QueryDict("")
    return follow_up


def _pick_verification_fetch_response(request, response, *, batch_id=None):
    """Return the unchanged verification result as JSON for an in-place scan."""
    if request.headers.get("X-Requested-With") != "fetch":
        return response

    location = response.get("Location")
    if location:
        target_path = urlsplit(location).path
        if (
            request.method == "POST"
            and batch_id is not None
            and target_path == request.path
        ):
            follow_up_request = _pick_verification_follow_up_request(
                request,
                location,
            )
            follow_up_response = _tsd_pick_verification_response(
                follow_up_request,
                batch_id,
            )
            return _pick_verification_fetch_response(
                follow_up_request,
                follow_up_response,
                batch_id=batch_id,
            )
        redirect_key = (
            "reload_url"
            if target_path == request.path
            else "redirect_url"
        )
        payload = {redirect_key: location}
        status = 200
    elif "text/html" in str(response.get("Content-Type") or "").casefold():
        payload = {
            "html": response.content.decode(response.charset or "utf-8", "replace"),
            "url": request.get_full_path(),
        }
        status = response.status_code
    else:
        message = response.content.decode(response.charset or "utf-8", "replace").strip()
        payload = {
            "error": message or "Сканирование отклонено сервером. Повторите попытку."
        }
        status = response.status_code

    result = JsonResponse(payload, status=status)
    result["Cache-Control"] = "no-store"
    return result


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_verification(request, batch_id: int):
    response = _tsd_pick_verification_response(request, batch_id)
    return _pick_verification_fetch_response(request, response, batch_id=batch_id)


@fbs_module_required
@role_required("fbs_controller")
@require_http_methods(["GET", "POST"])
def tsd_pick_verification_problem(request, allocation_id: int):
    allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation_id)
    task = allocation.pick_task
    if task is None:
        return HttpResponseForbidden("Позиция не относится к волне FBS")
    batch = task.batch
    if batch.verification_assigned_to_id != request.user.id:
        return HttpResponseForbidden("Проверку волны выполняет другой контролер")
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:controller_home")

    error = ""
    lookup_requested = request.GET.get("lookup") == "1"
    if request.method == "POST":
        action = str(request.POST.get("action") or "return_to_pick").strip()
        try:
            if action == "return_problem_item":
                from .services.problems import return_problem_tote_item_to_check_tote

                lookup_requested = True
                return_problem_tote_item_to_check_tote(
                    allocation_id=allocation.id,
                    problem_item_id=int(request.POST.get("problem_item_id") or 0),
                    problem_tote_scan=str(
                        request.POST.get("problem_tote_scan") or ""
                    ).strip(),
                    returned_by=request.user,
                )
            elif action == "return_to_pick":
                queue_verification_problem_restock(
                    allocation_id=allocation.id,
                    exception_type=str(
                        request.POST.get("exception_type") or ""
                    ).strip(),
                    reason=str(request.POST.get("reason") or "").strip(),
                    service_tote_scan=str(
                        request.POST.get("service_tote_scan") or ""
                    ).strip(),
                    reported_by=request.user,
                    order_scan=str(request.POST.get("order_scan") or "").strip(),
                )
            else:
                raise ValueError("Неизвестная команда проблемного товара.")
        except (FbsError, TypeError, ValueError) as exc:
            error = str(exc)
            allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation.id)
            if request.POST.get("modal_flow") == "1":
                problem_kind = str(
                    request.POST.get("problem_kind") or "canceled"
                ).strip()
                if problem_kind not in {"canceled", "label", "marking"}:
                    problem_kind = "canceled"
                prompt_reason = str(request.POST.get("reason") or "").strip()
                return render(
                    request,
                    "fbs/tsd_pick_verification.html",
                    _pick_verification_context(
                        request,
                        batch,
                        allocation,
                        error=error,
                        item_selected=True,
                        service_tote_prompt=_verification_service_tote_prompt(
                            batch=batch,
                            allocation=allocation,
                            controller=request.user,
                            problem_kind=problem_kind,
                            reason=prompt_reason,
                        ),
                    ),
                    status=409,
                )
        else:
            if action == "return_problem_item":
                return redirect(
                    f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}"
                    "?problem_returned=1"
                )
            return redirect(
                f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}?problem=1"
            )

    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if batch.workstation_id and batch.workstation.device_agent_id:
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    pick_tote_context = (
        FbsControllerPickTote.objects.select_related(
            "session__problem_tote",
            "session__canceled_tote",
        )
        .filter(
            pick_batch=batch,
            session__controller=request.user,
            session__status="active",
            status__in=(
                FbsControllerPickTote.STATUS_PROCESSING,
                FbsControllerPickTote.STATUS_AWAITING_EMPTY,
            ),
        )
        .order_by("-id")
        .first()
    )
    session = pick_tote_context.session if pick_tote_context else None
    problem_tote = session.problem_tote if session else None
    order_is_canceled = _marketplace_order_is_canceled(task.order)
    service_tote = (
        session.canceled_tote if order_is_canceled and session else None
    ) or (problem_tote if not order_is_canceled else None)
    problem_tote_lookup = None
    if lookup_requested:
        try:
            from .services.problems import lookup_problem_tote_item

            problem_tote_lookup = lookup_problem_tote_item(
                allocation_id=allocation.id,
                actor=request.user,
            )
        except FbsError as exc:
            if not error:
                error = str(exc)
    return render(
        request,
        "fbs/tsd_pick_verification_problem.html",
        _base_context(
            request,
            page_title=f"FBS · Проблема заказа {task.order.external_order_id}",
            back_url=reverse("fbs:tsd_pick_verification", kwargs={"batch_id": batch.id}),
            batch=batch,
            task=task,
            order=task.order,
            allocation=allocation,
            balance=allocation.balance,
            problem_tote=problem_tote,
            service_tote=service_tote,
            service_tote_kind=("canceled" if order_is_canceled else "problem"),
            service_tote_label=(
                "тара отмененных заказов"
                if order_is_canceled
                else "проблемная тара"
            ),
            problem_tote_lookup=problem_tote_lookup,
            exception_choices=FbsPickException.TYPE_CHOICES,
            agent_scan_poll_url=agent_scan_poll_url,
            agent_scan_event_id=agent_scan_event_id,
            error=error,
        ),
        status=409 if error else 200,
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_plan(request, plan_id: int):
    plan = get_object_or_404(_plan_queryset(), pk=plan_id)
    ok_message = ""
    if request.GET.get("confirmed") == "1":
        ok_message = "План подтвержден и передан ричтраку."
    elif request.GET.get("canceled") == "1":
        ok_message = "План отменен, резерв общего склада снят."
    elif request.GET.get("packed") == "1":
        ok_message = "QR FBS-короба создан. Распечатайте его и передайте короб водителю."
    return render(
        request,
        "fbs/tsd_storekeeper_plan.html",
        _storekeeper_plan_context(request, plan, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_confirm_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        confirm_replenishment_plan(plan_id=plan_id, confirmed_by=request.user)
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?confirmed=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_cancel_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        cancel_replenishment_plan(plan_id=plan_id, canceled_by=request.user)
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?canceled=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_pack_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        pack_staged_item_plan(
            plan_id=plan_id,
            packed_by=request.user,
            staging_container_scan=request.POST.get(
                "staging_container_scan", ""
            ),
        )
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?packed=1")


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_GET
def tsd_reachtruck(request):
    return redirect("/reachtruck/")


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_POST
def tsd_claim_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        plan = claim_replenishment_plan(plan_id=plan_id, assigned_to=request.user)
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_reachtruck_list.html",
            _reachtruck_list_context(request, error=str(exc)),
            status=409,
        )
    allocation = (
        _allocation_queryset()
        .filter(line__plan=plan, status__in=OPEN_ALLOCATION_STATUSES)
        .order_by("id")
        .first()
    )
    if allocation is None:
        return redirect("fbs:tsd_reachtruck")
    return redirect("fbs:tsd_allocation", allocation_id=allocation.id)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_POST
def tsd_claim_internal_movement(request, movement_id: int):
    get_object_or_404(FbsInternalMovement, pk=movement_id)
    try:
        claim_internal_movement(movement_id=movement_id, assigned_to=request.user)
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_reachtruck_list.html",
            _reachtruck_list_context(request, error=str(exc)),
            status=409,
        )
    return redirect("fbs:tsd_internal_movement", movement_id=movement_id)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_internal_movement(request, movement_id: int):
    movement = get_object_or_404(
        FbsInternalMovement.objects.select_related(
            "agency",
            "source_box__pallet__cell",
            "target_pallet__cell",
            "target_box",
            "source_balance",
            "assigned_to",
        ),
        pk=movement_id,
    )
    if movement.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Перемещение назначено другому сотруднику")
    if movement.status == FbsInternalMovement.STATUS_DONE:
        return redirect("fbs:tsd_reachtruck")
    context = _base_context(
        request,
        page_title=f"FBS · Перемещение #{movement.id}",
        back_url=reverse("fbs:tsd_reachtruck"),
        movement=movement,
    )
    if request.method == "GET":
        return render(request, "fbs/tsd_internal_movement.html", context)
    try:
        movement = scan_internal_movement(
            movement_id=movement.id,
            source_box_scan=request.POST.get("source_box_scan", ""),
            target_scan=request.POST.get("target_scan", ""),
            item_scan=request.POST.get("item_scan", ""),
            performed_by=request.user,
        )
    except FbsError as exc:
        context["error"] = str(exc)
        return render(request, "fbs/tsd_internal_movement.html", context, status=400)
    if movement.status == FbsInternalMovement.STATUS_DONE:
        return redirect(f"{reverse('fbs:tsd_reachtruck')}?movement_done=1")
    context["movement"] = movement
    context["ok_message"] = f"Перемещено {movement.moved_qty} из {movement.planned_qty}."
    return render(request, "fbs/tsd_internal_movement.html", context)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_allocation(request, allocation_id: int):
    allocation = get_object_or_404(_allocation_queryset(), pk=allocation_id)
    plan = allocation.line.plan
    if plan.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Задание назначено другому сотруднику")
    if allocation.status in {
        FbsReplenishmentAllocation.STATUS_STAGED,
        FbsReplenishmentAllocation.STATUS_DONE,
    }:
        return redirect("fbs:tsd_reachtruck")

    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_allocation.html",
            _allocation_context(request, allocation),
        )

    action = str(request.POST.get("action") or "").strip()
    if action == "scan_source":
        source_scan = str(request.POST.get("source_scan") or "").strip()
        try:
            validate_replenishment_source_scan(
                allocation_id=allocation.id,
                source_scan=source_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(_source_session_key(allocation.id), None)
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session[_source_session_key(allocation.id)] = source_scan
        return render(
            request,
            "fbs/tsd_allocation.html",
            _allocation_context(
                request,
                allocation,
                ok_message="Источник подтвержден.",
            ),
        )

    if action == "complete":
        source_scan = str(request.session.get(_source_session_key(allocation.id)) or "").strip()
        target_scan = str(
            request.POST.get("staging_scan")
            or request.POST.get("target_box_scan")
            or ""
        ).strip()
        if not source_scan:
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error="Сначала подтвердите источник."),
                status=400,
            )
        try:
            is_staging_transfer = bool(
                plan.client_movement_request_id
                and plan.mode == FbsReplenishmentPlan.MODE_ITEM
                and plan.target_box_id is None
                and plan.staging_location_id
            )
            if is_staging_transfer:
                stage_replenishment_allocation(
                    allocation_id=allocation.id,
                    source_scan=source_scan,
                    staging_container_scan=request.POST.get(
                        "staging_container_scan", ""
                    ),
                    staging_scan=target_scan,
                    performed_by=request.user,
                )
            else:
                complete_replenishment_allocation(
                    allocation_id=allocation.id,
                    source_scan=source_scan,
                    target_box_scan=target_scan,
                    performed_by=request.user,
                )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session.pop(_source_session_key(allocation.id), None)
        next_allocation = (
            _allocation_queryset()
            .filter(line__plan=plan, status__in=OPEN_ALLOCATION_STATUSES)
            .order_by("id")
            .first()
        )
        if next_allocation is not None:
            return redirect("fbs:tsd_allocation", allocation_id=next_allocation.id)
        result_flag = "staged" if is_staging_transfer else "done"
        return redirect(f"{reverse('fbs:tsd_reachtruck')}?{result_flag}=1")

    return render(
        request,
        "fbs/tsd_allocation.html",
        _allocation_context(request, allocation, error="Неизвестная команда."),
        status=400,
    )
