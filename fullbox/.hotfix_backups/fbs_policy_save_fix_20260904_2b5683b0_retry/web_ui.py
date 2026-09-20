"""Head manager UI views and marketplace helper logic."""

import json
import re
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Iterable
from urllib.parse import urlencode

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import requests
from agent.models import DeviceAgent
from django.db import IntegrityError, transaction
from django.db.models import Count, Prefetch, Q, Sum

from audit.models import OrderAuditEntry
from client_cabinet.services import build_client_cabinet_url, build_client_list_context, build_client_list_queryset
from django.conf import settings
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.utils import timezone
from django.views import View
from django.views.generic import CreateView, ListView, TemplateView, UpdateView

from client_cabinet.models import OtherRequest
from employees.access import RoleRequiredMixin, get_request_role
from employees.models import Employee
from fbs.exceptions import FbsEquipmentError, FbsError
from fbs.flags import feature_enabled, module_enabled as fbs_module_enabled
from fbs.integrations.warehouse_catalog import (
    fetch_client_warehouse_catalog,
    marketplace_code,
)
from fbs.models import (
    FbsClientStoragePolicy,
    FbsComplianceOverride,
    FbsControllerPolicy,
    FbsEquipmentCodeSequence,
    FbsIntegrationProfile,
    FbsHandoverBatch,
    FbsHandoverVerificationOverride,
    FbsInternalMovement,
    FbsInventorySession,
    FbsMarketplaceMetadataTransfer,
    FbsOrder,
    FbsOzonCredential,
    FbsPickBatch,
    FbsPickException,
    FbsPickTask,
    FbsPickingCart,
    FbsReplenishmentPolicy,
    FbsSyncCursor,
    FbsWorkstation,
)
from fbs.services.problems import client_operations_report
from fbs.services.handover import (
    approve_compliance_override,
    approve_handover_verification_override,
)
from fbs.services.client_profiles import (
    SelectedMarketplaceWarehouse,
    configure_client_fbs_profiles,
)
from fbs.services.sync import pull_profile_orders_manually
from fbs.services.equipment import (
    create_fbs_picking_carts,
    create_fbs_workstations,
    equipment_label_queryset,
    update_fbs_picking_cart,
    update_fbs_workstation,
)
from fullbox.order_numbers import format_order_number, order_number_sequence
from logistics.models import LogisticsTrip
from marking.models import MarkingCode
from receiving_cz.models import ReceivingCzUnit
from receiving_cz.services import normalize_marking_code
from shipping.models import ShippingOrder
from sklad.models import WarehouseContainer, WarehouseEvent, WarehouseLocation, WarehouseStockSnapshot
from sklad.services.warehouse_events import WarehouseEventType
from sklad.services.operational_locations import (
    create_operational_location,
    operational_location_occupancy,
    update_operational_location,
)
from sklad.services.warehouse_stock_rows import normalize_stock_row_from_snapshot
from sklad.ui_services import _journal_edit_lock_reason
from sku.models import Agency, Market, MarketCredential, SKU
from todo.models import Task, build_receiving_display_context, _shipping_task_title_with_display_id

from .forms import (
    CarrierForm,
    FbsPickingCartBatchCreateForm,
    FbsPickingCartUpdateForm,
    FbsWorkstationBatchCreateForm,
    FbsWorkstationUpdateForm,
    OwnCompanyForm,
)
from .employee_report import (
    warehouse_employee_metric_points,
    warehouse_employee_work_windows,
)
from .employee_activity_report import (
    EMPLOYEE_ACTIVITY_CONTOUR_LABELS,
    warehouse_employee_activity_rollups,
    warehouse_employee_activity_timeline,
)
from .inspection_ai import build_ai_inspection_report
from .models import Carrier, OwnCompany
from .services import (
    build_marketplace_warehouses_context,
    build_reference_form_context,
    build_reference_list_context,
    save_marketplace_warehouses_response,
    sync_marketplace_warehouses_response,
)

CLIENT_SORT_FIELDS = {
    "name": "agn_name",
    "short_name": "short_name",
    "pref": "pref",
    "inn": "inn",
    "email": "email",
    "phone": "phone",
    "id": "id",
}
CLIENT_FILTER_FIELDS = {
    "agn_name": "agn_name",
    "short_name": "short_name",
    "inn": "inn",
    "pref": "pref",
    "email": "email",
    "phone": "phone",
}


_MOVEMENT_REPORT_HEADERS = [
    "№",
    "Клиент",
    "Паллет №",
    "Короб №",
    "SKU",
    "Номенклатура",
    "Штрихкод",
    "Количество, шт",
    "Статус",
    "Ширина, мм",
    "Высота, мм",
    "Глубина, мм",
    "Объем, м³",
    "Вес, г",
    "Приход дата",
    "Номер приходного документа",
    "Расход дата",
    "Номер расходного документа",
]

_SKU_MOVEMENT_EXPORT_DEFAULT_DAYS = 30
_SKU_MOVEMENT_EXPORT_ROW_LIMIT = 5000

_CONTAINER_REPORTS = {
    "boxes": {
        "title": "Состав коробов",
        "subtitle": "Содержимое коробов: клиент, SKU, количество, ЧЗ и текущее место хранения.",
        "container_label": "Короб",
        "container_key": "box_code",
        "export_name": "boxes",
    },
    "pallets": {
        "title": "Состав паллет",
        "subtitle": "Содержимое паллет: короба, SKU, количество, ЧЗ и текущее место хранения.",
        "container_label": "Паллета",
        "container_key": "pallet_code",
        "export_name": "pallets",
    },
}

_CONTAINER_REPORT_HEADERS = [
    "Контейнер",
    "Паллета",
    "Короб",
    "Клиент",
    "SKU",
    "Наименование",
    "Размер",
    "ШК",
    "Кол-во",
    "Доступно",
    "Резерв обр.",
    "Резерв отг.",
    "ЧЗ",
    "Склад",
    "Зона",
    "Место",
    "Статус",
    "Заявка",
]

_FBS_EMPLOYEE_REPORT_HEADERS = [
    "Дата",
    "Сотрудник",
    "Роль",
    "Волн собрано",
    "Заказов собрано",
    "Единиц собрано",
    "Среднее время заказа, мин",
    "Волн проверено",
    "Заказов проверено",
    "Единиц проверено",
    "Отгрузок передано",
    "Коробов отгружено",
    "Заказов отгружено",
    "Единиц отгружено",
]

_WAREHOUSE_EMPLOYEE_REPORT_GROUPS = (
    {
        "title": "Приёмка",
        "metrics": (
            ("receiving_cz_units", "Принято ЧЗ", "Количество принятых кодов Честного знака."),
            ("palletization_positions", "Паллетизация, поз.", "Завершённые позиции паллетизации; стартовые события не учитываются."),
            ("palletization_units", "Паллетизация, шт.", "Количество единиц в завершённых позициях паллетизации."),
        ),
    },
    {
        "title": "Размещение и перемещения",
        "metrics": (
            ("placement_positions", "Размещено, поз.", "Завершённые позиции размещения."),
            ("placement_units", "Размещено, шт.", "Количество размещённых единиц."),
            ("reachtruck_tasks", "Ричтрак, заданий", "Завершённые задания водителя ричтрака."),
            ("reachtruck_units", "Ричтрак, шт.", "Фактическое количество по завершённым заданиям ричтрака."),
            ("fbs_replenishment_tasks", "Подсорт FBS, заданий", "Уникальные завершённые задания пополнения FBS."),
            ("fbs_replenishment_units", "Подсорт FBS, шт.", "Количество единиц в завершённых заданиях пополнения FBS."),
            ("movement_positions", "Перемещено, поз.", "Завершённые складские позиции перемещения."),
            ("movement_units", "Перемещено, шт.", "Количество перемещённых единиц."),
            ("otg_positions", "В OTG, поз.", "Позиции, фактически доставленные в зону отгрузки."),
            ("otg_units", "В OTG, шт.", "Количество единиц, доставленных в зону отгрузки."),
        ),
    },
    {
        "title": "FBS-отбор",
        "metrics": (
            ("pick_waves", "Волн", "Волны с завершённым отбором."),
            ("pick_orders", "Заказов", "Заказы в завершённых волнах отбора."),
            ("pick_units", "Единиц", "Отобранные единицы в завершённых волнах."),
            ("avg_pick_minutes", "Среднее, мин", "Средняя длительность завершённой волны отбора."),
        ),
    },
    {
        "title": "FBS-проверка",
        "metrics": (
            ("verified_waves", "Волн", "Волны с завершённой проверкой."),
            ("verified_orders", "Заказов", "Заказы в завершённых проверках."),
            ("verified_units", "Единиц", "Единицы в завершённых проверках."),
            ("kiz_scans_success", "КИЗ принят", "Успешные сканирования КИЗ на проверке."),
            ("kiz_scan_errors", "Ошибки КИЗ", "Ошибочные сканирования КИЗ на проверке."),
            ("barcode_scan_errors", "Ошибки ШК", "Ошибочные сканирования штрихкода товара."),
        ),
    },
    {
        "title": "Отгрузка",
        "metrics": (
            ("shipments", "FBS поставок", "FBS-поставки, переданные сотрудником."),
            ("shipment_boxes", "FBS коробов", "Короба в переданных FBS-поставках."),
            ("shipment_orders", "FBS заказов", "Заказы в переданных FBS-поставках."),
            ("shipment_units", "FBS единиц", "Единицы в переданных FBS-поставках."),
            ("shipping_orders", "Заявок", "Уникальные не-FBS заявки с завершённой складской отгрузкой."),
            ("shipping_positions", "Позиций", "Завершённые позиции не-FBS отгрузки."),
            ("shipping_units", "Единиц", "Количество единиц не-FBS отгрузки."),
            ("shipping_containers", "Контейнеров", "Уникальные контейнеры не-FBS отгрузки."),
        ),
    },
    {
        "title": "Обработка",
        "metrics": (
            ("processing_operations", "Операций", "Завершённые назначенные операции обработки."),
            ("processing_units", "Обработано, шт.", "Фактическое количество по завершённым операциям обработки."),
            ("processing_boxes", "Сформировано коробов", "Финальные короба, закрытые при завершении формирования."),
            ("processing_boxed_units", "В коробах, шт.", "Количество единиц в финально сформированных коробах."),
        ),
    },
    {
        "title": "Инвентаризация",
        "metrics": (
            ("inventory_lines", "Строк", "Фактически пересчитанные строки инвентаризации."),
            ("inventory_units", "Единиц", "Фактическое количество в пересчитанных строках."),
            ("reachtruck_inventory_tasks", "Заданий ричтрака", "Завершённые задания ричтрака по инвентаризации."),
        ),
    },
)

_WAREHOUSE_EMPLOYEE_REPORT_METRIC_KEYS = tuple(
    metric[0]
    for group in _WAREHOUSE_EMPLOYEE_REPORT_GROUPS
    for metric in group["metrics"]
)

_WAREHOUSE_EMPLOYEE_CONTOUR_LABELS = {
    "fbs_handover": "FBS-отгрузка",
    "fbs_picking": "FBS-отбор",
    "fbs_quality": "FBS-сканирование",
    "fbs_replenishment": "подсорт FBS",
    "fbs_verification": "FBS-проверка",
    "inventory": "инвентаризация",
    "movement": "перемещение",
    "otg": "доставка в OTG",
    "palletization": "паллетизация",
    "placement": "размещение",
    "reachtruck_inventory": "инвентаризация ричтраком",
    "reachtruck_move": "ричтрак",
    "receiving_cz": "приёмка ЧЗ",
    "shipping": "отгрузка",
    "processing": "обработка и формирование коробов",
}

_WAREHOUSE_EMPLOYEE_PROCESSING_ROLES = {"packer", "processing_worker"}
_WAREHOUSE_EMPLOYEE_ALLOWED_ROLES = ("head_manager", "director", "admin")

_SKU_MOVEMENT_REPORT_HEADERS = [
    "Дата",
    "Операция",
    "Направление",
    "Клиент",
    "SKU",
    "Наименование",
    "ШК",
    "ЧЗ",
    "Кол-во",
    "Откуда",
    "Куда",
    "Короб",
    "Паллета",
    "Документ",
    "Отгрузка",
    "Остаток сейчас",
    "Доступно",
    "В резерве",
]

_SKU_MOVEMENT_EVENT_LABELS = {
    WarehouseEventType.RECEIVING_ARRIVED.value: "Приход на приемку",
    WarehouseEventType.PLACEMENT_COMPLETED.value: "Приход / размещение",
    WarehouseEventType.PUTAWAY_REQUESTED.value: "Задание на размещение",
    WarehouseEventType.PUTAWAY_COMPLETED.value: "Размещено на складе",
    WarehouseEventType.MOVEMENT_REQUESTED.value: "Задание на перемещение",
    WarehouseEventType.MOVEMENT_STARTED.value: "Перемещение начато",
    WarehouseEventType.MOVEMENT_COMPLETED.value: "Перемещение выполнено",
    WarehouseEventType.MOVEMENT_CANCELED.value: "Перемещение отменено",
    WarehouseEventType.STOCK_RETURNED_TO_STORAGE.value: "Возврат в остатки",
    WarehouseEventType.STOCK_CORRECTED.value: "Корректировка остатка",
    WarehouseEventType.PROCESSING_RESERVED.value: "Резерв под обработку",
    WarehouseEventType.PROCESSING_RESERVE_RELEASED.value: "Резерв обработки снят",
    WarehouseEventType.PROCESSING_ZONE_ARRIVED.value: "Передано в обработку",
    WarehouseEventType.PROCESSING_STARTED.value: "Обработка начата",
    WarehouseEventType.PROCESSING_COMPLETED.value: "Обработка выполнена",
    WarehouseEventType.PROCESSING_CONSUMED.value: "Списано в обработку",
    WarehouseEventType.SHIPPING_RESERVED.value: "Резерв под отгрузку",
    WarehouseEventType.SHIPPING_RESERVE_RELEASED.value: "Резерв отгрузки снят",
    WarehouseEventType.OTG_REQUESTED.value: "Задание в зону отгрузки",
    WarehouseEventType.OTG_ARRIVED.value: "Прибыло в зону отгрузки",
    WarehouseEventType.PALLETIZATION_STARTED.value: "Паллетизация начата",
    WarehouseEventType.PALLETIZATION_COMPLETED.value: "Паллетизация выполнена",
    WarehouseEventType.READY_FOR_LOADING.value: "Готово к погрузке",
    WarehouseEventType.ASSIGNED_TO_TRIP.value: "Назначено на рейс",
    WarehouseEventType.LOADING_STARTED.value: "Погрузка начата",
    WarehouseEventType.LOADED_TO_VEHICLE.value: "Загружено в машину",
    WarehouseEventType.SHIPPED.value: "Отгружено",
    WarehouseEventType.WAREHOUSE_CONTEXT_CANCELED.value: "Складской контур отменен",
}

_SKU_MOVEMENT_IN_EVENTS = {
    WarehouseEventType.RECEIVING_ARRIVED.value,
    WarehouseEventType.PLACEMENT_COMPLETED.value,
    WarehouseEventType.PUTAWAY_COMPLETED.value,
    WarehouseEventType.STOCK_RETURNED_TO_STORAGE.value,
}

_SKU_MOVEMENT_OUT_EVENTS = {
    WarehouseEventType.PROCESSING_CONSUMED.value,
    WarehouseEventType.SHIPPING_RESERVED.value,
    WarehouseEventType.OTG_REQUESTED.value,
    WarehouseEventType.OTG_ARRIVED.value,
    WarehouseEventType.LOADED_TO_VEHICLE.value,
    WarehouseEventType.SHIPPED.value,
}

_EXCEL_FORBIDDEN_CHARS_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1C\x1E-\x1F]")


def _excel_safe_value(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    text = value.replace("\x1d", "<GS>")
    text = _EXCEL_FORBIDDEN_CHARS_RE.sub("", text)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text
_MOVEMENT_REPORT_RECEIVING_EVENTS = {
    WarehouseEventType.RECEIVING_ARRIVED.value,
    WarehouseEventType.PLACEMENT_COMPLETED.value,
    WarehouseEventType.PUTAWAY_COMPLETED.value,
}
_MOVEMENT_REPORT_SHIPPING_EVENTS = {
    WarehouseEventType.SHIPPED.value,
}
_MOVEMENT_REPORT_STATUS_LABELS = {
    "unknown": "Неизвестно",
    "received_unplaced": "Принят, не размещен",
    "placed_in_receiving": "В зоне приемки",
    "stored": "На складе",
    "reserved_for_processing": "В резерве обработки",
    "moving_to_processing": "Перемещается в обработку",
    "in_processing_zone": "В зоне обработки",
    "processing_in_progress": "В обработке",
    "processing_consumed": "Списан в обработку",
    "placed_after_processing": "После обработки",
    "reserved_for_shipping": "В резерве отгрузки",
    "moving_to_otg": "Перемещается в отгрузку",
    "in_otg": "В зоне отгрузки",
    "palletizing": "Паллетизация",
    "ready_for_loading": "Готов к погрузке",
    "assigned_to_trip": "Назначен рейс",
    "loading_in_progress": "Погрузка",
    "loaded_to_vehicle": "Погружен в машину",
    "shipped": "Отгружен",
    "partially_shipped": "Частично отгружен",
    "canceled": "Отменен",
}
_MOVEMENT_REPORT_PROCESSING_SOURCE_STATES = {
    "moving_to_processing",
    "in_processing_zone",
    "processing_in_progress",
}


def _movement_report_datetime(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")


def _movement_report_order_type(value: str) -> str:
    text = str(value or "").strip().lower()
    if "receiving" in text or text == "placement_act":
        return "receiving"
    if "shipping" in text:
        return "shipping"
    if "processing" in text:
        return "processing"
    return text


def _movement_report_document_number(event: WarehouseEvent | None, fallback_type: str, fallback_id: str) -> str:
    if event is not None:
        source_document_id = str(getattr(event, "source_document_id", "") or "").strip()
        event_context_type = str(getattr(event, "stock_context_type", "") or "").strip()
        if source_document_id:
            source_document_type = str(getattr(event, "source_document_type", "") or "").strip()
            document_type = _movement_report_order_type(event_context_type or source_document_type)
            return format_order_number(document_type, source_document_id)
        event_context_id = str(getattr(event, "stock_context_id", "") or "").strip()
        if event_context_id:
            return format_order_number(_movement_report_order_type(event_context_type), event_context_id)
    return format_order_number(_movement_report_order_type(fallback_type), fallback_id)


def _movement_report_client_label(snapshot: WarehouseStockSnapshot, row: dict) -> str:
    agency = snapshot.agency
    prefix = str(getattr(agency, "pref", "") or "").strip()
    if prefix:
        return prefix.upper()
    for key in ("pallet_code", "box_code"):
        code = str(row.get(key) or "").strip()
        if "-" in code:
            return code.split("-", 1)[0].strip().upper()
    return str(getattr(agency, "short_name", "") or getattr(agency, "agn_name", "") or "").strip()


def _movement_report_status(snapshot: WarehouseStockSnapshot, shipping_event: WarehouseEvent | None) -> str:
    if shipping_event is not None:
        return _MOVEMENT_REPORT_STATUS_LABELS["shipped"]
    state_code = str(snapshot.warehouse_state_code or "").strip().lower()
    return _MOVEMENT_REPORT_STATUS_LABELS.get(state_code, state_code or "")


def _movement_report_volume_m3(width, height, depth):
    if not width or not height or not depth:
        return ""
    return round((int(width) * int(height) * int(depth)) / 1_000_000_000, 6)


def _head_manager_shipping_number(number) -> str:
    sequence = order_number_sequence("shipping", str(number or "").strip())
    if sequence is not None:
        return f"{sequence}_OTG"
    return format_order_number("shipping", number)


def _head_manager_agency_label(agency) -> str:
    return str(getattr(agency, "short_name", "") or getattr(agency, "agn_name", "") or "-")


def _head_manager_marketplace_label(marketplace) -> str:
    return str(getattr(marketplace, "name", "") or getattr(marketplace, "title", "") or "-")


def _head_manager_shipping_meta(order: ShippingOrder | None) -> dict:
    if order is None:
        return {
            "client": "-",
            "warehouse": "-",
            "zone": "-",
            "marketplace": "-",
        }
    warehouse = (
        str(order.destination_warehouse or "").strip()
        or str(order.transit_address or "").strip()
        or str(order.destination_address or "").strip()
        or "-"
    )
    return {
        "client": _head_manager_agency_label(order.agency),
        "warehouse": warehouse,
        "zone": "OTG",
        "marketplace": _head_manager_marketplace_label(order.marketplace),
    }


def _head_manager_payload_value(payload: dict | None, *keys: str) -> str:
    if not isinstance(payload, dict):
        return ""
    for key in keys:
        value = str(payload.get(key) or "").strip()
        if value:
            return value
    return ""


def _head_manager_receiving_meta(entries: list[OrderAuditEntry]) -> dict:
    agency = next((entry.agency for entry in reversed(entries) if getattr(entry, "agency_id", None)), None)
    payload = entries[-1].payload if entries else {}
    return {
        "client": _head_manager_agency_label(agency),
        "warehouse": _head_manager_payload_value(
            payload,
            "warehouse",
            "warehouse_name",
            "warehouse_code",
            "destination_warehouse",
        )
        or "Приемка",
        "zone": _head_manager_payload_value(payload, "zone", "zone_code") or "PR",
        "marketplace": _head_manager_payload_value(payload, "marketplace", "marketplace_name", "market") or "-",
    }


def _head_manager_request_type(route: str, title: str = "") -> str:
    source = f"{route or ''} {title or ''}".lower()
    if "/logistics/trips/" in source or "рейс" in source:
        return "trip"
    if "/orders/receiving/" in source:
        return "receiving"
    if "/orders/processing/" in source or "/processing/" in source:
        return "processing"
    if "/shipping/" in source or "/orders/shipping/" in source or "отгруз" in source:
        return "shipping"
    if "/stockmap/" in source or "размещ" in source:
        return "placement"
    if "/orders/other/" in source or "проч" in source:
        return "other"
    return "task"


_HEAD_MANAGER_REQUEST_TYPES = {
    "all": "Все",
    "receiving": "Приемка",
    "processing": "Обработка",
    "shipping": "Отгрузка",
    "placement": "Размещение",
    "trip": "Рейс",
    "other": "Прочие",
    "task": "Задачи",
}

_HEAD_MANAGER_RESPONSIBILITY_FILTERS = {
    "warehouse": "Склад",
    "manager": "Менеджеры",
    "all": "Все",
}

_HEAD_MANAGER_MANAGER_TASK_TITLE_MARKERS = (
    "клиент подтвердил акт",
    "проверьте размещение",
    "подтвердите заявку",
)

_HEAD_MANAGER_STAGE_FILTERS = {
    "all": "Все статусы",
    "overdue": "Просрочено",
    "today": "Сегодня",
    "soon": "Скоро",
    "in_work": "В работе",
    "no_owner": "Без исполнителя",
    "done": "Готово",
}


def _head_manager_is_manager_responsibility_task(task: Task, request_type: str) -> bool:
    role = str(getattr(task.assigned_to, "role", "") or "").strip()
    if role == "manager":
        return True
    title = str(task.title or "").strip().lower()
    if request_type == "receiving" and any(marker in title for marker in _HEAD_MANAGER_MANAGER_TASK_TITLE_MARKERS):
        return True
    return False


def _head_manager_task_responsibility(task: Task, request_type: str) -> str:
    if _head_manager_is_manager_responsibility_task(task, request_type):
        return "manager"
    return "warehouse"

_HEAD_MANAGER_STOCK_UNITS = {
    "items": {"label": "Штуки", "suffix": "шт."},
    "boxes": {"label": "Короба", "suffix": "кор."},
    "pallets": {"label": "Паллеты", "suffix": "пал."},
}

_HEAD_MANAGER_NON_STORAGE_ZONES = ("PR", "OBR", "OTG")

_HEAD_MANAGER_SHIPPING_WATCH_STATUSES = (
    ShippingOrder.STATUS_SUBMITTED,
    ShippingOrder.STATUS_RESERVED,
    ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
    ShippingOrder.STATUS_PICKING,
    ShippingOrder.STATUS_PACKED,
)

_HEAD_MANAGER_SHIPPING_STAGE_LABELS = {
    ShippingOrder.STATUS_SUBMITTED: "Ждет принятия складом",
    ShippingOrder.STATUS_RESERVED: "Ждет отбора",
    ShippingOrder.STATUS_STOREKEEPER_ACCEPTED: "В работе",
    ShippingOrder.STATUS_PICKING: "В отборе",
    ShippingOrder.STATUS_PACKED: "Ждет погрузки",
}

_HEAD_MANAGER_REPORT_FEATURES = [
    "Поиск",
    "Сортировка",
    "Группировка",
    "Сохранение фильтров",
    "Быстрые пресеты",
    "Экспорт Excel",
    "Экспорт PDF",
    "Печать",
    "Настройка колонок",
    "Шаблоны пользователя",
    "Графики и диаграммы",
]

_HEAD_MANAGER_REPORT_FILTERS = [
    {"key": "period", "label": "Период", "kind": "period"},
    {"key": "warehouse", "label": "Склад", "kind": "select", "options": ["Все склады", "Коледино", "Тула", "Екатеринбург"]},
    {"key": "stock_type", "label": "FBS / FBO", "kind": "select", "options": ["Все", "FBS", "FBO"]},
    {"key": "client", "label": "Клиент", "kind": "search", "placeholder": "Название клиента"},
    {"key": "product", "label": "Товар", "kind": "search", "placeholder": "Название товара"},
    {"key": "sku", "label": "SKU", "kind": "search", "placeholder": "Артикул или SKU"},
    {"key": "supplier", "label": "Поставщик", "kind": "search", "placeholder": "Поставщик"},
    {"key": "employee", "label": "Сотрудник", "kind": "search", "placeholder": "Исполнитель"},
    {"key": "status", "label": "Статус", "kind": "search", "placeholder": "Статус"},
    {"key": "zone", "label": "Зона хранения", "kind": "search", "placeholder": "PR, OS, OTG"},
    {"key": "cell", "label": "Ячейка", "kind": "search", "placeholder": "Адрес ячейки"},
    {"key": "operation_type", "label": "Тип операции", "kind": "select", "options": ["Все", "Приемка", "Отгрузка", "Размещение", "Перемещение", "Инвентаризация"]},
]

_HEAD_MANAGER_REPORT_CATEGORIES = [
    {
        "slug": "products",
        "title": "Отчёты по товарам",
        "icon": "SKU",
        "tone": "blue",
        "description": "Остатки, движение, история SKU и товарная аналитика.",
        "reports": [
            {"slug": "movement", "title": "Движение товаров", "description": "По артикулу: когда пришел, куда ушел, по какой отгрузке и текущий остаток.", "implemented_url": "/head-manager/reports/products/sku-movement/"},
            {"slug": "movement-for-invoices", "title": "Движение товара для счетов", "description": "Старый отчет по приходам, расходам, документам и статусам.", "implemented_url": "/head-manager/report/"},
            {"slug": "sku-history", "title": "История SKU", "description": "Изменения карточек SKU и связанных параметров."},
            {"slug": "turnover", "title": "Оборачиваемость", "description": "Скорость оборота товаров за выбранный период."},
            {"slug": "no-movement", "title": "Товары без движения", "description": "SKU без операций за выбранный период."},
            {"slug": "shortage", "title": "Дефицит товаров", "description": "Позиции с риском нехватки."},
            {"slug": "surplus", "title": "Излишки товаров", "description": "Позиции с избыточным остатком."},
            {"slug": "abc", "title": "ABC-анализ", "description": "Классификация товаров по вкладу."},
            {"slug": "xyz", "title": "XYZ-анализ", "description": "Классификация по стабильности спроса."},
            {"slug": "top-products", "title": "ТОП товаров", "description": "Лидеры по операциям, остаткам и обороту."},
            {"slug": "product-changes", "title": "История изменений товара", "description": "Аудит корректировок товарных данных."},
        ],
    },
    {
        "slug": "warehouse",
        "title": "Складские отчёты",
        "icon": "WH",
        "tone": "green",
        "description": "Складские остатки, зоны, ячейки и перемещения.",
        "reports": [
            {"slug": "warehouse-stock", "title": "Остатки склада", "description": "Текущие остатки по складам, зонам и статусам.", "implemented_url": "/head-manager/stock-editor/"},
            {"slug": "fbs-stock", "title": "Остатки FBS", "description": "Остатки по FBS-контуру."},
            {"slug": "fbo-stock", "title": "Остатки FBO", "description": "Остатки по FBO-контуру."},
            {"slug": "fbs-fbo-compare", "title": "Сравнение остатков FBS/FBO", "description": "Расхождения между контурами хранения."},
            {"slug": "fbo-to-fbs", "title": "Подсорт FBO → FBS", "description": "Потребность в подсорте между контурами."},
            {"slug": "warehouse-load", "title": "Загруженность склада", "description": "Общая загрузка складских мощностей."},
            {"slug": "zone-load", "title": "Загруженность зон", "description": "Использование PR, OS, OTG и других зон."},
            {"slug": "cell-load", "title": "Загруженность ячеек", "description": "Свободные и занятые ячейки."},
            {"slug": "free-pallets", "title": "Свободные паллетоместа", "description": "Доступные паллетоместа по складам."},
            {"slug": "boxes", "title": "Состав коробов", "description": "Что находится в каждом коробе: SKU, количество, ЧЗ и место хранения.", "implemented_url": "/head-manager/reports/warehouse/boxes/"},
            {"slug": "pallets", "title": "Состав паллет", "description": "Что находится на каждой паллете: короба, SKU, количество, ЧЗ и место хранения.", "implemented_url": "/head-manager/reports/warehouse/pallets/"},
            {"slug": "warehouse-moves", "title": "Перемещения по складу", "description": "Все движения контейнеров и товаров.", "implemented_url": "/audit/moves/"},
            {"slug": "inventories", "title": "История инвентаризаций", "description": "Проведённые инвентаризации и итоги."},
            {"slug": "inventory-diff", "title": "Расхождения после инвентаризации", "description": "Проблемные позиции после сверки."},
        ],
    },
    {
        "slug": "receiving",
        "title": "Отчёты по приёмке",
        "icon": "IN",
        "tone": "sun",
        "description": "Приёмки, ошибки, расхождения и скорость обработки.",
        "reports": [
            {"slug": "all-receivings", "title": "Все приёмки", "description": "Единый список приёмок за период.", "implemented_url": "/head-manager/requests/?type=receiving"},
            {"slug": "unfinished-receivings", "title": "Незавершённые приёмки", "description": "Приёмки, которые требуют действий."},
            {"slug": "receiving-errors", "title": "Ошибки приёмки", "description": "Ошибки и нарушения процесса приёмки."},
            {"slug": "receiving-discrepancies", "title": "Расхождения", "description": "Факт против плана по приёмке."},
            {"slug": "receiving-speed", "title": "Скорость приёмки", "description": "Среднее время обработки приёмок."},
            {"slug": "receiving-by-employees", "title": "Приёмка по сотрудникам", "description": "Производительность сотрудников на приёмке."},
            {"slug": "receiving-by-clients", "title": "Приёмка по клиентам", "description": "Приёмки в разрезе клиентов."},
            {"slug": "receiving-by-suppliers", "title": "Приёмка по поставщикам", "description": "Статистика по поставщикам."},
        ],
    },
    {
        "slug": "shipping",
        "title": "Отчёты по отгрузке",
        "icon": "OT",
        "tone": "red",
        "description": "Отгрузки, просрочки, ошибки сборки и упаковки.",
        "reports": [
            {"slug": "all-shippings", "title": "Все отгрузки", "description": "Единый список отгрузок за период.", "implemented_url": "/head-manager/requests/?type=shipping"},
            {"slug": "completed-shippings", "title": "Выполненные", "description": "Закрытые отгрузки."},
            {"slug": "late-shippings", "title": "Просроченные", "description": "Отгрузки с нарушением сроков."},
            {"slug": "picking-errors", "title": "Ошибки сборки", "description": "Ошибки при комплектации заказов."},
            {"slug": "packing-errors", "title": "Ошибки упаковки", "description": "Ошибки и возвраты на этапе упаковки."},
            {"slug": "shipping-by-clients", "title": "Отгрузки по клиентам", "description": "Разрез отгрузок по клиентам."},
            {"slug": "shipping-by-employees", "title": "Отгрузки по сотрудникам", "description": "Производительность исполнителей."},
            {"slug": "picking-speed", "title": "Скорость комплектации", "description": "Среднее время комплектации."},
        ],
    },
    {
        "slug": "operations",
        "title": "Отчёты по операциям",
        "icon": "OP",
        "tone": "violet",
        "description": "Размещение, отбор, упаковка, перемещения и списания.",
        "reports": [
            {"slug": "placement", "title": "Размещение", "description": "Операции размещения товара."},
            {"slug": "picking", "title": "Отбор", "description": "Операции отбора и статусы."},
            {"slug": "assembly", "title": "Комплектация", "description": "Комплектация заказов."},
            {"slug": "packing", "title": "Упаковка", "description": "Операции упаковки."},
            {"slug": "moves", "title": "Перемещение", "description": "Перемещения внутри склада.", "implemented_url": "/audit/moves/"},
            {"slug": "returns", "title": "Возвраты", "description": "Возвратные операции."},
            {"slug": "write-offs", "title": "Списание", "description": "Списания товаров и причины."},
            {"slug": "mis-sort", "title": "Пересорт", "description": "Случаи пересорта и исправления."},
            {"slug": "inventory", "title": "Инвентаризация", "description": "Операции инвентаризации."},
        ],
    },
    {
        "slug": "employees",
        "title": "Отчёты по сотрудникам",
        "icon": "HR",
        "tone": "blue",
        "description": "Производительность, KPI, просрочки и рабочее время.",
        "reports": [
            {"slug": "activity", "title": "Карточки сотрудников", "description": "Все складские сотрудники с переходом в ежедневную ленту действий и сканирований.", "implemented_url": "/head-manager/reports/employees/activity/"},
            {"slug": "fbs-productivity", "title": "Операционный отчёт FBS", "description": "Сколько заказов поступило и собрано, состояние потока и результат каждого сотрудника.", "implemented_url": "/head-manager/reports/employees/fbs-productivity/"},
            {"slug": "warehouse-productivity", "title": "Работа склада по сотрудникам", "description": "Приёмка, размещение, перемещения, FBS, отгрузка и инвентаризация по дням.", "implemented_url": "/head-manager/reports/employees/warehouse-productivity/"},
            {"slug": "productivity", "title": "Производительность", "description": "Выполнение операций сотрудниками."},
            {"slug": "kpi", "title": "KPI сотрудников", "description": "Ключевые показатели сотрудников."},
            {"slug": "done-tasks", "title": "Выполненные задания", "description": "Закрытые задачи за период."},
            {"slug": "late-tasks", "title": "Просроченные задания", "description": "Нарушения сроков по исполнителям.", "implemented_url": "/head-manager/requests/?request_filter=overdue"},
            {"slug": "employee-errors", "title": "Ошибки сотрудников", "description": "Ошибки операций по исполнителям."},
            {"slug": "avg-operation-time", "title": "Среднее время операций", "description": "Скорость выполнения операций."},
            {"slug": "employee-load", "title": "Загруженность сотрудников", "description": "Текущая и историческая нагрузка."},
            {"slug": "work-time", "title": "Рабочее время", "description": "Время работы и смены."},
        ],
    },
    {
        "slug": "clients",
        "title": "Отчёты по клиентам",
        "icon": "CL",
        "tone": "green",
        "description": "Остатки, движение, SLA и активность клиентов.",
        "reports": [
            {"slug": "client-stock", "title": "Остатки клиента", "description": "Остатки по выбранному клиенту."},
            {"slug": "client-movement", "title": "Движение клиента", "description": "Операции клиента за период."},
            {"slug": "client-receivings", "title": "Приёмки клиента", "description": "Приёмки конкретного клиента."},
            {"slug": "client-shippings", "title": "Отгрузки клиента", "description": "Отгрузки конкретного клиента."},
            {"slug": "client-returns", "title": "Возвраты клиента", "description": "Возвраты и причины."},
            {"slug": "client-sla", "title": "SLA клиента", "description": "Сроки выполнения по клиенту."},
            {"slug": "client-history", "title": "История операций", "description": "Полная история операций клиента."},
            {"slug": "client-activity", "title": "Активность клиента", "description": "Динамика операций клиента."},
        ],
    },
    {
        "slug": "analytics",
        "title": "Аналитика и контроль",
        "icon": "AN",
        "tone": "red",
        "description": "KPI, SLA, проблемные зоны и аудит действий.",
        "reports": [
            {"slug": "warehouse-kpi", "title": "KPI склада", "description": "Ключевые показатели склада."},
            {"slug": "sla", "title": "SLA", "description": "SLA операций и нарушений.", "implemented_url": "/team-manager/chats/sla/"},
            {"slug": "load", "title": "Загруженность склада", "description": "Использование мощностей склада."},
            {"slug": "late-operations", "title": "Просроченные операции", "description": "Операции с нарушением сроков.", "implemented_url": "/head-manager/requests/?request_filter=overdue"},
            {"slug": "problem-products", "title": "Проблемные товары", "description": "Товары с ошибками, дефицитом или расхождениями."},
            {"slug": "problem-clients", "title": "Проблемные клиенты", "description": "Клиенты с высоким числом проблем."},
            {"slug": "user-audit", "title": "Аудит действий пользователей", "description": "Действия пользователей в системе.", "implemented_url": "/audit/orders/"},
            {"slug": "manual-adjustments", "title": "Ручные корректировки", "description": "Ручные изменения и причины."},
            {"slug": "excess-actions", "title": "Избыточные операции", "description": "Операции сверх обычного процесса.", "implemented_url": "/audit/overactions/"},
        ],
    },
    {
        "slug": "summary",
        "title": "Сводные отчёты",
        "icon": "SUM",
        "tone": "violet",
        "description": "Итоги дня, смены, недели, месяца и экспорт аналитики.",
        "reports": [
            {"slug": "daily", "title": "Итоги за день", "description": "Операционные итоги за день."},
            {"slug": "shift", "title": "Итоги за смену", "description": "Сводка по выбранной смене."},
            {"slug": "weekly", "title": "Итоги за неделю", "description": "Недельная динамика склада."},
            {"slug": "monthly", "title": "Итоги за месяц", "description": "Месячная статистика склада."},
            {"slug": "warehouse-stats", "title": "Общая статистика склада", "description": "Сводная статистика всех операций."},
            {"slug": "analytics-export", "title": "Экспорт аналитики", "description": "Подготовка выгрузок для анализа."},
            {"slug": "management", "title": "Отчёт для руководства", "description": "Управленческая сводка по складу."},
        ],
    },
]


def _head_manager_report_categories() -> list[dict]:
    categories: list[dict] = []
    for category in _HEAD_MANAGER_REPORT_CATEGORIES:
        reports = []
        for report in category["reports"]:
            href = report.get("implemented_url") or f"/head-manager/reports/{category['slug']}/{report['slug']}/"
            reports.append(
                {
                    **report,
                    "href": href,
                    "is_implemented": bool(report.get("implemented_url")),
                    "action_label": "Открыть" if report.get("implemented_url") else "Настроить",
                }
            )
        categories.append(
            {
                **{key: value for key, value in category.items() if key != "reports"},
                "href": f"/head-manager/reports/{category['slug']}/",
                "count": len(reports),
                "reports": reports,
            }
        )
    return categories


def _head_manager_report_category(slug: str) -> dict:
    for category in _head_manager_report_categories():
        if category["slug"] == slug:
            return category
    raise Http404("Категория отчетов не найдена")


def _head_manager_report(category_slug: str, report_slug: str) -> tuple[dict, dict]:
    category = _head_manager_report_category(category_slug)
    for report in category["reports"]:
        if report["slug"] == report_slug:
            return category, report
    raise Http404("Отчет не найден")


def _head_manager_request_dt(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")


def _head_manager_request_date(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y")


def _head_manager_deadline_state(*, status_key: str, deadline, updated_at, assigned: bool) -> tuple[str, str]:
    now = timezone.now()
    today = timezone.localdate()
    status = str(status_key or "").strip().lower()
    if status in {"done", "shipped", "partial_shipped", "closed", "canceled", "cancelled", "rejected"}:
        if updated_at and timezone.localdate(updated_at) == today:
            return "done", "Готово сегодня"
        return "done", "Готово"
    if not assigned:
        return "no_owner", "Без исполнителя"
    if status == "blocked":
        return "soon", "Скоро"
    if status in {"in_progress", "reserved", "storekeeper_accepted", "picking", "packed", "warehouse_accepted"}:
        return "in_work", "В работе"
    if deadline:
        deadline_dt = deadline
        if not hasattr(deadline_dt, "hour"):
            deadline_dt = timezone.make_aware(datetime.combine(deadline_dt, datetime.max.time()))
        if deadline_dt < now:
            delta = now - deadline_dt
            hours = int(delta.total_seconds() // 3600)
            if hours >= 24:
                return "overdue", f"Просрочка {hours // 24} дн."
            return "overdue", f"Просрочка {hours} ч"
        if timezone.localdate(deadline_dt) == today:
            return "today", "Срок сегодня"
    return "all", ""


def _head_manager_row_sort_priority(row: dict) -> tuple[int, object]:
    priority = {
        "overdue": 0,
        "no_owner": 1,
        "today": 2,
        "soon": 3,
        "in_work": 4,
        "done": 6,
    }.get(row.get("stage"), 5)
    value = row.get("deadline_at") or row.get("sort_at") or timezone.make_aware(datetime.min)
    if value and not hasattr(value, "hour"):
        value = timezone.make_aware(datetime.combine(value, datetime.min.time()))
    return priority, value


def _head_manager_ordered_audit_entries_by_order(order_type: str, order_ids: Iterable[str]) -> dict[str, list[OrderAuditEntry]]:
    normalized_ids = sorted({str(order_id or "").strip() for order_id in order_ids if str(order_id or "").strip()})
    if not normalized_ids:
        return {}
    entries_by_order: dict[str, list[OrderAuditEntry]] = {order_id: [] for order_id in normalized_ids}
    entries = (
        OrderAuditEntry.objects.filter(order_type=order_type, order_id__in=normalized_ids)
        .select_related("agency")
        .only("order_id", "payload", "created_at", "id", "agency__agn_name", "agency__short_name")
        .order_by("order_id", "created_at", "id")
    )
    for entry in entries:
        entries_by_order.setdefault(str(entry.order_id), []).append(entry)
    return entries_by_order


def _head_manager_task_display_title(
    task: Task,
    *,
    receiving_entries_by_order: dict[str, list[OrderAuditEntry]] | None = None,
    shipping_orders_by_pk: dict[int, ShippingOrder] | None = None,
) -> str:
    cached = getattr(task, "_display_title_cache", None)
    if cached:
        return cached
    route = str(task.route or "")
    receiving_match = re.search(r"/orders/receiving/([^/]+)/", route)
    if receiving_match and "/act/print" not in route:
        order_id = receiving_match.group(1)
        entries = (receiving_entries_by_order or {}).get(order_id, [])
        display_context = build_receiving_display_context(order_id, entries=entries)
        task._display_title_cache = display_context.get("title") or f"Заявка на приемку №{format_order_number('receiving', order_id)}"
        return task._display_title_cache
    shipping_match = re.search(r"/shipping/(\d+)/", route)
    if shipping_match:
        order = (shipping_orders_by_pk or {}).get(int(shipping_match.group(1)))
        if order is not None:
            task._display_title_cache = _shipping_task_title_with_display_id(
                task.title,
                order.number,
                delivery_type=getattr(order, "delivery_type", ""),
            )
            return task._display_title_cache
    return task.display_title()


def _matches_head_manager_request_filters(row: dict, filters: dict) -> bool:
    selected_responsibility = filters.get("responsibility") or "warehouse"
    if selected_responsibility != "all" and row.get("responsibility") != selected_responsibility:
        return False
    selected_type = filters["type"]
    if selected_type != "all" and row["type"] != selected_type:
        return False
    selected_stage = filters.get("stage") or "all"
    if selected_stage != "all" and row.get("stage") != selected_stage:
        return False
    selected_status = filters["status"]
    if selected_status and selected_status.lower() not in str(row["status_key"]).lower() and selected_status.lower() not in str(row["status"]).lower():
        return False
    client = filters["client"].lower()
    if client and client not in str(row["client"]).lower():
        return False
    employee = str(filters.get("employee") or "").lower()
    if employee and employee not in str(row.get("executor") or "").lower():
        return False
    warehouse = str(filters.get("warehouse_filter") or "").lower()
    if warehouse and warehouse != "all" and warehouse not in str(row.get("warehouse") or "").lower():
        return False
    zone = str(filters.get("zone") or "").lower()
    if zone and zone not in str(row.get("zone") or "").lower():
        return False
    marketplace = str(filters.get("marketplace") or "").lower()
    if marketplace and marketplace not in str(row.get("marketplace") or "").lower():
        return False
    document_status = str(filters.get("document_status") or "").lower()
    if document_status and document_status not in str(row.get("document_status") or "").lower():
        return False
    query = filters["q"].lower()
    if query:
        haystack = " ".join(
            str(row.get(key) or "")
            for key in (
                "number",
                "title",
                "client",
                "status",
                "type_label",
                "responsibility_label",
                "executor",
                "warehouse",
                "zone",
                "marketplace",
                "document_status",
            )
        ).lower()
        if query not in haystack:
            return False
    date_from = filters["date_from"]
    if date_from and row["sort_at"] and timezone.localdate(row["sort_at"]) < date_from:
        return False
    date_to = filters["date_to"]
    if date_to and row["sort_at"] and timezone.localdate(row["sort_at"]) > date_to:
        return False
    return True


def _head_manager_request_filter_context(request, *, default_stage: str = "all") -> dict:
    selected_type = str(request.GET.get("type") or "all").strip()
    if selected_type not in _HEAD_MANAGER_REQUEST_TYPES:
        selected_type = "all"
    selected_stage = str(request.GET.get("request_filter") or request.GET.get("stage") or default_stage).strip()
    if selected_stage not in _HEAD_MANAGER_STAGE_FILTERS:
        selected_stage = "all"
    selected_responsibility = "warehouse"

    def date_param(value: str):
        value = str(value or "").strip()
        if not value:
            return None
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None

    return {
        "type": selected_type,
        "stage": selected_stage,
        "responsibility": selected_responsibility,
        "status": str(request.GET.get("status") or "").strip(),
        "client": str(request.GET.get("client") or "").strip(),
        "employee": str(request.GET.get("employee") or "").strip(),
        "warehouse_filter": str(request.GET.get("warehouse_filter") or "").strip(),
        "zone": str(request.GET.get("zone") or "").strip(),
        "marketplace": str(request.GET.get("marketplace") or "").strip(),
        "document_status": str(request.GET.get("document_status") or "").strip(),
        "q": str(request.GET.get("q") or "").strip(),
        "date_from": date_param(request.GET.get("date_from")),
        "date_to": date_param(request.GET.get("date_to")),
    }


def _head_manager_requests_rows(filters: dict) -> tuple[list[dict], dict, dict, dict]:
    task_status_labels = dict(Task.STATUS_CHOICES)
    rows: list[dict] = []
    shipping_order_pks: set[int] = set()

    task_qs = (
        Task.objects.select_related("assigned_to")
        .filter(
            Q(route__contains="/orders/receiving/")
            | Q(route__contains="/orders/processing/")
            | Q(route__contains="/shipping/")
            | Q(route__contains="/orders/shipping/")
            | Q(route__contains="/orders/other/")
            | Q(route__contains="/logistics/trips/")
            | Q(route__contains="/stockmap/")
        )
        .order_by("-updated_at", "-created_at")[:500]
    )
    tasks = list(task_qs)
    receiving_order_ids: set[str] = set()
    task_shipping_order_pks: set[int] = set()
    for task in tasks:
        route = str(task.route or "")
        receiving_match = re.search(r"/orders/receiving/([^/]+)/", route)
        if receiving_match and "/act/print" not in route:
            receiving_order_ids.add(receiving_match.group(1))
        shipping_match = re.search(r"/shipping/(\d+)/", route)
        if shipping_match:
            task_shipping_order_pks.add(int(shipping_match.group(1)))
    receiving_entries_by_order = _head_manager_ordered_audit_entries_by_order("receiving", receiving_order_ids)
    task_shipping_orders_by_pk = {
        order.pk: order
        for order in ShippingOrder.objects.filter(pk__in=task_shipping_order_pks)
        .select_related("agency", "marketplace")
        .only(
            "pk",
            "number",
            "delivery_type",
            "destination_warehouse",
            "destination_address",
            "transit_address",
            "agency__agn_name",
            "agency__short_name",
            "marketplace__name",
        )
    }

    for task in tasks:
        route = str(task.route or "")
        request_type = _head_manager_request_type(route, task.title)
        responsibility = _head_manager_task_responsibility(task, request_type)
        shipping_match = re.search(r"/shipping/(\d+)/", route)
        if shipping_match:
            shipping_order_pks.add(int(shipping_match.group(1)))
        receiving_match = re.search(r"/orders/receiving/([^/]+)/", route)
        task_meta = {
            "client": "-",
            "warehouse": "-",
            "zone": "-",
            "marketplace": "-",
        }
        if shipping_match:
            task_meta = _head_manager_shipping_meta(task_shipping_orders_by_pk.get(int(shipping_match.group(1))))
        elif receiving_match and "/act/print" not in route:
            task_meta = _head_manager_receiving_meta(
                receiving_entries_by_order.get(receiving_match.group(1), [])
            )
        assigned = task.assigned_to_id is not None
        stage, deadline_label = _head_manager_deadline_state(
            status_key=task.status,
            deadline=task.due_date,
            updated_at=task.updated_at,
            assigned=assigned,
        )
        row = {
            "type": request_type,
            "type_label": _HEAD_MANAGER_REQUEST_TYPES.get(request_type, "Задача"),
            "responsibility": responsibility,
            "responsibility_label": _HEAD_MANAGER_RESPONSIBILITY_FILTERS[responsibility],
            "number": f"#{task.pk}",
            "title": _head_manager_task_display_title(
                task,
                receiving_entries_by_order=receiving_entries_by_order,
                shipping_orders_by_pk=task_shipping_orders_by_pk,
            ),
            "client": task_meta["client"],
            "warehouse": task_meta["warehouse"],
            "zone": task_meta["zone"],
            "marketplace": task_meta["marketplace"],
            "document_status": "По задаче",
            "status": task_status_labels.get(task.status, task.status),
            "status_key": task.status,
            "executor": getattr(task.assigned_to, "full_name", "") or "Без исполнителя",
            "created": _head_manager_request_dt(task.created_at),
            "date": _head_manager_request_dt(task.due_date),
            "deadline": _head_manager_request_dt(task.due_date),
            "deadline_at": task.due_date,
            "deadline_label": deadline_label,
            "stage": stage,
            "stage_label": _HEAD_MANAGER_STAGE_FILTERS.get(stage, ""),
            "updated": _head_manager_request_dt(task.updated_at),
            "href": route or f"/todo/{task.pk}/",
            "sort_at": task.updated_at or task.created_at,
            "comments_count": 0,
            "documents_count": 0,
            "photos_count": 0,
        }
        rows.append(row)

    shipping_status_labels = dict(ShippingOrder.STATUS_CHOICES)
    shipping_qs = (
        ShippingOrder.objects.exclude(pk__in=shipping_order_pks)
        .select_related("agency", "marketplace")
        .order_by("-updated_at", "-created_at")[:300]
    )
    for order in shipping_qs:
        order_meta = _head_manager_shipping_meta(order)
        deadline = order.eta_at or order.planned_ship_date or order.slot_date
        stage, deadline_label = _head_manager_deadline_state(
            status_key=order.status,
            deadline=deadline,
            updated_at=order.updated_at,
            assigned=True,
        )
        row = {
            "type": "shipping",
            "type_label": "Отгрузка",
            "responsibility": "warehouse",
            "responsibility_label": _HEAD_MANAGER_RESPONSIBILITY_FILTERS["warehouse"],
            "number": _head_manager_shipping_number(order.number or order.pk),
            "title": "Заявка на отгрузку",
            "client": order_meta["client"],
            "warehouse": order_meta["warehouse"],
            "zone": order_meta["zone"],
            "marketplace": order_meta["marketplace"],
            "document_status": "По заявке",
            "status": shipping_status_labels.get(order.status, order.status),
            "status_key": order.status,
            "executor": "По заявке",
            "created": _head_manager_request_dt(order.created_at),
            "date": order.planned_ship_date.strftime("%d.%m.%Y") if order.planned_ship_date else "",
            "deadline": _head_manager_request_dt(deadline) if hasattr(deadline, "hour") else (deadline.strftime("%d.%m.%Y") if deadline else ""),
            "deadline_at": deadline,
            "deadline_label": deadline_label,
            "stage": stage,
            "stage_label": _HEAD_MANAGER_STAGE_FILTERS.get(stage, ""),
            "updated": _head_manager_request_dt(order.updated_at),
            "href": f"/shipping/{order.pk}/",
            "sort_at": order.updated_at or order.created_at,
            "comments_count": 0,
            "documents_count": 0,
            "photos_count": 0,
        }
        rows.append(row)

    other_status_labels = dict(OtherRequest.STATUS_CHOICES)
    other_qs = (
        OtherRequest.objects.select_related("agency", "assignee", "category")
        .order_by("-updated_at", "-created_at")[:300]
    )
    for request_obj in other_qs:
        client_label = str(
            getattr(request_obj.agency, "short_name", "") or getattr(request_obj.agency, "agn_name", "") or "-"
        )
        assigned = request_obj.assignee_id is not None
        responsibility = (
            "manager"
            if request_obj.department == OtherRequest.DEPARTMENT_MANAGERS
            or str(getattr(request_obj.assignee, "role", "") or "") == "manager"
            else "warehouse"
        )
        stage, deadline_label = _head_manager_deadline_state(
            status_key=request_obj.status,
            deadline=request_obj.due_at,
            updated_at=request_obj.updated_at,
            assigned=assigned,
        )
        row = {
            "type": "other",
            "type_label": "Прочие",
            "responsibility": responsibility,
            "responsibility_label": _HEAD_MANAGER_RESPONSIBILITY_FILTERS[responsibility],
            "number": request_obj.public_number,
            "title": request_obj.display_title,
            "client": client_label,
            "warehouse": request_obj.department_label,
            "zone": str(getattr(request_obj.category, "title", "") or "-"),
            "marketplace": "-",
            "document_status": "По прочей заявке",
            "status": other_status_labels.get(request_obj.status, request_obj.status),
            "status_key": request_obj.status,
            "executor": getattr(request_obj.assignee, "full_name", "") or "Без исполнителя",
            "created": _head_manager_request_dt(request_obj.created_at),
            "date": _head_manager_request_dt(request_obj.due_at),
            "deadline": _head_manager_request_dt(request_obj.due_at),
            "deadline_at": request_obj.due_at,
            "deadline_label": deadline_label,
            "stage": stage,
            "stage_label": _HEAD_MANAGER_STAGE_FILTERS.get(stage, ""),
            "updated": _head_manager_request_dt(request_obj.updated_at),
            "href": request_obj.wms_url,
            "sort_at": request_obj.updated_at or request_obj.created_at,
            "comments_count": 0,
            "documents_count": 0,
            "photos_count": 0,
        }
        rows.append(row)

    responsibility_counts = {key: 0 for key in _HEAD_MANAGER_RESPONSIBILITY_FILTERS}
    for row in rows:
        responsibility_counts["all"] += 1
        responsibility_counts[row["responsibility"]] = responsibility_counts.get(row["responsibility"], 0) + 1

    selected_responsibility = filters.get("responsibility") or "warehouse"
    responsibility_rows = [
        row
        for row in rows
        if selected_responsibility == "all" or row.get("responsibility") == selected_responsibility
    ]

    stage_counts = {key: 0 for key in _HEAD_MANAGER_STAGE_FILTERS}
    for row in responsibility_rows:
        stage_counts["all"] += 1
        stage_counts[row["stage"]] = stage_counts.get(row["stage"], 0) + 1

    rows = [row for row in rows if _matches_head_manager_request_filters(row, filters)]
    rows.sort(key=_head_manager_row_sort_priority)
    counts = {key: 0 for key in _HEAD_MANAGER_REQUEST_TYPES}
    for row in rows:
        counts["all"] += 1
        counts[row["type"]] = counts.get(row["type"], 0) + 1
    return rows, counts, stage_counts, responsibility_counts


def _movement_report_rows(*, limit: int | None = None) -> list[dict]:
    snapshots_qs = (
        WarehouseStockSnapshot.objects.filter(qty__gt=0)
        .select_related(
            "agency",
            "sku_ref",
            "container",
            "parent_container",
            "location",
            "active_operation",
            "last_event",
        )
        .order_by("created_at", "id")
    )
    if limit:
        snapshots_qs = snapshots_qs[:limit]
    snapshots = list(snapshots_qs)
    container_ids = {
        int(snapshot.container_id or 0)
        for snapshot in snapshots
        if int(snapshot.container_id or 0) > 0
    }
    receiving_event_by_container: dict[int, WarehouseEvent] = {}
    shipping_event_by_container: dict[int, WarehouseEvent] = {}
    if container_ids:
        events = (
            WarehouseEvent.objects.filter(container_id__in=container_ids)
            .filter(
                Q(event_type__in=_MOVEMENT_REPORT_RECEIVING_EVENTS)
                | Q(event_type__in=_MOVEMENT_REPORT_SHIPPING_EVENTS)
            )
            .order_by("occurred_at", "id")
        )
        for event in events:
            container_id = int(event.container_id or 0)
            if event.event_type in _MOVEMENT_REPORT_RECEIVING_EVENTS and container_id not in receiving_event_by_container:
                receiving_event_by_container[container_id] = event
            if event.event_type in _MOVEMENT_REPORT_SHIPPING_EVENTS:
                shipping_event_by_container[container_id] = event

    rows = []
    for index, snapshot in enumerate(snapshots, start=1):
        row = normalize_stock_row_from_snapshot(snapshot) or {}
        container = snapshot.container
        container_id = int(snapshot.container_id or 0)
        state_code = str(snapshot.warehouse_state_code or "").strip().lower()
        receiving_event = receiving_event_by_container.get(container_id)
        shipping_event = shipping_event_by_container.get(container_id)
        width = getattr(container, "width_mm", None)
        height = getattr(container, "height_mm", None)
        depth = getattr(container, "depth_mm", None)
        rows.append(
            {
                "index": index,
                "client": _movement_report_client_label(snapshot, row),
                "pallet_code": row.get("pallet_code") or "",
                "box_code": row.get("box_code") or "",
                "sku": row.get("sku") or "",
                "name": row.get("name") or "",
                "barcode": row.get("barcode") or str(getattr(snapshot.sku_ref, "code", "") or "").strip(),
                "qty": int(row.get("qty") or 0),
                "status": _movement_report_status(snapshot, shipping_event),
                "width": width or "",
                "height": height or "",
                "depth": depth or "",
                "volume": _movement_report_volume_m3(width, height, depth),
                "weight": getattr(container, "gross_weight_g", None) or "",
                "receipt_date": _movement_report_datetime(getattr(receiving_event, "occurred_at", None) or snapshot.created_at),
                "receipt_document": _movement_report_document_number(
                    receiving_event,
                    str(snapshot.source_context_type or "").strip(),
                    str(snapshot.source_context_id or "").strip(),
                ),
                "expense_date": _movement_report_datetime(getattr(shipping_event, "occurred_at", None)),
                "expense_document": (
                    _movement_report_document_number(shipping_event, "", "")
                    if shipping_event is not None
                    else ""
                ),
                "is_processing_source": state_code in _MOVEMENT_REPORT_PROCESSING_SOURCE_STATES,
            }
        )
    return rows


def _movement_report_response() -> HttpResponse:
    rows = _movement_report_rows()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчет"
    sheet.append(_MOVEMENT_REPORT_HEADERS)
    header_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    processing_source_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R1"

    row_keys = [
        "index",
        "client",
        "pallet_code",
        "box_code",
        "sku",
        "name",
        "barcode",
        "qty",
        "status",
        "width",
        "height",
        "depth",
        "volume",
        "weight",
        "receipt_date",
        "receipt_document",
        "expense_date",
        "expense_document",
    ]
    for row in rows:
        sheet.append([row.get(key, "") for key in row_keys])
        if row.get("is_processing_source"):
            cell = sheet.cell(row=sheet.max_row, column=4)
            if str(cell.value or "").strip():
                cell.fill = processing_source_fill

    column_widths = {
        "A": 6,
        "B": 12,
        "C": 22,
        "D": 22,
        "E": 18,
        "F": 38,
        "G": 20,
        "H": 16,
        "I": 22,
        "J": 14,
        "K": 14,
        "L": 14,
        "M": 14,
        "N": 12,
        "O": 18,
        "P": 28,
        "Q": 18,
        "R": 28,
    }
    for column, width_value in column_widths.items():
        sheet.column_dimensions[column].width = width_value

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="head_manager_movement_report_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


def _head_manager_sku_movement_date(value: str):
    value = str(value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _head_manager_sku_movement_filters(request) -> dict:
    return {
        "q": str(request.GET.get("q") or "").strip(),
        "sku": str(request.GET.get("sku") or "").strip(),
        "client": str(request.GET.get("client") or "").strip(),
        "warehouse": str(request.GET.get("warehouse") or "").strip(),
        "zone": str(request.GET.get("zone") or "").strip(),
        "container": str(request.GET.get("container") or "").strip(),
        "order": str(request.GET.get("order") or "").strip(),
        "chz": str(request.GET.get("chz") or "").strip(),
        "date_from": str(request.GET.get("date_from") or "").strip(),
        "date_to": str(request.GET.get("date_to") or "").strip(),
    }


def _head_manager_sku_movement_stock_queryset(filters: dict):
    qs = (
        WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        .select_related("agency", "sku_ref", "container", "parent_container", "location", "last_event")
        .order_by("sku_code", "container__container_code", "id")
    )
    query = str(filters.get("q") or "").strip()
    if query:
        qs = qs.filter(
            Q(sku_code__icontains=query)
            | Q(name__icontains=query)
            | Q(barcode__icontains=query)
            | Q(marking_code__icontains=query)
            | Q(container_code__icontains=query)
            | Q(container__container_code__icontains=query)
            | Q(parent_container__container_code__icontains=query)
            | Q(agency__agn_name__icontains=query)
            | Q(agency__short_name__icontains=query)
        )
    sku = str(filters.get("sku") or "").strip()
    if sku:
        qs = qs.filter(Q(sku_code__icontains=sku) | Q(name__icontains=sku) | Q(barcode__icontains=sku))
    client = str(filters.get("client") or "").strip()
    if client:
        qs = qs.filter(Q(agency__agn_name__icontains=client) | Q(agency__short_name__icontains=client))
    warehouse = str(filters.get("warehouse") or "").strip()
    if warehouse:
        qs = qs.filter(location__warehouse_code__icontains=warehouse)
    zone = str(filters.get("zone") or "").strip()
    if zone:
        qs = qs.filter(Q(zone_code__icontains=zone) | Q(location__zone_code__icontains=zone))
    container = str(filters.get("container") or "").strip()
    if container:
        qs = qs.filter(
            Q(container_code__icontains=container)
            | Q(container__container_code__icontains=container)
            | Q(parent_container__container_code__icontains=container)
        )
    chz = str(filters.get("chz") or "").strip()
    if chz:
        qs = qs.filter(marking_code__icontains=chz)
    order = str(filters.get("order") or "").strip()
    if order:
        qs = qs.filter(Q(source_context_id__icontains=order) | Q(source_context_type__icontains=order))
    return qs


def _head_manager_sku_movement_location(location, zone_code: str = "") -> str:
    parts = []
    warehouse_code = str(getattr(location, "warehouse_code", "") or "").strip()
    zone = str(zone_code or getattr(location, "zone_code", "") or "").strip()
    display_name = str(getattr(location, "display_name", "") or "").strip()
    location_code = str(getattr(location, "location_code", "") or "").strip()
    if warehouse_code:
        parts.append(warehouse_code)
    if zone:
        parts.append(zone)
    if display_name and display_name not in parts:
        parts.append(display_name)
    elif location_code and location_code not in parts:
        parts.append(location_code)
    return " / ".join(parts) or "-"


def _head_manager_sku_movement_document(event: WarehouseEvent | None, fallback_type: str = "", fallback_id: str = "") -> str:
    if event is not None:
        source_type = str(event.source_document_type or event.stock_context_type or "").strip()
        source_id = str(event.source_document_id or event.stock_context_id or "").strip()
    else:
        source_type = str(fallback_type or "").strip()
        source_id = str(fallback_id or "").strip()
    if not source_id:
        return "-"
    return format_order_number(source_type, source_id) if source_type else source_id


def _head_manager_sku_movement_shipping_document(event: WarehouseEvent, reserve=None) -> str:
    candidates = [
        (str(event.stock_context_type or "").strip(), str(event.stock_context_id or "").strip()),
        (str(event.source_document_type or "").strip(), str(event.source_document_id or "").strip()),
    ]
    if reserve is not None:
        candidates.extend(
            [
                (str(getattr(reserve, "context_type", "") or "").strip(), str(getattr(reserve, "context_id", "") or "").strip()),
                (
                    str(getattr(reserve, "source_document_type", "") or "").strip(),
                    str(getattr(reserve, "source_document_id", "") or "").strip(),
                ),
            ]
        )
    for context_type, context_id in candidates:
        if context_id and context_type == "shipping":
            return format_order_number("shipping", context_id)
    return "-"


def _head_manager_sku_movement_current_payload(snapshot: WarehouseStockSnapshot) -> dict:
    normalized = normalize_stock_row_from_snapshot(snapshot) or {}
    current_reserved = (
        int(snapshot.processing_reserved_qty or 0)
        + int(snapshot.shipping_reserved_qty or 0)
        + int(snapshot.other_reserved_qty or 0)
    )
    return {
        "snapshot_id": int(snapshot.id or 0),
        "client": _movement_report_client_label(snapshot, normalized) or "-",
        "sku": normalized.get("sku") or snapshot.sku_code or "-",
        "name": normalized.get("name") or snapshot.name or "-",
        "barcode": normalized.get("barcode") or snapshot.barcode or str(getattr(snapshot.sku_ref, "code", "") or "").strip() or "-",
        "marking_code": str(snapshot.marking_code or "").strip(),
        "box_code": normalized.get("box_code") or snapshot.container_code or "-",
        "pallet_code": normalized.get("pallet_code") or str(getattr(snapshot.parent_container, "container_code", "") or "").strip() or "-",
        "current_qty": int(snapshot.qty or 0),
        "available_qty": int(snapshot.available_qty or 0),
        "reserved_qty": current_reserved,
        "warehouse": str(getattr(snapshot.location, "warehouse_code", "") or normalized.get("warehouse") or "").strip(),
        "zone": normalized.get("zone") or snapshot.zone_code or "",
        "location": _head_manager_sku_movement_location(snapshot.location, snapshot.zone_code),
    }


def _head_manager_sku_movement_payload_from_event(event: WarehouseEvent) -> dict:
    payload = event.payload if isinstance(event.payload, dict) else {}
    reserve = event.reserve
    container = event.container
    parent_container = getattr(container, "parent_container", None)
    return {
        "snapshot_id": int(payload.get("snapshot_id") or 0),
        "client": str(getattr(event.agency, "short_name", "") or getattr(event.agency, "agn_name", "") or "").strip() or "-",
        "sku": str(
            payload.get("sku_code")
            or payload.get("sku")
            or getattr(reserve, "sku_code", "")
            or ""
        ).strip() or "-",
        "name": str(payload.get("name") or "").strip() or "-",
        "barcode": str(payload.get("barcode") or getattr(reserve, "barcode", "") or "").strip() or "-",
        "marking_code": str(payload.get("marking_code") or getattr(reserve, "marking_code", "") or "").strip(),
        "box_code": str(
            payload.get("box_code")
            or getattr(container, "container_code", "")
            or ""
        ).strip() or "-",
        "pallet_code": str(getattr(parent_container, "container_code", "") or "").strip() or "-",
        "current_qty": 0,
        "available_qty": 0,
        "reserved_qty": 0,
        "warehouse": "",
        "zone": "",
        "location": "-",
    }


def _head_manager_sku_movement_matches_filters(row: dict, filters: dict) -> bool:
    haystack = " ".join(
        str(row.get(key) or "")
        for key in [
            "client",
            "sku",
            "name",
            "barcode",
            "marking_code",
            "box_code",
            "pallet_code",
            "document",
            "shipping_document",
            "from_location",
            "to_location",
            "operation",
        ]
    ).lower()
    query = str(filters.get("q") or "").strip().lower()
    if query and query not in haystack:
        return False
    for key in ["sku", "client", "container", "chz", "order"]:
        value = str(filters.get(key) or "").strip().lower()
        if not value:
            continue
        if key == "container":
            target = f"{row.get('box_code', '')} {row.get('pallet_code', '')}".lower()
        elif key == "chz":
            target = str(row.get("marking_code") or "").lower()
        elif key == "order":
            target = f"{row.get('document', '')} {row.get('shipping_document', '')}".lower()
        else:
            target = str(row.get(key) or "").lower()
            if key == "sku":
                target = f"{target} {row.get('name', '')} {row.get('barcode', '')}".lower()
        if value not in target:
            return False
    warehouse = str(filters.get("warehouse") or "").strip().lower()
    if warehouse and warehouse not in str(row.get("warehouse") or row.get("from_location") or row.get("to_location") or "").lower():
        return False
    zone = str(filters.get("zone") or "").strip().lower()
    if zone and zone not in str(row.get("zone") or row.get("from_location") or row.get("to_location") or "").lower():
        return False
    return True


def _head_manager_receiving_order_search_terms(value: str) -> set[str]:
    raw = str(value or "").strip()
    if not raw:
        return set()
    terms = {raw}
    sequence = order_number_sequence("receiving", raw)
    if sequence is None:
        suffix_match = re.fullmatch(r"0*(\d+)[_-]?PR", raw, re.IGNORECASE)
        if suffix_match:
            sequence = int(suffix_match.group(1))
    if sequence is not None:
        terms.update({str(sequence), f"PR-{sequence:06d}", f"{sequence}_PR"})
    return terms


def _head_manager_open_receiving_cz_order_ids(order_ids: Iterable[str]) -> set[str]:
    entries_by_order = _head_manager_ordered_audit_entries_by_order("receiving", order_ids)
    open_order_ids: set[str] = set()
    for order_id, entries in entries_by_order.items():
        if not entries:
            continue
        is_open = True
        for entry in reversed(entries):
            payload = entry.payload if isinstance(entry.payload, dict) else {}
            if payload.get("flow_reopened"):
                break
            if payload.get("flow_closed") or (
                payload.get("act") == "receiving" and payload.get("act_state") == "closed"
            ):
                is_open = False
                break
        if is_open:
            open_order_ids.add(order_id)
    return open_order_ids


def _head_manager_sku_movement_pending_receiving_rows(
    filters: dict,
    *,
    limit: int | None = None,
) -> tuple[list[dict], int]:
    units_qs = ReceivingCzUnit.objects.select_related("agency", "sku").order_by("-accepted_at", "-id")

    warehouse = str(filters.get("warehouse") or "").strip().lower()
    if warehouse and warehouse not in "msk":
        return [], 0
    zone = str(filters.get("zone") or "").strip().lower()
    if zone and zone not in "pr":
        return [], 0

    date_from = _head_manager_sku_movement_date(filters.get("date_from"))
    if date_from:
        units_qs = units_qs.filter(accepted_at__date__gte=date_from)
    date_to = _head_manager_sku_movement_date(filters.get("date_to"))
    if date_to:
        units_qs = units_qs.filter(accepted_at__date__lte=date_to)

    query = str(filters.get("q") or "").strip()
    if query:
        static_haystack = "msk pr зона приёмки чз отсканирован приёмка продолжается"
        if query.lower() not in static_haystack:
            query_filter = (
                Q(sku_code__icontains=query)
                | Q(name__icontains=query)
                | Q(barcode__icontains=query)
                | Q(marking_code__icontains=query)
                | Q(box_code__icontains=query)
                | Q(pallet_code__icontains=query)
                | Q(source_shipping_order_number__icontains=query)
                | Q(agency__agn_name__icontains=query)
                | Q(agency__short_name__icontains=query)
            )
            for order_term in _head_manager_receiving_order_search_terms(query):
                query_filter |= Q(order_id__icontains=order_term)
            units_qs = units_qs.filter(query_filter)

    sku = str(filters.get("sku") or "").strip()
    if sku:
        units_qs = units_qs.filter(Q(sku_code__icontains=sku) | Q(name__icontains=sku) | Q(barcode__icontains=sku))
    client = str(filters.get("client") or "").strip()
    if client:
        units_qs = units_qs.filter(Q(agency__agn_name__icontains=client) | Q(agency__short_name__icontains=client))
    container = str(filters.get("container") or "").strip()
    if container:
        units_qs = units_qs.filter(Q(box_code__icontains=container) | Q(pallet_code__icontains=container))
    chz = str(filters.get("chz") or "").strip()
    if chz:
        units_qs = units_qs.filter(marking_code__icontains=chz)
    order = str(filters.get("order") or "").strip()
    if order:
        order_filter = Q(source_shipping_order_number__icontains=order)
        for order_term in _head_manager_receiving_order_search_terms(order):
            order_filter |= Q(order_id__icontains=order_term)
        units_qs = units_qs.filter(order_filter)

    candidate_order_ids = list(units_qs.order_by().values_list("order_id", flat=True).distinct())
    open_order_ids = _head_manager_open_receiving_cz_order_ids(candidate_order_ids)
    if not open_order_ids:
        return [], 0

    units_qs = units_qs.filter(order_id__in=open_order_ids)
    pending_count = units_qs.count()
    if limit is not None:
        units_qs = units_qs[:limit]

    rows: list[dict] = []
    for unit in units_qs:
        agency = unit.agency
        shipping_document = str(unit.source_shipping_order_number or "").strip()
        row = {
            "snapshot_id": 0,
            "client": str(getattr(agency, "short_name", "") or getattr(agency, "agn_name", "") or "").strip() or "-",
            "sku": str(unit.sku_code or getattr(unit.sku, "sku_code", "") or "").strip() or "-",
            "name": str(unit.name or getattr(unit.sku, "name", "") or "").strip() or "-",
            "barcode": str(unit.barcode or "").strip() or "-",
            "marking_code": str(unit.marking_code or "").strip(),
            "box_code": str(unit.box_code or "").strip() or "-",
            "pallet_code": str(unit.pallet_code or "").strip() or "-",
            "current_qty": 0,
            "available_qty": 0,
            "reserved_qty": 0,
            "warehouse": "MSK",
            "zone": "PR",
            "location": "MSK / PR / Зона приёмки",
            "occurred_at": unit.accepted_at,
            "date": _movement_report_datetime(unit.accepted_at),
            "operation": "ЧЗ отсканирован — приёмка продолжается",
            "direction": "Приёмка",
            "qty": 1,
            "from_location": "-",
            "to_location": "MSK / PR / Зона приёмки (не завершена)",
            "document": _head_manager_sku_movement_document(None, "receiving", unit.order_id),
            "shipping_document": format_order_number("shipping", shipping_document) if shipping_document else "-",
            "pending_receiving": True,
            "_sort_id": int(unit.id or 0),
        }
        if _head_manager_sku_movement_matches_filters(row, filters):
            rows.append(row)
    return rows, pending_count


def _head_manager_sku_movement_rows(filters: dict, *, limit: int | None = 300) -> tuple[list[dict], dict]:
    stock_qs = _head_manager_sku_movement_stock_queryset(filters)
    stock_summary = stock_qs.aggregate(
        total_qty=Sum("qty"),
        available_qty=Sum("available_qty"),
        processing_reserved_qty=Sum("processing_reserved_qty"),
        shipping_reserved_qty=Sum("shipping_reserved_qty"),
        other_reserved_qty=Sum("other_reserved_qty"),
        boxes=Count("container_id", distinct=True),
        pallets=Count("parent_container_id", distinct=True),
        sku_count=Count("sku_code", distinct=True),
    )
    current_snapshots = list(stock_qs[:5000])
    current_by_snapshot_id = {
        int(snapshot.id): _head_manager_sku_movement_current_payload(snapshot)
        for snapshot in current_snapshots
    }
    current_by_container: dict[int, list[dict]] = {}
    for snapshot in current_snapshots:
        container_id = int(snapshot.container_id or 0)
        if container_id:
            current_by_container.setdefault(container_id, []).append(current_by_snapshot_id[int(snapshot.id)])

    events_qs = (
        WarehouseEvent.objects.select_related(
            "agency",
            "container",
            "container__parent_container",
            "reserve",
            "from_location",
            "to_location",
        )
        .order_by("-occurred_at", "-id")
    )
    date_from = _head_manager_sku_movement_date(filters.get("date_from"))
    if date_from:
        events_qs = events_qs.filter(occurred_at__date__gte=date_from)
    date_to = _head_manager_sku_movement_date(filters.get("date_to"))
    if date_to:
        events_qs = events_qs.filter(occurred_at__date__lte=date_to)
    client = str(filters.get("client") or "").strip()
    if client:
        events_qs = events_qs.filter(Q(agency__agn_name__icontains=client) | Q(agency__short_name__icontains=client))
    order = str(filters.get("order") or "").strip()
    if order:
        events_qs = events_qs.filter(
            Q(source_document_id__icontains=order)
            | Q(stock_context_id__icontains=order)
            | Q(reserve__context_id__icontains=order)
            | Q(reserve__source_document_id__icontains=order)
        )

    event_filter = Q()
    sku = str(filters.get("sku") or "").strip()
    if sku:
        event_filter |= (
            Q(reserve__sku_code__icontains=sku)
            | Q(reserve__barcode__icontains=sku)
            | Q(payload__sku_code__icontains=sku)
            | Q(payload__sku__icontains=sku)
            | Q(payload__barcode__icontains=sku)
        )
    container = str(filters.get("container") or "").strip()
    if container:
        event_filter |= Q(container__container_code__icontains=container) | Q(container__parent_container__container_code__icontains=container)
    current_container_ids = [container_id for container_id in current_by_container]
    if (sku or filters.get("q") or filters.get("warehouse") or filters.get("zone") or filters.get("chz")) and current_container_ids:
        event_filter |= Q(container_id__in=current_container_ids)
    if event_filter:
        events_qs = events_qs.filter(event_filter)

    scan_limit = None if limit is None else max(limit * 8, 1200)
    if scan_limit:
        events_qs = events_qs[:scan_limit]

    event_rows: list[dict] = []
    for event in events_qs:
        payload = event.payload if isinstance(event.payload, dict) else {}
        snapshot_id = int(payload.get("snapshot_id") or 0)
        candidates: list[dict] = []
        if snapshot_id and snapshot_id in current_by_snapshot_id:
            candidates.append(current_by_snapshot_id[snapshot_id])
        if event.reserve_id:
            candidates.append(_head_manager_sku_movement_payload_from_event(event))
        container_id = int(event.container_id or 0)
        if container_id in current_by_container:
            candidates.extend(current_by_container[container_id])
        if not candidates and (payload.get("sku_code") or payload.get("sku") or payload.get("barcode")):
            candidates.append(_head_manager_sku_movement_payload_from_event(event))

        seen_keys = set()
        for candidate in candidates:
            row_key = (
                candidate.get("snapshot_id"),
                candidate.get("sku"),
                candidate.get("barcode"),
                candidate.get("marking_code"),
                candidate.get("box_code"),
            )
            if row_key in seen_keys:
                continue
            seen_keys.add(row_key)
            event_type = str(event.event_type or "").strip()
            direction = "Движение"
            if event_type in _SKU_MOVEMENT_IN_EVENTS:
                direction = "Приход"
            elif event_type in _SKU_MOVEMENT_OUT_EVENTS:
                direction = "Расход"
            qty = int(event.qty or 0)
            row = {
                **candidate,
                "occurred_at": event.occurred_at,
                "date": _movement_report_datetime(event.occurred_at),
                "operation": _SKU_MOVEMENT_EVENT_LABELS.get(event_type, event_type or "-"),
                "direction": direction,
                "qty": qty,
                "from_location": _head_manager_sku_movement_location(event.from_location, event.from_zone_code),
                "to_location": _head_manager_sku_movement_location(event.to_location, event.to_zone_code),
                "document": _head_manager_sku_movement_document(event),
                "shipping_document": _head_manager_sku_movement_shipping_document(event, event.reserve),
            }
            if not _head_manager_sku_movement_matches_filters(row, filters):
                continue
            row["_sort_id"] = int(event.id or 0)
            event_rows.append(row)
            if limit is not None and len(event_rows) >= limit:
                break
        if limit is not None and len(event_rows) >= limit:
            break

    pending_rows, pending_receiving_count = _head_manager_sku_movement_pending_receiving_rows(
        filters,
        limit=limit,
    )
    rows = sorted(
        [*event_rows, *pending_rows],
        key=lambda row: (row.get("occurred_at") or timezone.make_aware(datetime.min), int(row.get("_sort_id") or 0)),
        reverse=True,
    )
    if limit is not None:
        rows = rows[:limit]

    inbound_qty = sum(int(row.get("qty") or 0) for row in event_rows if row.get("direction") == "Приход")
    outbound_qty = sum(int(row.get("qty") or 0) for row in event_rows if row.get("direction") == "Расход")
    shipping_documents = {
        str(row.get("shipping_document"))
        for row in event_rows
        if str(row.get("shipping_document") or "-") != "-"
    }

    reserved_qty = (
        int(stock_summary.get("processing_reserved_qty") or 0)
        + int(stock_summary.get("shipping_reserved_qty") or 0)
        + int(stock_summary.get("other_reserved_qty") or 0)
    )
    summary = {
        "current_qty": int(stock_summary.get("total_qty") or 0),
        "available_qty": int(stock_summary.get("available_qty") or 0),
        "reserved_qty": reserved_qty,
        "boxes": int(stock_summary.get("boxes") or 0),
        "pallets": int(stock_summary.get("pallets") or 0),
        "sku_count": int(stock_summary.get("sku_count") or 0),
        "inbound_qty": inbound_qty,
        "outbound_qty": outbound_qty,
        "shipping_documents": len(shipping_documents),
        "pending_receiving_count": pending_receiving_count,
        "shown": len(rows),
        "has_filters": any(str(value or "").strip() for value in filters.values()),
    }
    return rows, summary


def _head_manager_sku_movement_workbook(rows: list[dict], summary: dict) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Движение по SKU"
    sheet.append(_SKU_MOVEMENT_REPORT_HEADERS)
    header_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    sheet.freeze_panes = "A2"
    row_keys = [
        "date",
        "operation",
        "direction",
        "client",
        "sku",
        "name",
        "barcode",
        "marking_code",
        "qty",
        "from_location",
        "to_location",
        "box_code",
        "pallet_code",
        "document",
        "shipping_document",
        "current_qty",
        "available_qty",
        "reserved_qty",
    ]
    for row in rows:
        sheet.append([_excel_safe_value(row.get(key, "")) for key in row_keys])
    sheet.auto_filter.ref = f"A1:R{max(sheet.max_row, 1)}"
    for index, width in enumerate([18, 26, 14, 28, 22, 38, 22, 64, 12, 32, 32, 24, 24, 24, 24, 16, 14, 14], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Текущий остаток", summary.get("current_qty", 0)])
    summary_sheet.append(["Доступно", summary.get("available_qty", 0)])
    summary_sheet.append(["В резерве", summary.get("reserved_qty", 0)])
    summary_sheet.append(["SKU", summary.get("sku_count", 0)])
    summary_sheet.append(["Коробов", summary.get("boxes", 0)])
    summary_sheet.append(["Паллет", summary.get("pallets", 0)])
    summary_sheet.append(["ЧЗ в незавершённой приёмке", summary.get("pending_receiving_count", 0)])
    summary_sheet.append(["Приход по событиям", summary.get("inbound_qty", 0)])
    summary_sheet.append(["Расход по событиям", summary.get("outbound_qty", 0)])
    summary_sheet.append(["Отгрузок", summary.get("shipping_documents", 0)])
    summary_sheet.append(["Выгружено строк", len(rows)])
    if summary.get("export_truncated"):
        summary_sheet.append(
            [
                "Ограничение выгрузки",
                f"Показаны первые {_SKU_MOVEMENT_EXPORT_ROW_LIMIT} строк. Уточните период или фильтры.",
            ]
        )
    for cell in summary_sheet["A"]:
        cell.font = header_font
    summary_sheet.column_dimensions["A"].width = 26
    summary_sheet.column_dimensions["B"].width = 18
    return workbook


def _head_manager_sku_movement_export_response(filters: dict) -> HttpResponse:
    export_filters = dict(filters or {})
    if not export_filters.get("date_from") and not export_filters.get("date_to"):
        date_to = timezone.localdate()
        export_filters["date_to"] = date_to.isoformat()
        export_filters["date_from"] = (
            date_to - timedelta(days=_SKU_MOVEMENT_EXPORT_DEFAULT_DAYS - 1)
        ).isoformat()
    rows, summary = _head_manager_sku_movement_rows(
        export_filters,
        limit=_SKU_MOVEMENT_EXPORT_ROW_LIMIT,
    )
    summary = dict(summary or {})
    summary["export_truncated"] = len(rows) >= _SKU_MOVEMENT_EXPORT_ROW_LIMIT
    workbook = _head_manager_sku_movement_workbook(rows, summary)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="head_manager_sku_movement_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


def _head_manager_fbs_employee_report_date(value: str):
    value = str(value or "").strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _head_manager_fbs_employee_report_filters(request) -> dict:
    today = timezone.localdate()
    default_from = today
    date_from = _head_manager_fbs_employee_report_date(request.GET.get("date_from"))
    date_to = _head_manager_fbs_employee_report_date(request.GET.get("date_to"))
    if date_from is None:
        date_from = default_from
    if date_to is None:
        date_to = today
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    return {
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "employee": str(request.GET.get("employee") or "").strip(),
    }


def _head_manager_fbs_employee_report_period(date_from, date_to):
    current_timezone = timezone.get_current_timezone()
    start_at = timezone.make_aware(datetime.combine(date_from, datetime.min.time()), current_timezone)
    end_at = timezone.make_aware(
        datetime.combine(date_to + timedelta(days=1), datetime.min.time()),
        current_timezone,
    )
    return start_at, end_at


def _head_manager_fbs_employee_user_filter(field: str, employee: str) -> Q:
    if not employee:
        return Q()
    return (
        Q(**{f"{field}__employee_profile__full_name__icontains": employee})
        | Q(**{f"{field}__username__icontains": employee})
        | Q(**{f"{field}__first_name__icontains": employee})
        | Q(**{f"{field}__last_name__icontains": employee})
    )


def _head_manager_fbs_employee_identity(user) -> tuple[str, str]:
    if user is None:
        return "Не указан", "-"
    employee = getattr(user, "employee_profile", None)
    if employee is not None:
        name = str(employee.full_name or "").strip()
        role = str(employee.get_role_display() or "").strip()
        if name:
            return name, role or "-"
    name = str(user.get_full_name() or "").strip() or str(user.get_username() or "").strip()
    return name or "Не указан", "Пользователь"


def _head_manager_fbs_employee_report_rows(filters: dict) -> tuple[list[dict], dict]:
    date_from = _head_manager_fbs_employee_report_date(filters.get("date_from"))
    date_to = _head_manager_fbs_employee_report_date(filters.get("date_to"))
    period_start, period_end = _head_manager_fbs_employee_report_period(date_from, date_to)
    employee_filter = str(filters.get("employee") or "").strip()
    rows_by_key: dict[tuple, dict] = {}

    def report_row(user, occurred_at) -> dict:
        local_day = timezone.localtime(occurred_at).date()
        user_id = int(getattr(user, "pk", 0) or 0)
        key = (local_day, user_id)
        if key not in rows_by_key:
            employee_name, role = _head_manager_fbs_employee_identity(user)
            rows_by_key[key] = {
                "user_id": user_id,
                "date_value": local_day,
                "date": local_day.strftime("%d.%m.%Y"),
                "employee": employee_name,
                "role": role,
                "pick_waves": 0,
                "pick_orders": 0,
                "pick_units": 0,
                "pick_duration_total": 0.0,
                "pick_duration_count": 0,
                "avg_pick_minutes": 0,
                "verified_waves": 0,
                "verified_orders": 0,
                "verified_units": 0,
                "shipments": 0,
                "shipment_boxes": 0,
                "shipment_orders": 0,
                "shipment_units": 0,
                "_pick_wave_ids": set(),
                "_pick_order_ids": set(),
            }
        return rows_by_key[key]

    pick_tasks = (
        FbsPickTask.objects.filter(
            status=FbsPickTask.STATUS_PICKED,
            assigned_to__isnull=False,
            completed_at__isnull=False,
            completed_at__gte=period_start,
            completed_at__lt=period_end,
        )
        .filter(_head_manager_fbs_employee_user_filter("assigned_to", employee_filter))
        .select_related("assigned_to__employee_profile", "batch")
    )
    for task in pick_tasks:
        row = report_row(task.assigned_to, task.completed_at)
        row["_pick_wave_ids"].add(task.batch_id)
        row["_pick_order_ids"].add(task.order_id)
        row["pick_units"] += int(task.picked_qty or 0)
        started_at = task.claimed_at or task.batch.started_at or task.batch.claimed_at
        if started_at and task.completed_at >= started_at:
            row["pick_duration_total"] += (
                task.completed_at - started_at
            ).total_seconds() / 60
            row["pick_duration_count"] += 1

    verified_batches = (
        FbsPickBatch.objects.filter(
            completed_at__isnull=False,
            verification_assigned_to__isnull=False,
            completed_at__gte=period_start,
            completed_at__lt=period_end,
        )
        .filter(
            _head_manager_fbs_employee_user_filter(
                "verification_assigned_to", employee_filter
            )
        )
        .select_related("verification_assigned_to__employee_profile")
        .annotate(
            report_order_count=Count(
                "tasks__order_id",
                filter=Q(tasks__status=FbsPickTask.STATUS_PICKED),
                distinct=True,
            ),
            report_unit_count=Sum(
                "tasks__picked_qty",
                filter=Q(tasks__status=FbsPickTask.STATUS_PICKED),
            ),
        )
    )
    for batch in verified_batches:
        row = report_row(batch.verification_assigned_to, batch.completed_at)
        row["verified_waves"] += 1
        row["verified_orders"] += int(batch.report_order_count or 0)
        row["verified_units"] += int(batch.report_unit_count or 0)

    handover_batches = (
        FbsHandoverBatch.objects.filter(
            dispatched_at__isnull=False,
            dispatched_by__isnull=False,
            dispatched_at__gte=period_start,
            dispatched_at__lt=period_end,
        )
        .filter(_head_manager_fbs_employee_user_filter("dispatched_by", employee_filter))
        .select_related("dispatched_by__employee_profile")
        .annotate(
            report_box_count=Count("boxes", distinct=True),
            report_order_count=Count("boxes__orders", distinct=True),
            report_unit_count=Sum("boxes__orders__order__items__quantity"),
        )
    )
    for batch in handover_batches:
        row = report_row(batch.dispatched_by, batch.dispatched_at)
        row["shipments"] += 1
        row["shipment_boxes"] += int(batch.report_box_count or 0)
        row["shipment_orders"] += int(batch.report_order_count or 0)
        row["shipment_units"] += int(batch.report_unit_count or 0)

    rows = list(rows_by_key.values())
    for row in rows:
        row["pick_waves"] = len(row.pop("_pick_wave_ids"))
        row["pick_orders"] = len(row.pop("_pick_order_ids"))
        duration_count = int(row.pop("pick_duration_count") or 0)
        duration_total = float(row.pop("pick_duration_total") or 0)
        row["avg_pick_minutes"] = round(duration_total / duration_count, 1) if duration_count else 0
    rows.sort(key=lambda row: (-row["date_value"].toordinal(), row["employee"].lower()))
    summary = {
        "employees": len({row["user_id"] for row in rows}),
        "days": len({row["date_value"] for row in rows}),
        "pick_waves": sum(row["pick_waves"] for row in rows),
        "pick_orders": sum(row["pick_orders"] for row in rows),
        "pick_units": sum(row["pick_units"] for row in rows),
        "verified_waves": sum(row["verified_waves"] for row in rows),
        "verified_orders": sum(row["verified_orders"] for row in rows),
        "verified_units": sum(row["verified_units"] for row in rows),
        "shipments": sum(row["shipments"] for row in rows),
        "shipment_orders": sum(row["shipment_orders"] for row in rows),
        "shipment_units": sum(row["shipment_units"] for row in rows),
    }
    return rows, summary


def _head_manager_fbs_employee_report_overview(filters: dict) -> dict:
    date_from = _head_manager_fbs_employee_report_date(filters.get("date_from"))
    date_to = _head_manager_fbs_employee_report_date(filters.get("date_to"))
    period_start, period_end = _head_manager_fbs_employee_report_period(date_from, date_to)

    received_orders = FbsOrder.objects.filter(
        imported_at__gte=period_start,
        imported_at__lt=period_end,
    )
    received_totals = received_orders.aggregate(
        orders=Count("id", distinct=True),
        units=Sum("items__quantity"),
        clients=Count("profile__agency_id", distinct=True),
    )
    status_counts = {
        row["internal_status"]: int(row["orders"] or 0)
        for row in received_orders.values("internal_status")
        .annotate(orders=Count("id"))
        .order_by()
    }

    pick_totals = FbsPickTask.objects.filter(
        status=FbsPickTask.STATUS_PICKED,
        completed_at__isnull=False,
        completed_at__gte=period_start,
        completed_at__lt=period_end,
    ).aggregate(
        waves=Count("batch_id", distinct=True),
        orders=Count("order_id", distinct=True),
        units=Sum("picked_qty"),
        employees=Count("assigned_to_id", distinct=True),
    )

    picked_task_filter = Q(tasks__status=FbsPickTask.STATUS_PICKED)
    verification_totals = FbsPickBatch.objects.filter(
        completed_at__isnull=False,
        completed_at__gte=period_start,
        completed_at__lt=period_end,
    ).exclude(status=FbsPickBatch.STATUS_CANCELED).aggregate(
        waves=Count("id", distinct=True),
        orders=Count(
            "tasks__order_id",
            filter=picked_task_filter,
            distinct=True,
        ),
        units=Sum("tasks__picked_qty", filter=picked_task_filter),
        employees=Count("verification_assigned_to_id", distinct=True),
    )

    handover_totals = FbsHandoverBatch.objects.filter(
        dispatched_at__isnull=False,
        dispatched_at__gte=period_start,
        dispatched_at__lt=period_end,
    ).aggregate(
        shipment_count=Count("id", distinct=True),
        box_count=Count("boxes", distinct=True),
        order_count=Count("boxes__orders__order_id", distinct=True),
        unit_count=Sum("boxes__orders__order__items__quantity"),
        employee_count=Count("dispatched_by_id", distinct=True),
    )

    status_groups = (
        (
            "Ожидают подготовки",
            (
                FbsOrder.STATUS_RECEIVED,
                FbsOrder.STATUS_VALIDATION_FAILED,
                FbsOrder.STATUS_AWAITING_STOCK,
            ),
        ),
        (
            "Готовы / в очереди",
            (FbsOrder.STATUS_RESERVED, FbsOrder.STATUS_QUEUED_FOR_PICK),
        ),
        ("Собираются", (FbsOrder.STATUS_PICKING,)),
        (
            "Собраны / готовы к передаче",
            (FbsOrder.STATUS_PICKED, FbsOrder.STATUS_READY_FOR_HANDOVER),
        ),
        (
            "Переданы / доставлены",
            (FbsOrder.STATUS_HANDED_OVER, FbsOrder.STATUS_DELIVERED),
        ),
        (
            "Отменены / возврат",
            (
                FbsOrder.STATUS_CANCELLED,
                FbsOrder.STATUS_RETURN_PENDING,
                FbsOrder.STATUS_RETURNED,
            ),
        ),
        ("Проблема", (FbsOrder.STATUS_EXCEPTION,)),
    )
    order_status_rows = [
        {
            "title": title,
            "orders": sum(status_counts.get(status, 0) for status in statuses),
        }
        for title, statuses in status_groups
    ]
    unfinished_statuses = {
        FbsOrder.STATUS_RECEIVED,
        FbsOrder.STATUS_VALIDATION_FAILED,
        FbsOrder.STATUS_AWAITING_STOCK,
        FbsOrder.STATUS_RESERVED,
        FbsOrder.STATUS_QUEUED_FOR_PICK,
        FbsOrder.STATUS_PICKING,
        FbsOrder.STATUS_EXCEPTION,
    }
    unfinished_orders = sum(
        status_counts.get(status, 0) for status in unfinished_statuses
    )

    section_rows = [
        {
            "title": "Поступило",
            "description": "Заказы, загруженные в FBS за выбранный период.",
            "orders": int(received_totals["orders"] or 0),
            "units": int(received_totals["units"] or 0),
            "operations": int(received_totals["clients"] or 0),
            "operations_label": "клиентов",
            "employees": None,
        },
        {
            "title": "Сборка",
            "description": "Фактически завершённые задания отбора за период.",
            "orders": int(pick_totals["orders"] or 0),
            "units": int(pick_totals["units"] or 0),
            "operations": int(pick_totals["waves"] or 0),
            "operations_label": "волн",
            "employees": int(pick_totals["employees"] or 0),
        },
        {
            "title": "Проверка",
            "description": "Волны, у которых проверка завершена за период.",
            "orders": int(verification_totals["orders"] or 0),
            "units": int(verification_totals["units"] or 0),
            "operations": int(verification_totals["waves"] or 0),
            "operations_label": "волн",
            "employees": int(verification_totals["employees"] or 0),
        },
        {
            "title": "Передача",
            "description": "Отгрузки, переданные водителю за период.",
            "orders": int(handover_totals["order_count"] or 0),
            "units": int(handover_totals["unit_count"] or 0),
            "operations": int(handover_totals["shipment_count"] or 0),
            "operations_label": "отгрузок",
            "employees": int(handover_totals["employee_count"] or 0),
            "boxes": int(handover_totals["box_count"] or 0),
        },
    ]

    return {
        "orders_received": int(received_totals["orders"] or 0),
        "received_units": int(received_totals["units"] or 0),
        "orders_picked": int(pick_totals["orders"] or 0),
        "picked_units": int(pick_totals["units"] or 0),
        "pickers": int(pick_totals["employees"] or 0),
        "unfinished_orders": unfinished_orders,
        "verified_orders": int(verification_totals["orders"] or 0),
        "handover_orders": int(handover_totals["order_count"] or 0),
        "section_rows": section_rows,
        "order_status_rows": order_status_rows,
    }


def _head_manager_fbs_employee_report_workbook(
    rows: list[dict],
    summary: dict,
    overview: dict | None = None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FBS по сотрудникам"
    sheet.append(_FBS_EMPLOYEE_REPORT_HEADERS)
    header_fill = PatternFill(fill_type="solid", fgColor="F8B800")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    row_keys = [
        "date",
        "employee",
        "role",
        "pick_waves",
        "pick_orders",
        "pick_units",
        "avg_pick_minutes",
        "verified_waves",
        "verified_orders",
        "verified_units",
        "shipments",
        "shipment_boxes",
        "shipment_orders",
        "shipment_units",
    ]
    for row in rows:
        sheet.append([_excel_safe_value(row.get(key, "")) for key in row_keys])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{max(sheet.max_row, 1)}"
    widths = [14, 32, 22, 16, 19, 19, 24, 18, 21, 21, 20, 20, 21, 21]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    summary_sheet = workbook.create_sheet("Сводка")
    for label, key in (
        ("Сотрудников", "employees"),
        ("Рабочих дней", "days"),
        ("Волн собрано", "pick_waves"),
        ("Заказов собрано", "pick_orders"),
        ("Единиц собрано", "pick_units"),
        ("Отгрузок передано", "shipments"),
        ("Заказов отгружено", "shipment_orders"),
        ("Единиц отгружено", "shipment_units"),
    ):
        summary_sheet.append([label, summary.get(key, 0)])
    for cell in summary_sheet["A"]:
        cell.font = header_font
    summary_sheet.column_dimensions["A"].width = 26
    summary_sheet.column_dimensions["B"].width = 18

    if overview:
        stages_sheet = workbook.create_sheet("По этапам", 1)
        stages_sheet.append(["Этап", "Описание", "Заказов", "Единиц", "Операций", "Исполнителей"])
        for cell in stages_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in overview["section_rows"]:
            stages_sheet.append(
                [
                    row["title"],
                    row["description"],
                    row["orders"],
                    row["units"],
                    f'{row["operations"]} {row["operations_label"]}',
                    row["employees"] if row["employees"] is not None else "-",
                ]
            )
        stages_sheet.freeze_panes = "A2"
        for column, width in zip("ABCDEF", (20, 54, 14, 14, 20, 18)):
            stages_sheet.column_dimensions[column].width = width

        status_sheet = workbook.create_sheet("Статус новых заказов", 2)
        status_sheet.append(["Состояние заказов, поступивших за период", "Заказов"])
        for cell in status_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in overview["order_status_rows"]:
            status_sheet.append([row["title"], row["orders"]])
        status_sheet.column_dimensions["A"].width = 38
        status_sheet.column_dimensions["B"].width = 14
    return workbook


def _head_manager_fbs_employee_report_export_response(filters: dict) -> HttpResponse:
    rows, summary = _head_manager_fbs_employee_report_rows(filters)
    overview = _head_manager_fbs_employee_report_overview(filters)
    workbook = _head_manager_fbs_employee_report_workbook(rows, summary, overview)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="head_manager_fbs_employees_{filters["date_from"]}_{filters["date_to"]}.xlsx"'
    )
    return response


def _head_manager_warehouse_employee_report_filters(request) -> dict:
    filters = _head_manager_fbs_employee_report_filters(request)
    role = str(request.GET.get("role") or "").strip()
    valid_roles = {key for key, _label in Employee.ROLE_CHOICES}
    filters["role"] = role if role in valid_roles else ""
    return filters


def _head_manager_warehouse_employee_profile(actor):
    if isinstance(actor, Employee):
        return actor
    return getattr(actor, "employee_profile", None)


def _head_manager_warehouse_employee_identity_key(actor) -> str:
    employee = _head_manager_warehouse_employee_profile(actor)
    if employee is not None and getattr(employee, "pk", None):
        return f"employee:{int(employee.pk)}"
    return f"user:{int(getattr(actor, 'pk', 0) or 0)}"


def _head_manager_warehouse_employee_role_key(actor) -> str:
    employee = _head_manager_warehouse_employee_profile(actor)
    return str(getattr(employee, "role", "") or "").strip()


def _head_manager_warehouse_employee_blank_row(day, actor) -> dict:
    employee = _head_manager_warehouse_employee_profile(actor)
    if employee is not None:
        employee_name = str(employee.full_name or "").strip() or "Не указан"
        role = str(employee.get_role_display() or "").strip() or "-"
        employee_id = int(employee.pk or 0)
        user_id = int(employee.user_id or 0)
    else:
        employee_name, role = _head_manager_fbs_employee_identity(actor)
        employee_id = 0
        user_id = int(getattr(actor, "pk", 0) or 0)
    role_key = _head_manager_warehouse_employee_role_key(actor)
    processing_history_partial = (
        role_key in _WAREHOUSE_EMPLOYEE_PROCESSING_ROLES
        and day < datetime(2026, 8, 25).date()
    )
    row = {
        "user_id": user_id,
        "employee_id": employee_id,
        "identity_key": _head_manager_warehouse_employee_identity_key(actor),
        "date_value": day,
        "date": day.strftime("%d.%m.%Y"),
        "employee": employee_name,
        "role": role,
        "role_key": role_key,
        "source_status": (
            "Частично: история обработки до 25.08.2026 неполная"
            if processing_history_partial
            else "Есть данные"
        ),
        "source_missing": False,
        "first_action": None,
        "last_action": None,
        "first_action_label": "—",
        "last_action_label": "—",
        "activity_span_seconds": None,
        "activity_span_label": "—",
        "contours": (),
        "contours_label": "—",
        "has_activity_data": False,
    }
    for metric in _WAREHOUSE_EMPLOYEE_REPORT_METRIC_KEYS:
        row[metric] = 0
    return row


def _head_manager_warehouse_employee_time_label(value) -> str:
    if value is None:
        return "—"
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.strftime("%H:%M")


def _head_manager_warehouse_employee_span_label(seconds) -> str:
    if seconds is None:
        return "—"
    total_minutes = max(0, int(seconds)) // 60
    hours, minutes = divmod(total_minutes, 60)
    if hours and minutes:
        return f"{hours} ч {minutes} мин"
    if hours:
        return f"{hours} ч"
    return f"{minutes} мин"


def _head_manager_warehouse_employee_report_rows(filters: dict) -> tuple[list[dict], dict]:
    role_filter = str(filters.get("role") or "").strip()
    rows_by_key: dict[tuple, dict] = {}

    def include_actor(actor) -> bool:
        return not role_filter or _head_manager_warehouse_employee_role_key(actor) == role_filter

    def report_row(day, actor) -> dict:
        key = (day, _head_manager_warehouse_employee_identity_key(actor))
        if key not in rows_by_key:
            rows_by_key[key] = _head_manager_warehouse_employee_blank_row(day, actor)
        return rows_by_key[key]

    for day, actor, metric, value in warehouse_employee_metric_points(filters):
        if metric not in _WAREHOUSE_EMPLOYEE_REPORT_METRIC_KEYS or not include_actor(actor):
            continue
        row = report_row(day, actor)
        if metric == "avg_pick_minutes":
            row[metric] = value if value is not None else 0
        else:
            row[metric] = (row.get(metric) or 0) + (value or 0)
        row["has_activity_data"] = True

    for window in warehouse_employee_work_windows(filters):
        actor = window.get("user")
        day = window.get("date_value")
        if actor is None or day is None or not include_actor(actor):
            continue
        row = report_row(day, actor)
        first_action = window.get("first_action")
        last_action = window.get("last_action")
        if first_action is not None:
            row["first_action"] = min(
                value for value in (row.get("first_action"), first_action) if value is not None
            )
        if last_action is not None:
            row["last_action"] = max(
                value for value in (row.get("last_action"), last_action) if value is not None
            )
        row["contours"] = tuple(
            sorted(set(row.get("contours") or ()) | set(window.get("contours") or ()))
        )
        row["has_activity_data"] = True

    date_from = _head_manager_fbs_employee_report_date(filters.get("date_from"))
    date_to = _head_manager_fbs_employee_report_date(filters.get("date_to"))
    processing_employees = Employee.objects.filter(
        is_active=True,
        role__in=_WAREHOUSE_EMPLOYEE_PROCESSING_ROLES,
    ).select_related("user")
    if role_filter:
        processing_employees = processing_employees.filter(role=role_filter)
    employee_filter = str(filters.get("employee") or "").strip()
    if employee_filter:
        processing_employees = processing_employees.filter(
            Q(full_name__icontains=employee_filter)
            | Q(user__username__icontains=employee_filter)
            | Q(user__first_name__icontains=employee_filter)
            | Q(user__last_name__icontains=employee_filter)
        )
    if date_from and date_to:
        report_day = date_from
        while report_day <= date_to:
            for employee in processing_employees:
                key = (report_day, _head_manager_warehouse_employee_identity_key(employee))
                if key not in rows_by_key:
                    rows_by_key[key] = _head_manager_warehouse_employee_blank_row(
                        report_day,
                        employee,
                    )
            report_day += timedelta(days=1)

    rows = list(rows_by_key.values())
    for row in rows:
        if row["first_action"] is not None and row["last_action"] is not None:
            row["activity_span_seconds"] = max(
                0,
                int((row["last_action"] - row["first_action"]).total_seconds()),
            )
        row["first_action_label"] = _head_manager_warehouse_employee_time_label(row["first_action"])
        row["last_action_label"] = _head_manager_warehouse_employee_time_label(row["last_action"])
        row["activity_span_label"] = _head_manager_warehouse_employee_span_label(
            row["activity_span_seconds"]
        )
        row["contours_label"] = ", ".join(
            _WAREHOUSE_EMPLOYEE_CONTOUR_LABELS.get(contour, contour)
            for contour in row["contours"]
        ) or "—"
        row["metric_cells"] = [
            {"key": metric, "value": row.get(metric)}
            for metric in _WAREHOUSE_EMPLOYEE_REPORT_METRIC_KEYS
        ]
        row["detail_url"] = (
            "/head-manager/reports/employees/"
            f"{row['employee_id']}/?date_from={filters['date_from']}&date_to={filters['date_to']}"
            if row.get("employee_id")
            else ""
        )
        row.pop("has_activity_data", None)

    rows.sort(key=lambda row: (-row["date_value"].toordinal(), row["employee"].lower()))
    summary = {
        "employees": len({row["identity_key"] for row in rows}),
        "days": len({row["date_value"] for row in rows}),
        "rows": len(rows),
        "no_source_employees": len(
            {row["identity_key"] for row in rows if row["source_missing"]}
        ),
        "receiving_cz_units": sum(row.get("receiving_cz_units") or 0 for row in rows),
        "pick_units": sum(row.get("pick_units") or 0 for row in rows),
        "verified_units": sum(row.get("verified_units") or 0 for row in rows),
        "shipment_units": sum(row.get("shipment_units") or 0 for row in rows),
        "shipping_units": sum(row.get("shipping_units") or 0 for row in rows),
        "processing_units": sum(row.get("processing_units") or 0 for row in rows),
        "processing_boxes": sum(row.get("processing_boxes") or 0 for row in rows),
    }
    return rows, summary


def _head_manager_employee_activity_filters(request) -> dict:
    filters = _head_manager_fbs_employee_report_filters(request)
    date_from = _head_manager_fbs_employee_report_date(filters["date_from"])
    date_to = _head_manager_fbs_employee_report_date(filters["date_to"])
    max_period_days = 92
    period_limited = (date_to - date_from).days >= max_period_days
    if period_limited:
        date_from = date_to - timedelta(days=max_period_days - 1)
        filters["date_from"] = date_from.isoformat()
    role = str(request.GET.get("role") or "").strip()
    valid_roles = {key for key, _label in Employee.ROLE_CHOICES}
    status_filter = str(request.GET.get("status") or "active").strip().lower()
    if status_filter not in {"active", "inactive", "all"}:
        status_filter = "active"
    filters.update(
        {
            "role": role if role in valid_roles else "",
            "status": status_filter,
            "period_limited": period_limited,
        }
    )
    return filters


def _head_manager_employee_activity_contours_label(contours) -> str:
    return ", ".join(
        EMPLOYEE_ACTIVITY_CONTOUR_LABELS.get(contour, contour)
        for contour in contours or ()
    ) or "—"


def _head_manager_employee_activity_present_rollup(rollup: dict) -> dict:
    result = dict(rollup)
    result["first_action_label"] = _head_manager_warehouse_employee_time_label(
        rollup.get("first_action")
    )
    result["last_action_label"] = _head_manager_warehouse_employee_time_label(
        rollup.get("last_action")
    )
    result["contours_label"] = _head_manager_employee_activity_contours_label(
        rollup.get("contours")
    )
    daily_rows = []
    for item in rollup.get("daily") or ():
        row = dict(item)
        first_action = row.get("first_action")
        last_action = row.get("last_action")
        span_seconds = (
            max(0, int((last_action - first_action).total_seconds()))
            if first_action is not None and last_action is not None
            else None
        )
        row.update(
            {
                "date": row["date_value"].strftime("%d.%m.%Y"),
                "first_action_label": _head_manager_warehouse_employee_time_label(first_action),
                "last_action_label": _head_manager_warehouse_employee_time_label(last_action),
                "activity_span_label": _head_manager_warehouse_employee_span_label(span_seconds),
                "contours_label": _head_manager_employee_activity_contours_label(
                    row.get("contours")
                ),
            }
        )
        daily_rows.append(row)
    result["daily"] = daily_rows
    return result


def _head_manager_employee_activity_directory_rows(filters: dict) -> tuple[list[dict], dict]:
    queryset = Employee.objects.select_related("user").order_by("full_name", "id")
    status_filter = str(filters.get("status") or "active")
    if status_filter == "active":
        queryset = queryset.filter(is_active=True)
    elif status_filter == "inactive":
        queryset = queryset.filter(is_active=False)
    role_filter = str(filters.get("role") or "")
    if role_filter:
        queryset = queryset.filter(role=role_filter)
    employee_filter = str(filters.get("employee") or "").strip()
    if employee_filter:
        queryset = queryset.filter(
            Q(full_name__icontains=employee_filter)
            | Q(user__username__icontains=employee_filter)
            | Q(user__first_name__icontains=employee_filter)
            | Q(user__last_name__icontains=employee_filter)
        )
    employees = list(queryset)
    rollups = warehouse_employee_activity_rollups(employees, filters)
    query_string = urlencode(
        {
            "date_from": filters["date_from"],
            "date_to": filters["date_to"],
        }
    )
    rows = []
    for employee in employees:
        rollup = _head_manager_employee_activity_present_rollup(
            rollups.get(employee.pk) or {
                "employee_id": employee.pk,
                "actions": 0,
                "scans": 0,
                "scan_success": 0,
                "scan_errors": 0,
                "volume": 0,
                "active_days": 0,
                "first_action": None,
                "last_action": None,
                "contours": (),
                "daily": (),
            }
        )
        row = {
            **rollup,
            "employee": str(employee.full_name or "").strip() or "Не указан",
            "role": str(employee.get_role_display() or "").strip() or "—",
            "role_key": str(employee.role or ""),
            "login": (
                str(employee.user.get_username() or "").strip()
                if employee.user_id
                else "Нет учётной записи"
            ),
            "is_active": bool(employee.is_active),
            "status_label": "Активен" if employee.is_active else "Неактивен",
            "detail_url": (
                f"/head-manager/reports/employees/{employee.pk}/?{query_string}"
            ),
        }
        rows.append(row)
    summary = {
        "employees": len(rows),
        "with_activity": sum(1 for row in rows if row["actions"]),
        "without_activity": sum(1 for row in rows if not row["actions"]),
        "actions": sum(row["actions"] for row in rows),
        "scans": sum(row["scans"] for row in rows),
        "scan_errors": sum(row["scan_errors"] for row in rows),
    }
    return rows, summary


def _head_manager_employee_activity_detail(employee, filters: dict) -> dict:
    rollup = _head_manager_employee_activity_present_rollup(
        warehouse_employee_activity_rollups([employee], filters)[employee.pk]
    )
    timeline, timeline_truncated = warehouse_employee_activity_timeline(
        employee,
        filters,
        limit=500,
    )
    for entry in timeline:
        local_time = timezone.localtime(entry["occurred_at"])
        entry["date"] = local_time.strftime("%d.%m.%Y")
        entry["time"] = local_time.strftime("%H:%M:%S")
    rollup.update(
        {
            "timeline": timeline,
            "timeline_truncated": timeline_truncated,
        }
    )
    return rollup


def _head_manager_warehouse_employee_report_workbook(rows: list[dict], summary: dict) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Работа сотрудников"
    metric_headers = [
        label
        for group in _WAREHOUSE_EMPLOYEE_REPORT_GROUPS
        for _key, label, _help in group["metrics"]
    ]
    headers = [
        "Дата",
        "Сотрудник",
        "Роль",
        "Доступность данных",
        "Первое действие",
        "Последнее действие",
        "Интервал действий",
        "Контуры",
        *metric_headers,
    ]
    sheet.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="F8B800")
    header_font = Font(bold=True, color="303030")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        metric_values = [
            "нет источника" if row.get(metric) is None else row.get(metric, 0)
            for metric in _WAREHOUSE_EMPLOYEE_REPORT_METRIC_KEYS
        ]
        sheet.append(
            [
                _excel_safe_value(row.get("date", "")),
                _excel_safe_value(row.get("employee", "")),
                _excel_safe_value(row.get("role", "")),
                _excel_safe_value(row.get("source_status", "")),
                row.get("first_action_label", "—"),
                row.get("last_action_label", "—"),
                row.get("activity_span_label", "—"),
                _excel_safe_value(row.get("contours_label", "—")),
                *metric_values,
            ]
        )
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
    base_widths = [14, 32, 24, 38, 18, 18, 20, 42]
    for index, width in enumerate(base_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index in range(len(base_widths) + 1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 18

    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Показатель", "Значение"])
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for label, key in (
        ("Сотрудников", "employees"),
        ("Дней с данными", "days"),
        ("Строк отчёта", "rows"),
        ("Сотрудников без источника", "no_source_employees"),
        ("Принято ЧЗ", "receiving_cz_units"),
        ("FBS единиц отобрано", "pick_units"),
        ("FBS единиц проверено", "verified_units"),
        ("FBS единиц отгружено", "shipment_units"),
        ("Не-FBS единиц отгружено", "shipping_units"),
        ("Обработано единиц", "processing_units"),
        ("Сформировано коробов", "processing_boxes"),
    ):
        summary_sheet.append([label, summary.get(key, 0)])
    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = f"A1:B{summary_sheet.max_row}"
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 18
    return workbook


def _head_manager_warehouse_employee_report_export_response(filters: dict) -> HttpResponse:
    rows, summary = _head_manager_warehouse_employee_report_rows(filters)
    workbook = _head_manager_warehouse_employee_report_workbook(rows, summary)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="warehouse_employees_{filters["date_from"]}_{filters["date_to"]}.xlsx"'
    )
    return response


def _head_manager_container_report_config(kind: str) -> dict:
    config = _CONTAINER_REPORTS.get(str(kind or "").strip())
    if not config:
        raise Http404("Отчет не найден")
    return config


def _head_manager_container_filters(request) -> dict:
    return {
        "q": str(request.GET.get("q") or "").strip(),
        "container": str(request.GET.get("container") or "").strip(),
        "client": str(request.GET.get("client") or "").strip(),
        "sku": str(request.GET.get("sku") or "").strip(),
        "warehouse": str(request.GET.get("warehouse") or "").strip(),
        "zone": str(request.GET.get("zone") or "").strip(),
        "chz": str(request.GET.get("chz") or "").strip(),
    }


def _head_manager_container_report_rows(kind: str, filters: dict, *, limit: int | None = 300) -> list[dict]:
    config = _head_manager_container_report_config(kind)
    container_key = config["container_key"]
    rows: list[dict] = []
    total_seen = 0
    container_codes: set[str] = set()
    box_codes: set[str] = set()
    pallet_codes: set[str] = set()
    qty_sum = 0
    marked_count = 0
    qs = (
        WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        .select_related(
            "agency",
            "sku_ref",
            "container",
            "parent_container",
            "location",
            "active_operation",
            "last_event",
        )
        .order_by("parent_container__container_code", "container__container_code", "sku_code", "marking_code", "id")
    )
    query = str(filters.get("q") or "").strip()
    if query:
        qs = qs.filter(
            Q(sku_code__icontains=query)
            | Q(name__icontains=query)
            | Q(barcode__icontains=query)
            | Q(marking_code__icontains=query)
            | Q(container_code__icontains=query)
            | Q(container__container_code__icontains=query)
            | Q(parent_container__container_code__icontains=query)
            | Q(agency__agn_name__icontains=query)
            | Q(agency__short_name__icontains=query)
        )
    client = str(filters.get("client") or "").strip()
    if client:
        qs = qs.filter(Q(agency__agn_name__icontains=client) | Q(agency__short_name__icontains=client))
    sku = str(filters.get("sku") or "").strip()
    if sku:
        qs = qs.filter(Q(sku_code__icontains=sku) | Q(name__icontains=sku) | Q(barcode__icontains=sku))
    warehouse = str(filters.get("warehouse") or "").strip()
    if warehouse:
        qs = qs.filter(location__warehouse_code__icontains=warehouse)
    zone = str(filters.get("zone") or "").strip()
    if zone:
        qs = qs.filter(Q(zone_code__icontains=zone) | Q(location__zone_code__icontains=zone))
    chz = str(filters.get("chz") or "").strip()
    if chz:
        qs = qs.filter(marking_code__icontains=chz)

    container_filter = str(filters.get("container") or "").strip().lower()
    for snapshot in qs.iterator(chunk_size=1000):
        normalized = normalize_stock_row_from_snapshot(snapshot) or {}
        box_code = str(normalized.get("box_code") or "").strip()
        pallet_code = str(normalized.get("pallet_code") or "").strip()
        container_code = box_code if container_key == "box_code" else pallet_code
        if not container_code:
            continue
        if container_filter and container_filter not in container_code.lower():
            continue
        total_seen += 1
        container_codes.add(container_code)
        if box_code:
            box_codes.add(box_code)
        if pallet_code:
            pallet_codes.add(pallet_code)
        qty_sum += int(snapshot.qty or 0)
        if str(snapshot.marking_code or "").strip():
            marked_count += 1
        if limit is not None and len(rows) >= limit:
            continue
        source_type = str(snapshot.source_context_type or normalized.get("source_order_type") or "").strip()
        source_id = str(snapshot.source_context_id or normalized.get("source_order_id") or "").strip()
        warehouse_code = (
            str(getattr(snapshot.location, "warehouse_code", "") or "").strip()
            or str(normalized.get("warehouse") or "").strip()
            or "-"
        )
        state_code = str(snapshot.warehouse_state_code or "").strip().lower()
        rows.append(
            {
                "container_code": container_code,
                "pallet_code": pallet_code or "-",
                "box_code": box_code or "-",
                "client": _movement_report_client_label(snapshot, normalized) or "-",
                "sku": normalized.get("sku") or snapshot.sku_code or "-",
                "name": normalized.get("name") or snapshot.name or "-",
                "size": normalized.get("size") or snapshot.size or "-",
                "barcode": normalized.get("barcode") or snapshot.barcode or "-",
                "qty": int(snapshot.qty or 0),
                "available_qty": int(snapshot.available_qty or 0),
                "processing_reserved_qty": int(snapshot.processing_reserved_qty or 0),
                "shipping_reserved_qty": int(snapshot.shipping_reserved_qty or 0),
                "marking_code": str(snapshot.marking_code or "").strip(),
                "warehouse": warehouse_code,
                "zone": normalized.get("zone") or snapshot.zone_code or "-",
                "location": normalized.get("location") or "-",
                "status": _MOVEMENT_REPORT_STATUS_LABELS.get(state_code, state_code or "-"),
                "order": format_order_number(source_type, source_id) if source_type and source_id else "-",
                "total_seen": total_seen,
            }
        )
    if rows:
        rows[0]["total_seen"] = total_seen
        rows[0]["summary_total"] = {
            "containers": len(container_codes),
            "boxes": len(box_codes),
            "pallets": len(pallet_codes),
            "qty": qty_sum,
            "marked": marked_count,
            "shown": len(rows),
            "total_seen": total_seen,
        }
    return rows


def _head_manager_container_report_summary(rows: list[dict]) -> dict:
    if rows and isinstance(rows[0].get("summary_total"), dict):
        return rows[0]["summary_total"]
    total_seen = int(rows[0].get("total_seen") or len(rows)) if rows else 0
    return {
        "containers": len({row["container_code"] for row in rows if row.get("container_code")}),
        "boxes": len({row["box_code"] for row in rows if row.get("box_code") and row.get("box_code") != "-"}),
        "pallets": len({row["pallet_code"] for row in rows if row.get("pallet_code") and row.get("pallet_code") != "-"}),
        "qty": sum(int(row.get("qty") or 0) for row in rows),
        "marked": sum(1 for row in rows if row.get("marking_code")),
        "shown": len(rows),
        "total_seen": total_seen,
    }


def _head_manager_container_report_workbook(kind: str, rows: list[dict]) -> Workbook:
    config = _head_manager_container_report_config(kind)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config["title"][:31]
    sheet.append(_CONTAINER_REPORT_HEADERS)
    header_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    sheet.freeze_panes = "A2"

    row_keys = [
        "container_code",
        "pallet_code",
        "box_code",
        "client",
        "sku",
        "name",
        "size",
        "barcode",
        "qty",
        "available_qty",
        "processing_reserved_qty",
        "shipping_reserved_qty",
        "marking_code",
        "warehouse",
        "zone",
        "location",
        "status",
        "order",
    ]
    for row in rows:
        sheet.append([_excel_safe_value(row.get(key, "")) for key in row_keys])
    sheet.auto_filter.ref = f"A1:R{max(sheet.max_row, 1)}"
    for index, width in enumerate([24, 24, 24, 28, 22, 38, 14, 22, 12, 12, 14, 14, 64, 14, 12, 34, 22, 22], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    summary = _head_manager_container_report_summary(rows)
    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Отчет", config["title"]])
    summary_sheet.append(["Контейнеров", summary["containers"]])
    summary_sheet.append(["Коробов", summary["boxes"]])
    summary_sheet.append(["Паллет", summary["pallets"]])
    summary_sheet.append(["Количество, шт.", summary["qty"]])
    summary_sheet.append(["Строк с ЧЗ", summary["marked"]])
    for cell in summary_sheet["A"]:
        cell.font = header_font
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 32
    return workbook


def _head_manager_container_export_response(kind: str, filters: dict) -> HttpResponse:
    config = _head_manager_container_report_config(kind)
    rows = _head_manager_container_report_rows(kind, filters, limit=None)
    workbook = _head_manager_container_report_workbook(kind, rows)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="head_manager_{config["export_name"]}_chz_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


_HEAD_MANAGER_WAREHOUSE_CITY_LABELS = {
    "MSK": "Москва",
    "MOSCOW": "Москва",
    "SPB": "Санкт-Петербург",
    "LED": "Санкт-Петербург",
    "KLD": "Калининград",
    "TUL": "Тула",
    "TULA": "Тула",
    "EKB": "Екатеринбург",
    "KZN": "Казань",
    "NSK": "Новосибирск",
}


def _head_manager_compact_marking_code(value) -> str:
    text = normalize_marking_code(value)
    text = re.sub(r"_x001d_", "\x1d", text, flags=re.IGNORECASE)
    text = re.sub(r"<\s*gs\s*>", "\x1d", text, flags=re.IGNORECASE)
    text = text.replace("\\u001d", "\x1d").replace("\\x1d", "\x1d").replace("\\x001d", "\x1d")
    return re.sub(r"[\s\x1d]+", "", text)


def _head_manager_marking_display_code(value) -> str:
    text = normalize_marking_code(value)
    return text.replace("\x1d", "<GS>")


def _head_manager_marking_search_terms(value) -> list[str]:
    normalized = normalize_marking_code(value)
    compact = _head_manager_compact_marking_code(normalized)
    terms = [
        str(value or "").strip(),
        normalized,
        normalized.replace("\x1d", "<GS>"),
        compact,
    ]
    return [term for term in dict.fromkeys(terms) if term]


def _head_manager_marking_prefix(value) -> str:
    compact = _head_manager_compact_marking_code(value)
    return compact[:24] if len(compact) >= 24 else compact[:16]


def _head_manager_warehouse_city(warehouse_code: str) -> str:
    code = str(warehouse_code or "").strip()
    if not code:
        return "-"
    return _HEAD_MANAGER_WAREHOUSE_CITY_LABELS.get(code.upper(), code)


def _head_manager_location_label(location, zone_code: str = "") -> str:
    display_name = str(getattr(location, "display_name", "") or "").strip()
    location_code = str(getattr(location, "location_code", "") or "").strip()
    mojibake_markers = ("В·", "Р ", "РЎ", "РЇ", "Р—", "Рћ", "Рњ")
    if display_name and not any(marker in display_name for marker in mojibake_markers):
        return display_name
    zone = str(getattr(location, "zone_code", "") or zone_code or "").strip().upper()
    row = int(getattr(location, "row_no", 0) or 0)
    section = int(getattr(location, "section_no", 0) or 0)
    tier = int(getattr(location, "tier_no", 0) or 0)
    cell = int(getattr(location, "cell_no", 0) or 0)
    if zone == "OS" and all((row, section, tier, cell)):
        return f"OS · Ряд {row} · Секция {section} · Ярус {tier} · Ячейка {cell}"
    if zone == "OS" and row:
        return f"OS · Ряд {row}"
    zone_labels = {
        "PR": "PR · Зона приемки",
        "OBR": "OBR · Зона обработки",
        "OTG": "OTG · Зона отгрузки",
        "MR": f"MR · Между рядами · Ряд {row}" if row else "MR · Между рядами",
        "OS": "OS · Основной склад",
    }
    return zone_labels.get(zone) or location_code or zone or "-"


def _head_manager_snapshot_location(snapshot: WarehouseStockSnapshot | None) -> dict:
    if not snapshot:
        return {"city": "-", "warehouse": "-", "zone": "-", "location": "-"}
    location = snapshot.location
    warehouse_code = str(getattr(location, "warehouse_code", "") or "").strip()
    zone = str(getattr(location, "zone_code", "") or snapshot.zone_code or "").strip()
    location_label = _head_manager_location_label(location, zone)
    return {
        "city": _head_manager_warehouse_city(warehouse_code),
        "warehouse": warehouse_code or "-",
        "zone": zone or "-",
        "location": location_label,
    }


def _head_manager_user_label(user) -> str:
    if not user:
        return "-"
    full_name = ""
    if hasattr(user, "get_full_name"):
        full_name = str(user.get_full_name() or "").strip()
    return full_name or str(getattr(user, "username", "") or "").strip() or "-"


def _head_manager_marking_stock_candidates(query: str) -> list[WarehouseStockSnapshot]:
    compact_query = _head_manager_compact_marking_code(query)
    if not compact_query:
        return []
    terms = _head_manager_marking_search_terms(query)
    prefix = _head_manager_marking_prefix(query)
    filters = Q(marking_code__in=terms)
    if prefix:
        filters |= Q(marking_code__icontains=prefix)
    qs = (
        WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        .filter(filters)
        .select_related("agency", "sku_ref", "container", "parent_container", "location")
        .order_by("-updated_at", "-id")[:200]
    )
    return [
        snapshot
        for snapshot in qs
        if _head_manager_compact_marking_code(snapshot.marking_code) == compact_query
    ]


def _head_manager_marking_receiving_candidates(query: str) -> list[ReceivingCzUnit]:
    compact_query = _head_manager_compact_marking_code(query)
    if not compact_query:
        return []
    terms = _head_manager_marking_search_terms(query)
    prefix = _head_manager_marking_prefix(query)
    filters = Q(marking_code__in=terms)
    if prefix:
        filters |= Q(marking_code__icontains=prefix)
    qs = (
        ReceivingCzUnit.objects.filter(filters)
        .select_related("agency", "sku", "accepted_by")
        .order_by("-accepted_at", "-id")[:200]
    )
    return [unit for unit in qs if _head_manager_compact_marking_code(unit.marking_code) == compact_query]


def _head_manager_marking_registry_candidates(query: str) -> list[MarkingCode]:
    compact_query = _head_manager_compact_marking_code(query)
    if not compact_query:
        return []
    terms = _head_manager_marking_search_terms(query)
    prefix = _head_manager_marking_prefix(query)
    filters = Q(code__in=terms)
    if prefix:
        filters |= Q(code__icontains=prefix)
    qs = (
        MarkingCode.objects.filter(filters)
        .select_related("agency", "sku", "used_by")
        .order_by("-used_at", "-created_at", "-id")[:200]
    )
    return [code for code in qs if _head_manager_compact_marking_code(code.code) == compact_query]


def _head_manager_receiving_payload(unit: ReceivingCzUnit | None) -> dict:
    if not unit:
        return {
            "receiving_document": "-",
            "accepted_at": None,
            "accepted_by": "-",
            "receiving_box": "-",
        }
    return {
        "receiving_document": format_order_number("receiving", unit.order_id) if unit.order_id else "-",
        "accepted_at": unit.accepted_at,
        "accepted_by": _head_manager_user_label(unit.accepted_by),
        "receiving_box": unit.box_code or "-",
    }


def _head_manager_snapshot_payload(snapshot: WarehouseStockSnapshot, receiving_unit: ReceivingCzUnit | None = None) -> dict:
    normalized = normalize_stock_row_from_snapshot(snapshot) or {}
    location = _head_manager_snapshot_location(snapshot)
    box_code = str(normalized.get("box_code") or snapshot.container_code or getattr(snapshot.container, "container_code", "") or "").strip()
    pallet_code = str(normalized.get("pallet_code") or getattr(snapshot.parent_container, "container_code", "") or "").strip()
    reserved_qty = int(snapshot.processing_reserved_qty or 0) + int(snapshot.shipping_reserved_qty or 0) + int(snapshot.other_reserved_qty or 0)
    source_type = str(snapshot.source_context_type or normalized.get("source_order_type") or "").strip()
    source_id = str(snapshot.source_context_id or normalized.get("source_order_id") or "").strip()
    order_label = format_order_number(source_type, source_id) if source_type and source_id else "-"
    return {
        "source": "Складской остаток",
        "client": _movement_report_client_label(snapshot, normalized) or "-",
        "sku": normalized.get("sku") or snapshot.sku_code or "-",
        "name": normalized.get("name") or snapshot.name or "-",
        "size": normalized.get("size") or snapshot.size or "-",
        "barcode": normalized.get("barcode") or snapshot.barcode or "-",
        "marking_code": _head_manager_marking_display_code(snapshot.marking_code),
        "box": box_code or "-",
        "pallet": pallet_code or "-",
        "qty": int(snapshot.qty or 0),
        "available_qty": int(snapshot.available_qty or 0),
        "reserved_qty": reserved_qty,
        "order": order_label,
        "status": _MOVEMENT_REPORT_STATUS_LABELS.get(str(snapshot.warehouse_state_code or "").strip().lower(), snapshot.warehouse_state_code or "-"),
        "updated_at": snapshot.updated_at,
        **_head_manager_receiving_payload(receiving_unit),
        **location,
    }


def _head_manager_container_location_by_codes(box_code: str, pallet_code: str) -> dict:
    box = str(box_code or "").strip()
    pallet = str(pallet_code or "").strip()
    if not box and not pallet:
        return _head_manager_snapshot_location(None)
    filters = Q()
    if box:
        filters |= Q(container_code=box) | Q(container__container_code=box)
    if pallet:
        filters |= Q(parent_container__container_code=pallet)
    snapshot = (
        WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        .filter(filters)
        .select_related("location")
        .order_by("-updated_at", "-id")
        .first()
    )
    return _head_manager_snapshot_location(snapshot)


def _head_manager_marking_search_rows(query: str) -> list[dict]:
    compact_query = _head_manager_compact_marking_code(query)
    if not compact_query:
        return []

    rows: list[dict] = []
    seen_keys: set[tuple[str, str, str]] = set()
    receiving_units = _head_manager_marking_receiving_candidates(query)
    receiving_by_code = {
        _head_manager_compact_marking_code(unit.marking_code): unit
        for unit in receiving_units
    }
    stock_rows = [
        _head_manager_snapshot_payload(
            snapshot,
            receiving_by_code.get(_head_manager_compact_marking_code(snapshot.marking_code)),
        )
        for snapshot in _head_manager_marking_stock_candidates(query)
    ]
    for row in stock_rows:
        key = (row["source"], row["marking_code"], row["box"])
        if key not in seen_keys:
            seen_keys.add(key)
            rows.append(row)

    for unit in receiving_units:
        location = _head_manager_container_location_by_codes(unit.box_code, unit.pallet_code)
        row = {
            "source": "Приемка ЧЗ",
            "client": getattr(unit.agency, "agn_name", "") or getattr(unit.agency, "short_name", "") or "-",
            "sku": unit.sku_code or "-",
            "name": unit.name or "-",
            "size": unit.size or "-",
            "barcode": unit.barcode or "-",
            "marking_code": _head_manager_marking_display_code(unit.marking_code),
            "box": unit.box_code or "-",
            "pallet": unit.pallet_code or "-",
            "qty": 1,
            "available_qty": "-",
            "reserved_qty": "-",
            "order": format_order_number("receiving", unit.order_id) if unit.order_id else "-",
            "status": "Принят",
            "updated_at": unit.accepted_at,
            **_head_manager_receiving_payload(unit),
            **location,
        }
        key = (row["source"], row["marking_code"], row["box"])
        if key not in seen_keys:
            seen_keys.add(key)
            rows.append(row)

    for code in _head_manager_marking_registry_candidates(query):
        row = {
            "source": "Реестр маркировки",
            "client": getattr(code.agency, "agn_name", "") or getattr(code.agency, "short_name", "") or "-",
            "sku": code.sku_code or "-",
            "name": getattr(code.sku, "name", "") or "-",
            "size": code.size or "-",
            "barcode": code.barcode or "-",
            "marking_code": _head_manager_marking_display_code(code.code),
            "box": code.box_barcode or "-",
            "pallet": "-",
            "qty": 1,
            "available_qty": "-",
            "reserved_qty": "-",
            "order": format_order_number(code.order_type, code.order_id) if code.order_type and code.order_id else "-",
            "status": "Использован" if code.used_at else "В реестре",
            "updated_at": code.used_at or code.created_at,
            **_head_manager_receiving_payload(None),
            **_head_manager_container_location_by_codes(code.box_barcode, ""),
        }
        key = (row["source"], row["marking_code"], row["box"])
        if key not in seen_keys:
            seen_keys.add(key)
            rows.append(row)

    return rows


def _head_manager_user_name(request) -> str:
    employee = getattr(request.user, "employee_profile", None)
    if employee:
        return employee.full_name
    return request.user.get_full_name() or request.user.get_username()


def _fbs_order_status_class(status: str) -> str:
    if status in {
        FbsOrder.STATUS_VALIDATION_FAILED,
        FbsOrder.STATUS_EXCEPTION,
    }:
        return "danger"
    if status in {
        FbsOrder.STATUS_PICKING,
        FbsOrder.STATUS_PICKED,
        FbsOrder.STATUS_QUEUED_FOR_PICK,
    }:
        return "info"
    if status in {
        FbsOrder.STATUS_READY_FOR_HANDOVER,
        FbsOrder.STATUS_HANDED_OVER,
        FbsOrder.STATUS_DELIVERED,
    }:
        return "success"
    if status in {FbsOrder.STATUS_CANCELLED, FbsOrder.STATUS_RETURNED}:
        return "neutral"
    return "warning"


def _fbs_form_error_text(form) -> str:
    parts = []
    for field_name, errors in form.errors.items():
        label = form.fields.get(field_name).label if field_name in form.fields else "Ошибка"
        parts.extend(f"{label}: {error}" for error in errors)
    return " ".join(parts) or "Проверьте заполнение формы."


def _fbs_known_printers() -> list[str]:
    names = set()
    for agent in DeviceAgent.objects.only("meta"):
        meta = agent.meta if isinstance(agent.meta, dict) else {}
        for key in ("printers", "last_known_printers"):
            values = meta.get(key)
            if not isinstance(values, list):
                continue
            names.update(str(value).strip() for value in values if str(value).strip())
    return sorted(names, key=str.casefold)


class HeadManagerFbsOperationsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_operations.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        query = str(self.request.GET.get("q") or "").strip()
        status_filter = str(self.request.GET.get("status") or "").strip()
        marketplace_filter = str(self.request.GET.get("marketplace") or "").strip()

        pick_tasks = FbsPickTask.objects.select_related("batch", "assigned_to").order_by("-id")
        orders = (
            FbsOrder.objects.select_related("profile__agency")
            .prefetch_related(Prefetch("pick_tasks", queryset=pick_tasks, to_attr="overview_pick_tasks"))
            .annotate(item_count=Count("items", distinct=True), total_qty=Sum("items__quantity"))
        )
        if query:
            orders = orders.filter(
                Q(external_order_id__icontains=query)
                | Q(profile__agency__agn_name__icontains=query)
                | Q(items__external_sku__icontains=query)
                | Q(items__product_name__icontains=query)
            ).distinct()
        valid_statuses = {value for value, _ in FbsOrder.STATUS_CHOICES}
        if status_filter in valid_statuses:
            orders = orders.filter(internal_status=status_filter)
        else:
            status_filter = ""
        valid_marketplaces = {value for value, _ in FbsIntegrationProfile.MARKETPLACE_CHOICES}
        if marketplace_filter in valid_marketplaces:
            orders = orders.filter(profile__marketplace=marketplace_filter)
        else:
            marketplace_filter = ""

        page_obj = Paginator(orders.order_by("cutoff_at", "imported_at", "id"), 50).get_page(
            self.request.GET.get("page")
        )
        for order in page_obj.object_list:
            order.status_class = _fbs_order_status_class(order.internal_status)
            order.overview_task = order.overview_pick_tasks[0] if order.overview_pick_tasks else None
            order.is_overdue = bool(
                order.cutoff_at
                and order.cutoff_at < timezone.now()
                and order.internal_status not in {
                    FbsOrder.STATUS_HANDED_OVER,
                    FbsOrder.STATUS_DELIVERED,
                    FbsOrder.STATUS_CANCELLED,
                    FbsOrder.STATUS_RETURNED,
                }
            )

        all_orders = FbsOrder.objects.all()
        final_statuses = {
            FbsOrder.STATUS_DELIVERED,
            FbsOrder.STATUS_CANCELLED,
            FbsOrder.STATUS_RETURNED,
        }
        transfer_conflicts = (
            FbsOrder.objects.filter(
                items__metadata_transfers__status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                    FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                )
            )
            .distinct()
            .count()
        )
        active_batches = list(
            FbsPickBatch.objects.filter(
                status__in=(
                    FbsPickBatch.STATUS_QUEUED,
                    FbsPickBatch.STATUS_IN_PROGRESS,
                    FbsPickBatch.STATUS_VERIFICATION,
                )
            )
            .select_related("agency", "assigned_to", "workstation", "cart")
            .annotate(task_count=Count("tasks"), picked_total=Sum("tasks__picked_qty"))
            .order_by("created_at", "id")[:20]
        )
        overridden_pairs = set(
            FbsComplianceOverride.objects.filter(is_active=True).values_list(
                "order_id", "metadata_type"
            )
        )
        compliance_override_candidates = []
        seen_override_pairs = set()
        for transfer in (
            FbsMarketplaceMetadataTransfer.objects.select_related(
                "order_item__order__profile__agency"
            )
            .filter(
                is_required=True,
                status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                    FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                    FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED,
                ),
            )
            .order_by("updated_at", "id")[:100]
        ):
            pair = (transfer.order_item.order_id, transfer.metadata_type)
            if pair in overridden_pairs or pair in seen_override_pairs:
                continue
            seen_override_pairs.add(pair)
            compliance_override_candidates.append(transfer)
            if len(compliance_override_candidates) >= 30:
                break
        params = self.request.GET.copy()
        params.pop("page", None)
        handover_batches = list(
            FbsHandoverBatch.objects.select_related("profile__agency")
            .prefetch_related(
                Prefetch(
                    "verification_overrides",
                    queryset=FbsHandoverVerificationOverride.objects.filter(
                        is_active=True
                    ).select_related("approved_by", "used_by"),
                    to_attr="active_verification_overrides",
                )
            )
            .annotate(box_count=Count("boxes", distinct=True))
            .order_by("-created_at")[:20]
        )
        for handover in handover_batches:
            handover.active_verification_override = (
                handover.active_verification_overrides[0]
                if handover.active_verification_overrides
                else None
            )
            handover.can_request_verification_override = bool(
                handover.profile.marketplace
                == FbsIntegrationProfile.MARKETPLACE_WB
                and handover.status
                in (
                    FbsHandoverBatch.STATUS_OPEN,
                    FbsHandoverBatch.STATUS_READY,
                )
                and handover.active_verification_override is None
            )
        context.update(
            {
                "active_batches": active_batches,
                "client_operations": client_operations_report(),
                "open_pick_exceptions": FbsPickException.objects.select_related(
                    "task__order__profile__agency",
                    "allocation__balance__box__pallet__cell",
                    "created_by",
                ).filter(status=FbsPickException.STATUS_OPEN).order_by("created_at")[:30],
                "active_inventory_sessions": FbsInventorySession.objects.select_related(
                    "agency", "cell", "pallet", "box", "sku"
                ).exclude(
                    status__in=(
                        FbsInventorySession.STATUS_DONE,
                        FbsInventorySession.STATUS_CANCELED,
                    )
                ).order_by("created_at")[:20],
                "active_internal_movements": FbsInternalMovement.objects.select_related(
                    "agency", "source_box", "target_pallet__cell", "target_box", "assigned_to"
                ).filter(
                    status__in=(
                        FbsInternalMovement.STATUS_PROPOSED,
                        FbsInternalMovement.STATUS_IN_PROGRESS,
                        FbsInternalMovement.STATUS_BLOCKED,
                    )
                ).order_by("created_at")[:20],
                "handover_batches": handover_batches,
                "compliance_override_candidates": compliance_override_candidates,
                "fbs_enabled": fbs_module_enabled(),
                "marketplace_choices": FbsIntegrationProfile.MARKETPLACE_CHOICES,
                "marketplace_filter": marketplace_filter,
                "page_obj": page_obj,
                "query": query,
                "query_string": params.urlencode(),
                "sidebar_active": "fbs",
                "status_choices": FbsOrder.STATUS_CHOICES,
                "status_filter": status_filter,
                "summary": {
                    "total": all_orders.count(),
                    "open": all_orders.exclude(internal_status__in=final_statuses).count(),
                    "picking": all_orders.filter(
                        internal_status__in=(
                            FbsOrder.STATUS_QUEUED_FOR_PICK,
                            FbsOrder.STATUS_PICKING,
                            FbsOrder.STATUS_PICKED,
                        )
                    ).count(),
                    "ready": all_orders.filter(
                        internal_status=FbsOrder.STATUS_READY_FOR_HANDOVER
                    ).count(),
                    "conflicts": transfer_conflicts,
                    "overdue": all_orders.filter(
                        cutoff_at__lt=timezone.now()
                    ).exclude(internal_status__in=final_statuses).count(),
                    "problems": all_orders.filter(
                        internal_status=FbsOrder.STATUS_EXCEPTION
                    ).count(),
                },
                "user_display_name": _head_manager_user_name(self.request),
            }
        )
        return context


class HeadManagerFbsStaffView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_staff.html"
    allowed_roles = ("head_manager",)
    staff_kind = "controllers"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.staff_kind not in {"controllers", "pickers"}:
            raise Http404("Неизвестный раздел сотрудников FBS.")

        is_controllers = self.staff_kind == "controllers"
        role = "fbs_controller" if is_controllers else "picker"
        employees = list(
            Employee.objects.filter(role=role)
            .select_related("user")
            .order_by("-is_active", "full_name", "id")
        )
        user_ids = [employee.user_id for employee in employees if employee.user_id]
        active_statuses = (
            FbsPickBatch.STATUS_QUEUED,
            FbsPickBatch.STATUS_IN_PROGRESS,
            FbsPickBatch.STATUS_VERIFICATION,
        )
        assignee_field = (
            "verification_assigned_to_id" if is_controllers else "assigned_to_id"
        )
        current_by_user: dict[int, list[FbsPickBatch]] = {}
        if user_ids:
            for batch in (
                FbsPickBatch.objects.filter(
                    **{
                        f"{assignee_field}__in": user_ids,
                        "status__in": active_statuses,
                    }
                )
                .select_related("agency", "workstation", "cart")
                .order_by("created_at", "id")
            ):
                current_by_user.setdefault(
                    getattr(batch, assignee_field), []
                ).append(batch)

        today = timezone.localdate()
        completed_by_user: dict[int, dict[str, int]] = {}
        if user_ids:
            completed_filter = {
                f"{assignee_field}__in": user_ids,
                (
                    "completed_at__date"
                    if is_controllers
                    else "picking_completed_at__date"
                ): today,
            }
            if is_controllers:
                completed_filter["status"] = FbsPickBatch.STATUS_DONE
            for item in (
                FbsPickBatch.objects.filter(**completed_filter)
                .values(assignee_field)
                .annotate(wave_count=Count("id"), unit_count=Sum("picked_qty"))
            ):
                completed_by_user[item[assignee_field]] = {
                    "waves": int(item["wave_count"] or 0),
                    "units": int(item["unit_count"] or 0),
                }

        policies = {
            policy.controller_id: policy
            for policy in FbsControllerPolicy.objects.filter(
                controller_id__in=user_ids
            ).select_related("updated_by")
        }
        rows = []
        for employee in employees:
            completed = completed_by_user.get(employee.user_id, {})
            rows.append(
                {
                    "employee": employee,
                    "current_waves": current_by_user.get(employee.user_id, []),
                    "completed_waves": completed.get("waves", 0),
                    "completed_units": completed.get("units", 0),
                    "policy": policies.get(employee.user_id),
                    "fast_wb_enabled": bool(
                        is_controllers
                        and employee.user_id
                        and policies.get(employee.user_id)
                        and policies[employee.user_id].skip_repeat_wb_label_scan
                    ),
                }
            )

        context.update(
            {
                "is_controllers": is_controllers,
                "rows": rows,
                "sidebar_active": "fbs",
                "staff_kind": self.staff_kind,
                "summary": {
                    "total": len(rows),
                    "active": sum(row["employee"].is_active for row in rows),
                    "current_waves": sum(
                        len(row["current_waves"]) for row in rows
                    ),
                    "completed_waves": sum(
                        row["completed_waves"] for row in rows
                    ),
                    "fast_wb": sum(row["fast_wb_enabled"] for row in rows),
                },
                "user_display_name": _head_manager_user_name(self.request),
            }
        )
        return context


class HeadManagerFbsControllerPolicyView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, employee_id: int):
        with transaction.atomic():
            employee = get_object_or_404(
                Employee.objects.select_for_update().select_related("user"),
                pk=employee_id,
            )
            if (
                employee.role != "fbs_controller"
                or not employee.is_active
                or employee.user_id is None
            ):
                messages.error(
                    request,
                    "Настройку можно менять только активному FBS-контролёру с учётной записью.",
                )
                return redirect("head-manager-fbs-controllers")
            enabled = str(request.POST.get("skip_repeat_wb_label_scan") or "") in {
                "1",
                "true",
                "on",
                "yes",
            }
            policy = (
                FbsControllerPolicy.objects.select_for_update()
                .filter(controller_id=employee.user_id)
                .first()
            )
            if policy is None:
                policy = FbsControllerPolicy(controller=employee.user)
            policy.skip_repeat_wb_label_scan = enabled
            policy.updated_by = request.user
            policy.save()
        state = "включён" if enabled else "выключен"
        messages.success(
            request,
            f"Ускоренный контроль WB для {employee.full_name} {state}.",
        )
        return redirect("head-manager-fbs-controllers")


class HeadManagerFbsComplianceOverrideView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, order_id: int):
        try:
            approve_compliance_override(
                order_id=order_id,
                metadata_type=str(request.POST.get("metadata_type") or "").strip(),
                reason=str(request.POST.get("reason") or "").strip(),
                approved_by=request.user,
            )
        except FbsError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, "Разрешение отгрузки записано в журнал FBS.")
        return redirect("head-manager-fbs")


class HeadManagerFbsHandoverVerificationOverrideView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, batch_id: int):
        try:
            approve_handover_verification_override(
                batch_id=batch_id,
                reason=str(request.POST.get("reason") or "").strip(),
                approved_by=request.user,
            )
        except FbsError as exc:
            messages.error(request, str(exc))
            return redirect("head-manager-fbs")
        messages.success(
            request,
            "Аварийный допуск записан в аудит. Повторная проверка состава отключена "
            "только для текущего неизмененного состава отгрузки.",
        )
        return redirect(f"/fbs/tsd/storekeeper/handover/{batch_id}/")


class HeadManagerFbsSettingsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_settings.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        workstations = list(FbsWorkstation.objects.select_related("device_agent", "updated_by"))
        carts = list(FbsPickingCart.objects.select_related("updated_by"))
        agencies = list(
            Agency.objects.filter(fbs_integration_profiles__isnull=False)
            .distinct()
            .order_by("agn_name", "id")
        )
        storage_policy_by_agency = {
            policy.agency_id: policy
            for policy in FbsClientStoragePolicy.objects.filter(
                agency_id__in=[agency.id for agency in agencies]
            )
        }
        for agency in agencies:
            agency.fbs_storage_policy = storage_policy_by_agency.get(agency.id)
        fbs_client_rows = _head_manager_fbs_client_rows()
        context.update(
            {
                "agents": DeviceAgent.objects.order_by("name", "host", "agent_id"),
                "cart_create_form": FbsPickingCartBatchCreateForm(),
                "carts": carts,
                "known_printers": _fbs_known_printers(),
                "fbs_agencies": agencies,
                "fbs_client_rows": fbs_client_rows,
                "global_order_pull_enabled": feature_enabled("order_pull"),
                "billing_choices": FbsClientStoragePolicy.BILLING_CHOICES,
                "replenishment_policies": FbsReplenishmentPolicy.objects.select_related(
                    "agency", "sku"
                ).order_by("agency__agn_name", "sku__sku_code")[:500],
                "fbs_skus": SKU.objects.filter(
                    agency_id__in=[agency.id for agency in agencies],
                    deleted=False,
                ).select_related("agency").order_by("agency__agn_name", "sku_code")[:1000],
                "sidebar_active": "settings",
                "summary": {
                    "workstations": len(workstations),
                    "active_workstations": sum(item.is_active for item in workstations),
                    "carts": len(carts),
                    "active_carts": sum(item.is_active for item in carts),
                    "configured_printers": sum(bool(item.printer_name) for item in workstations),
                    "clients": len(fbs_client_rows),
                    "active_clients": sum(row["fbs_enabled"] for row in fbs_client_rows),
                },
                "user_display_name": _head_manager_user_name(self.request),
                "workstation_create_form": FbsWorkstationBatchCreateForm(),
                "workstations": workstations,
            }
        )
        return context


class HeadManagerWarehouseLocationsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/warehouse_locations.html"
    allowed_roles = ("head_manager",)

    def post(self, request, *args, **kwargs):
        try:
            create_operational_location(
                zone_code=request.POST.get("zone_code", ""),
                location_code=request.POST.get("location_code", ""),
                display_name=request.POST.get("display_name", ""),
                capacity_containers=int(request.POST.get("capacity_containers") or 0),
                is_fbs_visible=request.POST.get("is_fbs_visible") == "1",
            )
        except (ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            messages.success(request, "Дополнительное место создано. Можно печатать и сканировать QR.")
        return redirect("head-manager-warehouse-locations")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        rows = list(
            WarehouseLocation.objects.filter(
                warehouse_code="MSK",
                zone_code__in=("PR", "OTG"),
                is_topology_visible=False,
            ).order_by("zone_code", "location_code", "id")
        )
        for location in rows:
            occupancy = operational_location_occupancy(location)
            location.occupancy = occupancy
            location.qr_value = (
                f"LOC:{location.warehouse_code}:{location.location_code}"
            )
        selected_location = None
        print_location_id = str(self.request.GET.get("print") or "").strip()
        if print_location_id.isdigit():
            selected_location = next(
                (
                    location
                    for location in rows
                    if int(location.id) == int(print_location_id)
                ),
                None,
            )
        context.update(
            {
                "locations": rows,
                "pr_locations": [row for row in rows if row.zone_code == "PR"],
                "otg_locations": [row for row in rows if row.zone_code == "OTG"],
                "selected_location": selected_location,
                "sidebar_active": "settings",
                "summary": {
                    "total": len(rows),
                    "active": sum(row.is_active for row in rows),
                    "fbs": sum(row.is_fbs_visible for row in rows),
                    "occupied": sum(row.occupancy.occupied for row in rows),
                },
                "user_display_name": _head_manager_user_name(self.request),
            }
        )
        return context


class HeadManagerWarehouseLocationUpdateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, location_id: int):
        try:
            update_operational_location(
                location_id=location_id,
                display_name=request.POST.get("display_name", ""),
                capacity_containers=int(request.POST.get("capacity_containers") or 0),
                is_fbs_visible=request.POST.get("is_fbs_visible") == "1",
                is_active=request.POST.get("is_active") == "1",
            )
        except WarehouseLocation.DoesNotExist:
            raise Http404
        except (ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            messages.success(request, "Настройки места сохранены.")
        return redirect("head-manager-warehouse-locations")


def _head_manager_fbs_client_rows() -> list[dict]:
    credential_rows = list(
        MarketCredential.objects.select_related("agency", "market")
        .exclude(market_key__isnull=True)
        .exclude(market_key="")
        .order_by("agency__agn_name", "agency_id", "id")
    )
    agencies: dict[int, Agency] = {}
    marketplaces: dict[int, set[str]] = {}
    for credential in credential_rows:
        code = marketplace_code(credential.market.name)
        if not code or not str(credential.market_key or "").strip():
            continue
        agencies[credential.agency_id] = credential.agency
        marketplaces.setdefault(credential.agency_id, set()).add(code)
    for credential in (
        FbsOzonCredential.objects.select_related("agency")
        .order_by("agency__agn_name", "agency_id", "slot")
    ):
        agencies[credential.agency_id] = credential.agency
        marketplaces.setdefault(credential.agency_id, set()).add(
            FbsIntegrationProfile.MARKETPLACE_OZON
        )
    profile_rows = list(
        FbsIntegrationProfile.objects.filter(agency_id__in=agencies)
        .prefetch_related(
            Prefetch(
                "sync_cursors",
                queryset=FbsSyncCursor.objects.filter(
                    stream=FbsSyncCursor.STREAM_ORDERS,
                    cursor_key="default",
                ),
                to_attr="order_sync_cursors",
            )
        )
        .order_by("agency_id", "marketplace", "name", "id")
    )
    profiles_by_agency: dict[int, list[FbsIntegrationProfile]] = {}
    for profile in profile_rows:
        profiles_by_agency.setdefault(profile.agency_id, []).append(profile)
    labels = dict(FbsIntegrationProfile.MARKETPLACE_CHOICES)
    result = []
    for agency_id, agency in agencies.items():
        profiles = profiles_by_agency.get(agency_id, [])
        active_profiles = [
            profile for profile in profiles if profile.is_active and profile.order_pull_enabled
        ]
        sync_cursors = [
            cursor
            for profile in active_profiles
            for cursor in getattr(profile, "order_sync_cursors", ())
        ]
        last_success_at = max(
            (cursor.last_success_at for cursor in sync_cursors if cursor.last_success_at),
            default=None,
        )
        last_error = next(
            (cursor.last_error for cursor in sync_cursors if cursor.last_error),
            "",
        )
        result.append(
            {
                "agency": agency,
                "marketplaces": [
                    {"code": code, "label": labels[code]}
                    for code in sorted(marketplaces.get(agency_id, set()))
                ],
                "profile_count": len(profiles),
                "active_profile_count": len(active_profiles),
                "fbs_enabled": bool(active_profiles),
                "warehouse_names": [profile.name for profile in active_profiles],
                "last_order_sync_at": last_success_at,
                "last_order_sync_error": last_error,
            }
        )
    return result


def _head_manager_fbs_client_agency(agency_id: int) -> Agency:
    return get_object_or_404(
        Agency.objects.filter(
            Q(pk=agency_id),
            Q(fbs_integration_profiles__isnull=False)
            | Q(fbs_ozon_credentials__isnull=False)
            | Q(market_credentials__market_key__isnull=False),
        ).distinct(),
        pk=agency_id,
    )


def _head_manager_fbs_client_catalog(agency: Agency) -> dict[str, dict]:
    catalog = fetch_client_warehouse_catalog(agency)
    profiles = list(
        FbsIntegrationProfile.objects.filter(agency=agency).order_by(
            "marketplace", "name", "id"
        )
    )
    profiles_by_marketplace: dict[str, list[FbsIntegrationProfile]] = {}
    for profile in profiles:
        profiles_by_marketplace.setdefault(profile.marketplace, []).append(profile)
    for marketplace, entry in catalog.items():
        profile_rows = profiles_by_marketplace.get(marketplace, [])
        warehouse_map = {
            str(row.get("selection_key") or row["warehouse_id"]): row
            for row in entry.get("warehouses", [])
        }
        for profile in profile_rows:
            warehouse_id = str(profile.external_warehouse_id or "").strip()
            external_account_id = str(profile.external_account_id or "").strip()
            selection_key = (
                f"{external_account_id}:{warehouse_id}"
                if marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
                else warehouse_id
            )
            if warehouse_id and selection_key not in warehouse_map:
                warehouse_map[selection_key] = {
                    "marketplace": marketplace,
                    "warehouse_id": warehouse_id,
                    "name": profile.name,
                    "status": "Сохранен в Fullbox",
                    "external_account_id": external_account_id,
                    "credential_slot": 0,
                    "selection_key": selection_key,
                }
        selected_keys = {
            (
                f"{str(profile.external_account_id or '').strip()}:{str(profile.external_warehouse_id or '').strip()}"
                if marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
                else str(profile.external_warehouse_id or "").strip()
            )
            for profile in profile_rows
            if profile.is_active and profile.order_pull_enabled
        }
        stock_push_keys = {
            (
                f"{str(profile.external_account_id or '').strip()}:{str(profile.external_warehouse_id or '').strip()}"
                if marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
                else str(profile.external_warehouse_id or "").strip()
            )
            for profile in profile_rows
            if (
                profile.is_active
                and profile.order_pull_enabled
                and profile.stock_push_enabled
                and profile.stock_mode == FbsIntegrationProfile.STOCK_MODE_MANAGED
            )
        }
        entry["warehouses"] = sorted(
            (
                {
                    **row,
                    "selection_key": str(row.get("selection_key") or selection_key),
                    "selected": selection_key in selected_keys,
                    "stock_push_selected": selection_key in stock_push_keys,
                }
                for selection_key, row in warehouse_map.items()
            ),
            key=lambda row: (
                int(row.get("credential_slot") or 0),
                str(row.get("name") or "").lower(),
                row["warehouse_id"],
            ),
        )
        entry["configured_profiles"] = len(profile_rows)
        entry["active_profiles"] = len(selected_keys)
    return catalog


class HeadManagerFbsClientSettingsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_client_settings.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        agency = _head_manager_fbs_client_agency(kwargs["agency_id"])
        catalog = _head_manager_fbs_client_catalog(agency)
        active_profiles = FbsIntegrationProfile.objects.filter(
            agency=agency,
            is_active=True,
            order_pull_enabled=True,
        )
        fbs_enabled = active_profiles.exists()
        context.update(
            {
                "agency": agency,
                "catalog": [catalog[key] for key in ("wb", "ozon")],
                "fbs_enabled": fbs_enabled,
                "profile_count": sum(
                    entry["configured_profiles"] for entry in catalog.values()
                ),
                "global_order_pull_enabled": feature_enabled("order_pull"),
                "sidebar_active": "settings",
                "user_display_name": _head_manager_user_name(self.request),
            }
        )
        return context


class HeadManagerFbsClientSettingsUpdateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, agency_id: int):
        agency = _head_manager_fbs_client_agency(agency_id)
        action = str(request.POST.get("action") or "save").strip()
        if action == "sync":
            self._sync_orders(request, agency)
            if request.POST.get("return_to") == "settings":
                return redirect("/head-manager/fbs/settings/#clients")
            return redirect("head-manager-fbs-client-settings", agency_id=agency.id)
        enabled = request.POST.get("fbs_enabled") == "1"
        catalog = _head_manager_fbs_client_catalog(agency)
        selections = []
        stock_push_keys = set()
        for marketplace in ("wb", "ozon"):
            selected_warehouse_keys = {
                str(value or "").strip()
                for value in request.POST.getlist(f"warehouses_{marketplace}")
                if str(value or "").strip()
            }
            selected_stock_warehouse_keys = {
                str(value or "").strip()
                for value in request.POST.getlist(f"stock_push_warehouses_{marketplace}")
                if str(value or "").strip()
            }
            if selected_stock_warehouse_keys - selected_warehouse_keys:
                messages.error(
                    request,
                    "Выгрузку остатков можно включить только для выбранного FBS-склада.",
                )
                return redirect("head-manager-fbs-client-settings", agency_id=agency.id)
            if selected_warehouse_keys:
                entry = catalog[marketplace]
                if not entry["credentials_configured"]:
                    messages.error(request, "Для выбранного маркетплейса не настроен API-ключ.")
                    return redirect(
                        "head-manager-fbs-client-settings", agency_id=agency.id
                    )
                if marketplace == "ozon" and not entry["client_id_configured"]:
                    messages.error(request, "Для Ozon не настроен Client ID.")
                    return redirect(
                        "head-manager-fbs-client-settings", agency_id=agency.id
                    )
            warehouse_map = {
                str(row.get("selection_key") or row["warehouse_id"]): row
                for row in catalog[marketplace].get("warehouses", [])
            }
            for warehouse_key in selected_warehouse_keys:
                row = warehouse_map.get(str(warehouse_key).strip())
                if row is None:
                    messages.error(request, "Список складов изменился. Обновите страницу.")
                    return redirect(
                        "head-manager-fbs-client-settings", agency_id=agency.id
                    )
                external_account_id = str(row.get("external_account_id") or "").strip()
                selections.append(
                    SelectedMarketplaceWarehouse(
                        marketplace=marketplace,
                        warehouse_id=row["warehouse_id"],
                        name=row["name"],
                        external_account_id=external_account_id,
                    )
                )
                if warehouse_key in selected_stock_warehouse_keys:
                    if marketplace == FbsIntegrationProfile.MARKETPLACE_OZON:
                        stock_push_keys.add(
                            (marketplace, str(row["warehouse_id"]), external_account_id)
                        )
                    else:
                        stock_push_keys.add((marketplace, str(row["warehouse_id"])))
        try:
            result = configure_client_fbs_profiles(
                agency=agency,
                enabled=enabled,
                selections=tuple(selections),
                stock_push_keys=frozenset(stock_push_keys),
            )
        except FbsError as exc:
            messages.error(request, str(exc))
            return redirect("head-manager-fbs-client-settings", agency_id=agency.id)

        messages.success(
            request,
            "Настройки FBS сохранены: "
            f"активных складов {len(result.active_profile_ids)}, "
            f"создано {result.created}, обновлено {result.updated}, выключено {result.disabled}; "
            f"выгрузка остатков включена для {len(result.stock_push_profile_ids)} складов.",
        )
        if action == "save_sync" and result.active_profile_ids:
            self._sync_orders(request, agency)
        return redirect("head-manager-fbs-client-settings", agency_id=agency.id)

    @staticmethod
    def _sync_orders(request, agency: Agency) -> None:
        profiles = list(
            FbsIntegrationProfile.objects.filter(
                agency=agency,
                is_active=True,
                order_pull_enabled=True,
            ).order_by("marketplace", "name", "id")
        )
        if not profiles:
            messages.warning(
                request,
                "Сначала включите FBS и выберите хотя бы один склад клиента.",
            )
            return
        totals = {
            "received": 0,
            "created": 0,
            "updated": 0,
            "duplicate": 0,
            "skipped": 0,
        }
        failures = []
        for profile in profiles:
            try:
                sync_result = pull_profile_orders_manually(
                    profile_id=profile.id,
                    actor=request.user,
                    limit=1000,
                )
            except FbsError as exc:
                failures.append(f"{profile.get_marketplace_display()}: {exc}")
                continue
            for key in totals:
                totals[key] += getattr(sync_result, key)
        if failures:
            messages.warning(
                request,
                f"Обновление выполнено частично. Ошибок: {len(failures)}. {failures[0]}",
            )
            return
        messages.success(
            request,
            "Заказы обновлены вручную: "
            f"получено {totals['received']}, новых {totals['created']}, "
            f"обновлено {totals['updated']}, без изменений {totals['duplicate']}, "
            f"пропущено {totals['skipped']}.",
        )


class HeadManagerFbsStoragePolicyUpdateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request):
        agency = get_object_or_404(
            Agency.objects.filter(fbs_integration_profiles__isnull=False).distinct(),
            pk=request.POST.get("agency_id"),
        )
        billing_mode = str(request.POST.get("billing_mode") or "").strip()
        if billing_mode not in dict(FbsClientStoragePolicy.BILLING_CHOICES):
            messages.error(request, "Выберите режим тарификации хранения FBS.")
            return redirect("head-manager-fbs-settings")
        FbsClientStoragePolicy.objects.update_or_create(
            agency=agency,
            defaults={"billing_mode": billing_mode, "is_active": True},
        )
        messages.success(request, f"Тарификация FBS для {agency} сохранена.")
        return redirect("head-manager-fbs-settings")


class HeadManagerFbsReplenishmentPolicyUpdateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request):
        sku = get_object_or_404(
            SKU.objects.select_related("agency").filter(
                agency__fbs_integration_profiles__isnull=False,
                deleted=False,
            ),
            pk=request.POST.get("sku_id"),
        )
        try:
            minimum_qty = max(int(request.POST.get("minimum_qty") or 0), 0)
            target_qty = max(int(request.POST.get("target_qty") or 0), 0)
        except (TypeError, ValueError):
            messages.error(request, "Порог и целевой остаток должны быть целыми числами.")
            return redirect("head-manager-fbs-settings")
        policy = FbsReplenishmentPolicy(
            agency=sku.agency,
            sku=sku,
            minimum_qty=minimum_qty,
            target_qty=target_qty,
            auto_suggest=request.POST.get("auto_suggest") == "1",
            is_active=True,
        )
        try:
            policy.full_clean(exclude={"id"})
        except ValidationError as exc:
            messages.error(request, str(exc))
            return redirect("head-manager-fbs-settings")
        FbsReplenishmentPolicy.objects.update_or_create(
            agency=sku.agency,
            sku=sku,
            defaults={
                "minimum_qty": minimum_qty,
                "target_qty": target_qty,
                "auto_suggest": policy.auto_suggest,
                "is_active": True,
            },
        )
        messages.success(request, f"Порог подсорта для {sku.sku_code} сохранен.")
        return redirect("head-manager-fbs-settings")


class HeadManagerFbsEquipmentCreateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, kind: str):
        if kind == FbsEquipmentCodeSequence.KIND_WORKSTATION:
            form = FbsWorkstationBatchCreateForm(request.POST)
            create_service = create_fbs_workstations
        elif kind == FbsEquipmentCodeSequence.KIND_CART:
            form = FbsPickingCartBatchCreateForm(request.POST)
            create_service = create_fbs_picking_carts
        else:
            raise Http404
        if not form.is_valid():
            messages.error(request, _fbs_form_error_text(form))
            return redirect("head-manager-fbs-settings")
        try:
            created = create_service(actor=request.user, **form.cleaned_data)
        except FbsEquipmentError as exc:
            messages.error(request, str(exc))
            return redirect("head-manager-fbs-settings")
        messages.success(request, f"Создано объектов FBS: {len(created)}.")
        if request.POST.get("print_labels"):
            query = urlencode(
                {
                    "kind": kind,
                    "ids": ",".join(str(item.pk) for item in created),
                }
            )
            return redirect(f"/head-manager/fbs/settings/labels/?{query}")
        return redirect("head-manager-fbs-settings")


class HeadManagerFbsEquipmentUpdateView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, kind: str, pk: int):
        if kind == FbsEquipmentCodeSequence.KIND_WORKSTATION:
            instance = get_object_or_404(FbsWorkstation, pk=pk)
            form = FbsWorkstationUpdateForm(request.POST, instance=instance)
            update_service = update_fbs_workstation
            anchor = "workstations"
            service_kwargs = {"workstation": instance}
        elif kind == FbsEquipmentCodeSequence.KIND_CART:
            instance = get_object_or_404(FbsPickingCart, pk=pk)
            form = FbsPickingCartUpdateForm(request.POST, instance=instance)
            update_service = update_fbs_picking_cart
            anchor = "carts"
            service_kwargs = {"cart": instance}
        else:
            raise Http404
        if not form.is_valid():
            messages.error(request, _fbs_form_error_text(form))
            return redirect(f"/head-manager/fbs/settings/#{anchor}")
        try:
            update_service(actor=request.user, **service_kwargs, **form.cleaned_data)
        except FbsEquipmentError as exc:
            messages.error(request, str(exc))
            return redirect(f"/head-manager/fbs/settings/#{anchor}")
        messages.success(request, f"Настройки «{instance.name}» сохранены.")
        return redirect(f"/head-manager/fbs/settings/#{anchor}")


class HeadManagerFbsEquipmentLabelsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_equipment_labels.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        kind = str(self.request.GET.get("kind") or "").strip()
        raw_ids = str(self.request.GET.get("ids") or "").strip()
        try:
            ids = tuple(dict.fromkeys(int(value) for value in raw_ids.split(",") if value))
        except ValueError as exc:
            raise Http404 from exc
        if not ids or len(ids) > 100:
            raise Http404
        try:
            labels = list(equipment_label_queryset(kind=kind, ids=ids))
        except FbsEquipmentError as exc:
            raise Http404 from exc
        if len(labels) != len(ids):
            raise Http404
        context.update(
            {
                "kind": kind,
                "kind_title": (
                    "Рабочее место"
                    if kind == FbsEquipmentCodeSequence.KIND_WORKSTATION
                    else "Тележка"
                ),
                "labels": labels,
            }
        )
        return context


class HeadManagerAIInspectionView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/ai_inspection.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        report = build_ai_inspection_report()
        context.update(
            {
                "inspection_report": report,
                "incidents": report["incidents"],
                "inspection_sources": report["sources"],
                "inspection_summary": report["summary"],
                "sidebar_active": "inspection",
                "user_display_name": _head_manager_user_name(self.request),
            }
        )
        return context


class HeadManagerDashboard(RoleRequiredMixin, TemplateView):
    template_name = 'head_manager/dashboard.html'
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        now = timezone.now()
        local_now = timezone.localtime(now)
        today = timezone.localdate()
        selected_warehouse = str(self.request.GET.get("warehouse") or "all").strip() or "all"
        selected_shift = str(self.request.GET.get("shift") or "day").strip() or "day"

        warehouse_codes = list(
            WarehouseLocation.objects.exclude(warehouse_code="")
            .order_by("warehouse_code")
            .values_list("warehouse_code", flat=True)
            .distinct()
        )
        warehouses = [{"code": "all", "label": "Все склады"}]
        warehouses.extend({"code": code, "label": code} for code in warehouse_codes)
        if selected_warehouse not in {item["code"] for item in warehouses}:
            selected_warehouse = "all"
        shifts = [
            {"code": "day", "label": "1 (08:00 - 20:00)"},
            {"code": "night", "label": "2 (20:00 - 08:00)"},
            {"code": "all", "label": "Все смены"},
        ]
        if selected_shift not in {item["code"] for item in shifts}:
            selected_shift = "day"
        if local_now.hour < 12:
            greeting = "Доброе утро"
        elif local_now.hour < 18:
            greeting = "Добрый день"
        else:
            greeting = "Добрый вечер"
        user_display_name = "коллега"
        employee_profile = getattr(self.request.user, "employee_profile", None)
        if self.request.user.is_authenticated and employee_profile:
            user_display_name = employee_profile.full_name
        elif self.request.session.get("employee_name"):
            user_display_name = self.request.session.get("employee_name")
        elif self.request.user.is_authenticated:
            user_display_name = self.request.user.get_full_name() or self.request.user.username

        open_tasks = Task.objects.exclude(status="done")
        overdue_count = open_tasks.filter(due_date__lt=now).count()
        today_count = open_tasks.filter(due_date__date=today).count()
        in_progress_count = open_tasks.filter(status="in_progress").count()
        no_owner_count = open_tasks.filter(assigned_to__isnull=True).count()
        done_today_count = Task.objects.filter(status="done", updated_at__date=today).count()
        done_recent = list(Task.objects.filter(status="done", updated_at__date=today).only("created_at", "updated_at")[:200])
        if done_recent:
            avg_seconds = sum(max((task.updated_at - task.created_at).total_seconds(), 0) for task in done_recent) / len(done_recent)
            avg_hours = int(avg_seconds // 3600)
            avg_minutes = int((avg_seconds % 3600) // 60)
            avg_processing_label = f"{avg_hours} ч {avg_minutes} мин" if avg_hours else f"{avg_minutes} мин"
        else:
            avg_processing_label = "Нет данных"

        active_trip_statuses = [
            LogisticsTrip.STATUS_PLANNED,
            LogisticsTrip.STATUS_LOADING,
            LogisticsTrip.STATUS_DEPARTED,
        ]
        active_trips = LogisticsTrip.objects.filter(status__in=active_trip_statuses)
        today_trips = active_trips.filter(Q(trip_date=today) | Q(trip_date__isnull=True))

        stock_rows = WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        if selected_warehouse != "all":
            stock_rows = stock_rows.filter(location__warehouse_code=selected_warehouse)
        stock_unit = str(self.request.GET.get("stock_unit") or "items").strip()
        if stock_unit not in _HEAD_MANAGER_STOCK_UNITS:
            stock_unit = "items"
        non_storage_zone_filter = Q(zone_code__in=_HEAD_MANAGER_NON_STORAGE_ZONES) | Q(
            location__zone_code__in=_HEAD_MANAGER_NON_STORAGE_ZONES
        )
        reserved_filter = (
            Q(processing_reserved_qty__gt=0)
            | Q(shipping_reserved_qty__gt=0)
            | Q(other_reserved_qty__gt=0)
        )
        non_storage_available_filter = (
            non_storage_zone_filter
            & Q(available_qty__gt=0)
            & Q(processing_reserved_qty=0)
            & Q(shipping_reserved_qty=0)
            & Q(other_reserved_qty=0)
        )
        stock_summary = stock_rows.aggregate(
            total_qty_sum=Sum("qty"),
            available_qty_sum=Sum("available_qty"),
            pallets=Count("parent_container_id", distinct=True),
            boxes=Count("container_id", distinct=True),
            reserved_boxes=Count("container_id", filter=reserved_filter, distinct=True),
            reserved_pallets=Count("parent_container_id", filter=reserved_filter, distinct=True),
            available_boxes=Count("container_id", filter=Q(available_qty__gt=0), distinct=True),
            available_pallets=Count("parent_container_id", filter=Q(available_qty__gt=0), distinct=True),
            without_location_qty=Sum(
                "available_qty",
                filter=non_storage_available_filter,
            ),
            without_location_boxes=Count(
                "container_id",
                filter=non_storage_available_filter,
                distinct=True,
            ),
            without_location_pallets=Count(
                "parent_container_id",
                filter=non_storage_available_filter,
                distinct=True,
            ),
            receiving_qty=Sum("qty", filter=Q(zone_code="PR") | Q(location__zone_code="PR")),
            storage_qty=Sum("qty", filter=Q(zone_code="OS") | Q(location__zone_code="OS")),
            otg_qty=Sum("qty", filter=Q(zone_code="OTG") | Q(location__zone_code="OTG")),
        )
        total_qty = int(stock_summary.get("total_qty_sum") or 0)
        available_qty = int(stock_summary.get("available_qty_sum") or 0)
        reserved_qty = max(total_qty - available_qty, 0)
        stock_unit_suffix = _HEAD_MANAGER_STOCK_UNITS[stock_unit]["suffix"]
        if stock_unit == "boxes":
            stock_metric_values = {
                "total": stock_summary.get("boxes") or 0,
                "available": stock_summary.get("available_boxes") or 0,
                "reserved": stock_summary.get("reserved_boxes") or 0,
                "without_location": stock_summary.get("without_location_boxes") or 0,
            }
        elif stock_unit == "pallets":
            stock_metric_values = {
                "total": stock_summary.get("pallets") or 0,
                "available": stock_summary.get("available_pallets") or 0,
                "reserved": stock_summary.get("reserved_pallets") or 0,
                "without_location": stock_summary.get("without_location_pallets") or 0,
            }
        else:
            stock_metric_values = {
                "total": total_qty,
                "available": available_qty,
                "reserved": reserved_qty,
                "without_location": stock_summary.get("without_location_qty") or 0,
            }
        stock_without_location_count = stock_metric_values["without_location"]
        stock_unit_options = []
        for key, item in _HEAD_MANAGER_STOCK_UNITS.items():
            params = self.request.GET.copy()
            params["stock_unit"] = key
            stock_unit_options.append(
                {
                    "key": key,
                    "label": item["label"],
                    "href": f"?{params.urlencode()}#warehouse-stock",
                    "active": key == stock_unit,
                }
            )
        receiving_qty = stock_summary.get("receiving_qty") or 0
        storage_qty = stock_summary.get("storage_qty") or 0
        otg_qty = stock_summary.get("otg_qty") or 0
        documents_today_count = Task.objects.filter(updated_at__date=today).filter(
            Q(route__contains="/documents/") | Q(title__icontains="документ")
        ).count()

        request_filters = _head_manager_request_filter_context(self.request)
        request_view = str(self.request.GET.get("request_view") or "list").strip()
        if request_view not in {"list", "board"}:
            request_view = "list"
        requests_panel_open = any(
            [
                request_view != "list",
                request_filters["stage"] != "all",
                request_filters["type"] != "all",
                request_filters["q"],
                request_filters["client"],
                request_filters["employee"],
                request_filters["warehouse_filter"],
                request_filters["responsibility"] != "warehouse",
            ]
        )
        (
            request_rows,
            request_type_counts,
            request_stage_counts,
            request_responsibility_counts,
        ) = _head_manager_requests_rows(request_filters)
        request_responsibility_query = (
            ""
            if request_filters["responsibility"] == "warehouse"
            else f"&responsibility={request_filters['responsibility']}"
        )
        dashboard_request_rows = request_rows[:40]
        request_board_groups = []
        for key in ("overdue", "today", "soon", "in_work", "no_owner", "done"):
            group_rows = [row for row in request_rows if row.get("stage") == key]
            request_board_groups.append(
                {
                    "key": key,
                    "label": _HEAD_MANAGER_STAGE_FILTERS[key],
                    "count": len(group_rows),
                    "rows": group_rows[:6],
                    "more": max(len(group_rows) - 6, 0),
                    "href": f"/head-manager/?request_filter={key}&request_view=list{request_responsibility_query}#all-requests",
                }
            )

        status_cards = [
            {
                "label": "Просрочено",
                "value": request_stage_counts.get("overdue", overdue_count),
                "note": "Дедлайн уже прошел",
                "href": f"/head-manager/?request_filter=overdue{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "overdue",
                "tone": "danger",
                "icon": "◷",
            },
            {
                "label": "Сегодня",
                "value": request_stage_counts.get("today", today_count),
                "note": "Срок исполнения сегодня",
                "href": f"/head-manager/?request_filter=today{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "today",
                "tone": "sun",
                "icon": "□",
            },
            {
                "label": "Скоро",
                "value": request_stage_counts.get("soon", 0),
                "note": "Статус скоро",
                "href": f"/head-manager/?request_filter=soon{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "soon",
                "tone": "sun",
                "icon": "!",
            },
            {
                "label": "В работе",
                "value": request_stage_counts.get("in_work", in_progress_count),
                "note": "Активные задачи в работе",
                "href": f"/head-manager/?request_filter=in_work{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "in_work",
                "tone": "blue",
                "icon": "▣",
            },
            {
                "label": "Ожидают людей",
                "value": request_stage_counts.get("no_owner", no_owner_count),
                "note": "Заявки без исполнителя",
                "href": f"/head-manager/?request_filter=no_owner{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "no_owner",
                "tone": "violet",
                "icon": "○",
            },
            {
                "label": "Готово",
                "value": request_stage_counts.get("done", done_today_count),
                "note": "Закрыто сегодня",
                "href": f"/head-manager/?request_filter=done{request_responsibility_query}#all-requests",
                "active": request_filters["stage"] == "done",
                "tone": "green",
                "icon": "✓",
            },
            {
                "label": "Среднее время",
                "value": avg_processing_label,
                "note": "Среднее по закрытым сегодня",
                "href": f"/head-manager/?request_filter=done{request_responsibility_query}#all-requests",
                "active": False,
                "tone": "ink",
                "icon": "◴",
            },
        ]

        problems = [
            {
                "label": "Просроченные задачи",
                "value": overdue_count,
                "href": "/todo/?status=backlog",
                "tone": "danger",
            },
            {
                "label": "Задачи без исполнителя",
                "value": no_owner_count,
                "href": "/todo/",
                "tone": "sun",
            },
            {
                "label": "Остатки без места хранения",
                "value": stock_without_location_count,
                "href": "/head-manager/stock-editor/?f_location=-",
                "tone": "blue",
            },
            {
                "label": "Рейсы на сегодня",
                "value": today_trips.count(),
                "href": "/logistics/trips/",
                "tone": "green",
            },
            {
                "label": "Зона приемки загружена",
                "value": f"{receiving_qty} шт.",
                "href": "/head-manager/stock-editor/?f_location=PR",
                "tone": "violet",
            },
        ]

        trips_today = list(
            today_trips.select_related("assigned_logistician")
            .order_by("trip_date", "number")[:6]
        )
        shipping_order_qs = (
            ShippingOrder.objects.filter(status__in=_HEAD_MANAGER_SHIPPING_WATCH_STATUSES)
            .select_related("agency")
            .order_by("eta_at", "planned_ship_date", "slot_date", "-updated_at", "-pk")
        )
        if selected_warehouse != "all":
            shipping_order_qs = shipping_order_qs.filter(
                Q(destination_warehouse__icontains=selected_warehouse)
                | Q(transit_address__icontains=selected_warehouse)
                | Q(destination_address__icontains=selected_warehouse)
            )
        shipping_orders = list(shipping_order_qs[:8])
        shipping_status_labels = dict(ShippingOrder.STATUS_CHOICES)
        shipping_task_filter = Q()
        for order in shipping_orders:
            shipping_task_filter |= Q(route__contains=f"/shipping/{order.pk}/")
        shipping_tasks_by_order: dict[int, Task] = {}
        if shipping_orders and shipping_task_filter:
            for task in (
                Task.objects.exclude(status="done")
                .filter(shipping_task_filter)
                .select_related("assigned_to")
                .order_by("-updated_at", "due_date")
            ):
                match = re.search(r"/shipping/(\d+)/", str(task.route or ""))
                if match:
                    shipping_tasks_by_order.setdefault(int(match.group(1)), task)
        shipping_watch = []
        for order in shipping_orders:
            task = shipping_tasks_by_order.get(order.pk)
            planned_at = order.eta_at or order.planned_ship_date or order.slot_date
            if hasattr(planned_at, "hour"):
                due_label = timezone.localtime(planned_at).strftime("%d.%m.%Y %H:%M")
            elif planned_at:
                due_label = planned_at.strftime("%d.%m.%Y")
            else:
                due_label = ""
            shipping_watch.append(
                {
                    "title": f"Заявка на отгрузку №{_head_manager_shipping_number(order.number or order.pk)}",
                    "href": f"/shipping/{order.pk}/",
                    "documents": "По заявке",
                    "status": shipping_status_labels.get(order.status, order.status),
                    "stage_label": _HEAD_MANAGER_SHIPPING_STAGE_LABELS.get(order.status, "В контроле"),
                    "executor": getattr(getattr(task, "assigned_to", None), "full_name", "") or "По складской заявке",
                    "client": str(getattr(order.agency, "short_name", "") or getattr(order.agency, "agn_name", "") or "-"),
                    "due": due_label,
                }
            )
        shipping_overdue_count = 0
        for order in shipping_orders:
            planned_at = order.eta_at or order.planned_ship_date or order.slot_date
            if not planned_at:
                continue
            planned_dt = planned_at
            if not hasattr(planned_dt, "hour"):
                planned_dt = timezone.make_aware(datetime.combine(planned_dt, datetime.max.time()))
            if planned_dt < now:
                shipping_overdue_count += 1

        sla_summary = [
            {
                "label": "Приемка",
                "percent": 100 if today_count == 0 else max(0, min(100, round(((today_count - overdue_count) / max(today_count, 1)) * 100))),
                "delta": "по задачам",
                "href": "/orders/receiving/",
                "tone": "green",
            },
            {
                "label": "Отгрузка",
                "percent": 100 if not shipping_orders else max(0, min(100, round(((len(shipping_orders) - shipping_overdue_count) / max(len(shipping_orders), 1)) * 100))),
                "delta": "по контролю",
                "href": "/shipping/",
                "tone": "blue",
            },
            {
                "label": "Документы",
                "percent": 100 if documents_today_count == 0 else 91,
                "delta": "сегодня",
                "href": "/audit/orders/",
                "tone": "sun",
            },
        ]

        today_totals = [
            {"label": "Приемка", "value": f"{receiving_qty} шт.", "note": "в зоне приемки", "icon": "PR", "href": "/stockmap/pr/"},
            {"label": "Отгрузка", "value": f"{otg_qty} шт.", "note": "в зоне отгрузки", "icon": "OT", "href": "/shipping/"},
            {"label": "Размещение", "value": f"{storage_qty} шт.", "note": "в основном складе", "icon": "OS", "href": "/stockmap/"},
            {"label": "Документы", "value": str(documents_today_count), "note": "обновлены сегодня", "icon": "DOC", "href": "/audit/orders/"},
        ]

        reports = [
            {
                "title": "Движение товара для счетов",
                "description": "XLSX: приход, расход, документы, ОВХ и статусы по текущим остаткам.",
                "href": "/head-manager/report/",
                "tag": "Excel",
            },
            {
                "title": "Складские остатки",
                "description": "Рабочий журнал остатков с фильтрами, правками и выгрузкой в Excel.",
                "href": "/head-manager/stock-editor/",
                "tag": "Журнал",
            },
            {
                "title": "История заявок",
                "description": "Аудит изменений по заявкам приемки, обработки, отгрузки и прочим задачам.",
                "href": "/audit/orders/",
                "tag": "Аудит",
            },
            {
                "title": "Движения склада",
                "description": "Складские события и перемещения по контейнерам, зонам и операциям.",
                "href": "/audit/moves/",
                "tag": "Аудит",
            },
            {
                "title": "SLA чатов",
                "description": "Без ответа, просрочки, нагрузка и контроль клиентских коммуникаций.",
                "href": "/team-manager/chats/sla/",
                "tag": "SLA",
            },
            {
                "title": "Избыточные действия",
                "description": "Корректировки, ручные вмешательства и действия сверх обычного процесса.",
                "href": "/audit/overactions/",
                "tag": "Контроль",
            },
        ]

        ctx.update(
            {
                "today_label": today.strftime("%d.%m.%Y"),
                "dashboard_updated_at": local_now.strftime("%H:%M"),
                "greeting": greeting,
                "user_display_name": user_display_name,
                "warehouses": warehouses,
                "selected_warehouse": selected_warehouse,
                "shifts": shifts,
                "selected_shift": selected_shift,
                "status_cards": status_cards,
                "dashboard_problems": problems,
                "dashboard_problem_count": sum(
                    1
                    for item in problems
                    if int(re.sub(r"\D+", "", str(item.get("value") or "0")) or 0) > 0
                ),
                "trips_today": trips_today,
                "shipping_watch": shipping_watch,
                "sla_summary": sla_summary,
                "stock_summary": {
                    "total_value": stock_metric_values["total"],
                    "available_value": stock_metric_values["available"],
                    "reserved_value": stock_metric_values["reserved"],
                    "pallets": stock_summary.get("pallets") or 0,
                    "boxes": stock_summary.get("boxes") or 0,
                    "without_location": stock_without_location_count,
                    "unit": stock_unit,
                    "unit_suffix": stock_unit_suffix,
                    "unit_options": stock_unit_options,
                },
                "request_rows": dashboard_request_rows,
                "request_total_count": len(request_rows),
                "request_view": request_view,
                "requests_panel_open": requests_panel_open,
                "request_filters": request_filters,
                "request_types": [
                    {"key": key, "label": label, "count": request_type_counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_REQUEST_TYPES.items()
                    if key != "task" or request_type_counts.get(key, 0)
                ],
                "request_stage_filters": [
                    {"key": key, "label": label, "count": request_stage_counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_STAGE_FILTERS.items()
                ],
                "request_responsibility_filters": [
                    {"key": key, "label": label, "count": request_responsibility_counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_RESPONSIBILITY_FILTERS.items()
                ],
                "request_board_groups": request_board_groups,
                "today_totals": today_totals,
                "dashboard_reports": reports,
            }
        )
        return ctx


class HeadManagerReportsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/reports.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        categories = _head_manager_report_categories()
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "categories": categories,
                "total_reports": sum(len(category["reports"]) for category in categories),
            }
        )
        return ctx


class HeadManagerReportCategoryView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/report_category.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        category = _head_manager_report_category(self.kwargs["category_slug"])
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "category": category,
                "categories": _head_manager_report_categories(),
            }
        )
        return ctx


class HeadManagerReportDetailView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/report_detail.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        category, report = _head_manager_report(self.kwargs["category_slug"], self.kwargs["report_slug"])
        active_filters = {
            key: str(self.request.GET.get(key) or "").strip()
            for key in [item["key"] for item in _HEAD_MANAGER_REPORT_FILTERS]
        }
        filter_schema = [
            {
                **item,
                "value": active_filters.get(item["key"], ""),
            }
            for item in _HEAD_MANAGER_REPORT_FILTERS
        ]
        has_filters = any(active_filters.values())
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "category": category,
                "report": report,
                "filter_schema": filter_schema,
                "active_filters": active_filters,
                "has_filters": has_filters,
                "features": _HEAD_MANAGER_REPORT_FEATURES,
                "presets": ["Сегодня", "Текущая смена", "7 дней", "30 дней"],
                "sample_columns": ["Дата", "Склад", "Клиент", "SKU", "Операция", "Статус", "Количество"],
            }
        )
        return ctx


class HeadManagerSkuMovementReportView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/sku_movement_report.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        filters = _head_manager_sku_movement_filters(self.request)
        rows, summary = _head_manager_sku_movement_rows(filters)
        query_string = self.request.GET.urlencode()
        export_url = "/head-manager/reports/products/sku-movement/export/"
        if query_string:
            export_url = f"{export_url}?{query_string}"
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "filters": filters,
                "rows": rows,
                "summary": summary,
                "export_url": export_url,
                "category": _head_manager_report_category("products"),
                "report": {
                    "title": "Движение товаров",
                    "subtitle": "Отчет по артикулу: приход, перемещения, отгрузки и текущий остаток.",
                },
            }
        )
        return ctx


class HeadManagerSkuMovementReportExportView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def get(self, request, *args, **kwargs):
        filters = _head_manager_sku_movement_filters(request)
        return _head_manager_sku_movement_export_response(filters)


class HeadManagerFbsEmployeeReportView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/fbs_employee_report.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        filters = _head_manager_fbs_employee_report_filters(self.request)
        rows, summary = _head_manager_fbs_employee_report_rows(filters)
        overview = _head_manager_fbs_employee_report_overview(filters)
        query_string = self.request.GET.urlencode()
        export_url = "/head-manager/reports/employees/fbs-productivity/export/"
        if query_string:
            export_url = f"{export_url}?{query_string}"
        today = timezone.localdate()
        ctx.update(
            {
                "today_label": today.strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "filters": filters,
                "rows": rows,
                "summary": summary,
                "overview": overview,
                "export_url": export_url,
                "today_iso": today.isoformat(),
                "week_from_iso": (today - timedelta(days=6)).isoformat(),
                "month_from_iso": (today - timedelta(days=29)).isoformat(),
                "category": _head_manager_report_category("employees"),
                "report": {
                    "title": "Операционный отчёт FBS",
                    "subtitle": "Поступившие заказы, фактическая сборка и результат каждого сотрудника.",
                },
            }
        )
        return ctx


class HeadManagerFbsEmployeeReportExportView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def get(self, request, *args, **kwargs):
        filters = _head_manager_fbs_employee_report_filters(request)
        return _head_manager_fbs_employee_report_export_response(filters)


class HeadManagerWarehouseEmployeeReportView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/warehouse_employee_report.html"
    allowed_roles = _WAREHOUSE_EMPLOYEE_ALLOWED_ROLES

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        filters = _head_manager_warehouse_employee_report_filters(self.request)
        rows, summary = _head_manager_warehouse_employee_report_rows(filters)
        query_string = self.request.GET.urlencode()
        export_url = "/head-manager/reports/employees/warehouse-productivity/export/"
        if query_string:
            export_url = f"{export_url}?{query_string}"
        today = timezone.localdate()
        ctx.update(
            {
                "today_label": today.strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "filters": filters,
                "rows": rows,
                "summary": summary,
                "export_url": export_url,
                "today_iso": today.isoformat(),
                "week_from_iso": (today - timedelta(days=6)).isoformat(),
                "month_from_iso": (today - timedelta(days=29)).isoformat(),
                "metric_groups": _WAREHOUSE_EMPLOYEE_REPORT_GROUPS,
                "role_choices": Employee.ROLE_CHOICES,
                "category": _head_manager_report_category("employees"),
                "report": {
                    "title": "Работа склада по сотрудникам",
                    "subtitle": "Фактически завершённые действия по складским контурам и след активности в системе.",
                },
            }
        )
        return ctx


class HeadManagerWarehouseEmployeeReportExportView(RoleRequiredMixin, View):
    allowed_roles = _WAREHOUSE_EMPLOYEE_ALLOWED_ROLES

    def get(self, request, *args, **kwargs):
        filters = _head_manager_warehouse_employee_report_filters(request)
        return _head_manager_warehouse_employee_report_export_response(filters)


class HeadManagerEmployeeActivityDirectoryView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/employee_activity_directory.html"
    allowed_roles = _WAREHOUSE_EMPLOYEE_ALLOWED_ROLES

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        filters = _head_manager_employee_activity_filters(self.request)
        rows, summary = _head_manager_employee_activity_directory_rows(filters)
        today = timezone.localdate()
        ctx.update(
            {
                "today_label": today.strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "filters": filters,
                "rows": rows,
                "summary": summary,
                "today_iso": today.isoformat(),
                "week_from_iso": (today - timedelta(days=6)).isoformat(),
                "month_from_iso": (today - timedelta(days=29)).isoformat(),
                "role_choices": Employee.ROLE_CHOICES,
            }
        )
        return ctx


class HeadManagerEmployeeActivityDetailView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/employee_activity_detail.html"
    allowed_roles = _WAREHOUSE_EMPLOYEE_ALLOWED_ROLES

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        employee = get_object_or_404(
            Employee.objects.select_related("user"),
            pk=self.kwargs["employee_id"],
        )
        filters = _head_manager_employee_activity_filters(self.request)
        activity = _head_manager_employee_activity_detail(employee, filters)
        today = timezone.localdate()
        ctx.update(
            {
                "today_label": today.strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "employee": employee,
                "employee_role": str(employee.get_role_display() or "").strip() or "—",
                "employee_login": (
                    str(employee.user.get_username() or "").strip()
                    if employee.user_id
                    else "Нет учётной записи"
                ),
                "filters": filters,
                "activity": activity,
                "today_iso": today.isoformat(),
                "week_from_iso": (today - timedelta(days=6)).isoformat(),
                "month_from_iso": (today - timedelta(days=29)).isoformat(),
            }
        )
        return ctx


class HeadManagerContainerReportView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/container_report.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        kind = self.kwargs["kind"]
        report = _head_manager_container_report_config(kind)
        filters = _head_manager_container_filters(self.request)
        rows = _head_manager_container_report_rows(kind, filters)
        summary = _head_manager_container_report_summary(rows)
        query_string = self.request.GET.urlencode()
        export_url = f"/head-manager/reports/warehouse/{kind}/export/"
        if query_string:
            export_url = f"{export_url}?{query_string}"
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "kind": kind,
                "report": report,
                "filters": filters,
                "rows": rows,
                "summary": summary,
                "export_url": export_url,
                "category": _head_manager_report_category("warehouse"),
            }
        )
        return ctx


class HeadManagerContainerReportExportView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def get(self, request, *args, **kwargs):
        filters = _head_manager_container_filters(request)
        return _head_manager_container_export_response(kwargs["kind"], filters)


class HeadManagerRequestsView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/requests.html"
    allowed_roles = ("head_manager",)

    @staticmethod
    def _date_param(value: str):
        value = str(value or "").strip()
        if not value:
            return None
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        filters = _head_manager_request_filter_context(self.request)
        rows, counts, stage_counts, responsibility_counts = _head_manager_requests_rows(filters)
        paginator = Paginator(rows, 50)
        page_obj = paginator.get_page(self.request.GET.get("page") or 1)
        query_params = self.request.GET.copy()
        query_params.pop("page", None)
        status_choices = []
        seen_statuses = set()
        for value, label in list(Task.STATUS_CHOICES) + list(ShippingOrder.STATUS_CHOICES) + list(OtherRequest.STATUS_CHOICES):
            if value in seen_statuses:
                continue
            seen_statuses.add(value)
            status_choices.append((value, label))
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "rows": page_obj.object_list,
                "page_obj": page_obj,
                "request_types": [
                    {"key": key, "label": label, "count": counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_REQUEST_TYPES.items()
                    if key != "task" or counts.get(key, 0)
                ],
                "request_stage_filters": [
                    {"key": key, "label": label, "count": stage_counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_STAGE_FILTERS.items()
                ],
                "request_responsibility_filters": [
                    {"key": key, "label": label, "count": responsibility_counts.get(key, 0)}
                    for key, label in _HEAD_MANAGER_RESPONSIBILITY_FILTERS.items()
                ],
                "selected_type": filters["type"],
                "selected_stage": filters["stage"],
                "selected_responsibility": filters["responsibility"],
                "filters": {
                    "status": filters["status"],
                    "responsibility": filters["responsibility"],
                    "client": filters["client"],
                    "employee": filters["employee"],
                    "warehouse_filter": filters["warehouse_filter"],
                    "zone": filters["zone"],
                    "marketplace": filters["marketplace"],
                    "document_status": filters["document_status"],
                    "q": filters["q"],
                    "date_from": str(filters["date_from"] or ""),
                    "date_to": str(filters["date_to"] or ""),
                },
                "query_string": query_params.urlencode(),
                "status_choices": status_choices,
            }
        )
        return ctx


class HeadManagerMovementReportView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/movement_report.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        stock_rows = WarehouseStockSnapshot.objects.filter(qty__gt=0, is_archived=False)
        summary = stock_rows.aggregate(
            total_qty=Sum("qty"),
            pallets=Count("parent_container_id", distinct=True),
            boxes=Count("container_id", distinct=True),
        )
        preview_rows = _movement_report_rows(limit=80)
        ctx.update(
            {
                "today_label": timezone.localdate().strftime("%d.%m.%Y"),
                "updated_at": timezone.localtime(timezone.now()).strftime("%H:%M"),
                "export_url": "/head-manager/report/export/",
                "preview_rows": preview_rows,
                "total_rows": stock_rows.count(),
                "summary": {
                    "total_qty": summary.get("total_qty") or 0,
                    "pallets": summary.get("pallets") or 0,
                    "boxes": summary.get("boxes") or 0,
                },
            }
        )
        return ctx


class HeadManagerMovementReportExportView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def get(self, request, *args, **kwargs):
        return _movement_report_response()


class HeadManagerStockEditorView(RoleRequiredMixin, TemplateView):
    allowed_roles = ("head_manager",)
    _GOODS_TYPE_ALIASES = {
        "op": "op",
        "оптовый": "op",
        "gv": "gv",
        "готовый": "gv",
        "готовые": "gv",
        "br": "br",
        "брак": "br",
        "vz": "vz",
        "возврат": "vz",
        "rh": "rh",
        "расходный": "rh",
        "расходные": "rh",
        "no": "no",
        "не обработанный": "no",
        "необработанный": "no",
        "votg": "votg",
        "возврат с отгрузки": "votg",
    }
    _OS_ROW_SECTIONS = {
        1: 9,
        2: 9,
        3: 9,
        4: 9,
        5: 9,
        6: 8,
        7: 6,
        8: 6,
        9: 6,
        10: 6,
    }
    _OS_TIERS = 4
    _OS_CELLS_PER_TIER = 3
    _OS_TIER_OVERRIDES = {
        (4, 9): 5,
        (5, 9): 5,
    }
    _OS_PASSAGE_POSITIONS = {
        (4, section_number, tier_number)
        for section_number in range(1, 7)
        for tier_number in (1, 2)
    }

    def get(self, request, *args, **kwargs):
        from django.shortcuts import render
        from sklad.ui_services import build_inventory_journal_page

        if str(request.GET.get("export") or "").strip().lower() == "xlsx":
            return self._export_inventory_excel(request)
        page = build_inventory_journal_page(request=request)
        context = page["context"]
        paginator = Paginator(context.get("rows") or [], 100)
        page_obj = paginator.get_page(request.GET.get("page"))
        query_params = request.GET.copy()
        query_params.pop("page", None)
        context["rows"] = list(page_obj.object_list)
        context["page_obj"] = page_obj
        context["page_query"] = query_params.urlencode()
        export_params = request.GET.copy()
        export_params.pop("page", None)
        export_params["export"] = "xlsx"
        context["export_excel_url"] = f"{request.path}?{export_params.urlencode()}"
        return render(request, "head_manager/stock_editor.html", context)

    def _export_inventory_excel(self, request):
        from sklad.ui_services import build_inventory_journal_page

        page = build_inventory_journal_page(request=request)
        rows = list(page.get("context", {}).get("rows") or [])

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Остатки"
        processing_source_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
        sheet.append(
            [
                "Дата",
                "Паллета",
                "Короб",
                "Заявка",
                "Клиент",
                "SKU",
                "Наименование",
                "Размер",
                "Вес, г",
                "Тип товара",
                "Место",
                "Всего",
                "Основной",
                "Склад обр",
                "Склад отг",
                "Доступно",
                "Резерв обр",
                "В обработке",
                "Резерв отг",
            ]
        )
        for row in rows:
            created_at = row.get("created_at")
            created_at_text = (
                timezone.localtime(created_at).strftime("%d.%m.%Y %H:%M")
                if created_at
                else ""
            )
            sheet.append(
                [
                    created_at_text,
                    row.get("pallet_code") or "",
                    row.get("box_code") or "",
                    row.get("order_display") or "",
                    row.get("client_label") or "",
                    row.get("sku") or "",
                    row.get("name") or "",
                    row.get("box_size") or "",
                    row.get("box_weight") or "",
                    row.get("goods_type") or "",
                    row.get("location_short") or row.get("location") or "",
                    int(row.get("qty") or 0),
                    int(row.get("stock_main_qty") or 0),
                    int(row.get("stock_processing_qty") or 0),
                    int(row.get("stock_otg_qty") or 0),
                    int(row.get("available_qty") or 0),
                    int(row.get("processing_reserved_qty") or 0),
                    int(row.get("processing_in_progress_qty") or 0),
                    int(row.get("shipping_reserved_qty") or 0),
                ]
            )
            if row.get("is_processing_source_box"):
                cell = sheet.cell(row=sheet.max_row, column=3)
                if str(cell.value or "").strip():
                    cell.fill = processing_source_fill

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="head_manager_stock_{timezone.localdate().isoformat()}.xlsx"'
        )
        return response

    def post(self, request, *args, **kwargs):
        try:
            payload = json.loads(request.body.decode("utf-8"))
        except (TypeError, ValueError, json.JSONDecodeError):
            return JsonResponse({"ok": False, "message": "Не удалось прочитать данные редактирования."}, status=400)

        snapshot_id = int(payload.get("snapshot_id") or 0)
        field_name = str(payload.get("field") or "").strip()
        apply_mode = str(payload.get("apply_mode") or "single").strip().lower()
        raw_value = payload.get("value")

        allowed_fields = {
            "pallet_code",
            "move_box_to_pallet",
            "move_box_to_new_pallet",
            "merge_pallets",
            "rename_pallet",
            "box_code",
            "qty",
            "location",
            "box_metrics",
            "goods_type",
        }
        if not snapshot_id or field_name not in allowed_fields:
            return JsonResponse({"ok": False, "message": "Переданы неполные данные для редактирования."}, status=400)
        if field_name == "box_metrics" and apply_mode not in {"single", "analog"}:
            return JsonResponse({"ok": False, "message": "РџРµСЂРµРґР°РЅС‹ РЅРµРїРѕР»РЅС‹Рµ РґР°РЅРЅС‹Рµ РґР»СЏ СЂРµРґР°РєС‚РёСЂРѕРІР°РЅРёСЏ."}, status=400)

        success_message = "Изменения сохранены."
        try:
            with transaction.atomic():
                snapshot = (
                    WarehouseStockSnapshot.objects
                    .select_for_update(of=("self",))
                    .select_related("container", "parent_container", "location")
                    .get(pk=snapshot_id)
                )
                lock_reason = self._get_snapshot_edit_lock_reason(snapshot)
                if lock_reason:
                    return JsonResponse({"ok": False, "message": lock_reason}, status=409)

                updated_snapshot_ids = [snapshot.pk]
                if field_name in {"pallet_code", "move_box_to_pallet"}:
                    updated_snapshot_ids = self._move_box_to_pallet(snapshot, raw_value, request=request)
                    success_message = "Короб перенесён на выбранную паллету."
                elif field_name == "move_box_to_new_pallet":
                    updated_snapshot_ids = self._move_box_to_new_pallet(snapshot, raw_value, request=request)
                    success_message = "Новая паллета создана. Короб перенесён."
                elif field_name == "merge_pallets":
                    updated_snapshot_ids = self._merge_pallets(snapshot, raw_value, request=request)
                    success_message = "Паллеты объединены. Все короба и товар перенесены."
                elif field_name == "rename_pallet":
                    updated_snapshot_ids = self._rename_pallet(snapshot, raw_value, request=request)
                    success_message = "Номер паллеты изменён. Товар не перемещался."
                elif field_name == "box_code":
                    self._update_snapshot_box(snapshot, raw_value)
                elif field_name == "qty":
                    self._update_snapshot_qty(snapshot, raw_value)
                elif field_name == "box_metrics":
                    updated_snapshot_ids = self._update_snapshot_box_metrics(snapshot, raw_value, apply_mode=apply_mode)
                elif field_name == "goods_type":
                    self._update_snapshot_goods_type(snapshot, raw_value, request=request)
                    success_message = "Тип товара изменён. Запись добавлена в историю."
                else:
                    self._update_snapshot_location(snapshot, raw_value)
        except WarehouseStockSnapshot.DoesNotExist:
            return JsonResponse({"ok": False, "message": "Строка складских остатков не найдена."}, status=404)
        except IntegrityError:
            return JsonResponse(
                {"ok": False, "message": "Такой номер уже используется. Обновите страницу и повторите операцию."},
                status=409,
            )
        except ValueError as exc:
            return JsonResponse({"ok": False, "message": str(exc)}, status=400)

        refreshed_snapshots = list(
            WarehouseStockSnapshot.objects.select_related(
                "container",
                "parent_container",
                "location",
            ).filter(pk__in=updated_snapshot_ids).order_by("id")
        )
        if not refreshed_snapshots:
            return JsonResponse({"ok": False, "message": "Не удалось обновить строку."}, status=500)
        rows_payload = [self._serialize_editor_row(item) for item in refreshed_snapshots]
        return JsonResponse(
            {
                "ok": True,
                "message": success_message,
                "row": rows_payload[0],
                "rows": rows_payload,
            }
        )

    def _get_snapshot_edit_lock_reason(self, snapshot: WarehouseStockSnapshot) -> str:
        row = normalize_stock_row_from_snapshot(snapshot)
        return _journal_edit_lock_reason(row)

    def _serialize_editor_row(self, snapshot: WarehouseStockSnapshot):
        row = normalize_stock_row_from_snapshot(snapshot)
        state_code = str(row.get("warehouse_state_code") or "").strip().lower()
        box_code = row.get("box_code") or "-"
        is_processing_source_box = state_code in _MOVEMENT_REPORT_PROCESSING_SOURCE_STATES and box_code != "-"
        return {
            "snapshot_id": int(row.get("id") or 0),
            "container_id": int(row.get("container_id") or 0),
            "parent_container_id": int(row.get("parent_container_id") or 0),
            "pallet_code": row.get("pallet_code") or "-",
            "box_code": box_code,
            "is_processing_source_box": is_processing_source_box,
            "qty": row.get("qty") or 0,
            "available_qty": row.get("available_qty") or 0,
            "location": row.get("location") or "",
            "location_short": row.get("location_short") or "",
            "box_size": row.get("box_size") or "",
            "box_weight": row.get("box_weight") or "",
            "goods_type": row.get("goods_type") or "",
            "edit_lock_reason": self._get_snapshot_edit_lock_reason(snapshot),
        }

    def _normalize_box_metrics_payload(self, raw_value):
        if not isinstance(raw_value, dict):
            raise ValueError("Переданы неполные данные ОВХ.")

        normalized = {}
        field_map = {
            "gross_weight_g": "Вес",
            "width_mm": "Ширина",
            "height_mm": "Высота",
            "depth_mm": "Глубина",
        }
        for field_name, field_label in field_map.items():
            value = str(raw_value.get(field_name) or "").strip()
            if not value:
                raise ValueError(f"Поле «{field_label}» обязательно.")
            try:
                parsed_value = int(value)
            except (TypeError, ValueError):
                raise ValueError(f"Поле «{field_label}» должно быть числом.")
            if parsed_value <= 0:
                raise ValueError(f"Поле «{field_label}» должно быть больше нуля.")
            normalized[field_name] = parsed_value
        return normalized

    def _build_box_signature(self, snapshot_rows):
        signature = []
        for item in snapshot_rows:
            signature.append(
                (
                    str(item.sku_code or "").strip(),
                    str(item.name or "").strip(),
                    str(item.barcode or "").strip(),
                    str(item.goods_type or "").strip(),
                    int(item.qty or 0),
                )
            )
        signature.sort()
        return tuple(signature)

    def _find_analog_box_container_ids(self, snapshot: WarehouseStockSnapshot):
        container_id = int(snapshot.container_id or 0)
        parent_container_id = int(snapshot.parent_container_id or 0)
        if not container_id:
            return []
        if not parent_container_id:
            return [container_id]

        candidate_rows = (
            WarehouseStockSnapshot.objects.filter(
                agency=snapshot.agency,
                is_archived=False,
                parent_container_id=parent_container_id,
                container_id__isnull=False,
            )
            .exclude(container_id=0)
            .order_by("container_id", "id")
        )
        grouped_rows = {}
        for item in candidate_rows:
            current_container_id = int(item.container_id or 0)
            grouped_rows.setdefault(current_container_id, []).append(item)

        target_signature = self._build_box_signature(grouped_rows.get(container_id, []))
        if not target_signature:
            return [container_id]

        matched_ids = []
        for current_container_id, rows in grouped_rows.items():
            if self._build_box_signature(rows) == target_signature:
                matched_ids.append(current_container_id)
        return matched_ids or [container_id]

    def _update_snapshot_box_metrics(self, snapshot: WarehouseStockSnapshot, raw_value, apply_mode: str = "single"):
        container_id = int(snapshot.container_id or 0)
        if not container_id:
            raise ValueError("Для этой строки не найден короб.")

        box_metrics = self._normalize_box_metrics_payload(raw_value)
        target_container_ids = [container_id]
        if apply_mode == "analog":
            target_container_ids = self._find_analog_box_container_ids(snapshot)

        containers = list(
            WarehouseContainer.objects.filter(
                agency=snapshot.agency,
                id__in=target_container_ids,
                container_type=WarehouseContainer.TYPE_BOX,
            )
        )
        if not containers:
            raise ValueError("Не удалось найти короб для сохранения ОВХ.")

        now = timezone.now()
        for container in containers:
            container.gross_weight_g = box_metrics["gross_weight_g"]
            container.width_mm = box_metrics["width_mm"]
            container.height_mm = box_metrics["height_mm"]
            container.depth_mm = box_metrics["depth_mm"]
            container.updated_at = now
        WarehouseContainer.objects.bulk_update(
            containers,
            ["gross_weight_g", "width_mm", "height_mm", "depth_mm", "updated_at"],
        )

        return list(
            WarehouseStockSnapshot.objects.filter(
                agency=snapshot.agency,
                is_archived=False,
                container_id__in=[container.id for container in containers],
            ).values_list("id", flat=True)
        )

    def _is_valid_os_editor_cell(self, row_no: int, section_no: int, tier_no: int, cell_no: int) -> bool:
        row_no = int(row_no or 0)
        section_no = int(section_no or 0)
        tier_no = int(tier_no or 0)
        cell_no = int(cell_no or 0)
        sections_total = int(self._OS_ROW_SECTIONS.get(row_no) or 0)
        if row_no <= 0 or section_no <= 0 or tier_no <= 0 or cell_no <= 0:
            return False
        if section_no > sections_total:
            return False
        if cell_no > self._OS_CELLS_PER_TIER:
            return False
        max_tiers = int(self._OS_TIER_OVERRIDES.get((row_no, section_no)) or self._OS_TIERS)
        if tier_no > max_tiers:
            return False
        return (row_no, section_no, tier_no) not in self._OS_PASSAGE_POSITIONS

    def _resolve_editor_location(self, snapshot: WarehouseStockSnapshot, raw_value: str) -> WarehouseLocation | None:
        value = str(raw_value or "").strip()
        if not value:
            return None
        location = WarehouseLocation.objects.filter(
            Q(display_name__iexact=value)
            | Q(location_code__iexact=value)
            | Q(zone_code__iexact=value)
        ).first()
        if location:
            return location

        normalized = re.sub(r"\s+", "", value.upper().replace("\u00b7", "-"))
        normalized = re.sub(r"^OS[-:]*", "", normalized)
        match = re.match(r"^(?P<section>[0A-ZА-Я])-(?P<row>\d+)/(?P<tier>\d+)-(?P<cell>\d+)$", normalized)
        if not match:
            return None

        section_map = {
            "0": 1,
            "A": 2,
            "А": 2,
            "B": 3,
            "В": 3,
            "C": 4,
            "С": 4,
            "D": 5,
            "Д": 5,
            "E": 6,
            "Е": 6,
            "F": 7,
            "Ф": 7,
            "G": 8,
            "Г": 8,
            "I": 9,
            "И": 9,
        }
        section_no = section_map.get(match.group("section"))
        if not section_no:
            return None
        row_no = int(match.group("row") or 0)
        tier_no = int(match.group("tier") or 0)
        cell_no = int(match.group("cell") or 0)
        warehouse_code = str(getattr(snapshot.location, "warehouse_code", "") or "MSK").strip() or "MSK"

        location = WarehouseLocation.objects.filter(
            warehouse_code=warehouse_code,
            zone_code="OS",
            row_no=row_no,
            section_no=section_no,
            tier_no=tier_no,
            cell_no=cell_no,
        ).first()
        if location:
            return location
        if not self._is_valid_os_editor_cell(row_no, section_no, tier_no, cell_no):
            return None
        short_code = f"{match.group('section').upper()}-{row_no}/{tier_no}-{cell_no}"
        location, _ = WarehouseLocation.objects.get_or_create(
            warehouse_code=warehouse_code,
            zone_code="OS",
            row_no=row_no,
            section_no=section_no,
            tier_no=tier_no,
            cell_no=cell_no,
            defaults={
                "zone_kind": WarehouseLocation.ZONE_KIND_STORAGE,
                "location_code": short_code,
                "display_name": f"OS · {short_code}",
                "is_active": True,
                "is_pickable": True,
                "is_storage": True,
            },
        )
        return location

    def _update_snapshot_pallet(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите номер паллеты.")
        pallet = snapshot.parent_container
        if not pallet and snapshot.container and snapshot.container.container_type in {
            WarehouseContainer.TYPE_PALLET,
            WarehouseContainer.TYPE_MIXED_PALLET,
        }:
            pallet = snapshot.container
        if not pallet:
            raise ValueError("Для этой строки паллета не найдена.")
        pallet.container_code = value
        pallet.save(update_fields=["container_code", "updated_at"])

    def _update_snapshot_box(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите номер короба.")
        if not snapshot.container_id:
            raise ValueError("Для этой строки короб не найден.")
        snapshot.container.container_code = value
        snapshot.container.save(update_fields=["container_code", "updated_at"])

    def _update_snapshot_qty(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value_text = str(raw_value or "").strip()
        if not value_text:
            raise ValueError("Укажите количество в коробе.")
        try:
            value = int(value_text)
        except (TypeError, ValueError):
            raise ValueError("Количество в коробе должно быть целым числом.")
        if value < 0:
            raise ValueError("Количество в коробе не может быть отрицательным.")
        snapshot.qty = value
        snapshot.save(update_fields=["qty", "updated_at"])

    def _update_snapshot_location(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите место хранения.")
        location = WarehouseLocation.objects.filter(
            Q(display_name__iexact=value)
            | Q(location_code__iexact=value)
            | Q(zone_code__iexact=value)
        ).first()
        if not location:
            raise ValueError("Место хранения не найдено.")
        snapshot.location = location
        snapshot.save(update_fields=["location", "updated_at"])
        if snapshot.container_id:
            snapshot.container.current_location = location
            snapshot.container.save(update_fields=["current_location", "updated_at"])
        if snapshot.parent_container_id:
            snapshot.parent_container.current_location = location
            snapshot.parent_container.save(update_fields=["current_location", "updated_at"])

    def _find_active_pallet(self, agency, raw_value) -> WarehouseContainer:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите номер паллеты.")
        pallets = list(
            WarehouseContainer.objects
            .select_for_update(of=("self",))
            .select_related("current_location")
            .filter(
                agency=agency,
                container_code__iexact=value,
                container_type__in=[
                    WarehouseContainer.TYPE_PALLET,
                    WarehouseContainer.TYPE_MIXED_PALLET,
                ],
                status=WarehouseContainer.STATUS_ACTIVE,
            )
            .order_by("id")[:2]
        )
        if not pallets:
            raise ValueError("Паллета с таким номером не найдена.")
        if len(pallets) > 1:
            raise ValueError("Найдено несколько паллет с таким номером. Укажите точный регистр номера.")
        return pallets[0]

    def _get_source_pallet(self, snapshot: WarehouseStockSnapshot) -> WarehouseContainer:
        pallet_id = int(snapshot.parent_container_id or 0)
        if not pallet_id and snapshot.container_id and snapshot.container.container_type in {
            WarehouseContainer.TYPE_PALLET,
            WarehouseContainer.TYPE_MIXED_PALLET,
        }:
            pallet_id = int(snapshot.container_id)
        if not pallet_id:
            raise ValueError("Для этой строки паллета не найдена.")
        try:
            pallet = (
                WarehouseContainer.objects
                .select_for_update(of=("self",))
                .select_related("current_location")
                .get(pk=pallet_id, agency=snapshot.agency)
            )
        except WarehouseContainer.DoesNotExist:
            raise ValueError("Для этой строки паллета не найдена.")
        if pallet.container_type not in {
            WarehouseContainer.TYPE_PALLET,
            WarehouseContainer.TYPE_MIXED_PALLET,
        }:
            raise ValueError("Для этой строки паллета не найдена.")
        if pallet.status != WarehouseContainer.STATUS_ACTIVE:
            raise ValueError("Эта паллета уже не активна. Обновите страницу.")
        return pallet

    def _ensure_editor_snapshots_unlocked(self, snapshots) -> None:
        for item in snapshots:
            lock_reason = self._get_snapshot_edit_lock_reason(item)
            if lock_reason:
                code = str(getattr(item.container, "container_code", "") or item.container_code or item.pk).strip()
                raise ValueError(f"Операция отменена для позиции {code}: {lock_reason}")

    def _lock_pallet_snapshots(self, pallet: WarehouseContainer):
        return list(
            WarehouseStockSnapshot.objects
            .select_for_update(of=("self",))
            .select_related("container", "parent_container", "location", "active_operation", "last_event", "sku_ref")
            .filter(agency=pallet.agency, is_archived=False)
            .filter(Q(parent_container=pallet) | Q(container=pallet))
            .order_by("id")
        )

    def _record_stock_editor_event(
        self,
        *,
        request,
        snapshot: WarehouseStockSnapshot,
        event_type: str,
        container: WarehouseContainer,
        from_location,
        to_location,
        qty: int,
        payload: dict,
    ) -> WarehouseEvent:
        user = request.user if request.user.is_authenticated else None
        return WarehouseEvent.objects.create(
            agency=snapshot.agency,
            event_type=event_type,
            stock_context_type="stock_editor",
            stock_context_id=str(snapshot.pk),
            container=container,
            from_location=from_location,
            to_location=to_location,
            from_zone_code=str(getattr(from_location, "zone_code", "") or ""),
            to_zone_code=str(getattr(to_location, "zone_code", "") or ""),
            qty=max(0, int(qty or 0)),
            payload=payload,
            performed_by=user,
            performed_by_role=str(get_request_role(request) or ""),
            occurred_at=timezone.now(),
        )

    def _move_box_to_pallet(self, snapshot: WarehouseStockSnapshot, raw_value, *, request):
        if not snapshot.container_id:
            raise ValueError("Для этой строки короб не найден.")
        try:
            box = (
                WarehouseContainer.objects
                .select_for_update(of=("self",))
                .select_related("parent_container", "current_location")
                .get(pk=snapshot.container_id, agency=snapshot.agency)
            )
        except WarehouseContainer.DoesNotExist:
            raise ValueError("Для этой строки короб не найден.")
        if box.container_type != WarehouseContainer.TYPE_BOX:
            raise ValueError("Выберите строку конкретного короба.")

        target_pallet = self._find_active_pallet(snapshot.agency, raw_value)
        if not target_pallet.current_location_id:
            raise ValueError("У выбранной паллеты не указано место хранения.")
        if int(box.parent_container_id or 0) == int(target_pallet.pk):
            raise ValueError("Короб уже находится на этой паллете.")
        self._ensure_editor_snapshots_unlocked(self._lock_pallet_snapshots(target_pallet))

        box_snapshots = list(
            WarehouseStockSnapshot.objects
            .select_for_update(of=("self",))
            .select_related("container", "parent_container", "location", "active_operation", "last_event", "sku_ref")
            .filter(agency=snapshot.agency, container=box, is_archived=False)
            .order_by("id")
        )
        if not box_snapshots:
            raise ValueError("В выбранном коробе нет активного товара.")
        self._ensure_editor_snapshots_unlocked(box_snapshots)

        source_pallet = box.parent_container
        source_location = box.current_location
        snapshot_ids = [item.pk for item in box_snapshots]
        now = timezone.now()
        box.parent_container = target_pallet
        box.current_location = target_pallet.current_location
        box.save(update_fields=["parent_container", "current_location", "updated_at"])
        WarehouseStockSnapshot.objects.filter(pk__in=snapshot_ids).update(
            parent_container=target_pallet,
            location=target_pallet.current_location,
            zone_code=target_pallet.current_location.zone_code,
            zone_kind=target_pallet.current_location.zone_kind,
            updated_at=now,
        )
        self._record_stock_editor_event(
            request=request,
            snapshot=snapshot,
            event_type="stock_editor_box_moved_to_pallet",
            container=box,
            from_location=source_location,
            to_location=target_pallet.current_location,
            qty=sum(int(item.qty or 0) for item in box_snapshots),
            payload={
                "box_code": box.container_code,
                "source_pallet_id": int(getattr(source_pallet, "pk", 0) or 0),
                "source_pallet_code": str(getattr(source_pallet, "container_code", "") or ""),
                "target_pallet_id": int(target_pallet.pk),
                "target_pallet_code": target_pallet.container_code,
                "snapshot_ids": snapshot_ids,
            },
        )
        return snapshot_ids

    def _move_box_to_new_pallet(self, snapshot: WarehouseStockSnapshot, raw_value, *, request):
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите номер новой паллеты.")
        if not snapshot.container_id:
            raise ValueError("Для этой строки короб не найден.")
        try:
            box = (
                WarehouseContainer.objects
                .select_for_update(of=("self",))
                .select_related("parent_container", "current_location")
                .get(pk=snapshot.container_id, agency=snapshot.agency)
            )
        except WarehouseContainer.DoesNotExist:
            raise ValueError("Для этой строки короб не найден.")
        if box.container_type != WarehouseContainer.TYPE_BOX:
            raise ValueError("Выберите строку конкретного короба.")
        if WarehouseContainer.objects.filter(
            agency=snapshot.agency,
            container_code__iexact=value,
        ).exists():
            raise ValueError("Такой номер паллеты уже используется.")

        box_snapshots = list(
            WarehouseStockSnapshot.objects
            .select_for_update(of=("self",))
            .select_related("container", "parent_container", "location", "active_operation", "last_event", "sku_ref")
            .filter(agency=snapshot.agency, container=box, is_archived=False)
            .order_by("id")
        )
        if not box_snapshots:
            raise ValueError("В выбранном коробе нет активного товара.")
        self._ensure_editor_snapshots_unlocked(box_snapshots)

        source_pallet = box.parent_container
        source_location = box.current_location or snapshot.location or getattr(source_pallet, "current_location", None)
        warehouse_code = str(getattr(source_location, "warehouse_code", "") or "").strip()
        if not warehouse_code:
            raise ValueError("Не удалось определить склад выбранного короба.")
        target_location = (
            WarehouseLocation.objects
            .filter(warehouse_code=warehouse_code, zone_code__iexact="PR")
            .order_by("id")
            .first()
        )
        if not target_location:
            raise ValueError("Зона PR для этого склада не найдена.")

        user = request.user if request.user.is_authenticated else None
        target_pallet = WarehouseContainer.objects.create(
            agency=snapshot.agency,
            container_type=WarehouseContainer.TYPE_PALLET,
            container_code=value,
            current_location=target_location,
            source_context_type=str(snapshot.source_context_type or "stock_editor"),
            source_context_id=str(snapshot.source_context_id or snapshot.pk),
            created_by=user,
        )
        snapshot_ids = [item.pk for item in box_snapshots]
        now = timezone.now()
        box.parent_container = target_pallet
        box.current_location = target_location
        box.save(update_fields=["parent_container", "current_location", "updated_at"])
        WarehouseStockSnapshot.objects.filter(pk__in=snapshot_ids).update(
            parent_container=target_pallet,
            location=target_location,
            zone_code=target_location.zone_code,
            zone_kind=target_location.zone_kind,
            warehouse_state_code="placed_in_receiving",
            updated_at=now,
        )
        self._record_stock_editor_event(
            request=request,
            snapshot=snapshot,
            event_type="stock_editor_box_moved_to_new_pallet",
            container=box,
            from_location=source_location,
            to_location=target_location,
            qty=sum(int(item.qty or 0) for item in box_snapshots),
            payload={
                "box_id": int(box.pk),
                "box_code": box.container_code,
                "source_pallet_id": int(getattr(source_pallet, "pk", 0) or 0),
                "source_pallet_code": str(getattr(source_pallet, "container_code", "") or ""),
                "target_pallet_id": int(target_pallet.pk),
                "target_pallet_code": target_pallet.container_code,
                "snapshot_ids": snapshot_ids,
            },
        )
        return snapshot_ids

    def _merge_pallets(self, snapshot: WarehouseStockSnapshot, raw_value, *, request):
        source_pallet = self._get_source_pallet(snapshot)
        target_pallet = self._find_active_pallet(snapshot.agency, raw_value)
        if int(source_pallet.pk) == int(target_pallet.pk):
            raise ValueError("Исходная и целевая паллеты совпадают.")
        if not target_pallet.current_location_id:
            raise ValueError("У выбранной паллеты не указано место хранения.")
        self._ensure_editor_snapshots_unlocked(self._lock_pallet_snapshots(target_pallet))

        source_boxes = list(
            WarehouseContainer.objects
            .select_for_update()
            .filter(
                agency=snapshot.agency,
                parent_container=source_pallet,
                container_type=WarehouseContainer.TYPE_BOX,
                status=WarehouseContainer.STATUS_ACTIVE,
            )
            .order_by("id")
        )
        source_box_ids = [box.pk for box in source_boxes]
        source_snapshots = list(
            WarehouseStockSnapshot.objects
            .select_for_update(of=("self",))
            .select_related("container", "parent_container", "location", "active_operation", "last_event", "sku_ref")
            .filter(agency=snapshot.agency, is_archived=False)
            .filter(
                Q(parent_container=source_pallet)
                | Q(container=source_pallet)
                | Q(container_id__in=source_box_ids)
            )
            .order_by("id")
        )
        if not source_snapshots:
            raise ValueError("На исходной паллете нет активного товара.")
        self._ensure_editor_snapshots_unlocked(source_snapshots)

        snapshot_ids = [item.pk for item in source_snapshots]
        direct_snapshot_ids = [item.pk for item in source_snapshots if item.container_id == source_pallet.pk]
        boxed_snapshot_ids = [item.pk for item in source_snapshots if item.pk not in direct_snapshot_ids]
        source_location = source_pallet.current_location
        now = timezone.now()

        if source_box_ids:
            WarehouseContainer.objects.filter(pk__in=source_box_ids).update(
                parent_container=target_pallet,
                current_location=target_pallet.current_location,
                updated_at=now,
            )
        if boxed_snapshot_ids:
            WarehouseStockSnapshot.objects.filter(pk__in=boxed_snapshot_ids).update(
                parent_container=target_pallet,
                location=target_pallet.current_location,
                zone_code=target_pallet.current_location.zone_code,
                zone_kind=target_pallet.current_location.zone_kind,
                updated_at=now,
            )
        if direct_snapshot_ids:
            WarehouseStockSnapshot.objects.filter(pk__in=direct_snapshot_ids).update(
                container=target_pallet,
                container_code=target_pallet.container_code,
                parent_container=None,
                location=target_pallet.current_location,
                zone_code=target_pallet.current_location.zone_code,
                zone_kind=target_pallet.current_location.zone_kind,
                updated_at=now,
            )

        source_pallet.status = WarehouseContainer.STATUS_MERGED
        source_pallet.current_location = None
        source_pallet.save(update_fields=["status", "current_location", "updated_at"])
        self._record_stock_editor_event(
            request=request,
            snapshot=snapshot,
            event_type="stock_editor_pallets_merged",
            container=source_pallet,
            from_location=source_location,
            to_location=target_pallet.current_location,
            qty=sum(int(item.qty or 0) for item in source_snapshots),
            payload={
                "source_pallet_id": int(source_pallet.pk),
                "source_pallet_code": source_pallet.container_code,
                "target_pallet_id": int(target_pallet.pk),
                "target_pallet_code": target_pallet.container_code,
                "moved_box_ids": source_box_ids,
                "moved_box_codes": [box.container_code for box in source_boxes],
                "snapshot_ids": snapshot_ids,
            },
        )
        return snapshot_ids

    def _rename_pallet(self, snapshot: WarehouseStockSnapshot, raw_value, *, request):
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("Укажите новый номер паллеты.")
        source_pallet = self._get_source_pallet(snapshot)
        old_code = str(source_pallet.container_code or "").strip()
        if old_code == value:
            raise ValueError("Новый номер совпадает с текущим.")
        if (
            WarehouseContainer.objects
            .filter(agency=snapshot.agency, container_code__iexact=value)
            .exclude(pk=source_pallet.pk)
            .exists()
        ):
            raise ValueError("Такой номер паллеты уже используется.")

        pallet_snapshots = list(
            WarehouseStockSnapshot.objects
            .select_for_update(of=("self",))
            .select_related("container", "parent_container", "location", "active_operation", "last_event", "sku_ref")
            .filter(agency=snapshot.agency, is_archived=False)
            .filter(Q(parent_container=source_pallet) | Q(container=source_pallet))
            .order_by("id")
        )
        if not pallet_snapshots:
            raise ValueError("На выбранной паллете нет активного товара.")
        self._ensure_editor_snapshots_unlocked(pallet_snapshots)

        snapshot_ids = [item.pk for item in pallet_snapshots]
        direct_snapshot_ids = [item.pk for item in pallet_snapshots if item.container_id == source_pallet.pk]
        source_pallet.container_code = value
        source_pallet.save(update_fields=["container_code", "updated_at"])
        if direct_snapshot_ids:
            WarehouseStockSnapshot.objects.filter(pk__in=direct_snapshot_ids).update(
                container_code=value,
                updated_at=timezone.now(),
            )
        self._record_stock_editor_event(
            request=request,
            snapshot=snapshot,
            event_type="stock_editor_pallet_renamed",
            container=source_pallet,
            from_location=source_pallet.current_location,
            to_location=source_pallet.current_location,
            qty=sum(int(item.qty or 0) for item in pallet_snapshots),
            payload={
                "pallet_id": int(source_pallet.pk),
                "old_pallet_code": old_code,
                "new_pallet_code": value,
                "snapshot_ids": snapshot_ids,
            },
        )
        return snapshot_ids

    def _update_snapshot_box(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043d\u043e\u043c\u0435\u0440 \u043a\u043e\u0440\u043e\u0431\u0430.")
        if not snapshot.container_id:
            raise ValueError("\u0414\u043b\u044f \u044d\u0442\u043e\u0439 \u0441\u0442\u0440\u043e\u043a\u0438 \u043a\u043e\u0440\u043e\u0431 \u043d\u0435 \u043d\u0430\u0439\u0434\u0435\u043d.")
        snapshot.container.container_code = value
        snapshot.container.save(update_fields=["container_code", "updated_at"])
        WarehouseStockSnapshot.objects.filter(
            agency=snapshot.agency,
            container=snapshot.container,
            is_archived=False,
        ).update(container_code=value, updated_at=timezone.now())

    def _update_snapshot_goods_type(self, snapshot: WarehouseStockSnapshot, raw_value, *, request) -> None:
        raw_text = " ".join(str(raw_value or "").strip().lower().split())
        value = self._GOODS_TYPE_ALIASES.get(raw_text, "")
        if not value:
            raise ValueError(
                "Выберите допустимый тип товара: gv, no, op, br, vz, rh или votg."
            )

        old_value = str(snapshot.goods_type or "").strip().lower()
        if old_value == value:
            raise ValueError("Новый тип товара совпадает с текущим.")

        event = self._record_stock_editor_event(
            request=request,
            snapshot=snapshot,
            event_type=WarehouseEventType.STOCK_CORRECTED.value,
            container=snapshot.container or snapshot.parent_container,
            from_location=snapshot.location,
            to_location=snapshot.location,
            qty=int(snapshot.qty or 0),
            payload={
                "editor_action": "goods_type_changed",
                "snapshot_id": int(snapshot.pk),
                "sku": str(snapshot.sku_code or ""),
                "box_code": str(getattr(snapshot.container, "container_code", "") or snapshot.container_code or ""),
                "pallet_code": str(getattr(snapshot.parent_container, "container_code", "") or ""),
                "old_goods_type": old_value,
                "new_goods_type": value,
            },
        )
        snapshot.goods_type = value
        snapshot.last_event = event
        snapshot.snapshot_version = int(snapshot.snapshot_version or 0) + 1
        snapshot.save(update_fields=["goods_type", "last_event", "snapshot_version", "updated_at"])

    def _update_snapshot_location(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value = str(raw_value or "").strip()
        if not value:
            raise ValueError("\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043c\u0435\u0441\u0442\u043e \u0445\u0440\u0430\u043d\u0435\u043d\u0438\u044f.")
        location = self._resolve_editor_location(snapshot, value)
        if not location:
            raise ValueError("\u042f\u0447\u0435\u0439\u043a\u0430 \u043d\u0435 \u043d\u0430\u0439\u0434\u0435\u043d\u0430 \u0432 \u0441\u043f\u0440\u0430\u0432\u043e\u0447\u043d\u0438\u043a\u0435 \u043c\u0435\u0441\u0442.")
        if snapshot.container_id:
            container = snapshot.container
            now = timezone.now()
            target_pallets = WarehouseContainer.objects.filter(
                current_location=location,
                container_type__in=[
                    WarehouseContainer.TYPE_PALLET,
                    WarehouseContainer.TYPE_MIXED_PALLET,
                ],
                status=WarehouseContainer.STATUS_ACTIVE,
            )
            if target_pallets.exclude(agency=snapshot.agency).exists():
                raise ValueError("\u041d\u0430 \u044d\u0442\u043e\u043c \u043c\u0435\u0441\u0442\u0435 \u0443\u0436\u0435 \u0441\u0442\u043e\u0438\u0442 \u043f\u0430\u043b\u043b\u0435\u0442\u0430 \u0434\u0440\u0443\u0433\u043e\u0433\u043e \u043a\u043b\u0438\u0435\u043d\u0442\u0430.")

            same_agency_pallets = target_pallets.filter(agency=snapshot.agency).order_by("id")
            same_agency_count = same_agency_pallets.count()
            if same_agency_count > 1:
                raise ValueError("\u041d\u0430 \u044d\u0442\u043e\u043c \u043c\u0435\u0441\u0442\u0435 \u043d\u0435\u0441\u043a\u043e\u043b\u044c\u043a\u043e \u043f\u0430\u043b\u043b\u0435\u0442. \u0421\u043d\u0430\u0447\u0430\u043b\u0430 \u0443\u0442\u043e\u0447\u043d\u0438\u0442\u0435, \u0432 \u043a\u0430\u043a\u0443\u044e \u043f\u0430\u043b\u043b\u0435\u0442\u0443 \u043f\u0435\u0440\u0435\u043d\u0435\u0441\u0442\u0438 \u043a\u043e\u0440\u043e\u0431.")
            target_pallet = same_agency_pallets.first()

            if container.container_type == WarehouseContainer.TYPE_BOX and target_pallet:
                container.parent_container = target_pallet
                container.current_location = location
                container.save(update_fields=["parent_container", "current_location", "updated_at"])
                WarehouseStockSnapshot.objects.filter(
                    agency=snapshot.agency,
                    container=container,
                    is_archived=False,
                ).update(
                    parent_container=target_pallet,
                    location=location,
                    zone_code=location.zone_code,
                    zone_kind=location.zone_kind,
                    updated_at=now,
                )
                return

            if container.container_type == WarehouseContainer.TYPE_BOX and snapshot.parent_container_id:
                pallet = snapshot.parent_container
                pallet.current_location = location
                pallet.save(update_fields=["current_location", "updated_at"])
                WarehouseContainer.objects.filter(
                    agency=snapshot.agency,
                    parent_container=pallet,
                    status=WarehouseContainer.STATUS_ACTIVE,
                ).update(current_location=location, updated_at=now)
                WarehouseStockSnapshot.objects.filter(
                    agency=snapshot.agency,
                    is_archived=False,
                ).filter(
                    Q(parent_container=pallet) | Q(container=pallet)
                ).update(
                    location=location,
                    zone_code=location.zone_code,
                    zone_kind=location.zone_kind,
                    updated_at=now,
                )
                return

            container.current_location = location
            container.save(update_fields=["current_location", "updated_at"])
            WarehouseStockSnapshot.objects.filter(
                agency=snapshot.agency,
                container=container,
                is_archived=False,
            ).update(
                location=location,
                zone_code=location.zone_code,
                zone_kind=location.zone_kind,
                updated_at=now,
            )
        else:
            snapshot.location = location
            snapshot.zone_code = location.zone_code
            snapshot.zone_kind = location.zone_kind
            snapshot.save(update_fields=["location", "zone_code", "zone_kind", "updated_at"])

    def _update_snapshot_qty(self, snapshot: WarehouseStockSnapshot, raw_value) -> None:
        value_text = str(raw_value or "").strip()
        if not value_text:
            raise ValueError("\u0423\u043a\u0430\u0436\u0438\u0442\u0435 \u043a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e \u0432 \u043a\u043e\u0440\u043e\u0431\u0435.")
        try:
            value = int(value_text)
        except (TypeError, ValueError):
            raise ValueError("\u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e \u0432 \u043a\u043e\u0440\u043e\u0431\u0435 \u0434\u043e\u043b\u0436\u043d\u043e \u0431\u044b\u0442\u044c \u0446\u0435\u043b\u044b\u043c \u0447\u0438\u0441\u043b\u043e\u043c.")
        if value < 0:
            raise ValueError("\u041a\u043e\u043b\u0438\u0447\u0435\u0441\u0442\u0432\u043e \u0432 \u043a\u043e\u0440\u043e\u0431\u0435 \u043d\u0435 \u043c\u043e\u0436\u0435\u0442 \u0431\u044b\u0442\u044c \u043e\u0442\u0440\u0438\u0446\u0430\u0442\u0435\u043b\u044c\u043d\u044b\u043c.")
        snapshot.qty = value
        snapshot.available_qty = value
        snapshot.save(update_fields=["qty", "available_qty", "updated_at"])


class HeadManagerReferenceMixin(RoleRequiredMixin):
    allowed_roles = ("head_manager",)
    success_url = "/head-manager/"

    def get_success_url(self):
        return self.success_url


class HeadManagerClientsView(HeadManagerReferenceMixin, TemplateView):
    template_name = "head_manager/clients.html"
    paginate_by = 30

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        request = self.request
        from accountant.selectors import manager_visible_agencies

        queryset = build_client_list_queryset(
            request=request,
            base_queryset=manager_visible_agencies(Agency.objects.select_related("portal_user", "lifecycle")),
            sort_fields=CLIENT_SORT_FIELDS,
            filter_fields=CLIENT_FILTER_FIELDS,
            default_sort="name",
        )
        paginator = Paginator(queryset, self.paginate_by)
        page_obj = paginator.get_page(request.GET.get("page") or 1)
        list_ctx = build_client_list_context(
            request=request,
            items=page_obj.object_list,
            view_modes=("table",),
            sort_fields=CLIENT_SORT_FIELDS,
            default_sort="name",
        )
        client_rows = [
            {
                "agency": agency,
                "lk_url": build_client_cabinet_url(agency.id),
            }
            for agency in page_obj.object_list
        ]
        ctx.update(list_ctx)
        ctx.update(
            {
                "role": "head_manager",
                "title": "Клиенты",
                "active_nav": "references",
                "page_obj": page_obj,
                "client_rows": client_rows,
                "total_count": paginator.count,
            }
        )
        return ctx


class HeadManagerMarkingSearchView(HeadManagerReferenceMixin, TemplateView):
    template_name = "head_manager/marking_search.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        query = str(self.request.GET.get("q") or "").strip()
        rows = _head_manager_marking_search_rows(query) if query else []
        ctx.update(
            {
                "title": "Поиск Честного знака",
                "active_nav": "references",
                "query": query,
                "normalized_query": _head_manager_marking_display_code(query) if query else "",
                "compact_query": _head_manager_compact_marking_code(query) if query else "",
                "rows": rows,
                "total_count": len(rows),
            }
        )
        return ctx


class OwnCompanyListView(HeadManagerReferenceMixin, ListView):
    template_name = "head_manager/reference_list.html"
    context_object_name = "items"
    model = OwnCompany

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_list_context(
            title="Наши компании",
            subtitle="Юрлица и ИП FullBox для документов, отгрузок и транспортных накладных.",
            create_url="/head-manager/own-companies/new/",
            edit_base_url="/head-manager/own-companies",
            back_url="/head-manager/",
            type_label="компания",
            directory_mode="own_companies",
        ))
        return ctx


class OwnCompanyCreateView(HeadManagerReferenceMixin, CreateView):
    template_name = "head_manager/reference_form.html"
    form_class = OwnCompanyForm
    model = OwnCompany
    success_url = "/head-manager/own-companies/"

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, "Компания сохранена.")
        return response

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_form_context(
            title="Новая компания",
            subtitle="Заполни реквизиты своей компании для документов и ТН.",
            back_url="/head-manager/own-companies/",
            submit_label="Сохранить компанию",
        ))
        return ctx


class OwnCompanyUpdateView(HeadManagerReferenceMixin, UpdateView):
    template_name = "head_manager/reference_form.html"
    form_class = OwnCompanyForm
    model = OwnCompany
    success_url = "/head-manager/own-companies/"

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, "Компания обновлена.")
        return response

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_form_context(
            title="Редактирование компании",
            subtitle="Исправь реквизиты своей компании.",
            back_url="/head-manager/own-companies/",
            submit_label="Сохранить изменения",
        ))
        return ctx


class CarrierListView(HeadManagerReferenceMixin, ListView):
    template_name = "head_manager/reference_list.html"
    context_object_name = "items"
    model = Carrier

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_list_context(
            title="Перевозчики",
            subtitle="Справочник транспортных компаний и внешних перевозчиков.",
            create_url="/head-manager/carriers/new/",
            edit_base_url="/head-manager/carriers",
            back_url="/head-manager/",
            type_label="перевозчик",
            directory_mode="carriers",
        ))
        return ctx


class CarrierCreateView(HeadManagerReferenceMixin, CreateView):
    template_name = "head_manager/reference_form.html"
    form_class = CarrierForm
    model = Carrier
    success_url = "/head-manager/carriers/"

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, "Перевозчик сохранен.")
        return response

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_form_context(
            title="Новый перевозчик",
            subtitle="Заполни карточку транспортной компании или ИП.",
            back_url="/head-manager/carriers/",
            submit_label="Сохранить перевозчика",
        ))
        return ctx


class CarrierUpdateView(HeadManagerReferenceMixin, UpdateView):
    template_name = "head_manager/reference_form.html"
    form_class = CarrierForm
    model = Carrier
    success_url = "/head-manager/carriers/"

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, "Перевозчик обновлен.")
        return response

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(build_reference_form_context(
            title="Редактирование перевозчика",
            subtitle="Исправь реквизиты перевозчика.",
            back_url="/head-manager/carriers/",
            submit_label="Сохранить изменения",
        ))
        return ctx


def _marketplace_warehouses_path() -> Path:
    return settings.BASE_DIR.parent / "marketplace_warehouses.json"


def _normalize_lines(values: Iterable) -> list[str]:
    lines = []
    for value in values:
        text = str(value).strip()
        if text:
            lines.append(text)
    return list(dict.fromkeys(lines))


def _guess_marketplace_warehouse_type(name: str, address: str = "") -> str:
    haystack = f"{name} {address}".lower()
    if any(token in haystack for token in ("транзит", "ппп", "рцпп", "гольёво")):
        return "Транзитный"
    if any(token in haystack for token in ("сц", "сортиров")):
        return "Сортировочный"
    if any(token in haystack for token in ("рфц", "ффц", "фулфил", "фулфилл")):
        return "Фулфилмент"
    return "Обычный"


def _warehouse_row_from_legacy_line(value) -> dict | None:
    text = str(value or "").strip()
    if not text:
        return None
    name = text
    address = ""
    for separator in (" — ", " - ", " – "):
        if separator in text:
            left, right = text.split(separator, 1)
            if left.strip() and right.strip():
                name = left.strip()
                address = right.strip()
                break
    return {
        "type": _guess_marketplace_warehouse_type(name, address),
        "name": name,
        "address": address,
    }


def _normalize_marketplace_warehouse_rows(values) -> list[dict]:
    if isinstance(values, str):
        values = [values]
    if not isinstance(values, list):
        return []
    rows: list[dict] = []
    seen: set[tuple[str, str, str]] = set()
    for value in values:
        if isinstance(value, dict):
            row_type = str(value.get("type") or "").strip() or _guess_marketplace_warehouse_type(
                str(value.get("name") or ""),
                str(value.get("address") or ""),
            )
            name = str(value.get("name") or "").strip()
            address = str(value.get("address") or "").strip()
            if not name and not address:
                continue
            row = {"type": row_type, "name": name, "address": address}
        else:
            row = _warehouse_row_from_legacy_line(value)
            if row is None:
                continue
        dedupe_key = (row["type"].lower(), row["name"].lower(), row["address"].lower())
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        rows.append(row)
    return rows


def _warehouse_display_line(row: dict) -> str:
    warehouse_type = str(row.get("type") or "").strip()
    name = str(row.get("name") or "").strip()
    address = str(row.get("address") or "").strip()
    left = " · ".join(part for part in (warehouse_type, name) if part)
    if left and address:
        return f"{left} — {address}"
    return left or address


def _load_marketplace_warehouses() -> dict:
    base = {"wb": [], "ozon": [], "yandex": [], "sber": []}
    path = _marketplace_warehouses_path()
    if not path.exists():
        return base
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return base
    if not isinstance(data, dict):
        return base
    for key in base:
        values = data.get(key) or []
        base[key] = _normalize_marketplace_warehouse_rows(values)
    return base


def _save_marketplace_warehouses(data: dict, user: str = "", meta: dict | None = None) -> None:
    payload = {
        "wb": _normalize_marketplace_warehouse_rows(data.get("wb", [])),
        "ozon": _normalize_marketplace_warehouse_rows(data.get("ozon", [])),
        "yandex": _normalize_marketplace_warehouse_rows(data.get("yandex", [])),
        "sber": _normalize_marketplace_warehouse_rows(data.get("sber", [])),
        "meta": {
            "updated_at": timezone.localtime().isoformat(),
            "updated_by": user,
        },
    }
    if meta:
        payload["meta"].update(meta)
    path = _marketplace_warehouses_path()
    path.write_text(json.dumps(payload, ensure_ascii=True, indent=2), encoding="utf-8")


def _extract_value(item: dict, keys: Iterable[str]) -> str:
    for key in keys:
        value = item.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return ""


def _format_address_line(item: dict, name_keys: Iterable[str], address_keys: Iterable[str], city_keys: Iterable[str]):
    name = _extract_value(item, name_keys)
    address = _extract_value(item, address_keys)
    city = _extract_value(item, city_keys)
    address_parts = []
    if city and city.lower() not in address.lower():
        address_parts.append(city)
    if address:
        address_parts.append(address)
    address_text = ", ".join(address_parts).strip()
    return {
        "type": _guess_marketplace_warehouse_type(name, address_text),
        "name": name,
        "address": address_text,
    }


def _parse_items_payload(data) -> list[dict] | None:
    if isinstance(data, list):
        return data
    if not isinstance(data, dict):
        return None
    for key in ("result", "warehouses", "data"):
        value = data.get(key)
        if isinstance(value, list):
            return value
        if isinstance(value, dict):
            for nested_key in ("warehouses", "data", "result", "items"):
                nested_value = value.get(nested_key)
                if isinstance(nested_value, list):
                    return nested_value
    return None


def _fetch_wb_warehouses(token: str) -> tuple[list[dict], str | None]:
    endpoints = ["https://marketplace-api.wildberries.ru/api/v3/warehouses"]
    errors = []
    for url in endpoints:
        try:
            response = requests.get(url, headers={"Authorization": token}, timeout=20)
        except requests.RequestException as exc:
            errors.append(f"WB API недоступен ({url}): {exc}")
            continue
        if response.status_code != 200:
            detail = ""
            try:
                payload = response.json()
                detail = (payload.get("detail") or payload.get("title") or "").strip()
            except ValueError:
                detail = ""
            extra = f": {detail}" if detail else ""
            errors.append(f"WB API ошибка {response.status_code} ({url}){extra}")
            continue
        try:
            data = response.json()
        except ValueError:
            errors.append(f"WB API вернул некорректный JSON ({url}).")
            continue
        items = _parse_items_payload(data)
        if items is None:
            errors.append(f"WB API не вернул список складов ({url}).")
            continue
        if not items:
            errors.append("WB: список складов пуст.")
            continue
        rows = []
        for item in items:
            if not isinstance(item, dict):
                continue
            row = _format_address_line(
                item,
                name_keys=("name", "warehouseName", "officeName", "warehouse", "title"),
                address_keys=("address", "warehouseAddress", "officeAddress", "addr", "addressFull"),
                city_keys=("city", "town", "region"),
            )
            if row:
                rows.append(row)
        return _normalize_marketplace_warehouse_rows(rows), None
    if errors:
        return [], errors[0]
    return [], "WB API не отвечает."


def _normalize_ozon_client_id(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    if text.isdigit():
        return text
    match = re.fullmatch(r"(\d+)(?:\.0+)?", text)
    return match.group(1) if match else text


def _ozon_headers(client_id: str, api_key: str) -> dict:
    return {
        "Client-Id": client_id,
        "Api-Key": api_key,
        "Content-Type": "application/json",
    }


def _ozon_post(path: str, client_id: str, api_key: str, payload: dict, timeout: int = 30):
    url = f"https://api-seller.ozon.ru{path}"
    try:
        response = requests.post(
            url,
            headers=_ozon_headers(client_id, api_key),
            json=payload,
            timeout=timeout,
        )
    except requests.RequestException as exc:
        return None, f"Ozon API недоступен: {exc}"
    if response.status_code != 200:
        snippet = (response.text or "").strip()
        if len(snippet) > 200:
            snippet = f"{snippet[:200]}..."
        detail = f": {snippet}" if snippet else ""
        return None, f"Ozon API ошибка {response.status_code}{detail}"
    try:
        data = response.json()
    except ValueError:
        return None, "Ozon API вернул некорректный JSON."
    return data, None


def _parse_ozon_clusters(data) -> list[dict] | None:
    if not isinstance(data, dict):
        return None
    clusters = data.get("clusters")
    if not isinstance(clusters, list):
        return None
    items = []
    for cluster in clusters:
        if not isinstance(cluster, dict):
            continue
        cluster_name = str(cluster.get("name") or "").strip()
        logistic_clusters = cluster.get("logistic_clusters") or []
        if not isinstance(logistic_clusters, list):
            continue
        for log_cluster in logistic_clusters:
            if not isinstance(log_cluster, dict):
                continue
            warehouses = log_cluster.get("warehouses") or []
            if not isinstance(warehouses, list):
                continue
            for warehouse in warehouses:
                if not isinstance(warehouse, dict):
                    continue
                item = dict(warehouse)
                if cluster_name:
                    item["cluster_name"] = cluster_name
                items.append(item)
    return items


def _fetch_ozon_clusters(client_id: str, api_key: str) -> tuple[list[dict], str | None]:
    rows = []
    errors = []
    for cluster_type in (1, 2):
        data, error = _ozon_post(
            "/v1/cluster/list",
            client_id,
            api_key,
            {"limit": 200, "offset": 0, "cluster_type": cluster_type},
        )
        if error:
            errors.append(error)
            continue
        items = _parse_ozon_clusters(data or {})
        if items is None:
            errors.append("Ozon API не вернул список кластеров.")
            continue
        for item in items:
            name = str(item.get("name") or "").strip()
            cluster_name = str(item.get("cluster_name") or "").strip()
            if cluster_name and name and cluster_name not in name:
                rows.append(
                    {
                        "type": _guess_marketplace_warehouse_type(cluster_name, name),
                        "name": cluster_name,
                        "address": name,
                    }
                )
            elif name:
                rows.append(
                    {
                        "type": _guess_marketplace_warehouse_type(name, ""),
                        "name": name,
                        "address": "",
                    }
                )
    return _normalize_marketplace_warehouse_rows(rows), errors[0] if errors else None


def _fetch_ozon_warehouses(client_id: str, api_key: str) -> tuple[list[dict], str | None]:
    data, error = _ozon_post("/v1/warehouse/list", client_id, api_key, {})
    if error:
        data, error = _ozon_post("/v1/warehouse/list", client_id, api_key, {"limit": 200, "offset": 0})
    if error:
        return [], error
    items = _parse_items_payload(data or {})
    if items is None:
        return [], "Ozon API не вернул список складов."
    if not items:
        cluster_lines, cluster_error = _fetch_ozon_clusters(client_id, api_key)
        if cluster_lines:
            return cluster_lines, None
        return [], cluster_error or "Ozon: список складов пуст."
    rows = []
    for item in items:
        if not isinstance(item, dict):
            continue
        row = _format_address_line(
            item,
            name_keys=("name", "warehouse_name", "title"),
            address_keys=("address", "address_full", "warehouse_address", "address_text"),
            city_keys=("city", "region"),
        )
        if row:
            rows.append(row)
    return _normalize_marketplace_warehouse_rows(rows), None


def _find_agency(client_id: str | None, client_name: str | None) -> Agency | None:
    if client_id:
        return Agency.objects.filter(pk=client_id).first()
    tokens = []
    if client_name:
        tokens.append(client_name)
    tokens.extend(["кейзи", "keizi", "keyzi", "кейз"])
    query = Q()
    for token in tokens:
        token = (token or "").strip()
        if not token:
            continue
        query |= Q(agn_name__icontains=token) | Q(fio_agn__icontains=token)
    if query:
        agency = Agency.objects.filter(query).order_by("id").first()
        if agency:
            return agency
    case_query = Q()
    for token in tokens:
        token = (token or "").strip()
        if not token:
            continue
        variants = {token, token.lower(), token.upper()}
        for variant in variants:
            case_query |= Q(agn_name__contains=variant) | Q(fio_agn__contains=variant)
    if case_query:
        return Agency.objects.filter(case_query).order_by("id").first()
    return None


def _sync_marketplace_warehouses(agency: Agency) -> tuple[dict, list[str]]:
    data = _load_marketplace_warehouses()
    errors = []
    wb_market = Market.objects.filter(name__iexact="WB").first()
    ozon_market = Market.objects.filter(name__iexact="OZON").first()

    if wb_market:
        credential = MarketCredential.objects.filter(agency=agency, market=wb_market).first()
        token = (credential.market_key or "").strip() if credential else ""
        if token:
            wb_list, wb_error = _fetch_wb_warehouses(token)
            if wb_list:
                data["wb"] = wb_list
            else:
                errors.append(wb_error or "WB: список складов пуст.")
        else:
            errors.append("WB: не указан токен.")
    else:
        errors.append("WB: маркетплейс не найден.")

    if ozon_market:
        credential = MarketCredential.objects.filter(agency=agency, market=ozon_market).first()
        token = (credential.market_key or "").strip() if credential else ""
        client_id_value = _normalize_ozon_client_id(credential.client_id) if credential else ""
        if client_id_value and token:
            ozon_list, ozon_error = _fetch_ozon_warehouses(client_id_value, token)
            if ozon_list:
                data["ozon"] = ozon_list
            else:
                errors.append(ozon_error or "Ozon: список складов пуст.")
        else:
            if not client_id_value:
                errors.append("Ozon: не указан Client ID.")
            if not token:
                errors.append("Ozon: не указан API ключ.")
    else:
        errors.append("Ozon: маркетплейс не найден.")

    return data, errors


class MarketplaceWarehousesView(RoleRequiredMixin, TemplateView):
    template_name = "head_manager/marketplace_warehouses.html"
    allowed_roles = ("head_manager",)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(
            build_marketplace_warehouses_context(
                request=self.request,
                saved=kwargs.get("saved", False),
                error=kwargs.get("error", ""),
            )
        )
        return ctx

    def post(self, request, *args, **kwargs):
        response, error = save_marketplace_warehouses_response(request=request)
        if response is None:
            return self.render_to_response(self.get_context_data(error="Не удалось сохранить список."))
        return response


class MarketplaceWarehousesSyncView(RoleRequiredMixin, View):
    allowed_roles = ("head_manager",)

    def post(self, request, *args, **kwargs):
        return sync_marketplace_warehouses_response(request=request)
