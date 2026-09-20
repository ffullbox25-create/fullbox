from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from audit.models import OrderAuditEntry
from employees.models import Employee
from sku.models import Agency
from todo.models import Task

from .employee_report import processing_employee_fact_rows


User = get_user_model()


def _login_processing_head(client):
    user = User.objects.create_user(
        username="employee_report_head",
        password="pwd",
        first_name="Марина",
        last_name="Ларионова",
    )
    employee = Employee.objects.create(
        user=user,
        full_name="Ларионова Марина",
        role="processing_head",
        is_active=True,
    )
    client.force_login(user)
    session = client.session
    session["employee_id"] = employee.pk
    session["employee_role"] = employee.role
    session["employee_name"] = employee.full_name
    session.save()
    return user, employee


@override_settings(ALLOWED_HOSTS=["testserver"])
class ProcessingEmployeeReportTests(TestCase):
    def setUp(self):
        self.client = Client()
        self.head_user, self.head = _login_processing_head(self.client)
        self.worker_user = User.objects.create_user(
            username="employee_report_worker",
            password="pwd",
            first_name="Анна",
            last_name="Соколова",
        )
        self.worker = Employee.objects.create(
            user=self.worker_user,
            full_name="Соколова Анна",
            role="processing_worker",
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name='ООО "Клиент отчёта"')
        self.now = timezone.localtime().replace(microsecond=0)

    def _create_fact(
        self,
        *,
        order_id="EMPLOYEE-REPORT-1",
        processed_at=None,
        include_box=False,
        card_completed=True,
    ):
        processed_at = processed_at or self.now
        card = {
            "id": "card-a",
            "article": "OLD-SKU",
            "name": "Тестовый товар",
            "rows": [{"article": "OLD-SKU", "size": "42", "qty": "12"}],
        }
        if card_completed:
            card.update(
                {
                    "processed_done": True,
                    "processed_at": processed_at.isoformat(),
                    "processed_by": self.worker.full_name,
                }
            )
        payload = {
            "marketplace": "Wildberries",
            "marking_stickers": ["58x40"],
            "cards": [card],
            "processed_cards": ["card-a"] if card_completed else [],
            "processing_results": [
                {
                    "card_id": "card-a",
                    "source_article": "OLD-SKU",
                    "article": "NEW-SKU",
                    "product_name": "Тестовый товар",
                    "barcode": "460000000001",
                    "size": "42",
                    "processed": "12",
                    "defect": "2",
                    "shortage": "1",
                    "labels_printed": "10",
                    "tags_replaced": "8",
                }
            ],
        }
        if include_box:
            payload.update(
                {
                    "flow_closed": True,
                    "flow_closed_at": processed_at.isoformat(),
                    "act_state": "closed",
                    "act_boxes": [
                        {
                            "code": "BOX-REPORT-1",
                            "owner_user_id": self.worker_user.pk,
                            "owner_user_label": self.worker.full_name,
                            "direction": "Wildberries",
                            "items": [
                                {
                                    "sku": "NEW-SKU",
                                    "name": "Тестовый товар",
                                    "size": "42",
                                    "barcode": "460000000001",
                                    "qty": 7,
                                },
                                {
                                    "sku": "NEW-SKU-2",
                                    "name": "Второй товар",
                                    "size": "44",
                                    "barcode": "460000000002",
                                    "qty": 5,
                                },
                            ],
                        }
                    ],
                }
            )
        return OrderAuditEntry.objects.create(
            order_id=order_id,
            order_type="processing",
            action="update",
            agency=self.agency,
            user=self.worker_user,
            description="Фактический результат обработки",
            payload=payload,
            created_at=processed_at,
        )

    def _fact_rows(self):
        return processing_employee_fact_rows(
            self.now - timezone.timedelta(days=1),
            self.now + timezone.timedelta(days=1),
        )

    def test_report_uses_completed_product_card_fact(self):
        self._create_fact()

        rows = self._fact_rows()

        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["executor"], self.worker.full_name)
        self.assertEqual(row["employee_role"], "Обработчик")
        self.assertEqual(row["client"], self.agency.short_name)
        self.assertEqual(row["operation"], "Обработка товара")
        self.assertEqual(row["product"], "Тестовый товар")
        self.assertEqual(row["source_article"], "OLD-SKU")
        self.assertEqual(row["article"], "NEW-SKU")
        self.assertEqual(row["processed"], 12)
        self.assertEqual(row["defect"], 2)
        self.assertEqual(row["shortage"], 1)
        self.assertEqual(row["labels"], 10)
        self.assertEqual(row["tags"], 8)
        self.assertIn("Wildberries", row["details"])

    def test_assignment_without_completed_card_is_not_counted(self):
        Task.objects.create(
            title="Только назначенная заявка",
            route="/orders/processing/999/work/",
            assigned_to=self.worker,
            status="in_progress",
        )
        self._create_fact(card_completed=False)

        self.assertEqual(self._fact_rows(), [])

    def test_closed_boxes_are_attributed_without_double_counting_boxes(self):
        self._create_fact(include_box=True)

        rows = self._fact_rows()
        box_rows = [
            row for row in rows if row["operation"] == "Формирование коробов"
        ]

        self.assertEqual(len(box_rows), 2)
        self.assertEqual(sum(row["boxed"] for row in box_rows), 12)
        self.assertEqual(sum(row["boxes"] for row in box_rows), 1)
        self.assertTrue(
            all(row["executor"] == self.worker.full_name for row in box_rows)
        )
        self.assertTrue(
            all("BOX-REPORT-1" in row["details"] for row in box_rows)
        )

    def test_filters_and_summary_apply_to_fact_rows(self):
        self._create_fact(include_box=True)

        response = self.client.get(
            "/processing-head/reports/employees",
            {
                "period": "all",
                "client": "Клиент отчёта",
                "executor": "Соколова",
                "operation": "Формирование коробов",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["rows_total"], 2)
        summary = dict(response.context["summary"])
        self.assertEqual(summary["Сотрудников с результатом"], "1")
        self.assertEqual(summary["Клиентов"], "1")
        self.assertEqual(summary["Заявок"], "1")
        self.assertEqual(summary["Обработано единиц"], "0")
        self.assertEqual(summary["Сформировано в короба"], "12")
        self.assertEqual(summary["Сформировано коробов"], "1")
        self.assertEqual(
            response.context["employee_totals"],
            [
                {
                    "executor": self.worker.full_name,
                    "employee_role": "Обработчик",
                    "clients": 1,
                    "documents": 1,
                    "products": 2,
                    "operations": "Формирование коробов",
                    "processed": 0,
                    "boxed": 12,
                    "boxes": 1,
                    "defect": 0,
                    "labels": 0,
                }
            ],
        )
        self.assertContains(response, "Итоги по сотрудникам")
        self.assertContains(response, "Соколова Анна")
        self.assertContains(response, "Тестовый товар")
        self.assertNotContains(response, "Обработка завершена")

    def test_excel_export_contains_employee_product_and_quantities(self):
        self._create_fact(include_box=True)

        response = self.client.get(
            "/processing-head/reports/employees",
            {
                "period": "all",
                "executor": "Соколова",
                "export": "xlsx",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Отчёт", "По сотрудникам"])
        values = [
            str(cell or "")
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows(values_only=True)
            for cell in row
        ]
        exported = "\n".join(values)
        self.assertIn("Соколова Анна", exported)
        self.assertIn("Тестовый товар", exported)
        self.assertIn("Обработка товара", exported)
        self.assertIn("Формирование коробов", exported)
        self.assertIn("12", exported)

    def test_completed_fact_outside_period_is_excluded(self):
        self._create_fact(processed_at=self.now - timezone.timedelta(days=40))

        rows = processing_employee_fact_rows(
            self.now - timezone.timedelta(days=5),
            self.now + timezone.timedelta(days=1),
        )

        self.assertEqual(rows, [])
