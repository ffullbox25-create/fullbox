from io import BytesIO
from datetime import timedelta
import re

from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import Http404, HttpResponse, HttpResponseForbidden, HttpResponseNotAllowed, JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from audit.models import OrderAuditEntry
from employees.access import get_request_role, is_staff_role, role_required
from fullbox.order_numbers import format_order_number
from sku.models import Agency

from sklad.models import WarehouseContainer, WarehouseEvent, WarehouseStockSnapshot
from sklad.services.warehouse_write_path import WarehouseWritePathService
from sklad.services.missing_box_replacement import (
    mark_missing_box_not_found,
    resolve_found_missing_box,
    take_missing_box_for_check,
)
from sklad.services.warehouse_stock_rows import normalize_stock_row_from_snapshot

from .knowledge_base import get_knowledge_article, list_knowledge_articles
from .ui_services import (
    _agency_journal_label,
    _client_agency_for_request,
    _journal_goods_type_hint,
    _short_location_label,
    build_inventory_journal_page,
    build_sku_stock_page,
)

STOREKEEPER_KB_ROLES = ("storekeeper", "admin", "director", "developer", "head_manager")


def inventory_pagination_links(page_obj, *, radius: int = 2) -> list[dict]:
    total_pages = int(getattr(page_obj.paginator, "num_pages", 0) or 0)
    current = int(getattr(page_obj, "number", 1) or 1)
    if total_pages <= 1:
        return []

    pages = {1, total_pages}
    for number in range(max(1, current - radius), min(total_pages, current + radius) + 1):
        pages.add(number)

    links: list[dict] = []
    previous = 0
    for number in sorted(pages):
        if previous and number - previous > 1:
            links.append({"ellipsis": True})
        links.append({"number": number, "current": number == current})
        previous = number
    return links


@role_required("storekeeper")
def dashboard(request):
    return render(request, "sklad/dashboard.html")


@role_required(*STOREKEEPER_KB_ROLES)
def warehouse_box_check(request):
    if request.method != "GET":
        return HttpResponseNotAllowed(["GET"])
    return render(
        request,
        "sklad/warehouse_box_check.html",
        {"active_nav": "warehouse_box_check"},
    )


@role_required(*STOREKEEPER_KB_ROLES)
def knowledge_catalog(request):
    return render(
        request,
        "sklad/knowledge.html",
        {
            "active_nav": "knowledge",
            "kb_articles": list_knowledge_articles(),
        },
    )


@role_required(*STOREKEEPER_KB_ROLES)
def knowledge_article(request, slug: str):
    article = get_knowledge_article(slug)
    if article is None:
        raise Http404("Статья не найдена")
    return render(
        request,
        "sklad/knowledge_article.html",
        {
            "active_nav": "knowledge",
            "kb_article": article,
            "kb_articles": list_knowledge_articles(),
        },
    )


def inventory_journal(request):
    page = build_inventory_journal_page(request=request)
    if isinstance(page, HttpResponseForbidden):
        return page
    if str(request.GET.get("filter_options") or "").strip() == "1":
        filter_name = str(request.GET.get("filter_name") or "").strip()
        options_by_filter = page["context"].get("column_filter_options") or {}
        if filter_name not in options_by_filter:
            return JsonResponse({"options": []}, status=400)
        return JsonResponse({"options": options_by_filter.get(filter_name) or []})
    if str(request.GET.get("export") or "").strip().lower() == "xlsx":
        return _export_inventory_excel(page)
    context = page["context"]
    paginator = Paginator(context.get("rows") or [], 100)
    page_obj = paginator.get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    context["rows"] = list(page_obj.object_list)
    context["page_obj"] = page_obj
    context["page_query"] = query_params.urlencode()
    context["page_links"] = inventory_pagination_links(page_obj)
    export_params = request.GET.copy()
    export_params.pop("page", None)
    export_params["export"] = "xlsx"
    context["export_excel_url"] = f"{request.path}?{export_params.urlencode()}"
    return render(request, page["template_name"], context)


@role_required(*STOREKEEPER_KB_ROLES)
def problem_box_take(request, snapshot_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        WarehouseWritePathService.start_problem_box_return(
            snapshot_id=snapshot_id,
            expected_last_event_id=int(request.POST.get("last_event_id") or 0),
            performed_by=request.user,
        )
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(
            request,
            "Задание создано ричтраку. Остаток изменится только после скана короба и паллеты хранения.",
        )
    return redirect("sklad:inventory_journal")


@role_required(*STOREKEEPER_KB_ROLES)
def problem_box_return(request, snapshot_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        WarehouseWritePathService.complete_problem_box_return(
            operation_id=int(request.POST.get("operation_id") or 0),
            snapshot_id=snapshot_id,
            scanned_box_code=request.POST.get("box_code") or "",
            destination_pallet_code=request.POST.get("pallet_code") or "",
            performed_by=request.user,
        )
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Короб проверен и возвращен в доступный остаток.")
    return redirect("sklad:inventory_journal")


@role_required(*STOREKEEPER_KB_ROLES)
def problem_box_release(request, snapshot_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        WarehouseWritePathService.cancel_problem_box_return(
            operation_id=int(request.POST.get("operation_id") or 0),
            performed_by=request.user,
        )
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Проверка освобождена. Короб остался в списке проблемных.")
    return redirect("sklad:inventory_journal")


@role_required(*STOREKEEPER_KB_ROLES)
def missing_box_take(request, operation_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        take_missing_box_for_check(operation_id=operation_id, performed_by=request.user)
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Короб взят в проверку. Найдите короб или отметьте, что он не найден.")
    return redirect("sklad:inventory_journal")


@role_required(*STOREKEEPER_KB_ROLES)
def missing_box_found(request, operation_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        resolve_found_missing_box(
            operation_id=operation_id,
            scanned_box_code=request.POST.get("box_code") or "",
            destination_location_code=request.POST.get("location_code") or "",
            performed_by=request.user,
        )
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, "Короб найден: карантин снят, фактическая ячейка сохранена.")
    return redirect("sklad:inventory_journal")


@role_required(*STOREKEEPER_KB_ROLES)
def missing_box_not_found(request, operation_id: int):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    try:
        mark_missing_box_not_found(operation_id=operation_id, performed_by=request.user)
    except (TypeError, ValueError) as exc:
        messages.error(request, str(exc))
    else:
        messages.warning(
            request,
            "Результат сохранен: короб не найден. Остаток, место и карантинный резерв не изменены.",
        )
    return redirect("sklad:inventory_journal")


def sku_stock(request):
    page = build_sku_stock_page(request=request)
    if isinstance(page, HttpResponseForbidden):
        return page
    context = page["context"]
    paginator = Paginator(context.get("rows") or [], 100)
    page_obj = paginator.get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)
    context["rows"] = list(page_obj.object_list)
    context["page_obj"] = page_obj
    context["page_query"] = query_params.urlencode()
    context["page_links"] = inventory_pagination_links(page_obj)
    return render(request, page["template_name"], context)


def inventory_journal_pallet_boxes(request):
    pallet_code = str(request.GET.get("pallet") or "").strip()
    if not pallet_code or pallet_code == "-":
        return JsonResponse({"ok": False, "error": "pallet_required", "boxes": []}, status=400)

    agency_scope = _inventory_journal_agency_scope(request)
    if isinstance(agency_scope, HttpResponseForbidden):
        return agency_scope

    snapshots = (
        WarehouseStockSnapshot.objects.filter(
            is_archived=False,
            qty__gt=0,
            parent_container__container_code__iexact=pallet_code,
        )
        .select_related("agency", "container", "parent_container")
        .order_by("container__container_code", "id")
    )
    if agency_scope is not None:
        snapshots = snapshots.filter(agency=agency_scope)

    boxes = []
    seen_box_codes = set()
    for snapshot in snapshots.iterator(chunk_size=500):
        container = snapshot.container
        box_code = str(getattr(container, "container_code", "") or snapshot.container_code or "").strip()
        if not box_code or box_code == "-" or box_code in seen_box_codes:
            continue
        seen_box_codes.add(box_code)
        boxes.append(
            {
                "box_code": box_code,
                "pallet_code": pallet_code,
                "client_label": _agency_journal_label(snapshot.agency),
                "sequence_number": len(boxes) + 1,
            }
        )

    response_payload = {
        "ok": True,
        "pallet_code": pallet_code,
        "boxes": boxes,
        "count": len(boxes),
    }
    detail_requested = str(request.GET.get("detail") or "").strip().lower() in {"1", "true", "yes"}
    export_requested = str(request.GET.get("export") or "").strip().lower() in {"xlsx", "excel"}
    if detail_requested or export_requested:
        agency_ids = {agency_scope.id} if agency_scope is not None else set()
        detail_payload = _pallet_detail_payload(pallet_code, agency_ids)
        if export_requested:
            return _export_inventory_pallet_excel(pallet_code, detail_payload["rows"])
        response_payload.update(detail_payload)
        response_payload["boxes"] = boxes
        response_payload["count"] = len(boxes)
    return JsonResponse(response_payload)


def _inventory_journal_agency_scope(request):
    """Return None for unrestricted staff, or the only agency visible to the request."""
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Access denied")
    role = get_request_role(request)
    if request.user.is_staff or is_staff_role(role):
        client_id = request.GET.get("client") or request.GET.get("agency")
        if not client_id:
            return None
        agency = Agency.objects.filter(pk=client_id).first()
        return agency if agency is not None else HttpResponseForbidden("Access denied")
    agency = _client_agency_for_request(request)
    return agency if agency is not None else HttpResponseForbidden("Access denied")



def _safe_filename_part(value):
    text = str(value or '').strip()
    text = re.sub(r'[^A-Za-z0-9_.-]+', '_', text)
    return text.strip('._') or 'box'


def _positive_int(value):
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def _box_weight_kg_value(value):
    try:
        grams = float(value or 0)
    except (TypeError, ValueError):
        return None
    if grams <= 0:
        return None
    return round(grams / 1000, 3)


def _box_weight_kg_label(value):
    kg = _box_weight_kg_value(value)
    if kg is None:
        return '-'
    text = f'{kg:.3f}'.rstrip('0').rstrip('.')
    return text.replace('.', ',')


def _unique_nonempty(values):
    result = []
    seen = set()
    for value in values:
        text = str(value or '').strip()
        if not text or text == '-' or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result


def _box_order_display(snapshot, row):
    source_type = str(row.get('source_order_type') or getattr(snapshot, 'source_order_type', '') or '').strip()
    source_id = row.get('source_order_id') or getattr(snapshot, 'source_order_id', None)
    if source_type and source_id:
        try:
            return format_order_number(source_type, source_id)
        except Exception:
            return f'{source_id}_{source_type}'
    active_type = str(row.get('active_context_type') or '').strip()
    active_id = row.get('active_context_id')
    if active_type and active_id:
        try:
            return format_order_number(active_type, active_id)
        except Exception:
            return f'{active_id}_{active_type}'
    active_operation = getattr(snapshot, 'active_operation', None)
    if active_operation is not None:
        operation_type = str(getattr(active_operation, 'operation_type', '') or '').strip()
        operation_id = getattr(active_operation, 'operation_id', None)
        if operation_type and operation_id:
            try:
                return format_order_number(operation_type, operation_id)
            except Exception:
                return f'{operation_id}_{operation_type}'
    return '-'


def _snapshot_location_label(snapshot, row):
    location_obj = getattr(snapshot, 'location', None)
    if location_obj is not None:
        try:
            short_label = _short_location_label(
                zone=getattr(location_obj, 'zone_code', '') or getattr(snapshot, 'zone_code', '') or row.get('zone') or '',
                row=getattr(location_obj, 'row_no', 0),
                section=getattr(location_obj, 'section_no', 0),
                tier=getattr(location_obj, 'tier_no', 0),
                cell=getattr(location_obj, 'cell_no', 0),
                location=getattr(location_obj, 'display_name', '') or getattr(location_obj, 'location_code', ''),
            )
            if short_label and short_label != '-':
                return short_label
        except Exception:
            pass
    location = str(row.get('location_short') or row.get('location') or '').strip()
    if location and location != '-':
        return location
    try:
        return _short_location_label(
            zone=getattr(snapshot, 'zone_code', '') or row.get('zone') or '',
            location=getattr(snapshot, 'location', None),
        )
    except Exception:
        return str(getattr(snapshot, 'zone_code', '') or row.get('zone') or '-') or '-'


def _box_status_parts(snapshot, row):
    qty = _positive_int(row.get('qty'))
    if qty <= 0:
        return []
    state = str(getattr(snapshot, 'warehouse_state_code', '') or '').strip()
    location_label = _snapshot_location_label(snapshot, row)
    if state == 'processing_consumed':
        return [(qty, '\u0421\u043f\u0438\u0441\u0430\u043b\u0438', '\u0421\u043f\u0438\u0441\u0430\u043b\u0438')]

    parts = []
    available_qty = _positive_int(row.get('available_qty'))
    shipping_reserved_qty = _positive_int(row.get('shipping_reserved_qty'))
    processing_reserved_qty = _positive_int(row.get('processing_reserved_qty'))
    other_reserved_qty = _positive_int(row.get('other_reserved_qty'))

    if available_qty:
        parts.append((available_qty, '\u0414\u043e\u0441\u0442\u0443\u043f\u043d\u043e', location_label))
    if shipping_reserved_qty:
        parts.append((shipping_reserved_qty, '\u0420\u0435\u0437\u0435\u0440\u0432 OTG', location_label))
    if processing_reserved_qty:
        parts.append((processing_reserved_qty, '\u0420\u0435\u0437\u0435\u0440\u0432 OBR', location_label))
    if other_reserved_qty:
        parts.append((other_reserved_qty, '\u0420\u0435\u0437\u0435\u0440\u0432', location_label))

    used_qty = sum(part[0] for part in parts)
    remaining_qty = max(0, qty - used_qty)
    if remaining_qty:
        zone = str(row.get('zone') or getattr(snapshot, 'zone_code', '') or '').upper()
        if zone == 'OTG':
            status = '\u0421\u043a\u043b\u0430\u0434 OTG'
        elif zone == 'OBR':
            status = '\u0421\u043a\u043b\u0430\u0434 OBR'
        elif zone == 'PR':
            status = 'PR'
        else:
            status = '\u0414\u043e\u0441\u0442\u0443\u043f\u043d\u043e'
        parts.append((remaining_qty, status, location_label))
    return parts


def _box_detail_rows(box_code, agency_ids):
    snapshots = (
        WarehouseStockSnapshot.objects.filter(is_archived=False, qty__gt=0)
        .filter(
            Q(container_code=box_code)
            | Q(container__container_code=box_code)
            | Q(last_event__payload__source_box_code=box_code)
        )
        .select_related('agency', 'sku_ref', 'container', 'parent_container', 'location', 'active_operation', 'last_event')
        .order_by('container_code', 'sku_code', 'marking_code', 'id')
    )
    if agency_ids:
        snapshots = snapshots.filter(agency_id__in=agency_ids)

    details = []
    for snapshot in snapshots:
        row = normalize_stock_row_from_snapshot(snapshot)
        normalized_box = str(row.get('box_code') or '').strip()
        source_box = ''
        last_event = getattr(snapshot, 'last_event', None)
        payload = getattr(last_event, 'payload', None) if last_event is not None else None
        if isinstance(payload, dict):
            source_box = str(payload.get('source_box_code') or '').strip()
        if normalized_box != box_code and source_box != box_code:
            continue

        marking_code = str(getattr(snapshot, 'marking_code', '') or '').strip()
        parts = _box_status_parts(snapshot, row)
        if not parts:
            continue
        for qty, status_label, location_label in parts:
            details.append({
                'box_code': box_code,
                'pallet_code': str(row.get('pallet_code') or '-') or '-',
                'location': location_label or '-',
                'order_label': _box_order_display(snapshot, row),
                'client_label': _agency_journal_label(getattr(snapshot, 'agency', None)),
                'sku': str(row.get('sku') or '-'),
                'name': str(row.get('name') or '-'),
                'size': str(row.get('size') or '-'),
                'barcode': str(row.get('barcode') or '-'),
                'goods_type': str(row.get('goods_type') or '-'),
                'goods_type_label': _journal_goods_type_hint(row.get('goods_type')),
                'qty': qty,
                'status': status_label,
                'chz_label': '\u0441 \u0427\u0417' if marking_code else '\u0431\u0435\u0437 \u0427\u0417',
                'marking_code': marking_code,
                'box_size': str(row.get('box_size') or '-'),
                'box_weight_kg': _box_weight_kg_label(row.get('box_weight')),
                'box_weight_kg_value': _box_weight_kg_value(row.get('box_weight')),
            })

    details.sort(key=lambda item: (
        item.get('pallet_code') or '',
        item.get('location') or '',
        item.get('sku') or '',
        item.get('marking_code') or '',
        item.get('status') or '',
    ))
    return details


def _pallet_detail_rows(pallet_code, agency_ids):
    snapshots = (
        WarehouseStockSnapshot.objects.filter(is_archived=False, qty__gt=0)
        .filter(
            Q(parent_container__container_code__iexact=pallet_code)
            | Q(container__container_code__iexact=pallet_code)
        )
        .select_related('agency', 'sku_ref', 'container', 'parent_container', 'location', 'active_operation', 'last_event')
        .order_by('container__container_code', 'sku_code', 'marking_code', 'id')
    )
    if agency_ids:
        snapshots = snapshots.filter(agency_id__in=agency_ids)

    details = []
    for snapshot in snapshots:
        row = normalize_stock_row_from_snapshot(snapshot)
        normalized_pallet = str(row.get('pallet_code') or '').strip()
        container = getattr(snapshot, 'container', None)
        if normalized_pallet.lower() != pallet_code.lower() and str(getattr(container, 'container_code', '') or '').strip().lower() != pallet_code.lower():
            continue
        box_code = '-'
        if container is not None and getattr(container, 'container_type', '') == WarehouseContainer.TYPE_BOX:
            box_code = str(getattr(container, 'container_code', '') or row.get('box_code') or '-').strip() or '-'
        marking_code = str(getattr(snapshot, 'marking_code', '') or '').strip()
        parts = _box_status_parts(snapshot, row)
        for qty, status_label, location_label in parts:
            details.append({
                'box_code': box_code,
                'pallet_code': pallet_code,
                'location': location_label or '-',
                'order_label': _box_order_display(snapshot, row),
                'client_label': _agency_journal_label(getattr(snapshot, 'agency', None)),
                'sku': str(row.get('sku') or '-'),
                'name': str(row.get('name') or '-'),
                'size': str(row.get('size') or '-'),
                'barcode': str(row.get('barcode') or '-'),
                'goods_type': str(row.get('goods_type') or '-'),
                'goods_type_label': _journal_goods_type_hint(row.get('goods_type')),
                'qty': qty,
                'status': status_label,
                'chz_label': 'с ЧЗ' if marking_code else 'без ЧЗ',
                'marking_code': marking_code,
                'box_size': str(row.get('box_size') or '-'),
                'box_weight_kg': _box_weight_kg_label(row.get('box_weight')),
                'box_weight_kg_value': _box_weight_kg_value(row.get('box_weight')),
            })

    details.sort(key=lambda item: (
        item.get('box_code') or '',
        item.get('location') or '',
        item.get('sku') or '',
        item.get('marking_code') or '',
        item.get('status') or '',
    ))
    return details


WAREHOUSE_EVENT_LABELS = {
    'receiving_arrived': 'Товар прибыл на приемку',
    'placement_started': 'Размещение начато',
    'placement_completed': 'Размещение завершено',
    'putaway_requested': 'Запрошено размещение',
    'putaway_completed': 'Размещение в хранение завершено',
    'processing_requested': 'Запрошено перемещение в обработку',
    'processing_reserved': 'Создан резерв под обработку',
    'processing_reserve_released': 'Снят резерв обработки',
    'processing_zone_arrived': 'Доставлено в зону обработки',
    'processing_started': 'Обработка начата',
    'processing_completed': 'Обработка завершена',
    'processing_consumed': 'Товар использован в обработке',
    'shipping_reserved': 'Создан резерв под отгрузку',
    'shipping_reserve_released': 'Снят резерв отгрузки',
    'otg_requested': 'Запрошено перемещение в OTG',
    'otg_arrived': 'Доставлено в зону OTG',
    'palletization_started': 'Паллетизация начата',
    'palletization_completed': 'Паллетизация завершена',
    'ready_for_loading': 'Готово к погрузке',
    'assigned_to_trip': 'Назначено в рейс',
    'loading_started': 'Погрузка начата',
    'loaded_to_vehicle': 'Погружено в автомобиль',
    'shipped': 'Отгружено со склада',
    'movement_requested': 'Запрошено перемещение',
    'movement_task_created': 'Создано задание на перемещение',
    'movement_started': 'Перемещение начато',
    'movement_completed': 'Перемещение завершено',
    'movement_canceled': 'Перемещение отменено',
    'stock_returned_to_storage': 'Возвращено в хранение',
    'warehouse_context_canceled': 'Складская операция отменена',
    'missing_box_reported': 'Короб отмечен как отсутствующий',
}


WAREHOUSE_ROLE_LABELS = {
    'storekeeper': 'Кладовщик',
    'manager': 'Менеджер',
    'head_manager': 'Начальник склада',
    'reachtruck': 'Ричтрак',
    'reachtruck_driver': 'Водитель ричтрака',
    'admin': 'Администратор',
    'client': 'Клиент',
}


def _warehouse_reference_label(context_type, context_id):
    context_type = str(context_type or '').strip()
    context_id = str(context_id or '').strip()
    if not context_id:
        return '-'
    try:
        return format_order_number(context_type, context_id)
    except Exception:
        return context_id


def _warehouse_location_label(location, zone_code=''):
    if location is not None:
        location_text = str(
            getattr(location, 'display_name', '')
            or getattr(location, 'location_code', '')
            or ''
        ).strip()
        return _short_location_label(
            zone=getattr(location, 'zone_code', '') or zone_code,
            row=getattr(location, 'row_no', 0),
            section=getattr(location, 'section_no', 0),
            tier=getattr(location, 'tier_no', 0),
            cell=getattr(location, 'cell_no', 0),
            location=location_text,
        )
    return _short_location_label(zone=zone_code) if zone_code else '-'


def _warehouse_user_label(user):
    if user is not None:
        employee = getattr(user, 'employee_profile', None)
        employee_name = str(getattr(employee, 'full_name', '') or '').strip()
        full_name = str(user.get_full_name() or '').strip()
        username = str(user.get_username() or '').strip()
        if employee_name:
            return employee_name
        if full_name:
            return full_name
        if username:
            return username
    return '-'


def _warehouse_actor_label(event, fallback_label=''):
    user = getattr(event, 'performed_by', None)
    user_label = _warehouse_user_label(user)
    if user_label != '-':
        return user_label
    fallback_label = str(fallback_label or '').strip()
    if fallback_label:
        return fallback_label
    role = str(event.performed_by_role or '').strip()
    return WAREHOUSE_ROLE_LABELS.get(role, role) or '-'


def _warehouse_event_order_key(event):
    context_candidates = (
        getattr(event, 'stock_context_type', ''),
        getattr(event, 'source_document_type', ''),
    )
    order_type = ''
    for candidate in context_candidates:
        normalized = str(candidate or '').strip().lower()
        for supported in ('receiving', 'processing', 'shipping'):
            if normalized == supported or normalized.startswith(f'{supported}_'):
                order_type = supported
                break
        if order_type:
            break
    order_id = str(
        getattr(event, 'stock_context_id', '')
        or getattr(event, 'source_document_id', '')
        or ''
    ).strip()
    return order_type, order_id


def _warehouse_audit_actor_labels(events):
    labels = {}
    cache = {}
    for event in events:
        if getattr(event, 'performed_by_id', None):
            continue
        occurred_at = getattr(event, 'occurred_at', None)
        order_type, order_id = _warehouse_event_order_key(event)
        if not occurred_at or not order_type or not order_id:
            continue
        minute_key = occurred_at.replace(second=0, microsecond=0)
        cache_key = (order_type, order_id, minute_key)
        if cache_key not in cache:
            candidates = list(
                OrderAuditEntry.objects.filter(
                    order_type=order_type,
                    order_id=order_id,
                    user__isnull=False,
                    created_at__gte=occurred_at - timedelta(minutes=5),
                    created_at__lte=occurred_at + timedelta(minutes=5),
                )
                .select_related('user', 'user__employee_profile')
                .order_by('created_at', 'id')
            )
            closest = min(
                candidates,
                key=lambda entry: abs((entry.created_at - occurred_at).total_seconds()),
                default=None,
            )
            cache[cache_key] = _warehouse_user_label(getattr(closest, 'user', None))
        label = cache.get(cache_key) or '-'
        if label != '-':
            labels[event.id] = label
    return labels


def _box_container_and_history(box_code, agency_ids):
    containers = (
        WarehouseContainer.objects.filter(container_code=box_code)
        .select_related('agency', 'parent_container', 'current_location', 'created_by')
        .order_by('-updated_at', '-id')
    )
    if agency_ids:
        containers = containers.filter(agency_id__in=agency_ids)
    container_rows = list(containers)
    container_ids = [container.id for container in container_rows]

    snapshots = WarehouseStockSnapshot.objects.filter(
        Q(container_id__in=container_ids)
        | Q(container_code=box_code)
        | Q(last_event__payload__source_box_code=box_code)
    )
    if agency_ids:
        snapshots = snapshots.filter(agency_id__in=agency_ids)
    snapshot_links = list(snapshots.values_list('id', 'last_event_id'))
    snapshot_ids = [snapshot_id for snapshot_id, _last_event_id in snapshot_links]
    snapshot_keys = snapshot_ids + [str(snapshot_id) for snapshot_id in snapshot_ids]
    last_event_ids = [last_event_id for _snapshot_id, last_event_id in snapshot_links if last_event_id]

    event_match = (
        Q(container_id__in=container_ids)
        | Q(id__in=last_event_ids)
        | Q(payload__source_box_code=box_code)
        | Q(payload__target_box_code=box_code)
        | Q(payload__box_code=box_code)
        | Q(payload__container_code=box_code)
    )
    if snapshot_keys:
        event_match |= Q(payload__snapshot_id__in=snapshot_keys)

    events = (
        WarehouseEvent.objects.filter(event_match)
        .select_related('container', 'operation', 'from_location', 'to_location', 'performed_by', 'performed_by__employee_profile')
        .order_by('-occurred_at', '-id')
        .distinct()
    )
    if agency_ids:
        events = events.filter(agency_id__in=agency_ids)

    history_events = list(events[:5000])
    audit_actor_labels = _warehouse_audit_actor_labels(history_events)
    history = []
    for event in history_events:
        document_type = event.source_document_type or event.stock_context_type
        document_id = event.source_document_id or event.stock_context_id
        operation = getattr(event, 'operation', None)
        if not document_id and operation is not None:
            document_type = operation.source_document_type or operation.context_type
            document_id = operation.source_document_id or operation.context_id
        occurred_at = event.occurred_at
        if occurred_at and timezone.is_aware(occurred_at):
            occurred_at = timezone.localtime(occurred_at)
        history.append({
            'occurred_at': occurred_at.strftime('%d.%m.%Y %H:%M') if occurred_at else '-',
            'event_type': event.event_type,
            'event_label': WAREHOUSE_EVENT_LABELS.get(event.event_type, event.event_type.replace('_', ' ')),
            'document': _warehouse_reference_label(document_type, document_id),
            'container_code': str(getattr(event.container, 'container_code', '') or box_code or '-'),
            'qty': _positive_int(event.qty),
            'from_location': _warehouse_location_label(event.from_location, event.from_zone_code),
            'to_location': _warehouse_location_label(event.to_location, event.to_zone_code),
            'performed_by': _warehouse_actor_label(event, audit_actor_labels.get(event.id)),
            'role': WAREHOUSE_ROLE_LABELS.get(event.performed_by_role, event.performed_by_role) or '-',
            'event_id': event.id,
            'operation_id': event.operation_id,
        })

    container = container_rows[0] if container_rows else None
    if container is None:
        return None, history
    created_at = container.created_at
    if created_at and timezone.is_aware(created_at):
        created_at = timezone.localtime(created_at)
    created_by = getattr(container, 'created_by', None)
    created_by_label = '-'
    if created_by is not None:
        created_by_label = str(created_by.get_full_name() or created_by.get_username() or '-').strip() or '-'
    passport = {
        'client': _agency_journal_label(container.agency),
        'type': container.get_container_type_display(),
        'status': container.get_status_display(),
        'pallet': str(getattr(container.parent_container, 'container_code', '') or '-'),
        'location': _warehouse_location_label(container.current_location),
        'source': _warehouse_reference_label(container.source_context_type, container.source_context_id),
        'created_at': created_at.strftime('%d.%m.%Y %H:%M') if created_at else '-',
        'created_by': created_by_label,
        'updated_at': (
            timezone.localtime(container.updated_at).strftime('%d.%m.%Y %H:%M')
            if timezone.is_aware(container.updated_at)
            else container.updated_at.strftime('%d.%m.%Y %H:%M')
        ),
    }
    return passport, history


def _pallet_container_and_history(pallet_code, agency_ids):
    containers = (
        WarehouseContainer.objects.filter(container_code__iexact=pallet_code)
        .select_related('agency', 'current_location', 'created_by')
        .order_by('-updated_at', '-id')
    )
    if agency_ids:
        containers = containers.filter(agency_id__in=agency_ids)
    container_rows = list(containers)
    pallet_ids = [container.id for container in container_rows]

    snapshots = WarehouseStockSnapshot.objects.filter(
        Q(parent_container_id__in=pallet_ids)
        | Q(parent_container__container_code__iexact=pallet_code)
        | Q(container_id__in=pallet_ids)
        | Q(container__container_code__iexact=pallet_code)
    )
    if agency_ids:
        snapshots = snapshots.filter(agency_id__in=agency_ids)
    snapshot_links = list(snapshots.values_list('id', 'last_event_id', 'container_id'))
    snapshot_ids = [snapshot_id for snapshot_id, _last_event_id, _container_id in snapshot_links]
    snapshot_keys = snapshot_ids + [str(snapshot_id) for snapshot_id in snapshot_ids]
    last_event_ids = [last_event_id for _snapshot_id, last_event_id, _container_id in snapshot_links if last_event_id]
    child_container_ids = [
        container_id
        for _snapshot_id, _last_event_id, container_id in snapshot_links
        if container_id and container_id not in pallet_ids
    ]

    event_match = (
        Q(container_id__in=pallet_ids + child_container_ids)
        | Q(id__in=last_event_ids)
        | Q(payload__pallet_code=pallet_code)
        | Q(payload__source_pallet_code=pallet_code)
        | Q(payload__target_pallet_code=pallet_code)
        | Q(payload__container_code=pallet_code)
    )
    if snapshot_keys:
        event_match |= Q(payload__snapshot_id__in=snapshot_keys)

    events = (
        WarehouseEvent.objects.filter(event_match)
        .select_related('container', 'operation', 'from_location', 'to_location', 'performed_by', 'performed_by__employee_profile')
        .order_by('-occurred_at', '-id')
        .distinct()
    )
    if agency_ids:
        events = events.filter(agency_id__in=agency_ids)

    history_events = list(events[:5000])
    audit_actor_labels = _warehouse_audit_actor_labels(history_events)
    history = []
    for event in history_events:
        document_type = event.source_document_type or event.stock_context_type
        document_id = event.source_document_id or event.stock_context_id
        operation = getattr(event, 'operation', None)
        if not document_id and operation is not None:
            document_type = operation.source_document_type or operation.context_type
            document_id = operation.source_document_id or operation.context_id
        occurred_at = event.occurred_at
        if occurred_at and timezone.is_aware(occurred_at):
            occurred_at = timezone.localtime(occurred_at)
        history.append({
            'occurred_at': occurred_at.strftime('%d.%m.%Y %H:%M') if occurred_at else '-',
            'event_type': event.event_type,
            'event_label': WAREHOUSE_EVENT_LABELS.get(event.event_type, event.event_type.replace('_', ' ')),
            'document': _warehouse_reference_label(document_type, document_id),
            'container_code': str(getattr(event.container, 'container_code', '') or pallet_code or '-'),
            'qty': _positive_int(event.qty),
            'from_location': _warehouse_location_label(event.from_location, event.from_zone_code),
            'to_location': _warehouse_location_label(event.to_location, event.to_zone_code),
            'performed_by': _warehouse_actor_label(event, audit_actor_labels.get(event.id)),
            'role': WAREHOUSE_ROLE_LABELS.get(event.performed_by_role, event.performed_by_role) or '-',
            'event_id': event.id,
            'operation_id': event.operation_id,
        })

    container = container_rows[0] if container_rows else None
    if container is None:
        return None, history
    created_at = container.created_at
    if created_at and timezone.is_aware(created_at):
        created_at = timezone.localtime(created_at)
    created_by = getattr(container, 'created_by', None)
    created_by_label = '-'
    if created_by is not None:
        created_by_label = str(created_by.get_full_name() or created_by.get_username() or '-').strip() or '-'
    updated_at = container.updated_at
    if updated_at and timezone.is_aware(updated_at):
        updated_at = timezone.localtime(updated_at)
    passport = {
        'client': _agency_journal_label(container.agency),
        'type': container.get_container_type_display(),
        'status': container.get_status_display(),
        'location': _warehouse_location_label(container.current_location),
        'source': _warehouse_reference_label(container.source_context_type, container.source_context_id),
        'created_at': created_at.strftime('%d.%m.%Y %H:%M') if created_at else '-',
        'created_by': created_by_label,
        'updated_at': updated_at.strftime('%d.%m.%Y %H:%M') if updated_at else '-',
    }
    return passport, history


def _box_detail_payload(box_code, agency_ids):
    rows = _box_detail_rows(box_code, agency_ids)
    container, history = _box_container_and_history(box_code, agency_ids)
    total_qty = sum(_positive_int(row.get('qty')) for row in rows)
    with_chz_qty = sum(_positive_int(row.get('qty')) for row in rows if row.get('marking_code'))
    without_chz_qty = total_qty - with_chz_qty
    return {
        'box_code': box_code,
        'rows': rows,
        'total_qty': total_qty,
        'with_chz_qty': with_chz_qty,
        'without_chz_qty': without_chz_qty,
        'pallets': _unique_nonempty(row.get('pallet_code') for row in rows),
        'locations': _unique_nonempty(row.get('location') for row in rows),
        'orders': _unique_nonempty(row.get('order_label') for row in rows),
        'container': container,
        'history': history,
    }


def _pallet_detail_payload(pallet_code, agency_ids):
    rows = _pallet_detail_rows(pallet_code, agency_ids)
    container, history = _pallet_container_and_history(pallet_code, agency_ids)
    total_qty = sum(_positive_int(row.get('qty')) for row in rows)
    with_chz_qty = sum(_positive_int(row.get('qty')) for row in rows if row.get('marking_code'))
    boxes = _unique_nonempty(row.get('box_code') for row in rows)
    return {
        'pallet_code': pallet_code,
        'rows': rows,
        'total_qty': total_qty,
        'with_chz_qty': with_chz_qty,
        'without_chz_qty': total_qty - with_chz_qty,
        'boxes_count': len(boxes),
        'locations': _unique_nonempty(row.get('location') for row in rows),
        'orders': _unique_nonempty(row.get('order_label') for row in rows),
        'container': container,
        'history': history,
    }


@role_required('storekeeper', 'manager', 'admin')
def inventory_journal_box_info(request):
    box_code = str(request.GET.get('box') or '').strip()
    if not box_code or box_code == '-':
        raise Http404

    agency_scope = _inventory_journal_agency_scope(request)
    if isinstance(agency_scope, HttpResponseForbidden):
        return agency_scope
    agency_ids = {agency_scope.id} if agency_scope is not None else set()

    payload = _box_detail_payload(box_code, agency_ids)
    if str(request.GET.get('export') or '').lower() in {'xlsx', 'excel'}:
        return _export_inventory_box_excel(box_code, payload['rows'])
    return JsonResponse(payload)


def _excel_safe_marking_code(value):
    text = str(value or '')

    def replace_control_character(match):
        character = match.group(0)
        return '<GS>' if character == '\x1d' else f'\\x{ord(character):02X}'

    return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', replace_control_character, text)


def _export_inventory_box_excel(box_code, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Box'
    headers = [
        '\u041a\u043e\u0440\u043e\u0431',
        '\u041f\u0430\u043b\u043b\u0435\u0442\u0430',
        '\u041c\u0435\u0441\u0442\u043e',
        '\u0417\u0430\u044f\u0432\u043a\u0430',
        '\u041a\u043b\u0438\u0435\u043d\u0442',
        'SKU',
        '\u041d\u0430\u0438\u043c\u0435\u043d\u043e\u0432\u0430\u043d\u0438\u0435',
        '\u0420\u0430\u0437\u043c\u0435\u0440 \u0442\u043e\u0432\u0430\u0440\u0430',
        '\u0428\u041a',
        '\u0422\u0438\u043f \u0442\u043e\u0432\u0430\u0440\u0430',
        '\u041a\u043e\u043b-\u0432\u043e',
        '\u0421\u0442\u0430\u0442\u0443\u0441',
        '\u0427\u0417',
        '\u041a\u043e\u0434 \u0427\u0417',
        '\u0420\u0430\u0437\u043c\u0435\u0440 \u043a\u043e\u0440\u043e\u0431\u0430',
        '\u0412\u0435\u0441 \u043a\u043e\u0440\u043e\u0431\u0430, \u043a\u0433',
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row.get('box_code') or '-',
            row.get('pallet_code') or '-',
            row.get('location') or '-',
            row.get('order_label') or '-',
            row.get('client_label') or '-',
            row.get('sku') or '-',
            row.get('name') or '-',
            row.get('size') or '-',
            row.get('barcode') or '-',
            row.get('goods_type_label') or row.get('goods_type') or '-',
            row.get('qty') or 0,
            row.get('status') or '-',
            row.get('chz_label') or '-',
            _excel_safe_marking_code(row.get('marking_code')),
            row.get('box_size') or '-',
            row.get('box_weight_kg_value') if row.get('box_weight_kg_value') is not None else '',
        ])
    for column in sheet.columns:
        max_length = max(len(str(cell.value or '')) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_length + 2, 12), 45)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    filename = f'inventory_box_{_safe_filename_part(box_code)}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _export_inventory_pallet_excel(pallet_code, rows):
    response = _export_inventory_box_excel(pallet_code, rows)
    filename = f'inventory_pallet_{_safe_filename_part(pallet_code)}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response

def _export_inventory_excel(page):
    rows = list(page.get("context", {}).get("rows") or [])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Остатки"
    processing_source_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
    sheet.append(
        [
            "Дата",
            "Паллета",
            "Короб",
            "Заявка",
            "Клиент",
            "SKU",
            "Штрихкод",
            "Наименование",
            "Размер",
            "Вес, г",
            "Тип товара",
            "Место",
            "Всего",
            "ФБС",
            "Основной",
            "Склад обр",
            "Склад отг",
            "Доступно",
            "Резерв обр",
            "В обработке",
            "Резерв отг",
        ]
    )
    for row in rows:
        created_at = row.get("created_at")
        created_at_text = (
            timezone.localtime(created_at).strftime("%d.%m.%Y %H:%M")
            if created_at
            else ""
        )
        sheet.append(
            [
                created_at_text,
                row.get("pallet_code") or "",
                row.get("box_code") or "",
                row.get("order_display") or "",
                row.get("client_label") or "",
                row.get("sku") or "",
                row.get("barcode") or "",
                row.get("name") or "",
                row.get("box_size") or "",
                row.get("box_weight") or "",
                row.get("goods_type") or "",
                row.get("location_short") or row.get("location") or "",
                int(row.get("qty") or 0),
                int(row.get("fbs_qty") or 0),
                int(row.get("stock_main_qty") or 0),
                int(row.get("stock_processing_qty") or 0),
                int(row.get("stock_otg_qty") or 0),
                int(row.get("available_qty") or 0),
                int(row.get("processing_reserved_qty") or 0),
                int(row.get("processing_in_progress_qty") or 0),
                int(row.get("shipping_reserved_qty") or 0),
            ]
        )
        if row.get("is_processing_source_box"):
            cell = sheet.cell(row=sheet.max_row, column=3)
            if str(cell.value or "").strip():
                cell.fill = processing_source_fill

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="storekeeper_stock_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response
