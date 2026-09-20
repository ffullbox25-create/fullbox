"""Focused read-only tests for the client FBS shipped ЧЗ report."""

from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import RequestFactory, SimpleTestCase, override_settings
from django.urls import resolve
from django.utils import timezone
from openpyxl import load_workbook

from client_cabinet.fbs_lk import (
    api_fbs_marking_report_export,
    build_fbs_shipped_marking_workbook,
)


class ClientFbsMarkingReportTests(SimpleTestCase):
    @override_settings(ROOT_URLCONF="client_cabinet.urls")
    def test_export_route_is_client_cabinet_only(self):
        match = resolve("/api/v1/fbs/reports/marking/export/")
        self.assertEqual(match.url_name, "client-api-fbs-marking-report-export")

    @patch("client_cabinet.fbs_lk.FbsMarketplaceMetadataTransfer.objects.filter")
    @patch("client_cabinet.fbs_lk.OrderAuditEntry.objects.filter")
    def test_workbook_uses_dispatch_audit_and_transferred_marking_codes(
        self,
        audit_filter_mock,
        transfer_filter_mock,
    ):
        agency = SimpleNamespace(id=37)
        shipped_at = timezone.make_aware(datetime(2026, 9, 3, 21, 11))
        audit_values = MagicMock()
        audit_values.order_by.return_value = [("6865", shipped_at)]
        audit_filter_mock.return_value.values_list.return_value = audit_values

        profile = SimpleNamespace(marketplace="wb")
        order = SimpleNamespace(
            id=6865,
            profile=profile,
            external_order_id="WB-123",
            internal_status="handed_over",
        )
        item = SimpleNamespace(
            id=91,
            order_id=6865,
            order=order,
            external_sku="SKU-1",
            product_name="Товар",
            sku=None,
            barcode="2040000000001",
        )
        transfer = SimpleNamespace(
            id=44,
            order_item_id=91,
            order_item=item,
            value="0101234567890123" + chr(29) + "21ABC",
            status="confirmed",
        )
        transfer_queryset = MagicMock()
        transfer_queryset.select_related.return_value.order_by.return_value = [transfer]
        transfer_filter_mock.return_value = transfer_queryset

        workbook = build_fbs_shipped_marking_workbook(
            agency=agency,
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 4),
        )
        output = BytesIO()
        workbook.save(output)
        rows = list(load_workbook(BytesIO(output.getvalue())).active.iter_rows(values_only=True))

        audit_kwargs = audit_filter_mock.call_args.kwargs
        self.assertEqual(audit_kwargs["agency"], agency)
        self.assertEqual(
            audit_kwargs["payload__changes__internal_status__to"],
            "handed_over",
        )
        self.assertEqual(
            audit_kwargs["payload__source__startswith"],
            "dispatch_handover_batch",
        )
        transfer_kwargs = transfer_filter_mock.call_args.kwargs
        self.assertEqual(transfer_kwargs["order_item__order_id__in"], (6865,))
        self.assertEqual(transfer_kwargs["order_item__order__profile__agency"], agency)
        self.assertEqual(rows[1][1:5], ("03.09.2026 21:11", "Wildberries", "WB-123", "6865"))
        self.assertEqual(rows[1][8], "0101234567890123<GS>21ABC")
        self.assertEqual(rows[1][9:11], ("Подтверждено", "Передан"))

    @patch("client_cabinet.fbs_lk.build_fbs_shipped_marking_workbook")
    @patch("client_cabinet.fbs_lk._agency_or_response")
    def test_http_export_scopes_to_selected_agency_and_period(
        self,
        agency_mock,
        workbook_mock,
    ):
        agency = SimpleNamespace(id=37)
        agency_mock.return_value = (agency, None)
        workbook = MagicMock()
        workbook.save.side_effect = lambda target: target.write(b"PK-test")
        workbook_mock.return_value = workbook
        request = RequestFactory().get(
            "/client/api/v1/fbs/reports/marking/export/",
            {"date_from": "2026-09-01", "date_to": "2026-09-04", "client": "37"},
        )
        request.user = SimpleNamespace(is_authenticated=True)

        response = api_fbs_marking_report_export(request)

        self.assertEqual(response.status_code, 200)
        workbook_mock.assert_called_once_with(
            agency=agency,
            date_from=date(2026, 9, 1),
            date_to=date(2026, 9, 4),
        )
        self.assertIn(
            "fbs-chestny-znak-2026-09-01_2026-09-04.xlsx",
            response["Content-Disposition"],
        )

    @patch("client_cabinet.fbs_lk.build_fbs_shipped_marking_workbook")
    @patch("client_cabinet.fbs_lk._agency_or_response")
    def test_invalid_period_is_rejected(self, agency_mock, workbook_mock):
        agency_mock.return_value = (SimpleNamespace(id=37), None)
        request = RequestFactory().get(
            "/client/api/v1/fbs/reports/marking/export/",
            {"date_from": "2026-09-04", "date_to": "2026-09-01"},
        )
        request.user = SimpleNamespace(is_authenticated=True)

        response = api_fbs_marking_report_export(request)

        self.assertEqual(response.status_code, 400)
        workbook_mock.assert_not_called()
