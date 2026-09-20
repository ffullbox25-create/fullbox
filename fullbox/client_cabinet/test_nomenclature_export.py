"""Excel export of the client's actual nomenclature."""
from __future__ import annotations

from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from sku.models import Agency, Market, SKU, SKUBarcode
from .api_views import api_nomenclature_export


class NomenclatureExportTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username="nomenclature_export_client",
            password="pwd",
        )
        self.other_user = user_model.objects.create_user(
            username="nomenclature_export_other",
            password="pwd",
        )
        self.agency = Agency.objects.create(
            agn_name="Клиент выгрузки",
            portal_user=self.user,
        )
        self.other_agency = Agency.objects.create(
            agn_name="Чужой клиент",
            portal_user=self.other_user,
        )
        self.market = Market.objects.create(id=99101, name="Ozon")
        self.sku = SKU.objects.create(
            agency=self.agency,
            market=self.market,
            sku_code="ART-001",
            name="Товар клиента",
            brand="Бренд",
            size="M",
            color="Оранжевый",
        )
        SKUBarcode.objects.create(
            sku=self.sku,
            value="04601234567890",
            is_primary=True,
        )
        SKUBarcode.objects.create(
            sku=self.sku,
            value="2048123456789",
            is_primary=False,
        )
        self.other_sku = SKU.objects.create(
            agency=self.other_agency,
            sku_code="FOREIGN-001",
            name="Чужой товар",
        )
        self.factory = RequestFactory()

    def _export(self, agency_id):
        request = self.factory.get(
            f"/client/api/v1/nomenclature/export/?client={agency_id}"
        )
        request.user = self.user
        return api_nomenclature_export(request)

    def test_exports_actual_nomenclature_with_text_barcodes(self):
        response = self._export(self.agency.id)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        worksheet = workbook["Номенклатура"]

        self.assertEqual(worksheet["A2"].value, "ART-001")
        self.assertEqual(worksheet["B2"].value, "Товар клиента")
        self.assertEqual(worksheet["D2"].value, "Ozon")
        self.assertEqual(worksheet["F2"].value, "04601234567890")
        self.assertEqual(worksheet["F2"].number_format, "@")
        self.assertEqual(
            worksheet["G2"].value,
            "04601234567890\n2048123456789",
        )
        self.assertEqual(worksheet["G2"].number_format, "@")

    def test_portal_user_cannot_export_another_clients_nomenclature(self):
        response = self._export(self.other_agency.id)

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        worksheet = workbook["Номенклатура"]
        exported_articles = {
            worksheet.cell(row=row, column=1).value
            for row in range(2, worksheet.max_row + 1)
        }
        self.assertIn("ART-001", exported_articles)
        self.assertNotIn("FOREIGN-001", exported_articles)
