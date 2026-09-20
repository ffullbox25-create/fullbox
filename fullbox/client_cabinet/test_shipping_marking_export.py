"""Focused read-only tests for the client shipping ЧЗ document."""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.test import RequestFactory, SimpleTestCase, override_settings
from django.urls import resolve
from openpyxl import load_workbook

from client_cabinet.api_views import api_shipping_marking_export, build_shipping_marking_workbook
from client_cabinet.client_shipping_view import _build_documents


class ClientShippingMarkingDocumentTests(SimpleTestCase):
    def test_documents_include_separate_marking_register(self):
        order = SimpleNamespace(number="OTG-000187", pk=187)

        documents = _build_documents(order=order, stage_key="ready", agency=SimpleNamespace(id=3))

        document = next(row for row in documents if row["doc_type"] == "marking")
        self.assertEqual(document["title"], "Честный знак по отгрузке (Excel)")
        self.assertEqual(
            document["url"],
            "/client/api/v1/requests/shipping/OTG-000187/marking/export/",
        )
        self.assertTrue(document["download"])

    @override_settings(ROOT_URLCONF="client_cabinet.urls")
    def test_marking_export_route_is_client_cabinet_only(self):
        match = resolve("/api/v1/requests/shipping/OTG-000187/marking/export/")
        self.assertEqual(match.url_name, "client-api-shipping-marking-export")

    @patch("client_cabinet.api_views.WarehouseStockSnapshot.objects.filter")
    def test_workbook_reads_active_and_archived_snapshots_and_escapes_gs(self, filter_mock):
        agency = SimpleNamespace(id=3)
        pallet = SimpleNamespace(container_code="PAL-1")
        box = SimpleNamespace(container_code="BOX-1", parent_container=pallet)
        latest = SimpleNamespace(
            marking_code="0101234567890123" + chr(29) + "21ABC",
            container=box,
            container_code="BOX-1",
            parent_container=pallet,
            sku_code="SKU-1",
            name="Товар",
            size="M",
            barcode="2040000000001",
            warehouse_state_code="SHIPPED",
            is_archived=True,
        )
        duplicate = SimpleNamespace(**latest.__dict__)
        queryset = MagicMock()
        queryset.select_related.return_value.order_by.return_value = [latest, duplicate]
        filter_mock.return_value = queryset

        workbook = build_shipping_marking_workbook(agency=agency, order_number="OTG-000187")
        output = BytesIO()
        workbook.save(output)
        loaded = load_workbook(BytesIO(output.getvalue()))
        rows = list(loaded.active.iter_rows(values_only=True))

        filter_mock.assert_called_once_with(
            agency=agency,
            last_event__stock_context_type="shipping",
            last_event__stock_context_id="OTG-000187",
            marking_code__gt="",
        )
        self.assertNotIn("is_archived", filter_mock.call_args.kwargs)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2:5], ("PAL-1", "BOX-1", "SKU-1"))
        self.assertEqual(rows[1][8], "0101234567890123<GS>21ABC")
        self.assertEqual(rows[1][10], "Да")

    @patch("client_cabinet.api_views.build_shipping_marking_workbook")
    @patch("shipping.models.ShippingOrder.objects.filter")
    @patch("client_cabinet.api_views._ctx")
    def test_http_export_scopes_order_to_selected_client(self, ctx_mock, order_filter_mock, workbook_mock):
        agency = SimpleNamespace(id=3)
        request = RequestFactory().get(
            "/client/api/v1/requests/shipping/OTG-000187/marking/export/?client=3"
        )
        request.user = SimpleNamespace(is_authenticated=True)
        ctx_mock.return_value = (agency, True, True)
        order = SimpleNamespace(number="OTG-000187")
        order_filter_mock.return_value.only.return_value.first.return_value = order
        workbook = MagicMock()

        def save(target):
            target.write(b"PK-test")

        workbook.save.side_effect = save
        workbook_mock.return_value = workbook

        response = api_shipping_marking_export(request, "OTG-000187")

        self.assertEqual(response.status_code, 200)
        order_filter_mock.assert_called_once_with(agency=agency, number="OTG-000187")
        workbook_mock.assert_called_once_with(agency=agency, order_number="OTG-000187")
        self.assertIn("chestny-znak-OTG-000187.xlsx", response["Content-Disposition"])
