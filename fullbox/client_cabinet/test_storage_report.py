from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import json

from django.contrib.auth import get_user_model
from django.template.loader import get_template
from django.test import RequestFactory, TestCase
from django.utils import timezone
from openpyxl import load_workbook

from billing.models import BillingStorageDay, StorageSnapshotLine
from client_cabinet.models import AgencyPortalMember
from client_cabinet.portal_access import SECTION_REQUESTS
from client_cabinet.storage_report import api_storage_report, export_storage_report
from sku.models import Agency


User = get_user_model()


class ClientStorageReportTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username="storage_owner", password="pwd")
        self.agency = Agency.objects.create(agn_name="ООО Отчёт", portal_user=self.owner)
        self.other_owner = User.objects.create_user(username="storage_other", password="pwd")
        self.other_agency = Agency.objects.create(agn_name="ООО Другой", portal_user=self.other_owner)
        self.today = timezone.localdate()
        self.yesterday = self.today - timedelta(days=1)

        self.day = BillingStorageDay.objects.create(
            client=self.agency,
            day=self.today,
            pallet_count=2,
            box_count=3,
            sku_unit_count=15,
            physical_volume_m3=Decimal("1.250000"),
            billable_volume_m3=Decimal("1.500000"),
            amount=Decimal("150.00"),
            vat_amount=Decimal("30.00"),
            status=BillingStorageDay.STATUS_CALCULATED,
            payload={"weight_volume_m3": "1.100000"},
        )
        StorageSnapshotLine.objects.create(
            day=self.day,
            sku_code="ART-001",
            name="Тестовый товар",
            barcode="460000000001",
            box_code="BOX-1",
            pallet_code="PAL-1",
            quantity=Decimal("10.000"),
            total_volume_m3=Decimal("0.700000"),
            billable_volume_m3=Decimal("0.840000"),
        )
        StorageSnapshotLine.objects.create(
            day=self.day,
            sku_code="ART-001",
            name="Тестовый товар",
            barcode="460000000001",
            box_code="BOX-2",
            pallet_code="PAL-1",
            quantity=Decimal("5.000"),
            total_volume_m3=Decimal("0.550000"),
            billable_volume_m3=Decimal("0.660000"),
            status=StorageSnapshotLine.STATUS_NO_DIMS,
        )
        other_day = BillingStorageDay.objects.create(
            client=self.other_agency,
            day=self.today,
            sku_unit_count=99,
            physical_volume_m3=Decimal("9.000000"),
        )
        StorageSnapshotLine.objects.create(
            day=other_day,
            sku_code="FOREIGN-999",
            name="Чужой товар",
            quantity=Decimal("99.000"),
            total_volume_m3=Decimal("9.000000"),
        )
        self.factory = RequestFactory()

    def _request(self, path, params=None, *, user=None):
        request = self.factory.get(path, params or {})
        request.user = user or self.owner
        return request

    def test_api_returns_daily_and_article_volume_for_current_client_only(self):
        with self.assertNumQueries(4):
            response = api_storage_report(
                self._request(
                    "/client/api/v1/reports/storage/",
                    {"date_from": self.today.isoformat(), "date_to": self.today.isoformat()},
                )
            )
        self.assertEqual(response.status_code, 200, response.content)
        payload = json.loads(response.content)["data"]
        self.assertEqual(payload["summary"]["days_count"], 1)
        self.assertEqual(payload["summary"]["unit_days"], 15)
        self.assertEqual(payload["summary"]["physical_m3_days"], "1.250000")
        self.assertEqual(len(payload["articles"]), 1)
        row = payload["articles"][0]
        self.assertEqual(row["sku_code"], "ART-001")
        self.assertEqual(row["quantity"], "15.000")
        self.assertEqual(row["box_count"], 2)
        self.assertEqual(row["pallet_count"], 1)
        self.assertEqual(row["physical_volume_m3"], "1.250000")
        self.assertEqual(row["no_dims_count"], 1)
        self.assertNotContains(response, "FOREIGN-999")

    def test_api_filters_by_date_and_search(self):
        outside = BillingStorageDay.objects.create(client=self.agency, day=self.yesterday, sku_unit_count=7)
        StorageSnapshotLine.objects.create(
            day=outside,
            sku_code="OLD-001",
            name="Старый товар",
            quantity=Decimal("7.000"),
        )
        response = api_storage_report(
            self._request(
                "/client/api/v1/reports/storage/",
                {"date_from": self.today.isoformat(), "date_to": self.today.isoformat(), "q": "460000000001"},
            )
        )
        self.assertEqual(response.status_code, 200)
        rows = json.loads(response.content)["data"]["articles"]
        self.assertEqual([row["sku_code"] for row in rows], ["ART-001"])

    def test_excel_export_contains_article_and_daily_sheets(self):
        response = export_storage_report(
            self._request(
                "/client/api/v1/reports/storage/export/",
                {"date_from": self.today.isoformat(), "date_to": self.today.isoformat()},
            )
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["По артикулам", "По дням", "Пояснения"])
        article_rows = list(workbook["По артикулам"].iter_rows(values_only=True))
        self.assertEqual(article_rows[1][1], "ART-001")
        self.assertEqual(article_rows[1][4], 15)
        self.assertNotIn("FOREIGN-999", {row[1] for row in article_rows[1:]})

    def test_employee_without_reports_section_is_forbidden(self):
        employee = User.objects.create_user(username="storage_employee", password="pwd")
        AgencyPortalMember.objects.create(
            agency=self.agency,
            user=employee,
            last_name="Сотрудник",
            first_name="Без отчётов",
            email="storage-employee@example.com",
            role=AgencyPortalMember.ROLE_CUSTOM,
            sections=[SECTION_REQUESTS],
        )
        response = api_storage_report(
            self._request("/client/api/v1/reports/storage/", user=employee)
        )
        self.assertEqual(response.status_code, 403)
        self.assertIn("Недостаточно прав", json.loads(response.content)["error"])

    def test_client_template_contains_storage_report_route_and_lazy_loader(self):
        source = get_template("client_cabinet/dashboard_lk_react.html").template.source
        self.assertIn('data-route="/storage-report"', source)
        self.assertIn('data-page="/storage-report"', source)
        self.assertIn("__lkEnsureStorageReport", source)
        self.assertIn("/client/api/v1/reports/storage/", source)
