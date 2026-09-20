"""Read-only client report for daily storage volume by SKU."""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .portal_access import SECTION_REPORTS
from .portal_members import portal_sections_for_user
from .web_ui import _get_client_for_request


ZERO = Decimal("0")
MAX_REPORT_DAYS = 366
MAX_PAGE_SIZE = 200


def _json_error(message: str, status: int = 400) -> JsonResponse:
    return JsonResponse({"ok": False, "error": message}, status=status)


def _client_for_report(request):
    agency, _client_view, allowed = _get_client_for_request(request)
    if not allowed or not agency:
        return None, _json_error("Доступ запрещен", 403)
    sections = portal_sections_for_user(request.user, agency)
    if sections is not None and SECTION_REPORTS not in set(sections):
        return None, _json_error("Недостаточно прав для раздела «Отчёты».", 403)
    return agency, None


def _parse_date(value: str, *, fallback: date) -> date:
    try:
        return date.fromisoformat(str(value or "").strip())
    except (TypeError, ValueError):
        return fallback


def _report_range(request) -> tuple[date, date]:
    today = timezone.localdate()
    default_start = today.replace(day=1)
    start = _parse_date(request.GET.get("date_from"), fallback=default_start)
    end = _parse_date(request.GET.get("date_to"), fallback=today)
    if end < start:
        start, end = end, start
    if (end - start).days + 1 > MAX_REPORT_DAYS:
        start = end - timedelta(days=MAX_REPORT_DAYS - 1)
    return start, end


def _decimal(value) -> Decimal:
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return ZERO


def _decimal_text(value, places: str = "0.000000") -> str:
    return format(_decimal(value).quantize(Decimal(places), rounding=ROUND_HALF_UP), "f")


def _weight_volume(day_row) -> Decimal:
    payload = day_row.payload if isinstance(day_row.payload, dict) else {}
    return _decimal(payload.get("weight_volume_m3"))


def _storage_days(agency, start: date, end: date):
    from billing.models import BillingStorageDay

    return (
        BillingStorageDay.objects.filter(client=agency, day__range=(start, end))
        .exclude(status=BillingStorageDay.STATUS_CANCELLED)
        .only(
            "id",
            "day",
            "pallet_count",
            "box_count",
            "sku_unit_count",
            "physical_volume_m3",
            "billable_volume_m3",
            "amount",
            "vat_amount",
            "status",
            "payload",
        )
        .order_by("day", "id")
    )


def _article_rows(day_ids: list[int], search: str = ""):
    from billing.models import StorageSnapshotLine

    rows = StorageSnapshotLine.objects.filter(day_id__in=day_ids)
    if search:
        rows = rows.filter(
            Q(sku_code__icontains=search)
            | Q(name__icontains=search)
            | Q(barcode__icontains=search)
        )
    return (
        rows.values("day__day", "sku_code", "name", "barcode")
        .annotate(
            quantity=Sum("quantity"),
            physical_volume_m3=Sum("total_volume_m3"),
            line_billable_volume_m3=Sum("billable_volume_m3"),
            box_count=Count("box_code", distinct=True, filter=~Q(box_code="")),
            pallet_count=Count("pallet_code", distinct=True, filter=~Q(pallet_code="")),
            no_dims_count=Count("id", filter=Q(status=StorageSnapshotLine.STATUS_NO_DIMS)),
        )
        .order_by("day__day", "sku_code", "name", "barcode")
    )


def _serialize_day(day_row) -> dict:
    return {
        "date": day_row.day.isoformat(),
        "date_label": day_row.day.strftime("%d.%m.%Y"),
        "sku_unit_count": int(day_row.sku_unit_count or 0),
        "box_count": int(day_row.box_count or 0),
        "pallet_count": int(day_row.pallet_count or 0),
        "physical_volume_m3": _decimal_text(day_row.physical_volume_m3),
        "weight_volume_m3": _decimal_text(_weight_volume(day_row)),
        "billable_volume_m3": _decimal_text(day_row.billable_volume_m3),
        "amount": _decimal_text(day_row.amount, "0.01"),
        "vat_amount": _decimal_text(day_row.vat_amount, "0.01"),
        "total_amount": _decimal_text(_decimal(day_row.amount) + _decimal(day_row.vat_amount), "0.01"),
        "status": day_row.status,
        "status_label": day_row.get_status_display(),
    }


def _serialize_article(row: dict) -> dict:
    day_value = row.get("day__day")
    return {
        "date": day_value.isoformat() if day_value else "",
        "date_label": day_value.strftime("%d.%m.%Y") if day_value else "",
        "sku_code": str(row.get("sku_code") or ""),
        "name": str(row.get("name") or ""),
        "barcode": str(row.get("barcode") or ""),
        "quantity": _decimal_text(row.get("quantity"), "0.001"),
        "box_count": int(row.get("box_count") or 0),
        "pallet_count": int(row.get("pallet_count") or 0),
        "physical_volume_m3": _decimal_text(row.get("physical_volume_m3")),
        "line_billable_volume_m3": _decimal_text(row.get("line_billable_volume_m3")),
        "no_dims_count": int(row.get("no_dims_count") or 0),
    }


def _report_data(agency, request, *, paginate: bool) -> dict:
    start, end = _report_range(request)
    days = list(_storage_days(agency, start, end))
    day_ids = [row.id for row in days]
    search = str(request.GET.get("q") or "").strip()[:200]
    article_qs = _article_rows(day_ids, search=search) if day_ids else []

    page_number = 1
    pages_count = 1
    total_rows = 0
    if paginate:
        try:
            page_size = min(max(int(request.GET.get("page_size") or 100), 20), MAX_PAGE_SIZE)
        except (TypeError, ValueError):
            page_size = 100
        paginator = Paginator(article_qs, page_size)
        page = paginator.get_page(request.GET.get("page") or 1)
        article_rows = [_serialize_article(row) for row in page.object_list]
        page_number = page.number
        pages_count = paginator.num_pages
        total_rows = paginator.count
    else:
        article_rows = [_serialize_article(row) for row in article_qs]
        total_rows = len(article_rows)

    physical_total = sum((_decimal(row.physical_volume_m3) for row in days), ZERO)
    billable_total = sum((_decimal(row.billable_volume_m3) for row in days), ZERO)
    unit_days = sum(int(row.sku_unit_count or 0) for row in days)
    return {
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "search": search,
        "days": [_serialize_day(row) for row in days],
        "articles": article_rows,
        "summary": {
            "days_count": len(days),
            "article_rows_count": total_rows,
            "unit_days": unit_days,
            "physical_m3_days": _decimal_text(physical_total),
            "billable_m3_days": _decimal_text(billable_total),
        },
        "pagination": {
            "page": page_number,
            "pages": pages_count,
            "total": total_rows,
        },
        "read_only": True,
    }


@login_required
@require_GET
def api_storage_report(request):
    agency, error = _client_for_report(request)
    if error is not None:
        return error
    return JsonResponse({"ok": True, "data": _report_data(agency, request, paginate=True)})


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="F8B800")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="303030")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            value = str(cell.value or "")
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), 48)
    for index, width in widths.items():
        ws.column_dimensions[get_column_letter(index)].width = max(width, 12)


@login_required
@require_GET
def export_storage_report(request):
    agency, error = _client_for_report(request)
    if error is not None:
        return error
    data = _report_data(agency, request, paginate=False)

    workbook = Workbook()
    article_sheet = workbook.active
    article_sheet.title = "По артикулам"
    article_sheet.append([
        "Дата",
        "Артикул",
        "Наименование",
        "Штрих-код",
        "Количество, шт.",
        "Коробов",
        "Палет",
        "Физический объём, м³",
        "Объём строк с коэффициентом, м³",
        "Строк без габаритов",
    ])
    for row in data["articles"]:
        article_sheet.append([
            row["date_label"],
            row["sku_code"],
            row["name"],
            row["barcode"],
            float(row["quantity"]),
            row["box_count"],
            row["pallet_count"],
            float(row["physical_volume_m3"]),
            float(row["line_billable_volume_m3"]),
            row["no_dims_count"],
        ])
    for cell in article_sheet["D"][1:]:
        cell.number_format = "@"
    _style_sheet(article_sheet)

    day_sheet = workbook.create_sheet("По дням")
    day_sheet.append([
        "Дата",
        "Единиц товара",
        "Коробов",
        "Палет",
        "Физический объём, м³",
        "Объём по весу, м³",
        "Тарифицируемый объём дня, м³",
        "Статус расчёта",
        "Сумма без НДС",
        "НДС",
        "Итого",
    ])
    for row in data["days"]:
        day_sheet.append([
            row["date_label"],
            row["sku_unit_count"],
            row["box_count"],
            row["pallet_count"],
            float(row["physical_volume_m3"]),
            float(row["weight_volume_m3"]),
            float(row["billable_volume_m3"]),
            row["status_label"],
            float(row["amount"]),
            float(row["vat_amount"]),
            float(row["total_amount"]),
        ])
    _style_sheet(day_sheet)

    note_sheet = workbook.create_sheet("Пояснения")
    note_sheet.append(["Показатель", "Пояснение"])
    note_sheet.append(["Источник", "Сохранённые дневные снимки хранения Fullbox. Отчёт ничего не меняет на складе."])
    note_sheet.append(["Физический объём", "Сумма объёма товара по артикулу за выбранный день."])
    note_sheet.append(["Тарифицируемый объём дня", "Итог расчёта дня. Он может отличаться от суммы строк из-за правил тарифа и веса."])
    note_sheet.append(["Нет габаритов", "Для части товара физический объём не рассчитан; требуется заполнить размеры в номенклатуре."])
    _style_sheet(note_sheet)

    output = BytesIO()
    workbook.save(output)
    filename = f"storage-by-sku-{data['date_from']}-{data['date_to']}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
