from __future__ import annotations

import json
import time
from io import BytesIO
from typing import Any

from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, OuterRef, Q, Subquery, Sum
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from openpyxl import Workbook

from audit.models import OrderAuditEntry, log_sku_change
from marking.models import MarkingCode
from sklad.models import WarehouseStockSnapshot
from sku.models import SKU, SKUBarcode
from .finance_lk import build_billing_payload, build_finance_summary, list_finance_documents
from .lk_requests import (
    CLIENT_REQUEST_ORDER_TYPES,
    apply_client_request_action,
    build_action_urls,
    build_client_notifications,
    build_request_detail,
    build_request_excel_workbook,
    build_receiving_act_excel_workbook,
    canonical_request_order_id,
    canonical_shipping_order_id,
    lk_request_hash,
    receiving_attention_from_payload,
    request_attention,
    _merge_receiving_act_fields,
)
from .lk_status_map import resolve_lk_entry_status, resolve_lk_request_status
from .request_titles import build_request_display_title
from .marketplace_lk import (
    build_live_dashboard,
    build_live_history,
    build_marketplaces_payload,
    build_mp_stocks_cached_payload,
    build_mp_stocks_payload,
    build_supplies_payload,
    sync_marketplace,
)
from .home_progress import enrich_request_row
from .other_requests import OTHER_REQUEST_CATEGORIES, create_other_request, save_other_attachments
from .web_ui import _get_client_for_request

SHIPPING_DONE_STATUSES = {"shipped", "delivered", "completed", "closed", "done"}
SHIPPING_ACTIVE_STATUSES = {"draft", "reserved", "ready", "loading", "in_progress", "assigned", "created"}


def _chats_enabled() -> bool:
    return bool(getattr(settings, "FULLBOX_CHATS_ENABLED", True))


def _chat_disabled_response():
    return _json_error("Чаты временно отключены", 503)

# Short TTL: SSR «Мои заявки» + сразу же API refresh не должны дважды гонять audit.
_REQUEST_PAYLOADS_CACHE: dict[int, tuple[float, list[dict[str, Any]]]] = {}
_REQUEST_PAYLOADS_TTL_SEC = 25.0
_DASHBOARD_LITE_CACHE_TTL_SEC = 20
_DASHBOARD_LIVE_CACHE_TTL_SEC = 75
_RECEIVING_ACT_PAYLOAD_Q = (
    Q(payload__has_key="act_sent")
    | Q(payload__has_key="act_client_response")
    | Q(payload__has_key="act_client_response_at")
    | Q(payload__has_key="act_viewed")
    | Q(payload__has_key="act_viewed_at")
    | Q(payload__has_key="act_manager_signed")
    | Q(payload__has_key="act_storekeeper_signed")
    | Q(payload__has_key="act_logistician_signed")
)


def _cache_get(key: str):
    try:
        return cache.get(key)
    except Exception:
        return None


def _cache_set(key: str, value: Any, timeout: int) -> None:
    try:
        cache.set(key, value, timeout)
    except Exception:
        return None


def _cache_delete(key: str) -> None:
    try:
        cache.delete(key)
    except Exception:
        return None


def _dashboard_lite_cache_key(agency_id: int) -> str:
    return f"client-cabinet:dashboard-lite:v2:{int(agency_id)}"


def _dashboard_live_cache_key(agency_id: int) -> str:
    return f"client-cabinet:dashboard-live:v2:{int(agency_id)}"


def invalidate_dashboard_lite_cache(agency_id=None) -> None:
    """Drop short-lived client LK home caches."""
    if agency_id is not None:
        agency_id = int(agency_id)
        _cache_delete(_dashboard_lite_cache_key(agency_id))
        _cache_delete(_dashboard_live_cache_key(agency_id))
        return
    try:
        delete_pattern = getattr(cache, "delete_pattern", None)
        if callable(delete_pattern):
            delete_pattern("client-cabinet:dashboard-lite:v1:*")
            delete_pattern("client-cabinet:dashboard-live:v1:*")
            delete_pattern("client-cabinet:dashboard-lite:v2:*")
            delete_pattern("client-cabinet:dashboard-live:v2:*")
    except Exception:
        return None


def invalidate_request_payloads_cache(agency_id=None) -> None:
    """Drop cached «Мои заявки» rows (all clients or one)."""
    if agency_id is None:
        _REQUEST_PAYLOADS_CACHE.clear()
        invalidate_dashboard_lite_cache(None)
        return
    _REQUEST_PAYLOADS_CACHE.pop(int(agency_id), None)
    invalidate_dashboard_lite_cache(agency_id)


def _truthy_param(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y", "on"}


def _json_error(message: str, status: int = 400):
    return JsonResponse({"ok": False, "error": message}, status=status)


def _mojibake_score(value: str) -> int:
    markers = ("\u00d0", "\u00d1", "\ufffd", "\u0420", "\u0421", "\u0432\u0402")
    pairs = (
        "\u0420\u00b0", "\u0420\u00b5", "\u0420\u00b8", "\u0420\u00be", "\u0420\u00bd",
        "\u0420\u00ba", "\u0420\u00bb", "\u0420\u00bc", "\u0420\u00bf", "\u0420\u00a1",
        "\u0421\u0080", "\u0421\u0081", "\u0421\u0082", "\u0421\u008c", "\u0421\u008f",
        "\u0432\u0402", "\u0432\u040e",
    )
    return sum(value.count(marker) for marker in markers) + sum(value.count(pair) * 3 for pair in pairs)


def _repair_mojibake(value: str) -> str:
    if not value:
        return ""
    initial_score = _mojibake_score(value)
    if initial_score == 0:
        return value
    best = value
    best_score = initial_score
    for encoding in ("cp1251", "latin1"):
        try:
            candidate = value.encode(encoding).decode("utf-8")
        except UnicodeError:
            continue
        candidate_score = _mojibake_score(candidate)
        if candidate and candidate_score < best_score:
            best = candidate
            best_score = candidate_score
    return best


def _clean(value: Any) -> str:
    return "" if value is None else _repair_mojibake(str(value).strip())


def _local_dt(value: Any) -> str:
    if not value:
        return ""
    try:
        return timezone.localtime(value).isoformat()
    except Exception:
        try:
            return value.isoformat()
        except Exception:
            return _clean(value)


def _local_date(value: Any) -> str:
    if not value:
        return ""
    try:
        return timezone.localtime(value).strftime("%d.%m.%Y, %H:%M")
    except Exception:
        try:
            return value.strftime("%d.%m.%Y")
        except Exception:
            return _clean(value)


def _value(obj: Any, *names: str, default: Any = "") -> Any:
    for name in names:
        try:
            value = getattr(obj, name)
        except Exception:
            continue
        if value not in (None, "") and not callable(value):
            return value
    return default


def _ctx(request):
    selected_client, client_view, allowed = _get_client_for_request(request)
    if not allowed or not selected_client:
        return selected_client, client_view, False
    from .portal_access import section_for_api_url_name
    from .portal_members import portal_sections_for_user

    url_name = getattr(getattr(request, "resolver_match", None), "url_name", None)
    required_section = section_for_api_url_name(url_name)
    if required_section:
        sections = portal_sections_for_user(request.user, selected_client)
        if sections is not None and required_section not in set(sections):
            return selected_client, client_view, False
    return selected_client, client_view, True


def _client_payload(agency) -> dict[str, Any] | None:
    if not agency:
        return None
    return {
        "id": agency.id,
        "name": _clean(getattr(agency, "agn_name", "")) or f"Клиент {agency.id}",
        "short_name": _clean(getattr(agency, "short_name", "")),
        "inn": _clean(getattr(agency, "inn", "")),
        "phone": _clean(getattr(agency, "phone", "")),
        "email": _clean(getattr(agency, "email", "")),
        "address": _clean(getattr(agency, "fakt_adres", "")) or _clean(getattr(agency, "adres", "")),
    }


def _action_urls(agency) -> dict[str, str]:
    return build_action_urls(agency)


def _stock_qs(agency):
    return WarehouseStockSnapshot.objects.filter(agency=agency, is_archived=False)


def _stock_availability_rows(agency) -> list[dict[str, Any]]:
    """Rows with client-facing availability: physical stock minus active reserves."""
    from sklad.services.stock_availability import StockAvailabilityService

    return StockAvailabilityService.stock_rows_with_availability(agency=agency)


def _stock_row_visible_to_client(row: dict[str, Any], *, hide_consumed: bool = True) -> bool:
    """Single client LK rule for visible stock rows.

    The client sees only units that can be shipped now: positive available_qty,
    not consumed, and not in operational/hidden warehouse states such as
    processing or OTG.
    """
    from sklad.ui_services import _JOURNAL_HIDDEN_WAREHOUSE_STATES

    state_code = str(row.get("warehouse_state_code") or "").strip()
    if state_code in _JOURNAL_HIDDEN_WAREHOUSE_STATES:
        return False
    if int(row.get("available_qty") or 0) <= 0:
        return False
    if hide_consumed and int(row.get("qty") or 0) <= 0:
        return False
    return True


def _stock_summary(agency, availability_rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    qs = _stock_qs(agency)
    rows = availability_rows if availability_rows is not None else _stock_availability_rows(agency)
    visible_rows = [row for row in rows if _stock_row_visible_to_client(row)]
    available = sum(int(row.get("available_qty") or 0) for row in visible_rows)
    # ЛК клиента: на дашборде и в KPI только то, что можно отгрузить.
    # Резерв и выгрузка (available=0) клиенту не показываем.
    total_units = available
    reserved = sum(
        int(row.get("processing_reserved_qty") or 0)
        + int(row.get("shipping_reserved_qty") or 0)
        + int(row.get("other_reserved_qty") or 0)
        for row in rows
    )
    receiving = qs.filter(Q(zone_kind="receiving") | Q(warehouse_state_code__icontains="receiving")).aggregate(value=Sum("qty"))["value"] or 0
    return {
        "total_units": int(total_units or 0),
        "sku_count": SKU.objects.filter(agency=agency, deleted=False).count(),
        "available": int(available or 0),
        # Резерв в клиентском summary не отдаём — иначе снова «виден» в UI.
        "reserved": 0,
        "receiving": int(receiving or 0),
        "no_movement_30_days": 0,
        "reserved_internal": int(reserved or 0),
    }


def _barcode_values(sku) -> list[str]:
    values = []
    for item in SKUBarcode.objects.filter(sku=sku).order_by("-is_primary", "id")[:20]:
        value = _clean(item.value)
        if value:
            values.append(value)
    return values


def _barcode_map_for_skus(sku_ids: list[int]) -> dict[int, list[str]]:
    """Bulk load barcodes for SKU ids (avoids N+1 in stock payload)."""
    if not sku_ids:
        return {}
    result: dict[int, list[str]] = {int(sku_id): [] for sku_id in sku_ids}
    rows = (
        SKUBarcode.objects.filter(sku_id__in=sku_ids)
        .order_by("sku_id", "-is_primary", "id")
        .values_list("sku_id", "value")
    )
    for sku_id, value in rows:
        bucket = result.get(int(sku_id))
        if bucket is None or len(bucket) >= 20:
            continue
        cleaned = _clean(value)
        if cleaned:
            bucket.append(cleaned)
    return result


def _sku_payload(sku, *, barcodes: list[str] | None = None) -> dict[str, Any]:
    barcodes = list(barcodes) if barcodes is not None else _barcode_values(sku)
    image = _clean(getattr(sku, "img", ""))
    market_obj = getattr(sku, "market", None)
    market = _clean(getattr(market_obj, "name", None) or market_obj or "")
    market_upper = market.upper()
    marketplace_codes = []
    if "WB" in market_upper or "WILDBERRIES" in market_upper:
        marketplace_codes.append("WB")
    elif "OZON" in market_upper or "\u041e\u0417\u041e\u041d" in market_upper:
        marketplace_codes.append("OZON")
    elif market:
        marketplace_codes.append(market)
    size = _clean(sku.size)
    updated_at = getattr(sku, "updated_at", None) or getattr(sku, "modified_at", None) or getattr(sku, "created_at", None)
    return {
        "id": sku.id,
        "name": _clean(sku.name) or f"SKU {sku.id}",
        "title": _clean(sku.name) or f"SKU {sku.id}",
        "article": _clean(sku.code) or _clean(sku.sku_code),
        "sku": _clean(sku.sku_code),
        "sku_code": _clean(sku.sku_code),
        "barcode": barcodes[0] if barcodes else "",
        "barcodes": barcodes,
        "brand": _clean(sku.brand),
        "market": market,
        "marketplace": market,
        "marketplaces": [market] if market else [],
        "marketplace_codes": marketplace_codes,
        "category": _clean(sku.tovar_category) or _clean(sku.vid_tovar) or _clean(sku.type_tovar),
        "size": size,
        "sizes": [size] if size else [],
        "color": _clean(sku.color),
        "photo_url": image,
        "marketplace_url": "",
        "updated_at": _local_dt(updated_at),
        "honest_sign": bool(getattr(sku, "honest_sign", False)),
        "weight_kg": float(sku.weight_kg) if sku.weight_kg is not None else None,
        "volume": float(sku.volume) if sku.volume is not None else None,
        "edit_url": f"/client/{sku.agency_id}/sku/{sku.id}/edit/",
        "duplicate_url": f"/client/{sku.agency_id}/sku/{sku.id}/duplicate/",
        "delete_url": f"/client/api/v1/nomenclature/{sku.id}/delete/",
    }


def _stock_product_rows(agency, availability_rows: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    stock_rows = availability_rows if availability_rows is not None else _stock_availability_rows(agency)
    grouped: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for row in stock_rows:
        if not _stock_row_visible_to_client(row):
            continue
        available_qty = int(row.get("available_qty") or 0)
        key = (
            _clean(row.get("sku") or row.get("sku_code")),
            _clean(row.get("name")),
            _clean(row.get("barcode")),
            _clean(row.get("size")),
        )
        item = grouped.setdefault(
            key,
            {
                "sku_code": key[0],
                "name": key[1],
                "barcode": key[2],
                "size": key[3],
                "qty": 0,
                "available_qty": 0,
                "reserved_qty": 0,
            },
        )
        item["qty"] += available_qty
        item["available_qty"] += available_qty
    sorted_rows = sorted(
        grouped.values(),
        key=lambda row: ((_clean(row.get("name")) or "").lower(), (_clean(row.get("sku_code")) or "").lower()),
    )[:200]
    rows = []
    for idx, row in enumerate(sorted_rows, start=1):
        rows.append({
            "id": f"stock-{idx}",
            "name": _clean(row.get("name")),
            "title": _clean(row.get("name")),
            "sku": _clean(row.get("sku_code")),
            "sku_code": _clean(row.get("sku_code")),
            "barcode": _clean(row.get("barcode")),
            "barcodes": [_clean(row.get("barcode"))] if _clean(row.get("barcode")) else [],
            "size": _clean(row.get("size")),
            "qty": int(row.get("qty") or 0),
            "available_qty": int(row.get("available_qty") or 0),
            "reserved_qty": 0,
        })
    return rows


def _stock_payload(agency, availability_rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    availability_rows = availability_rows if availability_rows is not None else _stock_availability_rows(agency)
    summary = _stock_summary(agency, availability_rows=availability_rows)
    products = _stock_product_rows(agency, availability_rows=availability_rows)
    sku_qs = list(SKU.objects.filter(agency=agency, deleted=False).order_by("name", "id")[:120])
    barcode_map = _barcode_map_for_skus([int(sku.id) for sku in sku_qs])
    if not products:
        products = [
            _sku_payload(item, barcodes=barcode_map.get(int(item.id)) or [])
            for item in sku_qs[:80]
        ]
    without_barcode = []
    for sku in sku_qs:
        codes = barcode_map.get(int(sku.id)) or []
        if codes:
            continue
        without_barcode.append(_sku_payload(sku, barcodes=[]))
        if len(without_barcode) >= 20:
            break
    low_stock = [item for item in products if 0 < int(item.get("available_qty") or item.get("qty") or 0) <= 2][:20]
    return {
        "summary": summary,
        "critical_stock": low_stock,
        "products_without_barcode": without_barcode,
        "all_products": products,
    }


def _request_title(entry) -> str:
    payload = entry.payload if isinstance(entry.payload, dict) else {}
    for key in ("title", "order_title", "name"):
        if payload.get(key):
            return _clean(payload[key])
    if entry.description:
        return _clean(entry.description)
    return f"Заявка {entry.order_id}"


def _shipping_orders_for_titles(agency, order_ids: list[str]) -> dict[str, Any]:
    if not order_ids:
        return {}
    try:
        from shipping.models import ShippingOrder
    except Exception:
        return {}
    numbers = {canonical_shipping_order_id(value) for value in order_ids if str(value or "").strip()}
    numbers = {value for value in numbers if value}
    if not numbers:
        return {}
    try:
        qs = (
            ShippingOrder.objects.filter(agency=agency, number__in=numbers)
            .select_related("marketplace")
            .prefetch_related("destinations")
            .annotate(title_units=Sum("items__qty_requested"), title_items=Count("items", distinct=True))
        )
        return {str(order.number): order for order in qs}
    except Exception:
        return {}


def _request_payloads(agency) -> list[dict[str, Any]]:
    """Build request rows with the same status labels/buckets as the SSR kanban."""
    agency_id = getattr(agency, "id", None)
    if agency_id:
        cached = _REQUEST_PAYLOADS_CACHE.get(int(agency_id))
        if cached and (time.monotonic() - cached[0]) < _REQUEST_PAYLOADS_TTL_SEC:
            return cached[1]
    rows = _build_request_payloads(agency)
    if agency_id:
        _REQUEST_PAYLOADS_CACHE[int(agency_id)] = (time.monotonic(), rows)
    return rows


def _request_kpi_from_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    active_requests = [
        row
        for row in rows
        if row.get("bucket") != "done" and row.get("status") not in {"completed", "cancelled"}
    ]
    type_breakdown = {"receiving": 0, "processing": 0, "shipping": 0, "other": 0}
    for row in active_requests:
        key = str(row.get("type") or row.get("order_type") or "other")
        if key == "packing":
            key = "processing"
        if key not in type_breakdown:
            key = "other"
        type_breakdown[key] += 1
    return {
        "active_requests": len(active_requests),
        "active_by_type": type_breakdown,
        "in_processing": len([row for row in rows if row.get("bucket") == "warehouse"]),
        "waiting": len([row for row in rows if row.get("bucket") in {"client", "manager"}]),
    }


def _dashboard_lite_payload(agency, *, client_view: bool) -> dict[str, Any]:
    """Small payload for the first client LK paint.

    Full ``api_dashboard`` remains backward-compatible. The home screen only
    needs KPI, stock summary and LIVE strip; request rows, documents,
    marketplace details and notifications have their own endpoints.
    """
    agency_id = int(getattr(agency, "id", 0) or 0)
    cache_key = _dashboard_lite_cache_key(agency_id)
    cached = _cache_get(cache_key) if agency_id else None
    if isinstance(cached, dict):
        data = dict(cached)
        data["client_view"] = client_view
        return data

    availability_rows = _stock_availability_rows(agency) if agency else []
    stock_summary = _stock_summary(agency, availability_rows=availability_rows) if agency else {
        "total_units": 0,
        "sku_count": 0,
        "available": 0,
        "reserved": 0,
        "receiving": 0,
        "no_movement_30_days": 0,
        "reserved_internal": 0,
    }
    live = _live_dashboard(agency, availability_rows=availability_rows) if agency else {}
    request_rows = _request_payloads(agency) if agency else []
    request_kpi = _request_kpi_from_rows(request_rows) if agency else {
        "active_requests": 0,
        "active_by_type": {"receiving": 0, "processing": 0, "shipping": 0, "other": 0},
        "in_processing": 0,
        "waiting": 0,
    }
    active_requests_home = [
        enrich_request_row(row)
        for row in request_rows
        if row.get("bucket") != "done" and row.get("status") not in {"completed", "cancelled"}
    ][:5]
    packaging = (live or {}).get("packaging") or {}
    client_payload = _client_payload(agency)
    data = {
        "authenticated": True,
        "client": client_payload,
        "client_view": client_view,
        "company": client_payload,
        "stock": {
            "summary": stock_summary,
            "critical_stock": [],
            "products_without_barcode": [],
            "all_products": [],
        },
        "inventory_check": {
            "discrepancy": 0,
            "status_label": "Данные подключены к боевой базе",
            "status_tone": "ok",
            "processing_in_progress_total": stock_summary.get("reserved_internal") or 0,
        },
        "live_dashboard": live,
        "active_requests_home": active_requests_home,
        "orders_panel_total": request_kpi["active_requests"],
        "kpi": {
            "stock_units": stock_summary["total_units"],
            "sku_count": stock_summary["sku_count"],
            "box_count": int(packaging.get("boxes") or 0),
            "pallet_count": int(packaging.get("pallets") or 0),
            "active_requests": request_kpi["active_requests"],
            "active_by_type": request_kpi["active_by_type"],
            "in_processing": request_kpi["in_processing"],
            "waiting": request_kpi["waiting"],
            "completed_this_month": 0,
            "completed_delta": 0,
            "low_stock_count": 0,
            # Финансы грузятся своим разделом; первый экран не блокируем тяжёлым расчётом.
            "amount_due": None,
            "unpaid_invoices_count": None,
            "marking_free": 0,
            "marking_printed": 0,
            "marking_used": 0,
        },
        "action_urls": _action_urls(agency),
    }
    if agency_id:
        _cache_set(_dashboard_live_cache_key(agency_id), live, _DASHBOARD_LIVE_CACHE_TTL_SEC)
        cached_data = dict(data)
        cached_data.pop("client_view", None)
        _cache_set(cache_key, cached_data, _DASHBOARD_LITE_CACHE_TTL_SEC)
    return data


def _latest_client_request_entries(agency) -> list[OrderAuditEntry]:
    """Return one newest audit row per client request without a raw-event cap."""
    base = OrderAuditEntry.objects.filter(
        agency=agency,
        order_type__in=CLIENT_REQUEST_ORDER_TYPES,
    )
    latest_id = (
        base.filter(
            order_type=OuterRef("order_type"),
            order_id=OuterRef("order_id"),
        )
        .order_by("-created_at", "-id")
        .values("id")[:1]
    )
    candidates = (
        base.filter(pk=Subquery(latest_id))
        .only(
            "id",
            "order_id",
            "order_type",
            "action",
            "agency_id",
            "payload",
            "description",
            "created_at",
        )
        .order_by("-created_at", "-id")
    )
    candidate_entries = list(candidates)
    processing_status_fallbacks: dict[tuple[str, str], OrderAuditEntry] = {}
    sparse_processing_keys = {
        (str(entry.order_type or ""), str(entry.order_id or "").strip())
        for entry in candidate_entries
        if entry.order_type in {"processing", "packing"}
        and not _entry_has_request_status(entry)
        and str(entry.order_id or "").strip()
    }
    if sparse_processing_keys:
        processing_order_ids = {order_id for _order_type, order_id in sparse_processing_keys}
        latest_status_id = (
            base.filter(
                order_type=OuterRef("order_type"),
                order_id=OuterRef("order_id"),
                action="status",
            )
            .order_by("-created_at", "-id")
            .values("id")[:1]
        )
        status_entries = (
            base.filter(
                pk=Subquery(latest_status_id),
                order_type__in=("processing", "packing"),
                order_id__in=processing_order_ids,
            )
            .only(
                "id",
                "order_id",
                "order_type",
                "action",
                "agency_id",
                "payload",
                "description",
                "created_at",
            )
        )
        processing_status_fallbacks = {
            (str(entry.order_type or ""), str(entry.order_id or "").strip()): entry
            for entry in status_entries
        }

    latest_entries: list[OrderAuditEntry] = []
    seen: set[tuple[str, str]] = set()
    for entry in candidate_entries:
        if entry.order_type in {"processing", "packing"} and not _entry_has_request_status(entry):
            entry = processing_status_fallbacks.get(
                (str(entry.order_type or ""), str(entry.order_id or "").strip()),
                entry,
            )
        normalized_type = "processing" if entry.order_type == "packing" else str(entry.order_type or "")
        key = (normalized_type, str(entry.order_id or "").strip())
        if not key[1] or key in seen:
            continue
        seen.add(key)
        latest_entries.append(entry)
    return latest_entries


def _entry_has_request_status(entry: OrderAuditEntry) -> bool:
    payload = entry.payload if isinstance(entry.payload, dict) else {}
    return bool(
        entry.action == "status"
        or payload.get("status")
        or payload.get("status_label")
        or payload.get("submit_action")
        or payload.get("processing_stage")
    )


def _processing_client_list_payload(payload: dict[str, Any]) -> tuple[dict[str, Any], bool]:
    normalized = dict(payload or {})
    status_value = str(normalized.get("status") or normalized.get("submit_action") or "").strip().lower()
    if status_value in {"done", "completed", "closed", "finished"}:
        return normalized, False
    processing_act = normalized.get("processing_act")
    processing_act = dict(processing_act) if isinstance(processing_act, dict) else {}
    if int(processing_act.get("version") or 0) < 1:
        return normalized, False
    act_status = str(processing_act.get("status") or "").strip().lower()
    if act_status not in {"awaiting_confirmations", "confirmed", "dispute"}:
        return normalized, False
    manager_response = str(processing_act.get("manager_response") or "pending").strip().lower()
    if manager_response != "confirmed":
        normalized["status_label"] = "Обработка подтверждена, ожидает проверки менеджера"
        return normalized, True
    normalized["status_label"] = "Результат обработки подтвержден менеджером"
    return normalized, False


def _build_request_payloads(agency) -> list[dict[str, Any]]:
    from types import SimpleNamespace

    from . import web_ui as client_web_ui
    from .client_visibility import client_facing_status_label

    client_web_ui.clear_lk_shipping_status_caches()
    rows = []
    latest_entries = _latest_client_request_entries(agency)

    receiving_ids = [str(e.order_id) for e in latest_entries if e.order_type == "receiving"]
    receiving_by_id: dict[str, list] = {oid: [] for oid in receiving_ids}
    if receiving_ids:
        # Only act-related rows — not the full receiving audit trail (was the main lag).
        for entry in (
            OrderAuditEntry.objects.filter(
                agency=agency,
                order_type="receiving",
                order_id__in=receiving_ids,
            )
            .filter(_RECEIVING_ACT_PAYLOAD_Q)
            .only("id", "order_id", "payload", "created_at")
            .order_by("created_at", "id")
            .iterator(chunk_size=200)
        ):
            receiving_by_id.setdefault(str(entry.order_id), []).append(entry)

    shipping_numbers = [str(e.order_id) for e in latest_entries if e.order_type == "shipping"]
    client_web_ui.preload_lk_shipping_status_caches(getattr(agency, "id", None), shipping_numbers)
    shipping_orders_by_number = _shipping_orders_for_titles(agency, shipping_numbers)

    for entry in latest_entries:
        payload = entry.payload if isinstance(entry.payload, dict) else {}
        processing_waits_manager = False
        if entry.order_type == "receiving":
            payload = _merge_receiving_act_fields(receiving_by_id.get(str(entry.order_id)) or [], payload)
        elif entry.order_type in {"processing", "packing"}:
            payload, processing_waits_manager = _processing_client_list_payload(payload)

        synthetic = SimpleNamespace(
            order_type=entry.order_type,
            order_id=str(entry.order_id),
            agency=agency,
            agency_id=getattr(agency, "id", None),
            payload=payload,
            description=getattr(entry, "description", "") or "",
        )
        # List path: no live warehouse resolver (avoids N+1 stock snapshots).
        if entry.order_type == "shipping":
            status_label, bucket = client_web_ui._shipping_kanban_status(entry)
            status_label = client_web_ui._repair_mojibake_text(status_label)
        elif entry.order_type in {"processing", "packing"}:
            status_label = client_web_ui._repair_mojibake_text(client_web_ui._order_status_label(synthetic))
            bucket = "manager" if processing_waits_manager else client_web_ui._order_bucket(synthetic)
        else:
            mapped = resolve_lk_entry_status(
                synthetic,
                audience="client",
                use_live_warehouse=False,
            )
            status_label = client_web_ui._repair_mojibake_text(mapped.status_label)
            bucket = mapped.bucket

        if status_label == "Отменена":
            status_label = "ОТМЕНЕНА"
        awaiting_act = (
            entry.order_type == "receiving"
            and payload.get("act_sent")
            and str(payload.get("act_client_response") or "").lower() not in {"confirmed", "dispute"}
        )
        if awaiting_act:
            status_label = "Акт отправлен клиенту"
            bucket = "client"
        elif bucket == "done" and entry.order_type != "shipping" and status_label != "ОТМЕНЕНА":
            status_label = "Выполнена"

        if not processing_waits_manager:
            status_label = client_facing_status_label(status_label, bucket=bucket)
        if entry.order_type == "receiving":
            attention = receiving_attention_from_payload(payload)
        else:
            attention = request_attention(
                agency=agency,
                order_type=entry.order_type,
                order_id=str(entry.order_id),
                payload=payload,
            )
        display_status = resolve_lk_request_status(
            bucket=bucket,
            status_label=status_label,
            attention=attention,
            payload=payload,
        )
        order_type = _clean(entry.order_type)
        type_key = "processing" if order_type == "packing" else client_web_ui._lk_request_type_key(order_type)
        # Карточка всегда в ЛК (#/request/…); форма продолжения черновика — через wms_url / continue.
        is_draft = client_web_ui._is_draft_entry(entry)
        detail_url = lk_request_hash(type_key, entry.order_id)
        wms_url = ""
        if is_draft:
            from .client_drafts import continue_url_for_entry

            wms_url = continue_url_for_entry(
                agency_id=agency.id,
                order_type=entry.order_type,
                order_id=str(entry.order_id),
            ) or ""
        title = _request_title(entry)
        shipping_order = None
        if entry.order_type == "shipping":
            shipping_order = shipping_orders_by_number.get(canonical_shipping_order_id(str(entry.order_id)))
        display_title = build_request_display_title(
            order_type=entry.order_type,
            order_id=entry.order_id,
            payload=payload,
            base_title=title,
            shipping_order=shipping_order,
        )
        rows.append({
            "id": f"{entry.order_type}-{entry.order_id}",
            "order_id": _clean(entry.order_id),
            "number": _clean(entry.order_id),
            "order_type": order_type,
            "type": type_key,
            "type_label": client_web_ui._lk_request_type_label(order_type),
            "type_badge": client_web_ui._lk_request_type_badge(order_type),
            "title": title,
            "display_title": display_title,
            "subtitle": status_label,
            "status": display_status.filter_status,
            "status_label": status_label,
            "status_pill": display_status.status_pill,
            "status_tone": "danger" if status_label == "ОТМЕНЕНА" else "",
            "stage": display_status.bucket,
            "bucket": display_status.bucket,
            "cancel_policy": display_status.cancel_policy,
            "manager_status_label": display_status.manager_label,
            "attention": attention,
            "date": _local_dt(entry.created_at),
            "created_at": _local_dt(entry.created_at),
            "updated_at": _local_dt(entry.created_at),
            "detail_url": detail_url,
            "wms_url": wms_url,
            "is_draft": is_draft,
        })
        rows[-1] = enrich_request_row(rows[-1])
    return rows


@login_required
@require_GET
def api_requests(request):
    """Lightweight requests list for «Мои заявки» (no stock/finance/marking)."""
    agency, client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    rows = _request_payloads(agency)
    raw_page = request.GET.get("page")
    raw_page_size = request.GET.get("page_size")
    if raw_page is None and raw_page_size is None:
        page_rows = rows
        pagination = {
            "page": 1,
            "page_size": len(rows),
            "total_pages": 1,
            "total_count": len(rows),
            "has_next": False,
            "has_previous": False,
        }
    else:
        try:
            page_size = min(max(int(raw_page_size or 50), 1), 200)
        except (TypeError, ValueError):
            page_size = 50
        paginator = Paginator(rows, page_size)
        page = paginator.get_page(raw_page or 1)
        page_rows = list(page.object_list)
        pagination = {
            "page": page.number,
            "page_size": page_size,
            "total_pages": paginator.num_pages,
            "total_count": paginator.count,
            "has_next": page.has_next(),
            "has_previous": page.has_previous(),
        }
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "requests": page_rows,
                "client_view": client_view,
                "pagination": pagination,
            },
        }
    )


def _messages(
    agency,
    notification_rows: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    rows = []
    source_rows = (
        list(notification_rows)
        if notification_rows is not None
        else build_client_notifications(agency, limit=30)
    )
    for item in source_rows:
        created = item.get("created_at")
        rows.append({
            "id": item.get("id"),
            "title": item.get("title") or "",
            "message": item.get("message") or item.get("text") or "",
            "type": item.get("type") or "comment",
            "priority": item.get("priority") or "normal",
            "read": bool(item.get("read")),
            "created_at": _local_dt(created) if created else "",
            "detail_url": item.get("detail_url") or "",
        })
    return rows


def _live_dashboard(
    agency,
    availability_rows: list[dict[str, Any]] | None = None,
    marketplaces_rows: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    if not agency:
        return {}
    if availability_rows is not None:
        return build_live_dashboard(
            agency,
            availability_rows=availability_rows,
            marketplaces_rows=marketplaces_rows,
        )
    agency_id = int(getattr(agency, "id", 0) or 0)
    cache_key = _dashboard_live_cache_key(agency_id)
    cached = _cache_get(cache_key) if agency_id else None
    if isinstance(cached, dict):
        return cached
    payload = build_live_dashboard(agency)
    if agency_id:
        _cache_set(cache_key, payload, _DASHBOARD_LIVE_CACHE_TTL_SEC)
    return payload


def _marketplaces_payload(agency, *, check_auth: bool = False, force_auth_refresh: bool = False) -> list[dict[str, Any]]:
    return build_marketplaces_payload(
        agency,
        check_auth=check_auth,
        force_auth_refresh=force_auth_refresh,
    )


def _supplies_payload(agency) -> dict[str, Any]:
    return build_supplies_payload(agency)


def _marking_payload(agency) -> dict[str, int]:
    qs = MarkingCode.objects.filter(agency=agency)
    total = qs.count()
    used = qs.exclude(used_at__isnull=True).count()
    printed = qs.exclude(printed_at__isnull=True).count()
    free = max(total - used, 0)
    return {"total": total, "free": free, "reserved": printed, "used": used}


def _stock_journal_rows(agency, hide_consumed=True) -> list[dict[str, Any]]:
    """Остатки ЛК клиента: только свободные к отгрузке единицы.

    Резерв и выгрузка (available_qty <= 0) клиенту не показываются.
    """
    from sklad.services.stock_availability import StockAvailabilityService
    from sklad.ui_services import (
        _processing_in_progress_groups,
        _reserve_key,
    )

    raw_rows = StockAvailabilityService.stock_rows_with_availability(agency=agency)
    _, processing_in_progress_totals = _processing_in_progress_groups(agency)
    rows = []
    for item in raw_rows:
        available_qty = int(item.get("available_qty") or 0)
        # Клиент видит только то, что можно загрузить/отгрузить.
        if not _stock_row_visible_to_client(item, hide_consumed=hide_consumed):
            continue
        state_code = str(item.get("warehouse_state_code") or "").strip()
        agency_id = int(item.get("agency_id") or getattr(agency, "id", 0) or 0)
        sku = _clean(item.get("sku") or item.get("sku_code"))
        size = _clean(item.get("size"))
        goods_type = _clean(item.get("goods_type")) or "-"
        progress_key = _reserve_key(agency_id, sku, size, goods_type)
        processing_in_progress_qty = int(processing_in_progress_totals.get(progress_key, 0) or 0)
        box_code = _clean(item.get("box_code")) or _clean(item.get("container_code")) or "-"
        when_raw = item.get("updated_at") or item.get("created_at")
        rows.append({
            "id": int(item.get("id") or 0),
            "when": _local_dt(when_raw),
            "when_display": _local_date(when_raw) if when_raw else "-",
            "box_code": box_code,
            "sku": sku or "-",
            "name": _clean(item.get("name")) or "-",
            "size": size or "-",
            "goods_type": goods_type,
            "barcode": _clean(item.get("barcode")),
            # В ЛК qty = доступно к отгрузке (не физический резерв/OTG).
            "qty": available_qty,
            "stock_main_qty": int(item.get("stock_main_qty") or 0),
            "stock_processing_qty": 0,
            "stock_otg_qty": 0,
            "available_qty": available_qty,
            "processing_reserved_qty": 0,
            "processing_in_progress_qty": processing_in_progress_qty,
            "shipping_reserved_qty": 0,
            "location": _clean(item.get("location")),
            "container_code": box_code,
            "state_label": state_code,
        })
    rows.sort(key=lambda row: row.get("when") or "", reverse=True)
    return rows


@login_required
@require_GET
def api_dashboard(request):
    agency, client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if _truthy_param(request.GET.get("lite")):
        return JsonResponse({"ok": True, "data": _dashboard_lite_payload(agency, client_view=client_view)})
    # One availability scan shared by stock KPI and LIVE strip.
    availability_rows = _stock_availability_rows(agency) if agency else []
    stock = _stock_payload(agency, availability_rows=availability_rows)
    requests = _request_payloads(agency)
    notification_rows = build_client_notifications(agency, limit=30)
    messages = _messages(agency, notification_rows=notification_rows)
    marketplaces = _marketplaces_payload(agency)
    live = _live_dashboard(
        agency,
        availability_rows=availability_rows,
        marketplaces_rows=marketplaces,
    )
    marking = _marking_payload(agency)
    from .messaging_lk import unread_notifications_count
    from .home_progress import enrich_request_row

    chat_unread_count = 0
    if agency and _chats_enabled():
        from .messaging_lk import unread_chat_count_for_client

        chat_unread_count = unread_chat_count_for_client(agency)

    finance = build_finance_summary(agency) if agency else {
        "documents": [],
        "billing": None,
        "billing_available": False,
        "months": [],
    }
    billing = finance.get("billing") or {}
    documents = finance.get("documents") or []
    unpaid_invoices = [
        doc
        for doc in documents
        if str(doc.get("doc_kind") or "") == "invoice"
        and str(doc.get("status") or "").lower() not in {"paid", "оплачен", "cancelled", "canceled", "отменён", "отменен"}
        and "оплачен" not in str(doc.get("status_label") or "").lower()
    ]
    act_attention = _act_attention_payload(agency)
    request_kpi = _request_kpi_from_rows(requests)
    active_requests = [
        row
        for row in requests
        if row.get("bucket") != "done" and row.get("status") not in {"completed", "cancelled"}
    ]
    home_active = [enrich_request_row(row) for row in active_requests[:5]]
    packaging = (live or {}).get("packaging") or {}
    completed_month = OrderAuditEntry.objects.filter(agency=agency, created_at__date__gte=timezone.localdate().replace(day=1)).filter(Q(action__icontains="done") | Q(action__icontains="completed") | Q(action__icontains="готов") | Q(action__icontains="закрыт") | Q(action__icontains="выполн")).values("order_type", "order_id").distinct().count()
    data = {
        "authenticated": True,
        "client": _client_payload(agency),
        "client_view": client_view,
        "company": _client_payload(agency),
        "warehouses": [],
        "months": finance.get("months") or [],
        "billing": billing,
        "billing_available": bool(finance.get("billing_available")),
        "stock": stock,
        "requests": requests,
        "active_requests_home": home_active,
        "marketplaces": marketplaces,
        "documents": documents,
        "notifications": messages,
        "notifications_unread": (
            unread_notifications_count(agency, sync=False)
            if agency
            else 0
        ),
        "attention_events": act_attention,
        "act_attention": {
            "count": len(act_attention),
            "items": act_attention,
        },
        "live_history": (
            build_live_history(
                agency,
                limit=5,
                notification_rows=notification_rows,
            )
            if agency
            else []
        ),
        "inventory_check": {
            "discrepancy": stock["summary"]["total_units"] - stock["summary"]["available"],
            "status_label": "Данные подключены к боевой базе",
            "status_tone": "ok",
            "processing_in_progress_total": stock["summary"]["reserved"],
        },
        "live_dashboard": live,
        "marking": marking,
        "other_request_categories": OTHER_REQUEST_CATEGORIES,
        "orders_panel_total": len(requests),
        "chat_unread_count": chat_unread_count,
        "kpi": {
            "stock_units": stock["summary"]["total_units"],
            "sku_count": stock["summary"]["sku_count"],
            "box_count": int(packaging.get("boxes") or 0),
            "pallet_count": int(packaging.get("pallets") or 0),
            "active_requests": request_kpi["active_requests"],
            "active_by_type": request_kpi["active_by_type"],
            "in_processing": request_kpi["in_processing"],
            "waiting": request_kpi["waiting"],
            "completed_this_month": completed_month,
            "completed_delta": 0,
            "low_stock_count": len(stock["critical_stock"]),
            "amount_due": billing.get("total_due") or billing.get("open_amount") or 0,
            "unpaid_invoices_count": len(unpaid_invoices),
            "marking_free": (marking or {}).get("free") or 0,
            "marking_printed": (marking or {}).get("reserved") or 0,
            "marking_used": (marking or {}).get("used") or 0,
        },
        "action_urls": _action_urls(agency),
    }
    return JsonResponse({"ok": True, "data": data})


def _act_attention_payload(agency) -> list[dict[str, Any]]:
    """Mirror SSR act attention cards for API consumers."""
    if not agency:
        return []
    from . import web_ui as client_web_ui

    cards: list[dict[str, Any]] = []
    # Reuse same scan approach as build_dashboard_context (lightweight subset).
    qs = (
        OrderAuditEntry.objects.filter(agency=agency, order_type__in=("receiving", "shipping"))
        .order_by("-created_at")[:300]
    )
    seen = set()
    act_by_order: dict[tuple[str, str], Any] = {}
    status_by_order: dict[tuple[str, str], Any] = {}
    for entry in qs:
        key = (entry.order_type, str(entry.order_id))
        payload = entry.payload if isinstance(entry.payload, dict) else {}
        if entry.order_type == "receiving" and payload.get("act_sent") and key not in act_by_order:
            act_by_order[key] = entry
        if entry.order_type == "shipping" and (payload.get("act_sent") or payload.get("act_manager_signed")) and key not in act_by_order:
            act_by_order[key] = entry
        if client_web_ui._is_status_entry(entry) and key not in status_by_order:
            status_by_order[key] = entry
    for key, act_entry in act_by_order.items():
        status_entry = status_by_order.get(key)
        if key[0] == "receiving":
            if not client_web_ui._receiving_act_needs_client_attention(act_entry, status_entry):
                continue
        else:
            if not client_web_ui._shipping_act_needs_client_attention(act_entry, status_entry):
                continue
        if key in seen:
            continue
        seen.add(key)
        type_key = "receiving" if key[0] == "receiving" else "shipping"
        payload = act_entry.payload if isinstance(act_entry.payload, dict) else {}
        act_title = _clean(payload.get("act_sent") or "")
        if not act_title or act_title.lower() in {"true", "1", "yes"}:
            act_title = _request_title(act_entry) or f"Акт по заявке №{key[1]}"
        cards.append(
            {
                "order_id": key[1],
                "order_type": key[0],
                "type": type_key,
                "type_label": client_web_ui._lk_request_type_label(key[0]),
                "title": act_title,
                "status_label": "Акт ожидает действия",
                "detail_url": lk_request_hash(type_key, key[1]),
                "created_at": _local_dt(act_entry.created_at),
            }
        )
        if len(cards) >= 20:
            break
    return cards


@login_required
@require_GET
def api_dashboard_live(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    try:
        live = _live_dashboard(agency)
        return JsonResponse({"ok": True, "data": live})
    except Exception:
        return JsonResponse(
            {
                "ok": False,
                "error": "Сервер временно недоступен",
                "data": {"status": "offline", "updated_at": timezone.localtime().isoformat()},
            },
            status=503,
        )


@login_required
@require_GET
def api_dashboard_live_history(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    try:
        limit = int(request.GET.get("limit") or 5)
    except Exception:
        limit = 5
    return JsonResponse({"ok": True, "data": {"results": build_live_history(agency, limit=limit)}})


@login_required
@require_GET
def api_me(request):
    agency, client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    return JsonResponse({"ok": True, "data": {"user": request.user.username, "client": _client_payload(agency), "client_view": client_view}})


@login_required
@require_GET
def api_nomenclature(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    qs = (
        SKU.objects.filter(agency=agency, deleted=False)
        .select_related("market")
        .order_by("name", "id")
    )
    search = _clean(request.GET.get("search"))
    if search:
        qs = qs.filter(Q(name__icontains=search) | Q(sku_code__icontains=search) | Q(code__icontains=search) | Q(brand__icontains=search))
    brand = _clean(request.GET.get("filter_brand"))
    if brand:
        qs = qs.filter(brand=brand)
    market = _clean(request.GET.get("filter_market"))
    if market:
        if market.isdigit():
            qs = qs.filter(market_id=int(market))
        else:
            qs = qs.filter(market__name__iexact=market)
    size = _clean(request.GET.get("filter_size"))
    if size:
        qs = qs.filter(size=size)
    try:
        page_size = min(max(int(request.GET.get("page_size") or 50), 1), 200)
    except Exception:
        page_size = 50
    try:
        page_number = max(int(request.GET.get("page") or 1), 1)
    except Exception:
        page_number = 1
    paginator = Paginator(qs, page_size)
    page = paginator.get_page(page_number)
    page_items = list(page.object_list)
    barcode_map = _barcode_map_for_skus([item.id for item in page_items])
    return JsonResponse({"ok": True, "data": {
        "rows": [
            _sku_payload(item, barcodes=barcode_map.get(item.id, []))
            for item in page_items
        ],
        "pagination": {"page": page.number, "page_size": page_size, "total_pages": paginator.num_pages, "total_count": paginator.count, "has_next": page.has_next(), "has_previous": page.has_previous()},
        "counts": {"total": paginator.count, "with_barcode": SKUBarcode.objects.filter(sku__agency=agency).values("sku_id").distinct().count()},
        "filters": {
            "search": search,
            "values": {"filter_brand": brand, "filter_market": market, "filter_size": size},
            "brand_options": list(SKU.objects.filter(agency=agency, deleted=False).exclude(brand="").values_list("brand", flat=True).distinct().order_by("brand")[:100]),
            "market_options": list(
                SKU.objects.filter(agency=agency, deleted=False, market__isnull=False)
                .exclude(market__name="")
                .values_list("market__name", flat=True)
                .distinct()
                .order_by("market__name")[:100]
            ),
            "size_options": list(SKU.objects.filter(agency=agency, deleted=False).exclude(size="").values_list("size", flat=True).distinct().order_by("size")[:100]),
        },
        "action_urls": {
            "create": f"/client/{agency.id}/sku/new/",
            "legacy_list": f"/client/{agency.id}/sku/?view=table",
            "export": f"/client/api/v1/nomenclature/export/?client={agency.id}",
            "template": "/orders/templates/sku-upload/",
            "upload": "/client/api/v1/nomenclature/upload/",
        },
    }})


@login_required
@require_GET
def api_nomenclature_export(request):
    """Download the selected client's current nomenclature as an Excel file."""
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)

    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Номенклатура"
    headers = [
        "Артикул",
        "Наименование",
        "Бренд",
        "Маркетплейс",
        "Размер",
        "Основной штрих-код",
        "Все штрих-коды",
        "Цвет",
        "Категория",
        "Состав",
        "Пол",
        "Сезон",
        "Страна производства",
        "Вес нетто, кг",
        "Вес брутто, кг",
        "Ссылка на фото",
        "Обновлено",
    ]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="F89000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    queryset = (
        SKU.objects.filter(agency=agency, deleted=False)
        .select_related("market")
        .prefetch_related("barcodes")
        .order_by("sku_code", "id")
    )
    text_columns = {1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 16, 17}
    for sku in queryset:
        barcode_rows = list(sku.barcodes.all())
        barcode_values = [_clean(item.value) for item in barcode_rows if _clean(item.value)]
        primary_barcode = next(
            (_clean(item.value) for item in barcode_rows if item.is_primary and _clean(item.value)),
            barcode_values[0] if barcode_values else "",
        )
        category = (
            _clean(sku.tovar_category)
            or _clean(sku.vid_tovar)
            or _clean(sku.type_tovar)
        )
        updated_at = getattr(sku, "updated_at", None)
        updated_label = (
            timezone.localtime(updated_at).strftime("%d.%m.%Y %H:%M")
            if updated_at
            else ""
        )
        worksheet.append([
            _clean(sku.sku_code) or _clean(sku.code),
            _clean(sku.name),
            _clean(sku.brand),
            _clean(getattr(sku.market, "name", "")),
            _clean(sku.size),
            primary_barcode,
            "\n".join(barcode_values),
            _clean(sku.color),
            category,
            _clean(sku.composition),
            _clean(sku.gender),
            _clean(sku.season),
            _clean(sku.made_in),
            sku.weight_net_kg,
            sku.weight_gross_kg,
            _clean(sku.img),
            updated_label,
        ])
        row_number = worksheet.max_row
        for column_number in text_columns:
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.data_type = "s"
        worksheet.cell(row=row_number, column=6).number_format = "@"
        worksheet.cell(row=row_number, column=7).number_format = "@"
        worksheet.cell(row=row_number, column=7).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    column_widths = {
        "A": 22,
        "B": 42,
        "C": 22,
        "D": 18,
        "E": 14,
        "F": 24,
        "G": 30,
        "H": 18,
        "I": 22,
        "J": 28,
        "K": 14,
        "L": 16,
        "M": 22,
        "N": 16,
        "O": 16,
        "P": 42,
        "Q": 20,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="nomenclature-client-{agency.id}-'
        f'{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


def _nomenclature_delete_block_reason(agency, sku) -> str:
    sku_code = _clean(getattr(sku, "sku_code", ""))
    stock_filter = Q(sku_ref=sku)
    if sku_code:
        stock_filter |= Q(sku_code__iexact=sku_code)
    stock_row = (
        WarehouseStockSnapshot.objects.filter(agency=agency)
        .filter(stock_filter)
        .order_by("-updated_at", "-id")
        .first()
    )
    if not stock_row:
        return ""

    details = []
    source_type = _clean(getattr(stock_row, "source_context_type", ""))
    source_id = _clean(getattr(stock_row, "source_context_id", ""))
    if source_id:
        prefix = "по заявке"
        if source_type == "receiving":
            prefix = "по приёмке"
        elif source_type == "shipping":
            prefix = "по отгрузке"
        elif source_type == "processing":
            prefix = "по обработке"
        details.append(f"{prefix} {source_id}")
    container_code = _clean(getattr(stock_row, "container_code", ""))
    if container_code:
        details.append(f"короб/палета {container_code}")
    detail_text = f" ({'; '.join(details)})" if details else ""
    return (
        f"Нельзя удалить артикул {sku_code or sku.id}: по нему уже был приход товара "
        f"и есть складская история{detail_text}. Чтобы не нарушить остатки и документы, "
        "оставьте SKU в номенклатуре или обратитесь к менеджеру FullBox."
    )


@login_required
@require_http_methods(["POST", "DELETE"])
def api_nomenclature_delete(request, sku_id: int):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)

    sku = SKU.objects.filter(pk=sku_id, agency=agency, deleted=False).first()
    if not sku:
        return _json_error("SKU не найден в номенклатуре клиента или уже удалён.", 404)

    reason = _nomenclature_delete_block_reason(agency, sku)
    if reason:
        return _json_error(reason, 409)

    with transaction.atomic():
        sku.barcodes.all().delete()
        sku.deleted = True
        sku.save(update_fields=["deleted", "updated_at"])
        log_sku_change(
            "delete",
            sku,
            user=request.user if getattr(request.user, "is_authenticated", False) else None,
            description="Удаление SKU клиентом из ЛК клиента (складской истории нет).",
        )
    invalidate_dashboard_lite_cache(agency.id)
    return JsonResponse({
        "ok": True,
        "data": {
            "deleted": True,
            "sku_id": sku.id,
            "sku_code": _clean(sku.sku_code),
            "message": f"Артикул {_clean(sku.sku_code) or sku.id} удалён из номенклатуры.",
        },
    })


@login_required
@require_POST
def api_nomenclature_upload(request):
    """Парсит Excel и возвращает предпросмотр — в БД ничего не пишет."""
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    uploaded = (
        request.FILES.get("file")
        or request.FILES.get("template")
        or request.FILES.get("sku_template")
    )
    if not uploaded:
        return _json_error("Выберите Excel-файл шаблона номенклатуры")
    name = str(getattr(uploaded, "name", "") or "")
    ext = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in {".xlsx", ".xls"}:
        return _json_error("Нужен файл Excel (.xlsx или .xls)")
    from .nomenclature_upload import parse_sku_template_preview

    try:
        preview = parse_sku_template_preview(uploaded, agency)
    except ValueError as exc:
        return _json_error(str(exc) or "Ошибка чтения шаблона")
    except Exception:
        return _json_error("Не удалось прочитать Excel-файл")
    return JsonResponse({"ok": True, "data": preview})


@login_required
@require_POST
def api_nomenclature_commit(request):
    """Сохраняет строки из модалки предпросмотра в каталог клиента."""
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except Exception:
        return _json_error("Некорректный JSON")
    rows = payload.get("rows") if isinstance(payload, dict) else None
    if not isinstance(rows, list):
        return _json_error("Передайте массив rows для сохранения")
    from .nomenclature_upload import commit_sku_template_rows

    try:
        result = commit_sku_template_rows(agency, rows)
    except ValueError as exc:
        return _json_error(str(exc) or "Ошибка сохранения")
    except Exception:
        return _json_error("Не удалось сохранить номенклатуру")
    return JsonResponse({"ok": True, "data": result})


@login_required
@require_GET
def api_stock_journal(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    hide = _clean(request.GET.get("hide_consumed") or "1").lower() not in {"0", "false", "no", "off"}
    rows = _stock_journal_rows(agency, hide_consumed=hide)
    source_row_count = len(rows)
    grouped = request.GET.get("group") == "product"
    if grouped:
        rows = _group_stock_journal_rows(rows)
    return JsonResponse({"ok": True, "data": {"rows": rows, "summary": {"row_count": len(rows), "source_row_count": source_row_count, "total_qty": sum(row["qty"] for row in rows), "available_qty": sum(row["available_qty"] for row in rows)}, "hide_consumed": hide, "grouped": grouped}})


def _group_stock_journal_rows(rows):
    """Transport-only equivalent of LK groupRowsByBarcode; no stock writes.

    Preserve SKU/size/goods-type identity, input order, quantities, latest date
    and the legacy ungrouped API/export. All amounts come from warehouse truth.
    """
    from django.utils.dateparse import parse_datetime

    grouped = {}
    for row in rows:
        barcode = str(row.get("barcode") or "").strip()
        fields = (barcode, row.get("sku"), row.get("size"), row.get("goods_type")) if barcode else (
            row.get("sku"), row.get("name"), row.get("size"), row.get("goods_type")
        )
        key = (bool(barcode), *(str(value if value is not None else "").strip().lower() for value in fields))
        if key not in grouped:
            grouped[key] = {**row, "barcode": barcode, "available_qty": 0, "qty": 0, "box_count": 0}
        target = grouped[key]
        target["available_qty"] += int(row.get("available_qty") or 0)
        target["qty"] += int(row.get("qty") or 0)
        target["box_count"] += 1
        current = parse_datetime(str(row.get("when") or ""))
        previous = parse_datetime(str(target.get("when") or ""))
        if current is not None and (previous is None or current > previous):
            target["when"] = row.get("when")
            target["when_display"] = row.get("when_display")
    return list(grouped.values())


@login_required
@require_POST
def api_stock_journal_export(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    try:
        body = json.loads(request.body.decode("utf-8") or "{}")
    except Exception:
        body = {}
    rows = body.get("rows") if isinstance(body, dict) else None
    if not isinstance(rows, list):
        rows = _stock_journal_rows(agency)
    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.append([
        "Когда", "Короб", "SKU", "Штрих-код товара",
        "Наименование", "Размер", "Тип товара", "Доступно",
    ])
    for row in rows:
        if isinstance(row, dict):
            ws.append([
                row.get("when_display") or row.get("when", ""),
                row.get("box_code", ""),
                row.get("sku", ""),
                str(row.get("barcode") or ""),
                row.get("name", ""),
                row.get("size", ""),
                row.get("goods_type", ""),
                row.get("available_qty", 0),
            ])
    for cell in ws["D"][1:]:
        cell.number_format = "@"
    output = BytesIO()
    wb.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="stock-journal.xlsx"'
    return response


@login_required
@require_GET
def api_marketplaces(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    refresh = str(request.GET.get("refresh") or "").lower() in {"1", "true", "yes"}
    check_auth_param = request.GET.get("check_auth")
    check_auth = (
        refresh
        if check_auth_param is None
        else str(check_auth_param or "").lower() not in {"0", "false", "no"}
    )
    mp_stocks = (
        build_mp_stocks_payload(agency, force_refresh=True)
        if refresh
        else build_mp_stocks_cached_payload(agency)
    )
    return JsonResponse({
        "ok": True,
        "data": {
            "marketplaces": _marketplaces_payload(
                agency,
                check_auth=check_auth,
                force_auth_refresh=refresh,
            ),
            "live_dashboard": _live_dashboard(agency),
            "mp_stocks": mp_stocks,
            "supplies": _supplies_payload(agency),
        },
    })


@login_required
@require_POST
def api_marketplace_sync(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = {}
    marketplace = _clean((payload or {}).get("marketplace") or request.POST.get("marketplace") or "")
    return sync_marketplace(agency=agency, marketplace=marketplace)


@login_required
@require_GET
def api_marking(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    return JsonResponse({"ok": True, "data": _marking_payload(agency)})


@login_required
@require_http_methods(["GET", "POST"])
def api_request_detail(request, order_type: str, order_id: str):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    if request.method == "POST":
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
        except json.JSONDecodeError:
            payload = {}
        payload = payload if isinstance(payload, dict) else {}
        action = _clean(payload.get("action") or request.POST.get("action") or "")
        text = _clean(payload.get("text") or payload.get("comment") or request.POST.get("text") or "")
        ok, message = apply_client_request_action(
            agency=agency,
            order_type=order_type,
            order_id=order_id,
            action=action,
            user=request.user,
            text=text,
        )
        if not ok:
            return _json_error(message, 400)
        invalidate_request_payloads_cache(getattr(agency, "id", None))
        detail = build_request_detail(agency=agency, order_type=order_type, order_id=order_id)
        return JsonResponse({"ok": True, "message": message, "data": detail})

    detail = build_request_detail(agency=agency, order_type=order_type, order_id=order_id)
    if not detail:
        return _json_error("Заявка не найдена", 404)
    return JsonResponse({"ok": True, "data": detail})


@login_required
@require_GET
def api_request_detail_export(request, order_type: str, order_id: str):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    raw = _clean(order_type).lower()
    if raw == "packing":
        raw = "processing"
    if raw not in {"receiving", "shipping", "processing"}:
        return _json_error("Excel-выгрузка доступна для приёмки, отгрузки и обработки", 400)
    if raw == "shipping":
        from shipping.models import ShippingOrder
        from shipping.web_ui import _shipping_box_composition_safe_filename, _shipping_box_composition_workbook

        order_number = canonical_request_order_id("shipping", order_id)
        order = ShippingOrder.objects.filter(agency=agency, number=order_number).first()
        if order is None:
            return _json_error("Заявка не найдена", 404)
        workbook = _shipping_box_composition_workbook(order)
        filename = _shipping_box_composition_safe_filename(order)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
    workbook, filename = build_request_excel_workbook(
        agency=agency,
        order_type=raw,
        order_id=order_id,
    )
    if workbook is None:
        return _json_error("Заявка не найдена", 404)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
@require_GET
def api_receiving_act_excel(request, order_id: str):
    """Download the visible receiving act as Excel for the selected client only."""
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    workbook, filename = build_receiving_act_excel_workbook(
        agency=agency,
        order_id=order_id,
    )
    if workbook is None:
        return _json_error("Акт приёмки не найден или ещё не отправлен клиенту", 404)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _marking_excel_text(value: Any) -> str:
    """Make GS/control separators visible and safe for an XLSX XML cell."""
    text = _clean(value)
    safe: list[str] = []
    for char in text:
        code = ord(char)
        if code == 29:
            safe.append("<GS>")
        elif code < 32 and char not in {"\t", "\n", "\r"}:
            safe.append(f"<0x{code:02X}>")
        else:
            safe.append(char)
    return "".join(safe)


def build_shipping_marking_workbook(*, agency, order_number: str) -> Workbook:
    """Build a read-only ЧЗ register from snapshots linked to one shipping order."""
    from openpyxl.styles import Alignment, Font, PatternFill

    snapshots = (
        WarehouseStockSnapshot.objects.filter(
            agency=agency,
            last_event__stock_context_type="shipping",
            last_event__stock_context_id=order_number,
            marking_code__gt="",
        )
        .select_related("container", "container__parent_container", "parent_container", "last_event")
        .order_by("marking_code", "-updated_at", "-id")
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Честный знак"
    headers = [
        "№",
        "Заявка",
        "Паллета",
        "Короб",
        "Артикул",
        "Наименование",
        "Размер",
        "ШК товара",
        "Код ЧЗ (GS=<GS>)",
        "Состояние",
        "Архив",
    ]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="F89000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    seen_codes: set[str] = set()
    row_number = 0
    for snapshot in snapshots:
        raw_code = _clean(getattr(snapshot, "marking_code", ""))
        if not raw_code or raw_code in seen_codes:
            continue
        seen_codes.add(raw_code)
        row_number += 1

        container = getattr(snapshot, "container", None)
        box_code = _clean(getattr(snapshot, "container_code", "")) or _clean(
            getattr(container, "container_code", "")
        )
        pallet = getattr(snapshot, "parent_container", None) or getattr(container, "parent_container", None)
        pallet_code = _clean(getattr(pallet, "container_code", ""))
        worksheet.append(
            [
                row_number,
                order_number,
                pallet_code,
                box_code,
                _clean(getattr(snapshot, "sku_code", "")),
                _clean(getattr(snapshot, "name", "")),
                _clean(getattr(snapshot, "size", "")),
                _clean(getattr(snapshot, "barcode", "")),
                _marking_excel_text(raw_code),
                _clean(getattr(snapshot, "warehouse_state_code", "")),
                "Да" if getattr(snapshot, "is_archived", False) else "Нет",
            ]
        )

    if row_number == 0:
        worksheet.append(["", order_number, "", "", "", "Коды ЧЗ по этой отгрузке пока отсутствуют"])

    for row in worksheet.iter_rows(min_row=2, min_col=2, max_col=11):
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    widths = {
        "A": 8,
        "B": 20,
        "C": 24,
        "D": 24,
        "E": 22,
        "F": 40,
        "G": 16,
        "H": 24,
        "I": 58,
        "J": 22,
        "K": 12,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    return workbook


@login_required
@require_GET
def api_shipping_marking_export(request, order_id: str):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)

    from shipping.models import ShippingOrder

    order_number = canonical_request_order_id("shipping", order_id)
    order = ShippingOrder.objects.filter(agency=agency, number=order_number).only("number").first()
    if order is None:
        return _json_error("Заявка не найдена", 404)

    workbook = build_shipping_marking_workbook(agency=agency, order_number=order.number)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="chestny-znak-{order.number}.xlsx"'
    return response


@login_required
@require_GET
def api_shipping_return_act_doc(request, order_id: str):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)

    from shipping.models import ShippingOrder
    from shipping.return_act import render_return_act_doc, return_act_doc_filename

    order_number = canonical_request_order_id("shipping", order_id)
    order = (
        ShippingOrder.objects.filter(agency=agency, number=order_number)
        .select_related("agency", "created_by", "marketplace")
        .prefetch_related("items")
        .first()
    )
    if order is None:
        return _json_error("Заявка не найдена", 404)
    if order.status not in {
        ShippingOrder.STATUS_PACKED,
        ShippingOrder.STATUS_SHIPPED,
        ShippingOrder.STATUS_PARTIAL,
    }:
        return _json_error("Акт МХ-3 будет доступен после подготовки заявки складом", 400)

    doc_bytes = render_return_act_doc(order)
    response = HttpResponse(
        doc_bytes,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="{return_act_doc_filename(order)}"'
    return response


@login_required
@require_POST
def api_other_request_create(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    payload: dict[str, Any] = {}
    content_type = (request.content_type or "").lower()
    if "application/json" in content_type:
        try:
            parsed = json.loads(request.body.decode("utf-8") or "{}")
        except json.JSONDecodeError:
            parsed = {}
        payload = parsed if isinstance(parsed, dict) else {}
    category = _clean(payload.get("category") or request.POST.get("category") or "")
    description = _clean(payload.get("description") or request.POST.get("description") or "")
    source = _clean(payload.get("source") or request.POST.get("source") or "other_form") or "other_form"
    order_id = _clean(payload.get("order_id") or request.POST.get("order_id") or "")
    save_as_draft = bool(
        payload.get("save_as_draft")
        if "save_as_draft" in payload
        else request.POST.get("save_as_draft") in {"1", "true", "True", "on", "yes"}
    )
    if not category and not description and not request.FILES:
        return _json_error("Выберите тип задачи или опишите заявку")
    if not category:
        category = "custom"
    result = create_other_request(
        agency=agency,
        user=request.user,
        category=category,
        description=description,
        save_as_draft=save_as_draft,
        source=source,
        order_id=order_id or None,
    )
    invalidate_request_payloads_cache(getattr(agency, "id", None))
    files = list(request.FILES.getlist("attachments")) or list(request.FILES.getlist("files"))
    if not files and request.FILES.get("attachment"):
        files = [request.FILES["attachment"]]
    attachments = []
    if files:
        attachments = save_other_attachments(
            order_id=result["order_id"],
            agency=agency,
            user=request.user,
            files=files,
        )
    result = dict(result)
    result["attachments"] = attachments
    result["wms_url"] = result.get("detail_url") or ""
    result["detail_url"] = lk_request_hash("other", result["order_id"])
    return JsonResponse({"ok": True, "data": result, "order_id": result["order_id"]})


@login_required
@require_GET
def api_other_attachment(request, attachment_id: int):
    from django.http import FileResponse

    from .models import OtherRequestAttachment

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    attachment = OtherRequestAttachment.objects.filter(pk=attachment_id).select_related("agency").first()
    if not attachment:
        return _json_error("Файл не найден", 404)
    if agency and attachment.agency_id != agency.id:
        return _json_error("Access denied", 403)
    if not agency and not (request.user.is_staff or request.user.is_superuser):
        return _json_error("Access denied", 403)
    try:
        return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=attachment.filename)
    except Exception:
        return _json_error("Файл недоступен", 404)


def _resolve_client_chat_thread(request, agency, *, payload: dict | None = None):
    from .chats import ensure_client_general_thread, get_thread_for_client

    raw = ""
    if payload:
        raw = str(payload.get("thread") or payload.get("thread_id") or "").strip()
    if not raw:
        raw = str(request.GET.get("thread") or request.GET.get("thread_id") or "").strip()
    if not raw and request.method == "POST" and "application/json" not in (request.content_type or "").lower():
        raw = str(request.POST.get("thread") or request.POST.get("thread_id") or "").strip()
    if raw.isdigit():
        thread = get_thread_for_client(agency=agency, thread_id=int(raw))
        if not thread:
            return None, "Чат не найден или недоступен"
        return thread, None
    return ensure_client_general_thread(agency, user=request.user), None


@login_required
@require_GET
def api_chat_threads(request):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .chats import serialize_thread_card, threads_for_client
    from .messaging_lk import unread_chat_count_for_client

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    cards = [serialize_thread_card(t) for t in threads_for_client(agency)]
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "threads": cards,
                "unread_count": unread_chat_count_for_client(agency),
            },
        }
    )


@login_required
@require_http_methods(["GET", "POST"])
def api_chat_messages(request):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .messaging_lk import (
        list_chat_messages,
        mark_chat_read_for_client,
        post_chat_message,
        resolve_chat_author_role,
        serialize_chat_message,
        unread_chat_count_for_client,
    )

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    if request.method == "POST":
        content_type = (request.content_type or "").lower()
        text = ""
        payload = {}
        if "application/json" in content_type:
            try:
                payload = json.loads(request.body.decode("utf-8") or "{}")
            except json.JSONDecodeError:
                payload = {}
            payload = payload if isinstance(payload, dict) else {}
            text = _clean(payload.get("text") or payload.get("message") or "")
        else:
            text = _clean(request.POST.get("text") or request.POST.get("message") or "")
            payload = request.POST.dict()
        thread, err = _resolve_client_chat_thread(request, agency, payload=payload)
        if err:
            return _json_error(err, 404)
        files = list(request.FILES.getlist("attachments")) or list(request.FILES.getlist("files"))
        if request.FILES.get("attachment"):
            files.append(request.FILES["attachment"])
        idem = ""
        reply_to_id = None
        if isinstance(payload, dict):
            idem = str(payload.get("idempotency_key") or payload.get("client_key") or "").strip()
            raw_reply = payload.get("reply_to") or payload.get("reply_to_id")
            if raw_reply and str(raw_reply).isdigit():
                reply_to_id = int(raw_reply)
        if not idem:
            idem = str(request.POST.get("idempotency_key") or request.POST.get("client_key") or "").strip()
        try:
            message = post_chat_message(
                agency=agency,
                user=request.user,
                text=text,
                files=files,
                thread=thread,
                idempotency_key=idem,
                reply_to_id=reply_to_id,
            )
        except Exception as exc:
            from .chat_files import ChatFileError

            if isinstance(exc, ChatFileError):
                return _json_error(str(exc), 400)
            raise
        if not message:
            return _json_error("Введите текст или прикрепите файл")
        return JsonResponse({
            "ok": True,
            "data": {
                "message": serialize_chat_message(message, agency_id=agency.id),
                "messages": list_chat_messages(agency, thread=thread, for_client=True),
                "thread": {"id": thread.id} if thread else None,
                "unread_count": unread_chat_count_for_client(agency),
            },
        })
    thread, err = _resolve_client_chat_thread(request, agency)
    if err:
        return _json_error(err, 404)
    since_raw = str(request.GET.get("since_id") or "").strip()
    since_id = int(since_raw) if since_raw.isdigit() else None
    # Opening chat marks staff messages as read for portal client (не для poll since_id).
    if since_id is None and resolve_chat_author_role(request.user, agency) == "client":
        mark_chat_read_for_client(agency, thread=thread, user=request.user)
    messages = list_chat_messages(
        agency, thread=thread, for_client=True, since_id=since_id, limit=80
    )
    return JsonResponse({
        "ok": True,
        "data": {
            "messages": messages,
            "thread": {"id": thread.id} if thread else None,
            "unread_count": unread_chat_count_for_client(agency),
        },
    })


@login_required
@require_http_methods(["POST"])
def api_chat_reaction(request, message_id: int):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .messaging_lk import toggle_message_reaction
    from .models import ClientChatMessage

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    message = ClientChatMessage.objects.filter(pk=message_id, agency=agency).first()
    if not message or not message.client_can_see:
        return _json_error("Сообщение не найдено", 404)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = {}
    emoji = str((payload or {}).get("emoji") or request.POST.get("emoji") or "").strip()
    data = toggle_message_reaction(message=message, user=request.user, emoji=emoji)
    return JsonResponse({"ok": True, "data": {"message": data}})


@login_required
@require_http_methods(["POST"])
def api_chat_message_edit(request, message_id: int):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .chat_actions import ChatActionError, edit_message
    from .messaging_lk import serialize_chat_message
    from .models import ClientChatMessage

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    message = ClientChatMessage.objects.filter(pk=message_id, agency=agency).first()
    if not message:
        return _json_error("Сообщение не найдено", 404)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = {}
    text = _clean((payload or {}).get("text") or request.POST.get("text") or "")
    try:
        message = edit_message(message=message, user=request.user, text=text)
    except ChatActionError as exc:
        return _json_error(str(exc), 400)
    return JsonResponse({"ok": True, "data": {"message": serialize_chat_message(message, agency_id=agency.id)}})


@login_required
@require_http_methods(["POST"])
def api_chat_message_delete(request, message_id: int):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .chat_actions import ChatActionError, soft_delete_message
    from .messaging_lk import serialize_chat_message
    from .models import ClientChatMessage

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    message = ClientChatMessage.objects.filter(pk=message_id, agency=agency).first()
    if not message:
        return _json_error("Сообщение не найдено", 404)
    try:
        message = soft_delete_message(message=message, user=request.user)
    except ChatActionError as exc:
        return _json_error(str(exc), 400)
    return JsonResponse({"ok": True, "data": {"message": serialize_chat_message(message, agency_id=agency.id)}})


@login_required
@require_GET
def api_chat_mention_candidates(request):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .chat_actions import mention_candidates

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    q = str(request.GET.get("q") or "").strip()
    return JsonResponse({"ok": True, "data": {"candidates": mention_candidates(q=q, agency=agency)}})


@login_required
@require_POST
def api_chat_mark_all_read(request):
    if not _chats_enabled():
        return _chat_disabled_response()
    from .chats import threads_for_client
    from .messaging_lk import mark_chat_read_for_client, unread_chat_count_for_client

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    total = 0
    for thread in threads_for_client(agency, limit=200):
        total += mark_chat_read_for_client(agency, thread=thread, user=request.user)
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "marked": total,
                "unread_count": unread_chat_count_for_client(agency),
            },
        }
    )


@login_required
@require_http_methods(["GET", "POST"])
def api_chat_prefs(request):
    if not _chats_enabled():
        return _chat_disabled_response()
    from datetime import timedelta

    from django.utils import timezone

    from .messaging_lk import get_or_create_chat_prefs

    pref = get_or_create_chat_prefs(request.user)
    if not pref:
        return _json_error("Access denied", 403)
    if request.method == "POST":
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
        except json.JSONDecodeError:
            payload = {}
        if "sound_enabled" in payload:
            pref.sound_enabled = bool(payload.get("sound_enabled"))
        mute = str(payload.get("mute") or "").strip()
        now = timezone.now()
        if mute == "1h":
            pref.sound_muted_until = now + timedelta(hours=1)
        elif mute == "day":
            pref.sound_muted_until = now.replace(hour=23, minute=59, second=59, microsecond=0)
        elif mute == "tomorrow":
            pref.sound_muted_until = now + timedelta(days=1)
        elif mute == "off":
            pref.sound_muted_until = None
            pref.sound_enabled = False
        elif mute == "on":
            pref.sound_muted_until = None
            pref.sound_enabled = True
        if "browser_notifications" in payload:
            pref.browser_notifications = bool(payload.get("browser_notifications"))
        pref.save()
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "sound_enabled": pref.sound_enabled,
                "sound_active": pref.sound_is_active(),
                "sound_muted_until": pref.sound_muted_until.isoformat() if pref.sound_muted_until else "",
                "browser_notifications": pref.browser_notifications,
            },
        }
    )


@login_required
@require_GET
def api_chat_attachment(request, attachment_id: int):
    if not _chats_enabled():
        return _chat_disabled_response()
    from django.http import FileResponse

    from employees.access import get_request_role

    from .chat_files import is_image_filename
    from .messaging_lk import get_chat_attachment

    agency, _client_view, allowed = _ctx(request)
    role = get_request_role(request) or ""
    staff_chat_roles = {
        "manager",
        "logistician",
        "head_manager",
        "director",
        "admin",
        "developer",
        "storekeeper",
        "picker",
        "processing_worker",
        "processing_head",
        "reachtruck_driver",
        "super_car",
    }
    is_staff_chat = bool(
        request.user.is_staff
        or request.user.is_superuser
        or role in staff_chat_roles
    )
    if not allowed and not is_staff_chat:
        return _json_error("Access denied", 403)
    attachment = get_chat_attachment(attachment_id=attachment_id, agency=agency if allowed else None)
    if not attachment and is_staff_chat:
        attachment = get_chat_attachment(attachment_id=attachment_id, agency=None)
    if not attachment:
        return _json_error("Файл не найден", 404)
    if agency and allowed and attachment.message.agency_id != agency.id and not is_staff_chat:
        return _json_error("Access denied", 403)
    # Клиенту не отдаём internal-вложения.
    if allowed and agency and not is_staff_chat:
        if attachment.message.visibility == "internal":
            return _json_error("Access denied", 403)
    try:
        inline = is_image_filename(attachment.filename) or request.GET.get("inline") in {
            "1",
            "true",
            "True",
        }
        return FileResponse(
            attachment.file.open("rb"),
            as_attachment=not inline,
            filename=attachment.filename,
        )
    except Exception:
        return _json_error("Файл недоступен", 404)


@login_required
@require_http_methods(["GET", "POST"])
def api_notifications(request):
    from .messaging_lk import list_notifications, mark_notifications_read, unread_notifications_count

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    if not agency:
        return _json_error("Client is required", 400)
    if request.method == "POST":
        try:
            payload = json.loads(request.body.decode("utf-8") or "{}")
        except json.JSONDecodeError:
            payload = {}
        payload = payload if isinstance(payload, dict) else {}
        action = _clean(payload.get("action") or request.POST.get("action") or "mark_read").lower()
        if action in {"mark_read", "read"}:
            ids = payload.get("ids") or payload.get("id") or []
            if isinstance(ids, (int, str)):
                ids = [ids]
            id_list = []
            for value in ids:
                try:
                    id_list.append(int(value))
                except (TypeError, ValueError):
                    continue
            mark_all = bool(payload.get("all") or payload.get("mark_all"))
            # Без явного all/mark_all и без ids ничего не помечаем — иначе «прочитать всё» случайно.
            updated = mark_notifications_read(agency, ids=id_list, all_items=mark_all)
            return JsonResponse({
                "ok": True,
                "data": {
                    "updated": updated,
                    "unread_count": unread_notifications_count(agency),
                    "notifications": list_notifications(agency),
                },
            })
        return _json_error("Неизвестное действие")
    unread_only = str(request.GET.get("unread") or "").lower() in {"1", "true", "yes"}
    rows = list_notifications(agency, unread_only=unread_only)
    serialized = []
    for item in rows:
        created = item.get("created_at")
        serialized.append({
            **item,
            "created_at": _local_dt(created) if created else "",
            "created_at_display": timezone.localtime(created).strftime("%d.%m.%Y, %H:%M") if created else "",
        })
    return JsonResponse({
        "ok": True,
        "data": {
            "notifications": serialized,
            "unread_count": unread_notifications_count(agency),
        },
    })


@login_required
@require_GET
def api_finance(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    all_docs = list_finance_documents(agency)
    kind = _clean(request.GET.get("kind") or request.GET.get("doc_kind"))
    docs = [row for row in all_docs if row.get("doc_kind") == kind] if kind in {"invoice", "upd", "act"} else all_docs
    return JsonResponse({
        "ok": True,
        "data": {
            "documents": docs,
            "counts": {
                "all": len(all_docs),
                "invoice": len([d for d in all_docs if d.get("doc_kind") == "invoice"]),
                "upd": len([d for d in all_docs if d.get("doc_kind") == "upd"]),
                "act": len([d for d in all_docs if d.get("doc_kind") == "act"]),
            },
        },
    })


@login_required
@require_GET
def api_billing(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    period = _clean(request.GET.get("period") or request.GET.get("period_id"))
    billing = build_billing_payload(agency, period_id=period or None)
    from .finance_lk import list_billing_months

    return JsonResponse({
        "ok": True,
        "data": {
            "billing_available": True,
            "billing": billing,
            "months": list_billing_months(agency),
        },
    })


@login_required
@require_GET
def api_billing_storage(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from .finance_lk import build_storage_billing_payload

    period = _clean(request.GET.get("period") or request.GET.get("period_id"))
    day = _clean(request.GET.get("day"))
    return JsonResponse(
        {
            "ok": True,
            "data": build_storage_billing_payload(agency, period_id=period or None, day=day or None),
        }
    )


@login_required
@require_GET
def api_billing_tariffs(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from billing.client_tariffs import build_client_tariffs_payload

    return JsonResponse({"ok": True, "data": build_client_tariffs_payload(agency)})


@login_required
@require_GET
def api_tariff_version_current(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from billing.serializers import client_tariff_version_dict
    from billing.tariff_services import active_tariff_version

    version = active_tariff_version(agency)
    return JsonResponse({"ok": True, "data": client_tariff_version_dict(version, include_children=True) if version else None})


@login_required
@require_GET
def api_tariff_version_history(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from billing.serializers import client_tariff_version_dict
    from billing.tariff_services import tariff_history

    return JsonResponse({"ok": True, "data": [client_tariff_version_dict(version) for version in tariff_history(agency)[:50]]})


@login_required
@require_GET
def api_tariff_version_detail(request, version_id: int):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from billing.models import ClientTariffVersion
    from billing.serializers import client_tariff_version_dict

    version = ClientTariffVersion.objects.filter(pk=version_id, client=agency).first()
    if not version:
        return _json_error("Редакция тарифов не найдена", 404)
    return JsonResponse({"ok": True, "data": client_tariff_version_dict(version, include_children=True)})


@login_required
@require_GET
def api_tariff_version_download(request, version_id: int):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    from billing.models import ClientTariffVersion
    from billing.tariff_services import render_tariff_print_response

    version = ClientTariffVersion.objects.filter(pk=version_id, client=agency).first()
    if not version:
        return _json_error("Редакция тарифов не найдена", 404)
    return render_tariff_print_response(version, as_attachment=request.GET.get("download") == "1")


@login_required
@require_GET
def api_client_requisites(request):
    """Read-only реквизиты клиента для ЛК."""
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    lifecycle = None
    serving = None
    try:
        from accountant.selectors import ensure_lifecycle
        from billing.tariff_services import active_tariff_version

        lifecycle = ensure_lifecycle(agency)
        serving = lifecycle.serving_company
        tariff = active_tariff_version(agency)
    except Exception:
        tariff = None
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "legal_name": agency.agn_name or "",
                "short_name": agency.short_name or "",
                "inn": agency.inn or "",
                "kpp": agency.kpp or "",
                "ogrn": agency.ogrn or "",
                "contract_number": agency.contract_numb or "",
                "commercial_offer_number": getattr(lifecycle, "commercial_offer_number", "") if lifecycle else "",
                "serving_company": (
                    {
                        "name": serving.short_name or serving.name,
                        "vat_label": "Без НДС" if serving.tax_mode == "no_vat" else f"НДС {serving.vat_rate}%",
                    }
                    if serving
                    else None
                ),
                "vat_label": (
                    "Без НДС"
                    if serving and serving.tax_mode == "no_vat"
                    else (f"НДС {serving.vat_rate}%" if serving else "—")
                ),
                "active_tariff": (
                    {
                        "id": tariff.id,
                        "name": tariff.name,
                        "valid_from": tariff.valid_from.isoformat() if tariff.valid_from else None,
                        "status": tariff.status,
                        "status_label": tariff.get_status_display(),
                    }
                    if tariff
                    else None
                ),
                "read_only": True,
            },
        }
    )


@login_required
@require_POST
def api_requisites_change_request(request):
    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    try:
        import json

        from accountant.services import create_requisites_change_request

        data = json.loads(request.body.decode("utf-8")) if request.body else {}
        comment = str(data.get("comment") or "").strip()
        task = create_requisites_change_request(agency=agency, comment=comment, user=request.user)
        return JsonResponse({"ok": True, "data": {"task_id": task.id}})
    except Exception as exc:
        return _json_error(exc, 400)


@login_required
@require_GET
def api_client_my_tariffs(request):
    """Алиас спеки GET /client/api/tariffs/ — активные согласованные тарифы."""
    return api_tariff_version_current(request)


@login_required
@require_GET
def api_finance_document_file(request, document_id: int):
    from django.http import FileResponse

    from .models import ClientFinanceDocument

    agency, _client_view, allowed = _ctx(request)
    if not allowed:
        return _json_error("Access denied", 403)
    doc = ClientFinanceDocument.objects.filter(pk=document_id).select_related("agency").first()
    if not doc or not doc.file:
        return _json_error("Файл не найден", 404)
    if agency and doc.agency_id != agency.id:
        return _json_error("Access denied", 403)
    if not agency and not (request.user.is_staff or request.user.is_superuser):
        return _json_error("Access denied", 403)
    try:
        filename = doc.file.name.rsplit("/", 1)[-1]
        return FileResponse(doc.file.open("rb"), as_attachment=True, filename=filename)
    except Exception:
        return _json_error("Файл недоступен", 404)
