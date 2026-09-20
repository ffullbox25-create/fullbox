from __future__ import annotations

import json
import re
import uuid
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path

from audit.models import OrderAuditEntry
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Count, Q, Sum
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from employees.access import get_request_role
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fbs.client_portal import (
    client_movement_detail,
    client_movement_source_stock_payload,
    client_movements_payload,
    client_order_detail,
    client_orders_payload,
    client_overview_payload,
    client_reports_payload,
    client_stock_detail,
    client_stock_payload,
    create_client_movement_request,
    serialize_movement_request,
    validate_client_movement_lines,
)
from fbs.integrations.warehouse_catalog import (
    client_marketplace_credentials,
    fetch_client_warehouses,
)
from fbs.models import (
    FbsBox,
    FbsClientMovementRequest,
    FbsExternalIssue,
    FbsIntegrationProfile,
    FbsMarketplaceMetadataTransfer,
    FbsOrder,
    FbsPallet,
    FbsStockBalance,
)
from fbs.services.external_issues import (
    create_issue,
    issue_display_status,
    issue_shipping_details,
    with_manager_approval_state,
)
from fbs.services.inventory import filter_unlocked_balances
from fbs.services.physical_locations import fbs_box_physical_location_label
from fbs.services.client_profiles import is_client_fbs_enabled
from fbs.services.client_movements import cancel_client_movement_request
from fbs.exceptions import FbsError, FbsIntegrationError, FbsReplenishmentError
from sku.models import Market

from .portal_access import SECTION_REPORTS, SECTION_REQUESTS, SECTION_STOCK
from .portal_members import (
    portal_sections_for_user,
    portal_user_has_section,
    portal_user_is_admin,
    resolve_portal_agency_for_user,
)
from .web_ui import _get_client_for_request, _staff_preview_agency_allowed


_HEADER_ALIASES = {
    "barcode": {
        "barcode",
        "штрихкод",
        "штрихкодтовара",
        "шк",
        "баркод",
        "баркодтовара",
    },
    "qty": {
        "qty",
        "quantity",
        "количество",
        "количествошт",
        "количествотовара",
        "штук",
    },
    "units_per_box": {
        "unitsperbox",
        "кратность",
        "кратностькороба",
        "кратностькоробашт",
        "штуквкоробе",
        "количествовкоробе",
    },
    "box_count": {
        "boxcount",
        "boxes",
        "коробов",
        "количествокоробов",
        "коробовдляпоштучного",
    },
}


def _json_error(message: str, status: int = 400, *, details=None):
    payload = {"ok": False, "error": str(message)}
    if details:
        payload["details"] = list(details)
    return JsonResponse(payload, status=status)


def _client_can_create_fbs_outbound(request, agency) -> bool:
    direct_agency = resolve_portal_agency_for_user(
        request.user,
        required_section=SECTION_REQUESTS,
    )
    if bool(
        direct_agency
        and direct_agency.pk == agency.pk
        and portal_user_has_section(request.user, agency, SECTION_REQUESTS)
    ):
        return True
    selected_agency, client_view, allowed = _get_client_for_request(request)
    return bool(
        allowed
        and not client_view
        and selected_agency
        and selected_agency.pk == agency.pk
        and get_request_role(request) in {"manager", "head_manager", "director", "admin", "developer"}
        and _staff_preview_agency_allowed(agency.pk)
    )


def _display_datetime(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")


def _public_fbs_outbound_issue(issue) -> dict:
    requested_qty = int(getattr(issue, "requested_qty_total", 0) or 0)
    picked_qty = int(getattr(issue, "picked_qty_total", 0) or 0)
    returned_qty = int(getattr(issue, "returned_qty_total", 0) or 0)
    shipped_qty = int(getattr(issue, "shipped_qty_total", 0) or 0)
    status, status_label = issue_display_status(issue)
    shipping = issue_shipping_details(issue)
    return {
        "id": issue.pk,
        "number": issue.number,
        "basis": issue.basis,
        "shipping": shipping,
        "delivery_type": shipping["delivery_type"],
        "delivery_type_label": shipping["delivery_type_label"],
        "status": status,
        "status_label": status_label,
        "requested_qty": requested_qty,
        "picked_qty": picked_qty - returned_qty - shipped_qty,
        "shipped_qty": shipped_qty,
        "created_at": _display_datetime(issue.created_at),
        "completed_at": _display_datetime(issue.completed_at),
    }


def _fbs_outbound_documents(*, agency):
    return with_manager_approval_state(
        FbsExternalIssue.objects.filter(agency=agency)
        .annotate(
            requested_qty_total=Sum("lines__requested_qty"),
            picked_qty_total=Sum("lines__picked_qty"),
            returned_qty_total=Sum("lines__returned_qty"),
            shipped_qty_total=Sum("lines__shipped_qty"),
        )
    ).order_by("-id")


def _serialize_fbs_outbound_balance(balance, *, selected_qty: int = 0) -> dict:
    return {
        "balance_id": balance.pk,
        "sku_code": str(balance.sku_code or ""),
        "name": str(balance.name or ""),
        "size": str(balance.size or ""),
        "barcode": str(balance.barcode or ""),
        "marking_code": str(balance.marking_code or ""),
        "lot_code": str(balance.lot_code or ""),
        "expiry_date": balance.expiry_date.isoformat() if balance.expiry_date else "",
        "box_code": balance.box.box_code,
        "location": fbs_box_physical_location_label(balance.box),
        "available_qty": int(balance.available_qty or 0),
        "selected_qty": int(selected_qty or 0),
    }


def _fbs_outbound_available_balances(*, agency):
    balances = filter_unlocked_balances(
        FbsStockBalance.objects.filter(
            agency=agency,
            available_qty__gt=0,
            box__status=FbsBox.STATUS_ACTIVE,
            box__pallet__status=FbsPallet.STATUS_ACTIVE,
        ).select_related(
            "box__pallet__cell__location",
            "box__source_container__current_location",
        )
    )
    return balances


def _fbs_outbound_stock(*, agency, search: str = "") -> dict:
    balances = _fbs_outbound_available_balances(agency=agency)
    needle = str(search or "").strip()
    if needle:
        balances = balances.filter(
            Q(sku_code__icontains=needle)
            | Q(name__icontains=needle)
            | Q(barcode__icontains=needle)
            | Q(marking_code__icontains=needle)
            | Q(box__box_code__icontains=needle)
        )
    total = balances.count()
    available_qty = balances.aggregate(total=Sum("available_qty"))["total"] or 0
    rows = []
    for balance in balances.order_by(
        "sku_code",
        "barcode",
        "expiry_date",
        "box__box_code",
        "id",
    )[:500]:
        rows.append(_serialize_fbs_outbound_balance(balance))
    return {
        "results": rows,
        "total": total,
        "available_qty": int(available_qty),
        "truncated": total > len(rows),
    }


def _agency_or_response(request, *, manage_fbs: bool = False):
    agency, client_view, allowed = _get_client_for_request(request)
    if not allowed:
        return None, _json_error("Доступ запрещен.", 403)
    if agency is None:
        return None, _json_error("Клиент не выбран.", 400)
    sections = portal_sections_for_user(request.user, agency)
    if sections is not None and not (
        {SECTION_REQUESTS, SECTION_STOCK, SECTION_REPORTS} & set(sections)
    ):
        return None, _json_error("Нет доступа к разделу FBS.", 403)
    if not is_client_fbs_enabled(agency):
        return None, _json_error("FBS для клиента выключен начальником склада.", 403)
    if manage_fbs and client_view and not portal_user_is_admin(request.user, agency):
        return None, _json_error("Менять склады FBS может владелец или администратор кабинета.", 403)
    return agency, None


_PUBLIC_FBS_PICKED_STATUSES = {
    FbsOrder.STATUS_PICKED,
    FbsOrder.STATUS_READY_FOR_HANDOVER,
}
_PUBLIC_FBS_SHIPPED_STATUSES = {
    FbsOrder.STATUS_HANDED_OVER,
    FbsOrder.STATUS_DELIVERED,
    FbsOrder.STATUS_RETURN_PENDING,
    FbsOrder.STATUS_RETURNED,
}


def _public_fbs_order_status(internal_status: str) -> dict:
    """Collapse warehouse FBS states into the three client milestones."""
    if internal_status in _PUBLIC_FBS_SHIPPED_STATUSES:
        return {
            "status": "shipped",
            "status_label": "Заказ отгружен",
            "status_group": "completed",
        }
    if internal_status in _PUBLIC_FBS_PICKED_STATUSES:
        return {
            "status": "picked",
            "status_label": "Товар подобран",
            "status_group": "in_work",
        }
    return {
        "status": "received",
        "status_label": "Заказ поступил",
        "status_group": "not_started",
    }


def _public_fbs_order_payload(source: dict | None) -> dict:
    """Return only client-safe order fields; warehouse problems never leave this API."""
    source = source if isinstance(source, dict) else {}
    payload = {
        key: source.get(key)
        for key in (
            "id",
            "external_order_id",
            "marketplace",
            "marketplace_label",
            "warehouse",
            "ordered_at",
            "cutoff_at",
            "imported_at",
            "updated_at",
            "line_count",
            "unit_count",
        )
        if key in source
    }
    payload.update(_public_fbs_order_status(str(source.get("status") or "")))
    if "items" in source:
        payload["items"] = []
        for raw_item in source.get("items") or []:
            if not isinstance(raw_item, dict):
                continue
            requirements = raw_item.get("requirements")
            requirements = requirements if isinstance(requirements, dict) else {}
            payload["items"].append(
                {
                    key: raw_item.get(key)
                    for key in (
                        "id",
                        "external_sku",
                        "sku_code",
                        "barcode",
                        "product_name",
                        "quantity",
                    )
                    if key in raw_item
                }
                | {
                    "requirements": {
                        key: bool(requirements.get(key))
                        for key in ("marking_required", "expiry_required")
                        if requirements.get(key)
                    }
                }
            )
    return payload


def _public_fbs_order_counts(queryset) -> dict:
    shipped_statuses = tuple(_PUBLIC_FBS_SHIPPED_STATUSES)
    picked_statuses = tuple(_PUBLIC_FBS_PICKED_STATUSES)
    public_statuses = shipped_statuses + picked_statuses
    values = queryset.aggregate(
        total=Count("id"),
        received=Count("id", filter=~Q(internal_status__in=public_statuses)),
        picked=Count("id", filter=Q(internal_status__in=picked_statuses)),
        shipped=Count("id", filter=Q(internal_status__in=shipped_statuses)),
    )
    return {key: int(values.get(key) or 0) for key in ("total", "received", "picked", "shipped")}


def _public_fbs_reports_payload(*, agency) -> dict:
    """Build client reports without warehouse problem buckets or status details."""
    source = client_reports_payload(agency=agency)
    orders = FbsOrder.objects.filter(profile__agency=agency)
    marketplace_labels = dict(FbsIntegrationProfile.MARKETPLACE_CHOICES)
    marketplace_rows = []
    for marketplace in orders.values_list("profile__marketplace", flat=True).distinct():
        counts = _public_fbs_order_counts(
            orders.filter(profile__marketplace=marketplace)
        )
        marketplace_rows.append(
            {
                "marketplace": marketplace,
                "label": marketplace_labels.get(marketplace, marketplace),
                **counts,
            }
        )
    counts = _public_fbs_order_counts(orders)
    return {
        "summary": counts,
        "marketplaces": marketplace_rows,
        "statuses": [
            {"status": "received", "label": "Заказ поступил", "count": counts["received"]},
            {"status": "picked", "label": "Товар подобран", "count": counts["picked"]},
            {"status": "shipped", "label": "Заказ отгружен", "count": counts["shipped"]},
        ],
        "trend": [
            {"date": row.get("date"), "orders": int(row.get("orders") or 0)}
            for row in source.get("trend") or []
        ],
    }


def _ozon_warehouse_settings_payload(*, agency, warehouses=None, error: str = "") -> dict:
    credentials = client_marketplace_credentials(agency)
    credential = credentials.get(FbsIntegrationProfile.MARKETPLACE_OZON)
    client_id = str(getattr(credential, "client_id", "") or "").strip()
    if warehouses is None and credential is not None and client_id and not error:
        try:
            warehouses = fetch_client_warehouses(
                agency=agency,
                marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
            )
        except FbsIntegrationError as exc:
            error = str(exc)
    if credential is None and not error:
        error = "Для Ozon не настроен API-ключ."
    elif not client_id and not error:
        error = "Для Ozon не указан Client ID."
    selected_ids = set(
        FbsIntegrationProfile.objects.filter(
            agency=agency,
            marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
            is_active=True,
            order_pull_enabled=True,
        ).values_list("external_warehouse_id", flat=True)
    )
    warehouse_rows = []
    for warehouse in warehouses or ():
        warehouse_id = str(warehouse.warehouse_id or "").strip()
        warehouse_rows.append(
            {
                **warehouse.payload,
                "selected": warehouse_id in selected_ids,
            }
        )
    return {
        "marketplace": FbsIntegrationProfile.MARKETPLACE_OZON,
        "credentials_configured": credential is not None,
        "client_id_configured": bool(client_id),
        "warehouses": warehouse_rows,
        "selected_count": len(selected_ids),
        "error": error,
    }


def _save_ozon_warehouse_selection(*, agency, warehouse_ids: set[str], warehouses) -> None:
    credentials = client_marketplace_credentials(agency)
    credential = credentials.get(FbsIntegrationProfile.MARKETPLACE_OZON)
    if credential is None:
        raise FbsIntegrationError("Для Ozon не настроен API-ключ.")
    client_id = str(credential.client_id or "").strip()
    if not client_id:
        raise FbsIntegrationError("Для Ozon не указан Client ID.")
    warehouse_by_id = {
        str(warehouse.warehouse_id or "").strip(): warehouse
        for warehouse in warehouses
    }
    unknown_ids = warehouse_ids - set(warehouse_by_id)
    if unknown_ids:
        raise FbsIntegrationError("Выбран неизвестный или недоступный склад Ozon.")

    with transaction.atomic():
        existing = list(
            FbsIntegrationProfile.objects.select_for_update()
            .filter(
                agency=agency,
                marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
            )
            .order_by("id")
        )
        by_exact_key = {
            (str(profile.external_account_id or "").strip(), str(profile.external_warehouse_id or "").strip()): profile
            for profile in existing
        }
        by_warehouse: dict[str, list[FbsIntegrationProfile]] = {}
        for profile in existing:
            by_warehouse.setdefault(str(profile.external_warehouse_id or "").strip(), []).append(profile)
        selected_profile_ids: set[int] = set()
        for warehouse_id in sorted(warehouse_ids):
            warehouse = warehouse_by_id[warehouse_id]
            profile = by_exact_key.get((client_id, warehouse_id))
            if profile is None:
                profile = next(iter(by_warehouse.get(warehouse_id, ())), None)
            name = f"Ozon · {warehouse.name}"[:128]
            if profile is None:
                profile = FbsIntegrationProfile(
                    agency=agency,
                    marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
                    name=name,
                    external_account_id=client_id,
                    external_warehouse_id=warehouse_id,
                    stock_mode=FbsIntegrationProfile.STOCK_MODE_DISABLED,
                    is_active=True,
                    order_pull_enabled=True,
                    status_pull_enabled=False,
                    outbox_enabled=False,
                    marking_push_enabled=False,
                    stock_push_enabled=False,
                )
                profile.full_clean(exclude={"id"})
                profile.save()
            else:
                changed = []
                for field, value in (
                    ("name", name),
                    ("external_account_id", client_id),
                    ("is_active", True),
                    ("order_pull_enabled", True),
                ):
                    if getattr(profile, field) != value:
                        setattr(profile, field, value)
                        changed.append(field)
                if changed:
                    profile.full_clean(exclude={"id"})
                    profile.save(update_fields=[*changed, "updated_at"])
            selected_profile_ids.add(profile.id)

        for profile in existing:
            if profile.id in selected_profile_ids:
                continue
            keep_active = any(
                (
                    profile.stock_push_enabled,
                    profile.status_pull_enabled,
                    profile.outbox_enabled,
                    profile.marking_push_enabled,
                )
            )
            changed = []
            for field, value in (
                ("order_pull_enabled", False),
                ("is_active", keep_active),
            ):
                if getattr(profile, field) != value:
                    setattr(profile, field, value)
                    changed.append(field)
            if changed:
                profile.full_clean(exclude={"id"})
                profile.save(update_fields=[*changed, "updated_at"])


def _as_page(value) -> int:
    try:
        return max(int(value or 1), 1)
    except (TypeError, ValueError):
        return 1


def _json_body(request) -> dict:
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (UnicodeDecodeError, json.JSONDecodeError):
        raise ValidationError("Некорректный JSON.")
    if not isinstance(payload, dict):
        raise ValidationError("Ожидался объект JSON.")
    return payload


def _header_key(value) -> str:
    raw = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]", "", raw)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_movement_workbook(uploaded_file) -> list[dict]:
    if not uploaded_file:
        raise ValidationError("Выберите Excel-файл.")
    if int(getattr(uploaded_file, "size", 0) or 0) > 5 * 1024 * 1024:
        raise ValidationError("Размер Excel-файла не должен превышать 5 МБ.")
    filename = str(getattr(uploaded_file, "name", "") or "").lower()
    if not filename.endswith(".xlsx"):
        raise ValidationError("Поддерживается формат .xlsx.")
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as exc:
        raise ValidationError(f"Не удалось прочитать Excel-файл: {exc}")
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise ValidationError("Excel-файл пуст.")

    columns = {}
    for index, value in enumerate(header_row):
        normalized = _header_key(value)
        for key, aliases in _HEADER_ALIASES.items():
            if normalized in aliases and key not in columns:
                columns[key] = index
                break
    missing = [key for key in ("barcode", "qty") if key not in columns]
    if missing:
        raise ValidationError("В файле нужны столбцы «Штрихкод» и «Количество, шт.». ")

    result = []
    for row_number, values in enumerate(rows, start=2):
        barcode = _cell_text(values[columns["barcode"]] if columns["barcode"] < len(values) else "")
        qty = _cell_text(values[columns["qty"]] if columns["qty"] < len(values) else "")
        units_index = columns.get("units_per_box")
        units_per_box = _cell_text(
            values[units_index] if units_index is not None and units_index < len(values) else ""
        )
        boxes_index = columns.get("box_count")
        box_count = _cell_text(
            values[boxes_index] if boxes_index is not None and boxes_index < len(values) else ""
        )
        if not barcode and not qty and not units_per_box and not box_count:
            continue
        result.append(
            {
                "row_number": row_number,
                "barcode": barcode,
                "qty": qty,
                "units_per_box": units_per_box,
                "box_count": box_count,
            }
        )
        if len(result) > 2000:
            raise ValidationError("В одной заявке разрешено не более 2000 строк.")
    if not result:
        raise ValidationError("В Excel-файле нет заполненных строк.")
    return result


def _parse_outbound_workbook(uploaded_file) -> list[dict]:
    return [
        {
            "row_number": row["row_number"],
            "barcode": row["barcode"],
            "qty": row["qty"],
        }
        for row in _parse_movement_workbook(uploaded_file)
    ]


def _match_outbound_workbook(*, agency, raw_lines: list[dict]) -> dict:
    requested_by_barcode: dict[str, int] = {}
    first_row_by_barcode: dict[str, int] = {}
    errors = []
    for row in raw_lines:
        row_number = int(row.get("row_number") or 0)
        barcode = str(row.get("barcode") or "").strip()
        raw_qty = str(row.get("qty") or "").strip()
        if not barcode:
            errors.append(f"Строка {row_number}: укажите штрихкод.")
            continue
        try:
            qty = int(raw_qty)
        except (TypeError, ValueError):
            errors.append(f"Строка {row_number}: количество должно быть целым числом.")
            continue
        if qty < 1 or qty > 1_000_000 or str(qty) != raw_qty:
            errors.append(
                f"Строка {row_number}: количество должно быть целым числом от 1 до 1 000 000."
            )
            continue
        requested_by_barcode[barcode] = requested_by_barcode.get(barcode, 0) + qty
        first_row_by_barcode.setdefault(barcode, row_number)
    if errors:
        raise ValidationError(errors)

    balances = list(
        _fbs_outbound_available_balances(agency=agency)
        .filter(barcode__in=tuple(requested_by_barcode))
        .order_by("barcode", "expiry_date", "box__box_code", "id")
    )
    balances_by_barcode: dict[str, list[FbsStockBalance]] = {}
    for balance in balances:
        balances_by_barcode.setdefault(str(balance.barcode or "").strip(), []).append(balance)

    matched = []
    selected_rows = []
    for barcode, requested_qty in requested_by_barcode.items():
        barcode_balances = balances_by_barcode.get(barcode, [])
        sku_codes = {str(balance.sku_code or "").strip() for balance in barcode_balances}
        if not barcode_balances:
            errors.append(
                f"Строка {first_row_by_barcode[barcode]}: штрихкод {barcode} не найден в свободном остатке FBS."
            )
            continue
        if len(sku_codes) > 1:
            errors.append(
                f"Строка {first_row_by_barcode[barcode]}: штрихкод {barcode} относится к нескольким товарам. Выберите товар вручную."
            )
            continue
        available_qty = sum(int(balance.available_qty or 0) for balance in barcode_balances)
        if available_qty < requested_qty:
            errors.append(
                f"Строка {first_row_by_barcode[barcode]}: по штрихкоду {barcode} доступно {available_qty} шт., запрошено {requested_qty} шт."
            )
            continue
        remaining = requested_qty
        allocation_count = 0
        for balance in barcode_balances:
            selected_qty = min(remaining, int(balance.available_qty or 0))
            if selected_qty <= 0:
                continue
            selected_rows.append(
                _serialize_fbs_outbound_balance(
                    balance,
                    selected_qty=selected_qty,
                )
            )
            allocation_count += 1
            remaining -= selected_qty
            if remaining == 0:
                break
        matched.append(
            {
                "barcode": barcode,
                "sku_code": next(iter(sku_codes), ""),
                "name": str(barcode_balances[0].name or ""),
                "requested_qty": requested_qty,
                "available_qty": available_qty,
                "allocation_count": allocation_count,
            }
        )
    if errors:
        raise ValidationError(errors)
    return {
        "matched": matched,
        "requested_qty": sum(requested_by_barcode.values()),
        "stock": {
            "results": selected_rows,
            "total": len(selected_rows),
            "available_qty": sum(row["available_qty"] for row in selected_rows),
            "truncated": False,
        },
    }


@login_required
@require_GET
def api_fbs_overview(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_overview_payload(agency=agency)
    data["orders"] = _public_fbs_order_counts(
        FbsOrder.objects.filter(profile__agency=agency)
    )
    data["outbound_requests_open"] = FbsExternalIssue.objects.filter(
        agency=agency,
        status__in=("reserved", "picking"),
    ).count()
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_http_methods(["GET", "POST"])
def api_fbs_ozon_warehouses(request):
    agency, access_error = _agency_or_response(request, manage_fbs=True)
    if access_error:
        return access_error
    if request.method == "GET":
        return JsonResponse(
            {"ok": True, "data": _ozon_warehouse_settings_payload(agency=agency)}
        )
    try:
        payload = _json_body(request)
        raw_warehouse_ids = payload.get("warehouse_ids")
        if not isinstance(raw_warehouse_ids, list):
            raise ValidationError("Передайте список складов Ozon.")
        warehouse_ids = {
            str(warehouse_id or "").strip()
            for warehouse_id in raw_warehouse_ids
            if str(warehouse_id or "").strip()
        }
        warehouses = fetch_client_warehouses(
            agency=agency,
            marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
        )
        _save_ozon_warehouse_selection(
            agency=agency,
            warehouse_ids=warehouse_ids,
            warehouses=warehouses,
        )
    except ValidationError as exc:
        return _json_error("Склады Ozon не сохранены.", 400, details=exc.messages)
    except FbsIntegrationError as exc:
        return _json_error(str(exc), 409)
    return JsonResponse(
        {
            "ok": True,
            "data": _ozon_warehouse_settings_payload(
                agency=agency,
                warehouses=warehouses,
            ),
        }
    )


@login_required
@require_GET
def api_fbs_orders(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_orders_payload(
        agency=agency,
        group="all",
        marketplace=str(request.GET.get("marketplace") or "").lower(),
        search=str(request.GET.get("search") or ""),
        page=_as_page(request.GET.get("page")),
    )
    data["results"] = [
        _public_fbs_order_payload(row)
        for row in data.get("results") or []
    ]
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_GET
def api_fbs_order_detail(request, order_id: int):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_order_detail(agency=agency, order_id=order_id)
    if data is None:
        return _json_error("Заказ FBS не найден.", 404)
    return JsonResponse({"ok": True, "data": _public_fbs_order_payload(data)})


@login_required
@require_GET
def api_fbs_stock(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_stock_payload(agency=agency, search=request.GET.get("search") or "")
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_GET
def api_fbs_stock_detail(request, barcode: str):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_stock_detail(agency=agency, barcode=barcode)
    if data is None:
        return _json_error("Остаток по этому ШК не найден.", 404)
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_http_methods(["GET", "POST"])
def api_fbs_outbound(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    can_create = _client_can_create_fbs_outbound(request, agency)
    if request.method == "POST":
        if not can_create:
            return _json_error(
                "Создать заявку на вывоз может владелец или сотрудник клиента с доступом к заявкам.",
                403,
            )
        try:
            direct_agency = resolve_portal_agency_for_user(
                request.user,
                required_section=SECTION_REQUESTS,
            )
            staff_client_request = not bool(
                direct_agency and direct_agency.pk == agency.pk
            )
            payload = _json_body(request)
            request_key = str(payload.get("idempotency_key") or "").strip()
            uuid.UUID(request_key)
            issue = create_issue(
                user=request.user,
                agency_id=agency.pk,
                reference="",
                recipient="",
                purpose="",
                basis=payload.get("basis"),
                selections=payload.get("selections"),
                request_key=request_key,
                client_request=not staff_client_request,
                staff_client_request=staff_client_request,
                shipping_details=payload.get("shipping_details") or {},
            )
        except (FbsError, ValidationError, ValueError, TypeError) as exc:
            messages = getattr(exc, "messages", None)
            return _json_error(
                "Заявка на вывоз не создана.",
                400,
                details=messages or [str(exc)],
            )
        except IntegrityError:
            return _json_error(
                "Остаток или номер заявки изменился во время сохранения. Обновите страницу и повторите.",
                409,
            )
        issue = _fbs_outbound_documents(agency=agency).get(pk=issue.pk)
        return JsonResponse(
            {"ok": True, "data": _public_fbs_outbound_issue(issue)},
            status=201,
        )
    documents = _fbs_outbound_documents(agency=agency)
    stock = (
        _fbs_outbound_stock(agency=agency, search=request.GET.get("search") or "")
        if can_create
        else {"results": [], "total": 0, "available_qty": 0, "truncated": False}
    )
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "can_create": can_create,
                "stock": stock,
                "results": [
                    _public_fbs_outbound_issue(row)
                    for row in documents[:100]
                ],
                "marketplaces": list(
                    Market.objects.order_by("name").values("id", "name")
                ),
            },
        }
    )


@login_required
@require_POST
def api_fbs_outbound_import(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    if not _client_can_create_fbs_outbound(request, agency):
        return _json_error(
            "Загрузить шаблон может владелец или сотрудник клиента с доступом к заявкам.",
            403,
        )
    try:
        uploaded_file = request.FILES.get("file")
        data = _match_outbound_workbook(
            agency=agency,
            raw_lines=_parse_outbound_workbook(uploaded_file),
        )
    except ValidationError as exc:
        return _json_error("Файл не прошёл проверку.", 400, details=exc.messages)
    data["filename"] = str(getattr(uploaded_file, "name", "") or "")
    data["can_create"] = True
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_GET
def fbs_outbound_template(request):
    _, error = _agency_or_response(request)
    if error:
        return error
    template_path = (
        Path(__file__).resolve().parent
        / "static"
        / "client_cabinet"
        / "fbs-outbound-template.xlsx"
    )
    response = FileResponse(
        template_path.open("rb"),
        as_attachment=True,
        filename="fbs-outbound-template.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


@login_required
@require_GET
def api_fbs_movement_stock(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    data = client_movement_source_stock_payload(
        agency=agency,
        search=request.GET.get("search") or "",
        article=request.GET.get("article") or "",
        brand=request.GET.get("brand") or "",
        barcode=request.GET.get("barcode") or "",
        include_box_options=(request.GET.get("mode") == "box"),
    )
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_GET
def api_fbs_reports(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    return JsonResponse({"ok": True, "data": _public_fbs_reports_payload(agency=agency)})


def _fbs_marking_report_period(raw_from: str, raw_to: str) -> tuple[date, date]:
    if not raw_from or not raw_to:
        raise ValidationError("Укажите начало и окончание периода.")
    try:
        date_from = date.fromisoformat(str(raw_from).strip())
        date_to = date.fromisoformat(str(raw_to).strip())
    except (TypeError, ValueError) as exc:
        raise ValidationError("Период указан в неверном формате.") from exc
    if date_from > date_to:
        raise ValidationError("Начало периода не может быть позже окончания.")
    if (date_to - date_from).days > 365:
        raise ValidationError("Период отчёта не должен превышать 366 дней.")
    return date_from, date_to


def _fbs_marking_excel_text(value) -> str:
    """Keep a ЧЗ value readable while removing XML-forbidden controls."""
    safe = []
    for char in str(value or "").strip():
        code = ord(char)
        if code == 29:
            safe.append("<GS>")
        elif code < 32 and char not in {"\t", "\n", "\r"}:
            safe.append(f"<0x{code:02X}>")
        else:
            safe.append(char)
    return "".join(safe)


def build_fbs_shipped_marking_workbook(*, agency, date_from: date, date_to: date) -> Workbook:
    """Build a read-only register of ЧЗ codes from actually dispatched FBS orders."""
    current_timezone = timezone.get_current_timezone()
    start_at = timezone.make_aware(datetime.combine(date_from, time.min), current_timezone)
    end_at = timezone.make_aware(
        datetime.combine(date_to + timedelta(days=1), time.min),
        current_timezone,
    )
    shipment_rows = (
        OrderAuditEntry.objects.filter(
            agency=agency,
            order_type="fbs_order",
            action="status",
            created_at__gte=start_at,
            created_at__lt=end_at,
            payload__changes__internal_status__to=FbsOrder.STATUS_HANDED_OVER,
            payload__source__startswith="dispatch_handover_batch",
        )
        .values_list("order_id", "created_at")
        .order_by("created_at", "id")
    )
    shipped_at_by_order_id = {}
    for raw_order_id, shipped_at in shipment_rows:
        try:
            order_id = int(raw_order_id)
        except (TypeError, ValueError):
            continue
        shipped_at_by_order_id.setdefault(order_id, shipped_at)

    transfers = []
    if shipped_at_by_order_id:
        transfers = list(
            FbsMarketplaceMetadataTransfer.objects.filter(
                order_item__order_id__in=tuple(shipped_at_by_order_id),
                order_item__order__profile__agency=agency,
                metadata_type=FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE,
                value__gt="",
            )
            .select_related("order_item__order__profile", "order_item__sku")
            .order_by("order_item__order_id", "order_item_id", "id")
        )
        transfers.sort(
            key=lambda transfer: (
                shipped_at_by_order_id[transfer.order_item.order_id],
                transfer.order_item.order_id,
                transfer.order_item_id,
                transfer.id,
            )
        )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Честный знак FBS"
    headers = [
        "№",
        "Дата отгрузки",
        "Маркетплейс",
        "Заказ маркетплейса",
        "ID заказа Fullbox",
        "Артикул",
        "Наименование",
        "ШК товара",
        "Код ЧЗ (GS=<GS>)",
        "Статус передачи кода",
        "Текущий статус заказа",
    ]
    worksheet.append(headers)
    header_fill = PatternFill(fill_type="solid", fgColor="F89000")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    marketplace_labels = dict(FbsIntegrationProfile.MARKETPLACE_CHOICES)
    order_status_labels = dict(FbsOrder.STATUS_CHOICES)
    transfer_status_labels = dict(FbsMarketplaceMetadataTransfer.STATUS_CHOICES)
    for row_number, transfer in enumerate(transfers, start=1):
        item = transfer.order_item
        order = item.order
        shipped_at = shipped_at_by_order_id[order.id]
        shipped_at_display = timezone.localtime(shipped_at).strftime("%d.%m.%Y %H:%M")
        product_name = str(item.product_name or getattr(item.sku, "name", "") or "")
        worksheet.append(
            [
                row_number,
                shipped_at_display,
                marketplace_labels.get(order.profile.marketplace, order.profile.marketplace),
                order.external_order_id,
                str(order.id),
                item.external_sku,
                product_name,
                item.barcode,
                _fbs_marking_excel_text(transfer.value),
                transfer_status_labels.get(transfer.status, transfer.status),
                order_status_labels.get(order.internal_status, order.internal_status),
            ]
        )

    if not transfers:
        worksheet.append(
            [
                "",
                "",
                "",
                "",
                "",
                "",
                "Коды ЧЗ по отгруженным FBS-заказам за выбранный период отсутствуют",
            ]
        )

    for row in worksheet.iter_rows(min_row=2, min_col=2, max_col=len(headers)):
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 32
    for column, width in {
        "A": 8,
        "B": 20,
        "C": 18,
        "D": 24,
        "E": 18,
        "F": 24,
        "G": 42,
        "H": 24,
        "I": 58,
        "J": 25,
        "K": 24,
    }.items():
        worksheet.column_dimensions[column].width = width
    return workbook


@login_required
@require_GET
def api_fbs_marking_report_export(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    try:
        date_from, date_to = _fbs_marking_report_period(
            request.GET.get("date_from") or "",
            request.GET.get("date_to") or "",
        )
    except ValidationError as exc:
        return _json_error("; ".join(exc.messages), 400)

    workbook = build_fbs_shipped_marking_workbook(
        agency=agency,
        date_from=date_from,
        date_to=date_to,
    )
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="fbs-chestny-znak-{date_from.isoformat()}_{date_to.isoformat()}.xlsx"'
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


@login_required
@require_http_methods(["GET", "POST"])
def api_fbs_movements(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    if request.method == "GET":
        return JsonResponse({"ok": True, "data": client_movements_payload(agency=agency)})
    try:
        payload = _json_body(request)
        idempotency_key = str(payload.get("idempotency_key") or "").strip()
        if not idempotency_key:
            return _json_error(
                "Форма устарела. Обновите страницу перед отправкой заявки.",
                409,
            )
        row = create_client_movement_request(
            agency=agency,
            mode=str(payload.get("mode") or ""),
            raw_lines=payload.get("lines") or [],
            requested_by=request.user,
            comment=payload.get("comment") or "",
            source_file_name=payload.get("source_file_name") or "",
            requested_box_count=payload.get("requested_box_count"),
            requested_mixed_box_count=payload.get("requested_mixed_box_count", 0),
            mixed_box_codes=payload.get("mixed_box_codes") or [],
            idempotency_key=idempotency_key,
        )
    except ValidationError as exc:
        return _json_error("Заявка не создана.", 400, details=exc.messages)
    return JsonResponse(
        {"ok": True, "data": serialize_movement_request(row, include_lines=True)},
        status=201,
    )


@login_required
@require_http_methods(["GET", "POST"])
def api_fbs_movement_detail(request, request_id: int):
    agency, error = _agency_or_response(request)
    if error:
        return error
    if request.method == "POST":
        try:
            payload = _json_body(request)
            if str(payload.get("action") or "").strip() != "cancel":
                return _json_error("Неизвестное действие с заявкой FBS.", 400)
            cancel_client_movement_request(
                request_id=request_id,
                agency=agency,
                canceled_by=request.user,
            )
        except FbsClientMovementRequest.DoesNotExist:
            return _json_error("Заявка FBS не найдена.", 404)
        except FbsReplenishmentError as exc:
            return _json_error(str(exc), 409)
    data = client_movement_detail(agency=agency, request_id=request_id)
    if data is None:
        return _json_error("Заявка FBS не найдена.", 404)
    return JsonResponse({"ok": True, "data": data})


@login_required
@require_POST
def api_fbs_movement_import(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    mode = str(request.POST.get("mode") or "")
    try:
        uploaded_file = request.FILES.get("file")
        raw_lines = _parse_movement_workbook(uploaded_file)
        lines, errors = validate_client_movement_lines(
            agency=agency,
            mode=mode,
            raw_lines=raw_lines,
        )
    except ValidationError as exc:
        return _json_error("Файл не прошёл проверку.", 400, details=exc.messages)
    serializable = [
        {key: value for key, value in row.items() if key != "sku"}
        for row in lines
    ]
    if errors:
        return _json_error("Файл не прошёл проверку.", 400, details=errors)
    return JsonResponse(
        {
            "ok": True,
            "data": {
                "filename": str(getattr(uploaded_file, "name", "") or ""),
                "lines": serializable,
            },
        }
    )


@login_required
@require_GET
def fbs_movement_template(request):
    agency, error = _agency_or_response(request)
    if error:
        return error
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявка FBS"
    sheet.append(["Штрихкод", "Количество, шт.", "Кратность короба, шт.", "Коробов для поштучного"])
    sheet.append(["4600000000001", 24, 12, ""])
    sheet.append(["4600000000002", 5, "", 2])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="F89000")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 28
    instructions = workbook.create_sheet("Инструкция")
    instructions.append(["Правила заполнения"])
    instructions.append(["Сопоставление выполняется только по точному штрихкоду клиента."])
    instructions.append(["В поштучном режиме укажите количество новых FBS-коробов; кратность можно оставить пустой."])
    instructions.append(["В режиме коробами количество должно быть кратно количеству штук в коробе."])
    instructions.column_dimensions["A"].width = 95
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="fbs-movement-template.xlsx"'
    return response
