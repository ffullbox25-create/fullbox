import json
import uuid
from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accountant.models import ClientLifecycle
from employees.models import Employee
from fbs import test_inventory as fixtures
from fbs.exceptions import FbsError
from fbs.models import (
    FbsBox,
    FbsExternalIssue,
    FbsIntegrationProfile,
    FbsPallet,
    FbsStockBalance,
)
from fbs.services.external_issues import (
    create_issue,
    issue_command,
    review_issue_by_manager,
)
from sku.models import Agency, Market


@override_settings(
    FBS_MODULE_ENABLED=True,
    FBS_WAREHOUSE_WRITES_ENABLED=True,
    FBS_STOCK_PUSH_ENABLED=False,
)
class ClientFbsOutboundTests(TestCase):
    def setUp(self):
        fixtures.FbsInventoryScanModeTests.setUp(self)
        self.pallet.status = FbsPallet.STATUS_ACTIVE
        self.pallet.save(update_fields=["status"])
        self.box.status = FbsBox.STATUS_ACTIVE
        self.box.save(update_fields=["status"])
        self.balance = FbsStockBalance.objects.create(
            agency=self.agency,
            box=self.box,
            sku_ref=self.sku,
            identity_key="plain",
            sku_code=self.sku.sku_code,
            name=self.sku.name,
            barcode="4600000000401",
            qty=12,
            available_qty=9,
            reserved_qty=3,
        )
        FbsIntegrationProfile.objects.create(
            agency=self.agency,
            marketplace=FbsIntegrationProfile.MARKETPLACE_WB,
            name="Кабинет клиента",
            external_warehouse_id="client-outbound-test",
            is_active=True,
            order_pull_enabled=True,
        )
        self.portal_user = get_user_model().objects.create_user(
            username="client-outbound-owner",
            password="test-password",
        )
        self.agency.portal_user = self.portal_user
        self.agency.save(update_fields=["portal_user"])
        ClientLifecycle.objects.update_or_create(
            agency=self.agency,
            defaults={"status": ClientLifecycle.STATUS_ACTIVE},
        )
        self.client.force_login(self.portal_user)
        self.url = reverse("client-api-fbs-outbound")
        self.import_url = reverse("client-api-fbs-outbound-import")
        self.template_url = reverse("client-fbs-outbound-template")
        self.marketplace = Market.objects.create(id=904, name="OZON FBS OUTBOUND")
        self.storekeeper = get_user_model().objects.create_user(username="outbound-storekeeper")
        Employee.objects.create(
            user=self.storekeeper,
            full_name="Кладовщик вывоза",
            role="storekeeper",
            is_active=True,
        )
        self.manager = get_user_model().objects.create_user(username="outbound-manager")
        Employee.objects.create(
            user=self.manager,
            full_name="Менеджер вывоза",
            role="head_manager",
            is_active=True,
        )

    def outbound_workbook(self, rows, *, headers=("Штрихкод", "Количество, шт.")):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(headers))
        for row in rows:
            sheet.append(list(row))
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            "fbs-outbound.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def payload(self, *, request_key=None, **changes):
        data = {
            "basis": "Забрать товар со склада FBS",
            "shipping_details": {
                "delivery_type": "pickup",
                "eta_date": (timezone.localdate() + timedelta(days=1)).isoformat(),
                "vehicle_type": "client",
                "vehicle_number": "A123BC77",
                "driver_phone": "+7 900 000-00-00",
            },
            "selections": [{"balance_id": self.balance.pk, "qty": 4}],
            "idempotency_key": request_key or str(uuid.uuid4()),
        }
        data.update(changes)
        return data

    def test_client_sees_only_free_fbs_stock(self):
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 200)
        data = response.json()["data"]
        self.assertTrue(data["can_create"])
        self.assertEqual(data["stock"]["available_qty"], 9)
        self.assertEqual(data["stock"]["results"][0]["balance_id"], self.balance.pk)
        self.assertEqual(data["stock"]["results"][0]["available_qty"], 9)
        self.assertIn(
            self.marketplace.pk,
            [row["id"] for row in data["marketplaces"]],
        )

    def test_client_form_replaces_legacy_recipient_fields_with_shipping_parameters(self):
        html = render_to_string(
            "client_cabinet/_fbs_lk.html",
            {"selected_client": self.agency},
        )

        self.assertIn('id="fbs-outbound-delivery-type"', html)
        self.assertIn('id="fbs-outbound-marketplace"', html)
        self.assertIn('id="fbs-outbound-eta-date"', html)
        self.assertIn('id="fbs-outbound-file"', html)
        self.assertIn('id="fbs-outbound-file-button"', html)
        self.assertIn("fbs/outbound-template.xlsx", html)
        self.assertNotIn('id="fbs-outbound-recipient"', html)
        self.assertNotIn('id="fbs-outbound-purpose"', html)
        self.assertNotIn('id="fbs-outbound-reference"', html)

    def test_outbound_template_has_barcode_and_quantity_columns(self):
        response = self.client.get(self.template_url)

        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        self.assertEqual(
            [cell.value for cell in workbook.active[1]][:2],
            ["Штрихкод", "Количество, шт."],
        )
        self.assertIn("fbs-outbound-template.xlsx", response["Content-Disposition"])

    def test_outbound_upload_matches_stock_without_creating_reserve(self):
        response = self.client.post(
            self.import_url,
            {"file": self.outbound_workbook([(self.balance.barcode, 4)])},
        )

        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["requested_qty"], 4)
        self.assertEqual(data["matched"][0]["barcode"], self.balance.barcode)
        self.assertEqual(data["stock"]["results"][0]["balance_id"], self.balance.pk)
        self.assertEqual(data["stock"]["results"][0]["selected_qty"], 4)
        self.assertFalse(FbsExternalIssue.objects.exists())
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.available_qty, 9)
        self.assertEqual(self.balance.external_reserved_qty, 0)

    def test_outbound_upload_aggregates_duplicate_barcodes(self):
        response = self.client.post(
            self.import_url,
            {"file": self.outbound_workbook([(self.balance.barcode, 2), (self.balance.barcode, 3)])},
        )

        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()["data"]
        self.assertEqual(data["requested_qty"], 5)
        self.assertEqual(data["matched"][0]["requested_qty"], 5)
        self.assertEqual(data["stock"]["results"][0]["selected_qty"], 5)

    def test_outbound_upload_rejects_missing_or_insufficient_stock(self):
        response = self.client.post(
            self.import_url,
            {"file": self.outbound_workbook([(self.balance.barcode, 10), ("9999999999999", 1)])},
        )

        self.assertEqual(response.status_code, 400, response.content)
        details = " ".join(response.json()["details"])
        self.assertIn("доступно 9 шт., запрошено 10 шт.", details)
        self.assertIn("не найден в свободном остатке FBS", details)
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.external_reserved_qty, 0)

    def test_client_submission_requires_shipping_parameters(self):
        payload = self.payload()
        payload["shipping_details"] = {}

        response = self.client.post(
            self.url,
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Дата отгрузки", response.json()["details"][0])
        self.assertFalse(FbsExternalIssue.objects.exists())
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.external_reserved_qty, 0)

    def test_client_submission_creates_fbs_out_and_reserves_free_stock(self):
        response = self.client.post(
            self.url,
            data=json.dumps(self.payload()),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201, response.content)
        issue = FbsExternalIssue.objects.get()
        self.assertEqual(issue.agency, self.agency)
        self.assertEqual(issue.created_by, self.portal_user)
        self.assertTrue(issue.reference.startswith(f"LK-{self.agency.pk}-"))
        self.assertEqual(issue.recipient, self.agency.agn_name)
        self.assertEqual(issue.purpose, "owner")
        self.assertEqual(issue.shipping_details["delivery_type"], "pickup")
        self.assertEqual(issue.shipping_details["vehicle_type"], "client")
        self.assertEqual(issue.lines.get().requested_qty, 4)
        response_data = response.json()["data"]
        self.assertEqual(response_data["status"], "awaiting_manager")
        self.assertEqual(response_data["delivery_type_label"], "Самовывоз")
        self.assertNotIn("recipient", response_data)
        self.assertNotIn("purpose", response_data)
        self.assertNotIn("reference", response_data)
        self.assertEqual(
            response_data["status_label"],
            "Ожидает подтверждения менеджером",
        )
        self.assertEqual(issue.events.get(action="created").payload["source"], "client_portal")
        self.assertEqual(
            issue.events.get(action="created").payload["shipping_details"]["delivery_type"],
            "pickup",
        )
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.available_qty, 5)
        self.assertEqual(self.balance.reserved_qty, 3)
        self.assertEqual(self.balance.external_reserved_qty, 4)

    def test_marketplace_submission_keeps_shipping_parameters_without_shipping_order(self):
        payload = self.payload()
        payload["shipping_details"] = {
            "delivery_type": "marketplace",
            "marketplace_id": self.marketplace.pk,
            "shipping_barcode": "SHIP-001",
            "supply_number": "SUPPLY-001",
            "supply_type": "box",
            "destination_warehouse": "Ozon Хоругвино",
            "slot_date": (timezone.localdate() + timedelta(days=2)).isoformat(),
            "slot_time": "11:30",
            "eta_date": (timezone.localdate() + timedelta(days=1)).isoformat(),
            "vehicle_type": "fulfillment",
            "vehicle_number": "B456CD77",
            "driver_phone": "+7 901 000-00-00",
        }

        response = self.client.post(
            self.url,
            data=json.dumps(payload),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201, response.content)
        issue = FbsExternalIssue.objects.get()
        self.assertEqual(issue.purpose, "external")
        self.assertEqual(issue.shipping_details["delivery_type"], "marketplace")
        self.assertEqual(issue.shipping_details["marketplace_name"], self.marketplace.name)
        self.assertEqual(issue.shipping_details["shipping_barcode"], "SHIP-001")
        self.assertEqual(issue.shipping_details["destination_warehouse"], "Ozon Хоругвино")
        self.assertEqual(response.json()["data"]["delivery_type_label"], "Маркетплейс")

    def test_storekeeper_cannot_work_before_manager_approval(self):
        response = self.client.post(
            self.url,
            data=json.dumps(self.payload()),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        issue = FbsExternalIssue.objects.get()
        line = issue.lines.get()

        with self.assertRaisesMessage(FbsError, "ожидает подтверждения менеджером"):
            issue_command(
                user=self.storekeeper,
                issue_id=issue.pk,
                action="pick",
                request_key=str(uuid.uuid4()),
                line_id=line.pk,
                cell_scan=self.cell.cell_code,
                box_scan=self.box.box_code,
                item_scan=self.balance.barcode,
            )

        self.client.force_login(self.storekeeper)
        list_response = self.client.get(reverse("fbs:external_issues"), {"client": self.agency.pk})
        self.assertNotContains(list_response, issue.number)
        self.assertEqual(
            self.client.get(reverse("fbs:external_issue_detail", args=[issue.pk])).status_code,
            404,
        )

    def test_manager_approval_makes_request_visible_to_storekeeper(self):
        response = self.client.post(
            self.url,
            data=json.dumps(self.payload()),
            content_type="application/json",
        )
        issue = FbsExternalIssue.objects.get()

        self.client.force_login(self.manager)
        manager_page = self.client.get(reverse("team-manager-fbs-movements"))
        self.assertContains(manager_page, issue.number)
        self.assertContains(manager_page, "Ожидает подтверждения менеджером")
        manager_search = self.client.get(
            reverse("team-manager-fbs-movements"),
            {"q": issue.number},
        )
        self.assertContains(manager_search, issue.number)
        self.assertEqual(
            [row.pk for row in manager_search.context["external_rows"]],
            [issue.pk],
        )
        manager_queue = self.client.get(reverse("team-manager-queue"))
        self.assertContains(manager_queue, issue.number)
        queue_row = next(
            row for row in manager_queue.context["rows"] if row["number"] == issue.number
        )
        self.assertEqual(queue_row["state"], "to_confirm")
        self.assertEqual(queue_row["client_name"], self.agency.agn_name)
        self.assertIn(f"q={issue.number}", queue_row["open_url"])
        self.assertGreaterEqual(manager_queue.context["counts"]["to_confirm"], 1)
        approve = self.client.post(
            reverse("team-manager-fbs-movements"),
            {
                "action": "approve_outbound",
                "issue_id": issue.pk,
                "request_key": str(uuid.uuid4()),
            },
        )
        self.assertEqual(approve.status_code, 302)
        self.assertTrue(issue.events.filter(action="manager_approved").exists())
        queue_after_approval = self.client.get(reverse("team-manager-queue"))
        self.assertNotIn(
            issue.number,
            [row["number"] for row in queue_after_approval.context["rows"]],
        )

        self.client.force_login(self.storekeeper)
        list_response = self.client.get(reverse("fbs:external_issues"), {"client": self.agency.pk})
        self.assertContains(list_response, issue.number)
        self.assertEqual(
            self.client.get(reverse("fbs:external_issue_detail", args=[issue.pk])).status_code,
            200,
        )

    def test_manager_rejection_releases_only_outbound_reserve(self):
        response = self.client.post(
            self.url,
            data=json.dumps(self.payload()),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201, response.content)
        issue = FbsExternalIssue.objects.get()

        review_issue_by_manager(
            user=self.manager,
            issue_id=issue.pk,
            action="reject",
            request_key=str(uuid.uuid4()),
        )

        issue.refresh_from_db()
        self.balance.refresh_from_db()
        self.assertEqual(issue.status, "canceled")
        self.assertEqual(self.balance.available_qty, 9)
        self.assertEqual(self.balance.reserved_qty, 3)
        self.assertEqual(self.balance.external_reserved_qty, 0)
        self.assertTrue(issue.events.filter(action="manager_rejected").exists())

    def test_staff_created_client_request_still_requires_explicit_approval(self):
        issue = create_issue(
            user=self.manager,
            agency_id=self.agency.pk,
            reference="STAFF-CLIENT-1",
            recipient="Представитель владельца",
            purpose="owner",
            basis="Заявка клиента передана менеджеру",
            selections=[{"balance_id": self.balance.pk, "qty": 2}],
            request_key=str(uuid.uuid4()),
            staff_client_request=True,
        )
        self.assertEqual(
            issue.events.get(action="created").payload["source"],
            "staff_client_portal",
        )
        with self.assertRaisesMessage(FbsError, "ожидает подтверждения менеджером"):
            issue_command(
                user=self.storekeeper,
                issue_id=issue.pk,
                action="pick",
                request_key=str(uuid.uuid4()),
                line_id=issue.lines.get().pk,
                cell_scan=self.cell.cell_code,
                box_scan=self.box.box_code,
                item_scan=self.balance.barcode,
            )

    def test_retry_with_same_key_does_not_create_duplicate(self):
        key = str(uuid.uuid4())
        payload = self.payload(request_key=key)
        first = self.client.post(self.url, data=json.dumps(payload), content_type="application/json")
        second = self.client.post(self.url, data=json.dumps(payload), content_type="application/json")

        self.assertEqual(first.status_code, 201, first.content)
        self.assertEqual(second.status_code, 201, second.content)
        self.assertEqual(FbsExternalIssue.objects.count(), 1)
        self.balance.refresh_from_db()
        self.assertEqual(self.balance.external_reserved_qty, 4)

    def test_client_cannot_create_issue_for_another_agency(self):
        other = Agency.objects.create(agn_name="Другой клиент")
        with self.assertRaisesMessage(FbsError, "только для своего клиента"):
            create_issue(
                user=self.portal_user,
                agency_id=other.pk,
                reference="OTHER-1",
                recipient="Получатель",
                purpose="owner",
                basis="Попытка чужой выдачи",
                selections=[{"balance_id": self.balance.pk, "qty": 1}],
                request_key=str(uuid.uuid4()),
                client_request=True,
            )

    def test_client_cannot_register_historical_issue(self):
        with self.assertRaisesMessage(FbsError, "задним числом"):
            create_issue(
                user=self.portal_user,
                agency_id=self.agency.pk,
                reference="HIST-1",
                recipient="Получатель",
                purpose="owner",
                basis="Историческая выдача",
                selections=[{"balance_id": self.balance.pk, "qty": 1}],
                request_key=str(uuid.uuid4()),
                historical=True,
                client_request=True,
            )
