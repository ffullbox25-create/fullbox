"""Client LK request detail helpers (list/detail URLs, attention, actions)."""

from __future__ import annotations

import re
from typing import Any
from urllib.parse import unquote

from audit.models import OrderAuditEntry, log_order_action
from django.utils import timezone
from fullbox.order_numbers import build_public_order_number, format_order_number, order_number_sequence

from .other_requests import (
    ORDER_TYPE as OTHER_ORDER_TYPE,
    cancel_other_request,
    latest_other_entry,
    resolve_client_manager,
)
from .request_titles import build_request_display_title


CLIENT_REQUEST_ORDER_TYPES = frozenset({"receiving", "shipping", "processing", "packing", "other"})


def _user_full_name(user) -> str:
    if not user:
        return ""
    for attr in ("get_full_name", "get_username"):
        value = getattr(user, attr, None)
        if callable(value):
            text = str(value() or "").strip()
            if text:
                return text
    parts = [str(getattr(user, "first_name", "") or "").strip(), str(getattr(user, "last_name", "") or "").strip()]
    text = " ".join(part for part in parts if part).strip()
    if text:
        return text
    return str(getattr(user, "username", "") or "").strip()


def _request_responsibles(*, agency, entries: list) -> list[dict[str, str]]:
    client_manager = ""
    for entry in entries or []:
        client_manager = _user_full_name(getattr(entry, "user", None))
        if client_manager:
            break
    company_manager_obj = resolve_client_manager(agency)
    company_manager = str(getattr(company_manager_obj, "full_name", "") or "").strip()
    return [
        {"label": "Менеджер клиента", "value": client_manager or "Не указан"},
        {"label": "Менеджер Fullbox", "value": company_manager or "Не назначен"},
    ]


def lk_request_hash(order_type: str, order_id: str) -> str:
    raw_type = str(order_type or "").strip().lower() or "other"
    if raw_type == "packing":
        raw_type = "processing"
    safe_order_id = canonical_request_order_id(raw_type, order_id)
    return f"#/request/{raw_type}/{safe_order_id}"


def canonical_shipping_order_id(order_id: str) -> str:
    """Return the real ShippingOrder.number for LK/audit helper links.

    Some historical audit rows used technical ids like
    ``SO-000051#reopened#5891``. The warehouse order number is the prefix before
    ``#``; client cards and exports must read the live ShippingOrder by that
    canonical number. New LK shipping numbers are stored as ``OTG-000091`` and
    older client links may still use the display form ``91_OTG``.
    """

    value = unquote(str(order_id or "").strip())
    base_value = value.split("#", 1)[0].strip() if "#" in value else value
    if not base_value:
        return base_value
    if base_value.upper().startswith("SO-"):
        return base_value
    sequence = order_number_sequence("shipping", base_value)
    if sequence is not None:
        return build_public_order_number("shipping", sequence)
    return base_value


def canonical_request_order_id(order_type: str, order_id: str) -> str:
    raw_type = str(order_type or "").strip().lower()
    if raw_type == "packing":
        raw_type = "processing"
    if raw_type == "shipping":
        return canonical_shipping_order_id(order_id)
    return str(order_id or "").strip()


def notification_detail_url(*, order_type: str, order_id: str, agency_id=None) -> str:
    raw = str(order_type or "").strip().lower()
    if raw == "packing":
        raw = "processing"
    if raw in CLIENT_REQUEST_ORDER_TYPES or raw == "processing":
        return lk_request_hash(raw, order_id)
    if agency_id:
        return f"/orders/{raw}/{order_id}/?client={agency_id}"
    return f"/orders/{raw}/{order_id}/"


def build_client_notifications(agency, *, limit: int = 30) -> list[dict[str, Any]]:
    """LK notifications from ClientNotification (synced from audit)."""
    from .messaging_lk import list_notifications

    return list_notifications(agency, limit=limit)


def build_action_urls(agency) -> dict[str, str]:
    from .services import build_client_cabinet_url

    cid = agency.id if agency else ""
    suffix = f"?client={cid}" if cid else ""
    lk = build_client_cabinet_url(cid or None) if cid else "/client/dashboard/lk/"
    return {
        "receiving_new": f"/client/{cid}/receiving/new/" if cid else "#",
        "shipping_new": f"/client/{cid}/shipping/new/" if cid else "#",
        "processing_new": f"/client/{cid}/packing/new/" if cid else "#",
        "other_new": f"{lk}#/other-requests",
        "other_journal": f"/orders/other/{suffix}",
        "sku_new": f"/client/{cid}/sku/new/" if cid else "#",
        "sku_list": f"/client/{cid}/sku/?view=table" if cid else "#",
        "sku_export": f"/client/api/v1/nomenclature/export/?client={cid}" if cid else "#",
        "sku_template": "/orders/templates/sku-upload/",
        "sku_template_upload": "/client/api/v1/nomenclature/upload/",
        "stock_journal": f"/sklad/journal/{suffix}",
        "market_sync": f"/market-sync/{suffix}",
        "marking": f"/client/marking/{suffix}",
        "legacy_dashboard": lk,
        "dashboard": lk,
        "client_edit": f"/client/{cid}/edit/" if cid else "/client/",
        "orders_journal": f"/orders/{suffix}",
        "shipping_journal": f"/shipping/{suffix}",
    }


def _receiving_attention_for_order(agency, order_id: str) -> bool:
    from . import web_ui as client_web_ui

    entries = list(
        OrderAuditEntry.objects.filter(agency=agency, order_type="receiving", order_id=str(order_id))
        .order_by("created_at")
    )
    if not entries:
        return False
    act_entry = None
    status_entry = None
    for entry in reversed(entries):
        payload = entry.payload or {}
        if act_entry is None and payload.get("act_sent"):
            act_entry = entry
        if status_entry is None and client_web_ui._is_status_entry(entry):
            status_entry = entry
        if act_entry and status_entry:
            break
    if not act_entry:
        return False
    return client_web_ui._receiving_act_needs_client_attention(act_entry, status_entry or entries[-1])


def _merge_receiving_act_fields(entries, payload: dict | None = None) -> dict:
    """Overlay act_* fields from the full audit trail onto the latest status payload.

    Manager send / client response / viewed often live on different entries; the
    latest status row alone is not enough for client-facing act UX.
    """
    merged = dict(payload or {})
    for entry in entries or []:
        data = getattr(entry, "payload", None)
        if not isinstance(data, dict):
            continue
        if data.get("act_sent"):
            merged["act_sent"] = data.get("act_sent")
        for key in (
            "act_client_response",
            "act_client_response_at",
            "act_viewed",
            "act_viewed_at",
            "act_manager_signed",
            "act_storekeeper_signed",
            "act_logistician_signed",
        ):
            if key in data and data.get(key) not in (None, ""):
                merged[key] = data.get(key)
    return merged


def _receiving_act_block(*, agency, order_id: str, entries) -> dict[str, Any] | None:
    """Client-facing act card for request detail (open/download + confirm)."""
    from urllib.parse import quote, urlencode

    merged = _merge_receiving_act_fields(entries)
    act_sent = merged.get("act_sent")
    if not act_sent:
        return None
    response = str(merged.get("act_client_response") or "").strip().lower()
    viewed = bool(merged.get("act_viewed"))
    if response == "confirmed":
        status, status_label = "confirmed", "Подтверждён клиентом"
    elif response == "dispute":
        status, status_label = "dispute", "Разногласия по акту"
    elif viewed:
        status, status_label = "viewed", "Ожидает подтверждения"
    else:
        status, status_label = "sent", "Отправлен клиенту"
    needs_confirm = response not in {"confirmed", "dispute"}
    title = str(act_sent).strip()
    if not title or title.lower() in {"true", "1", "yes"}:
        title = f"Акт приёмки №{order_id}"
    cid = getattr(agency, "id", None)
    params = {}
    if cid:
        params["client"] = cid
        params["return"] = f"/client/dashboard/lk/?client={cid}{lk_request_hash('receiving', order_id)}"
    suffix = f"?{urlencode(params)}" if params else ""
    print_url = f"/orders/receiving/{order_id}/act/print/{suffix}"
    excel_params = {"client": cid} if cid else {}
    excel_suffix = f"?{urlencode(excel_params)}" if excel_params else ""
    excel_url = (
        f"/client/api/v1/requests/receiving/"
        f"{quote(str(order_id), safe='')}/act.xlsx{excel_suffix}"
    )
    return {
        "exists": True,
        "title": title,
        "status": status,
        "status_label": status_label,
        "needs_confirm": needs_confirm,
        "print_url": print_url,
        "open_url": print_url,
        "open_label": "Открыть / скачать акт",
        "excel_url": excel_url,
        "excel_label": "Скачать акт в Excel",
        "confirm_label": "Подтвердить акт" if needs_confirm else "",
        "hint": (
            "Менеджер подписал акт. Откройте документ, чтобы скачать и подтвердить приёмку."
            if needs_confirm
            else "Акт приёмки можно открыть и скачать."
        ),
    }


def _receiving_act_excel_text(value: Any) -> str:
    """Return a literal Excel value and neutralise formula-like source text."""
    text = str(value or "").strip()
    if text[:1] in {"=", "+", "-", "@"}:
        return f"'{text}"
    return text


def build_receiving_act_excel_workbook(*, agency, order_id: str):
    """Build the client-visible receiving act as a styled, auditable XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    detail = build_request_detail(
        agency=agency,
        order_type="receiving",
        order_id=order_id,
    )
    act = (detail or {}).get("act") or {}
    if not detail or detail.get("type") != "receiving" or not act.get("exists"):
        return None, None

    lines = list(detail.get("lines") or [])
    agency_name = str(
        getattr(agency, "agn_name", None)
        or getattr(agency, "short_name", None)
        or ""
    ).strip()
    request_number = str(detail.get("number") or order_id).strip()
    act_title = str(act.get("title") or f"Акт приёмки {request_number}").strip()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Акт приёмки"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    sheet.print_title_rows = "1:6"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.auto_filter.ref = "A6:K6"

    dark = "303030"
    orange = "F89000"
    pale_orange = "FFF1DA"
    pale_red = "FDE8E7"
    white = "FFFFFF"
    border_color = "D9D3C8"
    thin = Side(style="thin", color=border_color)
    row_border = Border(bottom=thin)

    sheet.merge_cells("A1:K1")
    sheet["A1"] = _receiving_act_excel_text(act_title)
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=dark)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 30

    metadata = (
        (2, "Клиент", agency_name),
        (3, "Заявка", request_number),
        (4, "Статус акта", act.get("status_label") or ""),
        (5, "Дата выгрузки", timezone.localdate()),
    )
    for row_index, label, value in metadata:
        sheet.cell(row=row_index, column=1, value=label)
        sheet.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=2,
        )
        sheet.cell(
            row=row_index,
            column=3,
            value=value if row_index == 5 else _receiving_act_excel_text(value),
        )
        sheet.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=11)
        sheet.cell(row=row_index, column=1).font = Font(name="Arial", size=10, bold=True, color=dark)
        sheet.cell(row=row_index, column=3).font = Font(name="Arial", size=10, color=dark)
        sheet.cell(row=row_index, column=3).alignment = Alignment(horizontal="left")
    sheet["C5"].number_format = "dd.mm.yyyy"

    headers = (
        "№",
        "Штрихкод",
        "Артикул",
        "Наименование",
        "Размер",
        "Тип товара",
        "План",
        "Факт",
        "Короба",
        "Расхождение",
        "Комментарий",
    )
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=6, column=column, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=orange)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[6].height = 30

    first_data_row = 7
    for index, line in enumerate(lines, start=1):
        row_index = first_data_row + index - 1
        actual = line.get("qty_actual")
        box_qty = line.get("box_qty")
        if box_qty in (None, "", "—"):
            box_qty = None
        row_values = (
            index,
            _receiving_act_excel_text(line.get("barcode")),
            _receiving_act_excel_text(line.get("sku_code")),
            _receiving_act_excel_text(line.get("name")),
            _receiving_act_excel_text(line.get("size")),
            _receiving_act_excel_text(
                line.get("goods_type_label") or line.get("goods_type")
            ),
            _as_int(line.get("qty_planned")),
            None if actual in (None, "") else _as_int(actual),
            None if box_qty is None else _as_int(box_qty),
            None,
            _receiving_act_excel_text(line.get("comment")),
        )
        for column, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Arial", size=9, color=dark)
            cell.border = row_border
            cell.alignment = Alignment(
                horizontal="right" if column in {1, 7, 8, 9, 10} else "left",
                vertical="top",
                wrap_text=column in {4, 6, 11},
            )
        sheet.cell(
            row=row_index,
            column=10,
            value=f'=IF(H{row_index}="","",H{row_index}-G{row_index})',
        )
        sheet.cell(row=row_index, column=10).number_format = "#,##0;[Red]-#,##0;—"
        if line.get("mismatch"):
            sheet.cell(row=row_index, column=10).fill = PatternFill("solid", fgColor=pale_red)

    last_data_row = first_data_row + len(lines) - 1
    total_row = max(first_data_row, last_data_row + 1)
    sheet.cell(row=total_row, column=1, value="ИТОГО")
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
    for column in range(1, 12):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = Font(name="Arial", size=10, bold=True, color=dark)
        cell.fill = PatternFill("solid", fgColor=pale_orange)
        cell.border = Border(top=Side(style="medium", color=orange))
        cell.alignment = Alignment(
            horizontal="right" if column >= 7 else "left",
            vertical="center",
        )
    if lines:
        sheet.cell(row=total_row, column=7, value=f"=SUM(G{first_data_row}:G{last_data_row})")
        sheet.cell(row=total_row, column=8, value=f"=SUM(H{first_data_row}:H{last_data_row})")
        sheet.cell(row=total_row, column=9, value=f"=SUM(I{first_data_row}:I{last_data_row})")
        sheet.cell(row=total_row, column=10, value=f"=H{total_row}-G{total_row}")
        sheet.auto_filter.ref = f"A6:K{last_data_row}"
    for column in range(7, 11):
        for row_index in range(first_data_row, total_row + 1):
            sheet.cell(row=row_index, column=column).number_format = "#,##0;[Red]-#,##0;—"

    widths = (6, 20, 20, 42, 12, 18, 12, 12, 12, 16, 34)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[total_row].height = 24
    sheet.print_area = f"A1:K{total_row}"

    safe_number = re.sub(r"[^A-Za-z0-9_-]+", "_", str(order_id)).strip("_")[:48]
    filename = f"receiving_act_{safe_number or 'export'}.xlsx"
    return workbook, filename


def receiving_attention_from_payload(payload: dict | None) -> bool:
    """True when receiving act awaits client confirm/dispute (no DB)."""
    data = payload if isinstance(payload, dict) else {}
    if not data.get("act_sent"):
        return False
    response = str(data.get("act_client_response") or "").strip().lower()
    return response not in {"confirmed", "dispute"}


def request_attention(*, agency, order_type: str, order_id: str, payload: dict | None = None) -> bool:
    raw = str(order_type or "").strip().lower()
    data = payload or {}
    if raw == "receiving":
        # Fast path when caller already merged act_* across the audit trail.
        if any(key in data for key in ("act_sent", "act_client_response", "act_viewed")):
            return receiving_attention_from_payload(data)
        return _receiving_attention_for_order(agency, order_id)
    label = str(data.get("status_label") or "").lower()
    return "уточн" in label or bool(data.get("needs_client_attention"))


def add_other_comment(*, order_id: str, agency, user=None, text: str = "") -> bool:
    message = str(text or "").strip()
    if not message:
        return False
    latest = latest_other_entry(order_id)
    if not latest or latest.agency_id != getattr(agency, "id", None):
        return False
    payload = dict(latest.payload or {})
    log_order_action(
        "comment",
        order_id=str(order_id),
        order_type=OTHER_ORDER_TYPE,
        user=user if getattr(user, "is_authenticated", False) else None,
        agency=agency,
        description=message,
        payload=payload,
    )
    return True


def can_client_cancel_other(payload: dict | None) -> bool:
    status = str((payload or {}).get("status") or "").lower()
    return status in {"draft", "submitted"}


def _request_status_value(order_type: str, payload: dict | None) -> str:
    data = payload or {}
    raw_type = str(order_type or "").strip().lower()
    if raw_type == "shipping":
        if str(data.get("status") or "").strip().lower() == "warehouse_cancel_requested":
            return "warehouse_cancel_requested"
        return str(data.get("shipping_state") or data.get("status") or "").strip().lower()
    return str(data.get("status") or data.get("submit_action") or "").strip().lower()


def _cancel_request_pending(payload: dict | None) -> bool:
    data = payload or {}
    status = str(data.get("status") or data.get("submit_action") or "").strip().lower()
    label = str(data.get("status_label") or "").strip().lower()
    return status in {"cancel_requested", "warehouse_cancel_requested"} or "отмена на согласовании" in label or "отмена ожидает подтверждения склада" in label


def _change_request_pending(payload: dict | None) -> bool:
    data = payload or {}
    if data.get("change_requested_by_client") and not data.get("change_resolved_at"):
        return True
    status = str(data.get("status") or "").strip().lower()
    return status == "change_requested"


def _client_cancel_locked_after_warehouse(status_label: str, payload: dict | None = None) -> bool:
    """После приёмки складом / на подписи — отмена только через менеджера, без кнопки в ЛК."""
    data = payload or {}
    if (
        (data.get("act_storekeeper_signed") or data.get("act_logistician_signed"))
        and not data.get("act_manager_signed")
        and not data.get("act_sent")
    ):
        return True
    label = str(status_label or "").strip().lower()
    return any(
        token in label
        for token in (
            "ожидает подписи",
            "отправлен менеджеру",
            "принято складом",
            "принята складом",
        )
    )


def _can_client_cancel_from_bucket(*, bucket: str, status_label: str, payload: dict | None) -> bool:
    if _cancel_request_pending(payload):
        return False
    label = str(status_label or "").strip().lower()
    completed_label = "выполн" in label and "выполняется" not in label
    if "отмен" in label or completed_label or "заверш" in label or "закрыт" in label:
        return False
    if _client_cancel_locked_after_warehouse(status_label, payload):
        return False
    return bucket in {"client", "manager", "warehouse"}


def _requires_manager_cancel_approval(bucket: str, status_label: str = "") -> bool:
    if str(bucket or "").strip().lower() == "warehouse":
        return True
    return False


def _order_route_for_manager(order_type: str, order_id: str, *, agency=None) -> str:
    raw = str(order_type or "").strip().lower()
    oid = canonical_request_order_id(raw, order_id)
    if raw == "shipping":
        try:
            from shipping.models import ShippingOrder

            qs = ShippingOrder.objects.filter(number=oid)
            if agency is not None:
                qs = qs.filter(agency=agency)
            order = qs.only("id").first()
            if order:
                return f"/shipping/{order.pk}/"
        except Exception:
            pass
        return "/shipping/"
    if raw in {"processing", "packing"}:
        return f"/orders/processing/{oid}/"
    if raw == "receiving":
        return f"/orders/receiving/{oid}/"
    if raw == OTHER_ORDER_TYPE:
        from .other_requests import other_detail_url

        return other_detail_url(oid)
    return f"/orders/{raw}/{oid}/"


def _close_open_request_tasks(*, order_type: str, order_id: str, agency=None) -> int:
    from todo.models import Task

    route = _order_route_for_manager(order_type, order_id, agency=agency)
    return Task.objects.filter(route=route).exclude(status="done").update(status="done")


def _active_manager_for_user_id(user_id):
    if not user_id:
        return None
    from employees.models import Employee

    return (
        Employee.objects.filter(
            user_id=user_id,
            role="manager",
            is_active=True,
            user__is_active=True,
        )
        .select_related("user")
        .first()
    )


def _configured_client_manager(agency):
    """Return only an explicitly configured manager, not the global fallback."""
    raw_manager_id = getattr(agency, "mened_user_id", None) if agency is not None else None
    try:
        manager_id = int(raw_manager_id) if raw_manager_id not in (None, "", 0, "0") else None
    except (TypeError, ValueError):
        manager_id = None
    if not manager_id:
        return None
    manager = resolve_client_manager(agency)
    if manager and manager_id in {manager.pk, manager.user_id}:
        return manager
    return None


def _request_owner_manager(*, agency, order_type: str, order_id: str):
    """Resolve the manager who actually owns the request before global fallback."""
    configured = _configured_client_manager(agency)
    if configured:
        return configured

    normalized_type = str(order_type or "").strip().lower()
    normalized_order_id = canonical_request_order_id(normalized_type, order_id)
    if normalized_type == "shipping":
        try:
            from shipping.models import ShippingOrder

            orders = ShippingOrder.objects.filter(number=normalized_order_id)
            if agency is not None:
                orders = orders.filter(agency=agency)
            order = orders.only("created_by_id").first()
            manager = _active_manager_for_user_id(getattr(order, "created_by_id", None))
            if manager:
                return manager
        except Exception:
            pass

    from employees.models import Employee

    active_manager_user_ids = Employee.objects.filter(
        role="manager",
        is_active=True,
        user__is_active=True,
    ).values_list("user_id", flat=True)
    manager_entry = (
        OrderAuditEntry.objects.filter(
            order_type=normalized_type,
            order_id=normalized_order_id,
            user_id__in=active_manager_user_ids,
        )
        .order_by("-created_at", "-id")
        .only("user_id")
        .first()
    )
    manager = _active_manager_for_user_id(getattr(manager_entry, "user_id", None))
    return manager or resolve_client_manager(agency)


def _create_manager_appeal_task(
    *,
    agency,
    order_type: str,
    order_id: str,
    title: str,
    user=None,
    reason: str = "",
):
    from datetime import timedelta

    from todo.models import Task

    manager = _request_owner_manager(
        agency=agency,
        order_type=order_type,
        order_id=order_id,
    )
    if not manager:
        return None
    route = _order_route_for_manager(order_type, order_id, agency=agency)
    existing = Task.objects.filter(route=route, assigned_to=manager, title=title).exclude(status="done").first()
    client_name = getattr(agency, "agn_name", None) or getattr(agency, "inn", None) or getattr(agency, "id", "")
    description = f"Клиент: {client_name}"
    reason_text = str(reason or "").strip()
    if reason_text:
        description += f"\nПричина: {reason_text}"
    due_date = timezone.localtime() + timedelta(days=1)
    if existing:
        existing.description = description
        existing.due_date = due_date
        existing.save(update_fields=["description", "due_date", "updated_at"])
        return existing
    return Task.objects.create(
        title=title,
        description=description,
        route=route,
        assigned_to=manager,
        created_by=user if getattr(user, "is_authenticated", False) else None,
        due_date=due_date,
    )


def _create_cancel_approval_task(*, agency, order_type: str, order_id: str, user=None, reason: str = ""):
    number = format_order_number(order_type, order_id)
    return _create_manager_appeal_task(
        agency=agency,
        order_type=order_type,
        order_id=order_id,
        title=f"Согласуйте отмену заявки №{number}",
        user=user,
        reason=reason,
    )


def _create_change_approval_task(*, agency, order_type: str, order_id: str, user=None, reason: str = ""):
    number = format_order_number(order_type, order_id)
    return _create_manager_appeal_task(
        agency=agency,
        order_type=order_type,
        order_id=order_id,
        title=f"Запрос на изменение заявки №{number}",
        user=user,
        reason=reason,
    )


def _request_manager_change_approval(*, latest, agency, user=None, reason: str = "") -> bool:
    """Client appeal only — does not mutate warehouse ShippingOrder facts."""
    if _change_request_pending(latest.payload):
        return False
    payload = dict(latest.payload or {})
    payload["change_requested_by_client"] = True
    payload["change_requested_at"] = timezone.localtime().isoformat()
    payload["change_status"] = "new"
    if reason:
        payload["change_reason"] = reason
    # Keep warehouse status untouched; only mark appeal on audit payload.
    log_order_action(
        "change_request",
        order_id=str(latest.order_id),
        order_type=latest.order_type,
        user=user if getattr(user, "is_authenticated", False) else None,
        agency=agency,
        description=reason or "Клиент запросил изменение заявки",
        payload=payload,
    )
    _create_change_approval_task(
        agency=agency,
        order_type=latest.order_type,
        order_id=str(latest.order_id),
        user=user,
        reason=reason,
    )
    return True


def _can_client_request_change(*, bucket: str, status_label: str, payload: dict | None, order_type: str) -> bool:
    if str(order_type or "").strip().lower() not in {"shipping", "processing", "packing"}:
        return False
    if _change_request_pending(payload) or _cancel_request_pending(payload):
        return False
    label = str(status_label or "").strip().lower()
    if "отмен" in label or "закрыт" in label or "выполн" in label or "черновик" in label:
        return False
    return bucket in {"client", "manager", "warehouse"}


def _log_generic_client_cancel(*, latest, agency, user=None, reason: str = "") -> bool:
    payload = dict(latest.payload or {})
    payload["status"] = "cancelled"
    payload["status_label"] = "Отменена"
    payload["submit_action"] = "cancel"
    payload["cancelled_by_client"] = True
    if reason:
        payload["cancel_reason"] = reason
    log_order_action(
        "status",
        order_id=str(latest.order_id),
        order_type=latest.order_type,
        user=user if getattr(user, "is_authenticated", False) else None,
        agency=agency,
        description=reason or "Заявка отменена клиентом",
        payload=payload,
    )
    _close_open_request_tasks(order_type=latest.order_type, order_id=str(latest.order_id), agency=agency)
    return True


def _request_manager_cancel_approval(*, latest, agency, user=None, reason: str = "") -> bool:
    if _cancel_request_pending(latest.payload):
        return False
    payload = dict(latest.payload or {})
    payload["status"] = "cancel_requested"
    payload["status_label"] = "Отмена на согласовании менеджера"
    payload["cancel_requested_by_client"] = True
    payload["cancel_requested_at"] = timezone.localtime().isoformat()
    if reason:
        payload["cancel_reason"] = reason
    log_order_action(
        "cancel_request",
        order_id=str(latest.order_id),
        order_type=latest.order_type,
        user=user if getattr(user, "is_authenticated", False) else None,
        agency=agency,
        description=reason or "Клиент запросил отмену заявки",
        payload=payload,
    )
    _create_cancel_approval_task(
        agency=agency,
        order_type=latest.order_type,
        order_id=str(latest.order_id),
        user=user,
        reason=reason,
    )
    return True


def can_client_comment_other(payload: dict | None) -> bool:
    status = str((payload or {}).get("status") or "").lower()
    return status not in {"completed", "done", "cancelled"}


def _as_int(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _text(value) -> str:
    return str(value or "").strip()


_OZON_GM_BLOCK_RE = re.compile(r"ШК\s*ГМ(?:\s*Ozon)?\s*:\s*([^\n]+)", re.IGNORECASE)
_OZON_GM_LINE_RE = re.compile(r"(?:^|\n)\s*-?\s*ГМ\s+([^\s,;→]+)", re.IGNORECASE)


def _append_unique(target: list[str], value: Any) -> None:
    text = _text(value).strip(" .;,\u00a0")
    if not text or text == "—" or text in target:
        return
    target.append(text)


def _extract_ozon_gm_codes(*values: Any) -> list[str]:
    """Read Ozon cargo-place barcodes from persisted comments/payloads."""

    result: list[str] = []

    def walk(value: Any) -> None:
        if value in (None, ""):
            return
        if isinstance(value, (list, tuple, set)):
            for item in value:
                walk(item)
            return
        if isinstance(value, dict):
            for key in ("gm_barcode", "ozon_gm_comment", "comment", "gm_barcodes", "ozon_gm_barcodes"):
                if key in value:
                    walk(value.get(key))
            return
        text = _text(value)
        if not text:
            return
        for match in _OZON_GM_BLOCK_RE.finditer(text):
            chunk = match.group(1)
            for token in re.split(r"[,;]", chunk):
                _append_unique(result, token)
        for match in _OZON_GM_LINE_RE.finditer(text):
            _append_unique(result, match.group(1))

    for value in values:
        walk(value)
    return result


def _join_codes(codes: list[str], *, limit: int = 80) -> str:
    visible = list(codes[:limit])
    suffix = f" +{len(codes) - limit}" if len(codes) > limit else ""
    return ", ".join(visible) + suffix


def _shipping_order_gm_codes(order) -> list[str]:
    values: list[Any] = [getattr(order, "comment", "")]
    try:
        values.extend(dest.comment for dest in order.destinations.all())
    except Exception:
        pass
    try:
        values.extend(item.comment for item in order.items.all())
    except Exception:
        pass
    return _extract_ozon_gm_codes(*values)


def _shipping_destination_summary(order) -> str:
    try:
        destinations = list(order.destinations.all())
    except Exception:
        destinations = []
    if not destinations:
        return ""
    parts: list[str] = []
    for dest in destinations[:8]:
        title = _text(getattr(dest, "warehouse_name", "")) or "Склад Ozon"
        boxes = _as_int(getattr(dest, "planned_boxes", 0))
        units = _as_int(getattr(dest, "planned_units", 0))
        codes = _extract_ozon_gm_codes(getattr(dest, "comment", ""))
        tail = f" — {boxes} кор. / {units} шт."
        if codes:
            tail += f" / ШК ГМ: {_join_codes(codes, limit=12)}"
        parts.append(title + tail)
    if len(destinations) > 8:
        parts.append(f"+{len(destinations) - 8} складов")
    return "; ".join(parts)


def _model_relation_exists(model, name: str) -> bool:
    try:
        model._meta.get_field(name)
        return True
    except Exception:
        return False


def _prefetches_if_present(model, *names: str) -> list[str]:
    return [name for name in names if _model_relation_exists(model, name)]


def _line_key(sku: str, size: str = "", goods_type: str = "") -> str:
    return f"{_text(sku).upper()}|{_text(size).upper()}|{_text(goods_type).lower()}"


def _receiving_goods_type_label(goods_type: str) -> str:
    from orders.services import ReceivingWorkflowService

    code = _text(goods_type).lower()
    return ReceivingWorkflowService.RECEIVING_GOODS_TYPE_LABELS.get(code, code)


def _receiving_detail_from_entries(entries: list) -> dict[str, Any]:
    """Build LK receiving composition from placement boxes (same truth as warehouse act)."""
    from orders.services import ReceivingWorkflowService

    planned_by_item: dict[str, dict[str, Any]] = {}
    planned_order: list[str] = []
    fact_rows: dict[str, dict[str, Any]] = {}
    fact_order: list[str] = []
    meta_src: dict[str, Any] = {}
    placement_payload: dict[str, Any] = {}
    act_payload: dict[str, Any] = {}

    for entry in entries:
        payload = entry.payload or {}
        if not isinstance(payload, dict):
            continue
        if not meta_src:
            meta_src = payload
        items = payload.get("items")
        if isinstance(items, list):
            for item in items:
                if not isinstance(item, dict):
                    continue
                sku = _text(item.get("sku_code") or item.get("sku"))
                size = _text(item.get("size"))
                if not sku:
                    continue
                item_key = _line_key(sku, size)
                qty = _as_int(item.get("qty") or item.get("planned_qty") or item.get("qty_planned"))
                if item_key not in planned_by_item:
                    planned_by_item[item_key] = {
                        "sku_code": sku,
                        "name": _text(item.get("name")),
                        "size": size,
                        "barcode": _text(item.get("barcode")),
                        "qty_planned": qty,
                    }
                    planned_order.append(item_key)
                else:
                    row = planned_by_item[item_key]
                    if qty:
                        row["qty_planned"] = qty
                    if item.get("name") and not row["name"]:
                        row["name"] = _text(item.get("name"))
                    if item.get("barcode") and not row["barcode"]:
                        row["barcode"] = _text(item.get("barcode"))

        if payload.get("act") == "placement" and (
            payload.get("act_boxes") or payload.get("act_items") or payload.get("act_pallets")
        ):
            placement_payload = payload
            meta_src = {**meta_src, **payload}
        if payload.get("act") == "receiving" and payload.get("act_items"):
            act_payload = payload
            meta_src = {**meta_src, **payload}

    # Prefer document placement selected the same way as warehouse act print.
    document_placement = ReceivingWorkflowService._find_document_act_entry(
        entries, "placement", "акт размещения"
    )
    if document_placement and isinstance(document_placement.payload, dict):
        placement_payload = document_placement.payload
        meta_src = {**meta_src, **placement_payload}

    def _add_fact(raw_item: dict[str, Any], *, goods_type_hint: str = "", count_box: bool = False):
        if not isinstance(raw_item, dict):
            return
        sku = _text(raw_item.get("sku_code") or raw_item.get("sku"))
        if not sku:
            return
        size = _text(raw_item.get("size"))
        goods_type = ReceivingWorkflowService._normalize_flow_goods_type(
            raw_item.get("goods_type"),
            goods_type_hint,
        )
        qty = _as_int(raw_item.get("qty") or raw_item.get("actual_qty") or raw_item.get("qty_actual"))
        if qty <= 0 and not count_box:
            return
        key = _line_key(sku, size, goods_type)
        row = fact_rows.get(key)
        if row is None:
            row = {
                "sku_code": sku,
                "name": _text(raw_item.get("name")),
                "size": size,
                "barcode": _text(raw_item.get("barcode")),
                "goods_type": goods_type,
                "qty_actual": 0,
                "box_qty": 0,
                "comment": _text(raw_item.get("comment")),
            }
            fact_rows[key] = row
            fact_order.append(key)
        if qty > 0:
            row["qty_actual"] = _as_int(row["qty_actual"]) + qty
        if count_box:
            row["box_qty"] = _as_int(row["box_qty"]) + 1
        if raw_item.get("name") and not row["name"]:
            row["name"] = _text(raw_item.get("name"))
        if raw_item.get("barcode") and not row["barcode"]:
            row["barcode"] = _text(raw_item.get("barcode"))

    boxes = placement_payload.get("act_boxes") if isinstance(placement_payload, dict) else None
    if isinstance(boxes, list) and boxes:
        for box in boxes:
            if not isinstance(box, dict):
                continue
            box_gt = ReceivingWorkflowService._normalize_flow_goods_type(box.get("goods_type"))
            items = [item for item in (box.get("items") or []) if isinstance(item, dict)]
            counted_keys: set[str] = set()
            for item in items:
                item_gt = ReceivingWorkflowService._normalize_flow_goods_type(item.get("goods_type"), box_gt)
                sku = _text(item.get("sku_code") or item.get("sku"))
                size = _text(item.get("size"))
                key = _line_key(sku, size, item_gt)
                _add_fact(item, goods_type_hint=box_gt, count_box=False)
                if key in fact_rows and key not in counted_keys:
                    fact_rows[key]["box_qty"] = _as_int(fact_rows[key]["box_qty"]) + 1
                    counted_keys.add(key)
        for pallet in placement_payload.get("act_pallets") or []:
            if not isinstance(pallet, dict):
                continue
            for item in pallet.get("items") or []:
                _add_fact(item, goods_type_hint=ReceivingWorkflowService._normalize_flow_goods_type(pallet.get("goods_type")))

    if not fact_rows:
        # Fallback: flat act_items from receiving/placement when boxes are absent.
        raw_act = []
        if isinstance(act_payload.get("act_items"), list):
            raw_act = act_payload.get("act_items") or []
        elif isinstance(placement_payload.get("act_items"), list):
            raw_act = placement_payload.get("act_items") or []
        for item in raw_act:
            if not isinstance(item, dict):
                continue
            sku = _text(item.get("sku_code") or item.get("sku"))
            size = _text(item.get("size"))
            goods_type = ReceivingWorkflowService._normalize_flow_goods_type(item.get("goods_type"))
            key = _line_key(sku, size, goods_type)
            actual = item.get("actual_qty")
            if actual is None:
                actual = item.get("qty_actual")
            if actual is None:
                actual = item.get("qty")
            fact_rows[key] = {
                "sku_code": sku,
                "name": _text(item.get("name")),
                "size": size,
                "barcode": _text(item.get("barcode")),
                "goods_type": goods_type,
                "qty_actual": _as_int(actual) if actual is not None else None,
                "box_qty": _as_int(item.get("box_qty")) if item.get("box_qty") not in (None, "") else None,
                "comment": _text(item.get("comment")),
            }
            fact_order.append(key)
            item_key = _line_key(sku, size)
            if item_key not in planned_by_item:
                planned = _as_int(item.get("planned_qty") or item.get("qty_planned") or item.get("qty"))
                planned_by_item[item_key] = {
                    "sku_code": sku,
                    "name": _text(item.get("name")),
                    "size": size,
                    "barcode": _text(item.get("barcode")),
                    "qty_planned": planned,
                }
                planned_order.append(item_key)

    # Assign plan to the goods-type row with the largest fact for that SKU/size
    # (matches printed act: plan on Готовый, брак gets plan 0).
    planned_assigned: dict[str, str] = {}
    for item_key, planned in planned_by_item.items():
        candidates = [
            key
            for key in fact_order
            if key.startswith(f"{_text(planned['sku_code']).upper()}|{_text(planned['size']).upper()}|")
        ]
        if not candidates:
            # No warehouse fact yet — show plan-only line.
            empty_key = _line_key(planned["sku_code"], planned["size"], "")
            fact_rows[empty_key] = {
                "sku_code": planned["sku_code"],
                "name": planned["name"],
                "size": planned["size"],
                "barcode": planned["barcode"],
                "goods_type": "",
                "qty_actual": None,
                "box_qty": None,
                "comment": "",
            }
            fact_order.append(empty_key)
            planned_assigned[empty_key] = item_key
            continue
        best_key = max(candidates, key=lambda k: _as_int(fact_rows[k].get("qty_actual")))
        planned_assigned[best_key] = item_key

    lines: list[dict[str, Any]] = []
    has_mismatch = False
    total_planned = 0
    total_actual = 0
    total_boxes = 0
    for key in fact_order:
        row = fact_rows[key]
        planned_key = planned_assigned.get(key)
        qty_planned = _as_int(planned_by_item[planned_key]["qty_planned"]) if planned_key else 0
        qty_actual = row["qty_actual"]
        box_qty = row.get("box_qty")
        display_name = _text(row.get("name")) or "-"
        if row.get("goods_type"):
            label = _receiving_goods_type_label(row["goods_type"])
            if label and f"({label})" not in display_name:
                display_name = f"{display_name} ({label})"
        mismatch = qty_actual is not None and qty_planned != _as_int(qty_actual)
        if mismatch:
            has_mismatch = True
        if qty_planned:
            total_planned += qty_planned
        if qty_actual is not None:
            total_actual += _as_int(qty_actual)
        if box_qty not in (None, ""):
            total_boxes += _as_int(box_qty)
        lines.append(
            {
                "sku_code": row["sku_code"],
                "name": display_name,
                "size": row["size"],
                "barcode": row.get("barcode") or "",
                "goods_type": row.get("goods_type") or "",
                "goods_type_label": _receiving_goods_type_label(row["goods_type"]) if row.get("goods_type") else "",
                "qty_planned": qty_planned,
                "qty_actual": qty_actual,
                "box_qty": box_qty if box_qty not in (None, "") else "—",
                "delta_qty": (_as_int(qty_actual) - qty_planned) if qty_actual is not None else None,
                "comment": row.get("comment") or "",
                "mismatch": mismatch,
            }
        )

    from .client_receiving_view import format_client_datetime, localize_place_type

    show_fact = any(row.get("qty_actual") is not None for row in lines)
    meta = []
    for label, key in (
        ("Ожидаемая дата", "eta_at"),
        ("Ожидаемые короба", "expected_boxes"),
        ("Тип места", "place_type"),
        ("Номер ТС", "vehicle_number"),
        ("Телефон водителя", "driver_phone"),
        ("Комментарий", "comment"),
    ):
        value = meta_src.get(key)
        if value in (None, ""):
            continue
        if key == "eta_at":
            value = format_client_datetime(value)
        elif key == "place_type":
            value = localize_place_type(value)
        meta.append({"label": label, "value": _text(value)})
    if show_fact:
        meta.append({"label": "Итого план", "value": str(total_planned)})
        meta.append({"label": "Итого факт", "value": str(total_actual)})
        if total_boxes:
            meta.append({"label": "Итого коробов", "value": str(total_boxes)})
        meta.append({"label": "Итого расхождение", "value": str(total_actual - total_planned)})

    columns = [
        {"key": "barcode", "label": "ШК"},
        {"key": "sku_code", "label": "Артикул"},
        {"key": "name", "label": "Наименование"},
        {"key": "size", "label": "Размер"},
        {"key": "qty_planned", "label": "План"},
    ]
    if show_fact:
        columns.extend(
            [
                {"key": "qty_actual", "label": "Факт"},
                {"key": "box_qty", "label": "Короба"},
                {"key": "delta_qty", "label": "Расхождение"},
            ]
        )

    return {
        "meta": meta,
        "lines": lines,
        "columns": columns,
        "discrepancies": [
            {
                "sku_code": row["sku_code"],
                "name": row["name"],
                "size": row["size"],
                "expected_qty": row["qty_planned"],
                "factual_qty": row["qty_actual"],
                "delta_qty": row["delta_qty"],
            }
            for row in lines
            if row.get("mismatch")
        ],
        "has_mismatch": has_mismatch,
        "show_fact_columns": show_fact,
        "lines_title": "Состав поставки",
        "empty_lines_hint": "Состав заявки пока не заполнен.",
    }


def _shipping_detail_from_db(*, agency, order_id: str, fallback_payload: dict | None = None) -> dict[str, Any]:
    from shipping.models import ShippingOrder

    from . import web_ui as client_web_ui
    from .client_shipping_view import build_shipping_client_view

    lookup_order_id = canonical_shipping_order_id(order_id)
    order = (
        ShippingOrder.objects.filter(agency=agency, number=lookup_order_id)
        .select_related("marketplace")
        .prefetch_related(*_prefetches_if_present(ShippingOrder, "items", "destinations"))
        .first()
    )
    meta = []
    lines: list[dict[str, Any]] = []
    shipping_view = None
    ozon_gm_structure: dict[str, Any] | None = None
    if order:
        trip_status = ""
        try:
            trip_status = client_web_ui._shipping_trip_status(
                getattr(order, "number", None) or lookup_order_id,
                agency_id=getattr(agency, "id", None),
            )
        except Exception:
            trip_status = ""
        shipping_view = build_shipping_client_view(agency=agency, order=order, trip_status=trip_status)
    if order:
        try:
            from shipping.services import _ozon_api_summary_for_form

            ozon_summary = _ozon_api_summary_for_form(order)
            if ozon_summary.get("show"):
                ozon_gm_structure = {
                    "supply_number": ozon_summary.get("supply_number") or "",
                    "shipping_barcode": ozon_summary.get("shipping_barcode") or "",
                    "boxes_total": ozon_summary.get("boxes_total") or "",
                    "cargoes": list(ozon_summary.get("gm_cargoes") or []),
                    "has_gm_codes": bool(ozon_summary.get("gm_barcodes")),
                }
        except Exception:
            ozon_gm_structure = None
        delivery_type = getattr(order, "delivery_type", "")
        supply_type_value = getattr(order, "supply_type", "")
        place_type_value = getattr(order, "place_type", "")
        delivery = order.get_delivery_type_display() if hasattr(order, "get_delivery_type_display") else delivery_type
        supply_type = order.get_supply_type_display() if hasattr(order, "get_supply_type_display") else supply_type_value
        place_type = order.get_place_type_display() if hasattr(order, "get_place_type_display") else place_type_value
        distribution_status = (
            order.get_distribution_status_display()
            if getattr(order, "distribution_status", "") and hasattr(order, "get_distribution_status_display")
            else ""
        )
        gm_codes = _shipping_order_gm_codes(order)
        destination_summary = _shipping_destination_summary(order)
        supply_number = getattr(order, "supply_number", "")
        wb_supply_barcode = getattr(order, "wb_supply_barcode", "")
        shipping_barcode = getattr(order, "shipping_barcode", "")
        transfer_destination = (
            getattr(order, "destination_address", "")
            if delivery_type == ShippingOrder.DELIVERY_TRANSFER
            else ""
        )
        for label, value in (
            ("Маркетплейс", order.marketplace.name if getattr(order, "marketplace_id", None) else ""),
            ("Тип доставки", delivery),
            ("Назначение перемещения", transfer_destination),
            ("Слот", order.slot_date.isoformat() if getattr(order, "slot_date", None) else ""),
            ("Время слота", order.slot_time.strftime("%H:%M") if getattr(order, "slot_time", None) else ""),
            ("Транзитный склад", getattr(order, "transit_address", "")),
            ("Склад назначения", getattr(order, "destination_warehouse", "")),
            ("Конечные склады Ozon", destination_summary),
            ("План. отгрузка", order.planned_ship_date.isoformat() if getattr(order, "planned_ship_date", None) else ""),
            ("Короба", getattr(order, "expected_boxes", "")),
            ("Тип поставки", supply_type),
            ("Тип мест", place_type),
            ("Номер поставки Ozon", supply_number),
            ("Номер ТС", getattr(order, "vehicle_number", "")),
            ("Телефон водителя", getattr(order, "driver_phone", "")),
            ("ШК поставки", wb_supply_barcode or shipping_barcode or supply_number),
            ("ШК ГМ Ozon", _join_codes(gm_codes) if gm_codes else ""),
            ("Статус распределения Ozon", distribution_status),
            ("Комментарий", getattr(order, "comment", "")),
        ):
            if value not in (None, ""):
                meta.append({"label": label, "value": _text(value)})
        try:
            order_items = order.items.all()
        except Exception:
            order_items = []
        if shipping_view and shipping_view.get("lines"):
            lines = list(shipping_view["lines"])
            for row, item in zip(lines, order_items):
                item_gm_codes = _extract_ozon_gm_codes(item.comment)
                if item_gm_codes:
                    row["gm_barcodes"] = _join_codes(item_gm_codes)
        else:
            for item in order_items:
                item_gm_codes = _extract_ozon_gm_codes(item.comment)
                qty_requested = _as_int(item.qty_requested)
                qty_shipped = _as_int(item.qty_shipped)
                lines.append(
                    {
                        "sku_code": _text(item.sku_code),
                        "name": _text(item.name),
                        "size": _text(item.size),
                        "barcode": _text(item.barcode),
                        "qty_requested": qty_requested,
                        "qty_reserved": _as_int(item.qty_reserved),
                        "qty_shipped": qty_shipped,
                        "delta_qty": qty_requested - qty_shipped,
                        "gm_barcodes": _join_codes(item_gm_codes) if item_gm_codes else "",
                        "comment": _text(item.comment),
                    }
                )
    else:
        payload = fallback_payload or {}
        if (
            str(payload.get("delivery_type") or "").strip().lower() == ShippingOrder.DELIVERY_TRANSFER
            and payload.get("destination_address") not in (None, "")
        ):
            meta.append(
                {
                    "label": "Назначение перемещения",
                    "value": _text(payload.get("destination_address")),
                }
            )
        payload_gm_codes = _extract_ozon_gm_codes(
            payload.get("comment"),
            payload.get("ozon_gm_comment"),
            payload.get("gm_barcodes"),
            payload.get("items"),
        )
        supply_type_labels = {
            "box": "Короб",
            "monopallet": "Монопаллет",
            "supersafe": "Супер сейф",
        }
        for label, key in (
            ("Маркетплейс", "marketplace"),
            ("Тип доставки", "delivery_type"),
            ("Слот", "slot_date"),
            ("Время слота", "slot_time"),
            ("Транзитный склад", "transit_address"),
            ("Склад назначения", "destination_warehouse"),
            ("Тип поставки", "supply_type_label"),
            ("Номер поставки Ozon", "supply_number"),
            ("ШК поставки", "shipping_barcode"),
            ("Комментарий", "comment"),
        ):
            value = payload.get(key)
            if key == "supply_type_label" and value in (None, ""):
                raw_supply_type = str(payload.get("supply_type") or "").strip()
                value = supply_type_labels.get(raw_supply_type, raw_supply_type)
            if value not in (None, ""):
                meta.append({"label": label, "value": _text(value)})
        if payload_gm_codes:
            meta.append({"label": "ШК ГМ Ozon", "value": _join_codes(payload_gm_codes)})
        for item in payload.get("items") or []:
            if not isinstance(item, dict):
                continue
            item_gm_codes = _extract_ozon_gm_codes(item)
            qty_requested = _as_int(item.get("qty_requested") or item.get("qty"))
            qty_shipped = _as_int(item.get("qty_shipped"))
            lines.append(
                {
                    "sku_code": _text(item.get("sku_code")),
                    "name": _text(item.get("name")),
                    "size": _text(item.get("size")),
                    "barcode": _text(item.get("barcode")),
                    "qty_requested": qty_requested,
                    "qty_reserved": _as_int(item.get("qty_reserved")),
                    "qty_shipped": qty_shipped,
                    "delta_qty": qty_requested - qty_shipped,
                    "gm_barcodes": _join_codes(item_gm_codes) if item_gm_codes else "",
                    "comment": _text(item.get("comment")),
                }
            )

    if shipping_view and shipping_view.get("lines"):
        columns = [
            {"key": "sku_code", "label": "Артикул"},
            {"key": "name", "label": "Наименование"},
            {"key": "size", "label": "Размер"},
            {"key": "barcode", "label": "ШК товара"},
            {"key": "qty_requested", "label": "Заявлено"},
            {"key": "qty_shipped", "label": "Отгружено"},
            {"key": "delta_qty", "label": "Расхождение"},
            {"key": "line_status", "label": "Статус позиции"},
        ]
    else:
        columns = [
            {"key": "sku_code", "label": "Артикул"},
            {"key": "name", "label": "Наименование"},
            {"key": "size", "label": "Размер"},
            {"key": "barcode", "label": "ШК товара"},
            {"key": "qty_requested", "label": "Запрошено"},
            {"key": "qty_shipped", "label": "Отгружено"},
        ]
    if any(row.get("gm_barcodes") for row in lines):
        columns.append({"key": "gm_barcodes", "label": "ШК ГМ Ozon"})

    discrepancies = list((shipping_view or {}).get("discrepancies") or [])
    return {
        "meta": meta,
        "lines": lines,
        "columns": columns,
        "discrepancies": discrepancies,
        "has_mismatch": bool((shipping_view or {}).get("has_mismatch") or discrepancies),
        "lines_title": "Состав отгрузки",
        "empty_lines_hint": "Позиции отгрузки появятся после согласования заявки.",
        "shipping_view": shipping_view,
        "ozon_gm_structure": ozon_gm_structure,
    }


def _processing_detail_from_entries(entries: list) -> dict[str, Any]:
    work: dict[str, Any] = {}
    for entry in reversed(entries):
        payload = entry.payload or {}
        if any(payload.get(key) for key in ("stock_rows", "size_rows", "product_name", "article", "discrepancy_items")):
            work = payload
            break
    if not work and entries:
        work = entries[-1].payload or {}

    meta = []
    for label, key in (
        ("Товар", "product_name"),
        ("Артикул", "article"),
        ("Бренд", "brand"),
        ("Маркетплейс", "marketplace"),
        ("Комментарий", "comments"),
        ("Статус расхождений", "discrepancy_status"),
    ):
        value = work.get(key)
        if value not in (None, ""):
            meta.append({"label": label, "value": _text(value)})

    lines: list[dict[str, Any]] = []
    raw_rows = work.get("stock_rows") if isinstance(work.get("stock_rows"), list) else None
    if not raw_rows:
        raw_rows = work.get("size_rows") if isinstance(work.get("size_rows"), list) else []
    for item in raw_rows or []:
        if not isinstance(item, dict):
            continue
        size_value = _text(item.get("size") or item.get("size_value"))
        qty_value = _as_int(item.get("qty") or item.get("recount_qty") or item.get("factual_qty"))
        lines.append(
            {
                "sku_code": _text(item.get("article") or item.get("sku_code") or item.get("sku") or item.get("size_no")),
                "name": _text(item.get("name") or item.get("product_name") or work.get("product_name")),
                "size": size_value if size_value and size_value not in {"0", "0.0"} else "Не указан",
                "barcode": _text(item.get("barcode")),
                "qty": f"{qty_value} шт." if qty_value else "—",
                "service": _text(item.get("service") or item.get("processing_service")),
            }
        )

    discrepancies = []
    for item in work.get("discrepancy_items") or []:
        if not isinstance(item, dict):
            continue
        discrepancies.append(
            {
                "sku_code": _text(item.get("sku_code") or item.get("article")),
                "name": _text(item.get("name")),
                "size": _text(item.get("size")),
                "expected_qty": _as_int(item.get("expected_qty")),
                "factual_qty": _as_int(item.get("factual_qty")),
                "delta_qty": _as_int(item.get("delta_qty")),
            }
        )

    return {
        "meta": meta,
        "lines": lines,
        "columns": [
            {"key": "sku_code", "label": "Артикул"},
            {"key": "name", "label": "Наименование"},
            {"key": "size", "label": "Размер"},
            {"key": "barcode", "label": "Штрихкод"},
            {"key": "qty", "label": "Количество"},
            {"key": "service", "label": "Услуга"},
        ],
        "discrepancies": discrepancies,
        "has_mismatch": bool(discrepancies or work.get("discrepancy_detected")),
        "lines_title": "Состав обработки",
        "empty_lines_hint": "Позиции появятся после запуска обработки на складе.",
    }


def _processing_client_act(payload: dict | None) -> dict[str, Any] | None:
    source = payload if isinstance(payload, dict) else {}
    processing_act = source.get("processing_act")
    processing_act = dict(processing_act) if isinstance(processing_act, dict) else {}
    if int(processing_act.get("version") or 0) < 1:
        return None
    act_status = str(processing_act.get("status") or "").strip().lower()
    if act_status not in {"awaiting_confirmations", "confirmed", "dispute"}:
        return None
    manager_response = str(processing_act.get("manager_response") or "pending").strip().lower()
    manager_status_label = (
        "Подтвержден менеджером"
        if manager_response == "confirmed"
        else "Ожидает проверки менеджером"
    )
    items = []
    for item in processing_act.get("items") or []:
        if not isinstance(item, dict):
            continue
        items.append(
            {
                "sku_code": _text(item.get("sku_code") or item.get("article")),
                "name": _text(item.get("name")),
                "size": _text(item.get("size")) or "Не указан",
                "barcode": _text(item.get("barcode")),
                "actual_qty": _as_int(item.get("actual_qty")),
                "box_qty": _as_int(item.get("box_qty")),
                "pallet_qty": _as_int(item.get("pallet_qty")),
                "condition_label": _text(item.get("condition_label")),
            }
        )
    services = []
    for service in processing_act.get("services") or []:
        if not isinstance(service, dict):
            continue
        services.append(
            {
                "operation_label": _text(service.get("operation_label")),
                "performer_name": _text(service.get("performer_name")),
                "planned_qty": _as_int(service.get("planned_qty")),
                "actual_qty": _as_int(service.get("actual_qty")),
                "comment": _text(service.get("comment")),
            }
        )
    declared_qty = _as_int(processing_act.get("declared_qty"))
    processed_qty = _as_int(processing_act.get("processed_qty"))
    boxed_qty = _as_int(processing_act.get("boxed_qty"))
    return {
        "version": int(processing_act.get("version") or 0),
        "status": act_status,
        "status_label": manager_status_label,
        "manager_response": manager_response,
        "declared_qty": declared_qty,
        "processed_qty": processed_qty,
        "boxed_qty": boxed_qty,
        "boxes_count": _as_int(processing_act.get("boxes_count")),
        "pallets_count": _as_int(processing_act.get("pallets_count")),
        "has_mismatch": bool(
            processing_act.get("discrepancy_items")
            or declared_qty != processed_qty
            or processed_qty != boxed_qty
        ),
        "confirmed_by": _text(processing_act.get("head_confirmed_by")),
        "confirmed_at": _text(processing_act.get("head_confirmed_at")),
        "items": items,
        "services": services,
    }


def _processing_client_stage(payload: dict | None) -> str:
    try:
        from processing_app.stages import (
            PROCESSING_STAGE_AWAITING_APPROVAL,
            PROCESSING_STAGE_CANCELLED,
            PROCESSING_STAGE_DONE,
            PROCESSING_STAGE_MANAGER_APPROVED,
            PROCESSING_STAGE_OBR_ARRIVED,
            PROCESSING_STAGE_OBR_MOVE_CREATED,
            PROCESSING_STAGE_QUALITY_APPROVED,
            PROCESSING_STAGE_QUALITY_CONTROL,
            PROCESSING_STAGE_RETURN_TO_STOCK_SENT,
            PROCESSING_STAGE_RETURNED_TO_STOCK,
            PROCESSING_STAGE_TAKEN_INTO_PROCESSING,
            PROCESSING_STAGE_UNBOXING_COMPLETED,
            PROCESSING_STAGE_UNBOXING_OPENED,
            processing_is_cancelled,
            processing_is_done,
            processing_stage_from_payload,
        )
    except Exception:
        return "manager_review"
    if processing_is_cancelled(payload):
        return PROCESSING_STAGE_CANCELLED
    if processing_is_done(payload):
        return PROCESSING_STAGE_DONE
    if _processing_client_act(payload):
        return "result_review"
    stage = processing_stage_from_payload(payload)
    if stage in {
        PROCESSING_STAGE_OBR_MOVE_CREATED,
        PROCESSING_STAGE_OBR_ARRIVED,
    }:
        return "handoff"
    if stage in {
        PROCESSING_STAGE_TAKEN_INTO_PROCESSING,
        PROCESSING_STAGE_UNBOXING_OPENED,
        PROCESSING_STAGE_UNBOXING_COMPLETED,
        PROCESSING_STAGE_QUALITY_CONTROL,
        PROCESSING_STAGE_QUALITY_APPROVED,
        PROCESSING_STAGE_RETURN_TO_STOCK_SENT,
        PROCESSING_STAGE_RETURNED_TO_STOCK,
    }:
        return "in_work"
    if stage == PROCESSING_STAGE_MANAGER_APPROVED:
        return "approved"
    if stage == PROCESSING_STAGE_AWAITING_APPROVAL:
        return "manager_review"
    if stage == PROCESSING_STAGE_DONE:
        return PROCESSING_STAGE_DONE
    return "manager_review"


def _processing_client_view(*, agency, order_id: str, entries: list, payload: dict, body: dict, responsibles: list | None = None) -> dict[str, Any]:
    stage = _processing_client_stage(payload)
    processing_act = _processing_client_act(payload)
    status_map = {
        "manager_review": ("Ожидает проверки менеджером", "Менеджер проверит состав обработки и подтвердит заявку."),
        "approved": ("Подтверждена", "Менеджер подтвердил заявку. Дальше заявка будет передана в обработку."),
        "handoff": ("Передана в обработку", "Заявка передана руководителю обработки."),
        "in_work": ("Обработка выполняется", "Внутренние операции обработки объединены в этот клиентский этап."),
        "result_review": (
            "Обработка подтверждена, ожидает проверки менеджера",
            "Руководитель обработки подтвердил результат. Менеджер проверяет акт обработки.",
        ),
        "done": ("Завершена", "Работы по обработке завершены."),
        "cancelled": ("Отменена", "Заявка отменена и сохранена в истории."),
    }
    status_label, status_hint = status_map.get(stage, status_map["manager_review"])
    if stage == "result_review" and processing_act and processing_act.get("manager_response") == "confirmed":
        status_label = "Результат обработки подтвержден менеджером"
        status_hint = "Акт подтвержден. Товар возвращается из зоны обработки на склад."
    steps_def = [
        ("submitted", "Заявка отправлена"),
        ("manager_review", "Проверка менеджером"),
        ("approved", "Подтверждена"),
        ("handoff", "Передана в обработку"),
        ("in_work", "Обработка выполняется"),
        ("result_review", "Проверка результата"),
        ("done", "Завершена"),
    ]
    rank = {key: index for index, (key, _label) in enumerate(steps_def)}
    current_rank = rank.get(stage, 1)
    is_cancelled = stage == "cancelled"
    created_at = getattr(entries[0], "created_at", None) if entries else None
    steps = []
    for key, label in steps_def:
        state = "future"
        if not is_cancelled:
            if key == stage:
                state = "current"
            elif rank[key] < current_rank:
                state = "done"
        if key == "submitted" and stage == "manager_review":
            state = "done"
        steps.append(
            {
                "key": key,
                "label": label,
                "state": state,
                "date": created_at.isoformat() if key == "submitted" and created_at else "",
            }
        )
    lines = list(body.get("lines") or [])
    sku_count = len({str(row.get("sku_code") or "").strip().lower() for row in lines if row.get("sku_code")})
    qty_total = 0
    for row in lines:
        match = re.search(r"\d+", str(row.get("qty") or ""))
        if match:
            qty_total += int(match.group(0))
    general = [
        {"label": "Тип заявки", "value": "Обработка"},
        {"label": "Маркетплейс", "value": _text(payload.get("marketplace"))},
        {"label": "Количество SKU", "value": str(sku_count) if sku_count else ""},
        {"label": "Общее количество товара", "value": f"{qty_total} шт." if qty_total else ""},
        *(responsibles or []),
        {"label": "Текущий этап", "value": status_label},
    ]
    tech = [
        {"label": "Товар", "value": _text(payload.get("product_name"))},
        {"label": "Артикул", "value": _text(payload.get("article"))},
        {"label": "Комментарий клиента", "value": _text(payload.get("comments"))},
    ]
    return {
        "enabled": True,
        "status_key": stage,
        "status_label": status_label,
        "status_hint": status_hint,
        "show_success": stage == "manager_review",
        "steps": steps,
        "general": [row for row in general if row["value"]],
        "technical": [row for row in tech if row["value"]],
        "act": processing_act,
        "created_at": created_at.isoformat() if created_at else "",
    }


def _enrich_request_body(*, agency, order_type: str, order_id: str, entries: list, payload: dict) -> dict[str, Any]:
    raw = str(order_type or "").strip().lower()
    if raw == "receiving":
        return _receiving_detail_from_entries(entries)
    if raw == "shipping":
        return _shipping_detail_from_db(agency=agency, order_id=order_id, fallback_payload=payload)
    if raw in {"processing", "packing"}:
        return _processing_detail_from_entries(entries)
    meta = []
    if payload.get("category_label") or payload.get("category"):
        meta.append({"label": "Категория", "value": _text(payload.get("category_label") or payload.get("category"))})
    if payload.get("description"):
        meta.append({"label": "Описание", "value": _text(payload.get("description"))})
    return {
        "meta": meta,
        "lines": [],
        "columns": [],
        "discrepancies": [],
        "has_mismatch": False,
        "lines_title": "Состав",
        "empty_lines_hint": "",
    }


def _other_execution_result(*, agency, order_id: str) -> dict[str, Any] | None:
    """Return the same completed-result facts that managers see, without prices."""
    from django.core.exceptions import ValidationError

    from billing.warehouse_services import facts_payload, list_facts

    from .models import OtherRequest
    from .other_request_workflow import result_outcome_label

    item = (
        OtherRequest.objects.filter(agency=agency, public_number=str(order_id or "").strip())
        .select_related("assignee")
        .first()
    )
    if item is None or not (item.result_outcome or item.result_comment or item.completed_at):
        return None

    service_facts = []
    try:
        service_facts = facts_payload(
            list_facts(
                client=agency,
                order_type="other",
                order_id=item.public_number,
            )
        )
    except ValidationError:
        service_facts = []

    def format_dt(value) -> str:
        if not value:
            return ""
        return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")

    return {
        "outcome_label": result_outcome_label(item.result_outcome) if item.result_outcome else "",
        "comment": str(item.result_comment or "").strip(),
        "quantity": str(item.result_qty) if item.result_qty is not None else "",
        "unit": str(item.result_unit or "").strip(),
        "services": [
            {
                "name": str(fact.get("name") or "").strip(),
                "quantity": str(fact.get("quantity") or "").strip(),
                "unit": str(fact.get("unit") or "").strip(),
                "reported_by": str(fact.get("reported_by") or "").strip(),
            }
            for fact in service_facts
        ],
        "assignee": str(getattr(item.assignee, "full_name", "") or "").strip(),
        "started_at": format_dt(item.started_at),
        "completed_at": format_dt(item.completed_at),
    }


def build_request_detail(*, agency, order_type: str, order_id: str) -> dict[str, Any] | None:
    from . import web_ui as client_web_ui

    raw_type = str(order_type or "").strip().lower()
    raw_order_id = str(order_id or "").strip()
    lookup_order_id = canonical_request_order_id(raw_type, raw_order_id)
    if raw_type == "processing":
        type_filter = ("processing", "packing")
    else:
        type_filter = (raw_type,)
    if raw_type not in CLIENT_REQUEST_ORDER_TYPES and raw_type != "processing":
        return None

    entry_order_ids = [lookup_order_id]
    if raw_order_id and raw_order_id != lookup_order_id:
        entry_order_ids.append(raw_order_id)

    entries = list(
        OrderAuditEntry.objects.filter(agency=agency, order_id__in=entry_order_ids, order_type__in=type_filter)
        .select_related("user")
        .order_by("created_at")
    )
    if not entries:
        return None

    latest = entries[-1]
    for entry in reversed(entries):
        if client_web_ui._is_status_entry(entry):
            latest = entry
            break

    payload = dict(latest.payload or {})
    detail_order_id = lookup_order_id if latest.order_type == "shipping" else raw_order_id
    act_block = None
    if latest.order_type == "receiving":
        payload = _merge_receiving_act_fields(entries, payload)
        act_block = _receiving_act_block(agency=agency, order_id=detail_order_id, entries=entries)
    status_label = client_web_ui._repair_mojibake_text(client_web_ui._order_status_label(latest))
    if status_label == "Отменена":
        status_label = "ОТМЕНЕНА"
    bucket = client_web_ui._order_bucket(latest)
    if latest.order_type == "receiving":
        try:
            from sklad.services.warehouse_state import WarehouseGoodsStateResolver

            resolved = WarehouseGoodsStateResolver.resolve_for_receiving_order(
                order_id=detail_order_id,
                agency=agency,
                payload=payload,
            )
            resolved_label = client_web_ui._repair_mojibake_text(resolved.label_for("client") or "")
            if resolved_label:
                status_label = resolved_label
            if status_label == "Отменена":
                status_label = "ОТМЕНЕНА"
            if resolved.is_terminal:
                bucket = "done"
            elif bucket == "client" and status_label and status_label != "Черновик":
                bucket = "warehouse"
        except Exception:
            pass
        # Акт отправлен клиенту — статус/корзина из act-aware map, не «Выполнена» со склада.
        if payload.get("act_sent") and str(payload.get("act_client_response") or "").lower() not in {
            "confirmed",
            "dispute",
        }:
            from types import SimpleNamespace

            from .lk_status_map import resolve_lk_entry_status

            mapped = resolve_lk_entry_status(
                SimpleNamespace(
                    order_type="receiving",
                    order_id=detail_order_id,
                    agency=agency,
                    agency_id=getattr(agency, "id", None),
                    payload=payload,
                ),
                audience="client",
            )
            status_label = mapped.status_label
            bucket = mapped.bucket
    if (
        bucket == "done"
        and latest.order_type != "shipping"
        and status_label != "ОТМЕНЕНА"
        and not (
            latest.order_type == "receiving"
            and payload.get("act_sent")
            and str(payload.get("act_client_response") or "").lower() not in {"confirmed", "dispute"}
        )
    ):
        status_label = "Выполнена"
    from .client_visibility import build_client_visible_history, client_facing_status_label

    body = _enrich_request_body(
        agency=agency,
        order_type=latest.order_type,
        order_id=detail_order_id,
        entries=entries,
        payload=payload,
    )
    responsibles = _request_responsibles(agency=agency, entries=entries)
    shipping_view = body.get("shipping_view") if latest.order_type == "shipping" else None
    receiving_view = None
    processing_view = None
    if latest.order_type == "receiving":
        from .client_receiving_view import build_receiving_client_view

        receiving_view = build_receiving_client_view(
            agency=agency,
            order_id=detail_order_id,
            entries=entries,
            payload=payload,
            body=body,
            status_label=status_label,
            bucket=bucket,
            created_at=getattr(entries[0], "created_at", None) if entries else None,
        )
    elif latest.order_type in {"processing", "packing"}:
        processing_view = _processing_client_view(
            agency=agency,
            order_id=detail_order_id,
            entries=entries,
            payload=payload,
            body=body,
            responsibles=responsibles,
        )
    if shipping_view and shipping_view.get("status_label"):
        # Client 6-step scale — do not collapse to generic «В работе».
        status_label = shipping_view["status_label"]
        stage_key = shipping_view.get("stage_key") or ""
        if stage_key == "canceled":
            status_label = "ОТМЕНЕНА"
            bucket = "done"
        elif stage_key == "closed":
            bucket = "done"
        elif stage_key in {"assembling", "ready"}:
            bucket = "warehouse"
        elif stage_key == "confirmed":
            bucket = "manager"
        elif stage_key == "not_delivered":
            bucket = "done"
        elif stage_key in {"in_transit", "delivered"}:
            bucket = "warehouse"
    elif receiving_view and receiving_view.get("status_label"):
        status_label = receiving_view["status_label"]
        stage_key = receiving_view.get("stage_key") or ""
        if stage_key == "canceled":
            status_label = "ОТМЕНЕНА"
            bucket = "done"
        elif stage_key == "done":
            bucket = "done"
        elif stage_key in {"receiving", "accepted"}:
            bucket = "warehouse"
        elif stage_key == "confirmed":
            bucket = "warehouse"
        elif stage_key == "pending":
            bucket = "manager"
        elif stage_key == "created":
            bucket = "client"
    elif processing_view and processing_view.get("status_label"):
        status_label = processing_view["status_label"]
        stage_key = processing_view.get("status_key") or ""
        if stage_key == "cancelled":
            status_label = "ОТМЕНЕНА"
            bucket = "done"
        elif stage_key == "done":
            bucket = "done"
        elif stage_key in {"handoff", "in_work"}:
            bucket = "warehouse"
        elif stage_key == "result_review":
            processing_act = processing_view.get("act") or {}
            bucket = "warehouse" if processing_act.get("manager_response") == "confirmed" else "manager"
        elif stage_key == "approved":
            bucket = "warehouse"
        else:
            bucket = "manager"
    else:
        status_label = client_facing_status_label(status_label, bucket=bucket)

    if _cancel_request_pending(payload):
        status_label = "Отмена на согласовании менеджера"
        bucket = "manager"

    attention = request_attention(
        agency=agency,
        order_type=latest.order_type,
        order_id=detail_order_id,
        payload=payload,
    )
    if shipping_view and shipping_view.get("attention"):
        attention = True
    if shipping_view is not None:
        appeal_attention = list(shipping_view.get("attention") or [])
        if _cancel_request_pending(payload):
            appeal_attention.append(
                {
                    "type": "cancel_request",
                    "title": "Запрос на отмену на рассмотрении",
                    "description": "Складская заявка не изменена до решения менеджера.",
                }
            )
            attention = True
        if _change_request_pending(payload):
            appeal_attention.append(
                {
                    "type": "change_request",
                    "title": "Запрос на изменение на рассмотрении",
                    "description": str(payload.get("change_reason") or "Ожидается решение менеджера."),
                }
            )
            attention = True
        shipping_view = dict(shipping_view)
        shipping_view["attention"] = appeal_attention
    filter_status = client_web_ui._lk_request_filter_status(
        bucket=bucket,
        status_label=status_label,
        attention=attention,
    )
    type_key = "processing" if latest.order_type == "packing" else client_web_ui._lk_request_type_key(latest.order_type)
    wms_url = client_web_ui._order_detail_url(latest, agency.id, client_view=True)
    history = build_client_visible_history(
        list(reversed(entries[-40:])),
        repair_text=client_web_ui._repair_mojibake_text,
        format_date=client_web_ui._lk_format_request_date,
    )

    can_cancel = False
    cancel_label = "Отменить заявку"
    can_comment = False
    can_continue_draft = False
    can_request_change = False
    can_edit = False
    can_accept_result = False
    can_return_to_work = False
    continue_url = ""
    if _can_client_cancel_from_bucket(bucket=bucket, status_label=status_label, payload=payload):
        can_cancel = True
        if _requires_manager_cancel_approval(bucket, status_label) or latest.order_type == "shipping":
            cancel_label = "Запросить отмену у менеджера"
    can_request_change = _can_client_request_change(
        bucket=bucket,
        status_label=status_label,
        payload=payload,
        order_type=latest.order_type,
    )
    if latest.order_type in {"processing", "packing"}:
        try:
            from processing_app.stages import processing_client_can_edit

            can_edit = processing_client_can_edit(payload)
        except Exception:
            can_edit = False
        if can_edit:
            continue_url = f"/orders/processing/?client={agency.id}&order={detail_order_id}&edit=1"
        elif not can_request_change and bucket in {"manager", "warehouse"} and not _change_request_pending(payload):
            label = str(status_label or "").lower()
            can_request_change = not any(token in label for token in ("отмен", "выполн", "заверш", "закрыт"))
    if latest.order_type == "other":
        can_comment = can_client_comment_other(payload)
        result_status = str(payload.get("client_result_status") or "").strip().lower()
        raw_status = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
        can_accept_result = raw_status in {"completed", "done"} and result_status != "accepted"
        can_return_to_work = can_accept_result
    if client_web_ui._is_draft_entry(latest):
        from .client_drafts import continue_url_for_entry

        can_continue_draft = True
        continue_url = continue_url_for_entry(
            agency_id=agency.id,
            order_type=latest.order_type,
            order_id=detail_order_id,
        )
        wms_url = continue_url or wms_url

    title = client_web_ui._repair_mojibake_text(
        payload.get("order_title")
        or payload.get("title")
        or payload.get("category_label")
        or latest.description
        or f"Заявка {detail_order_id}"
    )
    if latest.order_type == "shipping" and (not title or title == f"Заявка {detail_order_id}"):
        title = "Заявка на отгрузку"
    if latest.order_type == "receiving":
        if receiving_view and receiving_view.get("sku_count"):
            title = "Заявка на приёмку"
        elif title and "без указания" in title.lower():
            title = "Заявка на приёмку"
        elif not title or title == f"Заявка {detail_order_id}":
            title = "Заявка на приёмку"
    display_title = build_request_display_title(
        order_type=latest.order_type,
        order_id=detail_order_id,
        payload=payload,
        base_title=title,
        body=body,
        receiving_view=receiving_view,
        processing_view=processing_view,
        shipping_view=shipping_view,
    )
    attachments = []
    execution_result = None
    if latest.order_type == "other":
        from .other_requests import list_other_attachments

        attachments = list_other_attachments(order_id=str(order_id), agency=agency)
        execution_result = _other_execution_result(agency=agency, order_id=latest.order_id)
    elif latest.order_type == "receiving" and receiving_view:
        attachments = list(receiving_view.get("attachments") or [])

    detail_lines = body.get("lines") or []
    detail_columns = body.get("columns") or []
    if receiving_view:
        detail_lines = receiving_view.get("lines") or detail_lines
        detail_columns = receiving_view.get("columns") or detail_columns

    return {
        "order_id": detail_order_id,
        "order_type": latest.order_type,
        "type": type_key,
        "type_label": client_web_ui._lk_request_type_label(latest.order_type),
        "type_badge": client_web_ui._lk_request_type_badge(latest.order_type),
        "number": format_order_number(latest.order_type, detail_order_id),
        "title": title,
        "display_title": display_title,
        "description": _text(payload.get("description") or ""),
        "category": str(payload.get("category") or "").strip(),
        "category_label": str(payload.get("category_label") or "").strip(),
        "status": filter_status,
        "status_label": status_label,
        "status_pill": status_label if type_key == "receiving" else client_web_ui._lk_request_status_pill(filter_status),
        "cancel_reason": _text(payload.get("cancel_reason") or ""),
        "bucket": bucket,
        "attention": attention,
        "payload": {
            "status": payload.get("status"),
            "category": payload.get("category"),
            "category_label": payload.get("category_label"),
            "description": payload.get("description"),
            "cancel_reason": payload.get("cancel_reason"),
        },
        "attachments": attachments,
        "execution_result": execution_result,
        "responsibles": responsibles,
        "meta": body.get("meta") or [],
        "lines": detail_lines,
        "columns": detail_columns,
        "discrepancies": body.get("discrepancies") or [],
        "has_mismatch": bool(body.get("has_mismatch")),
        "lines_title": body.get("lines_title") or "Состав",
        "empty_lines_hint": body.get("empty_lines_hint") or "",
        "act": act_block,
        "shipping_view": shipping_view,
        "ozon_gm_structure": body.get("ozon_gm_structure"),
        "receiving_view": receiving_view,
        "processing_view": processing_view,
        "wms_url": wms_url,
        "lk_url": lk_request_hash(type_key, detail_order_id),
        "history": history,
        "actions": {
            "can_cancel": can_cancel,
            "cancel_label": cancel_label,
            "can_request_change": can_request_change,
            "change_label": "Запросить изменение",
            "can_edit": can_edit,
            "edit_url": continue_url if can_edit else "",
            "can_comment": can_comment,
            "can_accept_result": can_accept_result,
            "accept_result_label": "Принять результат",
            "can_return_to_work": can_return_to_work,
            "return_to_work_label": "Вернуть в работу",
            "can_continue_draft": can_continue_draft,
            "continue_url": continue_url,
            "open_wms": False,
            "can_export_excel": type_key in {"receiving", "shipping", "processing"} and not can_continue_draft,
        },
    }


def build_request_excel_workbook(*, agency, order_type: str, order_id: str):
    """Build Excel with article lines only (no request meta / discrepancies)."""
    from openpyxl import Workbook

    detail = build_request_detail(agency=agency, order_type=order_type, order_id=order_id)
    if not detail:
        return None, None
    type_key = str(detail.get("type") or "").lower()
    if type_key not in {"receiving", "shipping", "processing"}:
        return None, None

    columns = detail.get("columns") or []
    lines = detail.get("lines") or []
    # Export only article rows in the same shape as the LK card table.
    art_keys = {
        "receiving": (
            ("barcode", "ШК"),
            ("sku_code", "Артикул"),
            ("name", "Наименование"),
            ("size", "Размер"),
            ("qty_planned", "План"),
            ("qty_actual", "Факт"),
            ("box_qty", "Короба"),
            ("delta_qty", "Расхождение"),
        ),
        "shipping": (
            ("sku_code", "Артикул"),
            ("name", "Наименование"),
            ("size", "Размер"),
            ("barcode", "ШК"),
            ("qty_requested", "Заявлено"),
            ("qty_reserved", "Резерв"),
            ("qty_shipped", "Отгружено"),
            ("gm_barcodes", "ШК ГМ"),
            ("comment", "Комментарий"),
        ),
        "processing": (
            ("sku_code", "Артикул"),
            ("name", "Наименование"),
            ("size", "Размер"),
            ("barcode", "ШК"),
            ("qty", "Кол-во"),
        ),
    }
    preferred = art_keys.get(type_key) or ()
    if preferred:
        by_key = {str(col.get("key") or ""): col for col in columns}
        export_columns = []
        show_fact = bool((detail.get("receiving_view") or {}).get("show_fact_columns"))
        if type_key != "receiving":
            show_fact = True
        for key, default_label in preferred:
            if type_key == "receiving" and key in {"qty_actual", "box_qty", "delta_qty"} and not show_fact:
                continue
            col = by_key.get(key) or {"key": key, "label": default_label}
            if not col.get("label"):
                col = {**col, "label": default_label}
            # Prefer stable export labels for receiving fact columns.
            if type_key == "receiving" and key in {"qty_actual", "box_qty", "delta_qty", "barcode"}:
                col = {"key": key, "label": default_label}
            export_columns.append(col)
    else:
        export_columns = list(columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "Артикулы"

    if export_columns:
        ws.append([col.get("label") or col.get("key") or "" for col in export_columns])
        for line in lines:
            ws.append([
                "" if line.get(col.get("key")) is None else line.get(col.get("key"))
                for col in export_columns
            ])
    else:
        ws.append(["Артикул", "Наименование", "Размер", "План", "Факт"])
        for line in lines:
            ws.append([
                line.get("sku_code") or "",
                line.get("name") or "",
                line.get("size") if line.get("size") is not None else "",
                line.get("qty_planned") if line.get("qty_planned") is not None else line.get("qty_requested"),
                line.get("qty_actual") if line.get("qty_actual") is not None else line.get("qty_shipped", line.get("qty")),
            ])

    number = detail.get("number") or order_id
    safe_number = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(number))[:48]
    filename = f"art_{type_key}_{safe_number or order_id}.xlsx"
    return wb, filename


def apply_client_request_action(
    *,
    agency,
    order_type: str,
    order_id: str,
    action: str,
    user=None,
    text: str = "",
) -> tuple[bool, str]:
    raw = str(order_type or "").strip().lower()
    if raw == "packing":
        raw = "processing"
    action_name = str(action or "").strip().lower()
    if action_name == "cancel":
        if raw not in CLIENT_REQUEST_ORDER_TYPES and raw != "processing":
            return False, "Для этого типа заявки действие недоступно в ЛК"

        entries = list(
            OrderAuditEntry.objects.filter(agency=agency, order_id=str(order_id), order_type__in=("processing", "packing") if raw == "processing" else (raw,))
            .select_related("user")
            .order_by("created_at", "id")
        )
        if not entries:
            return False, "Заявка не найдена"

        from . import web_ui as client_web_ui

        latest = entries[-1]
        for entry in reversed(entries):
            if client_web_ui._is_status_entry(entry):
                latest = entry
                break
        payload = dict(latest.payload or {})
        status_label = client_web_ui._repair_mojibake_text(client_web_ui._order_status_label(latest))
        if status_label == "Отменена":
            status_label = "ОТМЕНЕНА"
        bucket = client_web_ui._order_bucket(latest)
        if latest.order_type == "receiving":
            try:
                from sklad.services.warehouse_state import WarehouseGoodsStateResolver

                resolved = WarehouseGoodsStateResolver.resolve_for_receiving_order(
                    order_id=str(order_id),
                    agency=agency,
                    payload=payload,
                )
                resolved_label = client_web_ui._repair_mojibake_text(resolved.label_for("client") or "")
                if resolved_label:
                    status_label = resolved_label
                if status_label == "Отменена":
                    status_label = "ОТМЕНЕНА"
                if resolved.is_terminal:
                    bucket = "done"
                elif bucket == "client" and status_label and status_label != "Черновик":
                    bucket = "warehouse"
            except Exception:
                pass
        if not _can_client_cancel_from_bucket(bucket=bucket, status_label=status_label, payload=payload):
            return False, "Отмена недоступна в текущем статусе"

        reason = str(text or "").strip() or "Отменено клиентом"
        # Shipping: never mutate warehouse from LK — always manager appeal.
        if latest.order_type == "shipping" or _requires_manager_cancel_approval(bucket, status_label):
            ok = _request_manager_cancel_approval(latest=latest, agency=agency, user=user, reason=reason)
            return (
                (True, "Запрос на отмену отправлен менеджеру")
                if ok
                else (False, "Запрос на отмену уже ожидает менеджера")
            )

        if latest.order_type == "other":
            ok = cancel_other_request(order_id=order_id, user=user, reason=reason)
            return (True, "Заявка отменена") if ok else (False, "Не удалось отменить")
        ok = _log_generic_client_cancel(latest=latest, agency=agency, user=user, reason=reason)
        return (True, "Заявка отменена") if ok else (False, "Не удалось отменить")

    if action_name == "change_request":
        if raw not in {"shipping", "processing"}:
            return False, "Запрос изменения доступен только для отгрузки и обработки"
        entries = list(
            OrderAuditEntry.objects.filter(
                agency=agency,
                order_id=str(order_id),
                order_type__in=("processing", "packing") if raw == "processing" else ("shipping",),
            )
            .select_related("user")
            .order_by("created_at", "id")
        )
        if not entries:
            return False, "Заявка не найдена"
        from . import web_ui as client_web_ui
        if raw == "shipping":
            from .client_shipping_view import resolve_client_shipping_stage

        latest = entries[-1]
        for entry in reversed(entries):
            if client_web_ui._is_status_entry(entry):
                latest = entry
                break
        payload = dict(latest.payload or {})
        status_label = client_web_ui._repair_mojibake_text(client_web_ui._order_status_label(latest))
        bucket = client_web_ui._order_bucket(latest)
        if raw == "shipping":
            try:
                from shipping.models import ShippingOrder

                order = ShippingOrder.objects.filter(agency=agency, number=str(order_id)).first()
                if order is not None:
                    trip_status = ""
                    try:
                        trip_status = client_web_ui._shipping_trip_status(
                            getattr(order, "number", None) or order_id,
                            agency_id=getattr(agency, "id", None),
                        )
                    except Exception:
                        trip_status = ""
                    stage_key, stage_label = resolve_client_shipping_stage(
                        order_status=getattr(order, "status", "") or "",
                        trip_status=trip_status,
                    )
                    status_label = stage_label or status_label
                    if stage_key == "canceled":
                        bucket = "done"
                    elif stage_key == "closed":
                        bucket = "done"
                    elif stage_key in {"assembling", "ready", "in_transit", "delivered"}:
                        bucket = "warehouse"
                    elif stage_key == "not_delivered":
                        bucket = "done"
                    elif stage_key == "confirmed":
                        bucket = "manager"
            except Exception:
                pass
        elif raw == "processing":
            stage_key = _processing_client_stage(payload)
            if stage_key == "cancelled":
                bucket = "done"
            elif stage_key == "done":
                bucket = "done"
            elif stage_key in {"approved", "handoff", "in_work"}:
                bucket = "warehouse"
            else:
                bucket = "manager"
        if not _can_client_request_change(
            bucket=bucket,
            status_label=status_label,
            payload=payload,
            order_type=raw,
        ):
            return False, "Запрос изменения недоступен в текущем статусе"
        reason = str(text or "").strip()
        if not reason:
            return False, "Опишите, что нужно изменить"
        ok = _request_manager_change_approval(latest=latest, agency=agency, user=user, reason=reason)
        return (
            (True, "Запрос на изменение отправлен менеджеру. Складская заявка не изменена.")
            if ok
            else (False, "Запрос на изменение уже ожидает менеджера")
        )

    if raw != "other":
        return False, "Для этого типа заявки действие пока недоступно в ЛК"
    latest = latest_other_entry(order_id)
    if not latest or latest.agency_id != getattr(agency, "id", None):
        return False, "Заявка не найдена"
    payload = dict(latest.payload or {})
    if action_name in {"accept", "accept_result"}:
        raw_status = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
        if raw_status not in {"completed", "done"}:
            return False, "Принять можно только выполненную заявку"
        if str(payload.get("client_result_status") or "").strip().lower() == "accepted":
            return False, "Результат уже принят"
        from .other_request_workflow import client_accept_result, ensure_other_request_from_audit

        req = ensure_other_request_from_audit(order_id)
        if req is None or req.agency_id != getattr(agency, "id", None):
            return False, "Заявка не найдена"
        client_accept_result(req, user=user)
        return True, "Результат принят"
    if action_name in {"return_to_work", "return_work", "rework"}:
        raw_status = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
        if raw_status not in {"completed", "done"}:
            return False, "Вернуть в работу можно только выполненную заявку"
        reason = str(text or "").strip()
        if not reason:
            return False, "Опишите, что нужно доработать"
        from .other_request_workflow import client_return_to_work, ensure_other_request_from_audit

        req = ensure_other_request_from_audit(order_id)
        if req is None or req.agency_id != getattr(agency, "id", None):
            return False, "Заявка не найдена"
        client_return_to_work(req, user=user, reason=reason)
        return True, "Заявка возвращена в работу"
    if action_name == "comment":
        if not can_client_comment_other(payload):
            return False, "Комментарий к закрытой заявке недоступен"
        if not str(text or "").strip():
            return False, "Введите текст комментария"
        ok = add_other_comment(order_id=order_id, agency=agency, user=user, text=text)
        return (True, "Комментарий добавлен") if ok else (False, "Не удалось сохранить комментарий")
    return False, "Неизвестное действие"
