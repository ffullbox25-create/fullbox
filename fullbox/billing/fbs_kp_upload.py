"""Read-only FBS KPI/KP calculation from uploaded acceptance acts and shipment report.

The module intentionally does not create billing applications, charges, acts,
invoices or warehouse events. It only parses manager-uploaded files and builds
a preview/export using the already configured FBS client rates.
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Iterable

from django.core.exceptions import ValidationError
from django.db.models import Q
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from sku.models import SKU

from .fbs_standard_shipping import fbs_rate_billable_quantity, sku_liters
from .models import ClientTariffVersion, FbsClientRate
from .services import calculate_amounts


ZERO = Decimal("0")
MONEY = Decimal("0.01")
QTY = Decimal("0.000")
LITERS = Decimal("0.001")
ORANGE = "F89000"
YELLOW = "F8B800"
PALE_ORANGE = "FCE6C4"
PALE_GREEN = "DDF3EE"
TEXT = "303030"
MUTED = "6B7280"
BORDER = "E5E7EB"

OPERATION_LABELS = dict(FbsClientRate.OPERATION_CHOICES)


@dataclass
class SourceRow:
    source: str
    operation: str
    source_number: str
    service_date: date | None
    name: str
    sku_code: str
    barcode: str
    quantity: Decimal


@dataclass
class CalculatedRow:
    source: str
    operation: str
    source_number: str
    service_date: date | None
    sku: SKU | None
    name: str
    sku_code: str
    barcode: str
    quantity: Decimal
    billing_quantity: Decimal
    liters: Decimal | None
    rate: FbsClientRate | None
    amount: Decimal
    vat_amount: Decimal
    total_amount: Decimal
    problem: str = ""

    @property
    def operation_label(self) -> str:
        return OPERATION_LABELS.get(self.operation, self.operation)


def _decimal(value) -> Decimal:
    if value is None:
        return ZERO
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, AttributeError, ValueError):
        return ZERO


def _money(value) -> Decimal:
    return Decimal(value or 0).quantize(MONEY, rounding=ROUND_HALF_UP)


def _qty(value) -> Decimal:
    return _decimal(value).quantize(QTY, rounding=ROUND_HALF_UP)


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw:
        return None
    parsed = parse_date(raw)
    if parsed:
        return parsed
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    return None


def _extract_sku_code(name: str) -> str:
    match = re.search(r"\(([^)]+)\)", name or "")
    if match:
        return match.group(1).strip()
    # Some acts contain only the article as the product name.
    return (name or "").strip().split(" ", 1)[0].strip()


def _pdf_text(file_obj) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValidationError(
            "На сервере не установлен PDF-парсер pypdf. Excel-отчет можно разобрать, но PDF-акты сейчас недоступны."
        ) from exc
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        reader = PdfReader(file_obj)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # pragma: no cover - depends on malformed external files.
        raise ValidationError(f"Не удалось прочитать PDF {getattr(file_obj, 'name', '')}: {exc}") from exc


def parse_receiving_act_text(text: str, *, file_name: str = "") -> list[SourceRow]:
    act_match = re.search(r"хранение\s+№\s*([0-9]+)", text, re.IGNORECASE)
    # The document header also contains a contract line like "от 01.01.1970".
    # The act date itself is printed in uppercase: "ОТ 04.08.2026".
    date_match = re.search(r"\bОТ\s+([0-9]{2}\.[0-9]{2}\.[0-9]{4})", text)
    task_match = re.search(r"Задача\s+на\s+прием\s+товара\s+№\s*([0-9]+)", text, re.IGNORECASE)
    act_number = act_match.group(1) if act_match else file_name
    act_date = _date(date_match.group(1)) if date_match else None
    task_number = task_match.group(1) if task_match else ""

    if "№ ТОВАР" in text and "ВСЕГО:" in text:
        body = text.split("№ ТОВАР", 1)[1].split("ВСЕГО:", 1)[0]
    else:
        body = text
    flat = re.sub(r"\s+", " ", body)
    pattern = re.compile(
        r"(\d+)\s+(.+?)\s+(\d{8,14})\s+(\d+)\s+(\d+(?:[,.]\d+)?)\s+(\d+(?:[,.]\d+)?)\s+Не указана",
        re.IGNORECASE,
    )
    rows: list[SourceRow] = []
    for _, name, barcode, _external_article, _plan_qty, fact_qty in pattern.findall(flat):
        quantity = _qty(fact_qty)
        if quantity <= 0:
            continue
        rows.append(
            SourceRow(
                source="PDF акт приемки",
                operation=FbsClientRate.OP_RECEIVING,
                source_number=f"ACT-{act_number}" + (f" / PR-{task_number}" if task_number else ""),
                service_date=act_date,
                name=name.strip(),
                sku_code=_extract_sku_code(name),
                barcode=str(barcode).strip(),
                quantity=quantity,
            )
        )
    if not rows:
        raise ValidationError(f"В PDF {file_name or act_number} не найдены строки товаров.")
    return rows


def parse_receiving_pdfs(files: Iterable) -> list[SourceRow]:
    rows: list[SourceRow] = []
    for file_obj in files:
        rows.extend(parse_receiving_act_text(_pdf_text(file_obj), file_name=getattr(file_obj, "name", "")))
    return rows


def _header_map(sheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in next(sheet.iter_rows(min_row=1, max_row=1)):
        key = str(cell.value or "").strip().lower()
        if key:
            headers[key] = cell.column
    return headers


def _cell(row, header: dict[str, int], *names):
    for name in names:
        idx = header.get(name.lower())
        if idx:
            return row[idx - 1]
    return None


def parse_shipping_report(file_obj, *, date_from: date | None = None, date_to: date | None = None) -> list[SourceRow]:
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        sheet = workbook.active
    except Exception as exc:
        raise ValidationError(f"Не удалось прочитать Excel-отчет отгрузок: {exc}") from exc
    headers = _header_map(sheet)
    required = ["дата", "товар", "артикул", "шк", "всего"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValidationError(f"В Excel-отчете нет колонок: {', '.join(missing)}.")

    rows: list[SourceRow] = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        shipped_at = _date(_cell(raw_row, headers, "дата"))
        if date_from and shipped_at and shipped_at < date_from:
            continue
        if date_to and shipped_at and shipped_at > date_to:
            continue
        total_qty = _qty(_cell(raw_row, headers, "всего"))
        if total_qty <= 0:
            continue
        sku_code = str(_cell(raw_row, headers, "артикул") or "").strip()
        name = str(_cell(raw_row, headers, "товар") or "").strip()
        barcode = str(_cell(raw_row, headers, "шк") or "").strip()
        source_number = f"Отгрузка {shipped_at:%d.%m.%Y}" if shipped_at else "Отгрузка из Excel"
        for operation in (FbsClientRate.OP_PICKING, FbsClientRate.OP_SHIPPING):
            rows.append(
                SourceRow(
                    source="Excel отгрузки",
                    operation=operation,
                    source_number=source_number,
                    service_date=shipped_at,
                    name=name,
                    sku_code=sku_code,
                    barcode=barcode,
                    quantity=total_qty,
                )
            )
        rows.append(
            SourceRow(
                source="Excel отгрузки",
                operation=FbsClientRate.OP_MARKING,
                source_number=source_number,
                service_date=shipped_at,
                name=name,
                sku_code=sku_code,
                barcode=barcode,
                quantity=total_qty,
            )
        )
    return rows


def _sku_maps(client, source_rows: Iterable[SourceRow]) -> tuple[dict[str, SKU], dict[str, SKU]]:
    codes = {row.sku_code.strip() for row in source_rows if row.sku_code}
    barcodes = {row.barcode.strip() for row in source_rows if row.barcode}
    skus = list(
        SKU.objects.filter(deleted=False, agency=client)
        .filter(Q(sku_code__in=codes) | Q(barcodes__value__in=barcodes))
        .prefetch_related("barcodes")
        .distinct()
    )
    by_code = {sku.sku_code.strip().lower(): sku for sku in skus if sku.sku_code}
    by_barcode: dict[str, SKU] = {}
    for sku in skus:
        for barcode in sku.barcodes.all():
            value = str(barcode.value or "").strip()
            if value:
                by_barcode[value] = sku
    return by_code, by_barcode


def _rate_for(*, client, operation: str, liters: Decimal | None, on_date: date) -> FbsClientRate | None:
    rates = (
        FbsClientRate.objects.filter(
            client=client,
            operation=operation,
            is_active=True,
            valid_from__lte=on_date,
        )
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=on_date))
    )
    if operation in {FbsClientRate.OP_MARKING, FbsClientRate.OP_STORAGE}:
        return rates.order_by("-valid_from", "liters_from", "-id").first()
    if liters is None:
        return None
    return (
        rates.filter(liters_from__lt=liters)
        .filter(Q(liters_to__isnull=True) | Q(liters_to__gte=liters))
        .order_by("-valid_from", "-liters_from", "-id")
        .first()
    )


def _rate_label(rate: FbsClientRate | None) -> str:
    if rate is None:
        return ""
    if rate.operation in {FbsClientRate.OP_MARKING, FbsClientRate.OP_STORAGE}:
        return "без диапазона"
    upper = f"≤ {rate.liters_to} л" if rate.liters_to is not None else "без верхней границы"
    return f"> {rate.liters_from} л — {upper}"


def calculate_kp_upload(*, client, receiving_files, shipping_report_file, date_from: date, date_to: date) -> dict:
    receiving_rows = parse_receiving_pdfs(receiving_files) if receiving_files else []
    shipping_rows = parse_shipping_report(shipping_report_file, date_from=date_from, date_to=date_to) if shipping_report_file else []
    source_rows = [
        row
        for row in receiving_rows + shipping_rows
        if row.service_date is None or date_from <= row.service_date <= date_to
    ]
    if not source_rows:
        raise ValidationError("В выбранных файлах не найдено строк для расчета за указанный период.")

    by_code, by_barcode = _sku_maps(client, source_rows)
    calculated: list[CalculatedRow] = []
    for row in source_rows:
        sku = by_barcode.get(row.barcode) or by_code.get(row.sku_code.lower())
        liters = sku_liters(sku) if sku is not None else None
        problem = ""
        rate = None
        service_date = row.service_date or date_from
        if sku is None:
            problem = "SKU не найден по штрихкоду или артикулу клиента."
        elif row.operation != FbsClientRate.OP_MARKING and liters is None:
            problem = "У SKU нет полных габаритов для расчета литража."
        else:
            rate = _rate_for(client=client, operation=row.operation, liters=liters, on_date=service_date)
            if rate is None:
                if row.operation == FbsClientRate.OP_MARKING:
                    problem = "Не найдена ставка КП для маркировки."
                else:
                    problem = f"Не найдена ставка КП для {liters.quantize(LITERS)} л."

        billing_quantity = row.quantity
        if rate is not None:
            billing_quantity = fbs_rate_billable_quantity(
                operation=row.operation,
                rate=rate,
                item_quantity=row.quantity,
                liters=liters,
            )
            amount, vat_amount, total_amount = calculate_amounts(
                billing_quantity,
                rate.price,
                rate.vat_rate,
                vat_type=rate.vat_type or ClientTariffVersion.VAT_EXTRA,
            )
        else:
            amount = vat_amount = total_amount = ZERO
        calculated.append(
            CalculatedRow(
                source=row.source,
                operation=row.operation,
                source_number=row.source_number,
                service_date=row.service_date,
                sku=sku,
                name=row.name,
                sku_code=row.sku_code,
                barcode=row.barcode,
                quantity=row.quantity,
                billing_quantity=billing_quantity,
                liters=liters,
                rate=rate,
                amount=_money(amount),
                vat_amount=_money(vat_amount),
                total_amount=_money(total_amount),
                problem=problem,
            )
        )

    groups = []
    grouped: dict[tuple, dict] = {}
    for row in calculated:
        if row.problem:
            continue
        key = (row.operation, row.rate.pk, _rate_label(row.rate))
        group = grouped.setdefault(
            key,
            {
                "operation": row.operation,
                "operation_label": OPERATION_LABELS.get(row.operation, row.operation),
                "rate_label": _rate_label(row.rate),
                "price": row.rate.price,
                "unit": row.rate.unit,
                "vat_rate": row.rate.vat_rate,
                "vat_type": row.rate.vat_type,
                "quantity": ZERO,
                "amount": ZERO,
                "vat_amount": ZERO,
                "total_amount": ZERO,
                "rows_count": 0,
            },
        )
        group["quantity"] += row.billing_quantity
        group["amount"] += row.amount
        group["vat_amount"] += row.vat_amount
        group["total_amount"] += row.total_amount
        group["rows_count"] += 1
    groups = sorted(grouped.values(), key=lambda item: (item["operation"], item["rate_label"]))
    for group in groups:
        group["amount"] = _money(group["amount"])
        group["vat_amount"] = _money(group["vat_amount"])
        group["total_amount"] = _money(group["total_amount"])

    errors = [row for row in calculated if row.problem]
    summary = {
        "source_rows": len(source_rows),
        "receiving_rows": len(receiving_rows),
        "shipping_rows": len(shipping_rows),
        "calculated_rows": len(calculated) - len(errors),
        "error_rows": len(errors),
        "quantity": sum((row.quantity for row in calculated if not row.problem), ZERO),
        "amount": _money(sum((row.amount for row in calculated), ZERO)),
        "vat_amount": _money(sum((row.vat_amount for row in calculated), ZERO)),
        "total_amount": _money(sum((row.total_amount for row in calculated), ZERO)),
    }
    return {
        "client": client,
        "date_from": date_from,
        "date_to": date_to,
        "rows": calculated,
        "groups": groups,
        "errors": errors,
        "summary": summary,
    }


def _set_title(sheet, title: str, last_column: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    cell = sheet["A1"]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=ORANGE)
    cell.font = Font(name="Manrope", size=14, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28


def _style_header(sheet, row: int, columns: int) -> None:
    thin = Side(style="thin", color=BORDER)
    for cells in sheet.iter_cols(min_col=1, max_col=columns, min_row=row, max_row=row):
        cell = cells[0]
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.font = Font(name="Manrope", size=9, bold=True, color=TEXT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[row].height = 30


def _format_rows(sheet, start_row: int, end_row: int, columns: int) -> None:
    thin = Side(style="thin", color=BORDER)
    for row in sheet.iter_rows(min_row=start_row, max_row=max(end_row, start_row), min_col=1, max_col=columns):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)


def _autosize(sheet) -> None:
    for column in range(1, sheet.max_column + 1):
        max_len = 10
        for cell in sheet.iter_cols(min_col=column, max_col=column, values_only=True):
            for value in cell:
                max_len = max(max_len, min(len(str(value or "")) + 2, 48))
        sheet.column_dimensions[get_column_letter(column)].width = max_len
    sheet.sheet_view.showGridLines = False


def _write_table(sheet, start_row: int, headers: list[str], rows: list[list]) -> int:
    sheet.append([])
    sheet.append(headers)
    header_row = start_row
    _style_header(sheet, header_row, len(headers))
    for row in rows:
        sheet.append(row)
    last_row = header_row + max(len(rows), 1)
    _format_rows(sheet, header_row + 1, last_row, len(headers))
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    return last_row


def build_kp_upload_workbook(result: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итоги"
    client = result["client"]
    _set_title(summary, f"FBS расчет по КП: {client}", "H")
    period = f"{result['date_from']:%d.%m.%Y} — {result['date_to']:%d.%m.%Y}"
    generated_at = timezone.localtime(timezone.now()).strftime("%d.%m.%Y %H:%M")
    summary_rows = [
        ["Период", period, "", "Сформировано", generated_at],
        ["Клиент", str(client), "", "Важно", "Файл не создает счета, акты, начисления или складские движения."],
        ["Строк исходных", result["summary"]["source_rows"], "", "Ошибок", result["summary"]["error_rows"]],
        ["Сумма без НДС", result["summary"]["amount"], "", "НДС", result["summary"]["vat_amount"]],
        ["Итого с НДС", result["summary"]["total_amount"], "", "", ""],
    ]
    for row in summary_rows:
        summary.append(row)
    for row in summary.iter_rows(min_row=2, max_row=6, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary.append([])
    _write_table(
        summary,
        8,
        ["Операция", "Диапазон", "Количество", "Ед.", "Цена", "Сумма без НДС", "НДС", "Итого"],
        [
            [
                group["operation_label"],
                group["rate_label"],
                group["quantity"],
                group["unit"],
                group["price"],
                group["amount"],
                group["vat_amount"],
                group["total_amount"],
            ]
            for group in result["groups"]
        ],
    )
    _autosize(summary)

    details = workbook.create_sheet("Детально")
    _set_title(details, f"FBS строки расчета: {client}", "O")
    detail_rows = []
    for row in result["rows"]:
        detail_rows.append(
            [
                row.source,
                OPERATION_LABELS.get(row.operation, row.operation),
                row.source_number,
                row.service_date.strftime("%d.%m.%Y") if row.service_date else "",
                row.sku.sku_code if row.sku else row.sku_code,
                row.name,
                row.barcode,
                row.quantity,
                row.liters.quantize(LITERS) if row.liters is not None else "",
                _rate_label(row.rate),
                row.rate.price if row.rate else "",
                row.amount,
                row.vat_amount,
                row.total_amount,
                row.problem,
            ]
        )
    _write_table(
        details,
        3,
        [
            "Источник",
            "Операция",
            "Документ/дата",
            "Дата",
            "SKU",
            "Наименование",
            "ШК",
            "Количество",
            "Литров 1 шт.",
            "Диапазон КП",
            "Цена",
            "Сумма без НДС",
            "НДС",
            "Итого",
            "Проблема",
        ],
        detail_rows,
    )
    _autosize(details)

    errors = workbook.create_sheet("Ошибки")
    _set_title(errors, "Строки, которые не попали в сумму", "J")
    error_rows = [
        [
            row.source,
            OPERATION_LABELS.get(row.operation, row.operation),
            row.source_number,
            row.service_date.strftime("%d.%m.%Y") if row.service_date else "",
            row.sku_code,
            row.name,
            row.barcode,
            row.quantity,
            row.liters.quantize(LITERS) if row.liters is not None else "",
            row.problem,
        ]
        for row in result["errors"]
    ]
    _write_table(
        errors,
        3,
        ["Источник", "Операция", "Документ", "Дата", "SKU", "Наименование", "ШК", "Количество", "Литров 1 шт.", "Ошибка"],
        error_rows,
    )
    _autosize(errors)

    rates = workbook.create_sheet("Ставки КП")
    _set_title(rates, f"Активные ставки FBS: {client}", "H")
    rate_rows = []
    for rate in FbsClientRate.objects.filter(client=client, is_active=True).order_by("operation", "valid_from", "liters_from", "id"):
        rate_rows.append(
            [
                rate.get_operation_display(),
                rate.liters_from,
                rate.liters_to if rate.liters_to is not None else "",
                rate.price,
                rate.unit,
                rate.valid_from.strftime("%d.%m.%Y"),
                rate.valid_to.strftime("%d.%m.%Y") if rate.valid_to else "",
                f"{rate.vat_rate}% / {rate.get_vat_type_display() if hasattr(rate, 'get_vat_type_display') else rate.vat_type}",
            ]
        )
    _write_table(rates, 3, ["Операция", "От л", "До л", "Цена", "Ед.", "С", "До", "НДС"], rate_rows)
    _autosize(rates)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, Decimal):
                    cell.value = float(cell.value)
                    cell.number_format = "#,##0.00"
        sheet.freeze_panes = "A4" if sheet.max_row > 4 else None

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
