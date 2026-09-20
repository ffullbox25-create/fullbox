"""Regression coverage for the manager-only storage Excel export."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client as DjangoClient, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from billing.manager_storage_export import build_manager_storage_export
from billing.models import ApplicationCharge, BillingApplication, BillingService, BillingStorageDay, StorageSnapshotLine
from sku.models import Agency
from sklad.models import WarehouseContainer


class ManagerStorageExportTests(TestCase):
    def setUp(self):
        self.client = Agency.objects.create(agn_name="Клиент отчёта хранения", short_name="Отчёт")
        self.service = BillingService.objects.create(code="storage_export_test", name="Хранение", unit="м3")
        self.application = BillingApplication.objects.create(
            application_type=BillingApplication.TYPE_STORAGE,
            application_id="STORAGE-EXPORT-TEST",
            client=self.client,
            legal_entity=self.client,
        )
        WarehouseContainer.objects.create(
            agency=self.client,
            container_type=WarehouseContainer.TYPE_BOX,
            container_code="BOX-REPORT-1",
            gross_weight_g=45000,
        )
        for offset, amount in enumerate((Decimal("30.00"), Decimal("40.00")), start=1):
            day = date(2026, 7, offset)
            charge = ApplicationCharge.objects.create(
                application=self.application,
                client=self.client,
                legal_entity=self.client,
                service=self.service,
                quantity=Decimal("1.000"),
                unit="м3",
                tariff=Decimal("30.0000"),
                amount=amount,
                vat_amount=amount * Decimal("0.05"),
                total_amount=amount * Decimal("1.05"),
                source_type=ApplicationCharge.SOURCE_STORAGE_DAY,
                source_key=f"storage:{self.client.id}:{day.isoformat()}",
                billing_period=day.replace(day=1),
            )
            storage_day = BillingStorageDay.objects.create(
                client=self.client,
                application=self.application,
                day=day,
                billing_mode="m3_day",
                physical_volume_m3=Decimal("1.000000"),
                billable_volume_m3=Decimal("1.000000"),
                amount=amount,
                vat_amount=amount * Decimal("0.05"),
                charge=charge,
                payload={"weight_volume_m3": "0.366667"},
            )
            StorageSnapshotLine.objects.create(
                day=storage_day,
                pallet_code="PAL-REPORT-1",
                box_code="BOX-REPORT-1",
                sku_code="SKU-REPORT-1",
                name="Тестовый товар",
                barcode="4600000000001",
                quantity=Decimal("10.000"),
                width_mm=100,
                height_mm=200,
                length_mm=300,
                total_volume_m3=Decimal("1.000000"),
                billable_volume_m3=Decimal("1.000000"),
            )

    def _create_storage_day(self, day: date, amount: Decimal = Decimal("30.00")) -> BillingStorageDay:
        charge = ApplicationCharge.objects.create(
            application=self.application,
            client=self.client,
            legal_entity=self.client,
            service=self.service,
            quantity=Decimal("1.000"),
            unit="м3",
            tariff=Decimal("30.0000"),
            amount=amount,
            vat_amount=amount * Decimal("0.05"),
            total_amount=amount * Decimal("1.05"),
            source_type=ApplicationCharge.SOURCE_STORAGE_DAY,
            source_key=f"storage:{self.client.id}:{day.isoformat()}",
            billing_period=day.replace(day=1),
        )
        storage_day = BillingStorageDay.objects.create(
            client=self.client,
            application=self.application,
            day=day,
            billing_mode="m3_day",
            physical_volume_m3=Decimal("1.000000"),
            billable_volume_m3=Decimal("1.000000"),
            amount=amount,
            vat_amount=amount * Decimal("0.05"),
            charge=charge,
            payload={"weight_volume_m3": "0.366667"},
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            pallet_code="PAL-REPORT-1",
            box_code=f"BOX-{day.isoformat()}",
            sku_code="SKU-REPORT-1",
            name="Тестовый товар",
            barcode="4600000000001",
            quantity=Decimal("10.000"),
            width_mm=100,
            height_mm=200,
            length_mm=300,
            total_volume_m3=Decimal("1.000000"),
            billable_volume_m3=Decimal("1.000000"),
        )
        return storage_day

    def test_calculation_sheet_reconciles_to_daily_billing(self):
        days = BillingStorageDay.objects.select_related("client", "application", "charge").order_by("day")
        response = build_manager_storage_export(
            client=self.client,
            start=date(2026, 7, 1),
            end=date(2026, 7, 2),
            storage_days=days,
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(response)
        self.assertEqual(workbook.sheetnames[0], "Расчёт по коробам")
        sheet = workbook["Расчёт по коробам"]
        self.assertEqual(sheet["A1"].value, "Расчёт хранения: Отчёт")
        self.assertEqual(sheet["K2"].value, 1)
        self.assertEqual(sheet["Q2"].value, Decimal("70.00"))
        self.assertEqual(sheet["V2"].value, Decimal("3.50"))
        self.assertEqual(sheet["Q3"].value, Decimal("73.50"))
        self.assertTrue(sheet.column_dimensions["B"].hidden)
        self.assertEqual(sheet["C7"].value, "BOX-REPORT-1")
        self.assertEqual(sheet["M7"].value, 45000)
        self.assertEqual(sheet["N7"].value, 0.3555555555555556)
        self.assertIsNone(sheet["A4"].value)
        self.assertIsNone(sheet["D4"].value)
        self.assertEqual(sheet["O7"].value, "Объём строки")
        self.assertEqual(sheet["R7"].value, 2)
        self.assertEqual(sheet["S7"].value, 35)
        self.assertEqual(sheet["T7"].value, "Объём")
        self.assertEqual(sheet["V7"].value, 70)
        self.assertEqual(sheet["W7"].value, 3.5)
        self.assertEqual(sheet["X7"].value, 73.5)
        self.assertIn("2 дн. хранения", sheet["Y7"].value)
        overview = workbook["Сводка по дням"]
        overview_headers = [cell.value for cell in overview[1]]
        self.assertNotIn("Палет", overview_headers)
        self.assertIn(sheet["B6"].value, ("", None))
        self.assertNotIn("Палет", [cell.value for cell in workbook["Хранение по артикулам"][1]])
        self.assertNotIn("Палета / откуда", [cell.value for cell in workbook["Движение товара"][1]])
        all_values = {
            str(cell.value)
            for report_sheet in workbook.worksheets
            for row in report_sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        self.assertNotIn("PAL-REPORT-1", all_values)
        self.assertFalse(any("Паллет" in value or "Палет" in value for value in all_values))

    def test_export_prefers_auditable_weight_value_from_current_cbm_rule(self):
        day = BillingStorageDay.objects.order_by("day").first()
        day.physical_volume_m3 = Decimal("0.100000")
        day.billable_volume_m3 = Decimal("0.800000")
        day.payload = {
            "weight_volume_m3": "9.999999",
            "billing_cbm_rule": {"weight_based_cbm": "0.800000"},
        }
        day.save(update_fields=["physical_volume_m3", "billable_volume_m3", "payload"])

        response = build_manager_storage_export(
            client=self.client,
            start=day.day,
            end=day.day,
            storage_days=BillingStorageDay.objects.filter(pk=day.pk).select_related("client", "application", "charge"),
        )

        workbook = load_workbook(response)
        overview = workbook["Сводка по дням"]
        self.assertEqual(overview["G2"].value, "0.800000")
        self.assertEqual(workbook["Расчёт по коробам"]["T7"].value, "Вес")
        notes = workbook["Пояснения"]
        note_titles = [row[0] for row in notes.iter_rows(values_only=True)]
        self.assertNotIn("Правило веса", note_titles)
        self.assertFalse(
            any(
                "450 кг = 1,6 м³" in str(cell.value)
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
                if cell.value is not None
            )
        )

    def test_pallet_report_keeps_pallet_columns(self):
        day = BillingStorageDay.objects.select_related("client", "application", "charge").order_by("day").first()
        day.billing_mode = "pallet_day"
        day.save(update_fields=["billing_mode"])
        response = build_manager_storage_export(
            client=self.client,
            start=day.day,
            end=day.day,
            storage_days=BillingStorageDay.objects.filter(pk=day.pk).select_related("client", "application", "charge"),
        )

        workbook = load_workbook(response)
        calculation = workbook["Расчёт по коробам"]
        self.assertFalse(calculation.column_dimensions["B"].hidden)
        self.assertEqual(calculation["B6"].value, "Паллет")
        self.assertEqual(calculation["B7"].value, "PAL-REPORT-1")
        self.assertIn("Палет", [cell.value for cell in workbook["Сводка по дням"][1]])
        self.assertIn("Палет", [cell.value for cell in workbook["Хранение по артикулам"][1]])
        self.assertIn("Палета / откуда", [cell.value for cell in workbook["Движение товара"][1]])

    def test_calculation_sheet_splits_same_box_by_daily_billing_basis(self):
        day = date(2026, 7, 3)
        charge = ApplicationCharge.objects.create(
            application=self.application,
            client=self.client,
            legal_entity=self.client,
            service=self.service,
            quantity=Decimal("0.367"),
            unit="м3",
            tariff=Decimal("30.0000"),
            amount=Decimal("11.00"),
            vat_amount=Decimal("0.55"),
            total_amount=Decimal("11.55"),
            source_type=ApplicationCharge.SOURCE_STORAGE_DAY,
            source_key=f"storage:{self.client.id}:{day.isoformat()}",
            billing_period=day.replace(day=1),
        )
        storage_day = BillingStorageDay.objects.create(
            client=self.client,
            application=self.application,
            day=day,
            billing_mode="m3_day",
            physical_volume_m3=Decimal("0.100000"),
            billable_volume_m3=Decimal("0.366667"),
            amount=Decimal("11.00"),
            vat_amount=Decimal("0.55"),
            charge=charge,
            payload={"weight_volume_m3": "0.366667"},
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            pallet_code="PAL-REPORT-1",
            box_code="BOX-REPORT-1",
            sku_code="SKU-REPORT-1",
            name="Тестовый товар",
            barcode="4600000000001",
            quantity=Decimal("10.000"),
            width_mm=100,
            height_mm=200,
            length_mm=300,
            total_volume_m3=Decimal("1.000000"),
            billable_volume_m3=Decimal("0.366667"),
        )

        days = BillingStorageDay.objects.select_related("client", "application", "charge").order_by("day")
        response = build_manager_storage_export(
            client=self.client,
            start=date(2026, 7, 1),
            end=date(2026, 7, 3),
            storage_days=days,
        )

        workbook = load_workbook(response)
        sheet = workbook["Расчёт по коробам"]
        bases = {sheet["T7"].value, sheet["T8"].value}
        self.assertEqual(bases, {"Вес", "Объём"})
        self.assertNotIn("Вес / Объём", bases)
        self.assertEqual(sheet["C7"].value, "BOX-REPORT-1")
        self.assertEqual(sheet["C8"].value, "BOX-REPORT-1")

    def test_export_endpoint_uses_selected_date_range(self):
        self._create_storage_day(date(2026, 8, 9), Decimal("50.00"))
        user = get_user_model().objects.create_superuser("storage-export-admin", "admin@example.com", "pwd")
        client = DjangoClient()
        client.force_login(user)

        response = client.get(
            reverse("billing-api-storage-export-xlsx"),
            {
                "client": str(self.client.id),
                "date_from": "2026-07-02",
                "date_to": "2026-08-09",
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        overview_dates = [row[0].value for row in workbook["Сводка по дням"].iter_rows(min_row=2)]
        self.assertNotIn("01.07.2026", overview_dates)
        self.assertIn("02.07.2026", overview_dates)
        self.assertIn("09.08.2026", overview_dates)
        notes = workbook["Пояснения"]
        period_row = next(row for row in notes.iter_rows(values_only=True) if row[0] == "Период")
        self.assertEqual(period_row[1], "02.07.2026 — 09.08.2026")

    def test_calculation_sheet_splits_same_box_by_tariff(self):
        day = date(2026, 7, 3)
        charge = ApplicationCharge.objects.create(
            application=self.application,
            client=self.client,
            legal_entity=self.client,
            service=self.service,
            quantity=Decimal("1.000"),
            unit="м3",
            tariff=Decimal("40.0000"),
            amount=Decimal("40.00"),
            vat_amount=Decimal("2.00"),
            total_amount=Decimal("42.00"),
            source_type=ApplicationCharge.SOURCE_STORAGE_DAY,
            source_key=f"storage:{self.client.id}:{day.isoformat()}",
            billing_period=day.replace(day=1),
        )
        storage_day = BillingStorageDay.objects.create(
            client=self.client,
            application=self.application,
            day=day,
            billing_mode="m3_day",
            physical_volume_m3=Decimal("1.000000"),
            billable_volume_m3=Decimal("1.000000"),
            amount=Decimal("40.00"),
            vat_amount=Decimal("2.00"),
            charge=charge,
            payload={"weight_volume_m3": "0.366667"},
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            pallet_code="PAL-REPORT-1",
            box_code="BOX-REPORT-1",
            sku_code="SKU-REPORT-1",
            name="Тестовый товар",
            barcode="4600000000001",
            quantity=Decimal("10.000"),
            width_mm=100,
            height_mm=200,
            length_mm=300,
            total_volume_m3=Decimal("1.000000"),
            billable_volume_m3=Decimal("1.000000"),
        )

        days = BillingStorageDay.objects.select_related("client", "application", "charge").order_by("day")
        response = build_manager_storage_export(
            client=self.client,
            start=date(2026, 7, 1),
            end=date(2026, 7, 3),
            storage_days=days,
        )

        workbook = load_workbook(response)
        sheet = workbook["Расчёт по коробам"]
        tariffs = {sheet["U7"].value, sheet["U8"].value}
        self.assertEqual(tariffs, {30, 40})
        self.assertNotIn("разные", tariffs)
        self.assertEqual(sheet["C7"].value, "BOX-REPORT-1")
        self.assertEqual(sheet["C8"].value, "BOX-REPORT-1")
