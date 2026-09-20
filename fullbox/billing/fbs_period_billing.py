"""Period FBS billing from actual warehouse scan and marking evidence.

The source side is read-only: scan events, labels, FBS orders, items and client
movement requests are never changed here. Only billing snapshots and draft
documents are created after an accountant explicitly requests the report.
"""
from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from types import SimpleNamespace

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import F, Q, Sum
from django.db.models.functions import Coalesce, TruncDate
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from accountant.models import ClientLifecycle
from fbs.models import (
    FbsClientStoragePolicy,
    FbsHandoverBatch,
    FbsHandoverOrderAssignment,
    FbsIntegrationProfile,
    FbsOrderItem,
    FbsOrderLabel,
    FbsOrderStockAllocation,
    FbsPickScanEvent,
    FbsOrderTraceability,
    FbsStockMovement,
    FbsStorageDailyUsage,
)
from marking.codes import marking_code_identity
from sku.models import SKU

from .fbs_standard_shipping import (
    FBS_MARKING_LABEL,
    fbs_rate_billable_quantity,
    sku_liters,
)
from .models import (
    ActStatus,
    BillingApplication,
    BillingService,
    ClientTariffVersion,
    FbsClientRate,
)
from .price_resolver import active_contract, apply_quantity_rules
from .services import BillingWorkflowService, calculate_amounts
from .tariff_services import get_company_tariff


ZERO = Decimal("0")
MONEY = Decimal("0.01")
LITERS = Decimal("0.001")
TITLE_FILL = "15524D"
NAVY = "2D827A"
ORANGE = "2D827A"
PALE_ORANGE = "E4F1EF"
PALE_GREEN = "F7FBFA"
PALE_GRAY = "F7FBFA"
TEXT = "303030"
MUTED = "6B7280"
BORDER = "D7DAE0"
OPERATION_ORDER = {
    FbsClientRate.OP_RECEIVING: 5,
    FbsClientRate.OP_PICKING: 10,
    FbsClientRate.OP_MARKING: 20,
    FbsClientRate.OP_CHZ_CHECK: 25,
    FbsClientRate.OP_SHIPPING: 30,
    FbsClientRate.OP_STORAGE: 40,
}
OPERATION_LABELS = {
    FbsClientRate.OP_RECEIVING: "Приемка FBS",
    FbsClientRate.OP_PICKING: "Подбор FBS",
    FbsClientRate.OP_MARKING: FBS_MARKING_LABEL,
    FbsClientRate.OP_CHZ_CHECK: "Проверка ЧЗ",
    FbsClientRate.OP_SHIPPING: "Доставка/отгрузка FBS",
    FbsClientRate.OP_STORAGE: "Хранение FBS",
}
SERVICE_BY_OPERATION = {
    FbsClientRate.OP_RECEIVING: "fbs_receiving_goods",
    FbsClientRate.OP_PICKING: "fbs_pick_item",
    FbsClientRate.OP_MARKING: "fbs_marking_label_58x40",
    FbsClientRate.OP_CHZ_CHECK: "fbs_honest_sign_check",
    FbsClientRate.OP_SHIPPING: "fbs_shipping_item",
    FbsClientRate.OP_STORAGE: "fbs_storage_liter_day",
}
FBS_STORAGE_LITER_SERVICE = "fbs_storage_liter_day"
FBS_STORAGE_PALLET_SERVICE = "fbs_storage_pallet_day"
OPERATION_UNITS = {
    FbsClientRate.OP_RECEIVING: "шт.",
    FbsClientRate.OP_PICKING: "шт.",
    FbsClientRate.OP_MARKING: "шт.",
    FbsClientRate.OP_CHZ_CHECK: "шт.",
    FbsClientRate.OP_SHIPPING: "шт.",
    FbsClientRate.OP_STORAGE: "литро-дн.",
}
MARKETPLACE_COLUMNS = {
    FbsIntegrationProfile.MARKETPLACE_WB: "wb_qty",
    FbsIntegrationProfile.MARKETPLACE_OZON: "ozon_qty",
}


def _money(value) -> Decimal:
    return Decimal(str(value or "0")).quantize(MONEY, rounding=ROUND_HALF_UP)


def _aware_midnight(value: date):
    return timezone.make_aware(datetime.combine(value, time.min))


def _local_date(value) -> date:
    if timezone.is_aware(value):
        return timezone.localtime(value).date()
    return value.date()


def _rate_label(rate: FbsClientRate) -> str:
    if rate.operation in (FbsClientRate.OP_MARKING, FbsClientRate.OP_CHZ_CHECK):
        return "за штуку"
    if rate.operation == FbsClientRate.OP_STORAGE:
        return "за паллето-день" if "пал" in str(getattr(rate, "unit", "")).lower() else "за литро-день"
    lower = f"{rate.liters_from:g}"
    upper = f"{rate.liters_to:g}" if rate.liters_to is not None else "∞"
    return f"{lower}-{upper} л"


def fbs_applied_label_quantity(*, order) -> Decimal:
    """Bill one 58×40 marking service for a confirmed order-label application."""
    return Decimal("1") if any(
        label.status == FbsOrderLabel.STATUS_APPLIED
        for label in order.marketplace_labels.all()
    ) else ZERO


def _unique_chz_quantities_by_item(rows) -> dict[int, int]:
    """Count one final KIZ check per order and serialized-item identity."""
    seen: set[tuple[int, str]] = set()
    quantities: dict[int, int] = defaultdict(int)
    for row in rows:
        identity = marking_code_identity(row["marking_code"])
        if not identity:
            continue
        key = (row["order_id"], identity)
        if key in seen:
            continue
        seen.add(key)
        quantities[row["order_item_id"]] += 1
    return dict(quantities)


def _first_unique_chz_scan_rows(rows) -> list[dict]:
    """Keep the first successful scan of each Honest Sign identity per order."""
    seen: set[tuple[int, str]] = set()
    unique_rows: list[dict] = []
    for row in rows:
        identity = marking_code_identity(row.get("scan_value"))
        if not identity:
            continue
        key = (row["order_id"], identity)
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)
    return unique_rows


def _missing_label_delivery_quantity(*, ordered_quantity, scanned_quantity) -> Decimal:
    """Return only the unrepresented quantity when an applied label proves packing."""
    return max(
        Decimal(str(ordered_quantity or 0)) - Decimal(str(scanned_quantity or 0)),
        ZERO,
    )


def _pick_scan_quantities(rows) -> tuple[dict[tuple[date, int], Decimal], dict[int, Decimal]]:
    """Count successful FBS controller item scans by service day and order item."""
    quantities_by_day_item: dict[tuple[date, int], Decimal] = defaultdict(lambda: ZERO)
    quantities_by_item: dict[int, Decimal] = defaultdict(lambda: ZERO)
    for row in rows:
        item_id = row.get("item_id")
        if not item_id:
            continue
        quantity = Decimal("1")
        quantities_by_day_item[(_local_date(row["created_at"]), item_id)] += quantity
        quantities_by_item[item_id] += quantity
    return dict(quantities_by_day_item), dict(quantities_by_item)


def _confirmed_chz_quantities_by_item(*, order_ids) -> dict[int, int]:
    if not order_ids:
        return {}
    rows = FbsOrderTraceability.objects.filter(
        allocation__order_item__order_id__in=order_ids,
        status=FbsOrderTraceability.STATUS_PICKED,
        allocation__status=FbsOrderStockAllocation.STATUS_PICKED,
    ).exclude(marking_code="").values(
        "marking_code",
        order_id=F("allocation__order_item__order_id"),
        order_item_id=F("allocation__order_item_id"),
    )
    return _unique_chz_quantities_by_item(rows)


def _rate_for(*, client, operation: str, liters: Decimal | None, on_date: date):
    rates = (
        FbsClientRate.objects.filter(
            client=client,
            operation=operation,
            is_active=True,
            valid_from__lte=on_date,
        )
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=on_date))
    )
    if operation in (
        FbsClientRate.OP_MARKING,
        FbsClientRate.OP_CHZ_CHECK,
        FbsClientRate.OP_STORAGE,
    ):
        return rates.order_by("-valid_from", "liters_from", "-id").first()
    if liters is None:
        return None
    return (
        rates.filter(liters_from__lt=liters)
        .filter(Q(liters_to__isnull=True) | Q(liters_to__gte=liters))
        .order_by("-valid_from", "-liters_from", "-id")
        .first()
    )


def _period_rate_resolver(client):
    """Reuse identical rate lookups while one period report is collected."""
    cache: dict[tuple[str, Decimal | None, date], FbsClientRate | None] = {}
    flat_rate_operations = {
        FbsClientRate.OP_MARKING,
        FbsClientRate.OP_CHZ_CHECK,
        FbsClientRate.OP_STORAGE,
    }

    def resolve(*, operation: str, liters: Decimal | None, on_date: date):
        rate_liters = None if operation in flat_rate_operations else liters
        key = (operation, rate_liters, on_date)
        if key not in cache:
            cache[key] = _rate_for(
                client=client,
                operation=operation,
                liters=rate_liters,
                on_date=on_date,
            )
        return cache[key]

    return resolve


def _tariff_version_for(*, client, on_date: date):
    return (
        ClientTariffVersion.objects.filter(
            client=client,
            status__in=(ClientTariffVersion.STATUS_ACTIVE, ClientTariffVersion.STATUS_SCHEDULED),
            valid_from__lte=on_date,
        )
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=on_date))
        .order_by("-valid_from", "-version_number", "-id")
        .first()
    )


def _client_vat_policy(*, client, on_date: date) -> dict[str, str]:
    """Return the single VAT mode configured on the client's accountant card."""
    lifecycle = (
        ClientLifecycle.objects.select_related("serving_company")
        .filter(agency=client)
        .first()
    )
    tariff_version = _tariff_version_for(client=client, on_date=on_date)
    contract = (
        tariff_version.contract
        if tariff_version is not None and tariff_version.contract_id
        else active_contract(client, on_date)
    )
    vat_type = str(
        getattr(lifecycle, "vat_type", "")
        or getattr(tariff_version, "vat_type", "")
        or ClientTariffVersion.VAT_NO
    ).strip()
    if vat_type not in {
        ClientTariffVersion.VAT_NO,
        ClientTariffVersion.VAT_WITH,
        ClientTariffVersion.VAT_EXTRA,
    }:
        raise ValidationError("В карточке клиента указан неизвестный режим НДС.")
    if vat_type == ClientTariffVersion.VAT_NO:
        return {"vat_type": vat_type, "vat_rate": "0", "caption": "Без НДС"}

    company = (
        getattr(lifecycle, "serving_company", None)
        or (contract.own_company if contract is not None else None)
    )
    if company is None:
        raise ValidationError(
            "Для клиента с НДС не указана обслуживающая компания со ставкой НДС."
        )
    raw_rate = str(getattr(company, "vat_rate", "") or "").strip().rstrip("%")
    try:
        numeric_rate = Decimal(raw_rate.replace(",", "."))
    except (ArithmeticError, ValueError):
        raise ValidationError("У обслуживающей компании клиента указана некорректная ставка НДС.")
    if numeric_rate <= ZERO:
        raise ValidationError("У обслуживающей компании клиента не указана ставка НДС.")
    vat_rate = format(numeric_rate.normalize(), "f")
    suffix = "сверху" if vat_type == ClientTariffVersion.VAT_EXTRA else "в том числе"
    return {
        "vat_type": vat_type,
        "vat_rate": vat_rate,
        "caption": f"{vat_rate}% {suffix}",
    }


def _storage_mode_for(client) -> str:
    policy = FbsClientStoragePolicy.objects.filter(agency=client, is_active=True).first()
    if policy is None:
        return FbsClientStoragePolicy.BILLING_LITERS
    return policy.billing_mode


def _pallet_storage_rate(*, client, on_date: date, quantity: Decimal):
    """Return an FbsClientRate-compatible snapshot from the agreed pallet tariff."""
    service = BillingService.objects.filter(code=FBS_STORAGE_PALLET_SERVICE, is_active=True).first()
    if service is None:
        return None
    tariff = get_company_tariff(client, service, on_date)
    if tariff is None:
        return None
    effective_price = apply_quantity_rules(
        tariff.price,
        quantity,
        coefficient=tariff.coefficient,
        minimum_amount=tariff.minimum_amount,
        calculation_type=tariff.item.calculation_type,
        minimum_quantity=tariff.item.minimum_quantity,
    )
    vat_type = tariff.version.vat_type
    vat_rate = "0" if vat_type == ClientTariffVersion.VAT_NO else str(service.vat_rate or "5")
    unit = tariff.unit.short_name or tariff.unit.name or "пал."
    return SimpleNamespace(
        pk=tariff.item.pk,
        id=tariff.item.pk,
        operation=FbsClientRate.OP_STORAGE,
        price=effective_price,
        unit=unit,
        liters_from=ZERO,
        liters_to=None,
        valid_from=tariff.version.valid_from,
        valid_to=tariff.version.valid_to,
        vat_rate=vat_rate,
        vat_type=vat_type,
        comment=tariff.conditions or tariff.item.description or "Согласованный тариф паллетного хранения FBS",
        tariff_version=tariff.version,
        tariff_item=tariff.item,
        service_code=FBS_STORAGE_PALLET_SERVICE,
    )


def _collect_pallet_storage(*, client, date_from: date, date_to: date) -> dict:
    """Build pallet/day storage from exact snapshots or physical stock events."""
    opening_in_rows = (
        FbsStockMovement.objects.filter(
            target_balance__agency=client,
            occurred_at__date__lt=date_from,
        )
        .values("target_balance__box__pallet_id", "target_balance__box__pallet__pallet_code")
        .annotate(quantity=Sum("qty"))
    )
    period_in_rows = (
        FbsStockMovement.objects.filter(
            target_balance__agency=client,
            occurred_at__date__gte=date_from,
            occurred_at__date__lte=date_to,
        )
        .annotate(event_date=TruncDate("occurred_at"))
        .values(
            "event_date",
            "target_balance__box__pallet_id",
            "target_balance__box__pallet__pallet_code",
        )
        .annotate(quantity=Sum("qty"))
    )
    outgoing = FbsOrderStockAllocation.objects.filter(
        balance__agency=client,
        qty_picked__gt=0,
    ).annotate(event_at=Coalesce("picked_at", "updated_at"))
    opening_out_rows = (
        outgoing.filter(event_at__date__lt=date_from)
        .values("balance__box__pallet_id", "balance__box__pallet__pallet_code")
        .annotate(quantity=Sum("qty_picked"))
    )
    period_out_rows = (
        outgoing.filter(event_at__date__gte=date_from, event_at__date__lte=date_to)
        .annotate(event_date=TruncDate("event_at"))
        .values("event_date", "balance__box__pallet_id", "balance__box__pallet__pallet_code")
        .annotate(quantity=Sum("qty_picked"))
    )
    snapshot_rows = (
        FbsStorageDailyUsage.objects.filter(
            agency=client,
            billing_mode=FbsClientStoragePolicy.BILLING_PALLETS,
            usage_date__gte=date_from,
            usage_date__lte=date_to,
        )
        .values("usage_date", "pallet_id", "pallet__pallet_code")
        .annotate(quantity=Sum("quantity"), pallet_places=Sum("pallet_places"))
    )

    pallet_codes: dict[int, str] = {}
    opening_in: dict[int, Decimal] = defaultdict(lambda: ZERO)
    opening_out: dict[int, Decimal] = defaultdict(lambda: ZERO)
    incoming_by_day: dict[tuple[date, int], Decimal] = defaultdict(lambda: ZERO)
    outgoing_by_day: dict[tuple[date, int], Decimal] = defaultdict(lambda: ZERO)
    snapshots_by_day: dict[date, dict[int, dict]] = defaultdict(dict)

    for row in opening_in_rows:
        pallet_id = row["target_balance__box__pallet_id"]
        if pallet_id is None:
            continue
        pallet_codes[pallet_id] = row["target_balance__box__pallet__pallet_code"] or f"PAL-{pallet_id}"
        opening_in[pallet_id] += Decimal(str(row["quantity"] or 0))
    for row in period_in_rows:
        pallet_id = row["target_balance__box__pallet_id"]
        if pallet_id is None:
            continue
        pallet_codes[pallet_id] = row["target_balance__box__pallet__pallet_code"] or f"PAL-{pallet_id}"
        incoming_by_day[(row["event_date"], pallet_id)] += Decimal(str(row["quantity"] or 0))
    for row in opening_out_rows:
        pallet_id = row["balance__box__pallet_id"]
        if pallet_id is None:
            continue
        pallet_codes[pallet_id] = row["balance__box__pallet__pallet_code"] or f"PAL-{pallet_id}"
        opening_out[pallet_id] += Decimal(str(row["quantity"] or 0))
    for row in period_out_rows:
        pallet_id = row["balance__box__pallet_id"]
        if pallet_id is None:
            continue
        pallet_codes[pallet_id] = row["balance__box__pallet__pallet_code"] or f"PAL-{pallet_id}"
        outgoing_by_day[(row["event_date"], pallet_id)] += Decimal(str(row["quantity"] or 0))
    for row in snapshot_rows:
        pallet_id = row["pallet_id"]
        if pallet_id is None:
            continue
        pallet_codes[pallet_id] = row["pallet__pallet_code"] or f"PAL-{pallet_id}"
        snapshots_by_day[row["usage_date"]][pallet_id] = {
            "quantity": Decimal(str(row["quantity"] or 0)),
            "pallet_places": Decimal(str(row["pallet_places"] or 0)),
        }

    pallet_ids = set(opening_in) | set(opening_out) | set(pallet_codes)
    balances = {
        pallet_id: max(opening_in[pallet_id] - opening_out[pallet_id], ZERO)
        for pallet_id in pallet_ids
    }
    storage_rows = []
    storage_daily = []
    storage_groups: dict[tuple[int, Decimal], dict] = {}
    errors = []
    day = date_from
    while day <= date_to:
        exact_snapshot = snapshots_by_day.get(day)
        day_rows = []
        for pallet_id in sorted(pallet_ids, key=lambda value: pallet_codes.get(value, "")):
            start_qty = balances.get(pallet_id, ZERO)
            incoming = incoming_by_day[(day, pallet_id)]
            shipped = outgoing_by_day[(day, pallet_id)]
            available = start_qty + incoming
            overdraw = max(shipped - available, ZERO)
            reconstructed_end = max(available - shipped, ZERO)
            if exact_snapshot is not None:
                snapshot = exact_snapshot.get(pallet_id)
                end_qty = snapshot["quantity"] if snapshot else ZERO
                pallet_places = snapshot["pallet_places"] if snapshot else ZERO
            else:
                end_qty = reconstructed_end
                pallet_places = Decimal("1.000") if end_qty > 0 else ZERO
            balances[pallet_id] = end_qty
            if any(value > 0 for value in (start_qty, incoming, shipped, end_qty, pallet_places, overdraw)):
                row = {
                    "date": day,
                    "pallet_id": pallet_id,
                    "pallet_code": pallet_codes.get(pallet_id, f"PAL-{pallet_id}"),
                    "name": "Паллетное хранение FBS",
                    "liters": ZERO,
                    "start_qty": start_qty,
                    "incoming": incoming,
                    "shipped": shipped,
                    "end_qty": end_qty,
                    "billable_quantity": pallet_places,
                    "liters_due": pallet_places,
                    "overdraw": overdraw,
                }
                day_rows.append(row)
                storage_rows.append(row)

        total_places = sum((row["billable_quantity"] for row in day_rows), ZERO).quantize(LITERS)
        storage_rate = None
        if total_places > 0:
            storage_rate = _pallet_storage_rate(client=client, on_date=day, quantity=total_places)
            if storage_rate is None:
                errors.append(f"{day:%d.%m.%Y}: нет согласованного тарифа хранения FBS за паллето-день.")
        for row in day_rows:
            row["rate"] = storage_rate
            row["amount"] = _money(row["billable_quantity"] * storage_rate.price) if storage_rate else ZERO
        day_amount = _money(total_places * storage_rate.price) if storage_rate else ZERO
        storage_daily.append(
            {
                "date": day,
                "incoming": sum((row["incoming"] for row in day_rows), ZERO),
                "shipped": sum((row["shipped"] for row in day_rows), ZERO),
                "end_qty": sum((row["end_qty"] for row in day_rows), ZERO),
                "liters": ZERO,
                "billable_quantity": total_places,
                "rate": storage_rate,
                "amount": day_amount,
                "source": "snapshot" if exact_snapshot is not None else "events",
            }
        )
        if storage_rate is not None and total_places > 0:
            group_key = (storage_rate.pk, storage_rate.price)
            group = storage_groups.setdefault(
                group_key,
                {
                    "operation": FbsClientRate.OP_STORAGE,
                    "operation_label": OPERATION_LABELS[FbsClientRate.OP_STORAGE],
                    "rate": storage_rate,
                    "rate_label": _rate_label(storage_rate),
                    "quantity": ZERO,
                    "amount": ZERO,
                },
            )
            group["quantity"] += total_places
            group["amount"] += day_amount
        day += timedelta(days=1)
    return {
        "rows": storage_rows,
        "daily": storage_daily,
        "groups": storage_groups,
        "errors": errors,
    }


@dataclass
class ShipmentRow:
    service_date: date
    marketplace: str
    external_sku: str
    product_name: str
    sku: SKU
    quantity: Decimal
    marking_quantity: Decimal = ZERO
    chz_quantity: Decimal = ZERO
    barcodes: set[str] = field(default_factory=set)
    liters: Decimal = ZERO
    picking_rate: FbsClientRate | None = None
    marking_rate: FbsClientRate | None = None
    chz_rate: FbsClientRate | None = None
    shipping_rate: FbsClientRate | None = None

    @property
    def amount(self) -> Decimal:
        picking_amount = (
            self.quantity * self.picking_rate.price
            if self.quantity > 0 and self.picking_rate is not None
            else ZERO
        )
        shipping_quantity = (
            fbs_rate_billable_quantity(
                operation=FbsClientRate.OP_SHIPPING,
                rate=self.shipping_rate,
                item_quantity=self.quantity,
                liters=self.liters,
            )
            if self.quantity > 0 and self.shipping_rate is not None
            else ZERO
        )
        shipping_amount = (
            shipping_quantity * self.shipping_rate.price
            if self.shipping_rate is not None
            else ZERO
        )
        marking_amount = (
            self.marking_quantity * self.marking_rate.price
            if self.marking_rate is not None
            else ZERO
        )
        chz_amount = (
            self.chz_quantity * self.chz_rate.price
            if self.chz_rate is not None
            else ZERO
        )
        return _money(
            picking_amount
            + marking_amount
            + chz_amount
            + shipping_amount
        )


def _sku_maps(*, client, items) -> tuple[dict[str, SKU], dict[str, SKU]]:
    external_skus = {str(item.external_sku or "").strip() for item in items if item.external_sku}
    barcodes = {str(item.barcode or "").strip() for item in items if item.barcode}
    queryset = (
        SKU.objects.filter(agency=client, deleted=False)
        .filter(Q(sku_code__in=external_skus) | Q(barcodes__value__in=barcodes))
        .prefetch_related("barcodes")
        .distinct()
    )
    by_code: dict[str, SKU] = {}
    by_barcode: dict[str, SKU] = {}
    for sku in queryset:
        by_code[str(sku.sku_code or "").strip().lower()] = sku
        for barcode in sku.barcodes.all():
            value = str(barcode.value or "").strip()
            if value:
                by_barcode[value] = sku
    return by_code, by_barcode


def _extend_sku_maps(
    *,
    client,
    by_code: dict[str, SKU],
    by_barcode: dict[str, SKU],
    sku_codes,
    barcodes,
) -> None:
    cleaned_codes = {str(value or "").strip() for value in sku_codes if str(value or "").strip()}
    cleaned_barcodes = {str(value or "").strip() for value in barcodes if str(value or "").strip()}
    if not cleaned_codes and not cleaned_barcodes:
        return
    queryset = (
        SKU.objects.filter(agency=client, deleted=False)
        .filter(Q(sku_code__in=cleaned_codes) | Q(barcodes__value__in=cleaned_barcodes))
        .prefetch_related("barcodes")
        .distinct()
    )
    for sku in queryset:
        by_code[str(sku.sku_code or "").strip().lower()] = sku
        for barcode in sku.barcodes.all():
            value = str(barcode.value or "").strip()
            if value:
                by_barcode[value] = sku


def _collect_period_data_legacy(*, client, date_from: date, date_to: date) -> dict:
    if date_from > date_to:
        raise ValidationError("Дата начала периода не может быть позже даты окончания.")
    batches = list(
        FbsHandoverBatch.objects.filter(
            profile__agency=client,
            dispatched_at__date__gte=date_from,
            dispatched_at__date__lte=date_to,
            status__in=(FbsHandoverBatch.STATUS_DISPATCHED, FbsHandoverBatch.STATUS_ACCEPTED),
        )
        .select_related("profile")
        .prefetch_related(
            "order_assignments__order__items__sku__barcodes",
            "order_assignments__order__marketplace_labels",
        )
        .order_by("dispatched_at", "id")
    )
    assignments = []
    for batch in batches:
        assignments.extend(
            assignment
            for assignment in batch.order_assignments.all()
            if assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        )
    if not assignments:
        raise ValidationError("За выбранный период нет подтверждённых передач FBS маркетплейсу.")

    items = [item for assignment in assignments for item in assignment.order.items.all()]
    by_code, by_barcode = _sku_maps(client=client, items=items)
    errors: list[str] = []
    details: dict[tuple, ShipmentRow] = {}
    assignment_batch = {assignment.order_id: assignment.batch for assignment in assignments}
    for assignment in assignments:
        order = assignment.order
        batch = assignment_batch[order.id]
        service_date = _local_date(batch.dispatched_at)
        marketplace = batch.profile.marketplace
        marking_quantity_remaining = fbs_applied_label_quantity(order=order)
        for item in order.items.all():
            quantity = Decimal(str(item.quantity or 0))
            if quantity <= 0:
                continue
            sku = item.sku or by_barcode.get(str(item.barcode or "").strip()) or by_code.get(
                str(item.external_sku or "").strip().lower()
            )
            item_label = item.external_sku or item.barcode or f"строка #{item.id}"
            if sku is None:
                errors.append(f"{service_date:%d.%m.%Y}, {item_label}: SKU клиента не найден.")
                continue
            liters = sku_liters(sku)
            if liters is None:
                errors.append(f"{service_date:%d.%m.%Y}, {item_label}: в карточке SKU нет полных габаритов.")
                continue
            marking_quantity = min(quantity, marking_quantity_remaining)
            picking_rate = _rate_for(
                client=client,
                operation=FbsClientRate.OP_PICKING,
                liters=liters,
                on_date=service_date,
            )
            marking_rate = (
                _rate_for(
                    client=client,
                    operation=FbsClientRate.OP_MARKING,
                    liters=None,
                    on_date=service_date,
                )
                if marking_quantity > 0
                else None
            )
            shipping_rate = _rate_for(
                client=client,
                operation=FbsClientRate.OP_SHIPPING,
                liters=liters,
                on_date=service_date,
            )
            missing = []
            if picking_rate is None:
                missing.append("подбор")
            if marking_quantity > 0 and marking_rate is None:
                missing.append(FBS_MARKING_LABEL)
            if shipping_rate is None:
                missing.append("доставка")
            if missing:
                errors.append(
                    f"{service_date:%d.%m.%Y}, {item_label}, {liters.quantize(LITERS)} л: "
                    f"нет ставки ({', '.join(missing)})."
                )
                continue
            key = (service_date, marketplace, str(item.external_sku or ""), sku.pk, item.product_name or sku.name)
            detail = details.get(key)
            if detail is None:
                detail = ShipmentRow(
                    service_date=service_date,
                    marketplace=marketplace,
                    external_sku=str(item.external_sku or ""),
                    product_name=item.product_name or sku.name,
                    sku=sku,
                    quantity=ZERO,
                    marking_quantity=ZERO,
                    liters=liters,
                    picking_rate=picking_rate,
                    marking_rate=marking_rate,
                    shipping_rate=shipping_rate,
                )
                details[key] = detail
            if marking_quantity > 0 and detail.marking_rate is None:
                detail.marking_rate = marking_rate
            detail.quantity += quantity
            detail.marking_quantity += marking_quantity
            marking_quantity_remaining -= marking_quantity
            if item.barcode:
                detail.barcodes.add(str(item.barcode).strip())
            detail.barcodes.update(
                str(barcode.value or "").strip()
                for barcode in sku.barcodes.all()
                if str(barcode.value or "").strip()
            )

    if errors:
        sample = "; ".join(errors[:8])
        suffix = f" Ещё ошибок: {len(errors) - 8}." if len(errors) > 8 else ""
        raise ValidationError(f"Отчёт и счёт не созданы: {sample}.{suffix}")

    detail_rows = sorted(details.values(), key=lambda row: (row.service_date, row.external_sku, row.sku.pk))
    groups: dict[tuple, dict] = {}
    for detail in detail_rows:
        operation_rows = [
            (FbsClientRate.OP_PICKING, detail.picking_rate),
            (FbsClientRate.OP_SHIPPING, detail.shipping_rate),
        ]
        if detail.marking_quantity > 0:
            operation_rows.append((FbsClientRate.OP_MARKING, detail.marking_rate))
        for operation, rate in operation_rows:
            key = (operation, rate.pk)
            group = groups.setdefault(
                key,
                {
                    "operation": operation,
                    "operation_label": OPERATION_LABELS[operation],
                    "rate": rate,
                    "rate_label": _rate_label(rate),
                    "quantity": ZERO,
                    "amount": ZERO,
                },
            )
            billable_quantity = fbs_rate_billable_quantity(
                operation=operation,
                rate=rate,
                item_quantity=(
                    detail.marking_quantity
                    if operation == FbsClientRate.OP_MARKING
                    else detail.quantity
                ),
                liters=detail.liters,
            )
            group["quantity"] += billable_quantity
            group["amount"] += billable_quantity * rate.price
    service_groups = sorted(
        groups.values(),
        key=lambda row: (OPERATION_ORDER[row["operation"]], row["rate"].liters_from, row["rate"].id),
    )
    for group in service_groups:
        group["amount"] = _money(group["amount"])

    operation_totals = {}
    for operation in OPERATION_ORDER:
        matching = [group for group in service_groups if group["operation"] == operation]
        operation_totals[operation] = {
            "quantity": sum((group["quantity"] for group in matching), ZERO),
            "amount": _money(sum((group["amount"] for group in matching), ZERO)),
            "unit": matching[0]["rate"].unit if matching else OPERATION_UNITS.get(operation, "шт."),
            "vat_rate": matching[0]["rate"].vat_rate if matching else "5",
            "vat_type": matching[0]["rate"].vat_type if matching else ClientTariffVersion.VAT_EXTRA,
        }
        if len({(group["rate"].vat_rate, group["rate"].vat_type) for group in matching}) > 1:
            raise ValidationError(
                f"Для операции «{OPERATION_LABELS[operation]}» в периоде заданы разные режимы НДС. Разделите период."
            )

    total_qty = sum((detail.quantity for detail in detail_rows), ZERO)
    subtotal = _money(sum((row["amount"] for row in operation_totals.values()), ZERO))
    vat_amount = _money(
        sum(
            (
                _money(row["amount"] * Decimal(str(row["vat_rate"] or "0")) / Decimal("100"))
                for row in operation_totals.values()
            ),
            ZERO,
        )
    )

    receiving_rows = []
    incoming_by_sku: dict[int, Decimal] = defaultdict(lambda: ZERO)
    movements = (
        FbsClientMovementRequest.objects.filter(
            agency=client,
            status=FbsClientMovementRequest.STATUS_COMPLETED,
            completed_at__date__gte=date_from,
            completed_at__date__lte=date_to,
        )
        .prefetch_related("lines__sku__barcodes")
        .order_by("completed_at", "id")
    )
    for movement in movements:
        for index, line in enumerate(movement.lines.all(), start=1):
            liters = sku_liters(line.sku)
            quantity = Decimal(str(line.requested_qty or 0))
            incoming_by_sku[line.sku_id] += quantity
            receiving_rows.append(
                {
                    "source": movement.number,
                    "date": _local_date(movement.completed_at),
                    "line": index,
                    "external_sku": "",
                    "sku": line.sku,
                    "barcode": line.barcode,
                    "name": line.product_name or line.sku.name,
                    "quantity": quantity,
                    "liters": liters,
                    "status": "Выполнено; справочно, без начисления",
                }
            )

    shipped_by_sku: dict[int, Decimal] = defaultdict(lambda: ZERO)
    sku_examples: dict[int, ShipmentRow] = {}
    for detail in detail_rows:
        shipped_by_sku[detail.sku.pk] += detail.quantity
        sku_examples[detail.sku.pk] = detail
    current_balances = {
        row["sku_ref_id"]: Decimal(str(row["quantity"] or 0))
        for row in (
            FbsStockBalance.objects.filter(agency=client, sku_ref__isnull=False)
            .values("sku_ref_id")
            .annotate(quantity=Sum("qty"))
        )
    }
    reconciliation = []
    for sku_id in sorted(set(incoming_by_sku) | set(shipped_by_sku)):
        example = sku_examples.get(sku_id)
        sku = example.sku if example else next(row["sku"] for row in receiving_rows if row["sku"].pk == sku_id)
        incoming = incoming_by_sku[sku_id]
        shipped = shipped_by_sku[sku_id]
        reconciliation.append(
            {
                "external_sku": example.external_sku if example else "",
                "name": example.product_name if example else sku.name,
                "barcode": ", ".join(sorted(example.barcodes)) if example else "",
                "sku": sku,
                "liters": sku_liters(sku),
                "incoming": incoming,
                "shipped": shipped,
                "period_balance": incoming - shipped,
                "current_fbs_balance": current_balances.get(sku_id, ZERO),
                "status": "OK" if incoming >= shipped else "Приходы периода справочные/неполные",
            }
        )

    rates = list(
        FbsClientRate.objects.filter(
            client=client,
            is_active=True,
            valid_from__lte=date_to,
        )
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=date_from))
        .order_by("operation", "valid_from", "liters_from", "id")
    )
    return {
        "client": client,
        "date_from": date_from,
        "date_to": date_to,
        "batches": batches,
        "orders_count": len(assignments),
        "details": detail_rows,
        "groups": service_groups,
        "operation_totals": operation_totals,
        "receiving_rows": receiving_rows,
        "reconciliation": reconciliation,
        "rates": rates,
        "quantity": total_qty,
        "subtotal": subtotal,
        "vat_amount": vat_amount,
        "total_amount": subtotal + vat_amount,
    }


def collect_period_data(*, client, date_from: date, date_to: date) -> dict:
    """Collect a client-facing FBS report without changing warehouse source data."""
    if date_from > date_to:
        raise ValidationError("Дата начала периода не может быть позже даты окончания.")

    client_vat = _client_vat_policy(client=client, on_date=date_to)
    storage_mode = _storage_mode_for(client)
    period_rate_for = _period_rate_resolver(client)
    end_exclusive = _aware_midnight(date_to + timedelta(days=1))
    successful_events = FbsPickScanEvent.objects.filter(
        allocation__order_item__order__profile__agency=client,
        result=FbsPickScanEvent.RESULT_SUCCESS,
        created_at__lt=end_exclusive,
    )
    pick_scan_rows = list(
        successful_events.filter(stage=FbsPickScanEvent.STAGE_PICK_ITEM)
        .order_by("created_at", "id")
        .values(
            "id",
            "created_at",
            item_id=F("allocation__order_item_id"),
            order_id=F("allocation__order_item__order_id"),
        )
    )
    chz_billing_enabled = (
        FbsClientRate.objects.filter(
            client=client,
            operation=FbsClientRate.OP_CHZ_CHECK,
            is_active=True,
            valid_from__lte=date_to,
        )
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=date_from))
        .exists()
    )
    chz_scan_rows = (
        _first_unique_chz_scan_rows(
            successful_events.filter(stage=FbsPickScanEvent.STAGE_VERIFY_MARKING)
            .order_by("created_at", "id")
            .values(
                "id",
                "created_at",
                "scan_value",
                item_id=F("allocation__order_item_id"),
                order_id=F("allocation__order_item__order_id"),
            )
        )
        if chz_billing_enabled
        else []
    )

    applied_labels = list(
        FbsOrderLabel.objects.filter(
            order__profile__agency=client,
            status=FbsOrderLabel.STATUS_APPLIED,
        )
        .filter(
            Q(applied_at__lt=end_exclusive)
            | Q(applied_at__isnull=True, updated_at__lt=end_exclusive)
        )
        .order_by("order_id", "applied_at", "updated_at", "id")
    )
    first_label_by_order: dict[int, FbsOrderLabel] = {}
    for label in applied_labels:
        current = first_label_by_order.get(label.order_id)
        label_moment = label.applied_at or label.updated_at or label.requested_at
        current_moment = (
            current.applied_at or current.updated_at or current.requested_at
            if current is not None
            else None
        )
        if current is None or label_moment < current_moment:
            first_label_by_order[label.order_id] = label

    event_item_ids = {
        row["item_id"] for row in pick_scan_rows + chz_scan_rows if row.get("item_id")
    }
    label_order_ids = set(first_label_by_order)
    items = list(
        FbsOrderItem.objects.filter(
            Q(pk__in=event_item_ids) | Q(order_id__in=label_order_ids)
        )
        .select_related("order__profile", "sku")
        .prefetch_related("sku__barcodes")
        .order_by("order_id", "id")
    )
    item_by_id = {item.pk: item for item in items}
    items_by_order: dict[int, list[FbsOrderItem]] = defaultdict(list)
    for item in items:
        items_by_order[item.order_id].append(item)

    evidence: dict[tuple[date, int], dict[str, Decimal]] = defaultdict(
        lambda: {
            "quantity": ZERO,
            "marking_quantity": ZERO,
            "chz_quantity": ZERO,
        }
    )
    pick_quantities_by_day_item, pick_counts_by_item = _pick_scan_quantities(pick_scan_rows)
    for key, quantity in pick_quantities_by_day_item.items():
        evidence[key]["quantity"] += quantity
        evidence[key]["marking_quantity"] += quantity

    for order_id, label in first_label_by_order.items():
        label_items = items_by_order.get(order_id, [])
        if not label_items:
            continue
        service_date = _local_date(label.applied_at or label.updated_at or label.requested_at)
        for item in label_items:
            missing_quantity = _missing_label_delivery_quantity(
                ordered_quantity=item.quantity,
                scanned_quantity=pick_counts_by_item[item.pk],
            )
            if missing_quantity > 0:
                evidence[(service_date, item.pk)]["quantity"] += missing_quantity

    for row in chz_scan_rows:
        item_id = row["item_id"]
        if item_id:
            evidence[(_local_date(row["created_at"]), item_id)]["chz_quantity"] += Decimal("1")

    selected_evidence = {
        key: quantities
        for key, quantities in evidence.items()
        if date_from <= key[0] <= date_to
    }
    selected_order_ids = {
        item_by_id[item_id].order_id
        for _service_date, item_id in selected_evidence
        if item_id in item_by_id
    }
    batches = list(
        FbsHandoverBatch.objects.filter(
            profile__agency=client,
            order_assignments__order_id__in=selected_order_ids,
        )
        .distinct()
        .order_by("id")
    )
    by_code, by_barcode = _sku_maps(client=client, items=items)
    errors: list[str] = []

    details: dict[tuple, ShipmentRow] = {}
    for (service_date, item_id), quantities in selected_evidence.items():
        item = item_by_id.get(item_id)
        if item is None:
            errors.append(f"{service_date:%d.%m.%Y}, строка #{item_id}: позиция заказа не найдена.")
            continue
        quantity = quantities["quantity"]
        marking_quantity = quantities["marking_quantity"]
        chz_quantity = quantities["chz_quantity"]
        sku = item.sku or by_barcode.get(str(item.barcode or "").strip()) or by_code.get(
            str(item.external_sku or "").strip().lower()
        )
        item_label = item.external_sku or item.barcode or f"строка #{item.id}"
        if sku is None:
            errors.append(f"{service_date:%d.%m.%Y}, {item_label}: SKU клиента не найден.")
            continue
        liters = sku_liters(sku)
        if liters is None:
            errors.append(f"{service_date:%d.%m.%Y}, {item_label}: в карточке SKU нет полных габаритов.")
            continue
        rated_operations = []
        if quantity > 0:
            rated_operations.extend((FbsClientRate.OP_PICKING, FbsClientRate.OP_SHIPPING))
        if marking_quantity > 0:
            rated_operations.append(FbsClientRate.OP_MARKING)
        if chz_quantity > 0:
            rated_operations.append(FbsClientRate.OP_CHZ_CHECK)
        rates = {
            operation: period_rate_for(operation=operation, liters=liters, on_date=service_date)
            for operation in rated_operations
        }
        missing = [OPERATION_LABELS[operation] for operation, rate in rates.items() if rate is None]
        if missing:
            errors.append(
                f"{service_date:%d.%m.%Y}, {item_label}, {liters.quantize(LITERS)} л: "
                f"нет ставки ({', '.join(missing)})."
            )
            continue
        key = (
            service_date,
            item.order.profile.marketplace,
            str(item.external_sku or ""),
            sku.pk,
            item.product_name or sku.name,
        )
        detail = details.get(key)
        if detail is None:
            detail = ShipmentRow(
                service_date=service_date,
                marketplace=item.order.profile.marketplace,
                external_sku=str(item.external_sku or ""),
                product_name=item.product_name or sku.name,
                sku=sku,
                quantity=ZERO,
                marking_quantity=ZERO,
                chz_quantity=ZERO,
                liters=liters,
                picking_rate=rates.get(FbsClientRate.OP_PICKING),
                marking_rate=rates.get(FbsClientRate.OP_MARKING),
                chz_rate=rates.get(FbsClientRate.OP_CHZ_CHECK),
                shipping_rate=rates.get(FbsClientRate.OP_SHIPPING),
            )
            details[key] = detail
        if quantity > 0:
            detail.picking_rate = detail.picking_rate or rates[FbsClientRate.OP_PICKING]
            detail.shipping_rate = detail.shipping_rate or rates[FbsClientRate.OP_SHIPPING]
        if marking_quantity > 0:
            detail.marking_rate = detail.marking_rate or rates[FbsClientRate.OP_MARKING]
        if chz_quantity > 0:
            detail.chz_rate = detail.chz_rate or rates[FbsClientRate.OP_CHZ_CHECK]
        detail.quantity += quantity
        detail.marking_quantity += marking_quantity
        detail.chz_quantity += chz_quantity
        if item.barcode:
            detail.barcodes.add(str(item.barcode).strip())
        detail.barcodes.update(
            str(barcode.value or "").strip()
            for barcode in sku.barcodes.all()
            if str(barcode.value or "").strip()
        )
    detail_rows = sorted(details.values(), key=lambda row: (row.service_date, row.sku.sku_code, row.external_sku))

    receiving_movements = list(
        FbsStockMovement.objects.filter(
            target_balance__agency=client,
            occurred_at__date__gte=date_from,
            occurred_at__date__lte=date_to,
        )
        .select_related(
            "target_balance__sku_ref",
            "allocation__line__client_movement_line__request",
        )
        .prefetch_related("target_balance__sku_ref__barcodes")
        .order_by("occurred_at", "id")
    )
    receiving_rows = []
    incoming_by_sku: dict[int, Decimal] = defaultdict(lambda: ZERO)
    incoming_by_day_sku: dict[tuple[date, int], Decimal] = defaultdict(lambda: ZERO)
    for movement in receiving_movements:
        sku = movement.target_balance.sku_ref
        service_date = _local_date(movement.occurred_at)
        if sku is None:
            errors.append(
                f"{service_date:%d.%m.%Y}, движение #{movement.pk}: FBS-остаток не связан с карточкой SKU."
            )
            continue
        liters = sku_liters(sku)
        if liters is None:
            errors.append(
                f"{service_date:%d.%m.%Y}, {sku.sku_code}: в карточке SKU нет полных габаритов."
            )
            continue
        rate = period_rate_for(
            operation=FbsClientRate.OP_RECEIVING,
            liters=liters,
            on_date=service_date,
        )
        if rate is None:
            errors.append(
                f"{service_date:%d.%m.%Y}, {sku.sku_code}, {liters.quantize(LITERS)} л: нет ставки приёмки."
            )
            continue
        quantity = Decimal(str(movement.qty or 0))
        movement_line = movement.allocation.line.client_movement_line
        request_row = movement_line.request if movement_line is not None else None
        barcodes = [
            str(barcode.value or "").strip()
            for barcode in sku.barcodes.all()
            if str(barcode.value or "").strip()
        ]
        receiving_rows.append(
            {
                "movement_id": movement.pk,
                "source": request_row.number if request_row is not None else f"FBS-MOVE-{movement.pk}",
                "date": service_date,
                "sku": sku,
                "barcode": movement.target_balance.barcode or (barcodes[0] if barcodes else ""),
                "name": movement.target_balance.name or sku.name,
                "quantity": quantity,
                "liters": liters,
                "rate": rate,
                "amount": _money(quantity * rate.price),
            }
        )
        incoming_by_sku[sku.pk] += quantity
        incoming_by_day_sku[(service_date, sku.pk)] += quantity

    if errors:
        sample = "; ".join(errors[:8])
        suffix = f" Ещё ошибок: {len(errors) - 8}." if len(errors) > 8 else ""
        raise ValidationError(f"Отчёт и счёт не созданы: {sample}.{suffix}")

    sku_cache: dict[int, SKU] = {detail.sku.pk: detail.sku for detail in detail_rows}
    sku_cache.update({row["sku"].pk: row["sku"] for row in receiving_rows})
    if storage_mode == FbsClientStoragePolicy.BILLING_PALLETS:
        prior_arrival_rows = []
        prior_shipping_rows = []
    else:
        prior_arrival_rows = list(
            FbsStockMovement.objects.filter(
                target_balance__agency=client,
                occurred_at__date__lt=date_from,
            )
            .values(
                "target_balance__sku_ref_id",
                "target_balance__sku_code",
                "target_balance__barcode",
            )
            .annotate(quantity=Sum("qty"))
        )
        prior_shipping_rows = [
            {
                "order__items__sku_id": item_by_id[item_id].sku_id,
                "order__items__external_sku": item_by_id[item_id].external_sku,
                "order__items__barcode": item_by_id[item_id].barcode,
                "quantity": quantities["quantity"],
            }
            for (service_date, item_id), quantities in evidence.items()
            if service_date < date_from
            and quantities["quantity"] > 0
            and item_id in item_by_id
        ]
    referenced_ids = {
        row["target_balance__sku_ref_id"] for row in prior_arrival_rows if row["target_balance__sku_ref_id"]
    } | {row["order__items__sku_id"] for row in prior_shipping_rows if row["order__items__sku_id"]}
    for sku in SKU.objects.filter(pk__in=referenced_ids).prefetch_related("barcodes"):
        sku_cache[sku.pk] = sku
        by_code[str(sku.sku_code or "").strip().lower()] = sku
        for barcode in sku.barcodes.all():
            if str(barcode.value or "").strip():
                by_barcode[str(barcode.value).strip()] = sku
    _extend_sku_maps(
        client=client,
        by_code=by_code,
        by_barcode=by_barcode,
        sku_codes=(
            [row["target_balance__sku_code"] for row in prior_arrival_rows]
            + [row["order__items__external_sku"] for row in prior_shipping_rows]
        ),
        barcodes=(
            [row["target_balance__barcode"] for row in prior_arrival_rows]
            + [row["order__items__barcode"] for row in prior_shipping_rows]
        ),
    )
    sku_cache.update({sku.pk: sku for sku in by_code.values()})
    sku_cache.update({sku.pk: sku for sku in by_barcode.values()})

    opening_arrivals: dict[int, Decimal] = defaultdict(lambda: ZERO)
    opening_shipments: dict[int, Decimal] = defaultdict(lambda: ZERO)
    history_errors: list[str] = []
    for row in prior_arrival_rows:
        sku = sku_cache.get(row["target_balance__sku_ref_id"] or 0)
        if sku is None:
            sku = by_barcode.get(str(row["target_balance__barcode"] or "").strip()) or by_code.get(
                str(row["target_balance__sku_code"] or "").strip().lower()
            )
        if sku is None:
            history_errors.append(
                f"приход до периода: {row['target_balance__sku_code'] or row['target_balance__barcode'] or 'без SKU'}"
            )
            continue
        sku_cache[sku.pk] = sku
        opening_arrivals[sku.pk] += Decimal(str(row["quantity"] or 0))
    for row in prior_shipping_rows:
        sku = sku_cache.get(row["order__items__sku_id"] or 0)
        if sku is None:
            sku = by_barcode.get(str(row["order__items__barcode"] or "").strip()) or by_code.get(
                str(row["order__items__external_sku"] or "").strip().lower()
            )
        if sku is None:
            history_errors.append(
                f"отгрузка до периода: {row['order__items__external_sku'] or row['order__items__barcode'] or 'без SKU'}"
            )
            continue
        sku_cache[sku.pk] = sku
        opening_shipments[sku.pk] += Decimal(str(row["quantity"] or 0))
    if history_errors:
        sample = "; ".join(history_errors[:8])
        suffix = f" Ещё ошибок: {len(history_errors) - 8}." if len(history_errors) > 8 else ""
        raise ValidationError(f"Нельзя безопасно восстановить начальный FBS-остаток: {sample}.{suffix}")

    shipped_by_sku: dict[int, Decimal] = defaultdict(lambda: ZERO)
    shipped_by_day_sku: dict[tuple[date, int], Decimal] = defaultdict(lambda: ZERO)
    for detail in detail_rows:
        shipped_by_sku[detail.sku.pk] += detail.quantity
        shipped_by_day_sku[(detail.service_date, detail.sku.pk)] += detail.quantity

    storage_sku_ids = set()
    if storage_mode == FbsClientStoragePolicy.BILLING_LITERS:
        storage_sku_ids = (
            set(opening_arrivals)
            | set(opening_shipments)
            | set(incoming_by_sku)
            | set(shipped_by_sku)
        )
    balances: dict[int, Decimal] = {}
    for sku_id in storage_sku_ids:
        opening = opening_arrivals[sku_id] - opening_shipments[sku_id]
        if opening < 0:
            sku = sku_cache.get(sku_id)
            raise ValidationError(
                f"Нельзя безопасно рассчитать хранение: по SKU {getattr(sku, 'sku_code', sku_id)} "
                f"до начала периода отгружено больше, чем принято ({-opening:g} шт.)."
            )
        balances[sku_id] = opening

    storage_rows = []
    storage_daily = []
    storage_groups: dict[int, dict] = {}
    day = date_from
    while day <= date_to:
        day_rows = []
        for sku_id in sorted(storage_sku_ids, key=lambda value: str(sku_cache[value].sku_code or "").lower()):
            sku = sku_cache[sku_id]
            liters = sku_liters(sku)
            start_qty = balances.get(sku_id, ZERO)
            incoming = incoming_by_day_sku[(day, sku_id)]
            shipped = shipped_by_day_sku[(day, sku_id)]
            available = start_qty + incoming
            overdraw = max(shipped - available, ZERO)
            end_qty = max(available - shipped, ZERO)
            if end_qty > 0 and liters is None:
                errors.append(f"{day:%d.%m.%Y}, {sku.sku_code}: нет полных габаритов для хранения.")
                liters_due = ZERO
            else:
                liters_due = (end_qty * (liters or ZERO)).quantize(LITERS)
            balances[sku_id] = end_qty
            row = {
                "date": day,
                "sku": sku,
                "liters": liters or ZERO,
                "start_qty": start_qty,
                "incoming": incoming,
                "shipped": shipped,
                "end_qty": end_qty,
                "liters_due": liters_due,
                "overdraw": overdraw,
            }
            day_rows.append(row)
            storage_rows.append(row)
        total_liters = sum((row["liters_due"] for row in day_rows), ZERO).quantize(LITERS)
        storage_rate = None
        if total_liters > 0:
            storage_rate = period_rate_for(
                operation=FbsClientRate.OP_STORAGE,
                liters=None,
                on_date=day,
            )
            if storage_rate is None:
                errors.append(f"{day:%d.%m.%Y}: нет ставки хранения FBS за литро-день.")
        for row in day_rows:
            row["rate"] = storage_rate
            row["amount"] = row["liters_due"] * storage_rate.price if storage_rate else ZERO
        day_amount = _money(total_liters * storage_rate.price) if storage_rate else ZERO
        storage_daily.append(
            {
                "date": day,
                "incoming": sum((row["incoming"] for row in day_rows), ZERO),
                "shipped": sum((row["shipped"] for row in day_rows), ZERO),
                "end_qty": sum((row["end_qty"] for row in day_rows), ZERO),
                "liters": total_liters,
                "billable_quantity": total_liters,
                "rate": storage_rate,
                "amount": day_amount,
            }
        )
        if storage_rate is not None and total_liters > 0:
            group = storage_groups.setdefault(
                storage_rate.pk,
                {
                    "operation": FbsClientRate.OP_STORAGE,
                    "operation_label": OPERATION_LABELS[FbsClientRate.OP_STORAGE],
                    "rate": storage_rate,
                    "rate_label": _rate_label(storage_rate),
                    "quantity": ZERO,
                    "amount": ZERO,
                },
            )
            group["quantity"] += total_liters
            group["amount"] += day_amount
        day += timedelta(days=1)
    if storage_mode == FbsClientStoragePolicy.BILLING_PALLETS:
        pallet_storage = _collect_pallet_storage(client=client, date_from=date_from, date_to=date_to)
        storage_rows = pallet_storage["rows"]
        storage_daily = pallet_storage["daily"]
        storage_groups = pallet_storage["groups"]
        errors.extend(pallet_storage["errors"])
    if errors:
        sample = "; ".join(errors[:8])
        suffix = f" Ещё ошибок: {len(errors) - 8}." if len(errors) > 8 else ""
        raise ValidationError(f"Отчёт и счёт не созданы: {sample}.{suffix}")

    groups: dict[tuple, dict] = {}

    def add_group(operation: str, rate: FbsClientRate, quantity: Decimal) -> None:
        key = (operation, rate.pk)
        group = groups.setdefault(
            key,
            {
                "operation": operation,
                "operation_label": OPERATION_LABELS[operation],
                "rate": rate,
                "rate_label": _rate_label(rate),
                "quantity": ZERO,
                "amount": ZERO,
            },
        )
        group["quantity"] += quantity
        group["amount"] += quantity * rate.price

    for detail in detail_rows:
        if detail.quantity > 0:
            add_group(FbsClientRate.OP_PICKING, detail.picking_rate, detail.quantity)
        if detail.marking_quantity > 0:
            add_group(
                FbsClientRate.OP_MARKING,
                detail.marking_rate,
                detail.marking_quantity,
            )
        if detail.chz_quantity > 0:
            add_group(
                FbsClientRate.OP_CHZ_CHECK,
                detail.chz_rate,
                detail.chz_quantity,
            )
        if detail.quantity > 0:
            add_group(
                FbsClientRate.OP_SHIPPING,
                detail.shipping_rate,
                fbs_rate_billable_quantity(
                    operation=FbsClientRate.OP_SHIPPING,
                    rate=detail.shipping_rate,
                    item_quantity=detail.quantity,
                    liters=detail.liters,
                ),
            )
    for row in receiving_rows:
        add_group(FbsClientRate.OP_RECEIVING, row["rate"], row["quantity"])
    groups.update(
        {
            (row["operation"], row["rate"].pk, row["rate"].price): row
            for row in storage_groups.values()
        }
    )
    service_groups = sorted(
        groups.values(),
        key=lambda row: (OPERATION_ORDER[row["operation"]], row["rate"].liters_from, row["rate"].id),
    )
    for group in service_groups:
        group["amount"] = _money(group["amount"])

    operation_totals = {}
    for operation in OPERATION_ORDER:
        matching = [group for group in service_groups if group["operation"] == operation]
        tariff_amount = _money(sum((group["amount"] for group in matching), ZERO))
        amount, operation_vat, total_amount = calculate_amounts(
            Decimal("1"),
            tariff_amount,
            client_vat["vat_rate"],
            vat_type=client_vat["vat_type"],
        )
        operation_totals[operation] = {
            "quantity": sum((group["quantity"] for group in matching), ZERO),
            "tariff_amount": tariff_amount,
            "amount": amount,
            "vat_amount": operation_vat,
            "total_amount": total_amount,
            "unit": matching[0]["rate"].unit if matching else OPERATION_UNITS.get(operation, "шт."),
            "vat_rate": client_vat["vat_rate"],
            "vat_type": client_vat["vat_type"],
        }

    total_qty = sum((detail.quantity for detail in detail_rows), ZERO)
    received_qty = sum((row["quantity"] for row in receiving_rows), ZERO)
    services_subtotal = _money(
        sum(
            (
                operation_totals[operation]["amount"]
                for operation in (
                    FbsClientRate.OP_RECEIVING,
                    FbsClientRate.OP_PICKING,
                    FbsClientRate.OP_MARKING,
                    FbsClientRate.OP_CHZ_CHECK,
                    FbsClientRate.OP_SHIPPING,
                )
            ),
            ZERO,
        )
    )
    storage_subtotal = operation_totals[FbsClientRate.OP_STORAGE]["amount"]
    subtotal = services_subtotal + storage_subtotal
    services_vat = _money(
        sum(
            (
                operation_totals[operation]["vat_amount"]
                for operation in (
                    FbsClientRate.OP_RECEIVING,
                    FbsClientRate.OP_PICKING,
                    FbsClientRate.OP_MARKING,
                    FbsClientRate.OP_CHZ_CHECK,
                    FbsClientRate.OP_SHIPPING,
                )
            ),
            ZERO,
        )
    )
    vat_amount = _money(
        sum(
            (total["vat_amount"] for total in operation_totals.values()),
            ZERO,
        )
    )

    reconciliation = []
    for sku_id in sorted(set(incoming_by_sku) | set(shipped_by_sku), key=lambda value: str(sku_cache[value].sku_code or "")):
        sku = sku_cache[sku_id]
        related_details = [detail for detail in detail_rows if detail.sku.pk == sku_id]
        related_receiving = [row for row in receiving_rows if row["sku"].pk == sku_id]
        reconciliation.append(
            {
                "sku": sku,
                "name": (related_details[0].product_name if related_details else related_receiving[0]["name"]),
                "liters": sku_liters(sku) or ZERO,
                "incoming": incoming_by_sku[sku_id],
                "shipped": shipped_by_sku[sku_id],
                "receiving_amount": _money(sum((row["quantity"] * row["rate"].price for row in related_receiving), ZERO)),
                "marking_amount": _money(
                    sum(
                        (
                            row.marking_quantity * row.marking_rate.price
                            for row in related_details
                            if row.marking_rate is not None
                        ),
                        ZERO,
                    )
                ),
                "chz_amount": _money(
                    sum(
                        (
                            row.chz_quantity * row.chz_rate.price
                            for row in related_details
                            if row.chz_rate is not None
                        ),
                        ZERO,
                    )
                ),
                "picking_amount": _money(
                    sum(
                        (
                            row.quantity * row.picking_rate.price
                            for row in related_details
                            if row.quantity > 0 and row.picking_rate is not None
                        ),
                        ZERO,
                    )
                ),
                "shipping_amount": _money(sum((
                    fbs_rate_billable_quantity(
                        operation=FbsClientRate.OP_SHIPPING,
                        rate=row.shipping_rate,
                        item_quantity=row.quantity,
                        liters=row.liters,
                    ) * row.shipping_rate.price
                    for row in related_details
                    if row.quantity > 0 and row.shipping_rate is not None
                ), ZERO)),
            }
        )

    if not detail_rows and not receiving_rows and not any(
        row["billable_quantity"] > 0 for row in storage_daily
    ):
        raise ValidationError(
            "За выбранный период нет приёмки, фактических сканов/маркировки или хранения FBS."
        )

    selected_pick_scan_event_ids = [
        row["id"] for row in pick_scan_rows if date_from <= _local_date(row["created_at"]) <= date_to
    ]
    selected_chz_scan_event_ids = [
        row["id"] for row in chz_scan_rows if date_from <= _local_date(row["created_at"]) <= date_to
    ]
    selected_label_ids = [
        label.pk
        for label in first_label_by_order.values()
        if date_from <= _local_date(label.applied_at or label.updated_at or label.requested_at) <= date_to
    ]

    rates = list(
        FbsClientRate.objects.filter(client=client, is_active=True, valid_from__lte=date_to)
        .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=date_from))
        .order_by("operation", "valid_from", "liters_from", "id")
    )
    return {
        "client": client,
        "date_from": date_from,
        "date_to": date_to,
        "batches": batches,
        "orders_count": len(selected_order_ids),
        "details": detail_rows,
        "groups": service_groups,
        "operation_totals": operation_totals,
        "receiving_rows": receiving_rows,
        "receiving_movement_ids": [row["movement_id"] for row in receiving_rows],
        "pick_scan_event_ids": selected_pick_scan_event_ids,
        "label_ids": selected_label_ids,
        "chz_scan_event_ids": selected_chz_scan_event_ids,
        "reconciliation": reconciliation,
        "storage_daily": storage_daily,
        "storage_rows": storage_rows,
        "storage_mode": storage_mode,
        "storage_service_code": (
            FBS_STORAGE_PALLET_SERVICE
            if storage_mode == FbsClientStoragePolicy.BILLING_PALLETS
            else FBS_STORAGE_LITER_SERVICE
        ),
        "storage_unit": (
            "паллето-дн."
            if storage_mode == FbsClientStoragePolicy.BILLING_PALLETS
            else "литро-дн."
        ),
        "client_vat": client_vat,
        "rates": rates,
        "quantity": total_qty,
        "marking_quantity": operation_totals[FbsClientRate.OP_MARKING]["quantity"],
        "chz_quantity": operation_totals[FbsClientRate.OP_CHZ_CHECK]["quantity"],
        "received_quantity": received_qty,
        "services_subtotal": services_subtotal,
        "services_vat": services_vat,
        "storage_subtotal": storage_subtotal,
        "storage_quantity_days": sum(
            (row["billable_quantity"] for row in storage_daily), ZERO
        ).quantize(LITERS),
        "storage_liter_days": (
            sum((row["liters"] for row in storage_daily), ZERO).quantize(LITERS)
            if storage_mode == FbsClientStoragePolicy.BILLING_LITERS
            else ZERO
        ),
        "subtotal": subtotal,
        "vat_amount": vat_amount,
        "total_amount": subtotal + vat_amount,
    }


def _application_key(client, date_from: date, date_to: date) -> str:
    return f"FBS-PERIOD-{client.pk}-{date_from:%Y%m%d}-{date_to:%Y%m%d}"


def _source_fingerprint(data: dict) -> str:
    payload = {
        "client_vat": data.get("client_vat", {}),
        "batch_ids": [batch.pk for batch in data["batches"]],
        "receiving_movement_ids": data.get("receiving_movement_ids", []),
        "pick_scan_event_ids": data.get("pick_scan_event_ids", []),
        "label_ids": data.get("label_ids", []),
        "chz_scan_event_ids": data.get("chz_scan_event_ids", []),
        "orders_count": data["orders_count"],
        "quantity": str(data["quantity"]),
        "marking_quantity": str(data.get("marking_quantity", ZERO)),
        "storage_mode": data.get("storage_mode", FbsClientStoragePolicy.BILLING_LITERS),
        "storage_quantity_days": str(data.get("storage_quantity_days", ZERO)),
        "groups": [
            {
                "operation": group["operation"],
                "rate_id": group["rate"].pk,
                "quantity": str(group["quantity"]),
                "amount": str(group["amount"]),
            }
            for group in data["groups"]
        ],
        "storage_daily": [
            {
                "date": row["date"].isoformat(),
                "incoming": str(row["incoming"]),
                "shipped": str(row["shipped"]),
                "end_qty": str(row["end_qty"]),
                "liters": str(row["liters"]),
                "billable_quantity": str(row.get("billable_quantity", row["liters"])),
                "rate_id": row["rate"].pk if row["rate"] else None,
                "amount": str(row["amount"]),
            }
            for row in data.get("storage_daily", [])
        ],
    }
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


@transaction.atomic
def create_period_invoice(*, data: dict, user=None):
    client = data["client"]
    date_from = data["date_from"]
    date_to = data["date_to"]
    application_id = _application_key(client, date_from, date_to)
    source_fingerprint = _source_fingerprint(data)
    application = BillingApplication.objects.filter(
        application_type=BillingApplication.TYPE_FBS,
        application_id=application_id,
        client=client,
    ).first()
    if application is not None:
        invoice = BillingWorkflowService.invoices_for_application(application).order_by("-id").first()
        if invoice is not None:
            previous_fingerprint = (application.source_payload or {}).get("source_fingerprint")
            if (
                invoice.subtotal != data["subtotal"]
                or invoice.vat_amount != data["vat_amount"]
                or invoice.total_amount != data["total_amount"]
                or (
                    previous_fingerprint
                    and previous_fingerprint != source_fingerprint
                )
            ):
                raise ValidationError(
                    f"По периоду уже существует счёт {invoice.number}, но текущая детализация изменилась. "
                    "Отмените старый счёт перед перерасчётом."
                )
            return application, invoice, False

    tariff_version = _tariff_version_for(client=client, on_date=date_to)
    if tariff_version is None:
        raise ValidationError("У клиента нет опубликованной редакции тарифа на конец выбранного периода.")
    contract = active_contract(client, date_to)
    application = BillingWorkflowService.sync_application_from_source(
        application_type=BillingApplication.TYPE_FBS,
        application_id=application_id,
        client=client,
        legal_entity=client,
        own_company=contract.own_company if contract else None,
        operational_status="period_closed",
        operational_status_label="FBS: отчёт за период сформирован",
        created_at_source=_aware_midnight(date_to),
        source_payload={
            "source": "fbs_period_billing",
            "date_from": date_from.isoformat(),
            "date_to": date_to.isoformat(),
            "handover_batch_ids": [batch.pk for batch in data["batches"]],
            "receiving_movement_ids": data.get("receiving_movement_ids", []),
            "pick_scan_event_ids": data.get("pick_scan_event_ids", []),
            "label_ids": data.get("label_ids", []),
            "chz_scan_event_ids": data.get("chz_scan_event_ids", []),
            "orders_count": data["orders_count"],
            "quantity": str(data["quantity"]),
            "marking_quantity": str(data.get("marking_quantity", ZERO)),
            "received_quantity": str(data.get("received_quantity", ZERO)),
            "storage_liter_days": str(data.get("storage_liter_days", ZERO)),
            "storage_mode": data.get("storage_mode", FbsClientStoragePolicy.BILLING_LITERS),
            "storage_quantity_days": str(data.get("storage_quantity_days", ZERO)),
            "source_fingerprint": source_fingerprint,
        },
        user=user,
    )
    active_acts = list(application.acts.exclude(status=ActStatus.CANCELLED).order_by("id"))
    if active_acts:
        if len(active_acts) != 1:
            raise ValidationError("По расчёту уже есть несколько действующих актов; требуется проверка бухгалтера.")
        act = active_acts[0]
    else:
        for operation in OPERATION_ORDER:
            total = data["operation_totals"][operation]
            if total["amount"] <= 0:
                continue
            service_code = (
                data.get("storage_service_code", FBS_STORAGE_LITER_SERVICE)
                if operation == FbsClientRate.OP_STORAGE
                else SERVICE_BY_OPERATION[operation]
            )
            operation_unit = (
                data.get("storage_unit", OPERATION_UNITS[operation])
                if operation == FbsClientRate.OP_STORAGE
                else total.get("unit", OPERATION_UNITS[operation])
            )
            service = BillingService.objects.filter(code=service_code, is_active=True).first()
            if service is None:
                raise ValidationError(f"В биллинге не настроена услуга {service_code}.")
            charge = BillingWorkflowService.create_or_update_charge(
                application,
                service=service,
                quantity=Decimal("1"),
                tariff=total.get("tariff_amount", total["amount"]),
                tariff_price=total.get("tariff_amount", total["amount"]),
                unit="услуга",
                vat_rate=total["vat_rate"],
                vat_type=total["vat_type"],
                source_type="fbs_period",
                source_id=application_id,
                source_key=f"fbs-period:{client.pk}:{date_from}:{date_to}:{operation}",
                performed_at=_aware_midnight(date_to),
                billing_period=date_to.replace(day=1),
                operation_type="fbs_period",
                operation_id=application_id,
                comment=(
                    f"{OPERATION_LABELS[operation]} за {date_from:%d.%m.%Y}–{date_to:%d.%m.%Y}: "
                    f"{total['quantity']:g} {operation_unit}; детализация в отчёте."
                ),
                user=user,
                client_tariff_version=tariff_version,
                tariff_source_label="Индивидуальные ставки FBS клиента",
                tariff_basis=(
                    f"Период {date_from:%d.%m.%Y}–{date_to:%d.%m.%Y}; "
                    f"{total['quantity']:g} {operation_unit}"
                ),
                service_name_snapshot=(
                    f"{OPERATION_LABELS[operation]} ({total['quantity']:g} {operation_unit})"
                ),
                resolve_from_agreed_tariff=False,
            )
            BillingWorkflowService.confirm_charge(
                charge,
                user=user,
                comment="Подтверждено при формировании отчёта и черновика счёта FBS за период.",
            )
        BillingWorkflowService.mark_operations_completed(application, completed_at=_aware_midnight(date_to), user=user)
        act = BillingWorkflowService.generate_act(application, user=user)
        if act.act_date != date_to:
            act.act_date = date_to
            act.save(update_fields=["act_date", "updated_at"])

    invoice = BillingWorkflowService.create_grouped_invoice(
        acts=[act],
        user=user,
        allow_unconfirmed_acts=True,
    )
    if invoice.subtotal != data["subtotal"]:
        raise ValidationError("Итог детализации не совпал с суммой черновика счёта; документы не сохранены.")
    return application, invoice, True


def generate_period_report(*, client, date_from: date, date_to: date) -> dict:
    """Build a current FBS Excel report without creating billing documents."""
    data = collect_period_data(client=client, date_from=date_from, date_to=date_to)
    data["workbook"] = build_period_workbook(data)
    return data


def generate_period_invoice(*, client, date_from: date, date_to: date, user=None) -> dict:
    """Create or reuse the period invoice from the same data used by the report."""
    data = collect_period_data(client=client, date_from=date_from, date_to=date_to)
    application, invoice, created = create_period_invoice(data=data, user=user)
    data["application"] = application
    data["invoice"] = invoice
    data["invoice_created"] = created
    return data


def generate_period_report_and_invoice(*, client, date_from: date, date_to: date, user=None) -> dict:
    """Backward-compatible combined operation for callers outside the FBS billing page."""
    with transaction.atomic():
        data = generate_period_invoice(
            client=client,
            date_from=date_from,
            date_to=date_to,
            user=user,
        )
        data["workbook"] = build_period_workbook(data)
        return data


def _title(sheet, value: str, last_column: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    cell = sheet["A1"]
    cell.value = value
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Manrope", size=15, bold=True, color="FFFFFF")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30


def _note(sheet, value: str, last_column: str) -> None:
    sheet.merge_cells(f"A3:{last_column}3")
    cell = sheet["A3"]
    cell.value = value
    cell.fill = PatternFill("solid", fgColor=PALE_ORANGE)
    cell.font = Font(name="Manrope", size=9, color=TEXT)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[3].height = 34


def _header(sheet, row: int, columns: int) -> None:
    side = Side(style="thin", color=BORDER)
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Manrope", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=side)
    sheet.row_dimensions[row].height = 34


def _body(sheet, start: int, end: int, columns: int) -> None:
    side = Side(style="thin", color=BORDER)
    for row_no in range(start, end + 1):
        fill = PatternFill("solid", fgColor=PALE_GRAY if row_no % 2 == 0 else "FFFFFF")
        for cell in sheet[row_no][:columns]:
            cell.fill = fill
            cell.font = Font(name="Manrope", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=side)


def _finish(sheet, *, freeze="A6") -> None:
    sheet.freeze_panes = freeze
    sheet.sheet_view.showGridLines = False
    for column in range(1, sheet.max_column + 1):
        width = 10
        for cells in sheet.iter_cols(min_col=column, max_col=column, values_only=True):
            for value in cells:
                width = max(width, min(len(str(value or "")) + 2, 48))
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.value = float(cell.value)
                cell.number_format = "#,##0.00"


def _append_total(sheet, label_column: int, value_column: int, label: str, value, *, strong=False) -> None:
    row = sheet.max_row + 1
    sheet.cell(row, label_column, label)
    sheet.cell(row, value_column, value)
    for cell in sheet[row]:
        cell.fill = PatternFill("solid", fgColor=ORANGE if strong else PALE_ORANGE)
        cell.font = Font(name="Manrope", size=10, bold=True, color="FFFFFF" if strong else TEXT)


def _format_columns(sheet, *, text_columns=(), date_columns=(), start_row=6) -> None:
    for column in text_columns:
        for row_no in range(start_row, sheet.max_row + 1):
            cell = sheet.cell(row_no, column)
            cell.number_format = "@"
            cell.quotePrefix = True
    for column in date_columns:
        for row_no in range(start_row, sheet.max_row + 1):
            cell = sheet.cell(row_no, column)
            if isinstance(cell.value, (date, datetime)):
                cell.value = cell.value.strftime("%d.%m.%Y")
            cell.number_format = "@"
            cell.quotePrefix = True


def _vat_caption(data: dict) -> str:
    """Describe the single VAT mode configured on the client's card."""
    client_vat = data.get("client_vat") or {}
    caption = str(client_vat.get("caption") or "").strip()
    if caption:
        return caption
    vat_type = str(client_vat.get("vat_type") or "").strip()
    vat_rate = str(client_vat.get("vat_rate") or "0").strip()
    if vat_type == ClientTariffVersion.VAT_NO or vat_rate == "0":
        return "Без НДС"
    suffix = "сверху" if vat_type == ClientTariffVersion.VAT_EXTRA else "в том числе"
    return f"{vat_rate}% {suffix}"


def _vat_total_label(data: dict) -> str:
    caption = _vat_caption(data)
    return caption if caption == "Без НДС" else f"НДС {caption}"


def _total_payment_label(data: dict, *, services: bool = False) -> str:
    if _vat_caption(data) == "Без НДС":
        return "Итого услуг" if services else "Итого к оплате"
    return "Итого услуг с НДС" if services else "Итого с НДС"


def _build_period_workbook_legacy(data: dict) -> bytes:
    client = data["client"]
    invoice = data["invoice"]
    period = f"{data['date_from']:%d.%m.%Y}–{data['date_to']:%d.%m.%Y}"
    vat_caption = _vat_caption(data)
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Итоги"
    _title(summary, f"FBS: подбор, сканирование ЧЗ и доставка · {client}", "H")
    _note(
        summary,
        f"Период передач маркетплейсу: {period}. Приходы показаны справочно и не включены в счёт. "
        f"Черновик счёта {invoice.number} сформирован одновременно с отчётом.",
        "H",
    )
    summary.append([])
    summary.append(["Клиент", str(client), "", "Счёт", invoice.number, "", "Период", period])
    summary.append(["Заявок FBS", data["orders_count"], "", "Передач", len(data["batches"]), "", "Статус", invoice.get_status_display()])
    summary.append([])
    summary.append(["Услуга", "Основание", "Количество, шт", "Сумма без налога", "Комментарий"])
    _header(summary, 8, 5)
    summary.append(["Приходы FBS (справочно)", "Заявки перемещения FBS", sum((row["quantity"] for row in data["receiving_rows"]), ZERO), ZERO, "Не включены в расчёт"])
    for operation in OPERATION_ORDER:
        total = data["operation_totals"][operation]
        summary.append([OPERATION_LABELS[operation], "Переданные заказы FBS", total["quantity"], total["amount"], f"{total['quantity']:g} шт."])
    _body(summary, 9, 12, 5)
    _append_total(summary, 3, 4, "Итого без налога", invoice.subtotal)
    _append_total(summary, 3, 4, _vat_total_label(data), invoice.vat_amount)
    _append_total(summary, 3, 4, "ИТОГО К ОПЛАТЕ", invoice.total_amount, strong=True)
    summary.append([])
    summary.append(["Контрольные проверки", "", "", ""])
    summary.append(["Приходы исключены из расчёта", "", "", "OK"])
    summary.append(["Сумма отчёта совпадает со счётом", "", "", "OK" if invoice.subtotal == data["subtotal"] else "ПРОВЕРИТЬ"])
    summary.append(["Все переданные SKU рассчитаны", "", "", "OK"])
    for row_no in range(summary.max_row - 2, summary.max_row + 1):
        summary.cell(row_no, 4).fill = PatternFill("solid", fgColor=PALE_GREEN)
        summary.cell(row_no, 4).font = Font(name="Manrope", bold=True, color=TEXT)
    _finish(summary, freeze="A8")

    services = workbook.create_sheet("Услуги по КП")
    _title(services, f"Расчёт услуг по тарифам · {client}", "E")
    _note(services, "Количество рассчитано по фактически переданным заказам FBS; ставки — из действующих ставок клиента.", "E")
    services.append([])
    services.append(["Вид услуги", "Диапазон объёма", "Количество, шт", "Ставка", "Сумма без налога"])
    _header(services, 5, 5)
    for group in data["groups"]:
        services.append([group["operation_label"], group["rate_label"], group["quantity"], group["rate"].price, group["amount"]])
    _body(services, 6, services.max_row, 5)
    _append_total(services, 4, 5, "Итого без налога", invoice.subtotal)
    _append_total(services, 4, 5, _vat_total_label(data), invoice.vat_amount)
    _append_total(services, 4, 5, "ИТОГО К ОПЛАТЕ", invoice.total_amount, strong=True)
    _finish(services)

    receiving = workbook.create_sheet("Приходы")
    _title(receiving, f"Приходы FBS · {client}", "R")
    _note(receiving, "Источник: завершённые заявки перемещения клиента в FBS-зону. Приёмка показана справочно и в счёт не включается.", "R")
    receiving.append([])
    receiving.append([
        "Источник", "Дата", "№", "Артикул МП", "SKU WMS", "Штрихкод", "Наименование", "Количество",
        "Длина, мм", "Ширина, мм", "Высота, мм", "Объём/шт, л", "Диапазон", "Приёмка в расчёте",
        "Сумма приёмки", "Статус", "", "",
    ])
    _header(receiving, 5, 18)
    for row in data["receiving_rows"]:
        sku = row["sku"]
        receiving.append([
            row["source"], row["date"], row["line"], row["external_sku"], sku.sku_code, row["barcode"], row["name"],
            row["quantity"], sku.length_mm or "", sku.width_mm or "", sku.height_mm or "", row["liters"] or "", "", "Нет", ZERO,
            row["status"], "", "",
        ])
    if data["receiving_rows"]:
        _body(receiving, 6, receiving.max_row, 18)
    _append_total(receiving, 7, 8, "ИТОГО", sum((row["quantity"] for row in data["receiving_rows"]), ZERO))
    _format_columns(receiving, text_columns=(4, 5, 6), date_columns=(2,))
    _finish(receiving)

    details = workbook.create_sheet("Отгрузка детально")
    _title(details, "Отгрузка FBS · детализация", "R")
    _note(details, f"Источник: подтверждённые передачи FBS маркетплейсу за {period}; заказов {data['orders_count']}, товаров {data['quantity']:g} шт.", "R")
    details.append([])
    details.append([
        "Дата", "Товар", "Артикул МП", "Штрихкоды", "WB", "OZON", "Yandex", "Всего", "SKU WMS",
        "Длина, мм", "Ширина, мм", "Высота, мм", "Объём/шт, л", "Диапазон", "Подбор, ₽/шт",
        "Маркировка ШК 58×40, ₽/шт", "Доставка, ₽/шт", "Сумма без налога",
    ])
    _header(details, 5, 18)
    for row in data["details"]:
        marketplace_values = {"wb_qty": ZERO, "ozon_qty": ZERO, "yandex_qty": ZERO}
        marketplace_values[MARKETPLACE_COLUMNS.get(row.marketplace, "yandex_qty")] += row.quantity
        details.append([
            row.service_date, row.product_name, row.external_sku, ", ".join(sorted(row.barcodes)), marketplace_values["wb_qty"],
            marketplace_values["ozon_qty"], marketplace_values["yandex_qty"], row.quantity, row.sku.sku_code,
            row.sku.length_mm, row.sku.width_mm, row.sku.height_mm, row.liters, _rate_label(row.picking_rate),
            row.picking_rate.price,
            row.marking_rate.price if row.marking_rate is not None else ZERO,
            row.shipping_rate.price,
            row.amount,
        ])
    _body(details, 6, details.max_row, 18)
    _append_total(details, 7, 8, "ИТОГО", data["quantity"])
    details.cell(details.max_row, 18, data["subtotal"])
    _format_columns(details, text_columns=(3, 4, 9), date_columns=(1,))
    _finish(details)

    reconciliation = workbook.create_sheet("Сверка SKU")
    _title(reconciliation, "Сверка SKU · приходы / отгрузки / FBS", "M")
    _note(reconciliation, "Приходы периода справочные; текущий остаток FBS взят из оперативного контура только для сверки и не изменяется отчётом.", "M")
    reconciliation.append([])
    reconciliation.append([
        "Артикул МП", "Наименование", "Штрихкод", "SKU WMS", "Длина, мм", "Ширина, мм", "Высота, мм",
        "Объём/шт, л", "Приход", "Отгружено", "Баланс периода", "Текущий FBS-остаток", "Статус",
    ])
    _header(reconciliation, 5, 13)
    for row in data["reconciliation"]:
        sku = row["sku"]
        reconciliation.append([
            row["external_sku"], row["name"], row["barcode"], sku.sku_code, sku.length_mm or "", sku.width_mm or "",
            sku.height_mm or "", row["liters"] or "", row["incoming"], row["shipped"], row["period_balance"],
            row["current_fbs_balance"], row["status"],
        ])
    _body(reconciliation, 6, reconciliation.max_row, 13)
    _append_total(reconciliation, 8, 9, "ИТОГО", sum((row["incoming"] for row in data["reconciliation"]), ZERO))
    reconciliation.cell(reconciliation.max_row, 10, sum((row["shipped"] for row in data["reconciliation"]), ZERO))
    reconciliation.cell(reconciliation.max_row, 11, sum((row["period_balance"] for row in data["reconciliation"]), ZERO))
    reconciliation.cell(reconciliation.max_row, 12, sum((row["current_fbs_balance"] for row in data["reconciliation"]), ZERO))
    _format_columns(reconciliation, text_columns=(1, 3, 4))
    _finish(reconciliation)

    rates = workbook.create_sheet("Ставки КП")
    _title(rates, f"Тарифы FBS · {client}", "J")
    _note(
        rates,
        "Действующие ставки клиента, пересекающие выбранный период. "
        "В расчёт вошли подбор, сканирование ЧЗ перед отгрузкой и доставка.",
        "J",
    )
    rates.append([])
    rates.append(["Операция", "Диапазон", "От, л", "До, л", "Ставка", "Ед.", "Действует с", "Действует до", "НДС", "Комментарий"])
    _header(rates, 5, 10)
    for rate in data["rates"]:
        rates.append([
            OPERATION_LABELS.get(rate.operation, rate.get_operation_display()),
            _rate_label(rate), rate.liters_from, rate.liters_to if rate.liters_to is not None else "",
            rate.price, rate.unit, rate.valid_from, rate.valid_to or "", f"{rate.vat_rate}% сверху", rate.comment,
        ])
    _body(rates, 6, rates.max_row, 10)
    _format_columns(rates, date_columns=(7, 8))
    _finish(rates)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_period_workbook(data: dict) -> bytes:
    """Build the six-sheet client report in the agreed FBS workbook format."""
    client = data["client"]
    pallet_storage = data.get("storage_mode") == FbsClientStoragePolicy.BILLING_PALLETS
    period = f"{data['date_from']:%Y-%m-%d} — {data['date_to']:%Y-%m-%d}"
    vat_caption = _vat_caption(data)
    workbook = Workbook()

    thin = Side(style="thin", color=BORDER)

    def style_header(sheet, columns: int) -> None:
        for cell in sheet[1][:columns]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Carlito", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}{max(sheet.max_row, 1)}"
        sheet.sheet_view.showGridLines = False

    def style_body(sheet, start_row: int, columns: int) -> None:
        for row_no in range(start_row, sheet.max_row + 1):
            fill = PatternFill("solid", fgColor=PALE_GRAY if row_no % 2 == 0 else "FFFFFF")
            for cell in sheet[row_no][:columns]:
                cell.fill = fill
                cell.font = Font(name="Carlito", size=11, color=TEXT)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, Decimal):
                    cell.value = float(cell.value)

    def fit_columns(sheet, widths: dict[int, float] | None = None) -> None:
        widths = widths or {}
        for column in range(1, sheet.max_column + 1):
            if column in widths:
                width = widths[column]
            else:
                width = 10
                for row in range(1, min(sheet.max_row, 300) + 1):
                    value = sheet.cell(row, column).value
                    width = max(width, min(len(str(value or "")) + 2, 42))
            sheet.column_dimensions[get_column_letter(column)].width = width

    summary = workbook.active
    summary.title = "Итоги"
    summary.merge_cells("A1:H1")
    summary["A1"] = (
        f"FBS: приемка, подбор, маркировка, Честный знак и доставка по фактическим сканам {client}"
    )
    summary["A1"].fill = PatternFill("solid", fgColor=TITLE_FILL)
    summary["A1"].font = Font(name="Carlito", size=15, bold=True, color="FFFFFF")
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30
    left_rows = [
        ("Клиент", str(client)),
        ("ИНН", getattr(client, "inn", "") or ""),
        ("Период", period),
        ("Заказов со сканами/маркировкой", data["orders_count"]),
        ("Связанных партий", len(data["batches"])),
        ("Принято, шт", data["received_quantity"]),
        ("Подобрано сканером FBS, шт", data["operation_totals"][FbsClientRate.OP_PICKING]["quantity"]),
        ("Маркировка по сканам FBS, шт", data["operation_totals"][FbsClientRate.OP_MARKING]["quantity"]),
        ("Проверено ЧЗ по сканам Data Matrix, шт", data["operation_totals"][FbsClientRate.OP_CHZ_CHECK]["quantity"]),
        ("В доставку по сканам/маркировке, шт", data["quantity"]),
        ("SKU в отчете", len(data["reconciliation"])),
        ("Проблем сопоставления", 0),
        ("НДС", vat_caption),
    ]
    for row_no, (label, value) in enumerate(left_rows, start=2):
        summary.cell(row_no, 1, label)
        summary.cell(row_no, 2, value)
    summary.append([])
    summary["D2"] = "Блок услуги"
    summary["E2"] = "Кол-во"
    summary["F2"] = "Сумма без НДС"
    service_summary = [
        (OPERATION_LABELS[FbsClientRate.OP_SHIPPING], data["operation_totals"][FbsClientRate.OP_SHIPPING]),
        (OPERATION_LABELS[FbsClientRate.OP_MARKING], data["operation_totals"][FbsClientRate.OP_MARKING]),
    ]
    service_summary.extend(
        [
            (OPERATION_LABELS[FbsClientRate.OP_CHZ_CHECK], data["operation_totals"][FbsClientRate.OP_CHZ_CHECK]),
            (OPERATION_LABELS[FbsClientRate.OP_PICKING], data["operation_totals"][FbsClientRate.OP_PICKING]),
            (OPERATION_LABELS[FbsClientRate.OP_RECEIVING], data["operation_totals"][FbsClientRate.OP_RECEIVING]),
        ]
    )
    summary_rows = [
        (row_no, label, total["quantity"], total["amount"])
        for row_no, (label, total) in enumerate(service_summary, start=3)
    ]
    total_start = 3 + len(service_summary)
    summary_rows.extend(
        [
            (total_start, "Итого услуг без НДС", "", data["services_subtotal"]),
            (total_start + 1, _vat_total_label(data), "", data["services_vat"]),
            (total_start + 2, _total_payment_label(data, services=True), "", data["services_subtotal"] + data["services_vat"]),
            (
                total_start + 3,
                "Хранение (паллеты)" if pallet_storage else "Хранение (литры)",
                data.get("storage_quantity_days", data.get("storage_liter_days", ZERO)),
                data["storage_subtotal"],
            ),
            (total_start + 4, "Итого услуг + хранение без НДС", "", data["subtotal"]),
            (total_start + 5, _vat_total_label(data), "", data["vat_amount"]),
            (total_start + 6, _total_payment_label(data), "", data["total_amount"]),
        ]
    )
    for row_no, label, quantity, amount in summary_rows:
        summary.cell(row_no, 4, label)
        summary.cell(row_no, 5, quantity)
        summary.cell(row_no, 6, amount)
    for cell in summary[2][3:6]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Carlito", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    strong_rows = (total_start, total_start + 2, total_start + 4, total_start + 6)
    last_summary_row = total_start + 6
    for row_no in range(2, last_summary_row + 1):
        for column in range(1, 7):
            cell = summary.cell(row_no, column)
            if cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor=PALE_GREEN if row_no in strong_rows else "FFFFFF")
            cell.font = cell.font.copy(name="Carlito", size=11, bold=(row_no in strong_rows or column in (1, 4)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, Decimal):
                cell.value = float(cell.value)
    for row_no in range(3, last_summary_row + 1):
        summary.cell(row_no, 6).number_format = '#,##0.00'
    summary.sheet_view.showGridLines = False
    fit_columns(summary, {1: 28, 2: 40, 4: 38, 5: 14, 6: 20})

    services = workbook.create_sheet("Услуги по КП")
    services.append(["Тип услуги", "Основание тарифа", "Количество", "Ставка", "Сумма без НДС"])
    report_order = [
        FbsClientRate.OP_SHIPPING,
        FbsClientRate.OP_MARKING,
        FbsClientRate.OP_CHZ_CHECK,
        FbsClientRate.OP_PICKING,
        FbsClientRate.OP_RECEIVING,
    ]
    for operation in report_order:
        operation_groups = sorted(
            (row for row in data["groups"] if row["operation"] == operation),
            key=lambda row: row["rate_label"],
        )
        for group in operation_groups:
            services.append(
                [group["operation_label"], group["rate_label"], group["quantity"], group["rate"].price, group["amount"]]
            )
    style_header(services, 5)
    style_body(services, 2, 5)
    for row_no in range(2, services.max_row + 1):
        services.cell(row_no, 3).number_format = '#,##0.###'
        services.cell(row_no, 4).number_format = '#,##0.####'
        services.cell(row_no, 5).number_format = '#,##0.00'
    fit_columns(services, {1: 34, 2: 22, 3: 18, 4: 14, 5: 20})

    details = workbook.create_sheet("Отгрузка детально")
    detail_header = [
        "Дата", "Товар", "SKU", "ШК", "WB", "OZON", "Yandex", "Всего, шт", "Артикул МП",
        "Длина, мм", "Ширина, мм", "Высота, мм", "Литров / шт", "Статус сопоставления",
        "Диапазон подбора", "Подбор, ставка", "Подбор, сумма",
    ]
    details.append(detail_header)
    for row in data["details"]:
        marketplace_values = {"wb_qty": ZERO, "ozon_qty": ZERO, "yandex_qty": ZERO}
        marketplace_values[MARKETPLACE_COLUMNS.get(row.marketplace, "yandex_qty")] += row.quantity
        detail_values = [
            row.service_date,
            row.product_name,
            row.sku.sku_code,
            ", ".join(sorted(row.barcodes)),
            marketplace_values["wb_qty"],
            marketplace_values["ozon_qty"],
            marketplace_values["yandex_qty"],
            row.quantity,
            row.external_sku,
            row.sku.length_mm,
            row.sku.width_mm,
            row.sku.height_mm,
            row.liters,
            "OK, SKU сопоставлен",
            _rate_label(row.picking_rate) if row.picking_rate is not None else "",
            row.picking_rate.price if row.picking_rate is not None else ZERO,
            _money(row.quantity * row.picking_rate.price) if row.picking_rate is not None else ZERO,
        ]
        details.append(detail_values)
    detail_columns = 17
    style_header(details, detail_columns)
    style_body(details, 2, detail_columns)
    for row_no in range(2, details.max_row + 1):
        details.cell(row_no, 1).number_format = 'DD.MM.YYYY'
        for column in (3, 4, 9):
            details.cell(row_no, column).number_format = '@'
            details.cell(row_no, column).quotePrefix = True
        for column in (5, 6, 7, 8, 10, 11, 12):
            details.cell(row_no, column).number_format = '#,##0'
        details.cell(row_no, 13).number_format = '#,##0.000'
        details.cell(row_no, 16).number_format = '#,##0.####'
        details.cell(row_no, 17).number_format = '#,##0.00'
    fit_columns(details, {1: 13, 2: 46, 3: 18, 4: 22, 9: 22, 14: 25, 15: 20})

    reconciliation = workbook.create_sheet("Сверка SKU")
    reconciliation_header = [
        "SKU", "Товар", "Литров / шт", "Принято, шт", "Отгружено, шт",
        "Приемка", FBS_MARKING_LABEL, "Проверка ЧЗ", "Подбор", "Доставка",
    ]
    reconciliation.append(reconciliation_header)
    for row in data["reconciliation"]:
        reconciliation_values = [
            row["sku"].sku_code,
            row["name"],
            row["liters"],
            row["incoming"],
            row["shipped"],
            row["receiving_amount"],
            row["marking_amount"],
            row["chz_amount"],
            row["picking_amount"],
            row["shipping_amount"],
        ]
        reconciliation.append(reconciliation_values)
    reconciliation_columns = 10
    style_header(reconciliation, reconciliation_columns)
    style_body(reconciliation, 2, reconciliation_columns)
    for row_no in range(2, reconciliation.max_row + 1):
        reconciliation.cell(row_no, 1).number_format = '@'
        reconciliation.cell(row_no, 1).quotePrefix = True
        reconciliation.cell(row_no, 3).number_format = '#,##0.000'
        reconciliation.cell(row_no, 4).number_format = '#,##0'
        reconciliation.cell(row_no, 5).number_format = '#,##0'
        for column in range(6, reconciliation_columns + 1):
            reconciliation.cell(row_no, column).number_format = '#,##0.00'
    fit_columns(reconciliation, {1: 20, 2: 48, 3: 16, 4: 16, 5: 18})

    storage = workbook.create_sheet("Хранение")
    storage_quantity_header = "Паллето-мест" if pallet_storage else "Хранимый объем, л"
    storage_rate_header = "Ставка, ₽/паллето-день" if pallet_storage else "Ставка, ₽/л/день"
    storage_total_label = "Паллето-дней" if pallet_storage else "Литро-дней"
    storage.append([
        "Дата", "Приемка, шт", "Отгрузка, шт", "Остаток, шт", storage_quantity_header,
        storage_rate_header, "Хранение за день, ₽", "", "Показатель", "Значение",
    ])
    for row in data["storage_daily"]:
        storage.append([
            row["date"], row["incoming"], row["shipped"], row["end_qty"],
            row.get("billable_quantity", row["liters"]),
            row["rate"].price if row["rate"] else ZERO, row["amount"], "", "", "",
        ])
    storage["I2"], storage["J2"] = "Период", period
    storage_rates = {row["rate"].price for row in data["storage_daily"] if row["rate"] is not None}
    storage["I3"], storage["J3"] = storage_rate_header, (next(iter(storage_rates)) if len(storage_rates) == 1 else "по дням")
    storage["I4"], storage["J4"] = storage_total_label, data.get(
        "storage_quantity_days", data.get("storage_liter_days", ZERO)
    )
    storage["I5"], storage["J5"] = "Итого хранение, ₽", data["storage_subtotal"]
    style_header(storage, 10)
    style_body(storage, 2, 10)
    for row_no in range(2, storage.max_row + 1):
        storage.cell(row_no, 1).number_format = 'DD.MM.YYYY'
        for column in (2, 3, 4):
            storage.cell(row_no, column).number_format = '#,##0'
        storage.cell(row_no, 5).number_format = '#,##0.000'
        storage.cell(row_no, 6).number_format = '#,##0.####'
        storage.cell(row_no, 7).number_format = '#,##0.00'
    storage["J4"].number_format = '#,##0.000'
    storage["J5"].number_format = '#,##0.00'
    fit_columns(storage, {1: 13, 2: 17, 3: 18, 4: 16, 5: 23, 6: 22, 7: 24, 8: 3, 9: 26, 10: 28})

    storage_details = workbook.create_sheet("Хранение детально")
    storage_identity_header = "Паллета" if pallet_storage else "SKU"
    storage_unit_header = "Единица хранения" if pallet_storage else "Литров / шт"
    storage_billable_header = "К оплате, паллето-мест" if pallet_storage else "К оплате, литров"
    storage_details.append([
        "Дата", storage_identity_header, "Товар / вид хранения", storage_unit_header,
        "Остаток начало, шт", "Приемка, шт",
        "Отгрузка, шт", "Остаток конец, шт", storage_billable_header, storage_rate_header,
        "Хранение за день, ₽", "Отгрузка сверх остатка, шт",
    ])
    for row in data["storage_rows"]:
        identity = row["pallet_code"] if pallet_storage else row["sku"].sku_code
        name = row["name"] if pallet_storage else row["sku"].name
        storage_unit_value = "паллето-место" if pallet_storage else row["liters"]
        storage_details.append([
            row["date"], identity, name, storage_unit_value, row["start_qty"],
            row["incoming"], row["shipped"], row["end_qty"],
            row.get("billable_quantity", row["liters_due"]),
            row["rate"].price if row["rate"] else ZERO, row["amount"], row["overdraw"],
        ])
    style_header(storage_details, 12)
    style_body(storage_details, 2, 12)
    for row_no in range(2, storage_details.max_row + 1):
        storage_details.cell(row_no, 1).number_format = 'DD.MM.YYYY'
        storage_details.cell(row_no, 2).number_format = '@'
        storage_details.cell(row_no, 2).quotePrefix = True
        storage_details.cell(row_no, 4).number_format = '@' if pallet_storage else '#,##0.000'
        for column in (5, 6, 7, 8, 12):
            storage_details.cell(row_no, column).number_format = '#,##0'
        storage_details.cell(row_no, 9).number_format = '#,##0.000'
        storage_details.cell(row_no, 10).number_format = '#,##0.####'
        storage_details.cell(row_no, 11).number_format = '#,##0.00####'
    fit_columns(storage_details, {1: 13, 2: 20, 3: 48, 4: 16, 5: 22, 6: 16, 7: 18, 8: 22, 9: 20, 10: 22, 11: 24, 12: 30})

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
