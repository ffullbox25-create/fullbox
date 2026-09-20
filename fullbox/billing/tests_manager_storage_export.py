from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from sku.models import Agency

from .manager_storage_export import build_manager_storage_export
from .models import BillingStorageDay, StorageSnapshotLine


class ManagerStoragePalletExportTests(TestCase):
    def test_pallet_report_allocates_by_pallet_and_shows_saved_box_counts(self):
        client = Agency.objects.create(agn_name="Тест паллетного хранения")
        storage_day = BillingStorageDay.objects.create(
            client=client,
            day=date(2026, 8, 28),
            pallet_count=2,
            box_count=15,
            billing_mode="pallet_day",
            amount=Decimal("140.00"),
            vat_amount=Decimal("0.00"),
            status=BillingStorageDay.STATUS_CALCULATED,
            payload={
                "charge_basis": "pallet",
                "pallet_box_counts": {"PALLET-TINY": 3, "PALLET-LARGE": 12},
            },
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            pallet_code="PALLET-TINY",
            zone_code="OTG",
            cell_code="OTG-01",
            quantity=1,
            total_volume_m3=Decimal("0.000001"),
            status=StorageSnapshotLine.STATUS_OK,
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            pallet_code="PALLET-LARGE",
            zone_code="OS",
            cell_code="OS-01",
            quantity=1,
            total_volume_m3=Decimal("6.000000"),
            status=StorageSnapshotLine.STATUS_OK,
        )

        response = build_manager_storage_export(
            client=client,
            start=storage_day.day,
            end=storage_day.day,
            storage_days=BillingStorageDay.objects.filter(pk=storage_day.pk),
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)

        self.assertEqual(workbook.sheetnames[0], "Расчёт по паллетам")
        self.assertIn("Паллеты по дням", workbook.sheetnames)
        self.assertNotIn("Хранение по артикулам", workbook.sheetnames)

        calculation = workbook["Расчёт по паллетам"]
        headers = [calculation.cell(row=6, column=column).value for column in range(1, 15)]
        self.assertIn("Коробов на паллете", headers)
        self.assertNotIn("Объём строки, м³", headers)

        rows = {
            calculation.cell(row=row, column=2).value: {
                "boxes": calculation.cell(row=row, column=3).value,
                "days": calculation.cell(row=row, column=8).value,
                "per_day": calculation.cell(row=row, column=10).value,
                "net": calculation.cell(row=row, column=11).value,
            }
            for row in range(7, calculation.max_row + 1)
        }
        self.assertEqual(rows["PALLET-TINY"], {"boxes": 3, "days": 1, "per_day": 70, "net": 70})
        self.assertEqual(rows["PALLET-LARGE"], {"boxes": 12, "days": 1, "per_day": 70, "net": 70})
        self.assertEqual(sum(row["net"] for row in rows.values()), 140)

        daily = workbook["Паллеты по дням"]
        daily_rows = {
            daily.cell(row=row, column=2).value: daily.cell(row=row, column=3).value
            for row in range(2, daily.max_row + 1)
        }
        self.assertEqual(daily_rows, {"PALLET-LARGE": 12, "PALLET-TINY": 3})

    def test_volume_report_keeps_volume_allocation_and_box_layout(self):
        client = Agency.objects.create(agn_name="Тест объёмного хранения")
        storage_day = BillingStorageDay.objects.create(
            client=client,
            day=date(2026, 8, 28),
            box_count=2,
            billing_mode="m3_day",
            amount=Decimal("100.00"),
            vat_amount=Decimal("0.00"),
            status=BillingStorageDay.STATUS_CALCULATED,
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            box_code="BOX-SMALL",
            quantity=1,
            total_volume_m3=Decimal("1.000000"),
            status=StorageSnapshotLine.STATUS_OK,
        )
        StorageSnapshotLine.objects.create(
            day=storage_day,
            box_code="BOX-LARGE",
            quantity=1,
            total_volume_m3=Decimal("3.000000"),
            status=StorageSnapshotLine.STATUS_OK,
        )

        response = build_manager_storage_export(
            client=client,
            start=storage_day.day,
            end=storage_day.day,
            storage_days=BillingStorageDay.objects.filter(pk=storage_day.pk),
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)

        self.assertEqual(workbook.sheetnames[0], "Расчёт по коробам")
        self.assertIn("Хранение по артикулам", workbook.sheetnames)
        calculation = workbook["Расчёт по коробам"]
        net_by_box = {
            calculation.cell(row=row, column=3).value: calculation.cell(row=row, column=22).value
            for row in range(7, calculation.max_row + 1)
        }
        self.assertEqual(net_by_box, {"BOX-LARGE": 75, "BOX-SMALL": 25})
