"""Manager-only Excel export for storage billing.

The report reads billing snapshots and the immutable warehouse event log only.
It never changes stock, containers, reservations, warehouse operations or
client-cabinet data.  The first sheet is deliberately suitable for sending to
a customer after the manager has checked it: its total is allocated from the
saved daily billing rows, so it always reconciles to billing.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import BillingStorageDay, StorageSnapshotLine


ZERO = Decimal("0")
MONEY_QUANTUM = Decimal("0.01")
WEIGHT_NORM_KG = Decimal("450")
WEIGHT_VOLUME_M3_PER_NORM = Decimal("1.6")

_DARK_TEAL = "164A4C"
_TEAL = "2B7A78"
_PALE_TEAL = "DDEFEF"
_PALE_YELLOW = "FFF2CC"
_BORDER = "C7D7D7"


def _decimal(value) -> Decimal:
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return ZERO


def _quantize_money(value: Decimal) -> Decimal:
    return _decimal(value).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)


def _number(value, pattern: str = "0.000000") -> str:
    return format(_decimal(value).quantize(Decimal(pattern), rounding=ROUND_HALF_UP), "f")


def _payload_value(payload, *keys) -> str:
    payload = payload if isinstance(payload, dict) else {}
    for key in keys:
        value = payload.get(key)
        if value not in (None, ""):
            return str(value)
    return ""


def _saved_weight_volume_m3(payload) -> Decimal:
    """Read the weight-equivalent CBM saved by the current billing rule.

    New snapshots keep the auditable value inside ``billing_cbm_rule``.  The
    legacy flat key remains a backwards-compatible fallback for older rows.
    """
    payload = payload if isinstance(payload, dict) else {}
    rule = payload.get("billing_cbm_rule")
    if isinstance(rule, dict) and rule.get("weight_based_cbm") not in (None, ""):
        return _decimal(rule.get("weight_based_cbm"))
    return _decimal(payload.get("weight_volume_m3"))


def _style_sheet(sheet) -> None:
    """Style the supporting sheets with the same readable export language."""
    fill = PatternFill("solid", fgColor="F8B800")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="303030")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(str(cell.value or "")) + 2), 48)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = max(width, 12)


def _style_calculation_sheet(
    sheet,
    *,
    hide_pallets: bool = False,
    table_last_column: str = "Y",
    table_column_count: int = 25,
    pallet_mode: bool = False,
) -> None:
    """Apply the visual structure of the supplied storage-report sample."""
    thin = Side(style="thin", color=_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor=_DARK_TEAL)
    summary_fill = PatternFill("solid", fgColor=_PALE_TEAL)
    header_fill = PatternFill("solid", fgColor=_TEAL)

    sheet.merge_cells("A1:Y1")
    title = sheet["A1"]
    title.fill = title_fill
    title.font = Font(bold=True, color="FFFFFF", size=14)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    for row in range(2, 5):
        for cell in sheet[row]:
            cell.fill = summary_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 22

    for column in range(1, table_column_count + 1):
        cell = sheet.cell(row=6, column=column)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[6].height = 33

    for row in sheet.iter_rows(min_row=7, max_col=table_column_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:{table_last_column}{max(sheet.max_row, 6)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:6"
    sheet.sheet_view.showGridLines = False

    if pallet_mode:
        widths = {
            "A": 7, "B": 25, "C": 19, "D": 12, "E": 18, "F": 16, "G": 16,
            "H": 14, "I": 21, "J": 24, "K": 20, "L": 15, "M": 18, "N": 48,
        }
    else:
        widths = {
            "A": 7, "B": 19, "C": 19, "D": 17, "E": 38, "F": 18, "G": 15, "H": 13,
            "I": 12, "J": 12, "K": 12, "L": 15, "M": 16, "N": 16, "O": 15, "P": 15,
            "Q": 13, "R": 17, "S": 19, "T": 24, "U": 20, "V": 18,
            "W": 15, "X": 18, "Y": 42,
        }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    if hide_pallets:
        sheet.column_dimensions["B"].hidden = True


def _is_m3_report(days: list[BillingStorageDay]) -> bool:
    return bool(days) and all((day.billing_mode or "") == "m3_day" for day in days)


def _is_pallet_billing_day(day: BillingStorageDay) -> bool:
    mode = (day.billing_mode or "").strip()
    payload = day.payload if isinstance(day.payload, dict) else {}
    return mode.startswith("pallet_") or payload.get("charge_basis") == "pallet"


def _is_pallet_report(days: list[BillingStorageDay]) -> bool:
    return bool(days) and all(_is_pallet_billing_day(day) for day in days)


def _pallet_box_counts(day: BillingStorageDay) -> dict[str, int]:
    payload = day.payload if isinstance(day.payload, dict) else {}
    raw_counts = payload.get("pallet_box_counts")
    if not isinstance(raw_counts, dict):
        return {}
    counts: dict[str, int] = {}
    for pallet_code, raw_count in raw_counts.items():
        try:
            counts[str(pallet_code)] = max(int(raw_count or 0), 0)
        except (TypeError, ValueError):
            continue
    return counts


def _article_rows(day_ids: list[int]):
    if not day_ids:
        return []
    return list(
        StorageSnapshotLine.objects.filter(day_id__in=day_ids)
        .values("day__day", "sku_code", "name", "barcode")
        .annotate(
            quantity=Sum("quantity"),
            box_count=Count("box_code", distinct=True, filter=~Q(box_code="")),
            pallet_count=Count("pallet_code", distinct=True, filter=~Q(pallet_code="")),
            physical_volume_m3=Sum("total_volume_m3"),
            billable_volume_m3=Sum("billable_volume_m3"),
            no_dims_count=Count("id", filter=Q(status=StorageSnapshotLine.STATUS_NO_DIMS)),
        )
        .order_by("day__day", "sku_code", "name", "barcode")
    )


def _pallet_day_rows(days: list[BillingStorageDay]):
    rows = []
    for day in days:
        box_counts = _pallet_box_counts(day)
        lines = list(day.lines.all().order_by("pallet_code", "id"))
        tariff = _decimal(getattr(day.charge, "tariff", None)) if day.charge_id else ZERO
        for line, net, vat in _allocate_day_amounts(day, lines):
            pallet_code = line.pallet_code or ""
            rows.append(
                {
                    "day": day.day,
                    "pallet": pallet_code,
                    "box_count": box_counts.get(pallet_code, 0),
                    "zone": line.zone_code or "",
                    "cell": line.cell_code or "",
                    "tariff": tariff if tariff else None,
                    "net": net,
                    "vat": vat,
                    "gross": net + vat,
                }
            )
    return rows


def _movement_rows(*, client_id: int, start: date, end: date, include_pallets: bool = True):
    """Return immutable warehouse events, enriched only by saved snapshot ids.

    A historical SKU is never guessed from a current box: a box can be reused
    after a shipment.  This keeps the movement sheet auditable.
    """
    from sklad.models import WarehouseEvent, WarehouseStockSnapshot

    start_at = timezone.make_aware(datetime.combine(start, time.min))
    end_at = timezone.make_aware(datetime.combine(end + timedelta(days=1), time.min))
    events = list(
        WarehouseEvent.objects.filter(agency_id=client_id, occurred_at__gte=start_at, occurred_at__lt=end_at)
        .select_related("agency", "container")
        .order_by("occurred_at", "id")[:50000]
    )
    snapshot_ids: set[int] = set()
    for event in events:
        raw_id = _payload_value(event.payload, "snapshot_id")
        if raw_id.isdigit():
            snapshot_ids.add(int(raw_id))
    snapshots = {
        row["id"]: row
        for row in WarehouseStockSnapshot.objects.filter(pk__in=snapshot_ids).values("id", "sku_code", "name", "barcode")
    }
    rows = []
    for event in events:
        payload = event.payload if isinstance(event.payload, dict) else {}
        raw_id = _payload_value(payload, "snapshot_id")
        # A historical event can retain an id of a snapshot that has since
        # been archived.  Leave SKU fields empty in that case; do not fail the
        # whole report or recover data from a current, possibly reused box.
        snapshot = (snapshots.get(int(raw_id)) or {}) if raw_id.isdigit() else {}
        values = [
            timezone.localtime(event.occurred_at).strftime("%d.%m.%Y %H:%M"),
            event.agency.short_name or event.agency.agn_name,
            event.event_type,
            event.qty,
            snapshot.get("sku_code", ""),
            snapshot.get("name", ""),
            snapshot.get("barcode", ""),
            _payload_value(payload, "box_code"),
        ]
        if include_pallets:
            values.extend([
                _payload_value(payload, "pallet_code", "source_pallet_code"),
                _payload_value(payload, "destination_pallet_code"),
            ])
        values.extend([
            event.from_zone_code or "",
            event.to_zone_code or "",
            f"{event.source_document_type}:{event.source_document_id}".strip(":"),
            event.stock_context_type + (f":{event.stock_context_id}" if event.stock_context_id else ""),
            _payload_value(payload, "reason", "note"),
        ])
        rows.append(values)
    return rows


def _billing_basis(day: BillingStorageDay) -> str:
    """Human-readable basis of a saved daily amount, without recalculating it."""
    if day.billing_mode != "m3_day":
        return "Палета" if "pallet" in (day.billing_mode or "") else (day.billing_mode or "По тарифу")
    payload = day.payload if isinstance(day.payload, dict) else {}
    weight_volume = _saved_weight_volume_m3(payload)
    return "Вес" if weight_volume > _decimal(day.physical_volume_m3) else "Объём"


def _weight_m3_from_g(weight_g) -> Decimal | None:
    if weight_g is None:
        return None
    return (_decimal(weight_g) / Decimal("1000") / WEIGHT_NORM_KG) * WEIGHT_VOLUME_M3_PER_NORM


def _line_basis_label(*, volume_m3, weight_g) -> str:
    """Explain an individual line without confusing it with daily billing."""
    weight_m3 = _weight_m3_from_g(weight_g)
    if weight_m3 is None:
        return "Нет веса"
    return "Вес строки" if weight_m3 > _decimal(volume_m3) else "Объём строки"


def _compact_date_list(days: set[date], *, limit: int = 6) -> str:
    ordered = sorted(days)
    if not ordered:
        return "—"
    if len(ordered) <= limit:
        return ", ".join(day.strftime("%d.%m") for day in ordered)
    head = ", ".join(day.strftime("%d.%m") for day in ordered[:3])
    tail = ", ".join(day.strftime("%d.%m") for day in ordered[-2:])
    return f"{head} … {tail}"


def _container_weights(*, client_id: int, box_codes: set[str]) -> dict[str, Decimal]:
    """Return current box weights for an audit reference, never as history."""
    if not box_codes:
        return {}
    from sklad.models import WarehouseContainer

    return {
        row["container_code"]: _decimal(row["gross_weight_g"])
        for row in WarehouseContainer.objects.filter(
            agency_id=client_id,
            container_type=WarehouseContainer.TYPE_BOX,
            container_code__in=box_codes,
        ).values("container_code", "gross_weight_g")
        if row["gross_weight_g"] is not None
    }


def _allocate_day_amounts(day: BillingStorageDay, lines: list[StorageSnapshotLine]):
    """Allocate a saved daily charge between the snapshot lines exactly.

    Billing is calculated at the daily snapshot level.  Lines are therefore
    an explanatory breakdown, and not a second billing algorithm.  The final
    line receives the rounding remainder so the Excel total equals billing.
    """
    candidates = [line for line in lines if line.status == StorageSnapshotLine.STATUS_OK]
    volume_total = sum((_decimal(line.total_volume_m3) for line in candidates), ZERO)
    quantity_total = sum((_decimal(line.quantity) for line in candidates), ZERO)
    if _is_pallet_billing_day(day) and candidates:
        # A pallet snapshot contains one line per occupied pallet.  The saved
        # daily pallet charge must therefore be split evenly by pallet, never
        # by the dimensions of the representative product on that pallet.
        bases = [Decimal("1") for _ in candidates]
        denominator = Decimal(len(candidates))
    elif volume_total > ZERO:
        bases = [_decimal(line.total_volume_m3) for line in candidates]
        denominator = volume_total
    elif quantity_total > ZERO:
        bases = [_decimal(line.quantity) for line in candidates]
        denominator = quantity_total
    else:
        candidates = list(lines)
        bases = [Decimal("1") for _ in candidates]
        denominator = Decimal(len(candidates) or 1)

    net_total = _quantize_money(day.amount)
    vat_total = _quantize_money(day.vat_amount)
    allocated_net = ZERO
    allocated_vat = ZERO
    result = []
    for index, (line, base) in enumerate(zip(candidates, bases)):
        if index == len(candidates) - 1:
            net = net_total - allocated_net
            vat = vat_total - allocated_vat
        else:
            ratio = base / denominator
            net = _quantize_money(net_total * ratio)
            vat = _quantize_money(vat_total * ratio)
            allocated_net += net
            allocated_vat += vat
        result.append((line, net, vat))
    return result


def _calculation_rows(days: list[BillingStorageDay], *, client_id: int):
    """Build period rows from immutable daily snapshots and allocated totals."""
    if not days:
        return []
    lines_by_day: dict[int, list[StorageSnapshotLine]] = defaultdict(list)
    for line in (
        StorageSnapshotLine.objects.filter(day_id__in=[day.id for day in days])
        .select_related("sku")
        .order_by("day__day", "id")
    ):
        lines_by_day[line.day_id].append(line)

    box_codes = {line.box_code for lines in lines_by_day.values() for line in lines if line.box_code}
    current_weights = _container_weights(client_id=client_id, box_codes=box_codes)
    groups: dict[tuple, dict] = {}
    for day in days:
        basis = _billing_basis(day)
        tariff = _decimal(getattr(day.charge, "tariff", None)) if day.charge_id else ZERO
        tariff_key = tariff if tariff else None
        pallet_mode = _is_pallet_billing_day(day)
        pallet_box_counts = _pallet_box_counts(day)
        for line, net, vat in _allocate_day_amounts(day, lines_by_day.get(day.id, [])):
            pallet_code = line.pallet_code or ""
            box_count = pallet_box_counts.get(pallet_code, 0) if pallet_mode else 0
            if pallet_mode:
                # A pallet can appear as a new report row when its saved box
                # count, zone or cell changes.  Product dimensions and the
                # representative SKU are deliberately excluded from the key.
                key = (
                    "pallet", pallet_code, box_count, line.zone_code or "", line.cell_code or "",
                    basis, tariff_key,
                )
            else:
                # Dimensions are in the key: a corrected historical snapshot
                # must not be silently merged into a different box state.
                key = (
                    "volume", pallet_code, line.box_code or "", line.sku_code or "", line.name or "",
                    line.barcode or "", line.length_mm or 0, line.width_mm or 0, line.height_mm or 0,
                    basis, tariff_key,
                )
            group = groups.get(key)
            if group is None:
                group = groups[key] = {
                    "pallet_mode": pallet_mode,
                    "pallet": pallet_code,
                    "box": line.box_code or "",
                    "box_count": box_count,
                    "zone": line.zone_code or "",
                    "cell": line.cell_code or "",
                    "sku": line.sku_code or "",
                    "name": line.name or "",
                    "barcode": line.barcode or "",
                    "quantity": _decimal(line.quantity),
                    "status": line.get_status_display(),
                    "width": line.width_mm,
                    "height": line.height_mm,
                    "depth": line.length_mm,
                    "volume": _decimal(line.total_volume_m3),
                    "weight_g": current_weights.get(line.box_code or ""),
                    "first_day": day.day,
                    "last_day": day.day,
                    "days": set(),
                    "bases": set(),
                    "tariffs": set(),
                    "net": ZERO,
                    "vat": ZERO,
                }
            group["quantity"] = _decimal(line.quantity)  # the latest saved daily quantity
            group["status"] = line.get_status_display()
            group["volume"] = _decimal(line.total_volume_m3)
            group["weight_g"] = current_weights.get(line.box_code or "")
            group["first_day"] = min(group["first_day"], day.day)
            group["last_day"] = max(group["last_day"], day.day)
            group["days"].add(day.day)
            group["bases"].add(basis)
            if tariff:
                group["tariffs"].add(tariff)
            group["net"] += net
            group["vat"] += vat

    rows = []
    for group in groups.values():
        bases = sorted(group.pop("bases"))
        tariffs = sorted(group.pop("tariffs"))
        days_count = len(group["days"])
        basis_label = " / ".join(bases) if bases else "—"
        group["basis"] = basis_label
        group["tariff"] = tariffs[0] if len(tariffs) == 1 else None
        group["gross"] = group["net"] + group["vat"]
        group["net_per_day"] = (group["net"] / Decimal(days_count)) if days_count else ZERO
        if group["pallet_mode"]:
            group["explanation"] = (
                f"Палета: {days_count} дн. хранения; коробов на паллете: {group['box_count']}; "
                f"даты: {_compact_date_list(group['days'])}. "
                "Сумма распределена поровну между учтёнными паллетами из дневных начислений биллинга."
            )
        else:
            group["explanation"] = (
                f"{basis_label}: {days_count} дн. хранения; даты: {_compact_date_list(group['days'])}. "
                "Сумма распределена из дневных начислений биллинга."
            )
        rows.append(group)
    return sorted(rows, key=lambda row: (row["pallet"], row["box"], row["sku"], row["name"], row["basis"], row["tariff"] or ZERO))


def _write_calculation_sheet(workbook: Workbook, *, client, start: date, end: date, days: list[BillingStorageDay]) -> None:
    sheet = workbook.active
    is_m3_report = _is_m3_report(days)
    is_pallet_report = _is_pallet_report(days)
    sheet.title = "Расчёт по паллетам" if is_pallet_report else "Расчёт по коробам"
    report_rows = _calculation_rows(days, client_id=client.id)
    boxes = {row["box"] for row in report_rows if row["box"]}
    pallets = {row["pallet"] for row in report_rows if row["pallet"]}
    by_volume = {row["box"] for row in report_rows if row["box"] and "Объём" in row["basis"]}
    by_weight = {row["box"] for row in report_rows if row["box"] and "Вес" in row["basis"]}
    net_total = sum((_decimal(day.amount) for day in days), ZERO)
    vat_total = sum((_decimal(day.vat_amount) for day in days), ZERO)
    total_snapshots = len(days)
    latest_box_count = max(days, key=lambda day: day.day).box_count if days else 0
    pallet_days = sum((day.pallet_count for day in days), 0)

    sheet["A1"] = f"Расчёт хранения: {client.short_name or client.agn_name}"
    common_summary = (
        ("A2", "A2:C2", "D2", "D2:E2", "Период", f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"),
        ("N2", "N2:P2", "Q2", "Q2:S2", "Сумма без НДС, ₽", _quantize_money(net_total)),
        ("T2", "T2:U2", "V2", "V2:V2", "НДС, ₽", _quantize_money(vat_total)),
        ("A3", "A3:C3", "D3", "D3:E3", "Клиент", client.short_name or client.agn_name),
        ("N3", "N3:P3", "Q3", "Q3:S3", "Итого с НДС, ₽", _quantize_money(net_total + vat_total)),
        ("T3", "T3:U3", "V3", "V3:V3", "Дневных снимков", total_snapshots),
        ("T4", "T4:U4", "V4", "V4:V4", "Строк в отчёте", len(report_rows)),
    )
    if is_pallet_report:
        summary_pairs = common_summary + (
            ("F2", "F2:J2", "K2", "K2:M2", "Коробов в последнем снимке", latest_box_count),
            ("F3", "F3:J3", "K3", "K3:M3", "Уникальных паллет за период", len(pallets)),
            ("F4", "F4:J4", "K4", "K4:M4", "Палето-дней", pallet_days),
            ("N4", "N4:P4", "Q4", "Q4:S4", "Единица расчёта", "палета / сутки"),
        )
    else:
        summary_pairs = common_summary + (
            ("F2", "F2:J2", "K2", "K2:M2", "Количество коробов", len(boxes)),
            ("F3", "F3:J3", "K3", "K3:M3", "Единица хранения", "м³" if is_m3_report else len(pallets)),
            ("F4", "F4:J4", "K4", "K4:M4", "Коробов с днями по объёму", len(by_volume)),
            ("N4", "N4:P4", "Q4", "Q4:S4", "Коробов с днями по весу", len(by_weight)),
        )
    for label_cell, label_range, value_cell, value_range, label, value in summary_pairs:
        if ":" in label_range:
            sheet.merge_cells(label_range)
        if ":" in value_range and value_range != f"{value_cell}:{value_cell}":
            sheet.merge_cells(value_range)
        sheet[label_cell] = label
        sheet[label_cell].font = Font(bold=True, color="164A4C")
        sheet[value_cell] = value
        sheet[value_cell].font = Font(bold=True, color="303030")

    if is_pallet_report:
        headers = [
            "№", "Паллет", "Коробов на паллете", "Зона", "Ячейка", "Первый снимок",
            "Последний снимок", "Дней хранения", "Тариф, ₽/пал./сутки",
            "Стоимость за день без НДС, ₽", "Стоимость без НДС, ₽", "НДС, ₽",
            "Итого, ₽", "Пояснение",
        ]
    else:
        headers = [
            "№", "" if is_m3_report else "Паллет", "Короб", "SKU", "Номенклатура", "Штрихкод", "Количество, шт", "Статус снимка",
            "Ширина, мм", "Высота, мм", "Глубина, мм", "Объём строки, м³", "Вес короба (текущий), г",
            "Вес в м³ (справочно)", "Расчёт строки (справочно)", "Первый снимок", "Последний снимок",
            "Дней хранения", "Стоимость за день без НДС, ₽", "Основание дневного начисления",
            "Тариф, ₽", "Стоимость без НДС, ₽", "НДС, ₽", "Итого, ₽", "Пояснение",
        ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=index, value=header)

    for number, row in enumerate(report_rows, start=1):
        if is_pallet_report:
            values = [
                number, row["pallet"] or "—", row["box_count"], row["zone"] or "—", row["cell"] or "—",
                row["first_day"].strftime("%d.%m.%Y"), row["last_day"].strftime("%d.%m.%Y"), len(row["days"]),
                float(row["tariff"]) if row["tariff"] is not None else "—",
                float(_quantize_money(row["net_per_day"])), float(_quantize_money(row["net"])),
                float(_quantize_money(row["vat"])), float(_quantize_money(row["gross"])), row["explanation"],
            ]
        else:
            weight_m3 = _weight_m3_from_g(row["weight_g"])
            line_basis = _line_basis_label(volume_m3=row["volume"], weight_g=row["weight_g"])
            values = [
                number, "" if is_m3_report else row["pallet"] or "—", row["box"] or "—", row["sku"], row["name"], row["barcode"], float(row["quantity"]), row["status"],
                row["width"] or None, row["height"] or None, row["depth"] or None, float(row["volume"]),
                float(row["weight_g"]) if row["weight_g"] is not None else None,
                float(weight_m3) if weight_m3 is not None else None,
                line_basis,
                row["first_day"].strftime("%d.%m.%Y"), row["last_day"].strftime("%d.%m.%Y"), len(row["days"]),
                float(_quantize_money(row["net_per_day"])),
                row["basis"], float(row["tariff"]) if row["tariff"] is not None else "—",
                float(_quantize_money(row["net"])), float(_quantize_money(row["vat"])), float(_quantize_money(row["gross"])),
                row["explanation"],
            ]
        sheet.append(values)

    if is_pallet_report:
        for row in range(7, sheet.max_row + 1):
            for column in (9, 10, 11, 12, 13):
                sheet.cell(row=row, column=column).number_format = '#,##0.00 "₽"'
        _style_calculation_sheet(
            sheet,
            table_last_column="N",
            table_column_count=14,
            pallet_mode=True,
        )
    else:
        for row in range(7, sheet.max_row + 1):
            for column in (7, 12, 14):
                sheet.cell(row=row, column=column).number_format = "0.000000"
            for column in (19, 21, 22, 23, 24):
                sheet.cell(row=row, column=column).number_format = '#,##0.00 "₽"'
        _style_calculation_sheet(sheet, hide_pallets=is_m3_report)


def build_manager_storage_export(*, client, start: date, end: date, storage_days) -> HttpResponse:
    """Build an xlsx report for one accessible client and selected period."""
    days = list(storage_days)
    workbook = Workbook()
    _write_calculation_sheet(workbook, client=client, start=start, end=end, days=days)
    is_m3_report = _is_m3_report(days)
    is_pallet_report = _is_pallet_report(days)

    overview = workbook.create_sheet("Сводка по дням")
    overview_headers = ["Дата", "Клиент", "Режим"]
    if not is_m3_report:
        overview_headers.append("Палет")
    overview_headers.extend([
        "Коробов", "Единиц товара", "Физический объём, м³", "Объём по весу, м³", "К начислению, м³",
        "Тариф, ₽", "Без НДС, ₽", "НДС, ₽", "Итого, ₽", "Статус", "Заявка", "Начисление",
    ])
    overview.append(overview_headers)
    for row in days:
        payload = row.payload if isinstance(row.payload, dict) else {}
        overview_row = [row.day.strftime("%d.%m.%Y"), client.short_name or client.agn_name, row.billing_mode or ""]
        if not is_m3_report:
            overview_row.append(row.pallet_count)
        overview_row.extend([
            row.box_count, row.sku_unit_count, _number(row.physical_volume_m3),
            _number(_saved_weight_volume_m3(payload)), _number(row.billable_volume_m3),
            _number(getattr(row.charge, "tariff", None), "0.0000"), _number(row.amount, "0.01"),
            _number(row.vat_amount, "0.01"), _number(_decimal(row.amount) + _decimal(row.vat_amount), "0.01"),
            row.get_status_display(), row.application_id or "", row.charge_id or "",
        ])
        overview.append(overview_row)
    _style_sheet(overview)

    if is_pallet_report:
        pallet_details = workbook.create_sheet("Паллеты по дням")
        pallet_details.append(
            [
                "Дата", "Паллет", "Коробов на паллете", "Зона", "Ячейка",
                "Тариф, ₽/пал./сутки", "Без НДС, ₽", "НДС, ₽", "Итого, ₽",
            ]
        )
        for row in _pallet_day_rows(days):
            pallet_details.append(
                [
                    row["day"].strftime("%d.%m.%Y"), row["pallet"] or "—", row["box_count"],
                    row["zone"] or "—", row["cell"] or "—",
                    float(row["tariff"]) if row["tariff"] is not None else "—",
                    float(_quantize_money(row["net"])), float(_quantize_money(row["vat"])),
                    float(_quantize_money(row["gross"])),
                ]
            )
        for row_index in range(2, pallet_details.max_row + 1):
            for column in (6, 7, 8, 9):
                pallet_details.cell(row=row_index, column=column).number_format = '#,##0.00 "₽"'
        _style_sheet(pallet_details)
    else:
        articles = workbook.create_sheet("Хранение по артикулам")
        article_headers = ["Дата", "Артикул", "Наименование", "Штрих-код", "Количество", "Коробов"]
        if not is_m3_report:
            article_headers.append("Палет")
        article_headers.extend(["Физический объём, м³", "К начислению, м³", "Строк без габаритов"])
        articles.append(article_headers)
        for row in _article_rows([row.id for row in days]):
            day_value = row.get("day__day")
            article_row = [
                day_value.strftime("%d.%m.%Y") if day_value else "", row.get("sku_code") or "", row.get("name") or "",
                row.get("barcode") or "", _number(row.get("quantity"), "0.001"), row.get("box_count") or 0,
            ]
            if not is_m3_report:
                article_row.append(row.get("pallet_count") or 0)
            article_row.extend([
                _number(row.get("physical_volume_m3")), _number(row.get("billable_volume_m3")),
                row.get("no_dims_count") or 0,
            ])
            articles.append(article_row)
        _style_sheet(articles)

    movements = workbook.create_sheet("Движение товара")
    movement_headers = [
        "Дата и время", "Клиент", "Событие", "Количество", "Артикул", "Наименование", "Штрих-код",
        "Короб",
    ]
    if not is_m3_report:
        movement_headers.extend(["Палета / откуда", "Палета / куда"])
    movement_headers.extend(["Зона откуда", "Зона куда", "Документ", "Контекст", "Примечание"])
    movements.append(movement_headers)
    for row in _movement_rows(client_id=client.id, start=start, end=end, include_pallets=not is_m3_report):
        movements.append(row)
    _style_sheet(movements)

    notes = workbook.create_sheet("Пояснения")
    notes.append(["Параметр", "Значение"])
    notes.append(["Клиент", client.short_name or client.agn_name])
    notes.append(["Период", f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"])
    notes.append(["Источник начислений", "Сохранённые дневные снимки биллинга; итог первого листа сверяется с ними до копейки."])
    if is_pallet_report:
        notes.append(["Единица расчёта", "Одна занятая паллета за одни сутки. Габариты и объём коробов в паллетном расчёте не участвуют."])
        notes.append(["Количество коробов", "Берётся из сохранённого дневного снимка отдельно для каждой паллеты. Если состав паллеты менялся, в расчёте появятся отдельные строки."])
        notes.append(["Распределение суммы", "Сохранённая дневная сумма распределяется поровну между всеми учтёнными паллетами этого дня и сверяется с биллингом до копейки."])
    else:
        notes.append(["Почему одинаковые строки могут иметь разную сумму", "Одинаковый товар/короб может храниться разное число дней. Сравнивайте колонку «Стоимость за день без НДС»: если она одинаковая, отличие итоговой суммы вызвано только количеством дней хранения."])
        notes.append(["Основание дневного начисления", "Отчёт разделяет одинаковый товар/короб по основанию дня и тарифу: отдельная строка для дней по весу, по объёму и для каждого тарифа. Основание выбирается по дневному остатку клиента целиком."])
        notes.append(["Расчёт строки", "Справочная проверка конкретного короба/SKU по текущему весу короба и объёму строки. Строковая сумма распределяется из дневных начислений биллинга и сверяется с ними до копейки."])
        notes.append(["Вес короба", "Текущее значение карточки короба. Исторический вес не хранится в строке снимка, поэтому текущий вес не подменяет историю."])
    notes.append(["Первый / последний снимок", "Дата первого и последнего сохранённого дневного снимка в выбранном периоде, не дата приёмки или отгрузки."])
    notes.append(["Движение", "Только чтение журнала склада. SKU указан только если событие сохранило ссылку на исторический снимок."])
    notes.append(["Ограничение движения", "В лист включается не более 50 000 событий за период."])
    _style_sheet(notes)

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = (
        f'attachment; filename="storage-manager-{client.id}-{start.isoformat()}-{end.isoformat()}.xlsx"'
    )
    return response
