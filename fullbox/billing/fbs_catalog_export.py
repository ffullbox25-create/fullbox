"""Manager-only XLSX export of an FBS client's catalogue and unit volume.

The export is intentionally read-only: it reads the client SKU catalogue and
FBS balance rows, but does not change FBS stock, pricing rules or billing.
The tariff-grid sheet is a template for agreeing ranges before an automated
FBS charging rule is introduced.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fbs.models import FbsStockBalance
from sku.models import SKU


ZERO = Decimal("0")
LITER_QUANTUM = Decimal("0.001")
ORANGE = "F89000"
YELLOW = "F8B800"
PALE_ORANGE = "FCE6C4"
TEXT = "303030"
MUTED = "6B7280"
BORDER = "E5E7EB"


def _decimal(value) -> Decimal:
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return ZERO


def _liters(sku: SKU) -> Decimal:
    """Return unit volume in litres from SKU dimensions recorded in millimetres."""
    length = _decimal(sku.length_mm)
    width = _decimal(sku.width_mm)
    height = _decimal(sku.height_mm)
    if min(length, width, height) <= 0:
        return ZERO
    return (length * width * height / Decimal("1000000")).quantize(
        LITER_QUANTUM, rounding=ROUND_HALF_UP
    )


def _primary_barcodes(skus) -> dict[int, str]:
    values: dict[int, list[str]] = defaultdict(list)
    for sku in skus:
        for barcode in sku.barcodes.all():
            if barcode.value:
                values[sku.id].append(str(barcode.value))
    return {sku_id: ", ".join(items) for sku_id, items in values.items()}


def _set_title(sheet, title: str, last_column: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    cell = sheet["A1"]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=ORANGE)
    cell.font = Font(name="Manrope", size=14, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28


def _style_header(sheet, row: int, columns: int) -> None:
    thin = Side(style="thin", color=BORDER)
    for cell in sheet.iter_cols(min_col=1, max_col=columns, min_row=row, max_row=row):
        target = cell[0]
        target.fill = PatternFill("solid", fgColor=YELLOW)
        target.font = Font(name="Manrope", size=9, bold=True, color=TEXT)
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.border = Border(bottom=thin)
    sheet.row_dimensions[row].height = 30


def _format_data_rows(sheet, start_row: int, end_row: int, columns: int) -> None:
    thin = Side(style="thin", color=BORDER)
    for row in sheet.iter_rows(min_row=start_row, max_row=max(end_row, start_row), min_col=1, max_col=columns):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)


def _autosize(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_view.showGridLines = False


def build_fbs_catalog_export(*, client) -> HttpResponse:
    """Build a catalogue report plus an unbound FBS litre-pricing template."""
    skus = list(
        SKU.objects.filter(agency=client, deleted=False)
        .select_related("market")
        .prefetch_related("barcodes")
        .order_by("sku_code", "id")
    )
    stock_by_sku = {
        row["sku_ref_id"]: row["quantity"] or 0
        for row in (
            FbsStockBalance.objects.filter(agency=client, sku_ref__isnull=False)
            .values("sku_ref_id")
            .annotate(quantity=Sum("qty"))
        )
    }
    barcodes = _primary_barcodes(skus)
    generated_at = timezone.localtime(timezone.now()).strftime("%d.%m.%Y %H:%M")

    workbook = Workbook()
    catalogue = workbook.active
    catalogue.title = "Номенклатура FBS"
    _set_title(catalogue, f"FBS · Номенклатура и литраж — {client}", "K")
    catalogue["A2"] = "Сформировано"
    catalogue["B2"] = generated_at
    catalogue["D2"] = "Формула литража"
    catalogue["E2"] = "Длина × ширина × высота / 1 000 000 (мм → л)"
    catalogue["A3"] = "Важно"
    catalogue["B3"] = (
        "Это справочный экспорт. Он не создаёт начисления и не меняет остатки, тарифы или счета FBS."
    )
    catalogue.merge_cells("B3:K3")
    for ref in ("A2", "D2", "A3"):
        catalogue[ref].font = Font(name="Manrope", bold=True, color=TEXT)
    for ref in ("B2", "E2", "B3"):
        catalogue[ref].font = Font(name="Manrope", color=MUTED)
    catalogue["B3"].alignment = Alignment(wrap_text=True, vertical="center")
    catalogue.row_dimensions[3].height = 30

    headers = [
        "№",
        "Артикул",
        "Наименование",
        "Штрихкоды",
        "Маркетплейс",
        "Длина, мм",
        "Ширина, мм",
        "Высота, мм",
        "Литраж 1 шт., л",
        "Текущий остаток FBS, шт.",
        "Литраж остатка FBS, л",
    ]
    catalogue.append([])
    catalogue.append(headers)
    _style_header(catalogue, 5, len(headers))

    complete_dimensions = 0
    total_liters_in_stock = ZERO
    for index, sku in enumerate(skus, start=1):
        liters = _liters(sku)
        quantity = _decimal(stock_by_sku.get(sku.id))
        total_liters = (liters * quantity).quantize(LITER_QUANTUM, rounding=ROUND_HALF_UP)
        if liters > 0:
            complete_dimensions += 1
        total_liters_in_stock += total_liters
        catalogue.append(
            [
                index,
                sku.sku_code,
                sku.name,
                barcodes.get(sku.id, ""),
                sku.market.name if sku.market_id else "",
                _decimal(sku.length_mm) or None,
                _decimal(sku.width_mm) or None,
                _decimal(sku.height_mm) or None,
                liters if liters > 0 else None,
                quantity or None,
                total_liters if liters > 0 and quantity > 0 else None,
            ]
        )
    last_catalogue_row = 5 + max(len(skus), 1)
    _format_data_rows(catalogue, 6, last_catalogue_row, len(headers))
    catalogue.freeze_panes = "A6"
    catalogue.auto_filter.ref = f"A5:K{last_catalogue_row}"
    for column in ("F", "G", "H", "I", "J", "K"):
        for cell in catalogue[column][5:]:
            cell.number_format = "#,##0.000"
    _autosize(
        catalogue,
        {
            "A": 7, "B": 18, "C": 42, "D": 28, "E": 18, "F": 13,
            "G": 13, "H": 13, "I": 18, "J": 23, "K": 23,
        },
    )

    settings = workbook.create_sheet("Тарифная сетка FBS")
    _set_title(settings, f"FBS · Черновик тарифной сетки — {client}", "F")
    settings.merge_cells("A2:F2")
    settings["A2"] = (
        "Заполните диапазоны и цену после согласования. Эта вкладка не применяется к биллингу автоматически "
        "и служит для подготовки правил."
    )
    settings["A2"].fill = PatternFill("solid", fgColor=PALE_ORANGE)
    settings["A2"].font = Font(name="Manrope", size=10, color=TEXT)
    settings["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    settings.row_dimensions[2].height = 32
    tariff_headers = [
        "От, л включительно",
        "До, л включительно",
        "Цена за 1 л, ₽",
        "Минимум за строку, ₽",
        "Период действия с",
        "Комментарий / правило",
    ]
    settings.append([])
    settings.append(tariff_headers)
    _style_header(settings, 4, len(tariff_headers))
    for _ in range(12):
        settings.append([None, None, None, None, None, None])
    _format_data_rows(settings, 5, 16, len(tariff_headers))
    for row in range(5, 17):
        settings.cell(row=row, column=1).number_format = "#,##0.000"
        settings.cell(row=row, column=2).number_format = "#,##0.000"
        settings.cell(row=row, column=3).number_format = "#,##0.00"
        settings.cell(row=row, column=4).number_format = "#,##0.00"
        settings.cell(row=row, column=5).number_format = "dd.mm.yyyy"
    settings.freeze_panes = "A5"
    _autosize(settings, {"A": 22, "B": 22, "C": 19, "D": 23, "E": 20, "F": 48})

    notes = workbook.create_sheet("Контроль")
    _set_title(notes, f"FBS · Контроль исходных данных — {client}", "B")
    notes.append(["Показатель", "Значение"])
    _style_header(notes, 2, 2)
    notes.append(["SKU в выгрузке", len(skus)])
    notes.append(["SKU с заполненными габаритами", complete_dimensions])
    notes.append(["SKU без полного литража", len(skus) - complete_dimensions])
    notes.append(["Литраж текущего остатка FBS, л", total_liters_in_stock.quantize(LITER_QUANTUM)])
    notes.append(["Автоматическое начисление FBS по сетке", "Не включено"])
    notes.append(["Следующий шаг", "Согласовать диапазоны литража и ставку за 1 л, затем добавить правило тарификации."])
    _format_data_rows(notes, 3, 9, 2)
    notes["B6"].number_format = "#,##0.000"
    _autosize(notes, {"A": 42, "B": 72})

    payload = BytesIO()
    workbook.save(payload)
    filename = f"fbs_nomenclature_liters_client_{client.id}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    response = HttpResponse(
        payload.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
