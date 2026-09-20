from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPORT_HEADERS = (
    "Отдел",
    "Должность",
    "Сотрудник",
    "Отработано, ч",
    "Опоздания, ч",
    "Стажировка, ч",
    "Почасовая ставка",
    "Месячный оклад",
    "Начислено",
    "Производительность, ед.",
    "Премия",
    "Средняя стоимость часа",
    "Выплаты по договорам",
    "Бонусы",
    "Удержания",
    "Итого к выплате",
    "Полнота данных",
)


def excel_safe_text(value):
    value = str(value or "")
    return "'" + value if value[:1] in {"=", "+", "-", "@"} else value


def _row_values(row):
    return [
        excel_safe_text(row["department"]),
        excel_safe_text(row["position"]),
        excel_safe_text(row["employee"]),
        float(row["worked_hours"]),
        float(row["late_hours"]),
        float(row["internship_hours"]),
        float(row["hourly_rate"]),
        float(row["monthly_salary"]),
        float(row["accrued"]),
        float(row["productivity_units"]),
        float(row["premium"]),
        float(row["average_hour_cost"]),
        float(row["contract_payments"]),
        float(row["bonuses"]),
        float(row["deductions"]),
        float(row["final_due"]),
        excel_safe_text(row["source_status"]),
    ]


def build_monthly_payroll_workbook(report) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Месячный отчёт"
    sheet.merge_cells("A1:Q1")
    sheet["A1"] = "Месячный отчёт по персоналу"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="303030")
    sheet["A2"] = "Период"
    sheet["B2"] = report["month"]
    sheet["B2"].number_format = "mmmm yyyy"
    sheet.append([])
    sheet.append(list(REPORT_HEADERS))

    header_fill = PatternFill(fill_type="solid", fgColor="F8B800")
    subtotal_fill = PatternFill(fill_type="solid", fgColor="FFF1CC")
    total_fill = PatternFill(fill_type="solid", fgColor="F89000")
    thin = Side(style="thin", color="DDD6CC")
    for cell in sheet[4]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="303030")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)

    data_start = 5
    for department in report["departments"]:
        for row in department["rows"]:
            sheet.append(_row_values(row))
        totals = department["totals"]
        sheet.append(
            [
                f"Итого: {excel_safe_text(department['name'])}",
                "",
                "",
                float(totals["worked_hours"]),
                float(totals["late_hours"]),
                float(totals["internship_hours"]),
                "",
                "",
                float(totals["accrued"]),
                float(totals["productivity_units"]),
                float(totals["premium"]),
                float(totals["average_hour_cost"]),
                float(totals["contract_payments"]),
                float(totals["bonuses"]),
                float(totals["deductions"]),
                float(totals["final_due"]),
                "",
            ]
        )
        for cell in sheet[sheet.max_row]:
            cell.fill = subtotal_fill
            cell.font = Font(name="Arial", bold=True, color="303030")

    totals = report["totals"]
    sheet.append(
        [
            "Итого по компании",
            "",
            "",
            float(totals["worked_hours"]),
            float(totals["late_hours"]),
            float(totals["internship_hours"]),
            "",
            "",
            float(totals["accrued"]),
            float(totals["productivity_units"]),
            float(totals["premium"]),
            float(totals["average_hour_cost"]),
            float(totals["contract_payments"]),
            float(totals["bonuses"]),
            float(totals["deductions"]),
            float(totals["final_due"]),
            "",
        ]
    )
    for cell in sheet[sheet.max_row]:
        cell.fill = total_fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")

    for row in sheet.iter_rows(min_row=data_start, max_row=sheet.max_row):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 3, 17})
        for column in (4, 5, 6, 10):
            row[column - 1].number_format = "#,##0.00"
        for column in (7, 8, 9, 11, 12, 13, 14, 15, 16):
            row[column - 1].number_format = '#,##0.00" ₽"'

    widths = (24, 26, 34, 16, 15, 16, 18, 18, 18, 22, 16, 23, 23, 16, 16, 20, 30)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[4].height = 42
    sheet.freeze_panes = "D5"
    if sheet.max_row >= data_start:
        sheet.auto_filter.ref = f"A4:Q{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:Q{sheet.max_row}"
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35

    method = workbook.create_sheet("Методика")
    method_rows = [
        ("Показатель", "Правило расчёта"),
        ("Отработано / опоздания / стажировка", "Сумма минут табеля за месяц, переведённая в часы."),
        ("Начислено", "Почасовая ставка × часы + оклад пропорционально плану + подтверждённые дополнительные начисления."),
        ("Производительность", "Сумма завершённых операционных единиц из WMS по сотруднику; складские данные используются только на чтение."),
        ("Средняя стоимость часа", "Начислено / отработанные часы."),
        ("Итого к выплате", "Начислено + премия + выплаты по договорам + бонусы − удержания."),
        ("Корректировки", "В отчёт входят только подтверждённые корректировки."),
        ("Образец структуры", "https://docs.google.com/spreadsheets/d/10pSxY-JxLZ8DYy2-V04aanaYkVuEGAOVMrE-MYK9Ssg/edit?gid=1789614734"),
    ]
    for values in method_rows:
        method.append(values)
    for cell in method[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="303030")
    method.column_dimensions["A"].width = 38
    method.column_dimensions["B"].width = 110
    for row in method.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    method.freeze_panes = "A2"
    method.sheet_view.showGridLines = False
    method.sheet_properties.pageSetUpPr.fitToPage = True
    method.page_setup.fitToWidth = 1
    method.page_setup.fitToHeight = 0
    return workbook


def workbook_bytes(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
