"""Client shipping Excel template: download, parse, stock/box validation."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from django.http import FileResponse, HttpResponse
from openpyxl import load_workbook


TEMPLATE_PATH = Path(__file__).resolve().parent / "data" / "shipping_import_template.xlsx"
DISCREPANCY_MARKER = "[РАСХОЖДЕНИЯ ШАБЛОНА]"

# Client fills product barcode and/or article (SKU) + qty in pieces (or Коробов).
# Box SHK is assigned by warehouse after pick.
TEMPLATE_HEADERS = ("Баркод товара", "Артикул", "Кол-во товаров", "Коробов", "Срок годности")


@dataclass
class TemplateLine:
    row_no: int
    barcode: str
    qty: int
    sku: str = ""
    box_barcode: str = ""
    expiry: str = ""

    def identity_label(self) -> str:
        parts = []
        if self.barcode:
            parts.append(f"ШК {self.barcode}")
        if self.sku:
            parts.append(f"артикул {self.sku}")
        return " / ".join(parts) if parts else "без ключа"


@dataclass
class AssembledLine:
    """What the client sees after template import (display only)."""

    row_no: int
    sku_code: str
    name: str
    size: str
    barcode: str
    requested_qty: int
    applied_qty: int
    boxes: int
    box_qty: int


@dataclass
class AssembledSkuSummary:
    """Per-article totals for the client LK summary under the form."""

    sku_code: str
    name: str
    boxes: int
    qty: int


@dataclass
class ImportResult:
    selected_boxes: dict[str, str] = field(default_factory=dict)
    partial_box_splits: list[dict[str, Any]] = field(default_factory=list)
    discrepancies: list[str] = field(default_factory=list)
    assembled_lines: list[AssembledLine] = field(default_factory=list)
    assembled_by_sku: list[AssembledSkuSummary] = field(default_factory=list)
    applied_lines: int = 0
    parse_errors: list[str] = field(default_factory=list)

    @property
    def has_discrepancies(self) -> bool:
        return bool(self.discrepancies)

    def discrepancy_comment_block(self) -> str:
        if not self.discrepancies:
            return ""
        lines = [DISCREPANCY_MARKER, "Требуется проверка менеджером:"]
        lines.extend(f"- {item}" for item in self.discrepancies)
        return "\n".join(lines)


def build_client_shipping_template_response() -> HttpResponse:
    """Prefer the saved file; if missing, build a minimal client template in memory."""
    if TEMPLATE_PATH.exists():
        response = FileResponse(
            TEMPLATE_PATH.open("rb"),
            as_attachment=True,
            filename="shk-excel.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return response
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Шаблон"
    ws.append(list(TEMPLATE_HEADERS))
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="shk-excel.xlsx"'
    return response


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _norm_key(value: Any) -> str:
    return _norm(value).casefold()


def _norm_barcode(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = str(value).strip()
    else:
        text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    # Excel sometimes returns scientific notation as text.
    try:
        as_float = float(text.replace(",", "."))
        if as_float.is_integer() and "e" in text.lower():
            return str(int(as_float))
    except ValueError:
        pass
    return text


def _cell_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError("дробное число")
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    if "." in text:
        num = float(text)
        if not num.is_integer():
            raise ValueError("дробное число")
        return int(num)
    return int(text)


def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    aliases = {
        # Product barcode (primary key for this template)
        "баркод товара": "barcode",
        "баркод": "barcode",
        "штрихкод": "barcode",
        "штрихкод товара": "barcode",
        "шк товара": "barcode",
        "шк": "barcode",
        "barcode": "barcode",
        # Quantity in pieces
        "кол-во товаров": "qty",
        "кол во товаров": "qty",
        "количество товаров": "qty",
        "количество, шт": "qty",
        "количество шт": "qty",
        "количество (шт)": "qty",
        "количество": "qty",
        "кол-во": "qty",
        "кол-во, шт": "qty",
        "кол-во шт": "qty",
        "штук": "qty",
        "шт": "qty",
        "qty": "qty",
        # Optional
        # Optional / ignored for matching: box SHK is produced by warehouse after assembly.
        "шк короба": "box_barcode",
        "штрихкод короба": "box_barcode",
        "срок годности": "expiry",
        # Article / SKU (alternative key to barcode)
        "артикул": "sku",
        "артикул товара": "sku",
        "артикул поставщика": "sku",
        "sku": "sku",
        "sku code": "sku",
        "код товара": "sku",
        "размер": "size",
        "коробов": "boxes",
        "короба": "boxes",
    }
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(header_row):
        key = aliases.get(_norm_key(raw))
        if key and key not in mapping:
            mapping[key] = idx
    return mapping


def parse_client_shipping_template(file_obj) -> tuple[list[TemplateLine], list[str]]:
    errors: list[str] = []
    try:
        if hasattr(file_obj, "seek"):
            try:
                file_obj.seek(0)
            except Exception:
                pass
        wb = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception:
        return [], ["Не удалось прочитать Excel. Загрузите файл по шаблону shk-excel."]
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["Файл пустой."]
    header_map = _header_map(tuple(rows[0]))
    if "barcode" not in header_map and "sku" not in header_map:
        return [], ["В первой строке нужен столбец «Баркод товара» или «Артикул»."]
    if "qty" not in header_map and "boxes" not in header_map:
        return [], ["Нужен столбец «Кол-во товаров»."]

    lines: list[TemplateLine] = []
    for row_no, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def get(name: str) -> Any:
            idx = header_map.get(name)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        barcode = _norm_barcode(get("barcode")) if "barcode" in header_map else ""
        sku = _norm(get("sku")) if "sku" in header_map else ""
        if not barcode and not sku:
            # Empty filler row — skip silently.
            continue
        try:
            qty = _cell_int(get("qty")) if "qty" in header_map else None
            boxes = _cell_int(get("boxes")) if "boxes" in header_map else None
        except ValueError:
            errors.append(f"Строка {row_no}: количество должно быть целым числом.")
            continue
        if qty is None and boxes is None:
            errors.append(f"Строка {row_no}: укажите «Кол-во товаров».")
            continue
        if qty is not None and qty <= 0:
            errors.append(f"Строка {row_no}: количество должно быть > 0.")
            continue
        if boxes is not None and boxes <= 0:
            errors.append(f"Строка {row_no}: число коробов должно быть > 0.")
            continue
        lines.append(
            TemplateLine(
                row_no=row_no,
                barcode=barcode,
                sku=sku,
                qty=int(qty) if qty is not None else -int(boxes or 0),  # negative qty marks legacy boxes-only
                # Box SHK from client file is ignored — warehouse issues it after pick/assembly.
                box_barcode="",
                expiry=_norm(get("expiry")),
            )
        )
    if not lines and not errors:
        errors.append("В файле нет строк с баркодом или артикулом товара.")
    return lines, errors


def _rows_for_barcode(barcode: str, stock_rows: list[dict], *, box_barcode: str = "") -> list[dict]:
    del box_barcode  # reserved; client does not supply box SHK
    matches = [
        row
        for row in stock_rows
        if _norm_barcode(row.get("barcode")) == barcode
    ]
    if not matches:
        return []
    plain = [row for row in matches if not row.get("is_mixed_box")]
    # Client does not supply box SHK — warehouse assigns it after assembly.
    return plain or matches


def _rows_for_sku(sku: str, stock_rows: list[dict]) -> list[dict]:
    key = _norm_key(sku)
    if not key:
        return []
    matches = [
        row
        for row in stock_rows
        if _norm_key(row.get("sku_code")) == key
    ]
    if not matches:
        return []
    plain = [row for row in matches if not row.get("is_mixed_box")]
    return plain or matches


def _pool_for_line(line: TemplateLine, stock_rows: list[dict]) -> list[dict]:
    """Match stock by barcode first (more specific), then by article."""
    if line.barcode:
        pool = _rows_for_barcode(line.barcode, stock_rows)
        if pool:
            return pool
        if not line.sku:
            return []
    if line.sku:
        return _rows_for_sku(line.sku, stock_rows)
    return []


def _available_for_row(row: dict, already_selected: dict[str, int], addition: dict[str, int]) -> int:
    key = str(row.get("key") or "").strip()
    if not key:
        return 0
    return max(
        int(row.get("available_boxes") or 0) - int(already_selected.get(key, 0)) - int(addition.get(key, 0)),
        0,
    )


def _allocate_exact_whole_boxes(
    *,
    qty: int,
    pool: list[dict],
    already_selected: dict[str, int],
) -> dict[str, int] | None:
    """Return an exact combination across all available box multiplicities."""
    target = max(int(qty or 0), 0)
    if target <= 0 or target > 200_000:
        return None

    whole_box_pool = sorted(
        [row for row in pool or [] if not row.get("partial_only")],
        key=lambda row: (
            0 if row.get("is_open_remainder") else 1,
            -max(int(row.get("box_qty") or 0), 0),
            str(row.get("key") or ""),
        ),
    )
    chunks: list[tuple[str, int, int]] = []
    for row in whole_box_pool:
        key = str(row.get("key") or "").strip()
        box_qty = max(int(row.get("box_qty") or 0), 0)
        available = _available_for_row(row, already_selected, {})
        if not key or box_qty <= 0 or available <= 0:
            continue
        step = 1
        left = available
        while left > 0:
            take = min(step, left)
            chunks.append((key, take, take * box_qty))
            left -= take
            step *= 2

    states: dict[int, dict[str, int]] = {0: {}}
    for key, boxes, value in chunks:
        for subtotal, allocation in list(states.items())[::-1]:
            candidate = subtotal + value
            if candidate > target or candidate in states:
                continue
            next_allocation = dict(allocation)
            next_allocation[key] = int(next_allocation.get(key, 0)) + boxes
            states[candidate] = next_allocation
        if target in states:
            return states[target]
    return None


def _piece_pick_split(
    *,
    pool: list[dict],
    remaining_qty: int,
    already_selected: dict[str, int],
    addition: dict[str, int],
) -> dict[str, Any] | None:
    """Create a client-form partial-box split for the non-multiple remainder."""
    pick_qty = max(int(remaining_qty or 0), 0)
    if pick_qty <= 0:
        return None

    candidates: list[tuple[int, int, str, dict]] = []
    for row in pool or []:
        key = str(row.get("key") or "").strip()
        box_qty = max(int(row.get("box_qty") or 0), 0)
        selected_boxes = max(
            int(already_selected.get(key, 0) or 0) + int(addition.get(key, 0) or 0),
            0,
        )
        available_boxes = max(int(row.get("available_boxes") or 0), 0) - selected_boxes
        box_codes = [
            str(code or "").strip()
            for code in (row.get("box_codes") or [])
            if str(code or "").strip()
        ]
        if (
            not key
            or row.get("is_mixed_box")
            or row.get("partial_only")
            or box_qty <= pick_qty
            or available_boxes <= 0
            or len(box_codes) <= selected_boxes
        ):
            continue
        candidates.append((box_qty, -available_boxes, key, row))

    if not candidates:
        return None

    box_qty, _negative_available, key, _row = sorted(candidates, key=lambda item: item[:3])[0]
    return {
        "identity": f"key:{key}",
        "group_key": f"partial:{key}",
        "row_key": key,
        "group": "",
        "boxes": 1,
        "items": [{"key": key, "qty": pick_qty}],
        "piece_qty": pick_qty,
        "source_box_qty": box_qty,
    }


def _allocate_pieces_across_pack_rows(
    *,
    identity: str,
    qty: int,
    pool: list[dict],
    row_no: int,
    already_selected: dict[str, int],
) -> tuple[dict[str, int], list[str], int]:
    """
    Client declares pieces and does not need to know box multiplicity.

    First find an exact combination across every available pack size. If no exact
    combination exists, consume opened remainders first, then full boxes without
    exceeding the requested quantity. Report the remaining shortfall so the
    caller can create one piece-pick split.
    Example: 445 pcs requested, pack 30, stock 15 boxes (450) → 14 boxes / 420 pcs,
    discrepancy 25 pcs.
    """
    discrepancies: list[str] = []
    sku_label = ", ".join(sorted({_norm(row.get("sku_code")) or "-" for row in pool})) or "-"
    # An opened-box remainder is a physical piece-pick source, not a whole box.
    # Its ``box_qty`` keeps the original pack size for audit purposes while
    # ``split_available_qty`` contains the actual remainder. Never feed such a
    # row into whole-box allocation: otherwise 3 remaining pieces from a
    # nominal 25-pack can be selected as one full 25-piece box.
    whole_box_pool = [row for row in pool if not row.get("partial_only")]
    pack_sizes = sorted(
        {
            max(int(row.get("box_qty") or 0), 0)
            for row in whole_box_pool
            if int(row.get("box_qty") or 0) > 0
        }
    )
    if not pack_sizes or int(qty) <= 0:
        discrepancies.append(
            f"Строка {row_no}: {identity} ({sku_label}): запрошено {qty} шт, нет данных по коробу."
        )
        return {}, discrepancies, 0

    exact = _allocate_exact_whole_boxes(
        qty=qty,
        pool=whole_box_pool,
        already_selected=already_selected,
    )
    if exact is not None:
        return exact, [], 1

    # Opened remainders are consumed before untouched full boxes.
    ranked = sorted(
        [row for row in whole_box_pool if max(int(row.get("box_qty") or 0), 0) > 0],
        key=lambda row: (
            0 if row.get("is_open_remainder") else 1,
            -max(int(row.get("box_qty") or 0), 0),
            -_available_for_row(row, already_selected, {}),
            str(row.get("key") or ""),
        ),
    )

    addition: dict[str, int] = {}
    remaining = int(qty)
    for row in ranked:
        if remaining <= 0:
            break
        key = str(row.get("key") or "").strip()
        box_qty = max(int(row.get("box_qty") or 0), 0)
        if not key or box_qty <= 0:
            continue
        available = _available_for_row(row, already_selected, addition)
        if available <= 0:
            continue
        take_boxes = min(available, remaining // box_qty)
        if take_boxes <= 0:
            continue
        addition[key] = int(addition.get(key, 0)) + take_boxes
        remaining -= take_boxes * box_qty

    applied_qty = int(qty) - remaining
    applied_boxes = sum(addition.values())
    if applied_boxes <= 0:
        packs_hint = ", ".join(str(s) for s in pack_sizes)
        min_pack = min(pack_sizes)
        discrepancies.append(
            f"Строка {row_no}: {identity} ({sku_label}): запрошено {qty} шт — "
            f"не удалось набрать целые короба [{packs_hint} шт/кор.]. "
            f"Минимум одного короба: {min_pack} шт."
        )
        return {}, discrepancies, 0

    if remaining > 0:
        box_qty_by_key = {
            str(row.get("key") or ""): max(int(row.get("box_qty") or 0), 0)
            for row in whole_box_pool
        }
        pack_detail = ", ".join(
            f"{boxes} кор.×{box_qty_by_key.get(key, 0)} шт" for key, boxes in addition.items()
        )
        discrepancies.append(
            f"Строка {row_no}: {identity} ({sku_label}): заявлено {qty} шт, "
            f"собрано {applied_boxes} кор. ({applied_qty} шт) [{pack_detail}], "
            f"расхождение {remaining} шт — можно подправить вручную."
        )
    return addition, discrepancies, 1


def _allocate_legacy_boxes(
    *,
    identity: str,
    boxes: int,
    pool: list[dict],
    row_no: int,
    already_selected: dict[str, int],
) -> tuple[dict[str, int], list[str], int]:
    """Legacy mode: quantity given as boxes; pick a unique stock row when possible."""
    whole_box_pool = [row for row in pool if not row.get("partial_only")]
    if not whole_box_pool:
        return {}, [
            f"Строка {row_no}: {identity}: доступны только остатки открытых коробов. "
            "Укажите количество в штуках для поштучного отбора."
        ], 0
    with_stock = [
        row
        for row in whole_box_pool
        if max(int(row.get("available_boxes") or 0) - int(already_selected.get(str(row.get("key") or ""), 0)), 0) > 0
    ]
    candidates = with_stock or whole_box_pool
    if len(candidates) != 1:
        # Prefer unique box_qty among candidates that can fulfill.
        fulfillable = [
            row
            for row in candidates
            if max(int(row.get("available_boxes") or 0) - int(already_selected.get(str(row.get("key") or ""), 0)), 0)
            >= boxes
        ]
        if len(fulfillable) == 1:
            candidates = fulfillable
        elif len({int(row.get("box_qty") or 0) for row in candidates}) == 1 and len(candidates) > 1:
            # Same pack size duplicates — take one with most remaining stock.
            candidates = [
                max(
                    candidates,
                    key=lambda row: max(
                        int(row.get("available_boxes") or 0)
                        - int(already_selected.get(str(row.get("key") or ""), 0)),
                        0,
                    ),
                )
            ]
        else:
            skus = ", ".join(sorted({_norm(row.get("sku_code")) or "-" for row in whole_box_pool}))
            packs = ", ".join(
                sorted({str(int(row.get("box_qty") or 0)) for row in whole_box_pool})
            )
            return {}, [
                f"Строка {row_no}: {identity} найден в нескольких остатках ({skus}, короба {packs} шт). "
                f"Укажите «Кол-во товаров» в штуках, кратное одному типу короба."
            ], 0

    row = candidates[0]
    key = str(row.get("key") or "").strip()
    if not key:
        return {}, [f"Строка {row_no}: {identity}: техническая ошибка ключа остатка."], 0
    available = max(int(row.get("available_boxes") or 0) - int(already_selected.get(key, 0)), 0)
    label = f"{row.get('sku_code')}/{row.get('size') or '-'} ({identity})"
    if available <= 0:
        return {}, [f"{label}: запрошено {boxes} кор., доступно 0 кор."], 0
    apply_boxes = boxes
    discrepancies: list[str] = []
    if boxes > available:
        discrepancies.append(f"{label}: запрошено {boxes} кор., доступно {available} кор.")
        apply_boxes = available
    return {key: apply_boxes}, discrepancies, 1


def apply_client_shipping_template(*, file_obj, stock_rows: list[dict]) -> ImportResult:
    result = ImportResult()
    lines, parse_errors = parse_client_shipping_template(file_obj)
    result.parse_errors = parse_errors
    if parse_errors:
        return result

    selected: dict[str, int] = {}
    occupied: dict[str, int] = {}
    stock_by_key = {str(row.get("key") or ""): row for row in stock_rows or []}

    for line in lines:
        identity = line.identity_label()
        pool = _pool_for_line(line, stock_rows)
        if not pool:
            result.discrepancies.append(
                f"Строка {line.row_no}: {identity} не найден в доступных остатках."
            )
            continue

        if line.qty < 0:
            addition, discrepancies, applied = _allocate_legacy_boxes(
                identity=identity,
                boxes=abs(line.qty),
                pool=pool,
                row_no=line.row_no,
                already_selected=occupied,
            )
            requested_pieces = None
            piece_split = None
        else:
            addition, discrepancies, applied = _allocate_pieces_across_pack_rows(
                identity=identity,
                qty=line.qty,
                pool=pool,
                row_no=line.row_no,
                already_selected=occupied,
            )
            requested_pieces = int(line.qty)
            whole_qty = 0
            for key, boxes in addition.items():
                row = stock_by_key.get(key) or {}
                whole_qty += int(boxes) * max(int(row.get("box_qty") or 0), 0)
            remaining_qty = max(int(line.qty) - whole_qty, 0)
            piece_split = _piece_pick_split(
                pool=pool,
                remaining_qty=remaining_qty,
                already_selected=occupied,
                addition=addition,
            )
            if piece_split:
                discrepancies = []
                applied = 1

        result.discrepancies.extend(discrepancies)
        line_parts: list[AssembledLine] = []
        for key, boxes in addition.items():
            selected[key] = int(selected.get(key, 0)) + int(boxes)
            occupied[key] = int(occupied.get(key, 0)) + int(boxes)
            row = stock_by_key.get(key) or {}
            box_qty = max(int(row.get("box_qty") or 0), 0)
            applied_qty = int(boxes) * box_qty
            line_parts.append(
                AssembledLine(
                    row_no=line.row_no,
                    sku_code=_norm(row.get("sku_code")),
                    name=_norm(row.get("name")),
                    size=_norm(row.get("size")),
                    barcode=_norm_barcode(row.get("barcode")) or line.barcode,
                    requested_qty=int(requested_pieces) if requested_pieces is not None else applied_qty,
                    applied_qty=applied_qty,
                    boxes=int(boxes),
                    box_qty=box_qty,
                )
            )
        if piece_split:
            split_key = str(piece_split.get("row_key") or "").strip()
            piece_qty = max(int(piece_split.get("piece_qty") or 0), 0)
            if split_key and piece_qty > 0:
                occupied[split_key] = int(occupied.get(split_key, 0)) + max(int(piece_split.get("boxes") or 1), 1)
                result.partial_box_splits.append(piece_split)
                row = stock_by_key.get(split_key) or {}
                line_parts.append(
                    AssembledLine(
                        row_no=line.row_no,
                        sku_code=_norm(row.get("sku_code")),
                        name=_norm(row.get("name")),
                        size=_norm(row.get("size")),
                        barcode=_norm_barcode(row.get("barcode")) or line.barcode,
                        requested_qty=int(requested_pieces) if requested_pieces is not None else piece_qty,
                        applied_qty=piece_qty,
                        boxes=max(int(piece_split.get("boxes") or 1), 1),
                        box_qty=max(int(row.get("box_qty") or 0), 0),
                    )
                )
        result.assembled_lines.extend(line_parts)
        result.applied_lines += applied

    for key, boxes in list(occupied.items()):
        row = stock_by_key.get(key)
        if row is None:
            continue
        available = max(int(row.get("available_boxes") or 0), 0)
        if boxes > available:
            label = f"{row.get('sku_code')}/{row.get('size') or '-'} (ШК {_norm_barcode(row.get('barcode'))})"
            result.discrepancies.append(
                f"{label}: суммарно по шаблону {boxes} кор., доступно {available}."
            )
            if selected.get(key, 0) > available:
                selected[key] = available

    result.selected_boxes = {key: str(boxes) for key, boxes in selected.items() if boxes > 0}

    # Aggregate boxes/qty by article for the client bottom summary.
    by_sku: dict[str, AssembledSkuSummary] = {}
    for line in result.assembled_lines:
        sku = _norm(line.sku_code) or "-"
        current = by_sku.get(sku)
        if current is None:
            by_sku[sku] = AssembledSkuSummary(
                sku_code=sku,
                name=_norm(line.name),
                boxes=int(line.boxes),
                qty=int(line.applied_qty),
            )
        else:
            current.boxes += int(line.boxes)
            current.qty += int(line.applied_qty)
            if not current.name and line.name:
                current.name = _norm(line.name)
    result.assembled_by_sku = sorted(by_sku.values(), key=lambda item: item.sku_code.casefold())
    return result


def merge_comment_with_discrepancies(comment: str, block: str) -> str:
    comment = str(comment or "").strip()
    block = str(block or "").strip()
    if not block:
        return comment
    if DISCREPANCY_MARKER in comment:
        head = comment.split(DISCREPANCY_MARKER, 1)[0].rstrip()
        comment = head
    if comment:
        return f"{comment}\n\n{block}"
    return block
