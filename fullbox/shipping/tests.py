from datetime import date, datetime, timedelta
from io import BytesIO
import json
from pathlib import Path
import shutil
from types import SimpleNamespace
import tempfile
from unittest.mock import patch
import zipfile

from django.contrib.auth import get_user_model
from django.contrib.messages.storage.fallback import FallbackStorage
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import models
from django.template.loader import render_to_string
from django.test import RequestFactory
from django.test import SimpleTestCase
from django.test import TestCase
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook

from audit.models import OrderAuditEntry
from billing.models import BillingService, WarehouseServiceFact
from employees.models import Employee
from head_manager.models import Carrier, OwnCompany
from logistics.models import LogisticsTrip, LogisticsTripOrder, ShippingRoutingState
from reachtruck.models import BoxClaim, MoveRequest, MoveRequestItem, MoveTask
from sklad.models import (
    WarehouseContainer,
    WarehouseEvent,
    WarehouseLocation,
    WarehouseReserve,
    WarehouseStockSnapshot,
)
from sklad.services import WarehouseGoodsStateResolver, WarehouseStateCode
from sklad.services.warehouse_write_path import WarehouseWritePathService
from sklad.test_utils import create_warehouse_snapshot_row
from sku.models import Agency, Market, SKU, SKUBarcode
from todo.models import Task

from .box_splits import encode_partial_box_split, extract_partial_box_split
from .forms import ShippingOrderForm
from .item_binding import (
    enforced_marketplace_item_binding_errors,
    enforced_ozon_request_binding_errors,
    shipping_final_truth_rows,
    validate_marketplace_item_binding,
    validate_ozon_request_binding,
)
from .discrepancy import (
    approve_shipping_discrepancy,
    confirm_shipping_discrepancy_pick,
    mark_supplement_awaiting_packing,
    reject_shipping_discrepancy,
    request_shipping_discrepancy,
    resolve_completed_discrepancy_supplement,
    shipping_discrepancy_context,
)
from .actions import ShippingDetailActionPermissions, handle_shipping_detail_action
from .models import ShippingOrder, ShippingOrderAttachment, ShippingOrderItem, ShippingTransportNote
from .packing import (
    _shipping_item_key_from_payload,
    _shipping_normalize_loose_packing_boxes,
    _shipping_packing_initial_state_from_draft,
    _shipping_repack_items,
    save_shipping_loose_packing,
)
from .ozon_supplies import (
    _ozon_label_groups,
    allocate_reserve_box_codes,
    apply_ozon_bundle_to_stock,
    fetch_ozon_gm_label_documents,
    get_ozon_supplies_batch_for_agency,
    lookup_existing_ozon_supply_shipments,
)
from .return_act import _item_rows
from .services import (
    _apply_planned_ship_date,
    _build_ozon_api_snapshot_payload,
    _comment_with_ozon_api_meta,
    _comment_with_reselected_box_codes,
    _format_shipping_reserve_shortage,
    _ozon_gm_rows_with_exact_fullbox_match,
    _ozon_gm_repack_context,
    _ozon_loose_packing_gm_plan,
    _ozon_meta_for_reserve,
    _reconcile_ozon_gm_meta_with_selected_rows,
    _saved_ozon_api_intake_meta,
    _ozon_reserved_fullbox_boxes,
    _ozon_supply_meta_from_request,
    _reselect_selected_boxes_locked,
    _storekeeper_marketplace_binding_alert,
    _strict_ozon_gm_intake_errors,
    _visible_shipping_integrity_warnings,
    build_shipping_detail_page_context,
    build_shipping_list_page_context,
    complete_transfer_without_trip,
    create_ozon_api_shipping_orders_batch,
    create_pick_tasks,
    ensure_storekeeper_task,
    handle_shipping_loose_packing_request,
    next_shipping_number,
    release_order_reserves,
    refresh_strict_ozon_gm_intake,
    reserve_order,
    shipping_pick_readiness,
    ship_order,
)
from .dispatch import confirm_dispatch_act_from_loading, shipping_dispatch_payload
from .workflow import cancel_order
from .truth import ShippingTruthService
from .web_ui import (
    _attach_ozon_stock_shortages,
    _ozon_piece_pick_tariff_confirmation_error,
    _selected_box_values_for_order,
)
from .views import (
    _can_cancel,
    _can_edit_items,
    _can_edit_order_form,
    _display_shipping_number,
    _order_box_count,
    _parse_selected_stock_items,
    _shipping_box_row_key,
    _shipping_boxes_from_packing_payload,
    _shipping_packing_summary,
    _shipping_packing_initial_state,
    _shipping_delivered_boxes,
    _shipping_reachtruck_metrics,
    _shipping_stock_picker_rows,
    save_shipping_packing,
)


class LoosePackingDesktopQrPrintTemplateTests(SimpleTestCase):
    def test_box_qr_uses_labels_renderer_and_desktop_direct_print(self):
        template_path = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "shipping"
            / "loose_packing.html"
        )
        source = template_path.read_text(encoding="utf-8")

        self.assertIn('id="labels-box-renderer"', source)
        self.assertIn("templateKey: 'box'", source)
        self.assertIn("window.fullboxDesktop.printPng", source)
        self.assertIn("desktopPrintRecordUrl", source)
        self.assertIn("printRenderedLabelInBrowser", source)
        self.assertNotIn("new QRCode(host,{text:${JSON.stringify(box.code)}", source)

    def test_large_marking_draft_is_not_sent_as_keepalive_request(self):
        template_path = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "shipping"
            / "loose_packing.html"
        )
        source = template_path.read_text(encoding="utf-8")
        flush_start = source.index("function flushPackingDraftToServer()")
        flush_end = source.index("function persistPackingDraft()", flush_start)
        flush_source = source[flush_start:flush_end]

        self.assertIn("fetch(packingDraftUrl", flush_source)
        self.assertNotIn("keepalive", flush_source)

    def test_physical_repack_requires_source_box_scans(self):
        template_path = (
            Path(__file__).resolve().parents[1]
            / "templates"
            / "shipping"
            / "loose_packing.html"
        )
        source = template_path.read_text(encoding="utf-8")

        self.assertIn('id="source-box-scan-input"', source)
        self.assertIn('name="repack_source_boxes_json"', source)
        self.assertIn("allRepackSourcesScanned()", source)


class ShippingOzonApiSnapshotTests(SimpleTestCase):
    def test_snapshot_is_normalized_complete_and_hash_is_stable(self):
        meta = {
            "payload": {"fetched_at": "2026-08-06T10:00:00+03:00"},
            "order_number": "200006189 7823",
            "boxes_total": 1,
            "composition": [{"barcode": "460000000001", "quantity": 20}],
            "destinations": ["СТАРАЯ_КУПАВНА_28"],
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-001",
                    "cargo_id": "cargo-1",
                    "items": [
                        {
                            "offer_id": "SKU-1",
                            "barcode": "460000000001",
                            "name": "Товар",
                            "quantity": 4,
                            "ozon_sku": "123456",
                        }
                    ],
                },
                *[
                    {
                        "gm_barcode": f"GM-00{index}",
                        "cargo_id": f"cargo-{index}",
                        "items": [
                            {
                                "offer_id": "SKU-1",
                                "barcode": "460000000001",
                                "name": "Товар",
                                "quantity": quantity,
                                "ozon_sku": "123456",
                            }
                        ],
                    }
                    for index, quantity in ((2, 4), (3, 3), (4, 3), (5, 3), (6, 3))
                ],
            ],
        }

        first = _build_ozon_api_snapshot_payload(meta)
        second = _build_ozon_api_snapshot_payload({**meta, "payload": {}})

        self.assertEqual(first["ozon_supply_id_normalized"], "2000061897823")
        self.assertTrue(first["ozon_snapshot_complete"])
        self.assertEqual(first["ozon_snapshot_schema_version"], 1)
        self.assertEqual(first["ozon_snapshot_source"], "ozon_api")
        self.assertEqual(first["gm_cargoes"][0]["items"][0]["quantity"], 20)
        self.assertEqual(len(first["ozon_snapshot_hash"]), 64)
        self.assertEqual(first["ozon_snapshot_hash"], second["ozon_snapshot_hash"])

    def test_snapshot_is_incomplete_without_exact_gm_item_composition(self):
        snapshot = _build_ozon_api_snapshot_payload(
            {
                "order_number": "2000061897823",
                "boxes_total": 1,
                "gm_cargoes": [{"gm_barcode": "GM-001", "items": []}],
            }
        )

        self.assertFalse(snapshot["ozon_snapshot_complete"])

    def test_combined_snapshot_keeps_all_supply_ids_and_more_than_fifty_gm_labels(self):
        gm_cargoes = [
            {
                "gm_barcode": f"GM-{index:03d}",
                "ozon_supply_number": "OZ-1" if index <= 30 else "OZ-2",
                "ozon_order_id": "101" if index <= 30 else "202",
                "items": [
                    {
                        "offer_id": f"SKU-{index}",
                        "barcode": f"46000000{index:04d}",
                        "quantity": 1,
                    }
                ],
            }
            for index in range(1, 61)
        ]

        snapshot = _build_ozon_api_snapshot_payload(
            {
                "order_number": "OZ-1 + ещё 1",
                "supply_numbers": ["OZ-1", "OZ-2"],
                "order_ids": ["101", "202"],
                "boxes_total": 60,
                "gm_cargoes": gm_cargoes,
            }
        )

        self.assertTrue(snapshot["ozon_snapshot_complete"])
        self.assertEqual(snapshot["ozon_supply_numbers"], ["OZ-1", "OZ-2"])
        self.assertEqual(snapshot["ozon_order_ids"], ["101", "202"])
        self.assertEqual(snapshot["ozon_supply_count"], 2)
        self.assertEqual(len(snapshot["gm_cargoes"]), 60)
        self.assertEqual(snapshot["gm_cargoes"][-1]["ozon_supply_number"], "OZ-2")

    def test_combined_request_meta_keeps_constituent_supplies_but_not_gm_box_binding_for_reserve(self):
        payload = {
            "order_id": "101",
            "order_number": "OZ-1 + ещё 1",
            "is_combined_ozon_supply": True,
            "supplies": [
                {"order_id": "101", "order_number": "OZ-1", "form": {}},
                {"order_id": "202", "order_number": "OZ-2", "form": {}},
            ],
            "ozon_boxes_total": 2,
            "gm_barcodes": ["GM-1", "GM-2"],
            "gm_cargoes": [
                {"gm_barcode": "GM-1", "ozon_supply_number": "OZ-1", "items": []},
                {"gm_barcode": "GM-2", "ozon_supply_number": "OZ-2", "items": []},
            ],
            "form": {
                "supply_number": "OZ-1 + ещё 1",
                "shipping_barcode": "OZ-1 + ещё 1",
            },
        }
        request = RequestFactory().post(
            "/shipping/new/",
            data={"ozon_supply_payload_json": json.dumps(payload)},
        )

        meta = _ozon_supply_meta_from_request(request)
        comment = _comment_with_ozon_api_meta("Комментарий клиента", meta)
        reserve_meta = _ozon_meta_for_reserve(meta)

        self.assertTrue(meta["is_combined"])
        self.assertEqual(meta["supply_numbers"], ["OZ-1", "OZ-2"])
        self.assertEqual(meta["order_ids"], ["101", "202"])
        self.assertIn("Поставок Ozon: 2", comment)
        self.assertIn("Номера поставок Ozon: OZ-1, OZ-2", comment)
        self.assertEqual(reserve_meta["gm_cargoes"], [])
        self.assertEqual(len(meta["gm_cargoes"]), 2)

    def test_removed_product_also_removes_its_exact_gm_cargo(self):
        cargoes = [
            {
                "gm_barcode": "GM-FLOWERS-16",
                "items": [
                    {
                        "offer_id": "flowers016",
                        "barcode": "2047519259817",
                        "quantity": 49,
                    }
                ],
            },
            {
                "gm_barcode": "GM-FLOWERS-11",
                "items": [
                    {
                        "offer_id": "flowers011",
                        "barcode": "2043986017738",
                        "quantity": 32,
                    }
                ],
            },
            {
                "gm_barcode": "GM-KNIFE-18",
                "ozon_order_id": "122340970",
                "items": [
                    {
                        "offer_id": "knife018",
                        "barcode": "2042857120430",
                        "quantity": 50,
                    }
                ],
            },
        ]
        meta = {
            "payload": {
                "gm_barcodes": [cargo["gm_barcode"] for cargo in cargoes],
                "gm_cargoes": cargoes,
                "ozon_boxes_total": 3,
                "form": {"ozon_gm_comment": "ШК ГМ Ozon: all"},
            },
            "gm_barcodes": [cargo["gm_barcode"] for cargo in cargoes],
            "gm_cargoes": cargoes,
            "boxes_total": 3,
            "ozon_gm_comment": "ШК ГМ Ozon: all",
        }

        reconciled, errors = _reconcile_ozon_gm_meta_with_selected_rows(
            agency=None,
            meta=meta,
            selected_rows=[
                {
                    "sku_code": "knife018",
                    "barcode": "2042857120430",
                    "qty_requested": 50,
                }
            ],
        )

        self.assertEqual(errors, [])
        self.assertEqual(reconciled["boxes_total"], 1)
        self.assertEqual(reconciled["gm_barcodes"], ["GM-KNIFE-18"])
        self.assertEqual(reconciled["gm_cargoes"][0]["gm_barcode"], "GM-KNIFE-18")
        self.assertTrue(reconciled["ozon_gm_reconciled_after_item_removal"])
        self.assertEqual(
            reconciled["removed_gm_barcodes"],
            ["GM-FLOWERS-16", "GM-FLOWERS-11"],
        )
        self.assertEqual(reconciled["order_ids"], ["122340970"])
        self.assertEqual(reconciled["payload"]["ozon_boxes_total"], 1)
        self.assertEqual(
            reconciled["payload"]["form"]["ozon_gm_comment"],
            "ШК ГМ Ozon: GM-KNIFE-18",
        )

    def test_mixed_gm_cargo_cannot_be_partially_removed(self):
        meta = {
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-MIXED",
                    "items": [
                        {"offer_id": "KEEP", "barcode": "KEEP-BC", "quantity": 5},
                        {"offer_id": "DELETE", "barcode": "DELETE-BC", "quantity": 7},
                    ],
                }
            ],
            "gm_barcodes": ["GM-MIXED"],
            "boxes_total": 1,
        }

        reconciled, errors = _reconcile_ozon_gm_meta_with_selected_rows(
            agency=None,
            meta=meta,
            selected_rows=[
                {"sku_code": "KEEP", "barcode": "KEEP-BC", "qty_requested": 5}
            ],
        )

        self.assertTrue(any("нельзя удалить частично" in error for error in errors))
        self.assertEqual(reconciled["gm_barcodes"], ["GM-MIXED"])
        self.assertEqual(reconciled["boxes_total"], 1)

    def test_quantity_change_must_match_remaining_gm_cargoes(self):
        meta = {
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-50",
                    "items": [{"offer_id": "KNIFE", "barcode": "KNIFE-BC", "quantity": 50}],
                }
            ],
            "gm_barcodes": ["GM-50"],
            "boxes_total": 1,
        }

        _reconciled, errors = _reconcile_ozon_gm_meta_with_selected_rows(
            agency=None,
            meta=meta,
            selected_rows=[
                {"sku_code": "KNIFE", "barcode": "KNIFE-BC", "qty_requested": 25}
            ],
        )

        self.assertTrue(any("в заявке 25 шт." in error for error in errors))

    def test_missing_gm_composition_blocks_item_removal_without_guessing(self):
        meta = {
            "gm_cargoes": [{"gm_barcode": "GM-PENDING", "items": []}],
            "gm_barcodes": ["GM-PENDING"],
            "boxes_total": 1,
        }

        with patch(
            "shipping.ozon_supplies.resolve_ozon_credential",
            return_value=(None, "credential unavailable"),
        ):
            reconciled, errors = _reconcile_ozon_gm_meta_with_selected_rows(
                agency=None,
                meta=meta,
                selected_rows=[],
            )

        self.assertTrue(any("невозможно однозначно удалить" in error for error in errors))
        self.assertEqual(reconciled["gm_barcodes"], ["GM-PENDING"])


class ShippingOzonLoosePackingPlanTests(SimpleTestCase):
    def setUp(self):
        self.order = SimpleNamespace(
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=SimpleNamespace(name="Ozon"),
        )

    @staticmethod
    def _loose_item(barcode: str, qty: int) -> dict:
        item = {
            "sku_code": f"SKU-{barcode}",
            "name": f"Товар {barcode}",
            "size": "0",
            "barcode": barcode,
            "goods_type": "gv",
            "qty": qty,
        }
        item["key"] = _shipping_item_key_from_payload(item)
        return item

    @staticmethod
    def _cargo(gm_barcode: str, barcode: str, qty: int) -> dict:
        return {
            "gm_barcode": gm_barcode,
            "supply_id": "SUPPLY-1",
            "bundle_id": f"BUNDLE-{gm_barcode}",
            "items": [{"barcode": barcode, "quantity": qty}],
        }

    def test_plan_allows_unambiguous_supplement_of_existing_box(self):
        loose_items = [
            self._loose_item("B23", 1),
            self._loose_item("B21", 25),
            self._loose_item("B24", 25),
        ]
        summary = {
            "gm_cargoes": [
                self._cargo("GM-23", "B23", 25),
                self._cargo("GM-22", "B22", 32),
                self._cargo("GM-21", "B21", 25),
                self._cargo("GM-24", "B24", 25),
            ]
        }
        delivered = [
            {"box_code": "BOX-23", "items": [{"barcode": "B23", "qty": 24}]},
            {"box_code": "BOX-22", "items": [{"barcode": "B22", "qty": 32}]},
        ]

        with patch("shipping.packing._shipping_delivered_boxes", return_value=delivered):
            with patch("shipping.services._ozon_api_summary_for_form", return_value=summary):
                plan = _ozon_loose_packing_gm_plan(self.order, loose_items=loose_items)

        self.assertTrue(plan["ready"], plan["errors"])
        self.assertEqual(plan["target_count"], 3)
        targets = {target["gm_barcode"]: target for target in plan["targets"]}
        self.assertEqual(targets["GM-23"]["existing_box_code"], "BOX-23")
        self.assertEqual(targets["GM-23"]["existing_qty"], 24)
        self.assertEqual(targets["GM-23"]["ozon_qty"], 25)
        self.assertEqual(targets["GM-23"]["items"][0]["qty"], 1)
        self.assertFalse(targets["GM-21"]["existing_box_code"])

    def test_plan_keeps_ambiguous_existing_box_match_blocked(self):
        loose_items = [self._loose_item("B23", 2)]
        summary = {
            "gm_cargoes": [
                self._cargo("GM-23-A", "B23", 25),
                self._cargo("GM-23-B", "B23", 25),
            ]
        }
        delivered = [
            {"box_code": "BOX-A", "items": [{"barcode": "B23", "qty": 24}]},
            {"box_code": "BOX-B", "items": [{"barcode": "B23", "qty": 24}]},
        ]

        with patch("shipping.packing._shipping_delivered_boxes", return_value=delivered):
            with patch("shipping.services._ozon_api_summary_for_form", return_value=summary):
                plan = _ozon_loose_packing_gm_plan(self.order, loose_items=loose_items)

        self.assertFalse(plan["ready"])
        self.assertTrue(any("физическая переупаковка" in error for error in plan["errors"]))

    def test_plan_rebuilds_all_source_boxes_into_ozon_gm_boxes(self):
        loose_items = [
            self._loose_item("ozn-a", 10),
            self._loose_item("OZN-B", 5),
        ]
        summary = {
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-1",
                    "supply_id": "SUPPLY-1",
                    "bundle_id": "BUNDLE-1",
                    "items": [
                        {"barcode": "OZN-A", "quantity": 4},
                        {"barcode": "ozn-b", "quantity": 1},
                    ],
                },
                self._cargo("GM-2", "OZN-A", 6),
                self._cargo("GM-3", "OZN-B", 4),
            ]
        }
        delivered = [
            {"box_code": "SOURCE-A", "items": [{"barcode": "ozn-a", "qty": 10}]},
            {"box_code": "SOURCE-B", "items": [{"barcode": "OZN-B", "qty": 5}]},
        ]

        with patch("shipping.packing._shipping_delivered_boxes", return_value=delivered):
            with patch("shipping.services._ozon_api_summary_for_form", return_value=summary):
                plan = _ozon_loose_packing_gm_plan(
                    self.order,
                    loose_items=loose_items,
                    repack_all_delivered=True,
                )

        self.assertTrue(plan["ready"], plan["errors"])
        self.assertEqual(plan["target_count"], 3)
        self.assertEqual(
            {target["gm_barcode"] for target in plan["targets"]},
            {"GM-1", "GM-2", "GM-3"},
        )
        self.assertTrue(all(not target["existing_box_code"] for target in plan["targets"]))
        self.assertEqual(
            sum(item["qty"] for target in plan["targets"] for item in target["items"]),
            15,
        )

    def test_normalizer_reuses_only_the_planned_existing_box(self):
        loose_item = self._loose_item("B23", 1)
        target = {
            "gm_barcode": "GM-23",
            "supply_id": "SUPPLY-1",
            "bundle_id": "BUNDLE-23",
            "existing_box_code": "BOX-23",
            "items": [dict(loose_item)],
        }
        boxes = [
            {
                "code": "BOX-23",
                "gm_barcode": "GM-23",
                "sealed": True,
                "items": [dict(loose_item)],
            }
        ]

        prepared, errors = _shipping_normalize_loose_packing_boxes(
            boxes,
            {loose_item["key"]: loose_item},
            gm_targets=[target],
        )

        self.assertEqual(errors, [])
        self.assertEqual(prepared[0]["existing_box_code"], "BOX-23")

        boxes[0]["code"] = "NEW-BOX"
        _prepared, wrong_box_errors = _shipping_normalize_loose_packing_boxes(
            boxes,
            {loose_item["key"]: loose_item},
            gm_targets=[target],
        )
        self.assertTrue(any("существующем коробе BOX-23" in error for error in wrong_box_errors))


class ShippingOzonLoosePackingAliasTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Ozon GM aliases")
        self.marketplace = Market.objects.create(id=1707, name="Ozon")
        self.order = ShippingOrder.objects.create(
            number="OTG-GM-ALIASES-1",
            agency=self.agency,
            marketplace=self.marketplace,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=7,
        )
        sku = SKU.objects.create(
            agency=self.agency,
            market=self.marketplace,
            sku_code="091124_redboardbear",
            name="Маркированный мольберт",
            honest_sign=True,
        )
        SKUBarcode.objects.create(
            sku=sku,
            value="4621320255022",
            size="0",
            is_primary=True,
        )
        SKUBarcode.objects.create(
            sku=sku,
            value="04621320255022",
            size="0",
        )
        SKUBarcode.objects.create(
            sku=sku,
            value="OZN1769869026",
            size="0",
        )

    @staticmethod
    def _loose_item(barcode: str, qty: int) -> dict:
        item = {
            "sku_code": "091124_redboardbear",
            "name": "Маркированный мольберт",
            "size": "0",
            "barcode": barcode,
            "goods_type": "gv",
            "qty": qty,
            "requires_marking_scan": True,
            "marking_preverified": False,
        }
        item["key"] = _shipping_item_key_from_payload(item)
        return item

    def test_repack_matches_gtin_and_ozon_aliases_but_keeps_kiz_scan(self):
        loose_items = [
            self._loose_item("4621320255022", 2),
            self._loose_item("OZN1769869026", 5),
        ]
        cargoes = [
            {
                "gm_barcode": f"GM-{index}",
                "supply_id": "SUPPLY-1",
                "bundle_id": f"BUNDLE-{index}",
                "items": [
                    {
                        "offer_id": "091124_redboardbear",
                        "barcode": "04621320255022",
                        "quantity": 1,
                    }
                ],
            }
            for index in range(1, 8)
        ]

        with patch("shipping.packing._shipping_delivered_boxes", return_value=[]), patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value={"gm_cargoes": cargoes},
        ):
            plan = _ozon_loose_packing_gm_plan(
                self.order,
                loose_items=loose_items,
                repack_all_delivered=True,
            )

        self.assertTrue(plan["ready"], plan["errors"])
        self.assertEqual(plan["target_count"], 7)
        target_key_totals = {}
        for target in plan["targets"]:
            for item in target["items"]:
                target_key_totals[item["key"]] = (
                    target_key_totals.get(item["key"], 0) + int(item["qty"])
                )
        self.assertEqual(
            target_key_totals,
            {item["key"]: item["qty"] for item in loose_items},
        )

        expected_by_key = {item["key"]: item for item in loose_items}
        boxes = [
            {
                "code": f"TARGET-{index}",
                "gm_barcode": target["gm_barcode"],
                "sealed": True,
                "items": [
                    {**item, "marking_units": []}
                    for item in target["items"]
                ],
            }
            for index, target in enumerate(plan["targets"], start=1)
        ]
        _prepared, errors = _shipping_normalize_loose_packing_boxes(
            boxes,
            expected_by_key,
            gm_targets=plan["targets"],
        )
        self.assertTrue(any("Data Matrix: 0 из 1" in error for error in errors))

    def test_plan_matches_unregistered_ozon_offer_by_unique_product_barcode(self):
        loose_items = [self._loose_item("4621320255022", 1)]
        summary = {
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-EXTERNAL-OFFER",
                    "supply_id": "SUPPLY-1",
                    "bundle_id": "BUNDLE-1",
                    "items": [
                        {
                            "offer_id": "ozon-article-not-in-client-catalog",
                            "barcode": "4621320255022",
                            "quantity": 1,
                        }
                    ],
                }
            ]
        }

        with patch("shipping.packing._shipping_delivered_boxes", return_value=[]), patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value=summary,
        ):
            plan = _ozon_loose_packing_gm_plan(self.order, loose_items=loose_items)

        self.assertTrue(plan["ready"], plan["errors"])
        self.assertEqual(plan["target_count"], 1)
        self.assertEqual(plan["targets"][0]["gm_barcode"], "GM-EXTERNAL-OFFER")

    def test_plan_keeps_known_conflicting_client_sku_blocked(self):
        conflicting_sku = SKU.objects.create(
            agency=self.agency,
            market=self.marketplace,
            sku_code="known-conflicting-offer",
            name="Другой товар клиента",
        )
        SKUBarcode.objects.create(
            sku=conflicting_sku,
            value="CONFLICTING-BARCODE",
            size="0",
            is_primary=True,
        )
        loose_items = [self._loose_item("4621320255022", 1)]
        summary = {
            "gm_cargoes": [
                {
                    "gm_barcode": "GM-CONFLICT",
                    "supply_id": "SUPPLY-1",
                    "bundle_id": "BUNDLE-1",
                    "items": [
                        {
                            "offer_id": "known-conflicting-offer",
                            "barcode": "4621320255022",
                            "quantity": 1,
                        }
                    ],
                }
            ]
        }

        with patch("shipping.packing._shipping_delivered_boxes", return_value=[]), patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value=summary,
        ):
            plan = _ozon_loose_packing_gm_plan(self.order, loose_items=loose_items)

        self.assertFalse(plan["ready"])
        self.assertTrue(any("не совпадает" in error for error in plan["errors"]))

    def test_exact_existing_boxes_only_pack_missing_loose_gm_box(self):
        loose_items = [self._loose_item("4621320255022", 1)]
        delivered = [
            {
                "box_code": f"SOURCE-{index}",
                "items": [
                    {
                        "sku_code": "091124_redboardbear",
                        "barcode": "4621320255022",
                        "qty": 1,
                    }
                ],
            }
            for index in range(1, 7)
        ]
        summary = {
            "gm_cargoes": [
                {
                    "gm_barcode": f"GM-{index}",
                    "supply_id": "SUPPLY-1",
                    "bundle_id": f"BUNDLE-{index}",
                    "items": [
                        {
                            "offer_id": "ozon-article-not-in-client-catalog",
                            "barcode": "4621320255022",
                            "quantity": 1,
                        }
                    ],
                }
                for index in range(1, 8)
            ]
        }

        with patch(
            "shipping.packing._shipping_delivered_boxes",
            return_value=delivered,
        ), patch(
            "shipping.packing._shipping_loose_items",
            return_value=loose_items,
        ), patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value=summary,
        ):
            repack = _ozon_gm_repack_context(
                self.order,
                delivered_boxes=delivered,
            )

        self.assertFalse(repack["required"])
        self.assertTrue(repack["gm_plan"]["ready"])


class ShippingOzonExistingBoxRepackTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="ozon_gm_repack_storekeeper",
            password="pwd",
        )
        self.agency = Agency.objects.create(agn_name="Ozon GM Repack")
        self.marketplace = Market.objects.create(id=707, name="Ozon")
        self.order = ShippingOrder.objects.create(
            number="OTG-REPACK-1",
            agency=self.agency,
            marketplace=self.marketplace,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=3,
        )
        self.location = WarehouseWritePathService.ensure_location(
            warehouse_code="MSK",
            zone_code="OTG",
        )
        self._snapshot("SOURCE-A", "SKU-A", "ozn-a", 10)
        self._snapshot("SOURCE-B", "SKU-B", "OZN-B", 5)
        self._loose_snapshot("SKU-A", "ozn-a", 2)

    def _snapshot(self, box_code: str, sku_code: str, barcode: str, qty: int):
        container = WarehouseContainer.objects.create(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_BOX,
            container_code=box_code,
            current_location=self.location,
        )
        event = WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="otg_arrived",
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            container=container,
            to_location=self.location,
            to_zone_code="OTG",
            qty=qty,
            occurred_at=timezone.now(),
        )
        WarehouseStockSnapshot.objects.create(
            agency=self.agency,
            source_context_type="receiving",
            source_context_id="PR-REPACK",
            sku_code=sku_code,
            name=sku_code,
            size="0",
            barcode=barcode,
            goods_type="gv",
            qty=qty,
            available_qty=0,
            container=container,
            container_code=box_code,
            location=self.location,
            zone_code="OTG",
            zone_kind=self.location.zone_kind,
            warehouse_state_code=WarehouseStateCode.IN_OTG.value,
            last_event=event,
        )

    def _loose_snapshot(self, sku_code: str, barcode: str, qty: int):
        event = WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="otg_arrived",
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            to_location=self.location,
            to_zone_code="OTG",
            qty=qty,
            occurred_at=timezone.now(),
        )
        WarehouseStockSnapshot.objects.create(
            agency=self.agency,
            source_context_type="receiving",
            source_context_id="PR-REPACK",
            sku_code=sku_code,
            name=sku_code,
            size="0",
            barcode=barcode,
            goods_type="gv",
            qty=qty,
            available_qty=0,
            location=self.location,
            zone_code="OTG",
            zone_kind=self.location.zone_kind,
            warehouse_state_code=WarehouseStateCode.IN_OTG.value,
            last_event=event,
        )

    def test_repack_replaces_source_containers_atomically_and_preserves_quantity(self):
        source_items = _shipping_repack_items(self.order)
        items = {item["barcode"].casefold(): item for item in source_items}
        cargoes = [
            {
                "gm_barcode": "GM-1",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-1",
                "items": [
                    {"barcode": "OZN-A", "quantity": 4},
                    {"barcode": "ozn-b", "quantity": 1},
                ],
            },
            {
                "gm_barcode": "GM-2",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-2",
                "items": [{"barcode": "OZN-A", "quantity": 8}],
            },
            {
                "gm_barcode": "GM-3",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-3",
                "items": [{"barcode": "OZN-B", "quantity": 4}],
            },
        ]
        targets = [
            {
                "gm_barcode": "GM-1",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-1",
                "existing_box_code": "",
                "items": [
                    {**items["ozn-a"], "qty": 4},
                    {**items["ozn-b"], "qty": 1},
                ],
            },
            {
                "gm_barcode": "GM-2",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-2",
                "existing_box_code": "",
                "items": [{**items["ozn-a"], "qty": 8}],
            },
            {
                "gm_barcode": "GM-3",
                "supply_id": "SUPPLY-1",
                "bundle_id": "BUNDLE-3",
                "existing_box_code": "",
                "items": [{**items["ozn-b"], "qty": 4}],
            },
        ]
        boxes = [
            {
                "code": f"TARGET-{index}",
                "gm_barcode": target["gm_barcode"],
                "sealed": True,
                "items": target["items"],
            }
            for index, target in enumerate(targets, start=1)
        ]

        with patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value={"gm_cargoes": cargoes},
        ):
            blocked = save_shipping_loose_packing(
                self.order,
                boxes_state=boxes,
                pallets_state=[],
                loose_items=source_items,
                gm_targets=targets,
                repack_existing_boxes=True,
                repack_source_box_codes=["SOURCE-A"],
                user=self.user,
            )
        self.assertFalse(blocked["saved"])
        self.assertFalse(
            WarehouseContainer.objects.filter(
                agency=self.agency,
                container_code__startswith="TARGET-",
            ).exists()
        )

        with patch(
            "shipping.services._ozon_api_summary_for_form",
            return_value={"gm_cargoes": cargoes},
        ):
            result = save_shipping_loose_packing(
                self.order,
                boxes_state=boxes,
                pallets_state=[],
                loose_items=source_items,
                gm_targets=targets,
                repack_existing_boxes=True,
                repack_source_box_codes=["SOURCE-A", "SOURCE-B"],
                user=self.user,
            )

        self.assertTrue(result["saved"], result["errors"])
        snapshots = WarehouseStockSnapshot.objects.filter(
            agency=self.agency,
            is_archived=False,
            qty__gt=0,
        )
        self.assertEqual(sum(snapshot.qty for snapshot in snapshots), 17)
        self.assertEqual(
            set(snapshots.values_list("container_code", flat=True)),
            {"TARGET-1", "TARGET-2", "TARGET-3"},
        )
        self.assertFalse(
            WarehouseContainer.objects.filter(
                agency=self.agency,
                container_code__in=["SOURCE-A", "SOURCE-B"],
                status=WarehouseContainer.STATUS_ACTIVE,
            ).exists()
        )
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)

class ShippingOzonWholeAndPiecePickTests(SimpleTestCase):
    @staticmethod
    def _stock_row(*, box_qty=20, available_boxes=3, box_codes=None, is_mixed_box=False):
        return {
            "key": "stock-ozon-1",
            "sku_code": "SKU-OZON-1",
            "name": "Товар Ozon",
            "size": "",
            "barcode": "460000000001",
            "goods_type": "gv",
            "box_qty": box_qty,
            "available_boxes": available_boxes,
            "available_qty": box_qty * available_boxes,
            "box_codes": box_codes or [f"BOX-OZON-{index}" for index in range(1, available_boxes + 1)],
            "is_mixed_box": is_mixed_box,
            "mixed_group": "mix:ozon" if is_mixed_box else "",
        }

    @staticmethod
    def _bundle(quantity):
        return [
            {
                "offer_id": "SKU-OZON-1",
                "barcode": "460000000001",
                "name": "Товар Ozon",
                "quantity": quantity,
            }
        ]

    def test_exact_quantity_uses_only_whole_boxes(self):
        result = apply_ozon_bundle_to_stock(self._bundle(40), [self._stock_row()])

        self.assertEqual(result["selected_boxes"], {"stock-ozon-1": "2"})
        self.assertEqual(result["partial_box_splits"], [])
        self.assertFalse(result["requires_piece_pick_confirmation"])
        self.assertEqual(result["composition"][0]["status"], "matched")
        self.assertEqual(result["composition"][0]["applied_qty"], 40)

    def test_non_multiple_quantity_stays_in_one_request_with_piece_pick(self):
        shared_selected = {}

        result = apply_ozon_bundle_to_stock(
            self._bundle(45),
            [self._stock_row()],
            already_selected=shared_selected,
        )

        self.assertEqual(result["selected_boxes"], {"stock-ozon-1": "2"})
        self.assertEqual(result["piece_pick_total_qty"], 5)
        self.assertEqual(result["piece_pick_positions"], 1)
        self.assertTrue(result["requires_piece_pick_confirmation"])
        self.assertEqual(result["partial_box_splits"][0]["items"], [{"key": "stock-ozon-1", "qty": 5}])
        self.assertEqual(result["composition"][0]["whole_box_qty"], 40)
        self.assertEqual(result["composition"][0]["piece_pick_qty"], 5)
        self.assertEqual(result["composition"][0]["applied_qty"], 45)
        self.assertEqual(result["composition"][0]["applied_boxes"], 3)
        self.assertEqual(result["composition"][0]["status"], "matched_with_piece_pick")
        self.assertTrue(result["composition"][0]["matched"])
        self.assertEqual(shared_selected, {"stock-ozon-1": 3})

    def test_quantity_below_box_multiplicity_uses_one_piece_pick_box(self):
        result = apply_ozon_bundle_to_stock(self._bundle(5), [self._stock_row(available_boxes=1)])

        self.assertEqual(result["selected_boxes"], {})
        self.assertEqual(result["piece_pick_total_qty"], 5)
        self.assertEqual(result["partial_box_splits"][0]["boxes"], 1)
        self.assertEqual(result["composition"][0]["status"], "matched_with_piece_pick")

    def test_piece_pick_is_not_claimed_without_a_known_free_physical_box(self):
        row = self._stock_row(available_boxes=2, box_codes=["BOX-OZON-1"])

        result = apply_ozon_bundle_to_stock(self._bundle(25), [row])

        self.assertEqual(result["selected_boxes"], {"stock-ozon-1": "1"})
        self.assertEqual(result["partial_box_splits"], [])
        self.assertFalse(result["requires_piece_pick_confirmation"])
        self.assertEqual(result["composition"][0]["status"], "needs_split")
        self.assertFalse(result["composition"][0]["matched"])

    def test_mixed_box_is_left_to_existing_manual_split_flow(self):
        result = apply_ozon_bundle_to_stock(
            self._bundle(5),
            [self._stock_row(available_boxes=1, is_mixed_box=True)],
        )

        self.assertEqual(result["partial_box_splits"], [])
        self.assertEqual(result["composition"][0]["status"], "needs_split")

    @staticmethod
    def _piece_pick_post(**extra):
        payload = {
            "ozon_supply_payload_json": json.dumps(
                {"order_id": 123, "piece_pick_total_qty": 5}
            ),
            "stock_partial_box_splits_json": json.dumps(
                [
                    {
                        "row_key": "stock-ozon-1",
                        "boxes": 1,
                        "items": [{"key": "stock-ozon-1", "qty": 5}],
                    }
                ]
            ),
        }
        payload.update(extra)
        return RequestFactory().post("/shipping/new/", data=payload)

    def test_piece_pick_requires_server_side_tariff_confirmation(self):
        request = self._piece_pick_post()

        error = _ozon_piece_pick_tariff_confirmation_error(request)

        self.assertIn("5 шт.", error)
        self.assertIn("Подтвердите отдельную тарификацию", error)

    def test_piece_pick_accepts_explicit_tariff_confirmation(self):
        request = self._piece_pick_post(piece_pick_tariff_confirmed="1")

        self.assertEqual(_ozon_piece_pick_tariff_confirmation_error(request), "")

    def test_piece_pick_rejects_composition_mismatch_even_when_confirmed(self):
        request = self._piece_pick_post(
            ozon_supply_payload_json=json.dumps(
                {"order_id": 123, "piece_pick_total_qty": 4}
            ),
            piece_pick_tariff_confirmed="1",
        )

        self.assertIn(
            "не совпадает с составом заявки Ozon",
            _ozon_piece_pick_tariff_confirmation_error(request),
        )

    def test_piece_pick_confirmation_is_not_required_for_draft(self):
        request = self._piece_pick_post(action="save_draft")

        self.assertEqual(_ozon_piece_pick_tariff_confirmation_error(request), "")

    def test_non_ozon_api_partial_split_keeps_existing_flow(self):
        request = self._piece_pick_post(ozon_supply_payload_json="")

        self.assertEqual(_ozon_piece_pick_tariff_confirmation_error(request), "")

    def test_shortage_diagnostics_show_article_barcode_and_blocking_request(self):
        payload = {
            "composition": [
                {
                    "offer_id": "нож040",
                    "barcode": "2050659560200",
                    "name": "Нож",
                    "quantity": 40,
                    "applied_qty": 0,
                    "status": "not_in_stock",
                }
            ]
        }

        _attach_ozon_stock_shortages(
            payload,
            reserve_rows=[
                {
                    "sku_code": "нож040",
                    "barcode": "2050659560200",
                    "context_id": "399",
                    "qty_reserved": 80,
                    "qty_satisfied": 10,
                }
            ],
        )

        self.assertEqual(
            payload["stock_shortages"],
            [
                {
                    "sku_code": "нож040",
                    "barcode": "2050659560200",
                    "name": "Нож",
                    "required_qty": 40,
                    "selected_qty": 0,
                    "shortage_qty": 40,
                    "reserved_qty": 70,
                    "blocking_reserves": [{"order": "399_OTG", "qty": 70}],
                    "blocking_orders": ["399_OTG"],
                }
            ],
        )

    def test_shortage_diagnostics_ignore_fully_satisfied_reserve(self):
        payload = {
            "composition": [
                {
                    "offer_id": "SKU-OZON-1",
                    "barcode": "460000000001",
                    "quantity": 20,
                    "applied_qty": 5,
                    "remaining_qty": 15,
                }
            ]
        }

        _attach_ozon_stock_shortages(
            payload,
            reserve_rows=[
                {
                    "sku_code": "SKU-OZON-1",
                    "barcode": "460000000001",
                    "context_id": "400",
                    "qty_reserved": 20,
                    "qty_satisfied": 20,
                }
            ],
        )

        self.assertEqual(payload["stock_shortages"][0]["shortage_qty"], 15)
        self.assertEqual(payload["stock_shortages"][0]["reserved_qty"], 0)
        self.assertEqual(payload["stock_shortages"][0]["blocking_reserves"], [])

    def test_submit_shortage_names_article_barcode_and_blocking_requests(self):
        item = SimpleNamespace(
            sku_code="N010",
            barcode="2046280833899",
        )

        error = _format_shipping_reserve_shortage(
            item,
            needed=15,
            available=0,
            blockers=[("OTG-000447", 10), ("OTG-000449", 5)],
        )

        self.assertIn("артикул N010", error)
        self.assertIn("ШК 2046280833899", error)
        self.assertIn("не хватает 15 шт.", error)
        self.assertIn("OTG-000447 — 10 шт.", error)
        self.assertIn("OTG-000449 — 5 шт.", error)

    def test_submit_shortage_explains_when_no_active_reserve_exists(self):
        item = SimpleNamespace(sku_code="N011", barcode="")

        error = _format_shipping_reserve_shortage(
            item,
            needed=20,
            available=12,
        )

        self.assertIn("не хватает 8 шт.", error)
        self.assertIn("В активных заявках этот товар не найден", error)

    @patch(
        "shipping.ozon_supplies.get_ozon_supply_for_agency",
        return_value={
            "ok": True,
            "order_id": 1,
            "stock_match": {"status": "partial"},
            "selected_boxes": {},
            "form": {},
        },
    )
    def test_lower_batch_service_accepts_ten_supplies(self, supply_mock):
        result = get_ozon_supplies_batch_for_agency(
            None,
            list(range(1, 11)),
            stock_rows=[],
        )

        self.assertTrue(result["ok"])
        self.assertEqual(result["selected_count"], 10)
        self.assertEqual(supply_mock.call_count, 10)

    def test_lower_batch_service_rejects_eleven_supplies(self):
        result = get_ozon_supplies_batch_for_agency(
            None,
            list(range(1, 12)),
            stock_rows=[],
        )

        self.assertFalse(result["ok"])
        self.assertEqual(
            result["error"],
            "За один раз можно выбрать не более 10 поставок Ozon.",
        )

    def test_batch_keeps_piece_pick_box_claim_between_selected_supplies(self):
        selected_states = []

        def fake_supply(_agency, order_id, **kwargs):
            shared = kwargs["already_selected"]
            selected_states.append(dict(shared))
            shared[f"stock-{order_id}"] = 1
            return {
                "ok": True,
                "order_id": order_id,
                "order_number": f"OZ-{order_id}",
                "stock_match": {"status": "full"},
                "selected_boxes": {},
                "partial_box_splits": [
                    {
                        "row_key": f"stock-{order_id}",
                        "boxes": 1,
                        "items": [{"key": f"stock-{order_id}", "qty": 1}],
                    }
                ],
                "piece_pick_total_qty": 1,
                "requires_piece_pick_confirmation": True,
                "form": {"slot_date": "2026-08-08", "destination_warehouse": "Ozon"},
            }

        with patch("shipping.ozon_supplies.get_ozon_supply_for_agency", side_effect=fake_supply):
            result = get_ozon_supplies_batch_for_agency(None, [101, 202], stock_rows=[])

        self.assertTrue(result["ok"])
        self.assertEqual(selected_states[0], {})
        self.assertEqual(selected_states[1], {"stock-101": 1})
        self.assertEqual(len(result["orders"]), 2)


class ShippingOzonReserveBoxReselectionTests(SimpleTestCase):
    item_key = ("barcode", "460000000001")

    def test_allocator_prefers_fully_free_box_before_partial_residual(self):
        result = allocate_reserve_box_codes(
            [
                {
                    "item_id": 1,
                    "item_key": self.item_key,
                    "qty": 20,
                    "preferred_box_codes": ["BOX-PARTIAL"],
                }
            ],
            [
                {
                    "item_key": self.item_key,
                    "box_code": "BOX-PARTIAL",
                    "qty": 50,
                    "available_qty": 35,
                },
                {
                    "item_key": self.item_key,
                    "box_code": "BOX-FREE",
                    "qty": 10,
                    "available_qty": 10,
                },
            ],
        )

        self.assertTrue(result["ok"])
        self.assertEqual(
            result["box_codes_by_item"][1],
            ["BOX-FREE", "BOX-PARTIAL"],
        )

    def test_allocator_never_uses_reserved_part_of_partial_box(self):
        stock_rows = [
            {
                "item_key": self.item_key,
                "box_code": "BOX-PARTIAL",
                "qty": 50,
                "available_qty": 35,
            }
        ]

        enough = allocate_reserve_box_codes(
            [{"item_id": 1, "item_key": self.item_key, "qty": 35}],
            stock_rows,
        )
        shortage = allocate_reserve_box_codes(
            [{"item_id": 1, "item_key": self.item_key, "qty": 36}],
            stock_rows,
        )

        self.assertTrue(enough["ok"])
        self.assertEqual(enough["box_codes_by_item"][1], ["BOX-PARTIAL"])
        self.assertFalse(shortage["ok"])
        self.assertEqual(shortage["shortages"][0]["available_qty"], 35)
        self.assertEqual(shortage["shortages"][0]["shortage_qty"], 1)

    def test_reselected_codes_preserve_regular_comment(self):
        updated = _comment_with_reselected_box_codes(
            "Коробов: 1; кратность: 10; короба: BOX-STALE; комментарий клиента",
            ["BOX-FREE"],
        )

        self.assertIn("короба: BOX-FREE", updated)
        self.assertNotIn("BOX-STALE", updated)
        self.assertIn("комментарий клиента", updated)

    def test_reselected_codes_update_partial_split_metadata(self):
        split_meta = {
            "source_box_codes": ["BOX-STALE"],
            "items": [{"key": "stock-row-1", "qty": 5}],
        }
        comment = (
            "Исходные короба: BOX-STALE; Разбить короб; "
            f"{encode_partial_box_split(split_meta)}"
        )

        updated = _comment_with_reselected_box_codes(comment, ["BOX-PARTIAL"])
        updated_meta = extract_partial_box_split(updated)

        self.assertIn("Исходные короба: BOX-PARTIAL", updated)
        self.assertNotIn("BOX-STALE", updated)
        self.assertEqual(updated_meta["source_box_codes"], ["BOX-PARTIAL"])
        self.assertEqual(updated_meta["items"], split_meta["items"])

    @patch("shipping.services.StockAvailabilityService.stock_rows_with_availability")
    def test_service_reselects_stale_hint_from_reservation_aware_rows(self, rows_mock):
        rows_mock.return_value = [
            {
                "sku": "N010",
                "size": "",
                "barcode": "460000000001",
                "goods_type": "gv",
                "box_code": "BOX-STALE",
                "qty": 50,
                "available_qty": 0,
            },
            {
                "sku": "N010",
                "size": "",
                "barcode": "460000000001",
                "goods_type": "gv",
                "box_code": "BOX-FREE",
                "qty": 10,
                "available_qty": 10,
            },
        ]

        class Item(SimpleNamespace):
            def save(self, *, update_fields):
                self.saved_fields = update_fields

        agency = object()
        order = SimpleNamespace(agency=agency, number="OTG-000447")
        item = Item(
            pk=1,
            sku_code="N010",
            size="",
            barcode="460000000001",
            goods_type="gv",
            qty_requested=10,
            comment="короба: BOX-STALE; клиентский комментарий",
        )

        changed = _reselect_selected_boxes_locked(order, [item])

        self.assertTrue(changed)
        self.assertIn("короба: BOX-FREE", item.comment)
        self.assertIn("клиентский комментарий", item.comment)
        self.assertEqual(item.saved_fields, ["comment", "updated_at"])
        rows_mock.assert_called_once_with(
            agency=agency,
            require_box=True,
            exclude_shipping_order_id="OTG-000447",
        )


class ShippingOzonApiBatchAuditTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="ozon_batch_snapshot", password="pwd")
        self.agency = Agency.objects.create(agn_name="Клиент Ozon snapshot")
        Market.objects.create(id=9911, name="Ozon")

    def test_legacy_batch_rejects_more_than_ten_supplies(self):
        with self.assertRaisesMessage(
            ValidationError,
            "За один раз можно отправить не более 10 поставок Ozon.",
        ):
            create_ozon_api_shipping_orders_batch(
                agency=self.agency,
                user=self.user,
                payloads=[{} for _index in range(11)],
                stock_rows=[],
                ship_date=timezone.localdate(),
            )

    @patch("shipping.services.apply_client_shipping_reserve_on_submit")
    @patch("shipping.ozon_template.items_from_selected_boxes")
    def test_batch_audit_keeps_full_gm_snapshot_but_reserve_does_not_receive_it(
        self,
        items_from_boxes_mock,
        reserve_mock,
    ):
        items_from_boxes_mock.return_value = (
            [
                {
                    "sku_code": "SKU-1",
                    "name": "Товар",
                    "size": "",
                    "barcode": "460000000001",
                    "goods_type": "gv",
                    "qty_requested": 20,
                    "comment": "Коробов: 1; кратность: 20",
                }
            ],
            8,
            [],
        )
        ship_date = timezone.localdate() + timedelta(days=1)
        payload = {
            "order_id": "ozon-order-1",
            "order_number": "200006189 7823",
            "fetched_at": "2026-08-06T10:00:00+03:00",
            "stock_match": {"status": "full"},
            "selected_boxes": {"stock-1": 1},
            "ozon_boxes_total": 6,
            "gm_barcodes": ["GM-001", "GM-002", "GM-003", "GM-004", "GM-005", "GM-006"],
            "gm_cargoes": [
                {
                    "gm_barcode": f"GM-{index:03d}",
                    "cargo_id": f"cargo-{index}",
                    "items": [
                        {
                            "offer_id": "SKU-1",
                            "barcode": "460000000001",
                            "name": "Товар",
                            "quantity": quantity,
                            "ozon_sku": "123456",
                        }
                    ],
                }
                for index, quantity in enumerate((4, 4, 3, 3, 3, 3), start=1)
            ],
            "composition": [{"barcode": "460000000001", "quantity": 20}],
            "destination_warehouses": ["СТАРАЯ_КУПАВНА_28"],
            "form": {
                "slot_date": ship_date.isoformat(),
                "slot_time": "18:00",
                "destination_warehouse": "СТАРАЯ_КУПАВНА_28",
                "shipping_barcode": "2000061897823",
            },
        }

        orders = create_ozon_api_shipping_orders_batch(
            agency=self.agency,
            user=self.user,
            payloads=[payload],
            stock_rows=[],
            ship_date=ship_date,
        )

        audit = OrderAuditEntry.objects.filter(
            order_type="shipping",
            order_id=orders[0].number,
            action="create",
        ).latest("id")
        self.assertEqual(orders[0].expected_boxes, 8)
        self.assertTrue(audit.payload["ozon_snapshot_complete"])
        self.assertEqual(audit.payload["ozon_boxes_total"], 6)
        self.assertEqual(audit.payload["ozon_supply_id_normalized"], "2000061897823")
        self.assertEqual(audit.payload["gm_cargoes"][0]["gm_barcode"], "GM-001")
        self.assertEqual(audit.payload["gm_cargoes"][0]["items"][0]["barcode"], "460000000001")
        reserve_meta = reserve_mock.call_args.kwargs["ozon_api_meta"]
        self.assertEqual(reserve_meta["gm_cargoes"], [])

    def test_combined_comment_maps_each_ozon_supply_to_the_same_fullbox_request(self):
        order = ShippingOrder.objects.create(
            number="OTG-900001",
            agency=self.agency,
            created_by=self.user,
            status=ShippingOrder.STATUS_SUBMITTED,
            supply_number="OZ-1 + ещё 1",
            comment=(
                "Данные Ozon API:\n"
                "Поставок Ozon: 2\n"
                "Номера поставок Ozon: OZ-1, OZ-2"
            ),
        )

        found = lookup_existing_ozon_supply_shipments(self.agency, ["OZ-1", "OZ-2"])

        self.assertEqual(found["OZ-1"]["id"], order.id)
        self.assertEqual(found["OZ-2"]["id"], order.id)


class ShippingDetailIntegrityWarningTests(SimpleTestCase):
    def test_resolved_supplement_hides_only_quantity_mismatch_warning(self):
        warnings = [
            "Расхождение отгрузки по ШК и количеству: ШК 2044440328704: в заявке 0 шт, факт OTG 30 шт.",
            "Неполная паллетизация: в OTG 6 коробов, в акте 5, не разложено 1.",
        ]

        visible = _visible_shipping_integrity_warnings(
            warnings,
            {
                "status": "resolved",
                "decision": "supplement",
                "supplement_fact_complete": True,
            },
        )

        self.assertEqual(
            visible,
            ["Неполная паллетизация: в OTG 6 коробов, в акте 5, не разложено 1."],
        )

    def test_resolved_supplement_hides_quantity_warning_by_status(self):
        warnings = [
            "Расхождение отгрузки по ШК и количеству: ШК 2044440328704: в заявке 0 шт, факт OTG 30 шт.",
        ]

        visible = _visible_shipping_integrity_warnings(
            warnings,
            {
                "status": "resolved",
                "decision": "supplement",
                "supplement_fact_complete": False,
            },
        )

        self.assertEqual(visible, [])

    def test_unconfirmed_supplement_keeps_quantity_mismatch_warning(self):
        warnings = [
            "Расхождение отгрузки по ШК и количеству: ШК 2044440328704: в заявке 0 шт, факт OTG 30 шт.",
        ]

        visible = _visible_shipping_integrity_warnings(
            warnings,
            {
                "status": "pickup_required",
                "decision": "supplement",
                "supplement_fact_complete": False,
            },
        )

        self.assertEqual(visible, warnings)


class ShippingFlowTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="shipping_manager", password="pwd")
        Employee.objects.create(
            full_name="Менеджер отгрузки",
            role="manager",
            user=self.user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Тест Клиент")
        self.order = ShippingOrder.objects.create(
            number="SO-000001",
            agency=self.agency,
            created_by=self.user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        self.item = ShippingOrderItem.objects.create(
            order=self.order,
            sku_code="SKU-001",
            name="Товар 1",
            size="42",
            barcode="200000000001",
            goods_type="Готовый",
            qty_requested=20,
        )
        self.shipping_service, _created = BillingService.objects.get_or_create(
            code="shipping_phase3_flow_test",
            defaults={"name": "Паллетизация отгрузки", "unit": "шт", "is_active": True},
        )
        self._create_snapshot_box()

    def _create_shipping_service_fact(self, order: ShippingOrder | None = None):
        order = order or self.order
        return WarehouseServiceFact.objects.get_or_create(
            client=order.agency,
            order_type=WarehouseServiceFact.ORDER_SHIPPING,
            order_id=order.number,
            service=self.shipping_service,
            defaults={
                "service_name_snapshot": self.shipping_service.name,
                "quantity": 1,
                "unit": "шт",
                "status": WarehouseServiceFact.STATUS_SENT_TO_BILLING,
                "reported_by": self.user,
            },
        )[0]

    def _create_snapshot_box(
        self,
        *,
        pallet_code: str = "PL-1",
        box_code: str = "BX-1",
        order_id: str = "R-SNAP",
        sku: str = "SKU-001",
        name: str = "Товар 1",
        size: str = "42",
        barcode: str = "200000000001",
        goods_type: str = "Готовый",
        qty: int = 100,
        zone: str = "OS",
        row: int = 1,
        section: int = 1,
        tier: int = 1,
        cell: int = 1,
        location_code: str = "OS-1-1-1-1",
    ) -> WarehouseStockSnapshot:
        location = WarehouseWritePathService.ensure_location(
            warehouse_code="MSK",
            zone_code=zone,
            row_no=row,
            section_no=section,
            tier_no=tier,
            cell_no=cell,
        )
        if location.location_code != location_code:
            location.location_code = location_code
            location.save(update_fields=["location_code", "updated_at"])
        pallet, _ = WarehouseContainer.objects.get_or_create(
            agency=self.agency,
            container_code=pallet_code,
            defaults={
                "container_type": WarehouseContainer.TYPE_PALLET,
                "current_location": location,
            },
        )
        if pallet.current_location_id != location.id:
            pallet.current_location = location
            pallet.save(update_fields=["current_location", "updated_at"])
        box, _ = WarehouseContainer.objects.get_or_create(
            agency=self.agency,
            container_code=box_code,
            defaults={
                "container_type": WarehouseContainer.TYPE_BOX,
                "parent_container": pallet,
                "current_location": location,
            },
        )
        changed_fields = []
        if box.parent_container_id != pallet.id:
            box.parent_container = pallet
            changed_fields.append("parent_container")
        if box.current_location_id != location.id:
            box.current_location = location
            changed_fields.append("current_location")
        if changed_fields:
            box.save(update_fields=[*changed_fields, "updated_at"])
        return WarehouseStockSnapshot.objects.create(
            agency=self.agency,
            source_context_type="receiving",
            source_context_id=order_id,
            sku_code=sku,
            name=name,
            size=size,
            barcode=barcode,
            goods_type=goods_type,
            qty=qty,
            available_qty=qty,
            container=box,
            container_code=box.container_code,
            parent_container=pallet,
            location=location,
            zone_code=location.zone_code,
            zone_kind=location.zone_kind,
            warehouse_state_code="stored",
        )

    def test_ozon_reserved_fullbox_boxes_exposes_exact_whole_box(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._create_snapshot_box(box_code="BX-OZON-GM-1", qty=20)
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            created_by=self.user,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
        )

        boxes = _ozon_reserved_fullbox_boxes(self.order)

        self.assertEqual(len(boxes), 1)
        self.assertEqual(boxes[0]["box_code"], "BX-OZON-GM-1")
        self.assertEqual(boxes[0]["qty"], 20)
        self.assertEqual(boxes[0]["items"][0]["barcode"], self.item.barcode)
        self.assertEqual(boxes[0]["items"][0]["qty"], 20)

    def test_ozon_reserved_fullbox_boxes_skips_partial_box(self):
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            created_by=self.user,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
        )

        self.assertEqual(_ozon_reserved_fullbox_boxes(self.order), [])

    def test_pickup_planned_ship_date_uses_submission_date_when_eta_is_empty(self):
        submitted_at = timezone.make_aware(datetime(2026, 7, 23, 15, 20))
        self.order.delivery_type = ShippingOrder.DELIVERY_PICKUP
        self.order.eta_at = None
        self.order.planned_ship_date = None

        _apply_planned_ship_date(self.order, submitted_at=submitted_at)

        self.assertEqual(self.order.planned_ship_date, submitted_at.date())

    def test_courier_keeps_explicit_eta_as_planned_ship_date(self):
        submitted_at = timezone.make_aware(datetime(2026, 7, 23, 15, 20))
        eta_at = timezone.make_aware(datetime(2026, 7, 25, 10, 0))
        self.order.delivery_type = ShippingOrder.DELIVERY_COURIER
        self.order.eta_at = eta_at
        self.order.planned_ship_date = None

        _apply_planned_ship_date(self.order, submitted_at=submitted_at)

        self.assertEqual(self.order.planned_ship_date, eta_at.date())

    def _load_snapshot_order_to_vehicle(self, *, trip_number: str = "TRIP-1", qty: int = 20) -> WarehouseStockSnapshot:
        snapshot = self._load_snapshot_order_ready_for_loading(qty=qty)
        trip = LogisticsTrip.objects.create(
            number=trip_number,
            status=LogisticsTrip.STATUS_LOADING,
            vehicle_type=LogisticsTrip.VEHICLE_FULFILLMENT,
            created_by=self.user,
        )
        LogisticsTripOrder.objects.create(trip=trip, shipping_order=self.order)
        WarehouseWritePathService.assign_to_trip(
            agency=self.agency,
            order_id=self.order.number,
            trip_id=trip.number,
            assigned_by=self.user,
        )
        loading = WarehouseWritePathService.start_loading(
            agency=self.agency,
            order_id=self.order.number,
            trip_id=trip.number,
            started_by=self.user,
        )
        WarehouseWritePathService.complete_loading(operation=loading, performed_by=self.user)
        snapshot.refresh_from_db()
        return snapshot

    def _load_snapshot_order_ready_for_loading(self, *, qty: int = 20) -> WarehouseStockSnapshot:
        snapshot = self._create_snapshot_box(qty=qty)
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": qty,
                }
            ],
            created_by=self.user,
        )
        move_op = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=self.order.number,
            requested_by=self.user,
        )
        WarehouseWritePathService.start_move_to_otg(operation=move_op, performed_by=self.user)
        WarehouseWritePathService.complete_move_to_otg(operation=move_op, performed_by=self.user)
        self._create_shipping_service_fact()
        packing_state = _shipping_packing_initial_state(self.order)
        row = packing_state["initial_boxes"][0]
        result = save_shipping_packing(
            self.order,
            boxes_state=[
                {
                    "row_key": row["row_key"],
                    "code": row["code"],
                    "qty": row["qty"],
                    "barcode_preview": row["barcode_preview"],
                    "pallet_code": "PAL-SHIP-FLOW",
                }
            ],
            pallets_state=[{"code": "PAL-SHIP-FLOW", "label": "1"}],
            delivered_boxes=packing_state["delivered_boxes"],
            initial_boxes=packing_state["initial_boxes"],
            user=self.user,
        )
        self.assertTrue(result["saved"], result.get("errors"))
        snapshot.refresh_from_db()
        return snapshot

    def _assign_ready_snapshot_to_trip(
        self,
        *,
        trip_number: str = "TRIP-1",
        qty: int = 20,
        start_loading: bool = False,
    ) -> tuple[WarehouseStockSnapshot, LogisticsTrip]:
        snapshot = self._load_snapshot_order_ready_for_loading(qty=qty)
        trip = LogisticsTrip.objects.create(
            number=trip_number,
            status=LogisticsTrip.STATUS_LOADING if start_loading else LogisticsTrip.STATUS_PLANNED,
            vehicle_type=LogisticsTrip.VEHICLE_FULFILLMENT,
            created_by=self.user,
        )
        LogisticsTripOrder.objects.create(trip=trip, shipping_order=self.order)
        WarehouseWritePathService.assign_to_trip(
            agency=self.agency,
            order_id=self.order.number,
            trip_id=trip.number,
            assigned_by=self.user,
        )
        if start_loading:
            WarehouseWritePathService.start_loading(
                agency=self.agency,
                order_id=self.order.number,
                trip_id=trip.number,
                started_by=self.user,
            )
        snapshot.refresh_from_db()
        return snapshot, trip

    def test_reserve_order_creates_reserve_rows(self):
        reserve_order(self.order, self.user)
        self.item.refresh_from_db()
        self.order.refresh_from_db()
        snapshot = WarehouseStockSnapshot.objects.get(agency=self.agency, sku_code="SKU-001", size="42")
        self.assertEqual(self.order.status, ShippingOrder.STATUS_RESERVED)
        self.assertEqual(self.item.qty_reserved, 20)
        self.assertEqual(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            ).count(),
            1,
        )
        self.assertEqual(snapshot.processing_reserved_qty, 0)
        self.assertEqual(snapshot.shipping_reserved_qty, 0)
        self.assertEqual(snapshot.available_qty, 100)
        reserve = WarehouseReserve.objects.get(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_id=self.order.number,
        )
        self.assertTrue(reserve.events.filter(payload__reserve_scope="pool").exists())

    def test_reserve_order_blocks_when_active_pool_reserves_exceed_locked_available(self):
        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_type="shipping",
            context_id="SO-OTHER",
            sku_code="SKU-001",
            size="42",
            barcode="200000000001",
            goods_type="Готовый",
            qty_reserved=90,
            status=WarehouseReserve.STATUS_ACTIVE,
        )

        with self.assertRaises(ValidationError):
            reserve_order(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertFalse(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            ).exists()
        )

    def test_reserve_order_can_keep_submitted_status_for_auto_reserve(self):
        reserve_order(
            self.order,
            self.user,
            target_status=ShippingOrder.STATUS_SUBMITTED,
            log_description="Авто-резерв при отправке клиента",
        )
        self.item.refresh_from_db()
        self.order.refresh_from_db()
        snapshot = WarehouseStockSnapshot.objects.get(agency=self.agency, sku_code="SKU-001", size="42")
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(self.item.qty_reserved, 20)
        self.assertEqual(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            ).count(),
            1,
        )
        self.assertEqual(snapshot.shipping_reserved_qty, 0)
        self.assertEqual(snapshot.available_qty, 100)

    def test_reserve_order_reuses_matching_pool_reserve_when_original_box_is_no_longer_available(self):
        self.item.comment = "Коробов: 1; кратность: 20; короба: BX-1"
        self.item.save(update_fields=["comment", "updated_at"])
        self._create_snapshot_box(box_code="BX-2", qty=20, order_id="R-SNAP-2")
        reserve_order(
            self.order,
            self.user,
            target_status=ShippingOrder.STATUS_SUBMITTED,
            log_description="Авто-резерв при отправке клиента",
        )
        own_reserve = WarehouseReserve.objects.get(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_id=self.order.number,
        )
        reserve_event = own_reserve.events.get(payload__reserve_scope="pool")
        reserve_event.payload = {**reserve_event.payload, "box_codes": []}
        reserve_event.save(update_fields=["payload"])
        self.assertEqual(
            reserve_event.payload.get("box_codes"),
            [],
        )
        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_type="shipping",
            context_id="SO-OTHER",
            sku_code=self.item.sku_code,
            size=self.item.size,
            barcode=self.item.barcode,
            goods_type=self.item.goods_type,
            qty_reserved=100,
            status=WarehouseReserve.STATUS_ACTIVE,
        )

        reserve_order(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_RESERVED)
        self.assertEqual(self.item.qty_reserved, 20)
        self.assertEqual(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
                status=WarehouseReserve.STATUS_ACTIVE,
            ).count(),
            1,
        )

    def test_release_order_reserves_restores_available_qty_without_global_refresh(self):
        reserve_order(self.order, self.user)

        release_order_reserves(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot = WarehouseStockSnapshot.objects.get(agency=self.agency, sku_code="SKU-001", size="42")
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertFalse(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            )
            .exclude(status__in=[WarehouseReserve.STATUS_RELEASED, WarehouseReserve.STATUS_CANCELED])
            .exists()
        )
        self.assertEqual(snapshot.shipping_reserved_qty, 0)
        self.assertEqual(snapshot.available_qty, 100)

    def test_release_order_reserves_releases_partially_satisfied_pool_reserve(self):
        from sklad.services.stock_availability import StockAvailabilityService

        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_type="shipping",
            context_id=self.order.number,
            sku_code=self.item.sku_code,
            size=self.item.size,
            barcode=self.item.barcode,
            goods_type=self.item.goods_type,
            qty_reserved=270,
            qty_satisfied=180,
            status=WarehouseReserve.STATUS_PARTIALLY_SATISFIED,
        )
        row_before = next(
            row
            for row in StockAvailabilityService.stock_rows_with_availability(agency=self.agency)
            if row["box_code"] == "BX-1"
        )
        self.assertEqual(row_before["available_qty"], 10)

        release_order_reserves(self.order, self.user)

        self.assertFalse(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            )
            .exclude(status__in=[WarehouseReserve.STATUS_RELEASED, WarehouseReserve.STATUS_CANCELED])
            .exists()
        )
        row_after = next(
            row
            for row in StockAvailabilityService.stock_rows_with_availability(agency=self.agency)
            if row["box_code"] == "BX-1"
        )
        self.assertEqual(row_after["available_qty"], 100)

    def test_cancel_order_releases_reserves(self):
        reserve_order(self.order, self.user)
        self.item.refresh_from_db()
        self.assertGreater(int(self.item.qty_reserved or 0), 0)

        loose_task = Task.objects.create(
            title=f"СРОЧНО: упаковать добор {self.order.number}",
            description="Добор доставлен в OTG; требуется упаковка.",
            route=f"/shipping/{self.order.pk}/packing/loose/",
            created_by=self.user,
            status="in_progress",
            priority="urgent",
        )

        cancel_order(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_CANCELED)
        self.assertIsNone(self.order.reserved_at)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertFalse(
            WarehouseReserve.objects.filter(
                agency=self.agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_id=self.order.number,
            )
            .exclude(status__in=[WarehouseReserve.STATUS_RELEASED, WarehouseReserve.STATUS_CANCELED])
            .exists()
        )
        loose_task.refresh_from_db()
        self.assertEqual(loose_task.status, "in_progress")
        self.assertEqual(
            loose_task.title,
            f"Упаковать возврат отмененной отгрузки №{self.order.number}",
        )
        self.assertIn("Повторная проверка ЧЗ не требуется", loose_task.description)

    def test_client_draft_does_not_keep_reserve_helpers(self):
        from .services import (
            apply_client_shipping_reserve_on_submit,
            release_client_shipping_reserve_on_cancel,
        )

        self.order.status = ShippingOrder.STATUS_DRAFT
        self.order.save(update_fields=["status", "updated_at"])
        apply_client_shipping_reserve_on_submit(self.order, self.user)
        self.item.refresh_from_db()
        self.assertEqual(int(self.item.qty_reserved or 0), 0)

        self.order.status = ShippingOrder.STATUS_SUBMITTED
        self.order.save(update_fields=["status", "updated_at"])
        apply_client_shipping_reserve_on_submit(self.order, self.user)
        self.item.refresh_from_db()
        self.assertGreater(int(self.item.qty_reserved or 0), 0)

        release_client_shipping_reserve_on_cancel(self.order, self.user)
        self.item.refresh_from_db()
        self.assertEqual(int(self.item.qty_reserved or 0), 0)

    def test_create_pick_tasks_creates_reachtruck_task(self):
        reserve_order(self.order, self.user)
        move_ids = create_pick_tasks(
            self.order,
            self.user,
            requested_by_name="Тест Менеджер",
            requested_by_role="manager",
        )
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        self.assertTrue(move_ids)
        self.assertEqual(MoveRequest.objects.count(), 1)
        self.assertEqual(MoveTask.objects.count(), 1)

    @override_settings(USE_OTG_REACHTRUCK_PLANNER=True)
    def test_create_pick_tasks_queues_temporarily_busy_source_box_for_retry(self):
        self.order.status = ShippingOrder.STATUS_STOREKEEPER_ACCEPTED
        self.order.save(update_fields=["status", "updated_at"])
        reason = "Нельзя переназначить занятый исходный короб для частичного отбора"

        with patch(
            "otg_reachtruck.services.shipping_queue_priority_wait_reason",
            return_value="",
        ):
            with patch(
                "otg_reachtruck.services.rebind_unavailable_shipping_source_boxes",
                side_effect=ValidationError(reason),
            ):
                with patch(
                    "otg_reachtruck.services.ensure_waiting_otg_shipping_pick_request",
                ) as ensure_waiting:
                    move_ids = create_pick_tasks(self.order, self.user)

        self.assertEqual(move_ids, [])
        ensure_waiting.assert_called_once()
        self.assertEqual(ensure_waiting.call_args.kwargs["waiting_message"], reason)

    @override_settings(USE_OTG_REACHTRUCK_PLANNER=True)
    def test_create_pick_tasks_queues_full_pick_shortage_for_retry(self):
        self.order.status = ShippingOrder.STATUS_STOREKEEPER_ACCEPTED
        self.order.save(update_fields=["status", "updated_at"])
        reason = "Нельзя создать полный подбор: не хватает 3 короб. / 40 шт."

        with patch(
            "otg_reachtruck.services.shipping_queue_priority_wait_reason",
            return_value="",
        ):
            with patch(
                "otg_reachtruck.services.rebind_unavailable_shipping_source_boxes",
                return_value=[],
            ):
                with patch(
                    "shipping.services.shipping_pick_readiness",
                    return_value={"can_pick": False, "reason": reason},
                ):
                    with patch(
                        "otg_reachtruck.services.ensure_waiting_otg_shipping_pick_request",
                    ) as ensure_waiting:
                        move_ids = create_pick_tasks(self.order, self.user)

        self.assertEqual(move_ids, [])
        ensure_waiting.assert_called_once()
        self.assertEqual(ensure_waiting.call_args.kwargs["waiting_message"], reason)

    def test_created_reachtruck_task_does_not_block_pool_pick_readiness(self):
        reserve_order(self.order, self.user)
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            agency=self.agency,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-1",
            legacy_order_id="BLOCK-CREATED",
            status=MoveTask.STATUS_CREATED,
        )

        readiness = shipping_pick_readiness(self.order)

        self.assertTrue(readiness["can_pick"])
        task = MoveTask.objects.first()
        self.assertEqual(task.to_zone, "OTG")
        self.assertEqual(task.request.items.count(), 1)

    def test_create_pick_tasks_groups_multiple_items_from_same_pallet_into_one_task(self):
        ShippingOrderItem.objects.create(
            order=self.order,
            sku_code="SKU-002",
            name="Товар 2",
            size="44",
            barcode="200000000002",
            goods_type="Готовый",
            qty_requested=10,
        )
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-002",
            name="Товар 2",
            size="44",
            barcode="200000000002",
            qty=30,
            box_code="BX-2",
            pallet_code="PL-1",
        )

        reserve_order(self.order, self.user)
        move_ids = create_pick_tasks(
            self.order,
            self.user,
            requested_by_name="Тест Менеджер",
            requested_by_role="manager",
        )

        self.assertEqual(len(move_ids), 1)
        self.assertEqual(MoveRequest.objects.count(), 1)
        self.assertEqual(MoveTask.objects.count(), 1)
        task = MoveTask.objects.first()
        self.assertEqual(task.pallet_code, "PL-1")
        self.assertEqual(task.request.items.count(), 2)
        self.assertEqual(task.qty_planned, 30)
        self.assertEqual(task.payload.get("task_kind_label"), "Частичный отбор с палеты для отгрузки")
        self.assertEqual(
            task.payload.get("requested_barcode_qty"),
            {"200000000001": 20, "200000000002": 10},
        )
        self.assertIn("Частичный отбор для отгрузки", task.payload.get("instruction") or "")
        self.assertIn("верни палету обратно", task.payload.get("instruction") or "")

    def test_create_pick_tasks_uses_full_pallet_when_request_covers_entire_pallet(self):
        self.item.qty_requested = 100
        self.item.save(update_fields=["qty_requested", "updated_at"])
        ShippingOrderItem.objects.create(
            order=self.order,
            sku_code="SKU-002",
            name="Товар 2",
            size="44",
            barcode="200000000002",
            goods_type="Готовый",
            qty_requested=30,
        )
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-002",
            name="Товар 2",
            size="44",
            barcode="200000000002",
            qty=30,
            box_code="BX-2",
            pallet_code="PL-1",
        )

        reserve_order(self.order, self.user)
        move_ids = create_pick_tasks(
            self.order,
            self.user,
            requested_by_name="Тест Менеджер",
            requested_by_role="manager",
        )

        self.assertEqual(len(move_ids), 1)
        task = MoveTask.objects.first()
        self.assertEqual(task.move_mode, MoveTask.MODE_PALLET_FULL)
        self.assertEqual(task.payload.get("task_kind_label"), "Паллета целиком")
        self.assertEqual(task.payload.get("pick_mode"), "full")
        self.assertEqual(task.payload.get("requested_qty"), "")
        self.assertEqual(task.payload.get("requested_barcode_qty"), {})
        self.assertIn("Возьми палету PL-1 целиком", task.payload.get("instruction") or "")
        self.assertNotIn("верни палету обратно", task.payload.get("instruction") or "")

    def test_warehouse_state_resolver_marks_picking_order_as_in_otg_after_delivery(self):
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-RES-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-RES-1",
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "200000000001",
                                "goods_type": "Готовый",
                                "qty": 20,
                            }
                        ],
                    }
                ],
            },
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-RES-1",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
                "receiving_order_id": "R-RES-1",
                "picked_boxes": ["BX-RES-1"],
                "picked_rows": [{"box_code": "BX-RES-1", "qty": 20}],
            },
        )

        result = WarehouseGoodsStateResolver.resolve_for_shipping_order(self.order)

        self.assertEqual(result.code, WarehouseStateCode.IN_OTG)
        self.assertEqual(result.label_for("default"), "Товар доставлен в OTG, ожидает паллетизации")

    def test_shipping_delivered_boxes_reads_boxes_from_all_placement_acts_of_same_receiving_order(self):
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-MULTI-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-MULTI-1",
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "200000000001",
                                "goods_type": "Готовый",
                                "qty": 10,
                            }
                        ],
                    }
                ],
            },
        )
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-MULTI-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-MULTI-2",
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "200000000001",
                                "goods_type": "Готовый",
                                "qty": 10,
                            }
                        ],
                    }
                ],
            },
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-MULTI-1",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
                "receiving_order_id": "R-MULTI-1",
                "picked_boxes": ["BX-MULTI-1", "BX-MULTI-2"],
            },
        )

        rows = _shipping_delivered_boxes(self.order)

        self.assertEqual(
            [(row["box_code"], row["qty"]) for row in rows],
            [("BX-MULTI-1", 10), ("BX-MULTI-2", 10)],
        )

    def test_shipping_delivered_boxes_skips_picked_boxes_without_placement_items(self):
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-MULTI-EMPTY-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-REAL-1",
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "200000000001",
                                "goods_type": "Готовый",
                                "qty": 10,
                            }
                        ],
                    }
                ],
            },
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-MULTI-EMPTY-1",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
                "receiving_order_id": "R-MULTI-EMPTY-1",
                "picked_boxes": ["BX-REAL-1", "BX-EMPTY-2"],
            },
        )

        rows = _shipping_delivered_boxes(self.order)

        self.assertEqual(
            [(row["box_code"], row["qty"]) for row in rows],
            [("BX-REAL-1", 10)],
        )

    def test_shipping_delivered_boxes_prefers_warehouse_snapshots_in_otg(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._create_snapshot_box(box_code="BX-WH-1", qty=20)
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
            created_by=self.user,
        )
        move_op = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=self.order.number,
            requested_by=self.user,
        )
        WarehouseWritePathService.start_move_to_otg(operation=move_op, performed_by=self.user)
        WarehouseWritePathService.complete_move_to_otg(operation=move_op, performed_by=self.user)

        rows = _shipping_delivered_boxes(self.order)

        self.assertEqual(
            rows,
            [
                {
                    "row_key": _shipping_box_row_key({"box_code": "BX-WH-1", "receiving_order_id": "R-SNAP"}, 1),
                    "box_code": "BX-WH-1",
                    "qty": 20,
                    "items": [
                        {
                            "sku_code": "SKU-001",
                            "name": "Товар 1",
                            "size": "42",
                            "barcode": "200000000001",
                            "goods_type": "Готовый",
                            "qty": 20,
                        }
                    ],
                    "barcode_preview": "200000000001 - 20 шт.",
                    "receiving_order_id": "R-SNAP",
                }
            ],
        )

    def test_save_shipping_packing_persists_shipping_pallets_in_warehouse(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._create_snapshot_box(box_code="BX-PACK-1", qty=20)
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
            created_by=self.user,
        )
        move_op = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=self.order.number,
            requested_by=self.user,
        )
        WarehouseWritePathService.start_move_to_otg(operation=move_op, performed_by=self.user)
        WarehouseWritePathService.complete_move_to_otg(operation=move_op, performed_by=self.user)

        packing_state = _shipping_packing_initial_state(self.order)
        row = packing_state["initial_boxes"][0]
        with self.assertRaisesMessage(ValidationError, "Перед завершением"):
            save_shipping_packing(
                self.order,
                boxes_state=[
                    {
                        "row_key": row["row_key"],
                        "code": row["code"],
                        "qty": row["qty"],
                        "barcode_preview": row["barcode_preview"],
                        "pallet_code": "PAL-1",
                    }
                ],
                pallets_state=[{"code": "PAL-1", "label": "1"}],
                delivered_boxes=packing_state["delivered_boxes"],
                initial_boxes=packing_state["initial_boxes"],
                user=self.user,
            )
        self.order.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        self.assertEqual(snapshot.warehouse_state_code, WarehouseStateCode.IN_OTG.value)
        self.assertFalse(
            WarehouseContainer.objects.filter(
                agency=self.agency,
                source_context_type="shipping",
                source_context_id=self.order.number,
            ).exists()
        )

        self._create_shipping_service_fact()
        result = save_shipping_packing(
            self.order,
            boxes_state=[
                {
                    "row_key": row["row_key"],
                    "code": row["code"],
                    "qty": row["qty"],
                    "barcode_preview": row["barcode_preview"],
                    "pallet_code": "PAL-1",
                }
            ],
            pallets_state=[{"code": "PAL-1", "label": "1"}],
            delivered_boxes=packing_state["delivered_boxes"],
            initial_boxes=packing_state["initial_boxes"],
            user=self.user,
        )

        self.assertTrue(result["saved"])
        self.order.refresh_from_db()
        snapshot.refresh_from_db()
        snapshot.container.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PACKED)
        self.assertEqual(snapshot.warehouse_state_code, "ready_for_loading")
        self.assertIsNone(snapshot.active_operation)
        self.assertIsNotNone(snapshot.parent_container)
        self.assertEqual(snapshot.parent_container.container_type, WarehouseContainer.TYPE_MIXED_PALLET)
        self.assertEqual(snapshot.parent_container.source_context_type, "shipping")
        self.assertEqual(snapshot.parent_container.source_context_id, self.order.number)
        self.assertTrue(str(snapshot.parent_container.container_code).strip())
        self.assertEqual(snapshot.container.parent_container_id, snapshot.parent_container_id)

    def test_save_shipping_packing_blocks_new_wb_item_barcode_mismatch_before_warehouse_sync(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._create_snapshot_box(box_code="BX-WB-MISMATCH", qty=20)
        market = Market.objects.create(id=9871, name="WB")
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.marketplace = market
        self.order.save(update_fields=["status", "marketplace", "updated_at"])
        ShippingOrder.objects.filter(pk=self.order.pk).update(
            created_at=datetime.fromisoformat("2026-07-31T00:00:00+00:00")
        )
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
            created_by=self.user,
        )
        move_op = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=self.order.number,
            requested_by=self.user,
        )
        WarehouseWritePathService.start_move_to_otg(operation=move_op, performed_by=self.user)
        WarehouseWritePathService.complete_move_to_otg(operation=move_op, performed_by=self.user)
        self.item.barcode = "200000000999"
        self.item.save(update_fields=["barcode", "updated_at"])
        self.order.refresh_from_db()

        packing_state = _shipping_packing_initial_state(self.order)
        row = packing_state["initial_boxes"][0]
        result = save_shipping_packing(
            self.order,
            boxes_state=[
                {
                    "row_key": row["row_key"],
                    "code": row["code"],
                    "qty": row["qty"],
                    "barcode_preview": row["barcode_preview"],
                    "pallet_code": "PAL-1",
                }
            ],
            pallets_state=[{"code": "PAL-1", "label": "1"}],
            delivered_boxes=packing_state["delivered_boxes"],
            initial_boxes=packing_state["initial_boxes"],
            user=self.user,
        )

        self.assertFalse(result["saved"])
        self.assertTrue(any("товарный ШК" in error for error in result["errors"]))
        self.order.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        self.assertEqual(snapshot.warehouse_state_code, WarehouseStateCode.IN_OTG.value)
        self.assertFalse(
            WarehouseContainer.objects.filter(
                agency=self.agency,
                source_context_type="shipping",
                source_context_id=self.order.number,
            ).exists()
        )

    def test_ship_order_blocks_new_wb_item_barcode_mismatch_before_stock_write(self):
        market = Market.objects.create(id=9874, name="WB")
        self.order.marketplace = market
        self.order.save(update_fields=["marketplace", "updated_at"])
        ShippingOrder.objects.filter(pk=self.order.pk).update(
            created_at=datetime.fromisoformat("2026-07-31T00:00:00+00:00")
        )
        self.order.refresh_from_db()
        snapshot = WarehouseStockSnapshot.objects.get(
            agency=self.agency,
            sku_code="SKU-001",
        )

        with patch(
            "shipping.packing._shipping_delivered_boxes",
            return_value=[
                {
                    "box_code": "BX-WB-OTHER",
                    "items": [
                        {
                            "sku_code": "SKU-OTHER",
                            "barcode": "200000000999",
                            "qty": 20,
                        }
                    ],
                }
            ],
        ):
            with self.assertRaises(ValidationError):
                ship_order(
                    self.order,
                    self.user,
                    shipped_qty_by_item={self.item.id: 20},
                )

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(self.item.qty_shipped, 0)
        self.assertEqual(snapshot.qty, 100)
        self.assertEqual(snapshot.available_qty, 100)

    def test_shipping_packing_summary_falls_back_to_warehouse_when_audit_missing(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._create_snapshot_box(box_code="BX-PACK-2", qty=20)
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=self.order.number,
            items=[
                {
                    "sku_code": self.item.sku_code,
                    "size": self.item.size,
                    "barcode": self.item.barcode,
                    "goods_type": self.item.goods_type,
                    "qty": 20,
                }
            ],
            created_by=self.user,
        )
        move_op = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=self.order.number,
            requested_by=self.user,
        )
        WarehouseWritePathService.start_move_to_otg(operation=move_op, performed_by=self.user)
        WarehouseWritePathService.complete_move_to_otg(operation=move_op, performed_by=self.user)
        packing_state = _shipping_packing_initial_state(self.order)
        row = packing_state["initial_boxes"][0]
        self._create_shipping_service_fact()
        save_shipping_packing(
            self.order,
            boxes_state=[
                {
                    "row_key": row["row_key"],
                    "code": row["code"],
                    "qty": row["qty"],
                    "barcode_preview": row["barcode_preview"],
                    "pallet_code": "PAL-1",
                }
            ],
            pallets_state=[{"code": "PAL-1", "label": "1"}],
            delivered_boxes=packing_state["delivered_boxes"],
            initial_boxes=packing_state["initial_boxes"],
            user=self.user,
        )
        OrderAuditEntry.objects.filter(
            agency=self.agency,
            order_type="shipping",
            order_id=self.order.number,
            payload__act="shipping_packing",
        ).delete()

        summary = _shipping_packing_summary(self.order)

        self.assertIsNotNone(summary)
        self.assertEqual(summary["pallet_count"], 1)
        self.assertEqual(summary["box_count"], 1)
        self.assertTrue(str(summary["pallets"][0]["code"] or "").strip())
        self.assertEqual(summary["pallets"][0]["boxes"][0]["box_code"], "BX-PACK-2")

    def test_shipping_box_composition_xlsx_matches_marketplace_box_template(self):
        OrderAuditEntry.objects.create(
            order_id=self.order.number,
            order_type="shipping",
            action="status",
            agency=self.agency,
            user=self.user,
            description="Упаковка закрыта",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-PACK-1",
                        "qty": 5,
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "barcode": "200000000001",
                                "qty": 5,
                            }
                        ],
                    }
                ],
                "act_pallets": [{"code": "PAL-1", "label": "1", "boxes": ["BX-PACK-1"]}],
            },
        )

        self.client.force_login(self.user)
        response = self.client.get(f"/shipping/{self.order.pk}/boxes.xlsx")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Баркод товара", "Кол-во товаров", "ШК короба", "Срок годности"))
        self.assertEqual(rows[1], ("200000000001", 5, "BX-PACK-1", None))
        self.assertEqual(workbook.active.max_column, 4)

    def test_warehouse_state_resolver_marks_packed_order_as_assigned_to_trip(self):
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["status", "updated_at"])

        result = WarehouseGoodsStateResolver.resolve_for_shipping_order(
            self.order,
            trip_status=LogisticsTrip.STATUS_PLANNED,
        )

        self.assertEqual(result.code, WarehouseStateCode.ASSIGNED_TO_TRIP)
        self.assertEqual(result.label_for("logistician"), "Подготовка к рейсу")

    def test_warehouse_state_resolver_marks_departed_packed_order_as_loaded(self):
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["status", "updated_at"])

        result = WarehouseGoodsStateResolver.resolve_for_shipping_order(
            self.order,
            trip_status=LogisticsTrip.STATUS_DEPARTED,
        )

        self.assertEqual(result.code, WarehouseStateCode.LOADED_TO_VEHICLE)
        self.assertEqual(result.label_for("storekeeper"), "Загружено в машину")

    def test_warehouse_state_resolver_prefers_ready_for_loading_snapshot_over_order_status(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._load_snapshot_order_ready_for_loading(qty=20)
        self.order.status = ShippingOrder.STATUS_SUBMITTED
        self.order.save(update_fields=["status", "updated_at"])

        result = WarehouseGoodsStateResolver.resolve_for_shipping_order(self.order)

        self.assertEqual(result.code, WarehouseStateCode.READY_FOR_LOADING)
        self.assertEqual(result.label_for("default"), "Подготовлена складом, ожидает логиста")

    def test_warehouse_state_resolver_prefers_loaded_snapshot_over_trip_status_fallback(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._load_snapshot_order_to_vehicle(qty=20)
        self.order.status = ShippingOrder.STATUS_SUBMITTED
        self.order.save(update_fields=["status", "updated_at"])

        result = WarehouseGoodsStateResolver.resolve_for_shipping_order(
            self.order,
            trip_status=LogisticsTrip.STATUS_PLANNED,
        )

        self.assertEqual(result.code, WarehouseStateCode.LOADED_TO_VEHICLE)
        self.assertEqual(result.label_for("default"), "Загружено в машину")

    def test_create_pick_tasks_prefers_single_pallet_when_it_covers_request(self):
        self.item.qty_requested = 110
        self.item.save(update_fields=["qty_requested", "updated_at"])
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-001",
            name="Товар 1",
            size="42",
            barcode="200000000001",
            qty=110,
            box_code="BX-2",
            pallet_code="PL-2",
            zone="MR",
            row=2,
            section=1,
            tier=1,
            cell=1,
            location_code="MR-2-1-1-1",
        )

        reserve_order(self.order, self.user)
        move_ids = create_pick_tasks(
            self.order,
            self.user,
            requested_by_name="Тест Менеджер",
            requested_by_role="manager",
        )

        self.assertEqual(len(move_ids), 1)
        self.assertEqual(MoveTask.objects.count(), 1)
        task = MoveTask.objects.first()
        self.assertEqual(task.pallet_code, "PL-2")
        self.assertEqual(task.move_mode, MoveTask.MODE_PALLET_FULL)
        self.assertEqual(task.payload.get("pick_mode"), "full")
        self.assertEqual(task.payload.get("requested_qty"), "")

    def test_ship_order_without_trip_records_shipped_event_before_archiving_snapshot(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._load_snapshot_order_ready_for_loading(qty=20)
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        shipped_event = WarehouseEvent.objects.get(
            agency=self.agency,
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            event_type="shipped",
        )
        self.assertEqual(shipped_event.qty, 20)
        self.assertIsNotNone(shipped_event.reserve_id)
        self.assertTrue(shipped_event.payload.get("direct_without_linked_trip"))
        self.assertEqual(snapshot.last_event_id, shipped_event.id)
        self.assertEqual(snapshot.qty, 0)
        self.assertEqual(snapshot.shipping_reserved_qty, 0)
        self.assertEqual(snapshot.available_qty, 0)
        self.assertEqual(snapshot.warehouse_state_code, "shipped")
        self.assertTrue(snapshot.is_archived)
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_shipped, 20)

    def test_complete_transfer_without_trip_writes_off_ready_stock(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._load_snapshot_order_ready_for_loading(qty=20)
        self.order.refresh_from_db()
        self.order.delivery_type = ShippingOrder.DELIVERY_TRANSFER
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["delivery_type", "status", "updated_at"])
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        completed = complete_transfer_without_trip(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        shipped_event = WarehouseEvent.objects.get(
            agency=self.agency,
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            event_type="shipped",
        )
        self.assertTrue(completed)
        self.assertFalse(LogisticsTripOrder.objects.filter(shipping_order=self.order).exists())
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertEqual(self.item.qty_shipped, 20)
        self.assertEqual(snapshot.qty, 0)
        self.assertEqual(snapshot.available_qty, 0)
        self.assertEqual(snapshot.warehouse_state_code, WarehouseStateCode.SHIPPED.value)
        self.assertTrue(snapshot.is_archived)
        self.assertTrue(shipped_event.payload.get("direct_without_linked_trip"))

    def test_complete_transfer_without_trip_ignores_and_releases_stale_source_reserve(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        shipped_snapshot = self._load_snapshot_order_ready_for_loading(qty=20)
        source_snapshot = self._create_snapshot_box(
            pallet_code="PL-STALE-SOURCE",
            box_code="BX-STALE-SOURCE",
            order_id="R-STALE-SOURCE",
            qty=20,
        )
        reserve = WarehouseReserve.objects.get(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_type="shipping",
            context_id=self.order.number,
        )
        reserve.qty_allocated = 20
        reserve.qty_satisfied = 20
        reserve.status = WarehouseReserve.STATUS_SATISFIED
        reserve.save(update_fields=["qty_allocated", "qty_satisfied", "status", "updated_at"])
        stale_event = WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="shipping_reserved",
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            container=source_snapshot.container,
            reserve=reserve,
            from_location=source_snapshot.location,
            to_location=source_snapshot.location,
            from_zone_code=source_snapshot.zone_code,
            to_zone_code=source_snapshot.zone_code,
            qty=20,
            performed_by=self.user,
            occurred_at=timezone.now(),
            payload={"snapshot_id": source_snapshot.id, "box_code": source_snapshot.container_code},
        )
        source_snapshot.available_qty = 0
        source_snapshot.shipping_reserved_qty = 20
        source_snapshot.warehouse_state_code = WarehouseStateCode.RESERVED_FOR_SHIPPING.value
        source_snapshot.last_event = stale_event
        source_snapshot.save(
            update_fields=[
                "available_qty",
                "shipping_reserved_qty",
                "warehouse_state_code",
                "last_event",
                "updated_at",
            ]
        )
        self.order.delivery_type = ShippingOrder.DELIVERY_TRANSFER
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["delivery_type", "status", "updated_at"])
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        completed = complete_transfer_without_trip(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        shipped_snapshot.refresh_from_db()
        source_snapshot.refresh_from_db()
        reserve.refresh_from_db()
        self.assertTrue(completed)
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertEqual(self.item.qty_shipped, 20)
        self.assertEqual(shipped_snapshot.qty, 0)
        self.assertTrue(shipped_snapshot.is_archived)
        self.assertEqual(source_snapshot.qty, 20)
        self.assertEqual(source_snapshot.available_qty, 20)
        self.assertEqual(source_snapshot.shipping_reserved_qty, 0)
        self.assertEqual(source_snapshot.warehouse_state_code, WarehouseStateCode.STORED.value)
        self.assertFalse(source_snapshot.is_archived)
        self.assertEqual(reserve.status, WarehouseReserve.STATUS_RELEASED)

    def test_complete_transfer_without_trip_rejects_linked_trip_without_stock_write(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._load_snapshot_order_ready_for_loading(qty=20)
        self.order.refresh_from_db()
        self.order.delivery_type = ShippingOrder.DELIVERY_TRANSFER
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["delivery_type", "status", "updated_at"])
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])
        trip = LogisticsTrip.objects.create(
            number="TRIP-TRANSFER-BLOCK",
            status=LogisticsTrip.STATUS_DRAFT,
            vehicle_type=LogisticsTrip.VEHICLE_FULFILLMENT,
            created_by=self.user,
        )
        LogisticsTripOrder.objects.create(trip=trip, shipping_order=self.order)

        with self.assertRaisesRegex(ValidationError, "связана с рейсом"):
            complete_transfer_without_trip(self.order, self.user)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PACKED)
        self.assertEqual(self.item.qty_reserved, 20)
        self.assertEqual(self.item.qty_shipped, 0)
        self.assertEqual(snapshot.qty, 20)
        self.assertEqual(snapshot.warehouse_state_code, WarehouseStateCode.READY_FOR_LOADING.value)
        self.assertFalse(
            WarehouseEvent.objects.filter(
                agency=self.agency,
                stock_context_type="shipping",
                stock_context_id=self.order.number,
                event_type="shipped",
            ).exists()
        )

    def test_complete_transfer_without_trip_is_idempotent_after_success(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._load_snapshot_order_ready_for_loading(qty=20)
        self.order.refresh_from_db()
        self.order.delivery_type = ShippingOrder.DELIVERY_TRANSFER
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["delivery_type", "status", "updated_at"])
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        self.assertTrue(complete_transfer_without_trip(self.order, self.user))
        self.assertFalse(complete_transfer_without_trip(self.order, self.user))
        self.assertEqual(
            WarehouseEvent.objects.filter(
                agency=self.agency,
                stock_context_type="shipping",
                stock_context_id=self.order.number,
                event_type="shipped",
            ).count(),
            1,
        )

    def test_ship_order_reloads_locked_order_before_closed_guard(self):
        ShippingOrder.objects.filter(pk=self.order.pk).update(status=ShippingOrder.STATUS_SHIPPED)

        with self.assertRaisesRegex(ValidationError, "закрытую заявку"):
            ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        self.assertFalse(
            WarehouseEvent.objects.filter(
                agency=self.agency,
                stock_context_type="shipping",
                stock_context_id=self.order.number,
                event_type="shipped",
            ).exists()
        )

    def test_ship_order_uses_warehouse_write_path_for_loaded_trip(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot = self._load_snapshot_order_to_vehicle(qty=20)
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_shipped, 20)
        self.assertEqual(self.item.qty_reserved, 0)
        self.assertEqual(snapshot.warehouse_state_code, "shipped")
        self.assertEqual(snapshot.shipping_reserved_qty, 0)
        self.assertTrue(snapshot.is_in_vehicle)
        self.assertTrue(snapshot.is_archived)
        shipped_event = WarehouseEvent.objects.get(
            agency=self.agency,
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            event_type="shipped",
        )
        self.assertEqual(shipped_event.qty, 20)
        self.assertIsNotNone(shipped_event.reserve_id)
        self.assertEqual(shipped_event.payload.get("trip_id"), "TRIP-1")
        self.assertEqual(snapshot.last_event_id, shipped_event.id)

    def test_ship_order_ignores_stored_source_remainder_with_shipping_event(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        loaded_snapshot = self._load_snapshot_order_to_vehicle(qty=20)
        source_remainder = self._create_snapshot_box(
            pallet_code="PL-SOURCE-REMAINDER",
            box_code="BX-SOURCE-REMAINDER",
            order_id="R-SOURCE-REMAINDER",
            qty=5,
        )
        remainder_event = WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="move_to_otg",
            stock_context_type="shipping",
            stock_context_id=self.order.number,
            container=source_remainder.container,
            from_location=source_remainder.location,
            to_location=source_remainder.location,
            from_zone_code=source_remainder.zone_code,
            to_zone_code=source_remainder.zone_code,
            qty=20,
            performed_by=self.user,
            occurred_at=timezone.now(),
            payload={
                "snapshot_id": source_remainder.id,
                "box_code": source_remainder.container_code,
            },
        )
        source_remainder.last_event = remainder_event
        source_remainder.save(update_fields=["last_event", "updated_at"])
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        loaded_snapshot.refresh_from_db()
        source_remainder.refresh_from_db()
        self.assertEqual(loaded_snapshot.warehouse_state_code, WarehouseStateCode.SHIPPED.value)
        self.assertTrue(loaded_snapshot.is_archived)
        self.assertEqual(source_remainder.warehouse_state_code, WarehouseStateCode.STORED.value)
        self.assertEqual(source_remainder.qty, 5)
        self.assertEqual(source_remainder.available_qty, 5)
        self.assertFalse(source_remainder.is_archived)

    def test_ship_order_auto_promotes_ready_for_loading_snapshot_to_loaded_vehicle(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot, trip = self._assign_ready_snapshot_to_trip(qty=20, start_loading=False)
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_shipped, 20)
        self.assertEqual(snapshot.current_trip_id, trip.number)
        self.assertEqual(snapshot.warehouse_state_code, "shipped")
        self.assertTrue(snapshot.is_in_vehicle)
        self.assertTrue(snapshot.is_archived)

    def test_ship_order_completes_loading_in_progress_snapshot_via_warehouse_path(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        snapshot, trip = self._assign_ready_snapshot_to_trip(qty=20, start_loading=True)
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])
        snapshot.refresh_from_db()
        self.assertEqual(snapshot.warehouse_state_code, "loading_in_progress")

        ship_order(self.order, self.user, shipped_qty_by_item={self.item.id: 20})

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_shipped, 20)
        self.assertEqual(snapshot.current_trip_id, trip.number)
        self.assertEqual(snapshot.warehouse_state_code, "shipped")
        self.assertTrue(snapshot.is_in_vehicle)
        self.assertTrue(snapshot.is_archived)

    def test_shipping_stock_picker_uses_warehouse_snapshot_when_legacy_rows_missing(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._create_snapshot_box()

        stock_rows = _shipping_stock_picker_rows(self.agency)

        self.assertEqual(len(stock_rows), 1)
        self.assertEqual(stock_rows[0]["sku_code"], "SKU-001")
        self.assertEqual(stock_rows[0]["available_boxes"], 1)
        self.assertEqual(stock_rows[0]["available_qty"], 100)

    def test_create_pick_tasks_uses_warehouse_snapshot_when_legacy_rows_missing(self):
        WarehouseStockSnapshot.objects.filter(agency=self.agency).delete()
        self._create_snapshot_box()
        self.item.qty_reserved = 20
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        move_ids = create_pick_tasks(
            self.order,
            self.user,
            requested_by_name="Тест Менеджер",
            requested_by_role="manager",
        )

        self.order.refresh_from_db()
        self.assertTrue(move_ids)
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        task = MoveTask.objects.get()
        self.assertEqual(task.pallet_code, "PL-1")
        self.assertEqual(task.to_zone, "OTG")

    def test_shipping_stock_picker_marks_mixed_boxes_by_barcodes(self):
        self._create_snapshot_box(
            order_id="R-1",
            sku="SKU-002",
            name="Товар 2",
            size="43",
            barcode="200000000002",
            qty=100,
            box_code="BX-1",
            pallet_code="PL-1",
        )
        stock_rows = _shipping_stock_picker_rows(self.agency)
        row_1 = next(row for row in stock_rows if row["sku_code"] == "SKU-001")
        row_2 = next(row for row in stock_rows if row["sku_code"] == "SKU-002")
        self.assertTrue(row_1["is_mixed_box"])
        self.assertTrue(row_2["is_mixed_box"])
        self.assertTrue(row_1["mixed_group"])
        self.assertEqual(row_1["mixed_group"], row_2["mixed_group"])

    def test_shipping_stock_picker_marks_mixed_boxes_without_barcodes_by_item_signature(self):
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-010",
            name="Товар без ШК",
            size="40",
            barcode="",
            qty=20,
            box_code="BX-NO-BC",
            pallet_code="PL-2",
            cell=2,
            location_code="OS-1-1-1-2",
        )
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-010",
            name="Товар без ШК",
            size="42",
            barcode="",
            qty=20,
            box_code="BX-NO-BC",
            pallet_code="PL-2",
            cell=2,
            location_code="OS-1-1-1-2",
        )
        stock_rows = _shipping_stock_picker_rows(self.agency)
        rows = [row for row in stock_rows if row["sku_code"] == "SKU-010"]
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row["is_mixed_box"] for row in rows))
        groups = {row["mixed_group"] for row in rows}
        self.assertEqual(len(groups), 1)

    def test_shipping_stock_picker_hides_box_with_partial_processing_reserve(self):
        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_PROCESSING,
            context_type="processing",
            context_id="P-1",
            sku_code="SKU-001",
            size="42",
            goods_type="Готовый",
            qty_reserved=20,
            status=WarehouseReserve.STATUS_ACTIVE,
        )

        stock_rows = _shipping_stock_picker_rows(self.agency)

        self.assertEqual(stock_rows, [])

    def test_shipping_stock_picker_keeps_only_fully_free_boxes(self):
        self._create_snapshot_box(
            order_id="R-2",
            sku="SKU-001",
            name="Товар 1",
            size="42",
            barcode="200000000001",
            qty=100,
            box_code="BX-2",
            pallet_code="PL-2",
            cell=2,
            location_code="OS-1-1-1-2",
        )
        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_PROCESSING,
            context_type="processing",
            context_id="P-2",
            sku_code="SKU-001",
            size="42",
            goods_type="Готовый",
            qty_reserved=20,
            status=WarehouseReserve.STATUS_ACTIVE,
        )

        stock_rows = _shipping_stock_picker_rows(self.agency)

        self.assertEqual(len(stock_rows), 1)
        self.assertEqual(stock_rows[0]["sku_code"], "SKU-001")
        self.assertEqual(stock_rows[0]["available_boxes"], 1)
        self.assertEqual(stock_rows[0]["available_qty"], 100)

    def test_shipping_stock_picker_exclude_order_restores_current_order_box(self):
        order = ShippingOrder.objects.create(
            number="SO-EXCLUDE-1",
            agency=self.agency,
            status=ShippingOrder.STATUS_RESERVED,
        )
        item = ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-001",
            name="Товар 1",
            size="42",
            barcode="200000000001",
            goods_type="Готовый",
            qty_requested=20,
            qty_reserved=20,
        )
        WarehouseReserve.objects.create(
            agency=self.agency,
            reserve_type=WarehouseReserve.TYPE_SHIPPING,
            context_type="shipping",
            context_id=order.number,
            sku_code="SKU-001",
            size="42",
            goods_type="Готовый",
            qty_reserved=20,
            status=WarehouseReserve.STATUS_ACTIVE,
        )

        self.assertEqual(_shipping_stock_picker_rows(self.agency), [])

        stock_rows = _shipping_stock_picker_rows(self.agency, exclude_order=order)

        self.assertEqual(len(stock_rows), 1)
        self.assertEqual(stock_rows[0]["available_boxes"], 1)
        self.assertEqual(stock_rows[0]["available_qty"], 100)

    def test_shipping_reachtruck_metrics_counts_done_otg_tasks(self):
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-A",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={"shipping_order_id": self.order.number, "requested_boxes": [{"box_code": "B1"}, {"box_code": "B2"}]},
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-B",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={"shipping_order_id": self.order.number, "requested_box": "B3"},
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-C",
            to_zone="OBR",
            status=MoveTask.STATUS_DONE,
            payload={"shipping_order_id": self.order.number, "requested_box": "B4"},
        )
        metrics = _shipping_reachtruck_metrics(self.order)
        self.assertEqual(metrics["pallet_count"], 2)
        self.assertEqual(metrics["box_count"], 3)

    def test_shipping_reachtruck_metrics_falls_back_for_done_task_with_stale_payload(self):
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-STUCK",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            qty_done=1090,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
                "status": "created",
                "requested_boxes": [],
            },
        )

        metrics = _shipping_reachtruck_metrics(self.order)

        self.assertEqual(metrics["pallet_count"], 1)
        self.assertGreater(metrics["box_count"], 0)

    def test_shipping_detail_hides_otg_packing_label_without_delivered_boxes(self):
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-A",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
            },
        )
        staff_user = get_user_model().objects.create_user(username="storekeeper_sync_shipping", password="pwd")
        Employee.objects.create(
            full_name="Кладовщик Синхронизатор",
            role="storekeeper",
            user=staff_user,
            is_active=True,
        )
        self.client.force_login(staff_user)

        response = self.client.get(f"/shipping/{self.order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        self.assertNotContains(response, "Товар доставлен в OTG, ожидает паллетизации")
        self.assertNotContains(response, "Разложить короба по новым паллетам")

    def test_shipping_detail_shows_otg_packing_label_when_delivered_boxes_resolve(self):
        self.order.status = ShippingOrder.STATUS_PICKING
        self.order.save(update_fields=["status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-1",
                        "qty": 20,
                        "items": [
                            {
                                "sku_code": "SKU-001",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "200000000001",
                                "goods_type": "Готовый",
                                "qty": 20,
                            }
                        ],
                    }
                ],
            },
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id="",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-A",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            payload={
                "shipping_order_id": self.order.number,
                "shipping_order_pk": self.order.pk,
                "receiving_order_id": "R-1",
                "picked_boxes": ["BX-1"],
                "picked_rows": [{"box_code": "BX-1", "qty": 20}],
            },
        )
        staff_user = get_user_model().objects.create_user(username="storekeeper_ready_shipping", password="pwd")
        Employee.objects.create(
            full_name="Кладовщик Готовый",
            role="storekeeper",
            user=staff_user,
            is_active=True,
        )
        self.client.force_login(staff_user)

        response = self.client.get(f"/shipping/{self.order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PICKING)
        self.assertContains(response, "Товар доставлен в OTG, ожидает паллетизации")
        self.assertContains(response, "Разложить короба по новым паллетам")


class ShippingReadModelServiceTests(TestCase):
    def setUp(self):
        self.request_factory = RequestFactory()
        user_model = get_user_model()
        self.manager_user = user_model.objects.create_user(username="shipping_read_manager", password="pwd")
        Employee.objects.create(
            full_name="Shipping Read Manager",
            role="manager",
            user=self.manager_user,
            is_active=True,
        )
        self.storekeeper_user = user_model.objects.create_user(username="shipping_read_storekeeper", password="pwd")
        Employee.objects.create(
            full_name="Shipping Read Storekeeper",
            role="storekeeper",
            user=self.storekeeper_user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Shipping Read Agency")
        self.market = Market.objects.create(id=303, name="Ozon")
        self.order = ShippingOrder.objects.create(
            number="SO-000777",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        ShippingOrderItem.objects.create(
            order=self.order,
            sku_code="SKU-LIST",
            name="List item",
            size="42",
            barcode="200000007777",
            goods_type="Готовый",
            qty_requested=10,
        )

    def test_build_shipping_list_page_context_applies_filters_and_display_fields(self):
        request = self.request_factory.get(
            reverse("shipping:list"),
            data={"client": str(self.agency.id), "status": ShippingOrder.STATUS_SUBMITTED},
        )
        request.user = self.manager_user

        context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="manager",
            client_agency=None,
        )

        orders = list(context["orders"])
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0].pk, self.order.pk)
        self.assertEqual(context["selected_client"], self.agency)
        self.assertEqual(context["client_filter"], str(self.agency.id))
        self.assertEqual(orders[0].display_number, _display_shipping_number(self.order.number))

    def test_shipping_list_includes_manager_approved_fbs_outbound_for_storekeeper(self):
        from uuid import uuid4

        from fbs.models import FbsExternalIssue, FbsExternalIssueEvent

        issue = FbsExternalIssue.objects.create(
            agency=self.agency,
            reference="POSITIVE-OUT-1",
            recipient="ООО ПОЗИТИВ",
            purpose="owner",
            basis="Возврат товара из FBS",
            shipping_details={
                "delivery_type": "pickup",
                "delivery_type_label": "Самовывоз",
                "eta_date": "2026-09-19",
                "vehicle_type": "client",
                "vehicle_type_label": "Транспорт клиента",
                "vehicle_number": "A123BC77",
            },
            request_key=uuid4(),
            request_hash="created-hash",
            created_by=self.manager_user,
        )
        FbsExternalIssueEvent.objects.create(
            issue=issue,
            request_key=uuid4(),
            request_hash="review-hash",
            action="manager_review_required",
            payload={},
            performed_by=self.manager_user,
        )
        FbsExternalIssueEvent.objects.create(
            issue=issue,
            request_key=uuid4(),
            request_hash="approved-hash",
            action="manager_approved",
            payload={},
            performed_by=self.manager_user,
        )
        request = self.request_factory.get(reverse("shipping:list"))
        request.user = self.storekeeper_user

        context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="storekeeper",
            client_agency=None,
        )

        self.assertEqual([row.pk for row in context["fbs_outbound_issues"]], [issue.pk])
        self.client.force_login(self.storekeeper_user)
        response = self.client.get(reverse("shipping:list"))
        self.assertContains(response, issue.number)
        self.assertContains(response, "Вывоз товара из FBS")
        self.assertContains(response, "Самовывоз")
        self.assertContains(response, "A123BC77")
        self.assertContains(response, reverse("fbs:external_issue_detail", args=[issue.pk]))

        manager_context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="manager",
            client_agency=None,
        )
        self.assertEqual(manager_context["fbs_outbound_issues"], [])

    def test_shipping_list_context_adds_problem_report_for_staff(self):
        blocker = ShippingOrder.objects.create(
            number="OTG-PROBLEM-BLOCKER",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PICKING,
            shipping_discrepancy_status="pending",
        )
        ozon_mismatch = ShippingOrder.objects.create(
            number="OTG-PROBLEM-OZON",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_SUBMITTED,
            marketplace=self.market,
            expected_boxes=1,
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=ozon_mismatch.number,
            action="update",
            agency=self.agency,
            user=self.manager_user,
            payload={"ozon_boxes_total": 3, "gm_barcodes": ["GM-1", "GM-2", "GM-3"]},
        )
        missing_fact = ShippingOrder.objects.create(
            number="OTG-PROBLEM-FACT",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
            expected_boxes=2,
            planned_ship_date=timezone.localdate(),
        )
        reachtruck_error = ShippingOrder.objects.create(
            number="OTG-PROBLEM-RT",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PICKING,
        )
        MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(reachtruck_error.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_BLOCKED,
            planning_error="HTTP 500 от ричтрака",
        )
        request = self.request_factory.get(
            reverse("shipping:list"),
            data={"client": str(self.agency.id)},
        )
        request.user = self.manager_user

        context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="manager",
            client_agency=None,
        )

        report = context["shipping_problem_report"]
        self.assertTrue(report["show"])
        self.assertTrue(report["has_rows"])
        self.assertEqual(report["totals"]["active_blockers"], 1)
        self.assertEqual(report["totals"]["box_imbalance"], 1)
        self.assertEqual(report["totals"]["missing_fact"], 1)
        self.assertEqual(report["totals"]["reachtruck_errors"], 1)
        self.assertIn(blocker.number, report["sections"][0]["rows"][0]["number"])
        self.assertIn("Ozon 3, Fullbox 1", report["sections"][1]["rows"][0]["detail"])
        self.assertIn("Факт OTG 0 из 2", report["sections"][2]["rows"][0]["detail"])
        self.assertIn("HTTP 500", report["sections"][3]["rows"][0]["detail"])

    def test_shipping_list_renders_problem_report_for_manager(self):
        ShippingOrder.objects.create(
            number="OTG-PROBLEM-VIEW",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
            expected_boxes=2,
            planned_ship_date=timezone.localdate(),
        )
        self.client.force_login(self.manager_user)

        response = self.client.get(reverse("shipping:list"), {"client": str(self.agency.id)})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Проблемные отгрузки сегодня")
        self.assertContains(response, "Недостающий факт в OTG")
        self.assertContains(response, "Факт OTG 0 из 2")

    def test_problem_report_ignores_old_active_order_without_today_schedule(self):
        old_order = ShippingOrder.objects.create(
            number="OTG-OLD-ACTIVE",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=4,
        )
        ShippingOrder.objects.filter(pk=old_order.pk).update(
            created_at=timezone.now() - timedelta(days=3),
        )
        request = self.request_factory.get(reverse("shipping:list"), data={"client": str(self.agency.id)})
        request.user = self.manager_user

        context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="manager",
            client_agency=None,
        )

        report = context["shipping_problem_report"]
        self.assertEqual(report["totals"]["missing_fact"], 0)

    def test_problem_report_ignores_old_reachtruck_error(self):
        order = ShippingOrder.objects.create(
            number="OTG-OLD-RT-ERROR",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PICKING,
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_BLOCKED,
            planning_error="Старая ошибка",
        )
        MoveRequest.objects.filter(pk=move_request.pk).update(
            updated_at=timezone.now() - timedelta(days=2),
        )
        request = self.request_factory.get(reverse("shipping:list"), data={"client": str(self.agency.id)})
        request.user = self.manager_user

        context = build_shipping_list_page_context(
            request=request,
            scope="staff",
            role="manager",
            client_agency=None,
        )

        report = context["shipping_problem_report"]
        self.assertEqual(report["totals"]["reachtruck_errors"], 0)

    def test_build_shipping_detail_page_context_adds_transport_note_and_pick_reason(self):
        self.order.status = ShippingOrder.STATUS_STOREKEEPER_ACCEPTED
        self.order.save(update_fields=["status", "updated_at"])
        request = self.request_factory.get(reverse("shipping:detail", args=[self.order.pk]))
        request.user = self.storekeeper_user

        context = build_shipping_detail_page_context(
            request=request,
            order=self.order,
            scope="staff",
            role="storekeeper",
        )

        self.assertEqual(context["order"].pk, self.order.pk)
        self.assertEqual(context["transport_note_url"], reverse("shipping:documents", args=[self.order.pk]))
        self.assertIn("pick_readiness", context)
        self.assertEqual(context["pick_readiness"], shipping_pick_readiness(self.order))
        self.assertIn("can_storekeeper_pick", context)
        self.assertIn("storekeeper_pick_unavailable_reason", context)

    def test_build_shipping_detail_page_context_allows_packed_transfer_completion(self):
        self.order.delivery_type = ShippingOrder.DELIVERY_TRANSFER
        self.order.status = ShippingOrder.STATUS_PACKED
        self.order.save(update_fields=["delivery_type", "status", "updated_at"])
        request = self.request_factory.get(reverse("shipping:detail", args=[self.order.pk]))
        request.user = self.storekeeper_user

        context = build_shipping_detail_page_context(
            request=request,
            order=self.order,
            scope="staff",
            role="storekeeper",
        )

        self.assertTrue(context["can_complete_transfer"])


class OzonGmLabelValidationTests(SimpleTestCase):
    def test_groups_only_exact_saved_gm_cargoes(self):
        groups, error = _ozon_label_groups(
            [
                {"gm_barcode": "GM-001", "supply_id": "8001", "cargo_id": "7001"},
                {"gm_barcode": "GM-002", "supply_id": "8001", "cargo_id": "7002"},
            ],
            expected_gm_barcodes=["GM-001", "GM-002"],
        )

        self.assertEqual(error, "")
        self.assertEqual(groups, {8001: [7001, 7002]})

    def test_rejects_incomplete_gm_cargoes_without_api_call(self):
        groups, error = _ozon_label_groups(
            [{"gm_barcode": "GM-001", "supply_id": "8001", "cargo_id": "7001"}],
            expected_gm_barcodes=["GM-001", "GM-002"],
        )

        self.assertEqual(groups, {})
        self.assertIn("неполный", error.lower())

    def test_rejects_gm_cargo_without_ozon_identifiers(self):
        groups, error = _ozon_label_groups(
            [{"gm_barcode": "GM-001", "supply_id": "", "cargo_id": ""}],
            expected_gm_barcodes=["GM-001"],
        )

        self.assertEqual(groups, {})
        self.assertIn("supply_id", error)
        self.assertIn("cargo_id", error)

    @patch("shipping.ozon_supplies.resolve_ozon_credential")
    @patch("shipping.ozon_supplies.ozon_get_pdf")
    @patch("shipping.ozon_supplies.ozon_post")
    def test_fetches_official_pdf_by_saved_supply_and_cargo_ids(
        self,
        ozon_post_mock,
        ozon_get_pdf_mock,
        resolve_credential_mock,
    ):
        resolve_credential_mock.return_value = (
            SimpleNamespace(client_id="12345", market_key="api-key"),
            "",
        )
        ozon_post_mock.side_effect = [
            ({"operation_id": "operation-001"}, None),
            (
                {
                    "status": "SUCCESS",
                    "result": {"file_guid": "12345678-abcd-1234-abcd-123456789012"},
                },
                None,
            ),
        ]
        ozon_get_pdf_mock.return_value = (b"%PDF-1.4\nlabel\n%%EOF", "")

        documents, error, pending = fetch_ozon_gm_label_documents(
            SimpleNamespace(id=987654),
            [{"gm_barcode": "GM-001", "supply_id": "8001", "cargo_id": "7001"}],
            expected_gm_barcodes=["GM-001"],
            poll_attempts=1,
        )

        self.assertEqual(error, "")
        self.assertFalse(pending)
        self.assertEqual(documents, [{"supply_id": 8001, "content": b"%PDF-1.4\nlabel\n%%EOF"}])
        self.assertEqual(ozon_post_mock.call_args_list[0].args[0], "/v1/cargoes-label/create")
        self.assertEqual(
            ozon_post_mock.call_args_list[0].args[3],
            {"supply_id": 8001, "cargoes": [{"cargo_id": 7001}]},
        )
        self.assertEqual(ozon_post_mock.call_args_list[1].args[0], "/v1/cargoes-label/get")
        ozon_get_pdf_mock.assert_called_once_with(
            "/v1/cargoes-label/file/12345678-abcd-1234-abcd-123456789012",
            "12345",
            "api-key",
            timeout=20,
        )


class ShippingManagerTaskFlowTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.client_user = user_model.objects.create_user(username="client_shipping", password="pwd")
        self.manager_user = user_model.objects.create_user(username="manager_shipping", password="pwd")
        self.storekeeper_user = user_model.objects.create_user(username="storekeeper_shipping", password="pwd")
        self.logistician_user = user_model.objects.create_user(username="logistician_shipping", password="pwd")
        self.manager = Employee.objects.create(
            full_name="Менеджер Проверяющий",
            role="manager",
            user=self.manager_user,
            is_active=True,
        )
        self.storekeeper = Employee.objects.create(
            full_name="Кладовщиков Алексей",
            role="storekeeper",
            user=self.storekeeper_user,
            is_active=True,
        )
        self.logistician = Employee.objects.create(
            full_name="Логист Петров",
            role="logistician",
            user=self.logistician_user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Клиент отгрузки", portal_user=self.client_user)
        self.market = Market.objects.create(id=301, name="Ozon")
        self.shipping_service, _created = BillingService.objects.get_or_create(
            code="shipping_phase3_manager_flow_test",
            defaults={"name": "Паллетизация отгрузки", "unit": "шт", "is_active": True},
        )
        create_warehouse_snapshot_row(
            agency=self.agency,
            order_type="receiving",
            order_id="R-SH-1",
            sku="SKU-SHIP-1",
            name="Товар отгрузки",
            size="44",
            barcode="300000000001",
            goods_type="Готовый",
            qty=24,
            box_code="BX-SH-1",
            pallet_code="PL-SH-1",
            zone="OS",
            row=1,
            section=1,
            tier=1,
            cell=1,
        )
        create_warehouse_snapshot_row(
            agency=self.agency,
            order_type="receiving",
            order_id="R-SH-2",
            sku="SKU-SHIP-1",
            name="Товар отгрузки",
            size="44",
            barcode="300000000001",
            goods_type="Готовый",
            qty=24,
            box_code="BX-SH-2",
            pallet_code="PL-SH-2",
            zone="OS",
            row=1,
            section=1,
            tier=1,
            cell=2,
        )

    def test_storekeeper_task_due_date_uses_planned_shipping_date(self):
        order = ShippingOrder.objects.create(
            number="OTG-STOREKEEPER-DUE",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_RESERVED,
            planned_ship_date=date(2026, 8, 3),
        )

        task = ensure_storekeeper_task(order, user=self.manager_user)

        self.assertIsNotNone(task)
        due_date = timezone.localtime(task.due_date)
        self.assertEqual(due_date.date(), date(2026, 8, 3))
        self.assertEqual((due_date.hour, due_date.minute), (23, 59))

    def test_manager_task_starts_unassigned_even_with_client_manager(self):
        from shipping.services import ensure_manager_review_task

        assigned_user = get_user_model().objects.create_user(
            username="assigned_client_manager", password="pwd"
        )
        Employee.objects.create(
            full_name="Назначенный менеджер клиента",
            role="manager",
            user=assigned_user,
            is_active=True,
        )
        self.agency.mened_user_id = assigned_user.id
        self.agency.save(update_fields=["mened_user_id"])
        order = ShippingOrder.objects.create(
            number="OTG-CLIENT-MANAGER",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )

        task = ensure_manager_review_task(order, user=self.client_user)

        self.assertIsNotNone(task)
        self.assertIsNone(task.assigned_to_id)

    def test_manager_task_due_date_is_24_hours_from_submission(self):
        from shipping.services import ensure_manager_review_task

        submitted_at = timezone.make_aware(datetime(2026, 8, 31, 10, 47))
        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-SLA-24",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )

        task = ensure_manager_review_task(
            order,
            user=self.client_user,
            submitted_at=submitted_at,
        )

        self.assertEqual(task.due_date - submitted_at, timedelta(hours=24))

    def test_repeated_manager_task_sync_preserves_manual_deadline(self):
        from shipping.services import ensure_manager_review_task

        submitted_at = timezone.make_aware(datetime(2026, 8, 31, 10, 47))
        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-MANUAL-DUE",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        task = ensure_manager_review_task(
            order,
            user=self.client_user,
            submitted_at=submitted_at,
        )
        manual_due_date = submitted_at + timedelta(days=3)
        Task.objects.filter(pk=task.pk).update(due_date=manual_due_date)

        repeated = ensure_manager_review_task(
            order,
            user=self.client_user,
            submitted_at=submitted_at + timedelta(hours=1),
        )

        self.assertEqual(repeated.due_date, manual_due_date)

    def test_manager_task_keeps_employee_who_claimed_it(self):
        from shipping.services import ensure_manager_review_task

        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-CLAIM",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        task = ensure_manager_review_task(order, user=self.client_user)
        task.assigned_to = self.manager
        task.status = "in_progress"
        task.save(update_fields=["assigned_to", "status", "updated_at"])

        repeated = ensure_manager_review_task(order, user=self.client_user)

        self.assertEqual(repeated.pk, task.pk)
        self.assertEqual(repeated.assigned_to_id, self.manager.id)
        self.assertEqual(Task.objects.filter(route=f"/shipping/{order.pk}/").count(), 1)

    def test_unassigned_manager_task_closes_after_approval(self):
        from shipping.services import close_manager_review_task, ensure_manager_review_task

        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-CLOSE-QUEUE",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        task = ensure_manager_review_task(order, user=self.client_user)

        close_manager_review_task(order)

        task.refresh_from_db()
        self.assertEqual(task.status, "done")

    def test_act_task_inherits_claimed_manager(self):
        from shipping.services import ensure_manager_review_task, ensure_shipping_act_manager_task

        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-ACT-OWNER",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SHIPPED,
        )
        review_task = ensure_manager_review_task(order, user=self.client_user)
        review_task.assigned_to = self.manager
        review_task.status = "done"
        review_task.save(update_fields=["assigned_to", "status", "updated_at"])

        act_task = ensure_shipping_act_manager_task(order, user=self.logistician_user)

        self.assertIsNotNone(act_task)
        self.assertEqual(act_task.assigned_to_id, self.manager.id)

    def test_unassigned_act_task_closes_after_signature(self):
        from shipping.services import close_shipping_act_manager_task, ensure_shipping_act_manager_task

        order = ShippingOrder.objects.create(
            number="OTG-MANAGER-ACT-CLOSE-QUEUE",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SHIPPED,
        )
        act_task = ensure_shipping_act_manager_task(order, user=self.logistician_user)

        close_shipping_act_manager_task(order)

        act_task.refresh_from_db()
        self.assertEqual(act_task.status, "done")

    def test_new_logistician_task_uses_least_loaded_employee(self):
        from shipping.services import ensure_logistician_task

        second_user = get_user_model().objects.create_user(
            username="logistician_shipping_second", password="pwd"
        )
        second_logistician = Employee.objects.create(
            full_name="Логист Свободный",
            role="logistician",
            user=second_user,
            is_active=True,
        )
        Task.objects.create(
            title="Другая активная заявка",
            route="/shipping/999999/",
            assigned_to=self.logistician,
            status="backlog",
        )
        order = ShippingOrder.objects.create(
            number="OTG-LEAST-LOADED",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
        )

        task = ensure_logistician_task(order, user=self.manager_user)

        self.assertIsNotNone(task)
        self.assertEqual(task.assigned_to_id, second_logistician.id)
        self.assertEqual(
            ShippingRoutingState.objects.get(shipping_order=order).status,
            ShippingRoutingState.STATUS_READY,
        )

    def test_logistician_task_and_routing_state_are_idempotent(self):
        from shipping.services import ensure_logistician_task

        order = ShippingOrder.objects.create(
            number="OTG-ROUTING-IDEMPOTENT",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
        )

        first = ensure_logistician_task(order, user=self.manager_user)
        second = ensure_logistician_task(order, user=self.manager_user)

        self.assertEqual(first.pk, second.pk)
        self.assertEqual(
            Task.objects.filter(
                route=f"/shipping/{order.pk}/",
                assigned_to__role="logistician",
            ).count(),
            1,
        )
        self.assertEqual(
            ShippingRoutingState.objects.filter(shipping_order=order).count(),
            1,
        )

    def test_new_logistician_task_rolls_back_when_routing_sync_fails(self):
        from shipping.services import RoutingStateSyncError, ensure_logistician_task

        order = ShippingOrder.objects.create(
            number="OTG-ROUTING-FAIL-NEW",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
        )
        route = f"/shipping/{order.pk}/"

        with patch(
            "logistics.routing_services.mark_ready_for_routing",
            side_effect=RuntimeError("routing unavailable"),
        ):
            with self.assertLogs("shipping.services", level="ERROR"):
                with self.assertRaises(RoutingStateSyncError):
                    ensure_logistician_task(order, user=self.manager_user)

        self.assertFalse(Task.objects.filter(route=route).exists())
        self.assertFalse(
            ShippingRoutingState.objects.filter(shipping_order=order).exists()
        )

    def test_existing_logistician_task_update_rolls_back_when_routing_sync_fails(self):
        from shipping.services import RoutingStateSyncError, ensure_logistician_task

        order = ShippingOrder.objects.create(
            number="OTG-ROUTING-FAIL-EXISTING",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
        )
        existing = Task.objects.create(
            title="Старая задача логиста",
            description="Старое описание",
            route=f"/shipping/{order.pk}/",
            assigned_to=self.logistician,
            status="done",
        )

        with patch(
            "logistics.routing_services.mark_ready_for_routing",
            side_effect=RuntimeError("routing unavailable"),
        ):
            with self.assertLogs("shipping.services", level="ERROR"):
                with self.assertRaises(RoutingStateSyncError):
                    ensure_logistician_task(order, user=self.manager_user)

        existing.refresh_from_db()
        self.assertEqual(existing.title, "Старая задача логиста")
        self.assertEqual(existing.description, "Старое описание")
        self.assertEqual(existing.status, "done")
        self.assertFalse(
            ShippingRoutingState.objects.filter(shipping_order=order).exists()
        )

    def test_transfer_does_not_create_logistician_task(self):
        from shipping.services import ensure_logistician_task

        order = ShippingOrder.objects.create(
            number="OTG-TRANSFER-NO-TRIP",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_TRANSFER,
        )
        existing = Task.objects.create(
            title=f"Заявка на отгрузку №{order.number}",
            route=f"/shipping/{order.pk}/",
            assigned_to=self.logistician,
            status="backlog",
        )

        task = ensure_logistician_task(order, user=self.manager_user)

        existing.refresh_from_db()
        self.assertIsNone(task)
        self.assertEqual(existing.status, "done")

    def test_return_act_groups_same_product_variant_into_one_row(self):
        order = ShippingOrder.objects.create(
            number="OTG-RETURN-ACT-1",
            agency=self.agency,
            created_by=self.storekeeper_user,
        )
        common_fields = {
            "order": order,
            "sku_code": "SKU-RETURN-1",
            "name": "Товар для акта",
            "size": "0",
            "barcode": "1600005000001",
            "goods_type": "gv",
        }
        ShippingOrderItem.objects.create(**common_fields, qty_requested=50)
        ShippingOrderItem.objects.create(**common_fields, qty_requested=49)
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-RETURN-2",
            name="Другой товар",
            size="0",
            barcode="1600005000002",
            goods_type="gv",
            qty_requested=20,
        )

        rows = _item_rows(order)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["index"], "1")
        self.assertEqual(rows[0]["code"], "SKU-RETURN-1")
        self.assertEqual(rows[0]["qty"], "99")
        self.assertEqual(rows[1]["index"], "2")
        self.assertEqual(rows[1]["code"], "SKU-RETURN-2")
        self.assertEqual(rows[1]["qty"], "20")

    def _submit_payload(self):
        picker_key = _shipping_stock_picker_rows(self.agency)[0]["key"]
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        return {
            "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
            "slot_date": eta.strftime("%Y-%m-%d"),
            "slot_time": "12:30",
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "shipping_barcode": "SHIP-CLIENT-1",
            "marketplace": str(self.market.id),
            "supply_number": "SUPPLY-1",
            "wb_supply_barcode": "SUPPLY-1",
            "destination_warehouse": "Склад маркетплейса",
            "supply_type": ShippingOrder.SUPPLY_BOX,
            "vehicle_type": ShippingOrder.VEHICLE_FULFILLMENT,
            "vehicle_number": "A123BC77",
            "driver_phone": "+7 900 000-00-00",
            "comment": "Клиентская отправка",
            "action": "submit",
            "stock_key_all[]": [picker_key],
            "stock_boxes[]": ["1"],
        }

    def _create_order_ready_for_packing(self, *, with_services: bool = True):
        order = ShippingOrder.objects.create(
            number="SO-000010",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=2,
            shipping_barcode="SHIP-PACK-1",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-SHIP-1",
            name="Товар отгрузки",
            size="44",
            barcode="300000000001",
            goods_type="Готовый",
            qty_requested=48,
            qty_reserved=48,
        )
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="receiving",
            order_id="R-SH-1",
            action="placement",
            payload={
                "act": "placement",
                "act_state": "closed",
                "act_boxes": [
                    {
                        "code": "BX-SH-1",
                        "qty": 24,
                        "items": [
                            {"sku_code": "SKU-SHIP-1", "name": "Товар отгрузки", "size": "44", "barcode": "300000000001", "goods_type": "Готовый", "qty": 24}
                        ],
                    },
                    {
                        "code": "BX-SH-2",
                        "qty": 24,
                        "items": [
                            {"sku_code": "SKU-SHIP-1", "name": "Товар отгрузки", "size": "44", "barcode": "300000000001", "goods_type": "Готовый", "qty": 24}
                        ],
                    },
                ],
            },
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=f"shipping:{order.pk}",
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        MoveTask.objects.create(
            request=move_request,
            pallet_code="PL-OTG-1",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
            qty_done=48,
            payload={
                "shipping_order_id": order.number,
                "shipping_order_pk": order.pk,
                "receiving_order_id": "R-SH-1",
                "picked_boxes": ["BX-SH-1", "BX-SH-2"],
                "picked_rows": [
                    {"box_code": "BX-SH-1", "qty": 24},
                    {"box_code": "BX-SH-2", "qty": 24},
                ],
            },
        )
        WarehouseWritePathService.reserve_for_shipping(
            agency=self.agency,
            order_id=order.number,
            items=[
                {
                    "sku_code": "SKU-SHIP-1",
                    "size": "44",
                    "barcode": "300000000001",
                    "goods_type": "Готовый",
                    "qty": 48,
                }
            ],
            created_by=self.manager_user,
        )
        move_operation = WarehouseWritePathService.request_move_to_otg(
            agency=self.agency,
            order_id=order.number,
            requested_by=self.manager_user,
        )
        WarehouseWritePathService.start_move_to_otg(
            operation=move_operation,
            performed_by=self.storekeeper_user,
        )
        WarehouseWritePathService.complete_move_to_otg(
            operation=move_operation,
            performed_by=self.storekeeper_user,
        )
        if with_services:
            WarehouseServiceFact.objects.create(
                client=self.agency,
                order_type=WarehouseServiceFact.ORDER_SHIPPING,
                order_id=order.number,
                service=self.shipping_service,
                service_name_snapshot=self.shipping_service.name,
                quantity=2,
                unit="шт",
                status=WarehouseServiceFact.STATUS_SENT_TO_BILLING,
                reported_by=self.storekeeper_user,
            )
        return order

    def _mark_transport_note_ready(
        self,
        order: ShippingOrder,
        *,
        vehicle_type: str = ShippingOrder.VEHICLE_CLIENT,
    ) -> ShippingOrder:
        order.status = ShippingOrder.STATUS_PACKED
        order.vehicle_type = vehicle_type
        order.save(update_fields=["status", "vehicle_type", "updated_at"])
        if vehicle_type == ShippingOrder.VEHICLE_FULFILLMENT:
            trip = LogisticsTrip.objects.create(
                number=f"TRIP-TN-{order.pk}",
                status=LogisticsTrip.STATUS_DEPARTED,
                vehicle_type=vehicle_type,
                created_by=self.logistician_user,
                assigned_logistician=self.logistician,
            )
            LogisticsTripOrder.objects.create(
                trip=trip,
                shipping_order=order,
                loading_sequence=1,
                delivery_sequence=1,
            )
        return order

    def _packing_rows_payload(self):
        return [
            {
                "row_key": _shipping_box_row_key({"code": "BX-SH-1", "receiving_order_id": "R-SH-1"}, 1),
                "code": "BX-SH-1",
                "qty": 24,
                "barcode_preview": "300000000001 - 24 шт.",
                "pallet_code": "PAL-1",
            },
            {
                "row_key": _shipping_box_row_key({"code": "BX-SH-2", "receiving_order_id": "R-SH-1"}, 2),
                "code": "BX-SH-2",
                "qty": 24,
                "barcode_preview": "300000000001 - 24 шт.",
                "pallet_code": "PAL-2",
            },
        ]

    def test_manager_draft_autosave_creates_draft_order(self):
        self.client.force_login(self.manager_user)

        response = self.client.post(
            "/shipping/new/",
            data={
                "agency": str(self.agency.id),
                "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
                "draft_autosave": "1",
                "action": "save_draft",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertTrue(data["ok"], data)
        order = ShippingOrder.objects.get(pk=data["order_id"])
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(order.agency_id, self.agency.id)
        self.assertEqual(data["draft_order_id"], order.pk)

    def test_manager_draft_autosave_updates_existing_draft(self):
        self.client.force_login(self.manager_user)
        create_response = self.client.post(
            "/shipping/new/",
            data={
                "agency": str(self.agency.id),
                "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
                "draft_autosave": "1",
                "comment": "Первый черновик",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(create_response.status_code, 200, create_response.content)
        order_id = create_response.json()["order_id"]

        response = self.client.post(
            f"/shipping/new/?order={order_id}&edit=1",
            data={
                "agency": str(self.agency.id),
                "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
                "draft_autosave": "1",
                "comment": "Обновлённый черновик",
                "edit_order_id": str(order_id),
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200, response.content)
        data = response.json()
        self.assertTrue(data["ok"], data)
        self.assertEqual(data["order_id"], order_id)
        order = ShippingOrder.objects.get(pk=order_id)
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(order.comment, "Обновлённый черновик")

    def test_client_draft_type_change_requires_explicit_reset_confirmation(self):
        wb_market = Market.objects.create(id=302, name="WB")
        order = ShippingOrder.objects.create(
            number="OTG-DRAFT-TYPE-GUARD",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=wb_market,
            shipping_barcode="OZON-GUARD",
            supply_number="OZON-GUARD",
            comment="Сохранить до подтверждения",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-SHIP-1",
            name="Товар отгрузки",
            size="44",
            barcode="300000000001",
            goods_type="Готовый",
            qty_requested=24,
        )
        self.client.force_login(self.client_user)

        response = self.client.post(
            f"/shipping/new/?client={self.agency.id}&order={order.pk}&edit=1",
            data={
                "agency": str(self.agency.id),
                "client": str(self.agency.id),
                "edit_order_id": str(order.pk),
                "edit": "1",
                "delivery_type": ShippingOrder.DELIVERY_TRANSFER,
                "action": "save_draft",
                "draft_autosave": "1",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn("Подтвердите изменение типа отгрузки", response.json()["error"])
        order.refresh_from_db()
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_MARKETPLACE)
        self.assertEqual(order.shipping_barcode, "OZON-GUARD")
        self.assertEqual(order.items.count(), 1)

    def test_confirmed_client_draft_type_change_clears_data_but_keeps_attachments(self):
        wb_market = Market.objects.create(id=302, name="WB")
        order = ShippingOrder.objects.create(
            number="OTG-DRAFT-TYPE-RESET",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=wb_market,
            slot_date=timezone.localdate() + timedelta(days=2),
            slot_time="12:30",
            eta_at=timezone.now() + timedelta(days=2),
            destination_address="Старый адрес",
            planned_ship_date=timezone.localdate() + timedelta(days=2),
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            vehicle_number="A123BC77",
            driver_phone="+7 900 000-00-00",
            shipping_barcode="OZON-RESET",
            wb_supply_barcode="WB-RESET",
            wb_transit_warehouse=True,
            transit_address="Старый транзит",
            destination_warehouse="Старый склад",
            supply_type=ShippingOrder.SUPPLY_BOX,
            supply_number="OZON-RESET",
            distribution_status=ShippingOrder.DISTRIBUTION_COMPLETE,
            expected_boxes=2,
            place_type=ShippingOrder.PLACE_TYPE_BOX,
            comment="Старый комментарий",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-SHIP-1",
            name="Товар отгрузки",
            size="44",
            barcode="300000000001",
            goods_type="Готовый",
            qty_requested=24,
        )
        order.destinations.create(
            warehouse_name="Старый склад Ozon",
            planned_units=24,
            planned_boxes=2,
        )
        attachment = ShippingOrderAttachment.objects.create(
            order=order,
            uploaded_by=self.client_user,
            file="shipping_attachments/test/keep-after-type-reset.pdf",
        )
        self.client.force_login(self.client_user)

        response = self.client.post(
            f"/shipping/new/?client={self.agency.id}&order={order.pk}&edit=1",
            data={
                "agency": str(self.agency.id),
                "client": str(self.agency.id),
                "edit_order_id": str(order.pk),
                "edit": "1",
                "delivery_type": ShippingOrder.DELIVERY_TRANSFER,
                "action": "save_draft",
                "draft_autosave": "1",
                "reset_draft_on_type_change": "1",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["ok"], response.json())
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_TRANSFER)
        self.assertIsNone(order.marketplace)
        self.assertIsNone(order.slot_date)
        self.assertIsNone(order.slot_time)
        self.assertIsNone(order.eta_at)
        self.assertIsNone(order.planned_ship_date)
        self.assertEqual(order.destination_address, "")
        self.assertEqual(order.vehicle_type, "")
        self.assertEqual(order.vehicle_number, "")
        self.assertEqual(order.driver_phone, "")
        self.assertEqual(order.shipping_barcode, "")
        self.assertEqual(order.wb_supply_barcode, "")
        self.assertFalse(order.wb_transit_warehouse)
        self.assertEqual(order.transit_address, "")
        self.assertEqual(order.destination_warehouse, "")
        self.assertEqual(order.supply_type, "")
        self.assertEqual(order.supply_number, "")
        self.assertEqual(order.distribution_status, "")
        self.assertEqual(order.expected_boxes, 0)
        self.assertEqual(order.place_type, "")
        self.assertEqual(order.comment, "")
        self.assertEqual(order.items.count(), 0)
        self.assertEqual(order.destinations.count(), 0)
        self.assertTrue(ShippingOrderAttachment.objects.filter(pk=attachment.pk).exists())

        page = self.client.get(
            f"/shipping/new/?client={self.agency.id}&order={order.pk}&edit=1"
        )
        self.assertEqual(page.status_code, 200, page.content)
        self.assertContains(page, "Вы уверены, что хотите изменить тип отгрузки?")

    def test_confirmed_transfer_draft_can_switch_to_marketplace_before_selection(self):
        order = ShippingOrder.objects.create(
            number="OTG-DRAFT-TRANSFER-TO-MARKETPLACE",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_TRANSFER,
            destination_address="Старое направление",
            comment="Старый комментарий",
        )
        ShippingOrderItem.objects.create(
            order=order, sku_code="SKU-SHIP-1", name="Товар отгрузки",
            size="44", barcode="300000000001", goods_type="Готовый", qty_requested=24,
        )
        attachment = ShippingOrderAttachment.objects.create(
            order=order, uploaded_by=self.client_user,
            file="shipping_attachments/test/keep-marketplace-reset.pdf",
        )
        self.client.force_login(self.client_user)
        payload = {
            "agency": str(self.agency.id), "client": str(self.agency.id),
            "edit_order_id": str(order.pk), "edit": "1",
            "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
            "action": "save_draft", "draft_autosave": "1",
        }
        url = f"/shipping/new/?client={self.agency.id}&order={order.pk}&edit=1"
        rejected = self.client.post(url, payload, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(rejected.status_code, 400, rejected.content)
        order.refresh_from_db()
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_TRANSFER)
        self.assertEqual(order.items.count(), 1)

        payload["reset_draft_on_type_change"] = "1"
        response = self.client.post(url, payload, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["ok"])
        self.assertEqual(response.json()["order_id"], order.pk)
        order.refresh_from_db()
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_MARKETPLACE)
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertIsNone(order.marketplace_id)
        self.assertEqual(order.items.count(), 0)
        self.assertEqual(order.reserves.count(), 0)
        self.assertEqual(order.destination_address, "")
        self.assertEqual(order.comment, "")
        self.assertTrue(ShippingOrderAttachment.objects.filter(pk=attachment.pk).exists())
        self.assertEqual(ShippingOrder.objects.count(), 1)

        # Further incomplete autosaves remain prohibited after the one-off reset.
        repeated = self.client.post(url, payload, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(repeated.status_code, 400, repeated.content)

    def test_client_ozon_draft_saves_are_rejected_without_creating_order(self):
        self.client.force_login(self.client_user)
        for autosave in (False, True):
            with self.subTest(autosave=autosave):
                draft_payload = self._submit_payload()
                draft_payload["action"] = "save_draft"
                if autosave:
                    draft_payload["draft_autosave"] = "1"

                draft_response = self.client.post(
                    "/shipping/new/",
                    data=draft_payload,
                    HTTP_X_REQUESTED_WITH="XMLHttpRequest",
                )

                self.assertEqual(draft_response.status_code, 400, draft_response.content)
                self.assertFalse(draft_response.json()["ok"])
                self.assertIn(
                    "сохранение черновиков отключено",
                    draft_response.json()["error"],
                )
                self.assertEqual(ShippingOrder.objects.count(), 0)

    def test_client_ozon_submit_saves_selected_documents(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["supply_number"] = "2000059517576"
        payload["shipping_barcode"] = "GM-001"
        payload["wb_supply_barcode"] = ""
        payload["ozon_supply_payload_json"] = json.dumps(
            {
                "order_id": 123456,
                "order_number": "2000059517576",
                "ozon_boxes_total": 1,
                "gm_barcodes": ["GM-001"],
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-001",
                        "items": [
                            {
                                "offer_id": "SKU-SHIP-1",
                                "barcode": "300000000001",
                                "name": "Товар отгрузки",
                                "quantity": 24,
                            }
                        ],
                    }
                ],
                "destination_warehouses": ["OZON ХОРУГВИНО"],
                "form": {
                    "supply_number": "2000059517576",
                    "shipping_barcode": "GM-001",
                    "ozon_gm_comment": "ШК ГМ Ozon: GM-001",
                },
                "composition": [
                    {
                        "offer_id": "SKU-SHIP-1",
                        "barcode": "300000000001",
                        "quantity": 24,
                        "ozon_boxes": 1,
                        "ozon_gm_barcodes": ["GM-001"],
                    }
                ],
            },
            ensure_ascii=False,
        )
        payload["ozon_gm_comment"] = "ШК ГМ Ozon: GM-001"
        payload["ozon_gm_barcodes"] = "GM-001"
        payload["ozon_boxes_total"] = "1"
        payload["documents"] = SimpleUploadedFile(
            "ozon-label.pdf",
            b"%PDF-1.4 selected",
            content_type="application/pdf",
        )

        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post("/shipping/new/", data=payload)

            self.assertEqual(response.status_code, 302, response.content)
            order = ShippingOrder.objects.latest("pk")
            attachments = list(
                ShippingOrderAttachment.objects.filter(order=order).values_list(
                    "file",
                    flat=True,
                )
            )
            self.assertTrue(
                any(str(name).endswith("ozon-label.pdf") for name in attachments),
                attachments,
            )

    def test_client_marketplace_draft_waits_for_marketplace_selection(self):
        self.client.force_login(self.client_user)
        draft_payload = self._submit_payload()
        draft_payload["delivery_type"] = ShippingOrder.DELIVERY_MARKETPLACE
        draft_payload["marketplace"] = ""
        draft_payload["action"] = "save_draft"
        draft_payload["draft_autosave"] = "1"

        response = self.client.post(
            "/shipping/new/",
            data=draft_payload,
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 400, response.content)
        self.assertFalse(response.json()["ok"])
        self.assertIn("Сначала выберите маркетплейс", response.json()["error"])
        self.assertEqual(ShippingOrder.objects.count(), 0)

    def test_client_non_ozon_draft_autosave_still_creates_draft(self):
        wb_market = Market.objects.create(id=303, name="WB")
        self.client.force_login(self.client_user)
        draft_payload = self._submit_payload()
        draft_payload["marketplace"] = str(wb_market.id)
        draft_payload["action"] = "save_draft"
        draft_payload["draft_autosave"] = "1"

        response = self.client.post(
            "/shipping/new/",
            data=draft_payload,
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["ok"], response.json())
        order = ShippingOrder.objects.get(pk=response.json()["order_id"])
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(order.agency_id, self.agency.id)
        self.assertEqual(order.marketplace_id, wb_market.id)

    def test_client_existing_ozon_draft_autosave_is_rejected_without_changes(self):
        order = ShippingOrder.objects.create(
            number="OTG-LEGACY-OZON-DRAFT",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            shipping_barcode="OZON-LEGACY",
            supply_number="OZON-LEGACY",
            comment="Старый черновик",
        )
        self.client.force_login(self.client_user)

        response = self.client.post(
            f"/shipping/new/?client={self.agency.id}&order={order.pk}&edit=1",
            data={
                "agency": str(self.agency.id),
                "client": str(self.agency.id),
                "edit_order_id": str(order.pk),
                "edit": "1",
                "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
                "action": "save_draft",
                "draft_autosave": "1",
                "comment": "Не должно сохраниться",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn("сохранение черновиков отключено", response.json()["error"])
        order.refresh_from_db()
        self.assertEqual(order.comment, "Старый черновик")
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)

    def test_client_ozon_submit_blocks_duplicate_active_supply(self):
        self.client.force_login(self.client_user)
        ShippingOrder.objects.create(
            number="OTG-000900",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            shipping_barcode="SHIP-CLIENT-1",
            supply_number="SUPPLY-1",
        )

        response = self.client.post("/shipping/new/", data=self._submit_payload())

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(ShippingOrder.objects.count(), 1)
        self.assertContains(response, "Поставка Ozon SUPPLY-1 уже загружена")

    def test_client_submit_creates_manager_review_task(self):
        self.client.force_login(self.client_user)

        response = self.client.post("/shipping/new/", data=self._submit_payload())

        self.assertEqual(response.status_code, 302)
        order = ShippingOrder.objects.get()
        task = Task.objects.get(route=f"/shipping/{order.pk}/")
        self.assertEqual(order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(order.expected_boxes, 1)
        self.assertEqual(str(order.slot_date), (timezone.localtime() + timedelta(days=2)).date().isoformat())
        self.assertEqual(task.assigned_to, self.manager)
        self.assertEqual(task.status, "backlog")
        self.assertIn(order.number, task.title)

    def test_client_transfer_submit_does_not_require_transport(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload.update(
            {
                "delivery_type": ShippingOrder.DELIVERY_TRANSFER,
                "destination_address": "Склад ФБС",
                "marketplace": "",
                "shipping_barcode": "",
                "supply_number": "",
                "wb_supply_barcode": "",
                "destination_warehouse": "",
                "supply_type": "",
                "slot_date": "",
                "slot_time": "",
                "vehicle_type": "",
                "vehicle_number": "",
                "driver_phone": "",
            }
        )

        response = self.client.post("/shipping/new/", data=payload)

        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_TRANSFER)
        self.assertEqual(order.destination_address, "Склад ФБС")
        self.assertEqual(order.vehicle_type, "")
        self.assertEqual(order.vehicle_number, "")
        self.assertEqual(order.driver_phone, "")
        self.assertEqual(order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(order.expected_boxes, 1)

    def test_client_pickup_submit_does_not_require_transport(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload.update(
            {
                "delivery_type": ShippingOrder.DELIVERY_PICKUP,
                "marketplace": "",
                "shipping_barcode": "",
                "supply_number": "",
                "wb_supply_barcode": "",
                "destination_warehouse": "",
                "supply_type": "",
                "slot_date": "",
                "slot_time": "",
                "vehicle_type": "",
                "vehicle_number": "",
                "driver_phone": "",
            }
        )

        response = self.client.post("/shipping/new/", data=payload)

        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_PICKUP)
        self.assertEqual(order.vehicle_type, "")
        self.assertEqual(order.vehicle_number, "")
        self.assertEqual(order.driver_phone, "")
        self.assertEqual(order.status, ShippingOrder.STATUS_SUBMITTED)
        self.assertEqual(order.expected_boxes, 1)

    def test_client_ozon_api_payload_is_saved_for_manager(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["supply_number"] = "2000059517576"
        payload["shipping_barcode"] = "GM-001"
        payload["wb_supply_barcode"] = ""
        payload["ozon_supply_payload_json"] = json.dumps(
            {
                "order_id": 123456,
                "order_number": "2000059517576",
                "ozon_boxes_total": 2,
                "gm_barcodes": ["GM-001", "GM-002"],
                "destination_warehouses": ["OZON ХОРУГВИНО"],
                "form": {
                    "supply_number": "2000059517576",
                    "shipping_barcode": "GM-001",
                    "ozon_gm_comment": "ШК ГМ Ozon: GM-001, GM-002",
                },
                "composition": [
                    {
                        "offer_id": "SKU-SHIP-1",
                        "barcode": "300000000001",
                        "quantity": 24,
                        "ozon_boxes": 2,
                        "ozon_gm_barcodes": ["GM-001", "GM-002"],
                    }
                ],
            },
            ensure_ascii=False,
        )
        payload["ozon_gm_comment"] = "ШК ГМ Ozon: GM-001, GM-002"
        payload["ozon_gm_barcodes"] = "GM-001, GM-002"
        payload["ozon_boxes_total"] = "2"

        response = self.client.post("/shipping/new/", data=payload)

        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()
        self.assertEqual(order.supply_number, "2000059517576")
        self.assertEqual(order.shipping_barcode, "GM-001")
        self.assertIn("Данные Ozon API:", order.comment)
        self.assertIn("Номер поставки Ozon: 2000059517576", order.comment)
        self.assertIn("ШК ГМ Ozon: GM-001, GM-002", order.comment)
        self.assertIn("Коробов Ozon: 2", order.comment)
        self.assertTrue(order.items.filter(comment__contains="ШК ГМ Ozon: GM-001, GM-002").exists())
        entry = OrderAuditEntry.objects.filter(order_type="shipping", order_id=order.number, action="create").latest("id")
        self.assertEqual(entry.payload["supply_number"], "2000059517576")
        self.assertEqual(entry.payload["gm_barcodes"], ["GM-001", "GM-002"])
        self.assertEqual(entry.payload["ozon_boxes_total"], 2)

    def test_client_ozon_submit_removes_gm_for_removed_item(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["supply_number"] = "2000061426070"
        payload["shipping_barcode"] = "1022095068167000"
        payload["wb_supply_barcode"] = ""
        payload["ozon_supply_payload_json"] = json.dumps(
            {
                "order_id": 777,
                "order_number": "2000061426070",
                "ozon_boxes_total": 2,
                "gm_barcodes": ["1022095068167000", "1022095068169000"],
                "form": {
                    "supply_number": "2000061426070",
                    "shipping_barcode": "1022095068167000",
                },
                "gm_cargoes": [
                    {
                        "gm_barcode": "1022095068167000",
                        "items": [
                            {
                                "offer_id": "SKU-SHIP-1",
                                "barcode": "300000000001",
                                "quantity": 24,
                            }
                        ],
                    },
                    {
                        "gm_barcode": "1022095068169000",
                        "items": [
                            {
                                "offer_id": "polka001",
                                "barcode": "2038273624348",
                                "quantity": 12,
                            }
                        ],
                    },
                ],
            },
            ensure_ascii=False,
        )
        payload["ozon_boxes_total"] = "2"
        payload["ozon_gm_barcodes"] = "1022095068167000, 1022095068169000"

        response = self.client.post("/shipping/new/", data=payload)

        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()
        self.assertEqual(order.expected_boxes, 1)
        self.assertEqual(order.items.count(), 1)
        entry = OrderAuditEntry.objects.filter(
            order_type="shipping",
            order_id=order.number,
            action="create",
        ).latest("id")
        self.assertEqual(entry.payload["gm_barcodes"], ["1022095068167000"])
        self.assertEqual(entry.payload["ozon_boxes_total"], 1)
        self.assertTrue(entry.payload["ozon_gm_reconciled_after_item_removal"])
        self.assertEqual(
            [row["gm_barcode"] for row in entry.payload["gm_cargoes"]],
            ["1022095068167000"],
        )

    def test_client_ozon_api_payload_is_visible_on_edit_form(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["supply_number"] = "2000059517576"
        payload["shipping_barcode"] = "GM-001"
        payload["wb_supply_barcode"] = ""
        payload["ozon_supply_payload_json"] = json.dumps(
            {
                "order_id": 123456,
                "order_number": "2000059517576",
                "ozon_boxes_total": 2,
                "gm_barcodes": ["GM-001", "GM-002"],
                "form": {
                    "supply_number": "2000059517576",
                    "shipping_barcode": "GM-001",
                    "ozon_gm_comment": "ШК ГМ Ozon: GM-001, GM-002",
                },
            },
            ensure_ascii=False,
        )
        payload["ozon_gm_comment"] = "ШК ГМ Ozon: GM-001, GM-002"
        payload["ozon_gm_barcodes"] = "GM-001, GM-002"
        payload["ozon_boxes_total"] = "2"
        response = self.client.post("/shipping/new/", data=payload)
        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()

        page = self.client.get(f"/shipping/new/?order={order.pk}&edit=1")

        self.assertEqual(page.status_code, 200, page.content)
        self.assertContains(page, "Данные Ozon по заявке")
        self.assertContains(page, "2000059517576")
        self.assertContains(page, "коробов Ozon")
        self.assertContains(page, "GM-001, GM-002")
        self.assertContains(page, 'name="ozon_gm_barcodes"')

    def test_client_ozon_edit_form_distributes_gm_codes_without_destination(self):
        self.client.force_login(self.client_user)
        order = ShippingOrder.objects.create(
            number="SO-000095",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            supply_number="2000060437104",
            shipping_barcode="2000060437104",
            expected_boxes=2,
            wb_transit_warehouse=True,
            transit_address="СТАРАЯ_КУПАВНА_28",
            destination_warehouse="",
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="полка027",
            name="Полка для ванной угловая",
            barcode="2046431529114",
            goods_type="gv",
            qty_requested=20,
            comment="Коробов: 1; кратность: 20",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="полка026",
            name="Полка для ванной",
            barcode="2046431528490",
            goods_type="gv",
            qty_requested=20,
            comment="Коробов: 1; кратность: 20",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            user=self.client_user,
            payload={
                "supply_number": "2000060437104",
                "shipping_barcode": "2000060437104",
                "ozon_boxes_total": 2,
                "ozon_gm_comment": "ШК ГМ Ozon: 1022092968613000, 1022092968612000",
                "gm_barcodes": ["1022092968613000", "1022092968612000"],
                "transit_address": "СТАРАЯ_КУПАВНА_28",
                "destination_warehouse": "",
            },
        )

        page = self.client.get(f"/shipping/new/?order={order.pk}&edit=1")

        self.assertEqual(page.status_code, 200, page.content)
        self.assertContains(page, "Данные Ozon по заявке")
        self.assertContains(page, "Ozon не вернул конечный склад")
        self.assertContains(page, '<td style="word-break:break-word;">1022092968613000</td>', html=False)
        self.assertContains(page, '<td style="word-break:break-word;">1022092968612000</td>', html=False)

    def test_client_ozon_edit_form_uses_summary_destination_as_initial_value(self):
        self.client.force_login(self.client_user)
        order = ShippingOrder.objects.create(
            number="SO-000096",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_DRAFT,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            supply_number="2000060437105",
            shipping_barcode="2000060437105",
            expected_boxes=1,
            wb_transit_warehouse=True,
            transit_address="СТАРАЯ_КУПАВНА_28",
            destination_warehouse="",
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            comment="Данные Ozon API:\nСклады Ozon: СТАРАЯ_КУПАВНА_28",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="полка027",
            name="Полка для ванной угловая",
            barcode="2046431529114",
            goods_type="gv",
            qty_requested=20,
            comment="Коробов: 1; кратность: 20; ШК ГМ Ozon: 1022092968613000",
        )

        page = self.client.get(f"/shipping/new/?order={order.pk}&edit=1")

        self.assertEqual(page.status_code, 200, page.content)
        self.assertContains(page, 'name="destination_warehouse" value="СТАРАЯ_КУПАВНА_28"', html=False)

    def test_manager_return_to_rework_allows_client_to_add_box(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["supply_number"] = "2000059517000"
        response = self.client.post("/shipping/new/", data=payload)
        self.assertEqual(response.status_code, 302, response.content)
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        response = self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})
        self.assertEqual(response.status_code, 302, response.content)
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_RESERVED)

        response = self.client.post(f"/shipping/{order.pk}/", data={"action": "release_reserve"})
        self.assertEqual(response.status_code, 302, response.content)
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertTrue(_can_edit_items("client", "client", order))

        picker_key = _shipping_stock_picker_rows(self.agency, exclude_order=order)[0]["key"]
        self.client.force_login(self.client_user)
        response = self.client.post(
            f"/shipping/{order.pk}/",
            data={
                "action": "add_item",
                "stock_key": picker_key,
                "stock_key_all[]": [picker_key],
                "stock_boxes[]": ["1"],
            },
        )
        self.assertEqual(response.status_code, 302, response.content)
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(order.items.count(), 2)
        self.assertEqual(order.expected_boxes, 2)

    def test_manager_approval_closes_review_task_and_creates_storekeeper_task(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()
        manager_task = Task.objects.get(route=f"/shipping/{order.pk}/", assigned_to=self.manager)

        self.client.force_login(self.manager_user)
        response = self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        manager_task.refresh_from_db()
        storekeeper_task = Task.objects.get(route=f"/shipping/{order.pk}/", assigned_to=self.storekeeper)
        self.assertEqual(order.status, ShippingOrder.STATUS_RESERVED)
        self.assertEqual(manager_task.status, "done")
        self.assertEqual(storekeeper_task.status, "backlog")
        self.assertEqual(storekeeper_task.title, f"Заявка на отгрузку №{order.number}")

    def test_storekeeper_can_accept_order_into_work(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.client.force_login(self.storekeeper_user)
        response = self.client.post(f"/shipping/{order.pk}/", data={"action": "accept_storekeeper"})

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        storekeeper_task = Task.objects.get(route=f"/shipping/{order.pk}/", assigned_to=self.storekeeper)
        self.assertEqual(order.status, ShippingOrder.STATUS_PICKING)
        self.assertEqual(storekeeper_task.status, "in_progress")
        self.assertTrue(MoveTask.objects.exists())

    def test_priority_immediately_dispatches_accepted_order_to_tsd(self):
        order = ShippingOrder.objects.create(
            number="OTG-PRIORITY-DISPATCH",
            agency=self.agency,
            created_by=self.storekeeper_user,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
        )
        Task.objects.create(
            title=f"Заявка на отгрузку №{order.number}",
            route=f"/shipping/{order.pk}/",
            assigned_to=self.storekeeper,
            created_by=self.storekeeper_user,
            status="in_progress",
            priority="normal",
        )
        request = RequestFactory().post(
            f"/shipping/{order.pk}/",
            {"action": "set_shipping_priority"},
        )
        request.user = self.storekeeper_user
        permissions = ShippingDetailActionPermissions(
            can_edit_items=False,
            can_submit_for_approval=False,
            can_manager_approve=False,
            can_manager_reopen=False,
            can_storekeeper_accept=False,
            can_storekeeper_pick=True,
            can_set_shipping_priority=True,
        )

        with patch(
            "shipping.actions.start_storekeeper_pick",
            return_value=["MOVE-PRIORITY"],
        ) as start_pick:
            result = handle_shipping_detail_action(
                action="set_shipping_priority",
                order=order,
                request=request,
                role="storekeeper",
                stock_rows_for_order=[],
                permissions=permissions,
                parse_selected_stock_items=lambda *_args, **_kwargs: ([], []),
                selected_stock_rows_with_boxes=lambda *_args, **_kwargs: [],
                parse_box_count_from_comment=lambda _comment: 0,
                log_update=lambda *_args, **_kwargs: None,
            )

        task = Task.objects.get(route=f"/shipping/{order.pk}/")
        self.assertEqual(task.priority, "urgent")
        start_pick.assert_called_once_with(
            order,
            self.storekeeper_user,
            requested_by_name=self.storekeeper_user.get_full_name()
            or self.storekeeper_user.username,
            requested_by_role="storekeeper",
        )
        self.assertEqual(
            result.messages,
            [
                (
                    "success",
                    "Срочный приоритет установлен. Задания направлены на ТСД: "
                    "MOVE-PRIORITY.",
                )
            ],
        )

    def test_storekeeper_detail_shows_accept_button_before_pick_tasks(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Принять в работу")
        self.assertContains(response, "const submitter = event.submitter;")
        self.assertContains(response, "submitActionCopy")
        self.assertContains(response, "button.disabled = true;")
        self.assertContains(response, "submitter.dataset.submittingLabel")
        self.assertNotContains(response, "Дать ричтраку доставку в OTG")

    def test_storekeeper_summary_shows_transit_warehouse(self):
        order = ShippingOrder.objects.create(
            number="OTG-TRANSIT-SUMMARY",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_RESERVED,
            wb_transit_warehouse=True,
            transit_address="МО_ЩЕРБИНКА_ХАБ",
            destination_warehouse="Новосибирск",
        )

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-testid="storekeeper-transit-warehouse"')
        self.assertContains(response, "Транзитный склад")
        self.assertContains(response, "МО_ЩЕРБИНКА_ХАБ")

    def test_storekeeper_summary_shows_saved_transit_address_without_legacy_flag(self):
        order = ShippingOrder.objects.create(
            number="OTG-TRANSIT-ADDRESS",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_RESERVED,
            wb_transit_warehouse=False,
            transit_address="КАЗАНЬ_ТРАНЗИТ",
            destination_warehouse="Екатеринбург",
        )

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-testid="storekeeper-transit-warehouse"')
        self.assertContains(response, "КАЗАНЬ_ТРАНЗИТ")

    def test_storekeeper_summary_shows_no_transit_for_regular_shipping(self):
        order = ShippingOrder.objects.create(
            number="OTG-WITHOUT-TRANSIT",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_RESERVED,
            wb_transit_warehouse=False,
            transit_address="",
            destination_warehouse="Москва",
        )

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-testid="storekeeper-transit-warehouse"')
        self.assertContains(
            response,
            '<span class="storekeeper-summary-value" data-testid="storekeeper-transit-warehouse-value" title="">Нет</span>',
            html=True,
        )

    def test_storekeeper_history_is_rendered_as_compact_rows(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/?tab=history")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="storekeeper-history-head"')
        self.assertContains(response, "Комментарий")
        self.assertContains(response, "Пользователь")
        self.assertContains(response, "storekeeper-history-row")

    def test_storekeeper_detail_hides_transport_note_link_before_loading(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, f"/shipping/{order.pk}/documents/")
        self.assertNotContains(response, "Сопроводительные документы")

    def test_storekeeper_detail_hides_pick_button_when_reserved_stock_has_no_pallets(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})
        self.client.force_login(self.storekeeper_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "accept_storekeeper"})
        WarehouseContainer.objects.filter(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_BOX,
        ).update(parent_container=None)
        WarehouseStockSnapshot.objects.filter(agency=self.agency).update(parent_container=None)

        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Дать ричтраку доставку в OTG")
        self.assertContains(response, "зарезервированный товар не размещен на паллетах")

    def test_storekeeper_pick_returns_specific_reason_when_stock_has_no_pallets(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})
        self.client.force_login(self.storekeeper_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "accept_storekeeper"})
        WarehouseContainer.objects.filter(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_BOX,
        ).update(parent_container=None)
        WarehouseStockSnapshot.objects.filter(agency=self.agency).update(parent_container=None)

        response = self.client.post(f"/shipping/{order.pk}/", data={"action": "create_pick_tasks"}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "зарезервированный товар не размещен на паллетах")
        self.assertEqual(MoveRequest.objects.count(), 0)

    def test_storekeeper_detail_shows_transport_note_link_after_loading_completed(self):
        order = self._create_order_ready_for_packing()
        order.status = ShippingOrder.STATUS_PACKED
        order.vehicle_type = ShippingOrder.VEHICLE_FULFILLMENT
        order.save(update_fields=["status", "vehicle_type", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 2,
                "delivered_box_count": 2,
                "act_pallets": [
                    {"label": "1", "code": "PAL-1", "boxes": ["BX-SH-1"]},
                    {"label": "2", "code": "PAL-2", "boxes": ["BX-SH-2"]},
                ],
                "act_boxes": [
                    {"code": "BX-SH-1", "qty": 24, "pallet_label": "1", "items": []},
                    {"code": "BX-SH-2", "qty": 24, "pallet_label": "2", "items": []},
                ],
            },
        )
        trip = LogisticsTrip.objects.create(
            number="TRIP-DOCS-0001",
            status=LogisticsTrip.STATUS_LOADING,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            created_by=self.logistician_user,
            assigned_logistician=self.logistician,
        )
        LogisticsTripOrder.objects.create(
            trip=trip,
            shipping_order=order,
            loading_sequence=1,
            delivery_sequence=1,
        )
        OrderAuditEntry.objects.create(
            order_id=str(trip.pk),
            order_type="logistics_trip",
            action="update",
            agency=self.agency,
            user=self.storekeeper_user,
            description="Все паллеты погружены",
            payload={
                "act": "trip_loading_progress",
                "trip_pk": trip.pk,
                "loaded_pallet_keys": ["PAL-1", "PAL-2"],
                "last_loaded_key": "PAL-2",
            },
        )

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f"/shipping/{order.pk}/documents/")
        self.assertContains(response, "Сопроводительные документы")

    def test_manager_detail_shows_edit_button_for_submitted_order(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f"/shipping/new/?order={order.pk}&amp;edit=1")
        self.assertContains(response, "Отредактировать заявку")
        self.assertNotContains(response, "Добавить позицию")
        self.assertContains(response, "Главная роли")
        self.assertNotContains(response, "Главная клиента")
        self.assertContains(response, "Текущий пользователь: Менеджер Проверяющий")
        self.assertNotContains(response, "Текущий пользователь: Клиент")

    def test_manager_detail_shows_ozon_api_summary_with_gm_rows(self):
        order = ShippingOrder.objects.create(
            number="SO-000095",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            supply_number="2000060437104",
            shipping_barcode="2000060437104",
            expected_boxes=2,
            wb_transit_warehouse=True,
            transit_address="СТАРАЯ_КУПАВНА_28",
            destination_warehouse="",
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            comment="Данные Ozon API:\nСклады Ozon: СТАРАЯ_КУПАВНА_28",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="полка027",
            name="Полка для ванной угловая",
            barcode="2046431529114",
            goods_type="gv",
            qty_requested=20,
            qty_reserved=20,
            comment="Коробов: 1; кратность: 20",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="полка026",
            name="Полка для ванной",
            barcode="2046431528490",
            goods_type="gv",
            qty_requested=20,
            qty_reserved=20,
            comment="Коробов: 1; кратность: 20",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            user=self.client_user,
            payload={
                "supply_number": "2000060437104",
                "shipping_barcode": "2000060437104",
                "ozon_boxes_total": 2,
                "gm_barcodes": ["1022092968613000", "1022092968612000"],
                "gm_cargoes": [
                    {
                        "gm_barcode": "1022092968613000",
                        "cargo_id": 7001,
                        "supply_id": 8001,
                        "items": [
                            {
                                "offer_id": "OZON-OFFER-SHELF-027",
                                "barcode": "2046431529114",
                                "quantity": 20,
                            }
                        ],
                    },
                    {
                        "gm_barcode": "1022092968612000",
                        "cargo_id": 7002,
                        "supply_id": 8001,
                        "items": [{"barcode": "2046431528490", "quantity": 20}],
                    },
                ],
                "transit_address": "СТАРАЯ_КУПАВНА_28",
            },
        )

        self.client.force_login(self.manager_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200, response.content)
        self.assertContains(response, "Данные Ozon для менеджера")
        self.assertContains(response, "2000060437104")
        self.assertContains(response, "СТАРАЯ_КУПАВНА_28")
        self.assertContains(response, "1022092968613000")
        self.assertContains(response, "1022092968612000")
        self.assertContains(response, "OZON-OFFER-SHELF-027")
        self.assertContains(response, "ШК товара:")
        self.assertContains(response, "2046431529114")
        self.assertContains(response, "ШК ГМ Ozon")
        self.assertContains(response, "Скачать этикетки ШК ГМ Ozon")
        self.assertContains(response, reverse("shipping:ozon-gm-labels", args=[order.pk]))

    def test_manager_detail_shows_ozon_gm_discrepancy_panel(self):
        order = ShippingOrder.objects.create(
            number="220_OTG",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_PICKING,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            expected_boxes=1,
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            comment="Данные Ozon API:\nКоробов Ozon: 2",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-IN-FULLBOX",
            name="Товар в заявке",
            barcode="2043986017000",
            goods_type="gv",
            qty_requested=10,
            qty_reserved=10,
            comment="Коробов: 1; кратность: 10",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            user=self.client_user,
            payload={
                "ozon_boxes_total": 2,
                "gm_barcodes": ["GM-220-1", "GM-220-2"],
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-220-1",
                        "items": [{"offer_id": "SKU-IN-FULLBOX", "barcode": "2043986017000", "quantity": 10}],
                    },
                    {
                        "gm_barcode": "GM-220-2",
                        "items": [{"offer_id": "SKU-MISSING", "barcode": "2043986017738", "quantity": 32}],
                    },
                ],
            },
        )

        self.client.force_login(self.manager_user)
        delivered_boxes = [
            {
                "box_code": "FB-221-1",
                "items": [{"sku_code": "SKU-IN-FULLBOX", "barcode": "2043986017000", "qty": 10}],
            },
            {
                "box_code": "FB-221-2",
                "items": [{"sku_code": "SKU-CLOSED", "barcode": "2043986017738", "qty": 32}],
            },
        ]
        with patch("shipping.views._shipping_delivered_boxes", return_value=delivered_boxes):
            response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200, response.content)
        self.assertContains(response, "Расхождения Ozon ГМ")
        self.assertContains(response, "Коробов Ozon")
        self.assertContains(response, "Коробов Fullbox")
        self.assertContains(response, "ШК товара есть в Ozon ГМ, но отсутствует в заявке Fullbox")
        self.assertContains(response, "2043986017738")
        self.assertContains(response, "SKU-MISSING")
        self.assertContains(response, "32 шт.")
        self.assertContains(response, "Расхождение не оформлялось")

    def test_manager_detail_marks_resolved_ozon_gm_mismatch_as_closed_supplement(self):
        order = ShippingOrder.objects.create(
            number="221_OTG",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            expected_boxes=2,
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            comment="Данные Ozon API:\nКоробов Ozon: 2",
            shipping_discrepancy_status="resolved",
            shipping_discrepancy_payload={
                "workflow_version": 2,
                "status": "resolved",
                "decision": "supplement",
                "additional_pick_move_ids": ["MOVE-221"],
            },
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-IN-FULLBOX",
            name="Товар в заявке",
            barcode="2043986017000",
            goods_type="gv",
            qty_requested=10,
            qty_reserved=10,
            comment="Коробов: 1; кратность: 10",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            user=self.client_user,
            payload={
                "ozon_boxes_total": 2,
                "gm_barcodes": ["GM-221-1", "GM-221-2"],
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-221-1",
                        "items": [{"offer_id": "SKU-IN-FULLBOX", "barcode": "2043986017000", "quantity": 10}],
                    },
                    {
                        "gm_barcode": "GM-221-2",
                        "items": [{"offer_id": "SKU-CLOSED", "barcode": "2043986017738", "quantity": 32}],
                    },
                ],
            },
        )

        self.client.force_login(self.manager_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200, response.content)
        self.assertContains(response, "Расхождения Ozon ГМ")
        self.assertContains(response, "ШК товара добраны и закрыты через расхождение")
        self.assertContains(response, "2043986017738")
        self.assertContains(response, "SKU-CLOSED")
        self.assertContains(response, "Закрыто добором")
        self.assertContains(response, "Добор закрыт через расхождение")
        self.assertNotContains(response, "ШК товара есть в Ozon ГМ, но отсутствует в заявке Fullbox")

    def test_storekeeper_binding_alert_marks_resolved_supplement_as_closed(self):
        order = ShippingOrder.objects.create(
            number="222_OTG",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_PACKED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            expected_boxes=2,
            supply_type=ShippingOrder.SUPPLY_BOX,
            vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT,
            shipping_discrepancy_status="resolved",
            shipping_discrepancy_payload={
                "workflow_version": 2,
                "status": "resolved",
                "decision": "supplement",
            },
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-IN-FULLBOX",
            name="Товар в заявке",
            barcode="2043986017000",
            goods_type="gv",
            qty_requested=10,
            qty_reserved=10,
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            user=self.client_user,
            payload={
                "ozon_boxes_total": 2,
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-222-1",
                        "items": [{"offer_id": "SKU-IN-FULLBOX", "barcode": "2043986017000", "quantity": 10}],
                    },
                    {
                        "gm_barcode": "GM-222-2",
                        "items": [{"offer_id": "SKU-CLOSED", "barcode": "2043986017738", "quantity": 32}],
                    },
                ],
            },
        )
        boxes = [
            {
                "box_code": "FB-222-1",
                "items": [{"sku_code": "SKU-IN-FULLBOX", "barcode": "2043986017000", "qty": 10}],
            },
            {
                "box_code": "FB-222-2",
                "items": [{"sku_code": "SKU-CLOSED", "barcode": "2043986017738", "qty": 32}],
            },
        ]

        alert = _storekeeper_marketplace_binding_alert(order, boxes=boxes)

        self.assertFalse(alert["show"])
        self.assertFalse(alert["blocking"])
        self.assertTrue(alert["closed_by_discrepancy"])
        self.assertTrue(alert["resolved_with_supplement"])
        self.assertEqual(len(alert["rows"]), 1)

    @patch("shipping.ozon_supplies.fetch_ozon_gm_label_documents")
    def test_manager_downloads_ozon_gm_label_pdf_without_changing_order(self, fetch_labels):
        order = ShippingOrder.objects.create(
            number="SO-OZON-GM-LABEL",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
            supply_number="2000060437104",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            user=self.client_user,
            payload={
                "gm_barcodes": ["1022092968613000"],
                "gm_cargoes": [
                    {
                        "gm_barcode": "1022092968613000",
                        "cargo_id": 7001,
                        "supply_id": 8001,
                        "items": [{"barcode": "2046431529114", "quantity": 20}],
                    }
                ],
            },
        )
        fetch_labels.return_value = (
            [{"supply_id": 8001, "content": b"%PDF-1.4\nmock label\n%%EOF"}],
            "",
            False,
        )
        original_status = order.status

        self.client.force_login(self.manager_user)
        response = self.client.post(reverse("shipping:ozon-gm-labels", args=[order.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("ozon_gm_labels_SO-OZON-GM-LABEL.pdf", response["Content-Disposition"])
        self.assertEqual(response.content, b"%PDF-1.4\nmock label\n%%EOF")
        order.refresh_from_db()
        self.assertEqual(order.status, original_status)
        fetch_labels.assert_called_once()

    @patch("shipping.ozon_supplies.fetch_ozon_gm_label_documents")
    def test_client_cannot_download_ozon_gm_labels(self, fetch_labels):
        order = ShippingOrder.objects.create(
            number="SO-OZON-GM-LABEL-CLIENT",
            agency=self.agency,
            created_by=self.client_user,
            status=ShippingOrder.STATUS_SUBMITTED,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            marketplace=self.market,
        )

        self.client.force_login(self.client_user)
        response = self.client.post(reverse("shipping:ozon-gm-labels", args=[order.pk]))

        self.assertEqual(response.status_code, 403)
        fetch_labels.assert_not_called()

    def test_manager_can_edit_submitted_order_in_create_like_form(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()
        row_key = _shipping_stock_picker_rows(self.agency, exclude_order=order)[0]["key"]
        edit_eta = (timezone.localtime() + timedelta(days=3)).replace(hour=10, minute=0, second=0, microsecond=0)

        self.client.force_login(self.manager_user)
        response = self.client.post(
            f"/shipping/new/?order={order.pk}&edit=1",
            data={
                "edit_order_id": str(order.pk),
                "edit": "1",
                "agency": str(self.agency.id),
                "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
                "slot_date": edit_eta.strftime("%Y-%m-%d"),
                "slot_time": "15:00",
                "eta_at": edit_eta.strftime("%Y-%m-%dT%H:%M"),
                "shipping_barcode": "SHIP-CLIENT-EDIT-1",
                "marketplace": str(self.market.id),
                "wb_supply_barcode": "SUPPLY-EDIT-1",
                "destination_warehouse": "Склад после редактирования",
                "supply_type": ShippingOrder.SUPPLY_BOX,
                "vehicle_type": ShippingOrder.VEHICLE_FULFILLMENT,
                "vehicle_number": "E555EE77",
                "driver_phone": "+7 911 000-00-00",
                "comment": "Менеджер отредактировал заявку",
                "action": "update",
                "stock_key_all[]": [row_key],
                "stock_boxes[]": ["2"],
            },
        )

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        item = order.items.get()
        task = Task.objects.get(route=f"/shipping/{order.pk}/")
        self.assertEqual(order.status, ShippingOrder.STATUS_SUBMITTED)
        # Ozon clears supply barcodes; slot time remains the source of truth.
        self.assertEqual(order.shipping_barcode, "")
        self.assertEqual(order.wb_supply_barcode, "")
        self.assertEqual(str(order.slot_date), edit_eta.date().isoformat())
        self.assertEqual(order.slot_time.strftime("%H:%M"), "15:00")
        self.assertEqual(order.destination_warehouse, "Склад после редактирования")
        self.assertEqual(order.vehicle_number, "E555EE77")
        self.assertEqual(order.expected_boxes, 2)
        self.assertEqual(item.qty_requested, 48)
        self.assertEqual(item.qty_reserved, 48)
        self.assertEqual(task.status, "backlog")

    def test_client_submit_persists_ozon_slot_time(self):
        self.client.force_login(self.client_user)

        response = self.client.post("/shipping/new/", data=self._submit_payload())

        self.assertEqual(response.status_code, 302)
        order = ShippingOrder.objects.get()
        self.assertEqual(order.marketplace, self.market)
        self.assertEqual(order.slot_time.strftime("%H:%M"), "12:30")

    def test_storekeeper_can_save_shipping_packing_act(self):
        order = self._create_order_ready_for_packing(with_services=False)

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/packing/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            f"Паллетизация отгрузки · заявка {_display_shipping_number(order.number)}",
        )
        self.assertContains(response, "BX-SH-1")
        self.assertContains(response, "BX-SH-2")
        self.assertContains(response, 'id="pallet-close-count-modal"', html=False)
        self.assertContains(response, "Сколько коробов в паллете")
        self.assertContains(response, "openPalletCloseConfirm")
        self.assertContains(response, "submittedCount !== expectedCount")
        self.assertNotContains(response, "Фактически оказанные услуги по отгрузке")
        self.assertNotContains(response, "FullboxWarehouseServicesModal")

        response = self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(
                    self._packing_rows_payload()
                ),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )

        response_messages = [str(message) for message in response.wsgi_request._messages]
        self.assertEqual(response.status_code, 302, response_messages)
        self.assertEqual(response.url, f"/shipping/{order.pk}/packing-slips/")
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_PACKED)
        self.assertFalse(
            WarehouseServiceFact.objects.filter(
                client=order.agency,
                order_type=WarehouseServiceFact.ORDER_SHIPPING,
                order_id=order.number,
            ).exists()
        )
        packing_entry = OrderAuditEntry.objects.filter(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            payload__act="shipping_packing",
        ).latest("created_at")
        self.assertEqual(packing_entry.payload["pallet_count"], 2)
        self.assertEqual(packing_entry.payload["delivered_box_count"], 2)
        self.assertEqual(
            [pallet["label"] for pallet in packing_entry.payload["act_pallets"]],
            ["1", "2"],
        )
        self.assertEqual(
            [box["code"] for box in packing_entry.payload["act_boxes"]],
            ["BX-SH-1", "BX-SH-2"],
        )

        detail_response = self.client.get(f"/shipping/{order.pk}/")
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, "Новые паллеты отгрузки")
        self.assertContains(detail_response, "SHIP-")
        self.assertContains(detail_response, "BX-SH-1")

    def test_storekeeper_direct_packing_post_is_blocked_without_service_facts(self):
        order = self._create_order_ready_for_packing(with_services=False)

        self.client.force_login(self.storekeeper_user)
        response = self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(self._packing_rows_payload()),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Перед завершением укажите хотя бы одну фактически оказанную услугу")
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_PICKING)
        self.assertFalse(
            OrderAuditEntry.objects.filter(
                agency=self.agency,
                order_type="shipping",
                order_id=order.number,
                payload__act="shipping_packing",
            ).exists()
        )

    def test_packing_slips_page_contains_preview_print_controls_and_qr(self):
        order = self._create_order_ready_for_packing()
        order.marketplace = self.market
        order.wb_supply_barcode = "SUPPLY-PACK-1"
        order.destination_warehouse = "Новосемейкино"
        order.wb_transit_warehouse = True
        order.transit_address = "Транзитный терминал Голёво"
        order.supply_type = ShippingOrder.SUPPLY_BOX
        order.slot_date = datetime(2026, 4, 7).date()
        order.slot_time = datetime.strptime("08:30", "%H:%M").time()
        order.save(
            update_fields=[
                "marketplace",
                "wb_supply_barcode",
                "destination_warehouse",
                "wb_transit_warehouse",
                "transit_address",
                "supply_type",
                "slot_date",
                "slot_time",
                "updated_at",
            ]
        )

        self.client.force_login(self.storekeeper_user)
        self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(self._packing_rows_payload()),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )

        packing_response = self.client.get(f"/shipping/{order.pk}/packing/")
        self.assertEqual(packing_response.status_code, 403)

        response = self.client.get(f"/shipping/{order.pk}/packing-slips/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Упаковочные листы 75x120")
        self.assertContains(response, "Предпросмотр")
        self.assertContains(response, "Печать всех")
        self.assertContains(response, 'data-label-printer', html=False)
        self.assertContains(response, "Модуль печати")
        self.assertContains(response, "Скачать Fullbox Agent")
        self.assertContains(response, "/static/vendor/qrcode.min.js")
        self.assertContains(response, "/static/vendor/html2canvas.min.js")
        self.assertContains(response, "На каждом листе есть QR-код паллеты.")
        self.assertContains(response, "Статус печати:")
        self.assertNotContains(response, "Автозапуск агента")
        self.assertNotContains(response, "Скрипт агента")
        self.assertNotContains(response, "Скачать агента")
        self.assertNotContains(response, "Синхронизировать")
        self.assertNotContains(response, "Остановить печать")
        self.assertNotContains(response, "Сбросить печать")
        self.assertNotContains(response, "Очистить очередь")
        self.assertEqual(len(response.context["packing_slips_data"]), 2)
        self.assertIn("print_status_line", response.context)
        self.assertIn("available_printers_meta", response.context)
        self.assertEqual(response.context["packing_slips_data"][0]["shipping_barcode"], "SHIP-PACK-1")
        self.assertEqual(response.context["packing_slips_data"][0]["supply_number"], "SUPPLY-PACK-1")
        self.assertEqual(response.context["packing_slips_data"][0]["destination_warehouse"], "Новосемейкино")
        self.assertEqual(response.context["packing_slips_data"][0]["transit_address"], "Транзитный терминал Голёво")
        self.assertEqual(response.context["packing_slips_data"][0]["qr_value"], "SO-000010::PAL-1")

    def test_client_detail_hides_pallet_and_box_breakdown_after_packing(self):
        order = self._create_order_ready_for_packing()

        self.client.force_login(self.storekeeper_user)
        self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(self._packing_rows_payload()),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )

        self.client.force_login(self.client_user)
        response = self.client.get(f"/shipping/{order.pk}/?client={self.agency.id}")

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "Новые паллеты отгрузки")
        self.assertNotContains(response, "BX-SH-1")
        self.assertNotContains(response, "PAL-1")

    def test_client_detail_uses_departed_trip_status_and_vehicle_data(self):
        order = self._create_order_ready_for_packing()
        order.status = ShippingOrder.STATUS_PACKED
        order.vehicle_type = ShippingOrder.VEHICLE_FULFILLMENT
        order.vehicle_number = ""
        order.driver_phone = ""
        order.save(update_fields=["status", "vehicle_type", "vehicle_number", "driver_phone", "updated_at"])

        trip = LogisticsTrip.objects.create(
            number="4_RS",
            status=LogisticsTrip.STATUS_DEPARTED,
            vehicle_type=LogisticsTrip.VEHICLE_FULFILLMENT,
            vehicle_number="Е777ЕЕ77",
            driver_name="Иван Петров",
            driver_phone="+7 911 222-33-44",
            created_by=self.logistician_user,
        )
        LogisticsTripOrder.objects.create(
            trip=trip,
            shipping_order=order,
            loading_sequence=1,
            delivery_sequence=1,
        )

        self.client.force_login(self.client_user)
        response = self.client.get(f"/shipping/{order.pk}/?client={self.agency.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Загружено в машину")
        self.assertContains(response, "Е777ЕЕ77")
        self.assertContains(response, "Иван Петров")
        self.assertContains(response, "+7 911 222-33-44")
        self.assertNotContains(response, "Логист сформировал рейс, ожидается погрузка")
        self.assertNotContains(response, "Открыть рейс 4_RS")
        self.assertContains(response, "Открыть акт отгрузки")
        self.assertContains(response, 'class="summary-table"', html=False)
        self.assertNotContains(response, 'class="chip', html=False)

    def test_storekeeper_can_open_and_save_transport_note(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})
        order.refresh_from_db()
        order.status = ShippingOrder.STATUS_PACKED
        order.vehicle_type = ShippingOrder.VEHICLE_CLIENT
        order.save(update_fields=["status", "vehicle_type", "updated_at"])

        self.client.force_login(self.storekeeper_user)
        response = self.client.get(f"/shipping/{order.pk}/documents/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Сопроводительные документы")
        self.assertNotContains(response, "/transport-note/pdf/")
        self.assertContains(response, "/transport-note/docx/")
        self.assertContains(response, "/transport-note/docx/?inline=1")
        self.assertContains(response, "/return-act/doc/")
        self.assertContains(response, "/return-act/doc/?inline=1")
        self.assertNotContains(response, "/transport-note/docx-preview/")
        self.assertContains(response, "Открыть DOCX")
        self.assertContains(response, "Скачать DOCX")
        self.assertContains(response, "Акт возврата")
        self.assertContains(response, "Открыть DOCX")
        self.assertContains(response, "Скачать DOCX")

        response = self.client.post(
            f"/shipping/{order.pk}/transport-note/",
            data={
                "document_number": "ТН-2026-001",
                "document_date": "2026-04-12",
                "shipper_name": "ООО Кейз",
                "shipper_inn": "7701234567",
                "shipper_address": "Москва, ул. Тестовая, 1",
                "shipper_phone": "+7 900 111-22-33",
                "consignee_name": "WB / Склад маркетплейса",
                "consignee_inn": "",
                "consignee_address": "МО, склад назначения",
                "consignee_phone": "",
                "carrier_name": "FullBox",
                "carrier_inn": "",
                "carrier_address": "Старая Купавна",
                "carrier_phone": "+7 499 450-35-55",
                "loading_address": "Склад FullBox",
                "unloading_address": "Склад WB",
                "cargo_name": "Тестовый груз",
                "cargo_package_count": "3",
                "cargo_package_type": "Короб",
                "cargo_weight_kg": "25.500",
                "cargo_declared_value": "150000.00",
                "accompanying_documents": "Заявка и УПД",
                "special_instructions": "Без повреждений",
                "transportation_conditions": "Доставка до склада",
                "delivery_notes": "Без замечаний",
                "driver_name": "Иванов И.И.",
                "driver_phone": "+7 900 123-45-67",
                "vehicle_number": "A123BC77",
                "trailer_number": "TR-01",
                "service_cost": "5000.00",
            },
        )

        self.assertEqual(response.status_code, 302)
        note = ShippingTransportNote.objects.get(order=order)
        self.assertEqual(note.document_number, "ТН-2026-001")
        self.assertEqual(note.cargo_package_count, 3)
        self.assertEqual(note.driver_name, "Иванов И.И.")

        pdf_response = self.client.get(f"/shipping/{order.pk}/transport-note/pdf/")
        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response["Content-Type"], "application/pdf")
        self.assertIn(b"%PDF", pdf_response.content[:8])

        docx_response = self.client.get(f"/shipping/{order.pk}/transport-note/docx/")
        self.assertEqual(docx_response.status_code, 200)
        self.assertEqual(
            docx_response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn(b"PK", docx_response.content[:4])
        self.assertIn("attachment", docx_response["Content-Disposition"].lower())

        inline_response = self.client.get(f"/shipping/{order.pk}/transport-note/docx/?inline=1")
        self.assertEqual(inline_response.status_code, 200)
        self.assertEqual(
            inline_response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("inline", inline_response["Content-Disposition"].lower())
        self.assertIn(b"PK", inline_response.content[:4])

        return_act_item = order.items.order_by("id").first()
        return_act_item.comment = "Коробов: 1; кратность: 50; короба: OSN-0807-000298-no"
        return_act_item.save(update_fields=["comment"])
        return_act_response = self.client.get(f"/shipping/{order.pk}/return-act/doc/")
        self.assertEqual(return_act_response.status_code, 200)
        self.assertEqual(
            return_act_response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("attachment", return_act_response["Content-Disposition"].lower())
        self.assertIn(b"PK", return_act_response.content[:4])
        with zipfile.ZipFile(BytesIO(return_act_response.content)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn(order.number, document_xml)
        self.assertIn(order.agency.agn_name, document_xml)
        self.assertIn("Товар отгрузки", document_xml)
        self.assertEqual(document_xml.count("Товар отгрузки"), 1)
        self.assertNotIn("Коробов: 1", document_xml)
        self.assertNotIn("кратность: 50", document_xml)
        self.assertNotIn("OSN-0807-000298-no", document_xml)
        self.assertEqual(len(Document(BytesIO(return_act_response.content)).tables[0].rows), 93)

        return_act_inline = self.client.get(f"/shipping/{order.pk}/return-act/doc/?inline=1")
        self.assertEqual(return_act_inline.status_code, 200)
        self.assertEqual(
            return_act_inline["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("inline", return_act_inline["Content-Disposition"].lower())
        self.assertIn(b"PK", return_act_inline.content[:4])

    def test_transport_note_is_forbidden_before_shipping_ready(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})

        self.client.force_login(self.storekeeper_user)
        detail_response = self.client.get(f"/shipping/{order.pk}/")
        documents_response = self.client.get(f"/shipping/{order.pk}/documents/")
        transport_note_response = self.client.get(f"/shipping/{order.pk}/transport-note/")
        docx_response = self.client.get(f"/shipping/{order.pk}/transport-note/docx/")
        return_act_response = self.client.get(f"/shipping/{order.pk}/return-act/doc/")

        self.assertEqual(detail_response.status_code, 200)
        self.assertNotContains(detail_response, "Сопроводительные документы")
        self.assertEqual(documents_response.status_code, 403)
        self.assertEqual(transport_note_response.status_code, 403)
        self.assertEqual(docx_response.status_code, 403)
        self.assertEqual(return_act_response.status_code, 403)

    def test_transport_note_uses_directories_for_shipper_carrier_customer_and_consignee(self):
        OwnCompany.objects.create(
            name='Общество с ограниченной ответственностью "ФуллБокс"',
            inn="5001149130",
            address="Московская область, Старая Купавна, Магистральная, 59",
            postal_address="Склад FullBox, Московская область, Старая Купавна, Магистральная, 59",
            phone="+7 977 808-83-86",
            is_default=True,
            is_active=True,
        )
        Carrier.objects.create(
            name="Индивидуальный предприниматель Касаев Абдурашид Метханович",
            inn="052903630023",
            address="Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.",
            postal_address="Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.",
            phone="+7 (927) 157-22-42",
            is_active=True,
        )
        self.agency.agn_name = 'Общество с ограниченной ответственностью "Кейзи"'
        self.agency.save(update_fields=["agn_name", "short_name"])
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["destination_warehouse"] = "Санкт_Петербург_РФЦ_1"
        with patch("shipping.transport_note.load_marketplace_warehouse_catalog") as catalog_mock:
            catalog_mock.return_value = {
                "wb": [],
                "ozon": [
                    {
                        "type": "Фулфилмент",
                        "name": "Санкт_Петербург_РФЦ_1",
                        "address": "Санкт-Петербург, склад ОЗОН",
                    }
                ],
                "yandex": [],
                "sber": [],
            }
            self.client.post("/shipping/new/", data=payload)
        order = ShippingOrder.objects.get()

        self.client.force_login(self.manager_user)
        self.client.post(f"/shipping/{order.pk}/", data={"action": "reserve"})
        self._mark_transport_note_ready(order, vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT)

        self.client.force_login(self.storekeeper_user)
        with patch("shipping.transport_note.load_marketplace_warehouse_catalog") as catalog_mock:
            catalog_mock.return_value = {
                "wb": [],
                "ozon": [
                    {
                        "type": "Фулфилмент",
                        "name": "Санкт_Петербург_РФЦ_1",
                        "address": "Санкт-Петербург, склад ОЗОН",
                    }
                ],
                "yandex": [],
                "sber": [],
            }
            response = self.client.get(f"/shipping/{order.pk}/transport-note/")

        note = ShippingTransportNote.objects.get(order=order)
        self.assertEqual(note.shipper_name, 'ООО "ФуллБокс"')
        self.assertEqual(note.shipper_inn, "5001149130")
        self.assertEqual(note.shipper_address, "Склад FullBox, Московская область, Старая Купавна, Магистральная, 59")
        self.assertEqual(note.shipper_phone, "+7 977 808-83-86")
        self.assertEqual(note.carrier_name, "ИП Касаев Абдурашид Метханович")
        self.assertEqual(note.carrier_inn, "052903630023")
        self.assertEqual(note.carrier_address, "Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.")
        self.assertEqual(note.carrier_phone, "+7 (927) 157-22-42")
        self.assertEqual(note.loading_address, "Склад FullBox, Московская область, Старая Купавна, Магистральная, 59")
        self.assertEqual(note.consignee_name, "Ozon / Санкт_Петербург_РФЦ_1")
        self.assertContains(response, "ООО &quot;Кейзи&quot;")
        self.assertContains(response, "ООО &quot;ФуллБокс&quot;")
        self.assertContains(response, "ИП Касаев Абдурашид Метханович")
        self.assertContains(response, "Ozon / Санкт_Петербург_РФЦ_1")

    def test_transport_note_refreshes_legacy_carrier_requisites_from_directory(self):
        OwnCompany.objects.create(
            name='Общество с ограниченной ответственностью "ФуллБокс"',
            inn="5001149130",
            postal_address="Склад FullBox, Московская область, Старая Купавна, Магистральная, 59",
            phone="+7 977 808-83-86",
            is_default=True,
            is_active=True,
        )
        Carrier.objects.create(
            name="Индивидуальный предприниматель Касаев Абдурашид Метханович",
            inn="052903630023",
            address="Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.",
            postal_address="Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.",
            phone="+7 (927) 157-22-42",
            is_active=True,
        )
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()
        self._mark_transport_note_ready(order, vehicle_type=ShippingOrder.VEHICLE_FULFILLMENT)
        ShippingTransportNote.objects.update_or_create(
            order=order,
            defaults={
                "document_number": order.number,
                "shipper_name": 'ООО "ФуллБокс"',
                "shipper_inn": "5001149130",
                "shipper_address": "Склад FullBox, Московская область, Старая Купавна, Магистральная, 59",
                "shipper_phone": "+7 977 808-83-86",
                "carrier_name": "FullBox",
                "carrier_inn": "5001149130",
                "carrier_address": "Склад FullBox, Московская область, Старая Купавна, Магистральная, 59",
                "carrier_phone": "+7 977 808-83-86",
            },
        )

        self.client.force_login(self.storekeeper_user)
        self.client.get(f"/shipping/{order.pk}/transport-note/")

        note = ShippingTransportNote.objects.get(order=order)
        self.assertEqual(note.carrier_name, "ИП Касаев Абдурашид Метханович")
        self.assertEqual(note.carrier_inn, "052903630023")
        self.assertEqual(note.carrier_address, "Дагестан Респ, Сулейман-Стальский м.р-н, село Орта-Стал с.п.")
        self.assertEqual(note.carrier_phone, "+7 (927) 157-22-42")

    def test_transport_note_switches_to_transit_address_when_order_gets_transit_route(self):
        self.client.force_login(self.client_user)
        payload = self._submit_payload()
        payload["transit_address"] = ""
        payload["wb_transit_warehouse"] = ""
        self.client.post("/shipping/new/", data=payload)
        order = ShippingOrder.objects.get()
        self._mark_transport_note_ready(order)

        self.client.force_login(self.storekeeper_user)
        self.client.get(f"/shipping/{order.pk}/transport-note/")

        note = ShippingTransportNote.objects.get(order=order)
        initial_address = note.consignee_address
        self.assertTrue(initial_address)
        self.assertEqual(note.unloading_address, initial_address)

        order.wb_transit_warehouse = True
        order.transit_address = "Москва, Транзитный терминал 1"
        order.save(update_fields=["wb_transit_warehouse", "transit_address", "updated_at"])

        self.client.get(f"/shipping/{order.pk}/transport-note/")

        note.refresh_from_db()
        self.assertEqual(note.consignee_address, "Москва, Транзитный терминал 1")
        self.assertEqual(note.unloading_address, "Москва, Транзитный терминал 1")

    def test_client_cannot_access_transport_note(self):
        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/", data=self._submit_payload())
        order = ShippingOrder.objects.get()

        response = self.client.get(f"/shipping/{order.pk}/transport-note/")

        self.assertEqual(response.status_code, 403)

        pdf_response = self.client.get(f"/shipping/{order.pk}/transport-note/pdf/")
        self.assertEqual(pdf_response.status_code, 403)

        docx_response = self.client.get(f"/shipping/{order.pk}/transport-note/docx/")
        self.assertEqual(docx_response.status_code, 403)
        inline_response = self.client.get(f"/shipping/{order.pk}/transport-note/docx/?inline=1")
        self.assertEqual(inline_response.status_code, 403)

    def test_shipping_packing_transfers_order_from_storekeeper_to_logistician(self):
        order = self._create_order_ready_for_packing()
        storekeeper_task = Task.objects.create(
            title=f"Заявка на отгрузку №{order.number}",
            route=f"/shipping/{order.pk}/",
            assigned_to=self.storekeeper,
            status="in_progress",
        )

        self.client.force_login(self.storekeeper_user)
        response = self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(
                    self._packing_rows_payload()
                ),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )

        self.assertEqual(response.status_code, 302)
        storekeeper_task.refresh_from_db()
        logistician_task = Task.objects.get(route=f"/shipping/{order.pk}/", assigned_to=self.logistician)
        self.assertEqual(storekeeper_task.status, "done")
        self.assertEqual(logistician_task.status, "backlog")
        self.assertIn("требуется погрузка и акт отгрузки", logistician_task.description.lower())

        detail_response = self.client.get(f"/shipping/{order.pk}/")
        self.assertNotContains(detail_response, "Фиксация отгрузки")
        self.assertContains(detail_response, "Логистика и закрытие заявки")

    def test_fulfillment_transport_dispatch_act_requires_trip(self):
        order = self._create_order_ready_for_packing()
        order.vehicle_type = ShippingOrder.VEHICLE_FULFILLMENT
        order.status = ShippingOrder.STATUS_PACKED
        order.save(update_fields=["vehicle_type", "status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 2,
                "delivered_box_count": 2,
                "act_pallets": [
                    {"label": "1", "code": "1", "boxes": ["BX-SH-1"]},
                    {"label": "2", "code": "2", "boxes": ["BX-SH-2"]},
                ],
                "act_boxes": [
                    {"code": "BX-SH-1", "qty": 24, "pallet_label": "1", "items": []},
                    {"code": "BX-SH-2", "qty": 24, "pallet_label": "2", "items": []},
                ],
            },
        )

        self.client.force_login(self.logistician_user)
        response = self.client.get(f"/shipping/{order.pk}/act/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "без рейса логист не может подписать акт")

    def test_shipping_draft_trip_does_not_count_as_formed_trip(self):
        order = self._create_order_ready_for_packing()
        order.vehicle_type = ShippingOrder.VEHICLE_FULFILLMENT
        order.status = ShippingOrder.STATUS_PACKED
        order.save(update_fields=["vehicle_type", "status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 1,
                "delivered_box_count": 1,
            },
        )
        trip = LogisticsTrip.objects.create(
            number="DRAFT-TEST-0001",
            status=LogisticsTrip.STATUS_DRAFT,
            created_by=self.manager_user,
        )
        LogisticsTripOrder.objects.create(
            trip=trip,
            shipping_order=order,
            loading_sequence=1,
            delivery_sequence=1,
        )

        self.client.force_login(self.logistician_user)
        response = self.client.get(f"/shipping/{order.pk}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Нужен рейс логиста")
        self.assertNotContains(response, "Логист сформировал рейс, ожидается погрузка")
        self.assertNotContains(response, "Логист подписывает акт")

    def test_dispatch_act_cannot_be_signed_before_confirmed_loading(self):
        order = self._create_order_ready_for_packing()
        order.vehicle_type = ShippingOrder.VEHICLE_CLIENT
        order.status = ShippingOrder.STATUS_PACKED
        order.save(update_fields=["vehicle_type", "status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 2,
                "delivered_box_count": 2,
                "act_pallets": [
                    {"label": "1", "code": "1", "boxes": ["BX-SH-1"]},
                    {"label": "2", "code": "2", "boxes": ["BX-SH-2"]},
                ],
                "act_boxes": [
                    {"code": "BX-SH-1", "qty": 24, "pallet_label": "1", "items": []},
                    {"code": "BX-SH-2", "qty": 24, "pallet_label": "2", "items": []},
                ],
            },
        )
        Task.objects.create(
            title=f"Заявка на отгрузку №{order.number}",
            route=f"/shipping/{order.pk}/",
            assigned_to=self.logistician,
            status="backlog",
        )

        self.client.force_login(self.logistician_user)
        response = self.client.post(f"/shipping/{order.pk}/act/sign-logistician/")

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_PACKED)
        self.assertFalse(
            Task.objects.filter(route=f"/shipping/{order.pk}/act/", assigned_to=self.manager).exists()
        )
        self.assertFalse(OrderAuditEntry.objects.filter(
            order_type="shipping",
            order_id=order.number,
            payload__act="shipping_dispatch_act",
        ).exists())

    def test_confirmed_dispatch_act_signatures_do_not_repeat_stock_shipping(self):
        order = self._create_order_ready_for_packing()
        order.vehicle_type = ShippingOrder.VEHICLE_CLIENT
        order.status = ShippingOrder.STATUS_SHIPPED
        order.shipped_at = timezone.now()
        order.save(update_fields=["vehicle_type", "status", "shipped_at", "updated_at"])
        item = order.items.get()
        item.qty_shipped = item.qty_requested
        item.qty_reserved = 0
        item.save(update_fields=["qty_shipped", "qty_reserved", "updated_at"])
        confirm_dispatch_act_from_loading(order, self.storekeeper_user)

        act_task = Task.objects.get(
            route=f"/shipping/{order.pk}/act/",
            assigned_to=self.logistician,
        )
        self.assertEqual(act_task.status, "backlog")

        self.client.force_login(self.logistician_user)
        response = self.client.post(f"/shipping/{order.pk}/act/sign-logistician/")
        self.assertEqual(response.status_code, 302)
        act_task.refresh_from_db()
        self.assertEqual(act_task.status, "done")
        manager_sign_task = (
            Task.objects.filter(route=f"/shipping/{order.pk}/act/")
            .exclude(assigned_to__role="logistician")
            .latest("created_at")
        )
        self.assertEqual(manager_sign_task.status, "backlog")

        self.client.force_login(self.manager_user)
        response = self.client.post(f"/shipping/{order.pk}/act/sign-manager/")

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        item = order.items.get()
        manager_sign_task.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(item.qty_shipped, 48)
        self.assertEqual(manager_sign_task.status, "done")
        dispatch_entry = OrderAuditEntry.objects.filter(
            order_type="shipping",
            order_id=order.number,
            payload__act="shipping_dispatch_act",
        ).latest("created_at")
        self.assertTrue(dispatch_entry.payload["act_manager_signed"])
        self.assertTrue(dispatch_entry.payload["act_sent"])
        self.assertTrue(dispatch_entry.payload["act_warehouse_confirmed"])

    def test_detail_does_not_count_packed_pallets_as_shipped(self):
        order = self._create_order_ready_for_packing()
        order.supply_type = ShippingOrder.SUPPLY_BOX
        order.status = ShippingOrder.STATUS_PACKED
        order.save(update_fields=["supply_type", "status", "updated_at"])
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 2,
                "delivered_box_count": 2,
                "act_pallets": [
                    {"label": "1", "code": "1", "boxes": ["BX-SH-1"]},
                    {"label": "2", "code": "2", "boxes": ["BX-SH-2"]},
                ],
                "act_boxes": [
                    {"code": "BX-SH-1", "qty": 24, "pallet_label": "1", "items": []},
                    {"code": "BX-SH-2", "qty": 24, "pallet_label": "2", "items": []},
                ],
            },
        )

        self.client.force_login(self.manager_user)
        detail_response = self.client.get(f"/shipping/{order.pk}/")

        self.assertContains(detail_response, "Отгружено паллет")
        self.assertContains(detail_response, 'id="shipped-pallet-count">0</td>', html=False)

    def test_storekeeper_can_reopen_saved_packing_when_otg_payload_lost(self):
        order = self._create_order_ready_for_packing()

        self.client.force_login(self.storekeeper_user)
        self.client.post(
            f"/shipping/{order.pk}/packing/",
            data={
                "boxes_json": json.dumps(
                    self._packing_rows_payload()
                ),
                "pallets_json": json.dumps(
                    [
                        {"code": "PAL-1", "label": "1"},
                        {"code": "PAL-2", "label": "2"},
                    ]
                ),
            },
        )
        move_task = MoveTask.objects.filter(
            payload__shipping_order_pk=order.pk,
            to_zone="OTG",
        ).first()
        self.assertIsNotNone(move_task)
        move_task.payload = {
            "shipping_order_id": order.number,
            "shipping_order_pk": order.pk,
            "receiving_order_id": "R-SH-1",
        }
        move_task.save(update_fields=["payload"])

        response = self.client.get(f"/shipping/{order.pk}/packing/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "BX-SH-1")
        self.assertContains(response, "BX-SH-2")

    def test_shipping_boxes_from_packing_payload_rebuilds_box_rows(self):
        rows = _shipping_boxes_from_packing_payload(
            {
                "act_boxes": [
                    {
                        "code": "BX-1",
                        "qty": 12,
                        "barcode_preview": "300000000001 - 12 шт.",
                        "items": [
                            {
                                "sku_code": "SKU-1",
                                "name": "Товар 1",
                                "size": "42",
                                "barcode": "300000000001",
                                "goods_type": "Готовый",
                                "qty": 12,
                            }
                        ],
                    }
                ]
            }
        )
        self.assertEqual(
            rows,
            [
                {
                    "row_key": _shipping_box_row_key({"code": "BX-1"}, 1),
                    "box_code": "BX-1",
                    "qty": 12,
                    "items": [
                        {
                            "sku_code": "SKU-1",
                            "name": "Товар 1",
                            "size": "42",
                            "barcode": "300000000001",
                            "goods_type": "Готовый",
                            "qty": 12,
                        }
                    ],
                    "barcode_preview": "300000000001 - 12 шт.",
                    "receiving_order_id": "",
                }
            ],
        )

    def test_shipping_boxes_from_packing_payload_keeps_rows_with_same_code_when_row_keys_differ(self):
        rows = _shipping_boxes_from_packing_payload(
            {
                "act_boxes": [
                    {
                        "row_key": "row-1",
                        "code": "BX-1",
                        "qty": 5,
                        "barcode_preview": "300000000001 - 5 шт.",
                        "items": [],
                    },
                    {
                        "row_key": "row-2",
                        "code": "BX-1",
                        "qty": 7,
                        "barcode_preview": "300000000001 - 7 шт.",
                        "items": [],
                    },
                ]
            }
        )

        self.assertEqual([row["row_key"] for row in rows], ["row-1", "row-2"])
        self.assertEqual([row["qty"] for row in rows], [5, 7])

    def test_shipping_packing_summary_ignores_zero_qty_boxes_from_payload(self):
        order = self._create_order_ready_for_packing()
        OrderAuditEntry.objects.create(
            agency=self.agency,
            order_type="shipping",
            order_id=order.number,
            action="status",
            user=self.storekeeper_user,
            description="Паллетизация завершена",
            payload={
                "act": "shipping_packing",
                "act_state": "closed",
                "pallet_count": 1,
                "delivered_box_count": 3,
                "act_pallets": [
                    {"label": "1", "code": "PAL-1", "boxes": ["BX-REAL-1", "BX-EMPTY-2", "BX-EMPTY-3"]},
                ],
                "act_boxes": [
                    {"code": "BX-REAL-1", "qty": 20, "barcode_preview": "2039754223166 - 20 шт.", "pallet_label": "1", "items": []},
                    {"code": "BX-EMPTY-2", "qty": 0, "barcode_preview": "-", "pallet_label": "1", "items": []},
                    {"code": "BX-EMPTY-3", "qty": 0, "barcode_preview": "-", "pallet_label": "1", "items": []},
                ],
            },
        )

        summary = _shipping_packing_summary(order)

        self.assertIsNotNone(summary)
        self.assertEqual(summary["pallet_count"], 1)
        self.assertEqual(summary["box_count"], 1)
        self.assertEqual(summary["pallets"][0]["box_count"], 1)
        self.assertEqual(summary["pallets"][0]["qty"], 20)
        self.assertEqual([box["box_code"] for box in summary["pallets"][0]["boxes"]], ["BX-REAL-1"])


class ShippingAttachmentFlowTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.mkdtemp(prefix="shipping-attachments-")
        self.media_override = override_settings(MEDIA_ROOT=self.temp_media)
        self.media_override.enable()

        user_model = get_user_model()
        self.client_user = user_model.objects.create_user(username="client_attach_shipping", password="pwd")
        self.manager_user = user_model.objects.create_user(username="manager_attach_shipping", password="pwd")
        self.other_client_user = user_model.objects.create_user(username="other_client_attach_shipping", password="pwd")

        Employee.objects.create(
            full_name="Менеджер вложений",
            role="manager",
            user=self.manager_user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Клиент с файлами", portal_user=self.client_user)
        Agency.objects.create(agn_name="Чужой клиент", portal_user=self.other_client_user)
        self.market = Market.objects.create(id=302, name="WB Attach")
        create_warehouse_snapshot_row(
            agency=self.agency,
            order_type="receiving",
            order_id="R-ATT-1",
            sku="SKU-ATT-1",
            name="Товар с файлами",
            size="44",
            barcode="355500000001",
            goods_type="Готовый",
            qty=24,
            box_code="BX-ATT-1",
            pallet_code="PL-ATT-1",
            zone="OS",
            row=1,
            section=1,
            tier=1,
            cell=1,
        )

    def tearDown(self):
        self.media_override.disable()
        shutil.rmtree(self.temp_media, ignore_errors=True)
        super().tearDown()

    def _submit_payload(self):
        picker_key = _shipping_stock_picker_rows(self.agency)[0]["key"]
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        return {
            "delivery_type": ShippingOrder.DELIVERY_MARKETPLACE,
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "slot_date": eta.strftime("%Y-%m-%d"),
            "slot_time": "12:30",
            "shipping_barcode": "SHIP-ATT-1",
            "marketplace": str(self.market.id),
            "wb_supply_barcode": "SUP-ATT-1",
            "destination_warehouse": "Склад с файлами",
            "supply_type": ShippingOrder.SUPPLY_BOX,
            "vehicle_type": ShippingOrder.VEHICLE_FULFILLMENT,
            "vehicle_number": "A321BC77",
            "driver_phone": "+7 900 123-45-67",
            "comment": "Файлы приложены",
            "action": "submit",
            "stock_key_all[]": [picker_key],
            "stock_boxes[]": ["1"],
        }

    def test_client_can_upload_attachment_on_create_and_manager_can_download_it(self):
        upload = SimpleUploadedFile("packing-note.txt", b"packing note", content_type="text/plain")
        payload = self._submit_payload()
        payload["documents"] = upload

        self.client.force_login(self.client_user)
        response = self.client.post("/shipping/new/?client=%s" % self.agency.id, data=payload)

        self.assertEqual(response.status_code, 302)
        order = ShippingOrder.objects.get()
        attachment = ShippingOrderAttachment.objects.get(order=order)
        self.assertEqual(attachment.filename, "packing-note.txt")

        self.client.force_login(self.manager_user)
        detail_response = self.client.get(f"/shipping/{order.pk}/")
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, "Вложения")
        self.assertContains(detail_response, "packing-note.txt")
        self.assertContains(detail_response, "Срок хранения на сервере: 60 дней")

        download_response = self.client.get(
            reverse("shipping:attachment-download", args=[order.pk, attachment.pk])
        )
        self.assertEqual(download_response.status_code, 200)
        self.assertIn("attachment;", download_response["Content-Disposition"])
        self.assertIn("packing-note.txt", download_response["Content-Disposition"])
        self.assertEqual(b"".join(download_response.streaming_content), b"packing note")

    def test_attachment_download_is_forbidden_for_other_client(self):
        upload = SimpleUploadedFile("invoice.pdf", b"%PDF-test", content_type="application/pdf")
        payload = self._submit_payload()
        payload["documents"] = upload

        self.client.force_login(self.client_user)
        self.client.post("/shipping/new/?client=%s" % self.agency.id, data=payload)
        order = ShippingOrder.objects.get()
        attachment = ShippingOrderAttachment.objects.get(order=order)

        self.client.force_login(self.other_client_user)
        response = self.client.get(
            reverse("shipping:attachment-download", args=[order.pk, attachment.pk])
        )

        self.assertEqual(response.status_code, 403)

    def test_create_form_handles_attachment_save_errors_without_500(self):
        payload = self._submit_payload()
        payload["documents"] = SimpleUploadedFile("broken.txt", b"oops", content_type="text/plain")

        self.client.force_login(self.client_user)
        with patch("shipping.views._save_shipping_attachments", side_effect=OSError("disk error")):
            response = self.client.post(f"/shipping/new/?client={self.agency.id}", data=payload)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Не удалось сохранить заявку: disk error")

    def test_client_attachment_runtime_keeps_file_selection_out_of_autosave(self):
        self.client.force_login(self.client_user)

        response = self.client.get(f"/shipping/new/?client={self.agency.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "target === documentsInput")
        self.assertContains(response, "target.closest('#files-section')")
        self.assertContains(response, "attachmentUploadInProgress || !draftAutosaveTouched")
        self.assertContains(response, "event.stopPropagation()")
        self.assertContains(response, "if (draftSavePromise)")
        self.assertContains(response, "response.status >= 400 && response.status < 500")
        self.assertContains(response, "draftAutosaveTouched = hasNewRevision")
        self.assertContains(response, "Date.now() < draftAutosaveRetryAt")
        self.assertContains(response, "120000")

    def test_client_can_upload_attachment_to_autosaved_draft(self):
        self.client.force_login(self.client_user)
        draft_payload = self._submit_payload()
        draft_payload["action"] = "save_draft"
        draft_payload["draft_autosave"] = "1"
        draft_payload.pop("stock_key_all[]", None)
        draft_payload.pop("stock_boxes[]", None)
        draft_response = self.client.post(
            f"/shipping/new/?client={self.agency.id}",
            data=draft_payload,
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(draft_response.status_code, 200, draft_response.content)
        order_id = draft_response.json()["draft_order_id"]

        upload_response = self.client.post(
            reverse("shipping:attachment-upload", args=[order_id]),
            data={
                "upload_id": "client-ozon-attachment-test",
                "documents": SimpleUploadedFile(
                    "ozon-label.pdf",
                    b"%PDF-1.4 test",
                    content_type="application/pdf",
                ),
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(upload_response.status_code, 200, upload_response.content)
        self.assertTrue(upload_response.json()["ok"])
        attachments = list(ShippingOrderAttachment.objects.filter(order_id=order_id))
        self.assertEqual([attachment.filename for attachment in attachments], ["ozon-label.pdf"])

    def test_create_form_hides_technical_reachtruck_error_from_client(self):
        payload = self._submit_payload()

        self.client.force_login(self.client_user)
        with patch(
            "shipping.services.apply_client_shipping_reserve_on_submit",
            side_effect=NameError("name 'exclude_shipping_order_id' is not defined"),
        ):
            response = self.client.post(f"/shipping/new/?client={self.agency.id}", data=payload)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Заявка уже передана в складскую работу.")
        self.assertContains(response, "Редактирование из ЛК клиента или менеджера недоступно.")
        self.assertNotContains(response, "exclude_shipping_order_id")

    def test_late_autosave_does_not_return_submitted_order_to_draft(self):
        draft_payload = self._submit_payload()
        draft_payload["action"] = "save_draft"
        draft_payload["draft_autosave"] = "1"

        self.client.force_login(self.client_user)
        draft_response = self.client.post(
            f"/shipping/new/?client={self.agency.id}",
            data=draft_payload,
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(draft_response.status_code, 200, draft_response.content)
        order_id = draft_response.json()["order_id"]

        stale_autosave_payload = self._submit_payload()
        stale_autosave_payload["action"] = "save_draft"
        stale_autosave_payload["draft_autosave"] = "1"
        stale_autosave_payload["edit_order_id"] = str(order_id)

        selected_rows = [
            {
                "sku_code": "SKU-ATT-1",
                "name": "Товар с файлами",
                "size": "44",
                "barcode": "355500000001",
                "goods_type": "Готовый",
                "qty_requested": 1,
                "comment": "",
            }
        ]

        def mark_order_submitted_before_save(*_args, **_kwargs):
            ShippingOrder.objects.filter(pk=order_id).update(status=ShippingOrder.STATUS_SUBMITTED)
            return selected_rows, []

        with patch("shipping.views._parse_selected_stock_items", side_effect=mark_order_submitted_before_save):
            response = self.client.post(
                f"/shipping/new/?client={self.agency.id}&order={order_id}&edit=1",
                data=stale_autosave_payload,
                HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            )

        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn("Автосохранение черновика остановлено", response.json()["error"])
        order = ShippingOrder.objects.get(pk=order_id)
        self.assertEqual(order.status, ShippingOrder.STATUS_SUBMITTED)

    def test_manager_can_delete_existing_attachment_upload_new_one_and_change_boxes_on_edit(self):
        create_warehouse_snapshot_row(
            agency=self.agency,
            order_type="receiving",
            order_id="R-ATT-2",
            sku="SKU-ATT-1",
            name="Товар с файлами",
            size="44",
            barcode="355500000001",
            goods_type="Готовый",
            qty=24,
            box_code="BX-ATT-2",
            pallet_code="PL-ATT-1",
            zone="OS",
            row=1,
            section=1,
            tier=1,
            cell=2,
        )
        payload = self._submit_payload()
        payload["stock_boxes[]"] = ["2"]
        payload["documents"] = SimpleUploadedFile("remove-me.txt", b"old", content_type="text/plain")

        self.client.force_login(self.client_user)
        response = self.client.post(f"/shipping/new/?client={self.agency.id}", data=payload)
        self.assertEqual(response.status_code, 302)
        order = ShippingOrder.objects.get()
        delete_attachment = ShippingOrderAttachment.objects.get(order=order, filename="remove-me.txt")
        keep_attachment = ShippingOrderAttachment.objects.create(
            order=order,
            uploaded_by=self.client_user,
            file=SimpleUploadedFile("keep-me.txt", b"keep", content_type="text/plain"),
        )

        row_key = _shipping_stock_picker_rows(self.agency, exclude_order=order)[0]["key"]
        edit_payload = {key: value for key, value in payload.items() if key != "documents"}
        edit_payload.update(
            {
                "edit_order_id": str(order.pk),
                "edit": "1",
                "agency": str(self.agency.id),
                "action": "update",
                "stock_key_all[]": [row_key],
                "stock_boxes[]": ["1"],
                "delete_attachment_ids": [str(delete_attachment.pk)],
                "documents": SimpleUploadedFile("new-file.txt", b"new", content_type="text/plain"),
            }
        )

        self.client.force_login(self.manager_user)
        response = self.client.post(f"/shipping/new/?order={order.pk}&edit=1", data=edit_payload)

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.expected_boxes, 1)
        self.assertEqual(order.items.count(), 1)
        self.assertEqual(order.items.get().qty_requested, 24)
        self.assertFalse(ShippingOrderAttachment.objects.filter(pk=delete_attachment.pk).exists())
        self.assertTrue(ShippingOrderAttachment.objects.filter(pk=keep_attachment.pk).exists())
        self.assertTrue(ShippingOrderAttachment.objects.filter(order=order, filename="new-file.txt").exists())


class ShippingClientPermissionsTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Клиент доступа")
        self.order = ShippingOrder.objects.create(
            number="SO-009999",
            agency=self.agency,
            status=ShippingOrder.STATUS_SUBMITTED,
        )

    def test_client_can_edit_submitted_order_before_manager_approval(self):
        self.assertTrue(_can_edit_items("client", "client", self.order))
        self.assertTrue(_can_edit_order_form("client", "client", self.order))

    def test_client_cannot_edit_reserved_order(self):
        self.order.status = ShippingOrder.STATUS_RESERVED
        self.order.save(update_fields=["status"])
        self.assertFalse(_can_edit_items("client", "client", self.order))
        self.assertFalse(_can_edit_order_form("client", "client", self.order))

    def test_client_can_cancel_submitted_order_before_manager_approval(self):
        self.assertTrue(_can_cancel("client", "client", self.order))


class ShippingWarehouseCancelApprovalTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.manager_user = user_model.objects.create_user(username="manager_cancel_review", password="pwd")
        self.storekeeper_user = user_model.objects.create_user(username="storekeeper_cancel_review", password="pwd")
        self.manager = Employee.objects.create(
            full_name="Менеджер отмены",
            role="manager",
            user=self.manager_user,
            is_active=True,
        )
        self.storekeeper = Employee.objects.create(
            full_name="Кладовщик отмены",
            role="storekeeper",
            user=self.storekeeper_user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Клиент отмены склада")
        self.order = ShippingOrder.objects.create(
            number="OTG-CANCEL-REVIEW-1",
            agency=self.agency,
            created_by=self.manager_user,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
        )

    def test_manager_requests_warehouse_confirmation_before_cancel(self):
        self.assertFalse(_can_cancel("staff", "manager", self.order))

        self.client.force_login(self.manager_user)
        response = self.client.post(
            reverse("shipping:detail", args=[self.order.pk]),
            {"action": "request_warehouse_cancel", "cancel_reason": "Клиент отменил поставку"},
        )

        self.assertEqual(response.status_code, 302)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_STOREKEEPER_ACCEPTED)
        latest = OrderAuditEntry.objects.filter(order_type="shipping", order_id=self.order.number).latest("id")
        self.assertEqual((latest.payload or {}).get("status"), "warehouse_cancel_requested")
        self.assertTrue(
            Task.objects.filter(
                route=f"/shipping/{self.order.pk}/",
                assigned_to=self.storekeeper,
                title__startswith="Подтвердите отмену отгрузки",
            ).exists()
        )

        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=self.order.number,
            action="status",
            description="Резерв снят",
            payload={},
            user=self.storekeeper_user,
        )

        self.client.force_login(self.storekeeper_user)
        detail = self.client.get(reverse("shipping:detail", args=[self.order.pk]))
        self.assertContains(detail, "Подтвердить отмену")
        self.assertContains(detail, "Отклонить отмену")
        response = self.client.post(
            reverse("shipping:detail", args=[self.order.pk]),
            {"action": "approve_warehouse_cancel", "cancel_reason": "Клиент отменил поставку"},
        )

        self.assertEqual(response.status_code, 302)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, ShippingOrder.STATUS_CANCELED)


class ShippingDisplayNumberTests(TestCase):
    def test_display_shipping_number_matches_public_shipping_order_number(self):
        self.assertEqual(_display_shipping_number("SO-000001"), "SO-000001")
        self.assertEqual(_display_shipping_number("OTG-000121"), "OTG-000121")
        self.assertEqual(_display_shipping_number("SO-000120"), "SO-000120")

    def test_next_shipping_number_uses_new_otg_prefix_after_legacy_so_numbers(self):
        agency = Agency.objects.create(agn_name="Клиент новой нумерации")
        ShippingOrder.objects.create(number="SO-000091", agency=agency)
        ShippingOrder.objects.create(number="OTG-000092", agency=agency)

        self.assertEqual(next_shipping_number(), "OTG-000093")


class ShippingBoxCountTests(TestCase):
    def test_order_box_count_prefers_expected_boxes(self):
        agency = Agency.objects.create(agn_name="Клиент коробов")
        order = ShippingOrder.objects.create(
            number="SO-123456",
            agency=agency,
            expected_boxes=7,
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-1",
            name="Товар",
            qty_requested=10,
            comment="Коробов: 2; кратность: 5",
        )

        self.assertEqual(_order_box_count(order), 7)

    def test_order_box_count_falls_back_to_item_comments(self):
        agency = Agency.objects.create(agn_name="Клиент коробов 2")
        order = ShippingOrder.objects.create(
            number="SO-123457",
            agency=agency,
            expected_boxes=0,
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-1",
            name="Товар 1",
            qty_requested=10,
            comment="Коробов: 2; кратность: 5",
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-2",
            name="Товар 2",
            qty_requested=12,
            comment="Коробов: 3; кратность: 4",
        )

        self.assertEqual(_order_box_count(order), 5)


class ShippingPickerParseTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.stock_rows = [
            {
                "key": "A",
                "sku_id": 1,
                "sku_code": "SKU-A",
                "name": "Товар A",
                "size": "42",
                "barcode": "111",
                "goods_type": "Готовый",
                "box_qty": 5,
                "available_boxes": 10,
                "available_qty": 50,
                "is_mixed_box": False,
                "mixed_group": "",
                "box_codes": ["BOX-A-1", "BOX-A-2", "BOX-A-3", "BOX-A-4", "BOX-A-5", "BOX-A-6", "BOX-A-7", "BOX-A-8", "BOX-A-9", "BOX-A-10"],
            },
            {
                "key": "B",
                "sku_id": 2,
                "sku_code": "SKU-B",
                "name": "Товар B",
                "size": "43",
                "barcode": "222",
                "goods_type": "Готовый",
                "box_qty": 10,
                "available_boxes": 4,
                "available_qty": 40,
                "is_mixed_box": False,
                "mixed_group": "",
                "box_codes": ["BOX-B-1", "BOX-B-2", "BOX-B-3", "BOX-B-4"],
            },
        ]

    def test_saved_whole_box_reselects_full_row_when_original_box_is_partial(self):
        item = SimpleNamespace(
            sku_code="S002",
            name="Сухоцветы",
            size="0",
            barcode="2049107189025",
            goods_type="gv",
            qty_requested=32,
            comment="Коробов: 1; кратность: 32; короба: BOX-OLD",
        )
        order = SimpleNamespace(
            items=SimpleNamespace(order_by=lambda *_args: [item]),
        )
        stock_rows = [
            {
                "key": "PARTIAL",
                "sku_code": "сухоцвет002",
                "name": "Сухоцветы",
                "size": "0",
                "barcode": "2049107189025",
                "goods_type": "gv",
                "box_qty": 32,
                "box_codes": ["BOX-OLD"],
                "partial_only": True,
                "is_mixed_box": False,
                "mixed_group": "partial-box:BOX-OLD",
            },
            {
                "key": "FULL-WRONG-SIZE",
                "sku_code": "сухоцвет002",
                "name": "Сухоцветы",
                "size": "0",
                "barcode": "2049107189025",
                "goods_type": "gv",
                "box_qty": 64,
                "box_codes": ["BOX-64"],
                "partial_only": False,
                "is_mixed_box": False,
                "mixed_group": "",
            },
            {
                "key": "FULL",
                "sku_code": "сухоцвет002",
                "name": "Сухоцветы",
                "size": "0",
                "barcode": "2049107189025",
                "goods_type": "gv",
                "box_qty": 32,
                "box_codes": ["BOX-NEW"],
                "partial_only": False,
                "is_mixed_box": False,
                "mixed_group": "",
            },
        ]

        selected = _selected_box_values_for_order(order, stock_rows)

        self.assertEqual(selected, {"FULL": "1"})

    def test_saved_whole_box_does_not_select_partial_row_without_replacement(self):
        item = SimpleNamespace(
            sku_code="S002",
            name="Сухоцветы",
            size="0",
            barcode="2049107189025",
            goods_type="gv",
            qty_requested=32,
            comment="Коробов: 1; кратность: 32; короба: BOX-OLD",
        )
        order = SimpleNamespace(
            items=SimpleNamespace(order_by=lambda *_args: [item]),
        )
        stock_rows = [
            {
                "key": "PARTIAL",
                "sku_code": "сухоцвет002",
                "name": "Сухоцветы",
                "size": "0",
                "barcode": "2049107189025",
                "goods_type": "gv",
                "box_qty": 32,
                "box_codes": ["BOX-OLD"],
                "partial_only": True,
                "is_mixed_box": False,
                "mixed_group": "partial-box:BOX-OLD",
            },
            {
                "key": "FULL-WRONG-SIZE",
                "sku_code": "сухоцвет002",
                "name": "Сухоцветы",
                "size": "0",
                "barcode": "2049107189025",
                "goods_type": "gv",
                "box_qty": 64,
                "box_codes": ["BOX-64"],
                "partial_only": False,
                "is_mixed_box": False,
                "mixed_group": "",
            },
        ]

        selected = _selected_box_values_for_order(order, stock_rows)

        self.assertEqual(selected, {})

    def test_parse_selected_stock_items_multiple_uses_boxes_for_selected_key(self):
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A", "B"],
                "stock_boxes[]": ["1", "3"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, self.stock_rows, multiple=True)
        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 2)
        selected_by_sku = {row["sku_code"]: row for row in selected}
        self.assertEqual(selected_by_sku["SKU-A"]["qty_requested"], 5)
        self.assertEqual(selected_by_sku["SKU-B"]["qty_requested"], 30)

    def test_parse_selected_stock_items_single_uses_boxes_for_selected_key(self):
        request = self.factory.post(
            "/shipping/1/",
            data={
                "stock_key": "B",
                "stock_key_all[]": ["A", "B"],
                "stock_boxes": ["1", "4"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, self.stock_rows, multiple=False)
        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]["sku_code"], "SKU-B")
        self.assertEqual(selected[0]["qty_requested"], 40)

    def test_parse_selected_stock_items_auto_adds_siblings_for_mixed_box(self):
        stock_rows = [
            {
                "key": "M1",
                "sku_id": 11,
                "sku_code": "SKU-MIX-1",
                "name": "Микс 1",
                "size": "41",
                "barcode": "9101",
                "goods_type": "Готовый",
                "box_qty": 6,
                "available_boxes": 3,
                "available_qty": 18,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
            },
            {
                "key": "M2",
                "sku_id": 12,
                "sku_code": "SKU-MIX-2",
                "name": "Микс 2",
                "size": "42",
                "barcode": "9102",
                "goods_type": "Готовый",
                "box_qty": 4,
                "available_boxes": 3,
                "available_qty": 12,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
            },
        ]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["M1", "M2"],
                "stock_boxes[]": ["2", "0"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)
        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 2)
        selected_by_sku = {row["sku_code"]: row for row in selected}
        self.assertEqual(selected_by_sku["SKU-MIX-1"]["qty_requested"], 12)
        self.assertEqual(selected_by_sku["SKU-MIX-2"]["qty_requested"], 8)

    def test_parse_selected_stock_items_uses_positive_boxes_without_checkbox(self):
        stock_rows = [
            {
                "key": "M1",
                "sku_id": 11,
                "sku_code": "SKU-MIX-1",
                "name": "Микс 1",
                "size": "41",
                "barcode": "9101",
                "goods_type": "Готовый",
                "box_qty": 6,
                "available_boxes": 3,
                "available_qty": 18,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
            },
            {
                "key": "M2",
                "sku_id": 12,
                "sku_code": "SKU-MIX-2",
                "name": "Микс 2",
                "size": "42",
                "barcode": "9102",
                "goods_type": "Готовый",
                "box_qty": 4,
                "available_boxes": 3,
                "available_qty": 12,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
            },
        ]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["M1", "M2"],
                "stock_boxes[]": ["1", "0"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)
        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 2)
        selected_by_sku = {row["sku_code"]: row for row in selected}
        self.assertEqual(selected_by_sku["SKU-MIX-1"]["qty_requested"], 6)
        self.assertEqual(selected_by_sku["SKU-MIX-2"]["qty_requested"], 4)

    def test_parse_selected_stock_items_multiple_ignores_zero_boxes(self):
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A", "B"],
                "stock_boxes[]": ["0", "2"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, self.stock_rows, multiple=True)
        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]["sku_code"], "SKU-B")
        self.assertEqual(selected[0]["qty_requested"], 20)

    def test_parse_selected_stock_items_requires_at_least_one_selected_item(self):
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A", "B"],
                "stock_boxes[]": ["0", "0"],
            },
        )
        selected, errors = _parse_selected_stock_items(request, self.stock_rows, multiple=True)
        self.assertEqual(selected, [])
        self.assertIn("Выберите минимум одну позицию", errors[0])

    def test_parse_selected_stock_items_adds_partial_box_split(self):
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A", "B"],
                "stock_boxes[]": ["0", "0"],
                "stock_partial_box_splits_json": json.dumps(
                    [
                        {
                            "identity": "key:A",
                            "group_key": "partial:A",
                            "row_key": "A",
                            "boxes": 1,
                            "items": [{"key": "A", "qty": 3}],
                        }
                    ]
                ),
            },
        )
        selected, errors = _parse_selected_stock_items(request, self.stock_rows, multiple=True)

        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]["sku_code"], "SKU-A")
        self.assertEqual(selected[0]["qty_requested"], 3)
        meta = extract_partial_box_split(selected[0]["comment"])
        self.assertEqual(meta["source_boxes"], 1)
        self.assertEqual(meta["item_pick_qty"], 3)
        self.assertEqual(meta["source_box_pattern"][0]["qty"], 5)

    def test_parse_selected_stock_items_allows_whole_box_plus_split_within_available_boxes(self):
        stock_rows = [dict(self.stock_rows[0], available_boxes=2, available_qty=10)]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A"],
                "stock_boxes[]": ["1"],
                "stock_partial_box_splits_json": json.dumps(
                    [
                        {
                            "identity": "key:A",
                            "group_key": "partial:A",
                            "row_key": "A",
                            "boxes": 1,
                            "items": [{"key": "A", "qty": 3}],
                        }
                    ]
                ),
            },
        )
        selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)

        self.assertEqual(errors, [])
        self.assertEqual(sum(item["qty_requested"] for item in selected), 8)
        self.assertEqual(len(selected), 2)

    def test_parse_selected_stock_items_rejects_whole_box_plus_split_over_available_boxes(self):
        stock_rows = [dict(self.stock_rows[0], available_boxes=2, available_qty=10)]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["A"],
                "stock_boxes[]": ["2"],
                "stock_partial_box_splits_json": json.dumps(
                    [
                        {
                            "identity": "key:A",
                            "group_key": "partial:A",
                            "row_key": "A",
                            "boxes": 1,
                            "items": [{"key": "A", "qty": 3}],
                        }
                    ]
                ),
            },
        )
        _selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)

        self.assertTrue(errors)
        self.assertIn("выбрано 3 короб", errors[0])

    def test_parse_selected_stock_items_adds_partial_mixed_box_split(self):
        stock_rows = [
            {
                "key": "M1",
                "sku_id": 11,
                "sku_code": "SKU-MIX-1",
                "name": "Микс 1",
                "size": "41",
                "barcode": "9101",
                "goods_type": "Готовый",
                "box_qty": 6,
                "available_boxes": 3,
                "available_qty": 18,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
                "box_codes": ["BOX-MIX-1", "BOX-MIX-2", "BOX-MIX-3"],
            },
            {
                "key": "M2",
                "sku_id": 12,
                "sku_code": "SKU-MIX-2",
                "name": "Микс 2",
                "size": "42",
                "barcode": "9102",
                "goods_type": "Готовый",
                "box_qty": 4,
                "available_boxes": 3,
                "available_qty": 12,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
                "box_codes": ["BOX-MIX-1", "BOX-MIX-2", "BOX-MIX-3"],
            },
        ]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["M1", "M2"],
                "stock_boxes[]": ["0", "0"],
                "stock_partial_box_splits_json": json.dumps(
                    [
                        {
                            "identity": "group:mix:g1",
                            "group_key": "partial:mix:g1",
                            "row_key": "M1",
                            "group": "mix:g1",
                            "boxes": 2,
                            "items": [{"key": "M1", "qty": 2}, {"key": "M2", "qty": 1}],
                        }
                    ]
                ),
            },
        )
        selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)

        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 2)
        selected_by_sku = {row["sku_code"]: row for row in selected}
        self.assertEqual(selected_by_sku["SKU-MIX-1"]["qty_requested"], 4)
        self.assertEqual(selected_by_sku["SKU-MIX-2"]["qty_requested"], 2)
        meta = extract_partial_box_split(selected_by_sku["SKU-MIX-1"]["comment"])
        self.assertEqual(meta["source_boxes"], 2)
        self.assertEqual(len(meta["source_box_pattern"]), 2)
        self.assertEqual(sum(row["qty"] for row in meta["pick_pattern"]), 3)
        self.assertEqual(meta["source_box_codes"], ["BOX-MIX-1", "BOX-MIX-2"])
        self.assertGreater(len(selected_by_sku["SKU-MIX-1"]["comment"]), 255)
        self.assertIsInstance(ShippingOrderItem._meta.get_field("comment"), models.TextField)

    def test_parse_selected_stock_items_keeps_split_and_whole_mixed_boxes_distinct(self):
        stock_rows = [
            {
                "key": "M1",
                "sku_id": 11,
                "sku_code": "1108_Beige",
                "name": "Пальто осеннее",
                "size": "M",
                "barcode": "2009891653454",
                "goods_type": "gv",
                "box_qty": 25,
                "available_boxes": 2,
                "available_qty": 50,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
                "box_codes": ["BOX-MIX-A", "BOX-MIX-B"],
            },
            {
                "key": "M2",
                "sku_id": 12,
                "sku_code": "1108_Beige",
                "name": "Пальто осеннее",
                "size": "XL",
                "barcode": "2009891653478",
                "goods_type": "gv",
                "box_qty": 25,
                "available_boxes": 2,
                "available_qty": 50,
                "is_mixed_box": True,
                "mixed_group": "mix:g1",
                "box_codes": ["BOX-MIX-A", "BOX-MIX-B"],
            },
        ]
        request = self.factory.post(
            "/shipping/new/",
            data={
                "stock_key_all[]": ["M1", "M2"],
                "stock_boxes[]": ["1", "0"],
                "stock_partial_box_splits_json": json.dumps(
                    [
                        {
                            "identity": "group:mix:g1",
                            "group_key": "partial:mix:g1",
                            "row_key": "M1",
                            "group": "mix:g1",
                            "boxes": 1,
                            "items": [{"key": "M1", "qty": 10}],
                        }
                    ]
                ),
            },
        )

        selected, errors = _parse_selected_stock_items(request, stock_rows, multiple=True)

        self.assertEqual(errors, [])
        self.assertEqual(len(selected), 3)
        loose_m = next(row for row in selected if row["size"] == "M" and "Разбить короб" in row["comment"])
        whole_rows = [row for row in selected if "микс-короб" in row["comment"]]
        self.assertEqual(loose_m["qty_requested"], 10)
        self.assertIn("короба: BOX-MIX-A", loose_m["comment"])
        self.assertEqual({row["size"] for row in whole_rows}, {"M", "XL"})
        self.assertTrue(all("короба: BOX-MIX-B" in row["comment"] for row in whole_rows))


class ShippingOrderFormValidationTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Client A")
        self.market = Market.objects.create(id=201, name="WB")
        self.ozon_market = Market.objects.create(id=202, name="Ozon")

    def _base_data(self, *, marketplace_id: str | None = None):
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        return {
            "agency": str(self.agency.id),
            "slot_date": eta.strftime("%Y-%m-%d"),
            "slot_time": "",
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "shipping_barcode": "SHIP-123",
            "marketplace": marketplace_id or str(self.market.id),
            "wb_supply_barcode": "SUP-123",
            "wb_transit_warehouse": "",
            "transit_address": "",
            "destination_warehouse": "Склад WB Коледино",
            "supply_type": ShippingOrder.SUPPLY_BOX,
            "vehicle_type": ShippingOrder.VEHICLE_FULFILLMENT,
            "vehicle_number": "A123BC77",
            "driver_phone": "+7 900 000-00-00",
            "comment": "",
        }

    def test_form_is_valid_with_receiving_like_fields(self):
        data = self._base_data()
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)

    def test_relax_for_autosave_allows_agency_only_draft(self):
        form = ShippingOrderForm(
            data={"agency": str(self.agency.id), "delivery_type": ""},
            relax_for_autosave=True,
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["agency"], self.agency)
        self.assertEqual(
            form.cleaned_data["delivery_type"],
            ShippingOrder.DELIVERY_MARKETPLACE,
        )

    def test_form_sets_marketplace_datalists_for_destination_fields(self):
        form = ShippingOrderForm()
        self.assertEqual(
            form.fields["destination_warehouse"].widget.attrs.get("list"),
            "destination-warehouse-options",
        )
        self.assertEqual(
            form.fields["transit_address"].widget.attrs.get("list"),
            "transit-address-options",
        )

    def test_invalid_driver_phone_is_rejected(self):
        data = self._base_data()
        data["driver_phone"] = "123"
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("driver_phone", form.errors)

    def test_driver_phone_and_vehicle_number_are_optional(self):
        data = self._base_data()
        data["driver_phone"] = ""
        data["vehicle_number"] = ""
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)

    def test_transport_is_required(self):
        data = self._base_data()
        data["vehicle_type"] = ""
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("vehicle_type", form.errors)

    def test_pickup_requires_shipping_date(self):
        data = self._base_data()
        data["delivery_type"] = ShippingOrder.DELIVERY_PICKUP
        data["eta_at"] = ""
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("eta_at", form.errors)

    def test_pickup_keeps_explicit_shipping_date(self):
        eta = (timezone.localtime() + timedelta(days=2)).replace(
            hour=0,
            minute=0,
            second=0,
            microsecond=0,
        )
        data = self._base_data()
        data["delivery_type"] = ShippingOrder.DELIVERY_PICKUP
        data["eta_at"] = eta.strftime("%Y-%m-%dT%H:%M")
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(
            timezone.localtime(form.cleaned_data["eta_at"]).date(),
            eta.date(),
        )
        self.assertEqual(form.cleaned_data["vehicle_type"], "")
        self.assertEqual(form.cleaned_data["vehicle_number"], "")
        self.assertEqual(form.cleaned_data["driver_phone"], "")

    def test_courier_form_hides_marketplace_requirements_and_allows_cdek(self):
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        data = {
            "agency": str(self.agency.id),
            "delivery_type": ShippingOrder.DELIVERY_COURIER,
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "vehicle_type": ShippingOrder.VEHICLE_CDEK,
            "vehicle_number": "",
            "driver_phone": "",
            "comment": "Курьер СДЭК",
            "marketplace": "",
            "shipping_barcode": "",
            "wb_supply_barcode": "",
            "destination_warehouse": "",
            "supply_type": "",
        }
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["marketplace"])
        self.assertEqual(form.cleaned_data["vehicle_type"], ShippingOrder.VEHICLE_CDEK)
        self.assertEqual(form.cleaned_data["destination_warehouse"], "")

    def test_courier_form_allows_yandex_vehicle(self):
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        data = {
            "agency": str(self.agency.id),
            "delivery_type": ShippingOrder.DELIVERY_COURIER,
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "vehicle_type": ShippingOrder.VEHICLE_YANDEX,
        }
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["vehicle_type"], ShippingOrder.VEHICLE_YANDEX)

    def test_transfer_form_hides_marketplace_and_transport_requirements(self):
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        data = {
            "agency": str(self.agency.id),
            "delivery_type": ShippingOrder.DELIVERY_TRANSFER,
            "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
            "destination_address": "Склад клиента в Подольске",
            "vehicle_type": "",
            "vehicle_number": "A123BC77",
            "driver_phone": "+7 900 000-00-00",
            "marketplace": "",
            "shipping_barcode": "",
            "wb_supply_barcode": "",
            "destination_warehouse": "",
            "supply_type": "",
        }
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["marketplace"])
        self.assertEqual(form.cleaned_data["vehicle_type"], "")
        self.assertEqual(form.cleaned_data["vehicle_number"], "")
        self.assertEqual(form.cleaned_data["driver_phone"], "")
        self.assertEqual(form.cleaned_data["destination_address"], "Склад клиента в Подольске")
        self.assertEqual(form.cleaned_data["destination_warehouse"], "")

    def test_transfer_form_requires_destination_address(self):
        eta = (timezone.localtime() + timedelta(days=2)).replace(hour=10, minute=0, second=0, microsecond=0)
        form = ShippingOrderForm(
            data={
                "agency": str(self.agency.id),
                "delivery_type": ShippingOrder.DELIVERY_TRANSFER,
                "eta_at": eta.strftime("%Y-%m-%dT%H:%M"),
                "vehicle_type": ShippingOrder.VEHICLE_CLIENT,
                "destination_address": "",
            }
        )
        self.assertFalse(form.is_valid())
        self.assertIn("destination_address", form.errors)

    def test_transit_address_is_required_when_transit_enabled(self):
        data = self._base_data()
        data["wb_transit_warehouse"] = "on"
        data["transit_address"] = ""
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("transit_address", form.errors)

    def test_transit_address_is_valid_when_transit_enabled(self):
        data = self._base_data()
        data["wb_transit_warehouse"] = "on"
        data["transit_address"] = "Москва, Транзитный терминал 1"
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)

    def test_ozon_transit_route_allows_empty_destination_warehouse(self):
        data = self._base_data(marketplace_id=str(self.ozon_market.id))
        data.update(
            {
                "slot_time": "13:30",
                "supply_number": "OZON-TRANSIT-001",
                "wb_transit_warehouse": "on",
                "transit_address": "СТАРАЯ_КУПАВНА_28",
                "destination_warehouse": "",
            }
        )
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["destination_warehouse"], "")

    def test_ozon_transit_route_still_requires_transit_address(self):
        data = self._base_data(marketplace_id=str(self.ozon_market.id))
        data.update(
            {
                "slot_time": "13:30",
                "supply_number": "OZON-TRANSIT-002",
                "wb_transit_warehouse": "on",
                "transit_address": "",
                "destination_warehouse": "",
            }
        )
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("transit_address", form.errors)
        self.assertNotIn("destination_warehouse", form.errors)

    def test_wb_transit_route_still_requires_destination_warehouse(self):
        data = self._base_data()
        data.update(
            {
                "wb_transit_warehouse": "on",
                "transit_address": "Москва, Транзитный терминал 1",
                "destination_warehouse": "",
            }
        )
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("destination_warehouse", form.errors)

    def test_core_fields_are_required(self):
        required_fields = [
            "shipping_barcode",
            "marketplace",
            "wb_supply_barcode",
            "destination_warehouse",
            "supply_type",
        ]
        for field_name in required_fields:
            data = self._base_data()
            data[field_name] = ""
            form = ShippingOrderForm(data=data)
            self.assertFalse(form.is_valid())
            self.assertIn(field_name, form.errors)

    def test_next_day_cutoff_after_11_is_rejected(self):
        fixed_now = datetime(2026, 3, 19, 11, 1, 0)
        data = self._base_data()
        data["eta_at"] = "2026-03-20T12:00"
        with patch("shipping.forms.timezone.localtime", return_value=fixed_now):
            form = ShippingOrderForm(data=data)
            self.assertFalse(form.is_valid())
            self.assertIn("eta_at", form.errors)

    def test_marketplace_before_11_allows_tomorrow_but_not_today(self):
        fixed_now = datetime(2026, 3, 19, 10, 59, 0)
        with patch("shipping.forms.timezone.localtime", return_value=fixed_now):
            today_data = self._base_data()
            today_data["eta_at"] = "2026-03-19T12:00"
            today_form = ShippingOrderForm(data=today_data)
            self.assertFalse(today_form.is_valid())
            self.assertIn("eta_at", today_form.errors)

            tomorrow_data = self._base_data()
            tomorrow_data["eta_at"] = "2026-03-20T12:00"
            tomorrow_form = ShippingOrderForm(data=tomorrow_data)
            self.assertTrue(tomorrow_form.is_valid(), tomorrow_form.errors)

    def test_marketplace_at_11_allows_only_day_after_tomorrow(self):
        fixed_now = datetime(2026, 3, 19, 11, 0, 0)
        with patch("shipping.forms.timezone.localtime", return_value=fixed_now):
            tomorrow_data = self._base_data()
            tomorrow_data["eta_at"] = "2026-03-20T12:00"
            tomorrow_form = ShippingOrderForm(data=tomorrow_data)
            self.assertFalse(tomorrow_form.is_valid())
            self.assertIn("eta_at", tomorrow_form.errors)

            later_data = self._base_data()
            later_data["eta_at"] = "2026-03-21T12:00"
            later_form = ShippingOrderForm(data=later_data)
            self.assertTrue(later_form.is_valid(), later_form.errors)

    def test_manager_emergency_override_allows_today(self):
        fixed_now = datetime(2026, 3, 19, 14, 0, 0)
        data = self._base_data()
        data["eta_at"] = "2026-03-19T12:00"
        with patch("shipping.forms.timezone.localtime", return_value=fixed_now):
            form = ShippingOrderForm(
                data=data,
                allow_marketplace_date_override=True,
            )
            self.assertTrue(form.is_valid(), form.errors)

    def test_marketplace_cutoff_does_not_apply_to_pickup(self):
        fixed_now = datetime(2026, 3, 19, 14, 0, 0)
        data = self._base_data()
        data["delivery_type"] = ShippingOrder.DELIVERY_PICKUP
        data["eta_at"] = "2026-03-19T12:00"
        with patch("shipping.forms.timezone.localtime", return_value=fixed_now):
            form = ShippingOrderForm(data=data)
            self.assertTrue(form.is_valid(), form.errors)

    def test_shipping_time_before_workday_is_rejected(self):
        data = self._base_data()
        eta_date = (timezone.localtime() + timedelta(days=2)).strftime("%Y-%m-%d")
        data["eta_at"] = f"{eta_date}T07:55"
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("eta_at", form.errors)

    def test_shipping_time_after_workday_is_rejected(self):
        data = self._base_data()
        eta_date = (timezone.localtime() + timedelta(days=2)).strftime("%Y-%m-%d")
        data["eta_at"] = f"{eta_date}T19:05"
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("eta_at", form.errors)

    def test_shipping_time_at_end_of_workday_is_allowed(self):
        data = self._base_data()
        eta_date = (timezone.localtime() + timedelta(days=2)).strftime("%Y-%m-%d")
        data["eta_at"] = f"{eta_date}T19:00"
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)

    def test_ozon_requires_slot_date_and_slot_time(self):
        data = self._base_data(marketplace_id=str(self.ozon_market.id))
        data["slot_date"] = ""
        data["slot_time"] = ""
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("slot_date", form.errors)
        self.assertIn("slot_time", form.errors)

    def test_ozon_accepts_slot_time(self):
        data = self._base_data(marketplace_id=str(self.ozon_market.id))
        data["slot_time"] = "13:30"
        form = ShippingOrderForm(data=data)
        self.assertTrue(form.is_valid(), form.errors)

    def test_slot_time_requires_slot_date(self):
        data = self._base_data()
        data["slot_date"] = ""
        data["slot_time"] = "13:30"
        form = ShippingOrderForm(data=data)
        self.assertFalse(form.is_valid())
        self.assertIn("slot_date", form.errors)


class ShippingMarketplaceWarehouseUiTests(TestCase):
    @patch(
        "shipping.views.load_marketplace_warehouse_catalog",
        return_value={
            "wb": [{"type": "Обычный", "name": "Коледино", "address": "Московская область"}],
            "ozon": [{"type": "Обычный", "name": "Санкт-Петербург_РФЦ", "address": "Софийская, 118"}],
            "yandex": [{"type": "Транзитный", "name": "Москва — Бутырский", "address": "Москва, ул. Руставели, 3"}],
            "sber": [],
        },
    )
    def test_create_view_exposes_marketplace_warehouse_catalog(self, _mock_loader):
        user_model = get_user_model()
        client_user = user_model.objects.create_user(username="client_ui", password="pwd")
        agency = Agency.objects.create(agn_name="Клиент UI", portal_user=client_user)

        self.client.force_login(client_user)
        response = self.client.get("/shipping/new/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.context["marketplace_warehouse_catalog"]["wb"],
            [{"type": "Обычный", "name": "Коледино", "address": "Московская область"}],
        )
        self.assertContains(response, "destination-warehouse-options")
        self.assertContains(response, "transit-address-options")
        self.assertContains(response, "marketplace-warehouse-catalog")
        self.assertContains(response, 'data-picker-target="destination"', html=False)
        self.assertContains(response, 'data-picker-target="transit"', html=False)
        self.assertContains(response, "warehouse-picker-overlay")
        self.assertContains(response, "Состав перемещения из Excel")
        self.assertContains(response, 'data-transfer-template-download="1"', count=1, html=False)
        self.assertContains(response, 'data-transfer-template-trigger="1"', count=2, html=False)
        self.assertContains(response, f'/shipping/template.xlsx?client={agency.id}', html=False)
        self.assertContains(response, 'for="shipping-template-input"', count=1, html=False)
        self.assertContains(response, 'data-marketplace-independent="1"', count=1, html=False)
        rendered = response.content.decode("utf-8")
        self.assertLess(
            rendered.index('id="shipping-template-input"'),
            rendered.index('class="controls" data-hide-for-marketplace="ozon"'),
        )
        self.assertContains(response, "templateInput.click()", count=1, html=False)
        self.assertContains(response, 'data-transfer-hide="1"', count=3, html=False)

    @patch(
        "shipping.views.load_marketplace_warehouse_catalog",
        return_value={"wb": [], "ozon": [], "yandex": [], "sber": []},
    )
    def test_client_create_view_shows_marketplace_shipping_date_rule(self, _mock_loader):
        user_model = get_user_model()
        client_user = user_model.objects.create_user(username="client_cutoff_ui", password="pwd")
        Agency.objects.create(agn_name="Клиент cutoff UI", portal_user=client_user)

        self.client.force_login(client_user)
        response = self.client.get("/shipping/new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "Заявки на маркетплейсы до 11:00 принимаются на отгрузку на завтра",
        )
        self.assertTrue(response.context["client_marketplace_min_ship_date"])

    @patch(
        "shipping.views.load_marketplace_warehouse_catalog",
        return_value={"wb": [], "ozon": [], "yandex": [], "sber": []},
    )
    def test_create_view_exposes_single_and_batch_ozon_api_picker(self, _mock_loader):
        user_model = get_user_model()
        client_user = user_model.objects.create_user(username="client_ozon_picker", password="pwd")
        Agency.objects.create(agn_name="Клиент Ozon picker", portal_user=client_user)

        self.client.force_login(client_user)
        response = self.client.get("/shipping/new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="ozon-pull-supply-btn"', html=False)
        self.assertContains(response, 'data-ozon-supplies-url="', html=False)
        self.assertContains(response, 'data-ozon-supplies-batch-url="', html=False)
        self.assertContains(response, 'data-ozon-supplies-detail-url="', html=False)
        self.assertContains(response, "/shipping/ozon-supplies/", html=False)
        self.assertContains(response, "/shipping/ozon-supplies/batch/", html=False)
        self.assertContains(response, "/shipping/ozon-supplies/{id}/", html=False)
        self.assertContains(response, "Выбрать заявки Ozon")
        self.assertContains(response, 'id="ozon-supply-picker-overlay"', html=False)
        self.assertNotContains(response, 'id="ozon-supply-picker-open-one"', html=False)
        self.assertContains(response, 'id="ozon-supply-picker-select-all"', html=False)
        self.assertContains(response, 'id="ozon-supply-picker-confirm"', html=False)
        self.assertContains(response, 'id="ozon-supply-picker-selected-numbers"', html=False)
        self.assertContains(response, 'name="documents"', html=False)
        self.assertContains(response, "Они будут прикреплены при отправке заявки Ozon")
        self.assertContains(response, "loadSelectedOzonSupply")
        self.assertContains(response, "previewSelectedOzonSupplies")
        self.assertNotContains(response, "submitReadyOzonSupplies")
        self.assertContains(response, "buildCombinedOzonPayload")
        self.assertContains(response, "is_combined_ozon_supply")
        self.assertContains(response, "piece_pick_total_qty")
        self.assertContains(response, "Поставка Ozon")
        self.assertContains(response, "field.readOnly = isOzon")
        self.assertContains(response, "Для Ozon ШК и номер поставки заполняются автоматически")
        self.assertContains(response, "Все выбранные поставки будут объединены в одну заявку Fullbox")
        self.assertContains(response, "Проверить и заполнить одну заявку")
        self.assertContains(response, "Выбрать до 10")
        self.assertContains(response, "ozonSupplySelectedOrderIds.size >= 10")
        self.assertContains(response, ".slice(0, 10)")
        self.assertContains(response, "renderOzonSelectedSupplyNumbers")
        self.assertContains(response, "Выбраны поставки Ozon:")
        self.assertContains(response, "applyOzonPartialSplits")
        self.assertContains(response, "matched_with_piece_pick")
        self.assertContains(response, "confirmOzonPiecePickTariff")
        self.assertContains(response, "Поштучный подбор тарифицируется отдельно")

    @patch(
        "shipping.ozon_supplies.get_ozon_supplies_batch_for_agency",
        return_value={"ok": True, "orders": [], "ready_order_ids": []},
    )
    @patch("shipping.web_ui._shipping_stock_picker_rows", return_value=[])
    def test_ozon_batch_accepts_ten_and_rejects_eleven(self, _mock_stock, mock_batch):
        user_model = get_user_model()
        client_user = user_model.objects.create_user(username="client_ozon_limit", password="pwd")
        Agency.objects.create(agn_name="Клиент Ozon limit", portal_user=client_user)
        self.client.force_login(client_user)

        for selected_count in (1, 8, 9, 10):
            mock_batch.reset_mock()
            accepted = self.client.post(
                "/shipping/ozon-supplies/batch/",
                data=json.dumps({"mode": "preview", "order_ids": list(range(1, selected_count + 1))}),
                content_type="application/json",
            )
            self.assertEqual(accepted.status_code, 200, msg=f"selected_count={selected_count}")
            mock_batch.assert_called_once()

        mock_batch.reset_mock()
        rejected = self.client.post(
            "/shipping/ozon-supplies/batch/",
            data=json.dumps({"mode": "preview", "order_ids": list(range(1, 12))}),
            content_type="application/json",
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(
            rejected.json().get("error"),
            "За один раз можно выбрать не более 10 поставок Ozon.",
        )
        mock_batch.assert_not_called()

    @patch(
        "shipping.views.load_marketplace_warehouse_catalog",
        return_value={
            "wb": [
                {"type": "Обычный", "name": "Коледино", "address": "Московская область"},
                {"type": "Транзитный / ППП", "name": "Гольёво", "address": "Московская область"},
                {"type": "Обычный + транзитный", "name": "Софьино", "address": "Московская область"},
            ],
            "ozon": [],
            "yandex": [],
            "sber": [],
        },
    )
    def test_create_view_splits_destination_and_transit_suggestions(self, _mock_loader):
        user_model = get_user_model()
        client_user = user_model.objects.create_user(username="client_ui_split", password="pwd")
        agency = Agency.objects.create(agn_name="Клиент UI split", portal_user=client_user)

        self.client.force_login(client_user)
        response = self.client.get("/shipping/new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "normalizeWarehouseRow")
        self.assertContains(response, "openWarehousePicker")
        self.assertContains(response, "обычных складов")
        self.assertContains(response, "транзитных складов")


class ShippingPackingSlipStatusEndpointTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="storekeeper_status", password="pwd")
        Employee.objects.create(
            full_name="Кладовщик статуса",
            role="storekeeper",
            user=self.user,
            is_active=True,
        )
        self.agency = Agency.objects.create(agn_name="Клиент статуса")
        self.order = ShippingOrder.objects.create(
            number="SO-STATUS-1",
            agency=self.agency,
            created_by=self.user,
            status=ShippingOrder.STATUS_PACKED,
        )
        self.client.force_login(self.user)

    @patch("shipping.views._shipping_packing_summary", return_value={"pallet_count": 1, "box_count": 2})
    @patch(
        "shipping.views.build_print_status_snapshot",
        return_value={
            "available_printers": ["Zebra GK420d"],
            "available_printers_meta": {},
            "print_status_line": "Готово к печати: 1 из 1",
            "print_agent_line": "Warehouse-PC-01 · 14.04.2026 16:00:00",
            "print_agent_online": True,
            "print_paused": False,
            "print_queue_pending": 2,
            "print_queue_printing": 1,
            "print_queue_failed": 0,
            "print_queue_stuck": 0,
            "print_queue_pending_jobs": ["#10 · Zebra GK420d"],
            "print_queue_printing_jobs": ["#11 · Zebra GK420d"],
            "print_queue_failed_jobs": [],
            "print_queue_stuck_jobs": [],
            "printer_statuses": [
                {
                    "name": "Zebra GK420d",
                    "state_key": "ready",
                    "state_label": "Готов",
                    "color": "green",
                    "jobs": 0,
                    "status": "",
                    "is_default": True,
                    "is_local": True,
                    "is_network": False,
                    "is_virtual": False,
                }
            ],
            "printer_details_updated_at": "14.04.2026 16:00:01",
        },
    )
    def test_get_status_endpoint_returns_snapshot(self, _snapshot, _packing_summary):
        response = self.client.get(reverse("shipping:packing-slips-status", args=[self.order.pk]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["print_queue_pending"], 2)
        self.assertEqual(payload["printer_statuses"][0]["state_key"], "ready")

    @patch("shipping.views._shipping_packing_summary", return_value={"pallet_count": 1, "box_count": 2})
    @patch(
        "shipping.views.refresh_print_agent_printers",
        return_value=(
            {
                "available_printers": ["Zebra GK420d"],
                "available_printers_meta": {},
                "print_status_line": "Есть ошибки печати: 1",
                "print_agent_line": "Warehouse-PC-01 · 14.04.2026 16:05:00",
                "print_agent_online": True,
                "print_paused": False,
                "print_queue_pending": 0,
                "print_queue_printing": 0,
                "print_queue_failed": 1,
                "print_queue_stuck": 0,
                "print_queue_pending_jobs": [],
                "print_queue_printing_jobs": [],
                "print_queue_failed_jobs": ["#12 · Zebra GK420d"],
                "print_queue_stuck_jobs": [],
                "printer_statuses": [],
                "printer_details_updated_at": "14.04.2026 16:05:01",
            },
            "Диагностика обновлена.",
        ),
    )
    def test_post_status_endpoint_refreshes_agent_state(self, _refresh, _packing_summary):
        response = self.client.post(reverse("shipping:packing-slips-status", args=[self.order.pk]))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["refresh_message"], "Диагностика обновлена.")
        self.assertEqual(payload["print_queue_failed"], 1)


class StrictOzonGmIntakeTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Strict Ozon client")
        self.ozon = Market.objects.create(id=9983, name="Ozon")

    def _order(
        self,
        *,
        number: str,
        expected_boxes: int = 1,
        created_at: str = "2026-08-11T06:05:00+00:00",
    ) -> ShippingOrder:
        order = ShippingOrder.objects.create(
            number=number,
            agency=self.agency,
            marketplace=self.ozon,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
            expected_boxes=expected_boxes,
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="knife018",
            name="Knife",
            barcode="2042857120430",
            qty_requested=50,
        )
        ShippingOrder.objects.filter(pk=order.pk).update(
            created_at=datetime.fromisoformat(created_at)
        )
        order.refresh_from_db()
        return order

    def _save_snapshot(
        self,
        order: ShippingOrder,
        *,
        boxes_total: int,
        cargoes: list[dict],
    ) -> None:
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            payload={
                "is_ozon_api_supply": True,
                "ozon_boxes_total": boxes_total,
                "gm_cargoes": cargoes,
            },
        )

    def test_incomplete_three_gm_snapshot_stays_blocked_without_source_box_comparison(self):
        order = self._order(number="OTG-STRICT-OZON-1")
        self._save_snapshot(
            order,
            boxes_total=3,
            cargoes=[
                {"gm_barcode": f"GM-{index}", "items": []}
                for index in range(1, 4)
            ],
        )

        errors = _strict_ozon_gm_intake_errors(order)

        self.assertTrue(any("полный состав" in error for error in errors))
        self.assertFalse(any("Ozon 3" in error and "Fullbox 1" in error for error in errors))

    def test_exact_gm_snapshot_matches_request_and_box_count(self):
        order = self._order(number="OTG-STRICT-OZON-2")
        self._save_snapshot(
            order,
            boxes_total=1,
            cargoes=[
                {
                    "gm_barcode": "GM-EXACT-1",
                    "items": [{"barcode": "2042857120430", "quantity": 50}],
                }
            ],
        )

        self.assertEqual(_strict_ozon_gm_intake_errors(order), [])

    def test_exact_gm_snapshot_allows_more_fullbox_source_boxes_than_ozon_cargoes(self):
        order = self._order(number="OTG-STRICT-OZON-CONSOLIDATED", expected_boxes=8)
        self._save_snapshot(
            order,
            boxes_total=6,
            cargoes=[
                {
                    "gm_barcode": f"GM-CONSOLIDATED-{index}",
                    "items": [
                        {
                            "barcode": "2042857120430",
                            "quantity": 10 if index == 6 else 8,
                        }
                    ],
                }
                for index in range(1, 7)
            ],
        )

        self.assertEqual(_strict_ozon_gm_intake_errors(order), [])

    @patch("shipping.ozon_supplies.enrich_gm_cargoes_with_items")
    @patch("shipping.ozon_supplies.fetch_ozon_gm_cargoes")
    @patch("shipping.ozon_supplies.resolve_ozon_credential")
    def test_refresh_fetches_cargoes_when_saved_snapshot_is_empty(
        self,
        resolve_credential,
        fetch_cargoes,
        enrich_cargoes,
    ):
        order = self._order(number="OTG-STRICT-OZON-EMPTY", expected_boxes=2)
        order.supply_number = "2000063050020"
        order.save(update_fields=["supply_number", "updated_at"])
        self._save_snapshot(order, boxes_total=0, cargoes=[])
        resolve_credential.return_value = (
            SimpleNamespace(client_id="123", market_key="secret"),
            "",
        )
        fetched = [
            {"gm_barcode": "GM-LIVE-1", "supply_id": 2000063050020, "cargo_id": 1},
            {"gm_barcode": "GM-LIVE-2", "supply_id": 2000063050020, "cargo_id": 2},
        ]
        fetch_cargoes.return_value = (fetched, "")
        enrich_cargoes.return_value = [
            {
                **cargo,
                "items": [{"barcode": "2042857120430", "quantity": 25}],
            }
            for cargo in fetched
        ]

        errors = refresh_strict_ozon_gm_intake(order)

        self.assertEqual(errors, [])
        fetch_cargoes.assert_called_once_with(
            "123",
            "secret",
            ["2000063050020"],
        )
        refreshed = _saved_ozon_api_intake_meta(order)
        self.assertEqual(refreshed["boxes_total"], 2)
        self.assertEqual(
            [row["gm_barcode"] for row in refreshed["gm_cargoes"]],
            ["GM-LIVE-1", "GM-LIVE-2"],
        )

    @patch("shipping.ozon_supplies.enrich_gm_cargoes_with_items")
    @patch("shipping.ozon_supplies.fetch_ozon_gm_cargoes")
    @patch("shipping.ozon_supplies.resolve_ozon_credential")
    def test_refresh_discards_snapshot_from_another_supply(
        self,
        resolve_credential,
        fetch_cargoes,
        enrich_cargoes,
    ):
        order = self._order(number="OTG-STRICT-OZON-STALE", expected_boxes=2)
        order.supply_number = "2000063050020"
        order.save(update_fields=["supply_number", "updated_at"])
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            payload={
                "is_ozon_api_supply": True,
                "ozon_supply_numbers": ["2000063999999"],
                "ozon_supply_id_normalized": "2000063999999",
                "ozon_boxes_total": 1,
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-STALE",
                        "supply_id": 2000063999999,
                        "cargo_id": 999,
                        "items": [{"barcode": "2042857120430", "quantity": 50}],
                    }
                ],
            },
        )
        resolve_credential.return_value = (
            SimpleNamespace(client_id="123", market_key="secret"),
            "",
        )
        fetched = [
            {"gm_barcode": "GM-LIVE-1", "supply_id": 2000063050020, "cargo_id": 1},
            {"gm_barcode": "GM-LIVE-2", "supply_id": 2000063050020, "cargo_id": 2},
        ]
        fetch_cargoes.return_value = (fetched, "")
        enrich_cargoes.return_value = [
            {
                **cargo,
                "items": [{"barcode": "2042857120430", "quantity": 25}],
            }
            for cargo in fetched
        ]

        errors = refresh_strict_ozon_gm_intake(order)

        self.assertEqual(errors, [])
        fetch_cargoes.assert_called_once_with(
            "123",
            "secret",
            ["2000063050020"],
        )
        refreshed = _saved_ozon_api_intake_meta(order)
        self.assertEqual(refreshed["supply_numbers"], ["2000063050020"])
        self.assertEqual(refreshed["boxes_total"], 2)
        self.assertEqual(
            [row["gm_barcode"] for row in refreshed["gm_cargoes"]],
            ["GM-LIVE-1", "GM-LIVE-2"],
        )

    def test_reconciled_snapshot_supersedes_initial_immutable_snapshot(self):
        order = self._order(number="OTG-STRICT-OZON-RECONCILED")
        self._save_snapshot(
            order,
            boxes_total=2,
            cargoes=[
                {
                    "gm_barcode": "GM-REMOVED",
                    "items": [{"barcode": "2043986017738", "quantity": 32}],
                },
                {
                    "gm_barcode": "GM-KEPT",
                    "items": [{"barcode": "2042857120430", "quantity": 50}],
                },
            ],
        )
        kept_cargo = {
            "gm_barcode": "GM-KEPT",
            "items": [{"barcode": "2042857120430", "quantity": 50}],
        }
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            payload={
                "act": "shipping_ozon_gm_reconcile_after_item_removal",
                "is_ozon_api_supply": True,
                "ozon_snapshot_complete": True,
                "ozon_boxes_total": 1,
                "gm_cargoes": [kept_cargo],
            },
        )

        meta = _saved_ozon_api_intake_meta(order)

        self.assertEqual(meta["boxes_total"], 1)
        self.assertEqual(meta["gm_cargoes"], [kept_cargo])
        self.assertEqual(_strict_ozon_gm_intake_errors(order), [])

    def test_editing_draft_prefers_complete_snapshot_matching_selected_rows(self):
        order = self._order(number="OTG-STRICT-OZON-EDIT", expected_boxes=2)
        order.status = ShippingOrder.STATUS_DRAFT
        order.supply_number = "2000063640896"
        order.save(update_fields=["status", "supply_number", "updated_at"])
        selected_rows = [
            {
                "sku_code": "knife018",
                "barcode": "2042857120430",
                "qty_requested": 50,
            }
        ]
        complete_cargoes = [
            {
                "gm_barcode": "GM-FULL-1",
                "items": [{"barcode": "2042857120430", "quantity": 20}],
            },
            {
                "gm_barcode": "GM-FULL-2",
                "items": [{"barcode": "2042857120430", "quantity": 30}],
            },
        ]
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            payload={
                "is_ozon_api_supply": True,
                "ozon_snapshot_source": "ozon_api",
                "ozon_snapshot_complete": True,
                "ozon_supply_id_normalized": "2000063640896",
                "ozon_supply_numbers": ["2000063640896"],
                "ozon_boxes_total": 2,
                "gm_cargoes": complete_cargoes,
            },
        )
        reduced_cargo = {
            "gm_barcode": "GM-REDUCED",
            "items": [{"barcode": "2042857120430", "quantity": 30}],
        }
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            payload={
                "act": "shipping_ozon_gm_reconcile_after_item_removal",
                "is_ozon_api_supply": True,
                "ozon_snapshot_source": "ozon_api",
                "ozon_snapshot_complete": True,
                "ozon_supply_id_normalized": "2000063640896",
                "ozon_supply_numbers": ["2000063640896"],
                "ozon_boxes_total": 1,
                "gm_cargoes": [reduced_cargo],
            },
        )

        latest_meta = _saved_ozon_api_intake_meta(order)
        exact_meta = _saved_ozon_api_intake_meta(
            order,
            selected_rows=selected_rows,
            prefer_exact_match=True,
        )
        reconciled_meta, reconcile_errors = _reconcile_ozon_gm_meta_with_selected_rows(
            agency=self.agency,
            meta=exact_meta,
            selected_rows=selected_rows,
        )

        self.assertEqual(latest_meta["boxes_total"], 1)
        self.assertTrue(exact_meta["_selected_rows_exact_match"])
        self.assertEqual(exact_meta["boxes_total"], 2)
        self.assertEqual(exact_meta["gm_cargoes"], complete_cargoes)
        self.assertEqual(reconcile_errors, [])
        self.assertEqual(reconciled_meta["boxes_total"], 2)

    def test_new_incomplete_snapshot_blocks_pick_without_warehouse_writes(self):
        order = self._order(number="OTG-STRICT-OZON-3")
        self._save_snapshot(
            order,
            boxes_total=3,
            cargoes=[
                {"gm_barcode": f"GM-BLOCK-{index}", "items": []}
                for index in range(1, 4)
            ],
        )
        move_requests_before = MoveRequest.objects.count()
        reserves_before = WarehouseReserve.objects.count()

        readiness = shipping_pick_readiness(order)

        self.assertFalse(readiness["can_pick"])
        self.assertIn("данные Ozon", readiness["reason"])
        self.assertEqual(MoveRequest.objects.count(), move_requests_before)
        self.assertEqual(WarehouseReserve.objects.count(), reserves_before)

    def test_orders_before_rollout_keep_previous_workflow(self):
        order = self._order(
            number="OTG-STRICT-OZON-OLD",
            created_at="2026-08-11T05:59:59+00:00",
        )
        self._save_snapshot(
            order,
            boxes_total=3,
            cargoes=[
                {"gm_barcode": f"GM-OLD-{index}", "items": []}
                for index in range(1, 4)
            ],
        )

        self.assertEqual(_strict_ozon_gm_intake_errors(order), [])


class MarketplaceItemBindingTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Binding Client")
        self.wb = Market.objects.create(id=9872, name="WB Binding")
        self.ozon = Market.objects.create(id=9873, name="Ozon Binding")

    def _order(self, *, number: str, market: Market, barcode: str = "460000000001") -> ShippingOrder:
        order = ShippingOrder.objects.create(
            number=number,
            agency=self.agency,
            marketplace=market,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            status=ShippingOrder.STATUS_PICKING,
        )
        ShippingOrderItem.objects.create(
            order=order,
            sku_code="SKU-BINDING",
            name="Binding item",
            barcode=barcode,
            qty_requested=20,
        )
        return order

    @staticmethod
    def _box(*, code: str, barcode: str = "460000000001", qty: int = 20) -> dict:
        return {
            "box_code": code,
            "items": [
                {
                    "sku_code": "ANY-TRANSPORT-BOX",
                    "barcode": barcode,
                    "qty": qty,
                }
            ],
        }

    def test_wb_binding_uses_product_barcode_and_quantity_not_box_number(self):
        order = self._order(number="OTG-BIND-WB-1", market=self.wb)

        result = validate_marketplace_item_binding(
            order,
            boxes=[self._box(code="FULLBOX-BOX-WITH-ANY-NUMBER")],
        )

        self.assertTrue(result["ready"], result["errors"])
        self.assertEqual(result["request_totals"], {"460000000001": 20})
        self.assertEqual(result["fact_totals"], {"460000000001": 20})

    def test_wb_binding_rejects_other_product_barcode_with_same_quantity(self):
        order = self._order(number="OTG-BIND-WB-2", market=self.wb)

        result = validate_marketplace_item_binding(
            order,
            boxes=[
                self._box(
                    code="FULLBOX-BOX-2",
                    barcode="460000000999",
                )
            ],
        )

        self.assertFalse(result["ready"])
        self.assertTrue(any("460000000001" in error for error in result["errors"]))
        self.assertTrue(any("460000000999" in error for error in result["errors"]))

    def test_ozon_binding_matches_request_gm_and_actual_box_by_barcode_quantity(self):
        order = self._order(number="OTG-BIND-OZON-1", market=self.ozon)
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            payload={
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-001",
                        "items": [
                            {
                                "offer_id": "SKU-BINDING",
                                "barcode": "460000000001",
                                "quantity": 20,
                            }
                        ],
                    }
                ]
            },
        )

        result = validate_marketplace_item_binding(
            order,
            boxes=[self._box(code="FULLBOX-BOX-77")],
        )

        self.assertTrue(result["ready"], result["errors"])
        self.assertEqual(result["gm_totals"], {"460000000001": 20})
        self.assertEqual(
            result["gm_box_assignments"],
            {"GM-001": "FULLBOX-BOX-77"},
        )

    def test_ozon_binding_does_not_replace_missing_product_barcode_with_offer_id(self):
        order = self._order(number="OTG-BIND-OZON-2", market=self.ozon)
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="create",
            agency=self.agency,
            payload={
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-002",
                        "items": [
                            {
                                "offer_id": "SKU-BINDING",
                                "barcode": "",
                                "quantity": 20,
                            }
                        ],
                    }
                ]
            },
        )

        result = validate_marketplace_item_binding(
            order,
            boxes=[self._box(code="FULLBOX-BOX-78")],
        )

        self.assertFalse(result["ready"])
        self.assertTrue(
            any("не указан товарный штрих-код" in error for error in result["errors"])
        )

    def test_ozon_binding_allows_same_totals_when_gm_box_compositions_differ(self):
        order = self._order(number="OTG-BIND-OZON-3", market=self.ozon)
        gm_cargoes = [
            {
                "gm_barcode": "GM-003",
                "items": [{"barcode": "460000000001", "quantity": 10}],
            },
            {
                "gm_barcode": "GM-004",
                "items": [{"barcode": "460000000001", "quantity": 10}],
            },
        ]

        result = validate_marketplace_item_binding(
            order,
            boxes=[self._box(code="FULLBOX-BOX-80", qty=20)],
            gm_cargoes=gm_cargoes,
        )

        self.assertEqual(result["request_totals"], result["gm_totals"])
        self.assertEqual(result["request_totals"], result["fact_totals"])
        self.assertTrue(result["ready"], result["errors"])
        self.assertEqual(result["gm_box_assignments"], {})

    def test_ozon_gm_visual_assignment_keeps_all_available_exact_pairs(self):
        order = self._order(number="OTG-BIND-OZON-PARTIAL", market=self.ozon)
        gm_cargoes = [
            {
                "gm_barcode": f"GM-{index}",
                "items": [{"barcode": "460000000001", "quantity": 5}],
            }
            for index in range(1, 4)
        ]
        fullbox_boxes = [
            self._box(code="FULLBOX-BOX-1", qty=5),
            self._box(code="FULLBOX-BOX-2", qty=5),
            self._box(code="FULLBOX-BOX-10", qty=10),
        ]

        rows = _ozon_gm_rows_with_exact_fullbox_match(
            order,
            gm_cargoes,
            fullbox_boxes=fullbox_boxes,
        )

        self.assertEqual(
            [row["fullbox_box_code"] for row in rows],
            ["FULLBOX-BOX-1", "FULLBOX-BOX-2", ""],
        )
        self.assertEqual(
            [row["assignment_status"] for row in rows],
            ["matched", "matched", "unassigned"],
        )

    def test_ozon_request_binding_blocks_missing_gm_product_before_reserve(self):
        order = self._order(number="OTG-BIND-OZON-REQ", market=self.ozon)
        ShippingOrder.objects.filter(pk=order.pk).update(
            created_at=datetime.fromisoformat("2026-07-31T08:00:00+00:00"),
            expected_boxes=1,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        order.refresh_from_db()
        gm_cargoes = [
            {
                "gm_barcode": "GM-REQ-1",
                "items": [{"offer_id": "SKU-BINDING", "barcode": "460000000001", "quantity": 20}],
            },
            {
                "gm_barcode": "GM-REQ-2",
                "items": [{"offer_id": "polka001", "barcode": "2038273624348", "quantity": 12}],
            },
        ]

        result = validate_ozon_request_binding(
            order,
            gm_cargoes=gm_cargoes,
            expected_boxes=order.expected_boxes,
            ozon_boxes_total=2,
        )

        self.assertFalse(result["ready"])
        self.assertTrue(any("Количество коробов не совпадает: Ozon 2" in error for error in result["errors"]))
        self.assertTrue(any("ШК 2038273624348" in error for error in result["errors"]))
        self.assertTrue(
            enforced_ozon_request_binding_errors(
                order,
                gm_cargoes=gm_cargoes,
                expected_boxes=order.expected_boxes,
                ozon_boxes_total=2,
            )
        )
        with self.assertRaisesMessage(ValidationError, "Нельзя подтвердить резерв"):
            reserve_order(
                order,
                ozon_api_meta={
                    "gm_cargoes": gm_cargoes,
                    "boxes_total": 2,
                },
            )
        order.items.get().refresh_from_db()
        self.assertEqual(order.items.get().qty_reserved, 0)
        self.assertFalse(WarehouseReserve.objects.filter(context_id=order.number).exists())

    def test_ozon_request_binding_can_ignore_source_box_count_but_not_item_content(self):
        order = self._order(number="OTG-BIND-OZON-SOURCE-BOXES", market=self.ozon)
        ShippingOrder.objects.filter(pk=order.pk).update(
            created_at=datetime.fromisoformat("2026-07-31T08:00:00+00:00"),
            expected_boxes=8,
            status=ShippingOrder.STATUS_SUBMITTED,
        )
        order.refresh_from_db()
        gm_cargoes = [
            {
                "gm_barcode": f"GM-SOURCE-{index}",
                "items": [
                    {
                        "offer_id": "SKU-BINDING",
                        "barcode": "460000000001",
                        "quantity": quantity,
                    }
                ],
            }
            for index, quantity in enumerate((4, 4, 3, 3, 3, 3), start=1)
        ]

        strict = validate_ozon_request_binding(
            order,
            gm_cargoes=gm_cargoes,
            expected_boxes=order.expected_boxes,
            ozon_boxes_total=6,
        )
        source_agnostic = validate_ozon_request_binding(
            order,
            gm_cargoes=gm_cargoes,
            expected_boxes=order.expected_boxes,
            ozon_boxes_total=6,
            compare_source_box_count=False,
        )

        self.assertFalse(strict["ready"])
        self.assertTrue(any("Количество коробов не совпадает" in error for error in strict["errors"]))
        self.assertTrue(source_agnostic["ready"], source_agnostic["errors"])

        gm_cargoes[0]["items"][0]["quantity"] = 5
        content_mismatch = validate_ozon_request_binding(
            order,
            gm_cargoes=gm_cargoes,
            expected_boxes=order.expected_boxes,
            ozon_boxes_total=6,
            compare_source_box_count=False,
        )
        self.assertFalse(content_mismatch["ready"])
        self.assertTrue(any("ШК 460000000001" in error for error in content_mismatch["errors"]))

    def test_ozon_request_binding_is_rechecked_before_pick_without_writes(self):
        order = self._order(number="OTG-BIND-OZON-PICK", market=self.ozon)
        ShippingOrder.objects.filter(pk=order.pk).update(
            created_at=datetime.fromisoformat("2026-07-31T08:00:00+00:00"),
            expected_boxes=1,
            status=ShippingOrder.STATUS_STOREKEEPER_ACCEPTED,
        )
        order.refresh_from_db()
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=order.number,
            action="update",
            agency=self.agency,
            payload={
                "ozon_boxes_total": 1,
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-PICK-1",
                        "items": [{"barcode": "2038273624348", "quantity": 20}],
                    }
                ],
            },
        )
        move_requests_before = MoveRequest.objects.count()
        reserves_before = WarehouseReserve.objects.count()

        readiness = shipping_pick_readiness(order)

        self.assertFalse(readiness["can_pick"])
        self.assertIn("ШК ГМ Ozon", readiness["reason"])
        self.assertIn("Остатки и задания не изменены", readiness["reason"])
        with self.assertRaisesMessage(ValidationError, "Нельзя начать отбор"):
            create_pick_tasks(order)
        order.refresh_from_db()
        self.assertEqual(order.status, ShippingOrder.STATUS_STOREKEEPER_ACCEPTED)
        self.assertEqual(MoveRequest.objects.count(), move_requests_before)
        self.assertEqual(WarehouseReserve.objects.count(), reserves_before)

    def test_enforcement_applies_only_to_orders_created_after_rollout(self):
        old_order = self._order(number="OTG-BIND-WB-OLD", market=self.wb)
        ShippingOrder.objects.filter(pk=old_order.pk).update(
            created_at=datetime.fromisoformat("2026-07-30T07:59:59+00:00")
        )
        old_order.refresh_from_db()
        new_order = self._order(number="OTG-BIND-WB-NEW", market=self.wb)
        ShippingOrder.objects.filter(pk=new_order.pk).update(
            created_at=datetime.fromisoformat("2026-07-30T08:00:00+00:00")
        )
        new_order.refresh_from_db()
        mismatched_box = self._box(
            code="FULLBOX-BOX-79",
            barcode="460000000999",
        )

        self.assertEqual(
            enforced_marketplace_item_binding_errors(
                old_order,
                boxes=[mismatched_box],
            ),
            [],
        )
        self.assertTrue(
            enforced_marketplace_item_binding_errors(
                new_order,
                boxes=[mismatched_box],
            )
        )


class ShippingTruthClaimIsolationTests(TestCase):
    def setUp(self):
        self.agency = Agency.objects.create(agn_name="Truth Claim Client")
        self.order = ShippingOrder.objects.create(
            number="OTG-TRUTH-1",
            agency=self.agency,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=2,
        )
        self.other_order = ShippingOrder.objects.create(
            number="OTG-TRUTH-2",
            agency=self.agency,
            status=ShippingOrder.STATUS_PICKING,
        )
        self.location = WarehouseWritePathService.ensure_location(
            warehouse_code="MSK",
            zone_code="OTG",
        )
        self.move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        self.move_task = MoveTask.objects.create(
            request=self.move_request,
            pallet_code="SOURCE-PALLET",
            to_zone="OTG",
            status=MoveTask.STATUS_DONE,
        )

    def _snapshot(self, *, box_code: str, context_type: str, context_id: str):
        container = WarehouseContainer.objects.create(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_BOX,
            container_code=box_code,
            current_location=self.location,
        )
        event = WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="otg_arrived",
            stock_context_type=context_type,
            stock_context_id=context_id,
            container=container,
            to_location=self.location,
            to_zone_code="OTG",
            qty=10,
            occurred_at=timezone.now(),
        )
        return WarehouseStockSnapshot.objects.create(
            agency=self.agency,
            source_context_type="receiving",
            source_context_id="PR-TRUTH",
            sku_code="SKU-TRUTH",
            barcode="460000000180",
            goods_type="gv",
            qty=10,
            available_qty=0,
            container=container,
            container_code=box_code,
            location=self.location,
            zone_code="OTG",
            warehouse_state_code=WarehouseStateCode.IN_OTG.value,
            last_event=event,
        )

    def _claim(self, *, box_code: str, claim_kind: str):
        return BoxClaim.objects.create(
            agency=self.agency,
            move_task=self.move_task,
            box_code=box_code,
            claim_kind=claim_kind,
            shipping_order_id=self.order.number,
            shipping_order_pk=self.order.pk,
            status=BoxClaim.STATUS_DELIVERED,
            delivered_at=timezone.now(),
        )

    def test_truth_excludes_partial_source_and_claim_reused_by_another_shipping(self):
        self._snapshot(
            box_code="PARTIAL-SOURCE",
            context_type="shipping",
            context_id=self.other_order.number,
        )
        self._claim(
            box_code="PARTIAL-SOURCE",
            claim_kind=BoxClaim.KIND_PARTIAL,
        )
        self._snapshot(
            box_code="FULL-REUSED",
            context_type="shipping",
            context_id=self.other_order.number,
        )
        self._claim(
            box_code="FULL-REUSED",
            claim_kind=BoxClaim.KIND_BOX,
        )
        self._snapshot(
            box_code="FULL-LEGACY",
            context_type="receiving",
            context_id="PR-TRUTH",
        )
        self._claim(
            box_code="FULL-LEGACY",
            claim_kind=BoxClaim.KIND_BOX,
        )
        self._snapshot(
            box_code="PARTIAL-RESULT",
            context_type="shipping",
            context_id=self.order.number,
        )

        truth = ShippingTruthService.for_order(self.order)

        self.assertEqual(
            truth.delivered_claim_box_codes,
            {"FULL-LEGACY", "FULL-REUSED"},
        )
        self.assertEqual(
            truth.warehouse_box_codes,
            {"full-legacy", "partial-result"},
        )


class ShippingPackingDraftNumberTests(SimpleTestCase):
    def test_restore_renumbers_used_ui_pallet_and_drops_empty_duplicate(self):
        delivered_boxes = [
            {"box_code": "BOX-1", "qty": 10},
            {"box_code": "BOX-2", "qty": 20},
        ]
        draft_state = {
            "boxes": [
                {"row_key": "ROW-1", "code": "BOX-1", "qty": 10, "pallet_code": "PAL-2"},
                {"row_key": "ROW-2", "code": "BOX-2", "qty": 20, "pallet_code": "PAL-2"},
            ],
            "pallets": [
                {"code": "PAL-2", "label": "Паллета 2", "sealed": True},
                {"code": "PAL-3", "label": "Паллета 2", "sealed": False},
            ],
        }

        state = _shipping_packing_initial_state_from_draft(delivered_boxes, draft_state)

        self.assertEqual(
            [box["pallet_code"] for box in state["initial_boxes"]],
            ["PAL-1", "PAL-1"],
        )
        self.assertEqual(
            state["initial_pallets"],
            [{"code": "PAL-1", "label": "Паллета 1", "sealed": True}],
        )


class ShippingDiscrepancyV2Tests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.storekeeper = user_model.objects.create_user(
            username="disc_storekeeper",
            password="pwd",
        )
        self.manager = user_model.objects.create_user(
            username="disc_manager",
            password="pwd",
        )
        self.head_manager = user_model.objects.create_user(
            username="disc_head_manager",
            password="pwd",
        )
        Employee.objects.create(
            full_name="Discrepancy Manager",
            user=self.manager,
            role="manager",
        )
        Employee.objects.create(
            full_name="Discrepancy Head",
            user=self.head_manager,
            role="head_manager",
        )
        self.agency = Agency.objects.create(agn_name="Discrepancy Client")
        self.ozon = Market.objects.create(id=9981, name="Ozon")
        self.order = ShippingOrder.objects.create(
            number="OTG-DISC-V2",
            agency=self.agency,
            marketplace=self.ozon,
            delivery_type=ShippingOrder.DELIVERY_MARKETPLACE,
            status=ShippingOrder.STATUS_PICKING,
            expected_boxes=3,
        )
        self.item = ShippingOrderItem.objects.create(
            order=self.order,
            sku_code="SKU-DISC",
            name="Discrepancy item",
            barcode="460000000170",
            qty_requested=30,
            comment="Коробов: 3; кратность: 10",
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=self.order.number,
            action="create",
            agency=self.agency,
            payload={
                "gm_cargoes": [
                    {
                        "gm_barcode": "GM-DISC",
                        "items": [
                            {
                                "barcode": "460000000170",
                                "quantity": 35,
                            }
                        ],
                    }
                ]
            },
        )
        self.boxes = [
            {
                "box_code": f"DISC-BOX-{index}",
                "items": [
                    {
                        "barcode": "460000000170",
                        "qty": 10,
                    }
                ],
            }
            for index in range(1, 4)
        ]
        self.truth = SimpleNamespace(
            warehouse_box_codes={row["box_code"] for row in self.boxes},
        )

    def test_approve_keeps_request_and_stock_plan_immutable(self):
        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=self.truth,
            ):
                payload = request_shipping_discrepancy(
                    self.order,
                    user=self.storekeeper,
                    storekeeper_comment="Ozon 35, OTG 30",
                )
                approved = approve_shipping_discrepancy(
                    self.order,
                    user=self.manager,
                )

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        self.assertEqual(payload["workflow_version"], 2)
        self.assertEqual(approved["decision"], "approve_mismatch")
        self.assertEqual(self.order.expected_boxes, 3)
        self.assertEqual(self.item.qty_requested, 30)
        self.assertEqual(self.order.shipping_discrepancy_status, "approved")
        self.assertEqual(Task.objects.filter(route=f"/shipping/{self.order.pk}/").count(), 2)
        self.assertFalse(
            Task.objects.filter(
                route=f"/shipping/{self.order.pk}/",
            ).exclude(status="done").exists()
        )
        self.assertEqual(
            enforced_marketplace_item_binding_errors(
                self.order,
                boxes=self.boxes,
            ),
            [],
        )

    def test_approved_shortage_closes_as_partial_and_persists_plan_fact(self):
        ShippingOrder.objects.filter(pk=self.order.pk).update(
            created_at=timezone.make_aware(datetime(2026, 7, 29, 8, 0)),
        )
        self.order.refresh_from_db()
        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=self.truth,
            ):
                request_shipping_discrepancy(
                    self.order,
                    user=self.storekeeper,
                    storekeeper_comment="Ozon 35, OTG 30",
                )
                approve_shipping_discrepancy(self.order, user=self.manager)

        self.order.refresh_from_db()
        self.item.qty_reserved = 30
        self.item.save(update_fields=["qty_reserved", "updated_at"])
        with patch(
            "shipping.packing._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            ship_order(self.order, self.manager)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        final_truth = self.order.shipping_discrepancy_payload["final_shipping_truth"]
        self.assertEqual(self.order.status, ShippingOrder.STATUS_PARTIAL)
        self.assertEqual(self.item.qty_requested, 30)
        self.assertEqual(self.item.qty_shipped, 30)
        self.assertEqual(final_truth["planned_total_qty"], 35)
        self.assertEqual(final_truth["fact_total_qty"], 30)
        self.assertEqual(final_truth["shortage_total_qty"], 5)
        self.assertFalse(final_truth["matches_plan"])
        self.assertEqual(
            shipping_final_truth_rows(self.order),
            [
                {
                    "sku_code": "SKU-DISC",
                    "name": "Discrepancy item",
                    "size": "",
                    "barcode": "460000000170",
                    "goods_type": "",
                    "qty_requested": 35,
                    "qty_reserved": 30,
                    "qty_shipped": 30,
                    "shortage_qty": 5,
                    "excess_qty": 0,
                }
            ],
        )
        dispatch_payload = shipping_dispatch_payload(
            self.order,
            payload={"act_items": [{"barcode": "STALE"}]},
        )
        self.assertEqual(dispatch_payload["act_items"][0]["qty_requested"], 35)
        self.assertEqual(dispatch_payload["act_items"][0]["qty_shipped"], 30)
        self.assertEqual(dispatch_payload["act_items"][0]["shortage_qty"], 5)

    def test_first_manager_decision_is_locked(self):
        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=self.truth,
            ):
                request_shipping_discrepancy(
                    self.order,
                    user=self.storekeeper,
                    storekeeper_comment="Ozon 35, OTG 30",
                )
                approve_shipping_discrepancy(self.order, user=self.manager)
                with self.assertRaisesMessage(
                    ValidationError,
                    "Решение по расхождению уже принято",
                ):
                    reject_shipping_discrepancy(
                        self.order,
                        user=self.head_manager,
                        correction_mode="piece",
                    )

    def test_piece_pick_preview_supports_barcode_missing_from_order_items(self):
        from otg_reachtruck.services import get_otg_discrepancy_pick_preview

        missing_barcode = "460000000999"
        catalog = {
            missing_barcode.casefold(): [
                {
                    "code": "DISC-MISSING-BOX",
                    "pallet_code": "DISC-MISSING-PALLET",
                    "location": {"zone": "OS", "row": 1, "section": 1, "tier": 1, "cell": 1},
                    "location_label": "OS-1/1-1-1",
                    "items": [
                        {
                            "sku": "SKU-MISSING",
                            "name": "Missing marketplace item",
                            "size": "0",
                            "barcode": missing_barcode,
                            "goods_type": "gv",
                            "qty": 64,
                        }
                    ],
                    "target_qty": 64,
                    "total_qty": 64,
                    "is_open": False,
                }
            ]
        }

        with patch(
            "otg_reachtruck.services._discrepancy_stock_catalog",
            return_value=catalog,
        ):
            preview = get_otg_discrepancy_pick_preview(
                order=self.order,
                shortage_rows=[{"barcode": missing_barcode, "missing_qty": 32}],
                correction_mode="piece",
            )

        self.assertTrue(preview["can_create"])
        self.assertEqual(preview["requested_qty"], 32)
        self.assertEqual(preview["requested_qty_by_item_id"], {})
        self.assertEqual(
            preview["detached_move_items"],
            [
                {
                    "sku_code": "SKU-MISSING",
                    "barcode": missing_barcode,
                    "goods_type": "gv",
                    "qty_requested": 32,
                }
            ],
        )
        self.assertEqual(preview["plan_rows"][0]["pick_qty"], 32)
        self.assertEqual(preview["demand_payloads"][0]["item_quantities_per_box"], {})

    def test_detached_discrepancy_item_creates_move_request_line(self):
        from otg_reachtruck.services import _request_items_for_move

        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
        )

        request_item_map = _request_items_for_move(
            move_request,
            self.order,
            items=[],
            detached_move_items=[
                {
                    "sku_code": "SKU-MISSING",
                    "barcode": "460000000999",
                    "goods_type": "gv",
                    "qty_requested": 32,
                }
            ],
        )

        self.assertEqual(request_item_map, {})
        move_item = MoveRequestItem.objects.get(request=move_request)
        self.assertEqual(move_item.sku_code, "SKU-MISSING")
        self.assertEqual(move_item.barcode, "460000000999")
        self.assertEqual(move_item.qty_requested, 32)
        self.assertEqual(move_item.qty_planned, 32)

    def _prepare_pick_confirmation(self, *, fingerprint="stored-plan"):
        payload = {
            "workflow_version": 2,
            "status": "pick_confirmation",
            "correction_mode": "piece",
            "snapshot": {
                "missing_rows": [
                    {
                        "barcode": "460000000170",
                        "goods_type": "gv",
                        "missing_qty": 10,
                    }
                ],
            },
            "correction_plan": {
                "can_create": True,
                "plan_fingerprint": fingerprint,
                "plan_rows": [{"box_code": "OLD-BOX", "pick_qty": 10}],
            },
        }
        self.order.shipping_discrepancy_status = "pick_confirmation"
        self.order.shipping_discrepancy_payload = payload
        self.order.save(
            update_fields=[
                "shipping_discrepancy_status",
                "shipping_discrepancy_payload",
                "updated_at",
            ]
        )
        task = Task.objects.create(
            title=f"СРОЧНО: подтвердить добор {self.order.number}",
            route=f"/shipping/{self.order.pk}/",
            created_by=self.storekeeper,
            status="in_progress",
        )
        return task

    def test_confirm_refreshes_changed_plan_without_creating_reachtruck_task(self):
        storekeeper_task = self._prepare_pick_confirmation()
        current_plan = {
            "can_create": True,
            "plan_fingerprint": "current-plan",
            "plan_rows": [{"box_code": "CURRENT-BOX", "pick_qty": 10}],
        }
        move_request_count = MoveRequest.objects.count()
        move_task_count = MoveTask.objects.count()
        reserve_count = WarehouseReserve.objects.count()
        stock_count = WarehouseStockSnapshot.objects.count()

        with patch(
            "otg_reachtruck.services.get_otg_discrepancy_pick_preview",
            return_value=current_plan,
        ):
            with patch(
                "otg_reachtruck.services.create_otg_shipping_discrepancy_pick",
            ) as create_pick:
                result = confirm_shipping_discrepancy_pick(
                    self.order,
                    user=self.storekeeper,
                    requested_by_role="storekeeper",
                )

        create_pick.assert_not_called()
        self.order.refresh_from_db()
        storekeeper_task.refresh_from_db()
        self.assertTrue(result["plan_refreshed"])
        self.assertEqual(self.order.shipping_discrepancy_status, "pick_confirmation")
        self.assertEqual(
            self.order.shipping_discrepancy_payload["correction_plan"],
            current_plan,
        )
        self.assertEqual(storekeeper_task.status, "in_progress")
        self.assertEqual(MoveRequest.objects.count(), move_request_count)
        self.assertEqual(MoveTask.objects.count(), move_task_count)
        self.assertEqual(WarehouseReserve.objects.count(), reserve_count)
        self.assertEqual(WarehouseStockSnapshot.objects.count(), stock_count)
        self.assertTrue(
            OrderAuditEntry.objects.filter(
                order_type="shipping",
                order_id=self.order.number,
                action="shipping_disc_plan_refreshed",
            ).exists()
        )

    def test_confirm_unchanged_plan_uses_existing_task_creation_path(self):
        storekeeper_task = self._prepare_pick_confirmation(fingerprint="current-plan")
        current_plan = {
            "can_create": True,
            "plan_fingerprint": "current-plan",
            "plan_rows": [{"box_code": "CURRENT-BOX", "pick_qty": 10}],
        }

        with patch(
            "otg_reachtruck.services.get_otg_discrepancy_pick_preview",
            return_value=current_plan,
        ):
            with patch(
                "otg_reachtruck.services.create_otg_shipping_discrepancy_pick",
                return_value=(SimpleNamespace(), ["MOVE-1"], current_plan),
            ) as create_pick:
                result = confirm_shipping_discrepancy_pick(
                    self.order,
                    user=self.storekeeper,
                    requested_by_role="storekeeper",
                )

        create_pick.assert_called_once()
        self.order.refresh_from_db()
        storekeeper_task.refresh_from_db()
        self.assertNotIn("plan_refreshed", result)
        self.assertEqual(result["additional_pick_move_ids"], ["MOVE-1"])
        self.assertEqual(self.order.shipping_discrepancy_status, "pickup_required")
        self.assertEqual(storekeeper_task.status, "done")

    def test_completed_supplement_move_creates_idempotent_packing_stage_without_stock_writes(self):
        employee = Employee.objects.create(
            full_name="Discrepancy Storekeeper",
            user=self.storekeeper,
            role="storekeeper",
            is_active=True,
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        move_task = MoveTask.objects.create(
            request=move_request,
            status=MoveTask.STATUS_DONE,
            qty_done=5,
        )
        self.order.shipping_discrepancy_status = "pickup_required"
        self.order.shipping_discrepancy_payload = {
            "workflow_version": 2,
            "status": "pickup_required",
            "decision": "supplement",
            "additional_pick_move_ids": [str(move_task.id)],
        }
        self.order.save(
            update_fields=[
                "shipping_discrepancy_status",
                "shipping_discrepancy_payload",
                "updated_at",
            ]
        )
        stock_count = WarehouseStockSnapshot.objects.count()
        reserve_count = WarehouseReserve.objects.count()

        with patch(
            "shipping.discrepancy._shipping_loose_items",
            return_value=[{"barcode": "460000000170", "qty": 5}],
        ):
            first = mark_supplement_awaiting_packing(
                completed_move_id=str(move_task.id),
                order_pk=self.order.pk,
                order_number=self.order.number,
                user=self.storekeeper,
            )
            second = mark_supplement_awaiting_packing(
                completed_move_id=str(move_task.id),
                order_pk=self.order.pk,
                order_number=self.order.number,
                user=self.storekeeper,
            )

        self.order.refresh_from_db()
        task = Task.objects.get(
            route=f"/shipping/{self.order.pk}/packing/loose/",
            title=f"СРОЧНО: упаковать добор {self.order.number}",
        )
        self.assertEqual(first["workflow_stage"], "awaiting_supplement_packing")
        self.assertEqual(second["workflow_stage"], "awaiting_supplement_packing")
        self.assertEqual(self.order.shipping_discrepancy_status, "pickup_required")
        self.assertEqual(self.order.shipping_discrepancy_payload["supplement_packing_qty"], 5)
        self.assertEqual(task.assigned_to, employee)
        self.assertEqual(task.priority, "urgent")
        self.assertEqual(task.status, "in_progress")
        self.assertIn("Количество к добору: 5 шт.", task.description)
        self.assertEqual(
            OrderAuditEntry.objects.filter(
                order_type="shipping",
                order_id=self.order.number,
                action="ship_disc_supp_awaiting_packing",
            ).count(),
            1,
        )
        self.assertEqual(WarehouseStockSnapshot.objects.count(), stock_count)
        self.assertEqual(WarehouseReserve.objects.count(), reserve_count)

        with patch(
            "shipping.discrepancy.validate_marketplace_item_binding",
            return_value={"kind": "ozon"},
        ):
            with patch(
                "shipping.discrepancy.shipping_supplement_fact_is_complete",
                return_value=False,
            ):
                context = shipping_discrepancy_context(self.order)
        self.assertTrue(context["is_awaiting_supplement_packing"])
        self.assertEqual(context["supplement_packing_qty"], 5)
        self.assertEqual(context["summary_title"], "Добор доставлен — требуется упаковка")

    def test_incomplete_or_unrelated_supplement_move_does_not_create_packing_stage(self):
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_IN_PROGRESS,
        )
        completed_task = MoveTask.objects.create(
            request=move_request,
            status=MoveTask.STATUS_DONE,
            qty_done=5,
        )
        pending_task = MoveTask.objects.create(
            request=move_request,
            status=MoveTask.STATUS_IN_PROGRESS,
        )
        self.order.shipping_discrepancy_status = "pickup_required"
        self.order.shipping_discrepancy_payload = {
            "workflow_version": 2,
            "status": "pickup_required",
            "decision": "supplement",
            "additional_pick_move_ids": [str(completed_task.id), str(pending_task.id)],
        }
        self.order.save(
            update_fields=[
                "shipping_discrepancy_status",
                "shipping_discrepancy_payload",
                "updated_at",
            ]
        )

        with patch(
            "shipping.discrepancy._shipping_loose_items",
            return_value=[{"barcode": "460000000170", "qty": 5}],
        ):
            incomplete = mark_supplement_awaiting_packing(
                completed_move_id=str(completed_task.id),
                order_pk=self.order.pk,
                user=self.storekeeper,
            )
            unrelated = mark_supplement_awaiting_packing(
                completed_move_id="999999",
                order_pk=self.order.pk,
                user=self.storekeeper,
            )

        self.order.refresh_from_db()
        self.assertEqual(incomplete, {})
        self.assertEqual(unrelated, {})
        self.assertNotIn("workflow_stage", self.order.shipping_discrepancy_payload)
        self.assertFalse(
            Task.objects.filter(
                route=f"/shipping/{self.order.pk}/packing/loose/",
            ).exists()
        )

    def test_confirm_rejects_unavailable_live_plan_without_writes(self):
        storekeeper_task = self._prepare_pick_confirmation()
        original_payload = dict(self.order.shipping_discrepancy_payload)

        with patch(
            "otg_reachtruck.services.get_otg_discrepancy_pick_preview",
            return_value={"can_create": False, "reason": "Подходящего короба больше нет."},
        ):
            with patch(
                "otg_reachtruck.services.create_otg_shipping_discrepancy_pick",
            ) as create_pick:
                with self.assertRaisesMessage(ValidationError, "Подходящего короба больше нет"):
                    confirm_shipping_discrepancy_pick(
                        self.order,
                        user=self.storekeeper,
                        requested_by_role="storekeeper",
                    )

        create_pick.assert_not_called()
        self.order.refresh_from_db()
        storekeeper_task.refresh_from_db()
        self.assertEqual(self.order.shipping_discrepancy_status, "pick_confirmation")
        self.assertEqual(self.order.shipping_discrepancy_payload, original_payload)
        self.assertEqual(storekeeper_task.status, "in_progress")

    def test_confirm_action_warns_when_plan_was_refreshed(self):
        request = RequestFactory().post(
            f"/shipping/{self.order.pk}/",
            {"action": "confirm_shipping_discrepancy_pick"},
        )
        request.user = self.storekeeper
        permissions = ShippingDetailActionPermissions(
            can_edit_items=False,
            can_submit_for_approval=False,
            can_manager_approve=False,
            can_manager_reopen=False,
            can_storekeeper_accept=False,
            can_storekeeper_pick=False,
            can_confirm_shipping_discrepancy_pick=True,
        )

        with patch(
            "shipping.actions.confirm_shipping_discrepancy_pick",
            return_value={"plan_refreshed": True},
        ):
            result = handle_shipping_detail_action(
                action="confirm_shipping_discrepancy_pick",
                order=self.order,
                request=request,
                role="storekeeper",
                stock_rows_for_order=[],
                permissions=permissions,
                parse_selected_stock_items=lambda _request: ([], []),
                selected_stock_rows_with_boxes=lambda *_args, **_kwargs: [],
                parse_box_count_from_comment=lambda _comment: 0,
                log_update=lambda *_args, **_kwargs: None,
            )

        self.assertEqual(
            result.messages,
            [
                (
                    "warning",
                    "План добора обновлён по текущему складу. "
                    "Проверьте новый короб и нажмите подтверждение повторно.",
                )
            ],
        )

    def test_pick_confirmation_page_marks_refreshed_plan_for_explicit_repeat_confirmation(self):
        html = render_to_string(
            "shipping/detail.html",
            {
                "order": self.order,
                "shipping_discrepancy": {
                    "is_pick_confirmation": True,
                    "correction_mode_label": "Поштучный добор",
                    "correction_plan": {
                        "requested_qty": 32,
                        "plan_rows": [
                            {
                                "barcode": "460000000170",
                                "box_code": "CURRENT-BOX",
                                "pallet_code": "CURRENT-PALLET",
                                "location": "OS-1/1-1-1",
                                "pick_qty": 32,
                                "is_open": True,
                            }
                        ],
                    },
                    "payload": {
                        "correction_plan_refreshed_at": "2026-07-31T19:35:37+03:00",
                    },
                },
                "can_write": True,
                "can_confirm_shipping_discrepancy_pick": True,
            },
        )

        self.assertIn("План добора изменился", html)
        self.assertIn("Задание ричтраку ещё не создано", html)
        self.assertIn("Итого к добору: 32 шт.", html)
        self.assertIn("CURRENT-BOX", html)
        self.assertIn(
            "Подтвердить обновлённый план и создать задание ричтраку",
            html,
        )

    def test_pick_confirmation_page_uses_clear_initial_confirmation_label(self):
        html = render_to_string(
            "shipping/detail.html",
            {
                "order": self.order,
                "shipping_discrepancy": {
                    "is_pick_confirmation": True,
                    "correction_mode_label": "Поштучный добор",
                    "correction_plan": {"plan_rows": []},
                    "payload": {},
                },
                "can_write": True,
                "can_confirm_shipping_discrepancy_pick": True,
            },
        )

        self.assertNotIn("План добора изменился", html)
        self.assertIn("Подтвердить план и создать задание ричтраку", html)
        self.assertIn('type="submit"', html)
        self.assertIn('data-submitting-label="Создаём задание ричтраку…"', html)

    def test_completed_supplement_resolves_discrepancy_and_updates_box_plan(self):
        final_boxes = [
            {
                "box_code": f"DISC-BOX-{index}",
                "items": [
                    {
                        "barcode": "460000000170",
                        "qty": 10,
                    }
                ],
            }
            for index in range(1, 4)
        ]
        final_boxes.append(
            {
                "box_code": "DISC-SUPP-BOX",
                "items": [
                    {
                        "barcode": "460000000170",
                        "qty": 5,
                    }
                ],
            }
        )
        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=self.order.number,
            action="create",
            agency=self.agency,
            payload={
                "gm_cargoes": [
                    {
                        "gm_barcode": f"GM-DISC-{index}",
                        "items": [
                            {
                                "barcode": "460000000170",
                                "quantity": qty,
                            }
                        ],
                    }
                    for index, qty in enumerate([10, 10, 10, 5], start=1)
                ]
            },
        )
        initial_truth = SimpleNamespace(
            warehouse_box_codes={row["box_code"] for row in self.boxes},
        )
        final_truth = SimpleNamespace(
            warehouse_box_codes={row["box_code"] for row in final_boxes},
        )
        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=initial_truth,
            ):
                payload = request_shipping_discrepancy(
                    self.order,
                    user=self.storekeeper,
                    storekeeper_comment="Ozon 35, OTG 30",
                )

        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        move_task = MoveTask.objects.create(
            request=move_request,
            status=MoveTask.STATUS_DONE,
            qty_done=5,
        )
        payload.update(
            {
                "status": "pickup_required",
                "decision": "supplement",
                "correction_mode": "piece",
                "additional_pick_move_ids": [str(move_task.id)],
                "workflow_stage": "awaiting_supplement_packing",
                "supplement_packing_qty": 5,
            }
        )
        self.order.shipping_discrepancy_status = "pickup_required"
        self.order.shipping_discrepancy_payload = payload
        self.order.save(update_fields=["shipping_discrepancy_status", "shipping_discrepancy_payload", "updated_at"])
        packing_task = Task.objects.create(
            title=f"СРОЧНО: упаковать добор {self.order.number}",
            description="Количество к добору: 5 шт.",
            route=f"/shipping/{self.order.pk}/packing/loose/",
            status="in_progress",
            priority="urgent",
        )

        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=final_boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=final_truth,
            ):
                resolved = resolve_completed_discrepancy_supplement(
                    self.order,
                    user=self.storekeeper,
                )
                self.order.refresh_from_db()
                context = shipping_discrepancy_context(self.order)

        self.assertEqual(resolved["status"], "resolved")
        self.assertEqual(resolved["final_otg_box_count"], 4)
        self.assertEqual(resolved["old_expected_boxes"], 3)
        self.assertEqual(resolved["new_expected_boxes"], 4)
        self.assertEqual(self.order.shipping_discrepancy_status, "resolved")
        self.assertEqual(self.order.shipping_discrepancy_payload["workflow_stage"], "resolved")
        self.assertEqual(self.order.expected_boxes, 4)
        packing_task.refresh_from_db()
        self.assertEqual(packing_task.status, "done")
        self.assertTrue(context["ready_for_palletization"])
        self.assertFalse(context["is_pickup_required"])
        self.assertEqual(
            enforced_marketplace_item_binding_errors(
                self.order,
                boxes=final_boxes,
            ),
            [],
        )

    def test_completed_supplement_closes_as_full_from_scanned_fact(self):
        final_boxes = [*self.boxes]
        final_boxes.append(
            {
                "box_code": "DISC-SUPP-BOX",
                "items": [
                    {
                        "sku_code": "SKU-DISC",
                        "name": "Discrepancy item",
                        "barcode": "460000000170",
                        "qty": 5,
                    }
                ],
            }
        )
        with patch(
            "shipping.discrepancy._shipping_delivered_boxes",
            return_value=self.boxes,
        ):
            with patch(
                "shipping.discrepancy.ShippingTruthService.for_order",
                return_value=self.truth,
            ):
                payload = request_shipping_discrepancy(
                    self.order,
                    user=self.storekeeper,
                    storekeeper_comment="Ozon 35, OTG 30",
                )

        OrderAuditEntry.objects.create(
            order_type="shipping",
            order_id=self.order.number,
            action="create",
            agency=self.agency,
            payload={
                "gm_cargoes": [
                    {
                        "gm_barcode": f"GM-DISC-{index}",
                        "items": [
                            {
                                "barcode": "460000000170",
                                "quantity": qty,
                            }
                        ],
                    }
                    for index, qty in enumerate([10, 10, 10, 5], start=1)
                ]
            },
        )
        move_request = MoveRequest.objects.create(
            context_type=MoveRequest.CONTEXT_MANUAL,
            context_id=str(self.order.pk),
            agency=self.agency,
            destination_zone="OTG",
            status=MoveRequest.STATUS_DONE,
        )
        move_task = MoveTask.objects.create(
            request=move_request,
            status=MoveTask.STATUS_DONE,
            qty_done=5,
        )
        payload.update(
            {
                "status": "resolved",
                "decision": "supplement",
                "correction_mode": "piece",
                "additional_pick_move_ids": [str(move_task.id)],
            }
        )
        self.order.shipping_discrepancy_status = "resolved"
        self.order.shipping_discrepancy_payload = payload
        self.order.expected_boxes = 4
        self.order.save(
            update_fields=[
                "shipping_discrepancy_status",
                "shipping_discrepancy_payload",
                "expected_boxes",
                "updated_at",
            ]
        )
        self.item.qty_reserved = 30
        self.item.save(update_fields=["qty_reserved", "updated_at"])

        with patch(
            "shipping.packing._shipping_delivered_boxes",
            return_value=final_boxes,
        ):
            ship_order(self.order, self.manager)

        self.order.refresh_from_db()
        self.item.refresh_from_db()
        final_truth = self.order.shipping_discrepancy_payload["final_shipping_truth"]
        self.assertEqual(self.order.status, ShippingOrder.STATUS_SHIPPED)
        self.assertEqual(self.item.qty_requested, 30)
        self.assertEqual(self.item.qty_shipped, 30)
        self.assertEqual(final_truth["planned_total_qty"], 35)
        self.assertEqual(final_truth["fact_total_qty"], 35)
        self.assertEqual(final_truth["shortage_total_qty"], 0)
        self.assertTrue(final_truth["matches_plan"])
        self.assertEqual(shipping_final_truth_rows(self.order)[0]["qty_requested"], 35)
        self.assertEqual(shipping_final_truth_rows(self.order)[0]["qty_shipped"], 35)

    def test_loose_packing_resolves_only_new_awaiting_supplement_workflow(self):
        self.order.shipping_discrepancy_status = "pickup_required"
        self.order.shipping_discrepancy_payload = {
            "workflow_version": 2,
            "status": "pickup_required",
            "decision": "supplement",
            "workflow_stage": "awaiting_supplement_packing",
        }
        self.order.save(
            update_fields=[
                "shipping_discrepancy_status",
                "shipping_discrepancy_payload",
                "updated_at",
            ]
        )
        request = RequestFactory().post(
            f"/shipping/{self.order.pk}/packing/loose/",
            {"boxes_json": "[]", "pallets_json": "[]"},
        )
        request.user = self.storekeeper
        request.session = {}
        request._messages = FallbackStorage(request)

        with patch("shipping.views._request_scope", return_value=("staff", "storekeeper", None)):
            with patch("shipping.views._can_storekeeper_manage_packing", return_value=True):
                with patch(
                    "shipping.views._shipping_loose_packing_initial_state",
                    return_value={
                        "loose_items": [{"barcode": "460000000170", "qty": 5}],
                        "existing_box_codes": [],
                        "box_code_base": "SHIP-BOX",
                    },
                ):
                    with patch(
                        "shipping.views.save_shipping_loose_packing",
                        return_value={"saved": True, "errors": [], "completed": False},
                    ):
                        with patch(
                            "shipping.discrepancy.resolve_completed_discrepancy_supplement",
                            return_value={"status": "resolved"},
                        ) as resolve_supplement:
                            with patch(
                                "shipping.services._ozon_loose_packing_gm_plan",
                                return_value={
                                    "required": False,
                                    "ready": True,
                                    "targets": [],
                                    "errors": [],
                                },
                            ):
                                response = handle_shipping_loose_packing_request(
                                    request=request,
                                    pk=self.order.pk,
                                )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, f"/shipping/{self.order.pk}/packing/")
        resolve_supplement.assert_called_once()

        self.order.shipping_discrepancy_payload = {
            "workflow_version": 2,
            "status": "pickup_required",
            "decision": "supplement",
        }
        self.order.save(update_fields=["shipping_discrepancy_payload", "updated_at"])
        old_request = RequestFactory().post(
            f"/shipping/{self.order.pk}/packing/loose/",
            {"boxes_json": "[]", "pallets_json": "[]"},
        )
        old_request.user = self.storekeeper
        old_request.session = {}
        old_request._messages = FallbackStorage(old_request)
        with patch("shipping.views._request_scope", return_value=("staff", "storekeeper", None)):
            with patch("shipping.views._can_storekeeper_manage_packing", return_value=True):
                with patch(
                    "shipping.views._shipping_loose_packing_initial_state",
                    return_value={
                        "loose_items": [{"barcode": "460000000170", "qty": 5}],
                        "existing_box_codes": [],
                        "box_code_base": "SHIP-BOX",
                    },
                ):
                    with patch(
                        "shipping.views.save_shipping_loose_packing",
                        return_value={"saved": True, "errors": [], "completed": False},
                    ):
                        with patch(
                            "shipping.discrepancy.resolve_completed_discrepancy_supplement",
                        ) as old_resolve:
                            with patch(
                                "shipping.services._ozon_loose_packing_gm_plan",
                                return_value={
                                    "required": False,
                                    "ready": True,
                                    "targets": [],
                                    "errors": [],
                                },
                            ):
                                old_response = handle_shipping_loose_packing_request(
                                    request=old_request,
                                    pk=self.order.pk,
                                )

        self.assertEqual(old_response.status_code, 302)
        old_resolve.assert_not_called()

    def test_awaiting_supplement_packing_page_has_mandatory_direct_action(self):
        html = render_to_string(
            "shipping/detail.html",
            {
                "order": self.order,
                "shipping_discrepancy": {
                    "is_awaiting_supplement_packing": True,
                    "supplement_packing_qty": 5,
                },
                "can_write": True,
                "can_storekeeper_manage_packing": True,
            },
        )

        self.assertIn("Обязательный этап: упаковать добор в короба", html)
        self.assertIn("5 шт.", html)
        self.assertIn(f'href="/shipping/{self.order.pk}/packing/loose/"', html)
        self.assertIn("Упаковать добор в короба", html)


class ShippingUploadPreviewUiTests(SimpleTestCase):
    def test_partial_mixed_box_split_uses_available_qty_per_exact_item(self):
        from django.template.loader import get_template

        html = render_to_string(
            "shipping/_stock_picker_table.html",
            {
                "stock_rows": [
                    {
                        "key": "MIX-ONE",
                        "sku_code": "SKU-MIX-ONE",
                        "name": "Товар из смешанного короба",
                        "size": "-",
                        "barcode": "460000000001",
                        "goods_type": "Готовый",
                        "box_qty": 10,
                        "available_boxes": 1,
                        "available_qty": 1,
                        "selected_boxes": 0,
                        "is_mixed_box": True,
                        "mixed_group": "partial-box:42",
                        "mixed_color": 0,
                        "partial_only": True,
                        "split_available_qty": 1,
                    }
                ]
            },
        )
        runtime_source = get_template("shipping/_shipping_form_runtime.html").template.source

        self.assertIn('data-partial-only="1"', html)
        self.assertIn('data-split-available-qty="1"', html)
        self.assertIn("lineRow.dataset.splitAvailableQty", runtime_source)
        self.assertIn("Доступно к отбору", runtime_source)
        self.assertIn("maxQty <= 0 ? 'disabled' : ''", runtime_source)

    def test_ozon_supply_number_rejects_browser_autofill_without_clearing_server_value(self):
        from django.template.loader import get_template

        form = ShippingOrderForm()
        attrs = form.fields["supply_number"].widget.attrs
        runtime_source = get_template("shipping/_shipping_form_runtime.html").template.source

        self.assertEqual(attrs.get("autocomplete"), "off")
        self.assertEqual(attrs.get("data-no-browser-autofill"), "1")
        self.assertIn("clearBrowserRestoredOzonSupplyNumber", runtime_source)
        self.assertIn("ozonSupplyNumberInput.defaultValue", runtime_source)
        self.assertIn("if (!event.persisted)", runtime_source)

    def test_client_shipping_form_groups_mixed_box_and_relaxes_ozon_transit_destination(self):
        from django.template.loader import get_template

        form_source = get_template("shipping/form_client_lk.html").template.source
        runtime_source = get_template("shipping/_shipping_form_runtime.html").template.source
        stock_source = get_template("shipping/_stock_picker_table.html").template.source

        self.assertIn('data-required-when="destination"', form_source)
        self.assertIn('data-ozon-transit-destination-hint', form_source)
        self.assertIn('@media (max-width: 1600px)', form_source)
        self.assertIn('const mixedSelectionGroups = new Map();', runtime_source)
        self.assertIn('destinationWarehouseRequired', runtime_source)
        self.assertIn('Физических коробов:', runtime_source)
        self.assertIn('Смешанный короб', stock_source)

    def test_shipping_detail_compacts_summary_before_history_stacks(self):
        from django.template.loader import get_template

        source = get_template("shipping/detail.html").template.source

        self.assertIn("@media (max-width: 1500px)", source)
        self.assertIn(".layout.with-history", source)
        self.assertIn(".summary-table tr", source)
        self.assertIn("grid-template-columns: minmax(120px, 0.8fr) minmax(0, 1.2fr)", source)

    def test_excel_import_preview_compares_requested_and_loaded_before_submit(self):
        from django.template.loader import get_template

        source = get_template("shipping/form_client_lk.html").template.source

        self.assertIn('id="shipping-import-preview"', source)
        self.assertIn('data-requested="{{ line.requested_qty }}"', source)
        self.assertIn('data-loaded="{{ line.applied_qty }}"', source)
        self.assertIn("const grouped = new Map();", source)
        self.assertIn("const differenceTotal = loadedTotal - requestedTotal;", source)
        self.assertIn('id="shipping-import-preview-edit"', source)
        self.assertIn("shipping-stock-picker-wrap", source)
        self.assertIn(
            '<button class="btn primary" type="button" id="shipping-import-preview-confirm">',
            source,
        )
