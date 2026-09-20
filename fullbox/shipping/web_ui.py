from __future__ import annotations

"""Shipping UI helpers and request handlers."""

from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
import logging
import re
import zipfile
from types import SimpleNamespace

from django.apps import apps
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.http import require_http_methods

from audit.models import OrderAuditEntry, log_order_action
from employees.access import resolve_cabinet_url
from fullbox.container_codes import issue_container_code, issue_pallet_code, normalize_container_kind
from fullbox.order_numbers import format_order_number
from labels.utils import build_print_status_snapshot, refresh_print_agent_printers
from sklad.services.stock_availability import StockAvailabilityService
from sklad.services.warehouse_write_path import WarehouseWritePathService
from sklad.models import WarehouseContainer, WarehouseReserve
from sku.models import Agency, SKU, SKUBarcode

from .actions import ShippingDetailActionPermissions, handle_shipping_detail_action
from .box_splits import encode_partial_box_split, extract_partial_box_split
from .dispatch import (
    shipping_dispatch_stage as _shipping_dispatch_stage,
    sign_dispatch_act_logistician,
    sign_dispatch_act_manager,
)
from .forms import (
    NEXT_DAY_DEADLINE_ERROR,
    NEXT_DAY_DEADLINE_HOUR,
    ShippingTransportNoteForm,
    WORKDAY_END_HOUR,
    WORKDAY_HOURS_ERROR,
    WORKDAY_START_HOUR,
    ShippingOrderForm,
)
from .item_binding import shipping_final_truth_rows
from .marketplace_warehouses import load_marketplace_warehouse_catalog
from .models import ShippingOrder, ShippingOrderAttachment, ShippingOrderItem
from .packing import (
    _shipping_box_row_key,
    _shipping_boxes_from_packing_payload,
    _shipping_delivered_boxes,
    _shipping_loose_packing_initial_state,
    _shipping_manageable_packing_boxes,
    _shipping_packing_initial_state,
    _shipping_packing_summary,
    request_shipping_packing_pallet_removal,
    resolve_shipping_packing_pallet_removal,
    shipping_packing_pending_pallet_removal,
    shipping_packing_slip_meta,
    shipping_packing_slips_data,
    save_shipping_loose_packing,
    save_shipping_packing,
    save_shipping_packing_draft,
)
from .return_act import render_return_act_doc, return_act_doc_filename
from .selectors import (
    active_shipping_attachments as _active_shipping_attachments,
    build_shipping_detail_context,
    build_shipping_dispatch_context,
    shipping_attachment_names as _shipping_attachment_names,
    shipping_expected_box_count,
    shipping_reachtruck_metrics as _shipping_reachtruck_metrics,
    shipping_ui_status_label as _shipping_ui_status_label,
)
from .services import (
    _ozon_api_summary_for_form,
    build_shipping_return_act_response,
    build_shipping_transport_note_docx_response,
    build_shipping_transport_note_pdf_response,
    build_shipping_detail_page_context,
    build_shipping_list_page_context,
    create_ozon_api_shipping_orders_batch,
    download_shipping_attachment,
    ensure_manager_review_task,
    handle_shipping_dispatch_act_request,
    handle_shipping_dispatch_sign_logistician_request,
    handle_shipping_dispatch_sign_manager_request,
    handle_shipping_create_request,
    handle_shipping_documents_request,
    handle_shipping_loose_packing_request,
    handle_shipping_packing_request,
    handle_shipping_packing_pallet_removal_request,
    handle_shipping_packing_pallet_removal_review_request,
    handle_shipping_packing_slips_request,
    handle_shipping_packing_slips_status_request,
    next_shipping_number,
    order_payload,
    reserve_order,
    shipping_pick_readiness,
    shipping_available_items,
    shipping_selectable_warehouse_state_codes,
)
from .transport_note import (
    build_transport_note_preview_context,
    can_access_transport_note,
    get_or_create_transport_note,
    render_transport_note_docx,
    render_transport_note_pdf,
    transport_note_docx_filename,
    transport_note_filename,
)
from .workflow import (
    can_access_order as _can_access_order,
    can_cancel as _can_cancel,
    can_edit_items as _can_edit_items,
    can_edit_order_form as _can_edit_order_form,
    can_manager_approve as _can_manager_approve,
    can_manager_reopen as _can_manager_reopen,
    can_set_shipping_priority as _can_set_shipping_priority,
    can_storekeeper_manage_packing,
    can_storekeeper_accept as _can_storekeeper_accept,
    can_storekeeper_pack as _can_storekeeper_pack,
    can_storekeeper_pick as _can_storekeeper_pick,
    can_submit_for_approval as _can_submit_for_approval,
    can_write as _can_write,
    is_logistician_role as _is_logistician_role,
    is_manager_role as _is_manager_role,
    is_storekeeper_role as _is_storekeeper_role,
    requires_warehouse_cancel_confirmation as _requires_warehouse_cancel_confirmation,
    request_scope as _request_scope,
    shipping_order_is_manual_priority as _shipping_order_is_manual_priority,
)

_BOX_COUNT_COMMENT_RE = re.compile(r"коробов:\s*(\d+)", re.IGNORECASE)
_BOX_CODES_COMMENT_RE = re.compile(r"короба:\s*([^;]+)", re.IGNORECASE)
logger = logging.getLogger(__name__)


def _to_int(value) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return 0


def _parse_box_count_from_comment(comment: str | None) -> int:
    match = _BOX_COUNT_COMMENT_RE.search(str(comment or ""))
    if not match:
        return 0
    return _to_int(match.group(1))


def _warehouse_cancel_request_payload(order: ShippingOrder) -> dict:
    entries = (
        OrderAuditEntry.objects.filter(order_id=order.number, order_type="shipping")
        .only("action", "payload")
        .order_by("-created_at", "-id")[:80]
    )
    for entry in entries:
        payload = entry.payload if isinstance(entry.payload, dict) else {}
        status = str(payload.get("status") or "").strip().lower()
        action = str(entry.action or "").strip().lower()
        if status == "warehouse_cancel_requested":
            return payload
        if status in {"cancelled", "canceled"} or action == "warehouse_cancel_rejected":
            return {}
    return {}


def _client_cancel_request_payload(order: ShippingOrder) -> dict:
    entries = (
        OrderAuditEntry.objects.filter(order_id=order.number, order_type="shipping")
        .only("action", "payload")
        .order_by("-created_at", "-id")[:80]
    )
    for entry in entries:
        payload = entry.payload if isinstance(entry.payload, dict) else {}
        status = str(payload.get("status") or "").strip().lower()
        action = str(entry.action or "").strip().lower()
        if status == "cancel_requested":
            return payload
        if status in {"warehouse_cancel_requested", "cancelled", "canceled"} or action in {
            "warehouse_cancel_request",
            "warehouse_cancel_rejected",
        }:
            return {}
    return {}


def _parse_box_codes_from_comment(comment: str | None) -> list[str]:
    match = _BOX_CODES_COMMENT_RE.search(str(comment or ""))
    if not match:
        return []
    codes: list[str] = []
    seen: set[str] = set()
    for raw_code in match.group(1).split(","):
        code = str(raw_code or "").strip()
        normalized = code.lower()
        if not code or normalized in seen:
            continue
        seen.add(normalized)
        codes.append(code)
    return codes


def _display_shipping_number(number: str | None) -> str:
    return format_order_number("shipping", number)


def _stock_row_identity(*, sku_code: str, name: str, size: str, barcode: str, goods_type: str) -> tuple[str, str, str, str, str]:
    return (
        str(sku_code or "").strip(),
        str(name or "").strip(),
        str(size or "").strip(),
        str(barcode or "").strip(),
        str(goods_type or "").strip(),
    )


def _selected_box_values_from_request(request) -> dict[str, str]:
    all_keys = request.POST.getlist("stock_key_all[]") or request.POST.getlist("stock_key_all")
    all_boxes = request.POST.getlist("stock_boxes[]") or request.POST.getlist("stock_boxes")
    values: dict[str, str] = {}
    for index, raw_key in enumerate(all_keys):
        key = str(raw_key or "").strip()
        if not key or key in values:
            continue
        values[key] = str(all_boxes[index] if index < len(all_boxes) else "").strip()
    return values


def _selected_box_values_for_order(order: ShippingOrder, stock_rows: list[dict]) -> dict[str, str]:
    rows_by_identity: dict[tuple[str, str, str, str, str], list[dict]] = defaultdict(list)
    rows_by_barcode: dict[str, list[dict]] = defaultdict(list)
    rows_by_sku: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for row in stock_rows:
        rows_by_identity[
            _stock_row_identity(
                sku_code=row.get("sku_code") or "",
                name=row.get("name") or "",
                size=row.get("size") or "",
                barcode=row.get("barcode") or "",
                goods_type=row.get("goods_type") or "",
            )
        ].append(row)
        barcode_key = str(row.get("barcode") or "").strip().casefold()
        if barcode_key:
            rows_by_barcode[barcode_key].append(row)
        sku_key = (
            str(row.get("sku_code") or "").strip().casefold(),
            str(row.get("size") or "").strip().casefold(),
            str(row.get("goods_type") or "").strip().casefold(),
        )
        if sku_key[0]:
            rows_by_sku[sku_key].append(row)

    selected_by_key: dict[str, str] = {}
    mixed_group_boxes: dict[str, int] = {}
    for item in order.items.order_by("id"):
        if extract_partial_box_split(item.comment):
            continue
        parsed_box_codes = _parse_box_codes_from_comment(item.comment)
        matches = rows_by_identity.get(
            _stock_row_identity(
                sku_code=item.sku_code,
                name=item.name,
                size=item.size,
                barcode=item.barcode,
                goods_type=item.goods_type,
            ),
            [],
        )
        if not matches:
            barcode_key = str(item.barcode or "").strip().casefold()
            if barcode_key:
                matches = rows_by_barcode.get(barcode_key, [])
        if not matches:
            matches = rows_by_sku.get(
                (
                    str(item.sku_code or "").strip().casefold(),
                    str(item.size or "").strip().casefold(),
                    str(item.goods_type or "").strip().casefold(),
                ),
                [],
            )
        if not matches:
            continue
        # A regular saved item represents whole boxes. Its original box may
        # have become a partial remainder after the order reserve was released.
        # Never restore that remainder into the whole-box input: prefer another
        # compatible whole-box row, or leave the item unselected when none exists.
        matches = [row for row in matches if not row.get("partial_only")]
        if not matches:
            continue
        parsed_boxes = _parse_box_count_from_comment(item.comment)
        selected_row = None
        if parsed_box_codes:
            parsed_box_code_set = {code.lower() for code in parsed_box_codes}
            for row in matches:
                row_codes = {str(code or "").strip().lower() for code in row.get("box_codes") or []}
                if parsed_box_code_set and parsed_box_code_set.issubset(row_codes):
                    parsed_boxes = len(parsed_box_codes)
                    selected_row = row
                    break
        for row in matches:
            if selected_row is not None:
                break
            box_qty = int(row.get("box_qty") or 0)
            if box_qty <= 0:
                continue
            # When saved physical box codes disappeared from the whole-box
            # rows, recalculate the count against each current package size.
            candidate_boxes = 0 if parsed_box_codes else parsed_boxes
            if candidate_boxes <= 0 and int(item.qty_requested or 0) % box_qty == 0:
                candidate_boxes = int(item.qty_requested or 0) // box_qty
            if candidate_boxes <= 0:
                continue
            if candidate_boxes * box_qty == int(item.qty_requested or 0):
                parsed_boxes = candidate_boxes
                selected_row = row
                break
            if selected_row is None:
                parsed_boxes = candidate_boxes
                selected_row = row
        if selected_row is None or parsed_boxes <= 0:
            continue
        mixed_group = str(selected_row.get("mixed_group") or "").strip()
        if selected_row.get("is_mixed_box") and mixed_group:
            mixed_group_boxes[mixed_group] = max(int(mixed_group_boxes.get(mixed_group) or 0), int(parsed_boxes))
        else:
            selected_by_key[str(selected_row["key"])] = str(int(parsed_boxes))

    for row in stock_rows:
        mixed_group = str(row.get("mixed_group") or "").strip()
        if row.get("is_mixed_box") and mixed_group and mixed_group in mixed_group_boxes:
            selected_by_key[str(row["key"])] = str(int(mixed_group_boxes[mixed_group]))
    return selected_by_key


def _partial_box_splits_from_request(request) -> list[dict]:
    raw = str(request.POST.get("stock_partial_box_splits_json") or "").strip()
    if not raw:
        return []
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    return [dict(item) for item in parsed if isinstance(item, dict)]


def _partial_box_splits_json_from_request(request) -> str:
    splits = _partial_box_splits_from_request(request)
    return json.dumps(splits, ensure_ascii=False, separators=(",", ":")) if splits else "[]"


def _partial_box_count_from_request(request) -> int:
    return sum(max(_to_int(split.get("boxes")), 0) for split in _partial_box_splits_from_request(request))


def _ozon_piece_pick_tariff_confirmation_error(request) -> str:
    """Require an explicit tariff confirmation for Ozon API piece picking.

    The check only validates the submitted client-form contract. It does not
    calculate billing and does not mutate stock, reserves or workflow status.
    """
    if str(request.POST.get("action") or "").strip() == "save_draft":
        return ""
    if str(request.POST.get("draft_autosave") or "").strip() == "1":
        return ""
    raw_payload = str(request.POST.get("ozon_supply_payload_json") or "").strip()
    if not raw_payload:
        return ""
    try:
        payload = json.loads(raw_payload)
    except json.JSONDecodeError:
        payload = {}

    piece_qty = 0
    for split in _partial_box_splits_from_request(request):
        boxes = max(_to_int(split.get("boxes")), 0)
        if boxes <= 0:
            continue
        for item in split.get("items") or []:
            if not isinstance(item, dict):
                continue
            piece_qty += boxes * max(_to_int(item.get("qty")), 0)
    if piece_qty <= 0:
        return ""
    expected_piece_qty = max(
        _to_int(payload.get("piece_pick_total_qty")) if isinstance(payload, dict) else 0,
        0,
    )
    if expected_piece_qty > 0 and piece_qty != expected_piece_qty:
        return (
            "Поштучный добор не совпадает с составом заявки Ozon. "
            "Обновите данные Ozon и повторите отправку."
        )
    if str(request.POST.get("piece_pick_tariff_confirmed") or "").strip() == "1":
        return ""
    return (
        f"В заявке Ozon есть поштучный подбор: {piece_qty} шт. "
        "Подтвердите отдельную тарификацию поштучного подбора перед отправкой заявки."
    )


def _partial_box_counts_from_request(
    request,
    stock_rows: list[dict],
) -> tuple[dict[str, int], dict[str, int]]:
    stock_map = {str(row.get("key") or "").strip(): row for row in stock_rows if str(row.get("key") or "").strip()}
    counts_by_key: dict[str, int] = defaultdict(int)
    counts_by_group: dict[str, int] = defaultdict(int)
    for split in _partial_box_splits_from_request(request):
        row_key = str(split.get("row_key") or "").strip()
        row = stock_map.get(row_key)
        boxes = max(_to_int(split.get("boxes")), 0)
        if row is None or boxes <= 0:
            continue
        group = str(row.get("mixed_group") or "").strip()
        if row.get("is_mixed_box") and group:
            counts_by_group[group] += boxes
        else:
            counts_by_key[row_key] += boxes
    return counts_by_key, counts_by_group


def _partial_box_splits_for_order(order: ShippingOrder, stock_rows: list[dict]) -> list[dict]:
    valid_keys = {str(row.get("key") or "").strip() for row in stock_rows}
    grouped: dict[str, dict] = {}
    for item in order.items.order_by("id"):
        meta = extract_partial_box_split(item.comment)
        if not meta:
            continue
        group_key = str(meta.get("group_key") or meta.get("row_key") or f"item:{item.id}").strip()
        row_key = str(meta.get("row_key") or "").strip()
        if row_key and valid_keys and row_key not in valid_keys:
            continue
        entry = grouped.setdefault(
            group_key,
            {
                "group_key": group_key,
                "row_key": row_key,
                "group": str(meta.get("source_group") or "").strip(),
                "boxes": max(_to_int(meta.get("source_boxes")), 0),
                "items": [],
            },
        )
        if not entry.get("row_key") and row_key:
            entry["row_key"] = row_key
        if not entry.get("group") and meta.get("source_group"):
            entry["group"] = str(meta.get("source_group") or "").strip()
        entry["boxes"] = max(_to_int(entry.get("boxes")), _to_int(meta.get("source_boxes")))
        pick_qty = max(_to_int(meta.get("item_pick_qty")), 0)
        if row_key and pick_qty > 0:
            entry["items"].append({"key": row_key, "qty": pick_qty})
    result: list[dict] = []
    for entry in grouped.values():
        boxes = _to_int(entry.get("boxes"))
        items = [
            {"key": str(item.get("key") or "").strip(), "qty": max(_to_int(item.get("qty")), 0)}
            for item in entry.get("items") or []
            if str(item.get("key") or "").strip() and _to_int(item.get("qty")) > 0
        ]
        if boxes <= 0 or not items:
            continue
        entry["boxes"] = boxes
        entry["items"] = items
        result.append(entry)
    return result


def _partial_box_splits_json_for_order(order: ShippingOrder, stock_rows: list[dict]) -> str:
    splits = _partial_box_splits_for_order(order, stock_rows)
    return json.dumps(splits, ensure_ascii=False, separators=(",", ":")) if splits else "[]"


def _with_selected_boxes(stock_rows: list[dict], selected_boxes: dict[str, str] | None = None) -> list[dict]:
    selected_boxes = selected_boxes or {}
    prepared: list[dict] = []
    for row in stock_rows:
        cloned = dict(row)
        selected_value = selected_boxes.get(str(row.get("key") or ""), "0")
        cloned["selected_boxes"] = str(selected_value or "0")
        prepared.append(cloned)
    return prepared


def _order_box_count(order: ShippingOrder) -> int:
    expected = shipping_expected_box_count(order)
    if expected > 0:
        return expected
    return sum(_parse_box_count_from_comment(item.comment) for item in order.items.all())


def _can_storekeeper_manage_packing(scope: str | None, role: str | None, order: ShippingOrder) -> bool:
    if order.status == ShippingOrder.STATUS_PACKED:
        return can_storekeeper_manage_packing(
            scope,
            role,
            order,
            has_manageable_boxes=True,
        )
    manageable_boxes = _shipping_manageable_packing_boxes(order)
    loose_state = _shipping_loose_packing_initial_state(order)
    has_loose_items = bool(loose_state.get("loose_items") if isinstance(loose_state, dict) else False)
    has_manageable_boxes = bool(manageable_boxes)
    if order.status != ShippingOrder.STATUS_PACKED:
        expected_boxes = _order_box_count(order)
        has_manageable_boxes = (expected_boxes > 0 and len(manageable_boxes) >= expected_boxes) or has_loose_items
        marketplace_name = str(getattr(getattr(order, "marketplace", None), "name", "") or "").strip().casefold()
        if (
            not has_manageable_boxes
            and "ozon" in marketplace_name
            and expected_boxes > 0
            and bool(manageable_boxes)
            and len(manageable_boxes) != expected_boxes
        ):
            # Ozon can prescribe a different number/composition of GM boxes than
            # the source Fullbox boxes physically delivered to OTG.  Let the
            # dedicated re-packing command decide whether the immutable GM plan
            # is complete; never bypass item/quantity validation here.
            from .truth import ShippingTruthService

            has_manageable_boxes = not bool(
                ShippingTruthService.for_order(order).quantity_mismatch_rows
            )
    return can_storekeeper_manage_packing(
        scope,
        role,
        order,
        has_manageable_boxes=has_manageable_boxes,
    )


def _save_shipping_attachments(order: ShippingOrder, request) -> list[str]:
    uploaded_names: list[str] = []
    for uploaded_file in request.FILES.getlist("documents"):
        if not uploaded_file or not getattr(uploaded_file, "name", ""):
            continue
        attachment = ShippingOrderAttachment.objects.create(
            order=order,
            uploaded_by=request.user if request.user.is_authenticated else None,
            file=uploaded_file,
        )
        uploaded_names.append(attachment.filename)
    return uploaded_names


def _delete_shipping_attachments(order: ShippingOrder, request) -> list[str]:
    raw_ids = request.POST.getlist("delete_attachment_ids") or request.POST.getlist("delete_attachment_ids[]")
    attachment_ids: list[int] = []
    for raw_id in raw_ids:
        try:
            attachment_ids.append(int(str(raw_id).strip()))
        except (TypeError, ValueError):
            continue
    if not attachment_ids:
        return []

    deleted_names: list[str] = []
    for attachment in ShippingOrderAttachment.objects.filter(order=order, pk__in=attachment_ids):
        deleted_names.append(attachment.filename)
        if attachment.file:
            attachment.file.delete(save=False)
        attachment.delete()
    return deleted_names


def _shipping_attachment_payload(order: ShippingOrder) -> list[dict]:
    return [
        {
            "id": attachment.pk,
            "name": attachment.filename,
            "size": int(getattr(attachment.file, "size", 0) or 0),
            "download_url": reverse(
                "shipping:attachment-download",
                kwargs={"pk": order.pk, "attachment_id": attachment.pk},
            ),
        }
        for attachment in _active_shipping_attachments(order)
    ]


def _uploaded_file_sha256(uploaded_file) -> str:
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    return digest.hexdigest()


def _stored_attachment_sha256(attachment: ShippingOrderAttachment) -> str:
    digest = hashlib.sha256()
    attachment.file.open("rb")
    try:
        for chunk in attachment.file.chunks():
            digest.update(chunk)
    finally:
        attachment.file.close()
    return digest.hexdigest()


def _can_manager_append_shipping_documents(scope, role, order):
    # Append only: no permission to edit the order or delete existing documents.
    return scope == "staff" and _is_manager_role(role) and not order.is_closed()


@login_required
@require_http_methods(["POST"])
@transaction.atomic
def shipping_attachment_upload(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    order = get_object_or_404(
        ShippingOrder.objects.select_for_update().select_related("agency"),
        pk=pk,
    )
    if (
        scope is None
        or not _can_access_order(scope, client_agency, order)
        or not (
            _can_edit_order_form(scope, role, order)
            or _can_manager_append_shipping_documents(scope, role, order)
        )
    ):
        return JsonResponse(
            {"ok": False, "error": "Добавление вложений недоступно для вашей роли или статуса заявки."},
            status=403,
        )
    upload_id = str(request.POST.get("upload_id") or "").strip()
    if not upload_id:
        return JsonResponse({"ok": False, "error": "Не указан идентификатор загрузки."}, status=400)
    cache_key = f"shipping-attachment-upload:{request.user.pk}:{order.pk}:{upload_id[:96]}"
    if cache.get(cache_key):
        return JsonResponse({"ok": True, "attachments": _shipping_attachment_payload(order)})

    files = [item for item in request.FILES.getlist("documents") if getattr(item, "name", "")]
    if not files:
        return JsonResponse({"ok": False, "error": "Выберите хотя бы один файл."}, status=400)
    existing = list(_active_shipping_attachments(order))
    if len(existing) + len(files) > 10:
        return JsonResponse({"ok": False, "error": "К одной заявке можно прикрепить не более 10 файлов."}, status=400)
    allowed_extensions = {".pdf", ".xls", ".xlsx", ".jpg", ".jpeg", ".png"}
    existing_hashes: set[tuple[str, int, str]] = set()
    for attachment in existing:
        try:
            existing_hashes.add(
                (
                    attachment.filename.lower(),
                    int(getattr(attachment.file, "size", 0) or 0),
                    _stored_attachment_sha256(attachment),
                )
            )
        except Exception:
            logger.warning("Cannot hash shipping attachment id=%s", attachment.pk)

    uploaded_names: list[str] = []
    skipped_names: list[str] = []
    prepared_files: list[tuple[object, tuple[str, int, str]]] = []
    for uploaded_file in files:
        filename = str(uploaded_file.name or "").strip()
        extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if extension not in allowed_extensions:
            return JsonResponse(
                {"ok": False, "error": f"{filename}: допустимы PDF, Excel, JPG и PNG."},
                status=400,
            )
        size = int(getattr(uploaded_file, "size", 0) or 0)
        if size > 25 * 1024 * 1024:
            return JsonResponse({"ok": False, "error": f"{filename}: файл больше 25 МБ."}, status=400)
        identity = (filename.lower(), size, _uploaded_file_sha256(uploaded_file))
        if identity in existing_hashes:
            skipped_names.append(filename)
            continue
        existing_hashes.add(identity)
        prepared_files.append((uploaded_file, identity))
    for uploaded_file, _identity in prepared_files:
        attachment = ShippingOrderAttachment.objects.create(
            order=order,
            uploaded_by=request.user,
            file=uploaded_file,
        )
        uploaded_names.append(attachment.filename)
    _log_update(
        order,
        request,
        "К заявке загружены вложения",
        extra={"uploaded_documents": uploaded_names, "skipped_duplicate_documents": skipped_names},
    )
    transaction.on_commit(lambda: cache.set(cache_key, True, timeout=24 * 60 * 60))
    return JsonResponse(
        {
            "ok": True,
            "attachments": _shipping_attachment_payload(order),
            "uploaded": uploaded_names,
            "skipped_duplicates": skipped_names,
        }
    )


@login_required
@require_http_methods(["POST"])
def shipping_attachment_delete(request, pk: int, attachment_id: int):
    scope, role, client_agency = _request_scope(request)
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency"),
        pk=pk,
    )
    if (
        scope is None
        or not _can_access_order(scope, client_agency, order)
        or not _can_edit_order_form(scope, role, order)
    ):
        return JsonResponse(
            {"ok": False, "error": "Вложения можно удалять только до передачи заявки на склад."},
            status=403,
        )
    attachment = get_object_or_404(ShippingOrderAttachment, pk=attachment_id, order=order)
    filename = attachment.filename
    if attachment.file:
        attachment.file.delete(save=False)
    attachment.delete()
    _log_update(order, request, "Из заявки удалено вложение", extra={"deleted_documents": [filename]})
    return JsonResponse({"ok": True, "attachments": _shipping_attachment_payload(order)})


def _log_update(order: ShippingOrder, request, description: str, *, action: str = "update", extra: dict | None = None) -> None:
    log_order_action(
        action=action,
        order_id=order.number,
        order_type="shipping",
        user=request.user if request.user.is_authenticated else None,
        agency=order.agency,
        description=description,
        payload=order_payload(order, extra=extra),
    )


def _compose_picker_key(
    *,
    sku_id: int,
    sku_code: str,
    name: str,
    size: str,
    goods_type: str,
    box_qty: int,
    barcode: str,
    mixed_group: str,
) -> str:
    return "|".join(
        [
            str(int(sku_id or 0)),
            str(sku_code or "").strip(),
            str(size or "").strip(),
            str(goods_type or "").strip(),
            str(int(box_qty or 0)),
            str(barcode or "").strip(),
            str(name or "").strip(),
            str(mixed_group or "").strip(),
        ]
    )


def _shipping_stock_picker_rows(
    agency: Agency | None,
    *,
    exclude_order: ShippingOrder | None = None,
    diagnostics: list[dict] | None = None,
) -> list[dict]:
    if not agency:
        return []
    rows = StockAvailabilityService.stock_rows_with_availability(
        agency=agency,
        require_box=True,
        exclude_shipping_order_id=str(exclude_order.number or "").strip() if exclude_order else None,
    )
    allowed_state_codes = set(
        shipping_selectable_warehouse_state_codes(include_reserved=exclude_order is not None)
    )
    partial_state_codes = set(shipping_selectable_warehouse_state_codes())
    rows = [
        row
        for row in rows
        if str(row.get("warehouse_state_code") or "").strip() in allowed_state_codes
    ]

    box_lines: dict[str, dict[tuple[int, str, str, str, str, str], int]] = defaultdict(lambda: defaultdict(int))
    box_available_lines: dict[str, dict[tuple[int, str, str, str, str, str], int]] = defaultdict(lambda: defaultdict(int))
    box_processing_reserved_qty: defaultdict[str, int] = defaultdict(int)
    box_shipping_reserved_qty: defaultdict[str, int] = defaultdict(int)
    box_unavailable_qty: defaultdict[str, int] = defaultdict(int)
    box_partial_state_allowed: defaultdict[str, bool] = defaultdict(lambda: True)
    box_open_remainder: defaultdict[str, bool] = defaultdict(bool)
    box_barcodes: dict[str, set[str]] = defaultdict(set)
    box_item_signatures: dict[str, set[tuple[str, str, str, str, str]]] = defaultdict(set)
    box_codes_by_id: dict[str, str] = {}
    for row in rows:
        sku_code = str(row.get("sku") or "").strip()
        if not sku_code:
            continue
        box_code = str(row.get("box_code") or "").strip()
        if not box_code:
            continue
        qty_in_box = int(row.get("qty") or 0)
        if qty_in_box <= 0:
            continue
        size = str(row.get("size") or "").strip()
        goods_type = str(row.get("goods_type") or "").strip()
        barcode = str(row.get("barcode") or "").strip()
        name = str(row.get("name") or "").strip()
        sku_id = int(row.get("sku_ref_id") or 0)
        line_key = (sku_id, sku_code, name, size, barcode, goods_type)
        box_id = f"{int(row.get('agency_id') or 0)}:{box_code}"
        box_codes_by_id[box_id] = box_code
        box_lines[box_id][line_key] += qty_in_box
        available_qty = max(int(row.get("available_qty") or 0), 0)
        box_available_lines[box_id][line_key] += available_qty
        box_processing_reserved_qty[box_id] += max(int(row.get("processing_reserved_qty") or 0), 0)
        box_shipping_reserved_qty[box_id] += max(int(row.get("shipping_reserved_qty") or 0), 0)
        box_unavailable_qty[box_id] += max(qty_in_box - available_qty, 0)
        box_open_remainder[box_id] = bool(
            box_open_remainder[box_id]
            or (
                str(row.get("last_event_type") or "").strip() == "movement_completed"
                and str(row.get("last_event_context_type") or "").strip()
                in {"shipping", "processing"}
                and available_qty > 0
            )
        )
        if str(row.get("warehouse_state_code") or "").strip() not in partial_state_codes:
            box_partial_state_allowed[box_id] = False
        item_signature = (
            sku_code.lower(),
            size.lower(),
            name.lower(),
            goods_type.lower(),
            barcode.lower(),
        )
        box_item_signatures[box_id].add(item_signature)
        normalized_barcode = barcode.lower()
        if normalized_barcode:
            box_barcodes[box_id].add(normalized_barcode)

    all_box_ids = set(box_item_signatures.keys()) | set(box_barcodes.keys())
    mixed_box_ids: set[str] = set()
    for box_id in all_box_ids:
        barcodes = box_barcodes.get(box_id, set())
        if len(barcodes) >= 2:
            mixed_box_ids.add(box_id)
            continue
        signatures = box_item_signatures.get(box_id, set())
        if len(signatures) >= 2:
            mixed_box_ids.add(box_id)

    fully_available_box_ids = {
        box_id
        for box_id, lines_map in box_lines.items()
        if lines_map
        and all(
            int(box_available_lines.get(box_id, {}).get(line_key, 0)) >= int(qty_in_box)
            for line_key, qty_in_box in lines_map.items()
        )
    }

    if diagnostics is not None:
        visible_by_line: defaultdict[tuple, int] = defaultdict(int)
        full_box_qty_by_line: defaultdict[tuple, int] = defaultdict(int)
        line_labels: dict[tuple, dict] = {}
        for box_id, lines_map in box_lines.items():
            for line_key, qty_in_box in lines_map.items():
                visible_by_line[line_key] += max(
                    int(box_available_lines.get(box_id, {}).get(line_key, 0)),
                    0,
                )
                if box_id in fully_available_box_ids:
                    full_box_qty_by_line[line_key] += int(qty_in_box)
                sku_id, sku_code, name, size, barcode, goods_type = line_key
                line_labels[line_key] = {
                    "sku_id": sku_id,
                    "sku_code": sku_code,
                    "name": name,
                    "size": size,
                    "barcode": barcode,
                    "goods_type": goods_type,
                }
        affected = [
            line_key
            for line_key, visible_qty in visible_by_line.items()
            if visible_qty > full_box_qty_by_line.get(line_key, 0)
        ]
        reserve_orders: defaultdict[tuple[str, str, str, str], set[str]] = defaultdict(set)
        reserve_qty_by_line: defaultdict[tuple[str, str, str, str], int] = defaultdict(int)
        if affected:
            active_statuses = {
                WarehouseReserve.STATUS_ACTIVE,
                WarehouseReserve.STATUS_PARTIALLY_ALLOCATED,
                WarehouseReserve.STATUS_ALLOCATED,
                WarehouseReserve.STATUS_PARTIALLY_SATISFIED,
            }
            affected_skus = {str(line_labels[key]["sku_code"] or "") for key in affected}
            reserve_qs = WarehouseReserve.objects.filter(
                agency=agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                status__in=active_statuses,
                sku_code__in=affected_skus,
            )
            if exclude_order is not None:
                reserve_qs = reserve_qs.exclude(
                    context_type="shipping",
                    context_id=str(exclude_order.number or ""),
                )
            for reserve in reserve_qs.only(
                "sku_code", "size", "barcode", "goods_type", "context_id", "qty_reserved", "qty_satisfied"
            ):
                reserve_key = (
                    str(reserve.sku_code or "").strip().lower(),
                    str(reserve.size or "").strip().lower(),
                    str(reserve.barcode or "").strip().lower(),
                    str(reserve.goods_type or "").strip().lower(),
                )
                context_id = str(reserve.context_id or "").strip()
                if context_id:
                    reserve_orders[reserve_key].add(_display_shipping_number(context_id))
                reserve_qty_by_line[reserve_key] += max(
                    int(reserve.qty_reserved or 0) - int(reserve.qty_satisfied or 0),
                    0,
                )
        for line_key in affected:
            label = line_labels[line_key]
            reserve_key = (
                str(label["sku_code"] or "").strip().lower(),
                str(label["size"] or "").strip().lower(),
                str(label["barcode"] or "").strip().lower(),
                str(label["goods_type"] or "").strip().lower(),
            )
            blocking_orders = sorted(reserve_orders.get(reserve_key, set()))
            visible_qty = int(visible_by_line[line_key])
            full_box_qty = int(full_box_qty_by_line.get(line_key, 0))
            reserved_qty = int(reserve_qty_by_line.get(reserve_key, 0))
            diagnostics.append(
                {
                    **label,
                    "visible_qty": visible_qty,
                    "full_box_qty": full_box_qty,
                    "reserved_qty": reserved_qty,
                    "blocking_orders": blocking_orders,
                    "blocking_orders_label": ", ".join(blocking_orders),
                    "reason": (
                        f"Товар есть на складе, но целый короб занят резервом заявки "
                        f"{blocking_orders[0]}. Это не физический дефицит."
                        if blocking_orders
                        else "Товар есть на складе, но сейчас нет свободного целого короба."
                    ),
                }
            )

    # Обычные (не mixed) короба агрегируем по строкам как раньше.
    regular_grouped: dict[tuple[int, str, str, str, str, str, int], set[str]] = defaultdict(set)
    for box_id, lines_map in box_lines.items():
        if box_id not in fully_available_box_ids:
            continue
        if box_id in mixed_box_ids:
            continue
        for line_key, qty_in_box in lines_map.items():
            sku_id, sku_code, name, size, barcode, goods_type = line_key
            key = (sku_id, sku_code, name, size, barcode, goods_type, int(qty_in_box))
            regular_grouped[key].add(box_id)

    # Mixed-короба группируем по составу, чтобы можно было выбирать количество коробов
    # и автоматически добавлять все позиции из того же короба/типа короба.
    mixed_compositions: dict[
        tuple[tuple[int, str, str, str, str, str, int], ...],
        list[str],
    ] = defaultdict(list)
    for box_id in sorted(mixed_box_ids & fully_available_box_ids):
        composition: list[tuple[int, str, str, str, str, str, int]] = []
        for line_key, qty_in_box in box_lines.get(box_id, {}).items():
            sku_id, sku_code, name, size, barcode, goods_type = line_key
            composition.append(
                (
                    int(sku_id),
                    str(sku_code),
                    str(name),
                    str(size),
                    str(barcode),
                    str(goods_type),
                    int(qty_in_box),
                )
            )
        if not composition:
            continue
        mixed_compositions[tuple(sorted(composition))].append(box_id)
    prepared = []
    for (sku_id, sku_code, name, size, barcode, goods_type, box_qty), boxes in regular_grouped.items():
        boxes_count = len(boxes)
        if boxes_count <= 0:
            continue
        ordered_box_ids = sorted(
            boxes,
            key=lambda box_id: (
                0 if box_open_remainder.get(box_id, False) else 1,
                box_codes_by_id.get(box_id, box_id.split(":", 1)[-1]).casefold(),
            ),
        )
        open_remainder_boxes = sum(
            1 for box_id in ordered_box_ids if box_open_remainder.get(box_id, False)
        )
        total_qty = boxes_count * int(box_qty)
        prepared.append(
            {
                "sku_code": sku_code,
                "sku_id": sku_id,
                "name": name,
                "size": size,
                "barcode": barcode,
                "goods_type": goods_type,
                "box_qty": int(box_qty),
                "boxes_count": boxes_count,
                "total_qty": total_qty,
                "box_codes": [box_codes_by_id.get(box_id, box_id.split(":", 1)[-1]) for box_id in ordered_box_ids],
                "is_mixed_box": False,
                "mixed_group": "",
                "is_open_remainder": bool(open_remainder_boxes),
                "open_remainder_boxes": int(open_remainder_boxes),
            }
        )

    for composition_index, (composition, box_ids) in enumerate(sorted(mixed_compositions.items(), key=lambda item: item[1][0])):
        boxes_count = len(box_ids)
        if boxes_count <= 0:
            continue
        mixed_group = f"mix:{composition_index}:{box_ids[0]}"
        for sku_id, sku_code, name, size, barcode, goods_type, box_qty in composition:
            total_qty = boxes_count * int(box_qty)
            prepared.append(
                {
                    "sku_code": sku_code,
                    "sku_id": int(sku_id),
                    "name": name,
                    "size": size,
                    "barcode": barcode,
                    "goods_type": goods_type,
                    "box_qty": int(box_qty),
                    "boxes_count": boxes_count,
                    "total_qty": total_qty,
                    "box_codes": [box_codes_by_id.get(box_id, box_id.split(":", 1)[-1]) for box_id in sorted(box_ids)],
                    "is_mixed_box": True,
                    "mixed_group": mixed_group,
                    "is_open_remainder": False,
                    "open_remainder_boxes": 0,
                }
            )

    prepared.sort(
        key=lambda row: (
            0 if row["is_mixed_box"] else 1,
            row["mixed_group"],
            -int(row["sku_id"] or 0),
            row["sku_code"].lower(),
            row["size"].lower(),
            StockAvailabilityService.normalize_goods_type(row["goods_type"]),
            -int(row["box_qty"] or 0),
        )
    )

    interim_rows: list[dict] = []
    for row in prepared:
        available_boxes = int(row["boxes_count"])
        interim_rows.append(
            {
                **row,
                "available_boxes": int(max(available_boxes, 0)),
                "available_qty": int(max(available_boxes, 0)) * int(row["box_qty"]),
            }
        )

    mixed_rows_by_group: dict[str, list[dict]] = defaultdict(list)
    for row in interim_rows:
        if row.get("is_mixed_box") and row.get("mixed_group"):
            mixed_rows_by_group[str(row["mixed_group"])].append(row)

    for group_rows in mixed_rows_by_group.values():
        group_available = min(int(item.get("available_boxes") or 0) for item in group_rows)
        for item in group_rows:
            item["available_boxes"] = max(group_available, 0)
            item["available_qty"] = max(group_available, 0) * int(item.get("box_qty") or 0)

    mixed_color_map: dict[str, int] = {}
    for index, group_key in enumerate(sorted(mixed_rows_by_group.keys())):
        mixed_color_map[group_key] = index % 6

    result: list[dict] = []
    for row in interim_rows:
        available_boxes = int(row.get("available_boxes") or 0)
        if available_boxes <= 0:
            continue
        available_qty = available_boxes * int(row["box_qty"])
        picker_key = _compose_picker_key(
            sku_id=row["sku_id"],
            sku_code=row["sku_code"],
            name=row["name"],
            size=row["size"],
            goods_type=row["goods_type"],
            box_qty=row["box_qty"],
            barcode=row["barcode"],
            mixed_group=row["mixed_group"],
        )
        result.append(
            {
                "key": picker_key,
                "sku_id": int(row["sku_id"] or 0),
                "sku_code": row["sku_code"],
                "name": row["name"],
                "size": row["size"],
                "barcode": row["barcode"],
                "goods_type": row["goods_type"],
                "box_qty": int(row["box_qty"]),
                "available_boxes": int(available_boxes),
                "available_qty": int(available_qty),
                "box_codes": list(row.get("box_codes") or []),
                "is_mixed_box": bool(row["is_mixed_box"]),
                "mixed_group": row["mixed_group"],
                "mixed_color": int(mixed_color_map.get(str(row["mixed_group"] or ""), -1)),
                "partial_only": False,
                "split_available_qty": int(row["box_qty"]),
                "is_open_remainder": bool(row.get("is_open_remainder")),
                "open_remainder_boxes": int(row.get("open_remainder_boxes") or 0),
            }
        )

    partial_box_ids = sorted(
        box_id
        for box_id in set(box_lines) - fully_available_box_ids
        if box_partial_state_allowed[box_id]
        and box_processing_reserved_qty[box_id] <= 0
        and box_shipping_reserved_qty[box_id] > 0
        and box_unavailable_qty[box_id] <= box_shipping_reserved_qty[box_id]
    )
    partial_color_offset = len(mixed_color_map)
    for partial_index, box_id in enumerate(partial_box_ids):
        lines_map = box_lines.get(box_id, {})
        if not lines_map:
            continue
        available_map = box_available_lines.get(box_id, {})
        if not any(max(int(available_map.get(line_key, 0)), 0) > 0 for line_key in lines_map):
            continue
        is_mixed_box = box_id in mixed_box_ids
        partial_group = f"partial-box:{box_id}"
        for line_key, qty_in_box in lines_map.items():
            sku_id, sku_code, name, size, barcode, goods_type = line_key
            split_available_qty = max(int(available_map.get(line_key, 0)), 0)
            if not is_mixed_box and split_available_qty <= 0:
                continue
            picker_key = _compose_picker_key(
                sku_id=sku_id,
                sku_code=sku_code,
                name=name,
                size=size,
                goods_type=goods_type,
                box_qty=qty_in_box,
                barcode=barcode,
                mixed_group=partial_group,
            )
            result.append(
                {
                    "key": picker_key,
                    "sku_id": int(sku_id or 0),
                    "sku_code": sku_code,
                    "name": name,
                    "size": size,
                    "barcode": barcode,
                    "goods_type": goods_type,
                    "box_qty": int(qty_in_box),
                    "available_boxes": 1,
                    "available_qty": int(split_available_qty),
                    "box_codes": [box_codes_by_id.get(box_id, box_id.split(":", 1)[-1])],
                    "is_mixed_box": bool(is_mixed_box),
                    "mixed_group": partial_group,
                    "mixed_color": (partial_color_offset + partial_index) % 6,
                    "partial_only": True,
                    "split_available_qty": int(split_available_qty),
                    "is_open_remainder": True,
                    "open_remainder_boxes": 1,
                }
            )
    return result


def _selected_stock_rows_with_boxes(
    request,
    stock_rows: list[dict],
    *,
    multiple: bool,
) -> tuple[list[tuple[dict, int]], int, list[str]]:
    stock_map = {row["key"]: row for row in stock_rows}
    mixed_group_rows: dict[str, list[dict]] = defaultdict(list)
    for row in stock_rows:
        if row.get("is_mixed_box") and row.get("mixed_group"):
            mixed_group_rows[str(row["mixed_group"])].append(row)
    errors: list[str] = []

    if multiple:
        keys = []
    else:
        keys = [request.POST.get("stock_key", "")]
    all_keys = request.POST.getlist("stock_key_all[]") or request.POST.getlist("stock_key_all")
    all_boxes = request.POST.getlist("stock_boxes[]") or request.POST.getlist("stock_boxes")
    boxes_by_key: dict[str, str] = {}
    for index, raw_key in enumerate(all_keys):
        key = str(raw_key or "").strip()
        if not key or key in boxes_by_key:
            continue
        boxes_by_key[key] = all_boxes[index] if index < len(all_boxes) else ""

    if multiple:
        positive_box_keys: list[str] = []
        seen_positive: set[str] = set()
        for key, raw_boxes in boxes_by_key.items():
            try:
                boxes = int(str(raw_boxes or "").strip())
            except ValueError:
                continue
            if boxes <= 0:
                continue
            if key not in seen_positive:
                seen_positive.add(key)
                positive_box_keys.append(key)
        keys = positive_box_keys

    explicit_rows: list[tuple[dict, int]] = []
    for raw_key in keys:
        key = str(raw_key or "").strip()
        if not key:
            continue
        row = stock_map.get(key)
        if row is None:
            errors.append("Выбрана неактуальная складская строка. Обновите страницу и повторите.")
            continue
        raw_boxes = boxes_by_key.get(key, "")
        try:
            boxes = int(str(raw_boxes or "").strip())
        except ValueError:
            errors.append(f"{row['sku_code']}: количество коробов должно быть целым числом.")
            continue
        if boxes <= 0:
            continue
        if row.get("partial_only"):
            errors.append(
                f"{row['sku_code']}/{row['size'] or '-'}: доступен только поштучный отбор из остатка короба."
            )
            continue
        explicit_rows.append((row, boxes))

    mixed_group_boxes: dict[str, int] = {}
    for row, boxes in explicit_rows:
        group = str(row.get("mixed_group") or "").strip()
        if not (row.get("is_mixed_box") and group):
            continue
        previous = mixed_group_boxes.get(group)
        if previous is None:
            mixed_group_boxes[group] = boxes
            continue
        if previous != boxes:
            errors.append(
                "Для одного микс-короба укажите одинаковое количество коробов по выбранным строкам."
            )

    expanded_rows: list[tuple[dict, int]] = []
    expanded_groups: set[str] = set()
    counted_groups: set[str] = set()
    total_box_count = 0
    for row, boxes in explicit_rows:
        group = str(row.get("mixed_group") or "").strip()
        if row.get("is_mixed_box") and group:
            if group not in counted_groups:
                counted_groups.add(group)
                total_box_count += boxes
            if group in expanded_groups:
                continue
            expanded_groups.add(group)
            group_boxes = mixed_group_boxes.get(group, boxes)
            for sibling in mixed_group_rows.get(group, []):
                expanded_rows.append((sibling, group_boxes))
        else:
            total_box_count += boxes
            expanded_rows.append((row, boxes))

    return expanded_rows, int(max(total_box_count, 0)), errors


def _stock_row_split_payload(row: dict) -> dict:
    return {
        "row_key": str(row.get("key") or "").strip(),
        "sku_code": str(row.get("sku_code") or "").strip(),
        "name": str(row.get("name") or "").strip(),
        "size": str(row.get("size") or "").strip(),
        "barcode": str(row.get("barcode") or "").strip(),
        "goods_type": str(row.get("goods_type") or "").strip(),
        "qty": int(row.get("box_qty") or 0),
    }


def _stock_row_item_identity(row: dict) -> tuple[str, str, str, str, str]:
    return (
        str(row.get("sku_code") or "").strip().lower(),
        str(row.get("name") or "").strip().lower(),
        str(row.get("size") or "").strip().lower(),
        str(row.get("barcode") or "").strip().lower(),
        str(row.get("goods_type") or "").strip().lower(),
    )


def _whole_box_exact_match_exists(stock_rows: list[dict], source_row: dict, target_qty: int) -> bool:
    target = int(target_qty or 0)
    if target <= 0 or source_row.get("is_mixed_box"):
        return False
    identity = _stock_row_item_identity(source_row)
    box_sizes: list[int] = []
    for row in stock_rows:
        if row.get("partial_only") or row.get("is_mixed_box"):
            continue
        if _stock_row_item_identity(row) != identity:
            continue
        box_qty = int(row.get("box_qty") or 0)
        available_boxes = int(row.get("available_boxes") or 0)
        if box_qty <= 0 or available_boxes <= 0:
            continue
        box_sizes.extend([box_qty] * available_boxes)
    if not box_sizes:
        return False

    reachable = {0}
    for box_qty in sorted(box_sizes, reverse=True):
        reachable |= {qty + box_qty for qty in reachable if qty + box_qty <= target}
        if target in reachable:
            return True
    return False


def _selected_partial_box_items(
    request,
    stock_rows: list[dict],
    whole_box_counts_by_key: dict[str, int],
    whole_box_counts_by_group: dict[str, int],
) -> tuple[list[dict], int, list[str]]:
    stock_map = {str(row.get("key") or "").strip(): row for row in stock_rows if str(row.get("key") or "").strip()}
    mixed_group_rows: dict[str, list[dict]] = defaultdict(list)
    for row in stock_rows:
        group = str(row.get("mixed_group") or "").strip()
        if row.get("is_mixed_box") and group:
            mixed_group_rows[group].append(row)

    selected_items: list[dict] = []
    errors: list[str] = []
    total_split_boxes = 0
    split_boxes_by_key: dict[str, int] = defaultdict(int)
    split_boxes_by_group: dict[str, int] = defaultdict(int)
    whole_qty_by_identity: defaultdict[tuple[str, str, str, str, str], int] = defaultdict(int)
    for stock_row in stock_rows:
        if stock_row.get("partial_only") or stock_row.get("is_mixed_box"):
            continue
        stock_key = str(stock_row.get("key") or "").strip()
        whole_boxes = int(whole_box_counts_by_key.get(stock_key, 0))
        if whole_boxes <= 0:
            continue
        whole_qty_by_identity[_stock_row_item_identity(stock_row)] += (
            whole_boxes * int(stock_row.get("box_qty") or 0)
        )

    for split_index, split in enumerate(_partial_box_splits_from_request(request), start=1):
        row_key = str(split.get("row_key") or "").strip()
        row = stock_map.get(row_key)
        if row is None:
            errors.append("Разбитый короб выбран из неактуальных остатков. Обновите страницу и повторите.")
            continue
        boxes = max(_to_int(split.get("boxes")), 0)
        if boxes <= 0:
            continue
        group = str(row.get("mixed_group") or "").strip()
        group_rows = list(mixed_group_rows.get(group) or []) if row.get("is_mixed_box") and group else [row]
        allowed_keys = {str(item.get("key") or "").strip() for item in group_rows}
        qty_by_key: dict[str, int] = {}
        for raw_item in split.get("items") or []:
            if not isinstance(raw_item, dict):
                continue
            item_key = str(raw_item.get("key") or "").strip()
            qty = max(_to_int(raw_item.get("qty")), 0)
            if not item_key or qty <= 0:
                continue
            if item_key not in allowed_keys:
                errors.append("В разбивке короба выбрана позиция из другого типа короба.")
                continue
            qty_by_key[item_key] = qty_by_key.get(item_key, 0) + qty
        if not qty_by_key:
            errors.append("Для разбитого короба укажите количество товара к отгрузке.")
            continue

        source_total_per_box = sum(max(int(item.get("box_qty") or 0), 0) for item in group_rows)
        pick_total_per_box = sum(qty_by_key.values())
        if pick_total_per_box >= source_total_per_box:
            errors.append("Если клиент забирает весь короб, используйте поле «К отгрузке (короба)».")
            continue

        row_by_key = {str(item.get("key") or "").strip(): item for item in group_rows}
        item_error_count = len(errors)
        for item_key, pick_qty in qty_by_key.items():
            source_row = row_by_key.get(item_key)
            if source_row is None:
                continue
            source_qty = max(int(source_row.get("box_qty") or 0), 0)
            split_available_value = source_row.get("split_available_qty")
            split_available_qty = max(
                int(source_qty if split_available_value is None else split_available_value),
                0,
            )
            if source_qty <= 0 or pick_qty > source_qty:
                errors.append(
                    f"{source_row['sku_code']}/{source_row['size'] or '-'}: в коробе только {source_qty} шт."
                )
                continue
            if pick_qty > split_available_qty:
                errors.append(
                    f"{source_row['sku_code']}/{source_row['size'] or '-'}: "
                    f"из остатка короба доступно только {split_available_qty} шт."
                )
                continue
            target_qty = int(whole_qty_by_identity[_stock_row_item_identity(source_row)]) + int(boxes) * int(pick_qty)
            if (
                not source_row.get("is_open_remainder")
                and _whole_box_exact_match_exists(stock_rows, source_row, target_qty)
            ):
                errors.append(
                    f"{source_row['sku_code']}/{source_row['size'] or '-'}: "
                    "это количество можно собрать целыми коробами. Выберите целые короба, "
                    "разбивка короба не требуется."
                )
                continue
        if len(errors) > item_error_count:
            continue

        source_box_pattern = [_stock_row_split_payload(source_row) for source_row in group_rows]
        pick_pattern = [
            {
                **_stock_row_split_payload(row_by_key[item_key]),
                "qty": int(qty),
            }
            for item_key, qty in qty_by_key.items()
            if item_key in row_by_key and int(qty or 0) > 0
        ]
        if not pick_pattern:
            continue

        group_key = str(split.get("group_key") or "").strip()
        if not group_key:
            group_key = f"partial:{group or row_key}:{split_index}"
        split_offset = int(split_boxes_by_group.get(group, 0)) if row.get("is_mixed_box") and group else int(split_boxes_by_key.get(row_key, 0))
        split_box_codes = [
            str(code or "").strip()
            for code in list(row.get("box_codes") or [])[split_offset:split_offset + boxes]
            if str(code or "").strip()
        ]
        if len(split_box_codes) < boxes:
            errors.append("Для разбитого короба не удалось определить исходные короба. Обновите страницу и повторите.")
            continue
        for item_key, pick_qty in qty_by_key.items():
            source_row = row_by_key.get(item_key)
            if source_row is None:
                continue
            meta = {
                "version": 1,
                "kind": "partial_box_split",
                "group_key": group_key,
                "row_key": item_key,
                "source_group": group,
                "source_boxes": int(boxes),
                "source_box_total_qty": int(source_total_per_box),
                "source_box_pattern": source_box_pattern,
                "source_box_codes": split_box_codes,
                "pick_pattern": pick_pattern,
                "item_pick_qty": int(pick_qty),
                "item_source_qty": int(source_row.get("box_qty") or 0),
                "is_mixed_box": bool(source_row.get("is_mixed_box")),
            }
            encoded_meta = encode_partial_box_split(meta)
            selected_items.append(
                {
                    "sku_code": str(source_row.get("sku_code") or "").strip(),
                    "name": str(source_row.get("name") or "").strip(),
                    "size": str(source_row.get("size") or "").strip(),
                    "barcode": str(source_row.get("barcode") or "").strip(),
                    "goods_type": str(source_row.get("goods_type") or "").strip(),
                    "qty_requested": int(boxes) * int(pick_qty),
                    "comment": (
                        f"Исходные короба: {', '.join(split_box_codes)}; "
                        f"Разбить короб: {boxes}; отбор из короба: {pick_qty} "
                        f"из {int(source_row.get('box_qty') or 0)}; {encoded_meta}"
                    ),
                }
            )

        if row.get("is_mixed_box") and group:
            split_boxes_by_group[group] += boxes
        else:
            split_boxes_by_key[row_key] += boxes
        total_split_boxes += boxes

    for key, split_boxes in split_boxes_by_key.items():
        row = stock_map.get(key)
        if row is None:
            continue
        selected_total = int(whole_box_counts_by_key.get(key, 0)) + int(split_boxes)
        if selected_total > int(row.get("available_boxes") or 0):
            errors.append(
                f"{row['sku_code']}/{row['size'] or '-'}: выбрано {selected_total} короб., доступно {row['available_boxes']}."
            )
    for group, split_boxes in split_boxes_by_group.items():
        rows = mixed_group_rows.get(group) or []
        if not rows:
            continue
        available = min(int(row.get("available_boxes") or 0) for row in rows)
        selected_total = int(whole_box_counts_by_group.get(group, 0)) + int(split_boxes)
        if selected_total > available:
            errors.append(f"Для микс-короба выбрано {selected_total} короб., доступно {available}.")

    return selected_items, int(max(total_split_boxes, 0)), errors


def _parse_selected_stock_items(
    request,
    stock_rows: list[dict],
    *,
    multiple: bool,
) -> tuple[list[dict], list[str]]:
    expanded_rows, _total_box_count, errors = _selected_stock_rows_with_boxes(
        request,
        stock_rows,
        multiple=multiple,
    )
    selected_map: dict[tuple[str, str, str, str, str, str, str], dict] = {}
    whole_box_counts_by_key: dict[str, int] = defaultdict(int)
    whole_box_counts_by_group: dict[str, int] = {}
    split_box_counts_by_key, split_box_counts_by_group = _partial_box_counts_from_request(request, stock_rows)

    for row, boxes in expanded_rows:
        if boxes > int(row["available_boxes"]):
            errors.append(
                f"{row['sku_code']}/{row['size'] or '-'}: доступно только {row['available_boxes']} короб."
            )
            continue
        group = str(row.get("mixed_group") or "").strip()
        if row.get("is_mixed_box") and group:
            whole_box_counts_by_group[group] = max(int(whole_box_counts_by_group.get(group) or 0), int(boxes))
        else:
            whole_box_counts_by_key[str(row.get("key") or "").strip()] += int(boxes)
        qty_requested = boxes * int(row["box_qty"])
        selected_code_offset = int(split_box_counts_by_group.get(group, 0)) if row.get("is_mixed_box") and group else int(split_box_counts_by_key.get(str(row.get("key") or "").strip(), 0))
        selected_box_codes = [
            str(code or "").strip()
            for code in list(row.get("box_codes") or [])[selected_code_offset:selected_code_offset + boxes]
            if str(code or "").strip()
        ]
        if row.get("box_codes") and len(selected_box_codes) < boxes:
            errors.append(
                f"{row['sku_code']}/{row['size'] or '-'}: не удалось определить выбранные короба."
            )
            continue
        item_key = (
            str(row.get("sku_code") or "").strip(),
            str(row.get("name") or "").strip(),
            str(row.get("size") or "").strip(),
            str(row.get("barcode") or "").strip(),
            str(row.get("goods_type") or "").strip(),
            str(int(row.get("box_qty") or 0)),
            ",".join(selected_box_codes),
        )
        existing = selected_map.get(item_key)
        comment = f"Коробов: {boxes}; кратность: {row['box_qty']}"
        if selected_box_codes:
            comment += f"; короба: {', '.join(selected_box_codes)}"
        if row.get("is_mixed_box"):
            comment += "; микс-короб"
        if existing:
            existing["qty_requested"] += qty_requested
            continue
        selected_map[item_key] = {
            "sku_code": item_key[0],
            "name": item_key[1],
            "size": item_key[2],
            "barcode": item_key[3],
            "goods_type": item_key[4],
            "qty_requested": qty_requested,
            "comment": comment,
        }
    partial_items, _split_box_count, partial_errors = _selected_partial_box_items(
        request,
        stock_rows,
        whole_box_counts_by_key,
        whole_box_counts_by_group,
    )
    errors.extend(partial_errors)
    tariff_confirmation_error = _ozon_piece_pick_tariff_confirmation_error(request)
    if tariff_confirmation_error:
        errors.append(tariff_confirmation_error)
    for item in partial_items:
        item_key = (
            str(item.get("sku_code") or "").strip(),
            str(item.get("name") or "").strip(),
            str(item.get("size") or "").strip(),
            str(item.get("barcode") or "").strip(),
            str(item.get("goods_type") or "").strip(),
            str(item.get("comment") or "").strip(),
            "partial_box_split",
        )
        selected_map[item_key] = item
    selected = list(selected_map.values())
    if not selected:
        errors.append("Выберите минимум одну позицию из складских остатков: целый короб или разбивку короба.")
    return selected, errors


@login_required
def shipping_list(request):
    scope, role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")
    context = build_shipping_list_page_context(
        request=request,
        scope=scope,
        role=role,
        client_agency=client_agency,
    )
    return render(request, "shipping/list.html", context)


@login_required
def shipping_create(request):
    scope, role, client_agency = _request_scope(request)
    if scope is None or not _can_write(scope, role):
        return HttpResponseForbidden("Доступ запрещен")
    return handle_shipping_create_request(
        request=request,
        scope=scope,
        role=role,
        client_agency=client_agency,
    )


@login_required
def shipping_client_template(request):
    """Download Excel template for client shipping import (client or staff in client LK)."""
    scope, _role, client_agency = _request_scope(request)
    if scope == "client" and client_agency is not None:
        from .client_template import build_client_shipping_template_response

        return build_client_shipping_template_response()
    if scope == "staff":
        # Manager opens template from client LK (`?client=`).
        client_id = str(request.GET.get("client") or "").strip()
        if client_id.isdigit() and Agency.objects.filter(id=int(client_id), archived=False).exists():
            from .client_template import build_client_shipping_template_response

            return build_client_shipping_template_response()
    return HttpResponseForbidden("Доступ запрещен")


@login_required
def shipping_ozon_client_template(request):
    """Download Ozon multi-warehouse Excel template for client/manager LK."""
    scope, _role, client_agency = _request_scope(request)
    if scope == "client" and client_agency is not None:
        from .ozon_template import build_ozon_shipping_template_response

        return build_ozon_shipping_template_response()
    if scope == "staff":
        client_id = str(request.GET.get("client") or "").strip()
        if client_id.isdigit() and Agency.objects.filter(id=int(client_id), archived=False).exists():
            from .ozon_template import build_ozon_shipping_template_response

            return build_ozon_shipping_template_response()
    return HttpResponseForbidden("Доступ запрещен")


def _resolve_form_agency_for_shipping(request, scope, client_agency) -> Agency | None:
    """Client agency for create-form API helpers (client self or staff ?client=)."""
    if scope == "client" and client_agency is not None:
        return client_agency
    if scope == "staff":
        raw = str(request.GET.get("client") or request.POST.get("client") or "").strip()
        if raw.isdigit():
            return Agency.objects.filter(id=int(raw), archived=False).first()
    return None


def _ozon_exclude_order_id(request) -> int | None:
    """Current shipping form order id — do not treat it as a duplicate upload."""
    raw = str(
        request.GET.get("exclude_order")
        or request.GET.get("edit_order_id")
        or request.POST.get("edit_order_id")
        or ""
    ).strip()
    if raw.isdigit():
        return int(raw)
    return None


def _attach_ozon_stock_shortages(
    payloads,
    *,
    agency: Agency | None = None,
    exclude_order_id: int | None = None,
    reserve_rows: list[dict] | None = None,
) -> None:
    """Attach read-only Ozon shortage details, including blocking shipping reserves."""
    if isinstance(payloads, dict):
        payload_list = [payloads]
    elif isinstance(payloads, (list, tuple)):
        payload_list = [row for row in payloads if isinstance(row, dict)]
    else:
        payload_list = []

    shortage_entries: list[tuple[dict, dict]] = []
    sku_codes: set[str] = set()
    barcodes: set[str] = set()
    for payload in payload_list:
        details: list[dict] = []
        for row in payload.get("composition") or []:
            if not isinstance(row, dict):
                continue
            required_qty = max(int(row.get("quantity") or 0), 0)
            selected_qty = max(int(row.get("applied_qty") or 0), 0)
            if row.get("remaining_qty") is None:
                shortage_qty = max(required_qty - selected_qty, 0)
            else:
                shortage_qty = max(int(row.get("remaining_qty") or 0), 0)
            if shortage_qty <= 0:
                continue
            sku_code = str(row.get("offer_id") or row.get("sku_code") or "").strip()
            barcode = str(row.get("barcode") or "").strip()
            detail = {
                "sku_code": sku_code,
                "barcode": barcode,
                "name": str(row.get("name") or "").strip(),
                "required_qty": required_qty,
                "selected_qty": selected_qty,
                "shortage_qty": shortage_qty,
                "reserved_qty": 0,
                "blocking_reserves": [],
                "blocking_orders": [],
            }
            details.append(detail)
            shortage_entries.append((payload, detail))
            if sku_code:
                sku_codes.add(sku_code)
            if barcode:
                barcodes.add(barcode)
        payload["stock_shortages"] = details

    if not shortage_entries:
        return

    if reserve_rows is None:
        if agency is None or (not sku_codes and not barcodes):
            reserve_rows = []
        else:
            active_statuses = {
                WarehouseReserve.STATUS_ACTIVE,
                WarehouseReserve.STATUS_PARTIALLY_ALLOCATED,
                WarehouseReserve.STATUS_ALLOCATED,
                WarehouseReserve.STATUS_PARTIALLY_SATISFIED,
            }
            identity_filter = Q()
            if sku_codes:
                identity_filter |= Q(sku_code__in=sku_codes)
            if barcodes:
                identity_filter |= Q(barcode__in=barcodes)
            reserve_qs = WarehouseReserve.objects.filter(
                identity_filter,
                agency=agency,
                reserve_type=WarehouseReserve.TYPE_SHIPPING,
                context_type="shipping",
                status__in=active_statuses,
            )
            excluded_context_ids: set[str] = set()
            if exclude_order_id:
                excluded_context_ids.add(str(exclude_order_id))
                current_number = (
                    ShippingOrder.objects.filter(pk=exclude_order_id, agency=agency)
                    .values_list("number", flat=True)
                    .first()
                )
                if current_number:
                    excluded_context_ids.add(str(current_number).strip())
            if excluded_context_ids:
                reserve_qs = reserve_qs.exclude(context_id__in=excluded_context_ids)
            reserve_rows = list(
                reserve_qs.values(
                    "sku_code",
                    "barcode",
                    "context_id",
                    "qty_reserved",
                    "qty_satisfied",
                )
            )

    normalized_reserves: list[dict] = []
    for reserve in reserve_rows or []:
        if isinstance(reserve, dict):
            value = reserve
        else:
            value = {
                "sku_code": getattr(reserve, "sku_code", ""),
                "barcode": getattr(reserve, "barcode", ""),
                "context_id": getattr(reserve, "context_id", ""),
                "qty_reserved": getattr(reserve, "qty_reserved", 0),
                "qty_satisfied": getattr(reserve, "qty_satisfied", 0),
            }
        outstanding_qty = max(
            int(value.get("qty_reserved") or 0) - int(value.get("qty_satisfied") or 0),
            0,
        )
        context_id = str(value.get("context_id") or "").strip()
        if outstanding_qty <= 0 or not context_id:
            continue
        normalized_reserves.append(
            {
                "sku_code": str(value.get("sku_code") or "").strip().lower(),
                "barcode": str(value.get("barcode") or "").strip().lower(),
                "context_id": context_id,
                "qty": outstanding_qty,
            }
        )

    for _payload, detail in shortage_entries:
        target_sku = str(detail["sku_code"] or "").strip().lower()
        target_barcode = str(detail["barcode"] or "").strip().lower()
        by_order: defaultdict[str, int] = defaultdict(int)
        for reserve in normalized_reserves:
            reserve_barcode = reserve["barcode"]
            if target_barcode and reserve_barcode:
                matches = target_barcode == reserve_barcode
            else:
                matches = bool(target_sku and target_sku == reserve["sku_code"])
            if matches:
                by_order[_display_shipping_number(reserve["context_id"])] += int(reserve["qty"])
        blocking_reserves = [
            {"order": order_number, "qty": qty}
            for order_number, qty in sorted(by_order.items())
        ]
        detail["blocking_reserves"] = blocking_reserves
        detail["blocking_orders"] = [row["order"] for row in blocking_reserves]
        detail["reserved_qty"] = sum(row["qty"] for row in blocking_reserves)


@login_required
def shipping_ozon_supplies_list(request):
    """JSON: list active Ozon FBO supply orders for the selected client."""
    scope, role, client_agency = _request_scope(request)
    if scope is None or not _can_write(scope, role):
        return JsonResponse({"ok": False, "error": "Доступ запрещен"}, status=403)
    agency = _resolve_form_agency_for_shipping(request, scope, client_agency)
    if agency is None:
        return JsonResponse({"ok": False, "error": "Выберите клиента."}, status=400)
    from .ozon_supplies import list_ozon_supplies_for_agency

    exclude_order_id = _ozon_exclude_order_id(request)
    try:
        result = list_ozon_supplies_for_agency(agency, exclude_order_id=exclude_order_id)
    except Exception:
        logger.exception("Ozon supply list failed for agency=%s", getattr(agency, "id", None))
        return JsonResponse(
            {
                "ok": False,
                "error": (
                    "Не удалось получить заявки Ozon. "
                    "Попробуйте ещё раз или загрузите Excel."
                ),
                "orders": [],
            },
            status=400,
        )
    if not isinstance(result, dict):
        return JsonResponse(
            {"ok": False, "error": "Не удалось получить заявки Ozon.", "orders": []},
            status=400,
        )
    status = 200 if result.get("ok") else 400
    return JsonResponse(result, status=status)


@login_required
def shipping_ozon_supplies_detail(request, order_id: int):
    """JSON: one Ozon supply order mapped to shipping form fields."""
    scope, role, client_agency = _request_scope(request)
    if scope is None or not _can_write(scope, role):
        return JsonResponse({"ok": False, "error": "Доступ запрещен"}, status=403)
    agency = _resolve_form_agency_for_shipping(request, scope, client_agency)
    if agency is None:
        return JsonResponse({"ok": False, "error": "Выберите клиента."}, status=400)
    from .ozon_supplies import get_ozon_supply_for_agency

    exclude_order_id = _ozon_exclude_order_id(request)
    exclude_order = None
    if exclude_order_id is not None:
        exclude_order = ShippingOrder.objects.filter(
            pk=exclude_order_id,
            agency=agency,
        ).first()
    try:
        stock_rows = _shipping_stock_picker_rows(agency, exclude_order=exclude_order)
    except Exception:
        logger.exception(
            "Ozon supply stock rows failed for agency=%s order_id=%s",
            getattr(agency, "id", None),
            order_id,
        )
        stock_rows = []
    try:
        result = get_ozon_supply_for_agency(
            agency,
            int(order_id),
            stock_rows=stock_rows,
            enrich_gm=False,
            gm_enrich_limit=8,
            gm_pause_seconds=1.05,
            api_timeout=10,
            bundle_max_pages=4,
            include_gm_cargoes=True,
            exclude_order_id=exclude_order_id,
        )
    except Exception:
        logger.exception("Ozon supply detail failed for agency=%s order_id=%s", getattr(agency, "id", None), order_id)
        return JsonResponse(
            {
                "ok": False,
                "error": (
                    "Не удалось разобрать заявку Ozon. "
                    "Попробуйте ещё раз или загрузите Excel."
                ),
            },
            status=400,
        )
    if not isinstance(result, dict):
        return JsonResponse(
            {"ok": False, "error": "Не удалось разобрать заявку Ozon."},
            status=400,
        )
    _attach_ozon_stock_shortages(
        result,
        agency=agency,
        exclude_order_id=exclude_order_id,
    )
    status = 200 if result.get("ok") else 400
    return JsonResponse(result, status=status)


@login_required
@require_http_methods(["POST"])
def shipping_ozon_gm_preload(request, order_id: int):
    """Load a bounded next chunk of exact Ozon GM composition into cache."""
    scope, role, client_agency = _request_scope(request)
    if scope is None or not _can_write(scope, role):
        return JsonResponse({"ok": False, "error": "Доступ запрещен"}, status=403)
    agency = _resolve_form_agency_for_shipping(request, scope, client_agency)
    if agency is None:
        return JsonResponse({"ok": False, "error": "Выберите клиента."}, status=400)

    from .ozon_supplies import preload_ozon_gm_for_agency

    exclude_order_id = _ozon_exclude_order_id(request)
    try:
        result = preload_ozon_gm_for_agency(
            agency, int(order_id), exclude_order_id=exclude_order_id,
        )
    except Exception:
        logger.exception(
            "Ozon GM preload failed agency=%s order_id=%s",
            getattr(agency, "id", None),
            order_id,
        )
        return JsonResponse(
            {
                "ok": False,
                "error": "Не удалось загрузить состав ШК ГМ из Ozon. Повторите проверку.",
            },
            status=503,
        )
    if not result.get("ok"):
        message = str(result.get("error") or "").lower()
        retryable = any(token in message for token in ("429", "502", "503", "504", "timeout", "timed out", "соединен", "сети"))
        return JsonResponse(result, status=503 if retryable else 400)
    return JsonResponse(result)


@login_required
@require_http_methods(["POST"])
def shipping_ozon_supplies_batch(request):
    """Preview or submit several Ozon supplies without changing warehouse logic."""
    scope, role, client_agency = _request_scope(request)
    if scope is None or not _can_write(scope, role):
        return JsonResponse({"ok": False, "error": "Доступ запрещен"}, status=403)
    agency = _resolve_form_agency_for_shipping(request, scope, client_agency)
    if agency is None:
        return JsonResponse({"ok": False, "error": "Выберите клиента."}, status=400)
    try:
        body = json.loads(request.body.decode("utf-8") or "{}")
    except (UnicodeDecodeError, ValueError):
        return JsonResponse({"ok": False, "error": "Некорректный запрос."}, status=400)
    mode = str(body.get("mode") or "preview").strip().lower()
    raw_ids = body.get("order_ids") if isinstance(body.get("order_ids"), list) else []
    order_ids: list[int] = []
    for value in raw_ids:
        try:
            order_id = int(value)
        except (TypeError, ValueError):
            continue
        if order_id > 0 and order_id not in order_ids:
            order_ids.append(order_id)
    if not order_ids:
        return JsonResponse({"ok": False, "error": "Выберите хотя бы одну поставку Ozon."}, status=400)
    if len(order_ids) > 10:
        return JsonResponse({"ok": False, "error": "За один раз можно выбрать не более 10 поставок Ozon."}, status=400)

    from .ozon_supplies import get_ozon_supplies_batch_for_agency

    try:
        exclude_order_id = _ozon_exclude_order_id(request)
        exclude_order = None
        if exclude_order_id is not None:
            exclude_order = ShippingOrder.objects.filter(
                pk=exclude_order_id,
                agency=agency,
            ).first()
        stock_rows = _shipping_stock_picker_rows(agency, exclude_order=exclude_order)
        result = get_ozon_supplies_batch_for_agency(
            agency,
            order_ids,
            stock_rows=stock_rows,
            exclude_order_id=exclude_order_id,
            combine_for_one_request=mode == "preview",
        )
        _attach_ozon_stock_shortages(
            result.get("orders") or [],
            agency=agency,
            exclude_order_id=exclude_order_id,
        )
        if not result.get("ok"):
            return JsonResponse(result, status=400)
        if mode == "preview":
            return JsonResponse(result)
        if mode != "submit":
            return JsonResponse({"ok": False, "error": "Неизвестный режим запроса."}, status=400)

        ship_date_raw = str(body.get("ship_date") or "").strip()
        try:
            ship_date = date.fromisoformat(ship_date_raw)
        except ValueError:
            return JsonResponse(
                {"ok": False, "error": "Укажите дату отгрузки со склада."},
                status=400,
            )
        payload_by_id = {
            int(row.get("order_id")): row
            for row in result.get("orders") or []
            if row.get("order_id")
        }
        payloads = [payload_by_id[order_id] for order_id in order_ids if order_id in payload_by_id]
        not_ready = [
            str(row.get("order_number") or row.get("order_id"))
            for row in payloads
            if not row.get("ready_for_batch")
        ]
        if len(payloads) != len(order_ids) or not_ready:
            return JsonResponse(
                {
                    "ok": False,
                    "error": (
                        "Не все выбранные поставки готовы целыми коробами. "
                        "Откройте подсвеченные поставки для поштучного подбора."
                    ),
                    "not_ready": not_ready,
                    **result,
                },
                status=409,
            )
        created = create_ozon_api_shipping_orders_batch(
            agency=agency,
            user=request.user,
            payloads=payloads,
            stock_rows=stock_rows,
            ship_date=ship_date,
            supply_type=str(body.get("supply_type") or ShippingOrder.SUPPLY_BOX),
            vehicle_type=str(body.get("vehicle_type") or ShippingOrder.VEHICLE_FULFILLMENT),
            comment=str(body.get("comment") or "").strip(),
        )
        return JsonResponse(
            {
                "ok": True,
                "created": [
                    {"id": order.pk, "number": order.number}
                    for order in created
                ],
                "redirect_url": (
                    reverse("shipping:detail", kwargs={"pk": created[0].pk})
                    + f"?client={int(agency.id)}"
                    if created
                    else ""
                ),
            }
        )
    except ValidationError as exc:
        return JsonResponse(
            {"ok": False, "error": " ".join(str(message) for message in exc.messages)},
            status=400,
        )
    except Exception:
        logger.exception("Ozon batch failed for agency=%s", getattr(agency, "id", None))
        return JsonResponse(
            {
                "ok": False,
                "error": "Не удалось обработать выбранные поставки Ozon. Повторите попытку.",
            },
            status=500,
        )


@login_required
def shipping_detail(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")

    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency", "created_by", "marketplace").prefetch_related(
            "items",
            "reserves",
        ),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")

    can_write = _can_write(scope, role)
    can_edit_items = _can_edit_items(scope, role, order)
    can_submit_for_approval = _can_submit_for_approval(scope, role, order)
    can_manager_approve = _can_manager_approve(scope, role, order)
    can_manager_reopen = _can_manager_reopen(scope, role, order)
    can_set_shipping_priority = _can_set_shipping_priority(scope, role, order)
    can_storekeeper_accept = _can_storekeeper_accept(scope, role, order)
    can_storekeeper_pick = _can_storekeeper_pick(scope, role, order)
    pick_readiness = shipping_pick_readiness(order) if can_storekeeper_pick else {"can_pick": False, "reason": ""}
    can_storekeeper_pack = _can_storekeeper_pack(scope, role, order)
    can_storekeeper_manage_packing = _can_storekeeper_manage_packing(scope, role, order)
    can_cancel = _can_cancel(scope, role, order)
    warehouse_cancel_request = _warehouse_cancel_request_payload(order)
    client_cancel_request = _client_cancel_request_payload(order)
    can_request_warehouse_cancel = (
        scope == "staff"
        and _is_manager_role(role)
        and _requires_warehouse_cancel_confirmation(order)
        and not warehouse_cancel_request
    )
    can_review_warehouse_cancel = (
        scope == "staff"
        and _is_storekeeper_role(role)
        and bool(warehouse_cancel_request)
    )
    can_edit_order_form = _can_edit_order_form(scope, role, order)
    can_request_shipping_discrepancy_action = (
        scope == "staff"
        and _is_storekeeper_role(role)
        and not order.is_closed()
        and not str(order.shipping_discrepancy_status or "").strip()
        and not bool(_shipping_loose_packing_initial_state(order).get("loose_items"))
    )
    from .discrepancy import (
        can_add_discrepancy_item,
        can_approve_shipping_discrepancy,
        can_confirm_shipping_discrepancy_pick,
    )

    can_approve_shipping_discrepancy_action = can_approve_shipping_discrepancy(
        scope,
        role,
        order,
    )
    can_add_discrepancy_item_action = can_add_discrepancy_item(
        scope,
        role,
        order,
    )
    can_confirm_shipping_discrepancy_pick_action = (
        can_confirm_shipping_discrepancy_pick(scope, role, order)
    )
    supplemental_pick_preview = {"can_create": False, "reason": ""}
    if can_storekeeper_pick:
        from otg_reachtruck.services import get_otg_supplemental_pick_preview

        supplemental_pick_preview = get_otg_supplemental_pick_preview(order)
    can_create_otg_supplemental_pick = bool(supplemental_pick_preview.get("can_create"))
    can_complete_transfer = (
        scope == "staff"
        and _is_storekeeper_role(role)
        and order.delivery_type == ShippingOrder.DELIVERY_TRANSFER
        and order.status == ShippingOrder.STATUS_PACKED
    )
    if request.method == "POST":
        if not can_write:
            return HttpResponseForbidden("Доступ запрещен")
        stock_rows_for_order = _shipping_stock_picker_rows(order.agency, exclude_order=order)
        action_result = handle_shipping_detail_action(
            action=(request.POST.get("action") or "").strip(),
            order=order,
            request=request,
            role=role,
            stock_rows_for_order=stock_rows_for_order,
            permissions=ShippingDetailActionPermissions(
                can_edit_items=can_edit_items,
                can_submit_for_approval=can_submit_for_approval,
                can_manager_approve=can_manager_approve,
                can_manager_reopen=can_manager_reopen,
                can_set_shipping_priority=can_set_shipping_priority,
                can_storekeeper_accept=can_storekeeper_accept,
                can_storekeeper_pick=can_storekeeper_pick,
                can_request_shipping_discrepancy=can_request_shipping_discrepancy_action,
                can_approve_shipping_discrepancy=can_approve_shipping_discrepancy_action,
                can_add_discrepancy_item=can_add_discrepancy_item_action,
                can_confirm_shipping_discrepancy_pick=can_confirm_shipping_discrepancy_pick_action,
                can_create_otg_supplemental_pick=can_create_otg_supplemental_pick,
                can_complete_transfer=can_complete_transfer,
                can_cancel=can_cancel,
                can_request_warehouse_cancel=can_request_warehouse_cancel,
                can_review_warehouse_cancel=can_review_warehouse_cancel,
            ),
            parse_selected_stock_items=_parse_selected_stock_items,
            selected_stock_rows_with_boxes=_selected_stock_rows_with_boxes,
            parse_box_count_from_comment=_parse_box_count_from_comment,
            log_update=_log_update,
        )
        for level, text in action_result.messages:
            getattr(messages, level)(request, text)
        if action_result.redirect_url:
            return redirect(action_result.redirect_url)
        if action_result.redirect_name == "shipping:packing":
            return redirect("shipping:packing", pk=order.pk)
        return redirect("shipping:detail", pk=order.pk)
    context = build_shipping_detail_page_context(
        request=request,
        order=order,
        scope=scope,
        role=role,
    )
    context.update(
        {
            "warehouse_cancel_request_pending": bool(warehouse_cancel_request),
            "warehouse_cancel_reason": str(warehouse_cancel_request.get("cancel_reason") or "").strip(),
            "client_cancel_request_pending": bool(client_cancel_request),
            "client_cancel_reason": str(client_cancel_request.get("cancel_reason") or "").strip(),
            "can_request_warehouse_cancel": can_request_warehouse_cancel,
            "can_review_warehouse_cancel": can_review_warehouse_cancel,
            "can_set_shipping_priority": can_set_shipping_priority,
            "can_append_shipping_documents": _can_manager_append_shipping_documents(scope, role, order),
            "shipping_order_is_priority": _shipping_order_is_manual_priority(order),
            "can_create_otg_supplemental_pick": can_create_otg_supplemental_pick,
            "supplemental_pick_preview": supplemental_pick_preview,
            "can_confirm_shipping_discrepancy_pick": can_confirm_shipping_discrepancy_pick_action,
        }
    )
    context["detail_otg_box_rows"] = _shipping_detail_otg_box_rows(order, context.get("packing_summary"))
    ozon_label_rows = (context.get("ozon_detail_summary") or {}).get("gm_cargoes") or []
    context["can_download_ozon_gm_labels"] = bool(
        scope == "staff"
        and ozon_label_rows
        and all(
            str(row.get("gm_barcode") or "").strip()
            and str(row.get("supply_id") or "").strip()
            and str(row.get("cargo_id") or "").strip()
            for row in ozon_label_rows
            if isinstance(row, dict)
        )
        and len([row for row in ozon_label_rows if isinstance(row, dict)]) == len(ozon_label_rows)
    )
    return render(request, "shipping/detail.html", context)


@login_required
@require_http_methods(["POST"])
def shipping_ozon_gm_labels(request, pk: int):
    scope, _role, client_agency = _request_scope(request)
    if scope != "staff":
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency", "marketplace"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    marketplace_name = str(getattr(order.marketplace, "name", "") or "").strip().casefold()
    if "ozon" not in marketplace_name:
        raise Http404("Этикетки ШК ГМ доступны только для заявок Ozon.")

    summary = _ozon_api_summary_for_form(order)
    gm_cargoes = summary.get("gm_cargoes") or []
    expected_gm_barcodes = summary.get("gm_barcodes") or []
    from .ozon_supplies import fetch_ozon_gm_label_documents

    documents, error, pending = fetch_ozon_gm_label_documents(
        order.agency,
        gm_cargoes,
        expected_gm_barcodes=expected_gm_barcodes,
    )
    if pending:
        messages.info(request, error)
        return redirect("shipping:detail", pk=order.pk)
    if error or not documents:
        messages.error(request, error or "Ozon не вернул этикетки ШК ГМ.")
        return redirect("shipping:detail", pk=order.pk)

    safe_order = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(order.number or order.pk)).strip("_")
    if len(documents) == 1:
        response = HttpResponse(documents[0]["content"], content_type="application/pdf")
        filename = f"ozon_gm_labels_{safe_order or order.pk}.pdf"
    else:
        archive = BytesIO()
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
            for document in documents:
                supply_id = int(document["supply_id"])
                bundle.writestr(f"ozon_gm_labels_supply_{supply_id}.pdf", document["content"])
        response = HttpResponse(archive.getvalue(), content_type="application/zip")
        filename = f"ozon_gm_labels_{safe_order or order.pk}.zip"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "private, no-store"
    response["X-Content-Type-Options"] = "nosniff"
    return response


def _shipping_detail_item_signature(item: dict | None) -> tuple[str, str]:
    if not isinstance(item, dict):
        return ("", "")
    sku_code = str(
        item.get("sku_code")
        or item.get("sku")
        or item.get("article")
        or item.get("requested_article")
        or ""
    ).strip().lower()
    barcode = str(item.get("barcode") or item.get("bar_code") or "").strip().lower()
    if not sku_code and not barcode:
        return ("", "")
    return (sku_code, barcode)


def _shipping_detail_loose_item_signatures(payload: dict) -> list[tuple[str, str]]:
    signatures: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for item in payload.get("otg_box_composition") or []:
        signature = _shipping_detail_item_signature(item if isinstance(item, dict) else None)
        if signature != ("", "") and signature not in seen:
            signatures.append(signature)
            seen.add(signature)
    return signatures


def _shipping_detail_loose_box_links(
    order: ShippingOrder,
    reachtruck_rows: list[dict],
    delivered_boxes: list[dict] | None = None,
) -> dict[str, list[dict]]:
    source_box_codes = {
        str(code or "").strip().lower()
        for row in reachtruck_rows
        for code in (row.get("box_codes") or [])
        if str(code or "").strip()
    }
    boxes_by_signature: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for box in delivered_boxes if delivered_boxes is not None else _shipping_delivered_boxes(order):
        box_code = str(box.get("box_code") or "").strip()
        if not box_code or box_code.lower() in source_box_codes:
            continue
        linked_box = {
            "box_code": box_code,
            "qty": int(box.get("qty") or 0),
            "composition": str(box.get("barcode_preview") or "-").strip() or "-",
        }
        for item in box.get("items") or []:
            signature = _shipping_detail_item_signature(item if isinstance(item, dict) else None)
            if signature == ("", ""):
                continue
            boxes_by_signature[signature].append(linked_box)

    links_by_source_code: dict[str, list[dict]] = defaultdict(list)
    for row in reachtruck_rows:
        status_code = str(row.get("task_status_code") or "").strip().lower()
        if status_code in {"canceled", "cancelled"}:
            continue
        if not row.get("is_loose_partial"):
            continue
        signatures = [signature for signature in (row.get("loose_item_signatures") or []) if signature != ("", "")]
        if not signatures:
            continue
        for raw_source_code in row.get("box_codes") or []:
            source_code = str(raw_source_code or "").strip().lower()
            if not source_code:
                continue
            seen_linked_codes = {
                str(link.get("box_code") or "").strip().lower()
                for link in links_by_source_code.get(source_code, [])
            }
            for signature in signatures:
                for linked_box in boxes_by_signature.get(signature) or []:
                    linked_code = str(linked_box.get("box_code") or "").strip().lower()
                    if linked_code and linked_code not in seen_linked_codes:
                        links_by_source_code[source_code].append(dict(linked_box))
                        seen_linked_codes.add(linked_code)
    return dict(links_by_source_code)


def _shipping_detail_reachtruck_row_priority(row: dict) -> int:
    status_code = str(row.get("task_status_code") or "").strip().lower()
    priority = 0
    if status_code in {"done", "completed"}:
        priority += 100
    elif status_code in {"canceled", "cancelled"}:
        priority -= 100
    if _to_int(row.get("done_qty")) > 0:
        priority += 20
    if row.get("completed_at"):
        priority += 10
    return priority


def _shipping_detail_delivered_box_facts(box: dict | None) -> dict:
    if not isinstance(box, dict):
        return {"qty": 0, "composition": "-"}
    return {
        "qty": int(box.get("qty") or 0),
        "composition": str(box.get("barcode_preview") or "-").strip() or "-",
    }


def _shipping_detail_otg_box_rows(order: ShippingOrder, packing_summary: dict | None) -> list[dict]:
    reachtruck_rows, boxes_by_code = _shipping_report_reachtruck_rows(order)
    delivered_boxes = _shipping_delivered_boxes(order)
    delivered_boxes_by_code = {
        str(box.get("box_code") or "").strip().lower(): box
        for box in delivered_boxes
        if str(box.get("box_code") or "").strip()
    }
    loose_links_by_source_code = _shipping_detail_loose_box_links(order, reachtruck_rows, delivered_boxes)
    linked_actual_box_codes = {
        str(link.get("box_code") or "").strip().lower()
        for links in loose_links_by_source_code.values()
        for link in links
        if str(link.get("box_code") or "").strip()
    }
    rows_by_code: dict[str, dict] = {}
    for box in delivered_boxes:
        box_code = str(box.get("box_code") or "").strip()
        if not box_code or box_code.lower() in linked_actual_box_codes:
            continue
        facts = _shipping_detail_delivered_box_facts(box)
        source = boxes_by_code.get(box_code.lower(), {})
        rows_by_code[box_code.lower()] = {
            "box_code": box_code,
            "qty": facts["qty"],
            "composition": facts["composition"],
            "source_pallet": source.get("source_pallet") or "-",
            "source_place": source.get("source_place") or "-",
            "final_pallet": str(box.get("pallet_code") or box.get("pallet_label") or "").strip() or "-",
            "driver": source.get("driver") or "-",
            "completed_at": source.get("completed_at"),
            "linked_box_rows": [],
            "_priority": 50,
        }
    if isinstance(packing_summary, dict):
        for pallet in packing_summary.get("pallets") or []:
            if not isinstance(pallet, dict):
                continue
            final_pallet = str(pallet.get("code") or pallet.get("label") or "").strip() or "-"
            for box in pallet.get("boxes") or []:
                if not isinstance(box, dict):
                    continue
                box_code = str(box.get("box_code") or box.get("code") or "").strip()
                if not box_code:
                    continue
                if box_code.lower() in linked_actual_box_codes:
                    continue
                source = boxes_by_code.get(box_code.lower(), {})
                facts = _shipping_detail_delivered_box_facts(delivered_boxes_by_code.get(box_code.lower()) or box)
                rows_by_code[box_code.lower()] = {
                    "box_code": box_code,
                    "qty": facts["qty"],
                    "composition": facts["composition"],
                    "source_pallet": source.get("source_pallet") or "-",
                    "source_place": source.get("source_place") or "-",
                    "final_pallet": final_pallet,
                    "driver": source.get("driver") or "-",
                    "completed_at": source.get("completed_at"),
                    "linked_box_rows": [],
                    "_priority": 200,
                }
    for task_row in reachtruck_rows:
        for box_code in task_row.get("box_codes") or []:
            code = str(box_code or "").strip()
            if not code:
                continue
            linked_boxes = list(loose_links_by_source_code.get(code.lower()) or [])
            linked_qty = sum(int(link.get("qty") or 0) for link in linked_boxes)
            linked_composition = "; ".join(
                str(link.get("composition") or "").strip()
                for link in linked_boxes
                if str(link.get("composition") or "").strip() and str(link.get("composition") or "").strip() != "-"
            )
            delivered_facts = _shipping_detail_delivered_box_facts(delivered_boxes_by_code.get(code.lower()))
            candidate = {
                "box_code": code,
                "qty": linked_qty if linked_boxes else delivered_facts["qty"],
                "composition": linked_composition or delivered_facts["composition"],
                "source_pallet": task_row.get("source_pallet") or "-",
                "source_place": task_row.get("source_place") or "-",
                "final_pallet": "-",
                "driver": task_row.get("driver") or "-",
                "completed_at": task_row.get("completed_at"),
                "linked_box_rows": linked_boxes,
                "_priority": _shipping_detail_reachtruck_row_priority(task_row),
            }
            existing = rows_by_code.get(code.lower())
            if existing is None or int(candidate.get("_priority") or 0) > int(existing.get("_priority") or 0):
                rows_by_code[code.lower()] = candidate
    rows = sorted(rows_by_code.values(), key=lambda row: str(row.get("box_code") or "").lower())
    for row in rows:
        row.pop("_priority", None)
    return rows


def _shipping_report_location_from_task(task) -> dict:
    payload = task.payload if isinstance(getattr(task, "payload", None), dict) else {}
    location = payload.get("from_location") if isinstance(payload.get("from_location"), dict) else {}
    if location:
        return dict(location)
    return {
        "zone": str(getattr(task, "from_zone", "") or "").strip(),
        "row": getattr(task, "from_row", None),
        "section": getattr(task, "from_section", None),
        "tier": getattr(task, "from_tier", None),
        "cell": getattr(task, "from_cell", None),
    }


def _shipping_report_location_label(location: dict | None) -> str:
    if not isinstance(location, dict):
        return "-"
    zone = str(location.get("zone") or "").strip()
    row = _to_int(location.get("row"))
    section = _to_int(location.get("section"))
    tier = _to_int(location.get("tier"))
    cell = _to_int(location.get("cell"))
    if not zone:
        return "-"
    if zone == "OS":
        line_label = _shipping_report_os_line_label(section)
        if line_label and row and tier and cell:
            return f"{line_label}-{row}/{tier}-{cell}"
        if line_label and row:
            return f"{line_label}-{row}"
        return "OS"
    if row and section and tier and cell:
        return f"{zone}-{row}/{section}-{tier}-{cell}"
    if row and section and tier:
        return f"{zone}-{row}/{section}-{tier}"
    if row and section:
        return f"{zone}-{row}/{section}"
    if row:
        return f"{zone}-{row}"
    return zone


def _shipping_report_os_line_label(section: int) -> str:
    labels = {
        1: "0",
        2: "A",
        3: "B",
        4: "C",
        5: "D",
        6: "E",
        7: "F",
        8: "G",
        9: "I",
    }
    return labels.get(_to_int(section), str(_to_int(section) or ""))


def _shipping_report_user_label(user, fallback: str = "") -> str:
    if user:
        full_name = str(user.get_full_name() or "").strip()
        if full_name:
            return full_name
        username = str(user.get_username() or "").strip()
        if username:
            return username
    return str(fallback or "").strip() or "-"


def _shipping_report_box_codes_from_payload(payload: dict, claims: list) -> list[str]:
    codes: list[str] = []
    scan_fact_mode = bool(isinstance(payload, dict) and payload.get("otg_scan_fact_mode") == "scan_facts_v1")
    actual_keys = ("picked_boxes", "shipping_arrived_box_codes")
    fallback_keys = ("planned_box_codes", "requested_boxes", "reserved_box_codes")
    for key in actual_keys:
        for raw_code in payload.get(key) or []:
            code = str(raw_code or "").strip()
            if code and code.lower() not in {existing.lower() for existing in codes}:
                codes.append(code)
    execution = payload.get("execution") if isinstance(payload.get("execution"), dict) else {}
    for raw_code in execution.get("boxes_scanned") or []:
        code = str(raw_code or "").strip()
        if code and code.lower() not in {existing.lower() for existing in codes}:
            codes.append(code)
    mobile_execution = payload.get("mobile_execution") if isinstance(payload.get("mobile_execution"), dict) else {}
    for raw_code in mobile_execution.get("boxes_scanned") or []:
        code = str(raw_code or "").strip()
        if code and code.lower() not in {existing.lower() for existing in codes}:
            codes.append(code)
    for claim in claims:
        code = str(getattr(claim, "box_code", "") or "").strip()
        if code and code.lower() not in {existing.lower() for existing in codes}:
            codes.append(code)
    if codes:
        return codes
    if scan_fact_mode:
        return []
    for key in fallback_keys:
        for raw_code in payload.get(key) or []:
            code = str(raw_code or "").strip()
            if code and code.lower() not in {existing.lower() for existing in codes}:
                codes.append(code)
    return codes


def _shipping_report_reachtruck_rows(order: ShippingOrder) -> tuple[list[dict], dict[str, dict]]:
    OtgDeliveryRequest = apps.get_model("otg_reachtruck", "OtgDeliveryRequest")
    BoxClaim = apps.get_model("reachtruck", "BoxClaim")
    requests = list(
        OtgDeliveryRequest.objects.select_related("move_request", "requested_by")
        .prefetch_related("pallet_plans__move_task", "pallet_plans__move_task__assigned_to")
        .filter(shipping_order=order)
        .order_by("created_at", "id")
    )
    task_ids = []
    for delivery_request in requests:
        for plan in delivery_request.pallet_plans.all():
            if plan.move_task_id:
                task_ids.append(plan.move_task_id)
    claims_by_task: dict[int, list] = defaultdict(list)
    if task_ids:
        for claim in BoxClaim.objects.filter(move_task_id__in=task_ids).order_by("box_code", "id"):
            claims_by_task[int(claim.move_task_id)].append(claim)

    rows: list[dict] = []
    boxes_by_code: dict[str, dict] = {}
    for delivery_request in requests:
        for plan in delivery_request.pallet_plans.all():
            task = plan.move_task
            if task is None:
                continue
            payload = task.payload if isinstance(task.payload, dict) else {}
            claims = claims_by_task.get(int(task.id), [])
            box_codes = _shipping_report_box_codes_from_payload(payload, claims)
            source_location = _shipping_report_location_from_task(task)
            source_place = _shipping_report_location_label(source_location)
            if source_place == "-":
                source_place = str(payload.get("from_label") or "").strip() or "-"
            row = {
                "request_id": delivery_request.id,
                "task_id": task.id,
                "task_status_code": str(task.status or "").strip(),
                "task_status": task.get_status_display(),
                "move_mode_code": str(task.move_mode or "").strip(),
                "move_mode": task.get_move_mode_display(),
                "is_loose_partial": bool(payload.get("ship_as_loose_units") or payload.get("picked_loose_units")),
                "loose_item_signatures": _shipping_detail_loose_item_signatures(payload),
                "source_pallet": str(plan.pallet_code or task.pallet_code or "").strip() or "-",
                "source_place": source_place,
                "target_place": str(task.to_zone or payload.get("to_zone") or "OTG").strip() or "OTG",
                "planned_boxes": int(plan.boxes_planned or len(plan.planned_box_codes or []) or 0),
                "done_qty": int(task.qty_done or 0),
                "box_codes": box_codes,
                "driver": _shipping_report_user_label(task.assigned_to, task.assigned_to_name),
                "created_at": task.created_at,
                "started_at": task.started_at,
                "completed_at": task.completed_at,
            }
            rows.append(row)
            for code in box_codes:
                boxes_by_code[code.lower()] = row
    return rows, boxes_by_code


def _shipping_report_pallet_rows(packing_summary: dict | None, boxes_by_code: dict[str, dict]) -> list[dict]:
    if not isinstance(packing_summary, dict):
        return []
    rows: list[dict] = []
    for pallet in packing_summary.get("pallets") or []:
        if not isinstance(pallet, dict):
            continue
        pallet_boxes = []
        for box in pallet.get("boxes") or []:
            if not isinstance(box, dict):
                continue
            box_code = str(box.get("box_code") or box.get("code") or "").strip()
            source = boxes_by_code.get(box_code.lower(), {})
            pallet_boxes.append(
                {
                    "box_code": box_code or "-",
                    "qty": int(box.get("qty") or 0),
                    "barcode_preview": str(box.get("barcode_preview") or "-").strip() or "-",
                    "source_pallet": source.get("source_pallet") or "-",
                    "source_place": source.get("source_place") or "-",
                    "driver": source.get("driver") or "-",
                    "completed_at": source.get("completed_at"),
                }
            )
        rows.append(
            {
                "label": str(pallet.get("label") or "-").strip() or "-",
                "code": str(pallet.get("code") or "").strip() or "-",
                "source_code": str(pallet.get("source_code") or "").strip() or "-",
                "box_count": int(pallet.get("box_count") or len(pallet_boxes) or 0),
                "qty": int(pallet.get("qty") or 0),
                "boxes": pallet_boxes,
            }
        )
    return rows


def _shipping_report_context(order: ShippingOrder, base_context: dict) -> dict:
    packing_summary = base_context.get("packing_summary")
    dispatch_stage = base_context.get("dispatch_stage") or _shipping_dispatch_stage(order, packing_summary=packing_summary)
    reachtruck_rows, boxes_by_code = _shipping_report_reachtruck_rows(order)
    payload = dispatch_stage.get("payload") if isinstance(dispatch_stage, dict) else {}
    return {
        "report_items": _shipping_report_item_rows(order),
        "report_reachtruck_rows": reachtruck_rows,
        "report_pallet_rows": _shipping_report_pallet_rows(packing_summary, boxes_by_code),
        "report_dispatch": {
            "trip_number": str(payload.get("trip_number") or "").strip() or "-",
            "vehicle_type": str(payload.get("vehicle_type_label") or "").strip() or "-",
            "vehicle_name": str(payload.get("trip_vehicle_name") or "").strip() or "-",
            "vehicle_number": str(payload.get("trip_vehicle_number") or payload.get("vehicle_number") or "").strip() or "-",
            "driver_name": str(payload.get("trip_driver_name") or payload.get("driver_name") or "").strip() or "-",
            "driver_phone": str(payload.get("trip_driver_phone") or payload.get("driver_phone") or "").strip() or "-",
            "destination": str(payload.get("destination_warehouse") or payload.get("destination_address") or "").strip() or "-",
            "loading_sequence": int(payload.get("loading_sequence") or 0),
            "delivery_sequence": int(payload.get("delivery_sequence") or 0),
            "act_logistician": dispatch_stage.get("logistician_signed_at_label") or "-",
            "act_manager": dispatch_stage.get("manager_signed_at_label") or "-",
            "act_sent": dispatch_stage.get("act_sent_at_label") or "-",
        },
    }


def _shipping_report_item_rows(order: ShippingOrder) -> list[SimpleNamespace]:
    final_rows = shipping_final_truth_rows(order)
    if final_rows:
        source_rows = final_rows
    else:
        source_rows = [
            {
                "sku_code": item.sku_code,
                "name": item.name,
                "size": item.size,
                "barcode": item.barcode,
                "goods_type": item.goods_type,
                "qty_requested": int(item.qty_requested or 0),
                "qty_reserved": int(item.qty_reserved or 0),
                "qty_shipped": int(item.qty_shipped or 0),
                "shortage_qty": max(
                    int(item.qty_requested or 0) - int(item.qty_shipped or 0),
                    0,
                ),
                "excess_qty": 0,
            }
            for item in order.items.all().order_by("sku_code", "barcode", "id")
        ]

    return [
        SimpleNamespace(
            sku_code=row.get("sku_code") or "-",
            article=row.get("sku_code") or "-",
            name=row.get("name") or "-",
            size=row.get("size") or "",
            barcode=row.get("barcode") or "-",
            goods_type=row.get("goods_type") or "",
            product_type=row.get("goods_type") or "",
            qty_requested=int(row.get("qty_requested") or 0),
            requested_qty=int(row.get("qty_requested") or 0),
            qty_reserved=int(row.get("qty_reserved") or 0),
            reserved_qty=int(row.get("qty_reserved") or 0),
            qty_shipped=int(row.get("qty_shipped") or 0),
            shipped_qty=int(row.get("qty_shipped") or 0),
            shortage_qty=int(row.get("shortage_qty") or 0),
            excess_qty=int(row.get("excess_qty") or 0),
        )
        for row in source_rows
        if isinstance(row, dict)
    ]


@login_required
def shipping_report(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency", "created_by", "marketplace").prefetch_related("items", "reserves"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    context = build_shipping_detail_page_context(
        request=request,
        order=order,
        scope=scope,
        role=role,
    )
    context.update(_shipping_report_context(order, context))
    return render(request, "shipping/report.html", context)


def _shipping_report_safe_filename(order: ShippingOrder) -> str:
    label = _display_shipping_number(order.number) or str(order.pk)
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(label or "")).strip("_")
    return f"shipping_report_{safe or order.pk}.xlsx"


def _shipping_box_composition_safe_filename(order: ShippingOrder) -> str:
    label = _display_shipping_number(order.number) or str(order.pk)
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(label or "")).strip("_")
    return f"shipping_boxes_{safe or order.pk}.xlsx"


def _shipping_packing_list_safe_filename(order: ShippingOrder) -> str:
    label = _display_shipping_number(order.number) or str(order.pk)
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(label or "")).strip("_")
    return f"packing_list_{safe or order.pk}.xlsx"


def _shipping_report_cell_value(value):
    if value is None:
        return ""
    return value


def _shipping_box_composition_item_label(item: dict) -> str:
    for key in ("barcode", "sku_code", "sku", "article", "name"):
        value = str(item.get(key) or "").strip()
        if value:
            return value
    return "-"


def _shipping_box_composition_qty(item: dict) -> int:
    for key in ("qty", "actual_qty", "count"):
        value = item.get(key)
        try:
            qty = int(value or 0)
        except (TypeError, ValueError):
            qty = 0
        if qty > 0:
            return qty
    return 0


def _shipping_box_composition_expiry(item: dict) -> str:
    for key in ("expiry_date", "expiration_date", "end_product_date", "best_before", "shelf_life"):
        value = str(item.get(key) or "").strip()
        if value:
            return value
    return ""


def _shipping_box_composition_rows(order: ShippingOrder) -> list[list[object]]:
    entry = (
        OrderAuditEntry.objects.filter(
            order_id=order.number,
            order_type="shipping",
            payload__act__in=["shipping_packing", "shipping_loose_packing"],
        )
        .order_by("-created_at")
        .first()
    )
    payload = entry.payload if entry and isinstance(entry.payload, dict) else {}
    rows: list[list[object]] = []
    if isinstance(payload, dict):
        for raw_box in payload.get("act_boxes") or []:
            if not isinstance(raw_box, dict):
                continue
            box_code = str(raw_box.get("code") or raw_box.get("box_code") or "").strip()
            if not box_code:
                continue
            for raw_item in raw_box.get("items") or []:
                if not isinstance(raw_item, dict):
                    continue
                qty = _shipping_box_composition_qty(raw_item)
                if qty <= 0:
                    continue
                rows.append([
                    _shipping_box_composition_item_label(raw_item),
                    qty,
                    box_code,
                    _shipping_box_composition_expiry(raw_item),
                ])
    if rows:
        return rows

    summary = _shipping_packing_summary(order)
    for pallet in (summary or {}).get("pallets") or []:
        if not isinstance(pallet, dict):
            continue
        for box in pallet.get("boxes") or []:
            if not isinstance(box, dict):
                continue
            box_code = str(box.get("box_code") or box.get("code") or "").strip()
            qty = int(box.get("qty") or 0)
            if box_code and qty > 0:
                rows.append([
                    str(box.get("barcode_preview") or "-").strip() or "-",
                    qty,
                    box_code,
                    _shipping_box_composition_expiry(box),
                ])
    return rows


def _shipping_box_composition_workbook(order: ShippingOrder) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Короба"
    ws.append(["Баркод товара", "Кол-во товаров", "ШК короба", "Срок годности"])
    for row in _shipping_box_composition_rows(order):
        ws.append(row)
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 48)
    return workbook


def _shipping_packing_list_entry_payload(order: ShippingOrder) -> dict:
    entry = (
        OrderAuditEntry.objects.filter(
            order_id=order.number,
            order_type="shipping",
            payload__act__in=["shipping_packing", "shipping_loose_packing"],
        )
        .order_by("-created_at", "-id")
        .first()
    )
    payload = entry.payload if entry and isinstance(entry.payload, dict) else {}
    return payload if isinstance(payload, dict) else {}


def _shipping_packing_list_pallet_by_box(payload: dict) -> dict[str, str]:
    pallet_by_box: dict[str, str] = {}
    for raw_pallet in payload.get("act_pallets") or []:
        if not isinstance(raw_pallet, dict):
            continue
        pallet_label = str(raw_pallet.get("label") or raw_pallet.get("code") or "").strip()
        for raw_box in raw_pallet.get("boxes") or []:
            if isinstance(raw_box, dict):
                box_code = str(raw_box.get("box_code") or raw_box.get("code") or "").strip()
            else:
                box_code = str(raw_box or "").strip()
            if box_code:
                pallet_by_box[box_code.lower()] = pallet_label or "-"
    return pallet_by_box


def _shipping_packing_list_box_rows(order: ShippingOrder, packing_summary: dict | None) -> list[dict]:
    payload = _shipping_packing_list_entry_payload(order)
    pallet_by_box = _shipping_packing_list_pallet_by_box(payload)
    rows: list[dict] = []
    for raw_box in payload.get("act_boxes") or []:
        if not isinstance(raw_box, dict):
            continue
        box_code = str(raw_box.get("code") or raw_box.get("box_code") or "").strip()
        if not box_code:
            continue
        items = [item for item in raw_box.get("items") or [] if isinstance(item, dict)]
        if not items:
            rows.append(
                {
                    "pallet": pallet_by_box.get(box_code.lower(), "-"),
                    "box": box_code,
                    "article": str(raw_box.get("barcode_preview") or "-").strip() or "-",
                    "name": "",
                    "barcode": "",
                    "qty": int(raw_box.get("qty") or 0),
                    "expiry": _shipping_box_composition_expiry(raw_box),
                }
            )
            continue
        for item in items:
            rows.append(
                {
                    "pallet": pallet_by_box.get(box_code.lower(), "-"),
                    "box": box_code,
                    "article": str(item.get("sku_code") or item.get("article") or "").strip(),
                    "name": str(item.get("name") or "").strip(),
                    "barcode": str(item.get("barcode") or item.get("bar_code") or "").strip(),
                    "qty": _shipping_box_composition_qty(item),
                    "expiry": _shipping_box_composition_expiry(item),
                }
            )
    if rows:
        return rows

    if isinstance(packing_summary, dict):
        for pallet in packing_summary.get("pallets") or []:
            if not isinstance(pallet, dict):
                continue
            pallet_label = str(pallet.get("label") or pallet.get("code") or "").strip() or "-"
            for box in pallet.get("boxes") or []:
                if not isinstance(box, dict):
                    continue
                box_code = str(box.get("box_code") or box.get("code") or "").strip()
                if not box_code:
                    continue
                rows.append(
                    {
                        "pallet": pallet_label,
                        "box": box_code,
                        "article": str(box.get("barcode_preview") or "-").strip() or "-",
                        "name": "",
                        "barcode": "",
                        "qty": int(box.get("qty") or 0),
                        "expiry": _shipping_box_composition_expiry(box),
                    }
                )
    return rows


def _shipping_packing_list_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _shipping_packing_list_enrich_box_rows(order: ShippingOrder, box_rows: list[dict]) -> list[dict]:
    box_codes = {
        str(row.get("box") or "").strip()
        for row in box_rows
        if str(row.get("box") or "").strip()
    }
    containers = {
        container.container_code.lower(): container
        for container in WarehouseContainer.objects.filter(
            agency=order.agency,
            container_code__in=box_codes,
        ).only(
            "container_code",
            "gross_weight_g",
            "width_mm",
            "height_mm",
            "depth_mm",
        )
    }

    articles = {
        str(row.get("article") or "").strip()
        for row in box_rows
        if str(row.get("article") or "").strip() not in {"", "-"}
    }
    barcodes = {
        str(row.get("barcode") or "").strip()
        for row in box_rows
        if str(row.get("barcode") or "").strip()
    }
    sku_by_article = {
        sku.sku_code.lower(): sku
        for sku in SKU.objects.filter(
            agency=order.agency,
            deleted=False,
            sku_code__in=articles,
        ).only("id", "sku_code", "weight_kg", "weight_net_kg", "weight_gross_kg")
    }
    sku_by_barcode = {
        barcode.value.strip().lower(): barcode.sku
        for barcode in SKUBarcode.objects.filter(
            sku__agency=order.agency,
            sku__deleted=False,
            value__in=barcodes,
        )
        .select_related("sku")
        .only(
            "value",
            "sku__id",
            "sku__sku_code",
            "sku__weight_kg",
            "sku__weight_net_kg",
            "sku__weight_gross_kg",
        )
    }

    enriched_rows: list[dict] = []
    for source_row in box_rows:
        row = dict(source_row)
        box_code = str(row.get("box") or "").strip()
        container = containers.get(box_code.lower())
        sku = sku_by_article.get(str(row.get("article") or "").strip().lower())
        if sku is None:
            sku = sku_by_barcode.get(str(row.get("barcode") or "").strip().lower())

        unit_weight_kg = None
        if sku is not None:
            unit_weight_kg = (
                _shipping_packing_list_decimal(sku.weight_net_kg)
                or _shipping_packing_list_decimal(sku.weight_kg)
                or _shipping_packing_list_decimal(sku.weight_gross_kg)
            )
        qty = int(row.get("qty") or 0)
        row["product_weight_kg"] = unit_weight_kg * qty if unit_weight_kg is not None else None

        gross_weight_g = getattr(container, "gross_weight_g", None) if container is not None else None
        row["box_gross_weight_kg"] = (
            Decimal(gross_weight_g) / Decimal("1000") if gross_weight_g is not None else None
        )
        dimensions = [
            getattr(container, "width_mm", None) if container is not None else None,
            getattr(container, "depth_mm", None) if container is not None else None,
            getattr(container, "height_mm", None) if container is not None else None,
        ]
        row["dimensions_mm"] = " x ".join(str(value) for value in dimensions) if all(dimensions) else ""
        row["volume_m3"] = (
            Decimal(dimensions[0] * dimensions[1] * dimensions[2]) / Decimal("1000000000")
            if all(dimensions)
            else None
        )
        enriched_rows.append(row)
    return enriched_rows


def _shipping_packing_list_box_summary_rows(box_rows: list[dict]) -> list[dict]:
    boxes: dict[str, dict] = {}
    for row in box_rows:
        box_code = str(row.get("box") or "").strip() or "-"
        key = box_code.lower()
        box = boxes.setdefault(
            key,
            {
                "pallet": row.get("pallet") or "-",
                "box": box_code,
                "contents": [],
                "qty": 0,
                "product_weight_kg": Decimal("0"),
                "product_weight_complete": True,
                "box_gross_weight_kg": row.get("box_gross_weight_kg"),
                "dimensions_mm": row.get("dimensions_mm") or "",
                "volume_m3": row.get("volume_m3"),
            },
        )
        qty = int(row.get("qty") or 0)
        box["qty"] += qty
        product_weight = row.get("product_weight_kg")
        if product_weight is None:
            box["product_weight_complete"] = False
        else:
            box["product_weight_kg"] += product_weight
        content_parts = [
            str(row.get("article") or "").strip(),
            str(row.get("name") or "").strip(),
            f"ШК {str(row.get('barcode') or '').strip()}" if str(row.get("barcode") or "").strip() else "",
            f"{qty} шт.",
        ]
        box["contents"].append(" | ".join(part for part in content_parts if part))

    result: list[dict] = []
    for box in boxes.values():
        box["contents"] = "\n".join(box["contents"])
        if not box.pop("product_weight_complete"):
            box["product_weight_kg"] = None
        result.append(box)
    return result


def _shipping_packing_list_append_kv(ws, row_number: int, label: str, value: object):
    ws.cell(row_number, 1, label)
    ws.cell(row_number, 2, _shipping_report_cell_value(value) or "-")
    ws.cell(row_number, 1).font = Font(bold=True, color="5B4B25")
    ws.cell(row_number, 1).fill = PatternFill("solid", fgColor="F8F1DF")
    ws.cell(row_number, 2).fill = PatternFill("solid", fgColor="FFFFFF")
    for column in (1, 2):
        cell = ws.cell(row_number, column)
        cell.border = _shipping_report_border()
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _shipping_packing_list_workbook(order: ShippingOrder, context: dict) -> Workbook:
    packing_summary = context.get("packing_summary") or _shipping_packing_summary(order) or {}
    report_context = _shipping_report_context(order, {**context, "packing_summary": packing_summary})
    item_rows = report_context.get("report_items") or _shipping_report_item_rows(order)
    box_rows = _shipping_packing_list_enrich_box_rows(
        order,
        _shipping_packing_list_box_rows(order, packing_summary),
    )
    box_summary_rows = _shipping_packing_list_box_summary_rows(box_rows)
    metrics = context.get("shipping_metrics") or {}
    dispatch = report_context.get("report_dispatch") or {}
    display_number = _display_shipping_number(order.number) or order.number
    client_name = getattr(getattr(order, "agency", None), "agn_name", "") or "-"
    marketplace_name = getattr(getattr(order, "marketplace", None), "name", "") or "-"
    destination = (
        str(getattr(order, "destination_warehouse", "") or "").strip()
        or str(getattr(order, "destination_address", "") or "").strip()
        or str(getattr(order, "transit_address", "") or "").strip()
        or str(dispatch.get("destination") or "").strip()
        or "-"
    )
    supply_number = (
        str(getattr(order, "supply_number", "") or "").strip()
        or str(getattr(order, "shipping_barcode", "") or "").strip()
        or str(getattr(order, "wb_supply_barcode", "") or "").strip()
        or "-"
    )
    pallet_count = int((packing_summary or {}).get("pallet_count") or metrics.get("pallet_count") or 0)
    box_count = int((packing_summary or {}).get("box_count") or metrics.get("box_count") or _order_box_count(order) or 0)
    known_box_weight_kg = sum(
        (row.get("box_gross_weight_kg") or Decimal("0"))
        for row in box_summary_rows
    )
    boxes_without_weight = sum(
        1 for row in box_summary_rows if row.get("box_gross_weight_kg") is None
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Упаковочный лист"
    ws.sheet_view.showGridLines = False
    ws.append(["УПАКОВОЧНЫЙ ЛИСТ / PACKING LIST"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
    ws.cell(1, 1).font = Font(bold=True, size=18, color="1F2933")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="F7E7B1")
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    _shipping_report_style_row(ws, 1, fill=PatternFill("solid", fgColor="F7E7B1"), height=30)

    document_date = date.today().strftime("%d.%m.%Y")
    created_at = order.created_at.strftime("%d.%m.%Y %H:%M") if order.created_at else "-"
    slot = "-"
    if order.slot_date and order.slot_time:
        slot = f"{order.slot_date.strftime('%d.%m.%Y')} {order.slot_time.strftime('%H:%M')}"
    elif order.slot_date:
        slot = order.slot_date.strftime("%d.%m.%Y")
    planned_ship = order.planned_ship_date.strftime("%d.%m.%Y") if order.planned_ship_date else "-"

    summary_rows = [
        ("Дата документа", document_date),
        ("Упаковочный лист", f"PL-{display_number}"),
        ("Заявка на отгрузку", display_number),
        ("Дата подачи заявки", created_at),
        ("Отправитель / клиент", client_name),
        ("Оператор склада", 'ООО "ФуллБокс"'),
        ("Получатель / направление", destination),
        ("Маркетплейс", marketplace_name),
        ("Номер поставки / ШК поставки", supply_number),
        ("Тип отгрузки", order.get_delivery_type_display()),
        ("Плановая отгрузка / слот", f"{planned_ship} / {slot}"),
        ("Тип поставки / мест", order.get_supply_type_display() if order.supply_type else order.get_place_type_display() or "-"),
        ("Паллет", pallet_count or "-"),
        ("Коробов", box_count or "-"),
        ("Вес коробов с товаром, кг", known_box_weight_kg if known_box_weight_kg else "-"),
        ("Коробов без указанного веса", boxes_without_weight or "-"),
        ("Номер авто", dispatch.get("vehicle_number") or getattr(order, "vehicle_number", "") or "-"),
        ("Водитель / телефон", " / ".join(part for part in [dispatch.get("driver_name"), dispatch.get("driver_phone") or getattr(order, "driver_phone", "")] if part and part != "-") or "-"),
        ("Комментарий", getattr(order, "comment", "") or "-"),
    ]
    row_no = 3
    for label, value in summary_rows:
        _shipping_packing_list_append_kv(ws, row_no, label, value)
        row_no += 1

    ws.append([])
    goods_title_row = ws.max_row + 1
    ws.append(["Состав товара"])
    ws.merge_cells(start_row=goods_title_row, start_column=1, end_row=goods_title_row, end_column=13)
    _shipping_report_style_row(ws, goods_title_row, fill=PatternFill("solid", fgColor="EFE7D8"), font=Font(bold=True, size=13, color="1F2933"), height=24)
    ws.append(["№", "Артикул", "Наименование", "Размер", "ШК", "Тип", "Кол-во", "Ед.", "Вес нетто, кг", "Вес брутто, кг", "Цена", "Сумма", "ТН ВЭД"])
    _shipping_report_style_row(ws, ws.max_row, fill=PatternFill("solid", fgColor="F4D47A"), font=Font(bold=True, color="4D3B12"), height=22)
    for index, item in enumerate(item_rows, start=1):
        qty = int(getattr(item, "shipped_qty", 0) or getattr(item, "reserved_qty", 0) or getattr(item, "requested_qty", 0) or 0)
        ws.append(
            [
                index,
                getattr(item, "article", "") or "-",
                getattr(item, "name", "") or "-",
                getattr(item, "size", "") or "",
                getattr(item, "barcode", "") or "",
                getattr(item, "product_type", "") or "",
                qty,
                "шт",
                "",
                "",
                "",
                "",
                "",
            ]
        )
        _shipping_report_style_row(ws, ws.max_row, fill=PatternFill("solid", fgColor="FFF8E5" if index % 2 else "FFFFFF"))

    _shipping_report_append_section(
        ws,
        "Короба, вес и состав",
        [
            "Паллета",
            "ШК короба",
            "Состав короба: артикул / наименование / ШК / количество",
            "Всего товара, шт.",
            "Расчетный вес товара, кг",
            "Вес короба с товаром, кг",
            "Габариты короба, мм",
            "Объем короба, м3",
        ],
        [
            [
                row.get("pallet") or "-",
                row.get("box") or "-",
                row.get("contents") or "-",
                int(row.get("qty") or 0),
                row.get("product_weight_kg"),
                row.get("box_gross_weight_kg"),
                row.get("dimensions_mm") or "",
                row.get("volume_m3"),
            ]
            for row in box_summary_rows
        ],
    )

    ws2 = workbook.create_sheet("Состав коробов")
    ws2.sheet_view.showGridLines = False
    ws2.append(
        [
            "Паллета",
            "ШК короба",
            "Артикул",
            "Наименование",
            "ШК товара",
            "Кол-во товара в коробе, шт.",
            "Расчетный вес товара, кг",
            "Вес короба с товаром, кг",
            "Срок годности",
            "Габариты короба, мм",
            "Объем короба, м3",
        ]
    )
    _shipping_report_style_row(ws2, 1, fill=PatternFill("solid", fgColor="F4D47A"), font=Font(bold=True, color="4D3B12"), height=22)
    for index, row in enumerate(box_rows, start=1):
        ws2.append(
            [
                row.get("pallet") or "-",
                row.get("box") or "-",
                row.get("article") or "-",
                row.get("name") or "",
                row.get("barcode") or "",
                int(row.get("qty") or 0),
                row.get("product_weight_kg"),
                row.get("box_gross_weight_kg"),
                row.get("expiry") or "",
                row.get("dimensions_mm") or "",
                row.get("volume_m3"),
            ]
        )
        _shipping_report_style_row(ws2, ws2.max_row, fill=PatternFill("solid", fgColor="FFF8E5" if index % 2 else "FFFFFF"))

    for sheet in workbook.worksheets:
        for column_index, column_cells in enumerate(sheet.columns, start=1):
            max_length = 0
            column_letter = get_column_letter(column_index)
            for cell in column_cells:
                max_length = max(max_length, len(str(cell.value or "")))
            sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 46)
        sheet.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 38
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 24
    ws2.column_dimensions["D"].width = 36
    for cell in ws["E"]:
        cell.number_format = "@"
    for column_letter in ("B", "C", "E"):
        for cell in ws2[column_letter]:
            cell.number_format = "@"
    for sheet in (ws, ws2):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and isinstance(cell.value, Decimal):
                    cell.number_format = "0.000"
    return workbook


def _shipping_report_border() -> Border:
    side = Side(style="thin", color="D7D0C3")
    return Border(left=side, right=side, top=side, bottom=side)


def _shipping_report_style_row(ws, row_number, fill=None, font=None, height=None):
    if height:
        ws.row_dimensions[row_number].height = height
    for cell in ws[row_number]:
        if isinstance(cell, MergedCell):
            continue
        cell.border = _shipping_report_border()
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font


def _shipping_report_append_section(ws, title, headers, rows):
    ws.append([])
    title_row = ws.max_row + 1
    ws.append([title])
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=max(len(headers), 2))
    title_cell = ws.cell(title_row, 1)
    title_cell.font = Font(bold=True, size=14, color="1F2933")
    title_cell.fill = PatternFill("solid", fgColor="EFE7D8")
    title_cell.alignment = Alignment(vertical="center")
    _shipping_report_style_row(ws, title_row, fill=PatternFill("solid", fgColor="EFE7D8"), height=24)
    if headers:
        ws.append(headers)
        _shipping_report_style_row(
            ws,
            ws.max_row,
            fill=PatternFill("solid", fgColor="F4D47A"),
            font=Font(bold=True, color="4D3B12"),
            height=22,
        )
    for row in rows:
        ws.append([_shipping_report_cell_value(value) for value in row])
        fill = PatternFill("solid", fgColor="FFF8E5" if ws.max_row % 2 else "FFFFFF")
        _shipping_report_style_row(ws, ws.max_row, fill=fill)


def _shipping_report_workbook(order: ShippingOrder, context: dict) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Отчет"
    ws.sheet_view.showGridLines = False

    report_title = f"Отчет об отгрузке {_display_shipping_number(order.number)}"
    client_name = getattr(getattr(order, "agency", None), "agn_name", "") or "-"
    marketplace_name = getattr(getattr(order, "marketplace", None), "name", "") or "-"
    metrics = context.get("shipping_metrics") or {}
    dispatch = context.get("report_dispatch") or {}

    ws.append([report_title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.cell(1, 1).font = Font(bold=True, size=18, color="222222")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="F7E7B1")
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    _shipping_report_style_row(ws, 1, fill=PatternFill("solid", fgColor="F7E7B1"), height=30)
    ws.append(["Клиент", client_name])
    ws.append(["Статус", getattr(order, "ui_status_label", "") or order.get_status_display()])
    ws.append(["Создана", order.created_at.strftime("%d.%m.%Y %H:%M") if order.created_at else ""])
    ws.append(["Склад назначения", getattr(order, "destination_warehouse", "") or "-"])
    ws.append(["Маркетплейс", marketplace_name])
    ws.append(["Коробов по заявке", metrics.get("box_count", "")])
    ws.append(["Итоговых паллет", metrics.get("pallet_count", "")])
    for row_number in range(2, 9):
        ws.cell(row_number, 1).font = Font(bold=True, color="75633A")
        ws.cell(row_number, 1).fill = PatternFill("solid", fgColor="F8F1DF")
        ws.cell(row_number, 2).fill = PatternFill("solid", fgColor="FFFFFF")
        ws.cell(row_number, 1).border = _shipping_report_border()
        ws.cell(row_number, 2).border = _shipping_report_border()
        ws.cell(row_number, 1).alignment = Alignment(vertical="center")
        ws.cell(row_number, 2).alignment = Alignment(vertical="center", wrap_text=True)

    _shipping_report_append_section(
        ws,
        "Машина и рейс",
        ["Поле", "Значение"],
        [
            ["Рейс", dispatch.get("trip", "-")],
            ["Транспорт", dispatch.get("transport", "-")],
            ["Номер авто", dispatch.get("vehicle_number", "-")],
            ["Водитель", dispatch.get("driver", "-")],
            ["Куда", dispatch.get("destination", "-")],
            ["Погрузка / доставка", dispatch.get("loading_delivery", "-")],
            ["Логист подписал", dispatch.get("logistician_signed", "-")],
            ["Менеджер закрыл", dispatch.get("manager_closed", "-")],
        ],
    )

    _shipping_report_append_section(
        ws,
        "Позиции заявки",
        [
            "Артикул",
            "Наименование",
            "Размер",
            "ШК",
            "Тип",
            "Запрошено",
            "Резерв",
            "Отгружено",
            "Недостача",
            "Излишек",
        ],
        [
            [
                getattr(item, "article", ""),
                getattr(item, "name", ""),
                getattr(item, "size", ""),
                getattr(item, "barcode", ""),
                getattr(item, "product_type", ""),
                getattr(item, "requested_qty", ""),
                getattr(item, "reserved_qty", ""),
                getattr(item, "shipped_qty", ""),
                getattr(item, "shortage_qty", ""),
                getattr(item, "excess_qty", ""),
            ]
            for item in context.get("report_items", [])
        ],
    )

    _shipping_report_append_section(
        ws,
        "Ричтрак: откуда привезли в OTG",
        ["Задание", "Статус", "Откуда", "Куда", "Паллета", "Короба", "Короба списком", "Исполнитель", "Обновлено"],
        [
            [
                row.get("label", ""),
                row.get("status", ""),
                row.get("source", ""),
                row.get("destination", ""),
                row.get("pallet", ""),
                row.get("box_count", ""),
                ", ".join(row.get("boxes", []) or []),
                row.get("assignee", ""),
                row.get("updated_at", ""),
            ]
            for row in context.get("report_reachtruck_rows", [])
        ],
    )

    _shipping_report_append_section(
        ws,
        "Итоговые паллеты отгрузки",
        ["Паллета", "Статус", "Коробов", "Количество", "Короба"],
        [
            [
                row.get("label", ""),
                row.get("status", ""),
                row.get("box_count", ""),
                row.get("qty", ""),
                ", ".join(row.get("boxes", []) or []),
            ]
            for row in context.get("report_pallet_rows", [])
        ],
    )

    _shipping_report_append_section(
        ws,
        "История",
        ["Дата", "Ответственный", "Роль", "Действие", "Комментарий"],
        [
            [
                entry.get("created_at").strftime("%d.%m.%Y %H:%M") if entry.get("created_at") else "",
                entry.get("actor_label", ""),
                entry.get("actor_role_label", ""),
                entry.get("action_label", ""),
                entry.get("text", ""),
            ]
            for entry in context.get("shipping_history", [])
        ],
    )

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 60)
    ws.freeze_panes = "A10"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 30)
    return workbook


@login_required
def shipping_report_xlsx(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency", "created_by", "marketplace").prefetch_related("items", "reserves"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    context = build_shipping_detail_page_context(
        order=order,
        scope=scope,
        role=role,
        request=request,
    )
    context.update(_shipping_report_context(order, context))
    workbook = _shipping_report_workbook(order, context)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_shipping_report_safe_filename(order)}"'
    return response


@login_required
def shipping_box_composition_xlsx(request, pk: int):
    scope, _role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    workbook = _shipping_box_composition_workbook(order)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_shipping_box_composition_safe_filename(order)}"'
    return response


@login_required
def shipping_packing_list_xlsx(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope is None:
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency", "created_by", "marketplace").prefetch_related("items", "reserves"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    context = build_shipping_detail_page_context(
        order=order,
        scope=scope,
        role=role,
        request=request,
    )
    workbook = _shipping_packing_list_workbook(order, context)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_shipping_packing_list_safe_filename(order)}"'
    return response


@login_required
def shipping_dispatch_act(request, pk: int):
    return handle_shipping_dispatch_act_request(request=request, pk=pk)


@login_required
def shipping_sign_dispatch_act_logistician(request, pk: int):
    return handle_shipping_dispatch_sign_logistician_request(request=request, pk=pk)


@login_required
def shipping_sign_dispatch_act_manager(request, pk: int):
    return handle_shipping_dispatch_sign_manager_request(request=request, pk=pk)


@login_required
def shipping_attachment_download(request, pk: int, attachment_id: int):
    return download_shipping_attachment(request=request, pk=pk, attachment_id=attachment_id)


def _shipping_documents_response(request, pk: int):
    return handle_shipping_documents_request(request=request, pk=pk)


@login_required
def shipping_documents(request, pk: int):
    return _shipping_documents_response(request, pk)


@login_required
def shipping_transport_note(request, pk: int):
    return _shipping_documents_response(request, pk)


@login_required
@xframe_options_sameorigin
def shipping_transport_note_pdf(request, pk: int):
    return build_shipping_transport_note_pdf_response(request=request, pk=pk)


@login_required
def shipping_transport_note_docx(request, pk: int):
    return build_shipping_transport_note_docx_response(request=request, pk=pk)


@login_required
def shipping_return_act_doc(request, pk: int):
    return build_shipping_return_act_response(request=request, pk=pk)


@login_required
def shipping_packing(request, pk: int):
    return handle_shipping_packing_request(request=request, pk=pk)


@login_required
@require_http_methods(["POST"])
def shipping_packing_draft(request, pk: int):
    order = get_object_or_404(ShippingOrder, pk=pk)
    scope, role, _client_agency = _request_scope(request)
    if not _can_storekeeper_manage_packing(scope, role, order):
        return JsonResponse({"saved": False, "error": "forbidden"}, status=403)
    try:
        boxes_state = json.loads(request.POST.get("boxes_json") or "[]")
        pallets_state = json.loads(request.POST.get("pallets_json") or "[]")
    except json.JSONDecodeError:
        return JsonResponse({"saved": False, "error": "bad_json"}, status=400)
    try:
        save_shipping_packing_draft(
            order,
            boxes_state=boxes_state,
            pallets_state=pallets_state,
            user=request.user,
        )
    except Exception:
        __import__("logging").getLogger(__name__).exception(
            "Failed to autosave shipping packing draft for order %s",
            order.pk,
        )
        return JsonResponse({"saved": False, "error": "save_failed"}, status=500)
    return JsonResponse({"saved": True})


@login_required
@require_http_methods(["POST"])
def shipping_packing_pallet_removal_request(request, pk: int):
    return handle_shipping_packing_pallet_removal_request(request=request, pk=pk)


@login_required
@require_http_methods(["GET", "POST"])
def shipping_packing_pallet_removal_review(request, pk: int):
    return handle_shipping_packing_pallet_removal_review_request(request=request, pk=pk)


@login_required
def shipping_loose_packing(request, pk: int):
    return handle_shipping_loose_packing_request(request=request, pk=pk)


def _shipping_loose_packing_draft_session_key(order: ShippingOrder) -> str:
    return f"shipping_loose_packing_draft:{int(order.pk)}"


@login_required
@require_http_methods(["POST"])
def shipping_loose_packing_draft(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope != "staff" or not _is_storekeeper_role(role):
        return JsonResponse({"saved": False, "error": "forbidden"}, status=403)
    order = get_object_or_404(ShippingOrder.objects.select_related("agency"), pk=pk)
    if not _can_access_order(scope, client_agency, order):
        return JsonResponse({"saved": False, "error": "forbidden"}, status=403)
    if len(request.body or b"") > 2_000_000:
        return JsonResponse({"saved": False, "error": "draft_too_large"}, status=413)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (UnicodeDecodeError, json.JSONDecodeError):
        return JsonResponse({"saved": False, "error": "bad_json"}, status=400)
    if not isinstance(payload, dict):
        return JsonResponse({"saved": False, "error": "bad_json"}, status=400)
    session_key = _shipping_loose_packing_draft_session_key(order)
    if str(payload.get("action") or "").strip() == "clear":
        request.session.pop(session_key, None)
        request.session.modified = True
        return JsonResponse({"saved": True, "cleared": True})
    draft = payload.get("draft")
    if not isinstance(draft, dict):
        return JsonResponse({"saved": False, "error": "bad_draft"}, status=400)
    boxes = draft.get("boxes")
    pallets = draft.get("pallets")
    if not isinstance(boxes, list) or not isinstance(pallets, list):
        return JsonResponse({"saved": False, "error": "bad_draft"}, status=400)
    if len(boxes) > 500 or len(pallets) > 500:
        return JsonResponse({"saved": False, "error": "draft_too_large"}, status=413)
    marking_units_count = sum(
        len(item.get("marking_units") or [])
        for box in boxes
        if isinstance(box, dict)
        for item in (box.get("items") or [])
        if isinstance(item, dict) and isinstance(item.get("marking_units") or [], list)
    )
    if marking_units_count > 10_000:
        return JsonResponse({"saved": False, "error": "draft_too_large"}, status=413)
    request.session[session_key] = {
        "signature": str(draft.get("signature") or "")[:20_000],
        "boxes": boxes,
        "pallets": pallets,
        "pendingMarkedItem": draft.get("pendingMarkedItem") if isinstance(draft.get("pendingMarkedItem"), dict) else None,
    }
    request.session.modified = True
    return JsonResponse({"saved": True, "marking_units": marking_units_count})


@login_required
@require_http_methods(["POST"])
def shipping_loose_packing_container_code(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope != "staff" or not _is_storekeeper_role(role):
        return HttpResponseForbidden("Р”РѕСЃС‚СѓРї Р·Р°РїСЂРµС‰РµРЅ")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Р”РѕСЃС‚СѓРї Р·Р°РїСЂРµС‰РµРЅ")
    if order.is_closed() or order.status not in {ShippingOrder.STATUS_PICKING, ShippingOrder.STATUS_PACKED}:
        return JsonResponse({"ok": False, "error": "not_allowed"}, status=403)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    goods_type = str(payload.get("goods_type") or "").strip()
    container_kind = normalize_container_kind(payload.get("kind"))
    if container_kind == "pallet":
        issued = issue_pallet_code(
            agency=order.agency,
            goods_type=goods_type,
            order_type="shipping",
            order_id=order.number,
        )
    else:
        issued = issue_container_code(
            agency=order.agency,
            goods_type=goods_type,
        )
    return JsonResponse({"ok": True, **issued})


@login_required
@require_http_methods(["POST"])
def shipping_loose_packing_marking_scan(request, pk: int):
    scope, role, client_agency = _request_scope(request)
    if scope != "staff" or not _is_storekeeper_role(role):
        return HttpResponseForbidden("Доступ запрещен")
    order = get_object_or_404(
        ShippingOrder.objects.select_related("agency"),
        pk=pk,
    )
    if not _can_access_order(scope, client_agency, order):
        return HttpResponseForbidden("Доступ запрещен")
    if order.is_closed() or order.status not in {ShippingOrder.STATUS_PICKING, ShippingOrder.STATUS_PACKED}:
        return JsonResponse({"ok": False, "error": "Упаковка этой заявки недоступна."}, status=403)
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    try:
        if payload.get("resolve_product_barcode") is True:
            from .marking_scan import resolve_shipping_product_barcode

            verified = resolve_shipping_product_barcode(
                order=order,
                barcode=payload.get("barcode"),
                candidate_keys=payload.get("candidate_keys"),
            )
        elif payload.get("resolve_from_gtin") is True:
            from .marking_scan import resolve_shipping_data_matrix

            verified = resolve_shipping_data_matrix(
                order=order,
                marking_code=str(payload.get("marking_code") or "").strip(),
                candidate_keys=payload.get("candidate_keys"),
            )
        else:
            verified = WarehouseWritePathService.validate_loose_shipping_marking_scan(
                agency=order.agency,
                order_id=order.number,
                barcode=str(payload.get("barcode") or "").strip(),
                marking_code=str(payload.get("marking_code") or "").strip(),
            )
    except ValueError as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=400)
    return JsonResponse({"ok": True, **verified})


@login_required
def shipping_packing_slips(request, pk: int):
    return handle_shipping_packing_slips_request(request=request, pk=pk)


@login_required
@require_http_methods(["GET", "POST"])
def shipping_packing_slips_status(request, pk: int):
    return handle_shipping_packing_slips_status_request(request=request, pk=pk)
