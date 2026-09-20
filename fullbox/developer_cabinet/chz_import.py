from __future__ import annotations

import json
import logging
import re
import uuid
from collections import Counter, defaultdict
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.conf import settings
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook

from fullbox.file_locks import acquire_file_lock_nonblocking, release_file_lock

from marking.codes import (
    MarkingCodeFormatError,
    marking_code_identity,
    marking_code_variants,
    normalize_marking_code,
    validate_import_marking_code,
)
from marking.models import MarkingCode
from sku.models import Agency, SKU, SKUBarcode, abbreviate_agency_name
from sklad.models import WarehouseContainer, WarehouseLocation, WarehouseStockSnapshot
from sklad.services.warehouse_write_path import WarehouseWritePathService


logger = logging.getLogger(__name__)


EXPECTED_HEADERS = [
    "Паллета",
    "Короб",
    "Клиент",
    "SKU",
    "Наименование",
    "Ширина, мм",
    "Высота, мм",
    "Глубина, мм",
    "Вес, г",
    "Тип товара",
    "Место",
    "Всего",
    "ШК",
    "ЧЗ",
]
AUTO_PRODUCT_HEADERS = [
    header
    for header in EXPECTED_HEADERS
    if header not in {"SKU", "Наименование"}
]
SUPPORTED_HEADERS = (EXPECTED_HEADERS, AUTO_PRODUCT_HEADERS)

GOODS_TYPE_MAP = {
    "gv": "gv",
    "готовый": "gv",
    "готовый товар": "gv",
    "br": "br",
    "брак": "br",
    "no": "no",
    "норма": "no",
    "обычный": "no",
    "rh": "rh",
    "расходник": "rh",
    "vz": "vz",
    "возврат": "vz",
    "votg": "votg",
    "возврат с отгрузки": "votg",
    "op": "op",
    "образец": "op",
}

KNOWN_ZONES = {"PR", "OS", "MR", "OBR", "OTG", "LOAD", "VEH"}
IMPORT_MODE_CHZ = "chz"
IMPORT_MODE_NO_CHZ = "no_chz"
IMPORT_MODE_LABELS = {
    IMPORT_MODE_CHZ: "Импорт товара с ЧЗ",
    IMPORT_MODE_NO_CHZ: "Загрузка остатков без ЧЗ",
}


def normalize_import_mode(value: Any) -> str:
    return IMPORT_MODE_NO_CHZ if str(value or "").strip() == IMPORT_MODE_NO_CHZ else IMPORT_MODE_CHZ


def validate_chz_import_file(uploaded_file, user=None, import_mode: str = IMPORT_MODE_CHZ) -> dict[str, Any]:
    import_mode = normalize_import_mode(import_mode)
    report: dict[str, Any] = {
        "ok": False,
        "token": "",
        "batch_id": "",
        "import_mode": import_mode,
        "import_mode_label": IMPORT_MODE_LABELS[import_mode],
        "summary": {},
        "errors": [],
        "warnings": [],
        "preview_rows": [],
    }

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        _add_error(report, None, "file", f"Не удалось открыть Excel-файл: {exc}")
        return _finalize_report(report)

    if "Остатки" not in workbook.sheetnames:
        _add_error(report, None, "sheet", "В файле нет листа «Остатки».")
        return _finalize_report(report)

    sheet = workbook["Остатки"]
    header_match = _find_header_row(sheet)
    if not header_match:
        _add_error(
            report,
            None,
            "headers",
            "Не найдена строка заголовков. Поддерживается исходный шаблон "
            "или шаблон без колонок «SKU» и «Наименование».",
        )
        return _finalize_report(report)

    header_row, headers = header_match
    raw_rows = _read_rows(sheet, header_row, headers)
    if not raw_rows:
        _add_error(report, None, "rows", "В файле нет строк для импорта.")
        return _finalize_report(report)

    agencies = _agency_lookup()
    parsed_rows: list[dict[str, Any]] = []
    all_chz_codes: list[str] = []
    requested_barcodes: set[str] = set()

    for raw in raw_rows:
        parsed = _parse_raw_row(raw, report, import_mode=import_mode)
        if not parsed:
            continue

        agency = agencies.get(_key(parsed["client_name"]))
        if not agency:
            _add_error(report, parsed["row_number"], "Клиент", f"Клиент не найден: {parsed['client_name']}")
            continue

        parsed["agency_id"] = agency.id
        parsed["agency_name"] = _agency_display_name(agency)
        requested_barcodes.add(_key(parsed["barcode"]))
        parsed_rows.append(parsed)
        all_chz_codes.extend(parsed["chz_codes"])

    sku_matches_by_barcode = _sku_lookup_by_barcode(requested_barcodes)
    locations_by_text: dict[str, dict[str, Any] | None] = {}

    if all_chz_codes:
        identity_counts = Counter(marking_code_identity(code) for code in all_chz_codes)
        duplicate_identities = {
            identity for identity, count in identity_counts.items() if count > 1
        }
        if duplicate_identities:
            for identity in sorted(duplicate_identities):
                _add_error(
                    report,
                    None,
                    "ЧЗ",
                    f"Код ЧЗ повторяется внутри файла: {identity}",
                )

        identities = set(identity_counts)
        variants = set()
        for code in all_chz_codes:
            variants.update(marking_code_variants(code))
        existing_marking_rows = MarkingCode.objects.filter(
            Q(identity_key__in=identities) | Q(code__in=variants)
        ).values_list("identity_key", "code")
        existing_identities = {
            identity or marking_code_identity(code)
            for identity, code in existing_marking_rows
        }
        existing_stock = set(
            WarehouseStockSnapshot.objects.filter(
                is_archived=False,
                marking_code__in=variants,
            ).values_list(
                "marking_code", flat=True
            )
        )
        existing_identities.update(marking_code_identity(code) for code in existing_stock)
        for identity in sorted(existing_identities & identities):
            _add_error(report, None, "ЧЗ", f"Код ЧЗ уже есть в системе: {identity}")

    box_to_pallets: dict[str, set[str]] = defaultdict(set)
    box_to_locations: dict[str, set[str]] = defaultdict(set)
    box_to_dimensions: dict[str, set[tuple[str, str, str, str]]] = defaultdict(set)
    pallet_to_locations: dict[str, set[str]] = defaultdict(set)

    import_items: list[dict[str, Any]] = []
    box_specs: dict[tuple[int, str], dict[str, Any]] = {}

    for row in parsed_rows:
        sku = _resolve_row_sku(row, sku_matches_by_barcode, report)
        if not sku:
            continue

        provided_sku_code = row["sku_code"]
        row["sku_id"] = sku.id
        row["sku_code"] = _clean_text(getattr(sku, "sku_code", ""))
        row["sku_name"] = getattr(sku, "name", "") or row["name"]
        row["size"] = _clean_text(getattr(sku, "size", "")) or row["size"]

        if provided_sku_code and _key(provided_sku_code) != _key(row["sku_code"]):
            _add_warning(
                report,
                row["row_number"],
                "SKU",
                f"В файле «{provided_sku_code}», по ШК найден SKU «{row['sku_code']}».",
            )

        if row["name"] and _key(row["name"]) != _key(row["sku_name"]):
            _add_warning(
                report,
                row["row_number"],
                "Наименование",
                f"В файле «{row['name']}», в карточке SKU «{row['sku_name']}».",
            )

        sku_has_honest_sign = bool(getattr(sku, "honest_sign", False))
        if import_mode == IMPORT_MODE_CHZ and not sku_has_honest_sign:
            _add_warning(report, row["row_number"], "SKU", "У SKU не включен признак «Честный знак».")
        if import_mode == IMPORT_MODE_NO_CHZ and sku_has_honest_sign:
            _add_warning(
                report,
                row["row_number"],
                "SKU",
                "У SKU включен признак «Честный знак», товар будет принят без кодов ЧЗ.",
            )

        location_key = _key(row["location_text"])
        if location_key not in locations_by_text:
            locations_by_text[location_key] = _resolve_location(row["location_text"])
        location = locations_by_text[location_key]
        if not location:
            _add_error(
                report,
                row["row_number"],
                "Место",
                f"Не удалось распознать место «{row['location_text']}». "
                "Укажите существующий код места или числовой адрес в формате "
                "OS·ряд·секция·ярус·ячейка.",
            )
            continue
        location_error = _location_coordinates_error(row["location_text"], location)
        if location_error:
            _add_error(report, row["row_number"], "Место", location_error)
            continue

        row["location"] = location
        box_to_pallets[row["box_code"]].add(row["pallet_code"] or "")
        box_to_locations[row["box_code"]].add(location["location_key"])
        box_to_dimensions[row["box_code"]].add(
            (row["width_mm"], row["height_mm"], row["depth_mm"], row["weight_g"])
        )
        if row["pallet_code"]:
            pallet_to_locations[row["pallet_code"]].add(location["location_key"])
        else:
            _add_warning(report, row["row_number"], "Паллета", "Паллета не указана, короб будет размещен без паллеты.")

        box_specs[(row["agency_id"], row["box_code"])] = {
            "agency_id": row["agency_id"],
            "code": row["box_code"],
            "width_mm": row["width_mm"],
            "height_mm": row["height_mm"],
            "depth_mm": row["depth_mm"],
            "gross_weight_g": row["weight_g"],
        }

        if import_mode == IMPORT_MODE_NO_CHZ:
            import_items.append(
                {
                    "row_number": row["row_number"],
                    "agency_id": row["agency_id"],
                    "agency_name": row["agency_name"],
                    "sku_id": row["sku_id"],
                    "sku_code": row["sku_code"],
                    "name": row["sku_name"],
                    "size": row["size"],
                    "barcode": row["barcode"],
                    "goods_type": row["goods_type"],
                    "qty": row["qty"],
                    "marking_code": "",
                    "box_code": row["box_code"],
                    "pallet_code": row["pallet_code"],
                    "location": location,
                }
            )
        else:
            for code in row["chz_codes"]:
                import_items.append(
                    {
                        "row_number": row["row_number"],
                        "agency_id": row["agency_id"],
                        "agency_name": row["agency_name"],
                        "sku_id": row["sku_id"],
                        "sku_code": row["sku_code"],
                        "name": row["sku_name"],
                        "size": row["size"],
                        "barcode": row["barcode"],
                        "goods_type": row["goods_type"],
                        "qty": 1,
                        "marking_code": code,
                        "box_code": row["box_code"],
                        "pallet_code": row["pallet_code"],
                        "location": location,
                    }
                )

    for box_code, pallets in sorted(box_to_pallets.items()):
        if len(pallets) > 1:
            _add_error(report, None, "Короб", f"Короб {box_code} указан на разных паллетах.")
    for box_code, locations in sorted(box_to_locations.items()):
        if len(locations) > 1:
            _add_error(report, None, "Короб", f"Короб {box_code} указан на разных местах хранения.")
    for box_code, dimensions in sorted(box_to_dimensions.items()):
        if len(dimensions) > 1:
            _add_error(report, None, "Короб", f"У короба {box_code} разные габариты/вес в разных строках.")
    for pallet_code, locations in sorted(pallet_to_locations.items()):
        if len(locations) > 1:
            _add_error(report, None, "Паллета", f"Паллета {pallet_code} указана на разных местах хранения.")

    allowed_existing_containers = _validate_existing_containers(report, import_items, import_mode=import_mode)
    if import_mode == IMPORT_MODE_NO_CHZ and allowed_existing_containers:
        box_specs = {
            key: spec
            for key, spec in box_specs.items()
            if str(spec.get("code") or "").strip() not in allowed_existing_containers
        }

    batch_prefix = "VMS-NOCHZ" if import_mode == IMPORT_MODE_NO_CHZ else "VMS"
    batch_id = timezone.localtime().strftime(f"{batch_prefix}-%Y%m%d-%H%M%S")
    report["batch_id"] = batch_id
    report["summary"] = _build_summary(raw_rows, import_items, parsed_rows, import_mode=import_mode)
    report["preview_rows"] = _build_preview_rows(parsed_rows)

    report = _finalize_report(report)
    if report["ok"]:
        token = uuid.uuid4().hex
        plan = {
            "token": token,
            "status": "ready",
            "batch_id": batch_id,
            "created_at": timezone.now().isoformat(),
            "created_by": getattr(user, "username", "") if user else "",
            "import_mode": import_mode,
            "import_mode_label": IMPORT_MODE_LABELS[import_mode],
            "summary": report["summary"],
            "items": import_items,
            "box_specs": list(box_specs.values()),
        }
        _save_plan(token, plan)
        report["token"] = token

    return report


def execute_chz_import(token: str, user=None) -> dict[str, Any]:
    result: dict[str, Any] = {
        "ok": False,
        "batch_id": "",
        "summary": {},
        "errors": [],
    }

    try:
        _plan_path(token)
    except Exception as exc:
        _add_error(result, None, "token", f"Не удалось открыть проверенный импорт: {exc}")
        return _finalize_report(result)

    with _plan_execution_lock(token) as acquired:
        if not acquired:
            _add_error(result, None, "token", "Этот импорт уже выполняется. Дождитесь завершения операции.")
            return _finalize_report(result)
        return _execute_chz_import_locked(token, user=user, result=result)


def _execute_chz_import_locked(token: str, *, user=None, result: dict[str, Any]) -> dict[str, Any]:
    try:
        plan = _load_plan(token)
    except Exception as exc:
        _add_error(result, None, "token", f"Не удалось открыть проверенный импорт: {exc}")
        return _finalize_report(result)

    if plan.get("status") != "ready":
        if plan.get("status") == "executing":
            message = "Этот импорт уже выполняется или был прерван после начала. Повторная запись заблокирована."
        else:
            message = "Этот импорт уже был выполнен или недоступен."
        _add_error(result, None, "token", message)
        return _finalize_report(result)

    items = plan.get("items") or []
    if not items:
        _add_error(result, None, "items", "В плане импорта нет строк.")
        return _finalize_report(result)

    execution_errors = _execution_item_errors(items)
    if execution_errors:
        result["errors"].extend(execution_errors)
        return _finalize_report(result)

    import_mode = normalize_import_mode(plan.get("import_mode"))
    batch_id = plan.get("batch_id") or timezone.localtime().strftime("VMS-%Y%m%d-%H%M%S")
    plan["status"] = "executing"
    plan["execution_started_at"] = timezone.now().isoformat()
    plan["execution_started_by"] = getattr(user, "username", "") if user else ""
    _save_plan(token, plan)

    try:
        service_summary = _execute_chz_import_plan(
            plan=plan,
            items=items,
            import_mode=import_mode,
            batch_id=batch_id,
            user=user,
            result=result,
        )
    except IntegrityError:
        plan["status"] = "ready"
        plan["last_execution_error"] = "concurrent_marking_duplicate"
        _save_plan(token, plan)
        _add_error(
            result,
            None,
            "ЧЗ",
            (
                "Импорт отменён: один из кодов ЧЗ был добавлен параллельно. "
                "Выполните проверку файла заново."
            ),
        )
        return _finalize_report(result)
    except Exception as exc:
        plan["status"] = "ready"
        plan["last_execution_error"] = " ".join(str(exc).split())[:1000]
        try:
            _save_plan(token, plan)
        except Exception:
            logger.exception("Could not restore CHZ import plan status token=%s", token)
        raise

    if service_summary is None:
        plan["status"] = "ready"
        _save_plan(token, plan)
        return _finalize_report(result)

    plan["status"] = "imported"
    plan["imported_at"] = timezone.now().isoformat()
    plan["imported_by"] = getattr(user, "username", "") if user else ""
    _save_plan(token, plan)

    result["ok"] = True
    result["batch_id"] = batch_id
    result["summary"] = {
        **(plan.get("summary") or {}),
        "events_created": service_summary["events"],
        "snapshots_created": service_summary["snapshots"],
    }
    return result


def _execute_chz_import_plan(*, plan, items, import_mode, batch_id, user, result):
    service_summary = {"events": 0, "snapshots": 0}
    with transaction.atomic():
        _lock_import_containers(items)
        race_errors = _race_check_before_import(items, import_mode=import_mode)
        if race_errors:
            result["errors"].extend(race_errors)
            return None

        if import_mode == IMPORT_MODE_CHZ:
            marking_objects = [
                MarkingCode(
                    order_type="other",
                    order_id=batch_id,
                    agency_id=item["agency_id"],
                    sku_id=item["sku_id"],
                    sku_code=item["sku_code"],
                    size=item.get("size", ""),
                    barcode=item.get("barcode", ""),
                    box_barcode=item["box_code"],
                    code=item["marking_code"],
                    source="import",
                    used_at=timezone.now(),
                )
                for item in items
                if str(item.get("marking_code") or "").strip()
            ]
            if marking_objects:
                MarkingCode.objects.bulk_create(marking_objects, batch_size=1000)

        items_by_agency: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for item in items:
            items_by_agency[item["agency_id"]].append(_warehouse_item(item))

        agency_cache = Agency.objects.in_bulk(items_by_agency.keys())
        for agency_id, agency_items in items_by_agency.items():
            agency = agency_cache[agency_id]
            response = WarehouseWritePathService.create_receiving_placement(
                agency=agency,
                order_id=batch_id,
                items=agency_items,
                performed_by=user if getattr(user, "is_authenticated", False) else None,
                warehouse_code="MSK",
                source_document_type="inventory_import_no_chz" if import_mode == IMPORT_MODE_NO_CHZ else "chz_inventory_import",
                source_document_id=batch_id,
                stock_context_type="inventory_import",
                respect_item_location=True,
            )
            service_summary["events"] += _count_result_part(response, "events")
            service_summary["snapshots"] += _count_result_part(response, "snapshots")

            agency_box_specs = [
                box
                for box in (plan.get("box_specs") or [])
                if int(box.get("agency_id") or 0) == agency_id
            ]
            if agency_box_specs:
                WarehouseWritePathService.sync_box_characteristics(
                    agency=agency,
                    order_id=batch_id,
                    order_type="other",
                    boxes=agency_box_specs,
                    performed_by=user if getattr(user, "is_authenticated", False) else None,
                )

    return service_summary


def _find_header_row(sheet) -> tuple[int, list[str]] | None:
    max_headers = max(len(headers) for headers in SUPPORTED_HEADERS)
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        values = [_clean_text(value) for value in row[:max_headers]]
        while values and not values[-1]:
            values.pop()
        for headers in SUPPORTED_HEADERS:
            if values == headers:
                return row_number, list(headers)
    return None


def _read_rows(sheet, header_row: int, headers: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        values = list(row[: len(headers)])
        if not any(_clean_text(value) for value in values):
            continue
        rows.append({"row_number": row_number, "values": dict(zip(headers, values))})
    return rows


def _parse_raw_row(raw: dict[str, Any], report: dict[str, Any], *, import_mode: str = IMPORT_MODE_CHZ) -> dict[str, Any] | None:
    import_mode = normalize_import_mode(import_mode)
    row_number = raw["row_number"]
    values = raw["values"]

    client_name = _clean_text(values.get("Клиент"))
    sku_code = _clean_text(values.get("SKU"))
    name = _clean_text(values.get("Наименование"))
    box_code = _clean_text(values.get("Короб"))
    pallet_code = _clean_text(values.get("Паллета"))
    barcode = _clean_text(values.get("ШК"))
    location_text = _clean_text(values.get("Место"))
    goods_type_raw = _clean_text(values.get("Тип товара"))
    goods_type = GOODS_TYPE_MAP.get(_key(goods_type_raw))
    quantity = _parse_positive_int(values.get("Всего"))
    chz_codes = _split_chz_codes(values.get("ЧЗ"))
    if import_mode == IMPORT_MODE_CHZ:
        validated_codes = []
        for code in chz_codes:
            try:
                validated_codes.append(
                    validate_import_marking_code(code, product_barcode=barcode)
                )
            except MarkingCodeFormatError as exc:
                _add_error(
                    report,
                    row_number,
                    "ЧЗ",
                    f"Некорректный код «{code}»: {exc}.",
                )
        chz_codes = validated_codes

    if not client_name:
        _add_error(report, row_number, "Клиент", "Клиент обязателен.")
    if not box_code:
        _add_error(report, row_number, "Короб", "Короб обязателен.")
    if not barcode:
        _add_error(report, row_number, "ШК", "ШК обязателен.")
    if not location_text:
        _add_error(report, row_number, "Место", "Место хранения обязательно.")
    if not goods_type:
        _add_error(report, row_number, "Тип товара", f"Неизвестный тип товара: {goods_type_raw}")
    if quantity is None:
        _add_error(report, row_number, "Всего", "Количество должно быть целым числом больше нуля.")
    elif import_mode == IMPORT_MODE_CHZ and len(chz_codes) != quantity:
        _add_error(report, row_number, "ЧЗ", f"Кодов ЧЗ {len(chz_codes)}, а в поле «Всего» указано {quantity}.")
    if import_mode == IMPORT_MODE_CHZ and not chz_codes:
        _add_error(report, row_number, "ЧЗ", "Коды ЧЗ обязательны.")
    if import_mode == IMPORT_MODE_NO_CHZ and chz_codes:
        _add_error(report, row_number, "ЧЗ", "В режиме «Остатки без ЧЗ» колонка ЧЗ должна быть пустой.")

    required_values = [client_name, box_code, barcode, location_text, goods_type, quantity]
    if import_mode == IMPORT_MODE_CHZ:
        required_values.append(chz_codes)
    if not all(required_values):
        return None

    return {
        "row_number": row_number,
        "client_name": client_name,
        "sku_code": sku_code,
        "name": name,
        "box_code": box_code,
        "pallet_code": pallet_code,
        "barcode": barcode,
        "location_text": location_text,
        "goods_type": goods_type,
        "qty": quantity,
        "chz_codes": chz_codes,
        "width_mm": _clean_number(values.get("Ширина, мм")),
        "height_mm": _clean_number(values.get("Высота, мм")),
        "depth_mm": _clean_number(values.get("Глубина, мм")),
        "weight_g": _clean_number(values.get("Вес, г")),
        "size": "",
    }


def _agency_lookup() -> dict[str, Agency]:
    lookup: dict[str, Agency] = {}
    for agency in Agency.objects.all():
        names = {
            _agency_display_name(agency),
            _clean_text(getattr(agency, "short_name", "")),
            _clean_text(getattr(agency, "name", "")),
        }
        for name in list(names):
            if name:
                try:
                    names.add(abbreviate_agency_name(name))
                except Exception:
                    pass
        for name in names:
            if name:
                lookup.setdefault(_key(name), agency)
    return lookup


def _agency_display_name(agency: Agency) -> str:
    return _clean_text(getattr(agency, "agn_name", "")) or _clean_text(getattr(agency, "name", "")) or str(agency)


def _sku_lookup_by_barcode(requested_barcodes: set[str]) -> dict[str, list[SKU]]:
    requested_keys = {_key(value) for value in requested_barcodes if _key(value)}
    if not requested_keys:
        return {}

    sku_ids_by_barcode: dict[str, set[int]] = defaultdict(set)
    sku_codes = SKU.objects.exclude(code__isnull=True).exclude(code="").values_list("id", "code")
    for sku_id, value in sku_codes.iterator():
        barcode_key = _key(value)
        if barcode_key in requested_keys:
            sku_ids_by_barcode[barcode_key].add(sku_id)

    barcode_rows = SKUBarcode.objects.values_list("sku_id", "value")
    for sku_id, value in barcode_rows.iterator():
        barcode_key = _key(value)
        if barcode_key in requested_keys:
            sku_ids_by_barcode[barcode_key].add(sku_id)

    sku_ids = {sku_id for values in sku_ids_by_barcode.values() for sku_id in values}
    skus_by_id = SKU.objects.in_bulk(sku_ids)
    return {
        barcode_key: [
            skus_by_id[sku_id]
            for sku_id in sorted(sku_ids)
            if sku_id in skus_by_id
        ]
        for barcode_key, sku_ids in sku_ids_by_barcode.items()
    }


def _resolve_row_sku(
    row: dict[str, Any],
    sku_matches_by_barcode: dict[str, list[SKU]],
    report: dict[str, Any],
) -> SKU | None:
    barcode = row["barcode"]
    matches = sku_matches_by_barcode.get(_key(barcode), [])
    client_matches = [
        sku
        for sku in matches
        if int(getattr(sku, "agency_id", 0) or 0) == int(row["agency_id"])
    ]

    if not client_matches:
        if matches:
            message = f"ШК {barcode} найден у другого клиента."
        else:
            message = f"ШК {barcode} не найден в справочнике SKU."
        _add_error(report, row["row_number"], "ШК", message)
        return None

    if len(client_matches) > 1:
        sku_codes = ", ".join(
            sorted({_clean_text(getattr(sku, "sku_code", "")) for sku in client_matches})
        )
        _add_error(
            report,
            row["row_number"],
            "ШК",
            f"ШК {barcode} неоднозначен для клиента, найдены SKU: {sku_codes}.",
        )
        return None

    sku = client_matches[0]
    if bool(getattr(sku, "deleted", False)):
        _add_error(
            report,
            row["row_number"],
            "ШК",
            f"ШК {barcode} связан с удаленным SKU {getattr(sku, 'sku_code', '')}.",
        )
        return None
    return sku


def _resolve_location(raw_text: str) -> dict[str, Any] | None:
    text = _clean_text(raw_text)
    if not text:
        return None

    try:
        location = WarehouseLocation.objects.filter(
            Q(location_code__iexact=text) | Q(display_name__iexact=text)
        ).first()
    except Exception:
        location = None
    if location:
        return _location_payload(location, text)

    parsed = _parse_location_text(text)
    if not parsed:
        return None

    try:
        location = WarehouseLocation.objects.filter(
            zone_code__iexact=parsed["zone_code"],
            row__iexact=parsed["row"],
            section__iexact=parsed["section"],
            tier__iexact=parsed["tier"],
            cell__iexact=parsed["cell"],
        ).first()
    except Exception:
        location = None
    if location:
        return _location_payload(location, text)
    return parsed


def _location_payload(location: WarehouseLocation, raw_text: str) -> dict[str, Any]:
    zone_code = _clean_text(getattr(location, "zone_code", "")) or _clean_text(getattr(location, "zone_kind", ""))
    row = _clean_text(getattr(location, "row", ""))
    section = _clean_text(getattr(location, "section", ""))
    tier = _clean_text(getattr(location, "tier", ""))
    cell = _clean_text(getattr(location, "cell", ""))
    location_code = _clean_text(getattr(location, "location_code", ""))
    display_name = _clean_text(getattr(location, "display_name", "")) or raw_text
    return {
        "id": location.id,
        "zone_code": zone_code,
        "zone_kind": zone_code,
        "row": row,
        "section": section,
        "tier": tier,
        "cell": cell,
        "location_code": location_code,
        "display_name": display_name,
        "raw": raw_text,
        "location_key": location_code or f"{zone_code}:{row}:{section}:{tier}:{cell}",
    }


def _parse_location_text(text: str) -> dict[str, Any] | None:
    normalized = (
        text.upper()
        .replace("·", " ")
        .replace("•", " ")
        .replace(",", " ")
        .replace(";", " ")
        .replace("-", " ")
    )
    zone_match = re.search(r"\b(PR|OS|MR|OBR|OTG|LOAD|VEH)\b", normalized)
    if not zone_match:
        return None
    zone_code = zone_match.group(1)

    row = _extract_location_part(normalized, ["РЯД", "ROW"])
    section = _extract_location_part(normalized, ["СЕКЦИЯ", "SECTION", "СЕК"])
    tier = _extract_location_part(normalized, ["ЯРУС", "TIER"])
    cell = _extract_location_part(normalized, ["ЯЧЕЙКА", "CELL", "МЕСТО"])

    if not all([row, section, tier]):
        tokens = [token for token in normalized.split() if token and token not in KNOWN_ZONES]
        if len(tokens) >= 3:
            row = row or tokens[0]
            section = section or tokens[1]
            tier = tier or tokens[2]
            cell = cell or (tokens[3] if len(tokens) > 3 else "1")

    if not all([row, section, tier]):
        return None
    cell = cell or "1"
    location_key = f"{zone_code}:{row}:{section}:{tier}:{cell}"
    return {
        "zone_code": zone_code,
        "zone_kind": zone_code,
        "row": row,
        "section": section,
        "tier": tier,
        "cell": cell,
        "location_code": "",
        "display_name": text,
        "raw": text,
        "location_key": location_key,
    }


def _extract_location_part(text: str, labels: list[str]) -> str:
    labels_re = "|".join(re.escape(label) for label in labels)
    match = re.search(rf"(?:{labels_re})\s+([A-ZА-ЯЁ0-9]+)", text)
    return match.group(1) if match else ""


def _container_codes_from_items(items: list[dict[str, Any]]) -> set[str]:
    codes = {str(item.get("box_code") or "").strip() for item in items if item.get("box_code")}
    codes.update(str(item.get("pallet_code") or "").strip() for item in items if item.get("pallet_code"))
    return {code for code in codes if code}


def _existing_containers(items: list[dict[str, Any]]) -> set[str]:
    codes = _container_codes_from_items(items)
    if not codes:
        return set()
    existing = set(WarehouseContainer.objects.filter(container_code__in=codes).values_list("container_code", flat=True))
    existing.update(
        WarehouseStockSnapshot.objects.filter(is_archived=False, container_code__in=codes).values_list(
            "container_code", flat=True
        )
    )
    return existing


def _validate_existing_containers(
    report: dict[str, Any],
    items: list[dict[str, Any]],
    *,
    import_mode: str = IMPORT_MODE_CHZ,
) -> set[str]:
    errors, warnings, allowed_codes = _existing_container_messages(items, import_mode=import_mode)
    for item in errors:
        _add_error(report, item.get("row") or None, item["field"], item["message"])
    for item in warnings:
        _add_warning(report, item.get("row") or None, item["field"], item["message"])
    return allowed_codes


def _existing_container_messages(
    items: list[dict[str, Any]],
    *,
    import_mode: str = IMPORT_MODE_CHZ,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], set[str]]:
    codes = _container_codes_from_items(items)
    if not codes:
        return [], [], set()

    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    allowed_codes: set[str] = set()
    usage: dict[str, dict[str, Any]] = {}
    for item in items:
        agency_id = int(item.get("agency_id") or 0)
        location = item.get("location") or {}
        box_code = str(item.get("box_code") or "").strip()
        pallet_code = str(item.get("pallet_code") or "").strip()
        if box_code:
            entry = usage.setdefault(box_code, {"box": True, "pallet": False, "agency_ids": set(), "locations": set(), "pallets": set()})
            entry["box"] = True
            entry["agency_ids"].add(agency_id)
            entry["locations"].add(_payload_location_key(location))
            if pallet_code:
                entry["pallets"].add(pallet_code)
        if pallet_code:
            entry = usage.setdefault(pallet_code, {"box": False, "pallet": True, "agency_ids": set(), "locations": set(), "pallets": set()})
            entry["pallet"] = True
            entry["agency_ids"].add(agency_id)
            entry["locations"].add(_payload_location_key(location))

    containers_by_code: dict[str, list[WarehouseContainer]] = defaultdict(list)
    for container in WarehouseContainer.objects.filter(container_code__in=codes).select_related(
        "agency",
        "current_location",
        "parent_container",
    ):
        containers_by_code[str(container.container_code or "").strip()].append(container)

    snapshots_by_code: dict[str, list[WarehouseStockSnapshot]] = defaultdict(list)
    snapshot_query = (
        Q(container_code__in=codes)
        | Q(container__container_code__in=codes)
        | Q(parent_container__container_code__in=codes)
    )
    for snapshot in WarehouseStockSnapshot.objects.filter(snapshot_query, is_archived=False).select_related(
        "location",
        "container",
        "parent_container",
        "agency",
    ):
        snapshot_codes = {
            str(snapshot.container_code or "").strip(),
            str(getattr(snapshot.container, "container_code", "") or "").strip(),
            str(getattr(snapshot.parent_container, "container_code", "") or "").strip(),
        }
        for code in snapshot_codes:
            if code in codes:
                snapshots_by_code[code].append(snapshot)

    for code, entry in sorted(usage.items()):
        entry["locations"].discard("")
        agency_ids = {int(value) for value in entry["agency_ids"] if int(value or 0)}
        if len(agency_ids) != 1:
            errors.append(_message(None, "Короб/паллета", f"Контейнер {code} указан для разных клиентов в файле."))
            continue
        agency_id = next(iter(agency_ids))
        expected_location = next(iter(entry["locations"])) if len(entry["locations"]) == 1 else ""
        if entry["box"] and entry["pallet"]:
            errors.append(_message(None, "Короб/паллета", f"Код {code} указан и как короб, и как паллета."))
            continue

        containers = containers_by_code.get(code) or []
        snapshots = snapshots_by_code.get(code) or []
        if not containers and snapshots:
            errors.append(
                _message(
                    None,
                    "Короб/паллета",
                    f"Контейнер {code} есть в остатках, но карточка контейнера не найдена. Нужна ручная проверка.",
                )
            )
            continue
        if not containers:
            continue

        if import_mode == IMPORT_MODE_CHZ and entry["box"]:
            errors.append(
                _message(
                    None,
                    "Короб",
                    f"Короб {code} уже есть в складском учете. "
                    "В режиме ЧЗ можно использовать повторно только существующую паллету.",
                )
            )
            continue

        code_allowed = True
        for container in containers:
            if int(container.agency_id or 0) != agency_id:
                errors.append(_message(None, "Короб/паллета", f"Контейнер {code} уже принадлежит другому клиенту."))
                code_allowed = False
                continue
            if str(container.status or "") != WarehouseContainer.STATUS_ACTIVE:
                errors.append(_message(None, "Короб/паллета", f"Контейнер {code} не активен: {container.status}."))
                code_allowed = False
            if entry["box"] and str(container.container_type or "") != WarehouseContainer.TYPE_BOX:
                errors.append(_message(None, "Короб/паллета", f"Код {code} уже существует не как короб."))
                code_allowed = False
            if entry["pallet"] and str(container.container_type or "") not in {
                WarehouseContainer.TYPE_PALLET,
                WarehouseContainer.TYPE_MIXED_PALLET,
            }:
                errors.append(_message(None, "Короб/паллета", f"Код {code} уже существует не как паллета."))
                code_allowed = False
            if expected_location and not _location_matches(container.current_location, expected_location):
                errors.append(
                    _message(
                        None,
                        "Короб/паллета",
                        f"Контейнер {code} стоит на другом месте: {_location_label(container.current_location)}.",
                    )
                )
                code_allowed = False
            if entry["box"] and entry["pallets"] and container.parent_container_id:
                current_pallet = str(getattr(container.parent_container, "container_code", "") or "").strip()
                if current_pallet and current_pallet not in entry["pallets"]:
                    errors.append(
                        _message(
                            None,
                            "Короб",
                            f"Короб {code} уже привязан к паллете {current_pallet}, в файле указана другая паллета.",
                        )
                    )
                    code_allowed = False

        for snapshot in snapshots:
            if int(snapshot.agency_id or 0) != agency_id:
                errors.append(_message(None, "Короб/паллета", f"По контейнеру {code} есть остатки другого клиента."))
                code_allowed = False
            if expected_location and not _location_matches(snapshot.location, expected_location):
                errors.append(
                    _message(
                        None,
                        "Короб/паллета",
                        f"По контейнеру {code} есть остатки на другом месте: {_location_label(snapshot.location)}.",
                    )
                )
                code_allowed = False

        if import_mode == IMPORT_MODE_CHZ and entry["pallet"]:
            has_reserve = any(
                sum(
                    int(getattr(snapshot, field, 0) or 0)
                    for field in (
                        "processing_reserved_qty",
                        "shipping_reserved_qty",
                        "other_reserved_qty",
                    )
                )
                > 0
                for snapshot in snapshots
            )
            has_active_operation = any(
                int(getattr(snapshot, "active_operation_id", 0) or 0)
                or str(getattr(snapshot, "active_operation_type", "") or "").strip()
                for snapshot in snapshots
            )
            if has_reserve:
                errors.append(
                    _message(
                        None,
                        "Паллета",
                        f"По паллете {code} есть активный резерв. Добавление новых коробов заблокировано.",
                    )
                )
                code_allowed = False
            if has_active_operation:
                errors.append(
                    _message(
                        None,
                        "Паллета",
                        f"По паллете {code} выполняется складская операция. "
                        "Добавление новых коробов заблокировано.",
                    )
                )
                code_allowed = False

        if code_allowed:
            allowed_codes.add(code)
            warnings.append(
                _message(
                    None,
                    "Короб/паллета",
                    f"Контейнер {code} уже есть у этого клиента, товар будет добавлен в него без смены места.",
                )
            )

    return errors, warnings, allowed_codes


def _payload_location_key(location: dict[str, Any]) -> str:
    if not location:
        return ""
    location_id = str(location.get("id") or "").strip()
    if location_id:
        return f"id:{location_id}"
    location_code = str(location.get("location_code") or "").strip()
    if location_code:
        return f"code:{location_code.casefold()}"
    return str(location.get("location_key") or "").strip().casefold()


def _warehouse_location_keys(location: WarehouseLocation | None) -> set[str]:
    if location is None:
        return set()
    keys: set[str] = {f"id:{location.id}"}
    location_code = str(getattr(location, "location_code", "") or "").strip()
    if location_code:
        keys.add(f"code:{location_code.casefold()}")
    row = str(getattr(location, "row", getattr(location, "row_no", "")) or "").strip()
    section = str(getattr(location, "section", getattr(location, "section_no", "")) or "").strip()
    tier = str(getattr(location, "tier", getattr(location, "tier_no", "")) or "").strip()
    cell = str(getattr(location, "cell", getattr(location, "cell_no", "")) or "").strip()
    zone = str(getattr(location, "zone_code", "") or "").strip()
    if zone and row and section and tier and cell:
        keys.add(f"{zone}:{row}:{section}:{tier}:{cell}".casefold())
    display_name = str(getattr(location, "display_name", "") or "").strip()
    if display_name:
        keys.add(display_name.casefold())
    return keys


def _location_matches(location: WarehouseLocation | None, expected_key: str) -> bool:
    if not expected_key:
        return True
    return expected_key in _warehouse_location_keys(location)


def _location_label(location: WarehouseLocation | None) -> str:
    if location is None:
        return "место не указано"
    return (
        str(getattr(location, "display_name", "") or "").strip()
        or str(getattr(location, "location_code", "") or "").strip()
        or f"{getattr(location, 'zone_code', '')} {getattr(location, 'row_no', '')}-{getattr(location, 'section_no', '')}-{getattr(location, 'tier_no', '')}-{getattr(location, 'cell_no', '')}"
    )


def _location_coordinates_error(raw_text: str, location: dict[str, Any] | None) -> str:
    if not isinstance(location, dict):
        return ""

    invalid_coordinates = []
    for key, fallback_key, label in (
        ("row", "row_no", "ряд"),
        ("section", "section_no", "секция"),
        ("tier", "tier_no", "ярус"),
        ("cell", "cell_no", "ячейка"),
    ):
        value = _clean_text(location.get(key) or location.get(fallback_key))
        if value and not re.fullmatch(r"\d+", value):
            invalid_coordinates.append(f"{label} «{value}»")

    if not invalid_coordinates:
        return ""

    location_text = _clean_text(raw_text) or _clean_text(location.get("display_name"))
    return (
        f"Место «{location_text}» содержит нечисловые координаты: "
        f"{', '.join(invalid_coordinates)}. Укажите существующий код места "
        "или числовой адрес в формате OS·ряд·секция·ярус·ячейка."
    )


def _execution_item_errors(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    errors = []
    for item in items:
        location = item.get("location") if isinstance(item.get("location"), dict) else {}
        message = _location_coordinates_error(
            str(location.get("raw") or location.get("display_name") or ""),
            location,
        )
        if message:
            errors.append(_message(item.get("row_number") or None, "Место", message))
        raw_code = item.get("marking_code")
        if raw_code:
            try:
                normalized = validate_import_marking_code(
                    raw_code,
                    product_barcode=item.get("barcode"),
                )
            except MarkingCodeFormatError as exc:
                errors.append(_message(item.get("row_number") or None, "ЧЗ", str(exc)))
            else:
                if normalized != raw_code:
                    errors.append(
                        _message(
                            item.get("row_number") or None,
                            "ЧЗ",
                            (
                                "План импорта содержит неканонический код. "
                                "Выполните проверку файла заново."
                            ),
                        )
                    )
    return errors


def _lock_import_containers(items: list[dict[str, Any]]) -> None:
    codes = _container_codes_from_items(items)
    if not codes:
        return

    list(
        WarehouseContainer.objects.select_for_update()
        .filter(container_code__in=codes)
        .order_by("id")
        .values_list("id", flat=True)
    )
    snapshot_query = (
        Q(container_code__in=codes)
        | Q(container__container_code__in=codes)
        | Q(parent_container__container_code__in=codes)
    )
    snapshot_ids = sorted(
        set(
            WarehouseStockSnapshot.objects.filter(snapshot_query, is_archived=False)
            .order_by("id")
            .values_list("id", flat=True)
        )
    )
    if not snapshot_ids:
        return

    list(
        WarehouseStockSnapshot.objects.select_for_update()
        .filter(id__in=snapshot_ids)
        .order_by("id")
        .values_list("id", flat=True)
    )


def _race_check_before_import(items: list[dict[str, Any]], *, import_mode: str = IMPORT_MODE_CHZ) -> list[dict[str, Any]]:
    errors: list[dict[str, Any]] = []
    chz_codes = [code for code in (str(item.get("marking_code") or "").strip() for item in items) if code]
    if chz_codes:
        identities = {marking_code_identity(code) for code in chz_codes}
        variants = set()
        for code in chz_codes:
            variants.update(marking_code_variants(code))
        existing_marking_rows = MarkingCode.objects.filter(
            Q(identity_key__in=identities) | Q(code__in=variants)
        ).values_list("identity_key", "code")
        existing_identities = {
            identity or marking_code_identity(code)
            for identity, code in existing_marking_rows
        }
        existing_stock = set(
            WarehouseStockSnapshot.objects.filter(
                is_archived=False,
                marking_code__in=variants,
            ).values_list(
                "marking_code", flat=True
            )
        )
        existing_identities.update(marking_code_identity(code) for code in existing_stock)
        for identity in sorted(existing_identities & identities):
            errors.append(
                _message(
                    None,
                    "ЧЗ",
                    f"Код ЧЗ уже появился в системе после проверки: {identity}",
                )
            )
    container_errors, _warnings, _allowed = _existing_container_messages(items, import_mode=import_mode)
    for item in container_errors:
        errors.append(
            _message(
                item.get("row") or None,
                item.get("field") or "Короб/паллета",
                item.get("message") or "Контейнер недоступен для импорта.",
            )
        )
    return errors


def _warehouse_item(item: dict[str, Any]) -> dict[str, Any]:
    location = item["location"]
    return {
        "sku_id": item["sku_id"],
        "sku_code": item["sku_code"],
        "name": item["name"],
        "size": item.get("size", ""),
        "barcode": item.get("barcode", ""),
        "goods_type": item.get("goods_type", "gv"),
        "qty": item.get("qty", 1),
        "marking_code": str(item.get("marking_code") or "").strip(),
        "box_code": item["box_code"],
        "container_code": item["box_code"],
        "pallet_code": item.get("pallet_code", ""),
        "parent_container_code": item.get("pallet_code", ""),
        "location": location,
        "zone_code": location.get("zone_code", ""),
        "zone_kind": location.get("zone_kind", ""),
        "row": location.get("row", ""),
        "section": location.get("section", ""),
        "tier": location.get("tier", ""),
        "cell": location.get("cell", ""),
    }


def _build_summary(
    raw_rows: list[dict[str, Any]],
    items: list[dict[str, Any]],
    parsed_rows: list[dict[str, Any]],
    *,
    import_mode: str = IMPORT_MODE_CHZ,
) -> dict[str, Any]:
    return {
        "import_mode": normalize_import_mode(import_mode),
        "import_mode_label": IMPORT_MODE_LABELS[normalize_import_mode(import_mode)],
        "rows_total": len(raw_rows),
        "sku_rows": len(parsed_rows),
        "marking_count": sum(len(row.get("chz_codes") or []) for row in parsed_rows),
        "qty_total": sum(int(row.get("qty") or 0) for row in parsed_rows),
        "agency_count": len({item["agency_id"] for item in items}),
        "box_count": len({item["box_code"] for item in items}),
        "pallet_count": len({item["pallet_code"] for item in items if item.get("pallet_code")}),
    }


def _build_preview_rows(parsed_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    preview = []
    for row in parsed_rows:
        preview.append(
            {
                "row_number": row["row_number"],
                "client": row.get("agency_name") or row["client_name"],
                "sku": row["sku_code"],
                "name": row.get("sku_name") or row["name"],
                "goods_type": row["goods_type"],
                "qty": row["qty"],
                "chz_count": len(row["chz_codes"]),
                "box": row["box_code"],
                "pallet": row["pallet_code"] or "-",
                "location": row["location_text"],
            }
        )
    return preview


def _split_chz_codes(value: Any) -> list[str]:
    text = _clean_text(value)
    if not text:
        return []
    return [
        normalize_marking_code(part)
        for part in re.split(r"[\r\n]+", text)
        if normalize_marking_code(part)
    ]


def _parse_positive_int(value: Any) -> int | None:
    text = _clean_text(value).replace(" ", "")
    if not text:
        return None
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation:
        return None
    if number <= 0 or number != number.to_integral_value():
        return None
    return int(number)


def _clean_number(value: Any) -> str:
    text = _clean_text(value).replace(" ", "")
    if not text:
        return ""
    try:
        number = Decimal(text.replace(",", "."))
    except InvalidOperation:
        return text
    if number == number.to_integral_value():
        return str(int(number))
    return str(number.normalize())


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _key(value: Any) -> str:
    return re.sub(r"\s+", " ", _clean_text(value)).casefold()


def _message(row: int | None, field: str, message: str) -> dict[str, Any]:
    return {"row": row or "", "field": field, "message": message}


def _add_error(report: dict[str, Any], row: int | None, field: str, message: str) -> None:
    report.setdefault("errors", []).append(_message(row, field, message))


def _add_warning(report: dict[str, Any], row: int | None, field: str, message: str) -> None:
    report.setdefault("warnings", []).append(_message(row, field, message))


def _finalize_report(report: dict[str, Any]) -> dict[str, Any]:
    report["ok"] = not report.get("errors")
    report.setdefault("warnings", [])
    report.setdefault("errors", [])
    report.setdefault("summary", {})
    report["summary"]["errors_count"] = len(report["errors"])
    report["summary"]["warnings_count"] = len(report["warnings"])
    return report


def _plan_dir() -> Path:
    path = Path(settings.MEDIA_ROOT) / "developer_chz_imports"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _plan_path(token: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", token or ""):
        raise ValueError("Некорректный token.")
    return _plan_dir() / f"{token}.json"


def _save_plan(token: str, plan: dict[str, Any]) -> None:
    path = _plan_path(token)
    temp_path = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        temp_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)


def _load_plan(token: str) -> dict[str, Any]:
    return json.loads(_plan_path(token).read_text(encoding="utf-8"))


@contextmanager
def _plan_execution_lock(token: str):
    lock_path = _plan_path(token).with_suffix(".lock")
    handle = lock_path.open("a+")
    acquired = False
    try:
        try:
            acquire_file_lock_nonblocking(handle)
        except BlockingIOError:
            yield False
            return
        acquired = True
        yield True
    finally:
        if acquired:
            release_file_lock(handle)
        handle.close()


def _count_result_part(response: Any, key: str) -> int:
    if not response:
        return 0
    if isinstance(response, dict):
        value = response.get(key)
    else:
        value = getattr(response, key, None)
    try:
        return len(value or [])
    except TypeError:
        return int(value or 0)
