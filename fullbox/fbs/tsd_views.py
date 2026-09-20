from collections import defaultdict
from copy import copy
from datetime import date, timedelta
from decimal import Decimal
from functools import wraps
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlencode, urlsplit
import re
import unicodedata
import uuid

from django.contrib import messages
from django.core import signing
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import (
    Case,
    Count,
    Exists,
    F,
    IntegerField,
    Max,
    Min,
    OuterRef,
    Prefetch,
    Q,
    Sum,
    Value,
    When,
)
from django.db.models.functions import Coalesce
from django.http import (
    FileResponse,
    Http404,
    HttpResponse,
    HttpResponseForbidden,
    JsonResponse,
    QueryDict,
)
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from openpyxl import Workbook

from agent.models import AgentEvent
from employees.access import (
    get_request_employee,
    get_request_role,
    get_request_roles,
    resolve_cabinet_url,
    role_required,
)
from employees.models import Employee
from sku.models import Agency, SKU, SKUBarcode
from processing_app.models import ProcessingPrintJob
from sklad.models import WarehouseLocation, WarehouseOperation
from sklad.services.operational_locations import normalize_operational_location_scan
from sklad.topology import os_location_code, os_location_label

from .exceptions import (
    FbsError,
    FbsIntegrationError,
    FbsLabelError,
    FbsMarkingAlreadyUsedError,
    FbsMovementError,
    FbsScanMismatchError,
)
from .flags import feature_enabled, module_enabled
from .controller_session import get_controller_workstation
from .services.printing import (
    fbs_desktop_print_lease_seconds,
    recover_stale_fbs_order_label_print_job,
    render_ozon_handover_internal_label,
)
from .services.ozon_handover_documents import (
    download_exact_ozon_handover_document,
)
from .integrations.ozon import (
    OZON_HANDOVER_DOCUMENT_BARCODE,
    OZON_HANDOVER_DOCUMENT_PDF,
)
from .services.wave_policy import (
    MAX_WAVE_SIZE,
    MIN_WAVE_SIZE,
    configured_or_default_wave_limits,
)
from .services.wave_queue import (
    current_wave_queue,
    decorate_wave_queue_rows,
    marketplace_queue_priority_rows,
    order_wave_queue_queryset,
    reorder_queued_wave,
    update_marketplace_queue_priorities,
)
from .services.pick_restock import (
    confirm_marketplace_rejected_order_return,
    confirm_ozon_canceled_order_return,
    order_is_client_canceled_by_marketplace,
)
from .services.physical_locations import (
    VIRTUAL_FBS_PLAN_PREFIX,
    fbs_box_physical_location,
    fbs_box_physical_location_code,
    fbs_box_physical_location_label,
)
from .services.picking import pick_route_ordering
from .services.stock_compare import build_agency_stock_comparison
from .workspace import workspace_shell_enabled
from .models import (
    FbsControllerCheckTote,
    FbsControllerPickTote,
    FbsControllerSession,
    FbsControllerToteOrder,
    FbsOrder,
    FbsHandoverBatch,
    FbsHandoverBox,
    FbsHandoverOrder,
    FbsHandoverOrderAssignment,
    FbsIntegrationProfile,
    FbsInventoryScan,
    FbsInventorySession,
    FbsMarketplaceMetadataTransfer,
    FbsMarketplaceCommand,
    FbsOrderLabel,
    FbsOrderItem,
    FbsOrderStockAllocation,
    FbsInternalMovement,
    FbsPickBatch,
    FbsPickException,
    FbsPickRestockRequest,
    FbsPickRestockScan,
    FbsPickScanEvent,
    FbsPickTask,
    FbsPickVerificationProgress,
    FbsPickingCart,
    FbsReplenishmentAllocation,
    FbsReplenishmentLine,
    FbsReplenishmentPlan,
    FbsReplenishmentPreparedBox,
    FbsRackContentMovement,
    FbsRackStagingBox,
    FbsStorageCell,
    FbsPallet,
    FbsBox,
    FbsStockBalance,
    FbsStorekeeperAlertAcknowledgement,
    FbsStorekeeperResponsible,
    FbsToteBinding,
    FbsToteMovement,
    FbsWorkstation,
    FbsWavePolicy,
)
from .services import (
    ACTIVE_PICK_RESTOCK_STATUSES,
    HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES,
    activate_drained_inventory,
    add_handover_box,
    add_wb_handover_boxes,
    add_order_to_handover_box,
    approve_inventory,
    confirm_inventory_discrepancies,
    analyze_replenishment_demand,
    cancel_replenishment_plan,
    claim_pick_batch,
    claim_pick_batch_verification,
    claim_next_pick_batch,
    claim_internal_movement,
    claim_pick_task,
    claim_pick_restock_request,
    claim_replenishment_plan,
    close_handover_box,
    complete_replenishment_allocation,
    complete_pick_allocation,
    confirm_order_label_scan,
    confirm_replenishment_plan,
    create_handover_batch,
    create_inventory_session,
    create_order_pick_restock_request,
    dispatch_handover_batch,
    handover_pick_batch_for_verification,
    finish_inventory_count,
    generate_replenishment_plans,
    move_all_box_contents_to_rack_cell,
    move_scanned_box_item_to_rack_cell,
    pack_staged_item_plan,
    prepare_pick_queue,
    pick_missing_group_quantity,
    configured_fbs_print_workstations,
    queue_fbs_handover_box_label_print,
    queue_fbs_handover_supply_label_print,
    queue_fbs_order_label_print,
    queue_verification_problem_restock,
    request_invalid_kiz_reroute,
    record_inventory_scan,
    refresh_handover_acceptance,
    request_wb_handover_delivery,
    report_pick_missing_quantity,
    release_order_pick_restock_to_queue,
    retry_order_pick_restock_marketplace,
    scan_internal_movement,
    scan_handover_box,
    scan_pick_restock,
    stage_replenishment_allocation,
    validate_replenishment_source_scan,
    validate_pick_box_scan,
    validate_pick_cell_scan,
    verify_handover_order_label,
    verify_pick_allocation_unit,
    wb_handover_uses_marketplace_boxes,
    pick_restock_state,
)
from .services.picking import (
    pick_requires_box_scan,
    ACTIVE_TASK_STATUSES,
    active_partial_ozon_verification_task,
    assign_pick_handover_workstation,
    find_pick_verification_allocation,
    format_queue_rejections,
    resolve_verification_item_scan,
    wb_optional_marking_available,
)
from .services.floor_replenishment import (
    analyze_floor_replenishment_needs,
    create_floor_replenishment_movement,
)
from .services.handover import (
    approve_handover_verification_override,
    archive_handover_batch,
    get_ready_handover_box_for_order,
    handover_archive_readiness,
    handover_composition_readiness,
    handover_verification_override_status,
    has_handover_verification_override,
)
from .services.marketplace import (
    is_final_wb_marking_rejection,
    retry_wb_marking_code,
)
from .services.scanning import record_order_label_scan_events
from .services.labels import is_preloaded_ozon_order_label
from .services.printing import queue_fbs_preloaded_ozon_order_label_print
from .services.sync import marketplace_terminal_order_q
from .services.totes import (
    auto_release_pick_tote_if_complete,
    confirm_order_label_to_check_tote,
    controller_pick_tote_for_batch,
    mark_pick_tote_awaiting_empty,
    record_extra_problem_tote_item,
    transfer_controller_pick_tote,
)
from .services.controller_shift import controller_shift_is_live
from .staging import movement_staging_container_code


FULLBOX_DESKTOP_DIRECT_PRINT_MIN_VERSION = (1, 0, 12)


def _fullbox_desktop_version(request) -> tuple[int, int, int] | None:
    user_agent = str(request.META.get("HTTP_USER_AGENT") or "")
    match = re.search(r"(?:^|\s)FullboxDesktop/(\d+)\.(\d+)\.(\d+)(?:\s|$)", user_agent)
    if match is None:
        return None
    return tuple(int(value) for value in match.groups())


def _fullbox_desktop_supports_direct_print(request) -> bool:
    version = _fullbox_desktop_version(request)
    return version is not None and version >= FULLBOX_DESKTOP_DIRECT_PRINT_MIN_VERSION


def _desktop_label_print_target(request, *, shipping: bool = False) -> dict:
    if str(request.POST.get("fullbox_desktop") or "").strip() != "1":
        return {}
    if not _fullbox_desktop_supports_direct_print(request):
        return {}
    workstation_id = str(request.POST.get("desktop_workstation_id") or "").strip()
    printer_field = "desktop_shipping_printer" if shipping else "desktop_label_printer"
    fallback_field = "desktop_label_printer" if shipping else "desktop_shipping_printer"
    printer_name = str(request.POST.get(printer_field) or "").strip()
    if not printer_name:
        # Наклейки отгрузки и заказа одного размера, поэтому если нужный принтер
        # в Fullbox Desktop не выбран, берем второй выбранный.
        printer_name = str(request.POST.get(fallback_field) or "").strip()
    if not workstation_id or len(workstation_id) > 80 or not re.fullmatch(r"[A-Za-z0-9._-]+", workstation_id):
        return {}
    if not printer_name:
        # Ни один принтер не выбран. Отдаем печать рабочему месту вместо отказа:
        # у него принтер и агент настроены, иначе оператор видел бы только
        # ошибку и не мог напечатать вообще ничего.
        return {}
    return {
        "desktop_agent_id": f"desktop:{workstation_id}",
        "desktop_printer_name": printer_name,
    }


STOREKEEPER_ROLES = ("storekeeper", "head_manager", "director", "admin")
CONTROLLER_ROLES = ("fbs_controller", *STOREKEEPER_ROLES)
INVENTORY_MANAGER_ROLES = ("head_manager", "director", "admin")
REACHTRUCK_ROLES = ("reachtruck_driver", "director", "admin")


def _controller_service_binding_ready(*, session, binding, tote_id: int | None) -> bool:
    if session is None or binding is None or tote_id is None:
        return False
    allowed_states = {FbsToteBinding.STATE_AT_CONTROL}
    if tote_id == session.unknown_tote_id:
        allowed_states.add(FbsToteBinding.STATE_UNKNOWN)
    return bool(
        binding.state in allowed_states
        and binding.workstation_id == session.workstation_id
        and binding.controller_session_id == session.id
    )


DEMAND_DISPATCH_ORDER_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
)
DEMAND_MAPPING_ORDER_STATUSES = (
    *DEMAND_DISPATCH_ORDER_STATUSES,
    FbsOrder.STATUS_VALIDATION_FAILED,
)
DEMAND_MAPPING_STATUS_LABELS = {
    "barcode_missing": "Маркетплейс не передал ШК.",
    "barcode_ambiguous": "Маркетплейс передал несколько ШК; нужен один точный ШК.",
    "barcode_not_found": "ШК не найден в номенклатуре клиента.",
}
DEMAND_STATE_CHOICES = (
    ("action", "Требуют действия"),
    ("mapping", "Ошибка сопоставления"),
    ("ready", "Готовы к плану"),
    ("in_plan", "Уже в планах"),
    ("blocked", "Есть блокировка"),
    ("covered", "Потребность закрыта"),
)
PICKER_ROLES = ("picker", "processing_worker", "storekeeper", "head_manager", "director", "admin")
TSD_ROLES = tuple(
    dict.fromkeys((*CONTROLLER_ROLES, *REACHTRUCK_ROLES, *PICKER_ROLES))
)
OPEN_PLAN_STATUSES = (
    FbsReplenishmentPlan.STATUS_CONFIRMED,
    FbsReplenishmentPlan.STATUS_IN_PROGRESS,
)
STOREKEEPER_OPEN_PLAN_STATUSES = (
    *OPEN_PLAN_STATUSES,
    FbsReplenishmentPlan.STATUS_AWAITING_PACK,
)
OPEN_ALLOCATION_STATUSES = (
    FbsReplenishmentAllocation.STATUS_RESERVED,
    FbsReplenishmentAllocation.STATUS_IN_PROGRESS,
)
ACTIVE_ORDER_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_VALIDATION_FAILED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
    FbsOrder.STATUS_PICKING,
    FbsOrder.STATUS_PICKED,
    FbsOrder.STATUS_READY_FOR_HANDOVER,
    FbsOrder.STATUS_EXCEPTION,
)
STOREKEEPER_WORK_ORDER_STATUSES = tuple(
    status for status in ACTIVE_ORDER_STATUSES if status != FbsOrder.STATUS_EXCEPTION
)
ORDER_ATTENTION_STALE_HOURS = 2
ORDER_ATTENTION_NOT_STARTED_STATUSES = (
    FbsOrder.STATUS_RECEIVED,
    FbsOrder.STATUS_VALIDATION_FAILED,
    FbsOrder.STATUS_AWAITING_STOCK,
    FbsOrder.STATUS_RESERVED,
    FbsOrder.STATUS_EXCEPTION,
)
ORDER_ATTENTION_NOT_PICKED_STATUSES = (
    *ORDER_ATTENTION_NOT_STARTED_STATUSES,
    FbsOrder.STATUS_QUEUED_FOR_PICK,
    FbsOrder.STATUS_PICKING,
)
ACTIVE_PICK_BATCH_STATUSES = (
    FbsPickBatch.STATUS_QUEUED,
    FbsPickBatch.STATUS_IN_PROGRESS,
    FbsPickBatch.STATUS_VERIFICATION,
)
SEPARATED_PICK_RESTOCK_STATUSES = (
    FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
    FbsPickRestockRequest.STATUS_QUEUED,
    FbsPickRestockRequest.STATUS_IN_PROGRESS,
    FbsPickRestockRequest.STATUS_COMPLETED,
    FbsPickRestockRequest.STATUS_FAILED,
)
PICK_EQUIPMENT_SESSION_KEY = "fbs_tsd_pick_equipment"
PICK_NEW_BOX_NOTICE_SESSION_KEY = "fbs_tsd_pick_new_box_notice"
FBS_STOCK_SORT_FIELDS = {
    "client": "agency__agn_name",
    "sku": "sku_code",
    "name": "name",
    "barcode": "barcode",
    "qty": "qty",
    "available": "available_qty",
    "reserved": "reserved_qty",
    "box": "box__box_code",
    "pallet": "box__pallet__pallet_code",
    "cell": "box__pallet__cell__cell_code",
    "expiry": "expiry_date",
    "updated": "updated_at",
}
FBS_STOCK_DESCENDING_DEFAULTS = {"qty", "available", "reserved", "updated"}
FBS_STOCK_BOX_SORT_FIELDS = {
    "client": "agency__agn_name",
    "sku": "group_sku",
    "name": "group_name",
    "barcode": "group_barcode",
    "qty": "group_qty",
    "available": "group_available_qty",
    "reserved": "group_reserved_qty",
    "box": "box__box_code",
    "pallet": "box__pallet__pallet_code",
    "cell": "box__pallet__cell__cell_code",
    "expiry": "group_expiry",
    "updated": "group_updated",
}


def fbs_module_required(view_func):
    @wraps(view_func)
    def wrapped(request, *args, **kwargs):
        if not module_enabled():
            raise Http404
        return view_func(request, *args, **kwargs)

    return wrapped


def _base_context(request, **extra):
    route_name = str(getattr(getattr(request, "resolver_match", None), "url_name", "") or "")
    role = get_request_role(request)
    if "controller_settings" in route_name:
        operator_section = "controller_settings"
    elif "controller" in route_name:
        operator_section = "controller"
    elif "wave_settings" in route_name:
        operator_section = "wave_settings"
    elif "storekeeper_stock" in route_name:
        operator_section = "stock"
    elif route_name == "tsd_storekeeper":
        operator_section = "overview"
    elif "inventory" in route_name:
        operator_section = "inventory"
    elif "handover_reports" in route_name:
        operator_section = "handover_reports"
    elif "handover" in route_name:
        operator_section = "handover"
    elif "demand" in route_name or "plan" in route_name:
        operator_section = "demand"
    elif "metadata" in route_name:
        operator_section = "metadata"
    elif "label" in route_name:
        operator_section = "labels"
    elif "picking" in route_name or "pick_" in route_name:
        operator_section = "waves"
    else:
        operator_section = ""
    context = {
        "inventory_assignment_count": FbsInventorySession.objects.filter(
            Q(assigned_to=request.user, status__in=("planned", "draining", "counting"))
            | Q(recount_assigned_to=request.user, status="recount")
        ).count() if request.user.is_authenticated else 0,
        "employee": get_request_employee(request),
        "request_role": role,
        "workspace_shell": workspace_shell_enabled(role),
        "cabinet_url": resolve_cabinet_url(role),
        "warehouse_writes_enabled": feature_enabled("warehouse_writes"),
        "operator_section": operator_section,
    }
    context.update(extra)
    workspace_section = str(context.get("operator_section") or "")
    workspace_presentation = {
        "overview": (
            "Пульт смены FBS",
            "Всё, что требует решения кладовщика прямо сейчас",
            "Пульт смены",
        ),
        "stock": (
            "Остатки FBS",
            "Физический, доступный, зарезервированный и заблокированный",
            "Остатки FBS",
        ),
        "inventory": ("Инвентаризация", "", "Инвентаризация"),
        "handover": (
            "Отгрузки и короба",
            "Набор коробов, сканирование, передача водителю",
            "Отгрузки",
        ),
        "handover_reports": (
            "Отчёты по отгрузкам",
            "Отгрузки, заказы и товары отдельного контура FBS",
            "Отгрузки / Отчёты",
        ),
        "demand": ("Подсорт", "Что просит сборка", "Подсорт"),
        "labels": ("Этикетки", "", "Этикетки"),
        "wave_settings": (
            "Настройка волн",
            "Максимальный размер новых волн по клиентам и кабинетам",
            "Настройка волн",
        ),
        "controller_settings": (
            "FBS-контроллеры",
            "Волны, сборщики, контроллеры, рабочие столы и тара проверки",
            "Настройка контроллеров",
        ),
    }
    workspace_title, workspace_subtitle, workspace_crumb = workspace_presentation.get(
        workspace_section,
        (str(context.get("page_title") or "FBS"), "", "FBS"),
    )
    context.setdefault("workspace_section", workspace_section)
    context.setdefault("workspace_title", workspace_title)
    context.setdefault("workspace_subtitle", workspace_subtitle)
    context.setdefault("workspace_crumb", workspace_crumb)
    return context


def _plan_queryset():
    lines = FbsReplenishmentLine.objects.select_related(
        "sku_ref",
        "source_container",
        "target_box",
    ).order_by("id")
    return (
        FbsReplenishmentPlan.objects.select_related(
            "agency",
            "target_cell__location",
            "target_pallet",
            "target_box",
            "staging_location",
            "assigned_to",
            "confirmed_by",
        )
        .prefetch_related(Prefetch("lines", queryset=lines))
        .order_by("-created_at", "-id")
    )


def _order_attention_started_before_q(moment) -> Q:
    return Q(ordered_at__lte=moment) | Q(
        ordered_at__isnull=True,
        imported_at__lte=moment,
    )


def _storekeeper_responsible_rows():
    selected_user_ids = set(
        FbsStorekeeperResponsible.objects.filter(is_active=True).values_list(
            "user_id", flat=True
        )
    )
    rows = list(
        Employee.objects.filter(
            role="storekeeper",
            is_active=True,
            user__isnull=False,
            user__is_active=True,
        )
        .select_related("user")
        .order_by("full_name", "id")
    )
    for employee in rows:
        employee.is_fbs_responsible = employee.user_id in selected_user_ids
    return rows


def _storekeeper_list_context(request, *, error="", ok_message=""):
    plan_summaries = FbsReplenishmentPlan.objects.select_related("agency").order_by(
        "-created_at",
        "-id",
    )
    now = timezone.now()
    today = timezone.localdate()
    active_orders_queryset = FbsOrder.objects.filter(
        internal_status__in=STOREKEEPER_WORK_ORDER_STATUSES
    ).exclude(marketplace_terminal_order_q())
    sla_counts = active_orders_queryset.aggregate(
        tracked=Count("id", filter=Q(cutoff_at__isnull=False)),
        on_time=Count("id", filter=Q(cutoff_at__gte=now)),
        urgent=Count(
            "id",
            filter=Q(cutoff_at__gte=now, cutoff_at__lte=now + timedelta(hours=2)),
        ),
        overdue=Count("id", filter=Q(cutoff_at__lt=now)),
        without_cutoff=Count("id", filter=Q(cutoff_at__isnull=True)),
    )
    tracked_count = int(sla_counts["tracked"] or 0)
    overdue_count = int(sla_counts["overdue"] or 0)
    sla_counts["risk_percent"] = (
        round(overdue_count * 100 / tracked_count) if tracked_count else 0
    )
    sla_counts["next_cutoff"] = (
        active_orders_queryset.filter(cutoff_at__gte=now)
        .order_by("cutoff_at", "id")
        .values_list("cutoff_at", flat=True)
        .first()
    )
    attention_orders_queryset = (
        FbsOrder.objects.filter(
            internal_status__in=ORDER_ATTENTION_NOT_PICKED_STATUSES,
        )
        .exclude(marketplace_terminal_order_q())
        .select_related("profile__agency")
        .annotate(
            has_active_pick_task=Exists(
                FbsPickTask.objects.filter(
                    order_id=OuterRef("pk"),
                    status__in=ACTIVE_TASK_STATUSES,
                )
            ),
        )
    )
    stale_attention_orders = attention_orders_queryset.filter(
        internal_status__in=ORDER_ATTENTION_NOT_STARTED_STATUSES,
        has_active_pick_task=False,
    ).filter(
        _order_attention_started_before_q(
            now - timedelta(hours=ORDER_ATTENTION_STALE_HOURS)
        )
    )
    local_midnight = timezone.localtime(now).replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )
    previous_day_attention_orders = attention_orders_queryset.filter(
        _order_attention_started_before_q(local_midnight)
    )
    stale_attention_counts = stale_attention_orders.aggregate(
        total=Count("id"),
        awaiting_stock=Count(
            "id",
            filter=Q(internal_status=FbsOrder.STATUS_AWAITING_STOCK),
        ),
        without_cutoff=Count("id", filter=Q(cutoff_at__isnull=True)),
    )
    order_attention_counts = {
        "stale_unstarted": int(stale_attention_counts["total"] or 0),
        "previous_day_unpicked": previous_day_attention_orders.count(),
        "awaiting_stock": int(stale_attention_counts["awaiting_stock"] or 0),
        "without_cutoff": int(stale_attention_counts["without_cutoff"] or 0),
    }
    daily_flow = FbsOrder.objects.aggregate(
        received=Count("id", filter=Q(imported_at__date=today)),
        dispatched=Count(
            "id",
            filter=Q(
                handover_order__status=FbsHandoverOrder.STATUS_ACTIVE,
                handover_order__box__batch__dispatched_at__date=today,
            ),
            distinct=True,
        ),
        accepted=Count(
            "id",
            filter=Q(
                handover_order__status=FbsHandoverOrder.STATUS_ACTIVE,
                handover_order__box__batch__accepted_at__date=today,
            ),
            distinct=True,
        ),
        canceled=Count(
            "id",
            filter=Q(
                internal_status=FbsOrder.STATUS_CANCELLED,
                updated_at__date=today,
            ),
        ),
    )
    active_batches_queryset = (
        FbsPickBatch.objects.select_related(
            "agency",
            "assigned_to",
            "workstation",
            "cart",
        )
        .filter(status__in=ACTIVE_PICK_BATCH_STATUSES)
        .annotate(order_count=Count("tasks"))
        .order_by("created_at", "id")
    )
    active_batch_count = active_batches_queryset.count()
    active_batches = list(active_batches_queryset[:8])

    open_exceptions = FbsPickException.objects.filter(
        status=FbsPickException.STATUS_OPEN
    )
    recent_exceptions = list(
        open_exceptions.select_related(
            "task__order__profile__agency",
            "allocation__order_item",
            "created_by",
        ).order_by("-created_at", "-id")[:6]
    )

    label_counts = FbsOrderLabel.objects.aggregate(
        requested=Count("id", filter=Q(status=FbsOrderLabel.STATUS_REQUESTED)),
        ready=Count("id", filter=Q(status=FbsOrderLabel.STATUS_READY)),
        errors=Count("id", filter=Q(status=FbsOrderLabel.STATUS_ERROR)),
    )
    from .controller_shipment_ui import shipment_stage_counts
    shipment_counts = shipment_stage_counts()
    handover_counts = FbsHandoverBox.objects.aggregate(
        waiting_scan=Count("id", filter=Q(status=FbsHandoverBox.STATUS_CLOSED)),
        ready=Count("id", filter=Q(status=FbsHandoverBox.STATUS_SCANNED)),
        problems=Count("id", filter=Q(status=FbsHandoverBox.STATUS_PROBLEM)),
    )
    metadata_counts = FbsMarketplaceMetadataTransfer.objects.aggregate(
        waiting=Count(
            "id",
            filter=Q(
                status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
                    FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
                    FbsMarketplaceMetadataTransfer.STATUS_SENT,
                    FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                )
            ),
        ),
        conflicts=Count(
            "id",
            filter=Q(
                status__in=(
                    FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                    FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                )
            ),
        ),
    )
    order_counts = FbsOrder.objects.exclude(marketplace_terminal_order_q()).aggregate(
        received=Count("id", filter=Q(internal_status=FbsOrder.STATUS_RECEIVED)),
        awaiting_stock=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_AWAITING_STOCK)
        ),
        reserved=Count("id", filter=Q(internal_status=FbsOrder.STATUS_RESERVED)),
        queued=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_QUEUED_FOR_PICK)
        ),
        picking=Count("id", filter=Q(internal_status=FbsOrder.STATUS_PICKING)),
        picked=Count("id", filter=Q(internal_status=FbsOrder.STATUS_PICKED)),
        ready_for_handover=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_READY_FOR_HANDOVER)
        ),
        validation_failed=Count(
            "id", filter=Q(internal_status=FbsOrder.STATUS_VALIDATION_FAILED)
        ),
        exception=Count("id", filter=Q(internal_status=FbsOrder.STATUS_EXCEPTION)),
        active=Count(
            "id", filter=Q(internal_status__in=STOREKEEPER_WORK_ORDER_STATUSES)
        ),
        problem=Count(
            "id",
            filter=Q(
                internal_status__in=(
                    FbsOrder.STATUS_VALIDATION_FAILED,
                    FbsOrder.STATUS_EXCEPTION,
                )
            ),
        ),
    )

    workstations = list(
        FbsWorkstation.objects.filter(is_active=True)
        .select_related("device_agent", "shift_controller")
        .order_by("name", "id")
    )
    online_threshold = now - timedelta(seconds=60)
    for workstation in workstations:
        agent = workstation.device_agent
        workstation.agent_online = bool(
            agent and agent.last_seen and agent.last_seen >= online_threshold
        )
        controller = workstation.shift_controller
        workstation.shift_controller_name = (
            str(controller.get_full_name() or controller.get_username())
            if controller is not None
            else ""
        )
    active_cart_count = FbsPickingCart.objects.filter(is_active=True).count()
    online_workstation_count = sum(
        1 for workstation in workstations if workstation.agent_online
    )

    active_controller_sessions = list(
        FbsControllerSession.objects.select_related("controller", "workstation")
        .filter(
            status=FbsControllerSession.STATUS_ACTIVE,
            workstation__is_active=True,
        )
        .order_by("workstation__name", "workstation_id", "id")
    )
    live_controller_sessions = [
        session
        for session in active_controller_sessions
        if session.workstation.shift_controller_id == session.controller_id
        and controller_shift_is_live(session.workstation, now=now)
    ]
    controller_pick_totes = list(
        FbsControllerPickTote.objects.select_related(
            "tote",
            "session__controller",
            "session__workstation",
            "check_tote__profile__agency",
            "pick_batch",
        )
        .filter(
            status__in=(
                FbsControllerPickTote.STATUS_PROCESSING,
                FbsControllerPickTote.STATUS_AWAITING_EMPTY,
            ),
            pick_batch__status=FbsPickBatch.STATUS_VERIFICATION,
            pick_batch__cart_released_at__isnull=True,
            pick_batch__completed_at__isnull=True,
        )
        .order_by(
            "session__workstation__name",
            "attached_at",
            "id",
        )
    )
    controller_tote_transfer_rows = []
    for pick_tote in controller_pick_totes:
        transfer_targets = [
            session
            for session in live_controller_sessions
            if session.id != pick_tote.session_id
            and session.controller_id != pick_tote.session.controller_id
        ]
        controller_tote_transfer_rows.append(
            {
                "pick_tote": pick_tote,
                "targets": transfer_targets,
            }
        )

    open_exception_count = open_exceptions.count()
    problem_count = (
        int(order_counts["problem"] or 0)
        + open_exception_count
        + int(label_counts["errors"] or 0)
        + int(handover_counts["problems"] or 0)
        + int(metadata_counts["conflicts"] or 0)
    )
    profile_counts = FbsIntegrationProfile.objects.aggregate(
        total=Count("id"),
        active=Count("id", filter=Q(is_active=True)),
    )
    proposed_plans = plan_summaries.filter(
        status=FbsReplenishmentPlan.STATUS_PROPOSED
    )[:30]
    active_plans = plan_summaries.filter(
        status__in=STOREKEEPER_OPEN_PLAN_STATUSES
    )[:30]
    return _base_context(
        request,
        page_title="FBS · Смена",
        operator_section="overview",
        back_url="/sklad/",
        proposed_plans=proposed_plans,
        active_plans=active_plans,
        active_batches=active_batches,
        recent_exceptions=recent_exceptions,
        order_counts=order_counts,
        label_counts=label_counts,
        handover_counts=handover_counts,
        shipment_counts=shipment_counts,
        metadata_counts=metadata_counts,
        profile_counts=profile_counts,
        sla_counts=sla_counts,
        order_attention_counts=order_attention_counts,
        daily_flow=daily_flow,
        active_inventory_count=FbsInventorySession.objects.exclude(
            status__in=(
                FbsInventorySession.STATUS_DONE,
                FbsInventorySession.STATUS_CANCELED,
            )
        ).count(),
        equipment_summary={
            "workstations": len(workstations),
            "workstations_online": online_workstation_count,
            "carts": active_cart_count,
        },
        workstations=workstations[:4],
        controller_tote_transfer_rows=controller_tote_transfer_rows,
        live_controller_session_count=len(live_controller_sessions),
        fbs_responsible_rows=_storekeeper_responsible_rows(),
        can_manage_fbs_responsibles=(
            get_request_role(request) in INVENTORY_MANAGER_ROLES
        ),
        dashboard_counts={
            "new": int(order_counts["received"] or 0)
            + int(order_counts["awaiting_stock"] or 0),
            "to_pick": int(order_counts["reserved"] or 0)
            + int(order_counts["queued"] or 0),
            "waves": active_batch_count,
            "problems": problem_count,
        },
        problem_breakdown={
            "orders": int(order_counts["problem"] or 0),
            "pick": open_exception_count,
            "labels": int(label_counts["errors"] or 0),
            "handover": int(handover_counts["problems"] or 0),
            "metadata": int(metadata_counts["conflicts"] or 0),
        },
        error=error,
        ok_message=ok_message,
    )


def _controller_settings_user_name(user) -> str:
    if user is None:
        return "Система"
    employee = getattr(user, "employee_profile", None)
    return (
        str(getattr(employee, "full_name", "") or "").strip()
        or str(user.get_full_name() or "").strip()
        or str(user.get_username() or "Система")
    )


def _controller_settings_context(request, *, error=""):
    now = timezone.now()
    active_statuses = (
        FbsControllerPickTote.STATUS_PROCESSING,
        FbsControllerPickTote.STATUS_AWAITING_EMPTY,
    )
    active_sessions = list(
        FbsControllerSession.objects.select_related(
            "controller__employee_profile",
            "workstation",
            "unknown_tote",
        )
        .filter(status=FbsControllerSession.STATUS_ACTIVE)
        .annotate(
            active_tote_count=Count(
                "pick_totes",
                filter=Q(pick_totes__status__in=active_statuses),
                distinct=True,
            )
        )
        .order_by("workstation__name", "workstation_id", "id")
    )
    live_sessions = []
    for session in active_sessions:
        session.controller_name = _controller_settings_user_name(session.controller)
        session.is_live = bool(
            session.workstation.shift_controller_id == session.controller_id
            and controller_shift_is_live(session.workstation, now=now)
        )
        if session.is_live:
            live_sessions.append(session)

    query = str(request.GET.get("q") or "").strip()
    selected_status = str(request.GET.get("status") or "all").strip()
    if selected_status not in {"all", "active", "closed", "problem"}:
        selected_status = "all"
    try:
        page_size = int(request.GET.get("page_size") or 50)
    except (TypeError, ValueError):
        page_size = 50
    if page_size not in (25, 50, 100):
        page_size = 50

    tote_queryset = FbsControllerPickTote.objects.select_related(
        "session__controller__employee_profile",
        "session__workstation",
        "check_tote",
        "pick_batch__agency",
        "pick_batch__assigned_to__employee_profile",
        "pick_batch__created_by__employee_profile",
        "tote",
    )
    if selected_status == "active":
        tote_queryset = tote_queryset.filter(status__in=active_statuses)
    elif selected_status == "closed":
        tote_queryset = tote_queryset.filter(status=FbsControllerPickTote.STATUS_CLOSED)
    elif selected_status == "problem":
        tote_queryset = tote_queryset.filter(status=FbsControllerPickTote.STATUS_PROBLEM)
    if query:
        query_filter = (
            Q(tote__barcode__icontains=query)
            | Q(tote__name__icontains=query)
            | Q(pick_batch__agency__agn_name__icontains=query)
            | Q(session__workstation__name__icontains=query)
            | Q(session__controller__username__icontains=query)
            | Q(session__controller__employee_profile__full_name__icontains=query)
            | Q(pick_batch__assigned_to__username__icontains=query)
            | Q(pick_batch__assigned_to__employee_profile__full_name__icontains=query)
            | Q(pick_batch__created_by__username__icontains=query)
            | Q(pick_batch__created_by__employee_profile__full_name__icontains=query)
        )
        if query.isdigit():
            query_filter |= Q(pick_batch_id=int(query)) | Q(pk=int(query))
        tote_queryset = tote_queryset.filter(query_filter)

    tote_queryset = tote_queryset.annotate(
        active_sort=Case(
            When(status__in=active_statuses, then=Value(0)),
            default=Value(1),
            output_field=IntegerField(),
        )
    ).order_by("active_sort", "-updated_at", "-id")
    page = Paginator(tote_queryset, page_size).get_page(request.GET.get("page"))
    rows = list(page.object_list)
    transferable_ids = set(
        FbsControllerPickTote.objects.filter(
            pk__in=[row.pk for row in rows],
            status__in=active_statuses,
            pick_batch__status=FbsPickBatch.STATUS_VERIFICATION,
            pick_batch__cart_released_at__isnull=True,
            pick_batch__completed_at__isnull=True,
        ).values_list("pk", flat=True)
    )
    for row in rows:
        row.controller_name = _controller_settings_user_name(row.session.controller)
        row.picker_name = _controller_settings_user_name(row.pick_batch.assigned_to)
        row.created_by_name = _controller_settings_user_name(row.pick_batch.created_by)
        row.is_active = row.status in active_statuses
        row.transfer_targets = [
            session
            for session in live_sessions
            if row.pk in transferable_ids
            and session.id != row.session_id
            and session.controller_id != row.session.controller_id
        ]

    waiting_queryset = FbsToteBinding.objects.filter(
        state=FbsToteBinding.STATE_WAITING_CONTROL, workstation__isnull=False,
        pick_batch__isnull=False,
    ).select_related("tote", "workstation", "pick_batch__agency", "pick_batch__assigned_to")
    if query:
        waiting_filter = (Q(tote__barcode__icontains=query) | Q(tote__name__icontains=query)
            | Q(workstation__name__icontains=query) | Q(pick_batch__agency__agn_name__icontains=query)
            | Q(pick_batch__assigned_to__username__icontains=query)
            | Q(pick_batch__assigned_to__employee_profile__full_name__icontains=query))
        if query.isdigit():
            waiting_filter |= Q(pick_batch_id=int(query))
        waiting_queryset = waiting_queryset.filter(waiting_filter)
    waiting_rows = list(waiting_queryset.order_by("workstation_id", "updated_at", "id")) if selected_status in {"all", "active"} else []
    for binding in waiting_rows:
        binding.picker_name = _controller_settings_user_name(binding.pick_batch.assigned_to)
        binding.transfer_targets = [session for session in live_sessions
            if session.workstation_id != binding.workstation_id]

    all_totes = FbsControllerPickTote.objects.aggregate(
        total=Count("id"),
        active=Count("id", filter=Q(status__in=active_statuses)),
        closed=Count("id", filter=Q(status=FbsControllerPickTote.STATUS_CLOSED)),
        problem=Count("id", filter=Q(status=FbsControllerPickTote.STATUS_PROBLEM)),
    )
    transfer_events = list(
        FbsToteMovement.objects.select_related(
            "tote",
            "pick_batch",
            "performed_by__employee_profile",
            "controller_session__controller__employee_profile",
        )
        .filter(details__operation="controller_tote_transfer")
        .order_by("-created_at", "-id")[:30]
    )
    source_session_ids = {
        int(event.details[key]) for event in transfer_events
        for key in ("source_session_id", "target_session_id")
        if str(event.details.get(key) or "").isdigit()
    }
    source_sessions = {
        session.id: session
        for session in FbsControllerSession.objects.select_related(
            "controller__employee_profile", "workstation"
        ).filter(pk__in=source_session_ids)
    }
    for event in transfer_events:
        source_session_value = event.details.get("source_session_id")
        source_session = source_sessions.get(
            int(source_session_value)
            if str(source_session_value or "").isdigit()
            else None
        )
        event.source_controller_name = (
            _controller_settings_user_name(source_session.controller)
            if source_session
            else "—"
        )
        audit_target_session = event.controller_session or source_sessions.get(
            int(event.details.get("target_session_id") or 0))
        event.target_controller_name = (
            _controller_settings_user_name(audit_target_session.controller)
            if audit_target_session
            else "—"
        )
        event.performed_by_name = _controller_settings_user_name(event.performed_by)

    filter_query = urlencode(
        {"q": query, "status": selected_status, "page_size": page_size}
    )
    return _base_context(
        request,
        page_title="FBS-контроллеры",
        operator_section="controller_settings",
        active_sessions=active_sessions,
        live_session_count=len(live_sessions),
        tote_rows=rows,
        waiting_tote_rows=waiting_rows,
        tote_page=page,
        tote_summary={key: int(value or 0) for key, value in all_totes.items()},
        transfer_events=transfer_events,
        query=query,
        selected_status=selected_status,
        page_size=page_size,
        page_sizes=(25, 50, 100),
        filter_query=filter_query,
        error=error,
    )


def _storekeeper_plan_context(request, plan, *, error="", ok_message=""):
    staged_qty = int(
        FbsReplenishmentAllocation.objects.filter(line__plan=plan).aggregate(
            total=Sum("qty_staged")
        )["total"]
        or 0
    )
    boxes = FbsBox.objects.filter(
        agency=plan.agency,
        status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        pallet__status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        pallet__cell__is_active=True,
    ).select_related("pallet__cell").order_by(
        "pallet__cell__cell_code", "pallet__pallet_code", "box_code", "id"
    )
    pallets = FbsPallet.objects.filter(
        agency=plan.agency,
        status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        cell__is_active=True,
    ).select_related("cell").annotate(
        active_box_count=Count(
            "boxes",
            filter=Q(boxes__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE)),
        )
    ).order_by("cell__cell_code", "pallet_code", "id")
    prepared_boxes = list(
        FbsReplenishmentPreparedBox.objects.filter(plan=plan)
        .select_related("fbs_box__pallet__cell")
        .order_by("sequence_no", "id")
    )
    return _base_context(
        request,
        page_title=f"План FBS #{plan.id}",
        back_url=reverse("fbs:tsd_storekeeper"),
        plan=plan,
        staging_container_code=(
            movement_staging_container_code(plan.client_movement_request_id)
            if plan.client_movement_request_id
            and plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and not prepared_boxes
            else ""
        ),
        staged_qty=staged_qty,
        can_pack=(
            not prepared_boxes
            and plan.status == FbsReplenishmentPlan.STATUS_AWAITING_PACK
        ),
        prepared_boxes=prepared_boxes,
        prepared_scanned_qty=sum(int(box.scanned_qty or 0) for box in prepared_boxes),
        prepared_placed_qty=sum(int(box.placed_qty or 0) for box in prepared_boxes),
        dynamic_destination=(
            plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and plan.status != FbsReplenishmentPlan.STATUS_DONE
        ),
        awaiting_placement=(
            not prepared_boxes
            and
            plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and plan.status == FbsReplenishmentPlan.STATUS_IN_PROGRESS
            and plan.target_box_id is not None
            and int(plan.moved_qty or 0) < int(plan.planned_qty or 0)
        ),
        boxes=boxes,
        pallets=pallets,
        error=error,
        ok_message=ok_message,
    )


def _demand_row_state(row) -> tuple[str, str, str]:
    if row.general_shortage_qty and row.proposed_qty:
        return "partial", "Частично готово", "conflict"
    if row.general_shortage_qty:
        return "blocked_stock", "Нет общего остатка", "failed"
    if row.destination_blocked_qty:
        return "blocked_destination", "Нет FBS-короба", "failed"
    if row.proposed_qty:
        return "ready", "Готово к плану", "confirmed"
    if row.open_plan_qty:
        return "in_plan", "Уже в планах", "in_progress"
    return "covered", "Потребность закрыта", "done"


def _demand_mapping_issue(item, valid_barcode_pairs) -> tuple[str, str]:
    requirements = item.requirements if isinstance(item.requirements, dict) else {}
    mapping_status = str(requirements.get("sku_mapping") or "").strip()
    imported_reason = DEMAND_MAPPING_STATUS_LABELS.get(mapping_status, "")
    barcode = str(item.barcode or "").strip()
    if item.sku_id is None:
        return mapping_status or "sku_missing", imported_reason or "SKU не привязан."
    if not barcode:
        return mapping_status or "barcode_missing", imported_reason or "ШК не указан."
    if item.sku.agency_id != item.order.profile.agency_id:
        return "foreign_sku", "SKU принадлежит другому клиенту."
    if item.sku.deleted:
        return "deleted_sku", "SKU удален из номенклатуры клиента."
    if (item.sku_id, barcode) not in valid_barcode_pairs:
        return "barcode_not_bound", "ШК не привязан к SKU клиента."
    return "", ""


def _demand_context(request, *, error="", ok_message=""):
    analysis = analyze_replenishment_demand()
    floor_analysis = analyze_floor_replenishment_needs()
    params = request.GET if request.method == "GET" else request.POST
    query = str(params.get("q") or "").strip()
    agency_filter = str(params.get("agency") or "").strip()
    state_filter = str(params.get("state") or "").strip()
    deadline_filter = str(params.get("deadline") or "").strip()
    mapping_items = list(
        FbsOrderItem.objects.filter(
            order__internal_status__in=DEMAND_MAPPING_ORDER_STATUSES,
        )
        .select_related("order__profile__agency", "sku")
        .order_by("order__cutoff_at", "order_id", "id")
    )
    agency_ids = {row.agency_id for row in analysis.rows} | {
        item.order.profile.agency_id for item in mapping_items
    }
    sku_ids = {row.sku_id for row in analysis.rows}
    analysis_barcodes = {
        str(getattr(row, "barcode", "") or "").strip()
        for row in analysis.rows
        if str(getattr(row, "barcode", "") or "").strip()
    }
    if not agency_filter.isdigit() or int(agency_filter) not in agency_ids:
        agency_filter = ""
    if state_filter not in {value for value, _ in DEMAND_STATE_CHOICES}:
        state_filter = ""
    if deadline_filter not in {"overdue", "urgent", "today", "none"}:
        deadline_filter = ""

    order_items = (
        FbsOrderItem.objects.filter(
            order__internal_status__in=DEMAND_DISPATCH_ORDER_STATUSES,
            order__profile__agency_id__in=agency_ids,
        )
        .filter(Q(sku_id__in=sku_ids) | Q(barcode__in=analysis_barcodes))
        .select_related("order__profile", "order__profile__agency")
        .order_by("order__cutoff_at", "order_id", "id")
    )
    orders_by_sku = defaultdict(dict)
    orders_by_barcode = defaultdict(dict)
    barcodes_by_sku = defaultdict(set)
    now = timezone.now()
    today = timezone.localdate()

    def add_order(target, key, item):
        entry = target[key].get(item.order_id)
        if entry is None:
            cutoff = item.order.cutoff_at
            entry = SimpleNamespace(
                id=item.order_id,
                external_order_id=item.order.external_order_id,
                marketplace=item.order.profile.get_marketplace_display(),
                quantity=0,
                cutoff_at=cutoff,
                is_overdue=bool(cutoff and cutoff < now),
                is_urgent=bool(cutoff and now <= cutoff <= now + timedelta(hours=2)),
                is_today=bool(cutoff and timezone.localtime(cutoff).date() == today),
            )
            target[key][item.order_id] = entry
        entry.quantity += int(item.quantity or 0)

    for item in order_items:
        agency_id = item.order.profile.agency_id
        barcode = str(item.barcode or "").strip()
        if item.sku_id:
            sku_key = (agency_id, item.sku_id)
            add_order(orders_by_sku, sku_key, item)
            if barcode:
                barcodes_by_sku[sku_key].add(barcode)
        if barcode:
            add_order(orders_by_barcode, (agency_id, barcode), item)

    mapping_sku_ids = {item.sku_id for item in mapping_items if item.sku_id}
    mapping_barcodes = {
        str(item.barcode or "").strip()
        for item in mapping_items
        if str(item.barcode or "").strip()
    }
    valid_mapping_pairs = set(
        SKUBarcode.objects.filter(
            sku_id__in=mapping_sku_ids,
            value__in=mapping_barcodes,
        ).values_list("sku_id", "value")
    )
    all_mapping_rows = []
    for item in mapping_items:
        reason_code, reason = _demand_mapping_issue(item, valid_mapping_pairs)
        if not reason_code:
            continue
        requirements = item.requirements if isinstance(item.requirements, dict) else {}
        marketplace_barcodes = [
            str(value or "").strip()
            for value in requirements.get("marketplace_barcodes", ())
            if str(value or "").strip()
        ]
        barcode = str(item.barcode or "").strip()
        display_barcode = barcode or ", ".join(marketplace_barcodes[:3]) or "Не передан"
        if len(marketplace_barcodes) > 3 and not barcode:
            display_barcode += f" +{len(marketplace_barcodes) - 3}"
        cutoff = item.order.cutoff_at
        all_mapping_rows.append(
            SimpleNamespace(
                id=item.id,
                order_id=item.order_id,
                external_order_id=item.order.external_order_id,
                agency_id=item.order.profile.agency_id,
                agency_name=str(item.order.profile.agency),
                marketplace=item.order.profile.get_marketplace_display(),
                external_sku=str(item.external_sku or "").strip() or "Без артикула",
                product_name=(
                    str(item.product_name or "").strip()
                    or str(getattr(item.sku, "name", "") or "").strip()
                    or "Товар не определен"
                ),
                display_barcode=display_barcode,
                quantity=int(item.quantity or 0),
                reason_code=reason_code,
                reason=reason,
                cutoff_at=cutoff,
                is_overdue=bool(cutoff and cutoff < now),
                is_urgent=bool(cutoff and now <= cutoff <= now + timedelta(hours=2)),
                is_today=bool(cutoff and timezone.localtime(cutoff).date() == today),
            )
        )

    mapping_rows = []
    for row in all_mapping_rows:
        if agency_filter and row.agency_id != int(agency_filter):
            continue
        query_folded = query.casefold()
        if query_folded and not any(
            query_folded in value.casefold()
            for value in (
                row.external_order_id,
                row.agency_name,
                row.external_sku,
                row.product_name,
                row.display_barcode,
                row.reason,
            )
        ):
            continue
        if state_filter not in {"", "action", "blocked", "mapping"}:
            continue
        if deadline_filter == "overdue" and not row.is_overdue:
            continue
        if deadline_filter == "urgent" and not row.is_urgent:
            continue
        if deadline_filter == "today" and not row.is_today:
            continue
        if deadline_filter == "none" and row.cutoff_at is not None:
            continue
        mapping_rows.append(row)
    mapping_rows.sort(
        key=lambda row: (
            0 if row.is_overdue else 1 if row.is_urgent else 2,
            row.cutoff_at is None,
            row.cutoff_at or now + timedelta(days=36500),
            row.agency_name.casefold(),
            row.external_order_id,
            row.id,
        )
    )

    demand_rows = []
    for row in analysis.rows:
        row_barcode = str(getattr(row, "barcode", "") or "").strip()
        sku_key = (row.agency_id, row.sku_id)
        if row_barcode:
            order_rows = list(
                orders_by_barcode[(row.agency_id, row_barcode)].values()
            )
            barcode_label = row_barcode
            order_search = row_barcode
        else:
            order_rows = list(orders_by_sku[sku_key].values())
            row_barcodes = sorted(barcodes_by_sku[sku_key])
            barcode_label = ", ".join(row_barcodes[:3]) or "не указан"
            if len(row_barcodes) > 3:
                barcode_label += f" +{len(row_barcodes) - 3}"
            order_search = row.sku_code
        order_rows.sort(
            key=lambda order: (
                order.cutoff_at is None,
                order.cutoff_at or now + timedelta(days=36500),
                order.id,
            )
        )
        if deadline_filter == "overdue":
            visible_order_rows = [order for order in order_rows if order.is_overdue]
        elif deadline_filter == "urgent":
            visible_order_rows = [order for order in order_rows if order.is_urgent]
        elif deadline_filter == "today":
            visible_order_rows = [order for order in order_rows if order.is_today]
        elif deadline_filter == "none":
            visible_order_rows = [
                order for order in order_rows if order.cutoff_at is None
            ]
        else:
            visible_order_rows = order_rows
        state, state_label, state_css = _demand_row_state(row)
        nearest_cutoff = next(
            (
                order.cutoff_at
                for order in visible_order_rows
                if order.cutoff_at is not None
            ),
            None,
        )
        decorated = SimpleNamespace(
            **row.__dict__,
            barcode_label=barcode_label,
            orders=visible_order_rows[:8],
            order_ids={order.id for order in visible_order_rows},
            order_count=len(visible_order_rows),
            hidden_order_count=max(len(visible_order_rows) - 8, 0),
            overdue_order_ids={
                order.id for order in visible_order_rows if order.is_overdue
            },
            urgent_order_ids={
                order.id for order in visible_order_rows if order.is_urgent
            },
            today_order_ids={
                order.id for order in visible_order_rows if order.is_today
            },
            no_cutoff_order_ids={
                order.id for order in visible_order_rows if order.cutoff_at is None
            },
            nearest_cutoff=nearest_cutoff,
            state=state,
            state_label=state_label,
            state_css=state_css,
            orders_url=(
                reverse("fbs:operator_orders")
                + "?"
                + urlencode(
                    {"agency": row.agency_id, "q": order_search, "page_size": 100}
                )
            ),
        )
        if agency_filter and row.agency_id != int(agency_filter):
            continue
        query_folded = query.casefold()
        if query_folded and not any(
            query_folded in value.casefold()
            for value in (
                row.agency_name,
                row.sku_code,
                row.sku_name,
                barcode_label,
                " ".join(order.external_order_id for order in order_rows),
            )
        ):
            continue
        if state_filter == "action" and not (
            row.proposed_qty or row.general_shortage_qty or row.destination_blocked_qty
        ):
            continue
        if state_filter == "mapping":
            continue
        if state_filter == "ready" and not (
            row.proposed_qty
            and not row.general_shortage_qty
            and not row.destination_blocked_qty
        ):
            continue
        if state_filter == "in_plan" and not row.open_plan_qty:
            continue
        if state_filter == "blocked" and not (
            row.general_shortage_qty or row.destination_blocked_qty
        ):
            continue
        if state_filter == "covered" and row.uncovered_qty:
            continue
        if deadline_filter and not decorated.order_ids:
            continue
        demand_rows.append(decorated)

    state_priority = {
        "blocked_stock": 0,
        "blocked_destination": 0,
        "partial": 1,
        "ready": 2,
        "in_plan": 3,
        "covered": 4,
    }
    demand_rows.sort(
        key=lambda row: (
            0 if row.overdue_order_ids else 1 if row.urgent_order_ids else 2,
            state_priority[row.state],
            row.nearest_cutoff is None,
            row.nearest_cutoff or now + timedelta(days=36500),
            row.agency_name.casefold(),
            row.sku_code.casefold(),
            row.barcode_label,
        )
    )
    affected_order_ids = (
        set().union(*(row.order_ids for row in demand_rows)) if demand_rows else set()
    )
    overdue_order_ids = (
        set().union(*(row.overdue_order_ids for row in demand_rows)) if demand_rows else set()
    )
    urgent_order_ids = (
        set().union(*(row.urgent_order_ids for row in demand_rows)) if demand_rows else set()
    )
    return _base_context(
        request,
        page_title="FBS · Потребность",
        back_url=reverse("fbs:tsd_storekeeper"),
        analysis=analysis,
        floor_analysis=floor_analysis,
        demand_rows=demand_rows,
        demand_agencies=Agency.objects.filter(id__in=agency_ids).order_by("agn_name", "id"),
        demand_state_choices=DEMAND_STATE_CHOICES,
        demand_filters={
            "q": query,
            "agency": agency_filter,
            "state": state_filter,
            "deadline": deadline_filter,
        },
        dispatch_summary={
            "rows": len(demand_rows),
            "orders": len(affected_order_ids),
            "overdue": len(overdue_order_ids),
            "urgent": len(urgent_order_ids),
        },
        mapping_rows=mapping_rows[:100],
        mapping_summary={
            "total_lines": len(all_mapping_rows),
            "total_orders": len({row.order_id for row in all_mapping_rows}),
            "total_qty": sum(row.quantity for row in all_mapping_rows),
            "filtered_lines": len(mapping_rows),
            "filtered_orders": len({row.order_id for row in mapping_rows}),
            "filtered_qty": sum(row.quantity for row in mapping_rows),
            "hidden_lines": max(len(mapping_rows) - 100, 0),
        },
        error=error,
        ok_message=ok_message,
    )


def _reachtruck_list_context(request, *, error="", ok_message=""):
    plans = (
        _plan_queryset()
        .filter(status__in=OPEN_PLAN_STATUSES)
        .filter(Q(assigned_to__isnull=True) | Q(assigned_to=request.user))
    )
    completed = _plan_queryset().filter(
        status=FbsReplenishmentPlan.STATUS_DONE,
        assigned_to=request.user,
    )[:10]
    return _base_context(
        request,
        page_title="FBS · Задания",
        back_url="/reachtruck/",
        available_plans=plans[:30],
        available_movements=FbsInternalMovement.objects.select_related(
            "agency",
            "source_box",
            "target_pallet__cell",
            "target_box",
            "assigned_to",
        ).filter(
            status__in=(
                FbsInternalMovement.STATUS_PROPOSED,
                FbsInternalMovement.STATUS_IN_PROGRESS,
            )
        ).filter(
            Q(assigned_to__isnull=True) | Q(assigned_to=request.user)
        ).order_by("created_at", "id")[:30],
        completed_plans=completed,
        error=error,
        ok_message=ok_message,
    )


def _allocation_queryset():
    return FbsReplenishmentAllocation.objects.select_related(
        "line__plan__agency",
        "line__plan__assigned_to",
        "line__plan__target_cell__location",
        "line__plan__target_pallet",
        "line__source_container",
        "source_snapshot__container",
        "source_snapshot__location",
        "source_snapshot__sku_ref",
        "target_box",
        "warehouse_task",
    )


def _pick_task_queryset():
    return FbsPickTask.objects.select_related(
        "batch__agency",
        "order__profile__agency",
        "assigned_to",
    ).order_by("batch_id", "sort_order", "id")


def _pick_allocation_queryset():
    return FbsOrderStockAllocation.objects.select_related(
        "pick_task__batch__agency",
        "pick_task__order__profile__agency",
        "pick_task__assigned_to",
        "order_item__sku",
        "order_item__order__profile",
        "balance__box__pallet__cell__location",
        "balance__box__source_container__current_location",
        "balance__sku_ref",
        "traceability",
        "verification_progress",
    )


def _picker_cart(request):
    stored = request.session.get(PICK_EQUIPMENT_SESSION_KEY)
    if not isinstance(stored, dict):
        return None
    try:
        cart_id = int(stored.get("cart_id"))
    except (TypeError, ValueError):
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        return None
    cart = FbsPickingCart.objects.filter(pk=cart_id, is_active=True).first()
    if cart is None:
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        return None
    return cart


def _store_picker_cart(request, cart) -> None:
    request.session[PICK_EQUIPMENT_SESSION_KEY] = {
        "cart_id": cart.id,
    }


def _open_pick_allocations(batch):
    return (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            status__in=(
                FbsOrderStockAllocation.STATUS_RESERVED,
                FbsOrderStockAllocation.STATUS_PICKING,
            ),
        )
        .order_by(
            *pick_route_ordering("balance__box__pallet__cell__location__"),
            "balance__box__box_code",
            "balance__sku_code",
            "pick_task__sort_order",
            "id",
        )
    )


def _next_pick_allocation_id(batch) -> int | None:
    return _open_pick_allocations(batch).values_list("id", flat=True).first()


def _picking_context(request, *, error="", ok_message=""):
    now = timezone.now()
    overdue_boundary = now - timedelta(hours=15)
    batches = (
        FbsPickBatch.objects.select_related(
            "agency",
            "assigned_to",
            "workstation",
            "cart",
        )
        .filter(
            status__in=(
                FbsPickBatch.STATUS_QUEUED,
                FbsPickBatch.STATUS_IN_PROGRESS,
                FbsPickBatch.STATUS_VERIFICATION,
            )
        )
        .filter(Q(assigned_to__isnull=True) | Q(assigned_to=request.user))
        .exclude(pick_restock_request__status__in=ACTIVE_PICK_RESTOCK_STATUSES)
        .annotate(
            oldest_order_at=Min(
                Coalesce("tasks__order__ordered_at", "tasks__order__imported_at")
            )
        )
        .order_by("created_at", "id")
    )
    my_active_batches = list(
        batches.filter(assigned_to=request.user).order_by("created_at", "id")
    )
    free_batches = list(
        order_wave_queue_queryset(batches.filter(assigned_to__isnull=True))[:50]
    )
    decorate_wave_queue_rows(free_batches)
    available_batches = [*my_active_batches, *free_batches]
    for batch in available_batches:
        batch.order_count = batch.tasks.count()
        batch.next_allocation_id = _next_pick_allocation_id(batch)
        batch.is_overdue = bool(
            batch.oldest_order_at
            and batch.oldest_order_at <= overdue_boundary
        )
        batch.overdue_age_hours = (
            int((now - batch.oldest_order_at).total_seconds() // 3600)
            if batch.is_overdue
            else 0
        )
    cart = _picker_cart(request)
    free_batch_count = FbsPickBatch.objects.filter(
        status=FbsPickBatch.STATUS_QUEUED,
        assigned_to__isnull=True,
    ).count()
    active_wave_count = FbsPickBatch.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsPickBatch.STATUS_IN_PROGRESS,
            FbsPickBatch.STATUS_VERIFICATION,
        ),
        picking_completed_at__isnull=True,
    ).count()
    cart_busy = bool(
        cart
        and FbsPickBatch.objects.filter(
            cart=cart,
            status__in=(
                FbsPickBatch.STATUS_IN_PROGRESS,
                FbsPickBatch.STATUS_VERIFICATION,
            ),
            cart_released_at__isnull=True,
        ).exists()
    )
    request_role = get_request_role(request)
    restock_requests = list(
        FbsPickRestockRequest.objects.select_related(
            "batch__workstation",
            "assigned_to",
            "source_tote",
            "quarantine_box__pallet__cell__location",
        )
        .filter(
            Q(
                status=FbsPickRestockRequest.STATUS_QUEUED,
                assigned_to__isnull=True,
            )
            | Q(
                status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
                assigned_to=request.user,
            )
        )
        .filter(Q(order__isnull=True) | Q(source_tote__isnull=False))
        .annotate(source_box_count=Count("lines__source_box", distinct=True))
        .order_by("created_at", "id")[:20]
    )
    overdue_unassembled = (
        FbsOrder.objects.select_related("profile__agency")
        .filter(
            profile__is_active=True,
            internal_status__in=(
                FbsOrder.STATUS_RECEIVED,
                FbsOrder.STATUS_VALIDATION_FAILED,
                FbsOrder.STATUS_AWAITING_STOCK,
                FbsOrder.STATUS_RESERVED,
                FbsOrder.STATUS_QUEUED_FOR_PICK,
                FbsOrder.STATUS_PICKING,
            ),
        )
        .annotate(age_started_at=Coalesce("ordered_at", "imported_at"))
        .filter(age_started_at__lte=overdue_boundary)
        .order_by("age_started_at", "id")
    )
    overdue_oldest = overdue_unassembled.first()
    overdue_client_rows = list(
        overdue_unassembled.annotate(
            client_name=F("profile__agency__agn_name")
        )
        .values("client_name")
        .annotate(order_count=Count("id"))
        .order_by("-order_count", "client_name")
    )
    return _base_context(
        request,
        page_title="Сборка FBS" if request_role == "picker" else "FBS · Отбор",
        back_url=(
            reverse("fbs:tsd_home")
            if request_role == "picker"
            else "/sklad/" if request_role in STOREKEEPER_ROLES else ""
        ),
        available_batches=available_batches,
        my_active_batches=my_active_batches,
        free_batch_preview=free_batches[:3],
        completed_batches=FbsPickBatch.objects.select_related("agency").filter(
            status=FbsPickBatch.STATUS_DONE,
            assigned_to=request.user,
        ).order_by("-completed_at", "-id")[:10],
        received_count=FbsOrder.objects.filter(
            internal_status__in=(
                FbsOrder.STATUS_RECEIVED,
                FbsOrder.STATUS_AWAITING_STOCK,
            )
        ).count(),
        reserved_count=FbsOrder.objects.filter(
            internal_status=FbsOrder.STATUS_RESERVED
        ).count(),
        can_prepare_queue=(
            request_role in STOREKEEPER_ROLES
            or request_role == "developer"
        ),
        picker_cart=cart,
        free_batch_count=free_batch_count,
        active_wave_count=active_wave_count,
        picker_cart_busy=cart_busy,
        restock_requests=restock_requests,
        restock_request_count=len(restock_requests),
        overdue_unassembled_count=overdue_unassembled.count(),
        overdue_unassembled_oldest_hours=(
            int((now - overdue_oldest.age_started_at).total_seconds() // 3600)
            if overdue_oldest is not None
            else 0
        ),
        overdue_unassembled_clients=overdue_client_rows,
        can_claim_next=bool(
            cart
            and free_batch_count
            and not cart_busy
        ),
        error=error,
        ok_message=ok_message,
    )


def _picker_home_context(request):
    active_wave_count = FbsPickBatch.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsPickBatch.STATUS_IN_PROGRESS,
            FbsPickBatch.STATUS_VERIFICATION,
        ),
        picking_completed_at__isnull=True,
    ).count()
    free_batch_count = FbsPickBatch.objects.filter(
        status=FbsPickBatch.STATUS_QUEUED,
        assigned_to__isnull=True,
    ).count()
    restock_request_count = FbsPickRestockRequest.objects.filter(
        Q(
            status=FbsPickRestockRequest.STATUS_QUEUED,
            assigned_to__isnull=True,
        )
        | Q(
            status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
            assigned_to=request.user,
        )
    ).filter(Q(order__isnull=True) | Q(source_tote__isnull=False)).count()
    assigned_movement_count = FbsInternalMovement.objects.filter(
        assigned_to=request.user,
        status__in=(
            FbsInternalMovement.STATUS_PROPOSED,
            FbsInternalMovement.STATUS_IN_PROGRESS,
        ),
    ).count()
    rack_staging_count = FbsRackStagingBox.objects.filter(
        status=FbsRackStagingBox.STATUS_AWAITING,
    ).count()
    return _base_context(
        request,
        page_title="Главное меню",
        back_url="",
        active_wave_count=active_wave_count,
        free_batch_count=free_batch_count,
        task_menu_count=active_wave_count + free_batch_count,
        restock_request_count=restock_request_count,
        movement_task_count=assigned_movement_count + rack_staging_count,
        rack_staging_count=rack_staging_count,
    )


def _normalized_picker_location_scan(value):
    normalized = unicodedata.normalize(
        "NFKC",
        normalize_operational_location_scan(value),
    ).strip().upper()
    normalized = normalized.translate(
        str.maketrans(
            {
                "А": "A",
                "В": "B",
                "С": "C",
                "Е": "E",
                "Н": "H",
                "К": "K",
                "М": "M",
                "О": "O",
                "Р": "P",
                "Т": "T",
                "Х": "X",
                "У": "Y",
            }
        )
    )
    normalized = re.sub(r"^FBS\s*[@:/-]?\s*", "", normalized)
    return "-".join(re.findall(r"[A-Z]+|\d+", normalized))


def _normalized_picker_box_scan(value):
    normalized = str(value or "").strip()
    normalized = re.sub(r"^\][A-Z0-9]{2}", "", normalized).strip()
    return normalized.upper()


def _picker_location_candidates(cell):
    location = cell.location
    candidates = {
        str(cell.cell_code or ""),
        str(location.location_code or ""),
    }
    if all(
        int(getattr(location, field, 0) or 0) > 0
        for field in ("row_no", "section_no", "tier_no", "cell_no")
    ):
        candidates.add(
            os_location_code(
                row=location.row_no,
                section=location.section_no,
                tier=location.tier_no,
                cell=location.cell_no,
            )
        )
    return {
        _normalized_picker_location_scan(candidate)
        for candidate in candidates
        if candidate
    }


def _picker_location_label(cell):
    location = cell.location
    if all(
        int(getattr(location, field, 0) or 0) > 0
        for field in ("row_no", "section_no", "tier_no", "cell_no")
    ):
        return os_location_label(
            row=location.row_no,
            section=location.section_no,
            tier=location.tier_no,
            cell=location.cell_no,
        )
    return str(location.display_name or location.location_code or cell.cell_code)


def _picker_location_balance_queryset(cell):
    # Individual box moves do not change their logical pallet. Resolve the
    # physical place with the same warehouse-aware helper as barcode lookup.
    candidates = FbsBox.objects.select_related(
        "source_container__current_location", "pallet__cell__location",
    ).filter(
        Q(source_container__current_location_id=cell.location_id)
        | Q(pallet__cell__location_id=cell.location_id),
        status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
    )
    box_ids = [
        box.pk for box in candidates
        if getattr(fbs_box_physical_location(box), "pk", None) == cell.location_id
    ]
    return FbsStockBalance.objects.select_related("box", "agency").filter(
        box_id__in=box_ids, qty__gt=0,
    )


def _picker_product_balance_queryset(scan):
    normalized_scan = _normalized_picker_box_scan(scan)
    if not normalized_scan:
        return FbsStockBalance.objects.none()

    # A SKU may contain several sizes, each with its own barcode.
    # Alternative barcodes must retain both the owner and the size.
    variants = list(
        SKUBarcode.objects.filter(value__iexact=normalized_scan).values_list(
            "sku__agency_id", "sku_id", "size"
        ).distinct()
    )
    balances = _fbs_stock_base_queryset()
    lookup = Q(barcode__iexact=normalized_scan)
    for agency_id, sku_id, size in variants:
        lookup |= Q(agency_id=agency_id, sku_ref_id=sku_id, size__iexact=size or "")
    if variants or balances.filter(barcode__iexact=normalized_scan).exists():
        return balances.filter(lookup)
    return balances.filter(sku_code__iexact=normalized_scan)


def _picker_box_location_code(box):
    physical_location = fbs_box_physical_location(box)
    return str(
        getattr(physical_location, "location_code", "")
        or fbs_box_physical_location_code(box)
    ).strip()


def _picker_agency_short_label(agency):
    label = " ".join(
        str(
            getattr(agency, "short_name", "")
            or getattr(agency, "agn_name", "")
            or agency
            or ""
        ).split()
    )
    person_name = label
    for prefix in ("Индивидуальный предприниматель ", "ИП "):
        if label.casefold().startswith(prefix.casefold()):
            person_name = label[len(prefix):].strip()
            break
    else:
        return label

    parts = person_name.split()
    if len(parts) < 3:
        return f"ИП {person_name}".strip()
    return f"ИП {parts[0]} {parts[1][0]}. {parts[2][0]}."


def _picker_product_placements(balances):
    grouped = {}
    for balance in balances:
        key = (
            balance.agency_id,
            balance.sku_ref_id,
            str(balance.sku_code or ""),
            str(balance.name or ""),
            str(balance.size or ""),
            balance.box_id,
        )
        row = grouped.get(key)
        if row is None:
            row = SimpleNamespace(
                agency=balance.agency,
                agency_label=_picker_agency_short_label(balance.agency),
                name=str(balance.name or balance.sku_code or "Товар"),
                sku_code=str(balance.sku_code or ""),
                size=str(balance.size or ""),
                barcode=str(balance.barcode or ""),
                box=balance.box,
                pallet=balance.box.pallet,
                location_code=_picker_box_location_code(balance.box),
                location_label=fbs_box_physical_location_label(balance.box),
                qty=0,
                available_qty=0,
                reserved_qty=0,
                external_reserved_qty=0,
            )
            grouped[key] = row
        row.qty += int(balance.qty or 0)
        row.available_qty += int(balance.available_qty or 0)
        row.reserved_qty += int(balance.reserved_qty or 0)
        row.external_reserved_qty += int(balance.external_reserved_qty or 0)
    return sorted(
        grouped.values(),
        key=lambda row: (
            str(row.location_code or row.location_label).casefold(),
            str(row.box.box_code).casefold(),
            str(row.agency).casefold(),
            str(row.sku_code).casefold(),
        ),
    )


def _pick_box_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_pick_box_{allocation_id}"


def _pick_cell_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_pick_cell_{allocation_id}"


def _pop_pick_new_box_notice(request, allocation_id: int) -> dict:
    notice = request.session.pop(PICK_NEW_BOX_NOTICE_SESSION_KEY, None)
    if not isinstance(notice, dict):
        return {}
    try:
        notice_allocation_id = int(notice.get("allocation_id") or 0)
    except (TypeError, ValueError):
        return {}
    if notice_allocation_id != int(allocation_id):
        return {}
    return notice


def _pick_location_label(cell) -> str:
    location = getattr(cell, "location", None)
    if location is not None and str(location.zone_code or "").strip().upper() == "OS":
        coordinates = (
            int(location.row_no or 0),
            int(location.section_no or 0),
            int(location.tier_no or 0),
            int(location.cell_no or 0),
        )
        if all(value > 0 for value in coordinates):
            return os_location_label(
                row=coordinates[0],
                section=coordinates[1],
                tier=coordinates[2],
                cell=coordinates[3],
            )
    return str(
        (location.display_name if location else "")
        or (location.location_code if location else "")
        or cell.cell_code
    )


def _pick_box_number_suffix(box_code: str, *, length: int = 4) -> str:
    digits = "".join(re.findall(r"\d", str(box_code or "")))
    return digits[-length:] if digits else ""


def _pick_allocation_context(
    request,
    allocation,
    *,
    error="",
    ok_message="",
    scan_error="",
    new_box_notice=None,
):
    task = allocation.pick_task
    open_count = FbsOrderStockAllocation.objects.filter(
        pick_task=task,
        status__in=(
            FbsOrderStockAllocation.STATUS_RESERVED,
            FbsOrderStockAllocation.STATUS_PICKING,
        ),
    ).count()
    balance = allocation.balance
    expected_item_code = balance.barcode or allocation.order_item.barcode
    photo = balance.sku_ref.photos.order_by("sort_order", "id").first() if balance.sku_ref_id else None
    batch = task.batch
    pick_location_label = fbs_box_physical_location_label(balance.box)
    box_progress = FbsOrderStockAllocation.objects.filter(
        pick_task__batch=batch,
        balance__box_id=balance.box_id,
    ).aggregate(
        planned_qty=Sum("qty_reserved"),
        picked_qty=Sum("qty_picked"),
        order_count=Count("pick_task__order_id", distinct=True),
    )
    box_planned_qty = int(box_progress["planned_qty"] or 0)
    box_picked_qty = int(box_progress["picked_qty"] or 0)
    requires_box_scan = pick_requires_box_scan(balance.box)
    return _base_context(
        request,
        page_title=f"FBS · Заказ {task.order.external_order_id}",
        back_url=reverse("fbs:tsd_picking"),
        allocation=allocation,
        task=task,
        order=task.order,
        balance=balance,
        open_count=open_count,
        cell_confirmed=bool(request.session.get(_pick_cell_session_key(allocation.id))),
        box_confirmed=bool(request.session.get(_pick_box_session_key(allocation.id))),
        requires_box_scan=requires_box_scan,
        requires_cell_scan=not requires_box_scan,
        allocation_remaining=max(
            int(allocation.qty_reserved or 0) - int(allocation.qty_picked or 0), 0
        ),
        missing_group_qty=pick_missing_group_quantity(
            allocation_id=allocation.id,
            assigned_to=request.user,
        ),
        position_picked_qty=int(allocation.qty_picked or 0),
        source_box_code=balance.box.box_code if requires_box_scan else "",
        box_planned_qty=box_planned_qty,
        box_picked_qty=box_picked_qty,
        box_remaining_qty=max(box_planned_qty - box_picked_qty, 0),
        box_order_count=int(box_progress["order_count"] or 0),
        wave_picked_qty=int(batch.picked_qty or 0),
        wave_planned_qty=int(batch.planned_qty or 0),
        recent_scan_events=FbsPickScanEvent.objects.filter(batch=batch)
        .select_related("created_by")
        .order_by("-created_at", "-id")[:12],
        expected_item_code=expected_item_code,
        pick_location_label=pick_location_label,
        item_scan_label="Штрихкод товара",
        product_photo_url=photo.url if photo else "",
        scan_error=scan_error,
        new_box_notice=new_box_notice or {},
        error=error,
        ok_message=ok_message,
    )


def _next_verification_allocation(batch):
    allocations = (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            pick_task__status=FbsPickTask.STATUS_PICKED,
            status=FbsOrderStockAllocation.STATUS_PICKED,
        )
        .exclude(
            pick_task__order__pick_restock_requests__status__in=(
                SEPARATED_PICK_RESTOCK_STATUSES
            )
        )
        .filter(
            Q(verification_progress__isnull=True)
            | Q(verification_progress__qty_verified__lt=F("qty_picked"))
        )
        .order_by("pick_task__sort_order", "pick_task_id", "id")
    )
    active_bundle_task = active_partial_ozon_verification_task(batch)
    if active_bundle_task is not None:
        allocations = allocations.filter(pick_task=active_bundle_task)
    return allocations.first()


def _next_verification_label(batch):
    tasks = list(
        batch.tasks.filter(status=FbsPickTask.STATUS_PICKED)
        .exclude(
            order__pick_restock_requests__status__in=SEPARATED_PICK_RESTOCK_STATUSES
        )
        .annotate(
            ui_picked_qty=Sum("allocations__qty_picked"),
            ui_verified_qty=Sum("allocations__verification_progress__qty_verified"),
        )
        .order_by("sort_order", "id")
    )
    if not tasks:
        return None
    active_pick_context = (
        FbsControllerPickTote.objects.filter(
            pick_batch=batch,
            status=FbsControllerPickTote.STATUS_PROCESSING,
        )
        .order_by("id")
        .first()
    )
    current_tote_order_ids = set()
    if active_pick_context is not None:
        current_tote_order_ids = set(
            active_pick_context.orders.exclude(
                status=FbsControllerToteOrder.STATUS_REMOVED,
            ).values_list("order_id", flat=True)
        )
    scanned_label_order_ids = FbsPickScanEvent.objects.filter(
        batch=batch,
        task__isnull=False,
        stage=FbsPickScanEvent.STAGE_ORDER_LABEL,
        result=FbsPickScanEvent.RESULT_SUCCESS,
    )
    if batch.verification_started_at is not None:
        scanned_label_order_ids = scanned_label_order_ids.filter(
            created_at__gte=batch.verification_started_at,
        )
    scanned_label_order_ids = set(
        scanned_label_order_ids.values_list("task__order_id", flat=True)
    )
    latest_labels = {}
    for label in (
        FbsOrderLabel.objects.filter(
            order_id__in=[task.order_id for task in tasks]
        )
        .exclude(status=FbsOrderLabel.STATUS_CANCELED)
        .order_by("order_id", "-requested_at", "-id")
    ):
        latest_labels.setdefault(label.order_id, label)
    for task in tasks:
        if task.order_id in scanned_label_order_ids:
            continue
        picked_qty = int(task.ui_picked_qty or 0)
        verified_qty = int(task.ui_verified_qty or 0)
        label = latest_labels.get(task.order_id)
        if (
            picked_qty
            and verified_qty >= picked_qty
            and label is not None
            and (
                active_pick_context is None
                or task.order_id not in current_tote_order_ids
            )
        ):
            return label
    return None


def _next_incomplete_verification_allocation(batch):
    return (
        _pick_allocation_queryset()
        .filter(
            pick_task__batch=batch,
            pick_task__status=FbsPickTask.STATUS_PICKED,
            status=FbsOrderStockAllocation.STATUS_PICKED,
            qty_picked__gt=0,
            pick_task__exceptions__exception_type=FbsPickException.TYPE_NOT_FOUND,
            pick_task__exceptions__status=FbsPickException.STATUS_OPEN,
        )
        .exclude(
            pick_task__order__pick_restock_requests__status__in=(
                SEPARATED_PICK_RESTOCK_STATUSES
            )
        )
        .order_by("pick_task__sort_order", "pick_task_id", "id")
        .first()
    )


def _open_pick_shortage_issue(task):
    return (
        FbsPickException.objects.filter(
            task=task,
            exception_type=FbsPickException.TYPE_NOT_FOUND,
            status=FbsPickException.STATUS_OPEN,
        )
        .order_by("id")
        .first()
    )


def _verification_order_allocations(task, *, current_allocation_id=None):
    rows = list(
        _pick_allocation_queryset()
        .filter(pick_task=task)
        .order_by("id")
    )
    for row in rows:
        row_progress = getattr(row, "verification_progress", None)
        row.verified_qty = int(getattr(row_progress, "qty_verified", 0) or 0)
        row.is_verified = row.verified_qty >= int(row.qty_picked or 0)
        row.is_current = row.id == current_allocation_id
    return rows


def _verification_order_unit_rows(rows, *, current_allocation_id=None):
    """Expand allocation progress into a stable, operator-facing unit checklist."""
    units = []
    for row in rows:
        planned_qty = max(int(row.qty_picked or 0), 0)
        verified_qty = min(
            max(int(getattr(row, "verified_qty", 0) or 0), 0),
            planned_qty,
        )
        item = row.order_item
        article = str(item.external_sku or row.balance.sku_code or "").strip()
        name = str(item.product_name or row.balance.name or "Товар").strip()
        barcode = str(item.barcode or row.balance.barcode or "").strip()
        for unit_number in range(1, planned_qty + 1):
            is_verified = unit_number <= verified_qty
            units.append(
                SimpleNamespace(
                    position=len(units) + 1,
                    total=0,
                    allocation_id=row.id,
                    article=article,
                    name=name,
                    barcode=barcode,
                    unit_number=unit_number,
                    planned_item_qty=planned_qty,
                    is_verified=is_verified,
                    is_current_scan=bool(
                        is_verified
                        and row.id == current_allocation_id
                        and unit_number == verified_qty
                    ),
                    is_next=False,
                )
            )
    total = len(units)
    next_pending_marked = False
    for unit in units:
        unit.total = total
        if not unit.is_verified and not next_pending_marked:
            unit.is_next = True
            next_pending_marked = True
    return units


def _ozon_multi_item_hold_prompt(allocation):
    """Describe the remaining composition after one accepted Ozon unit scan."""
    task = allocation.pick_task
    if (
        task is None
        or task.order.profile.marketplace
        != FbsIntegrationProfile.MARKETPLACE_OZON
    ):
        return None
    rows = _verification_order_allocations(
        task,
        current_allocation_id=allocation.id,
    )
    unit_rows = _verification_order_unit_rows(
        rows,
        current_allocation_id=allocation.id,
    )
    planned_qty = sum(int(row.qty_picked or 0) for row in rows)
    verified_qty = sum(int(row.verified_qty or 0) for row in rows)
    remaining_qty = max(planned_qty - verified_qty, 0)
    if planned_qty <= 1 or remaining_qty <= 0:
        return None

    remaining_by_item = {}
    for row in rows:
        remaining = max(int(row.qty_picked or 0) - int(row.verified_qty or 0), 0)
        if not remaining:
            continue
        item = row.order_item
        item_summary = remaining_by_item.setdefault(
            item.id,
            {
                "article": str(item.external_sku or row.balance.sku_code or "").strip(),
                "name": str(item.product_name or row.balance.name or "Товар").strip(),
                "barcode": str(item.barcode or row.balance.barcode or "").strip(),
                "quantity": 0,
            },
        )
        item_summary["quantity"] += remaining

    scanned_item = allocation.order_item
    next_unit = next((unit for unit in unit_rows if unit.is_next), None)
    return SimpleNamespace(
        order_number=task.order.external_order_id,
        planned_qty=planned_qty,
        verified_qty=verified_qty,
        remaining_qty=remaining_qty,
        scanned_article=str(
            scanned_item.external_sku or allocation.balance.sku_code or ""
        ).strip(),
        scanned_name=str(
            scanned_item.product_name or allocation.balance.name or "Товар"
        ).strip(),
        scanned_barcode=str(
            scanned_item.barcode or allocation.balance.barcode or ""
        ).strip(),
        next_position=(next_unit.position if next_unit is not None else verified_qty + 1),
        unit_rows=unit_rows,
        remaining_items=[SimpleNamespace(**item) for item in remaining_by_item.values()],
    )


def _verification_wave_rows(batch, *, current_task_id=None):
    tasks = list(
        batch.tasks.filter(status=FbsPickTask.STATUS_PICKED)
        .exclude(
            order__pick_restock_requests__status__in=SEPARATED_PICK_RESTOCK_STATUSES
        )
        .select_related("order")
        .annotate(
            ui_verified_qty=Sum("allocations__verification_progress__qty_verified")
        )
        .order_by("sort_order", "id")
    )
    order_ids = [task.order_id for task in tasks]
    active_pick_context = (
        FbsControllerPickTote.objects.filter(
            pick_batch=batch,
            status=FbsControllerPickTote.STATUS_PROCESSING,
        )
        .order_by("id")
        .first()
    )
    current_tote_order_ids = set()
    if active_pick_context is not None:
        current_tote_order_ids = set(
            active_pick_context.orders.exclude(
                status=FbsControllerToteOrder.STATUS_REMOVED,
            ).values_list("order_id", flat=True)
        )
    shortage_issues = {}
    for issue in FbsPickException.objects.filter(
        task_id__in=[task.id for task in tasks],
        exception_type=FbsPickException.TYPE_NOT_FOUND,
        status=FbsPickException.STATUS_OPEN,
    ).order_by("task_id", "id"):
        shortage_issues.setdefault(issue.task_id, issue)
    task_barcodes = defaultdict(list)
    for task_id, balance_barcode, item_barcode in (
        FbsOrderStockAllocation.objects.filter(
            pick_task_id__in=[task.id for task in tasks],
            qty_picked__gt=0,
        )
        .values_list("pick_task_id", "balance__barcode", "order_item__barcode")
        .order_by("pick_task_id", "id")
    ):
        barcode = str(balance_barcode or item_barcode or "").strip()
        if barcode and barcode not in task_barcodes[task_id]:
            task_barcodes[task_id].append(barcode)
    latest_labels = {}
    for label in (
        FbsOrderLabel.objects.filter(order_id__in=order_ids)
        .exclude(status=FbsOrderLabel.STATUS_CANCELED)
        .order_by("order_id", "-requested_at", "-id")
    ):
        latest_labels.setdefault(label.order_id, label)
    handover_links = {
        link.order_id: link
        for link in FbsHandoverOrder.objects.filter(
            order_id__in=order_ids,
            status=FbsHandoverOrder.STATUS_ACTIVE,
        )
        .select_related("box")
        .order_by("id")
    }
    handover_assignments = {
        assignment.order_id: assignment
        for assignment in FbsHandoverOrderAssignment.objects.filter(
            order_id__in=order_ids,
        )
        .exclude(status=FbsHandoverOrderAssignment.STATUS_CANCELED)
        .select_related("batch")
        .order_by("id")
    }
    target_handover_batch = (
        FbsHandoverBatch.objects.filter(
            controller_check_tote__pick_totes__pick_batch=batch,
        )
        .order_by("id")
        .first()
    )

    rows = []
    for task in tasks:
        shortage_issue = shortage_issues.get(task.id)
        is_incomplete = shortage_issue is not None
        verified_qty = int(task.ui_verified_qty or 0)
        planned_qty = int(task.planned_qty or 0)
        label = latest_labels.get(task.order_id)
        handover_link = handover_links.get(task.order_id)
        handover_assignment = handover_assignments.get(task.order_id)
        row_handover_batch = (
            handover_assignment.batch
            if handover_assignment is not None
            else target_handover_batch
        )
        product_complete = bool(planned_qty and verified_qty >= planned_qty)
        label_complete = bool(
            label is not None and label.status == FbsOrderLabel.STATUS_APPLIED
            and (
                active_pick_context is None
                or task.order_id in current_tote_order_ids
            )
        )
        handover_added = handover_assignment is not None
        if is_incomplete:
            state = "problem"
            state_label = "Недокомплект · передать в проблемные"
        elif handover_link is not None:
            state = "done"
            state_label = f"В коробе {handover_link.box.qr_code}"
        elif (
            task.order_id in current_tote_order_ids
            and label is not None
            and label.status == FbsOrderLabel.STATUS_REQUESTED
        ):
            state = "in_progress"
            state_label = "Заказ закреплён · Ozon проверяет"
        elif (
            task.order_id in current_tote_order_ids
            and label is not None
            and label.status == FbsOrderLabel.STATUS_READY
        ):
            state = "in_progress"
            state_label = "Этикетка Ozon готова · ожидает упаковки"
        elif label_complete and handover_assignment is not None:
            state = "done"
            state_label = f"В отгрузке #{handover_assignment.batch_id} · короб ожидается"
        elif label_complete:
            state = "in_progress"
            state_label = "QR отсканирован · отгрузка создается"
        elif label is not None and label.status in {
            FbsOrderLabel.STATUS_READY,
            FbsOrderLabel.STATUS_APPLIED,
        }:
            state = "in_progress"
            state_label = "QR готов, требуется скан"
        elif label is not None and label.status == FbsOrderLabel.STATUS_REQUESTED:
            state = "in_progress"
            state_label = "Готовится QR заказа"
        elif label is not None and label.status == FbsOrderLabel.STATUS_ERROR:
            state = "problem"
            state_label = "Ошибка этикетки"
        elif product_complete:
            state = "in_progress"
            state_label = "QR заказа не создан"
        elif verified_qty:
            state = "in_progress"
            state_label = f"Осталось проверить {max(planned_qty - verified_qty, 0)} шт."
        else:
            state = "pending"
            state_label = f"Остался товар {planned_qty} шт."
        rows.append(
            SimpleNamespace(
                task=task,
                order=task.order,
                verified_qty=verified_qty,
                planned_qty=planned_qty,
                label=label,
                handover_link=handover_link,
                handover_assignment=handover_assignment,
                target_handover_batch=row_handover_batch,
                is_product_complete=product_complete,
                is_label_complete=label_complete,
                is_handover_added=handover_added,
                is_verification_complete=bool(
                    product_complete
                    and label_complete
                    and handover_added
                    and not is_incomplete
                ),
                is_incomplete=is_incomplete,
                problem_reason=(
                    str(task.order.problem_reason or "").strip()
                    if is_incomplete
                    else ""
                ),
                state=state,
                state_label=state_label,
                is_current=task.id == current_task_id,
                item_barcodes=tuple(task_barcodes.get(task.id, ())),
            )
        )
    return rows


def _verification_wave_context(batch, *, current_task_id=None):
    rows = _verification_wave_rows(batch, current_task_id=current_task_id)
    pending_rows = [row for row in rows if not row.is_verification_complete]
    handover_batches = []
    seen_handover_ids = set()
    for row in rows:
        handover_batch = row.target_handover_batch
        if handover_batch is None or handover_batch.id in seen_handover_ids:
            continue
        seen_handover_ids.add(handover_batch.id)
        handover_batches.append(handover_batch)
    return {
        "verification_wave_rows": rows,
        "verification_pending_rows": pending_rows,
        "verification_pending_count": len(pending_rows),
        "verification_handover_batches": handover_batches,
        "verification_primary_handover": handover_batches[0] if handover_batches else None,
    }


def _verification_success_events(batch):
    return (
        FbsPickScanEvent.objects.filter(
            batch=batch,
            stage=FbsPickScanEvent.STAGE_VERIFY_ITEM,
            result=FbsPickScanEvent.RESULT_SUCCESS,
        )
        .select_related("task__order", "created_by")
        .order_by("-created_at", "-id")[:30]
    )


def _marketplace_order_is_canceled(order) -> bool:
    return order_is_client_canceled_by_marketplace(order)


_CANCELED_INITIAL_SCAN_SALT = "fbs.canceled-order-initial-product-scan.v1"
_CANCELED_INITIAL_SCAN_MAX_AGE_SECONDS = 15 * 60


def _canceled_initial_scan_token(*, batch, allocation, controller, item_scan: str) -> str:
    clean_scan = str(item_scan or "").strip()
    if not clean_scan:
        return ""
    from .services.picking import resolve_verification_item_scan

    matched_barcode, marking_scan = resolve_verification_item_scan(allocation, clean_scan)
    if not matched_barcode:
        return ""
    return signing.dumps(
        {
            "batch_id": int(batch.id),
            "allocation_id": int(allocation.id),
            "controller_id": int(controller.id),
            "item_scan": clean_scan,
        },
        salt=_CANCELED_INITIAL_SCAN_SALT,
        compress=True,
    )


def _canceled_initial_scan_from_token(*, token: str, batch, allocation, controller) -> str:
    try:
        payload = signing.loads(
            str(token or "").strip(),
            salt=_CANCELED_INITIAL_SCAN_SALT,
            max_age=_CANCELED_INITIAL_SCAN_MAX_AGE_SECONDS,
        )
    except signing.BadSignature as exc:
        raise ValueError(
            "Подтверждение первого товара устарело. Отсканируйте товары заказа повторно."
        ) from exc
    if not isinstance(payload, dict):
        raise ValueError(
            "Подтверждение товара повреждено. Отсканируйте товары повторно."
        )
    expected = (
        int(batch.id),
        int(allocation.id),
        int(controller.id),
    )
    actual = (
        int(payload.get("batch_id") or 0),
        int(payload.get("allocation_id") or 0),
        int(payload.get("controller_id") or 0),
    )
    item_scan = str(payload.get("item_scan") or "").strip()
    if actual != expected or not item_scan:
        raise ValueError(
            "Подтверждение товара относится к другой проверке. Отсканируйте товары повторно."
        )
    return item_scan


def _canceled_product_scan_rows(*, task, confirmed_allocation_id=None):
    rows = []
    confirmed_skipped = False
    allocations = (
        FbsOrderStockAllocation.objects.select_related("balance", "order_item")
        .filter(
            pick_task=task,
            status=FbsOrderStockAllocation.STATUS_PICKED,
            qty_picked__gt=0,
        )
        .order_by("id")
    )
    for row in allocations:
        quantity = int(row.qty_picked or 0)
        for unit_no in range(1, quantity + 1):
            if row.id == confirmed_allocation_id and not confirmed_skipped:
                confirmed_skipped = True
                continue
            rows.append(
                SimpleNamespace(
                    allocation_id=row.id,
                    unit_no=unit_no,
                    quantity=quantity,
                    sku_code=str(row.balance.sku_code or "").strip(),
                    name=str(
                        row.balance.name or row.order_item.product_name or "Товар"
                    ).strip(),
                    barcode=str(row.balance.barcode or row.order_item.barcode or "").strip(),
                    scan_label="ШК",
                )
            )
    return rows


def _label_requires_problem_tote(label) -> bool:
    if label.status == FbsOrderLabel.STATUS_ERROR:
        return True
    if label.status != FbsOrderLabel.STATUS_REQUESTED:
        return False
    profile = label.order.profile
    return not profile.is_active or not profile.outbox_enabled


def _verification_service_tote_prompt(
    *,
    batch,
    allocation,
    controller,
    problem_kind: str,
    reason: str,
    restock_request=None,
    initial_item_scan: str = "",
):
    pick_tote_context = (
        FbsControllerPickTote.objects.select_related(
            "session__problem_tote",
            "session__canceled_tote",
        )
        .filter(
            pick_batch=batch,
            session__controller=controller,
            session__status="active",
            status__in=(
                FbsControllerPickTote.STATUS_PROCESSING,
                FbsControllerPickTote.STATUS_AWAITING_EMPTY,
            ),
        )
        .order_by("-id")
        .first()
    )
    session = pick_tote_context.session if pick_tote_context else None
    is_canceled = problem_kind == "canceled"
    is_marketplace_rejected = problem_kind == "marketplace_rejected"
    uses_canceled_tote = is_canceled or is_marketplace_rejected
    service_tote = (
        session.canceled_tote if uses_canceled_tote and session else None
    ) or (session.problem_tote if not uses_canceled_tote and session else None)
    if is_canceled:
        title = "Заказ отменен"
    elif is_marketplace_rejected:
        title = "Маркетплейс отказал в приеме заказа"
    elif problem_kind == "label":
        title = "Этикетка заказа не создана"
    else:
        title = "Честный знак не принят"
    if is_canceled:
        exception_type = FbsPickException.TYPE_NOT_FOUND
    elif is_marketplace_rejected:
        exception_type = FbsPickException.TYPE_OTHER
    elif problem_kind == "label":
        exception_type = FbsPickException.TYPE_OTHER
    else:
        exception_type = FbsPickException.TYPE_BARCODE
    marketplace_check_pending = bool(
        is_marketplace_rejected
        and restock_request is not None
        and restock_request.status
        in {
            FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
            FbsPickRestockRequest.STATUS_FAILED,
        }
    )
    applied_label = (
        FbsOrderLabel.objects.filter(
            order_id=allocation.pick_task.order_id,
            status=FbsOrderLabel.STATUS_APPLIED,
        )
        .order_by("-applied_at", "-id")
        .first()
        if is_canceled
        else None
    )
    late_applied_label = applied_label is not None
    initial_scan_token = (
        _canceled_initial_scan_token(
            batch=batch,
            allocation=allocation,
            controller=controller,
            item_scan=initial_item_scan,
        )
        if is_canceled and not late_applied_label and initial_item_scan
        else ""
    )
    canceled_product_scan_rows = (
        _canceled_product_scan_rows(
            task=allocation.pick_task,
            confirmed_allocation_id=(allocation.id if initial_scan_token else None),
        )
        if is_canceled and not late_applied_label
        else []
    )
    canceled_product_total = (
        len(canceled_product_scan_rows) + (1 if initial_scan_token else 0)
    )
    return SimpleNamespace(
        kind=problem_kind,
        title=title,
        reason=str(reason or "").strip(),
        restock_request=restock_request,
        ready_to_scan=bool(
            not late_applied_label
            and (
                restock_request is None
                or restock_request.status == FbsPickRestockRequest.STATUS_QUEUED
                or marketplace_check_pending
            )
        ),
        marketplace_check_pending=marketplace_check_pending,
        marketplace_check_failed=bool(
            restock_request is not None
            and restock_request.status == FbsPickRestockRequest.STATUS_FAILED
        ),
        requires_order_scan=False,
        requires_product_scans=is_canceled and not late_applied_label,
        late_applied_label=late_applied_label,
        applied_label_barcode=(str(applied_label.barcode or "").strip() if applied_label else ""),
        uses_canceled_tote=uses_canceled_tote,
        order_scan_code=str(allocation.pick_task.order.external_order_id or "").strip(),
        initial_product_confirmed=bool(initial_scan_token),
        initial_product_scan_token=initial_scan_token,
        canceled_product_scan_rows=canceled_product_scan_rows,
        canceled_product_total=canceled_product_total,
        canceled_product_confirmed=(1 if initial_scan_token else 0),
        service_tote=service_tote,
        service_tote_label=(
            "тара отмененных заказов" if uses_canceled_tote else "проблемная тара"
        ),
        exception_type=exception_type,
        action_url=reverse(
            "fbs:tsd_pick_verification",
            kwargs={"batch_id": batch.id},
        ) if is_marketplace_rejected else reverse(
            "fbs:tsd_pick_verification_problem",
            kwargs={"allocation_id": allocation.id},
        ),
    )


def _pick_verification_base_template(request) -> str:
    if (
        get_request_role(request) == "fbs_controller"
        and request.headers.get("X-Requested-With") == "fetch"
    ):
        return "fbs/tsd_verification_fragment_base.html"
    return "fbs/tsd_base.html"


def _pick_verification_context(
    request,
    batch,
    allocation,
    *,
    error="",
    ok_message="",
    submitted_item_scan="",
    submitted_marking_scan="",
    retry_marking=False,
    item_selected=False,
    used_marking_info=None,
    unknown_scan_info=None,
    service_tote_prompt=None,
    multi_item_hold_prompt=None,
):
    task = allocation.pick_task
    progress = getattr(allocation, "verification_progress", None)
    verified_qty = int(getattr(progress, "qty_verified", 0) or 0)
    wave_verified_qty = int(
        FbsPickVerificationProgress.objects.filter(
            allocation__pick_task__batch=batch
        ).aggregate(total=Sum("qty_verified"))["total"]
        or 0
    )
    balance = allocation.balance
    from .services.traceability import (
        controller_legacy_marking_exception_allowed,
        controller_marking_scan_required,
        metadata_requirements,
    )

    metadata = metadata_requirements(allocation.order_item)
    trace_expiry = allocation.traceability.expiry_date
    raw_min_expiry = (
        allocation.order_item.requirements.get("min_expiry_date")
        if isinstance(allocation.order_item.requirements, dict)
        else None
    )
    try:
        minimum_expiry = date.fromisoformat(str(raw_min_expiry)) if raw_min_expiry else None
    except ValueError:
        minimum_expiry = None
    minimum_expiry = max(
        value for value in (timezone.localdate(), minimum_expiry) if value is not None
    )
    photo = balance.sku_ref.photos.order_by("sort_order", "id").first() if balance.sku_ref_id else None
    order_allocations = _verification_order_allocations(
        task,
        current_allocation_id=allocation.id,
    )
    order_verified_qty = sum(row.verified_qty for row in order_allocations)
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if (
        get_request_role(request) == "fbs_controller"
        and batch.workstation_id
        and batch.workstation.device_agent_id
    ):
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    verification_wave_context = _verification_wave_context(
        batch,
        current_task_id=task.id,
    )
    current_wave_row = next(
        (
            row
            for row in verification_wave_context["verification_wave_rows"]
            if row.task.id == task.id
        ),
        None,
    )
    order_incomplete = (
        bool(current_wave_row.is_incomplete)
        if current_wave_row is not None
        else _open_pick_shortage_issue(task) is not None
    )
    marking_required = controller_marking_scan_required(allocation)
    legacy_marking_exception_allowed = (
        controller_legacy_marking_exception_allowed(allocation)
    )
    marking_optional = bool(
        not marking_required and wb_optional_marking_available(allocation.order_item)
    )
    active_bundle_task = active_partial_ozon_verification_task(batch)
    active_multi_item_order = (
        active_bundle_task.order
        if active_bundle_task is not None and active_bundle_task.id == task.id
        else None
    )
    return _base_context(
        request,
        verification_base_template=_pick_verification_base_template(request),
        page_title=f"FBS · Проверка волны #{batch.id}",
        back_url=(
            reverse("fbs:controller_home")
            if get_request_role(request) == "fbs_controller"
            else reverse("fbs:operator_wave_detail", kwargs={"batch_id": batch.id})
        ),
        batch=batch,
        task=task,
        order=task.order,
        order_incomplete=order_incomplete,
        order_problem_reason=(
            str(task.order.problem_reason or "").strip()
            if order_incomplete
            else ""
        ),
        allocation=allocation,
        balance=balance,
        verified_qty=verified_qty,
        allocation_remaining=max(int(allocation.qty_picked or 0) - verified_qty, 0),
        wave_verified_qty=wave_verified_qty,
        wave_planned_qty=int(batch.planned_qty or 0),
        order_verified_qty=order_verified_qty,
        order_planned_qty=int(task.planned_qty or 0),
        order_allocations=order_allocations,
        **verification_wave_context,
        verified_product_events=_verification_success_events(batch),
        expected_item_code=balance.barcode or allocation.order_item.barcode,
        expected_marking_code=balance.marking_code,
        marking_required=marking_required,
        legacy_marking_exception_allowed=legacy_marking_exception_allowed,
        marking_optional=marking_optional,
        submitted_item_scan=submitted_item_scan,
        submitted_marking_scan=submitted_marking_scan,
        retry_marking=retry_marking,
        item_selected=item_selected,
        used_marking_info=used_marking_info,
        unknown_scan_info=unknown_scan_info,
        service_tote_prompt=service_tote_prompt,
        multi_item_hold_prompt=multi_item_hold_prompt,
        active_multi_item_order=active_multi_item_order,
        expiry_required=metadata.expiry_required,
        expiry_entry_required=metadata.expiry_required and trace_expiry is None,
        trace_expiry_date=trace_expiry,
        minimum_expiry_date=minimum_expiry,
        item_scan_label="Штрихкод товара",
        agent_scan_poll_url=agent_scan_poll_url,
        agent_scan_event_id=agent_scan_event_id,
        product_photo_url=photo.url if photo else "",
        recent_scan_events=FbsPickScanEvent.objects.filter(batch=batch)
        .select_related("created_by", "task__order")
        .order_by("-created_at", "-id")[:12],
        error=error,
        ok_message=ok_message,
    )


def _pick_verification_label_context(
    request,
    batch,
    label,
    *,
    error="",
    ok_message="",
):
    context = _label_context(request, label, error=error, ok_message=ok_message)
    context["verification_base_template"] = _pick_verification_base_template(
        request
    )
    task = (
        batch.tasks.filter(
            order_id=label.order_id,
            status=FbsPickTask.STATUS_PICKED,
        )
        .order_by("sort_order", "id")
        .first()
    )
    if task is None:
        return context
    shortage_issue = _open_pick_shortage_issue(task)
    order_allocations = _verification_order_allocations(task)
    order_unit_rows = _verification_order_unit_rows(order_allocations)
    workstation = (
        FbsWorkstation.objects.select_related(
            "active_handover_box__batch__profile__agency"
        )
        .filter(pk=batch.workstation_id, is_active=True)
        .first()
        if batch.workstation_id
        else None
    )
    assignment = (
        FbsHandoverOrderAssignment.objects.select_related("batch__profile__agency")
        .filter(order_id=label.order_id)
        .first()
    )
    active_box = workstation.active_handover_box if workstation is not None else None
    active_box_compatible = bool(
        active_box is not None
        and assignment is not None
        and active_box.batch_id == assignment.batch_id
        and active_box.status == FbsHandoverBox.STATUS_OPEN
        and active_box.batch.status == FbsHandoverBatch.STATUS_OPEN
    )
    packing_box = (
        active_box
        if active_box_compatible
        else get_ready_handover_box_for_order(order_id=label.order_id)
    )
    controller_pick_tote = controller_pick_tote_for_batch(batch_id=batch.id)
    service_tote_prompt = None
    problem_allocation = next(
        (row for row in order_allocations if int(row.qty_picked or 0) > 0),
        None,
    )
    if (
        get_request_role(request) == "fbs_controller"
        and problem_allocation is not None
        and _label_requires_problem_tote(label)
    ):
        service_tote_prompt = _verification_service_tote_prompt(
            batch=batch,
            allocation=problem_allocation,
            controller=request.user,
            problem_kind="label",
            reason=(
                str(label.error or "").strip()
                or "Этикетку заказа не удалось получить или напечатать."
            ),
        )
    verification_wave_context = _verification_wave_context(
        batch,
        current_task_id=task.id,
    )
    context.update(
        {
            "page_title": f"FBS · Проверка волны #{batch.id}",
            "back_url": reverse("fbs:controller_home"),
            "verification_label_phase": True,
            "preloaded_ozon_order_label": is_preloaded_ozon_order_label(label),
            "task": task,
            "order_incomplete": shortage_issue is not None,
            "order_problem_reason": (
                str(task.order.problem_reason or "").strip()
                if shortage_issue is not None
                else ""
            ),
            "order_allocations": order_allocations,
            "order_unit_rows": order_unit_rows,
            "multi_item_order": len(order_unit_rows) > 1,
            **verification_wave_context,
            "verified_product_events": _verification_success_events(batch),
            "order_verified_qty": sum(row.verified_qty for row in order_allocations),
            "order_planned_qty": int(task.planned_qty or 0),
            "wave_verified_qty": int(
                FbsPickVerificationProgress.objects.filter(
                    allocation__pick_task__batch=batch,
                    allocation__pick_task__status=FbsPickTask.STATUS_PICKED,
                ).aggregate(total=Sum("qty_verified"))["total"]
                or 0
            ),
            "wave_planned_qty": int(
                batch.tasks.filter(status=FbsPickTask.STATUS_PICKED).aggregate(
                    total=Sum("planned_qty")
                )["total"]
                or 0
            ),
            "verification_workstation": workstation,
            "handover_assignment": assignment,
            "handover_assignment_confirmed": bool(
                assignment is not None
                and assignment.status
                == FbsHandoverOrderAssignment.STATUS_CONFIRMED
            ),
            "active_handover_box": active_box,
            "active_box_compatible": active_box_compatible,
            "packing_handover_box": packing_box,
            "automatic_box_selection": bool(
                packing_box is not None and not active_box_compatible
            ),
            "controller_pick_tote": controller_pick_tote,
            "controller_check_tote": (
                controller_pick_tote.check_tote
                if controller_pick_tote is not None
                else None
            ),
            "service_tote_prompt": service_tote_prompt,
        }
    )
    return context


def _label_queryset():
    return FbsOrderLabel.objects.select_related(
        "order__profile__agency",
        "requested_by",
        "applied_by",
    ).order_by("requested_at", "id")


def _accessible_label_queryset(request):
    labels = _label_queryset()
    if get_request_role(request) == "fbs_controller":
        labels = labels.filter(
            order__pick_tasks__batch__verification_assigned_to=request.user
        ).distinct()
    return labels


def _labels_context(request, *, error="", ok_message=""):
    labels = _accessible_label_queryset(request)
    return _base_context(
        request,
        page_title="FBS · Этикетки заказов",
        back_url=(
            reverse("fbs:controller_home")
            if get_request_role(request) == "fbs_controller"
            else reverse("fbs:tsd_picking")
        ),
        requested_labels=labels.filter(status=FbsOrderLabel.STATUS_REQUESTED)[:50],
        ready_labels=labels.filter(status=FbsOrderLabel.STATUS_READY)[:50],
        applied_labels=labels.filter(status=FbsOrderLabel.STATUS_APPLIED).order_by(
            "-applied_at", "-id"
        )[:10],
        error=error,
        ok_message=ok_message,
    )


def _label_context(request, label, *, error="", ok_message=""):
    batch = (
        FbsPickBatch.objects.filter(
            tasks__order=label.order,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .select_related("workstation__device_agent")
        .order_by("created_at", "id")
        .first()
    )
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if (
        get_request_role(request) == "fbs_controller"
        and batch is not None
        and batch.workstation_id
        and batch.workstation.device_agent_id
    ):
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    preloaded_ozon_order_label = is_preloaded_ozon_order_label(label)
    print_prefix = (
        f"fbs:ozon-order-qr:{label.id}:"
        if preloaded_ozon_order_label
        else f"fbs:order-label:{label.id}:"
    )
    print_jobs = ProcessingPrintJob.objects.filter(card_id__startswith=print_prefix)
    if batch is not None and batch.verification_started_at is not None:
        print_jobs = print_jobs.filter(
            created_at__gte=batch.verification_started_at,
        )
    if get_request_role(request) == "fbs_controller":
        recover_stale_fbs_order_label_print_job(
            label_id=label.id,
            not_before=(
                batch.verification_started_at
                if batch is not None
                else None
            ),
        )
    print_job = print_jobs.order_by("-created_at", "-id").first()
    print_status_labels = {
        ProcessingPrintJob.STATUS_PENDING: "В очереди",
        ProcessingPrintJob.STATUS_PRINTING: "Печатается",
        ProcessingPrintJob.STATUS_PRINTED: "Напечатана",
        ProcessingPrintJob.STATUS_FAILED: "Ошибка печати",
    }
    print_job_delay_warning = ""
    print_job_retry_ready = False
    if print_job is not None and print_job.status in {
        ProcessingPrintJob.STATUS_PENDING,
        ProcessingPrintJob.STATUS_PRINTING,
    }:
        job_age_seconds = max(
            int((timezone.now() - print_job.updated_at).total_seconds()),
            0,
        )
        lease_seconds = fbs_desktop_print_lease_seconds()
        if (
            print_job.status == ProcessingPrintJob.STATUS_PRINTING
            and job_age_seconds >= lease_seconds
        ):
            print_job_retry_ready = True
            print_job_delay_warning = (
                "Fullbox Desktop не подтвердил печать за установленный срок. "
                "Задание доступно для безопасного повторного запуска."
            )
        elif job_age_seconds >= 20:
            if print_job.status == ProcessingPrintJob.STATUS_PENDING:
                print_job_delay_warning = (
                    "Задание больше 20 секунд ожидает Fullbox Desktop. "
                    "Проверьте, что Desktop запущен и выбран правильный принтер."
                )
            else:
                print_job_delay_warning = (
                    "Fullbox Desktop забрал задание, но ещё не подтвердил печать. "
                    "Проверьте принтер и не запускайте вторую копию до окончания ожидания."
                )
    return _base_context(
        request,
        page_title=f"FBS · Этикетка {label.order.external_order_id}",
        back_url=(
            reverse("fbs:tsd_pick_verification", kwargs={"batch_id": batch.id})
            if batch
            else reverse("fbs:tsd_labels")
        ),
        label=label,
        order=label.order,
        batch=batch,
        print_job=print_job,
        print_status_label=(print_status_labels.get(print_job.status, "") if print_job else ""),
        print_job_delay_warning=print_job_delay_warning,
        print_job_retry_ready=print_job_retry_ready,
        preloaded_ozon_order_label=preloaded_ozon_order_label,
        agent_scan_poll_url=agent_scan_poll_url,
        agent_scan_event_id=agent_scan_event_id,
        error=error,
        ok_message=ok_message,
    )


def _source_session_key(allocation_id: int) -> str:
    return f"fbs_tsd_source_scan_{allocation_id}"


def _query_int(request, key: str) -> int:
    try:
        raw_value = request.POST.get(key) if key in request.POST else request.GET.get(key)
        return max(int(raw_value or 0), 0)
    except (TypeError, ValueError):
        return 0


def _allocation_context(request, allocation, *, error="", ok_message=""):
    source_snapshot = allocation.source_snapshot
    source_code = (
        str(source_snapshot.container_code or "").strip()
        or str(getattr(source_snapshot.container, "container_code", "") or "").strip()
        or str(getattr(source_snapshot.location, "location_code", "") or "").strip()
    )
    source_confirmed = bool(request.session.get(_source_session_key(allocation.id)))
    open_count = FbsReplenishmentAllocation.objects.filter(
        line__plan=allocation.line.plan,
        status__in=OPEN_ALLOCATION_STATUSES,
    ).count()
    return _base_context(
        request,
        page_title=f"FBS · Задание #{allocation.id}",
        back_url=reverse("fbs:tsd_reachtruck"),
        allocation=allocation,
        plan=allocation.line.plan,
        source_code=source_code,
        source_confirmed=source_confirmed,
        open_count=open_count,
        is_staging_transfer=bool(
            allocation.line.plan.client_movement_request_id
            and allocation.line.plan.mode == FbsReplenishmentPlan.MODE_ITEM
            and allocation.line.plan.target_box_id is None
            and allocation.line.plan.staging_location_id
        ),
        staging_container_code=(
            movement_staging_container_code(
                allocation.line.plan.client_movement_request_id
            )
            if allocation.line.plan.client_movement_request_id
            and allocation.line.plan.mode == FbsReplenishmentPlan.MODE_ITEM
            else ""
        ),
        error=error,
        ok_message=ok_message,
    )


@fbs_module_required
@role_required(*TSD_ROLES)
@require_GET
def tsd_home(request):
    role = get_request_role(request)
    if role == "reachtruck_driver":
        return redirect("/reachtruck/")
    if role == "picker":
        return render(request, "fbs/tsd_picker_home.html", _picker_home_context(request))
    if role == "processing_worker":
        return redirect("fbs:tsd_picking")
    if role == "fbs_controller":
        return redirect("fbs:controller_home")
    if role in STOREKEEPER_ROLES or role == "developer":
        return redirect("fbs:tsd_storekeeper")
    if role in REACHTRUCK_ROLES:
        return redirect("/reachtruck/")
    if role in PICKER_ROLES:
        return redirect("fbs:tsd_picking")
    return HttpResponseForbidden("Доступ запрещен")


def _picker_free_box_move(request):
    """Move one scanned unit straight from one FBS box into another.

    The neighbouring form places stock into an addressed PR rack cell, which
    only works once a head manager has configured those cells.  Warehouse work
    also needs the plainer move the floor actually does: find a box, scan it,
    scan the box it goes into, scan the unit.  No task, no reservation and no
    cell registry -- the destination box already stands somewhere, and the
    movement service keeps the client, stock and reservation checks.

    The scan resolvers are shared with the rack form on purpose: marked goods
    must be identified by their Data Matrix in both, and duplicate barcodes must
    fail the same way, or the two forms would disagree about the same scan.
    """

    from .services.movements import (
        claim_internal_movement,
        create_internal_movement,
        scan_internal_movement,
    )
    from .services.racks import _resolve_scanned_balance, _resolve_source_box

    source_scan = str(request.POST.get("source_box_scan") or "").strip()
    target_scan = str(request.POST.get("target_box_scan") or "").strip()
    item_scan = str(request.POST.get("item_scan") or "").strip()
    request.session["fbs_free_move_source_box"] = source_scan
    request.session["fbs_free_move_target_box"] = target_scan
    try:
        with transaction.atomic():
            if not target_scan:
                raise FbsMovementError("Отсканируйте короб, в который перекладываете товар.")
            source_box = _resolve_source_box(source_scan)
            target_box = _resolve_source_box(target_scan)
            balance = _resolve_scanned_balance(
                balances=list(source_box.stock_balances.select_related("sku_ref").all()),
                item_scan=item_scan,
            )
            movement = create_internal_movement(
                mode=FbsInternalMovement.MODE_ITEM,
                source_box=source_box,
                target_pallet=target_box.pallet,
                target_box=target_box,
                source_balance=balance,
                qty=1,
                requested_by=request.user,
                comment="Свободное перемещение FBS: короб в короб.",
            )
            claim_internal_movement(movement_id=movement.id, assigned_to=request.user)
            scan_internal_movement(
                movement_id=movement.id,
                source_box_scan=source_box.box_code,
                target_scan=target_box.box_code,
                item_scan=item_scan,
                performed_by=request.user,
            )
    except (FbsError, ValidationError, ValueError) as exc:
        message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
        messages.error(request, message)
    else:
        remaining = int(
            source_box.stock_balances.aggregate(total=Sum("qty"))["total"] or 0
        )
        if remaining <= 0:
            request.session.pop("fbs_free_move_source_box", None)
        messages.success(
            request,
            f"Перемещена 1 шт. в короб {target_box.box_code}. "
            f"В исходном коробе осталось {remaining} шт.",
        )
    return redirect("fbs:tsd_picker_move_box_box")


@fbs_module_required
@role_required("picker")
@require_GET
def tsd_picker_movements(request):
    return render(
        request,
        "fbs/tsd_picker_movements.html",
        _base_context(
            request,
            page_title="Перемещение",
            back_url=reverse("fbs:tsd_home"),
        ),
    )


def _picker_location_display_code(value: str) -> str:
    """Return the short human-facing code while keeping scanner formats accepted."""
    return normalize_operational_location_scan(value)


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_move_box_cell(request):
    if request.method == "POST":
        source_box_scan = str(request.POST.get("source_box_scan") or "").strip()
        target_cell_scan = str(request.POST.get("target_cell_scan") or "").strip()
        request.session["fbs_pr_move_source_box"] = source_box_scan
        request.session["fbs_pr_move_target_cell"] = target_cell_scan
        try:
            movement = move_scanned_box_item_to_rack_cell(
                source_box_scan=source_box_scan,
                target_cell_scan=target_cell_scan,
                item_scan=request.POST.get("item_scan", ""),
                performed_by=request.user,
                idempotency_key=request.POST.get("idempotency_key", ""),
            )
        except (FbsError, ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            remaining_qty = int(getattr(movement, "source_remaining_qty", 0) or 0)
            if remaining_qty <= 0:
                request.session.pop("fbs_pr_move_source_box", None)
            messages.success(
                request,
                f"Перемещена 1 шт. в ячейку {movement.target_cell.cell_code}. "
                f"В коробе осталось {remaining_qty} шт.",
            )
        return redirect("fbs:tsd_picker_move_box_cell")

    return render(
        request,
        "fbs/tsd_picker_move_box_cell.html",
        _base_context(
            request,
            page_title="Короб → ячейка",
            back_url=reverse("fbs:tsd_picker_movements"),
            selected_source_box=str(request.session.get("fbs_pr_move_source_box") or ""),
            selected_target_cell=_picker_location_display_code(
                request.session.get("fbs_pr_move_target_cell") or ""
            ),
            idempotency_key=uuid.uuid4().hex,
        ),
    )


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_move_cell_cell(request):
    if request.method == "POST":
        from .services.racks import move_scanned_rack_item_to_rack_cell

        source_cell_scan = str(request.POST.get("source_cell_scan") or "").strip()
        target_cell_scan = str(request.POST.get("target_cell_scan") or "").strip()
        request.session["fbs_pr_move_source_cell"] = source_cell_scan
        request.session["fbs_pr_move_cell_target"] = target_cell_scan
        try:
            movement = move_scanned_rack_item_to_rack_cell(
                source_cell_scan=source_cell_scan,
                target_cell_scan=target_cell_scan,
                item_scan=request.POST.get("item_scan", ""),
                performed_by=request.user,
                idempotency_key=request.POST.get("idempotency_key", ""),
            )
        except (FbsError, ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            remaining_qty = int(getattr(movement, "source_cell_remaining_qty", 0) or 0)
            if remaining_qty <= 0:
                request.session.pop("fbs_pr_move_source_cell", None)
            messages.success(
                request,
                f"Перемещена 1 шт. из ячейки {movement.source_cell_code} "
                f"в ячейку {movement.target_cell.cell_code}. "
                f"В исходной ячейке осталось {remaining_qty} шт.",
            )
        return redirect("fbs:tsd_picker_move_cell_cell")

    return render(
        request,
        "fbs/tsd_picker_move_cell_cell.html",
        _base_context(
            request,
            page_title="Ячейка → ячейка",
            back_url=reverse("fbs:tsd_picker_movements"),
            selected_source_cell=_picker_location_display_code(
                request.session.get("fbs_pr_move_source_cell") or ""
            ),
            selected_target_cell=_picker_location_display_code(
                request.session.get("fbs_pr_move_cell_target") or ""
            ),
            idempotency_key=uuid.uuid4().hex,
        ),
    )


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_move_box_box(request):
    if request.method == "POST":
        return _picker_free_box_move(request)
    return render(
        request,
        "fbs/tsd_picker_move_box_box.html",
        _base_context(
            request,
            page_title="Короб → короб",
            back_url=reverse("fbs:tsd_picker_movements"),
            selected_free_source_box=str(
                request.session.get("fbs_free_move_source_box") or ""
            ),
            selected_free_target_box=str(
                request.session.get("fbs_free_move_target_box") or ""
            ),
        ),
    )


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_move_full_box(request):
    box_scan = str(request.POST.get("box_scan") or "").strip()
    destination_location_scan = str(
        request.POST.get("destination_location_scan") or ""
    ).strip()
    move_scope = str(request.POST.get("move_scope") or "box").strip().lower()
    if request.method == "POST":
        from reachtruck_free.box_relocation import (
            complete_fbs_box_relocation,
            inspect_fbs_box,
            start_fbs_box_relocation,
        )

        request.session["fbs_picker_full_box_scan"] = box_scan
        request.session[
            "fbs_picker_full_box_destination"
        ] = destination_location_scan
        try:
            with transaction.atomic():
                if move_scope == "pallet":
                    from .services.free_relocation import place_fbs_pallet_box_group

                    destination = _picker_os_location_from_scan(
                        destination_location_scan
                    )
                    operation = place_fbs_pallet_box_group(
                        source_scan=box_scan,
                        destination_location_id=int(destination.id),
                        performed_by=request.user,
                        performed_by_role="picker",
                    )
                    task = operation.tasks.order_by("id").first()
                    group_payload = dict(task.payload or {}) if task else {}
                elif move_scope == "box":
                    inspected = inspect_fbs_box(box_scan)
                    if not inspected.get("found"):
                        raise FbsMovementError(
                            "Активный FBS-короб по этому QR не найден."
                        )
                    blockers = list(inspected.get("blockers") or ())
                    if not inspected.get("ok") or not inspected.get("can_move"):
                        raise FbsMovementError(
                            "; ".join(blockers)
                            or "FBS-короб сейчас недоступен для перемещения."
                        )
                    operation = start_fbs_box_relocation(
                        box_id=int(inspected["fbs_box_id"]),
                        performed_by=request.user,
                        expected_location_id=int(inspected["location_id"]),
                        performed_by_role="picker",
                    )
                    operation = complete_fbs_box_relocation(
                        operation_id=operation.id,
                        destination_scan=destination_location_scan,
                        performed_by=request.user,
                        performed_by_role="picker",
                    )
                    destination = operation.destination_location
                    group_payload = {}
                else:
                    raise FbsMovementError("Неизвестный режим перемещения.")
        except (FbsError, ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            request.session.pop("fbs_picker_full_box_scan", None)
            request.session.pop("fbs_picker_full_box_destination", None)
            if move_scope == "pallet":
                messages.success(
                    request,
                    f"FBS-паллета {group_payload.get('fbs_pallet_code') or ''}: "
                    f"размещено коробов {int(group_payload.get('moved_box_count') or 0)}, "
                    f"уже было на адресе {int(group_payload.get('already_at_destination_count') or 0)}. "
                    f"Итого {int(group_payload.get('box_count') or 0)} коробов в "
                    f"{destination.location_code or destination.display_name}.",
                )
            else:
                messages.success(
                    request,
                    f"FBS-короб {inspected['box_code']} целиком перемещён в "
                    f"{destination.location_code or destination.display_name}.",
                )
        return redirect("fbs:tsd_picker_move_full_box")

    return render(
        request,
        "fbs/tsd_picker_move_full_box.html",
        _base_context(
            request,
            page_title="Короб / паллета → место",
            back_url=reverse("fbs:tsd_picker_movements"),
            selected_full_box=str(
                request.session.get("fbs_picker_full_box_scan") or ""
            ),
            selected_destination_location=str(
                request.session.get("fbs_picker_full_box_destination") or ""
            ),
        ),
    )


def _picker_os_location_from_scan(scan_value: str) -> WarehouseLocation:
    from reachtruck_free.services import parse_destination_scan

    normalized = normalize_operational_location_scan(scan_value)
    if not normalized:
        raise FbsMovementError("Отсканируйте ячейку OS.")
    queryset = WarehouseLocation.objects.filter(
        warehouse_code="MSK",
        zone_code__iexact="OS",
        is_active=True,
        is_storage=True,
    )
    location = queryset.filter(location_code__iexact=normalized).first()
    if location is not None:
        return location
    parsed, error = parse_destination_scan(normalized)
    if error or not parsed or str(parsed.get("zone") or "").upper() != "OS":
        raise FbsMovementError(error or "Отсканируйте полную ячейку OS.")
    location = queryset.filter(
        row_no=int(parsed.get("row") or 0),
        section_no=int(parsed.get("section") or 0),
        tier_no=int(parsed.get("tier") or 0),
        cell_no=int(parsed.get("cell") or 0),
    ).first()
    if location is None:
        raise FbsMovementError("Ячейка не найдена в действующей топологии OS.")
    return location


def _picker_fbs_location_from_scan(scan_value: str) -> WarehouseLocation:
    from sklad.services.operational_locations import (
        resolve_operational_location_scan,
    )

    normalized = normalize_operational_location_scan(scan_value)
    if not normalized:
        raise FbsMovementError("Отсканируйте точное FBS-место OS, PR или OTG.")
    operational_location = WarehouseLocation.objects.filter(
        warehouse_code="MSK",
        location_code__iexact=normalized,
        zone_code__in=("PR", "OTG"),
    ).first()
    if operational_location is None:
        return _picker_os_location_from_scan(scan_value)
    try:
        return resolve_operational_location_scan(
            scan_value,
            expected_zone=str(operational_location.zone_code).upper(),
            require_fbs=True,
            required_slots=1,
            lock=True,
        )
    except ValidationError as exc:
        raise FbsMovementError("; ".join(exc.messages)) from exc


def _picker_move_pallet_to_location(
    *,
    pallet_scan: str,
    destination: WarehouseLocation,
    performed_by,
) -> tuple[WarehouseOperation, str, bool]:
    from reachtruck_free.services import inspect_pallet, snapshots_agency
    from sklad.services.warehouse_write_path import WarehouseWritePathService

    from .services.free_relocation import (
        complete_fbs_free_relocation,
        inspect_fbs_pallet,
        start_fbs_free_relocation,
    )

    inspected = inspect_fbs_pallet(pallet_scan)
    if inspected.get("found"):
        blockers = list(inspected.get("blockers") or ())
        if not inspected.get("ok") or not inspected.get("can_move"):
            raise FbsMovementError(
                "; ".join(blockers)
                or "FBS-паллета сейчас недоступна для перемещения."
            )
        operation = start_fbs_free_relocation(
            pallet_id=int(inspected["fbs_pallet_id"]),
            performed_by=performed_by,
            expected_location_id=int(inspected["location_id"]),
            performed_by_role="picker",
        )
        operation = complete_fbs_free_relocation(
            operation_id=operation.id,
            destination_row_no=destination.row_no,
            destination_section_no=destination.section_no,
            destination_tier_no=destination.tier_no,
            destination_cell_no=destination.cell_no,
            destination_zone_code=str(destination.zone_code).upper(),
            destination_location_id=int(destination.id),
            performed_by=performed_by,
            performed_by_role="picker",
        )
        return operation, str(inspected["pallet_code"]), True

    inspected = inspect_pallet(pallet_scan)
    if not inspected.get("found"):
        raise FbsMovementError(
            "Активная FBS-паллета или физическая складская паллета "
            "по этому QR не найдена."
        )
    blockers = list(inspected.get("blockers") or ())
    # A sealed receiving-flow pallet is deliberately marked can_move=False by
    # the generic inspector because its normal route is driver takeout.  This
    # dedicated picker command may still relocate it to OS; the warehouse
    # write path below remains the authority for stock state and reservations.
    if not inspected.get("ok") or blockers:
        raise FbsMovementError(
            "; ".join(str(item) for item in blockers if str(item).strip())
            or "Физическая складская паллета сейчас недоступна для перемещения."
        )
    if str(destination.zone_code or "").strip().upper() != "OS":
        raise FbsMovementError(
            "Физическую паллету приёмки этим экраном можно переместить "
            "только в точное место OS."
        )
    agency_id = int(inspected.get("agency_id") or 0)
    if not agency_id:
        raise FbsMovementError("Не найден клиент физической складской паллеты.")
    pallet_code = str(inspected.get("pallet_code") or pallet_scan).strip()
    operation = WarehouseWritePathService.start_free_pallet_relocation(
        agency=snapshots_agency(agency_id),
        pallet_code=pallet_code,
        performed_by=performed_by,
        expected_location_id=int(inspected.get("location_id") or 0) or None,
    )
    operation = WarehouseWritePathService.complete_free_pallet_relocation(
        operation_id=operation.id,
        destination_zone_code="OS",
        destination_location_code=str(destination.location_code or "").strip(),
        destination_row_no=int(destination.row_no or 0),
        destination_section_no=int(destination.section_no or 0),
        destination_tier_no=int(destination.tier_no or 0),
        destination_cell_no=int(destination.cell_no or 0),
        performed_by=performed_by,
    )
    return operation, pallet_code, False


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_move_pallet(request):
    pallet_scan = str(request.POST.get("pallet_scan") or "").strip()
    destination_location_scan = str(
        request.POST.get("destination_location_scan") or ""
    ).strip()
    if request.method == "POST":
        request.session["fbs_picker_pallet_scan"] = pallet_scan
        request.session["fbs_picker_pallet_destination"] = destination_location_scan
        try:
            with transaction.atomic():
                destination = _picker_fbs_location_from_scan(destination_location_scan)
                _operation, moved_pallet_code, moved_fbs_pallet = (
                    _picker_move_pallet_to_location(
                        pallet_scan=pallet_scan,
                        destination=destination,
                        performed_by=request.user,
                    )
                )
        except (FbsError, ValidationError, ValueError) as exc:
            message = "; ".join(getattr(exc, "messages", ()) or (str(exc),))
            messages.error(request, message)
        else:
            for key in (
                "fbs_picker_pallet_scan",
                "fbs_picker_pallet_source",
                "fbs_picker_pallet_destination",
            ):
                request.session.pop(key, None)
            messages.success(
                request,
                (
                    f"FBS-паллета {moved_pallet_code} перемещена в "
                    if moved_fbs_pallet
                    else f"Физическая паллета {moved_pallet_code} перемещена в "
                )
                + f"{destination.location_code or destination.display_name}. "
                + (
                    ""
                    if moved_fbs_pallet
                    else "Количество товара и резервы не изменены."
                ),
            )
        return redirect("fbs:tsd_picker_move_pallet")

    return render(
        request,
        "fbs/tsd_picker_move_pallet.html",
        _base_context(
            request,
            page_title="FBS-паллета → место",
            back_url=reverse("fbs:tsd_picker_movements"),
            selected_pallet=str(request.session.get("fbs_picker_pallet_scan") or ""),
            selected_destination_location=str(
                request.session.get("fbs_picker_pallet_destination") or ""
            ),
        ),
    )


@fbs_module_required
@role_required("picker")
@require_GET
def tsd_picker_returns(request):
    restock_requests = list(
        FbsPickRestockRequest.objects.select_related(
            "batch__workstation",
            "assigned_to",
            "source_tote",
            "quarantine_box__pallet__cell__location",
        )
        .filter(
            Q(
                status=FbsPickRestockRequest.STATUS_QUEUED,
                assigned_to__isnull=True,
            )
            | Q(
                status=FbsPickRestockRequest.STATUS_IN_PROGRESS,
                assigned_to=request.user,
            )
        )
        .filter(Q(order__isnull=True) | Q(source_tote__isnull=False))
        .annotate(source_box_count=Count("lines__source_box", distinct=True))
        .order_by("created_at", "id")[:50]
    )
    return render(
        request,
        "fbs/tsd_picker_returns.html",
        _base_context(
            request,
            page_title="Возврат отборов",
            ok_message="Возврат завершен. Выберите следующее задание." if request.GET.get("done") == "1" else "",
            back_url=reverse("fbs:tsd_home"),
            restock_requests=restock_requests,
        ),
    )


@fbs_module_required
@role_required("picker")
@require_http_methods(["GET", "POST"])
def tsd_picker_location(request):
    scan = str(request.POST.get("scan") or "").strip() if request.method == "POST" else ""
    error = ""
    cell = None
    box = None
    pallet = None
    result_kind = ""
    balances = []
    product_placements = []
    totals = {
        "qty": 0,
        "available": 0,
        "reserved": 0,
        "boxes": 0,
        "pallets": 0,
        "locations": 0,
        "positions": 0,
    }
    if request.method == "POST":
        normalized_location_scan = _normalized_picker_location_scan(scan)
        normalized_box_scan = _normalized_picker_box_scan(scan)
        if not normalized_location_scan and not normalized_box_scan:
            error = "Скан не получен. Повторите сканирование места, палеты, короба или товара."
        else:
            cells = FbsStorageCell.objects.select_related("location").filter(
                is_active=True,
                location__is_active=True,
            )
            cell = next(
                (
                    candidate
                    for candidate in cells
                    if normalized_location_scan in _picker_location_candidates(candidate)
                ),
                None,
            )
            if cell is not None:
                result_kind = "cell"
                balance_queryset = _picker_location_balance_queryset(cell)
            else:
                box_matches = list(
                    FbsBox.objects.select_related(
                        "agency",
                        "pallet__cell__location",
                    )
                    .filter(
                        box_code__iexact=normalized_box_scan,
                        status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
                    )
                    .order_by("id")[:2]
                )
                if len(box_matches) > 1:
                    error = (
                        "Этот QR короба используется у нескольких клиентов. "
                        "Нужен уникальный QR короба."
                    )
                elif box_matches:
                    box = box_matches[0]
                    cell = box.pallet.cell
                    result_kind = "box"
                    balance_queryset = _fbs_stock_base_queryset().filter(
                        box=box,
                    )
                else:
                    pallet_matches = list(
                        FbsPallet.objects.select_related(
                            "agency",
                            "cell__location",
                        )
                        .filter(
                            pallet_code__iexact=normalized_box_scan,
                            status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
                        )
                        .order_by("id")[:2]
                    )
                    if len(pallet_matches) > 1:
                        error = (
                            "Этот QR палеты используется у нескольких клиентов. "
                            "Нужен уникальный QR палеты."
                        )
                    elif pallet_matches:
                        pallet = pallet_matches[0]
                        cell = pallet.cell
                        result_kind = "pallet"
                        balance_queryset = _fbs_stock_base_queryset().filter(
                            box__pallet=pallet,
                        )
                    else:
                        product_balances = list(
                            _picker_product_balance_queryset(normalized_box_scan).order_by(
                                "box__pallet__cell__cell_code",
                                "box__box_code",
                                "agency_id",
                                "sku_code",
                                "id",
                            )
                        )
                        if product_balances:
                            result_kind = "product"
                            product_placements = _picker_product_placements(product_balances)
                            locations = {
                                row.location_code or row.location_label
                                for row in product_placements
                            }
                            totals = {
                                "qty": sum(int(balance.qty or 0) for balance in product_balances),
                                "available": sum(
                                    int(balance.available_qty or 0)
                                    for balance in product_balances
                                ),
                                "reserved": sum(
                                    int(balance.reserved_qty or 0)
                                    + int(balance.external_reserved_qty or 0)
                                    for balance in product_balances
                                ),
                                "boxes": len({row.box.id for row in product_placements}),
                                "pallets": len({row.pallet.id for row in product_placements}),
                                "locations": len(locations),
                                "positions": len(product_placements),
                            }
                        else:
                            error = (
                                "На FBS не найдены место, палета, короб или остаток товара "
                                "по этому коду. Проверьте скан и повторите."
                            )
            if not error and result_kind != "product":
                summary = balance_queryset.aggregate(
                    qty=Sum("qty"),
                    available=Sum("available_qty"),
                    reserved=Sum(F("reserved_qty") + F("external_reserved_qty")),
                    boxes=Count("box_id", distinct=True),
                    pallets=Count("box__pallet_id", distinct=True),
                    positions=Count("id"),
                )
                balances = list(
                    balance_queryset.order_by("box__box_code", "sku_code", "id")[:200]
                )
                totals = {
                    "qty": int(summary["qty"] or 0),
                    "available": int(summary["available"] or 0),
                    "reserved": int(summary["reserved"] or 0),
                    "boxes": int(summary["boxes"] or 0),
                    "pallets": int(summary["pallets"] or 0),
                    "locations": 1,
                    "positions": int(summary["positions"] or 0),
                }
    return render(
        request,
        "fbs/tsd_picker_location.html",
        _base_context(
            request,
            page_title="Информация о месте",
            back_url=reverse("fbs:tsd_home"),
            scan=scan,
            cell=cell,
            box=box,
            pallet=pallet,
            result_kind=result_kind,
            location_label=_picker_location_label(cell) if cell else "",
            box_location_label=(fbs_box_physical_location_label(box) if box else ""),
            box_location_code=(_picker_box_location_code(box) if box else ""),
            balances=balances,
            product_placements=product_placements,
            totals=totals,
            error=error,
        ),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_storekeeper(request):
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        if action == "save_alert_responsibles":
            if get_request_role(request) not in INVENTORY_MANAGER_ROLES:
                return HttpResponseForbidden(
                    "Назначать ответственных FBS может только руководитель."
                )
            selected_ids = []
            for raw_value in request.POST.getlist("responsible_user_ids"):
                try:
                    selected_ids.append(int(raw_value))
                except (TypeError, ValueError):
                    continue
            selected_ids = list(dict.fromkeys(selected_ids))
            employees = list(
                Employee.objects.filter(
                    role="storekeeper",
                    is_active=True,
                    user_id__in=selected_ids,
                    user__is_active=True,
                ).select_related("user")
            )
            if not employees:
                return render(
                    request,
                    "fbs/tsd_storekeeper_list.html",
                    _storekeeper_list_context(
                        request,
                        error="Выберите хотя бы одного активного кладовщика FBS.",
                    ),
                    status=400,
                )
            valid_user_ids = {employee.user_id for employee in employees}
            with transaction.atomic():
                FbsStorekeeperResponsible.objects.exclude(
                    user_id__in=valid_user_ids
                ).update(is_active=False, assigned_by=request.user)
                for user_id in valid_user_ids:
                    FbsStorekeeperResponsible.objects.update_or_create(
                        user_id=user_id,
                        defaults={"is_active": True, "assigned_by": request.user},
                    )
            names = ", ".join(sorted(employee.full_name for employee in employees))
            messages.success(request, f"Ответственные FBS сохранены: {names}.")
            return redirect("fbs:tsd_storekeeper")
        if action != "transfer_controller_tote":
            return render(
                request,
                "fbs/tsd_storekeeper_list.html",
                _storekeeper_list_context(request, error="Неизвестная операция."),
                status=400,
            )
        try:
            pick_context = transfer_controller_pick_tote(
                pick_tote_id=int(request.POST.get("pick_tote_id") or 0),
                target_session_id=int(request.POST.get("target_session_id") or 0),
                performed_by=request.user,
            )
        except (TypeError, ValueError, FbsError) as exc:
            return render(
                request,
                "fbs/tsd_storekeeper_list.html",
                _storekeeper_list_context(request, error=str(exc)),
                status=400,
            )
        pick_context.refresh_from_db(
            fields=["session_id", "tote_id", "updated_at"]
        )
        target_session = FbsControllerSession.objects.select_related(
            "controller", "workstation"
        ).get(pk=pick_context.session_id)
        controller_name = (
            target_session.controller.get_full_name()
            or target_session.controller.get_username()
        )
        messages.success(
            request,
            f"Тара {pick_context.tote.barcode} передана на "
            f"{target_session.workstation.name}, контролер {controller_name}. "
            "Прогресс проверки сохранен.",
        )
        return redirect("fbs:tsd_storekeeper")
    return render(
        request,
        "fbs/tsd_storekeeper_list.html",
        _storekeeper_list_context(request),
    )


FBS_STOREKEEPER_ALERT_SNOOZE_SESSION_KEY = "fbs_storekeeper_alert_snooze_until"
FBS_STOREKEEPER_ALERT_SNOOZE_HOURS = 1


def _storekeeper_alert_payload_for_request(request, payload):
    payload = dict(payload)
    now = timezone.now()
    try:
        snooze_until_timestamp = float(
            request.session.get(FBS_STOREKEEPER_ALERT_SNOOZE_SESSION_KEY) or 0
        )
    except (TypeError, ValueError):
        snooze_until_timestamp = 0
    latest_acknowledgement = (
        FbsStorekeeperAlertAcknowledgement.objects.filter(
            responsible=request.user,
            acknowledged_at__gt=now
            - timedelta(hours=FBS_STOREKEEPER_ALERT_SNOOZE_HOURS),
        )
        .order_by("-acknowledged_at")
        .values_list("acknowledged_at", flat=True)
        .first()
    )
    if latest_acknowledgement is not None:
        acknowledgement_snooze_until = latest_acknowledgement + timedelta(
            hours=FBS_STOREKEEPER_ALERT_SNOOZE_HOURS
        )
        snooze_until_timestamp = max(
            snooze_until_timestamp,
            acknowledgement_snooze_until.timestamp(),
        )
    snooze_is_active = snooze_until_timestamp > now.timestamp()
    if snooze_until_timestamp and not snooze_is_active:
        request.session.pop(FBS_STOREKEEPER_ALERT_SNOOZE_SESSION_KEY, None)

    pending_due_count = max(0, int(payload.get("due_count") or 0))
    payload["pending_due_count"] = pending_due_count
    payload["due_count"] = 0 if snooze_is_active else pending_due_count
    payload["modal_due_count"] = payload["due_count"]
    payload["control_snoozed"] = bool(snooze_is_active and pending_due_count)
    payload["snoozed_due_count"] = pending_due_count if snooze_is_active else 0
    payload["snooze_hours"] = FBS_STOREKEEPER_ALERT_SNOOZE_HOURS
    if snooze_is_active:
        seconds_left = max(1, int(snooze_until_timestamp - now.timestamp()))
        snooze_until = now + timedelta(seconds=seconds_left)
        payload["snooze_minutes_left"] = max(1, (seconds_left + 59) // 60)
        payload["snoozed_until"] = timezone.localtime(snooze_until).isoformat()
        payload["snoozed_until_label"] = timezone.localtime(snooze_until).strftime(
            "%H:%M"
        )
    return payload


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_storekeeper_alerts(request):
    from .services.storekeeper_alerts import (
        acknowledge_storekeeper_alerts,
        build_storekeeper_alert_payload,
        is_storekeeper_alert_responsible,
    )

    if not is_storekeeper_alert_responsible(request.user):
        return JsonResponse(
            {"enabled": False, "detail": "Сотрудник не назначен ответственным FBS."},
        )
    if request.method == "POST":
        try:
            acknowledged_count = acknowledge_storekeeper_alerts(
                user=request.user,
                alert_keys=request.POST.getlist("alert_keys"),
            )
        except PermissionError as exc:
            return JsonResponse({"enabled": False, "detail": str(exc)}, status=403)
        if acknowledged_count:
            request.session[FBS_STOREKEEPER_ALERT_SNOOZE_SESSION_KEY] = int(
                (
                    timezone.now()
                    + timedelta(hours=FBS_STOREKEEPER_ALERT_SNOOZE_HOURS)
                ).timestamp()
            )
        payload = _storekeeper_alert_payload_for_request(
            request,
            build_storekeeper_alert_payload(),
        )
        payload["acknowledged_count"] = acknowledged_count
        return JsonResponse(payload)
    return JsonResponse(
        _storekeeper_alert_payload_for_request(
            request,
            build_storekeeper_alert_payload(),
        )
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_controller_settings(request):
    redirect_url = reverse("fbs:tsd_controller_settings")
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        if action == "transfer_waiting_controller_tote":
            from .services.controller_tote_transfer import transfer_waiting_controller_tote
            try:
                binding = transfer_waiting_controller_tote(
                    binding_id=int(request.POST.get("binding_id") or 0),
                    expected_workstation_id=int(request.POST.get("source_workstation_id") or 0),
                    target_session_id=int(request.POST.get("target_session_id") or 0),
                    performed_by=request.user,
                )
            except (TypeError, ValueError, FbsError) as exc:
                return render(request, "fbs/tsd_controller_settings.html",
                    _controller_settings_context(request, error=str(exc)), status=400)
            messages.success(request, f"Тара {binding.tote.barcode}, волна #{binding.pick_batch_id}, "
                f"передана на {binding.workstation.name}. Ожидает приема контролером.")
            return redirect(f"{redirect_url}#controller-work")
        if action != "transfer_controller_tote":
            return render(
                request,
                "fbs/tsd_controller_settings.html",
                _controller_settings_context(request, error="Неизвестная операция."),
                status=400,
            )
        try:
            pick_context = transfer_controller_pick_tote(
                pick_tote_id=int(request.POST.get("pick_tote_id") or 0),
                target_session_id=int(request.POST.get("target_session_id") or 0),
                performed_by=request.user,
            )
        except (TypeError, ValueError, FbsError) as exc:
            return render(
                request,
                "fbs/tsd_controller_settings.html",
                _controller_settings_context(request, error=str(exc)),
                status=400,
            )
        pick_context.refresh_from_db(fields=["session_id", "tote_id", "updated_at"])
        target_session = FbsControllerSession.objects.select_related(
            "controller__employee_profile", "workstation"
        ).get(pk=pick_context.session_id)
        messages.success(
            request,
            f"Тара {pick_context.tote.barcode} и волна #{pick_context.pick_batch_id} "
            f"переданы на {target_session.workstation.name}, контролер "
            f"{_controller_settings_user_name(target_session.controller)}. "
            "Прогресс проверки сохранен.",
        )
        return redirect(f"{redirect_url}#controller-work")
    return render(
        request,
        "fbs/tsd_controller_settings.html",
        _controller_settings_context(request),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_wave_settings(request):
    redirect_url = reverse("fbs:tsd_wave_settings")
    if request.method == "POST":
        action = str(request.POST.get("action") or "save").strip()
        if action == "reorder_queue":
            try:
                batch_id = int(request.POST.get("batch_id") or 0)
                reorder_queued_wave(
                    batch_id=batch_id,
                    action=request.POST.get("direction", ""),
                    actor=request.user,
                )
            except (TypeError, ValueError, FbsError) as exc:
                messages.error(request, str(exc))
            else:
                messages.success(request, f"Положение волны #{batch_id} сохранено на сегодня.")
            return redirect(f"{redirect_url}#queue-order")

        if action == "save_marketplace_priorities":
            try:
                effective_from = date.fromisoformat(
                    str(request.POST.get("effective_from") or "").strip()
                )
                update_marketplace_queue_priorities(
                    priorities={
                        marketplace: request.POST.get(f"priority_{marketplace}")
                        for marketplace, _label in FbsIntegrationProfile.MARKETPLACE_CHOICES
                    },
                    effective_from=effective_from,
                    actor=request.user,
                )
            except (TypeError, ValueError, ValidationError) as exc:
                messages.error(request, str(exc))
            else:
                messages.success(
                    request,
                    f"Приоритет маркетплейсов сохранён с {effective_from:%d.%m.%Y}.",
                )
            return redirect(f"{redirect_url}#marketplace-priority")

        try:
            profile_id = int(request.POST.get("profile_id") or 0)
        except (TypeError, ValueError):
            profile_id = 0
        if profile_id <= 0:
            messages.error(request, "Кабинет FBS не выбран.")
            return redirect(redirect_url)

        with transaction.atomic():
            profile = get_object_or_404(
                FbsIntegrationProfile.objects.select_for_update(),
                pk=profile_id,
            )
            if action == "reset":
                FbsWavePolicy.objects.filter(profile=profile).delete()
                messages.success(
                    request,
                    f"Для {profile} восстановлен системный размер волны.",
                )
                return redirect(redirect_url)

            try:
                max_orders = int(request.POST.get("max_orders_per_wave") or 0)
                max_units = int(request.POST.get("max_units_per_wave") or 0)
            except (TypeError, ValueError):
                max_orders = max_units = 0
            if not (
                MIN_WAVE_SIZE <= max_orders <= MAX_WAVE_SIZE
                and MIN_WAVE_SIZE <= max_units <= MAX_WAVE_SIZE
            ):
                messages.error(
                    request,
                    "Укажите от 1 до 100 заказов и от 1 до 100 единиц товара.",
                )
                return redirect(redirect_url)

            policy = (
                FbsWavePolicy.objects.select_for_update()
                .filter(profile=profile)
                .first()
            )
            if policy is None:
                policy = FbsWavePolicy(profile=profile)
            policy.max_orders_per_wave = max_orders
            policy.max_units_per_wave = max_units
            policy.updated_by = request.user
            policy.full_clean(exclude={"id"})
            policy.save()
        messages.success(
            request,
            f"Размер новых волн для {profile} сохранён: "
            f"до {max_orders} заказов и до {max_units} единиц.",
        )
        return redirect(redirect_url)

    profiles = list(
        FbsIntegrationProfile.objects.select_related("agency", "wave_policy")
        .order_by("agency__agn_name", "marketplace", "name", "id")
    )
    for profile in profiles:
        limits = configured_or_default_wave_limits(profile)
        profile.wave_max_orders = limits.max_orders
        profile.wave_max_units = limits.max_units
        profile.wave_policy_configured = limits.configured
    marketplace_priorities = marketplace_queue_priority_rows()
    return render(
        request,
        "fbs/tsd_wave_settings.html",
        _base_context(
            request,
            page_title="Настройка волн",
            wave_profiles=profiles,
            wave_queue=current_wave_queue(),
            marketplace_priorities=marketplace_priorities,
            marketplace_priority_date=(
                marketplace_priorities[0].effective_from
                if marketplace_priorities
                else timezone.localdate()
            ),
        ),
    )


def _fbs_stock_base_queryset():
    return FbsStockBalance.objects.filter(
        qty__gt=0,
        box__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        box__pallet__status__in=(FbsPallet.STATUS_PLANNED, FbsPallet.STATUS_ACTIVE),
        box__pallet__cell__is_active=True,
    ).select_related(
        "agency",
        "sku_ref",
        "box__pallet__cell__location",
        "box__source_container__current_location",
    )


def _fbs_stock_physical_location_filter(value: str) -> Q:
    source_prefix = "box__source_container__current_location"
    source_is_physical = (
        Q(**{f"{source_prefix}__isnull": False})
        & ~Q(**{f"{source_prefix}__zone_kind__iexact": WarehouseLocation.ZONE_KIND_VIRTUAL})
        & ~Q(**{f"{source_prefix}__location_code__istartswith": VIRTUAL_FBS_PLAN_PREFIX})
    )
    source_matches = (
        Q(**{f"{source_prefix}__location_code__icontains": value})
        | Q(**{f"{source_prefix}__display_name__icontains": value})
    )
    source_needs_fallback = (
        Q(**{f"{source_prefix}__isnull": True})
        | Q(**{f"{source_prefix}__zone_kind__iexact": WarehouseLocation.ZONE_KIND_VIRTUAL})
        | Q(**{f"{source_prefix}__location_code__istartswith": VIRTUAL_FBS_PLAN_PREFIX})
    )
    fallback_matches = (
        Q(box__pallet__cell__cell_code__icontains=value)
        | Q(box__pallet__cell__location__location_code__icontains=value)
        | Q(box__pallet__cell__location__display_name__icontains=value)
    )
    return (source_is_physical & source_matches) | (source_needs_fallback & fallback_matches)


def _fbs_stock_apply_filters(request, queryset, *, include_sku=True):
    q = str(request.GET.get("q") or "").strip()
    sku = str(request.GET.get("sku") or "").strip()
    client_id = str(request.GET.get("client") or "").strip()
    stock_status = str(request.GET.get("stock_status") or "").strip()
    cell = str(request.GET.get("cell") or "").strip()
    box = str(request.GET.get("box") or "").strip()
    pallet = str(request.GET.get("pallet") or "").strip()

    if q:
        exact_code_queryset = queryset.filter(
            Q(sku_code=q)
            | Q(barcode=q)
            | Q(marking_code=q)
            | Q(lot_code=q)
            | Q(box__box_code=q)
            | Q(box__pallet__pallet_code=q)
            | Q(box__pallet__cell__cell_code=q)
        )
        exact_code_found = (
            len(q) >= 4
            and not any(character.isspace() for character in q)
            and exact_code_queryset.exists()
        )
        if exact_code_found:
            queryset = exact_code_queryset
        else:
            queryset = queryset.filter(
                Q(agency__agn_name__icontains=q)
                | Q(agency__short_name__icontains=q)
                | Q(sku_code__icontains=q)
                | Q(name__icontains=q)
                | Q(size__icontains=q)
                | Q(barcode__icontains=q)
                | Q(marking_code__icontains=q)
                | Q(lot_code__icontains=q)
                | Q(box__box_code__icontains=q)
                | Q(box__pallet__pallet_code__icontains=q)
                | Q(box__pallet__cell__cell_code__icontains=q)
            )
    if include_sku and sku:
        queryset = queryset.filter(
            Q(sku_code__icontains=sku) | Q(name__icontains=sku)
        )
    if client_id.isdigit():
        queryset = queryset.filter(agency_id=int(client_id))
    if stock_status == "available":
        queryset = queryset.filter(available_qty__gt=0)
    elif stock_status == "reserved":
        queryset = queryset.filter(Q(reserved_qty__gt=0) | Q(external_reserved_qty__gt=0))
    elif stock_status == "unavailable":
        queryset = queryset.filter(available_qty=0)
    else:
        stock_status = ""
    if cell:
        queryset = queryset.filter(_fbs_stock_physical_location_filter(cell))
    if box:
        queryset = queryset.filter(box__box_code__icontains=box)
    if pallet:
        queryset = queryset.filter(box__pallet__pallet_code__icontains=pallet)

    return queryset, {
        "q": q,
        "sku": sku,
        "client_id": client_id if client_id.isdigit() else "",
        "stock_status": stock_status,
        "cell": cell,
        "box": box,
        "pallet": pallet,
    }


def _fbs_stock_sort(request, queryset):
    sort_key = str(request.GET.get("sort") or "client").strip()
    if sort_key not in FBS_STOCK_SORT_FIELDS:
        sort_key = "client"
    default_dir = "desc" if sort_key in FBS_STOCK_DESCENDING_DEFAULTS else "asc"
    sort_dir = str(request.GET.get("dir") or default_dir).strip().lower()
    if sort_dir not in {"asc", "desc"}:
        sort_dir = default_dir
    field = FBS_STOCK_SORT_FIELDS[sort_key]
    ordering = f"-{field}" if sort_dir == "desc" else field
    return queryset.order_by(ordering, "agency_id", "sku_code", "id"), sort_key, sort_dir


def _fbs_stock_grouped_sort(request, queryset):
    sort_key = str(request.GET.get("sort") or "client").strip()
    if sort_key not in FBS_STOCK_BOX_SORT_FIELDS:
        sort_key = "client"
    default_dir = "desc" if sort_key in FBS_STOCK_DESCENDING_DEFAULTS else "asc"
    sort_dir = str(request.GET.get("dir") or default_dir).strip().lower()
    if sort_dir not in {"asc", "desc"}:
        sort_dir = default_dir

    grouped_queryset = queryset.values(
        "box_id",
        "agency_id",
        "agency__agn_name",
        "agency__short_name",
        "box__box_code",
        "box__pallet_id",
        "box__pallet__pallet_code",
        "box__pallet__cell_id",
        "box__pallet__cell__cell_code",
    ).annotate(
        group_sku=Min("sku_code"),
        group_name=Min("name"),
        group_barcode=Min("barcode"),
        group_expiry=Min("expiry_date"),
        group_updated=Max("updated_at"),
        group_qty=Sum("qty"),
        group_available_qty=Sum("available_qty"),
        group_reserved_qty=Sum("reserved_qty"),
    )
    field = FBS_STOCK_BOX_SORT_FIELDS[sort_key]
    ordering = f"-{field}" if sort_dir == "desc" else field
    return (
        grouped_queryset.order_by(ordering, "box__box_code", "box_id"),
        sort_key,
        sort_dir,
    )


def _fbs_stock_box_groups(page_rows, balances) -> list[dict]:
    balances_by_box = defaultdict(list)
    for balance in balances:
        balances_by_box[balance.box_id].append(balance)

    groups = []
    for page_row in page_rows:
        box_balances = balances_by_box.get(page_row["box_id"], [])
        if not box_balances:
            continue

        items_by_identity = {}
        for balance in box_balances:
            identity = (
                balance.sku_ref_id,
                balance.sku_code,
                balance.name,
                balance.size,
                balance.barcode,
                balance.goods_type,
                balance.lot_code,
                balance.expiry_date,
            )
            item = items_by_identity.setdefault(
                identity,
                {
                    "balance": balance,
                    "qty": 0,
                    "available_qty": 0,
                    "reserved_qty": 0,
                    "external_reserved_qty": 0,
                    "marking_codes": [],
                    "balance_count": 0,
                },
            )
            item["qty"] += int(balance.qty or 0)
            item["available_qty"] += int(balance.available_qty or 0)
            item["reserved_qty"] += int(balance.reserved_qty or 0)
            item["external_reserved_qty"] += int(balance.external_reserved_qty or 0)
            item["balance_count"] += 1
            marking_code = str(balance.marking_code or "").strip()
            if marking_code:
                item["marking_codes"].append(marking_code)

        items = sorted(
            items_by_identity.values(),
            key=lambda item: (
                str(item["balance"].sku_code or "").casefold(),
                str(item["balance"].name or "").casefold(),
                str(item["balance"].size or "").casefold(),
                str(item["balance"].barcode or "").casefold(),
            ),
        )
        for item in items:
            item["marking_codes"].sort()
            item["marking_count"] = len(item["marking_codes"])

        first_balance = box_balances[0]
        physical_location_code = fbs_box_physical_location_code(first_balance.box)
        physical_location_label = fbs_box_physical_location_label(first_balance.box)
        groups.append(
            {
                "box_id": first_balance.box_id,
                "box": first_balance.box,
                "agency": first_balance.agency,
                "items": items,
                "item_count": len(items),
                "balance_count": len(box_balances),
                "marking_count": sum(item["marking_count"] for item in items),
                "qty": sum(item["qty"] for item in items),
                "available_qty": sum(item["available_qty"] for item in items),
                "reserved_qty": sum(item["reserved_qty"] for item in items),
                "external_reserved_qty": sum(item["external_reserved_qty"] for item in items),
                "physical_location_code": physical_location_code,
                "physical_location_label": physical_location_label,
                "updated_at": max(
                    (
                        balance.updated_at
                        for balance in box_balances
                        if balance.updated_at is not None
                    ),
                    default=None,
                ),
            }
        )
    return groups


def _fbs_stock_sort_urls(request, current_sort: str, current_dir: str) -> dict[str, str]:
    urls = {}
    for key in FBS_STOCK_SORT_FIELDS:
        params = request.GET.copy()
        params.pop("page", None)
        params["sort"] = key
        params["dir"] = (
            "desc"
            if key == current_sort and current_dir == "asc"
            else "asc"
        )
        urls[key] = f"{request.path}?{params.urlencode()}"
    return urls


def _fbs_stock_pagination_links(page_obj, *, radius: int = 2) -> list[dict]:
    total_pages = int(page_obj.paginator.num_pages or 0)
    current = int(page_obj.number or 1)
    pages = {1, total_pages}
    pages.update(range(max(1, current - radius), min(total_pages, current + radius) + 1))
    links = []
    previous = 0
    for number in sorted(page for page in pages if page > 0):
        if previous and number - previous > 1:
            links.append({"ellipsis": True})
        links.append({"number": number, "current": number == current})
        previous = number
    return links


def _fbs_stock_client_options():
    return Agency.objects.filter(
        fbs_stock_balances__qty__gt=0,
        fbs_stock_balances__box__status__in=(FbsBox.STATUS_PLANNED, FbsBox.STATUS_ACTIVE),
        fbs_stock_balances__box__pallet__status__in=(
            FbsPallet.STATUS_PLANNED,
            FbsPallet.STATUS_ACTIVE,
        ),
        fbs_stock_balances__box__pallet__cell__is_active=True,
    ).distinct().order_by("short_name", "agn_name", "id")


def _fbs_stock_sku_options(queryset, *, limit: int = 1000) -> list[dict[str, str]]:
    options = []
    seen_values = set()
    rows = (
        queryset.values_list("sku_code", "name")
        .order_by("sku_code", "name")
        .distinct()[:limit]
    )
    for sku_code, name in rows:
        sku_value = str(sku_code or "").strip()
        name_value = str(name or "").strip()
        value = sku_value or name_value
        normalized_value = value.casefold()
        if not value or normalized_value in seen_values:
            continue
        seen_values.add(normalized_value)
        options.append(
            {
                "value": value,
                "label": " · ".join(part for part in (sku_value, name_value) if part),
            }
        )
    return options


def _fbs_stock_summary(queryset) -> dict:
    totals = queryset.aggregate(
        row_count=Count("id"),
        client_count=Count("agency_id", distinct=True),
        box_count=Count("box_id", distinct=True),
        pallet_count=Count("box__pallet_id", distinct=True),
        qty=Sum("qty"),
        available_qty=Sum("available_qty"),
        reserved_qty=Sum("reserved_qty"),
        external_reserved_qty=Sum("external_reserved_qty"),
    )
    totals["sku_count"] = queryset.values(
        "agency_id", "sku_code", "size", "barcode", "goods_type"
    ).distinct().count()
    summary = {key: int(value or 0) for key, value in totals.items()}
    # Explain the stored availability without deriving a new warehouse balance.
    summary["unavailable_unreserved_qty"] = (
        summary["qty"] - summary["available_qty"] - summary["reserved_qty"] - summary["external_reserved_qty"]
    )
    return summary


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_stock(request):
    base_queryset = _fbs_stock_base_queryset()
    option_queryset, _ = _fbs_stock_apply_filters(
        request,
        base_queryset,
        include_sku=False,
    )
    queryset, filters = _fbs_stock_apply_filters(request, base_queryset)
    summary = _fbs_stock_summary(queryset)
    grouped_queryset, sort_key, sort_dir = _fbs_stock_grouped_sort(request, queryset)
    paginator = Paginator(grouped_queryset, 100)
    page_obj = paginator.get_page(request.GET.get("page"))
    page_rows = list(page_obj.object_list)
    page_box_ids = [row["box_id"] for row in page_rows]
    page_balances = list(
        queryset.filter(box_id__in=page_box_ids).order_by(
            "box_id",
            "sku_code",
            "name",
            "size",
            "barcode",
            "goods_type",
            "lot_code",
            "expiry_date",
            "marking_code",
            "id",
        )
    )

    page_params = request.GET.copy()
    page_params.pop("page", None)
    export_params = request.GET.copy()
    export_params.pop("page", None)
    export_url = reverse("fbs:tsd_storekeeper_stock_export")
    if export_params:
        export_url = f"{export_url}?{export_params.urlencode()}"

    marketplace_comparison = None
    marketplace_refresh_url = ""
    if filters["client_id"]:
        comparison_params = request.GET.copy()
        comparison_params.pop("page", None)
        comparison_params["refresh_marketplace"] = "1"
        marketplace_refresh_url = f"{request.path}?{comparison_params.urlencode()}"
        comparison_cache_key = f"fbs:stock-comparison:v1:{filters['client_id']}"
        if str(request.GET.get("refresh_marketplace") or "") == "1":
            cache.delete(comparison_cache_key)
        marketplace_comparison = cache.get(comparison_cache_key)
        if marketplace_comparison is None:
            marketplace_comparison = build_agency_stock_comparison(
                agency_id=int(filters["client_id"]),
            )
            cache.set(comparison_cache_key, marketplace_comparison, timeout=120)

    return render(
        request,
        "fbs/tsd_stock_list.html",
        _base_context(
            request,
            page_title="FBS · Остатки",
            operator_section="stock",
            back_url=reverse("fbs:tsd_storekeeper"),
            box_groups=_fbs_stock_box_groups(page_rows, page_balances),
            page_obj=page_obj,
            page_links=_fbs_stock_pagination_links(page_obj),
            page_query=page_params.urlencode(),
            client_options=_fbs_stock_client_options(),
            sku_options=_fbs_stock_sku_options(option_queryset),
            summary=summary,
            stock_sort_key=sort_key,
            stock_sort_dir=sort_dir,
            stock_sort_urls=_fbs_stock_sort_urls(request, sort_key, sort_dir),
            stock_groups_expanded=bool(filters["q"] or filters["sku"] or filters["box"]),
            export_url=export_url,
            reset_url=reverse("fbs:tsd_storekeeper_stock"),
            marketplace_comparison=marketplace_comparison,
            marketplace_refresh_url=marketplace_refresh_url,
            **filters,
        ),
    )


_XLSX_ILLEGAL_CONTROL_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _xlsx_safe_value(value):
    if not isinstance(value, str):
        return value
    return _XLSX_ILLEGAL_CONTROL_RE.sub(
        lambda match: f"\\x{ord(match.group(0)):02X}",
        value,
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_stock_export(request):
    queryset, _ = _fbs_stock_apply_filters(request, _fbs_stock_base_queryset())
    queryset, _, _ = _fbs_stock_sort(request, queryset)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title="Остатки FBS")
    sheet.append(
        [
            "Клиент",
            "SKU",
            "Наименование",
            "Размер",
            "Штрихкод",
            "Тип товара",
            "Честный знак",
            "Партия",
            "Срок годности",
            "Всего",
            "Доступно",
            "В резерве",
            "Резерв выдачи",
            "Короб FBS",
            "Паллета FBS",
            "Ячейка FBS",
            "Обновлено",
        ]
    )
    for balance in queryset.iterator(chunk_size=1000):
        updated_at = (
            timezone.localtime(balance.updated_at).strftime("%d.%m.%Y %H:%M")
            if balance.updated_at
            else ""
        )
        row = [
            balance.agency.short_name or balance.agency.agn_name or str(balance.agency),
            balance.sku_code,
            balance.name,
            balance.size,
            balance.barcode,
            balance.goods_type,
            balance.marking_code,
            balance.lot_code,
            balance.expiry_date.strftime("%d.%m.%Y") if balance.expiry_date else "",
            int(balance.qty or 0),
            int(balance.available_qty or 0),
            int(balance.reserved_qty or 0),
            int(balance.external_reserved_qty or 0),
            balance.box.box_code,
            balance.box.pallet.pallet_code,
            balance.box.pallet.cell.cell_code,
            updated_at,
        ]
        sheet.append([_xlsx_safe_value(value) for value in row])
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="fbs_stock_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response


def _inventory_queryset():
    return FbsInventorySession.objects.select_related(
        "agency",
        "cell",
        "pallet",
        "box__pallet__cell__location",
        "box__source_container__current_location",
        "sku",
        "first_counter",
        "second_counter",
        "approved_by",
        "assigned_to__employee_profile",
        "recount_assigned_to__employee_profile",
    ).prefetch_related("lines__balance__box__pallet__cell")


def _inventory_line_group_key(line, *, scan_mode: str):
    balance = line.balance
    barcode = str(balance.barcode or "").strip()
    if scan_mode == FbsInventorySession.SCAN_MODE_KIZ:
        return ("line", int(line.id))
    if barcode:
        return ("barcode", barcode.casefold())
    return (
        "item",
        str(balance.sku_code or "").strip().casefold(),
        str(balance.size or "").strip().casefold(),
        str(balance.name or "").strip().casefold(),
    )


def _inventory_approval_groups(lines, *, scan_mode: str) -> list[dict]:
    """Combine barcode-count approval rows without losing their source lines."""
    grouped = {}
    for line in lines:
        balance = line.balance
        group_key = _inventory_line_group_key(line, scan_mode=scan_mode)
        group = grouped.setdefault(
            group_key,
            {
                "key": int(line.id),
                "balance": balance,
                "lines": [],
                "conflict_lines": [],
                "expected_qty": 0,
                "first_count_qty": 0,
                "second_count_qty": 0,
                "fixed_qty": 0,
            },
        )
        first_qty = int(line.first_count_qty or 0)
        second_qty = (
            int(line.second_count_qty)
            if line.second_count_qty is not None
            else first_qty
        )
        is_conflict = (
            line.second_count_qty is not None
            and second_qty != first_qty
        )
        group["lines"].append(line)
        group["expected_qty"] += int(line.expected_qty or 0)
        group["first_count_qty"] += first_qty
        group["second_count_qty"] += second_qty
        if is_conflict:
            group["conflict_lines"].append(line)
        else:
            group["fixed_qty"] += first_qty
    return [group for group in grouped.values() if group["conflict_lines"]]


def _inventory_group_final_counts(group: dict, target_qty: int) -> dict[int, int]:
    """Map one visible barcode total back to line totals deterministically."""
    target_qty = int(target_qty)
    if target_qty < 0:
        raise ValueError("Итоговое количество не может быть отрицательным.")
    editable_target = target_qty - int(group["fixed_qty"])
    if editable_target < 0:
        raise ValueError(
            "Итог по штрих-коду меньше количества в совпавших строках."
        )
    conflict_lines = list(group["conflict_lines"])
    quantities = [int(line.second_count_qty or 0) for line in conflict_lines]
    delta = editable_target - sum(quantities)
    if delta > 0:
        quantities[0] += delta
    elif delta < 0:
        remaining = -delta
        for index in range(len(quantities) - 1, -1, -1):
            reduction = min(quantities[index], remaining)
            quantities[index] -= reduction
            remaining -= reduction
            if not remaining:
                break
        if remaining:
            raise ValueError("Не удалось распределить итог по штрих-коду.")
    return {
        int(line.id): quantity
        for line, quantity in zip(conflict_lines, quantities)
    }


@fbs_module_required
@role_required(*STOREKEEPER_ROLES, "picker", "fbs_controller")
@require_http_methods(["GET", "POST"])
def tsd_inventory(request):
    from .services.inventory_workflow import (
        AGENCY_BATCH_MAX_PALLETS,
        OPEN_STATUSES,
        create_agency_inventory_batch,
        has_inventory_role,
        start_self_kiz_inventory,
    )
    can_dispatch = has_inventory_role(request.user)
    role = get_request_role(request)
    error = ""
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        if action == "self_start":
            if role != "picker":
                return HttpResponseForbidden(
                    "Самостоятельную инвентаризацию ЧЗ может начать только подборщик"
                )
            try:
                session = start_self_kiz_inventory(
                    counted_by=request.user,
                    target_type=request.POST.get("target_type", ""),
                    target_scan=request.POST.get("target_scan", ""),
                )
                return redirect("fbs:tsd_inventory_detail", session_id=session.id)
            except (FbsError, ValueError) as exc:
                error = str(exc)
        elif action == "agency_batch":
            if not can_dispatch:
                return HttpResponseForbidden("Создать инвентаризацию может кладовщик")
            try:
                agency = Agency.objects.get(pk=int(request.POST.get("agency") or 0))
                sessions_created, skipped_busy, skipped_empty = create_agency_inventory_batch(
                    agency=agency,
                    scan_mode=str(
                        request.POST.get("scan_mode") or FbsInventorySession.SCAN_MODE_BARCODE
                    ).strip(),
                    created_by=request.user,
                    limit=request.POST.get("pallet_limit") or 10,
                )
            except (FbsError, ValueError, Agency.DoesNotExist) as exc:
                error = str(exc) or "Не удалось создать задание по клиенту."
            else:
                return redirect(
                    f"{reverse('fbs:tsd_inventory')}?state=unassigned"
                    f"&batch_created={len(sessions_created)}"
                    f"&batch_busy={skipped_busy}&batch_empty={skipped_empty}"
                )
        else:
            if not can_dispatch:
                return HttpResponseForbidden("Создать инвентаризацию может кладовщик")
            target_value = str(request.POST.get("target") or "").strip()
            mode = str(request.POST.get("mode") or "").strip()
            scan_mode = str(
                request.POST.get("scan_mode") or FbsInventorySession.SCAN_MODE_BARCODE
            ).strip()
            try:
                scope_type, raw_object_id = target_value.split(":", 1)
                object_id = int(raw_object_id or 0)
            except (TypeError, ValueError):
                scope_type = ""
                object_id = 0
            kwargs = {}
            model_by_scope = {
                FbsInventorySession.SCOPE_AGENCY: (Agency, "agency"),
                FbsInventorySession.SCOPE_CELL: (FbsStorageCell, "cell"),
                FbsInventorySession.SCOPE_PALLET: (FbsPallet, "pallet"),
                FbsInventorySession.SCOPE_BOX: (FbsBox, "box"),
                FbsInventorySession.SCOPE_SKU: (SKU, "sku"),
            }
            try:
                if not has_inventory_role(request.user, approve=True) and scope_type not in ("box", "cell", "pallet"):
                    raise ValueError("Кладовщик может назначить проверку короба, паллеты или ячейки.")
                if not has_inventory_role(request.user, approve=True) and mode != "drain":
                    raise ValueError("Для кладовщика доступна проверка после завершения активных операций.")
                if scope_type == FbsInventorySession.SCOPE_ALL:
                    pass
                elif scope_type in model_by_scope:
                    model, key = model_by_scope[scope_type]
                    kwargs[key] = model.objects.get(pk=object_id)
                    if scope_type == FbsInventorySession.SCOPE_SKU:
                        kwargs["agency"] = kwargs["sku"].agency
                else:
                    raise ValueError("Выберите поддерживаемую область инвентаризации.")
                session = create_inventory_session(
                    scope_type=scope_type,
                    mode=mode,
                    managed_workflow=mode == "drain" and scope_type in ("box", "cell", "pallet"),
                    scan_mode=scan_mode,
                    created_by=request.user,
                    **kwargs,
                )
                return redirect("fbs:tsd_inventory_detail", session_id=session.id)
            except (
                FbsError,
                ValueError,
                Agency.DoesNotExist,
                FbsStorageCell.DoesNotExist,
                FbsPallet.DoesNotExist,
                FbsBox.DoesNotExist,
                SKU.DoesNotExist,
            ) as exc:
                error = str(exc)
    sessions = _inventory_queryset().annotate(
        issue_count=Count("pick_issues", distinct=True),
        order_count=Count("pick_issues__task__order_id", distinct=True),
        nearest_cutoff=Min("pick_issues__task__order__cutoff_at"),
    )
    if not can_dispatch:
        sessions = sessions.filter(Q(assigned_to=request.user) | Q(recount_assigned_to=request.user))
    inventory_filter = request.GET.get("state", "open")
    if inventory_filter == "closed":
        sessions = sessions.exclude(status__in=OPEN_STATUSES)
    else:
        sessions = sessions.filter(status__in=OPEN_STATUSES)
        if inventory_filter == "unassigned":
            sessions = sessions.filter(Q(assigned_to__isnull=True, status__in=("planned", "draining", "counting")) | Q(recount_assigned_to__isnull=True, status="recount"))
        elif inventory_filter in ("recount", "approval"):
            sessions = sessions.filter(status=inventory_filter)
    context = _base_context(
        request,
        page_title="FBS · Инвентаризация",
        back_url=(
            reverse("fbs:tsd_home")
            if role == "picker"
            else reverse("fbs:tsd_storekeeper")
        ),
        sessions=sessions.order_by(F("nearest_cutoff").asc(nulls_last=True), "created_at")[:100],
        agencies=Agency.objects.filter(fbs_integration_profiles__isnull=False).distinct().order_by("agn_name"),
        cells=FbsStorageCell.objects.filter(is_active=True).order_by("cell_code"),
        pallets=FbsPallet.objects.exclude(status=FbsPallet.STATUS_ARCHIVED).order_by("pallet_code"),
        boxes=FbsBox.objects.exclude(status=FbsBox.STATUS_ARCHIVED).order_by("box_code"),
        skus=SKU.objects.filter(
            agency__fbs_integration_profiles__isnull=False,
            deleted=False,
        ).select_related("agency").distinct().order_by("agency__agn_name", "sku_code")[:1000],
        mode_choices=FbsInventorySession.MODE_CHOICES,
        scan_mode_choices=FbsInventorySession.SCAN_MODE_CHOICES,
        agency_batch_max=AGENCY_BATCH_MAX_PALLETS,
        batch_created=request.GET.get("batch_created") or "",
        batch_busy=request.GET.get("batch_busy") or "",
        batch_empty=request.GET.get("batch_empty") or "",
        can_manage=can_dispatch,
        can_self_start=role == "picker",
        can_approve=has_inventory_role(request.user, approve=True),
        error=error,
    )
    return render(request, "fbs/tsd_inventory_list.html", context, status=400 if error else 200)


@fbs_module_required
@role_required(*STOREKEEPER_ROLES, "picker", "fbs_controller")
@require_http_methods(["GET", "POST"])
def tsd_inventory_detail(request, session_id: int):
    from django.contrib.auth import get_user_model
    from employees.models import Employee
    from .services.inventory_workflow import (
        has_inventory_role, assign_inventory, start_inventory_count,
        ASSIGNABLE_COUNTER_ROLES, ASSIGNABLE_INVENTORY_SCOPES,
    )
    session = get_object_or_404(_inventory_queryset(), pk=session_id)
    can_dispatch = has_inventory_role(request.user)
    if not can_dispatch and request.user.id not in (session.assigned_to_id, session.recount_assigned_to_id):
        return HttpResponseForbidden("Проверка назначена другому сотруднику")
    error = ""
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        try:
            if action == "assign":
                employee_user = get_user_model().objects.filter(pk=int(request.POST.get("assigned_to") or 0)).first()
                assign_inventory(session_id=session.id, assigned_to=employee_user, assigned_by=request.user)
            elif action == "start":
                start_inventory_count(session_id=session.id, counted_by=request.user,
                                      place_scan=request.POST.get("place_scan", ""), box_scan=request.POST.get("box_scan", ""),
                                      container_absent=request.POST.get("container_absent") == "1")
            elif action == "activate":
                if session.managed_workflow:
                    raise ValueError("Начните назначенную проверку сканированием адреса.")
                activate_drained_inventory(session_id=session.id)
            elif action == "scan":
                record_inventory_scan(
                    session_id=session.id,
                    scan_code=request.POST.get("scan_code", ""),
                    counted_by=request.user,
                    box_scan=request.POST.get("box_scan", ""),
                )
            elif action == "finish":
                finish_inventory_count(session_id=session.id, counted_by=request.user,
                                       confirm_complete=request.POST.get("confirm_complete") == "1")
            elif action == "approve":
                if not has_inventory_role(request.user, approve=True):
                    return HttpResponseForbidden("Утверждение доступно только начальнику склада")
                final_counts = {
                    line.id: int(request.POST[f"final_{line.id}"])
                    for line in session.lines.all()
                    if str(request.POST.get(f"final_{line.id}") or "").strip()
                }
                approval_lines = list(
                    session.lines.select_related("balance__box").all()
                )
                for group in _inventory_approval_groups(
                    approval_lines,
                    scan_mode=session.scan_mode,
                ):
                    field_name = f"final_group_{group['key']}"
                    raw_target = str(request.POST.get(field_name) or "").strip()
                    if not raw_target:
                        continue
                    final_counts.update(
                        _inventory_group_final_counts(group, int(raw_target))
                    )
                approve_inventory(
                    session_id=session.id,
                    approved_by=request.user,
                    final_counts=final_counts,
                )
            elif action == "confirm_discrepancies":
                confirm_inventory_discrepancies(
                    session_id=session.id,
                    confirmed_by=request.user,
                    responsibility_acknowledged=(
                        request.POST.get("responsibility_acknowledged") == "1"
                    ),
                )
            else:
                raise ValueError("Неизвестная команда инвентаризации.")
            return redirect("fbs:tsd_inventory_detail", session_id=session.id)
        except (FbsError, ValueError) as exc:
            error = str(exc)
        session = get_object_or_404(_inventory_queryset(), pk=session_id)
    role = get_request_role(request)
    picker_inventory_agency = (
        session.agency
        or (session.box.agency if session.box_id else None)
        or (session.pallet.agency if session.pallet_id else None)
    )
    picker_inventory_client_label = "—"
    if picker_inventory_agency is not None:
        picker_inventory_client_label = (
            str(picker_inventory_agency.short_name or "").strip()
            or str(picker_inventory_agency.agn_name or "").strip()
            or str(picker_inventory_agency)
        )
    latest_inventory_scan = None
    latest_inventory_count = 0
    picker_inventory_scanned_total = 0
    picker_inventory_first_total = 0
    if role == "picker":
        picker_inventory_lines = list(session.lines.all())
        picker_inventory_first_total = sum(
            int(line.first_count_qty or 0)
            for line in picker_inventory_lines
        )
        count_round = (
            FbsInventoryScan.ROUND_SECOND
            if session.status == FbsInventorySession.STATUS_RECOUNT
            or (
                session.status == FbsInventorySession.STATUS_APPROVAL
                and session.second_counter_id == request.user.id
            )
            else FbsInventoryScan.ROUND_FIRST
        )
        latest_inventory_scan = (
            FbsInventoryScan.objects.filter(
                line__session=session,
                counted_by=request.user,
                count_round=count_round,
            )
            .select_related("line__balance")
            .order_by("-created_at", "-id")
            .first()
        )
        if latest_inventory_scan is not None:
            latest_inventory_count = int(
                (
                    latest_inventory_scan.line.second_count_qty
                    if count_round == FbsInventoryScan.ROUND_SECOND
                    else latest_inventory_scan.line.first_count_qty
                )
                or 0
            )
        count_field = (
            "second_count_qty"
            if count_round == FbsInventoryScan.ROUND_SECOND
            else "first_count_qty"
        )
        picker_inventory_scanned_total = sum(
            int(getattr(line, count_field) or 0)
            for line in picker_inventory_lines
        )
    inventory_result_rows = []
    inventory_result_summary = None
    inventory_lines = list(session.lines.all())
    inventory_approval_groups = _inventory_approval_groups(
        inventory_lines,
        scan_mode=session.scan_mode,
    )
    first_result_ready = bool(inventory_lines) and all(
        line.first_count_qty is not None for line in inventory_lines
    )
    if can_dispatch and first_result_ready:
        second_result_ready = all(
            line.second_count_qty is not None for line in inventory_lines
        )
        final_result_ready = all(
            line.final_qty is not None for line in inventory_lines
        )
        for line in inventory_lines:
            expected_qty = int(line.expected_qty or 0)
            first_qty = int(line.first_count_qty or 0)
            second_qty = (
                int(line.second_count_qty)
                if line.second_count_qty is not None
                else None
            )
            final_qty = (
                int(line.final_qty)
                if line.final_qty is not None
                else None
            )
            inventory_result_rows.append(
                {
                    "line": line,
                    "expected_qty": expected_qty,
                    "first_qty": first_qty,
                    "first_delta": first_qty - expected_qty,
                    "second_qty": second_qty,
                    "second_delta": (
                        second_qty - expected_qty
                        if second_qty is not None
                        else None
                    ),
                    "final_qty": final_qty,
                    "final_delta": (
                        final_qty - expected_qty
                        if final_qty is not None
                        else None
                    ),
                }
            )
        if session.scan_mode != FbsInventorySession.SCAN_MODE_KIZ:
            grouped_result_rows = {}
            for row in inventory_result_rows:
                group_key = _inventory_line_group_key(
                    row["line"],
                    scan_mode=session.scan_mode,
                )
                grouped_row = grouped_result_rows.get(group_key)
                if grouped_row is None:
                    grouped_row = dict(row)
                    grouped_row["line_count"] = 1
                    grouped_result_rows[group_key] = grouped_row
                    continue
                grouped_row["line_count"] += 1
                for field in ("expected_qty", "first_qty", "first_delta"):
                    grouped_row[field] += row[field]
                for field in ("second_qty", "second_delta", "final_qty", "final_delta"):
                    if grouped_row[field] is not None and row[field] is not None:
                        grouped_row[field] += row[field]
                    else:
                        grouped_row[field] = None
            inventory_result_rows = list(grouped_result_rows.values())
        expected_total = sum(row["expected_qty"] for row in inventory_result_rows)
        first_total = sum(row["first_qty"] for row in inventory_result_rows)
        second_total = (
            sum(row["second_qty"] for row in inventory_result_rows)
            if second_result_ready
            else None
        )
        final_total = (
            sum(row["final_qty"] for row in inventory_result_rows)
            if final_result_ready
            else None
        )
        inventory_result_summary = {
            "expected_total": expected_total,
            "first_total": first_total,
            "first_delta": first_total - expected_total,
            "second_total": second_total,
            "second_delta": (
                second_total - expected_total
                if second_total is not None
                else None
            ),
            "final_total": final_total,
            "final_delta": (
                final_total - expected_total
                if final_total is not None
                else None
            ),
        }
    return render(
        request,
        "fbs/tsd_inventory_detail.html",
        _base_context(
            request,
            page_title=f"FBS · Инвентаризация #{session.id}",
            back_url=reverse("fbs:tsd_inventory"),
            inventory=session,
            picker_inventory_client_label=picker_inventory_client_label,
            picker_inventory_box_label=(
                session.box.box_code
                if session.box_id
                else session.workflow_target_label
            ),
            latest_inventory_scan=latest_inventory_scan,
            latest_inventory_count=latest_inventory_count,
            picker_inventory_scanned_total=picker_inventory_scanned_total,
            picker_inventory_first_total=picker_inventory_first_total,
            inventory_result_rows=inventory_result_rows,
            inventory_result_summary=inventory_result_summary,
            inventory_approval_groups=inventory_approval_groups,
            hide_employee_name=role == "picker",
            hide_inventory_assignment_link=role == "picker",
            inventory_place_label=(
                fbs_box_physical_location_label(session.box)
                if session.box_id
                else session.workflow_place_label
            ),
            can_manage=has_inventory_role(request.user, approve=True),
            can_dispatch=can_dispatch,
            counter_employees=Employee.objects.filter(
                is_active=True,
                user__is_active=True,
                role__in=ASSIGNABLE_COUNTER_ROLES,
            ).select_related("user").order_by("full_name", "user__username") if can_dispatch else [],
            assigned_counter=session.recount_assigned_to if session.status == "recount" else session.assigned_to,
            count_started=session.recount_started_at if session.status == "recount" else session.count_started_at,
            can_assign_inventory=(
                can_dispatch
                and session.scope_type in ASSIGNABLE_INVENTORY_SCOPES
                and session.status in ("planned", "draining", "counting", "recount")
                and not (
                    session.recount_started_at
                    if session.status == "recount"
                    else session.count_started_at
                )
                and (session.managed_workflow or not session.first_counter_id)
            ),
            assignment_scope_supported=(
                session.scope_type in ASSIGNABLE_INVENTORY_SCOPES
            ),
            is_counter=request.user.id == (session.recount_assigned_to_id if session.status == "recount" else session.assigned_to_id),
            self_started_inventory=(
                session.managed_workflow
                and session.scan_mode == FbsInventorySession.SCAN_MODE_KIZ
                and session.created_by_id == session.assigned_to_id
                and not session.pick_issues.exists()
            ),
            inventory_issues=session.pick_issues.select_related("task__order", "allocation__balance", "created_by").order_by("created_at"),
            inventory_events=session.work_events.select_related("actor").order_by("-created_at")[:50] if can_dispatch else [],
            error=error,
        ),
        status=400 if error else 200,
    )


def _handover_queryset():
    list_items = FbsOrderItem.objects.select_related(
        "sku", "order__profile"
    ).prefetch_related(
        Prefetch(
            "metadata_transfers",
            queryset=FbsMarketplaceMetadataTransfer.objects.select_related(
                "order_item__order__profile"
            ).only(
                "id",
                "order_item_id",
                "metadata_type",
                "status",
                "is_required",
                "external_status",
                "last_error",
                "order_item__order__profile__marketplace",
            ),
        )
    )
    list_assignments = (
        FbsHandoverOrderAssignment.objects.exclude(
            status=FbsHandoverOrderAssignment.STATUS_CANCELED
        )
        .select_related("order")
        .prefetch_related(
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.only("id", "order_id", "status"),
            ),
            Prefetch("order__items", queryset=list_items),
        )
        .order_by("id")
    )
    list_box_orders = (
        FbsHandoverOrder.objects.filter(status=FbsHandoverOrder.STATUS_ACTIVE)
        .select_related("order")
        .prefetch_related(
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.only("id", "order_id", "status"),
            ),
            Prefetch("order__items", queryset=list_items),
        )
        .order_by("id")
    )
    list_boxes = FbsHandoverBox.objects.prefetch_related(
        Prefetch("orders", queryset=list_box_orders)
    ).order_by("id")
    list_commands = FbsMarketplaceCommand.objects.only(
        "id", "handover_batch_id", "order_id", "status"
    ).order_by("id")
    return FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by", "dispatch_location"
    ).prefetch_related(
        Prefetch("boxes", queryset=list_boxes, to_attr="list_boxes"),
        Prefetch(
            "order_assignments",
            queryset=list_assignments,
            to_attr="list_assignments",
        ),
        Prefetch(
            "marketplace_commands",
            queryset=list_commands,
            to_attr="list_commands",
        ),
    )


def _handover_relevant_commands(commands, *, order_ids):
    current_order_ids = set(order_ids)
    return [
        command
        for command in commands
        if command.order_id is None or command.order_id in current_order_ids
    ]


def _prepare_handover_list_rows(batches):
    from .handover_presentation import handover_metadata_issues
    from .services.traceability import (
        marketplace_metadata_transfer_resolved,
        metadata_requirements,
        ozon_marking_transfer_not_required,
    )

    for batch in batches:
        boxes = list(batch.list_boxes)
        assignments = list(batch.list_assignments)
        commands = list(batch.list_commands)
        boxed_links = [link for box in boxes for link in box.orders.all()]
        orders_by_id = {
            assignment.order_id: assignment.order for assignment in assignments
        }
        orders_by_id.update({link.order_id: link.order for link in boxed_links})
        batch.operator_metadata_issues = handover_metadata_issues(orders_by_id.values())
        batch.list_unit_count = 0
        batch.list_weight_kg = Decimal("0")
        batch.list_weight_has_value = False
        batch.list_weight_complete = bool(orders_by_id)
        for order in orders_by_id.values():
            for item in order.items.all():
                quantity = int(item.quantity or 0)
                batch.list_unit_count += quantity
                unit_weight = None
                if item.sku is not None:
                    unit_weight = item.sku.weight_gross_kg or item.sku.weight_kg
                if unit_weight is None:
                    batch.list_weight_complete = False
                    continue
                batch.list_weight_has_value = True
                batch.list_weight_kg += Decimal(unit_weight) * quantity
        assigned_order_ids = {assignment.order_id for assignment in assignments}
        boxed_order_ids = {link.order_id for link in boxed_links}
        transfers = [
            transfer
            for order in orders_by_id.values()
            for item in order.items.all()
            for transfer in item.metadata_transfers.all()
        ]
        metadata_missing_count = 0
        metadata_unconfirmed_count = 0
        metadata_blocked_order_numbers = []
        for order in orders_by_id.values():
            order_has_unconfirmed_metadata = False
            for item in order.items.all():
                transfer_by_type = {
                    transfer.metadata_type: transfer
                    for transfer in item.metadata_transfers.all()
                }
                requirements = metadata_requirements(item)
                required_types = set()
                marking_transfer = transfer_by_type.get(
                    FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                )
                if (
                    requirements.marketplace_marking_required
                    or (
                        marking_transfer is not None
                        and marking_transfer.is_required
                        and not ozon_marking_transfer_not_required(
                            marking_transfer
                        )
                    )
                ):
                    required_types.add(
                        FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                    )
                if requirements.expiry_required:
                    required_types.add(
                        FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION
                    )
                for metadata_type in required_types:
                    transfer = transfer_by_type.get(metadata_type)
                    if marketplace_metadata_transfer_resolved(transfer):
                        continue
                    metadata_unconfirmed_count += 1
                    metadata_missing_count += int(transfer is None)
                    order_has_unconfirmed_metadata = True
            if order_has_unconfirmed_metadata:
                metadata_blocked_order_numbers.append(order.external_order_id)
        batch.box_count = len(boxes)
        batch.scanned_box_count = sum(
            box.status
            in {
                FbsHandoverBox.STATUS_SCANNED,
                FbsHandoverBox.STATUS_DISPATCHED,
                FbsHandoverBox.STATUS_ACCEPTED,
            }
            for box in boxes
        )
        batch.assigned_order_count = len(assigned_order_ids)
        batch.boxed_order_count = len(boxed_order_ids)
        batch.label_ready_count = sum(
            any(
                label.status
                in {FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED}
                for label in order.marketplace_labels.all()
            )
            for order in orders_by_id.values()
        )
        batch.metadata_waiting_count = sum(
            transfer.status
            in {
                FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
                FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
                FbsMarketplaceMetadataTransfer.STATUS_SENT,
                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
            }
            for transfer in transfers
        )
        batch.metadata_problem_count = sum(
            transfer.status
            in {
                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED,
                FbsMarketplaceMetadataTransfer.STATUS_CANCELED,
            }
            and not marketplace_metadata_transfer_resolved(transfer)
            for transfer in transfers
        )
        batch.metadata_missing_count = metadata_missing_count
        batch.metadata_unconfirmed_count = metadata_unconfirmed_count
        batch.metadata_blocked_order_numbers = tuple(
            dict.fromkeys(metadata_blocked_order_numbers)
        )
        batch.supply_print_blocked = bool(
            batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and batch.metadata_unconfirmed_count
        )
        batch.supply_print_block_reason = (
            "КИЗ/срок не подтверждены: "
            f"{batch.metadata_unconfirmed_count}"
            + (
                " · заказы "
                + ", ".join(batch.metadata_blocked_order_numbers[:3])
                if batch.metadata_blocked_order_numbers
                else ""
            )
        )
        relevant_commands = _handover_relevant_commands(
            commands,
            order_ids=orders_by_id,
        )
        batch.command_active_count = sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_PENDING,
                FbsMarketplaceCommand.STATUS_SENT,
                FbsMarketplaceCommand.STATUS_RETRY,
            }
            for command in relevant_commands
        )
        batch.command_problem_count = sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_FAILED,
                FbsMarketplaceCommand.STATUS_CONFLICT,
            }
            for command in relevant_commands
        )
        batch.list_order_count = max(
            int(batch.assigned_order_count or 0),
            int(batch.boxed_order_count or 0),
        )
        batch.supply_label_available = bool(batch.supply_label_file) or bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and batch.list_order_count
            and batch.box_count
            and batch.status != FbsHandoverBatch.STATUS_OPEN
        )
        stages = _handover_stage_rows(batch)
        batch.current_stage = next(
            (stage for stage in stages if stage.state in {"in_progress", "problem"}),
            stages[-1],
        )
        if batch.metadata_missing_count:
            batch.metadata_state = "problem"
            batch.metadata_label = (
                f"Не подтверждены: {batch.metadata_unconfirmed_count}"
            )
        elif batch.metadata_problem_count:
            batch.metadata_state = "problem"
            batch.metadata_label = f"Ошибок: {batch.metadata_problem_count}"
        elif batch.metadata_waiting_count:
            batch.metadata_state = "in_progress"
            batch.metadata_label = f"Ожидают: {batch.metadata_waiting_count}"
        else:
            batch.metadata_state = "done" if batch.list_order_count else "pending"
            batch.metadata_label = "Без ошибок" if batch.list_order_count else "Нет заказов"
        batch.has_problem = bool(
            batch.status != FbsHandoverBatch.STATUS_ARCHIVED
            and (
                batch.status == FbsHandoverBatch.STATUS_PROBLEM
                or batch.marketplace_state == FbsHandoverBatch.MARKETPLACE_ERROR
                or batch.metadata_missing_count
                or batch.metadata_problem_count
                or batch.command_problem_count
            )
        )
    from .controller_shipment_ui import decorate_controller_shipments
    return decorate_controller_shipments(batches)


HANDOVER_REPORT_CHOICES = (
    ("status", "Статус сборки"),
    ("count", "Количество заказов"),
    ("orders", "Отгруженные заказы"),
    ("goods", "Отгруженные товары"),
)
HANDOVER_REPORT_ROW_LIMIT = 50_000


def _handover_report_filter_context(request):
    today = timezone.localdate()
    report_type = str(request.GET.get("report") or "status").strip()
    if report_type not in dict(HANDOVER_REPORT_CHOICES):
        report_type = "status"
    agency_filter = str(request.GET.get("agency_id") or "").strip()
    marketplace_filter = str(request.GET.get("marketplace") or "").strip()
    date_from_filter = str(request.GET.get("date_from") or today.replace(day=1).isoformat()).strip()
    date_to_filter = str(request.GET.get("date_to") or today.isoformat()).strip()
    filter_error = ""
    try:
        date_from = date.fromisoformat(date_from_filter)
        date_to = date.fromisoformat(date_to_filter)
    except ValueError:
        date_from = today.replace(day=1)
        date_to = today
        filter_error = "Укажите даты в корректном формате."
    if date_from > date_to:
        filter_error = "Дата начала не может быть позже даты окончания."
    if not agency_filter.isdigit():
        agency_filter = ""
    if marketplace_filter not in dict(FbsIntegrationProfile.MARKETPLACE_CHOICES):
        marketplace_filter = ""
    return {
        "report_type": report_type,
        "agency_filter": agency_filter,
        "marketplace_filter": marketplace_filter,
        "date_from_filter": date_from_filter,
        "date_to_filter": date_to_filter,
        "date_from": date_from,
        "date_to": date_to,
        "filter_error": filter_error,
    }


def _handover_report_xlsx(*, title, headers, rows, warning=""):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=title[:31])
    if warning:
        sheet.append((warning,))
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        'attachment; filename="fullbox-fbs-shipments-report.xlsx"'
    )
    return response


def _handover_detail_queryset():
    order_items = FbsOrderItem.objects.select_related("sku").prefetch_related(
        Prefetch(
            "metadata_transfers",
            queryset=FbsMarketplaceMetadataTransfer.objects.select_related(
                "traceability__allocation__pick_task__batch__workstation"
            ).order_by("metadata_type", "id"),
        )
    )
    handover_orders = (
        FbsHandoverOrder.objects.select_related(
            "order", "added_by", "excluded_by", "verified_by", "verified_label"
        )
        .prefetch_related(
            Prefetch("order__items", queryset=order_items),
            Prefetch(
                "order__marketplace_labels",
                queryset=FbsOrderLabel.objects.order_by("-requested_at", "-id"),
            ),
        )
        .order_by("id")
    )
    boxes = (
        FbsHandoverBox.objects.select_related("scanned_by")
        .prefetch_related(Prefetch("orders", queryset=handover_orders))
        .order_by("id")
    )
    return FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by"
    ).prefetch_related(
        Prefetch("boxes", queryset=boxes),
        Prefetch(
            "order_assignments",
            queryset=FbsHandoverOrderAssignment.objects.select_related(
                "order", "assigned_by"
            )
            .prefetch_related(
                Prefetch("order__items", queryset=order_items),
                Prefetch(
                    "order__marketplace_labels",
                    queryset=FbsOrderLabel.objects.order_by("-requested_at", "-id"),
                ),
            )
            .order_by("id"),
        ),
        Prefetch(
            "marketplace_commands",
            queryset=FbsMarketplaceCommand.objects.select_related(
                "order", "handover_box", "requested_by"
            ).order_by("-created_at", "-id"),
        ),
    )


def _handover_stage_rows(batch: FbsHandoverBatch) -> tuple[SimpleNamespace, ...]:
    from .controller_shipment_ui import controller_shipment_stage
    stage = controller_shipment_stage(batch)
    current = {
        'checking': 2, 'checked': 3, 'ready': 4, 'transit': 5, 'accepted': 6,
        'problem': 4 if batch.supply_label_file else 3,
    }.get(stage.key, 1)
    labels = (
        "Создана",
        "Проверяется",
        "Проверено",
        "Готово к отгрузке",
        "В пути",
        "Принято маркетплейсом",
    )
    rows = []
    for position, label in enumerate(labels, start=1):
        if batch.status == FbsHandoverBatch.STATUS_ARCHIVED:
            state = "done" if position == 1 else "pending"
            state_label = "Архив" if position == 1 else "Не выполняется"
        elif batch.status == FbsHandoverBatch.STATUS_PROBLEM and position == current:
            state, state_label = "problem", "Требует решения"
        elif position < current or (
            batch.status == FbsHandoverBatch.STATUS_ACCEPTED and position == current
        ):
            state, state_label = "done", "Готово"
        elif position == current:
            state, state_label = "in_progress", "Текущий этап"
        else:
            state, state_label = "pending", "Ожидает"
        rows.append(
            SimpleNamespace(
                number=position,
                label=label,
                state=state,
                state_label=state_label,
            )
        )
    return tuple(rows)


def _handover_actor_name(user) -> str:
    if user is None:
        return "Система"
    full_name = str(user.get_full_name() or "").strip()
    return full_name or str(user.get_username() or "Система")


def _handover_history_rows(
    batch: FbsHandoverBatch,
    *,
    boxes,
    assignments,
    commands,
) -> tuple[tuple[SimpleNamespace, ...], int]:
    rows = []
    background_reads = {}

    def add(occurred_at, title, detail="", *, user=None, state="pending"):
        if occurred_at is None:
            return
        rows.append(
            SimpleNamespace(
                occurred_at=occurred_at,
                title=title,
                detail=detail,
                actor=_handover_actor_name(user),
                state=state,
            )
        )

    add(
        batch.created_at,
        "Создана отгрузка",
        f"{batch.profile.get_marketplace_display()} · склад "
        f"{batch.profile.external_warehouse_id or 'не указан'}",
        user=batch.created_by,
        state="done",
    )
    if batch.archived_at:
        add(
            batch.archived_at,
            "Отгрузка перенесена в архив",
            batch.archive_reason,
            user=batch.archived_by,
            state="done",
        )
    for assignment in assignments:
        order_number = assignment.order.external_order_id
        add(
            assignment.created_at,
            f"Заказ {order_number} добавлен в состав",
            assignment.get_status_display(),
            user=assignment.assigned_by,
            state="done" if assignment.status == assignment.STATUS_CONFIRMED else "in_progress",
        )
        if assignment.confirmed_at:
            add(
                assignment.confirmed_at,
                f"Marketplace подтвердил заказ {order_number}",
                "Заказ включен в поставку",
                state="done",
            )
        elif assignment.status in {
            assignment.STATUS_ERROR,
            assignment.STATUS_CANCELED,
        }:
            add(
                assignment.updated_at,
                f"Заказ {order_number}: {assignment.get_status_display()}",
                assignment.error,
                user=assignment.assigned_by,
                state="problem",
            )

    for box in boxes:
        add(
            box.created_at,
            f"Создан короб {box.qr_code}",
            box.external_box_id,
            state="done",
        )
        if box.label_ready_at:
            add(
                box.label_ready_at,
                f"Готова этикетка короба {box.qr_code}",
                box.label_format.upper() if box.label_format else "",
                state="done",
            )
        if box.scanned_at:
            add(
                box.scanned_at,
                f"Проверен короб {box.qr_code}",
                "Состав подтвержден сканированием",
                user=box.scanned_by,
                state="done",
            )
        if box.accepted_at:
            add(
                box.accepted_at,
                f"Marketplace принял короб {box.qr_code}",
                state="done",
            )
        for link in box.orders.all():
            order_number = link.order.external_order_id
            add(
                link.added_at,
                f"Заказ {order_number} помещен в короб",
                box.qr_code,
                user=link.added_by,
                state="done",
            )
            if link.verified_at:
                add(
                    link.verified_at,
                    f"Проверен заказ {order_number}",
                    "Этикетка заказа подтверждена сканированием",
                    user=link.verified_by,
                    state="done",
                )
            if link.excluded_at:
                add(
                    link.excluded_at,
                    f"Заказ {order_number} исключен",
                    link.exclusion_reason,
                    user=link.excluded_by,
                    state="problem",
                )

    for command in commands:
        if command.status not in {
            command.STATUS_CONFIRMED,
            command.STATUS_FAILED,
            command.STATUS_CONFLICT,
            command.STATUS_CANCELLED,
        }:
            continue
        order_suffix = (
            f" · заказ {command.order.external_order_id}" if command.order_id else ""
        )
        occurred_at = command.confirmed_at or command.updated_at
        if (
            command.status == command.STATUS_CONFIRMED
            and str(command.command_type or "").startswith(("wb_read_", "ozon_read_"))
        ):
            key = (command.command_type, command.order_id)
            group = background_reads.setdefault(
                key,
                {
                    "count": 0,
                    "occurred_at": occurred_at,
                    "command": command,
                    "order_suffix": order_suffix,
                },
            )
            group["count"] += 1
            if occurred_at and (
                group["occurred_at"] is None or occurred_at > group["occurred_at"]
            ):
                group["occurred_at"] = occurred_at
                group["command"] = command
                group["order_suffix"] = order_suffix
            continue
        add(
            occurred_at,
            f"Marketplace: {command.get_status_display()}",
            f"{command.command_type}{order_suffix}"
            + (f" · {command.error}" if command.error else ""),
            user=command.requested_by,
            state=(
                "done"
                if command.status == command.STATUS_CONFIRMED
                else "problem"
            ),
        )

    for group in background_reads.values():
        command = group["command"]
        count = int(group["count"] or 0)
        detail = f"{command.command_type}{group['order_suffix']}"
        if count > 1:
            detail += f" · {count} фоновых проверок"
        add(
            group["occurred_at"],
            "Marketplace: фоновая проверка подтверждена",
            detail,
            user=command.requested_by,
            state="done",
        )

    if batch.dispatched_at:
        add(
            batch.dispatched_at,
            "Отгрузка передана водителю",
            batch.external_supply_id,
            user=batch.dispatched_by,
            state="done",
        )
    if batch.accepted_at:
        add(
            batch.accepted_at,
            "Отгрузка принята маркетплейсом",
            batch.external_supply_id,
            state="done",
        )
    elif batch.status == FbsHandoverBatch.STATUS_PROBLEM:
        add(
            batch.updated_at,
            "Отгрузка требует решения",
            batch.get_marketplace_state_display(),
            state="problem",
        )

    rows.sort(key=lambda row: row.occurred_at, reverse=True)
    total = len(rows)
    return tuple(rows[:300]), total


def _handover_next_action(
    batch: FbsHandoverBatch,
    *,
    boxes,
    assignment_count: int,
    missing_order_count: int,
    physical_missing_order_count: int = 0,
    unverified_order_count: int = 0,
    all_assignments_confirmed: bool,
    all_orders_ready: bool,
    composition_ready: bool,
    blocking_return_count: int = 0,
    blocking_return_order_numbers=(),
    blocking_return_statuses=(),
    background_return_count: int = 0,
    background_return_order_numbers=(),
) -> SimpleNamespace:
    open_boxes = [box for box in boxes if box.status == FbsHandoverBox.STATUS_OPEN]
    filled_open_boxes = [box for box in open_boxes if box.order_count]
    empty_open_boxes = [box for box in open_boxes if not box.order_count]
    closed_boxes = [box for box in boxes if box.status == FbsHandoverBox.STATUS_CLOSED]
    blocking_return_order_numbers = tuple(
        number
        for number in dict.fromkeys(
            str(number).strip() for number in blocking_return_order_numbers
        )
        if number
    )
    blocking_return_statuses = {
        str(status).strip() for status in blocking_return_statuses if str(status).strip()
    }
    background_return_order_numbers = tuple(
        number
        for number in dict.fromkeys(
            str(number).strip() for number in background_return_order_numbers
        )
        if number
    )

    def background_return_copy():
        numbers = ", ".join(background_return_order_numbers[:5])
        suffix = f" Заказы: {numbers}." if numbers else ""
        return (
            "Проверка продолжена — ждём сверку WB",
            (
                "Контролёр может перейти к следующей работе. Финальная отправка "
                "поставки автоматически останется закрыта до подтверждения WB."
                f"{suffix}"
            ),
        )

    def blocking_return_copy(*, ready: bool):
        if FbsPickRestockRequest.STATUS_FAILED in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Положите товар в назначенную тару отменённых заказов, "
                    "отсканируйте её и продолжайте проверку. Финальная отправка "
                    "останется заблокирована до успешной сверки WB."
                ),
            )
        if FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Положите товар в назначенную тару отменённых заказов и "
                    "отсканируйте её. Сверка WB продолжится в фоне."
                ),
            )
        if FbsPickRestockRequest.STATUS_QUEUED in blocking_return_statuses:
            if len(blocking_return_order_numbers) == 1:
                title = f"Снимите заказ {blocking_return_order_numbers[0]}"
            else:
                title = "Снимите проблемные заказы"
            return (
                title,
                (
                    "Контролёру: положите весь заказ в тару отменённых заказов "
                    "и подтвердите передачу подборщику."
                ),
            )
        if len(blocking_return_order_numbers) == 1:
            order_number = blocking_return_order_numbers[0]
            title = f"Завершить возврат заказа {order_number}"
            if ready:
                description = (
                    f"После физического возврата заказа {order_number} "
                    "он перестанет блокировать состав."
                )
            else:
                description = (
                    f"Подборщик должен вернуть товар заказа {order_number} "
                    "в исходный FBS-короб."
                )
            return title, description
        if blocking_return_order_numbers:
            visible_numbers = ", ".join(blocking_return_order_numbers[:5])
            remaining_count = len(blocking_return_order_numbers) - 5
            suffix = f" и ещё {remaining_count}" if remaining_count > 0 else ""
            return (
                f"Завершить возврат проблемных заказов: {len(blocking_return_order_numbers)}",
                f"Заказы: {visible_numbers}{suffix}. Верните товар в исходные FBS-короба.",
            )
        return (
            "Завершить возврат проблемного заказа",
            (
                "После физического возврата заказ перестанет блокировать состав."
                if ready
                else "Подборщик должен вернуть товар в исходный FBS-короб."
            ),
        )

    def action(code, step, title, description, *, box=None, blocked=False):
        return SimpleNamespace(
            code=code,
            step=step,
            title=title,
            description=description,
            box=box,
            blocked=blocked,
        )

    if batch.status == FbsHandoverBatch.STATUS_ARCHIVED:
        return action(
            "archived",
            1,
            "Отгрузка в архиве",
            "Активных действий нет. История marketplace сохранена.",
        )

    if batch.status == FbsHandoverBatch.STATUS_ACCEPTED:
        return action(
            "complete",
            6,
            "Отгрузка завершена",
            "Marketplace подтвердил приемку. Дополнительных действий не требуется.",
        )
    if batch.status in {
        FbsHandoverBatch.STATUS_DISPATCHED,
        FbsHandoverBatch.STATUS_PROBLEM,
    }:
        return action(
            "refresh",
            6,
            "Проверить приемку marketplace",
            "Обновите статусы после передачи. Ошибки будут показаны отдельно.",
        )
    if batch.status == FbsHandoverBatch.STATUS_READY:
        if blocking_return_count:
            title, description = blocking_return_copy(ready=True)
            return action(
                "await_return",
                2,
                title,
                description,
                blocked=True,
            )
        if background_return_count:
            title, description = background_return_copy()
            return action(
                "wait_marketplace_exclusion",
                3,
                title,
                description,
                blocked=True,
            )
        if not composition_ready:
            return action(
                "resolve_composition",
                2,
                "Устранить ошибки состава",
                "Проверка откроется после подтверждения всех заказов, этикеток, КИЗов и сроков.",
                blocked=True,
            )
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and unverified_order_count
        ):
            return action(
                "verify_orders",
                2,
                "Проверить состав сканированием",
                f"Отсканируйте WB-этикетки оставшихся заказов: {unverified_order_count}.",
            )
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and batch.marketplace_state != FbsHandoverBatch.MARKETPLACE_COMPLETE
        ):
            return action(
                "wait_marketplace_delivery",
                4,
                "Автоматическая отправка в WB",
                "Поставка проверена. Система автоматически передает ее в доставку Wildberries.",
            )
        if batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
            if not batch.supply_label_file:
                return action(
                    "wait_supply_label",
                    4,
                    "Дождаться ШК поставки",
                    "Поставка проверена. Статус «Готово к отгрузке» появится после получения ШК от WB.",
                    blocked=True,
                )
            return action(
                "print_supply_label",
                4,
                "Распечатать ШК поставки",
                "Печать нужна для внутренней работы и не подтверждает передачу товара.",
            )
        return action(
            "print_supply_label",
            4,
            "Распечатать ШК поставки",
            "Печать нужна для внутренней работы и не подтверждает передачу товара.",
        )

    if batch.status != FbsHandoverBatch.STATUS_OPEN:
        return action(
            "refresh_page",
            2,
            "Обновить состояние отгрузки",
            "Статус изменился. Обновите экран перед следующим действием.",
        )

    if blocking_return_count:
        title, description = blocking_return_copy(ready=False)
        return action(
            "await_return",
            2,
            title,
            description,
            blocked=True,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
        and assignment_count
        and not physical_missing_order_count
        and not composition_ready
    ):
        return action(
            "resolve_composition",
            2,
            "Ozon проверяет заказы в фоне",
            (
                "Заказы уже добавлены в транспортный короб по первичному скану. "
                "Повторно сканировать QR не нужно; дождитесь официальных этикеток "
                "или устраните показанную ошибку Ozon."
            ),
            blocked=True,
        )

    if assignment_count and missing_order_count:
        target_box = filled_open_boxes[0] if filled_open_boxes else (
            open_boxes[0] if open_boxes else None
        )
        return action(
            "pack_orders",
            2,
            f"Контрольный скан заказов: {missing_order_count}",
            (
                "Откройте тару проверки. Сканируйте WB-этикетку каждого "
                "действующего заказа и сразу перекладывайте товар в "
                "транспортный короб. После последнего скана закройте тару."
                if batch.profile.marketplace
                == FbsIntegrationProfile.MARKETPLACE_WB
                else
                "Для заказов, подтверждённых до включения одного скана, "
                "отсканируйте QR заказа Ozon при фактическом помещении в "
                "транспортный короб. Новые заказы повторного скана не требуют."
            ),
            box=target_box,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and unverified_order_count
    ):
        if not composition_ready and not background_return_count:
            return action(
                "resolve_composition",
                2,
                "Устранить ошибки состава",
                "Проверка откроется после подтверждения всех заказов, этикеток, КИЗов и сроков.",
                blocked=True,
            )
        return action(
            "verify_orders",
            2,
            "Проверить состав сканированием",
            f"Отсканируйте WB-этикетки оставшихся заказов: {unverified_order_count}.",
        )

    if empty_open_boxes and assignment_count and not missing_order_count:
        box = empty_open_boxes[0]
        return action(
            "blocked_empty_box",
            2,
            "Пустой короб блокирует отгрузку",
            f"Короб {box.qr_code} создан, но в нем нет заказа. Уберите лишний короб до отправки.",
            box=box,
            blocked=True,
        )

    if (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and boxes
        and all_assignments_confirmed
        and all_orders_ready
        and not background_return_count
    ):
        return action(
            "wait_marketplace_delivery",
            3,
            "Автоматическая отправка в WB",
            "Состав проверен. Система автоматически передает поставку в доставку Wildberries.",
        )

    if filled_open_boxes:
        box = filled_open_boxes[0]
        return action(
            "close_box",
            2,
            "Закрыть заполненный короб",
            f"Заказы уже находятся в коробе {box.qr_code}. Закройте его, чтобы открыть подтверждение к отгрузке.",
            box=box,
        )

    if background_return_count:
        title, description = background_return_copy()
        return action(
            "wait_marketplace_exclusion",
            3,
            title,
            description,
            blocked=True,
        )

    if closed_boxes:
        box = closed_boxes[0]
        return action(
            "scan_box",
            3,
            "Подтвердить закрытый короб",
            f"Короб {box.qr_code} закрыт. Отсканируйте его QR для подтверждения к отгрузке.",
            box=box,
        )

    if empty_open_boxes:
        box = empty_open_boxes[0]
        return action(
            "pack_orders",
            1,
            "Добавить заказ в открытый короб",
            f"Короб {box.qr_code} открыт. Отсканируйте этикетку следующего заказа.",
            box=box,
        )

    if not boxes:
        if (
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and wb_handover_uses_marketplace_boxes(batch)
        ):
            if all_assignments_confirmed and all_orders_ready:
                return action(
                    "create_wb_boxes",
                    1,
                    "Получить QR коробов WB",
                    "Состав поставки подтвержден. Укажите нужное количество транспортных коробов.",
                )
            return action(
                "wait_marketplace",
                1,
                "Дождаться подтверждения состава",
                "Короба можно создать после подтверждения всех заказов marketplace.",
            )
        if batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
            return action(
                "wait_marketplace",
                1,
                "Дождаться транспортного короба",
                "Для поставки на склад или СЦ короб Fullbox появится после подготовки заказа.",
            )
        return action(
            "add_box",
            1,
            "Добавить транспортный короб",
            "Сканируйте QR пустого короба, в который будут складываться проверенные заказы.",
        )

    return action(
        "refresh_page",
        3,
        "Обновить состояние отгрузки",
        "Все доступные короба обработаны. Обновите экран для следующего шага.",
    )


def _handover_restock_blocks_handover(request) -> bool:
    if request.source_tote_id is not None:
        return False
    return bool(
        request.status in HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES
        or (
            request.status == FbsPickRestockRequest.STATUS_QUEUED
            and request.source_tote_id is None
        )
    )


def _handover_assignment_display_priority(assignment, exclusion_by_order_id):
    exclusion_request = exclusion_by_order_id.get(assignment.order_id)
    if exclusion_request is not None and _handover_restock_blocks_handover(
        exclusion_request
    ):
        return (0, exclusion_request.id)
    return (1, assignment.id)


def _handover_exclusion_check_copy(exclusion_request):
    if exclusion_request.status == FbsPickRestockRequest.STATUS_FAILED:
        return (
            "Снимите заказ и освободите проверку",
            (
                "Контролёру: положить весь заказ в назначенную тару отменённых "
                "заказов, отсканировать её и продолжить проверку. Сверка WB "
                "останется отдельным финальным барьером."
            ),
        )
    if exclusion_request.status == FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE:
        return (
            "Снимите заказ и продолжите проверку",
            (
                "Контролёру: положить весь заказ в назначенную тару отменённых "
                "заказов и отсканировать её. Сверка WB продолжится в фоне."
            ),
        )
    if (
        exclusion_request.status == FbsPickRestockRequest.STATUS_QUEUED
        and exclusion_request.source_tote_id is None
    ):
        return (
            "Передать в возврат отбора",
            (
                "Контролёру: положить весь заказ в тару отменённых заказов "
                "и подтвердить передачу подборщику."
            ),
        )
    return (
        exclusion_request.get_status_display(),
        exclusion_request.reason or "Заказ исключается из поставки.",
    )


def _handover_detail_summary(
    batch: FbsHandoverBatch,
    *,
    compact_controller: bool = False,
) -> dict:
    from billing.models import WarehouseServiceFact
    from .handover_presentation import handover_metadata_issues

    from .services.handover import (
        MARKETPLACE_PROBLEM_STATUSES,
        handover_order_has_accepted_marketplace_status,
    )
    from .services.traceability import (
        marketplace_metadata_transfer_resolved,
        metadata_requirements,
        ozon_marking_transfer_not_required,
    )

    boxes = list(batch.boxes.all())
    all_assignments = list(batch.order_assignments.all())
    assignments = [
        assignment
        for assignment in all_assignments
        if assignment.status != FbsHandoverOrderAssignment.STATUS_CANCELED
    ]
    operator_metadata_issues = handover_metadata_issues(
        [assignment.order for assignment in assignments]
        + [link.order for box in boxes for link in box.orders.all()
           if link.status == FbsHandoverOrder.STATUS_ACTIVE]
    )
    excluded_assignments = [
        assignment
        for assignment in all_assignments
        if assignment.status == FbsHandoverOrderAssignment.STATUS_CANCELED
    ]
    exclusion_requests = list(
        FbsPickRestockRequest.objects.filter(
            handover_assignment__batch=batch,
            order__isnull=False,
        )
        .select_related("order", "assigned_to", "created_by", "source_tote")
        .order_by("created_at", "id")
    )
    exclusion_by_order_id = {
        request.order_id: request for request in exclusion_requests
    }
    assignments.sort(
        key=lambda assignment: _handover_assignment_display_priority(
            assignment,
            exclusion_by_order_id,
        )
    )
    commands = list(batch.marketplace_commands.all())
    relevant_commands = _handover_relevant_commands(
        commands,
        order_ids=(assignment.order_id for assignment in assignments),
    )
    order_ids = []
    counted_order_ids = set()
    total_units = 0
    total_lines = 0
    known_weight = Decimal("0")
    weight_complete = True
    label_document_count = 0
    metadata_problem_count = 0
    metadata_waiting_count = 0
    accepted_order_count = 0
    rejected_order_count = 0
    handover_links_by_order_id = {}
    acceptance_is_active = batch.status in {
        FbsHandoverBatch.STATUS_DISPATCHED,
        FbsHandoverBatch.STATUS_PROBLEM,
        FbsHandoverBatch.STATUS_ACCEPTED,
    }

    problem_transfer_statuses = {
        FbsMarketplaceMetadataTransfer.STATUS_FAILED,
        FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
        FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED,
        FbsMarketplaceMetadataTransfer.STATUS_CANCELED,
    }
    waiting_transfer_statuses = {
        FbsMarketplaceMetadataTransfer.STATUS_PREPARED,
        FbsMarketplaceMetadataTransfer.STATUS_QUEUED,
        FbsMarketplaceMetadataTransfer.STATUS_SENT,
        FbsMarketplaceMetadataTransfer.STATUS_RETRY,
    }

    def prepare_order(order):
        nonlocal total_units, total_lines, known_weight, weight_complete
        nonlocal label_document_count, metadata_problem_count, metadata_waiting_count
        nonlocal accepted_order_count, rejected_order_count

        items = list(order.items.all())
        labels = list(order.marketplace_labels.all())
        order_units = 0
        order_weight = Decimal("0")
        order_weight_complete = True
        required_transfer_count = 0
        confirmed_required_count = 0
        order_problem = False
        order_problem_label = ""
        order_has_invalid_kiz = False
        invalid_kiz_workstation_id = None
        order_waiting = False

        for item in items:
            item.photo_url = str(getattr(item.sku, "img", "") or "")
            item.transfers = list(item.metadata_transfers.all())
            requirements = metadata_requirements(item)
            transfer_by_type = {
                transfer.metadata_type: transfer for transfer in item.transfers
            }
            item.metadata_checks = []
            for metadata_type, label, required in (
                (
                    FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE,
                    "КИЗ",
                    requirements.marketplace_marking_required,
                ),
                (
                    FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION,
                    "Срок годности",
                    requirements.expiry_required,
                ),
            ):
                transfer = transfer_by_type.get(metadata_type)
                effective_required = bool(
                    required
                    or (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and transfer is not None
                        and transfer.is_required
                        and not ozon_marking_transfer_not_required(transfer)
                    )
                )
                if (compact_controller and transfer is not None
                        and transfer.status in problem_transfer_statuses
                        and not marketplace_metadata_transfer_resolved(transfer)):
                    state, state_label = "problem", "Ошибка передачи данных"
                elif not effective_required:
                    state, state_label = "pending", "Не требуется"
                elif marketplace_metadata_transfer_resolved(transfer):
                    state, state_label = "done", "Подтверждено"
                elif transfer is None or transfer.status in waiting_transfer_statuses:
                    state, state_label = "in_progress", "Ожидает подтверждения"
                elif transfer.status in problem_transfer_statuses:
                    state, state_label = "problem", transfer.get_status_display()
                else:
                    state, state_label = "in_progress", transfer.get_status_display()
                item.metadata_checks.append(
                    SimpleNamespace(
                        label=label,
                        required=effective_required,
                        state=state,
                        state_label=state_label,
                        transfer=transfer,
                        can_retry=(
                            metadata_type
                            == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                            and transfer is not None
                            and transfer.status
                            in {
                                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                            }
                            and batch.status
                            in {
                                FbsHandoverBatch.STATUS_OPEN,
                                FbsHandoverBatch.STATUS_READY,
                            }
                            and batch.profile.marketplace
                            == FbsIntegrationProfile.MARKETPLACE_WB
                            and str(order.marketplace_status or "").strip().casefold()
                            == "confirm"
                        ),
                        retry_block_reason=(
                            ""
                            if metadata_type
                            != FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                            or transfer is None
                            or transfer.status
                            not in {
                                FbsMarketplaceMetadataTransfer.STATUS_RETRY,
                                FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                                FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                            }
                            else (
                                "Повтор доступен только в открытой отгрузке."
                                if batch.status
                                not in {
                                    FbsHandoverBatch.STATUS_OPEN,
                                    FbsHandoverBatch.STATUS_READY,
                                }
                                else (
                                    "Повтор доступен только для WB."
                                    if batch.profile.marketplace
                                    != FbsIntegrationProfile.MARKETPLACE_WB
                                    else (
                                        "WB разрешает повтор КИЗа только в статусе confirm; "
                                        f"сейчас {order.marketplace_status or 'статус не получен'}."
                                        if str(order.marketplace_status or "")
                                        .strip()
                                        .casefold()
                                        != "confirm"
                                        else ""
                                    )
                                )
                            )
                        ),
                    )
                )
            required_types = set()
            marking_transfer = transfer_by_type.get(
                FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
            )
            if requirements.marketplace_marking_required or (
                marking_transfer is not None
                and marking_transfer.is_required
                and not ozon_marking_transfer_not_required(marking_transfer)
            ):
                required_types.add(FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE)
            if requirements.expiry_required:
                required_types.add(FbsMarketplaceMetadataTransfer.TYPE_EXPIRATION)
            required_transfer_count += len(required_types)
            for metadata_type in required_types:
                transfer = transfer_by_type.get(metadata_type)
                if marketplace_metadata_transfer_resolved(transfer):
                    confirmed_required_count += 1
                elif transfer is None or transfer.status in waiting_transfer_statuses:
                    order_waiting = True
                elif transfer.status in problem_transfer_statuses:
                    order_problem = True
                    if (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and is_final_wb_marking_rejection(transfer)
                    ):
                        allocation = (
                            transfer.traceability.allocation
                            if transfer.traceability_id
                            else None
                        )
                        if (
                            allocation is not None
                            and allocation.status
                            == FbsOrderStockAllocation.STATUS_PICKED
                            and allocation.pick_task_id
                        ):
                            order_has_invalid_kiz = True
                            invalid_kiz_workstation_id = (
                                allocation.pick_task.batch.workstation_id
                            )
                    if (
                        metadata_type
                        == FbsMarketplaceMetadataTransfer.TYPE_MARKING_CODE
                        and is_final_wb_marking_rejection(transfer)
                    ):
                        order_problem_label = "Невалиден"
                    elif transfer.status in {
                        FbsMarketplaceMetadataTransfer.STATUS_FAILED,
                        FbsMarketplaceMetadataTransfer.STATUS_CONFLICT,
                    }:
                        order_problem_label = "Ошибка проверки WB"
                    elif (
                        not order_problem_label
                        and transfer.status
                        == FbsMarketplaceMetadataTransfer.STATUS_UNSUPPORTED
                    ):
                        order_problem_label = "Проверка недоступна"
                    elif not order_problem_label:
                        order_problem_label = transfer.get_status_display()

            quantity = int(item.quantity or 0)
            order_units += quantity
            unit_weight = None
            if item.sku is not None:
                unit_weight = item.sku.weight_gross_kg or item.sku.weight_kg
            if unit_weight is None:
                order_weight_complete = False
            else:
                order_weight += Decimal(unit_weight) * quantity

        if order_problem:
            order.metadata_state = "problem"
            order.metadata_state_label = order_problem_label or "Есть ошибка"
        elif order_waiting:
            order.metadata_state = "in_progress"
            order.metadata_state_label = "Ожидает"
        elif required_transfer_count:
            order.metadata_state = "done"
            order.metadata_state_label = (
                f"Подтверждено {confirmed_required_count} из {required_transfer_count}"
            )
        else:
            order.metadata_state = "pending"
            order.metadata_state_label = "Не требуется"

        order.ui_items = items
        order.has_invalid_kiz = order_has_invalid_kiz
        order.invalid_kiz_workstation_id = invalid_kiz_workstation_id
        order.ui_labels = labels
        order.latest_label = labels[0] if labels else None
        kiz_retry_label_number = ""
        if order.latest_label is not None:
            kiz_retry_label_number = (
                str(order.latest_label.external_label_id or "").strip()
                or str(order.latest_label.barcode or "").strip()
            )
        for item in items:
            item.kiz_retry_label_number = kiz_retry_label_number
            item.kiz_retry_product_barcode = str(item.barcode or "").strip()
        normalized_marketplace_status = str(order.marketplace_status or "").strip().casefold()
        if acceptance_is_active and normalized_marketplace_status in MARKETPLACE_PROBLEM_STATUSES:
            order.acceptance_state = "problem"
            order.acceptance_label = "Отклонен"
        elif acceptance_is_active and handover_order_has_accepted_marketplace_status(
            order,
            marketplace=batch.profile.marketplace,
        ):
            order.acceptance_state = "done"
            order.acceptance_label = "Принят"
        elif acceptance_is_active:
            order.acceptance_state = "in_progress"
            order.acceptance_label = "Не принят"
        else:
            order.acceptance_state = "pending"
            order.acceptance_label = "Не передан"
        order.unit_count = order_units
        order.line_count = len(items)
        order.weight_kg = order_weight
        order.weight_complete = order_weight_complete

        if order.id not in counted_order_ids:
            counted_order_ids.add(order.id)
            order_ids.append(order.id)
            total_units += order_units
            total_lines += len(items)
            known_weight += order_weight
            weight_complete = weight_complete and order_weight_complete
            label_document_count += sum(bool(label.document_url) for label in labels)
            metadata_problem_count += int(order_problem)
            metadata_waiting_count += int(order_waiting and not order_problem)
            accepted_order_count += int(order.acceptance_state == "done")
            rejected_order_count += int(order.acceptance_state == "problem")
        return order_units, order_weight, order_weight_complete

    for assignment in assignments:
        prepare_order(assignment.order)
        assignment.latest_label = assignment.order.latest_label
        label_payload = (
            assignment.latest_label.payload
            if assignment.latest_label is not None
            and isinstance(assignment.latest_label.payload, dict)
            else {}
        )
        assignment.uses_preconfirmed_ozon_qr = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and assignment.latest_label is not None
            and label_payload.get("_preloaded_ozon_order_barcode") is True
            and label_payload.get("_preconfirmed_order_barcode_at")
        )

    retryable_kiz_items = []
    retryable_kiz_item_ids = set()
    for assignment in assignments:
        for item in assignment.order.ui_items:
            if item.id in retryable_kiz_item_ids:
                continue
            if not any(
                check.label == "КИЗ" and check.can_retry
                for check in item.metadata_checks
            ):
                continue
            retryable_kiz_item_ids.add(item.id)
            retryable_kiz_items.append(
                SimpleNamespace(
                    order_item_id=item.id,
                    order_number=assignment.order.external_order_id,
                    product_name=item.product_name or item.external_sku,
                    label_number=item.kiz_retry_label_number,
                    product_barcode=item.kiz_retry_product_barcode,
                )
            )

    invalid_kiz_workstation_ids = {
        assignment.order.invalid_kiz_workstation_id
        for assignment in assignments
        if assignment.order.has_invalid_kiz
        and assignment.order.invalid_kiz_workstation_id
    }
    invalid_kiz_sessions_by_workstation = {
        session.workstation_id: session
        for session in FbsControllerSession.objects.select_related(
            "problem_tote__binding", "workstation"
        ).filter(
            workstation_id__in=invalid_kiz_workstation_ids,
            status=FbsControllerSession.STATUS_ACTIVE,
            problem_tote__isnull=False,
        )
    }

    for box in boxes:
        box_units = 0
        box_weight = Decimal("0")
        box_weight_complete = True
        links = [
            link
            for link in box.orders.all()
            if link.status == FbsHandoverOrder.STATUS_ACTIVE
        ]
        for link in links:
            handover_links_by_order_id[link.order_id] = link
            order_units, order_weight, order_weight_complete = prepare_order(link.order)
            box_units += order_units
            box_weight += order_weight
            box_weight_complete = box_weight_complete and order_weight_complete

        box.ui_orders = links
        box.order_count = len(links)
        box.is_empty_history = not links and box.orders.count() > 0
        box.unit_count = box_units
        box.weight_kg = box_weight
        box.weight_complete = box_weight_complete

    verification_rows = []
    requires_composition_scan = (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
    )
    verification_override_active = has_handover_verification_override(batch)
    for assignment in assignments:
        order = assignment.order
        order.display_created_at = order.ordered_at or order.imported_at
        exclusion_request = exclusion_by_order_id.get(order.id)
        link = handover_links_by_order_id.get(order.id)
        assignment.exclusion_request = exclusion_request
        assignment.can_controller_stage_exclusion = bool(
            exclusion_request is not None
            and exclusion_request.source_tote_id is None
            and exclusion_request.status
            in {
                FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
                FbsPickRestockRequest.STATUS_QUEUED,
                FbsPickRestockRequest.STATUS_FAILED,
            }
        )
        # ``order_is_client_canceled_by_marketplace`` understands WB just as
        # well as Ozon, but this screen used to consult it only for Ozon.  A WB
        # order cancelled by the customer therefore rendered exactly like a live
        # one: the only trace was the raw ``canceled_by_client`` string inside a
        # technical column, so a controller working through the list packed it
        # into the transport box like everything else.
        assignment.client_canceled = order_is_client_canceled_by_marketplace(order)
        assignment.ozon_client_canceled = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and assignment.client_canceled
        )
        assignment.can_return_ozon_canceled = bool(
            assignment.ozon_client_canceled
            and batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and assignment.status
            in {
                FbsHandoverOrderAssignment.STATUS_PENDING,
                FbsHandoverOrderAssignment.STATUS_CONFIRMED,
                FbsHandoverOrderAssignment.STATUS_ERROR,
            }
            and (
                exclusion_request is None
                or (
                    exclusion_request.status == FbsPickRestockRequest.STATUS_QUEUED
                    and exclusion_request.source_tote_id is None
                )
            )
        )
        assignment.can_exclude = (
            batch.status
            in {
                FbsHandoverBatch.STATUS_OPEN,
                FbsHandoverBatch.STATUS_READY,
            }
            and exclusion_request is None
            and link is not None
            and assignment.status
            in {
                FbsHandoverOrderAssignment.STATUS_CONFIRMED,
                FbsHandoverOrderAssignment.STATUS_ERROR,
            }
        )
        assignment.exclusion_requires_seller_cancel = (
            assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        )
        assignment.invalid_kiz_route_available = bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
            and batch.status
            in {FbsHandoverBatch.STATUS_OPEN, FbsHandoverBatch.STATUS_READY}
            and assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
            and order.has_invalid_kiz
            and str(order.marketplace_status or "").strip().casefold() == "confirm"
        )
        invalid_kiz_session = invalid_kiz_sessions_by_workstation.get(
            order.invalid_kiz_workstation_id
        )
        invalid_kiz_problem_tote = (
            invalid_kiz_session.problem_tote if invalid_kiz_session is not None else None
        )
        invalid_kiz_binding = (
            getattr(invalid_kiz_problem_tote, "binding", None)
            if invalid_kiz_problem_tote is not None
            else None
        )
        assignment.invalid_kiz_problem_tote = invalid_kiz_problem_tote
        invalid_kiz_binding_state_ready = _controller_service_binding_ready(
            session=invalid_kiz_session,
            binding=invalid_kiz_binding,
            tote_id=(
                invalid_kiz_session.problem_tote_id
                if invalid_kiz_session is not None
                else None
            ),
        )
        assignment.can_reroute_invalid_kiz = bool(
            assignment.invalid_kiz_route_available
            and invalid_kiz_binding_state_ready
        )
        label = (
            link.verified_label
            if link is not None and link.verified_label_id
            else assignment.latest_label
        )
        label_scanned = bool(
            label is not None and label.status == FbsOrderLabel.STATUS_APPLIED
        )
        composition_verified = bool(
            link is not None
            and (
                link.verified_at is not None or verification_override_active
                if requires_composition_scan
                else label_scanned
            )
        )
        if (
            exclusion_request is not None
            and exclusion_request.status != FbsPickRestockRequest.STATUS_COMPLETED
        ):
            state = "problem"
            state_label, state_detail = _handover_exclusion_check_copy(
                exclusion_request
            )
        elif assignment.status in {
            FbsHandoverOrderAssignment.STATUS_ERROR,
            FbsHandoverOrderAssignment.STATUS_CANCELED,
        }:
            state = "problem"
            state_label = "Не ОК"
            state_detail = assignment.error or assignment.get_status_display()
        elif assignment.status == FbsHandoverOrderAssignment.STATUS_PENDING:
            state = "in_progress"
            state_label = "Сверяется с WB"
            state_detail = assignment.error or "Ожидается подтверждение состава поставки WB."
        elif composition_verified:
            state = "done"
            if link.verified_at is not None:
                state_label = "ОК · Проверен"
                state_detail = f"Проверен в коробе {link.box.qr_code}."
            else:
                state_label = "ОК · Разрешено"
                state_detail = (
                    "Первичный скан подтвержден; повторную проверку разрешил "
                    f"начальник склада · короб {link.box.qr_code}."
                )
        elif label_scanned and link is not None:
            state = "in_progress"
            state_label = "Ожидает проверки"
            state_detail = f"Нужен контрольный скан · короб {link.box.qr_code}."
        elif label_scanned:
            state = "in_progress"
            state_label = "Ожидает упаковки"
            state_detail = "Этикетка отсканирована, короб не назначен."
        else:
            state = "problem"
            state_label = "Не проверен"
            state_detail = "Заказ не отсканирован."
        assignment.check_state = state
        assignment.check_label = state_label
        assignment.check_detail = state_detail
        assignment.handover_link = link
        verification_rows.append(
            SimpleNamespace(
                assignment=assignment,
                order=order,
                label=label,
                handover_link=link,
                state=state,
                state_label=state_label,
                state_detail=state_detail,
                is_scanned=composition_verified,
                exclusion_request=exclusion_request,
                can_exclude=assignment.can_exclude,
            )
        )

    verified_order_count = sum(row.is_scanned for row in verification_rows)
    missing_order_count = len(verification_rows) - verified_order_count
    unboxed_order_count = sum(
        row.handover_link is None or row.label is None
        or row.label.status != FbsOrderLabel.STATUS_APPLIED
        for row in verification_rows
    )
    physical_missing_order_count = sum(
        row.handover_link is None for row in verification_rows
    )
    assignment_confirmed_count = sum(
        assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        for assignment in assignments
    )
    all_assignments_confirmed = bool(assignments) and all(
        assignment.status == FbsHandoverOrderAssignment.STATUS_CONFIRMED
        for assignment in assignments
    )
    all_orders_ready = bool(assignments) and all(
        assignment.order.internal_status == FbsOrder.STATUS_READY_FOR_HANDOVER
        for assignment in assignments
    )
    active_boxes = [box for box in boxes if box.order_count]
    active_check_tote = (
        FbsControllerCheckTote.objects.select_related(
            "tote", "session__workstation"
        )
        .filter(
            handover_batch=batch,
            status__in=(
                FbsControllerCheckTote.STATUS_OPEN,
                FbsControllerCheckTote.STATUS_WAITING_KIZ,
                FbsControllerCheckTote.STATUS_READY,
                FbsControllerCheckTote.STATUS_COMPOSITION,
            ),
        )
        .first()
    )
    if compact_controller and batch.status in {
        FbsHandoverBatch.STATUS_ACCEPTED,
        FbsHandoverBatch.STATUS_ARCHIVED,
    }:
        # The authoritative composition gate is only actionable before dispatch.
        # Re-running it for a finished controller screen performs one query per
        # order and cannot change the next action or any available operation.
        composition_readiness = SimpleNamespace(ready=True, reasons=())
    else:
        composition_readiness = handover_composition_readiness(batch)
    archive_readiness = (
        SimpleNamespace(ready=False, reasons=())
        if compact_controller
        else handover_archive_readiness(batch)
    )
    blocking_return_requests = [
        request
        for request in exclusion_requests
        if _handover_restock_blocks_handover(request)
    ]
    blocking_return_count = len(blocking_return_requests)
    background_return_requests = [
        request
        for request in exclusion_requests
        if request.source_tote_id is not None
        and request.status in HANDOVER_BLOCKING_PICK_RESTOCK_STATUSES
    ]
    background_return_count = len(background_return_requests)
    blocking_return_order_numbers = tuple(
        str(request.order.external_order_id or request.order_id).strip()
        for request in blocking_return_requests
    )
    next_action = _handover_next_action(
        batch,
        boxes=active_boxes,
        assignment_count=len(assignments),
        missing_order_count=unboxed_order_count,
        physical_missing_order_count=physical_missing_order_count,
        unverified_order_count=missing_order_count,
        all_assignments_confirmed=all_assignments_confirmed,
        all_orders_ready=all_orders_ready,
        composition_ready=composition_readiness.ready,
        blocking_return_count=blocking_return_count,
        blocking_return_order_numbers=blocking_return_order_numbers,
        blocking_return_statuses=tuple(
            request.status for request in blocking_return_requests
        ),
        background_return_count=background_return_count,
        background_return_order_numbers=tuple(
            str(request.order.external_order_id or request.order_id).strip()
            for request in background_return_requests
        ),
    )

    box_count = len(active_boxes)
    accepted_boxes = sum(
        box.status == FbsHandoverBox.STATUS_ACCEPTED for box in active_boxes
    )
    problem_boxes = sum(
        box.status == FbsHandoverBox.STATUS_PROBLEM for box in active_boxes
    )
    unique_order_count = len(counted_order_ids)
    if unique_order_count and accepted_order_count == unique_order_count:
        acceptance_label, acceptance_class = "Принята полностью", "done"
    elif accepted_order_count:
        acceptance_label, acceptance_class = "Принята частично", "in_progress"
    elif rejected_order_count or problem_boxes or batch.status == FbsHandoverBatch.STATUS_PROBLEM:
        acceptance_label, acceptance_class = "Отклонена", "problem"
    elif batch.status == FbsHandoverBatch.STATUS_ACCEPTED and box_count:
        acceptance_label, acceptance_class = "Принята полностью", "done"
    elif accepted_boxes:
        acceptance_label, acceptance_class = "Принята частично", "in_progress"
    elif batch.status == FbsHandoverBatch.STATUS_DISPATCHED:
        acceptance_label, acceptance_class = "Ожидает приемки", "in_progress"
    else:
        acceptance_label, acceptance_class = "Не передана", "pending"

    billing_facts = []
    handover_history = ()
    handover_history_total = 0
    if not compact_controller:
        billing_order_ids = [f"FBS-ORDER-{order_id}" for order_id in set(order_ids)]
        billing_order_ids.append(f"FBS-DELIVERY-{batch.id}")
        billing_facts = list(
            WarehouseServiceFact.objects.filter(
                client=batch.profile.agency,
                order_type=WarehouseServiceFact.ORDER_FBS,
                order_id__in=billing_order_ids,
            )
            .select_related("service", "charge")
            .order_by("performed_at", "id")
        )
        for fact in billing_facts:
            fact.display_name = fact.service_name_snapshot or fact.service.name

        handover_history, handover_history_total = _handover_history_rows(
            batch,
            boxes=boxes,
            assignments=all_assignments,
            commands=commands,
        )
    pending_verification_rows = []
    related_pick_batches = (
        FbsPickBatch.objects.filter(
            controller_pick_tote__check_tote__handover_batch=batch,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .order_by("id")
        .distinct()
    )
    assigned_order_ids = {assignment.order_id for assignment in assignments}
    for pick_batch in related_pick_batches:
        for row in _verification_wave_rows(pick_batch):
            if row.order.id in assigned_order_ids:
                continue
            row.pick_batch = pick_batch
            if not row.is_product_complete:
                row.missing_reason = (
                    f"Товар проверен {row.verified_qty} из {row.planned_qty} шт."
                )
            elif not row.is_label_complete:
                row.missing_reason = "QR заказа не создан или не отсканирован."
            else:
                row.missing_reason = "QR отсканирован, заказ еще не добавлен в отгрузку."
            pending_verification_rows.append(row)

    reconciliation_available = (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and batch.status
        in {
            FbsHandoverBatch.STATUS_OPEN,
            FbsHandoverBatch.STATUS_READY,
        }
        and bool(verification_rows or pending_verification_rows)
    )
    reconciliation_missing_count = missing_order_count + len(pending_verification_rows)

    # A controller scans through a check tote, and ``active_check_tote`` only
    # finds one once the wave has been taken into verification.  A shipment
    # whose wave was never started therefore has no tote, and the screen used to
    # offer nothing but a generic link to the controller home -- it could not
    # say which wave was holding the shipment.  Resolve that wave here so the
    # next step can name it.
    handover_pending_pick_batch = None
    if active_check_tote is None and missing_order_count:
        handover_pending_pick_batch = (
            FbsPickBatch.objects.filter(
                tasks__order_id__in=assigned_order_ids,
                status=FbsPickBatch.STATUS_VERIFICATION,
            )
            .order_by("id")
            .distinct()
            .first()
        )

    return {
        "handover_boxes": active_boxes,
        "handover_check_tote": active_check_tote,
        "handover_pending_pick_batch": handover_pending_pick_batch,
        "handover_excluded_boxes": [box for box in boxes if not box.order_count],
        "handover_assignments": assignments,
        "handover_excluded_assignments": excluded_assignments,
        "handover_exclusion_requests": exclusion_requests,
        "handover_kiz_retry_items": retryable_kiz_items,
        "handover_kiz_retry_count": len(retryable_kiz_items),
        "handover_blocking_return_count": blocking_return_count,
        "handover_verification_rows": verification_rows,
        "handover_verified_order_count": verified_order_count,
        "handover_missing_order_count": missing_order_count,
        "handover_unboxed_order_count": unboxed_order_count,
        "handover_physical_missing_order_count": physical_missing_order_count,
        "handover_reconciliation_available": reconciliation_available,
        "handover_reconciliation_total_count": (
            len(verification_rows) + len(pending_verification_rows)
        ),
        "handover_reconciliation_missing_count": reconciliation_missing_count,
        "handover_scan_complete": bool(verification_rows)
        and not reconciliation_missing_count,
        "handover_composition_ready": composition_readiness.ready,
        "handover_composition_blockers": composition_readiness.reasons,
        "handover_operator_metadata_issues": operator_metadata_issues,
        "handover_archive_ready": archive_readiness.ready,
        "handover_archive_blockers": archive_readiness.reasons,
        "handover_exclusion_reason_choices": FbsPickRestockRequest.REASON_CHOICES,
        "handover_assignment_confirmed_count": assignment_confirmed_count,
        "handover_printable_order_label_count": sum(
            bool(
                assignment.latest_label
                and assignment.latest_label.status
                in {FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED}
                and assignment.latest_label.file
                and assignment.latest_label.label_format
                in {FbsOrderLabel.FORMAT_PNG, FbsOrderLabel.FORMAT_PDF}
            )
            for assignment in assignments
        ),
        "handover_all_assignments_confirmed": all_assignments_confirmed,
        "handover_all_orders_ready": all_orders_ready,
        "handover_next_action": next_action,
        "handover_commands": commands,
        "handover_commands_active": any(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_PENDING,
                FbsMarketplaceCommand.STATUS_SENT,
                FbsMarketplaceCommand.STATUS_RETRY,
            }
            for command in relevant_commands
        ),
        "handover_command_problem_count": sum(
            command.status
            in {
                FbsMarketplaceCommand.STATUS_FAILED,
                FbsMarketplaceCommand.STATUS_CONFLICT,
            }
            for command in relevant_commands
        ),
        "handover_stages": _handover_stage_rows(batch),
        "handover_order_count": unique_order_count,
        "handover_unit_count": total_units,
        "handover_line_count": total_lines,
        "handover_weight_kg": known_weight,
        "handover_weight_complete": weight_complete and bool(order_ids),
        "handover_accepted_boxes": accepted_boxes,
        "handover_problem_boxes": problem_boxes,
        "handover_accepted_orders": accepted_order_count,
        "handover_rejected_orders": rejected_order_count,
        "handover_acceptance_label": acceptance_label,
        "handover_acceptance_class": acceptance_class,
        "handover_label_document_count": label_document_count,
        "handover_metadata_problem_count": metadata_problem_count,
        "handover_metadata_waiting_count": metadata_waiting_count,
        "handover_driver_manifest_ready": bool(boxes)
        and batch.status != FbsHandoverBatch.STATUS_OPEN,
        "handover_supply_label_available": bool(batch.supply_label_file)
        or bool(
            batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
            and order_ids
            and active_boxes
            and batch.status != FbsHandoverBatch.STATUS_OPEN
        ),
        "handover_billing_facts": billing_facts,
        "handover_billing_charged_count": sum(
            fact.status == WarehouseServiceFact.STATUS_CHARGED for fact in billing_facts
        ),
        "handover_print_workstations": configured_fbs_print_workstations(),
        "handover_uses_marketplace_boxes": wb_handover_uses_marketplace_boxes(batch),
        "handover_history": handover_history,
        "handover_history_total": handover_history_total,
        "handover_pending_verification_rows": pending_verification_rows,
        "handover_pending_verification_count": len(pending_verification_rows),
        "handover_fbs_otg_locations_configured": WarehouseLocation.objects.filter(
            warehouse_code="MSK",
            zone_code="OTG",
            is_active=True,
            is_topology_visible=False,
            is_fbs_visible=True,
        ).exists(),
    }


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_handover(request):
    error = ""
    filter_error = ""
    if request.method == "POST":
        try:
            action = str(request.POST.get("action") or "").strip()
            if action == "bulk_print_supply_labels":
                raw_batch_ids = request.POST.getlist("batch_ids")
                try:
                    batch_ids = list(dict.fromkeys(int(value) for value in raw_batch_ids))
                except (TypeError, ValueError):
                    raise ValueError("Некорректный список отгрузок для печати.")
                if not batch_ids:
                    raise ValueError("Выберите хотя бы одну отгрузку.")
                if len(batch_ids) > 100:
                    raise ValueError("За один раз можно напечатать не более 100 ШК поставок.")
                selected_batches = list(
                    FbsHandoverBatch.objects.select_related("profile")
                    .filter(pk__in=batch_ids)
                    .only("id", "status", "supply_label_file", "profile__marketplace")
                    .order_by("id")
                )
                if len(selected_batches) != len(batch_ids):
                    raise ValueError("Одна из выбранных отгрузок не найдена.")
                missing_labels = [
                    batch.id
                    for batch in selected_batches
                    if not batch.supply_label_file
                    and batch.profile.marketplace != FbsIntegrationProfile.MARKETPLACE_OZON
                ]
                if missing_labels:
                    missing_text = ", ".join(f"#{batch_id}" for batch_id in missing_labels[:10])
                    raise ValueError(f"ШК поставки еще не готов: {missing_text}.")
                workstation_id = int(request.POST.get("workstation_id") or 0)
                desktop_target = _desktop_label_print_target(request, shipping=True)
                with transaction.atomic():
                    for selected_batch in selected_batches:
                        queue_fbs_handover_supply_label_print(
                            batch_id=selected_batch.id,
                            workstation_id=workstation_id,
                            requested_by=request.user,
                            **desktop_target,
                        )
                redirect_query = request.GET.copy()
                redirect_query["printed_supplies"] = str(len(selected_batches))
                redirect_url = reverse("fbs:tsd_handover")
                if redirect_query:
                    redirect_url = f"{redirect_url}?{redirect_query.urlencode()}"
                return redirect(redirect_url)
            profile = FbsIntegrationProfile.objects.get(pk=int(request.POST.get("profile_id") or 0))
            batch = create_handover_batch(
                profile=profile,
                external_supply_id=request.POST.get("external_supply_id", ""),
                created_by=request.user,
            )
            return redirect("fbs:tsd_handover_detail", batch_id=batch.id)
        except (FbsError, ValueError, FbsIntegrationProfile.DoesNotExist) as exc:
            error = str(exc)
    from .controller_shipment_ui import (SHIPMENT_STAGE_CHOICES, shipment_stage_filters, shipment_stage_counts)
    stages = shipment_stage_filters()
    queryset = _handover_queryset()
    search = str(request.GET.get("q") or "").strip()
    status_filter = str(request.GET.get("status") or "").strip()
    marketplace_filter = str(request.GET.get("marketplace") or "").strip()
    agency_filter = str(request.GET.get("agency_id") or "").strip()
    date_field_filter = str(request.GET.get("date_field") or "created").strip()
    date_from_filter = str(request.GET.get("date_from") or "").strip()
    date_to_filter = str(request.GET.get("date_to") or "").strip()
    date_field_choices = (
        ("created", "Дата создания"),
        ("verified", "Дата проверки"),
        ("dispatched", "Дата отправки"),
    )
    if date_field_filter not in dict(date_field_choices):
        date_field_filter = "created"
    if search:
        search_filter = (
            Q(external_supply_id__icontains=search)
            | Q(external_name__icontains=search)
            | Q(profile__agency__agn_name__icontains=search)
            | Q(profile__external_warehouse_id__icontains=search)
        )
        if search.isdigit():
            search_filter |= Q(id=int(search))
        queryset = queryset.filter(search_filter)
    if marketplace_filter in dict(FbsIntegrationProfile.MARKETPLACE_CHOICES):
        queryset = queryset.filter(profile__marketplace=marketplace_filter)
    if agency_filter.isdigit():
        queryset = queryset.filter(profile__agency_id=int(agency_filter))

    parsed_date_from = None
    parsed_date_to = None
    try:
        if date_from_filter:
            parsed_date_from = date.fromisoformat(date_from_filter)
        if date_to_filter:
            parsed_date_to = date.fromisoformat(date_to_filter)
    except ValueError:
        filter_error = "Укажите даты в формате ДД.ММ.ГГГГ."
    if parsed_date_from and parsed_date_to and parsed_date_from > parsed_date_to:
        filter_error = "Дата начала не может быть позже даты окончания."
    if not filter_error and (parsed_date_from or parsed_date_to):
        date_lookup = {
            "created": "created_at__date",
            "verified": "boxes__orders__verified_at__date",
            "dispatched": "dispatched_at__date",
        }[date_field_filter]
        if parsed_date_from:
            queryset = queryset.filter(**{f"{date_lookup}__gte": parsed_date_from})
        if parsed_date_to:
            queryset = queryset.filter(**{f"{date_lookup}__lte": parsed_date_to})
        if date_field_filter == "verified":
            queryset = queryset.distinct()
    handover_totals = shipment_stage_counts(queryset)
    handover_counter_urls = {}
    for stage_key in stages:
        params = request.GET.copy()
        params['status'] = stage_key
        params.pop('page', None)
        params.pop('scope', None)
        handover_counter_urls[stage_key] = '?' + params.urlencode()
    status_filter = {"open": "checking", "dispatched": "transit"}.get(status_filter, status_filter)
    if status_filter in stages:
        queryset = queryset.filter(stages[status_filter])
    elif not (get_request_role(request) == "fbs_controller" and request.GET.get("scope") == "all"):
        queryset = queryset.exclude(status=FbsHandoverBatch.STATUS_ARCHIVED)
    controller_queue_context = {}
    if get_request_role(request) == 'fbs_controller':
        from .controller_shipment_ui import controller_queue
        queryset, controller_queue_context = controller_queue(queryset, request, get_controller_workstation(request))
    page_size_choices = (25, 50, 100, 200, 500)
    try:
        page_size = int(request.GET.get("page_size") or 50)
    except (TypeError, ValueError):
        page_size = 50
    if page_size not in page_size_choices:
        page_size = 50
    paginator = Paginator(queryset.order_by("-created_at", "-id"), page_size)
    page_obj = paginator.get_page(request.GET.get("page"))
    batches = _prepare_handover_list_rows(list(page_obj.object_list))
    handover_totals['dispatched'] = handover_totals['transit']
    client_profiles = (
        FbsIntegrationProfile.objects.filter(handover_batches__isnull=False)
        .select_related("agency")
        .order_by("agency__agn_name", "agency_id", "id")
        .distinct()
    )
    handover_clients = []
    seen_agency_ids = set()
    for profile in client_profiles:
        if profile.agency_id in seen_agency_ids:
            continue
        seen_agency_ids.add(profile.agency_id)
        profile.agency.handover_selected = agency_filter == str(profile.agency_id)
        handover_clients.append(profile.agency)
    stale_batches = list(
        queryset.filter(stages['active'], created_at__lt=timezone.now()-timedelta(days=4))
        .order_by('created_at', 'id')[:20]
    )
    query_without_page = request.GET.copy()
    query_without_page.pop("page", None)
    query_without_page_encoded = query_without_page.urlencode()
    return render(
        request,
        "fbs/tsd_handover_list.html",
        _base_context(
            request,
            page_title="FBS · Отгрузки",
            back_url=(
                reverse("fbs:controller_home")
                if get_request_role(request) == "fbs_controller"
                else reverse("fbs:tsd_storekeeper")
            ),
            batches=batches,
            page_obj=page_obj,
            handover_page_size=page_size,
            handover_page_size_choices=page_size_choices,
            handover_query_without_page=query_without_page_encoded,
            handover_totals=handover_totals,
            handover_counter_urls=handover_counter_urls,
            **controller_queue_context,
            handover_search=search,
            handover_status_filter=status_filter,
            handover_marketplace_filter=marketplace_filter,
            handover_agency_filter=agency_filter,
            handover_date_field_filter=date_field_filter,
            handover_date_from_filter=date_from_filter,
            handover_date_to_filter=date_to_filter,
            handover_status_choices=(("active", "На складе"),) + SHIPMENT_STAGE_CHOICES,
            handover_marketplace_choices=FbsIntegrationProfile.MARKETPLACE_CHOICES,
            handover_date_field_choices=date_field_choices,
            handover_clients=handover_clients,
            handover_stale_batches=stale_batches,
            handover_print_workstations=configured_fbs_print_workstations(),
            ok_message=(
                f"В очередь печати добавлено ШК поставок: {request.GET.get('printed_supplies')}."
                if str(request.GET.get("printed_supplies") or "").isdigit()
                else ""
            ),
            profiles=FbsIntegrationProfile.objects.filter(
                is_active=True,
                marketplace=FbsIntegrationProfile.MARKETPLACE_OZON,
            ).select_related("agency"),
            error=error,
            filter_error=filter_error,
        ),
        status=400 if error else 200,
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_reports(request):
    filters = _handover_report_filter_context(request)
    report_type = filters["report_type"]
    report_rows = []
    report_totals = {"shipments": 0, "orders": 0, "units": 0, "boxes": 0}
    report_truncation_warning = ""
    marketplace_labels = dict(FbsIntegrationProfile.MARKETPLACE_CHOICES)

    batch_queryset = FbsHandoverBatch.objects.select_related(
        "profile__agency", "created_by", "dispatched_by"
    ).filter(
        created_at__date__gte=filters["date_from"],
        created_at__date__lte=filters["date_to"],
    )
    if filters["agency_filter"]:
        batch_queryset = batch_queryset.filter(
            profile__agency_id=int(filters["agency_filter"])
        )
    if filters["marketplace_filter"]:
        batch_queryset = batch_queryset.filter(
            profile__marketplace=filters["marketplace_filter"]
        )

    if filters["filter_error"]:
        pass
    elif report_type == "status":
        aggregates = {
            row["status"]: row
            for row in batch_queryset.values("status").annotate(
                shipment_count=Count("id", distinct=True),
                order_count=Count("order_assignments__order_id", distinct=True),
                box_count=Count("boxes__id", distinct=True),
            )
        }
        for status, label in FbsHandoverBatch.STATUS_CHOICES:
            values = aggregates.get(status, {})
            report_rows.append(
                SimpleNamespace(
                    status=status,
                    label=label,
                    shipment_count=int(values.get("shipment_count") or 0),
                    order_count=int(values.get("order_count") or 0),
                    box_count=int(values.get("box_count") or 0),
                )
            )
        report_totals["shipments"] = sum(row.shipment_count for row in report_rows)
        report_totals["orders"] = sum(row.order_count for row in report_rows)
        report_totals["boxes"] = sum(row.box_count for row in report_rows)
    else:
        item_queryset = FbsOrderItem.objects.select_related("sku").order_by("id")
        link_queryset = (
            FbsHandoverOrder.objects.filter(
                status=FbsHandoverOrder.STATUS_ACTIVE,
                box__batch__dispatched_at__isnull=False,
                box__batch__dispatched_at__date__gte=filters["date_from"],
                box__batch__dispatched_at__date__lte=filters["date_to"],
            )
            .select_related(
                "order",
                "box",
                "box__batch",
                "box__batch__profile",
                "box__batch__profile__agency",
            )
            .prefetch_related(Prefetch("order__items", queryset=item_queryset))
            .order_by("-box__batch__dispatched_at", "-id")
        )
        if filters["agency_filter"]:
            link_queryset = link_queryset.filter(
                box__batch__profile__agency_id=int(filters["agency_filter"])
            )
        if filters["marketplace_filter"]:
            link_queryset = link_queryset.filter(
                box__batch__profile__marketplace=filters["marketplace_filter"]
            )
        link_total_count = link_queryset.count()
        links = list(link_queryset[:HANDOVER_REPORT_ROW_LIMIT])
        if link_total_count > HANDOVER_REPORT_ROW_LIMIT:
            report_truncation_warning = (
                f"Показаны первые {HANDOVER_REPORT_ROW_LIMIT} строк из "
                f"{link_total_count}. Сузьте период или фильтры."
            )
        shipment_ids = set()
        box_ids = set()
        total_units = 0
        for link in links:
            link.report_items = list(link.order.items.all())
            link.report_unit_count = sum(int(item.quantity or 0) for item in link.report_items)
            shipment_ids.add(link.box.batch_id)
            box_ids.add(link.box_id)
            total_units += link.report_unit_count
        report_totals = {
            "shipments": len(shipment_ids),
            "orders": len(links),
            "units": total_units,
            "boxes": len(box_ids),
        }

        if report_type == "count":
            grouped = {}
            for link in links:
                profile = link.box.batch.profile
                key = (profile.agency_id, profile.marketplace)
                row = grouped.setdefault(
                    key,
                    SimpleNamespace(
                        agency=profile.agency,
                        marketplace=marketplace_labels.get(
                            profile.marketplace, profile.marketplace
                        ),
                        shipment_ids=set(),
                        order_count=0,
                        unit_count=0,
                    ),
                )
                row.shipment_ids.add(link.box.batch_id)
                row.order_count += 1
                row.unit_count += link.report_unit_count
            report_rows = sorted(
                grouped.values(),
                key=lambda row: (str(row.agency), row.marketplace),
            )
            for row in report_rows:
                row.shipment_count = len(row.shipment_ids)
        elif report_type == "orders":
            report_rows = links
        else:
            for link in links:
                for item in link.report_items:
                    report_rows.append(
                        SimpleNamespace(
                            link=link,
                            item=item,
                            dispatched_at=link.box.batch.dispatched_at,
                        )
                    )

    if (
        not filters["filter_error"]
        and str(request.GET.get("export") or "").strip() == "xlsx"
    ):
        if report_type == "status":
            headers = ("Статус", "Отгрузки", "Заказы", "Короба")
            rows = (
                (row.label, row.shipment_count, row.order_count, row.box_count)
                for row in report_rows
            )
        elif report_type == "count":
            headers = ("Клиент", "Площадка", "Отгрузки", "Заказы", "Товары")
            rows = (
                (
                    str(row.agency),
                    row.marketplace,
                    row.shipment_count,
                    row.order_count,
                    row.unit_count,
                )
                for row in report_rows
            )
        elif report_type == "orders":
            headers = (
                "Дата отправки",
                "Отгрузка",
                "Клиент",
                "Площадка",
                "Заказ",
                "Короб",
                "Товаров",
                "Статус Fullbox",
                "Статус marketplace",
            )
            rows = (
                (
                    timezone.localtime(row.box.batch.dispatched_at).strftime("%d.%m.%Y %H:%M"),
                    row.box.batch_id,
                    str(row.box.batch.profile.agency),
                    row.box.batch.profile.get_marketplace_display(),
                    row.order.external_order_id,
                    row.box.qr_code,
                    row.report_unit_count,
                    row.order.get_internal_status_display(),
                    row.order.marketplace_status,
                )
                for row in report_rows
            )
        else:
            headers = (
                "Дата отправки",
                "Отгрузка",
                "Клиент",
                "Площадка",
                "Заказ",
                "Артикул",
                "Штрихкод",
                "Товар",
                "Количество",
            )
            rows = (
                (
                    timezone.localtime(row.dispatched_at).strftime("%d.%m.%Y %H:%M"),
                    row.link.box.batch_id,
                    str(row.link.box.batch.profile.agency),
                    row.link.box.batch.profile.get_marketplace_display(),
                    row.link.order.external_order_id,
                    row.item.external_sku,
                    row.item.barcode,
                    row.item.product_name,
                    int(row.item.quantity or 0),
                )
                for row in report_rows
            )
        return _handover_report_xlsx(
            title=dict(HANDOVER_REPORT_CHOICES)[report_type],
            headers=headers,
            rows=rows,
            warning=report_truncation_warning,
        )

    page_size_choices = (25, 50, 100, 200, 500)
    try:
        page_size = int(request.GET.get("page_size") or 100)
    except (TypeError, ValueError):
        page_size = 100
    if page_size not in page_size_choices:
        page_size = 100
    page_obj = Paginator(report_rows, page_size).get_page(request.GET.get("page"))
    query_without_page = request.GET.copy()
    query_without_page.pop("page", None)
    export_query = request.GET.copy()
    export_query["export"] = "xlsx"
    clients = list(
        Agency.objects.filter(fbs_integration_profiles__handover_batches__isnull=False)
        .order_by("agn_name", "id")
        .distinct()
    )
    return render(
        request,
        "fbs/tsd_handover_reports.html",
        _base_context(
            request,
            page_title="FBS · Отчеты по отгрузкам",
            back_url=reverse("fbs:tsd_handover"),
            handover_report_choices=HANDOVER_REPORT_CHOICES,
            handover_report_type=report_type,
            handover_report_title=dict(HANDOVER_REPORT_CHOICES)[report_type],
            handover_report_period_label=(
                "Дата создания отгрузки"
                if report_type == "status"
                else "Дата отправки"
            ),
            handover_report_rows=page_obj.object_list,
            handover_report_totals=report_totals,
            handover_report_truncation_warning=report_truncation_warning,
            handover_report_clients=clients,
            handover_report_agency_filter=filters["agency_filter"],
            handover_report_marketplace_filter=filters["marketplace_filter"],
            handover_report_date_from=filters["date_from_filter"],
            handover_report_date_to=filters["date_to_filter"],
            handover_report_marketplace_choices=FbsIntegrationProfile.MARKETPLACE_CHOICES,
            handover_report_page_size=page_size,
            handover_report_page_size_choices=page_size_choices,
            handover_report_query_without_page=query_without_page.urlencode(),
            handover_report_export_query=export_query.urlencode(),
            page_obj=page_obj,
            filter_error=filters["filter_error"],
        ),
    )


def _handover_agent_scan_context(request) -> dict:
    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if get_request_role(request) == "fbs_controller":
        workstation = get_controller_workstation(request)
        if workstation is not None and workstation.device_agent_id:
            agent_scan_poll_url = reverse("fbs:controller_scan_events")
            agent_scan_event_id = int(
                AgentEvent.objects.filter(
                    agent_id=workstation.device_agent.agent_id,
                    event_type=AgentEvent.EVENT_SCAN,
                )
                .order_by("-id")
                .values_list("id", flat=True)
                .first()
                or 0
            )
    return {
        "agent_scan_poll_url": agent_scan_poll_url,
        "agent_scan_event_id": agent_scan_event_id,
    }


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_handover_detail(request, batch_id: int):
    batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    request_role = get_request_role(request)
    error = ""
    action = ""
    kiz_rescan_all_requested = (
        str(request.GET.get("kiz_rescan_all") or "").strip() == "1"
    )
    verified_order = str(request.GET.get("verified_order") or "").strip()
    verified_box = str(request.GET.get("verified_box") or "").strip()
    handover_check_message = (
        f"Заказ {verified_order} подтвержден в коробе {verified_box}."
        if verified_order
        else ""
    )
    ok_message = {
        "box": "QR короба отправлен в очередь печати.",
        "supply": "QR поставки отправлен в очередь печати.",
        "order": "Этикетка заказа отправлена в очередь печати.",
        "ozon_qr": "Рабочий QR Ozon отправлен на повторную печать.",
    }.get(str(request.GET.get("printed") or "").strip(), "")
    printed_orders = str(request.GET.get("printed_orders") or "").strip()
    skipped_orders = str(request.GET.get("skipped_orders") or "").strip()
    if printed_orders.isdigit():
        ok_message = f"В очередь печати добавлено этикеток заказов: {printed_orders}."
        if skipped_orders.isdigit() and int(skipped_orders):
            ok_message += f" Пропущено без готовой этикетки: {skipped_orders}."
    if str(request.GET.get("ozon_return") or "").strip() == "queued":
        ok_message = (
            "Отмененный клиентом заказ убран из отгрузки и помещен в тару "
            "возврата. Остаток восстановится после физического возврата товара."
        )
    retried_order = str(request.GET.get("kiz_requeued") or "").strip()
    if retried_order:
        ok_message = f"Новый КИЗ заказа {retried_order} поставлен в очередь WB."
    kiz_route = str(request.GET.get("kiz_route") or "").strip()
    target_handover = str(request.GET.get("target_handover") or "").strip()
    problem_tote = str(request.GET.get("problem_tote") or "").strip()
    if kiz_route in {"rewave", "quarantine"} and target_handover:
        ok_message = (
            f"Товар зарегистрирован в проблемной таре {problem_tote}. "
            "Заказ не отменен в WB. Запущен безопасный перенос в отгрузку "
            f"№{target_handover}; после подтверждения WB "
            + (
                "система создаст новую волну."
                if kiz_route == "rewave"
                else "задача останется в карантине кладовщика."
            )
        )
    if str(request.GET.get("exclusion") or "").strip() == "queued":
        ok_message = (
            "Комментарий сохранен. Система сверяет состав с WB; после подтверждения "
            "товар появится в разделе «Возврат отбора»."
        )
    if str(request.GET.get("exclusion") or "").strip() == "return_queued":
        ok_message = (
            "Заказ исключен из отгрузки и передан в «Возврат отбора». "
            "FBS-остаток восстановится только после физического возврата товара."
        )
    if str(request.GET.get("exclusion") or "").strip() == "controller_removed":
        ok_message = (
            "Заказ снят с текущей проверки и помещён в назначенную тару "
            "отменённых заказов. Проверяйте следующий заказ; сверка WB "
            "продолжится отдельно, без изменения остатков и резервов."
        )
    if str(request.GET.get("verification_override") or "").strip() == "approved":
        ok_message = (
            "Начальник склада разрешил передачу без повторной проверки. "
            "Состав, короба, этикетки, КИЗ и сроки остаются обязательными."
        )
    if request.method == "POST":
        action = str(request.POST.get("action") or "").strip()
        kiz_rescan_all_requested = (
            str(request.POST.get("kiz_rescan_all") or "").strip() == "1"
        )
        try:
            if (
                batch.status == FbsHandoverBatch.STATUS_ARCHIVED
                and action != "archive"
            ):
                raise ValueError("Отгрузка находится в архиве.")
            if action == "archive":
                archive_handover_batch(
                    batch_id=batch.id,
                    archived_by=request.user,
                    reason=request.POST.get("archive_reason", ""),
                )
            elif action == "approve_verification_override":
                approve_handover_verification_override(
                    batch_id=batch.id,
                    reason=request.POST.get("override_reason", ""),
                    approved_by=request.user,
                )
            elif action == "add_box":
                if batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
                    raise ValueError("Для WB создайте транспортные короба через API.")
                add_handover_box(
                    batch_id=batch.id,
                    qr_code=request.POST.get("box_qr", ""),
                    external_box_id=request.POST.get("external_box_id", ""),
                )
            elif action == "add_order":
                add_order_to_handover_box(
                    box_id=int(request.POST.get("box_id") or 0),
                    order_label_scan=request.POST.get("order_label_scan", ""),
                    added_by=request.user,
                )
            elif action == "verify_order":
                link = verify_handover_order_label(
                    batch_id=batch.id,
                    order_label_scan=request.POST.get("order_label_scan", ""),
                    verified_by=request.user,
                )
            elif action == "exclude_order":
                order_id = int(request.POST.get("order_id") or 0)
                pick_batch_id = (
                    FbsPickTask.objects.filter(
                        order_id=order_id,
                        batch__status=FbsPickBatch.STATUS_VERIFICATION,
                        batch__picking_completed_at__isnull=False,
                    )
                    .order_by("-batch_id")
                    .values_list("batch_id", flat=True)
                    .first()
                )
                if not pick_batch_id:
                    raise ValueError("Завершенная волна заказа не найдена.")
                create_order_pick_restock_request(
                    batch_id=pick_batch_id,
                    handover_batch_id=batch.id,
                    order_id=order_id,
                    reason_code=FbsPickRestockRequest.REASON_OTHER,
                    reason=request.POST.get("reason", ""),
                    confirm_seller_cancel=True,
                    created_by=request.user,
                )
            elif action == "retry_exclusion":
                retry_order_pick_restock_marketplace(
                    request_id=int(request.POST.get("request_id") or 0),
                    requested_by=request.user,
                )
            elif action == "resolve_problem_order":
                if request_role != "fbs_controller":
                    raise ValueError("Действие доступно контролёру этой смены.")
                from .services.controller_problem_resolution import confirm_problem_order_to_tote
                confirm_problem_order_to_tote(
                    batch_id=batch.id, order_id=int(request.POST.get("order_id") or 0),
                    order_scan=request.POST.get("order_scan", ""),
                    tote_scan=request.POST.get("canceled_tote_scan", ""), actor=request.user,
                )
            elif action == "resolve_problem_order_without_label":
                if request_role != "fbs_controller":
                    raise ValueError("Действие доступно контролёру этой смены.")
                from .services.controller_problem_resolution import (
                    confirm_problem_order_to_tote_without_label,
                )
                confirm_problem_order_to_tote_without_label(
                    batch_id=batch.id,
                    order_id=int(request.POST.get("order_id") or 0),
                    product_scans=request.POST.getlist("product_scan"),
                    tote_scan=request.POST.get("canceled_tote_scan", ""),
                    actor=request.user,
                )
            elif action == "release_exclusion":
                if request_role != "fbs_controller":
                    raise ValueError("Снять заказ может только FBS-контролёр этой смены.")
                release_order_pick_restock_to_queue(
                    request_id=int(request.POST.get("request_id") or 0),
                    comment=request.POST.get("reason", ""),
                    confirm_physical=True,
                    canceled_tote_scan=request.POST.get("canceled_tote_scan", ""),
                    handover_batch_id=batch.id,
                    performed_by=request.user,
                )
            elif action == "confirm_ozon_canceled_return":
                if request_role != "fbs_controller":
                    raise ValueError("Возврат отмененного заказа выполняет FBS-контролер.")
                confirm_ozon_canceled_order_return(
                    handover_batch_id=batch.id,
                    order_id=int(request.POST.get("order_id") or 0),
                    order_scan=request.POST.get("order_scan", ""),
                    canceled_tote_scan=request.POST.get("canceled_tote_scan", ""),
                    comment=request.POST.get("reason", ""),
                    performed_by=request.user,
                )
            elif action == "close_box":
                close_handover_box(box_id=int(request.POST.get("box_id") or 0))
            elif action == "scan_box":
                scan_handover_box(
                    batch_id=batch.id,
                    box_qr_scan=request.POST.get("box_qr_scan", ""),
                    scanned_by=request.user,
                )
            elif action == "create_wb_boxes":
                add_wb_handover_boxes(
                    batch_id=batch.id,
                    amount=int(request.POST.get("box_count") or 0),
                    requested_by=request.user,
                )
            elif action == "deliver_marketplace":
                request_wb_handover_delivery(
                    batch_id=batch.id,
                    requested_by=request.user,
                )
            elif action == "print_box_label":
                queue_fbs_handover_box_label_print(
                    batch_id=batch.id,
                    box_id=int(request.POST.get("box_id") or 0),
                    workstation_id=int(request.POST.get("workstation_id") or 0),
                    requested_by=request.user,
                    **_desktop_label_print_target(request, shipping=True),
                )
            elif action == "print_supply_label":
                queue_fbs_handover_supply_label_print(
                    batch_id=batch.id,
                    workstation_id=int(request.POST.get("workstation_id") or 0),
                    requested_by=request.user,
                    **_desktop_label_print_target(request, shipping=True),
                )
            elif action == "print_order_label":
                label_id = int(request.POST.get("label_id") or 0)
                if not batch.order_assignments.exclude(
                    status=FbsHandoverOrderAssignment.STATUS_CANCELED
                ).filter(order__marketplace_labels__id=label_id).exists():
                    raise ValueError("Этикетка не относится к этой отгрузке.")
                queue_fbs_order_label_print(
                    label_id=label_id,
                    requested_by=request.user,
                    force=True,
                    **_desktop_label_print_target(request),
                )
            elif action == "reprint_ozon_order_qr":
                if batch.profile.marketplace != FbsIntegrationProfile.MARKETPLACE_OZON:
                    raise ValueError("Рабочий QR Ozon доступен только для отгрузки Ozon.")
                label_id = int(request.POST.get("label_id") or 0)
                if not batch.order_assignments.exclude(
                    status=FbsHandoverOrderAssignment.STATUS_CANCELED
                ).filter(order__marketplace_labels__id=label_id).exists():
                    raise ValueError("Этикетка не относится к этой отгрузке.")
                queue_fbs_preloaded_ozon_order_label_print(
                    label_id=label_id,
                    requested_by=request.user,
                    force=True,
                    **_desktop_label_print_target(request),
                )
            elif action == "print_all_order_labels":
                raise ValueError(
                    "Массовая повторная печать этикеток Ozon отключена. Рабочий QR "
                    "печатается и сканируется один раз во время контроля заказа."
                )
            elif action == "dispatch":
                dispatch_handover_batch(
                    batch_id=batch.id,
                    dispatched_by=request.user,
                    dispatch_location_scan=request.POST.get("dispatch_location_scan", ""),
                )
            elif action == "refresh":
                refresh_handover_acceptance(batch_id=batch.id)
            elif action == "retry_kiz":
                command = retry_wb_marking_code(
                    batch_id=batch.id,
                    order_item_id=int(request.POST.get("order_item_id") or 0),
                    marking_scan=request.POST.get("marking_scan", ""),
                    requested_by=request.user,
                )
            elif action == "reroute_invalid_kiz":
                reroute = request_invalid_kiz_reroute(
                    handover_batch_id=batch.id,
                    order_id=int(request.POST.get("order_id") or 0),
                    route=request.POST.get("route", ""),
                    problem_tote_scan=request.POST.get("problem_tote_scan", ""),
                    requested_by=request.user,
                )
            else:
                raise ValueError("Неизвестная команда передачи.")
            detail_url = reverse("fbs:tsd_handover_detail", kwargs={"batch_id": batch.id})
            if request.POST.get("problem_resolution") == "1" and action in {
                "resolve_problem_order", "resolve_problem_order_without_label",
                "retry_kiz", "reroute_invalid_kiz", "confirm_ozon_canceled_return",
            }:
                return redirect(f"{detail_url}?problems=1&problem_saved=1")
            if action == "archive":
                archive_url = reverse("fbs:tsd_handover")
                return redirect(f"{archive_url}?status={FbsHandoverBatch.STATUS_ARCHIVED}")
            if action == "approve_verification_override":
                return redirect(f"{detail_url}?verification_override=approved")
            if action == "verify_order":
                query = urlencode(
                    {
                        "check": "1",
                        "verified_order": link.order.external_order_id,
                        "verified_box": link.box.qr_code,
                    }
                )
                return redirect(f"{detail_url}?{query}")
            if action in {"exclude_order", "retry_exclusion"}:
                return redirect(f"{detail_url}?exclusion=queued")
            if action == "release_exclusion":
                return redirect(f"{detail_url}?exclusion=controller_removed")
            if action == "confirm_ozon_canceled_return":
                return redirect(f"{detail_url}?ozon_return=queued")
            if action == "print_box_label":
                return redirect(f"{detail_url}?printed=box")
            if action == "print_supply_label":
                return redirect(f"{detail_url}?printed=supply")
            if action == "print_order_label":
                return redirect(f"{detail_url}?printed=order")
            if action == "reprint_ozon_order_qr":
                return redirect(f"{detail_url}?printed=ozon_qr")
            if action == "retry_kiz":
                query = {"kiz_requeued": command.order.external_order_id}
                if kiz_rescan_all_requested:
                    query["kiz_rescan_all"] = "1"
                return redirect(
                    f"{detail_url}?" + urlencode(query)
                )
            if action == "reroute_invalid_kiz":
                return redirect(
                    f"{detail_url}?"
                    + urlencode(
                        {
                            "kiz_route": reroute.route,
                            "target_handover": reroute.target_batch_id,
                            "problem_tote": request.POST.get(
                                "problem_tote_scan", ""
                            ),
                        }
                    )
                )
            return redirect(detail_url)
        except (FbsError, ValueError) as exc:
            error = str(exc)
        batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    from .controller_shipment_ui import controller_shipment_stage
    batch.movement_stage = controller_shipment_stage(batch)
    detail_context = _handover_detail_summary(
        batch,
        compact_controller=request_role == "fbs_controller",
    )
    problem_context = {}
    if request_role == "fbs_controller":
        from .controller_problems import controller_problem_context
        problem_context = controller_problem_context(batch, detail_context, request.user)
        if request.GET.get("problem_status") == "1":
            return JsonResponse({"ok": True,
                "fingerprint": problem_context.get("controller_problem_fingerprint", batch.status + batch.marketplace_state),
                "problem_count": problem_context.get("controller_problem_count", 0)})
        problem_context.update(
            controller_problem_force_open=request.GET.get("problems") == "1" or request.POST.get("problem_resolution") == "1",
            controller_problem_error=error if request.POST.get("problem_resolution") == "1" else "",
            controller_problem_saved=request.GET.get("problem_saved") == "1",
        )
    retryable_kiz_items = detail_context["handover_kiz_retry_items"]
    selected_kiz_retry = None
    if action == "retry_kiz":
        submitted_item_id = str(request.POST.get("order_item_id") or "").strip()
        selected_kiz_retry = next(
            (
                item
                for item in retryable_kiz_items
                if str(item.order_item_id) == submitted_item_id
            ),
            None,
        )
    elif kiz_rescan_all_requested and retryable_kiz_items:
        selected_kiz_retry = retryable_kiz_items[0]
    if kiz_rescan_all_requested and not retryable_kiz_items and not error:
        ok_message = "Все проблемные КИЗы отгрузки отсканированы и поставлены в очередь WB."
    can_manage_verification_override = bool(
        get_request_roles(request).intersection(
            {"head_manager", "director", "admin"}
        )
    )
    active_verification_override = None
    verification_override_blocker = ""
    if (
        can_manage_verification_override
        and batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB
        and batch.status
        in (FbsHandoverBatch.STATUS_OPEN, FbsHandoverBatch.STATUS_READY)
    ):
        (
            active_verification_override,
            verification_override_blocker,
        ) = handover_verification_override_status(batch)
    controller_canceled_tote = None
    if request_role == "fbs_controller":
        controller_session = (
            FbsControllerSession.objects.select_related("canceled_tote")
            .filter(
                controller=request.user,
                status=FbsControllerSession.STATUS_ACTIVE,
                check_totes__handover_batch=batch,
            )
            .order_by("-id")
            .first()
        )
        controller_canceled_tote = (
            controller_session.canceled_tote if controller_session else None
        )
    return render(
        request,
        "fbs/tsd_handover_detail.html",
        _base_context(
            request,
            page_title=f"FBS · Передача #{batch.id}",
            back_url=reverse("fbs:tsd_handover"),
            handover=batch,
            error=error,
            ok_message=ok_message,
            handover_check_open=(
                detail_context["handover_reconciliation_available"]
                and (
                    str(request.GET.get("check") or "") == "1"
                    or action == "verify_order"
                )
            ),
            handover_check_message=handover_check_message,
            handover_check_error=error if action == "verify_order" else "",
            kiz_retry_open=(bool(error and action == "retry_kiz")
            or bool(kiz_rescan_all_requested and selected_kiz_retry)) and request.POST.get("problem_resolution") != "1",
            kiz_retry_item_id=(
                str(selected_kiz_retry.order_item_id)
                if selected_kiz_retry is not None
                else str(request.POST.get("order_item_id") or "")
                if action == "retry_kiz"
                else ""
            ),
            kiz_retry_caption=(
                f"Заказ {selected_kiz_retry.order_number} · "
                f"{selected_kiz_retry.product_name}"
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_retry_label_number=(
                selected_kiz_retry.label_number
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_retry_product_barcode=(
                selected_kiz_retry.product_barcode
                if selected_kiz_retry is not None
                else ""
            ),
            kiz_rescan_all_requested=kiz_rescan_all_requested,
            kiz_retry_remaining_count=len(retryable_kiz_items),
            kiz_retry_error=error if action == "retry_kiz" else "",
            can_manage_verification_override=can_manage_verification_override,
            active_verification_override=active_verification_override,
            verification_override_blocker=verification_override_blocker,
            handover_controller_canceled_tote=controller_canceled_tote,
            **_handover_agent_scan_context(request),
            **detail_context,
            **problem_context,
        ),
        status=400 if error else 200,
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_box_label(request, batch_id: int, box_id: int):
    box = get_object_or_404(
        FbsHandoverBox.objects.select_related("batch__profile__agency"),
        pk=box_id,
        batch_id=batch_id,
    )
    if not box.label_file:
        if box.batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_WB:
            raise Http404
        response = render(
            request,
            "fbs/tsd_handover_internal_box_label.html",
            {"box": box, "handover": box.batch},
        )
        response["Cache-Control"] = "private, no-store"
        response["X-Robots-Tag"] = "noindex, nofollow"
        return response
    response = FileResponse(
        box.label_file.open("rb"),
        as_attachment=False,
        filename=Path(box.label_file.name).name,
        content_type="image/png",
    )
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_supply_label(request, batch_id: int):
    batch = get_object_or_404(
        FbsHandoverBatch.objects.select_related("profile__agency"),
        pk=batch_id,
    )
    is_internal_ozon = (
        batch.profile.marketplace == FbsIntegrationProfile.MARKETPLACE_OZON
        and not batch.supply_label_file
        and batch.status != FbsHandoverBatch.STATUS_OPEN
    )
    if not batch.supply_label_file and not is_internal_ozon:
        raise Http404
    if request.GET.get("raw") == "1":
        if is_internal_ozon:
            try:
                content = render_ozon_handover_internal_label(batch)
            except FbsLabelError as exc:
                response = HttpResponse(
                    f"Этикетка Ozon недоступна: {exc}",
                    content_type="text/plain; charset=utf-8",
                    status=409,
                )
                response["Cache-Control"] = "private, no-store"
                return response
            response = HttpResponse(content, content_type="image/png")
            response["Content-Disposition"] = (
                f'inline; filename="ozon-fbs-{batch.id}-internal.png"'
            )
        else:
            response = FileResponse(
                batch.supply_label_file.open("rb"),
                as_attachment=False,
                filename=Path(batch.supply_label_file.name).name,
                content_type="image/png",
            )
        response["X-Content-Type-Options"] = "nosniff"
        response["Cache-Control"] = "private, no-store"
        return response
    response = render(
        request,
        "fbs/tsd_handover_supply_label_print.html",
        {
            "handover": batch,
            "label_image_url": f"{request.path}?raw=1",
        },
    )
    response["Cache-Control"] = "private, no-store"
    response["X-Robots-Tag"] = "noindex, nofollow"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_ozon_document(request, batch_id: int, document_kind: str):
    if document_kind not in {
        OZON_HANDOVER_DOCUMENT_BARCODE,
        OZON_HANDOVER_DOCUMENT_PDF,
    }:
        raise Http404
    try:
        document = download_exact_ozon_handover_document(
            batch_id=batch_id,
            document_kind=document_kind,
        )
    except FbsIntegrationError as exc:
        response = HttpResponse(
            f"Документ Ozon недоступен: {exc}",
            content_type="text/plain; charset=utf-8",
            status=409,
        )
        response["Cache-Control"] = "private, no-store"
        return response
    response = HttpResponse(document.content, content_type=document.content_type)
    response["Content-Disposition"] = f'attachment; filename="{document.filename}"'
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_handover_driver_manifest(request, batch_id: int):
    batch = get_object_or_404(_handover_detail_queryset(), pk=batch_id)
    detail_context = _handover_detail_summary(batch)
    return render(
        request,
        "fbs/tsd_handover_driver_manifest.html",
        {
            "handover": batch,
            "generated_at": timezone.now(),
            **detail_context,
        },
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_demand(request):
    created = _query_int(request, "created")
    skipped = _query_int(request, "skipped")
    ok_message = ""
    if created:
        ok_message = f"Создано планов: {created}."
        if skipped:
            ok_message += f" Клиентов с блокировками: {skipped}."
    elif request.GET.get("generated") == "1":
        ok_message = "Новая потребность для планирования отсутствует."
        if skipped:
            ok_message = f"Планы не созданы. Клиентов с блокировками: {skipped}."
    floor_movement_id = _query_int(request, "floor_movement")
    if floor_movement_id:
        ok_message = (
            f"Заявка #{floor_movement_id} на перемещение товара на первый ярус "
            "передана водителю ричтрака."
        )
    return render(
        request,
        "fbs/tsd_demand.html",
        _demand_context(request, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_create_floor_movement(request):
    try:
        source_box_id = int(request.POST.get("source_box_id") or 0)
        target_pallet_id = int(request.POST.get("target_pallet_id") or 0)
        needed_qty = int(request.POST.get("needed_qty") or 0)
    except (TypeError, ValueError):
        source_box_id = target_pallet_id = needed_qty = 0
    try:
        result = create_floor_replenishment_movement(
            source_box_id=source_box_id,
            target_pallet_id=target_pallet_id,
            barcode=str(request.POST.get("barcode") or ""),
            needed_qty=needed_qty,
            requested_by=request.user,
            comment=str(request.POST.get("comment") or ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_demand.html",
            _demand_context(request, error=str(exc)),
            status=409,
        )
    return redirect(
        f"{reverse('fbs:tsd_demand')}?floor_movement={result.movement.id}"
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_generate_plans(request):
    try:
        result = generate_replenishment_plans(
            requested_by=request.user,
            requested_by_role=get_request_role(request),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_demand.html",
            _demand_context(request, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_demand")
    return redirect(
        f"{url}?generated=1&created={len(result.plans)}"
        f"&skipped={len(result.skipped_agencies)}"
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_GET
def tsd_picking(request):
    ok_message = ""
    error = ""
    missing_qty = _query_int(request, "missing_qty")
    missing_orders = _query_int(request, "missing_orders")
    if missing_qty:
        ok_message = (
            f"Снято с текущей волны: {missing_qty} шт. "
            f"Затронуто заказов: {missing_orders}. "
            "Проверка остатка передана кладовщику. Продолжайте отбор остальных позиций."
        )
        if request.GET.get("source_quarantined") == "1":
            ok_message += " Остаток в указанном коробе закрыт до пересчета."
    elif request.GET.get("equipment") == "1":
        ok_message = "ТСД привязан к рабочему месту и таре."
    elif request.GET.get("done") == "1":
        ok_message = "Заказ отобран."
    elif request.GET.get("released") == "1":
        ok_message = (
            "Заказ исключен из текущей волны и возвращен в ожидание товара."
        )
    elif request.GET.get("shortage") == "1":
        ok_message = (
            "Позиция отмечена как отсутствующая. "
            "Заказ поедет на проверку неполным."
        )
    elif request.GET.get("restock_done") == "1":
        ok_message = "Ошибочный отбор полностью возвращен в исходные FBS-короба."
    elif request.GET.get("prepared") == "1":
        tasks = _query_int(request, "tasks")
        waiting = _query_int(request, "waiting")
        failed = _query_int(request, "failed")
        ok_message = f"Создано заданий: {tasks}."
        if waiting or failed:
            ok_message += f" Ожидают товар: {waiting}; ошибки SKU: {failed}."
    if request.GET.get("wave_stopped") == "1":
        error = (
            "Задание этой волны остановлено. Список обновлен; "
            "выберите доступную волну."
        )
    return render(
        request,
        "fbs/tsd_picking_list.html",
        _picking_context(request, error=error, ok_message=ok_message),
    )


def _pick_restock_context(request, restock, *, error="", ok_message=""):
    state = pick_restock_state(restock.id)
    line = state["line"]
    recent_scans = restock.scans.select_related("line").order_by("-created_at", "-id")[:8]
    stage = state["stage"]
    stage_copy = {
        FbsPickRestockScan.STAGE_WORKSTATION: (
            "Забрать из тары" if state["source_tote"] else "Забрать со стола",
            (
                "Подойдите к указанной служебной таре и подтвердите ее QR."
                if state["source_tote"]
                else "Подойдите к указанному рабочему столу и подтвердите его QR."
            ),
            "QR служебной тары" if state["source_tote"] else "QR рабочего стола",
        ),
        FbsPickRestockScan.STAGE_PICKUP_ITEM: (
            "Забрать товар из тары" if state["source_tote"] else "Снять товар со стола",
            (
                "Сканируйте каждую единицу, которую забираете из тары. "
                "Перейти к размещению можно только после полного сканирования."
                if state["source_tote"]
                else "Сканируйте каждую единицу, которую забираете со стола. "
                "Перейти к размещению можно только после полного сканирования."
            ),
            "Штрихкод товара из тары" if state["source_tote"] else "Штрихкод товара со стола",
        ),
        FbsPickRestockScan.STAGE_CELL: (
            "Карантинная ячейка" if state["is_quarantine"] else "Исходная ячейка",
            (
                "Отсканируйте QR карантинного места, указанного контролером."
                if state["is_quarantine"]
                else "Отсканируйте QR места, из которого товар был отобран."
            ),
            "QR ячейки",
        ),
        FbsPickRestockScan.STAGE_BOX: (
            "Карантинный короб" if state["is_quarantine"] else "Исходный короб",
            (
                "Отсканируйте QR карантинного FBS-короба."
                if state["is_quarantine"]
                else "Отсканируйте QR исходного короба. Все товары этого короба вернутся одним проходом."
            ),
            "QR короба",
        ),
        FbsPickRestockScan.STAGE_ITEM: (
            "Товар",
            "Сканируйте штрихкод каждой единицы. КИЗ для этой позиции не требуется. После последней единицы система сама даст следующий короб.",
            "Штрихкод товара",
        ),
    }
    if state.get("destination_event_id"):
        stage_copy[FbsPickRestockScan.STAGE_CELL] = ("Место возврата", "Отсканируйте QR места выбранного короба.", "QR места")
        stage_copy[FbsPickRestockScan.STAGE_BOX] = ("Короб возврата", "Отсканируйте QR выбранного короба клиента.", "QR короба")
    title, instruction, input_label = stage_copy.get(
        stage,
        ("Завершено", "Все единицы возвращены.", "Скан"),
    )
    show_alternatives = request.GET.get("choose_box") == "1" or request.POST.get("action") == "select_destination"
    alternative_boxes = []
    if show_alternatives:
        from .services.restock_destinations import destination_choices
        alternative_boxes = destination_choices(state)
    destination_cell = state.get("destination_cell")
    destination_box = state.get("destination_box")
    location_label = _pick_location_label(destination_cell) if destination_cell else ""
    if stage == FbsPickRestockScan.STAGE_BOX and destination_box is not None:
        box_number_suffix = _pick_box_number_suffix(destination_box.box_code)
        if box_number_suffix:
            title = f"{title} · {box_number_suffix}"
    return _base_context(
        request,
        page_title=f"Возврат отбора #{restock.id}",
        back_url=reverse("fbs:tsd_picker_returns" if get_request_role(request) == "picker" else "fbs:tsd_picking"),
        alternative_boxes=alternative_boxes,
        show_alternatives=show_alternatives,
        restock=state["request"],
        restock_state=state,
        source_workstation=state["source_workstation"],
        source_tote=state["source_tote"],
        destination_cell=destination_cell,
        destination_box=destination_box,
        is_quarantine=state["is_quarantine"],
        line=line,
        stage=stage,
        stage_title=title,
        stage_instruction=instruction,
        scan_input_label=input_label,
        source_location_label=location_label,
        selected_pickup_line_id=str(request.POST.get("pickup_line_id") or ""),
        request_token=str(uuid.uuid4()),
        recent_scans=recent_scans,
        error=error,
        ok_message=ok_message,
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_restock(request, request_id: int):
    restock = get_object_or_404(
        FbsPickRestockRequest.objects.select_related("batch__workstation", "assigned_to"),
        pk=request_id,
    )
    if restock.assigned_to_id not in {None, request.user.id}:
        return HttpResponseForbidden("Задание возврата уже выполняет другой подборщик.")
    if request.method == "POST":
        action = str(request.POST.get("action") or "scan").strip()
        try:
            if action == "claim":
                claim_pick_restock_request(
                    request_id=restock.id,
                    assigned_to=request.user,
                )
                return redirect("fbs:tsd_pick_restock", request_id=restock.id)
            if action == "select_destination":
                from .services.restock_destinations import select_destination
                select_destination(request_id=restock.pk, line_id=request.POST.get("line_id"),
                    box_scan=request.POST.get("box_scan"), performed_by=request.user,
                    space_confirmed=request.POST.get("space_confirmed") == "1")
                return redirect("fbs:tsd_pick_restock", request_id=restock.pk)
            if action != "scan":
                raise FbsError("Неизвестная команда возврата отбора.")
            result = scan_pick_restock(
                request_id=restock.id,
                stage=request.POST.get("stage", ""),
                scan_value=request.POST.get("scan_value", ""),
                request_token=request.POST.get("request_token", ""),
                performed_by=request.user,
                destination_event_id=request.POST.get("destination_event_id"),
                pickup_line_id=request.POST.get("pickup_line_id"),
            )
        except FbsError as exc:
            restock.refresh_from_db()
            return render(
                request,
                "fbs/tsd_pick_restock_detail.html",
                _pick_restock_context(request, restock, error=str(exc)),
                status=409,
            )
        if result.request.status == FbsPickRestockRequest.STATUS_COMPLETED:
            if get_request_role(request) == "picker":
                return redirect(f"{reverse('fbs:tsd_picker_returns')}?done=1")
            return redirect(f"{reverse('fbs:tsd_picking')}?restock_done=1")
        restock.refresh_from_db()
        return render(
            request,
            "fbs/tsd_pick_restock_detail.html",
            _pick_restock_context(
                request,
                restock,
                ok_message=result.event.message,
            ),
        )
    return render(
        request,
        "fbs/tsd_pick_restock_detail.html",
        _pick_restock_context(request, restock),
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_bind_pick_equipment(request):
    try:
        batch = claim_next_pick_batch(
            assigned_to=request.user,
            workstation_scan="",
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    _store_picker_cart(request, batch.cart)
    if batch.status == FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:tsd_pick_handover", batch_id=batch.id)
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_clear_pick_equipment(request):
    request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
    return redirect("fbs:tsd_picking")


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_next_pick_batch(request):
    cart = _picker_cart(request)
    if cart is None:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(
                request,
                error="Сначала отсканируйте свободную тару.",
            ),
            status=409,
        )
    try:
        batch = claim_next_pick_batch(
            assigned_to=request.user,
            workstation_scan="",
            cart_scan=cart.barcode,
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    if batch.status == FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:tsd_pick_handover", batch_id=batch.id)
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_labels(request):
    ok_message = ""
    if request.GET.get("requested") == "1":
        ok_message = "Заказ отобран. Этикетка маркетплейса запрошена."
    elif request.GET.get("done") == "1":
        ok_message = "Этикетка подтверждена. Заказ готов к передаче."
    return render(
        request,
        "fbs/tsd_label_list.html",
        _labels_context(request, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_GET
def tsd_label_file(request, label_id: int):
    label = get_object_or_404(
        _accessible_label_queryset(request),
        pk=label_id,
        status__in=(FbsOrderLabel.STATUS_READY, FbsOrderLabel.STATUS_APPLIED),
    )
    if not label.file:
        raise Http404
    content_type = {
        FbsOrderLabel.FORMAT_PDF: "application/pdf",
        FbsOrderLabel.FORMAT_PNG: "image/png",
        FbsOrderLabel.FORMAT_ZPL: "application/octet-stream",
    }.get(label.label_format, "application/octet-stream")
    response = FileResponse(
        label.file.open("rb"),
        as_attachment=False,
        filename=Path(label.file.name).name,
        content_type=content_type,
    )
    response["X-Content-Type-Options"] = "nosniff"
    response["Cache-Control"] = "private, no-store"
    return response


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_label_detail(request, label_id: int):
    label = get_object_or_404(_accessible_label_queryset(request), pk=label_id)
    if get_request_role(request) == "fbs_controller":
        controller_batch = (
            FbsPickBatch.objects.filter(
                tasks__order_id=label.order_id,
                verification_assigned_to=request.user,
                status=FbsPickBatch.STATUS_VERIFICATION,
            )
            .order_by("verification_started_at", "id")
            .first()
        )
        if controller_batch is not None:
            return redirect(
                "fbs:tsd_pick_verification",
                batch_id=controller_batch.id,
            )
        if request.method == "POST":
            return render(
                request,
                "fbs/tsd_label_detail.html",
                _label_context(
                    request,
                    label,
                    error=(
                        "Контролер подтверждает этикетку только внутри активной "
                        "проверки волны. Вернитесь на рабочий стол и отсканируйте тару."
                    ),
                ),
                status=409,
            )
    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(
                request,
                label,
                ok_message=(
                    "Этикетка отправлена в очередь печати."
                    if request.GET.get("printed") == "1"
                    else ""
                ),
            ),
        )
    action = str(request.POST.get("action") or "confirm").strip()
    if action == "print":
        try:
            queue_fbs_order_label_print(
                label_id=label.id,
                requested_by=request.user,
                force=True,
                **_desktop_label_print_target(request),
            )
        except FbsError as exc:
            label.refresh_from_db()
            return render(
                request,
                "fbs/tsd_label_detail.html",
                _label_context(request, label, error=str(exc)),
                status=400,
            )
        return redirect(f"{reverse('fbs:tsd_label_detail', kwargs={'label_id': label.id})}?printed=1")
    if action != "confirm":
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(request, label, error="Неизвестная команда этикетки."),
            status=400,
        )
    label_scan = str(request.POST.get("label_scan") or "").strip()
    try:
        confirm_order_label_scan(
            label_id=label.id,
            label_scan=label_scan,
            performed_by=request.user,
        )
    except FbsError as exc:
        record_order_label_scan_events(
            order_id=label.order_id,
            scan_value=label_scan,
            expected_value=label.barcode,
            result=FbsPickScanEvent.RESULT_ERROR,
            message=str(exc),
            created_by=request.user,
        )
        label.refresh_from_db()
        return render(
            request,
            "fbs/tsd_label_detail.html",
            _label_context(request, label, error=str(exc)),
            status=400,
        )
    next_batch = (
        FbsPickBatch.objects.filter(
            tasks__order_id=label.order_id,
            verification_assigned_to=request.user,
            status=FbsPickBatch.STATUS_VERIFICATION,
        )
        .order_by("created_at", "id")
        .first()
    )
    if next_batch is not None:
        return redirect("fbs:tsd_pick_verification", batch_id=next_batch.id)
    completed_batch = (
        FbsPickBatch.objects.filter(
            tasks__order_id=label.order_id,
            verification_assigned_to=request.user,
            status=FbsPickBatch.STATUS_DONE,
        )
        .order_by("-completed_at", "-id")
        .first()
    )
    if completed_batch is not None:
        if get_request_role(request) == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect(
            f"{reverse('fbs:operator_wave_detail', kwargs={'batch_id': completed_batch.id})}?done=1"
        )
    return redirect(f"{reverse('fbs:tsd_labels')}?done=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_prepare_picking(request):
    try:
        max_orders_per_batch = int(request.POST.get("wave_size") or 50)
    except (TypeError, ValueError):
        max_orders_per_batch = 0
    try:
        result = prepare_pick_queue(
            max_orders_per_batch=max_orders_per_batch,
            max_units_per_batch=100,
            performed_by=request.user,
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    rejection_message = format_queue_rejections(
        getattr(result, "rejected_orders", ())
    )
    if rejection_message and not result.tasks_created:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(
                request,
                error=f"Волна не создана. {rejection_message}",
            ),
            status=409,
        )
    if rejection_message:
        messages.warning(request, rejection_message)
    url = reverse("fbs:tsd_picking")
    return redirect(
        f"{url}?prepared=1&tasks={result.tasks_created}"
        f"&waiting={result.awaiting_stock_orders}"
        f"&failed={result.validation_failed_orders}"
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_pick_task(request, task_id: int):
    get_object_or_404(FbsPickTask, pk=task_id)
    try:
        task = claim_pick_task(
            task_id=task_id,
            assigned_to=request.user,
            workstation_scan=request.POST.get("workstation_scan", ""),
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    allocation = (
        _pick_allocation_queryset()
        .filter(
            pick_task=task,
            status__in=(
                FbsOrderStockAllocation.STATUS_RESERVED,
                FbsOrderStockAllocation.STATUS_PICKING,
            ),
        )
        .order_by("balance__box__pallet__cell__cell_code", "id")
        .first()
    )
    if allocation is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation.id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_POST
def tsd_claim_pick_batch(request, batch_id: int):
    get_object_or_404(FbsPickBatch, pk=batch_id)
    # Очередь строгая: волны выдаются по возрасту, сборщик их не выбирает.
    # В интерфейсе кнопки взятия конкретной волны нет, но адрес открыт
    # напрямую, поэтому порядок проверяем здесь. Свою уже начатую волну
    # продолжить можно всегда.
    own_active = FbsPickBatch.objects.filter(
        pk=batch_id,
        assigned_to=request.user,
        picking_completed_at__isnull=True,
    ).exists()
    if not own_active:
        oldest_free_id = order_wave_queue_queryset().values_list(
            "id", flat=True
        ).first()
        if oldest_free_id != batch_id:
            return render(
                request,
                "fbs/tsd_picking_list.html",
                _picking_context(
                    request,
                    error=(
                        "Волны выдаются по очереди, с самой старой. "
                        "Нажмите «Получить следующую волну»."
                    ),
                ),
                status=409,
            )
    try:
        batch = claim_pick_batch(
            batch_id=batch_id,
            assigned_to=request.user,
            workstation_scan=request.POST.get("workstation_scan", ""),
            cart_scan=request.POST.get("cart_scan", ""),
        )
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_picking_list.html",
            _picking_context(request, error=str(exc)),
            status=409,
        )
    allocation_id = _next_pick_allocation_id(batch)
    if allocation_id is None:
        return redirect("fbs:tsd_picking")
    return redirect("fbs:tsd_pick_allocation", allocation_id=allocation_id)


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_allocation(request, allocation_id: int):
    allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation_id)
    task = allocation.pick_task
    if task is None:
        return HttpResponseForbidden("Позиция не включена в задание")
    if task.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Задание назначено другому сборщику")
    if allocation.status == FbsOrderStockAllocation.STATUS_PICKED:
        return redirect("fbs:tsd_picking")
    if (
        task.status != FbsPickTask.STATUS_IN_PROGRESS
        or task.batch.status != FbsPickBatch.STATUS_IN_PROGRESS
        or allocation.status != FbsOrderStockAllocation.STATUS_PICKING
    ):
        request.session.pop(_pick_cell_session_key(allocation.id), None)
        request.session.pop(_pick_box_session_key(allocation.id), None)
        return redirect(f"{reverse('fbs:tsd_picking')}?wave_stopped=1")
    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(
                request,
                allocation,
                new_box_notice=_pop_pick_new_box_notice(request, allocation.id),
            ),
        )

    action = str(request.POST.get("action") or "").strip()
    cell_session_key = _pick_cell_session_key(allocation.id)
    session_key = _pick_box_session_key(allocation.id)
    requires_box_scan = pick_requires_box_scan(allocation.balance.box)
    if action == "not_found":
        if request.POST.get("confirm_not_found") != "1":
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request,
                    allocation,
                    error="Подтвердите, что указанной позиции нет на месте.",
                ),
                status=400,
            )
        try:
            missing_qty = int(request.POST.get("missing_qty") or 0)
        except (TypeError, ValueError):
            missing_qty = 0
        try:
            result = report_pick_missing_quantity(
                allocation_id=allocation.id,
                missing_qty=missing_qty,
                reason=request.POST.get("reason", ""),
                reported_by=request.user,
            )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session.pop(cell_session_key, None)
        request.session.pop(session_key, None)
        query = urlencode(
            {
                "missing_qty": result.released_qty,
                "missing_orders": result.affected_order_count,
                "source_quarantined": "1" if result.source_quarantined else "0",
            }
        )
        return redirect(f"{reverse('fbs:tsd_picking')}?{query}")
    if action == "scan_cell":
        cell_scan = str(request.POST.get("cell_scan") or "").strip()
        try:
            if requires_box_scan:
                raise FbsError("Отсканируйте короб")
            validate_pick_cell_scan(
                allocation_id=allocation.id,
                cell_scan=cell_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(cell_session_key, None)
            request.session.pop(session_key, None)
            original_error = str(exc)
            scan_error = (
                "Неверно отсканирована ячейка. "
                "Отсканируйте указанную ячейку еще раз."
                if "Скан ячейки FBS не совпадает с маршрутом" in original_error
                else original_error
            )
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request,
                    allocation,
                    error=scan_error,
                    scan_error=scan_error,
                ),
                status=400,
            )
        request.session[cell_session_key] = cell_scan
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(request, allocation, ok_message="Ячейка подтверждена."),
        )
    if action == "scan_box":
        if not requires_box_scan and not str(request.session.get(cell_session_key) or "").strip():
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error="Сначала подтвердите ячейку."),
                status=400,
            )
        box_scan = str(request.POST.get("box_scan") or "").strip()
        try:
            validate_pick_box_scan(
                allocation_id=allocation.id,
                box_scan=box_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(session_key, None)
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session[session_key] = box_scan
        return render(
            request,
            "fbs/tsd_pick_allocation.html",
            _pick_allocation_context(request, allocation, ok_message="Короб подтвержден."),
        )

    if action == "complete":
        cell_scan = (
            "" if requires_box_scan
            else str(request.session.get(cell_session_key) or "").strip()
        )
        box_scan = str(request.session.get(session_key) or "").strip()
        item_scan = str(request.POST.get("item_scan") or "").strip()
        if (requires_box_scan and not box_scan) or (not requires_box_scan and not cell_scan):
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(
                    request, allocation, error=(
                        "Сначала подтвердите короб."
                        if requires_box_scan else "Сначала подтвердите ячейку."
                    )
                ),
                status=400,
            )
        try:
            complete_pick_allocation(
                allocation_id=allocation.id,
                cell_scan=cell_scan,
                box_scan=box_scan,
                item_scan=item_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_pick_allocation.html",
                _pick_allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        allocation.refresh_from_db()
        if allocation.status == FbsOrderStockAllocation.STATUS_PICKED:
            next_allocation = _open_pick_allocations(task.batch).first()
            request.session.pop(cell_session_key, None)
            request.session.pop(session_key, None)
            if next_allocation is not None:
                current_cell_id = allocation.balance.box.pallet.cell_id
                next_cell_id = next_allocation.balance.box.pallet.cell_id
                current_box_id = allocation.balance.box_id
                next_box_id = next_allocation.balance.box_id
                next_requires_box_scan = pick_requires_box_scan(next_allocation.balance.box)
                if (
                    not requires_box_scan
                    and not next_requires_box_scan
                    and current_cell_id == next_cell_id
                ):
                    request.session[
                        _pick_cell_session_key(next_allocation.id)
                    ] = cell_scan
                if (
                    requires_box_scan
                    and next_requires_box_scan
                    and current_box_id == next_box_id
                    and current_cell_id == next_cell_id
                ):
                    request.session[
                        _pick_box_session_key(next_allocation.id)
                    ] = box_scan
                elif not _open_pick_allocations(task.batch).filter(
                    balance__box_id=current_box_id,
                ).exists():
                    request.session[PICK_NEW_BOX_NOTICE_SESSION_KEY] = {
                        "allocation_id": next_allocation.id,
                        "completed_box_code": allocation.balance.box.box_code,
                        "next_box_code": next_allocation.balance.box.box_code,
                        "completed_direct_cell": not requires_box_scan,
                        "completed_location_label": fbs_box_physical_location_label(allocation.balance.box),
                        "next_direct_cell": not next_requires_box_scan,
                        "next_location_label": fbs_box_physical_location_label(
                            next_allocation.balance.box
                        ),
                        "same_cell": current_cell_id == next_cell_id,
                    }
                return redirect(
                    "fbs:tsd_pick_allocation",
                    allocation_id=next_allocation.id,
                )
        else:
            return redirect("fbs:tsd_pick_allocation", allocation_id=allocation.id)
        task.batch.refresh_from_db()
        if task.batch.status == FbsPickBatch.STATUS_VERIFICATION:
            return redirect("fbs:tsd_pick_handover", batch_id=task.batch_id)
        return redirect("fbs:tsd_picking")

    return render(
        request,
        "fbs/tsd_pick_allocation.html",
        _pick_allocation_context(request, allocation, error="Неизвестная команда."),
        status=400,
    )


@fbs_module_required
@role_required(*PICKER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_handover(request, batch_id: int):
    batch = get_object_or_404(
        FbsPickBatch.objects.select_related("agency", "assigned_to", "workstation", "cart"),
        pk=batch_id,
    )
    if batch.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Волна назначена другому сборщику")
    handover_redirect_url = f"{reverse('fbs:tsd_picking')}?handed_over=1"
    if batch.picking_completed_at is not None:
        request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
        if request.GET.get("handover_status") == "1":
            response = JsonResponse(
                {
                    "handed_over": True,
                    "redirect_url": handover_redirect_url,
                }
            )
            response["Cache-Control"] = "no-store"
            return response
        return redirect(handover_redirect_url)
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return HttpResponseForbidden("Волна еще не собрана полностью")
    if request.GET.get("handover_status") == "1":
        response = JsonResponse({"handed_over": False})
        response["Cache-Control"] = "no-store"
        return response
    error = ""
    if request.method == "POST":
        try:
            handover_pick_batch_for_verification(
                batch_id=batch.id,
                workstation_scan=request.POST.get("workstation_scan", ""),
                performed_by=request.user,
            )
            request.session.pop(PICK_EQUIPMENT_SESSION_KEY, None)
            return redirect(handover_redirect_url)
        except FbsError as exc:
            error = str(exc)
    workstation_recommendation = None
    try:
        workstation_recommendation = assign_pick_handover_workstation(
            batch_id=batch.id
        )
        batch.refresh_from_db(fields=["workstation", "updated_at"])
    except FbsError as exc:
        if not error:
            error = str(exc)
    return render(
        request,
        "fbs/tsd_pick_handover.html",
        _base_context(
            request,
            page_title=f"FBS · Сдать волну #{batch.id}",
            back_url=reverse("fbs:tsd_picking"),
            batch=batch,
            workstation_recommendation=workstation_recommendation,
            error=error,
        ),
        status=400 if error else 200,
    )


def _tsd_pick_verification_response(request, batch_id: int):
    batch = get_object_or_404(
        FbsPickBatch.objects.select_related("agency", "assigned_to", "workstation", "cart"),
        pk=batch_id,
    )
    request_role = get_request_role(request)
    if batch.verification_assigned_to_id is None:
        if request_role == "fbs_controller":
            workstation = get_controller_workstation(request)
            if workstation is None:
                return redirect("fbs:controller_home")
            if batch.workstation_id != workstation.id:
                return HttpResponseForbidden(
                    "Волна доставлена на другое рабочее место контролера"
                )
            return redirect("fbs:controller_home")
        if request.method == "POST" and request.POST.get("action") == "claim":
            try:
                claim_pick_batch_verification(batch_id=batch.id, assigned_to=request.user)
            except FbsError as exc:
                return render(
                    request,
                    "fbs/tsd_pick_verification_claim.html",
                    _base_context(
                        request,
                        batch=batch,
                        back_url=(
                            reverse("fbs:controller_home")
                            if request_role == "fbs_controller"
                            else reverse("fbs:operator_waves")
                        ),
                        error=str(exc),
                    ),
                    status=409,
                )
            return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
        return render(
            request,
            "fbs/tsd_pick_verification_claim.html",
            _base_context(
                request,
                batch=batch,
                back_url=(
                    reverse("fbs:controller_home")
                    if request_role == "fbs_controller"
                    else reverse("fbs:operator_waves")
                ),
            ),
        )
    if batch.verification_assigned_to_id != request.user.id:
        return HttpResponseForbidden("Проверку волны выполняет другой оператор")
    marketplace_rejection_return = None
    marketplace_rejection_allocation = None
    if request_role == "fbs_controller":
        marketplace_rejection_return = (
            FbsPickRestockRequest.objects.select_related(
                "order",
                "handover_assignment",
            )
            .filter(
                batch=batch,
                reason_code=FbsPickRestockRequest.REASON_MARKETPLACE,
                status__in=(
                    FbsPickRestockRequest.STATUS_WAITING_MARKETPLACE,
                    FbsPickRestockRequest.STATUS_QUEUED,
                    FbsPickRestockRequest.STATUS_FAILED,
                ),
                source_tote__isnull=True,
            )
            .order_by("created_at", "id")
            .first()
        )
        if marketplace_rejection_return is not None:
            marketplace_rejection_allocation = (
                _pick_allocation_queryset()
                .filter(
                    pick_task__batch=batch,
                    pick_task__order_id=marketplace_rejection_return.order_id,
                    status=FbsOrderStockAllocation.STATUS_PICKED,
                    qty_picked__gt=0,
                )
                .order_by("id")
                .first()
            )
    if (
        marketplace_rejection_return is not None
        and marketplace_rejection_allocation is not None
    ):
        error = ""
        if (
            request.method == "POST"
            and str(request.POST.get("action") or "").strip()
            == "confirm_marketplace_return_tote"
        ):
            try:
                confirm_marketplace_rejected_order_return(
                    request_id=marketplace_rejection_return.id,
                    canceled_tote_scan=str(
                        request.POST.get("service_tote_scan") or ""
                    ).strip(),
                    performed_by=request.user,
                )
            except FbsError as exc:
                error = str(exc)
                marketplace_rejection_return.refresh_from_db()
            else:
                return redirect(
                    f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}"
                    "?problem=1"
                )
        prompt = _verification_service_tote_prompt(
            batch=batch,
            allocation=marketplace_rejection_allocation,
            controller=request.user,
            problem_kind="marketplace_rejected",
            reason=marketplace_rejection_return.reason,
            restock_request=marketplace_rejection_return,
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                marketplace_rejection_allocation,
                error=error,
                item_selected=True,
                service_tote_prompt=prompt,
            ),
            status=409 if error else 200,
        )
    if batch.status == FbsPickBatch.STATUS_DONE:
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return HttpResponseForbidden("Волна еще не готова к проверке")
    label = _next_verification_label(batch)
    allocation = None if label is not None else _next_verification_allocation(batch)
    if allocation is None and label is None:
        allocation = _next_incomplete_verification_allocation(batch)
    if allocation is None:
        if label is not None:
            if request_role == "fbs_controller":
                error = ""
                ok_message = (
                    "Этикетка отправлена в очередь печати."
                    if request.GET.get("printed") == "1"
                    else ""
                )
                if request.method == "GET" and not _fullbox_desktop_supports_direct_print(request):
                    # Экран контролера дошел до шага этикетки — ставим печать
                    # сразу, не дожидаясь фонового обработчика очереди.
                    # Дедупликация по card_id делает вызов идемпотентным:
                    # если задание уже создано, вернется существующее.
                    try:
                        if is_preloaded_ozon_order_label(label):
                            queue_fbs_preloaded_ozon_order_label_print(label_id=label.id)
                        else:
                            queue_fbs_order_label_print(label_id=label.id)
                    except FbsError:
                        pass
                if request.method == "POST":
                    action = str(request.POST.get("action") or "confirm_label").strip()
                    if action == "print_label":
                        try:
                            print_mode = str(
                                request.POST.get("print_mode") or "automatic"
                            ).strip()
                            print_scope = ""
                            if print_mode != "manual":
                                verification_started = (
                                    batch.verification_started_at.isoformat()
                                    if batch.verification_started_at is not None
                                    else "not-started"
                                )
                                print_scope = (
                                    f"verification:{batch.id}:{verification_started}"
                                )
                            if is_preloaded_ozon_order_label(label):
                                queue_fbs_preloaded_ozon_order_label_print(
                                    label_id=label.id,
                                    requested_by=request.user,
                                    force=True,
                                    **_desktop_label_print_target(request),
                                )
                            else:
                                queue_fbs_order_label_print(
                                    label_id=label.id,
                                    requested_by=request.user,
                                    force=True,
                                    print_scope=print_scope,
                                    **_desktop_label_print_target(request),
                                )
                        except FbsError as exc:
                            error = str(exc)
                        else:
                            return redirect(
                                f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}?printed=1"
                            )
                    elif action == "confirm_label":
                        label_scan = str(request.POST.get("label_scan") or "").strip()
                        try:
                            if batch.workstation_id is None:
                                raise ValueError("У волны не указано рабочее место контролера.")
                            pick_tote_context = controller_pick_tote_for_batch(
                                batch_id=batch.id
                            )
                            if pick_tote_context is None:
                                raise ValueError(
                                    "Сначала примите тару подбора и выберите тару проверки."
                                )
                            confirm_order_label_to_check_tote(
                                label_id=label.id,
                                label_scan=label_scan,
                                pick_batch_id=batch.id,
                                performed_by=request.user,
                            )
                        except (FbsError, ValueError) as exc:
                            record_order_label_scan_events(
                                order_id=label.order_id,
                                scan_value=label_scan,
                                expected_value=label.barcode,
                                result=FbsPickScanEvent.RESULT_ERROR,
                                message=str(exc),
                                created_by=request.user,
                            )
                            error = str(exc)
                        else:
                            # Вся волна отсканирована — тара освобождается сама,
                            # нажимать ничего не нужно. Если освободить нельзя
                            # (выключен флаг, чужой контролер), остается прежний
                            # путь с подтверждением пустоты на рабочем столе.
                            released_tote = auto_release_pick_tote_if_complete(
                                pick_batch_id=batch.id
                            )
                            if released_tote is not None:
                                released_query = urlencode(
                                    {
                                        "tote_released": "1",
                                        "tote": released_tote.tote.barcode,
                                    }
                                )
                                return redirect(
                                    f"{reverse('fbs:controller_home')}?{released_query}"
                                )
                            if mark_pick_tote_awaiting_empty(
                                pick_batch_id=batch.id
                            ):
                                return redirect(
                                    f"{reverse('fbs:controller_home')}?empty={batch.id}"
                                )
                            return redirect(
                                "fbs:tsd_pick_verification",
                                batch_id=batch.id,
                            )
                    else:
                        error = "Неизвестная команда этикетки."
                    label.refresh_from_db()
                return render(
                    request,
                    "fbs/tsd_pick_verification.html",
                    _pick_verification_label_context(
                        request,
                        batch,
                        label,
                        error=error,
                        ok_message=ok_message,
                    ),
                    status=400 if error else 200,
                )
            return redirect("fbs:tsd_label_detail", label_id=label.id)
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if request.method == "GET":
        ok_message = ""
        if request.GET.get("problem") == "1":
            ok_message = (
                "Проблемный заказ исключен из проверки и передан комплектовщику "
                "для размещения в карантин."
            )
        elif request.GET.get("problem_returned") == "1":
            ok_message = (
                "Товар возвращен из проблемной тары в тару проверки. "
                "Отсканируйте товар для продолжения проверки."
            )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                allocation,
                ok_message=ok_message,
            ),
        )
    action = str(request.POST.get("action") or "select_item").strip()
    if action not in {"select_item", "verify"}:
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                allocation,
                error="Неизвестная команда проверки.",
            ),
            status=400,
        )
    item_scan = str(request.POST.get("item_scan") or "").strip()
    marking_scan = str(request.POST.get("marking_scan") or "").strip()
    legacy_marking_absent = (
        str(request.POST.get("legacy_marking_absent") or "").strip() == "1"
    )
    expiry_date_value = str(request.POST.get("expiry_date") or "").strip()
    selected_allocation = allocation
    inferred_marking_scan = ""
    if action == "verify":
        try:
            selected_allocation_id = int(request.POST.get("allocation_id") or 0)
        except (TypeError, ValueError):
            selected_allocation_id = 0
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(),
            pk=selected_allocation_id,
            pick_task__batch=batch,
        )
        if request_role == "fbs_controller" and _marketplace_order_is_canceled(
            selected_allocation.pick_task.order
        ):
            canceled_reason = "Маркетплейс отменил этот заказ."
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    error=canceled_reason,
                    submitted_item_scan=item_scan,
                    item_selected=True,
                    service_tote_prompt=_verification_service_tote_prompt(
                        batch=batch,
                        allocation=selected_allocation,
                        controller=request.user,
                        problem_kind="canceled",
                        reason=canceled_reason,
                        initial_item_scan=item_scan,
                    ),
                ),
                status=409,
            )
    else:
        try:
            selected_allocation = find_pick_verification_allocation(
                batch_id=batch.id,
                item_scan=item_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            error_text = str(exc)
            unknown_scan_info = None
            if (
                request_role == "fbs_controller"
                and item_scan
                and "отсутствует в текущей таре" in error_text.casefold()
            ):
                try:
                    problem_item = record_extra_problem_tote_item(
                        pick_batch_id=batch.id,
                        scanned_value=item_scan,
                        performed_by=request.user,
                    )
                except FbsError as problem_exc:
                    error_text = str(problem_exc)
                else:
                    unknown_scan_info = SimpleNamespace(
                        scanned_value=item_scan,
                        tote_name=problem_item.problem_tote.name,
                        tote_barcode=problem_item.problem_tote.barcode,
                    )
                    error_text = (
                        f"Товара со штрихкодом {item_scan} в этой таре нет. "
                        "Скан не засчитан; положите лишний товар в проблемную тару "
                        f"{problem_item.problem_tote.name} "
                        f"({problem_item.problem_tote.barcode})."
                    )
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    allocation,
                    error=error_text,
                    unknown_scan_info=unknown_scan_info,
                ),
                status=400,
            )

        if request_role == "fbs_controller" and _marketplace_order_is_canceled(
            selected_allocation.pick_task.order
        ):
            canceled_reason = "Маркетплейс отменил этот заказ."
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    error=canceled_reason,
                    submitted_item_scan=item_scan,
                    item_selected=True,
                    service_tote_prompt=_verification_service_tote_prompt(
                        batch=batch,
                        allocation=selected_allocation,
                        controller=request.user,
                        problem_kind="canceled",
                        reason=canceled_reason,
                        initial_item_scan=item_scan,
                    ),
                ),
                status=409,
            )

        matched_item_scan, inferred_marking_scan = resolve_verification_item_scan(
            selected_allocation,
            item_scan,
        )
        if inferred_marking_scan:
            item_scan = matched_item_scan
            marking_scan = inferred_marking_scan

        from .services.traceability import (
            controller_marking_scan_required,
            metadata_requirements,
        )

        requirements = metadata_requirements(selected_allocation.order_item)
        optional_wb_marking = wb_optional_marking_available(
            selected_allocation.order_item
        )
        needs_follow_up = (
            controller_marking_scan_required(selected_allocation)
            or requirements.expiry_required
            or optional_wb_marking
        )
        if needs_follow_up and (not inferred_marking_scan or requirements.expiry_required):
            return render(
                request,
                "fbs/tsd_pick_verification.html",
                _pick_verification_context(
                    request,
                    batch,
                    selected_allocation,
                    submitted_item_scan=item_scan,
                    submitted_marking_scan=inferred_marking_scan,
                    item_selected=True,
                ),
            )
    try:
        verify_pick_allocation_unit(
            allocation_id=selected_allocation.id,
            item_scan=item_scan,
            marking_scan=marking_scan,
            expiry_date_value=expiry_date_value,
            legacy_marking_absent=legacy_marking_absent,
            performed_by=request.user,
        )
    except FbsScanMismatchError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=str(exc),
                submitted_item_scan=item_scan,
                retry_marking=True,
                item_selected=True,
            ),
            status=400,
        )
    except FbsMarkingAlreadyUsedError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        error_text = str(exc)
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=error_text,
                submitted_item_scan=item_scan,
                retry_marking=True,
                item_selected=True,
            ),
            status=409,
        )
    except FbsError as exc:
        selected_allocation = get_object_or_404(
            _pick_allocation_queryset(), pk=selected_allocation.id
        )
        expected_item_code = str(
            selected_allocation.balance.barcode
            or selected_allocation.order_item.barcode
            or ""
        ).strip()
        error_text = str(exc)
        marking_error = any(
            marker in error_text.casefold()
            for marker in ("киз", "чз", "data matrix", "маркиров")
        )
        retry_marking = bool(
            marking_error
            and expected_item_code
            and item_scan.casefold() == expected_item_code.casefold()
        )
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                selected_allocation,
                error=error_text,
                submitted_item_scan=item_scan,
                retry_marking=retry_marking,
                item_selected=True,
            ),
            status=400,
        )
    selected_allocation.refresh_from_db()
    next_allocation = _next_verification_allocation(batch)
    multi_item_hold_prompt = None
    if request_role == "fbs_controller":
        multi_item_hold_prompt = _ozon_multi_item_hold_prompt(selected_allocation)
    if multi_item_hold_prompt is not None:
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                next_allocation or selected_allocation,
                ok_message=(
                    "Товар принят. Заказ состоит из нескольких товаров — "
                    "отложите эту единицу отдельно до завершения заказа."
                ),
                multi_item_hold_prompt=multi_item_hold_prompt,
            ),
        )
    # Finish the checked order before offering another unit from this wave.
    # Fetch renders in place, so it must preserve the GET path's label priority.
    label = _next_verification_label(batch)
    if label is not None:
        if request_role == "fbs_controller":
            return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
        return redirect("fbs:tsd_label_detail", label_id=label.id)
    if next_allocation is None:
        if _next_incomplete_verification_allocation(batch) is not None:
            return redirect("fbs:tsd_pick_verification", batch_id=batch.id)
        if request_role == "fbs_controller":
            return redirect(f"{reverse('fbs:controller_home')}?done=1")
        return redirect("fbs:operator_wave_detail", batch_id=batch.id)
    if (
        request_role == "fbs_controller"
        and request.headers.get("X-Requested-With") == "fetch"
    ):
        return render(
            request,
            "fbs/tsd_pick_verification.html",
            _pick_verification_context(
                request,
                batch,
                next_allocation,
                ok_message="Товар принят. Сканируйте следующую единицу.",
            ),
        )
    return redirect("fbs:tsd_pick_verification", batch_id=batch.id)


def _pick_verification_follow_up_request(request, location):
    """Build a GET copy for a same-screen post/redirect/get response."""
    target = urlsplit(location)
    follow_up = copy(request)
    follow_up.method = "GET"
    follow_up.path = target.path
    follow_up.path_info = target.path
    follow_up.META = request.META.copy()
    follow_up.META["REQUEST_METHOD"] = "GET"
    follow_up.META["QUERY_STRING"] = target.query
    follow_up.GET = QueryDict(target.query)
    follow_up.POST = QueryDict("")
    return follow_up


def _pick_verification_fetch_response(request, response, *, batch_id=None):
    """Return the unchanged verification result as JSON for an in-place scan."""
    if request.headers.get("X-Requested-With") != "fetch":
        return response

    location = response.get("Location")
    if location:
        target_path = urlsplit(location).path
        if (
            request.method == "POST"
            and batch_id is not None
            and target_path == request.path
        ):
            follow_up_request = _pick_verification_follow_up_request(
                request,
                location,
            )
            follow_up_response = _tsd_pick_verification_response(
                follow_up_request,
                batch_id,
            )
            return _pick_verification_fetch_response(
                follow_up_request,
                follow_up_response,
                batch_id=batch_id,
            )
        redirect_key = (
            "reload_url"
            if target_path == request.path
            else "redirect_url"
        )
        payload = {redirect_key: location}
        status = 200
    elif "text/html" in str(response.get("Content-Type") or "").casefold():
        payload = {
            "html": response.content.decode(response.charset or "utf-8", "replace"),
            "url": request.get_full_path(),
        }
        status = response.status_code
    else:
        message = response.content.decode(response.charset or "utf-8", "replace").strip()
        payload = {
            "error": message or "Сканирование отклонено сервером. Повторите попытку."
        }
        status = response.status_code

    result = JsonResponse(payload, status=status)
    result["Cache-Control"] = "no-store"
    return result


@fbs_module_required
@role_required(*CONTROLLER_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_pick_verification(request, batch_id: int):
    response = _tsd_pick_verification_response(request, batch_id)
    return _pick_verification_fetch_response(request, response, batch_id=batch_id)


@fbs_module_required
@role_required("fbs_controller")
@require_http_methods(["GET", "POST"])
def tsd_pick_verification_problem(request, allocation_id: int):
    allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation_id)
    task = allocation.pick_task
    if task is None:
        return HttpResponseForbidden("Позиция не относится к волне FBS")
    batch = task.batch
    if batch.verification_assigned_to_id != request.user.id:
        return HttpResponseForbidden("Проверку волны выполняет другой контролер")
    if batch.status != FbsPickBatch.STATUS_VERIFICATION:
        return redirect("fbs:controller_home")

    error = ""
    lookup_requested = request.GET.get("lookup") == "1"
    if request.method == "POST":
        action = str(request.POST.get("action") or "return_to_pick").strip()
        try:
            if action == "return_problem_item":
                from .services.problems import return_problem_tote_item_to_check_tote

                lookup_requested = True
                return_problem_tote_item_to_check_tote(
                    allocation_id=allocation.id,
                    problem_item_id=int(request.POST.get("problem_item_id") or 0),
                    problem_tote_scan=str(
                        request.POST.get("problem_tote_scan") or ""
                    ).strip(),
                    returned_by=request.user,
                )
            elif action == "return_to_pick":
                product_scans = [
                    str(value or "").strip()
                    for value in request.POST.getlist("product_scan")
                    if str(value or "").strip()
                ]
                initial_scan_token = str(
                    request.POST.get("initial_product_scan_token") or ""
                ).strip()
                if initial_scan_token:
                    product_scans.insert(
                        0,
                        _canceled_initial_scan_from_token(
                            token=initial_scan_token,
                            batch=batch,
                            allocation=allocation,
                            controller=request.user,
                        ),
                    )
                queue_verification_problem_restock(
                    allocation_id=allocation.id,
                    exception_type=str(
                        request.POST.get("exception_type") or ""
                    ).strip(),
                    reason=str(request.POST.get("reason") or "").strip(),
                    service_tote_scan=str(
                        request.POST.get("service_tote_scan") or ""
                    ).strip(),
                    reported_by=request.user,
                    product_scans=product_scans,
                )
            else:
                raise ValueError("Неизвестная команда проблемного товара.")
        except (FbsError, TypeError, ValueError) as exc:
            error = str(exc)
            allocation = get_object_or_404(_pick_allocation_queryset(), pk=allocation.id)
            if request.POST.get("modal_flow") == "1":
                problem_kind = str(
                    request.POST.get("problem_kind") or "canceled"
                ).strip()
                if problem_kind not in {"canceled", "label", "marking"}:
                    problem_kind = "canceled"
                prompt_reason = str(request.POST.get("reason") or "").strip()
                return render(
                    request,
                    "fbs/tsd_pick_verification.html",
                    _pick_verification_context(
                        request,
                        batch,
                        allocation,
                        error=error,
                        item_selected=True,
                        service_tote_prompt=_verification_service_tote_prompt(
                            batch=batch,
                            allocation=allocation,
                            controller=request.user,
                            problem_kind=problem_kind,
                            reason=prompt_reason,
                        ),
                    ),
                    status=409,
                )
        else:
            if action == "return_problem_item":
                return redirect(
                    f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}"
                    "?problem_returned=1"
                )
            return redirect(
                f"{reverse('fbs:tsd_pick_verification', kwargs={'batch_id': batch.id})}?problem=1"
            )

    agent_scan_poll_url = ""
    agent_scan_event_id = 0
    if batch.workstation_id and batch.workstation.device_agent_id:
        agent_scan_poll_url = reverse("fbs:controller_scan_events")
        agent_scan_event_id = int(
            AgentEvent.objects.filter(
                agent_id=batch.workstation.device_agent.agent_id,
                event_type=AgentEvent.EVENT_SCAN,
            )
            .order_by("-id")
            .values_list("id", flat=True)
            .first()
            or 0
        )
    pick_tote_context = (
        FbsControllerPickTote.objects.select_related(
            "session__problem_tote",
            "session__canceled_tote",
        )
        .filter(
            pick_batch=batch,
            session__controller=request.user,
            session__status="active",
            status__in=(
                FbsControllerPickTote.STATUS_PROCESSING,
                FbsControllerPickTote.STATUS_AWAITING_EMPTY,
            ),
        )
        .order_by("-id")
        .first()
    )
    session = pick_tote_context.session if pick_tote_context else None
    problem_tote = session.problem_tote if session else None
    order_is_canceled = _marketplace_order_is_canceled(task.order)
    canceled_label_applied = bool(
        order_is_canceled
        and FbsOrderLabel.objects.filter(
            order_id=task.order_id,
            status=FbsOrderLabel.STATUS_APPLIED,
        ).exists()
    )
    service_tote = (
        session.canceled_tote if order_is_canceled and session else None
    ) or (problem_tote if not order_is_canceled else None)
    problem_tote_lookup = None
    if lookup_requested:
        try:
            from .services.problems import lookup_problem_tote_item

            problem_tote_lookup = lookup_problem_tote_item(
                allocation_id=allocation.id,
                actor=request.user,
            )
        except FbsError as exc:
            if not error:
                error = str(exc)
    return render(
        request,
        "fbs/tsd_pick_verification_problem.html",
        _base_context(
            request,
            page_title=f"FBS · Проблема заказа {task.order.external_order_id}",
            back_url=reverse("fbs:tsd_pick_verification", kwargs={"batch_id": batch.id}),
            batch=batch,
            task=task,
            order=task.order,
            allocation=allocation,
            balance=allocation.balance,
            problem_tote=problem_tote,
            service_tote=service_tote,
            service_tote_kind=("canceled" if order_is_canceled else "problem"),
            canceled_label_applied=canceled_label_applied,
            service_tote_label=(
                "тара отмененных заказов"
                if order_is_canceled
                else "проблемная тара"
            ),
            canceled_product_scan_rows=(
                _canceled_product_scan_rows(task=task)
                if order_is_canceled and not canceled_label_applied
                else []
            ),
            canceled_product_total=(
                sum(
                    int(row.qty_picked or 0)
                    for row in FbsOrderStockAllocation.objects.filter(
                        pick_task=task,
                        status=FbsOrderStockAllocation.STATUS_PICKED,
                        qty_picked__gt=0,
                    )
                )
                if order_is_canceled
                else 0
            ),
            problem_tote_lookup=problem_tote_lookup,
            exception_choices=FbsPickException.TYPE_CHOICES,
            agent_scan_poll_url=agent_scan_poll_url,
            agent_scan_event_id=agent_scan_event_id,
            error=error,
        ),
        status=409 if error else 200,
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_GET
def tsd_storekeeper_plan(request, plan_id: int):
    plan = get_object_or_404(_plan_queryset(), pk=plan_id)
    ok_message = ""
    if request.GET.get("confirmed") == "1":
        ok_message = "План подтвержден и передан ричтраку."
    elif request.GET.get("canceled") == "1":
        ok_message = "План отменен, резерв общего склада снят."
    elif request.GET.get("packed") == "1":
        ok_message = "QR FBS-короба создан. Распечатайте его и передайте короб водителю."
    return render(
        request,
        "fbs/tsd_storekeeper_plan.html",
        _storekeeper_plan_context(request, plan, ok_message=ok_message),
    )


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_confirm_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        confirm_replenishment_plan(plan_id=plan_id, confirmed_by=request.user)
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?confirmed=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_cancel_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        cancel_replenishment_plan(plan_id=plan_id, canceled_by=request.user)
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?canceled=1")


@fbs_module_required
@role_required(*STOREKEEPER_ROLES)
@require_POST
def tsd_pack_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        pack_staged_item_plan(
            plan_id=plan_id,
            packed_by=request.user,
            staging_container_scan=request.POST.get(
                "staging_container_scan", ""
            ),
        )
    except FbsError as exc:
        plan = get_object_or_404(_plan_queryset(), pk=plan_id)
        return render(
            request,
            "fbs/tsd_storekeeper_plan.html",
            _storekeeper_plan_context(request, plan, error=str(exc)),
            status=409,
        )
    url = reverse("fbs:tsd_storekeeper_plan", kwargs={"plan_id": plan_id})
    return redirect(f"{url}?packed=1")


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_GET
def tsd_reachtruck(request):
    return redirect("/reachtruck/")


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_POST
def tsd_claim_plan(request, plan_id: int):
    get_object_or_404(FbsReplenishmentPlan, pk=plan_id)
    try:
        plan = claim_replenishment_plan(plan_id=plan_id, assigned_to=request.user)
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_reachtruck_list.html",
            _reachtruck_list_context(request, error=str(exc)),
            status=409,
        )
    allocation = (
        _allocation_queryset()
        .filter(line__plan=plan, status__in=OPEN_ALLOCATION_STATUSES)
        .order_by("id")
        .first()
    )
    if allocation is None:
        return redirect("fbs:tsd_reachtruck")
    return redirect("fbs:tsd_allocation", allocation_id=allocation.id)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_POST
def tsd_claim_internal_movement(request, movement_id: int):
    get_object_or_404(FbsInternalMovement, pk=movement_id)
    try:
        claim_internal_movement(movement_id=movement_id, assigned_to=request.user)
    except FbsError as exc:
        return render(
            request,
            "fbs/tsd_reachtruck_list.html",
            _reachtruck_list_context(request, error=str(exc)),
            status=409,
        )
    return redirect("fbs:tsd_internal_movement", movement_id=movement_id)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_internal_movement(request, movement_id: int):
    movement = get_object_or_404(
        FbsInternalMovement.objects.select_related(
            "agency",
            "source_box__pallet__cell",
            "target_pallet__cell",
            "target_box",
            "source_balance",
            "assigned_to",
        ),
        pk=movement_id,
    )
    if movement.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Перемещение назначено другому сотруднику")
    if movement.status == FbsInternalMovement.STATUS_DONE:
        return redirect("fbs:tsd_reachtruck")
    context = _base_context(
        request,
        page_title=f"FBS · Перемещение #{movement.id}",
        back_url=reverse("fbs:tsd_reachtruck"),
        movement=movement,
    )
    if request.method == "GET":
        return render(request, "fbs/tsd_internal_movement.html", context)
    try:
        movement = scan_internal_movement(
            movement_id=movement.id,
            source_box_scan=request.POST.get("source_box_scan", ""),
            target_scan=request.POST.get("target_scan", ""),
            item_scan=request.POST.get("item_scan", ""),
            performed_by=request.user,
        )
    except FbsError as exc:
        context["error"] = str(exc)
        return render(request, "fbs/tsd_internal_movement.html", context, status=400)
    if movement.status == FbsInternalMovement.STATUS_DONE:
        return redirect(f"{reverse('fbs:tsd_reachtruck')}?movement_done=1")
    context["movement"] = movement
    context["ok_message"] = f"Перемещено {movement.moved_qty} из {movement.planned_qty}."
    return render(request, "fbs/tsd_internal_movement.html", context)


@fbs_module_required
@role_required(*REACHTRUCK_ROLES)
@require_http_methods(["GET", "POST"])
def tsd_allocation(request, allocation_id: int):
    allocation = get_object_or_404(_allocation_queryset(), pk=allocation_id)
    plan = allocation.line.plan
    if plan.assigned_to_id != request.user.id:
        return HttpResponseForbidden("Задание назначено другому сотруднику")
    if allocation.status in {
        FbsReplenishmentAllocation.STATUS_STAGED,
        FbsReplenishmentAllocation.STATUS_DONE,
    }:
        return redirect("fbs:tsd_reachtruck")

    if request.method == "GET":
        return render(
            request,
            "fbs/tsd_allocation.html",
            _allocation_context(request, allocation),
        )

    action = str(request.POST.get("action") or "").strip()
    if action == "scan_source":
        source_scan = str(request.POST.get("source_scan") or "").strip()
        try:
            validate_replenishment_source_scan(
                allocation_id=allocation.id,
                source_scan=source_scan,
                performed_by=request.user,
            )
        except FbsError as exc:
            request.session.pop(_source_session_key(allocation.id), None)
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session[_source_session_key(allocation.id)] = source_scan
        return render(
            request,
            "fbs/tsd_allocation.html",
            _allocation_context(
                request,
                allocation,
                ok_message="Источник подтвержден.",
            ),
        )

    if action == "complete":
        source_scan = str(request.session.get(_source_session_key(allocation.id)) or "").strip()
        target_scan = str(
            request.POST.get("staging_scan")
            or request.POST.get("target_box_scan")
            or ""
        ).strip()
        if not source_scan:
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error="Сначала подтвердите источник."),
                status=400,
            )
        try:
            is_staging_transfer = bool(
                plan.client_movement_request_id
                and plan.mode == FbsReplenishmentPlan.MODE_ITEM
                and plan.target_box_id is None
                and plan.staging_location_id
            )
            if is_staging_transfer:
                stage_replenishment_allocation(
                    allocation_id=allocation.id,
                    source_scan=source_scan,
                    staging_container_scan=request.POST.get(
                        "staging_container_scan", ""
                    ),
                    staging_scan=target_scan,
                    performed_by=request.user,
                )
            else:
                complete_replenishment_allocation(
                    allocation_id=allocation.id,
                    source_scan=source_scan,
                    target_box_scan=target_scan,
                    performed_by=request.user,
                )
        except FbsError as exc:
            return render(
                request,
                "fbs/tsd_allocation.html",
                _allocation_context(request, allocation, error=str(exc)),
                status=400,
            )
        request.session.pop(_source_session_key(allocation.id), None)
        next_allocation = (
            _allocation_queryset()
            .filter(line__plan=plan, status__in=OPEN_ALLOCATION_STATUSES)
            .order_by("id")
            .first()
        )
        if next_allocation is not None:
            return redirect("fbs:tsd_allocation", allocation_id=next_allocation.id)
        result_flag = "staged" if is_staging_transfer else "done"
        return redirect(f"{reverse('fbs:tsd_reachtruck')}?{result_flag}=1")

    return render(
        request,
        "fbs/tsd_allocation.html",
        _allocation_context(request, allocation, error="Неизвестная команда."),
        status=400,
    )
