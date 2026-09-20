from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from fbs.tsd_views import _xlsx_safe_value


class StockExportXlsxValueTests(SimpleTestCase):
    def test_illegal_xml_control_characters_are_escaped(self):
        value = "0104681026565474\x1d91EE11\x0092TEST"

        self.assertEqual(
            _xlsx_safe_value(value),
            "0104681026565474\\x1D91EE11\\x0092TEST",
        )

    def test_escaped_marking_code_can_be_written_and_read(self):
        output = BytesIO()
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(title="Остатки FBS")
        sheet.append([_xlsx_safe_value("code\x1dpart")])
        workbook.save(output)

        output.seek(0)
        restored = load_workbook(output, read_only=True)
        try:
            self.assertEqual(
                restored["Остатки FBS"].cell(row=1, column=1).value,
                "code\\x1Dpart",
            )
        finally:
            restored.close()
