"""Автотесты шаблона и проверки Excel (без склада)."""
from __future__ import annotations

from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import Workbook

from receiving_distribution.services.excel import (
    TEMPLATE_HEADERS,
    assert_template_healthy,
    build_template_workbook,
    ensure_template_file,
    parse_distribution_excel,
)
from receiving_distribution.services.validate import (
    DistributionDraft,
    has_blocking_errors,
    validate_distribution_draft,
)


def _xlsx(rows, sheet_title="Распределение"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "test.xlsx"
    return buf


HEADERS = list(TEMPLATE_HEADERS)


class TemplateHealthTests(SimpleTestCase):
    def test_built_template_is_healthy(self):
        path = ensure_template_file(force_rebuild=True)
        ok, reason = assert_template_healthy(path)
        self.assertTrue(ok, reason)

    def test_workbook_has_instruction_sheet(self):
        wb = build_template_workbook()
        self.assertIn("Инструкция", wb.sheetnames)
        self.assertIn("Распределение", wb.sheetnames)


class ExcelParseTests(SimpleTestCase):
    def test_correct_file(self):
        f = _xlsx(
            [
                HEADERS,
                ["A", "1", "T", 100, "WB", "K", "S1", 60, "", "", ""],
                ["A", "1", "T", 100, "Хранение FullBox", "", "", 40, "", "", ""],
            ]
        )
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])
        self.assertEqual(preview.stats["shipping_drafts_count"], 1)
        self.assertEqual(preview.stats["sku_count"], 1)

    def test_empty_file(self):
        f = _xlsx([])
        preview = parse_distribution_excel(f)
        self.assertFalse(preview.ok)

    def test_missing_required_column(self):
        f = _xlsx([["Артикул", "Количество по направлению"], ["A", 10]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertFalse(preview.ok)
        self.assertTrue(any(i.code == "bad_headers" for i in preview.issues))

    def test_header_extra_space_still_maps(self):
        headers = list(HEADERS)
        headers[0] = " Артикул "
        f = _xlsx([headers, ["A", "1", "T", 10, "WB", "K", "S1", 10, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])

    def test_columns_reordered(self):
        headers = [
            "Количество по направлению",
            "Артикул",
            "Общее количество",
            "Маркетплейс",
            "Склад назначения",
            "Номер поставки",
            "Штрихкод товара",
            "Наименование",
            "Дата поставки",
            "Таймслот",
            "Комментарий",
        ]
        f = _xlsx([headers, [10, "A", 10, "WB", "K", "S1", "1", "T", "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])

    def test_duplicate_headers(self):
        headers = list(HEADERS) + ["Артикул"]
        f = _xlsx([headers, ["A", "1", "T", 10, "WB", "K", "S1", 10, "", "", "", "A"]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(any(i.code == "dup_header" for i in preview.issues))

    def test_bad_extension(self):
        buf = BytesIO(b"not-excel")
        buf.name = "x.csv"
        buf.size = 9
        preview = parse_distribution_excel(buf)
        self.assertTrue(any(i.code == "bad_extension" for i in preview.issues))

    def test_broken_excel(self):
        buf = BytesIO(b"PK\x03\x04broken")
        buf.name = "x.xlsx"
        preview = parse_distribution_excel(buf)
        self.assertTrue(any(i.code == "bad_file" for i in preview.issues))

    def test_blank_rows_between_data(self):
        f = _xlsx(
            [
                HEADERS,
                ["A", "1", "T", 10, "WB", "K", "S1", 10, "", "", ""],
                [None] * 11,
                ["B", "2", "T2", 5, "Ozon", "X", "S2", 5, "", "", ""],
            ]
        )
        preview = parse_distribution_excel(f, known_sku_codes={"a", "b"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])
        self.assertEqual(preview.stats["sku_count"], 2)

    def test_formula_result_as_number(self):
        # data_only=True: без Excel engine формула может быть None — проверяем целое float
        f = _xlsx([HEADERS, ["A", "1", "T", 10.0, "WB", "K", "S1", 10.0, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])


class DistributionDataTests(SimpleTestCase):
    def test_sum_less(self):
        draft = DistributionDraft(
            items=[{"sku_code": "SKU-123", "qty_total": 1000}],
            directions=[{"key": "wb", "kind": "marketplace", "marketplace_name": "WB", "destination_warehouse": "K"}],
            allocations=[{"sku_code": "SKU-123", "direction_key": "wb", "qty": 900}],
        )
        issues = validate_distribution_draft(draft)
        self.assertTrue(any("не распределено 100" in i.message for i in issues))

    def test_sum_more(self):
        draft = DistributionDraft(
            items=[{"sku_code": "SKU-123", "qty_total": 100}],
            directions=[{"key": "wb", "kind": "marketplace", "marketplace_name": "WB", "destination_warehouse": "K"}],
            allocations=[{"sku_code": "SKU-123", "direction_key": "wb", "qty": 150}],
        )
        issues = validate_distribution_draft(draft)
        self.assertTrue(any("на 50 единиц больше" in i.message for i in issues))

    def test_negative_qty(self):
        f = _xlsx([HEADERS, ["A", "1", "T", 10, "WB", "K", "S1", -5, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertFalse(preview.ok)

    def test_fraction_qty(self):
        f = _xlsx([HEADERS, ["A", "1", "T", 10, "WB", "K", "S1", 1.5, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertFalse(preview.ok)

    def test_unknown_sku(self):
        f = _xlsx([HEADERS, ["UNKNOWN", "1", "T", 10, "WB", "K", "S1", 10, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(any(i.code == "unknown_sku" for i in preview.issues))

    def test_storage_fullbox(self):
        f = _xlsx([HEADERS, ["A", "1", "T", 10, "Хранение FullBox", "", "", 10, "", "", ""]])
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])
        self.assertEqual(preview.stats["shipping_drafts_count"], 0)

    def test_three_directions_one_sku(self):
        f = _xlsx(
            [
                HEADERS,
                ["A", "1", "T", 1000, "WB", "K1", "S1", 600, "", "", ""],
                ["A", "1", "T", 1000, "Ozon", "K2", "S2", 300, "", "", ""],
                ["A", "1", "T", 1000, "Хранение FullBox", "", "", 100, "", "", ""],
            ]
        )
        preview = parse_distribution_excel(f, known_sku_codes={"a"})
        self.assertTrue(preview.ok, [i.message for i in preview.issues])
        self.assertEqual(preview.stats["directions_count"], 3)
        self.assertEqual(preview.stats["shipping_drafts_count"], 2)

    def test_warnings_do_not_block_has_blocking(self):
        issues = validate_distribution_draft(
            DistributionDraft(
                items=[{"sku_code": "A", "qty_total": 10}],
                directions=[
                    {
                        "key": "wb",
                        "kind": "marketplace",
                        "marketplace_name": "WB",
                        "destination_warehouse": "K",
                        # нет supply → warning
                    }
                ],
                allocations=[{"sku_code": "A", "direction_key": "wb", "qty": 10}],
            )
        )
        # missing_supply is warning — create may proceed if only warnings
        self.assertTrue(any(i.code == "missing_supply" for i in issues))
        # has_blocking_errors ignores warnings
        self.assertFalse(has_blocking_errors([i for i in issues if i.severity == "warning"]))
