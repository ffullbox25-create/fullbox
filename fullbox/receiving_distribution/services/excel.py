"""Один Excel-шаблон для приёмки с распределением."""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from .validate import (
    DistributionDraft,
    ValidationIssue,
    has_blocking_errors,
    validate_distribution_draft,
)

logger = logging.getLogger(__name__)

TEMPLATE_HEADERS = (
    "Артикул",
    "Штрихкод товара",
    "Наименование",
    "Общее количество",
    "Маркетплейс",
    "Склад назначения",
    "Номер поставки",
    "Количество по направлению",
    "Дата поставки",
    "Таймслот",
    "Комментарий",
)

REQUIRED_LOGICAL = ("sku", "qty_dir", "qty_total", "marketplace")

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "data" / "receiving_distribution_template.xlsx"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024

STORAGE_ALIASES = {
    "хранение",
    "хранение fullbox",
    "хранение full box",
    "хранение fullbox",
    "fullbox",
    "storage",
    "storage_fullbox",
}

HEADER_ALIASES = {
    "артикул": "sku",
    "sku": "sku",
    "штрихкод товара": "barcode",
    "штрихкод": "barcode",
    "баркод": "barcode",
    "наименование": "name",
    "общее количество": "qty_total",
    "кол-во": "qty_total",
    "количество": "qty_total",
    "маркетплейс": "marketplace",
    "тип направления": "direction_type",
    "склад назначения": "warehouse",
    "склад": "warehouse",
    "номер поставки": "supply",
    "количество по направлению": "qty_dir",
    "дата поставки": "slot_date",
    "дата": "slot_date",
    "таймслот": "slot_time",
    "комментарий": "comment",
}


@dataclass
class ExcelPreview:
    draft: DistributionDraft
    issues: list[ValidationIssue] = field(default_factory=list)
    warnings: list[ValidationIssue] = field(default_factory=list)
    row_count: int = 0
    preview_rows: list[dict[str, Any]] = field(default_factory=list)
    stats: dict[str, Any] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not has_blocking_errors(self.issues)

    @property
    def has_warnings(self) -> bool:
        return bool(self.warnings) or any(i.severity == "warning" for i in self.issues)

    def to_payload(self) -> dict[str, Any]:
        errors = [i for i in self.issues if i.severity != "warning"]
        warnings = [i for i in self.issues if i.severity == "warning"] + list(self.warnings)
        return {
            "ok": self.ok,
            "row_count": self.row_count,
            "stats": self.stats,
            "issues": [i.to_dict() for i in errors],
            "warnings": [i.to_dict() for i in warnings],
            "preview_rows": self.preview_rows,
            "draft": {
                "items": self.draft.items,
                "directions": self.draft.directions,
                "allocations": self.draft.allocations,
                "meta": self.draft.meta,
            }
            if self.ok
            else None,
        }


def _norm(value: Any) -> str:
    return str(value or "").strip()


def _norm_key(value: Any) -> str:
    return _norm(value).casefold()


def _cell_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    if "." in text:
        num = float(text)
        if not num.is_integer():
            raise ValueError("fraction")
        return int(num)
    return int(text)


def _header_map(header_row: tuple[Any, ...]) -> tuple[dict[str, int], list[ValidationIssue]]:
    issues: list[ValidationIssue] = []
    mapping: dict[str, int] = {}
    seen_labels: dict[str, int] = {}
    for idx, raw in enumerate(header_row):
        label = _norm(raw)
        if not label:
            continue
        key = HEADER_ALIASES.get(_norm_key(label))
        if label.casefold() in seen_labels:
            issues.append(
                ValidationIssue(
                    "dup_header",
                    f"Дублирующаяся колонка «{label}».",
                    column=label,
                    recommendation="Удалите дублирующие заголовки.",
                )
            )
            continue
        seen_labels[label.casefold()] = idx
        if key and key not in mapping:
            mapping[key] = idx
    return mapping, issues


def assert_template_healthy(path: Path | None = None) -> tuple[bool, str]:
    """Проверка системного шаблона перед выдачей клиенту."""
    target = path or TEMPLATE_PATH
    if not target.is_file():
        return False, "Файл шаблона отсутствует."
    try:
        wb = load_workbook(target, read_only=True, data_only=True)
    except Exception as exc:
        logger.exception("receiving_distribution template open failed: %s", exc)
        return False, "Файл шаблона повреждён."
    try:
        if "Распределение" in wb.sheetnames:
            ws = wb["Распределение"]
        else:
            ws = wb.active
        rows = list(ws.iter_rows(values_only=True, max_row=1))
        if not rows:
            return False, "В шаблоне нет заголовков."
        mapping, header_issues = _header_map(tuple(rows[0]))
        if header_issues:
            return False, header_issues[0].message
        for logical in ("sku", "qty_dir", "qty_total", "marketplace"):
            if logical not in mapping:
                return False, f"В шаблоне нет обязательной колонки ({logical})."
        expected_labels = {_norm_key(h) for h in TEMPLATE_HEADERS}
        actual = {_norm_key(c) for c in rows[0] if _norm(c)}
        if not expected_labels.issubset(actual):
            missing = expected_labels - actual
            return False, f"В шаблоне не хватает колонок: {', '.join(sorted(missing))}."
    finally:
        wb.close()
    return True, ""


def _write_instruction_sheet(wb: Workbook) -> None:
    if "Инструкция" in wb.sheetnames:
        del wb["Инструкция"]
    ws = wb.create_sheet("Инструкция", 0)
    rows = [
        ["Инструкция по заполнению шаблона «Приёмка с распределением»"],
        [],
        ["Не изменяйте названия и порядок обязательных колонок на листе «Распределение»."],
        [],
        ["Колонка", "Обязательность", "Формат", "Пример", "Пояснение"],
        ["Артикул", "Да", "Текст", "ART-001", "Артикул из номенклатуры клиента"],
        ["Штрихкод товара", "Нет", "Текст/число", "2000000000001", "Штрихкод SKU клиента"],
        ["Наименование", "Нет", "Текст", "Товар пример", "Для удобства"],
        ["Общее количество", "Да", "Целое > 0", "1000", "Сколько единиц придёт всего по артикулу"],
        ["Маркетплейс", "Да*", "Текст", "WB / Ozon / Яндекс Маркет / Хранение FullBox", "* Для хранения — «Хранение FullBox»"],
        ["Склад назначения", "Для МП", "Текст", "Коледино", "Не обязателен для хранения"],
        ["Номер поставки", "Для МП", "Текст", "SUP-1", "Не обязателен для хранения"],
        ["Количество по направлению", "Да", "Целое > 0", "600", "Часть общего количества на это направление"],
        ["Дата поставки", "Нет", "Дата", "2026-08-01", "Плановая дата отгрузки направления"],
        ["Таймслот", "Нет", "Текст", "10:00-12:00", "Окно поставки"],
        ["Комментарий", "Нет", "Текст ≤500", "", "Комментарий к строке"],
        [],
        ["Правило распределения:"],
        ["Сумма «Количество по направлению» по артикулу = «Общее количество»."],
        ["Пример: 1000 = 600 (WB) + 250 (WB другой склад) + 150 (Хранение FullBox)."],
        [],
        ["Допустимые маркетплейсы:"],
        ["Wildberries / WB", "Ozon", "Яндекс Маркет", "Хранение FullBox"],
        [],
        ["Один артикул можно разнести на несколько строк — по одной на каждое направление."],
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 48


def build_template_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Распределение"
    ws.append(list(TEMPLATE_HEADERS))
    ws.append(["ART-001", "2000000000001", "Товар пример", 1000, "WB", "Коледино", "SUP-1", 600, "", "", ""])
    ws.append(["ART-001", "2000000000001", "Товар пример", 1000, "WB", "Электросталь", "SUP-2", 250, "", "", ""])
    ws.append(["ART-001", "2000000000001", "Товар пример", 1000, "Хранение FullBox", "", "", 150, "", "", ""])
    _write_instruction_sheet(wb)
    # Инструкция первой, Распределение — рабочей
    wb.move_sheet("Распределение", offset=len(wb.sheetnames))
    return wb


def ensure_template_file(*, force_rebuild: bool = False) -> Path:
    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    if TEMPLATE_PATH.is_file() and not force_rebuild:
        ok, _ = assert_template_healthy(TEMPLATE_PATH)
        if ok:
            return TEMPLATE_PATH
    wb = build_template_workbook()
    wb.save(TEMPLATE_PATH)
    return TEMPLATE_PATH


def build_template_response() -> HttpResponse:
    ensure_template_file()
    ok, reason = assert_template_healthy()
    if not ok:
        logger.error("receiving_distribution template unhealthy: %s", reason)
        raise ValueError(reason)
    data = TEMPLATE_PATH.read_bytes()
    response = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="receiving-distribution.xlsx"'
    return response


def _is_storage(marketplace: str, warehouse: str, direction_type: str = "") -> bool:
    return (
        _norm_key(marketplace) in STORAGE_ALIASES
        or _norm_key(warehouse) in STORAGE_ALIASES
        or _norm_key(direction_type) in STORAGE_ALIASES
    )


def _build_stats(draft: DistributionDraft, row_count: int) -> dict[str, Any]:
    units = sum(int(i.get("qty_total") or 0) for i in draft.items)
    shipping_dirs = [
        d
        for d in draft.directions
        if _norm(d.get("kind")) != "storage_fullbox"
    ]
    return {
        "row_count": row_count,
        "sku_count": len(draft.items),
        "units_total": units,
        "directions_count": len(draft.directions),
        "shipping_drafts_count": len(shipping_dirs),
        "storage_directions": len(draft.directions) - len(shipping_dirs),
    }


def parse_distribution_excel(
    file_obj,
    *,
    known_sku_codes: set[str] | None = None,
    known_barcodes: set[str] | None = None,
    max_bytes: int = MAX_UPLOAD_BYTES,
) -> ExcelPreview:
    issues: list[ValidationIssue] = []

    size = getattr(file_obj, "size", None)
    if size is not None and int(size) > max_bytes:
        return ExcelPreview(
            draft=DistributionDraft(),
            issues=[
                ValidationIssue(
                    "file_too_large",
                    f"Файл больше {max_bytes // (1024 * 1024)} МБ.",
                    recommendation="Разбейте файл или уменьшите объём.",
                )
            ],
        )

    name = _norm(getattr(file_obj, "name", "")).casefold()
    if name and not (name.endswith(".xlsx") or name.endswith(".xls")):
        return ExcelPreview(
            draft=DistributionDraft(),
            issues=[
                ValidationIssue(
                    "bad_extension",
                    "Допустимы только файлы .xlsx или .xls.",
                    recommendation="Скачайте шаблон и заполните его.",
                )
            ],
        )

    try:
        if hasattr(file_obj, "seek"):
            try:
                file_obj.seek(0)
            except Exception:
                pass
        wb = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception:
        return ExcelPreview(
            draft=DistributionDraft(),
            issues=[
                ValidationIssue(
                    "bad_file",
                    "Не удалось прочитать Excel. Используйте шаблон.",
                    recommendation="Скачайте актуальный шаблон и заполните заново.",
                )
            ],
        )

    try:
        ws = wb["Распределение"] if "Распределение" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return ExcelPreview(
            draft=DistributionDraft(),
            issues=[ValidationIssue("empty", "Файл пустой.", recommendation="Заполните лист «Распределение».")],
        )

    header_map, header_issues = _header_map(tuple(rows[0]))
    issues.extend(header_issues)
    for logical, title in (
        ("sku", "Артикул"),
        ("qty_dir", "Количество по направлению"),
        ("qty_total", "Общее количество"),
        ("marketplace", "Маркетплейс"),
    ):
        if logical not in header_map:
            issues.append(
                ValidationIssue(
                    "bad_headers",
                    f"Нужна колонка «{title}».",
                    column=title,
                    recommendation="Не меняйте названия колонок шаблона.",
                )
            )
    if has_blocking_errors(issues) and ("sku" not in header_map or "qty_dir" not in header_map):
        return ExcelPreview(draft=DistributionDraft(), issues=issues)

    items_map: dict[str, dict[str, Any]] = {}
    directions_map: dict[str, dict[str, Any]] = {}
    allocations: list[dict[str, Any]] = []
    preview_rows: list[dict[str, Any]] = []
    row_count = 0
    seen_row_keys: dict[tuple, int] = {}

    for row_no, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def get(logical: str) -> Any:
            idx = header_map.get(logical)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        sku = _norm(get("sku"))
        barcode = _norm(get("barcode"))
        if not sku and not barcode:
            issues.append(
                ValidationIssue(
                    "no_sku_or_barcode",
                    "Укажите артикул или штрихкод.",
                    row_no=row_no,
                    column="Артикул",
                    recommendation="Заполните артикул из номенклатуры.",
                )
            )
            continue
        if not sku:
            issues.append(
                ValidationIssue(
                    "empty_sku",
                    "Укажите артикул.",
                    row_no=row_no,
                    column="Артикул",
                    value=barcode,
                    recommendation="Артикул обязателен для создания заявки.",
                )
            )
            continue

        row_count += 1
        try:
            qty_total = _cell_int(get("qty_total")) if "qty_total" in header_map else None
            qty_dir = _cell_int(get("qty_dir"))
        except ValueError:
            issues.append(
                ValidationIssue(
                    "bad_qty",
                    "Количество должно быть целым.",
                    row_no=row_no,
                    sku_code=sku,
                    column="Количество по направлению",
                    value=str(get("qty_dir") or get("qty_total") or ""),
                    recommendation="Укажите положительное целое число без дробей.",
                )
            )
            continue

        if qty_dir is None:
            issues.append(
                ValidationIssue(
                    "bad_qty_dir",
                    "Укажите количество по направлению.",
                    row_no=row_no,
                    sku_code=sku,
                    column="Количество по направлению",
                    recommendation="Укажите положительное целое число.",
                )
            )
            continue
        if qty_dir <= 0:
            issues.append(
                ValidationIssue(
                    "bad_qty_dir",
                    f"Указано значение «{qty_dir}».",
                    row_no=row_no,
                    sku_code=sku,
                    column="Количество по направлению",
                    value=str(qty_dir),
                    recommendation="Укажите положительное целое число.",
                )
            )
            continue

        marketplace = _norm(get("marketplace"))
        warehouse = _norm(get("warehouse"))
        supply = _norm(get("supply"))
        direction_type = _norm(get("direction_type"))
        name_val = _norm(get("name"))
        comment = _norm(get("comment"))
        if len(comment) > 500:
            issues.append(
                ValidationIssue(
                    "comment_long",
                    "Комментарий длиннее 500 символов.",
                    row_no=row_no,
                    sku_code=sku,
                    column="Комментарий",
                    recommendation="Сократите комментарий до 500 символов.",
                )
            )
            comment = comment[:500]

        if known_barcodes is not None and barcode and barcode.casefold() not in known_barcodes:
            issues.append(
                ValidationIssue(
                    "unknown_barcode",
                    f"Штрихкод {barcode} не найден у клиента.",
                    severity="warning",
                    row_no=row_no,
                    sku_code=sku,
                    column="Штрихкод товара",
                    value=barcode,
                    recommendation="Проверьте штрихкод в номенклатуре.",
                )
            )

        is_storage = _is_storage(marketplace, warehouse, direction_type)
        if is_storage:
            dkey = "storage_fullbox"
            directions_map.setdefault(
                dkey,
                {
                    "key": dkey,
                    "kind": "storage_fullbox",
                    "title": "Хранение FullBox",
                    "marketplace_name": "",
                    "destination_warehouse": "",
                    "supply_number": "",
                },
            )
            direction_label = "Хранение FullBox"
        else:
            if not marketplace:
                issues.append(
                    ValidationIssue(
                        "missing_marketplace",
                        "Укажите маркетплейс или «Хранение FullBox».",
                        row_no=row_no,
                        sku_code=sku,
                        column="Маркетплейс",
                        recommendation="Для хранения укажите «Хранение FullBox».",
                    )
                )
            dkey = f"{marketplace.casefold()}|{warehouse.casefold()}|{supply.casefold()}"
            directions_map.setdefault(
                dkey,
                {
                    "key": dkey,
                    "kind": "marketplace",
                    "title": " — ".join(p for p in (marketplace, warehouse) if p) or dkey,
                    "marketplace_name": marketplace,
                    "destination_warehouse": warehouse,
                    "supply_number": supply,
                    "slot_date": _norm(get("slot_date")),
                    "slot_time": _norm(get("slot_time")),
                },
            )
            direction_label = directions_map[dkey]["title"]

        dup_key = (sku.casefold(), barcode.casefold(), marketplace.casefold(), warehouse.casefold(), supply.casefold())
        if dup_key in seen_row_keys:
            issues.append(
                ValidationIssue(
                    "dup_row",
                    f"Повторяющаяся строка (как строка {seen_row_keys[dup_key]}).",
                    severity="warning",
                    row_no=row_no,
                    sku_code=sku,
                    recommendation="Объедините количества или удалите дубль.",
                )
            )
        else:
            seen_row_keys[dup_key] = row_no

        key = sku.casefold()
        if key not in items_map:
            items_map[key] = {
                "sku_code": sku,
                "barcode": barcode,
                "name": name_val,
                "qty_total": int(qty_total or 0),
                "comment": comment,
            }
        else:
            existing_total = int(items_map[key].get("qty_total") or 0)
            if qty_total and existing_total and int(qty_total) != existing_total:
                issues.append(
                    ValidationIssue(
                        "qty_total_conflict",
                        f"Разное «Общее количество» для {sku}.",
                        row_no=row_no,
                        sku_code=sku,
                        column="Общее количество",
                        value=str(qty_total),
                        recommendation="Для одного артикула общее количество должно совпадать во всех строках.",
                    )
                )
            if qty_total and not existing_total:
                items_map[key]["qty_total"] = int(qty_total)
            if barcode and not items_map[key].get("barcode"):
                items_map[key]["barcode"] = barcode
            if name_val and not items_map[key].get("name"):
                items_map[key]["name"] = name_val

        allocations.append(
            {
                "sku_code": sku,
                "direction_key": dkey,
                "qty": int(qty_dir),
                "row_no": row_no,
                "barcode": barcode,
                "name": name_val or items_map[key].get("name") or "",
                "marketplace_name": "" if is_storage else marketplace,
                "destination_warehouse": "" if is_storage else warehouse,
                "supply_number": "" if is_storage else supply,
                "slot_date": _norm(get("slot_date")),
                "slot_time": _norm(get("slot_time")),
                "direction_label": direction_label,
            }
        )
        preview_rows.append(
            {
                "row_no": row_no,
                "sku_code": sku,
                "barcode": barcode,
                "name": name_val or items_map[key].get("name") or "",
                "qty_total": int(items_map[key].get("qty_total") or qty_total or 0),
                "direction_key": dkey,
                "direction_label": direction_label,
                "marketplace_name": "" if is_storage else marketplace,
                "destination_warehouse": "" if is_storage else warehouse,
                "supply_number": "" if is_storage else supply,
                "qty": int(qty_dir),
                "slot_date": _norm(get("slot_date")),
                "slot_time": _norm(get("slot_time")),
                "status": "ok",
            }
        )

    for key, item in items_map.items():
        if int(item.get("qty_total") or 0) <= 0:
            item["qty_total"] = sum(
                int(a["qty"]) for a in allocations if _norm(a["sku_code"]).casefold() == key
            )

    draft = DistributionDraft(
        items=list(items_map.values()),
        directions=list(directions_map.values()),
        allocations=allocations,
    )
    issues.extend(
        validate_distribution_draft(draft, allow_incomplete=False, known_sku_codes=known_sku_codes)
    )

    # дубли артикула по одному направлению + суммы по SKU
    pair_counts: dict[tuple[str, str], list[int]] = {}
    sku_dir_sums: dict[str, int] = {}
    for prow in preview_rows:
        sku_key = str(prow.get("sku_code") or "").casefold()
        dkey = str(prow.get("direction_key") or "").casefold()
        pair_counts.setdefault((sku_key, dkey), []).append(int(prow.get("row_no") or 0))
        sku_dir_sums[sku_key] = sku_dir_sums.get(sku_key, 0) + int(prow.get("qty") or 0)
    item_totals = {
        str(it.get("sku_code") or "").casefold(): int(it.get("qty_total") or 0) for it in draft.items
    }

    # пометить строки preview с ошибками по sku
    error_skus = {i.sku_code.casefold() for i in issues if i.sku_code and i.severity != "warning"}
    error_rows = {i.row_no for i in issues if i.row_no and i.severity != "warning"}
    for prow in preview_rows:
        sku_key = str(prow.get("sku_code") or "").casefold()
        dkey = str(prow.get("direction_key") or "").casefold()
        peer_rows = pair_counts.get((sku_key, dkey), [])
        is_dup = len(peer_rows) > 1
        prow["is_duplicate"] = is_dup
        prow["duplicate_rows"] = [r for r in peer_rows if r != prow.get("row_no")]
        total = item_totals.get(sku_key, int(prow.get("qty_total") or 0))
        allocated = sku_dir_sums.get(sku_key, 0)
        prow["sku_allocated"] = allocated
        prow["qty_mismatch"] = allocated != total
        prow["qty_delta"] = allocated - total
        if is_dup:
            prow["status"] = "duplicate"
        elif prow["row_no"] in error_rows or sku_key in error_skus or prow["qty_mismatch"]:
            prow["status"] = "error"
        elif any(
            i.severity == "warning"
            and (i.row_no == prow["row_no"] or (i.sku_code and i.sku_code.casefold() == sku_key))
            for i in issues
        ):
            prow["status"] = "warning"

    stats = _build_stats(draft, row_count)
    stats["duplicate_rows"] = sum(1 for p in preview_rows if p.get("is_duplicate"))
    stats["duplicate_sku_dirs"] = sum(1 for peers in pair_counts.values() if len(peers) > 1)
    return ExcelPreview(
        draft=draft,
        issues=issues,
        row_count=row_count,
        preview_rows=preview_rows,
        stats=stats,
    )


def build_error_report_response(issues: list[ValidationIssue]) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"
    ws.append(["Строка", "Колонка", "Значение", "Тип", "Ошибка", "Как исправить", "Артикул", "Код"])
    for issue in issues:
        ws.append(
            [
                issue.row_no or "",
                issue.column,
                issue.value,
                "Предупреждение" if issue.severity == "warning" else "Ошибка",
                issue.message,
                issue.recommendation,
                issue.sku_code,
                issue.code,
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="receiving-distribution-errors.xlsx"'
    return response
