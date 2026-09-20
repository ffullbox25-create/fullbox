import json
from io import BytesIO
from datetime import datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from django.contrib.auth import get_user_model
from django.test import RequestFactory
from django.test import TestCase
from django.utils import timezone

from audit.models import AuditEntry
from employees.models import Employee
from fbs.models import (
    FbsBox,
    FbsInventorySession,
    FbsHandoverBatch,
    FbsOrder,
    FbsPallet,
    FbsPickBatch,
    FbsPickRestockRequest,
    FbsPickingCart,
    FbsStockBalance,
    FbsStorageCell,
    FbsToteBinding,
)
from marking.models import MarkingCode
from logistics.models import LogisticsTrip
from shipping.models import ShippingOrder
from sklad.models import WarehouseContainer, WarehouseEvent, WarehouseLocation, WarehouseStockSnapshot
from sklad.services.warehouse_events import WarehouseEventType
from sku.models import Agency, SKU
from todo.models import Task
from warehouse_goods.models import GoodsExtraFieldDefinition, GoodsExtraFieldValue
from wms_new.models import (
    WmsNewAcceptance,
    WmsNewAssemblyLine,
    WmsNewAssemblySession,
    WmsNewBox,
    WmsNewBoxItem,
    WmsNewBundleComponent,
    WmsNewBundleOperation,
    WmsNewBundleStock,
    WmsNewEvent,
    WmsNewExtraFieldDefinition,
    WmsNewExtraFieldValue,
    WmsNewInventorySession,
    WmsNewMarkingCode,
    WmsNewLogisticsManifest,
    WmsNewLogisticsOrder,
    WmsNewLogisticsPackage,
    WmsNewLogisticsRouteRule,
    WmsNewMovement,
    WmsNewOrder,
    WmsNewOrderItem,
    WmsNewProduct,
    WmsNewRecord,
    WmsNewReturn,
    WmsNewShipment,
    WmsNewShipmentBox,
    WmsNewShipmentOrder,
    WmsNewTask,
    WmsNewWave,
    WmsNewWaveAllocation,
    WmsNewWaveOrder,
    WmsNewWavePickLine,
)

from .employee_report import (
    _work_windows_from_activity_points,
    fbs_metric_points,
    warehouse_palletization_metric_points,
    warehouse_placement_metric_points,
)
from .models import Carrier, OwnCompany
from .services import (
    build_marketplace_warehouses_context,
    build_reference_form_context,
    build_reference_list_context,
)
from .views import _normalize_marketplace_warehouse_rows
from .web_ui import (
    _SKU_MOVEMENT_EXPORT_ROW_LIMIT,
    _head_manager_report_categories,
    _head_manager_sku_movement_export_response,
    _head_manager_warehouse_employee_report_rows,
    _head_manager_warehouse_employee_report_workbook,
    _movement_report_response,
    _movement_report_rows,
)


class HeadManagerReportStockNavigationTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username="hm_report_stock_navigation",
            password="pwd",
        )
        Employee.objects.create(
            full_name="Старший менеджер",
            role="head_manager",
            user=self.user,
            is_active=True,
        )
        self.client.force_login(self.user)

    def test_stock_is_available_only_from_warehouse_reports(self):
        categories = {
            category["slug"]: category
            for category in _head_manager_report_categories()
        }
        product_stock_links = [
            report["href"]
            for report in categories["products"]["reports"]
            if report["href"] == "/head-manager/stock-editor/"
        ]
        warehouse_stock_reports = [
            report
            for report in categories["warehouse"]["reports"]
            if report["href"] == "/head-manager/stock-editor/"
        ]

        self.assertEqual(product_stock_links, [])
        self.assertEqual(len(warehouse_stock_reports), 1)
        self.assertEqual(warehouse_stock_reports[0]["title"], "Остатки склада")

        reports_response = self.client.get("/head-manager/reports/")
        self.assertEqual(reports_response.status_code, 200)
        self.assertNotContains(
            reports_response,
            '<span class="nav-title">Остатки</span>',
            html=False,
        )

        warehouse_response = self.client.get("/head-manager/reports/warehouse/")
        self.assertEqual(warehouse_response.status_code, 200)
        self.assertContains(warehouse_response, "Остатки склада")
        self.assertContains(warehouse_response, "/head-manager/stock-editor/")

        stock_response = self.client.get("/head-manager/stock-editor/")
        self.assertEqual(stock_response.status_code, 200)
        self.assertContains(stock_response, "Складские отчёты")
        self.assertContains(stock_response, "/head-manager/reports/warehouse/")


class HeadManagerSkuMovementExportTests(TestCase):
    @patch("head_manager.web_ui._head_manager_sku_movement_rows")
    def test_export_defaults_to_last_30_days_and_caps_rows(self, rows_mock):
        rows_mock.return_value = ([], {})

        response = _head_manager_sku_movement_export_response({})

        self.assertEqual(response.status_code, 200)
        export_filters = rows_mock.call_args.args[0]
        self.assertEqual(export_filters["date_to"], timezone.localdate().isoformat())
        self.assertEqual(
            export_filters["date_from"],
            (timezone.localdate() - timedelta(days=29)).isoformat(),
        )
        self.assertEqual(rows_mock.call_args.kwargs["limit"], _SKU_MOVEMENT_EXPORT_ROW_LIMIT)

    @patch("head_manager.web_ui._head_manager_sku_movement_rows")
    def test_export_preserves_explicit_period(self, rows_mock):
        rows_mock.return_value = ([], {})

        _head_manager_sku_movement_export_response(
            {"date_from": "2026-08-01", "date_to": "2026-08-02"}
        )

        export_filters = rows_mock.call_args.args[0]
        self.assertEqual(export_filters["date_from"], "2026-08-01")
        self.assertEqual(export_filters["date_to"], "2026-08-02")


class HeadManagerMovementReportExportTests(TestCase):
    def test_rows_exclude_archived_stock(self):
        agency = Agency.objects.create(agn_name="Клиент отчета")
        active = WarehouseStockSnapshot.objects.create(
            agency=agency,
            sku_code="ACTIVE-SKU",
            name="Активный товар",
            qty=2,
            available_qty=2,
        )
        WarehouseStockSnapshot.objects.create(
            agency=agency,
            sku_code="ARCHIVED-SKU",
            name="Архивный товар",
            qty=3,
            available_qty=0,
            is_archived=True,
        )

        rows = _movement_report_rows()

        self.assertEqual([row["sku"] for row in rows], [active.sku_code])

    @patch("head_manager.web_ui._movement_report_rows")
    def test_export_is_valid_xlsx_and_cleans_unsafe_text(self, rows_mock):
        rows_mock.return_value = [
            {
                "index": 1,
                "client": "Клиент",
                "pallet_code": "PALLET",
                "box_code": "BOX",
                "sku": "SKU",
                "name": "=опасная\x00строка",
                "barcode": "123",
                "qty": 1,
                "status": "На складе",
                "width": "",
                "height": "",
                "depth": "",
                "volume": "",
                "weight": "",
                "receipt_date": "",
                "receipt_document": "",
                "expense_date": "",
                "expense_document": "",
                "is_processing_source": True,
            }
        ]

        response = _movement_report_response()
        workbook = load_workbook(BytesIO(response.content), read_only=False)
        sheet = workbook["Отчет"]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["F2"].value, "'=опаснаястрока")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["D2"].fill.fgColor.rgb[-6:], "DDEBF7")


class HeadManagerStockEditorPalletActionsTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="hm_stock_editor", password="pwd")
        Employee.objects.create(full_name="Главный менеджер склада", role="head_manager", user=self.user, is_active=True)
        self.client.force_login(self.user)
        self.agency = Agency.objects.create(agn_name="Клиент редактора остатков")
        self.source_location = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="OS",
            zone_kind=WarehouseLocation.ZONE_KIND_STORAGE,
            row_no=1,
            section_no=1,
            tier_no=1,
            cell_no=1,
            location_code="0-1/1-1",
            display_name="OS · 0-1/1-1",
            is_storage=True,
        )
        self.target_location = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="OS",
            zone_kind=WarehouseLocation.ZONE_KIND_STORAGE,
            row_no=1,
            section_no=1,
            tier_no=1,
            cell_no=2,
            location_code="0-1/1-2",
            display_name="OS · 0-1/1-2",
            is_storage=True,
        )
        self.receiving_location = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="PR",
            zone_kind=WarehouseLocation.ZONE_KIND_RECEIVING,
            location_code="PR",
            display_name="PR · Зона приемки",
            is_storage=False,
        )
        self.source_pallet = self._create_pallet("SRC-PALLET", self.source_location)
        self.target_pallet = self._create_pallet("DST-PALLET", self.target_location)

    def _create_pallet(self, code, location):
        return WarehouseContainer.objects.create(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_PALLET,
            container_code=code,
            current_location=location,
        )

    def _create_box(self, code, pallet=None, location=None):
        return WarehouseContainer.objects.create(
            agency=self.agency,
            container_type=WarehouseContainer.TYPE_BOX,
            container_code=code,
            parent_container=pallet or self.source_pallet,
            current_location=location or self.source_location,
        )

    def _create_snapshot(self, code, *, box=None, pallet=None, qty=1, state=""):
        container = box or pallet
        parent = box.parent_container if box else None
        location = container.current_location
        return WarehouseStockSnapshot.objects.create(
            agency=self.agency,
            sku_code=code,
            name=code,
            qty=qty,
            available_qty=qty,
            container=container,
            container_code=container.container_code,
            parent_container=parent,
            location=location,
            zone_code=location.zone_code,
            zone_kind=location.zone_kind,
            warehouse_state_code=state,
        )

    def _post_action(self, snapshot, field, value):
        return self.client.post(
            "/head-manager/stock-editor/",
            data=json.dumps({"snapshot_id": snapshot.pk, "field": field, "value": value}),
            content_type="application/json",
        )

    def test_context_menus_expose_separate_box_and_pallet_actions(self):
        response = self.client.get("/head-manager/stock-editor/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-action="move-box-to-pallet"', html=False)
        self.assertContains(response, "Перенести короб на существующую паллету")
        self.assertContains(response, 'data-action="move-box-to-new-pallet"', html=False)
        self.assertContains(response, "Поставить короб на новую паллету")
        self.assertContains(response, 'data-action="merge-pallets"', html=False)
        self.assertContains(response, "Объединить паллеты")
        self.assertContains(response, 'data-action="rename-pallet"', html=False)
        self.assertContains(response, "Изменить номер паллеты")

    def test_goods_type_can_be_changed_from_cell_and_is_audited(self):
        box = self._create_box("BOX-GOODS-TYPE")
        snapshot = self._create_snapshot("SKU-GOODS-TYPE", box=box, qty=4)
        snapshot.goods_type = "no"
        snapshot.save(update_fields=["goods_type", "updated_at"])

        template_source = (
            Path(__file__).resolve().parent / "templates" / "head_manager" / "stock_editor.html"
        ).read_text(encoding="utf-8")
        self.assertIn('data-edit-field="goods_type"', template_source)
        self.assertIn('id="stockGoodsTypeOptions"', template_source)

        response = self._post_action(snapshot, "goods_type", "Готовый")

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["row"]["goods_type"], "gv")
        snapshot.refresh_from_db()
        self.assertEqual(snapshot.goods_type, "gv")
        self.assertIsNotNone(snapshot.last_event_id)
        event = WarehouseEvent.objects.get(pk=snapshot.last_event_id)
        self.assertEqual(event.event_type, WarehouseEventType.STOCK_CORRECTED.value)
        self.assertEqual(event.performed_by_id, self.user.pk)
        self.assertEqual(event.performed_by_role, "head_manager")
        self.assertEqual(event.payload["editor_action"], "goods_type_changed")
        self.assertEqual(event.payload["old_goods_type"], "no")
        self.assertEqual(event.payload["new_goods_type"], "gv")

    def test_goods_type_rejects_unknown_value(self):
        box = self._create_box("BOX-BAD-GOODS-TYPE")
        snapshot = self._create_snapshot("SKU-BAD-GOODS-TYPE", box=box)
        snapshot.goods_type = "gv"
        snapshot.save(update_fields=["goods_type", "updated_at"])

        response = self._post_action(snapshot, "goods_type", "неизвестный")

        self.assertEqual(response.status_code, 400, response.content)
        snapshot.refresh_from_db()
        self.assertEqual(snapshot.goods_type, "gv")
        self.assertFalse(
            WarehouseEvent.objects.filter(
                event_type=WarehouseEventType.STOCK_CORRECTED.value,
                stock_context_id=str(snapshot.pk),
            ).exists()
        )

    def test_move_selected_box_moves_all_its_rows_only(self):
        selected_box = self._create_box("BOX-SELECTED")
        other_box = self._create_box("BOX-OTHER")
        first_row = self._create_snapshot("SKU-1", box=selected_box, qty=2)
        second_row = self._create_snapshot("SKU-2", box=selected_box, qty=3)
        other_row = self._create_snapshot("SKU-3", box=other_box, qty=4)

        response = self._post_action(first_row, "move_box_to_pallet", self.target_pallet.container_code)

        self.assertEqual(response.status_code, 200, response.content)
        selected_box.refresh_from_db()
        other_box.refresh_from_db()
        first_row.refresh_from_db()
        second_row.refresh_from_db()
        other_row.refresh_from_db()
        self.assertEqual(selected_box.parent_container_id, self.target_pallet.pk)
        self.assertEqual(selected_box.current_location_id, self.target_location.pk)
        self.assertEqual(first_row.parent_container_id, self.target_pallet.pk)
        self.assertEqual(second_row.parent_container_id, self.target_pallet.pk)
        self.assertEqual(other_box.parent_container_id, self.source_pallet.pk)
        self.assertEqual(other_row.parent_container_id, self.source_pallet.pk)
        event = WarehouseEvent.objects.get(event_type="stock_editor_box_moved_to_pallet")
        self.assertEqual(event.performed_by_id, self.user.pk)
        self.assertEqual(event.performed_by_role, "head_manager")
        self.assertEqual(event.qty, 5)

    def test_move_selected_box_to_new_pallet_creates_it_in_pr(self):
        selected_box = self._create_box("BOX-NEW-PALLET")
        other_box = self._create_box("BOX-STAYS-SOURCE")
        first_row = self._create_snapshot("SKU-NEW-1", box=selected_box, qty=2)
        second_row = self._create_snapshot("SKU-NEW-2", box=selected_box, qty=3)
        other_row = self._create_snapshot("SKU-STAYS", box=other_box, qty=4)

        response = self._post_action(first_row, "move_box_to_new_pallet", "NEW-PR-PALLET")

        self.assertEqual(response.status_code, 200, response.content)
        new_pallet = WarehouseContainer.objects.get(container_code="NEW-PR-PALLET")
        selected_box.refresh_from_db()
        other_box.refresh_from_db()
        first_row.refresh_from_db()
        second_row.refresh_from_db()
        other_row.refresh_from_db()
        self.assertEqual(new_pallet.container_type, WarehouseContainer.TYPE_PALLET)
        self.assertEqual(new_pallet.current_location_id, self.receiving_location.pk)
        self.assertEqual(new_pallet.created_by_id, self.user.pk)
        self.assertEqual(selected_box.container_code, "BOX-NEW-PALLET")
        self.assertEqual(selected_box.parent_container_id, new_pallet.pk)
        self.assertEqual(selected_box.current_location_id, self.receiving_location.pk)
        self.assertEqual(first_row.parent_container_id, new_pallet.pk)
        self.assertEqual(second_row.parent_container_id, new_pallet.pk)
        self.assertEqual(first_row.location_id, self.receiving_location.pk)
        self.assertEqual(second_row.location_id, self.receiving_location.pk)
        self.assertEqual(first_row.zone_code, "PR")
        self.assertEqual(second_row.zone_code, "PR")
        self.assertEqual(first_row.warehouse_state_code, "placed_in_receiving")
        self.assertEqual(second_row.warehouse_state_code, "placed_in_receiving")
        self.assertEqual(other_box.parent_container_id, self.source_pallet.pk)
        self.assertEqual(other_row.parent_container_id, self.source_pallet.pk)
        event = WarehouseEvent.objects.get(event_type="stock_editor_box_moved_to_new_pallet")
        self.assertEqual(event.performed_by_id, self.user.pk)
        self.assertEqual(event.qty, 5)
        self.assertEqual(event.payload["source_pallet_code"], self.source_pallet.container_code)
        self.assertEqual(event.payload["target_pallet_code"], "NEW-PR-PALLET")

    def test_move_box_to_new_pallet_rejects_used_number(self):
        selected_box = self._create_box("BOX-DUPLICATE-PALLET")
        snapshot = self._create_snapshot("SKU-DUPLICATE-PALLET", box=selected_box, qty=2)

        response = self._post_action(
            snapshot,
            "move_box_to_new_pallet",
            self.target_pallet.container_code,
        )

        self.assertEqual(response.status_code, 400, response.content)
        selected_box.refresh_from_db()
        self.assertEqual(selected_box.parent_container_id, self.source_pallet.pk)
        self.assertFalse(
            WarehouseEvent.objects.filter(event_type="stock_editor_box_moved_to_new_pallet").exists()
        )

    def test_merge_pallets_moves_every_box_and_direct_stock(self):
        first_box = self._create_box("BOX-MERGE-1")
        second_box = self._create_box("BOX-MERGE-2")
        first_row = self._create_snapshot("SKU-MERGE-1", box=first_box, qty=2)
        second_row = self._create_snapshot("SKU-MERGE-2", box=second_box, qty=3)
        direct_row = self._create_snapshot("SKU-DIRECT", pallet=self.source_pallet, qty=4)

        response = self._post_action(first_row, "merge_pallets", self.target_pallet.container_code)

        self.assertEqual(response.status_code, 200, response.content)
        self.source_pallet.refresh_from_db()
        first_box.refresh_from_db()
        second_box.refresh_from_db()
        first_row.refresh_from_db()
        second_row.refresh_from_db()
        direct_row.refresh_from_db()
        self.assertEqual(self.source_pallet.status, WarehouseContainer.STATUS_MERGED)
        self.assertIsNone(self.source_pallet.current_location_id)
        self.assertEqual(first_box.parent_container_id, self.target_pallet.pk)
        self.assertEqual(second_box.parent_container_id, self.target_pallet.pk)
        self.assertEqual(first_row.parent_container_id, self.target_pallet.pk)
        self.assertEqual(second_row.parent_container_id, self.target_pallet.pk)
        self.assertEqual(direct_row.container_id, self.target_pallet.pk)
        self.assertEqual(direct_row.container_code, self.target_pallet.container_code)
        event = WarehouseEvent.objects.get(event_type="stock_editor_pallets_merged")
        self.assertEqual(event.performed_by_id, self.user.pk)
        self.assertEqual(event.qty, 9)
        self.assertCountEqual(event.payload["moved_box_codes"], ["BOX-MERGE-1", "BOX-MERGE-2"])

    def test_rename_pallet_changes_only_its_free_number(self):
        box = self._create_box("BOX-RENAME")
        snapshot = self._create_snapshot("SKU-RENAME", box=box, qty=6)
        original_location_id = self.source_pallet.current_location_id

        response = self._post_action(snapshot, "rename_pallet", "RENAMED-PALLET")

        self.assertEqual(response.status_code, 200, response.content)
        self.source_pallet.refresh_from_db()
        box.refresh_from_db()
        snapshot.refresh_from_db()
        self.assertEqual(self.source_pallet.container_code, "RENAMED-PALLET")
        self.assertEqual(self.source_pallet.current_location_id, original_location_id)
        self.assertEqual(self.source_pallet.status, WarehouseContainer.STATUS_ACTIVE)
        self.assertEqual(box.parent_container_id, self.source_pallet.pk)
        self.assertEqual(snapshot.parent_container_id, self.source_pallet.pk)
        event = WarehouseEvent.objects.get(event_type="stock_editor_pallet_renamed")
        self.assertEqual(event.performed_by_id, self.user.pk)
        self.assertEqual(event.payload["old_pallet_code"], "SRC-PALLET")
        self.assertEqual(event.payload["new_pallet_code"], "RENAMED-PALLET")

    def test_merge_is_rejected_when_any_source_position_is_locked(self):
        first_box = self._create_box("BOX-UNLOCKED")
        second_box = self._create_box("BOX-LOCKED")
        first_row = self._create_snapshot("SKU-UNLOCKED", box=first_box)
        self._create_snapshot("SKU-LOCKED", box=second_box, state="reserved_for_shipping")

        response = self._post_action(first_row, "merge_pallets", self.target_pallet.container_code)

        self.assertEqual(response.status_code, 400, response.content)
        self.source_pallet.refresh_from_db()
        first_box.refresh_from_db()
        second_box.refresh_from_db()
        self.assertEqual(self.source_pallet.status, WarehouseContainer.STATUS_ACTIVE)
        self.assertEqual(first_box.parent_container_id, self.source_pallet.pk)
        self.assertEqual(second_box.parent_container_id, self.source_pallet.pk)
        self.assertFalse(WarehouseEvent.objects.filter(event_type="stock_editor_pallets_merged").exists())


class HeadManagerFbsToteJournalTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="hm_tote_journal", password="pwd")
        Employee.objects.create(
            full_name="Начальник склада",
            role="head_manager",
            user=self.user,
            is_active=True,
        )
        self.client.force_login(self.user)

    def test_settings_show_tote_status_location_and_active_wave(self):
        user_model = get_user_model()
        picker = user_model.objects.create_user(username="picker_tote_journal", password="pwd")
        Employee.objects.create(
            full_name="Сборщик Тары",
            role="picker",
            user=picker,
            is_active=True,
        )
        agency = Agency.objects.create(agn_name="Клиент журнала тары")
        free_cart = FbsPickingCart.objects.create(
            barcode="FBS-CART-910201",
            name="Свободная тара",
        )
        busy_cart = FbsPickingCart.objects.create(
            barcode="FBS-CART-910202",
            name="Тара в работе",
        )
        wave = FbsPickBatch.objects.create(
            agency=agency,
            status=FbsPickBatch.STATUS_IN_PROGRESS,
            assigned_to=picker,
            cart=busy_cart,
        )
        FbsToteBinding.objects.create(
            tote=busy_cart,
            state=FbsToteBinding.STATE_PICKING,
            employee=picker,
            pick_batch=wave,
            updated_by=self.user,
        )

        response = self.client.get("/head-manager/fbs/settings/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Журнал тары")
        self.assertContains(response, free_cart.barcode)
        self.assertContains(response, "Свободна 1")
        self.assertContains(response, busy_cart.barcode)
        self.assertContains(response, "В работе 1")
        self.assertContains(response, "У сборщика")
        self.assertContains(response, "У сотрудника: Сборщик Тары")
        self.assertContains(response, f"Волна #{wave.id}")

    def test_create_tote_returns_to_tote_journal(self):
        response = self.client.post(
            "/head-manager/fbs/settings/equipment/cart/create/",
            {"count": 1, "name_prefix": "Новая тара", "notes": ""},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], "/head-manager/fbs/settings/#tote-journal")
        self.assertTrue(FbsPickingCart.objects.filter(name__startswith="Новая тара").exists())


class HeadManagerDirectoryTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="hm_test", password="pwd")
        self.employee = Employee.objects.create(
            full_name="Главный менеджер",
            role="head_manager",
            user=self.user,
            is_active=True,
        )
        self.client.force_login(self.user)

    def test_manager_with_head_manager_access_can_open_dashboard(self):
        self.employee.role = "manager"
        self.employee.access_roles = ["head_manager"]
        self.employee.save(update_fields=["role", "access_roles"])

        response = self.client.get("/head-manager/")

        self.assertEqual(response.status_code, 200)

    def test_dashboard_groups_reference_links(self):
        response = self.client.get("/head-manager/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Справочники")
        self.assertContains(response, "/head-manager/own-companies/")
        self.assertContains(response, "/head-manager/carriers/")

    def test_dashboard_keeps_old_fbs_and_adds_separate_fbs_new_entry(self):
        response = self.client.get("/head-manager/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'href="/head-manager/fbs/">FBS</a>', html=False)
        self.assertContains(response, 'href="/head-manager/fbs-new/"', html=False)
        self.assertContains(response, '<span class="nav-title">FBS-NEW</span>', html=False)

    def test_fbs_new_opens_parallel_full_wms_shell(self):
        response = self.client.get("/head-manager/fbs-new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'aria-label="Навигация FBS-NEW"', html=False)
        self.assertContains(response, 'href="/head-manager/"', html=False)
        self.assertNotContains(response, 'href="/head-manager/" aria-current="page"', html=False)
        for label in (
            "Главная",
            "Проблемы",
            "Задачи",
            "FBS",
            "Склад",
            "Логистика",
            "Документы",
            "Отчеты",
            "Аналитика",
            "Партнеры",
            "CRM",
            "Настройки",
            "Помощь",
        ):
            self.assertContains(response, label)
        self.assertContains(response, "Параллельный тестовый контур")

    def test_fbs_new_section_has_own_active_navigation_and_rejects_unknown_section(self):
        response = self.client.get("/head-manager/fbs-new/warehouse/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'href="/head-manager/fbs-new/warehouse-goods/" aria-current="page"',
            html=False,
        )
        self.assertContains(response, '<h2 class="card-title">Склад</h2>', html=False)
        self.assertEqual(self.client.get("/head-manager/fbs-new/not-a-section/").status_code, 404)

    def test_fbs_new_uses_measured_ts_wms_design_tokens(self):
        template_path = Path(__file__).resolve().parents[1] / "templates/head_manager/fbs_new.html"
        template_text = template_path.read_text(encoding="utf-8")

        self.assertIn("--ts-body: #f7f7f9", template_text)
        self.assertIn("--ts-primary: #666cff", template_text)
        self.assertIn("width: 260px", template_text)
        self.assertIn("min-height: 40px", template_text)
        self.assertIn("--ts-radius: 8px", template_text)
        self.assertIn("0 2px 10px rgba(76, 78, 100, .22)", template_text)
        self.assertIn("font-family: Inter", template_text)

    def test_fbs_new_settings_and_help_render_donor_sections(self):
        expected = {
            "settings-billing": "Оплата по счету",
            "settings-users": "Добавить пользователя",
            "settings-places": "Добавить место",
            "settings-printers": "РАБОЧИЕ МЕСТА",
            "settings-directories": "Единицы измерения",
            "settings-system": "Общие настройки",
            "settings-history": "Дата и время",
            "settings-labels": "Редактор шаблонов этикеток",
            "settings-matching": "Каталог WMS",
            "support": "Техническая поддержка",
            "instructions": "ПЕРВИЧНАЯ НАСТРОЙКА",
            "development": "Предложить задачу",
        }
        for slug, label in expected.items():
            with self.subTest(slug=slug):
                response = self.client.get(f"/head-manager/fbs-new/{slug}/")
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, label)

    def test_fbs_new_settings_actions_only_write_shadow_records(self):
        agency = Agency.objects.create(agn_name="Партнер настроек")
        product = WmsNewProduct.objects.create(
            agency=agency, name="Товар настроек", article="SETTINGS-1"
        )
        location = WarehouseLocation.objects.create(
            warehouse_code="MSK", zone_code="SET", location_code="SET-1"
        )
        source_employee_state = (self.employee.role, self.employee.is_active)
        source_location_state = (location.display_name, location.is_active)
        source_product_state = (product.name, product.article)

        responses = (
            self.client.post("/head-manager/fbs-new/settings/system/action/", {
                "acceptances_per_page": "300", "documents_per_page": "100",
                "products_per_page": "100", "session_timeout": "2",
                "mixed_partner_pallets": "no", "tab_direction": "vertical", "table_mode": "full",
            }),
            self.client.post("/head-manager/fbs-new/settings/users/action/", {
                "employee_id": self.employee.id, "notifications": "1", "active": "1",
            }),
            self.client.post("/head-manager/fbs-new/settings/places/action/", {
                "source_key": location.id, "title": "Пилотное место", "code": "SET-1",
                "warehouse": "MSK", "zone": "SET", "active": "0",
            }),
            self.client.post("/head-manager/fbs-new/settings/matching/action/", {
                "product_id": product.id, "ozon": "OZ-1", "wb": "WB-1",
            }),
            self.client.post("/head-manager/fbs-new/settings/printers/action/", {
                "action": "test", "title": "Тест этикетки",
            }),
            self.client.post("/head-manager/fbs-new/help/action/", {
                "title": "Проверить новый экран", "description": "Только FBS-NEW",
            }),
        )
        self.assertTrue(all(response.status_code == 302 for response in responses))
        self.employee.refresh_from_db()
        location.refresh_from_db()
        product.refresh_from_db()
        self.assertEqual((self.employee.role, self.employee.is_active), source_employee_state)
        self.assertEqual((location.display_name, location.is_active), source_location_state)
        self.assertEqual((product.name, product.article), source_product_state)
        self.assertTrue(WmsNewRecord.objects.filter(module="settings", entity_type="system").exists())
        self.assertTrue(WmsNewRecord.objects.filter(module="settings", entity_type="goods_mapping").exists())
        self.assertTrue(WmsNewRecord.objects.filter(module="help", entity_type="development_task").exists())
        self.assertGreaterEqual(WmsNewEvent.objects.count(), 6)

    def test_fbs_new_all_sections_render_read_only_fullbox_data_views(self):
        response = self.client.get("/head-manager/fbs-new/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "SLA приемки поставок")
        self.assertContains(response, "Таймер дедлайна заказов")
        self.assertNotContains(response, "Каркас")

        for slug in (
            "problems",
            "tasks",
            "fbs",
            "fbs-orders",
            "fbs-queue",
            "fbs-waves",
            "fbs-assembly",
            "fbs-shipments",
            "fbs-returns",
            "warehouse",
            "logistics",
            "documents",
            "reports",
            "analytics",
            "partners",
            "crm",
            "settings",
            "help",
            "support",
            "instructions",
            "development",
        ):
            with self.subTest(slug=slug):
                response = self.client.get(f"/head-manager/fbs-new/{slug}/")
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context["section_data"].get("title"))
                if slug == "fbs-assembly":
                    self.assertContains(response, "Для начала сортировки отсканируйте место")
                elif slug == "instructions":
                    self.assertContains(response, 'class="instruction-item', html=False)
                else:
                    self.assertContains(response, 'class="module-table', html=False)

    def test_fbs_new_queue_matches_ts_wms_queue_structure(self):
        response = self.client.get("/head-manager/fbs-new/fbs-queue/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "/ FBS")
        self.assertContains(response, "Добавить заказ")
        self.assertContains(response, "Действия с выбранными")
        for label in (
            "Номер заказа",
            "Все партнеры",
            "Все доставки",
            "Все интеграции",
            "Все статусы",
            "Все зоны",
            "Любое кол-во товара",
            "Все наличие",
            "Дата создания",
            "Очистить фильтры",
            "Последнее обновление",
            "На странице",
            "ID▼",
            "Партнёр",
            "Интеграция",
            "Тип доставки",
            "Трэк номер",
            "Кол-во товаров",
            "Статус наличия",
            "Наличие",
        ):
            self.assertContains(response, label)
        self.assertNotContains(
            response,
            "Новый интерфейс работает параллельно. Старый FBS и production write path не переключены.",
        )

    def test_fbs_new_orders_matches_ts_wms_order_list_and_uses_detail_routes(self):
        agency = Agency.objects.create(agn_name="Партнер списка заказов")
        order = WmsNewOrder.objects.create(
            source_order_id=321001,
            agency=agency,
            marketplace="wb",
            integration_name="Интеграция списка",
            external_order_id="FBN-LIST-1",
            delivery_type="Wildberries FBS",
            status=WmsNewOrder.STATUS_NEW,
            source_created_at=timezone.now(),
        )
        WmsNewOrderItem.objects.create(
            order=order,
            external_line_id="list-line-1",
            quantity=1,
        )

        response = self.client.get("/head-manager/fbs-new/fbs-orders/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["section_data"]["kind"], "fbs_orders_list")
        for label in (
            "/ FBS",
            "Заказы",
            "Добавить заказ",
            "Действия с выбранными",
            "Все партнеры",
            "Все доставки",
            "Все интеграции",
            "Все статусы",
            "Любое кол-во товара",
            "Дата создания",
            "ID▼",
            "Трэк номер",
            "Кол-во товаров",
            "321001",
            "FBN-LIST-1",
        ):
            self.assertContains(response, label)
        self.assertContains(
            response,
            f'data-row-url="/head-manager/fbs-new/fbs-orders/{order.id}/"',
            html=False,
        )
        self.assertNotContains(response, "/head-manager/fbs/?")

    def test_fbs_new_queue_preserves_filters_and_page_size(self):
        response = self.client.get(
            "/head-manager/fbs-new/fbs-queue/",
            {"q": "TEST-ORDER", "quantity": "many", "page_size": "50"},
        )

        self.assertEqual(response.status_code, 200)
        queue = response.context["section_data"]
        self.assertEqual(queue["filters"]["q"], "TEST-ORDER")
        self.assertEqual(queue["filters"]["quantity"], "many")
        self.assertEqual(queue["page_size"], 50)
        self.assertContains(response, 'value="TEST-ORDER"', html=False)
        self.assertContains(response, 'value="many" selected', html=False)

    def test_fbs_new_queue_row_opens_independent_order_detail(self):
        agency = Agency.objects.create(agn_name="Партнер карточки заказа")
        order = WmsNewOrder.objects.create(
            source_order_id=987654,
            agency=agency,
            marketplace="wb",
            integration_name="Wildberries · Фулбокс",
            external_order_id="FBN-DETAIL-1",
            delivery_type="Wildberries FBS",
            warehouse_code="Основной склад",
            status=WmsNewOrder.STATUS_NEW,
            source_created_at=timezone.now(),
        )
        WmsNewOrderItem.objects.create(
            order=order,
            external_line_id="detail-line-1",
            external_sku="ARTICLE-1",
            barcode="BARCODE-1",
            product_name="Тестовый товар карточки",
            quantity=2,
            source_snapshot={"convertedPrice": 123400},
        )

        queue_response = self.client.get("/head-manager/fbs-new/fbs-queue/")
        detail_url = f"/head-manager/fbs-new/fbs-orders/{order.id}/"
        self.assertContains(
            queue_response,
            f'data-row-url="{detail_url}" tabindex="0"',
            html=False,
        )

        detail_response = self.client.get(detail_url)
        self.assertEqual(detail_response.status_code, 200)
        self.assertEqual(detail_response.context["section_data"]["kind"], "fbs_order_detail")
        for label in (
            "Заказ #987654",
            "Тип доставки: Wildberries FBS",
            "№ заказа у партнера:",
            "FBN-DETAIL-1",
            "Сменить статус заказа",
            "История изменений",
            "Получатель заказа",
            "Комментарии к заказу",
            "Товары",
            "ARTICLE-1",
            "BARCODE-1",
            "Тестовый товар карточки",
        ):
            self.assertContains(detail_response, label)
        self.assertNotContains(detail_response, "/head-manager/fbs/?")

    def test_fbs_new_order_detail_returns_404_for_unknown_order(self):
        response = self.client.get("/head-manager/fbs-new/fbs-orders/999999/")

        self.assertEqual(response.status_code, 404)

    def test_fbs_new_waves_match_donor_list_and_open_independent_detail(self):
        agency = Agency.objects.create(agn_name="Партнер волны FBS-NEW")
        order = WmsNewOrder.objects.create(
            source_order_id=70201,
            agency=agency,
            external_order_id="FBN-WAVE-ORDER-1",
            status=WmsNewOrder.STATUS_QUEUED,
            source_created_at=timezone.now(),
        )
        WmsNewOrderItem.objects.create(
            order=order,
            source_item_id=90301,
            external_line_id="wave-detail-line-1",
            external_sku="WAVE-ARTICLE-1",
            barcode="WAVE-BARCODE-1",
            product_name="Товар в карточке волны",
            quantity=2,
        )
        wave = WmsNewWave.objects.create(
            source_batch_id=1027,
            number="FB-1027",
            agency=agency,
            planned_orders=1,
            planned_units=2,
            source_created_at=timezone.now(),
            last_synced_at=timezone.now(),
            is_manual=True,
        )
        WmsNewWaveOrder.objects.create(wave=wave, order=order, sequence=1)
        pick_line = WmsNewWavePickLine.objects.create(
            wave=wave,
            order=order,
            order_item=order.items.get(),
            planned_quantity=2,
            sequence=1,
        )
        WmsNewWaveAllocation.objects.create(
            pick_line=pick_line,
            source_place_code="A-01-01-01-01",
            source_place_name="A-01-01-01-01",
            reserved_quantity=2,
            sequence=1,
        )

        response = self.client.get("/head-manager/fbs-new/fbs-waves/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["section_data"]["kind"], "fbs_waves_list")
        for label in (
            "/ FBS",
            "Волны",
            "Все статусы",
            "Все пользователи",
            "Дата создания",
            "Время начала подбора",
            "Время окончания подбора",
            "Место подбора",
            "Кол-во заказов",
            "Кол-во товаров",
            "1027",
        ):
            self.assertContains(response, label)
        detail_url = f"/head-manager/fbs-new/fbs-waves/{wave.id}/"
        self.assertContains(response, f'data-row-url="{detail_url}" tabindex="0"', html=False)

        detail = self.client.get(detail_url)
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.context["section_data"]["kind"], "fbs_wave_detail")
        for label in (
            "Волны № 1027",
            "Способ создания: Вручную",
            "Действия с волной",
            "Статус волны",
            "Лист подбора: Скачать",
            "ID товара",
            "FBN-WAVE-ORDER-1",
            "Товар в карточке волны",
            "WAVE-BARCODE-1",
            "WAVE-ARTICLE-1",
        ):
            self.assertContains(detail, label)
        self.assertNotContains(detail, "/head-manager/fbs/")

        collect = self.client.get(
            f"/head-manager/fbs-new/fbs-waves/{wave.id}/collect/"
        )
        self.assertEqual(collect.status_code, 200)
        self.assertEqual(collect.context["section_data"]["kind"], "fbs_wave_collect")
        self.assertContains(collect, "Для начала подбора волны отсканируйте место подбора")
        self.assertContains(
            collect,
            f'/head-manager/fbs-new/fbs-waves/{wave.id}/action/',
            html=False,
        )
        self.assertNotContains(collect, "/head-manager/fbs/")

        export = self.client.get(
            f"/head-manager/fbs-new/fbs-waves/{wave.id}/collecting-list/"
        )
        self.assertEqual(export.status_code, 200)
        self.assertEqual(export["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("FBN-WAVE-ORDER-1", export.content.decode("utf-8-sig"))

    def test_fbs_new_wave_place_update_changes_only_isolated_wave(self):
        agency = Agency.objects.create(agn_name="Партнер места волны")
        FbsPickingCart.objects.create(
            barcode="FBS-CART-910099",
            name="Стол пилота",
        )
        wave = WmsNewWave.objects.create(
            number="FBN-PLACE-1",
            agency=agency,
            planned_orders=0,
            planned_units=0,
            is_manual=True,
        )
        old_wave_count = FbsPickBatch.objects.count()

        response = self.client.post(
            f"/head-manager/fbs-new/fbs-waves/{wave.id}/update/",
            {"status": "", "place_name": "Стол пилота"},
        )

        self.assertEqual(response.status_code, 302)
        wave.refresh_from_db()
        self.assertEqual(wave.cart_name, "Стол пилота")
        self.assertEqual(wave.pilot_revision, 1)
        self.assertEqual(FbsPickBatch.objects.count(), old_wave_count)
        self.assertTrue(
            WmsNewEvent.objects.filter(
                entity_type="wave", entity_id=wave.id, action="assign_place"
            ).exists()
        )

    def test_fbs_new_wave_detail_returns_404_for_unknown_wave(self):
        response = self.client.get("/head-manager/fbs-new/fbs-waves/999999/")

        self.assertEqual(response.status_code, 404)

    def test_fbs_new_assembly_matches_donor_scanner_and_runs_independently(self):
        agency = Agency.objects.create(agn_name="Партнер экрана сборки")
        cart = FbsPickingCart.objects.create(
            barcode="FBS-CART-910001",
            name="Место экрана сборки",
        )
        order = WmsNewOrder.objects.create(
            agency=agency,
            marketplace="wb",
            external_order_id="ASSEMBLY-SCREEN-ORDER",
            tracking_number="ASSEMBLY-SCREEN-TRACK",
            status=WmsNewOrder.STATUS_PICKED,
        )
        item = WmsNewOrderItem.objects.create(
            order=order,
            external_line_id="assembly-screen-line",
            external_sku="ASSEMBLY-SCREEN-SKU",
            barcode="ASSEMBLY-SCREEN-BARCODE",
            product_name="Товар экрана сборки",
            quantity=1,
        )
        wave = WmsNewWave.objects.create(
            number="FBN-ASSEMBLY-SCREEN",
            agency=agency,
            status=WmsNewWave.STATUS_VERIFICATION,
            planned_orders=1,
            planned_units=1,
            picked_units=1,
            cart_name=cart.name,
        )
        WmsNewWaveOrder.objects.create(wave=wave, order=order, sequence=1)
        legacy_wave_count = FbsPickBatch.objects.count()

        start_page = self.client.get("/head-manager/fbs-new/fbs-assembly/")
        self.assertEqual(start_page.status_code, 200)
        self.assertEqual(start_page.context["section_data"]["kind"], "fbs_assembly_start")
        self.assertContains(start_page, "/ FBS")
        self.assertContains(start_page, "Для начала сортировки отсканируйте место с товаром.")
        self.assertContains(start_page, "ID места:")
        self.assertContains(start_page, 'name="place_scan"', html=False)

        started = self.client.post(
            "/head-manager/fbs-new/fbs-assembly/start/",
            {"place_scan": cart.barcode},
        )
        self.assertEqual(started.status_code, 302)
        session = WmsNewAssemblySession.objects.get()
        work_url = f"/head-manager/fbs-new/fbs-assembly/?session={session.id}"
        work_page = self.client.get(work_url)
        self.assertEqual(work_page.context["section_data"]["kind"], "fbs_assembly_work")
        for label in (
            "Товар:",
            "Отсканируйте ШК товара",
            "Товары места сборки",
            item.barcode,
            order.external_order_id,
        ):
            self.assertContains(work_page, label)

        product_scan = self.client.post(
            f"/head-manager/fbs-new/fbs-assembly/{session.id}/scan/",
            {"action": "product", "scan": item.barcode},
        )
        self.assertEqual(product_scan.status_code, 302)
        line = WmsNewAssemblyLine.objects.get(session=session)
        self.assertEqual(line.status, WmsNewAssemblyLine.STATUS_AWAITING_ORDER)
        order_page = self.client.get(work_url)
        self.assertContains(order_page, "Отсканируйте ШК заказа")

        complete = self.client.post(
            f"/head-manager/fbs-new/fbs-assembly/{session.id}/scan/",
            {"action": "order", "scan": order.tracking_number},
        )
        self.assertEqual(complete.status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.status, WmsNewAssemblySession.STATUS_COMPLETED)
        self.assertEqual(FbsPickBatch.objects.count(), legacy_wave_count)
        self.assertNotContains(self.client.get(work_url), "/head-manager/fbs/")

    def test_fbs_new_shipments_match_donor_list_and_open_independent_detail(self):
        agency = Agency.objects.create(agn_name="Партнер отгрузки FBS-NEW")
        order = WmsNewOrder.objects.create(
            agency=agency,
            external_order_id="SHIPMENT-ORDER-1",
            tracking_number="SHIPMENT-STICKER-1",
            delivery_type="Wildberries FBS",
            status=WmsNewOrder.STATUS_READY,
            source_created_at=timezone.now(),
        )
        WmsNewOrderItem.objects.create(
            order=order,
            external_line_id="shipment-screen-line",
            product_name="Товар отгрузки",
            quantity=2,
        )
        shipment = WmsNewShipment.objects.create(
            source_batch_id=605,
            agency=agency,
            delivery_type="Wildberries FBS",
            integration_name="WB · Фулбокс",
            status=WmsNewShipment.STATUS_NEW,
            order_count=1,
            item_count=2,
            source_created_at=timezone.now(),
        )
        box = WmsNewShipmentBox.objects.create(
            shipment=shipment,
            qr_code="SHIPMENT-BOX-1",
            external_box_id="1",
        )
        WmsNewShipmentOrder.objects.create(
            shipment=shipment,
            order=order,
            box=box,
            sticker_number=order.tracking_number,
        )

        response = self.client.get("/head-manager/fbs-new/fbs-shipments/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["section_data"]["kind"], "fbs_shipments_list")
        for label in (
            "/ FBS",
            "Отгрузки",
            "Создать отгрузку",
            "Отправить все отгрузки",
            "Документы для отгрузок",
            "№ ЗАКАЗА",
            "Дата создания",
            "Все партнеры",
            "Все доставки",
            "Кроме принятых",
            "Кол-во заказов",
            "Кол-во товаров",
            "605",
        ):
            self.assertContains(response, label)
        detail_url = f"/head-manager/fbs-new/fbs-shipments/{shipment.id}/"
        self.assertContains(response, f'data-row-url="{detail_url}" tabindex="0"', html=False)

        detail = self.client.get(detail_url)
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.context["section_data"]["kind"], "fbs_shipment_detail")
        for label in (
            "Отгрузка #605",
            "Тип доставки",
            "Проверить отгрузку",
            "История изменений",
            "КОРОБ #1",
            "SHIPMENT-ORDER-1",
            "SHIPMENT-STICKER-1",
            "УСЛУГИ ПО ОТГРУЗКЕ",
            "Добавить услугу",
            "Завершить тарификацию по отгрузке",
        ):
            self.assertContains(detail, label)
        self.assertNotContains(detail, "/head-manager/fbs/")

    def test_fbs_new_shipment_create_writes_only_isolated_tables(self):
        agency = Agency.objects.create(agn_name="Пилотный партнер отгрузки")
        legacy_count = FbsHandoverBatch.objects.count()

        response = self.client.post(
            "/head-manager/fbs-new/fbs-shipments/create/",
            {
                "agency_id": agency.id,
                "delivery_type": "Wildberries FBS",
                "integration_name": "Пилот WB",
            },
        )

        self.assertEqual(response.status_code, 302)
        shipment = WmsNewShipment.objects.get(agency=agency)
        self.assertTrue(shipment.is_manual)
        self.assertEqual(FbsHandoverBatch.objects.count(), legacy_count)
        self.assertTrue(
            WmsNewEvent.objects.filter(
                entity_type="shipment", entity_id=shipment.id, action="create"
            ).exists()
        )

    def test_fbs_new_returns_match_donor_and_write_only_isolated_tables(self):
        agency = Agency.objects.create(agn_name="Партнер возвратов FBS-NEW")
        order = WmsNewOrder.objects.create(
            source_order_id=7654,
            agency=agency,
            external_order_id="RETURN-SCREEN-1",
            tracking_number="RETURN-TRACK-SCREEN-1",
            delivery_type="Wildberries FBS",
            status=WmsNewOrder.STATUS_DONE,
            source_created_at=timezone.now(),
        )
        WmsNewOrderItem.objects.create(
            order=order,
            external_line_id="return-screen-line",
            product_name="Возвращаемый товар",
            quantity=1,
        )
        item = WmsNewReturn.objects.create(
            source_request_id=91,
            order=order,
            status=WmsNewReturn.STATUS_COMPLETED,
            planned_qty=1,
            source_created_at=timezone.now(),
        )

        response = self.client.get("/head-manager/fbs-new/fbs-returns/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["section_data"]["kind"], "fbs_returns_list")
        for label in (
            "/ FBS",
            "Возвраты",
            "Добавить возврат",
            "Экспортировать",
            "Действия с выбранными",
            "Все партнеры",
            "Все доставки",
            "Дата создания",
            "Дата отгрузки",
            "Дата возврата",
            "ID заказа",
            "Номер задачи на возврат",
            "Статус задачи на возврат",
            "RETURN-SCREEN-1",
            "Возвращено",
        ):
            self.assertContains(response, label)
        self.assertContains(
            response,
            f'data-row-url="/head-manager/fbs-new/fbs-orders/{order.id}/" tabindex="0"',
            html=False,
        )
        export = self.client.get("/head-manager/fbs-new/fbs-returns/export/")
        self.assertEqual(export.status_code, 200)
        self.assertIn("RETURN-SCREEN-1", export.content.decode("utf-8-sig"))

        legacy_count = FbsPickRestockRequest.objects.count()
        created = self.client.post(
            "/head-manager/fbs-new/fbs-returns/create/",
            {"scan": order.tracking_number},
        )
        self.assertEqual(created.status_code, 302)
        self.assertEqual(FbsPickRestockRequest.objects.count(), legacy_count)
        self.assertEqual(WmsNewReturn.objects.filter(order=order).count(), 2)
        order.refresh_from_db()
        self.assertEqual(order.status, WmsNewOrder.STATUS_RETURN_PENDING)

    def test_fbs_new_bundles_match_donor_and_run_independently(self):
        agency = Agency.objects.create(agn_name="Партнер наборов FBS-NEW")
        bundle = WmsNewProduct.objects.create(
            agency=agency,
            name="Набор для экрана",
            article="SCREEN-BUNDLE",
        )
        component = WmsNewProduct.objects.create(
            agency=agency,
            name="Компонент набора",
            article="SCREEN-COMPONENT",
            stock_on_hand=5,
            stock_free=5,
        )
        location = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="KIT",
            row_no=8,
            section_no=1,
            tier_no=1,
            cell_no=1,
            location_code="KIT-SCREEN-1",
            is_active=True,
            is_storage=True,
        )
        old_sku_count = SKU.objects.count()

        create_page = self.client.get(
            "/head-manager/fbs-new/warehouse-bundles/",
            {"mode": "create", "partner": agency.id},
        )
        self.assertEqual(create_page.status_code, 200)
        self.assertEqual(create_page.context["section_data"]["kind"], "warehouse_bundles")
        for label in (
            "/ Склад",
            "Наборы",
            "Собрать набор",
            "Разобрать набор",
            "Создать новый набор",
            "01",
            "Выбор товара",
            "02",
            "Состав",
            "Товар, который будет преобразован в набор",
            "После сохранения состав будет изменить невозможно",
        ):
            self.assertContains(create_page, label)

        defined = self.client.post(
            "/head-manager/fbs-new/warehouse-bundles/action/",
            {
                "action": "define",
                "bundle_id": bundle.id,
                "component_id": [component.id],
                "component_quantity": [2],
            },
        )
        self.assertEqual(defined.status_code, 302)
        assembled = self.client.post(
            "/head-manager/fbs-new/warehouse-bundles/action/",
            {
                "action": "assemble",
                "bundle_id": bundle.id,
                "quantity": 2,
                "location_id": location.id,
            },
        )
        self.assertEqual(assembled.status_code, 302)
        bundle.refresh_from_db()
        component.refresh_from_db()
        self.assertTrue(bundle.is_bundle)
        self.assertEqual(bundle.stock_on_hand, 2)
        self.assertEqual(component.stock_on_hand, 1)
        self.assertEqual(WmsNewBundleComponent.objects.count(), 1)
        self.assertEqual(WmsNewBundleStock.objects.get().quantity, 2)
        self.assertEqual(WmsNewBundleOperation.objects.count(), 2)
        self.assertEqual(SKU.objects.count(), old_sku_count)

    def test_fbs_new_movement_and_history_match_donor_and_are_independent(self):
        agency = Agency.objects.create(agn_name="Партнер перемещения FBS-NEW")
        source = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="SCREEN-MOVE",
            row_no=9,
            section_no=1,
            tier_no=1,
            cell_no=1,
            location_code="SCREEN-MOVE-SOURCE",
            is_active=True,
            is_storage=True,
        )
        target = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="SCREEN-MOVE",
            row_no=9,
            section_no=1,
            tier_no=1,
            cell_no=2,
            location_code="SCREEN-MOVE-TARGET",
            is_active=True,
            is_storage=True,
        )
        product = WmsNewProduct.objects.create(
            agency=agency,
            name="Товар экрана перемещения",
            article="SCREEN-MOVE-1",
            stock_on_hand=6,
            stock_free=6,
        )
        box = WmsNewBox.objects.create(
            agency=agency,
            code="SCREEN-MOVE-BOX",
            location=source,
            location_code=source.location_code,
            zone_code=source.zone_code,
            stock_on_hand=6,
            stock_free=6,
            sku_count=1,
        )
        item = WmsNewBoxItem.objects.create(
            box=box,
            product=product,
            agency=agency,
            sku_code=product.article,
            product_name=product.name,
            qty=6,
            available_qty=6,
        )
        legacy_snapshot_count = WarehouseStockSnapshot.objects.count()

        page = self.client.get(
            "/head-manager/fbs-new/warehouse-movement/",
            {"partner": agency.id, "product": product.id, "source": source.id},
        )
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["section_data"]["kind"], "warehouse_movement")
        for label in (
            "/ Склад",
            "Переместить товар",
            "Переместить весь товар с места",
            "Партнер",
            "Товар",
            "Откуда",
            "Единицы",
            "Коробы",
            "На какой склад",
            "Куда",
            "Основной склад",
            box.code,
        ):
            self.assertContains(page, label)

        all_page = self.client.get(
            "/head-manager/fbs-new/warehouse-movement/", {"mode": "all"}
        )
        self.assertContains(all_page, source.location_code)
        self.assertContains(all_page, "Переместить все товары")

        response = self.client.post(
            "/head-manager/fbs-new/warehouse-movement/action/",
            {
                "action": "product",
                "product_id": product.id,
                "source_location_id": source.id,
                "target_location_id": target.id,
                "units": 2,
            },
        )
        self.assertEqual(response.status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.qty, 4)
        self.assertEqual(WarehouseStockSnapshot.objects.count(), legacy_snapshot_count)

        history = self.client.get("/head-manager/fbs-new/warehouse-history/")
        self.assertEqual(history.status_code, 200)
        self.assertEqual(history.context["section_data"]["kind"], "warehouse_history")
        for label in (
            "Экспорт в Excel",
            "Все партнеры",
            "Все действия",
            "Изменение остатка",
            "Поступление товара",
            "Списание",
            "Перемещение",
            "Название",
            "Дата операции",
            "Место Источник",
            "Место размещения",
            "Информация",
            "Пользователь",
            product.name,
        ):
            self.assertContains(history, label)
        export = self.client.get("/head-manager/fbs-new/warehouse-history/export/")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(export.content.startswith(b"PK"))

    def test_fbs_new_logistics_four_sections_match_donor_and_are_independent(self):
        agency = Agency.objects.create(agn_name="Партнер логистики FBS-NEW")
        order = WmsNewLogisticsOrder.objects.create(
            agency=agency,
            number="LOG-SCREEN-1",
            tracking_number="LOG-TRACK-1",
            source_type=WmsNewLogisticsOrder.SOURCE_MANUAL,
            source_label="Wildberries FBS",
            status=WmsNewLogisticsOrder.STATUS_UNMEASURED,
            item_count=2,
            is_manual=True,
        )
        WmsNewLogisticsRouteRule.objects.create(delivery_service="Wildberries FBS")
        legacy_shipping_count = ShippingOrder.objects.count()
        legacy_trip_count = LogisticsTrip.objects.count()

        orders_page = self.client.get("/head-manager/fbs-new/logistics-orders/")
        self.assertEqual(orders_page.context["section_data"]["kind"], "logistics_orders")
        for label in (
            "/ Логистика",
            "Отправления",
            "Загрузить из FBS",
            "Принять отправление",
            "Статус: все",
            "Грузоместо: все",
            "Номер отправления",
            "Габариты Ш × В × Г, мм",
            "Статус доставки",
            "Последняя проверка",
            order.number,
        ):
            self.assertContains(orders_page, label)
        accepted = self.client.post(
            "/head-manager/fbs-new/logistics-orders/create/", {"number": "LOG-MANUAL-2"}
        )
        self.assertEqual(accepted.status_code, 302)
        manual = WmsNewLogisticsOrder.objects.get(number="LOG-MANUAL-2")
        measured = self.client.post(
            f"/head-manager/fbs-new/logistics-orders/{manual.id}/action/",
            {
                "action": "measure",
                "weight_g": 700,
                "width_mm": 100,
                "height_mm": 200,
                "depth_mm": 300,
            },
        )
        self.assertEqual(measured.status_code, 302)
        manual.refresh_from_db()
        self.assertEqual(manual.status, WmsNewLogisticsOrder.STATUS_OK)

        create_package_page = self.client.get(
            "/head-manager/fbs-new/logistics-shipments/", {"create": 1}
        )
        self.assertEqual(
            create_package_page.context["section_data"]["kind"], "logistics_packages"
        )
        for label in (
            "Добавить грузоместо",
            "Транспортная компания",
            "Источник",
            "Печать ярлыков",
            "В рейс",
            "Отгрузить",
            "Название рейса",
        ):
            self.assertContains(create_package_page, label)
        package_response = self.client.post(
            "/head-manager/fbs-new/logistics-shipments/create/",
            {
                "carrier_code": WmsNewLogisticsPackage.CARRIER_GBS,
                "delivery_service": "Wildberries FBS",
                "warehouse_code": "Основной склад",
            },
        )
        self.assertEqual(package_response.status_code, 302)
        package = WmsNewLogisticsPackage.objects.get()
        linked = self.client.post(
            "/head-manager/fbs-new/logistics-shipments/action/",
            {"action": "add_orders", "package_id": package.id, "order_ids": [manual.id]},
        )
        self.assertEqual(linked.status_code, 302)

        manifest_page = self.client.get(
            "/head-manager/fbs-new/logistics-trips/", {"create": 1}
        )
        self.assertEqual(
            manifest_page.context["section_data"]["kind"], "logistics_manifests"
        )
        for label in (
            "Добавление рейса",
            "Укажите название рейса",
            "Аэропорт вылета",
            "Ожидаемая дата прилета",
            "Логистическая компания",
            "Всего товаров",
            "Общий вес",
        ):
            self.assertContains(manifest_page, label)
        created_manifest = self.client.post(
            "/head-manager/fbs-new/logistics-trips/create/",
            {"name": "LOG-MANIFEST-1", "total_weight_g": 700},
        )
        self.assertEqual(created_manifest.status_code, 302)
        self.assertTrue(WmsNewLogisticsManifest.objects.filter(name="LOG-MANIFEST-1").exists())

        settings_page = self.client.get("/head-manager/fbs-new/logistics-settings/")
        self.assertEqual(settings_page.context["section_data"]["kind"], "logistics_settings")
        for label in (
            "Настройка маршрутов",
            "Служба доставки",
            "Международная ТК",
            "Локальная ТК",
            "Активно",
            "GBS",
            "Почта России",
            "СДЭК",
        ):
            self.assertContains(settings_page, label)
        saved = self.client.post(
            "/head-manager/fbs-new/logistics-settings/save/",
            {
                "delivery_service": ["Wildberries FBS"],
                "international_carrier": ["gbs"],
                "local_carrier": ["cdek"],
                "active": ["Wildberries FBS"],
            },
        )
        self.assertEqual(saved.status_code, 302)
        rule = WmsNewLogisticsRouteRule.objects.get(delivery_service="Wildberries FBS")
        self.assertEqual(rule.international_carrier, "gbs")
        self.assertEqual(rule.local_carrier, "cdek")
        self.assertEqual(ShippingOrder.objects.count(), legacy_shipping_count)
        self.assertEqual(LogisticsTrip.objects.count(), legacy_trip_count)

    def test_fbs_new_queue_writes_only_to_isolated_tables(self):
        agency = Agency.objects.create(agn_name="Пилотный партнер")
        old_order_count = FbsOrder.objects.count()
        response = self.client.post(
            "/head-manager/fbs-new/fbs-orders/create/",
            {
                "agency_id": agency.id,
                "external_order_id": "FBN-TEST-1",
                "cutoff_date": timezone.localdate().isoformat(),
                "tracking_number": "TRACK-NEW",
            },
        )

        self.assertEqual(response.status_code, 302)
        order = WmsNewOrder.objects.get(external_order_id="FBN-TEST-1")
        self.assertTrue(order.is_manual)
        self.assertEqual(FbsOrder.objects.count(), old_order_count)
        self.assertTrue(
            WmsNewEvent.objects.filter(
                entity_type="order", entity_id=order.id, action="create_manual"
            ).exists()
        )

        order.availability_state = "ready"
        order.save(update_fields=["availability_state"])
        old_wave_count = FbsPickBatch.objects.count()
        response = self.client.post(
            "/head-manager/fbs-new/fbs-queue/action/",
            {"action": "launch_wave", "order_ids": [order.id]},
        )

        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.status, WmsNewOrder.STATUS_QUEUED)
        self.assertEqual(WmsNewWave.objects.count(), 1)
        self.assertEqual(FbsPickBatch.objects.count(), old_wave_count)

    def test_fbs_new_tasks_are_independent_and_match_donor_list(self):
        agency = Agency.objects.create(agn_name="Партнер задач FBS-NEW")
        old_task_count = Task.objects.count()
        response = self.client.post(
            "/head-manager/fbs-new/tasks/create/",
            {
                "agency_id": agency.id,
                "workflow_type": WmsNewTask.TYPE_ACCEPTANCE,
                "due_date": timezone.localdate().isoformat(),
                "title": "Проверить приемку",
                "description": "Тест отдельного контура",
                "priority": WmsNewTask.PRIORITY_HIGH,
            },
        )

        self.assertEqual(response.status_code, 302)
        task = WmsNewTask.objects.get(title="Проверить приемку")
        self.assertEqual(Task.objects.count(), old_task_count)
        self.assertTrue(task.is_manual)

        response = self.client.get("/head-manager/fbs-new/tasks-list/")
        self.assertEqual(response.status_code, 200)
        for label in (
            "Добавить задачу",
            "№ задачи",
            "Товар или название задачи",
            "Без выполненных",
            "Плановая дата",
            "Статус тарификации",
            "Подтверждение клиента",
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'href="/todo/', html=False)

    def test_fbs_new_goods_match_donor_and_write_only_new_products(self):
        agency = Agency.objects.create(agn_name="Партнер товаров FBS-NEW")
        legacy_count = SKU.objects.count()
        response = self.client.post(
            "/head-manager/fbs-new/warehouse-goods/create/",
            {
                "agency_id": agency.id,
                "name": "Новый товар",
                "article": "NEW-GOODS-1",
                "barcode": "4600000000001",
                "weight_grams": "125",
                "width_cm": "10",
                "depth_cm": "20",
                "height_cm": "30",
                "category": "Категория по умолчанию",
            },
        )

        self.assertEqual(response.status_code, 302)
        product = WmsNewProduct.objects.get(article="NEW-GOODS-1")
        self.assertEqual(SKU.objects.count(), legacy_count)
        self.assertTrue(product.is_manual)
        self.assertTrue(
            WmsNewEvent.objects.filter(
                entity_type="product", entity_id=product.id, action="create"
            ).exists()
        )

        response = self.client.get("/head-manager/fbs-new/warehouse-goods/")
        self.assertEqual(response.status_code, 200)
        for label in (
            "Добавить товар",
            "Экспорт в EXCEL",
            "Импорт из EXCEL",
            "Объединить товары",
            "ID, Название, Артикул, ШК",
            "Без нулевых остатков",
            "Только с резервами",
            "Картинка",
            "КИЗов",
            "FBO (резерв)",
            "FBS (резерв)",
            "Ожидается",
            "Заметки",
        ):
            self.assertContains(response, label)
        self.assertContains(response, "/head-manager/fbs-new/warehouse-goods/export/")
        self.assertNotContains(response, 'href="/head-manager/goods/', html=False)
        self.assertNotContains(response, 'href="/sku/', html=False)

        export = self.client.get("/head-manager/fbs-new/warehouse-goods/export/")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_fbs_new_acceptances_match_donor_and_have_no_legacy_actions(self):
        agency = Agency.objects.create(agn_name="Партнер приемки FBS-NEW")
        WmsNewAcceptance.objects.create(
            source_order_key="RCV-PAGE-1",
            agency=agency,
            task_number="618",
            title="Тестовая приемка",
            received_qty=32,
            expected_qty=32,
            acceptance_type=WmsNewAcceptance.TYPE_SCAN,
            status=WmsNewAcceptance.STATUS_DONE,
            source_created_at=timezone.now(),
            completed_at=timezone.now(),
        )

        response = self.client.get("/head-manager/fbs-new/warehouse-acceptances/")

        self.assertEqual(response.status_code, 200)
        for label in (
            "/ Склад /",
            "Приемки",
            "Дата создания",
            "Дата выполнения",
            "Все партнеры",
            "Все статусы",
            "Применить",
            "Партнер",
            "Задача",
            "Принято",
            "Тип",
            "Создана",
            "Выполнена",
            "Статус",
            "Сканирование",
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'href="/orders/receiving/', html=False)

    def test_fbs_new_marking_matches_donor_and_writes_only_new_registry(self):
        agency = Agency.objects.create(agn_name="Партнер маркировки FBS-NEW")
        product = WmsNewProduct.objects.create(
            agency=agency,
            name="Товар маркировки",
            article="MARK-PAGE-1",
            barcode="4600000000004",
        )
        WmsNewMarkingCode.objects.create(
            agency=agency,
            product=product,
            product_name=product.name,
            article=product.article,
            barcode=product.barcode,
            code="010460000000000421PAGE1",
            source=WmsNewMarkingCode.SOURCE_MANUAL,
            received_at=timezone.now(),
        )
        legacy_before = MarkingCode.objects.count()

        response = self.client.get("/head-manager/fbs-new/warehouse-marking/")

        self.assertEqual(response.status_code, 200)
        for label in (
            "Добавить КИЗы",
            "Распечатать КИЗы",
            "Экспорт в Excel",
            "Вернуть отмеченные",
            "Код маркировки",
            "Все партнеры",
            "Все товары",
            "Тип кода",
            "Поступил",
            "Выбыл",
            "Распечатан",
            "Печать",
            "Файл",
            "Обработано",
            "Добавлен",
        ):
            self.assertContains(response, label)
        self.assertContains(response, "010460000000000421PAGE1")
        self.assertNotContains(response, 'href="/marking/', html=False)

        response = self.client.post(
            "/head-manager/fbs-new/warehouse-marking/create/",
            {
                "agency_id": agency.id,
                "product_id": product.id,
                "code_type": WmsNewMarkingCode.TYPE_UNIT,
                "codes": "CODE-PAGE-NEW-1\nCODE-PAGE-NEW-2",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MarkingCode.objects.count(), legacy_before)
        self.assertTrue(WmsNewMarkingCode.objects.filter(code="CODE-PAGE-NEW-1").exists())

        export = self.client.get("/head-manager/fbs-new/warehouse-marking/export/")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(
            export["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_fbs_new_extra_fields_are_independent_and_functional(self):
        agency = Agency.objects.create(agn_name="Партнер дополнительных полей")
        product = WmsNewProduct.objects.create(
            agency=agency,
            name="Товар дополнительных полей",
            article="EXTRA-PAGE-1",
        )
        legacy_before = (
            GoodsExtraFieldDefinition.objects.count(),
            GoodsExtraFieldValue.objects.count(),
        )

        response = self.client.post(
            "/head-manager/fbs-new/warehouse-extra-fields/create/",
            {
                "name": "Срок серии",
                "code": "series-date",
                "field_type": WmsNewExtraFieldDefinition.TYPE_DATE,
                "sort_order": "20",
            },
        )
        self.assertEqual(response.status_code, 302)
        definition = WmsNewExtraFieldDefinition.objects.get(code="series-date")
        self.assertEqual(
            legacy_before,
            (
                GoodsExtraFieldDefinition.objects.count(),
                GoodsExtraFieldValue.objects.count(),
            ),
        )

        response = self.client.post(
            f"/head-manager/fbs-new/warehouse-extra-fields/{definition.id}/value/",
            {"product_id": product.id, "value": "2026-09-01"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            WmsNewExtraFieldValue.objects.filter(
                definition=definition,
                product=product,
                value="2026-09-01",
            ).exists()
        )

        response = self.client.get(
            f"/head-manager/fbs-new/warehouse-extra-fields/?field_id={definition.id}"
        )
        self.assertEqual(response.status_code, 200)
        for label in (
            "Добавить дополнительное поле",
            "Название или код",
            "Все типы",
            "Активные",
            "Название",
            "Код",
            "Тип",
            "Используется",
            "Статус",
            "Порядок",
            "Срок серии",
            "2026-09-01",
            "Сохранить значение",
            "Отключить поле",
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'href="/warehouse-goods/', html=False)

    def test_fbs_new_inventories_are_independent_and_open_full_workflow(self):
        agency = Agency.objects.create(agn_name="Партнер инвентаризации страницы")
        sku = SKU.objects.create(
            agency=agency,
            sku_code="INV-PAGE-1",
            name="Товар инвентаризации страницы",
        )
        WmsNewProduct.objects.create(
            source_sku_id=sku.id,
            agency=agency,
            article=sku.sku_code,
            name=sku.name,
            barcode="4600000000123",
            stock_on_hand=3,
            stock_free=3,
        )
        location = WarehouseLocation.objects.create(
            zone_code="FBS",
            location_code="INV-PAGE-CELL",
            display_name="Ячейка страницы инвентаризации",
            is_storage=True,
        )
        cell = FbsStorageCell.objects.create(cell_code="INV-PAGE-CELL", location=location)
        pallet = FbsPallet.objects.create(
            agency=agency,
            pallet_code="INV-PAGE-PALLET",
            cell=cell,
        )
        box = FbsBox.objects.create(
            agency=agency,
            pallet=pallet,
            box_code="INV-PAGE-BOX",
        )
        balance = FbsStockBalance.objects.create(
            agency=agency,
            box=box,
            sku_ref=sku,
            identity_key="inv-page-balance",
            sku_code=sku.sku_code,
            name=sku.name,
            barcode="4600000000123",
            qty=3,
            available_qty=3,
        )
        legacy_count = FbsInventorySession.objects.count()

        response = self.client.post(
            "/head-manager/fbs-new/warehouse-inventories/create/",
            {
                "scope_type": WmsNewInventorySession.SCOPE_BOX,
                "scope_object_id": box.id,
                "mode": WmsNewInventorySession.MODE_AUDIT,
                "scan_mode": WmsNewInventorySession.SCAN_MODE_BARCODE,
            },
        )

        self.assertEqual(response.status_code, 302)
        session = WmsNewInventorySession.objects.get()
        self.assertEqual(FbsInventorySession.objects.count(), legacy_count)
        balance.refresh_from_db()
        self.assertEqual(balance.qty, 3)

        response = self.client.get(
            f"/head-manager/fbs-new/warehouse-inventories/?inventory_id={session.id}"
        )
        self.assertEqual(response.status_code, 200)
        for label in (
            "Начать инвентаризацию",
            "Активные",
            "На пересчете",
            "На утверждении",
            "Номер, товар или место",
            "Все партнеры",
            "Все статусы",
            "Все области",
            "Контрольный пересчет",
            "Первый пересчет",
            "Отсканируйте код",
            "Принять скан",
            "Завершить пересчет",
            "INV-PAGE-BOX",
            "Товар инвентаризации страницы",
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'href="/fbs/tsd/', html=False)

    def test_fbs_new_boxes_use_new_domain_and_real_locations(self):
        agency = Agency.objects.create(agn_name="Партнер коробов страницы")
        location = WarehouseLocation.objects.create(
            zone_code="OS",
            location_code="BOX-PAGE-LOCATION",
            display_name="Место короба страницы",
            is_storage=True,
        )
        legacy_count = WarehouseContainer.objects.filter(
            container_type=WarehouseContainer.TYPE_BOX
        ).count()

        response = self.client.post(
            "/head-manager/fbs-new/warehouse-boxes/create/",
            {
                "agency_id": agency.id,
                "code": "WMS-NEW-PAGE-BOX",
                "location_id": location.id,
                "gross_weight_g": "1200",
                "width_mm": "300",
                "height_mm": "200",
                "depth_mm": "400",
            },
        )

        self.assertEqual(response.status_code, 302)
        box = WmsNewBox.objects.get(code="WMS-NEW-PAGE-BOX")
        self.assertTrue(box.is_manual)
        self.assertEqual(
            WarehouseContainer.objects.filter(container_type=WarehouseContainer.TYPE_BOX).count(),
            legacy_count,
        )

        response = self.client.get(
            f"/head-manager/fbs-new/warehouse-boxes/?box_id={box.id}"
        )
        self.assertEqual(response.status_code, 200)
        for label in (
            "Добавить короб",
            "Всего коробов",
            "С товаром",
            "Единиц товара",
            "Короб, товар, ШК или место",
            "Все зоны",
            "Только с товаром",
            "Фактические короба Fullbox в WMS NEW",
            "WMS-NEW-PAGE-BOX",
            "Место короба страницы",
            "Сохранить в WMS NEW",
            "Архивировать",
            "Короб пуст",
        ):
            self.assertContains(response, label)
        self.assertNotContains(response, 'href="/sklad/journal/', html=False)

    def test_fbs_new_pilot_access_keeps_primary_role_and_old_cabinet(self):
        pilot_user = get_user_model().objects.create_user(username="pilot_storekeeper", password="pwd")
        pilot = Employee.objects.create(
            full_name="Пилотный кладовщик",
            role="storekeeper",
            user=pilot_user,
            is_active=True,
        )
        self.client.force_login(pilot_user)

        self.assertEqual(self.client.get("/head-manager/fbs-new/").status_code, 403)

        pilot.access_roles = ["fbs_new"]
        pilot.save(update_fields=["access_roles"])
        response = self.client.get("/head-manager/fbs-new/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'href="/sklad/"', html=False)
        self.assertContains(response, "Ваш основной кабинет и прежние рабочие процессы сохранены")
        self.assertEqual(self.client.get("/head-manager/").status_code, 403)

    def test_head_manager_can_toggle_fbs_new_pilot_with_audit(self):
        target_user = get_user_model().objects.create_user(username="pilot_manager", password="pwd")
        target = Employee.objects.create(
            full_name="Пилотный менеджер",
            role="manager",
            access_roles=["accountant"],
            user=target_user,
            is_active=True,
        )

        response = self.client.post(
            f"/head-manager/fbs-new/pilot/{target.id}/",
            {"action": "enable"},
        )

        self.assertRedirects(
            response,
            "/head-manager/fbs-new/settings/#pilot-access",
            fetch_redirect_response=False,
        )
        target.refresh_from_db()
        self.assertEqual(target.role, "manager")
        self.assertEqual(target.normalized_access_roles(), ("accountant", "fbs_new"))
        event = AuditEntry.objects.get(journal__code="fbs_new_pilot")
        self.assertEqual(event.user_id, self.user.id)
        self.assertEqual(event.snapshot["employee_id"], target.id)
        self.assertEqual(event.snapshot["pilot_action"], "enable")

        response = self.client.post(
            f"/head-manager/fbs-new/pilot/{target.id}/",
            {"action": "disable"},
        )
        self.assertEqual(response.status_code, 302)
        target.refresh_from_db()
        self.assertEqual(target.role, "manager")
        self.assertEqual(target.normalized_access_roles(), ("accountant",))
        self.assertEqual(AuditEntry.objects.filter(journal__code="fbs_new_pilot").count(), 2)

    def test_dashboard_task_panel_shows_logistics_tab_for_head_manager(self):
        agency = Agency.objects.create(agn_name="Клиент рейса ГМ")
        logistician = Employee.objects.create(full_name="Логистов Сергей", role="logistician", is_active=True)
        storekeeper = Employee.objects.create(full_name="Кладовщиков Алексей", role="storekeeper", is_active=True)
        order = ShippingOrder.objects.create(number="SO-000150", agency=agency)
        trip = LogisticsTrip.objects.create(
            number="7_RS",
            status=LogisticsTrip.STATUS_LOADING,
            assigned_logistician=logistician,
        )
        Task.objects.create(
            title="Рейс №7_RS",
            route=f"/logistics/trips/{trip.pk}/",
            status="backlog",
            assigned_to=storekeeper,
        )
        Task.objects.create(
            title="Отгрузка",
            route=f"/shipping/{order.pk}/",
            status="backlog",
            assigned_to=logistician,
        )

        response = self.client.get("/head-manager/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'value="logistics"', html=False)
        self.assertContains(response, "Рейсы")
        self.assertContains(response, "Рейс №7_RS")
        self.assertContains(response, "Заявка на отгрузку №150_OTG")

    def test_create_own_company_generates_short_name(self):
        response = self.client.post(
            "/head-manager/own-companies/new/",
            data={
                "name": 'Общество с ограниченной ответственностью "ФуллБокс Логистика"',
                "inn": "7701234567",
                "kpp": "770101001",
                "ogrn": "1234567890123",
                "address": "Москва, тестовый адрес",
                "postal_address": "Москва, а/я 1",
                "phone": "+7 900 000-00-00",
                "email": "test@example.com",
                "director_name": "Петров П.П.",
                "director_basis": "Устав",
                "bank_name": "АО Тест Банк",
                "bank_bik": "044525000",
                "settlement_account": "40702810900000000001",
                "correspondent_account": "30101810400000000000",
                "bank_address": "Москва, банковский адрес",
                "edo_operator": "СБИС",
                "edo_id": "2BM-TEST",
                "is_default": "on",
                "is_active": "on",
                "comment": "Основная компания",
            },
        )

        self.assertRedirects(response, "/head-manager/own-companies/")
        company = OwnCompany.objects.get()
        self.assertEqual(company.short_name, 'ООО "ФуллБокс Логистика"')
        self.assertTrue(company.is_default)
        self.assertEqual(company.postal_address, "Москва, а/я 1")
        self.assertEqual(company.director_name, "Петров П.П.")
        self.assertEqual(company.bank_name, "АО Тест Банк")
        self.assertEqual(company.edo_operator, "СБИС")

    def test_create_carrier_generates_short_name(self):
        response = self.client.post(
            "/head-manager/carriers/new/",
            data={
                "name": "Индивидуальный предприниматель Иванов Иван Иванович",
                "inn": "770123456789",
                "kpp": "",
                "ogrn": "123456789012345",
                "address": "Тула, тестовый адрес",
                "postal_address": "Тула, а/я 1",
                "phone": "+7 900 111-22-33",
                "email": "carrier@example.com",
                "contact_person": "Иванов И.И.",
                "bank_name": "АО Банк",
                "bank_bik": "044525999",
                "settlement_account": "40802810900000000001",
                "correspondent_account": "30101810400000000001",
                "bank_address": "Москва",
                "is_active": "on",
                "comment": "Пул внешних перевозчиков",
            },
        )

        self.assertRedirects(response, "/head-manager/carriers/")
        carrier = Carrier.objects.get()
        self.assertEqual(carrier.short_name, "ИП Иванов Иван Иванович")
        self.assertEqual(carrier.bank_name, "АО Банк")
        self.assertEqual(carrier.postal_address, "Тула, а/я 1")

    def test_own_companies_list_has_fullscreen_company_columns(self):
        OwnCompany.objects.create(
            name='Общество с ограниченной ответственностью "ФуллБокс"',
            inn="5001149130",
            kpp="503101001",
            ogrn="1225000116604",
            address="Юридический адрес",
            postal_address="Почтовый адрес",
            bank_name="Банк ВТБ",
            bank_bik="044525411",
            settlement_account="40702810700250001553",
            correspondent_account="30101810145250000411",
            email="info@fullbox.ru",
            phone="+7 977 808-83-86",
            edo_operator="Контур Диадок",
            edo_id="2BM-test",
            is_default=True,
            is_active=True,
        )

        response = self.client.get("/head-manager/own-companies/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="wrap fullscreen"', html=False)
        self.assertContains(response, "Юр. адрес")
        self.assertContains(response, "Почтовый адрес")
        self.assertContains(response, "Банковские реквизиты")

    def test_reference_services_build_expected_context(self):
        list_context = build_reference_list_context(
            title="Компании",
            subtitle="Описание",
            create_url="/new/",
            edit_base_url="/edit",
            back_url="/back/",
            type_label="компания",
            directory_mode="own_companies",
        )
        form_context = build_reference_form_context(
            title="Новая компания",
            subtitle="Форма",
            back_url="/back/",
            submit_label="Сохранить",
        )

        self.assertEqual(list_context["directory_mode"], "own_companies")
        self.assertEqual(list_context["create_url"], "/new/")
        self.assertEqual(form_context["submit_label"], "Сохранить")

    def test_carriers_list_has_fullscreen_carrier_columns(self):
        Carrier.objects.create(
            name="Индивидуальный предприниматель Касаев Абдурашид Метханович",
            inn="052903630023",
            address="Юридический адрес перевозчика",
            postal_address="Почтовый адрес перевозчика",
            bank_name='ООО "Банк Точка"',
            bank_bik="044525104",
            settlement_account="40802810201500324670",
            phone="+7 (927) 157-22-42",
            email="kasaev1968@mail.ru",
            is_active=True,
        )

        response = self.client.get("/head-manager/carriers/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="wrap fullscreen"', html=False)
        self.assertContains(response, "Юр. адрес")
        self.assertContains(response, "Почтовый адрес")
        self.assertContains(response, "Банковские реквизиты")


class MarketplaceWarehouseDirectoryTests(TestCase):
    def test_normalizes_legacy_line_into_structured_row(self):
        rows = _normalize_marketplace_warehouse_rows(["Транзитный / ППП · Гольёво — Московская область, Красногорск"])

        self.assertEqual(
            rows,
            [
                {
                    "type": "Транзитный",
                    "name": "Транзитный / ППП · Гольёво",
                    "address": "Московская область, Красногорск",
                }
            ],
        )

    def test_shipping_loader_reads_structured_json(self):
        from shipping import marketplace_warehouses as shipping_marketplaces

        with TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "marketplace_warehouses.json"
            path.write_text(
                json.dumps(
                    {
                        "wb": [
                            {"type": "Транзитный", "name": "Гольёво", "address": "МО, Красногорск"},
                            {"type": "Обычный", "name": "Коледино", "address": "МО, Подольск"},
                        ],
                        "ozon": [],
                        "yandex": [],
                        "sber": [],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.object(shipping_marketplaces, "marketplace_warehouses_path", return_value=path):
                data = shipping_marketplaces.load_marketplace_warehouses()

        self.assertEqual(
            data["wb"],
            [
                "Транзитный · Гольёво — МО, Красногорск",
                "Обычный · Коледино — МО, Подольск",
            ],
        )

    def test_marketplace_warehouse_context_reads_session_flags(self):
        user = get_user_model().objects.create_user(username="hm_service_ctx", password="pwd")
        Employee.objects.create(full_name="Главный менеджер", role="head_manager", user=user, is_active=True)
        request = RequestFactory().get("/head-manager/marketplace-warehouses/")
        request.user = user
        session = self.client.session
        session["marketplace_sync_info"] = "sync ok"
        session["marketplace_sync_errors"] = ["warn"]
        session["marketplace_sync_agency"] = "Agency"
        session.save()
        request.session = session

        context = build_marketplace_warehouses_context(request=request, saved=True, error="")

        self.assertTrue(context["saved"])
        self.assertEqual(context["sync_info"], "sync ok")
        self.assertEqual(context["sync_errors"], ["warn"])
        self.assertEqual(context["sync_agency"], "Agency")


class WarehouseEmployeeMetricAggregationTests(TestCase):
    def setUp(self):
        self.day = timezone.localdate()
        self.occurred_at = timezone.now()
        self.filters = {
            "date_from": self.day.isoformat(),
            "date_to": self.day.isoformat(),
            "employee": "",
        }
        self.agency = Agency.objects.create(agn_name="Клиент отчета по сотрудникам")

    def _user(self, suffix: str):
        user = get_user_model().objects.create_user(username=f"warehouse_report_{suffix}")
        Employee.objects.create(
            full_name=f"Сотрудник отчета {suffix}",
            role="storekeeper",
            user=user,
            is_active=True,
        )
        return user

    @staticmethod
    def _values(points):
        return {metric: value for _day, _user, metric, value in points}

    def test_fbs_metrics_reuse_existing_report_rows(self):
        user = self._user("fbs")
        report_row = {
            "user_id": user.pk,
            "date_value": self.day,
            "pick_waves": 2,
            "pick_orders": 123,
            "pick_units": 123,
            "avg_pick_minutes": 9.2,
            "verified_waves": 0,
            "verified_orders": 0,
            "verified_units": 0,
            "shipments": 0,
            "shipment_boxes": 0,
            "shipment_orders": 0,
            "shipment_units": 0,
        }
        with patch(
            "head_manager.web_ui._head_manager_fbs_employee_report_rows",
            return_value=([report_row], {}),
        ) as report_rows_mock:
            values = self._values(fbs_metric_points(self.filters))

        report_rows_mock.assert_called_once_with(self.filters)
        self.assertEqual(values["pick_waves"], 2)
        self.assertEqual(values["pick_orders"], 123)
        self.assertEqual(values["pick_units"], 123)

    def test_palletization_counts_only_completed_events(self):
        user = self._user("palletization")
        for event_type in ("palletization_started", "palletization_completed"):
            WarehouseEvent.objects.create(
                agency=self.agency,
                event_type=event_type,
                qty=4562,
                performed_by=user,
                occurred_at=self.occurred_at,
            )

        values = self._values(warehouse_palletization_metric_points(self.filters))

        self.assertEqual(values["palletization_positions"], 1)
        self.assertEqual(values["palletization_units"], 4562)

    def test_query_count_does_not_grow_with_employee_count(self):
        first_user = self._user("query_01")
        WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="placement_completed",
            qty=1,
            performed_by=first_user,
            occurred_at=self.occurred_at,
        )
        with self.assertNumQueries(2):
            first_points = warehouse_placement_metric_points(self.filters)
        self.assertEqual(len(first_points), 2)

        for index in range(2, 21):
            user = self._user(f"query_{index:02d}")
            WarehouseEvent.objects.create(
                agency=self.agency,
                event_type="placement_completed",
                qty=index,
                performed_by=user,
                occurred_at=self.occurred_at,
            )
        with self.assertNumQueries(2):
            all_points = warehouse_placement_metric_points(self.filters)
        self.assertEqual(len(all_points), 40)

    def test_employee_filter_uses_employee_full_name(self):
        matching_user = self._user("Иванов Иван")
        other_user = self._user("Петров Петр")
        for user in (matching_user, other_user):
            WarehouseEvent.objects.create(
                agency=self.agency,
                event_type="placement_completed",
                qty=1,
                performed_by=user,
                occurred_at=self.occurred_at,
            )
        filters = {**self.filters, "employee": "Иванов"}

        points = warehouse_placement_metric_points(filters)

        self.assertEqual(
            {user.pk for _day, user, _metric, _value in points},
            {matching_user.pk},
        )

    def test_work_window_combines_first_and_last_actions_across_contours(self):
        user = self._user("activity_window")
        current_timezone = timezone.get_current_timezone()
        first_action = timezone.make_aware(
            datetime.combine(self.day, datetime.min.time()).replace(hour=8, minute=15),
            current_timezone,
        )
        last_action = first_action + timedelta(hours=8, minutes=30)
        points = [
            (self.day, user.pk, "receiving_cz", first_action, first_action),
            (self.day, user.pk, "placement", last_action, last_action),
        ]

        with self.assertNumQueries(1):
            windows = _work_windows_from_activity_points(points)

        self.assertEqual(len(windows), 1)
        self.assertEqual(windows[0]["first_action"], first_action)
        self.assertEqual(windows[0]["last_action"], last_action)
        self.assertEqual(windows[0]["activity_span_seconds"], 8 * 3600 + 30 * 60)
        self.assertEqual(windows[0]["activity_span_hours"], 8.5)
        self.assertEqual(windows[0]["contours"], ("placement", "receiving_cz"))


class WarehouseEmployeeReportPresentationTests(TestCase):
    def setUp(self):
        self.day = timezone.localdate()
        self.filters = {
            "date_from": self.day.isoformat(),
            "date_to": self.day.isoformat(),
            "employee": "",
            "role": "",
        }

    def _user(self, suffix: str, role: str):
        user = get_user_model().objects.create_user(
            username=f"warehouse_presentation_{suffix}",
            password="pwd",
        )
        Employee.objects.create(
            full_name=f"Сотрудник презентации {suffix}",
            role=role,
            user=user,
            is_active=True,
        )
        return user

    def test_rows_merge_metrics_and_activity_window(self):
        user = self._user("storekeeper", "storekeeper")
        first_action = timezone.now().replace(hour=8, minute=10, second=0, microsecond=0)
        last_action = first_action + timedelta(hours=7, minutes=25)
        metric_points = [
            (self.day, user, "receiving_cz_units", 12),
            (self.day, user, "placement_units", 30),
            (self.day, user, "pick_units", 4),
        ]
        windows = [
            {
                "date_value": self.day,
                "user": user,
                "first_action": first_action,
                "last_action": last_action,
                "activity_span_seconds": 7 * 3600 + 25 * 60,
                "contours": ("placement", "receiving_cz"),
            }
        ]
        with (
            patch(
                "head_manager.web_ui.warehouse_employee_metric_points",
                return_value=metric_points,
            ),
            patch(
                "head_manager.web_ui.warehouse_employee_work_windows",
                return_value=windows,
            ),
        ):
            rows, summary = _head_manager_warehouse_employee_report_rows(self.filters)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["receiving_cz_units"], 12)
        self.assertEqual(rows[0]["placement_units"], 30)
        self.assertEqual(rows[0]["pick_units"], 4)
        self.assertEqual(rows[0]["activity_span_label"], "7 ч 25 мин")
        self.assertEqual(rows[0]["contours_label"], "размещение, приёмка ЧЗ")
        self.assertEqual(summary["employees"], 1)

    def test_role_filter_keeps_only_selected_role(self):
        storekeeper = self._user("role_storekeeper", "storekeeper")
        picker = self._user("role_picker", "picker")
        filters = {**self.filters, "role": "storekeeper"}
        with (
            patch(
                "head_manager.web_ui.warehouse_employee_metric_points",
                return_value=[
                    (self.day, storekeeper, "placement_units", 2),
                    (self.day, picker, "pick_units", 3),
                ],
            ),
            patch(
                "head_manager.web_ui.warehouse_employee_work_windows",
                return_value=[],
            ),
        ):
            rows, _summary = _head_manager_warehouse_employee_report_rows(filters)

        self.assertEqual([row["user_id"] for row in rows], [storekeeper.pk])

    def test_processing_employee_uses_journal_source(self):
        packer = self._user("processing_source", "packer")
        filters = {
            **self.filters,
            "employee": packer.employee_profile.full_name,
            "role": "packer",
        }
        with (
            patch("head_manager.web_ui.warehouse_employee_metric_points", return_value=[]),
            patch("head_manager.web_ui.warehouse_employee_work_windows", return_value=[]),
        ):
            rows, summary = _head_manager_warehouse_employee_report_rows(filters)

        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]["source_missing"])
        self.assertEqual(rows[0]["source_status"], "Есть данные")
        self.assertEqual(rows[0]["receiving_cz_units"], 0)
        self.assertEqual(summary["no_source_employees"], 0)

    def test_processing_employee_without_user_is_in_report(self):
        employee = Employee.objects.create(
            full_name="Упаковщица без логина",
            role="packer",
            user=None,
            is_active=True,
        )
        filters = {
            **self.filters,
            "employee": employee.full_name,
            "role": "packer",
        }
        with (
            patch("head_manager.web_ui.warehouse_employee_metric_points", return_value=[]),
            patch("head_manager.web_ui.warehouse_employee_work_windows", return_value=[]),
        ):
            rows, summary = _head_manager_warehouse_employee_report_rows(filters)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["employee"], employee.full_name)
        self.assertEqual(rows[0]["identity_key"], f"employee:{employee.pk}")
        self.assertEqual(rows[0]["processing_operations"], 0)
        self.assertEqual(summary["employees"], 1)

    def test_workbook_contains_same_metric_values_and_summary_sheet(self):
        user = self._user("excel", "storekeeper")
        with (
            patch(
                "head_manager.web_ui.warehouse_employee_metric_points",
                return_value=[(self.day, user, "receiving_cz_units", 17)],
            ),
            patch("head_manager.web_ui.warehouse_employee_work_windows", return_value=[]),
        ):
            rows, summary = _head_manager_warehouse_employee_report_rows(self.filters)

        workbook = _head_manager_warehouse_employee_report_workbook(rows, summary)
        sheet = workbook["Работа сотрудников"]
        headers = [cell.value for cell in sheet[1]]
        receiving_column = headers.index("Принято ЧЗ") + 1

        self.assertEqual(sheet.cell(row=2, column=receiving_column).value, 17)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "F8B800")
        self.assertIn("Сводка", workbook.sheetnames)

    def test_report_access_is_limited_to_management_roles(self):
        summary = {
            "employees": 0,
            "days": 0,
            "rows": 0,
            "no_source_employees": 0,
            "receiving_cz_units": 0,
            "pick_units": 0,
            "verified_units": 0,
            "shipment_units": 0,
            "shipping_units": 0,
        }
        route = "/head-manager/reports/employees/warehouse-productivity/"
        with patch(
            "head_manager.web_ui._head_manager_warehouse_employee_report_rows",
            return_value=([], summary),
        ):
            for role in ("head_manager", "director", "admin"):
                user = self._user(f"allowed_{role}", role)
                self.client.force_login(user)
                self.assertEqual(self.client.get(route).status_code, 200, role)
                self.client.logout()

            for role in ("storekeeper", "picker", "fbs_controller"):
                user = self._user(f"denied_{role}", role)
                self.client.force_login(user)
                self.assertEqual(self.client.get(route).status_code, 403, role)
                self.client.logout()

            client_user = get_user_model().objects.create_user(
                username="warehouse_presentation_client",
                password="pwd",
            )
            self.client.force_login(client_user)
            self.assertEqual(self.client.get(route).status_code, 403)
