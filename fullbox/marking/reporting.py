from __future__ import annotations

from datetime import datetime, time
from io import BytesIO
from django.core.paginator import Paginator
from django.db.models import Count, Q, QuerySet
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from sku.models import Agency

from .codes import marking_code_identity, normalize_marking_code
from .models import MarkingCode


REPORT_STATUS_CHOICES = (
    ("", "Все состояния"),
    ("uploaded", "Загружены из файла"),
    ("free", "Свободны"),
    ("queued", "В очереди печати"),
    ("printed", "Распечатаны, не использованы"),
    ("used", "Использованы"),
)
REPORT_SOURCE_CHOICES = (
    ("", "Все источники"),
    ("import", "Импорт из файла"),
    ("scan", "Сканирование"),
)
REPORT_ORDER_TYPE_CHOICES = (("", "Все процессы"), *MarkingCode.ORDER_TYPE_CHOICES)


def _clean(value) -> str:
    return str(value or "").strip()


def _aware_day(value: str, *, end: bool = False):
    try:
        parsed = datetime.strptime(_clean(value), "%Y-%m-%d").date()
    except ValueError:
        return None
    point = datetime.combine(parsed, time.max if end else time.min)
    return timezone.make_aware(point, timezone.get_current_timezone())


def marking_report_filters(request) -> dict[str, str]:
    return {
        key: _clean(request.GET.get(key))
        for key in (
            "status",
            "source",
            "agency",
            "order_type",
            "order_id",
            "sku",
            "barcode",
            "size",
            "box",
            "operator",
            "q",
            "date_from",
            "date_to",
        )
    }


def _apply_common_filters(queryset: QuerySet, filters: dict[str, str]) -> QuerySet:
    agency_id = filters["agency"]
    if agency_id.isdigit():
        queryset = queryset.filter(agency_id=int(agency_id))
    if filters["source"] in {"import", "scan"}:
        queryset = queryset.filter(source=filters["source"])
    if filters["order_type"] in dict(MarkingCode.ORDER_TYPE_CHOICES):
        queryset = queryset.filter(order_type=filters["order_type"])
    if filters["order_id"]:
        queryset = queryset.filter(order_id__icontains=filters["order_id"])
    if filters["sku"]:
        queryset = queryset.filter(
            Q(sku_code__icontains=filters["sku"])
            | Q(sku__name__icontains=filters["sku"])
        )
    if filters["barcode"]:
        queryset = queryset.filter(barcode__icontains=filters["barcode"])
    if filters["size"]:
        queryset = queryset.filter(size__icontains=filters["size"])
    if filters["box"]:
        queryset = queryset.filter(box_barcode__icontains=filters["box"])
    if filters["operator"]:
        operator = filters["operator"]
        queryset = queryset.filter(
            Q(created_by__username__icontains=operator)
            | Q(created_by__employee_profile__full_name__icontains=operator)
            | Q(printed_by__username__icontains=operator)
            | Q(printed_by__employee_profile__full_name__icontains=operator)
            | Q(used_by__username__icontains=operator)
            | Q(used_by__employee_profile__full_name__icontains=operator)
        )
    date_from = _aware_day(filters["date_from"])
    date_to = _aware_day(filters["date_to"], end=True)
    if date_from:
        queryset = queryset.filter(created_at__gte=date_from)
    if date_to:
        queryset = queryset.filter(created_at__lte=date_to)
    if filters["q"]:
        query = normalize_marking_code(filters["q"])
        identity = marking_code_identity(query)
        queryset = queryset.filter(
            Q(code__icontains=query)
            | Q(identity_key__icontains=identity)
            | Q(order_id__icontains=query)
            | Q(sku_code__icontains=query)
            | Q(sku__name__icontains=query)
            | Q(barcode__icontains=query)
            | Q(box_barcode__icontains=query)
            | Q(agency__agn_name__icontains=query)
            | Q(agency__short_name__icontains=query)
        )
    return queryset


def _apply_status_filter(queryset: QuerySet, status: str) -> QuerySet:
    if status == "uploaded":
        return queryset.filter(source="import")
    if status == "free":
        return queryset.filter(
            used_at__isnull=True,
            printed_at__isnull=True,
            print_job_id__isnull=True,
        )
    if status == "queued":
        return queryset.filter(
            used_at__isnull=True,
            printed_at__isnull=True,
            print_job_id__isnull=False,
        )
    if status == "printed":
        return queryset.filter(used_at__isnull=True, printed_at__isnull=False)
    if status == "used":
        return queryset.filter(used_at__isnull=False)
    return queryset


def marking_report_querysets(request):
    filters = marking_report_filters(request)
    base = (
        MarkingCode.objects.select_related(
            "agency",
            "sku",
            "created_by",
            "created_by__employee_profile",
            "printed_by",
            "printed_by__employee_profile",
            "used_by",
            "used_by__employee_profile",
        )
        .order_by("-created_at", "-id")
    )
    base = _apply_common_filters(base, filters)
    rows = _apply_status_filter(base, filters["status"])
    return filters, base, rows


def marking_report_metrics(queryset: QuerySet) -> dict[str, int]:
    return queryset.aggregate(
        total=Count("id"),
        uploaded=Count("id", filter=Q(source="import")),
        free=Count(
            "id",
            filter=Q(
                used_at__isnull=True,
                printed_at__isnull=True,
                print_job_id__isnull=True,
            ),
        ),
        queued=Count(
            "id",
            filter=Q(
                used_at__isnull=True,
                printed_at__isnull=True,
                print_job_id__isnull=False,
            ),
        ),
        printed=Count(
            "id",
            filter=Q(used_at__isnull=True, printed_at__isnull=False),
        ),
        used=Count("id", filter=Q(used_at__isnull=False)),
    )


def marking_status(code: MarkingCode) -> tuple[str, str]:
    if code.used_at:
        return "used", "Использован"
    if code.printed_at:
        return "printed", "Распечатан"
    if code.print_job_id:
        return "queued", "В очереди"
    return "free", "Свободен"


def _user_label(user) -> str:
    if not user:
        return "—"
    employee = getattr(user, "employee_profile", None)
    if employee and _clean(employee.full_name):
        return _clean(employee.full_name)
    return _clean(user.get_full_name()) or _clean(user.username) or "—"


def _local_datetime(value) -> str:
    if not value:
        return "—"
    return timezone.localtime(value).strftime("%d.%m.%Y %H:%M")


def _mask_code(value: str) -> str:
    identity = marking_code_identity(value)
    if len(identity) <= 24:
        return identity
    return f"{identity[:18]}…{identity[-6:]}"


def marking_report_row(code: MarkingCode) -> dict:
    status_code, status_label = marking_status(code)
    event_user = code.used_by if code.used_at else code.printed_by if code.printed_at else code.created_by
    event_at = code.used_at or code.printed_at or code.print_reserved_at or code.created_at
    return {
        "id": code.id,
        "code_masked": _mask_code(code.code),
        "client": _clean(getattr(code.agency, "short_name", "")) or _clean(getattr(code.agency, "agn_name", "")) or "—",
        "order_type": code.get_order_type_display(),
        "order_id": code.order_id or "—",
        "sku_code": code.sku_code or "—",
        "product_name": _clean(getattr(code.sku, "name", "")) or "—",
        "size": code.size or "—",
        "barcode": code.barcode or "—",
        "box_barcode": code.box_barcode or "—",
        "source": code.get_source_display(),
        "status_code": status_code,
        "status_label": status_label,
        "created_at": _local_datetime(code.created_at),
        "printed_at": _local_datetime(code.printed_at),
        "used_at": _local_datetime(code.used_at),
        "operator": _user_label(event_user),
        "event_at": _local_datetime(event_at),
    }


def _page_url(request, page_number: int) -> str:
    params = request.GET.copy()
    params.pop("export", None)
    params["page"] = str(page_number)
    return f"{request.path}?{params.urlencode()}"


def marking_report_context(request, *, role: str) -> dict:
    filters, base_queryset, rows_queryset = marking_report_querysets(request)
    page_obj = Paginator(rows_queryset, 100).get_page(request.GET.get("page") or 1)
    rows = [marking_report_row(code) for code in page_obj.object_list]
    export_params = request.GET.copy()
    export_params.pop("page", None)
    export_params["export"] = "xlsx"
    role_context = {
        "processing_head": ("Руководитель обработки", "/processing-head/"),
        "head_manager": ("Начальник склада", "/head-manager/"),
        "manager": ("Менеджер", "/team-manager/"),
        "director": ("Директор", "/cabinet/director/"),
        "admin": ("Администратор", "/admin/"),
    }
    role_label, home_url = role_context.get(role, ("Сотрудник", "/"))
    agencies = Agency.objects.filter(marking_codes__isnull=False).order_by("agn_name", "id").distinct()
    return {
        "report_role": role,
        "report_role_label": role_label,
        "report_home_url": home_url,
        "filters": filters,
        "status_choices": REPORT_STATUS_CHOICES,
        "source_choices": REPORT_SOURCE_CHOICES,
        "order_type_choices": REPORT_ORDER_TYPE_CHOICES,
        "agencies": agencies,
        "metrics": marking_report_metrics(base_queryset),
        "rows": rows,
        "rows_total": page_obj.paginator.count,
        "page_obj": page_obj,
        "previous_url": _page_url(request, page_obj.previous_page_number()) if page_obj.has_previous() else "",
        "next_url": _page_url(request, page_obj.next_page_number()) if page_obj.has_next() else "",
        "export_url": f"{request.path}?{export_params.urlencode()}",
    }


def _excel_value(value):
    text = str(value or "")
    return f"'{text}" if text[:1] in {"=", "+", "-", "@"} else text


def marking_report_export_response(request) -> HttpResponse:
    _filters, _base_queryset, rows_queryset = marking_report_querysets(request)
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "Отчет ЧЗ"
    headers = [
        "ID",
        "Код ЧЗ",
        "Состояние",
        "Клиент",
        "Процесс",
        "Заявка",
        "SKU",
        "Товар",
        "Размер",
        "ШК",
        "Короб",
        "Источник",
        "Загружен",
        "Распечатан",
        "Использован",
        "Ответственный",
    ]
    sheet.append(headers)
    fill = PatternFill(fill_type="solid", fgColor="F89000")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    sheet.freeze_panes = "A2"
    for code in rows_queryset.iterator(chunk_size=1000):
        row = marking_report_row(code)
        sheet.append(
            [
                code.id,
                _excel_value(normalize_marking_code(code.code).replace("\x1d", "<GS>")),
                row["status_label"],
                row["client"],
                row["order_type"],
                row["order_id"],
                row["sku_code"],
                row["product_name"],
                row["size"],
                _excel_value(row["barcode"]),
                _excel_value(row["box_barcode"]),
                row["source"],
                row["created_at"],
                row["printed_at"],
                row["used_at"],
                row["operator"],
            ]
        )
    widths = [10, 58, 20, 28, 18, 18, 20, 36, 12, 22, 24, 18, 20, 20, 20, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.auto_filter.ref = f"A1:P{max(sheet.max_row, 1)}"

    summary = workbook.create_sheet("Сводка")
    metrics = marking_report_metrics(_base_queryset)
    for label, key in (
        ("Всего кодов", "total"),
        ("Загружено из файлов", "uploaded"),
        ("Свободно", "free"),
        ("В очереди печати", "queued"),
        ("Распечатано, не использовано", "printed"),
        ("Использовано", "used"),
    ):
        summary.append([label, metrics[key]])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 16

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="marking_report_{timezone.localdate().isoformat()}.xlsx"'
    )
    return response
