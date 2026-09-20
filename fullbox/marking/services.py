from __future__ import annotations

import base64
import binascii
import hashlib
import logging
import os
import re
from datetime import datetime, timedelta
from pathlib import Path

import requests
from django.core.cache import cache
from django.core import signing
from django.db import IntegrityError, transaction
from django.db.models import Count, Q
from django.http import HttpResponseBadRequest, JsonResponse
from django.utils import timezone
from openpyxl import load_workbook

from audit.models import AuditEntry, AuditJournal, OrderAuditEntry
from fullbox.order_numbers import build_public_order_number, format_order_number
from labels.utils import (
    get_effective_label_template,
    load_processing_param_template_bindings,
    resolve_processing_param_template_key,
)
from orders.title_truth import build_order_runtime_payload
from sku.models import SKUBarcode

from .codes import (
    MarkingCodeFormatError,
    marking_code_identity,
    marking_code_variants,
    normalize_marking_code,
    validate_import_marking_code,
)
from .integrity import (
    marking_conflict_message,
    occupied_marking_identities,
    printed_marking_storage_conflict,
)
from .models import MarkingCode
from .utils import extract_processing_items


logger = logging.getLogger(__name__)

TRUE_API_DEFAULT_BASE_URL = "https://markirovka.crpt.ru/api/v3/true-api"
TRUE_API_STATUS_LABELS = {
    "INTRODUCED": "В обороте",
    "RETIRED": "Выбыл из оборота",
    "WRITTEN_OFF": "Списан",
    "APPLIED": "Нанесен, но не введен в оборот",
    "EMITTED": "Эмитирован, но не введен в оборот",
    "APPLIED_NOT_PAID": "Нанесен, но не оплачен",
}
DUPLICATE_TOKEN_SALT = "marking.honest-sign-duplicate"
DUPLICATE_TOKEN_MAX_AGE_SECONDS = 10 * 60
DUPLICATE_MAX_LABEL_BASE64_SIZE = 2_500_000
FREE_PRINT_MAX_BATCH = 50
FREE_PRINT_MAX_COPIES = 100
FREE_PRINT_MAX_LABELS_PER_OPERATION = 100
FREE_PRINT_MAX_TOTAL_BASE64_SIZE = 25_000_000
FREE_PRINT_PROCESSING_PARAM_KEY = "free_marking_print"
FREE_PRINT_CONFIRMED_PARAM_KEY = "free_marking_confirmed"
FREE_PRINT_RETURNED_PARAM_KEY = "free_marking_returned"
FREE_PRINT_RESERVATION_PARAM_KEY = "free_marking_reservation"
FREE_PRINT_RESERVATION_TTL = timedelta(hours=2)


class TrueApiError(Exception):
    error_code = "true_api_error"
    http_status = 503


class TrueApiNotConfigured(TrueApiError):
    error_code = "not_configured"


class TrueApiUnavailable(TrueApiError):
    error_code = "unavailable"


class InvalidHonestSignCode(TrueApiError):
    error_code = "invalid_code"
    http_status = 400


def _strip_aim_prefix(value: str) -> str:
    value = value.strip("\r\n\t ")
    if re.match(r"^\][A-Za-z0-9]{2}", value):
        return value[3:]
    return value


def _normalize_gs_tokens(value: str) -> str:
    code = _strip_aim_prefix(str(value or ""))
    code = re.sub(r"^(?:\{FNC1\}|\^FNC1)", "", code, flags=re.IGNORECASE)
    return re.sub(
        r"(?:\\u001d|\\x001d|\\x1d|_x001d_|\{FNC1\}|\^FNC1)",
        "\x1d",
        code,
        flags=re.IGNORECASE,
    )


def normalize_honest_sign_code(raw_code: str) -> str:
    """Return a True API CIS without changing the scanned warehouse state."""
    code = _normalize_gs_tokens(raw_code)
    code = re.sub(r"\((01|21)\)", r"\1", code)
    if not code:
        raise InvalidHonestSignCode("Код Честного знака не указан.")

    # A scanner normally keeps the GS separator before crypto AIs 91/92.
    # cises/info needs the identification part (01 + GTIN + 21 + serial).
    if "\x1d" in code:
        code = code.split("\x1d", 1)[0]
    else:
        crypto_start = None
        for index in range(18, max(18, len(code) - 7)):
            if code[index : index + 2] == "91" and code[index + 6 : index + 8] in {"92", "93"}:
                crypto_start = index
                break
        if crypto_start is not None:
            code = code[:crypto_start]

    if not code.startswith("01") or len(code) < 18 or code[16:18] != "21":
        raise InvalidHonestSignCode("Не распознан код 01 + GTIN + 21 + серийный номер.")
    if not code[2:16].isdigit():
        raise InvalidHonestSignCode("GTIN в коде Честного знака указан неверно.")
    if not 18 <= len(code) <= 74:
        raise InvalidHonestSignCode("Длина кода Честного знака должна быть от 18 до 74 символов.")
    if any(ord(char) < 32 for char in code):
        raise InvalidHonestSignCode("Код содержит неподдерживаемые управляющие символы.")
    return code


def _read_true_api_token() -> str:
    token = str(os.getenv("CRPT_TRUE_API_TOKEN") or "").strip()
    token_file = str(os.getenv("CRPT_TRUE_API_TOKEN_FILE") or "").strip()
    if not token and token_file:
        try:
            token = Path(token_file).read_text(encoding="utf-8").strip()
        except (OSError, UnicodeError):
            logger.warning("CRPT True API token file is unavailable")
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    return token


def is_true_api_configured() -> bool:
    return bool(_read_true_api_token())


def _true_api_timeout() -> float:
    try:
        value = float(os.getenv("CRPT_TRUE_API_TIMEOUT_SECONDS") or "8")
    except (TypeError, ValueError):
        value = 8.0
    return min(max(value, 1.0), 20.0)


def _parse_iso_datetime(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def _masked_cis(cis: str) -> str:
    if len(cis) <= 22:
        return cis
    return f"{cis[:18]}...{cis[-4:]}"


def _normalize_true_api_result(*, cis: str, payload) -> dict:
    entries = payload if isinstance(payload, list) else [payload]
    entry = entries[0] if entries and isinstance(entries[0], dict) else {}
    error_code = str(entry.get("errorCode") or entry.get("error_code") or "").strip()
    error_message = str(
        entry.get("errorMessage")
        or entry.get("error_message")
        or (payload.get("error_message") if isinstance(payload, dict) else "")
        or ""
    ).strip()
    info = entry.get("cisInfo") if isinstance(entry.get("cisInfo"), dict) else {}

    gtin = str(info.get("gtin") or cis[2:16]).strip()
    status = str(info.get("status") or "").strip().upper()
    status_ex = str(info.get("statusEx") or info.get("statusExt") or "").strip().upper()
    mark_withdraw = bool(info.get("markWithdraw"))
    expiration_raw = info.get("expirationDate")
    expiration = _parse_iso_datetime(expiration_raw)
    now = timezone.now()
    expired = bool(expiration and expiration < now)
    special_state = status_ex not in {"", "EMPTY"}
    in_circulation = status == "INTRODUCED"
    accepted = bool(in_circulation and not special_state and not mark_withdraw and not expired)

    if error_code or not status:
        message = error_message or "Код не найден в ГИС МТ."
        return {
            "ok": True,
            "found": False,
            "accepted": False,
            "verdict": "Код не подтвержден",
            "tone": "danger",
            "message": message,
            "error_code": error_code or "not_found",
            "code_masked": _masked_cis(cis),
            "gtin": gtin,
            "crypto_check": "Не выполнялась",
        }

    if accepted:
        verdict = "Код действующий: в обороте"
        tone = "success"
    elif status in {"RETIRED", "WRITTEN_OFF"}:
        verdict = TRUE_API_STATUS_LABELS[status]
        tone = "danger"
    elif in_circulation:
        verdict = "Код в обороте с ограничением"
        tone = "warning"
    else:
        verdict = "Код не готов к обороту"
        tone = "warning"

    return {
        "ok": True,
        "found": True,
        "accepted": accepted,
        "verdict": verdict,
        "tone": tone,
        "message": "Статус получен из ГИС МТ.",
        "code_masked": _masked_cis(cis),
        "gtin": gtin,
        "product_name": str(info.get("productName") or "").strip(),
        "brand": str(info.get("brand") or "").strip(),
        "product_group": str(info.get("productGroup") or "").strip(),
        "status": status,
        "status_label": TRUE_API_STATUS_LABELS.get(status, status),
        "status_ex": status_ex or "EMPTY",
        "mark_withdraw": mark_withdraw,
        "withdraw_reason": str(
            info.get("withdrawReason") or info.get("eliminationReason") or ""
        ).strip(),
        "expiration_date": str(expiration_raw or "").strip(),
        "expired": expired,
        "owner_name": str(info.get("ownerName") or "").strip(),
        # cises/info confirms registration and circulation status, not crypto validity.
        "crypto_check": "Не выполнялась",
    }


def check_honest_sign_status(raw_code: str) -> dict:
    cis = normalize_honest_sign_code(raw_code)
    token = _read_true_api_token()
    if not token:
        raise TrueApiNotConfigured(
            "Проверка ГИС МТ не подключена: требуется действующий токен True API."
        )

    cache_key = f"marking:true-api:status:{hashlib.sha256(cis.encode('utf-8')).hexdigest()}"
    cached = cache.get(cache_key)
    if isinstance(cached, dict):
        return cached

    base_url = str(os.getenv("CRPT_TRUE_API_BASE_URL") or TRUE_API_DEFAULT_BASE_URL).rstrip("/")
    try:
        response = requests.post(
            f"{base_url}/cises/info",
            json=[cis],
            headers={
                "Accept": "application/json",
                "Content-Type": "application/json",
                "Authorization": f"Bearer {token}",
            },
            timeout=_true_api_timeout(),
        )
    except requests.RequestException as exc:
        logger.warning("CRPT True API request failed: %s", exc.__class__.__name__)
        raise TrueApiUnavailable("ГИС МТ временно недоступна. Повторите проверку.") from exc

    try:
        payload = response.json()
    except ValueError as exc:
        raise TrueApiUnavailable("ГИС МТ вернула некорректный ответ.") from exc
    if response.status_code in {401, 403}:
        raise TrueApiNotConfigured("Токен True API недействителен или доступ запрещен.")
    if response.status_code >= 500:
        raise TrueApiUnavailable("ГИС МТ временно недоступна. Повторите проверку.")
    if response.status_code >= 400 and not isinstance(payload, list):
        message = str(payload.get("error_message") or "Ошибка проверки в ГИС МТ.")
        raise TrueApiUnavailable(message)

    result = _normalize_true_api_result(cis=cis, payload=payload)
    cache.set(cache_key, result, timeout=5)
    logger.info(
        "CRPT status checked code_hash=%s gtin=%s status=%s accepted=%s",
        hashlib.sha256(cis.encode("utf-8")).hexdigest()[:12],
        result.get("gtin") or "",
        result.get("status") or result.get("error_code") or "",
        result.get("accepted"),
    )
    return result


def honest_sign_status_check_response(*, request):
    ok, response = _views()._require_status_check_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    try:
        result = check_honest_sign_status(data.get("code") or "")
    except TrueApiError as exc:
        return JsonResponse(
            {"ok": False, "error_code": exc.error_code, "error": str(exc)},
            status=exc.http_status,
        )
    return JsonResponse(result)


def _decode_duplicate_scan(raw_code: str) -> str:
    code = normalize_marking_code(_normalize_gs_tokens(raw_code))
    if not code:
        raise InvalidHonestSignCode("Код Честного знака не указан.")
    normalize_honest_sign_code(code)
    return code


def _barcode_as_gtin14(value: str) -> str:
    digits = str(value or "").strip()
    if not digits.isdigit() or not 8 <= len(digits) <= 14:
        return ""
    return digits.zfill(14)


def _duplicate_code_variants(code: str) -> set[str]:
    variants = marking_code_variants(code) | {code, normalize_honest_sign_code(code)}
    if "\x1d" in code:
        variants.add(code.replace("\x1d", "_x001D_"))
        variants.add(code.replace("\x1d", "{FNC1}"))
        variants.add("{FNC1}" + code.replace("\x1d", "{FNC1}"))
    return {value for value in variants if value}


def _duplicate_label_template() -> dict:
    label_key = "item_cz"
    template_key = resolve_processing_param_template_key(
        label_key,
        load_processing_param_template_bindings(),
    )
    template = get_effective_label_template(template_key) or get_effective_label_template(label_key)
    return {
        "label_key": label_key,
        "template_key": str(template.get("key") or template_key or label_key).strip(),
        "title": str(template.get("template_title") or template.get("title") or "Товар ЧЗ").strip(),
        "width_mm": float(template.get("width_mm") or 58),
        "height_mm": float(template.get("height_mm") or 40),
    }


def _label_date(value) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def _validate_duplicate_pair(*, product_barcode: str, raw_code: str) -> dict:
    barcode = str(product_barcode or "").strip()
    if not barcode:
        raise InvalidHonestSignCode("Сначала отсканируйте штрихкод товара.")
    barcode_rows = list(
        SKUBarcode.objects.select_related("sku", "sku__agency")
        .filter(value=barcode, sku__deleted=False)
        .order_by("id")[:2]
    )
    if len(barcode_rows) > 1:
        raise InvalidHonestSignCode(
            "Штрихкод используется у нескольких клиентов. "
            "Откройте печать из карточки товара нужного клиента."
        )
    barcode_row = barcode_rows[0] if barcode_rows else None
    if not barcode_row or not barcode_row.sku:
        raise InvalidHonestSignCode("Штрихкод товара не найден в номенклатуре.")
    sku = barcode_row.sku

    code = _decode_duplicate_scan(raw_code)
    identity_code = normalize_honest_sign_code(code)
    gtin = identity_code[2:16]

    status_result = None
    if is_true_api_configured():
        status_result = check_honest_sign_status(code)
        if not status_result.get("accepted"):
            raise InvalidHonestSignCode(
                f"Печать запрещена: {status_result.get('verdict') or 'код не подтвержден в ГИС МТ'}."
            )

    code_hash = hashlib.sha256(code.encode("utf-8")).hexdigest()
    return {
        "barcode": barcode,
        "barcode_size": str(barcode_row.size or sku.size or "").strip(),
        "code": code,
        "code_hash": code_hash,
        "code_masked": _masked_cis(identity_code),
        "gtin": gtin,
        "sku": sku,
        "sku_code": str(sku.sku_code or "").strip(),
        "product_name": str(sku.name or "").strip(),
        "agency_name": str(getattr(sku.agency, "agn_name", "") or "").strip(),
        "brand": str(sku.brand or "").strip(),
        "subject": str(sku.vid_tovar or sku.tovar_category or "").strip(),
        "color": str(sku.color or "").strip(),
        "composition": str(sku.composition or "").strip(),
        "supplier": str(
            getattr(sku.agency, "short_name", "")
            or getattr(sku.agency, "agn_name", "")
            or ""
        ).strip(),
        "country": str(sku.made_in or "").strip(),
        "manufacture_date": _label_date(sku.cr_product_date),
        "expiration_date": _label_date(sku.end_product_date),
        "status_result": status_result,
    }


def honest_sign_duplicate_validate_response(*, request):
    ok, response = _views()._require_duplicate_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    try:
        pair = _validate_duplicate_pair(
            product_barcode=data.get("barcode") or "",
            raw_code=data.get("code") or "",
        )
    except TrueApiError as exc:
        return JsonResponse(
            {"ok": False, "error_code": exc.error_code, "error": str(exc)},
            status=exc.http_status,
        )
    token = signing.dumps(
        {"barcode": pair["barcode"], "code_hash": pair["code_hash"]},
        salt=DUPLICATE_TOKEN_SALT,
        compress=True,
    )
    status_result = pair.get("status_result") or {}
    return JsonResponse(
        {
            "ok": True,
            "barcode": pair["barcode"],
            "matrix_code": pair["code"],
            "code_masked": pair["code_masked"],
            "gtin": pair["gtin"],
            "sku_code": pair["sku_code"],
            "product_name": pair["product_name"],
            "agency_name": pair["agency_name"],
            "size": pair["barcode_size"],
            "brand": pair["brand"],
            "subject": pair["subject"],
            "color": pair["color"],
            "composition": pair["composition"],
            "supplier": pair["supplier"],
            "country": pair["country"],
            "manufacture_date": pair["manufacture_date"],
            "expiration_date": pair["expiration_date"],
            "label_template": _duplicate_label_template(),
            "status_label": status_result.get("status_label") or "Структура кода проверена",
            "status_verified": bool(status_result),
            "token": token,
        }
    )


def _validate_duplicate_png(value: str) -> str:
    encoded = str(value or "").strip()
    if not encoded:
        raise InvalidHonestSignCode("Изображение этикетки не сформировано.")
    if len(encoded) > DUPLICATE_MAX_LABEL_BASE64_SIZE:
        raise InvalidHonestSignCode("Изображение этикетки превышает допустимый размер.")
    try:
        decoded = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise InvalidHonestSignCode("Изображение этикетки повреждено.") from exc
    if not decoded.startswith(b"\x89PNG\r\n\x1a\n"):
        raise InvalidHonestSignCode("Для печати принимается только PNG-этикетка.")
    return encoded


def honest_sign_duplicate_print_response(*, request):
    ok, response = _views()._require_duplicate_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    try:
        pair = _validate_duplicate_pair(
            product_barcode=data.get("barcode") or "",
            raw_code=data.get("code") or "",
        )
        token_payload = signing.loads(
            str(data.get("token") or ""),
            salt=DUPLICATE_TOKEN_SALT,
            max_age=DUPLICATE_TOKEN_MAX_AGE_SECONDS,
        )
        if (
            token_payload.get("barcode") != pair["barcode"]
            or token_payload.get("code_hash") != pair["code_hash"]
        ):
            raise InvalidHonestSignCode("Данные сканирования изменились. Отсканируйте товар повторно.")
        label_png = _validate_duplicate_png(data.get("label_png_base64") or "")
    except signing.SignatureExpired:
        return JsonResponse(
            {"ok": False, "error": "Проверка устарела. Отсканируйте товар повторно."},
            status=409,
        )
    except signing.BadSignature:
        return JsonResponse({"ok": False, "error": "Проверка сканов не подтверждена."}, status=400)
    except TrueApiError as exc:
        return JsonResponse(
            {"ok": False, "error_code": exc.error_code, "error": str(exc)},
            status=exc.http_status,
        )

    printer_name = str(data.get("printer_name") or "").strip()
    if not printer_name:
        return JsonResponse({"ok": False, "error": "Выберите принтер этикеток."}, status=400)

    from processing_app.services import ProcessingWorkflowService

    label_template = _duplicate_label_template()
    print_result = ProcessingWorkflowService.enqueue_processing_print_job(
        request=request,
        data={
            "order_id": "",
            "card_id": "honest-sign-duplicate",
            "article": pair["sku_code"],
            "barcode": pair["barcode"],
            "size": pair["barcode_size"],
            "printer_name": printer_name,
            "agent_id": str(data.get("agent_id") or "").strip(),
            "copies": 1,
            "label_png_base64": label_png,
            "template_key": label_template["template_key"],
            "label_width_mm": label_template["width_mm"],
            "label_height_mm": label_template["height_mm"],
        },
    )
    if print_result.http_status != 200 or not print_result.payload.get("ok"):
        return JsonResponse(print_result.payload, status=print_result.http_status)

    journal, _created = AuditJournal.objects.get_or_create(
        code="honest_sign_duplicate",
        defaults={
            "name": "Дублирование Честного знака",
            "description": "Печать копии существующего кода ЧЗ для наружной упаковки.",
        },
    )
    AuditEntry.objects.create(
        journal=journal,
        action="clone",
        sku=pair["sku"],
        agency=pair["sku"].agency,
        user=request.user if request.user.is_authenticated else None,
        description=f"Напечатана копия ЧЗ для {pair['sku_code']}",
        snapshot={
            "barcode": pair["barcode"],
            "gtin": pair["gtin"],
            "code_masked": pair["code_masked"],
            "code_hash": pair["code_hash"],
            "printer_name": printer_name,
            "agent_id": str(data.get("agent_id") or "").strip(),
            "print_job_ids": print_result.payload.get("job_ids") or [],
        },
    )
    return JsonResponse(
        {
            **print_result.payload,
            "message": "Этикетка поставлена в очередь печати.",
        }
    )


def return_printed_marking_response(*, request):
    ok, response = _views()._require_return_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    try:
        scanned_code = _decode_duplicate_scan(
            data.get("code") or data.get("marking_code") or ""
        )
    except TrueApiError as exc:
        return JsonResponse(
            {"ok": False, "error_code": exc.error_code, "error": str(exc)},
            status=exc.http_status,
        )

    identity = marking_code_identity(scanned_code)
    variants = marking_code_variants(scanned_code)
    with transaction.atomic():
        marking = (
            MarkingCode.objects.select_for_update(of=("self",))
            .select_related("agency", "sku", "used_by")
            .filter(Q(identity_key=identity) | Q(code__in=variants))
            .order_by("id")
            .first()
        )
        if not marking:
            return JsonResponse(
                {"ok": False, "error_code": "not_found", "error": "Код ЧЗ не найден в базе."},
                status=404,
            )
        print_job = None
        print_completed_pending_confirmation = False
        print_failed_needs_review = False
        if marking.print_job_id:
            from processing_app.models import ProcessingPrintJob

            print_job = (
                ProcessingPrintJob.objects.select_for_update()
                .filter(id=marking.print_job_id)
                .first()
            )
            print_completed_pending_confirmation = bool(
                not marking.printed_at
                and print_job
                and print_job.status == ProcessingPrintJob.STATUS_PRINTED
                and print_job.processing_param_key == FREE_PRINT_PROCESSING_PARAM_KEY
            )
            print_failed_needs_review = bool(
                not marking.printed_at
                and print_job
                and print_job.status == ProcessingPrintJob.STATUS_FAILED
                and print_job.processing_param_key == FREE_PRINT_PROCESSING_PARAM_KEY
            )
            if (
                not marking.printed_at
                and print_job
                and print_job.status
                in {
                    ProcessingPrintJob.STATUS_PENDING,
                    ProcessingPrintJob.STATUS_PRINTING,
                }
            ):
                return JsonResponse(
                    {
                        "ok": False,
                        "error_code": "print_in_progress",
                        "error": (
                            "Задание печати еще выполняется. Дождитесь результата агента, "
                            "затем повторно отсканируйте ЧЗ для возврата."
                        ),
                    },
                    status=409,
                )
        if (
            not marking.printed_at
            and not print_completed_pending_confirmation
            and not print_failed_needs_review
        ):
            return JsonResponse(
                {
                    "ok": False,
                    "error_code": "not_printed",
                    "error": "Код найден, но не отмечен как распечатанный. Возвращать нечего.",
                },
                status=409,
            )

        conflict = printed_marking_storage_conflict(marking)
        if conflict:
            return JsonResponse(
                {
                    "ok": False,
                    "error_code": "code_in_box",
                    "error": marking_conflict_message(conflict),
                    "details": conflict,
                },
                status=409,
            )

        previous_print_job_id = marking.print_job_id
        previous_printed_at = marking.printed_at
        marking.printed_at = None
        marking.printed_by = None
        marking.print_job_id = None
        marking.print_reserved_at = None
        marking.save(
            update_fields=[
                "printed_at",
                "printed_by",
                "print_job_id",
                "print_reserved_at",
            ]
        )

        if (print_completed_pending_confirmation or print_failed_needs_review) and print_job:
            print_job.processing_param_key = FREE_PRINT_RETURNED_PARAM_KEY
            if print_failed_needs_review:
                print_job.error = (
                    "КИЗ из неясного результата печати проверен и возвращен руководителем."
                )
            else:
                print_job.error = (
                    "Фактически распечатанный КИЗ возвращен руководителем после потери "
                    "подтверждения браузера."
                )
            print_job.save(update_fields=["processing_param_key", "error", "updated_at"])
        elif previous_print_job_id:
            ProcessingPrintJob.objects.filter(
                id=previous_print_job_id,
                status__in=[
                    ProcessingPrintJob.STATUS_PENDING,
                    ProcessingPrintJob.STATUS_PRINTING,
                ],
            ).update(
                status=ProcessingPrintJob.STATUS_FAILED,
                error="Код ЧЗ возвращен руководителем обработки как неиспользованный.",
                updated_at=timezone.now(),
            )

        journal, _created = AuditJournal.objects.get_or_create(
            code="marking_return",
            defaults={
                "name": "Возврат ЧЗ",
                "description": "Возврат распечатанных, но неиспользованных кодов ЧЗ.",
            },
        )
        AuditEntry.objects.create(
            journal=journal,
            action="update",
            sku=marking.sku,
            agency=marking.agency,
            user=request.user if request.user.is_authenticated else None,
            description=f"Код ЧЗ возвращен в доступные: {marking.sku_code or '-'}",
            snapshot={
                "marking_id": marking.id,
                "code_hash": hashlib.sha256(marking.code.encode("utf-8")).hexdigest(),
                "code_masked": _masked_cis(marking_code_identity(marking.code)),
                "order_type": marking.order_type,
                "order_id": marking.order_id,
                "sku_code": marking.sku_code,
                "size": marking.size,
                "barcode": marking.barcode,
                "previous_print_job_id": previous_print_job_id,
                "previous_printed_at": (
                    previous_printed_at.isoformat() if previous_printed_at else None
                ),
                "print_completed_pending_confirmation": (
                    print_completed_pending_confirmation
                ),
                "print_failed_needs_review": print_failed_needs_review,
            },
        )

    return JsonResponse(
        {
            "ok": True,
            "message": (
                "Код ЧЗ возвращен в базу как неиспользованный и снова доступен для печати. "
                "Уничтожьте отсканированную этикетку."
            ),
            "code_masked": _masked_cis(marking_code_identity(marking.code)),
            "sku_code": marking.sku_code,
            "size": marking.size,
            "order_id": marking.order_id,
        }
    )


def _views():
    from . import views as marking_views

    return marking_views


def _free_print_receiving_order_candidates(order_id: str) -> list[str]:
    normalized_id = str(order_id or "").strip()
    candidates = [normalized_id] if normalized_id else []
    suffix_match = re.fullmatch(r"(\d+)_PR", normalized_id, flags=re.IGNORECASE)
    if suffix_match:
        sequence = int(suffix_match.group(1))
        candidates.extend((build_public_order_number("receiving", sequence), str(sequence)))
    elif normalized_id.isdigit():
        candidates.append(build_public_order_number("receiving", int(normalized_id)))
    return list(dict.fromkeys(candidate for candidate in candidates if candidate))


def _free_print_order_context(*, order_type: str, order_id: str):
    normalized_type = str(order_type or "").strip().lower()
    normalized_id = str(order_id or "").strip()
    if not normalized_id:
        return None, {}, None, [], "", "Укажите номер заявки."
    if normalized_type == "processing":
        latest, payload, agency = _views()._get_processing_order(normalized_id)
        items = extract_processing_items(payload) if latest else []
        resolved_id = normalized_id
    elif normalized_type == "receiving":
        entries = []
        for candidate in _free_print_receiving_order_candidates(normalized_id):
            entries = list(
                OrderAuditEntry.objects.filter(order_id=candidate, order_type="receiving")
                .select_related("agency")
                .order_by("created_at", "id")
            )
            if entries:
                break
        latest = entries[-1] if entries else None
        payload = build_order_runtime_payload("receiving", entries=entries) if latest else {}
        agency = latest.agency if latest else None
        items = _views()._extract_receiving_items(payload) if latest else []
        resolved_id = latest.order_id if latest else normalized_id
    else:
        return None, {}, None, [], "", "Выберите заявку на приемку или обработку."
    if not latest:
        return None, {}, None, [], normalized_id, "Заявка не найдена."
    return latest, payload, agency, items, resolved_id, ""


def _free_print_selected_item(*, data: dict):
    order_type = str(data.get("order_type") or "").strip().lower()
    order_id = str(data.get("order_id") or "").strip()
    latest, payload, agency, items, resolved_id, error = _free_print_order_context(
        order_type=order_type,
        order_id=order_id,
    )
    if error:
        return None, JsonResponse({"ok": False, "error": error}, status=400)
    sku_code = str(data.get("sku_code") or "").strip()
    size = str(data.get("size") or "").strip()
    barcode = str(data.get("barcode") or "").strip()
    selected = None
    for item in items:
        item_sku = str(item.get("sku_code") or "").strip()
        item_size = str(item.get("size") or "").strip()
        item_barcode = str(item.get("barcode") or "").strip()
        if item_sku != sku_code or item_size != size:
            continue
        if barcode and item_barcode and barcode != item_barcode:
            continue
        selected = {
            **item,
            "sku_code": item_sku,
            "size": item_size,
            "barcode": item_barcode or barcode,
        }
        break
    if not selected:
        return None, JsonResponse(
            {"ok": False, "error": "Выбранная позиция не найдена в заявке."},
            status=400,
        )
    sku = _views()._resolve_sku(agency, selected["sku_code"])
    if sku and agency and sku.agency_id and sku.agency_id != agency.id:
        return None, JsonResponse(
            {"ok": False, "error": "Товар относится к другому клиенту."},
            status=409,
        )
    selected["sku"] = sku
    selected["product_name"] = str(getattr(sku, "name", "") or "").strip()
    return {
        "order_type": order_type,
        "order_id": resolved_id,
        "latest": latest,
        "payload": payload,
        "agency": agency,
        "items": items,
        "item": selected,
    }, None


def _free_print_import_items(context: dict) -> list[dict]:
    """Return unique order items that may receive codes from one import file."""
    result = []
    seen = set()
    agency = context.get("agency")
    for source_item in context.get("items") or []:
        sku_code = str(source_item.get("sku_code") or "").strip()
        size = str(source_item.get("size") or "").strip()
        barcode = str(source_item.get("barcode") or "").strip()
        key = (sku_code, size, barcode)
        if key in seen:
            continue
        seen.add(key)
        sku = _views()._resolve_sku(agency, sku_code)
        result.append(
            {
                "sku": sku,
                "sku_code": sku_code,
                "size": size,
                "barcode": barcode,
                "gtin14": _barcode_as_gtin14(barcode),
            }
        )
    return result


def _free_print_import_item_label(item: dict) -> str:
    parts = [item.get("sku_code") or "без SKU"]
    if item.get("size"):
        parts.append(f"размер {item['size']}")
    if item.get("barcode"):
        parts.append(f"ШК {item['barcode']}")
    return ", ".join(parts)


def _free_print_import_matches(
    *,
    items: list[dict],
    file_barcode: str = "",
    code_gtin: str = "",
) -> list[dict]:
    barcode = str(file_barcode or "").strip()
    if barcode:
        exact = [item for item in items if item["barcode"] == barcode]
        if exact:
            return exact
        barcode_gtin = _barcode_as_gtin14(barcode)
        if barcode_gtin:
            return [item for item in items if item["gtin14"] == barcode_gtin]
        return []
    if code_gtin:
        return [item for item in items if item["gtin14"] == code_gtin]
    return []


def _free_print_scope(context: dict):
    item = context["item"]
    scope = MarkingCode.objects.filter(
        order_type=context["order_type"],
        order_id=context["order_id"],
        sku_code=item["sku_code"],
        size=item["size"],
    )
    if context.get("agency"):
        scope = scope.filter(agency=context["agency"])
    if item.get("barcode"):
        scope = scope.filter(barcode=item["barcode"])
    return scope


def _free_print_job_ids(data: dict) -> list[int]:
    raw_values = data.get("job_ids")
    if raw_values is None:
        raw_values = data.get("jobIds")
    if not isinstance(raw_values, (list, tuple)):
        raw_values = []
    result: list[int] = []
    for raw_value in raw_values[:100]:
        try:
            value = int(raw_value)
        except (TypeError, ValueError):
            continue
        if value > 0 and value not in result:
            result.append(value)
    return result


def _sync_free_print_job_state(scope) -> None:
    job_ids = list(
        scope.filter(print_job_id__isnull=False)
        .values_list("print_job_id", flat=True)
        .distinct()
    )
    if not job_ids:
        return
    from processing_app.models import ProcessingPrintJob

    jobs = {
        int(row["id"]): row
        for row in ProcessingPrintJob.objects.filter(id__in=job_ids).values(
            "id", "status", "updated_at", "processing_param_key"
        )
    }
    pending_confirmation_job_ids = [
        job_id
        for job_id, row in jobs.items()
        if row["status"] == ProcessingPrintJob.STATUS_PRINTED
        and row["processing_param_key"] == FREE_PRINT_PROCESSING_PARAM_KEY
    ]
    reservation_job_ids = [
        job_id
        for job_id, row in jobs.items()
        if row["processing_param_key"] == FREE_PRINT_RESERVATION_PARAM_KEY
    ]
    stale_reservation_job_ids = [
        job_id
        for job_id in reservation_job_ids
        if jobs[job_id]["updated_at"]
        and jobs[job_id]["updated_at"] <= timezone.now() - FREE_PRINT_RESERVATION_TTL
    ]
    printed_job_ids = [
        job_id
        for job_id, row in jobs.items()
        if row["status"] == ProcessingPrintJob.STATUS_PRINTED
        and row["processing_param_key"]
        not in {FREE_PRINT_PROCESSING_PARAM_KEY, FREE_PRINT_RESERVATION_PARAM_KEY}
    ]
    released_job_ids = [
        job_id
        for job_id, row in jobs.items()
        if row["status"] == ProcessingPrintJob.STATUS_FAILED
        and row["processing_param_key"] != FREE_PRINT_PROCESSING_PARAM_KEY
    ]
    released_job_ids.extend(stale_reservation_job_ids)
    released_job_ids.extend(job_id for job_id in job_ids if int(job_id) not in jobs)
    if printed_job_ids:
        scope.filter(
            print_job_id__in=printed_job_ids,
            printed_at__isnull=True,
        ).update(printed_at=timezone.now())
    if pending_confirmation_job_ids:
        scope.filter(
            print_job_id__in=pending_confirmation_job_ids,
            used_at__isnull=True,
        ).update(printed_at=None)
    if released_job_ids:
        scope.filter(
            print_job_id__in=released_job_ids,
            printed_at__isnull=True,
        ).update(
            print_job_id=None,
            print_reserved_at=None,
            printed_by=None,
        )


def _free_print_stats(context: dict) -> dict:
    scope = _free_print_scope(context)
    _sync_free_print_job_state(scope)
    return {
        "loaded": scope.count(),
        "available": scope.filter(
            used_at__isnull=True,
            printed_at__isnull=True,
            print_job_id__isnull=True,
        ).count(),
        "queued": scope.filter(
            used_at__isnull=True,
            printed_at__isnull=True,
            print_job_id__isnull=False,
        ).count(),
        "printed": scope.filter(printed_at__isnull=False).count(),
        "used": scope.filter(used_at__isnull=False).count(),
    }


def free_marking_print_page_context(*, order_type: str, order_id: str) -> dict:
    normalized_type = str(order_type or "processing").strip().lower()
    if normalized_type not in {"processing", "receiving"}:
        normalized_type = "processing"
    normalized_id = str(order_id or "").strip()
    latest, _payload, agency, items, resolved_id, error = _free_print_order_context(
        order_type=normalized_type,
        order_id=normalized_id,
    )
    result_items = []
    if latest:
        order_scope = MarkingCode.objects.filter(
            order_type=normalized_type,
            order_id=resolved_id,
        )
        if agency:
            order_scope = order_scope.filter(agency=agency)
        _sync_free_print_job_state(order_scope)
        stats_map = {
            (row["sku_code"], row["size"] or "", row["barcode"] or ""): {
                "loaded": row["loaded"],
                "available": row["available"],
                "queued": row["queued"],
                "printed": row["printed"],
                "used": row["used"],
            }
            for row in order_scope.values("sku_code", "size", "barcode").annotate(
                loaded=Count("id"),
                available=Count(
                    "id",
                    filter=Q(
                        used_at__isnull=True,
                        printed_at__isnull=True,
                        print_job_id__isnull=True,
                    ),
                ),
                queued=Count(
                    "id",
                    filter=Q(
                        used_at__isnull=True,
                        printed_at__isnull=True,
                        print_job_id__isnull=False,
                    ),
                ),
                printed=Count("id", filter=Q(printed_at__isnull=False)),
                used=Count("id", filter=Q(used_at__isnull=False)),
            )
        }
        client_free_map = {}
        if agency:
            client_free_scope = (
                MarkingCode.objects.filter(
                    agency=agency,
                    used_at__isnull=True,
                    printed_at__isnull=True,
                    print_job_id__isnull=True,
                )
                .filter(Q(order_id__isnull=True) | Q(order_id=""))
            )
            client_free_map = {
                (row["sku_code"], row["size"] or "", row["barcode"] or ""): row["count"]
                for row in client_free_scope.values("sku_code", "size", "barcode").annotate(
                    count=Count("id")
                )
            }
        for item in items:
            sku_code = str(item.get("sku_code") or "").strip()
            size = str(item.get("size") or "").strip()
            barcode = str(item.get("barcode") or "").strip()
            sku = _views()._resolve_sku(agency, sku_code)
            key = (sku_code, size, barcode)
            stats = stats_map.get(
                key,
                {"loaded": 0, "available": 0, "queued": 0, "printed": 0, "used": 0},
            )
            client_free = client_free_map.get(key, 0)
            known_count = stats["loaded"] + client_free
            try:
                required_count = max(int(item.get("qty") or 0), 0)
            except (TypeError, ValueError):
                required_count = 0
            if known_count <= 0:
                marking_status = "missing"
                marking_status_label = "ЧЗ нет"
            elif required_count and known_count < required_count:
                marking_status = "partial"
                marking_status_label = "ЧЗ частично"
            else:
                marking_status = "complete"
                marking_status_label = "ЧЗ есть"
            result_items.append(
                {
                    **item,
                    "sku_code": sku_code,
                    "size": size,
                    "barcode": barcode,
                    "product_name": str(getattr(sku, "name", "") or "").strip(),
                    "stats": {**stats, "client_free": client_free},
                    "marking_status": marking_status,
                    "marking_status_label": marking_status_label,
                    "marking_known_count": known_count,
                    "marking_required_count": required_count,
                    "marking_client_free_count": client_free,
                }
            )
    display_id = format_order_number(normalized_type, resolved_id) if latest else normalized_id
    return {
        "selected_order_type": normalized_type,
        "selected_order_id": display_id,
        "selected_order_key": resolved_id or normalized_id,
        "selected_order": latest,
        "selected_agency": agency,
        "selected_items": result_items,
        "selection_error": error if normalized_id else "",
    }


def free_marking_import_response(*, request):
    ok, response = _views()._require_free_print_role(request)
    if not ok:
        return response
    context, error_response = _free_print_selected_item(data=request.POST)
    if error_response:
        return error_response
    uploaded = request.FILES.get("file")
    if not uploaded:
        return JsonResponse({"ok": False, "error": "Выберите файл .xlsx."}, status=400)
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception:
        return JsonResponse({"ok": False, "error": "Не удалось прочитать .xlsx файл."}, status=400)

    selected_barcode = context["item"].get("barcode") or ""
    order_items = _free_print_import_items(context)
    if not order_items:
        return JsonResponse(
            {"ok": False, "error": "В заявке нет товаров, к которым можно привязать КИЗ."},
            status=400,
        )
    rows = []
    errors = []
    first_by_identity = {}
    for index, row in enumerate(workbook.active.iter_rows(values_only=True), start=1):
        first = _views()._normalize_cell(row[0]) if row and len(row) > 0 else ""
        second = _views()._normalize_cell(row[1]) if row and len(row) > 1 else ""
        if index == 1:
            header = f"{first} {second}".lower()
            if any(token in header for token in ("код", "киз", "чест", "barcode", "штрих")):
                continue
        file_barcode = first if second else ""
        raw_code = second or first
        if not raw_code:
            continue
        try:
            code = validate_import_marking_code(
                raw_code,
                product_barcode=file_barcode or selected_barcode,
            )
            code = _decode_duplicate_scan(code)
        except (MarkingCodeFormatError, TrueApiError) as exc:
            errors.append(f"строка {index}: {exc}")
            continue
        identity = marking_code_identity(code)
        if identity in first_by_identity:
            errors.append(
                f"строка {index}: код ЧЗ уже указан в строке {first_by_identity[identity]}"
            )
            continue
        first_by_identity[identity] = index
        code_gtin = (
            identity[2:16]
            if identity.startswith("01") and identity[16:18] == "21"
            else ""
        )
        rows.append(
            {
                "row_number": index,
                "file_barcode": file_barcode,
                "code": code,
                "identity": identity,
                "code_gtin": code_gtin,
            }
        )

    if not rows:
        return JsonResponse(
            {
                "ok": False,
                "error": "В файле нет корректных кодов КИЗ.",
                "errors": errors,
            },
            status=400,
        )
    identities = {row["identity"] for row in rows}
    variants = set()
    for row in rows:
        variants.update(marking_code_variants(row["code"]))
    existing_identities = occupied_marking_identities(identities, variants)
    for row in rows:
        if row["identity"] in existing_identities:
            errors.append(
                f"строка {row['row_number']}: код ЧЗ уже существует в системе"
            )

    selected_key = (
        context["item"]["sku_code"],
        context["item"]["size"],
        selected_barcode,
    )
    selected_item = next(
        (
            item
            for item in order_items
            if (item["sku_code"], item["size"], item["barcode"]) == selected_key
        ),
        None,
    )
    for row in rows:
        matches = _free_print_import_matches(
            items=order_items,
            file_barcode=row["file_barcode"],
            code_gtin=row["code_gtin"],
        )
        if len(matches) == 1:
            row["item"] = matches[0]
            continue
        if len(matches) > 1:
            labels = "; ".join(
                _free_print_import_item_label(item) for item in matches[:3]
            )
            errors.append(
                f"строка {row['row_number']}: товар определён неоднозначно ({labels})"
            )
            continue
        if row["file_barcode"]:
            errors.append(
                f"строка {row['row_number']}: ШК {row['file_barcode']} "
                "не найден среди товаров заявки"
            )
            continue
        if len(order_items) > 1:
            if row["code_gtin"]:
                errors.append(
                    f"строка {row['row_number']}: GTIN {row['code_gtin']} "
                    "не найден среди товаров заявки"
                )
            else:
                errors.append(
                    f"строка {row['row_number']}: товар нельзя определить по КИЗ; "
                    "добавьте в файл колонку «ШК товара»"
                )
            continue
        if selected_item:
            row["item"] = selected_item
        else:
            errors.append(
                f"строка {row['row_number']}: не удалось определить товар из заявки"
            )
    if errors:
        preview = "; ".join(errors[:5])
        if len(errors) > 5:
            preview += f"; и ещё {len(errors) - 5}"
        return JsonResponse(
            {
                "ok": False,
                "error": f"Импорт отменён. Исправьте файл: {preview}",
                "errors": errors,
                "added": 0,
            },
            status=400,
        )

    to_create = [
        MarkingCode(
            order_type=context["order_type"],
            order_id=context["order_id"],
            agency=context.get("agency"),
            sku=row["item"].get("sku"),
            sku_code=row["item"]["sku_code"],
            size=row["item"]["size"],
            barcode=row["item"]["barcode"],
            code=row["code"],
            source="import",
            created_by=request.user if request.user.is_authenticated else None,
        )
        for row in rows
    ]
    try:
        with transaction.atomic():
            created = MarkingCode.objects.bulk_create(to_create, batch_size=500)
    except IntegrityError:
        return JsonResponse(
            {
                "ok": False,
                "error": (
                    "Импорт отменён: один из кодов ЧЗ был добавлен параллельно. "
                    "Проверьте файл повторно."
                ),
                "added": 0,
            },
            status=409,
        )
    added_by_item = {}
    for row in rows:
        item = row["item"]
        key = (item["sku_code"], item["size"], item["barcode"])
        if key not in added_by_item:
            added_by_item[key] = {
                "sku_code": item["sku_code"],
                "size": item["size"],
                "barcode": item["barcode"],
                "added": 0,
            }
        added_by_item[key]["added"] += 1
    return JsonResponse(
        {
            "ok": True,
            "added": len(created),
            "items_added": list(added_by_item.values()),
            "duplicates": 0,
            "invalid_rows": 0,
            "mismatched_barcodes": 0,
            "stats": _free_print_stats(context),
        }
    )


def free_marking_candidates_response(*, request):
    ok, response = _views()._require_free_print_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    context, error_response = _free_print_selected_item(data=data)
    if error_response:
        return error_response
    try:
        qty = int(data.get("qty") or 0)
    except (TypeError, ValueError):
        qty = 0
    reserve = data.get("reserve") is True
    max_qty = FREE_PRINT_MAX_LABELS_PER_OPERATION if reserve else FREE_PRINT_MAX_BATCH
    if qty <= 0 or qty > max_qty:
        return JsonResponse(
            {"ok": False, "error": f"Укажите количество от 1 до {max_qty}."},
            status=400,
        )
    scope = _free_print_scope(context)
    _sync_free_print_job_state(scope)
    item = context["item"]
    sku = item.get("sku")
    agency = context.get("agency")
    reservation_job_id = None
    if reserve:
        request_id = str(data.get("request_id") or data.get("requestId") or "").strip()
        if not re.fullmatch(r"[A-Za-z0-9_-]{16,64}", request_id):
            return JsonResponse(
                {"ok": False, "error": "Некорректный идентификатор партии печати."},
                status=400,
            )
        from processing_app.models import ProcessingPrintJob

        with transaction.atomic():
            prior_audit = (
                AuditEntry.objects.select_for_update()
                .filter(
                    journal__code="free_marking_print",
                    action="reserve_print",
                    snapshot__request_id=request_id,
                )
                .order_by("-id")
                .first()
            )
            if prior_audit:
                snapshot = prior_audit.snapshot if isinstance(prior_audit.snapshot, dict) else {}
                reservation_job_id = int(snapshot.get("reservation_job_id") or 0)
                same_request = all(
                    (
                        str(snapshot.get("order_type") or "") == context["order_type"],
                        str(snapshot.get("order_id") or "") == context["order_id"],
                        str(snapshot.get("sku_code") or "") == item["sku_code"],
                        str(snapshot.get("size") or "") == item["size"],
                        str(snapshot.get("barcode") or "") == (item.get("barcode") or ""),
                        int(snapshot.get("count") or 0) == qty,
                    )
                )
                reserved_rows = list(
                    scope.select_for_update()
                    .filter(
                        print_job_id=reservation_job_id,
                        used_at__isnull=True,
                        printed_at__isnull=True,
                    )
                    .order_by("created_at", "id")
                )
                if not same_request or not reservation_job_id or len(reserved_rows) != qty:
                    return JsonResponse(
                        {
                            "ok": False,
                            "error": (
                                "Сохранённая партия печати уже изменилась. "
                                "Обновите страницу и повторите операцию."
                            ),
                        },
                        status=409,
                    )
            else:
                reserved_rows = list(
                    scope.select_for_update()
                    .filter(
                        used_at__isnull=True,
                        printed_at__isnull=True,
                        print_job_id__isnull=True,
                    )
                    .order_by("created_at", "id")[:qty]
                )
                if len(reserved_rows) < qty:
                    return JsonResponse(
                        {
                            "ok": False,
                            "error": (
                                f"Недостаточно свободных КИЗ: доступно {len(reserved_rows)}, "
                                f"запрошено {qty}. Печать не запущена."
                            ),
                            "available": len(reserved_rows),
                            "stats": _free_print_stats(context),
                        },
                        status=409,
                    )
                reservation_job = ProcessingPrintJob.objects.create(
                    status=ProcessingPrintJob.STATUS_PRINTED,
                    order_id=context["order_id"],
                    article=item["sku_code"],
                    barcode=item.get("barcode") or item["sku_code"],
                    size=item["size"],
                    copies_count=1,
                    processing_param_key=FREE_PRINT_RESERVATION_PARAM_KEY,
                    requested_by=(
                        request.user.get_username() if request.user.is_authenticated else ""
                    ),
                    error="Служебный резерв полной партии КИЗ до передачи агенту печати.",
                )
                reservation_job_id = reservation_job.id
                now = timezone.now()
                for marking_code in reserved_rows:
                    marking_code.print_job_id = reservation_job_id
                    marking_code.print_reserved_at = now
                    marking_code.printed_by = (
                        request.user if request.user.is_authenticated else None
                    )
                MarkingCode.objects.bulk_update(
                    reserved_rows,
                    ["print_job_id", "print_reserved_at", "printed_by"],
                )
                journal, _created = AuditJournal.objects.get_or_create(
                    code="free_marking_print",
                    defaults={
                        "name": "Свободная печать КИЗ",
                        "description": "Предварительная печать КИЗ с привязкой к заявке и товару.",
                    },
                )
                AuditEntry.objects.create(
                    journal=journal,
                    action="reserve_print",
                    sku=item.get("sku"),
                    agency=agency,
                    user=request.user if request.user.is_authenticated else None,
                    description=(
                        f"Зарезервирована полная партия из {qty} КИЗ до начала печати: "
                        f"{context['order_type']} {context['order_id']} / {item['sku_code']}"
                    ),
                    snapshot={
                        "order_type": context["order_type"],
                        "order_id": context["order_id"],
                        "sku_code": item["sku_code"],
                        "size": item["size"],
                        "barcode": item.get("barcode") or "",
                        "count": qty,
                        "request_id": request_id,
                        "reservation_job_id": reservation_job_id,
                    },
                )
            codes = [row.code for row in reserved_rows]
    else:
        codes = list(
            scope.filter(
                used_at__isnull=True,
                printed_at__isnull=True,
                print_job_id__isnull=True,
            )
            .order_by("created_at", "id")
            .values_list("code", flat=True)[:qty]
        )
        if len(codes) < qty:
            return JsonResponse(
                {
                    "ok": False,
                    "error": "Недостаточно свободных КИЗ для печати.",
                    "available": len(codes),
                    "stats": _free_print_stats(context),
                },
                status=409,
            )
    agency_name = str(getattr(agency, "agn_name", "") or "").strip()
    supplier_name = str(
        getattr(agency, "short_name", "")
        or getattr(agency, "agn_name", "")
        or ""
    ).strip()
    return JsonResponse(
        {
            "ok": True,
            "codes": codes,
            "sku_code": item["sku_code"],
            "size": item["size"],
            "barcode": item.get("barcode") or "",
            "product_name": item.get("product_name") or "",
            "agency_name": agency_name,
            "brand": str(getattr(sku, "brand", "") or "").strip(),
            "subject": str(
                getattr(sku, "vid_tovar", "") or getattr(sku, "tovar_category", "") or ""
            ).strip(),
            "color": str(getattr(sku, "color", "") or "").strip(),
            "composition": str(getattr(sku, "composition", "") or "").strip(),
            "supplier": supplier_name,
            "country": str(getattr(sku, "made_in", "") or "").strip(),
            "manufacture_date": _label_date(getattr(sku, "cr_product_date", None)),
            "expiration_date": _label_date(getattr(sku, "end_product_date", None)),
            "label_template": _duplicate_label_template(),
            "order_type": context["order_type"],
            "order_id": context["order_id"],
            "reservation_job_id": reservation_job_id,
        }
    )


def free_marking_queue_response(*, request):
    ok, response = _views()._require_free_print_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    context, error_response = _free_print_selected_item(data=data)
    if error_response:
        return error_response
    try:
        reservation_job_id = int(data.get("reservation_job_id") or 0)
    except (TypeError, ValueError):
        reservation_job_id = 0
    if reservation_job_id <= 0:
        return JsonResponse(
            {"ok": False, "error": "Сначала зарезервируйте всю партию КИЗ."},
            status=409,
        )
    request_id = str(data.get("request_id") or data.get("requestId") or "").strip()
    if request_id and not re.fullmatch(r"[A-Za-z0-9_-]{16,64}", request_id):
        return JsonResponse(
            {"ok": False, "error": "Некорректный идентификатор операции печати."},
            status=400,
        )
    printer_name = str(data.get("printer_name") or "").strip()
    agent_id = str(data.get("agent_id") or "").strip()
    if not printer_name or not agent_id:
        return JsonResponse(
            {"ok": False, "error": "Выберите активный агент и принтер."},
            status=400,
        )
    raw_labels = data.get("labels")
    if not isinstance(raw_labels, list) or not raw_labels or len(raw_labels) > FREE_PRINT_MAX_BATCH:
        return JsonResponse(
            {"ok": False, "error": f"Передайте от 1 до {FREE_PRINT_MAX_BATCH} этикеток."},
            status=400,
        )
    try:
        copies = int(data.get("copies") or 1)
    except (TypeError, ValueError):
        copies = 0
    if copies < 1 or copies > FREE_PRINT_MAX_COPIES:
        return JsonResponse(
            {"ok": False, "error": f"Укажите количество копий от 1 до {FREE_PRINT_MAX_COPIES}."},
            status=400,
        )
    print_label_count = len(raw_labels) * copies
    if print_label_count > FREE_PRINT_MAX_LABELS_PER_OPERATION:
        return JsonResponse(
            {
                "ok": False,
                "error": (
                    f"За одну операцию можно передать не более "
                    f"{FREE_PRINT_MAX_LABELS_PER_OPERATION} этикеток с учетом копий."
                ),
            },
            status=400,
        )
    labels = []
    seen_codes = set()
    try:
        for row in raw_labels:
            if not isinstance(row, dict):
                raise InvalidHonestSignCode("Некорректный состав этикеток.")
            code = _decode_duplicate_scan(row.get("code") or "")
            if code in seen_codes:
                raise InvalidHonestSignCode("Один КИЗ передан в очередь дважды.")
            seen_codes.add(code)
            base_image = _validate_duplicate_png(row.get("label_png_base64") or "")
            raw_copy_images = row.get("copy_label_png_base64_list")
            if raw_copy_images is None:
                copy_images = [base_image for _copy_number in range(copies)]
            else:
                if not isinstance(raw_copy_images, list) or len(raw_copy_images) != copies:
                    raise InvalidHonestSignCode(
                        "Количество подготовленных этикеток не совпало с количеством копий."
                    )
                copy_images = [
                    _validate_duplicate_png(image or "")
                    for image in raw_copy_images
                ]
            labels.append((code, base_image, copy_images))
    except TrueApiError as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=exc.http_status)
    if (
        sum(len(image) for _code, _base_image, copy_images in labels for image in copy_images)
        > FREE_PRINT_MAX_TOTAL_BASE64_SIZE
    ):
        return JsonResponse({"ok": False, "error": "Пакет этикеток слишком большой."}, status=400)
    print_images = [
        image
        for _code, _base_image, copy_images in labels
        for image in copy_images
    ]
    label_template = _duplicate_label_template()

    from processing_app.services import ProcessingWorkflowService

    scope = _free_print_scope(context)
    _sync_free_print_job_state(scope)
    with transaction.atomic():
        from processing_app.models import ProcessingPrintJob

        reservation_job = (
            ProcessingPrintJob.objects.select_for_update()
            .filter(
                id=reservation_job_id,
                order_id=context["order_id"],
                article=context["item"]["sku_code"],
                barcode=(context["item"].get("barcode") or context["item"]["sku_code"]),
                size=context["item"]["size"],
                processing_param_key=FREE_PRINT_RESERVATION_PARAM_KEY,
            )
            .first()
        )
        if not reservation_job:
            return JsonResponse(
                {"ok": False, "error": "Резерв полной партии КИЗ не найден или истёк."},
                status=409,
            )
        locked = {
            row.code: row
            for row in scope.select_for_update().filter(
                code__in=[code for code, _base_image, _copy_images in labels]
            )
        }
        if request_id:
            prior_audit = (
                AuditEntry.objects.filter(
                    journal__code="free_marking_print",
                    action="print",
                    snapshot__request_id=request_id,
                )
                .order_by("-id")
                .first()
            )
            if prior_audit:
                snapshot = (
                    prior_audit.snapshot if isinstance(prior_audit.snapshot, dict) else {}
                )
                requested_code_hashes = [
                    hashlib.sha256(code.encode("utf-8")).hexdigest()
                    for code, _base_image, _copy_images in labels
                ]
                same_request = all(
                    (
                        str(snapshot.get("order_type") or "") == context["order_type"],
                        str(snapshot.get("order_id") or "") == context["order_id"],
                        str(snapshot.get("sku_code") or "") == context["item"]["sku_code"],
                        str(snapshot.get("size") or "") == context["item"]["size"],
                        str(snapshot.get("barcode") or "")
                        == (context["item"].get("barcode") or ""),
                        int(snapshot.get("count") or 0) == len(labels),
                        int(snapshot.get("copies_per_kiz") or 0) == copies,
                        str(snapshot.get("printer_name") or "") == printer_name,
                        str(snapshot.get("agent_id") or "") == agent_id,
                        list(snapshot.get("code_hashes") or []) == requested_code_hashes,
                    )
                )
                job_ids = [
                    int(value)
                    for value in (snapshot.get("print_job_ids") or [])
                    if str(value).isdigit() and int(value) > 0
                ]
                jobs = list(ProcessingPrintJob.objects.filter(id__in=job_ids))
                if not same_request or not job_ids or len(jobs) != len(set(job_ids)):
                    return JsonResponse(
                        {
                            "ok": False,
                            "error": (
                                "Идентификатор операции печати уже использован с другими "
                                "параметрами. Обновите страницу и повторите операцию."
                            ),
                        },
                        status=409,
                    )
                non_failed_job_ids = {
                    job.id
                    for job in jobs
                    if job.status != ProcessingPrintJob.STATUS_FAILED
                }
                linked_codes = set(
                    scope.filter(print_job_id__in=non_failed_job_ids).values_list(
                        "code", flat=True
                    )
                )
                requested_codes = {
                    code for code, _base_image, _copy_images in labels
                }
                if non_failed_job_ids and linked_codes != requested_codes:
                    return JsonResponse(
                        {
                            "ok": False,
                            "error": "Сохраненная операция печати больше не соответствует выбранным КИЗ.",
                        },
                        status=409,
                    )
                return JsonResponse(
                    {
                        "ok": True,
                        "job_id": job_ids[0],
                        "job_ids": job_ids,
                        "queued_count": int(snapshot.get("print_label_count") or 0),
                        "label_count": len(labels),
                        "agent_id": agent_id,
                        "kiz_count": len(labels),
                        "copies_per_kiz": copies,
                        "recovered": True,
                        "message": "Задание уже было создано; восстановлен его результат.",
                        "stats": _free_print_stats(context),
                    }
                )
        if len(locked) != len(labels) or any(
            locked[code].used_at
            or locked[code].printed_at
            or locked[code].print_job_id != reservation_job_id
            for code, _base_image, _copy_images in labels
            if code in locked
        ):
            return JsonResponse(
                {
                    "ok": False,
                    "error": (
                        "Часть КИЗ уже не относится к зарезервированной партии. "
                        "Новые задания печати не создавались."
                    ),
                },
                status=409,
            )
        print_result = ProcessingWorkflowService.enqueue_processing_print_job(
            request=request,
            data={
                "order_id": context["order_id"],
                "card_id": "",
                "article": context["item"]["sku_code"],
                "barcode": context["item"].get("barcode") or context["item"]["sku_code"],
                "size": context["item"]["size"],
                "printer_name": printer_name,
                "agent_id": agent_id,
                "copies": 1,
                "processing_param_key": FREE_PRINT_PROCESSING_PARAM_KEY,
                "label_png_base64_list": print_images,
                "label_width_mm": label_template["width_mm"],
                "label_height_mm": label_template["height_mm"],
            },
        )
        if print_result.http_status != 200 or not print_result.payload.get("ok"):
            return JsonResponse(print_result.payload, status=print_result.http_status)
        if int(print_result.payload.get("queued_count") or 0) != print_label_count:
            raise RuntimeError("Сервис печати вернул неверное количество копий этикеток.")
        job_ids = [int(value) for value in (print_result.payload.get("job_ids") or [])]
        if not job_ids:
            raise RuntimeError("Сервис печати не вернул номер задания.")
        if len(job_ids) == 1:
            code_job_ids = job_ids * len(labels)
        elif len(job_ids) == len(labels):
            code_job_ids = job_ids
        else:
            raise RuntimeError("Количество заданий печати не совпало с количеством КИЗ.")
        now = timezone.now()
        updated_codes = []
        for (code, _base_image, _copy_images), job_id in zip(labels, code_job_ids):
            marking_code = locked[code]
            marking_code.print_job_id = job_id
            marking_code.print_reserved_at = now
            marking_code.printed_by = request.user if request.user.is_authenticated else None
            updated_codes.append(marking_code)
        MarkingCode.objects.bulk_update(
            updated_codes,
            ["print_job_id", "print_reserved_at", "printed_by"],
        )

        journal, _created = AuditJournal.objects.get_or_create(
            code="free_marking_print",
            defaults={
                "name": "Свободная печать КИЗ",
                "description": "Предварительная печать КИЗ с привязкой к заявке и товару.",
            },
        )
        AuditEntry.objects.create(
            journal=journal,
            action="print",
            sku=context["item"].get("sku"),
            agency=context.get("agency"),
            user=request.user if request.user.is_authenticated else None,
            description=(
                f"В очередь передано {len(labels)} КИЗ, копий каждого: {copies}: "
                f"{context['order_type']} {context['order_id']} / {context['item']['sku_code']}"
            ),
            snapshot={
                "order_type": context["order_type"],
                "order_id": context["order_id"],
                "sku_code": context["item"]["sku_code"],
                "size": context["item"]["size"],
                "barcode": context["item"].get("barcode") or "",
                "count": len(labels),
                "copies_per_kiz": copies,
                "print_label_count": print_label_count,
                "printer_name": printer_name,
                "agent_id": agent_id,
                "print_job_ids": job_ids,
                "request_id": request_id,
                "reservation_job_id": reservation_job_id,
                "code_hashes": [
                    hashlib.sha256(code.encode("utf-8")).hexdigest()
                    for code, _base_image, _copy_images in labels
                ],
            },
        )
    return JsonResponse(
        {
            **print_result.payload,
            "kiz_count": len(labels),
            "copies_per_kiz": copies,
            "message": (
                f"В очередь передано {len(labels)} КИЗ, "
                f"этикеток с учетом копий: {print_label_count}."
            ),
            "stats": _free_print_stats(context),
        }
    )


def free_marking_batch_status_response(*, request):
    ok, response = _views()._require_free_print_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    context, error_response = _free_print_selected_item(data=data)
    if error_response:
        return error_response
    job_ids = _free_print_job_ids(data)
    if not job_ids:
        return JsonResponse({"ok": False, "error": "Не указана партия печати."}, status=400)

    from processing_app.models import ProcessingPrintJob

    rows = list(
        ProcessingPrintJob.objects.filter(id__in=job_ids)
        .order_by("id")
        .values(
            "id",
            "status",
            "error",
            "order_id",
            "article",
            "barcode",
            "size",
            "processing_param_key",
        )
    )
    if len(rows) != len(job_ids):
        return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
    linked_job_ids = set(
        _free_print_scope(context)
        .filter(print_job_id__in=job_ids)
        .values_list("print_job_id", flat=True)
    )
    item = context["item"]
    allowed_failed_job_ids = {
        int(row["id"])
        for row in rows
        if row["status"] == ProcessingPrintJob.STATUS_FAILED
        and str(row["order_id"] or "") == context["order_id"]
        and str(row["article"] or "") == item["sku_code"]
        and str(row["barcode"] or "") == (item.get("barcode") or item["sku_code"])
        and str(row["size"] or "") == item["size"]
    }
    if linked_job_ids | allowed_failed_job_ids != set(job_ids) or any(
        row["processing_param_key"] != FREE_PRINT_PROCESSING_PARAM_KEY for row in rows
    ):
        return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
    terminal_statuses = {ProcessingPrintJob.STATUS_PRINTED, ProcessingPrintJob.STATUS_FAILED}
    return JsonResponse(
        {
            "ok": True,
            "terminal": all(row["status"] in terminal_statuses for row in rows),
            "failed": sum(row["status"] == ProcessingPrintJob.STATUS_FAILED for row in rows),
            "jobs": rows,
            "stats": _free_print_stats(context),
        }
    )


def free_marking_confirm_printed_response(*, request):
    ok, response = _views()._require_free_print_role(request)
    if not ok:
        return response
    data = _views()._parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    context, error_response = _free_print_selected_item(data=data)
    if error_response:
        return error_response
    job_ids = _free_print_job_ids(data)
    if not job_ids:
        return JsonResponse({"ok": False, "error": "Не указана партия печати."}, status=400)
    printed = data.get("printed") is True or str(data.get("printed") or "").strip().lower() in {
        "1", "true", "yes", "да",
    }
    uncertain = data.get("uncertain") is True or str(
        data.get("uncertain") or ""
    ).strip().lower() in {"1", "true", "yes", "да"}
    if printed and uncertain:
        return JsonResponse(
            {"ok": False, "error": "Нельзя одновременно подтвердить и оспорить печать."},
            status=400,
        )

    from processing_app.models import ProcessingPrintJob

    with transaction.atomic():
        jobs = list(
            ProcessingPrintJob.objects.select_for_update()
            .filter(id__in=job_ids)
            .order_by("id")
        )
        if len(jobs) != len(job_ids):
            return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
        reservation_jobs = [
            job
            for job in jobs
            if job.processing_param_key == FREE_PRINT_RESERVATION_PARAM_KEY
        ]
        if reservation_jobs:
            if len(reservation_jobs) != len(jobs) or printed:
                return JsonResponse(
                    {"ok": False, "error": "Некорректное подтверждение резерва печати."},
                    status=409,
                )
            codes = (
                _free_print_scope(context)
                .select_for_update()
                .filter(print_job_id__in=job_ids, used_at__isnull=True)
            )
            confirmed_count = codes.update(
                printed_at=None,
                printed_by=None,
                print_job_id=None,
                print_reserved_at=None,
            )
            now = timezone.now()
            ProcessingPrintJob.objects.filter(id__in=job_ids).update(
                status=ProcessingPrintJob.STATUS_FAILED,
                processing_param_key=FREE_PRINT_RETURNED_PARAM_KEY,
                error="Не переданная агенту часть резерва печати освобождена.",
                updated_at=now,
            )
            journal, _created = AuditJournal.objects.get_or_create(
                code="free_marking_print",
                defaults={
                    "name": "Свободная печать КИЗ",
                    "description": "Предварительная печать КИЗ с привязкой к заявке и товару.",
                },
            )
            AuditEntry.objects.create(
                journal=journal,
                action="release_reservation",
                sku=context["item"].get("sku"),
                agency=context.get("agency"),
                user=request.user if request.user.is_authenticated else None,
                description=(
                    f"Освобождена не переданная агенту часть резерва: {confirmed_count} КИЗ; "
                    f"{context['order_type']} {context['order_id']} / "
                    f"{context['item']['sku_code']}"
                ),
                snapshot={
                    "order_type": context["order_type"],
                    "order_id": context["order_id"],
                    "sku_code": context["item"]["sku_code"],
                    "size": context["item"]["size"],
                    "barcode": context["item"].get("barcode") or "",
                    "count": confirmed_count,
                    "printed": False,
                    "reservation_job_ids": job_ids,
                },
            )
            return JsonResponse(
                {
                    "ok": True,
                    "printed": False,
                    "count": confirmed_count,
                    "reservation_released": True,
                    "stats": _free_print_stats(context),
                }
            )
        if any(job.processing_param_key != FREE_PRINT_PROCESSING_PARAM_KEY for job in jobs):
            return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
        if any(job.status != ProcessingPrintJob.STATUS_PRINTED for job in jobs):
            return JsonResponse(
                {"ok": False, "error": "Агент печати еще не подтвердил завершение партии."},
                status=409,
            )
        codes = _free_print_scope(context).select_for_update().filter(print_job_id__in=job_ids)
        linked_job_ids = set(codes.values_list("print_job_id", flat=True))
        if linked_job_ids != set(job_ids):
            return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
        # Once the agent has reported a print job as completed, the server can no
        # longer prove that no physical label came out.  Never release those KIZ
        # back into the printable pool from the operator confirmation dialog:
        # even an explicit "none printed" response must quarantine the batch
        # until a manager scans each physical label in the Return KIZ workflow.
        if not printed:
            uncertain = True
        now = timezone.now()
        if printed:
            confirmed_count = codes.update(
                printed_at=now,
                printed_by=request.user if request.user.is_authenticated else None,
            )
            ProcessingPrintJob.objects.filter(id__in=job_ids).update(
                processing_param_key=FREE_PRINT_CONFIRMED_PARAM_KEY,
                updated_at=now,
            )
        else:
            confirmed_count = codes.count()
            ProcessingPrintJob.objects.filter(id__in=job_ids).update(
                status=ProcessingPrintJob.STATUS_FAILED,
                error=(
                    "Оператор сообщил о неполной или неопределённой печати. "
                    "КИЗ заблокированы до проверки руководителем."
                ),
                updated_at=now,
            )

        journal, _created = AuditJournal.objects.get_or_create(
            code="free_marking_print",
            defaults={
                "name": "Свободная печать КИЗ",
                "description": "Предварительная печать КИЗ с привязкой к заявке и товару.",
            },
        )
        AuditEntry.objects.create(
            journal=journal,
            action="confirm_print" if printed else "uncertain_print",
            sku=context["item"].get("sku"),
            agency=context.get("agency"),
            user=request.user if request.user.is_authenticated else None,
            description=(
                f"Оператор подтвердил: КИЗ "
                f"{'напечатались' if printed else 'требуют проверки'}; "
                f"{context['order_type']} {context['order_id']} / {context['item']['sku_code']}"
            ),
            snapshot={
                "order_type": context["order_type"],
                "order_id": context["order_id"],
                "sku_code": context["item"]["sku_code"],
                "size": context["item"]["size"],
                "barcode": context["item"].get("barcode") or "",
                "count": confirmed_count,
                "printed": printed,
                "uncertain": uncertain,
                "print_job_ids": job_ids,
            },
        )
    return JsonResponse(
        {
            "ok": True,
            "printed": printed,
            "uncertain": uncertain,
            "count": confirmed_count,
            "stats": _free_print_stats(context),
        }
    )


def processing_marking_summary_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, _agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    rows = (
        MarkingCode.objects.filter(order_type="processing", order_id=order_id, used_at__isnull=False)
        .values("sku_code", "size")
        .annotate(count=Count("id"))
    )
    items = [
        {"sku_code": row["sku_code"], "size": row["size"] or "", "count": row["count"]}
        for row in rows
    ]
    total_count = sum(row["count"] for row in rows)
    return JsonResponse({"ok": True, "items": items, "total_count": total_count})


@transaction.atomic
def processing_marking_scan_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    code = normalize_marking_code(data.get("code") or "")
    sku_code = (data.get("sku_code") or "").strip()
    size = (data.get("size") or "").strip()
    barcode = (data.get("barcode") or "").strip()
    box_barcode = (data.get("box_barcode") or "").strip()
    if not code:
        return JsonResponse({"ok": False, "error": "Код ЧЗ не указан"}, status=400)
    if not sku_code:
        return JsonResponse({"ok": False, "error": "Артикул не указан"}, status=400)
    items = extract_processing_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    if not size:
        matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
        if len(matched) == 1:
            size = matched[0][1]
    if (sku_code, size) not in allowed_pairs:
        if size and (sku_code, "") in allowed_pairs:
            size = ""
        else:
            return JsonResponse(
                {"ok": False, "error": "Позиция не найдена в заявке."},
                status=400,
            )
    now = timezone.localtime()
    existing = (
        MarkingCode.objects.select_for_update()
        .select_related("agency", "sku")
        .filter(
            Q(identity_key=marking_code_identity(code))
            | Q(code__in=marking_code_variants(code))
        )
        .first()
    )
    if existing:
        if existing.used_at:
            return JsonResponse({"ok": False, "error": "Код уже использован."}, status=409)
        if agency and existing.agency_id and existing.agency_id != agency.id:
            return JsonResponse({"ok": False, "error": "Код принадлежит другому клиенту."}, status=409)
        if existing.order_type and existing.order_type != "processing":
            return JsonResponse({"ok": False, "error": "Код закреплен в другом процессе."}, status=409)
        if existing.order_id and existing.order_id != order_id:
            return JsonResponse({"ok": False, "error": "Код закреплен за другой заявкой."}, status=409)
        if existing.sku_code and existing.sku_code != sku_code:
            return JsonResponse({"ok": False, "error": "Код относится к другому артикулу."}, status=409)
        if existing.size and size and existing.size != size:
            return JsonResponse({"ok": False, "error": "Код относится к другому размеру."}, status=409)
        if existing.box_barcode and box_barcode and existing.box_barcode != box_barcode:
            return JsonResponse({"ok": False, "error": "Код закреплен за другим коробом."}, status=409)
        update_fields = []
        if not existing.order_id:
            existing.order_id = order_id
            update_fields.append("order_id")
        if not existing.order_type:
            existing.order_type = "processing"
            update_fields.append("order_type")
        if not existing.size and size:
            existing.size = size
            update_fields.append("size")
        if not existing.barcode and barcode:
            existing.barcode = barcode
            update_fields.append("barcode")
        if box_barcode and not existing.box_barcode:
            existing.box_barcode = box_barcode
            update_fields.append("box_barcode")
        if not existing.sku:
            sku = _views()._resolve_sku(agency, sku_code)
            existing.sku = sku
            update_fields.append("sku")
        existing.used_at = now
        existing.used_by = request.user if request.user.is_authenticated else None
        update_fields.extend(["used_at", "used_by"])
        existing.save(update_fields=update_fields)
    else:
        sku = _views()._resolve_sku(agency, sku_code)
        try:
            MarkingCode.objects.create(
                order_type="processing",
                order_id=order_id,
                agency=agency,
                sku=sku,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                box_barcode=box_barcode,
                code=code,
                source="scan",
                created_by=request.user if request.user.is_authenticated else None,
                used_at=now,
                used_by=request.user if request.user.is_authenticated else None,
            )
        except IntegrityError:
            return JsonResponse({"ok": False, "error": "Код уже учтен."}, status=409)
    count = MarkingCode.objects.filter(
        order_type="processing",
        order_id=order_id,
        sku_code=sku_code,
        size=size,
        used_at__isnull=False,
    ).count()
    total_count = MarkingCode.objects.filter(
        order_type="processing",
        order_id=order_id,
        used_at__isnull=False,
    ).count()
    return JsonResponse(
        {
            "ok": True,
            "sku_code": sku_code,
            "size": size,
            "count": count,
            "total_count": total_count,
        }
    )


@transaction.atomic
def receiving_marking_scan_response(*, request, order_id: str):
    ok, response = _views()._require_receiving_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_receiving_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    code = normalize_marking_code(data.get("code") or "")
    sku_code = (data.get("sku_code") or "").strip()
    size = (data.get("size") or "").strip()
    barcode = (data.get("barcode") or "").strip()
    box_barcode = (data.get("box_barcode") or "").strip()
    if not code:
        return JsonResponse({"ok": False, "error": "Код ЧЗ не указан"}, status=400)
    if not sku_code:
        return JsonResponse({"ok": False, "error": "Артикул не указан"}, status=400)
    if not box_barcode:
        return JsonResponse({"ok": False, "error": "Откройте короб перед сканированием ЧЗ."}, status=400)

    items = _views()._extract_receiving_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    if not size:
        matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
        if len(matched) == 1:
            size = matched[0][1]
    if (sku_code, size) not in allowed_pairs:
        if size and (sku_code, "") in allowed_pairs:
            size = ""
        else:
            return JsonResponse(
                {"ok": False, "error": "Позиция не найдена в заявке."},
                status=400,
            )

    sku = _views()._resolve_sku(agency, sku_code)
    order_cz_mode = str(payload.get("receiving_mode") or "").strip().lower() == "cz"
    if not order_cz_mode and (not sku or not sku.honest_sign):
        return JsonResponse(
            {"ok": False, "error": "Для этой позиции не включен режим Честного знака."},
            status=400,
        )

    now = timezone.localtime()
    existing = (
        MarkingCode.objects.select_for_update()
        .select_related("agency", "sku")
        .filter(
            Q(identity_key=marking_code_identity(code))
            | Q(code__in=marking_code_variants(code))
        )
        .first()
    )
    if existing:
        if existing.used_at:
            return JsonResponse({"ok": False, "error": "Код уже использован."}, status=409)
        if agency and existing.agency_id and existing.agency_id != agency.id:
            return JsonResponse({"ok": False, "error": "Код принадлежит другому клиенту."}, status=409)
        if existing.order_type and existing.order_type != "receiving":
            return JsonResponse({"ok": False, "error": "Код закреплен в другом процессе."}, status=409)
        if existing.order_id and existing.order_id != order_id:
            return JsonResponse({"ok": False, "error": "Код закреплен за другой заявкой."}, status=409)
        if existing.sku_code and existing.sku_code != sku_code:
            return JsonResponse({"ok": False, "error": "Код относится к другому артикулу."}, status=409)
        if existing.size and size and existing.size != size:
            return JsonResponse({"ok": False, "error": "Код относится к другому размеру."}, status=409)
        update_fields = []
        if not existing.order_id:
            existing.order_id = order_id
            update_fields.append("order_id")
        if not existing.order_type:
            existing.order_type = "receiving"
            update_fields.append("order_type")
        if not existing.size and size:
            existing.size = size
            update_fields.append("size")
        if not existing.barcode and barcode:
            existing.barcode = barcode
            update_fields.append("barcode")
        if existing.box_barcode != box_barcode:
            existing.box_barcode = box_barcode
            update_fields.append("box_barcode")
        if not existing.sku:
            existing.sku = sku
            update_fields.append("sku")
        existing.used_at = now
        existing.used_by = request.user if request.user.is_authenticated else None
        update_fields.extend(["used_at", "used_by"])
        existing.save(update_fields=update_fields)
    else:
        try:
            MarkingCode.objects.create(
                order_type="receiving",
                order_id=order_id,
                agency=agency,
                sku=sku,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                box_barcode=box_barcode,
                code=code,
                source="scan",
                created_by=request.user if request.user.is_authenticated else None,
                used_at=now,
                used_by=request.user if request.user.is_authenticated else None,
            )
        except IntegrityError:
            return JsonResponse({"ok": False, "error": "Код уже учтен."}, status=409)

    count = MarkingCode.objects.filter(
        order_type="receiving",
        order_id=order_id,
        sku_code=sku_code,
        size=size,
        used_at__isnull=False,
    ).count()
    total_count = MarkingCode.objects.filter(
        order_type="receiving",
        order_id=order_id,
        used_at__isnull=False,
    ).count()
    return JsonResponse(
        {
            "ok": True,
            "sku_code": sku_code,
            "size": size,
            "count": count,
            "total_count": total_count,
        }
    )


def _processing_marking_print_scope(*, order_id: str, agency, barcode: str, size: str):
    scope = MarkingCode.objects.filter(
        order_type="processing",
        used_at__isnull=True,
        printed_at__isnull=True,
        print_job_id__isnull=True,
        barcode=barcode,
    )
    if agency:
        scope = scope.filter(agency=agency)
    if size:
        scope = scope.filter(size=size)
    return scope.filter(Q(order_id=order_id) | Q(order_id="") | Q(order_id__isnull=True))


def _processing_print_job_ids(data: dict) -> list[int]:
    raw_values = data.get("job_ids")
    if raw_values is None:
        raw_values = data.get("jobIds")
    if not isinstance(raw_values, (list, tuple)):
        raw_values = []
    result: list[int] = []
    for raw_value in raw_values[:100]:
        try:
            value = int(raw_value)
        except (TypeError, ValueError):
            continue
        if value > 0 and value not in result:
            result.append(value)
    return result


def processing_marking_print_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    qty = data.get("qty")
    try:
        qty_value = int(qty)
    except (TypeError, ValueError):
        qty_value = 0
    if qty_value <= 0:
        return JsonResponse({"ok": False, "error": "Количество для печати не указано."}, status=400)
    barcode = (data.get("barcode") or "").strip()
    size = (data.get("size") or "").strip()
    if not barcode:
        return JsonResponse({"ok": False, "error": "ШК не указан."}, status=400)
    base_qs = _processing_marking_print_scope(
        order_id=order_id,
        agency=agency,
        barcode=barcode,
        size=size,
    )
    assigned = list(
        base_qs.filter(order_id=order_id)
        .order_by("created_at", "id")
        .values_list("code", flat=True)[:qty_value]
    )
    remaining = qty_value - len(assigned)
    extra = []
    if remaining > 0:
        extra = list(
            base_qs.filter(Q(order_id__isnull=True) | Q(order_id=""))
            .order_by("created_at", "id")
            .values_list("code", flat=True)[:remaining]
        )
    codes = assigned + extra
    if len(codes) < qty_value:
        return JsonResponse(
            {"ok": False, "error": "Недостаточно свободных кодов ЧЗ.", "available": len(codes)},
            status=409,
        )
    return JsonResponse({"ok": True, "codes": codes, "count": len(codes)})


def processing_marking_queue_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    barcode = str(data.get("barcode") or "").strip()
    size = str(data.get("size") or "").strip()
    article = str(data.get("article") or "").strip()
    printer_name = str(data.get("printer_name") or "").strip()
    agent_id = str(data.get("agent_id") or "").strip()
    raw_labels = data.get("labels")
    if not barcode:
        return JsonResponse({"ok": False, "error": "ШК не указан."}, status=400)
    if not printer_name or not agent_id:
        return JsonResponse({"ok": False, "error": "Выберите активный агент и принтер."}, status=400)
    if not isinstance(raw_labels, list) or not raw_labels or len(raw_labels) > FREE_PRINT_MAX_BATCH:
        return JsonResponse(
            {"ok": False, "error": f"Передайте от 1 до {FREE_PRINT_MAX_BATCH} этикеток."},
            status=400,
        )

    labels: list[tuple[str, str]] = []
    seen_codes: set[str] = set()
    try:
        for row in raw_labels:
            if not isinstance(row, dict):
                raise InvalidHonestSignCode("Некорректный состав этикеток.")
            code = _decode_duplicate_scan(row.get("code") or "")
            if code in seen_codes:
                raise InvalidHonestSignCode("Один код ЧЗ передан в партию дважды.")
            seen_codes.add(code)
            labels.append((code, _validate_duplicate_png(row.get("label_png_base64") or "")))
    except TrueApiError as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=exc.http_status)
    if sum(len(image) for _code, image in labels) > FREE_PRINT_MAX_TOTAL_BASE64_SIZE:
        return JsonResponse({"ok": False, "error": "Пакет этикеток слишком большой."}, status=400)

    from processing_app.services import ProcessingWorkflowService

    scope = _processing_marking_print_scope(
        order_id=order_id,
        agency=agency,
        barcode=barcode,
        size=size,
    )
    with transaction.atomic():
        locked = {
            row.code: row
            for row in scope.select_for_update().filter(code__in=[code for code, _image in labels])
        }
        if len(locked) != len(labels) or any(code not in locked for code, _image in labels):
            return JsonResponse(
                {
                    "ok": False,
                    "error": (
                        "Часть кодов ЧЗ уже занята другой партией, напечатана, "
                        "использована или относится к другому товару. Повторите печать."
                    ),
                },
                status=409,
            )
        print_result = ProcessingWorkflowService.enqueue_processing_print_job(
            request=request,
            data={
                "order_id": order_id,
                "card_id": str(data.get("card_id") or "").strip(),
                "article": article,
                "barcode": barcode,
                "size": size,
                "printer_name": printer_name,
                "agent_id": agent_id,
                "copies": 1,
                "processing_param_key": str(data.get("processing_param_key") or "").strip(),
                "template_key": str(data.get("template_key") or "item_cz").strip() or "item_cz",
                "label_png_base64_list": [image for _code, image in labels],
                "label_width_mm": data.get("label_width_mm") or 58,
                "label_height_mm": data.get("label_height_mm") or 40,
            },
        )
        if print_result.http_status != 200 or not print_result.payload.get("ok"):
            return JsonResponse(print_result.payload, status=print_result.http_status)
        job_ids = [int(value) for value in (print_result.payload.get("job_ids") or [])]
        if len(job_ids) != len(labels):
            raise RuntimeError("Количество заданий печати не совпало с количеством кодов ЧЗ.")
        now = timezone.now()
        updated_codes = []
        for (code, _image), job_id in zip(labels, job_ids):
            marking_code = locked[code]
            marking_code.order_id = order_id
            marking_code.print_job_id = job_id
            marking_code.print_reserved_at = now
            marking_code.printed_by = request.user if request.user.is_authenticated else None
            updated_codes.append(marking_code)
        MarkingCode.objects.bulk_update(
            updated_codes,
            ["order_id", "print_job_id", "print_reserved_at", "printed_by"],
        )
    return JsonResponse(
        {
            **print_result.payload,
            "ok": True,
            "job_ids": job_ids,
            "count": len(labels),
            "message": f"В очередь передано {len(labels)} этикеток ЧЗ.",
        }
    )


def processing_marking_batch_status_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, _agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    job_ids = _processing_print_job_ids(data)
    if not job_ids:
        return JsonResponse({"ok": False, "error": "Не указана партия печати."}, status=400)

    from processing_app.models import ProcessingPrintJob

    rows = list(
        ProcessingPrintJob.objects.filter(id__in=job_ids, order_id=order_id)
        .order_by("id")
        .values("id", "status", "error")
    )
    if len(rows) != len(job_ids):
        return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
    terminal_statuses = {ProcessingPrintJob.STATUS_PRINTED, ProcessingPrintJob.STATUS_FAILED}
    return JsonResponse(
        {
            "ok": True,
            "terminal": all(row["status"] in terminal_statuses for row in rows),
            "failed": sum(row["status"] == ProcessingPrintJob.STATUS_FAILED for row in rows),
            "jobs": rows,
        }
    )


def processing_marking_confirm_printed_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    job_ids = _processing_print_job_ids(data)
    if not job_ids:
        return JsonResponse({"ok": False, "error": "Не указана партия печати."}, status=400)
    printed = data.get("printed") is True or str(data.get("printed") or "").strip().lower() in {
        "1", "true", "yes", "да",
    }

    from processing_app.models import ProcessingPrintJob

    with transaction.atomic():
        jobs = list(
            ProcessingPrintJob.objects.select_for_update()
            .filter(id__in=job_ids, order_id=order_id)
            .order_by("id")
        )
        if len(jobs) != len(job_ids):
            return JsonResponse({"ok": False, "error": "Партия печати не найдена."}, status=404)
        codes = MarkingCode.objects.select_for_update().filter(
            order_type="processing",
            order_id=order_id,
            print_job_id__in=job_ids,
        )
        if agency:
            codes = codes.filter(agency=agency)
        now = timezone.now()
        if printed:
            confirmed_count = codes.update(
                printed_at=now,
                printed_by=request.user if request.user.is_authenticated else None,
            )
            ProcessingPrintJob.objects.filter(id__in=job_ids).update(
                status=ProcessingPrintJob.STATUS_PRINTED,
                error="",
                updated_at=now,
            )
        else:
            confirmed_count = codes.filter(used_at__isnull=True).update(
                printed_at=None,
                printed_by=None,
                print_job_id=None,
                print_reserved_at=None,
            )
            ProcessingPrintJob.objects.filter(id__in=job_ids).update(
                status=ProcessingPrintJob.STATUS_FAILED,
                error="Оператор подтвердил, что партия этикеток не распечаталась.",
                updated_at=now,
            )
    return JsonResponse({"ok": True, "printed": printed, "count": confirmed_count})


def processing_marking_reset_printed_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    barcode = (data.get("barcode") or "").strip()
    size = (data.get("size") or "").strip()
    raw_codes = data.get("codes")
    exact_codes = []
    if isinstance(raw_codes, (list, tuple)):
        for raw_code in raw_codes[:FREE_PRINT_MAX_BATCH]:
            try:
                code = _decode_duplicate_scan(raw_code or "")
            except TrueApiError:
                continue
            if code and code not in exact_codes:
                exact_codes.append(code)
    qty = data.get("qty")
    try:
        qty_value = int(qty)
    except (TypeError, ValueError):
        qty_value = 0
    if not barcode:
        return JsonResponse({"ok": False, "error": "ШК не указан."}, status=400)
    with transaction.atomic():
        qs = MarkingCode.objects.select_for_update().filter(
            order_type="processing",
            order_id=order_id,
            used_at__isnull=True,
            barcode=barcode,
        ).filter(Q(printed_at__isnull=False) | Q(print_job_id__isnull=False))
        if agency:
            qs = qs.filter(agency=agency)
        if size:
            qs = qs.filter(size=size)
        if exact_codes:
            ids = list(qs.filter(code__in=exact_codes).values_list("id", flat=True))
        elif qty_value > 0:
            ids = list(
                qs.order_by("-printed_at", "-created_at")
                .values_list("id", flat=True)[:qty_value]
            )
        else:
            ids = list(qs.values_list("id", flat=True))
        if ids:
            from processing_app.models import ProcessingPrintJob

            job_ids = list(
                MarkingCode.objects.filter(id__in=ids, print_job_id__isnull=False)
                .values_list("print_job_id", flat=True)
            )
            MarkingCode.objects.filter(id__in=ids).update(
                printed_at=None,
                printed_by=None,
                print_job_id=None,
                print_reserved_at=None,
            )
            if job_ids:
                ProcessingPrintJob.objects.filter(
                    id__in=job_ids,
                    status__in=[
                        ProcessingPrintJob.STATUS_PENDING,
                        ProcessingPrintJob.STATUS_PRINTING,
                    ],
                ).update(
                    status=ProcessingPrintJob.STATUS_FAILED,
                    error="Печать кодов ЧЗ отменена оператором.",
                    updated_at=timezone.now(),
                )
    return JsonResponse({"ok": True, "count": len(ids)})


def processing_marking_import_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    file = request.FILES.get("file")
    if not file:
        return JsonResponse({"ok": False, "error": "Файл не выбран."}, status=400)
    from processing_app.web_ui import _import_marking_code_files

    imported, result = _import_marking_code_files(
        [file],
        payload,
        order_id,
        agency,
        request.user,
    )
    return JsonResponse({"ok": imported, **result}, status=200 if imported else 400)
