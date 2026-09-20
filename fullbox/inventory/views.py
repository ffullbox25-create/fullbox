from io import BytesIO

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db.models import Count, Prefetch, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.views import View
from django.views.generic import FormView, TemplateView
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from employees.access import RoleRequiredMixin
from sku.models import SKU

from .forms import InventoryCreateForm, SkuChoiceField
from .models import Inventory, InventoryLine, InventoryLocation
from .services import cancel_inventory, transfer_to_work


MANAGER_ROLES = ("head_manager", "director", "admin")


class InventoryManagerMixin(RoleRequiredMixin):
    allowed_roles = MANAGER_ROLES

    def base_context(self) -> dict:
        return {
            "sidebar_active": "inventory",
            "user_display_name": self.request.user.get_full_name() or self.request.user.username,
        }


class InventoryListView(InventoryManagerMixin, TemplateView):
    template_name = "inventory/list.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from reachtruck_inventory.services import release_expired_tasks

        release_expired_tasks()
        queryset = Inventory.objects.select_related("agency", "sku", "created_by").annotate(
            location_count=Count("scope_locations", distinct=True),
            task_count=Count("execution_tasks", distinct=True),
        )
        status_filter = str(self.request.GET.get("status") or "").strip()
        if status_filter in dict(Inventory.STATUS_CHOICES):
            queryset = queryset.filter(status=status_filter)
        context.update(self.base_context())
        context.update(
            {
                "inventories": queryset[:200],
                "status_filter": status_filter,
                "status_choices": Inventory.STATUS_CHOICES,
                "created_count": Inventory.objects.filter(status=Inventory.STATUS_CREATED).count(),
                "pending_count": Inventory.objects.filter(status=Inventory.STATUS_PENDING).count(),
                "in_progress_count": Inventory.objects.filter(status=Inventory.STATUS_IN_PROGRESS).count(),
                "completed_count": Inventory.objects.filter(status=Inventory.STATUS_COMPLETED).count(),
            }
        )
        return context


class InventoryCreateView(InventoryManagerMixin, FormView):
    template_name = "inventory/form.html"
    form_class = InventoryCreateForm

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.base_context())
        form = context["form"]
        selected_location_ids = {str(value) for value in (form["locations"].value() or [])}
        location_field = form.fields["locations"]
        context["location_options"] = [
            {
                "id": location.id,
                "key": ":".join(
                    [
                        str(location.zone_code or "").strip().upper(),
                        str(location.row_no or 0),
                        str(location.section_no or 0),
                        str(location.tier_no or 0),
                        str(location.cell_no or 0),
                    ]
                ),
                "label": location_field.label_from_instance(location),
                "selected": str(location.id) in selected_location_ids,
            }
            for location in location_field.queryset
        ]
        return context

    def form_valid(self, form):
        inventory = form.save(commit=False)
        inventory.created_by = self.request.user
        inventory.save()
        locations = list(form.cleaned_data.get("locations") or [])
        if locations:
            InventoryLocation.objects.bulk_create(
                [InventoryLocation(inventory=inventory, location=location) for location in locations],
                ignore_conflicts=True,
            )
        messages.success(self.request, f"Инвентаризация №{inventory.id} создана.")
        return redirect("inventory:detail", pk=inventory.pk)


class InventorySkuSearchView(InventoryManagerMixin, View):
    def get(self, request, *args, **kwargs):
        query = str(request.GET.get("q") or "").strip()
        if len(query) < 2:
            return JsonResponse({"items": []})
        queryset = (
            SKU.objects.filter(deleted=False)
            .filter(
                Q(sku_code__icontains=query)
                | Q(name__icontains=query)
                | Q(size__icontains=query)
                | Q(code__icontains=query)
                | Q(agency__agn_name__icontains=query)
                | Q(barcodes__value__icontains=query)
            )
            .select_related("agency")
            .distinct()
            .order_by("sku_code", "size", "id")[:20]
        )
        label_field = SkuChoiceField(queryset=SKU.objects.none())
        return JsonResponse(
            {
                "items": [
                    {
                        "id": sku.id,
                        "label": label_field.label_from_instance(sku),
                        "code": sku.sku_code,
                        "name": sku.name,
                    }
                    for sku in queryset
                ]
            }
        )


class InventoryDetailView(InventoryManagerMixin, TemplateView):
    template_name = "inventory/detail.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from reachtruck_inventory.services import release_expired_tasks

        release_expired_tasks(inventory_id=kwargs["pk"])
        inventory = get_object_or_404(
            Inventory.objects.select_related("agency", "sku", "created_by").prefetch_related(
                Prefetch(
                    "scope_locations",
                    queryset=InventoryLocation.objects.select_related("location").order_by(
                        "location__zone_code",
                        "location__row_no",
                        "location__section_no",
                        "location__tier_no",
                        "location__cell_no",
                    ),
                ),
                "execution_tasks__assigned_to",
            ),
            pk=kwargs["pk"],
        )
        lines = inventory.lines.select_related("location", "agency", "sku_ref", "counted_by")
        discrepancy_count = (
            inventory.execution_tasks.filter(actual_box_count__isnull=False)
            .exclude(actual_box_count=models_f("planned_box_count"))
            .count()
        )
        context.update(self.base_context())
        context.update(
            {
                "inventory": inventory,
                "inventory_lines": lines,
                "discrepancy_count": discrepancy_count,
                "can_transfer": inventory.status == Inventory.STATUS_CREATED,
                "can_cancel": inventory.status not in (Inventory.STATUS_COMPLETED, Inventory.STATUS_CANCELED),
            }
        )
        return context


def models_f(field_name):
    from django.db.models import F

    return F(field_name)


class InventoryActionView(InventoryManagerMixin, View):
    def post(self, request, pk, *args, **kwargs):
        inventory = get_object_or_404(Inventory, pk=pk)
        action = str(request.POST.get("action") or "").strip()
        try:
            if action == "transfer":
                transfer_to_work(inventory, requested_by=request.user)
                messages.success(request, "Инвентаризация передана ричтракерам.")
            elif action == "cancel":
                cancel_inventory(inventory)
                messages.success(request, "Инвентаризация отменена.")
            else:
                messages.error(request, "Неизвестное действие.")
        except ValidationError as exc:
            messages.error(request, "; ".join(exc.messages))
        return redirect("inventory:detail", pk=inventory.pk)


class InventoryExportView(InventoryManagerMixin, View):
    def get(self, request, pk, *args, **kwargs):
        inventory = get_object_or_404(Inventory, pk=pk)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = f"Инвентаризация {inventory.id}"
        headers = ["Место", "Партнер", "Артикул", "Товар", "Размер", "Штрихкод", "План", "Факт", "Разница"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="F89000")
        for line in inventory.lines.select_related("location", "agency").order_by("location_id", "agency_id", "sku_code"):
            location = line.location_display
            sheet.append(
                [
                    location,
                    line.agency.agn_name,
                    line.sku_code,
                    line.name,
                    line.size,
                    line.barcode,
                    line.planned_qty,
                    line.actual_qty,
                    line.difference,
                ]
            )
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = width
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="inventory-{inventory.id}.xlsx"'
        return response
