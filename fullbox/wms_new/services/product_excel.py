from __future__ import annotations

from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook

from sku.models import Agency

from ..models import WmsNewProduct
from .products import ProductOperationError, create_product, update_product


HEADERS = (
    "ID FBS-NEW",
    "Партнер ID",
    "Партнер",
    "Название",
    "Артикул",
    "Штрихкод",
    "Цвет",
    "Вес, г",
    "Размер",
    "Ширина, см",
    "Глубина, см",
    "Высота, см",
    "Категория",
    "Набор",
    "Изображение",
    "Внутренние заметки",
    "Описание",
)


def export_products(queryset) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товары FBS-NEW"
    sheet.append(HEADERS)
    for product in queryset.select_related("agency").order_by("id"):
        sheet.append(
            (
                product.id,
                product.agency_id,
                str(product.agency),
                product.name,
                product.article,
                product.barcode,
                product.color,
                float(product.weight_grams),
                product.size,
                float(product.width_cm),
                float(product.depth_cm),
                float(product.height_cm),
                product.category,
                "Да" if product.is_bundle else "Нет",
                product.image_url,
                product.internal_notes,
                product.description,
            )
        )
    sheet.freeze_panes = "A2"
    for column, width in {
        "A": 14,
        "B": 13,
        "C": 28,
        "D": 48,
        "E": 22,
        "F": 22,
        "G": 16,
        "H": 12,
        "I": 12,
        "M": 24,
        "O": 36,
        "P": 36,
        "Q": 48,
    }.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _value(row: tuple, index: int):
    return row[index] if index < len(row) else None


def _boolean(value) -> bool:
    return str(value or "").strip().lower() in {"1", "да", "yes", "true", "набор"}


def import_products(uploaded_file, *, actor=None) -> dict[str, int]:
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ProductOperationError("Не удалось прочитать Excel-файл.") from exc
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = tuple(str(value or "").strip() for value in next(rows, ()))
    if header[: len(HEADERS)] != HEADERS:
        raise ProductOperationError("Формат файла не совпадает с экспортом FBS-NEW.")

    created = 0
    updated = 0
    with transaction.atomic():
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            product_id = str(_value(row, 0) or "").strip()
            agency_id = str(_value(row, 1) or "").strip()
            if not agency_id.isdigit():
                raise ProductOperationError(f"Строка {row_number}: не указан Партнер ID.")
            agency = Agency.objects.filter(pk=int(agency_id), archived=False).first()
            if not agency:
                raise ProductOperationError(f"Строка {row_number}: партнер не найден.")
            values = {
                "name": _value(row, 3),
                "article": _value(row, 4),
                "barcode": _value(row, 5),
                "color": _value(row, 6),
                "weight_grams": _value(row, 7),
                "size": _value(row, 8),
                "width_cm": _value(row, 9),
                "depth_cm": _value(row, 10),
                "height_cm": _value(row, 11),
                "category": _value(row, 12),
                "is_bundle": _boolean(_value(row, 13)),
                "image_url": _value(row, 14),
                "internal_notes": _value(row, 15),
                "description": _value(row, 16),
            }
            product = None
            if product_id:
                if not product_id.isdigit():
                    raise ProductOperationError(f"Строка {row_number}: некорректный ID FBS-NEW.")
                product = WmsNewProduct.objects.filter(pk=int(product_id), is_archived=False).first()
                if not product:
                    raise ProductOperationError(f"Строка {row_number}: товар FBS-NEW не найден.")
                if product.agency_id != agency.id:
                    raise ProductOperationError(f"Строка {row_number}: нельзя изменить партнера товара.")
            if product is None:
                product = WmsNewProduct.objects.filter(
                    agency=agency,
                    article=str(values["article"] or "").strip(),
                    is_archived=False,
                ).first()
            if product:
                update_product(product_id=product.id, actor=actor, **values)
                updated += 1
            else:
                create_product(agency=agency, actor=actor, **values)
                created += 1
    return {"created": created, "updated": updated}
