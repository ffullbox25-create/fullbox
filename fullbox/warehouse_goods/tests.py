import base64
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import connection
from django.test import Client, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from employees.models import Employee
from processing_app.models import ProcessingPrintJob
from shipping.models import ShippingOrder, ShippingOrderItem
from sklad.models import WarehouseEvent, WarehouseLocation, WarehouseStockSnapshot
from sku.models import Agency, SKU, SKUBarcode

from .models import GoodsProfile
from .selectors import goods_list_queryset, location_rows, stock_totals
from .services import build_label_pdf, queue_label


ONE_PIXEL_PNG = (
    "data:image/png;base64,"
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


class WarehouseGoodsManagerAccessTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(
            username="warehouse-goods-manager",
            password="test",
        )
        Employee.objects.create(
            user=cls.user,
            full_name="Тестовый менеджер товаров",
            role="manager",
            is_active=True,
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_manager_opens_the_same_goods_registry_and_create_form(self):
        goods_list = self.client.get(reverse("warehouse_goods:list"))
        self.assertEqual(goods_list.status_code, 200)
        self.assertContains(goods_list, "Добавить товар")
        self.assertContains(goods_list, "Экспорт в EXCEL")

        create_form = self.client.get(reverse("warehouse_goods:create"))
        self.assertEqual(create_form.status_code, 200)


class WarehouseGoodsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user, _ = get_user_model().objects.get_or_create(username="dev")
        cls.user.set_password("test")
        cls.user.save(update_fields=["password"])
        cls.agency = Agency.objects.create(agn_name="Тестовый клиент")
        cls.sku = SKU.objects.create(
            agency=cls.agency,
            sku_code="TEST-001",
            name="Тестовый товар",
            weight_kg=Decimal("1.250"),
        )
        SKUBarcode.objects.create(sku=cls.sku, value="460000000001", is_primary=True)
        WarehouseStockSnapshot.objects.create(
            agency=cls.agency,
            sku_ref=cls.sku,
            sku_code=cls.sku.sku_code,
            name=cls.sku.name,
            qty=12,
            available_qty=7,
            processing_reserved_qty=3,
            shipping_reserved_qty=2,
        )
        cls.client = Client()

    def setUp(self):
        self.client.force_login(self.user)

    def test_list_uses_warehouse_snapshot_truth(self):
        response = self.client.get(reverse("warehouse_goods:list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Тестовый товар")
        row = list(goods_list_queryset(search="TEST-001"))[0]
        self.assertEqual(row.stock_total, 12)
        self.assertEqual(row.stock_available, 7)
        self.assertEqual(row.stock_processing, 3)
        self.assertEqual(row.stock_shipping, 2)
        self.assertEqual(stock_totals(self.sku)["total"], 12)

    def test_list_query_count_does_not_grow_per_sku(self):
        SKU.objects.bulk_create(
            [SKU(agency=self.agency, sku_code=f"TEST-{index:03}", name=f"Товар {index}") for index in range(2, 42)]
        )
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(reverse("warehouse_goods:list"), {"page_size": 50})
        self.assertEqual(response.status_code, 200)
        self.assertLessEqual(len(queries), 12)

    def test_detail_tabs_and_location_selector(self):
        self.assertFalse(GoodsProfile.objects.filter(sku=self.sku).exists())
        detail = self.client.get(reverse("warehouse_goods:detail", args=[self.sku.pk]))
        self.assertEqual(detail.status_code, 200)
        responses = {
            tab: self.client.get(reverse("warehouse_goods:tab", args=[self.sku.pk, tab]))
            for tab in (
                "about",
                "characteristics",
                "locations",
                "identifiers",
                "services",
                "orders",
                "files",
                "marking",
                "extra",
                "tasks",
            )
        }
        for response in responses.values():
            self.assertEqual(response.status_code, 200)
        identifiers = responses["identifiers"]
        self.assertContains(identifiers, "460000000001")
        self.assertEqual(list(location_rows(self.sku))[0]["qty"], 12)
        self.assertFalse(GoodsProfile.objects.filter(sku=self.sku).exists())

    def test_partial_profile_updates_preserve_other_tabs(self):
        profile = GoodsProfile.objects.create(
            sku=self.sku,
            internal_notes="Заметка",
            receiving_task_template="Принимать внимательно",
            shipping_task_template="Отгружать внимательно",
            fbs_shelf_life_required=True,
        )
        response = self.client.post(
            reverse("warehouse_goods:profile-update", args=[self.sku.pk]),
            {"return_tab": "about", "internal_notes": "Новая заметка"},
        )
        self.assertEqual(response.status_code, 302)
        profile.refresh_from_db()
        self.assertEqual(profile.internal_notes, "Новая заметка")
        self.assertEqual(profile.receiving_task_template, "Принимать внимательно")
        self.assertTrue(profile.fbs_shelf_life_required)

        self.client.post(
            reverse("warehouse_goods:profile-update", args=[self.sku.pk]),
            {"return_tab": "characteristics"},
        )
        profile.refresh_from_db()
        self.assertFalse(profile.fbs_shelf_life_required)
        self.assertEqual(profile.shipping_task_template, "Отгружать внимательно")

    def test_transfer_creates_standard_shipping_draft(self):
        response = self.client.post(
            reverse("warehouse_goods:transfer", args=[self.sku.pk]),
            {"qty": 2, "comment": "На другой склад"},
        )
        self.assertEqual(response.status_code, 302)
        order = ShippingOrder.objects.get()
        item = ShippingOrderItem.objects.get(order=order)
        self.assertEqual(order.delivery_type, ShippingOrder.DELIVERY_TRANSFER)
        self.assertEqual(order.status, ShippingOrder.STATUS_DRAFT)
        self.assertEqual(item.sku_id, self.sku.pk)
        self.assertEqual(item.qty_requested, 2)

    def test_label_pdf_and_print_queue(self):
        pdf = build_label_pdf(ONE_PIXEL_PNG, width_mm=58, height_mm=40, copies=2)
        self.assertTrue(pdf.startswith(b"%PDF"))
        job = queue_label(
            self.sku,
            data_url=ONE_PIXEL_PNG,
            template_key="item",
            width_mm=58,
            height_mm=40,
            copies=3,
            user=self.user,
        )
        self.assertEqual(job.status, ProcessingPrintJob.STATUS_PENDING)
        self.assertEqual(job.copies_count, 3)
        self.assertEqual(base64.b64decode(job.label_png_base64)[:8], b"\x89PNG\r\n\x1a\n")

    def test_history_renders_event_without_user(self):
        source = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="OS",
            row_no=3,
            section_no=2,
            tier_no=1,
            cell_no=1,
            location_code="A-3/1-1",
            display_name="OS · Ряд 3 · Секция 2 · Ярус 1 · Ячейка 1",
        )
        destination = WarehouseLocation.objects.create(
            warehouse_code="MSK",
            zone_code="OS",
            row_no=10,
            section_no=2,
            tier_no=1,
            cell_no=2,
            location_code="A-10/1-2",
            display_name="OS · Ряд 10 · Секция 2 · Ярус 1 · Ячейка 2",
        )
        WarehouseEvent.objects.create(
            agency=self.agency,
            event_type="fbs_replenishment_completed",
            from_location=source,
            to_location=destination,
            qty=1,
            payload={
                "sku_code": self.sku.sku_code,
                "target_box_code": "TEST-BOX-001",
                "marking_code": "SECRET-MARKING-CODE",
                "allocation_id": 999,
            },
            performed_by=None,
            occurred_at=timezone.make_aware(datetime(2026, 8, 11, 18, 44, 15)),
        )
        response = self.client.get(reverse("warehouse_goods:history", args=[self.sku.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "История движения")
        self.assertContains(response, "Подсорт FBS выполнен")
        self.assertContains(response, "A-3/1-1")
        self.assertContains(response, "A-10/1-2")
        self.assertContains(response, "Короб: TEST-BOX-001")
        self.assertContains(response, "11.08.2026 18:44")
        self.assertNotContains(response, "11.08.2026 18:44:15")
        self.assertNotContains(response, "fbs_replenishment_completed")
        self.assertNotContains(response, "SECRET-MARKING-CODE")
        self.assertNotContains(response, "allocation_id")

        export = self.client.get(
            reverse("warehouse_goods:history-export", args=[self.sku.pk])
        )
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(BytesIO(export.content), read_only=True).active
        exported_row = list(
            sheet.iter_rows(min_row=2, max_row=2, values_only=True)
        )[0]
        self.assertEqual(exported_row[0], "11.08.2026 18:44")
        self.assertEqual(exported_row[1], "Подсорт FBS выполнен")
        self.assertEqual(exported_row[2], "A-3/1-1")
        self.assertEqual(exported_row[3], "A-10/1-2")
        self.assertEqual(exported_row[6], "Короб: TEST-BOX-001")
