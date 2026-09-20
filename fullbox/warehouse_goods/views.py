from io import BytesIO

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils import timezone
from django.views import View
from openpyxl import Workbook

from employees.access import RoleRequiredMixin, get_request_role
from labels.utils import LABEL_TEMPLATES
from shipping.models import ShippingOrder
from sklad.models import WarehouseLocation
from sku.models import Agency, MarketplaceBinding, SKU, SKUBarcode

from .forms import (
    GoodsBarcodeForm,
    GoodsDefaultServiceForm,
    GoodsExtraDefinitionForm,
    GoodsFileForm,
    GoodsMarketplaceBindingForm,
    GoodsProfileForm,
    GoodsSKUForm,
)
from .models import (
    GoodsDefaultService,
    GoodsExtraFieldDefinition,
    GoodsExtraFieldValue,
    GoodsFile,
    GoodsProfile,
)
from .history_presenter import present_history_events
from .selectors import (
    goods_list_queryset,
    location_rows,
    marking_rows,
    movement_rows,
    order_rows,
    stock_totals,
)
from .services import (
    add_barcode,
    build_label_pdf,
    create_move_request,
    create_transfer_draft,
    delete_barcode,
    log_goods_action,
    queue_label,
    save_sku_form,
)


class GoodsAccessMixin(RoleRequiredMixin):
    allowed_roles = ("head_manager", "manager", "storekeeper", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        request.warehouse_goods_role = get_request_role(request)
        return super().dispatch(request, *args, **kwargs)


class GoodsHistoryAccessMixin(GoodsAccessMixin):
    allowed_roles = (
        "head_manager",
        "manager",
        "storekeeper",
        "processing_head",
        "processing_worker",
    )


def _goods_or_404(pk):
    return get_object_or_404(
        SKU.objects.filter(deleted=False).select_related("agency", "market", "color_ref"), pk=pk
    )


def _profile_for_sku(sku):
    return GoodsProfile.objects.filter(sku=sku).first() or GoodsProfile(sku=sku)


def _redirect_tab(sku, tab):
    return redirect(f"{reverse('warehouse_goods:detail', args=[sku.pk])}?tab={tab}#{tab}")


def _int_between(value, minimum=1, maximum=100000):
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError("Укажите корректное количество.")
    if number < minimum or number > maximum:
        raise ValueError(f"Количество должно быть от {minimum} до {maximum}.")
    return number


class GoodsListView(GoodsAccessMixin, View):
    template_name = "warehouse_goods/goods_list.html"

    def get(self, request):
        search = request.GET.get("q", "")
        agency_id = request.GET.get("agency", "")
        stock_state = request.GET.get("stock", "")
        category = request.GET.get("category", "")
        bundle = request.GET.get("bundle", "")
        without_zero = request.GET.get("without_zero") == "1"
        reserves_only = request.GET.get("reserves_only") == "1"
        sort = request.GET.get("sort", "name")
        page_size = request.GET.get("page_size", "100")
        if page_size not in {"25", "50", "100", "200", "500"}:
            page_size = "100"
        queryset = goods_list_queryset(
            search=search,
            agency_id=agency_id,
            stock_state="positive" if without_zero else stock_state,
            sort=sort,
            category=category,
            reserves_only=reserves_only,
        )
        page_obj = Paginator(queryset, int(page_size)).get_page(request.GET.get("page"))
        query = request.GET.copy()
        query.pop("page", None)
        context = {
            "page_obj": page_obj,
            "agencies": Agency.objects.filter(archived=False).order_by("agn_name"),
            "categories": SKU.objects.filter(deleted=False)
            .exclude(tovar_category__isnull=True)
            .exclude(tovar_category="")
            .values_list("tovar_category", flat=True)
            .distinct()
            .order_by("tovar_category"),
            "filters": {
                "q": search,
                "agency": str(agency_id or ""),
                "stock": stock_state,
                "category": category,
                "bundle": bundle,
                "without_zero": without_zero,
                "reserves_only": reserves_only,
                "sort": sort,
                "page_size": page_size,
            },
            "query_without_page": query.urlencode(),
        }
        return render(request, self.template_name, context)


class GoodsExportView(GoodsAccessMixin, View):
    def get(self, request):
        queryset = goods_list_queryset(
            search=request.GET.get("q", ""),
            agency_id=request.GET.get("agency", ""),
            stock_state=request.GET.get("stock", ""),
            sort=request.GET.get("sort", "name"),
        )
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Товары")
        sheet.append(
            [
                "ID",
                "Клиент",
                "Артикул",
                "Наименование",
                "Штрихкод",
                "Всего",
                "Свободно",
                "Резерв обработки",
                "Резерв отгрузки",
            ]
        )
        for sku in queryset.iterator(chunk_size=1000):
            sheet.append(
                [
                    sku.pk,
                    sku.agency.agn_name if sku.agency else "",
                    sku.sku_code,
                    sku.name,
                    sku.display_barcode or "",
                    sku.stock_total,
                    sku.stock_available,
                    sku.stock_processing,
                    sku.stock_shipping,
                ]
            )
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="warehouse-goods.xlsx"'
        return response


class GoodsCreateView(GoodsAccessMixin, View):
    template_name = "warehouse_goods/sku_form.html"

    def get(self, request):
        return render(request, self.template_name, {"form": GoodsSKUForm(), "mode": "create"})

    def post(self, request):
        form = GoodsSKUForm(request.POST)
        if form.is_valid():
            sku = save_sku_form(form, request.user)
            messages.success(request, "Товар создан. Теперь можно добавить штрихкоды и складские настройки.")
            return redirect("warehouse_goods:detail", pk=sku.pk)
        return render(request, self.template_name, {"form": form, "mode": "create"}, status=400)


class GoodsEditView(GoodsAccessMixin, View):
    template_name = "warehouse_goods/sku_form.html"

    def get(self, request, pk):
        sku = _goods_or_404(pk)
        return render(
            request,
            self.template_name,
            {"form": GoodsSKUForm(instance=sku), "sku": sku, "mode": "edit"},
        )

    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsSKUForm(request.POST, instance=sku)
        if form.is_valid():
            save_sku_form(form, request.user)
            messages.success(request, "Изменения товара сохранены.")
            return redirect("warehouse_goods:detail", pk=sku.pk)
        return render(
            request,
            self.template_name,
            {"form": form, "sku": sku, "mode": "edit"},
            status=400,
        )


class GoodsDetailView(GoodsAccessMixin, View):
    template_name = "warehouse_goods/goods_detail.html"

    def get(self, request, pk):
        sku = _goods_or_404(pk)
        profile = _profile_for_sku(sku)
        labels = [item for item in LABEL_TEMPLATES if item.get("category") == "product"]
        locations = WarehouseLocation.objects.filter(
            is_active=True, zone_code__in=["PR", "OTG", "MR", "OS", "OBR"]
        ).order_by("zone_code", "row_no", "section_no", "tier_no", "cell_no")
        context = {
            "sku": sku,
            "profile": profile,
            "totals": stock_totals(sku),
            "primary_barcode": sku.barcodes.order_by("-is_primary", "id").first(),
            "initial_tab": request.GET.get("tab", "about"),
            "label_templates": labels,
            "move_locations": locations,
        }
        return render(request, self.template_name, context)


def _tab_context(sku, tab):
    context = {"sku": sku, "tab": tab}
    if tab == "about":
        profile = _profile_for_sku(sku)
        context.update(profile=profile, totals=stock_totals(sku))
    elif tab == "characteristics":
        context["profile"] = _profile_for_sku(sku)
    elif tab == "locations":
        context["locations"] = location_rows(sku)
    elif tab == "identifiers":
        context["barcodes"] = sku.barcodes.all()
        context["bindings"] = MarketplaceBinding.objects.filter(sku=sku).order_by("marketplace")
        context["barcode_form"] = GoodsBarcodeForm()
        context["binding_form"] = GoodsMarketplaceBindingForm()
    elif tab == "services":
        context["default_services"] = GoodsDefaultService.objects.filter(sku=sku).select_related("service")
        context["service_form"] = GoodsDefaultServiceForm()
    elif tab == "orders":
        context["orders"] = order_rows(sku)
    elif tab == "files":
        context["files"] = GoodsFile.objects.filter(sku=sku).select_related("uploaded_by")
        context["file_form"] = GoodsFileForm()
    elif tab == "marking":
        rows = marking_rows(sku)
        context["marking_rows"] = rows[:250]
        context["marking_summary"] = {
            "total": rows.count(),
            "unprinted": rows.filter(printed_at__isnull=True).count(),
            "unused": rows.filter(used_at__isnull=True).count(),
        }
    elif tab == "extra":
        definitions = list(GoodsExtraFieldDefinition.objects.filter(is_active=True))
        values = {
            row.definition_id: row.value
            for row in GoodsExtraFieldValue.objects.filter(sku=sku, definition__in=definitions)
        }
        context["extra_fields"] = [(definition, values.get(definition.pk)) for definition in definitions]
        context["definition_form"] = GoodsExtraDefinitionForm()
    elif tab == "tasks":
        context["profile"] = _profile_for_sku(sku)
    else:
        context["tab"] = "about"
        context["profile"] = _profile_for_sku(sku)
        context["totals"] = stock_totals(sku)
    return context


class GoodsTabView(GoodsAccessMixin, View):
    valid_tabs = {
        "about",
        "characteristics",
        "locations",
        "identifiers",
        "services",
        "orders",
        "files",
        "marking",
        "extra",
        "tasks",
    }

    def get(self, request, pk, tab):
        sku = _goods_or_404(pk)
        if tab not in self.valid_tabs:
            tab = "about"
        return render(request, "warehouse_goods/_tab_content.html", _tab_context(sku, tab))


class GoodsProfileUpdateView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        profile, _ = GoodsProfile.objects.get_or_create(sku=sku)
        data = request.POST.copy()
        for field_name in GoodsProfileForm.Meta.fields:
            if field_name not in data:
                current = getattr(profile, field_name)
                if field_name == "fbs_shelf_life_required" and request.POST.get("return_tab") == "characteristics":
                    data[field_name] = ""
                else:
                    data[field_name] = "on" if current is True else ("" if current is False else current)
        form = GoodsProfileForm(data, instance=profile)
        if form.is_valid():
            profile = form.save(commit=False)
            profile.updated_by = request.user
            profile.save()
            log_goods_action(sku, "profile_update", request.user)
            messages.success(request, "Складские настройки сохранены.")
        else:
            messages.error(request, "Не удалось сохранить настройки.")
        return _redirect_tab(sku, request.POST.get("return_tab", "about"))


class GoodsBarcodeAddView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsBarcodeForm(request.POST)
        if form.is_valid():
            try:
                add_barcode(sku, user=request.user, **form.cleaned_data)
                messages.success(request, "Штрихкод добавлен.")
            except IntegrityError:
                messages.error(request, "Такой штрихкод уже используется.")
        else:
            messages.error(request, "Проверьте значение штрихкода.")
        return _redirect_tab(sku, "identifiers")


class GoodsBarcodeDeleteView(GoodsAccessMixin, View):
    def post(self, request, pk, barcode_id):
        sku = _goods_or_404(pk)
        barcode = get_object_or_404(SKUBarcode, pk=barcode_id, sku=sku)
        delete_barcode(barcode, request.user)
        messages.success(request, "Штрихкод удален.")
        return _redirect_tab(sku, "identifiers")


class GoodsBindingAddView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsMarketplaceBindingForm(request.POST)
        if form.is_valid():
            binding = form.save(commit=False)
            binding.sku = sku
            try:
                binding.save()
                log_goods_action(
                    sku,
                    "marketplace_binding_add",
                    request.user,
                    marketplace=binding.marketplace,
                    external_id=binding.external_id,
                )
                messages.success(request, "Внешний артикул добавлен.")
            except IntegrityError:
                messages.error(request, "Такая привязка маркетплейса уже существует.")
        else:
            messages.error(request, "Проверьте маркетплейс и внешний артикул.")
        return _redirect_tab(sku, "identifiers")


class GoodsBindingDeleteView(GoodsAccessMixin, View):
    def post(self, request, pk, binding_id):
        sku = _goods_or_404(pk)
        binding = get_object_or_404(MarketplaceBinding, pk=binding_id, sku=sku)
        external_id = binding.external_id
        binding.delete()
        log_goods_action(sku, "marketplace_binding_delete", request.user, external_id=external_id)
        messages.success(request, "Внешний артикул удален.")
        return _redirect_tab(sku, "identifiers")


class GoodsFileUploadView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsFileForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded = form.save(commit=False)
            uploaded.sku = sku
            uploaded.original_name = request.FILES["file"].name[:255]
            uploaded.uploaded_by = request.user
            uploaded.save()
            log_goods_action(sku, "file_upload", request.user, filename=uploaded.original_name)
            messages.success(request, "Файл загружен.")
        else:
            messages.error(request, "Выберите файл для загрузки.")
        return _redirect_tab(sku, "files")


class GoodsFileDeleteView(GoodsAccessMixin, View):
    def post(self, request, pk, file_id):
        sku = _goods_or_404(pk)
        attachment = get_object_or_404(GoodsFile, pk=file_id, sku=sku)
        name = attachment.original_name
        attachment.file.delete(save=False)
        attachment.delete()
        log_goods_action(sku, "file_delete", request.user, filename=name)
        messages.success(request, "Файл удален.")
        return _redirect_tab(sku, "files")


class GoodsExtraDefinitionAddView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsExtraDefinitionForm(request.POST)
        if form.is_valid():
            definition = form.save()
            log_goods_action(sku, "extra_definition_add", request.user, definition=definition.name)
            messages.success(request, "Дополнительное поле создано.")
        else:
            messages.error(request, "Не удалось создать дополнительное поле.")
        return _redirect_tab(sku, "extra")


class GoodsExtraFieldsUpdateView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        definitions = GoodsExtraFieldDefinition.objects.filter(is_active=True)
        with transaction.atomic():
            for definition in definitions:
                raw = request.POST.get(f"field_{definition.pk}")
                if definition.field_type == definition.TYPE_BOOLEAN:
                    value = raw == "on"
                elif raw in (None, ""):
                    value = None
                elif definition.field_type == definition.TYPE_NUMBER:
                    try:
                        value = float(raw.replace(",", "."))
                    except (TypeError, ValueError):
                        messages.error(request, f"Поле «{definition.name}» должно быть числом.")
                        return _redirect_tab(sku, "extra")
                else:
                    value = raw
                GoodsExtraFieldValue.objects.update_or_create(
                    sku=sku,
                    definition=definition,
                    defaults={"value": value, "updated_by": request.user},
                )
        log_goods_action(sku, "extra_fields_update", request.user)
        messages.success(request, "Дополнительные поля сохранены.")
        return _redirect_tab(sku, "extra")


class GoodsServiceAddView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        form = GoodsDefaultServiceForm(request.POST)
        if form.is_valid():
            service = form.save(commit=False)
            service.sku = sku
            try:
                service.save()
                log_goods_action(sku, "default_service_add", request.user, service=service.service.name)
                messages.success(request, "Типовая услуга добавлена.")
            except IntegrityError:
                messages.error(request, "Эта услуга уже добавлена для выбранного этапа.")
        else:
            messages.error(request, "Проверьте параметры услуги.")
        return _redirect_tab(sku, "services")


class GoodsServiceDeleteView(GoodsAccessMixin, View):
    def post(self, request, pk, service_id):
        sku = _goods_or_404(pk)
        service = get_object_or_404(GoodsDefaultService, pk=service_id, sku=sku)
        name = service.service.name
        service.delete()
        log_goods_action(sku, "default_service_delete", request.user, service=name)
        messages.success(request, "Типовая услуга удалена.")
        return _redirect_tab(sku, "services")


class GoodsHistoryView(GoodsHistoryAccessMixin, View):
    def get(self, request, pk):
        sku = _goods_or_404(pk)
        barcode = str(request.GET.get("barcode") or "").strip()
        if barcode and not sku.barcodes.filter(value=barcode).exists():
            raise Http404("Штрихкод товара не найден")
        back_url = str(request.GET.get("next") or "").strip()
        back_to_stock = url_has_allowed_host_and_scheme(
            back_url,
            allowed_hosts={request.get_host()},
            require_https=request.is_secure(),
        )
        if not back_to_stock:
            back_url = reverse("warehouse_goods:detail", args=[sku.pk])
        rows = present_history_events(movement_rows(sku, limit=500, barcode=barcode))
        return render(
            request,
            "warehouse_goods/goods_history.html",
            {
                "sku": sku,
                "rows": rows,
                "selected_barcode": barcode,
                "back_url": back_url,
                "back_label": "Вернуться к остаткам" if back_to_stock else "К товару",
            },
        )


class GoodsHistoryExportView(GoodsHistoryAccessMixin, View):
    def get(self, request, pk):
        sku = _goods_or_404(pk)
        barcode = str(request.GET.get("barcode") or "").strip()
        if barcode and not sku.barcodes.filter(value=barcode).exists():
            raise Http404("Штрихкод товара не найден")
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("История движения")
        sheet.append(["Дата", "Действие", "Откуда", "Куда", "Количество", "Пользователь", "Информация"])
        for row in present_history_events(movement_rows(sku, limit=5000, barcode=barcode)):
            sheet.append(
                [
                    timezone.localtime(row.occurred_at).strftime("%d.%m.%Y %H:%M")
                    if row.occurred_at
                    else "",
                    row.action,
                    row.from_location,
                    row.to_location,
                    row.qty,
                    row.user_name,
                    row.information,
                ]
            )
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="goods-{sku.pk}-history.xlsx"'
        return response


class GoodsMoveView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        try:
            qty = _int_between(request.POST.get("qty"))
            location = get_object_or_404(WarehouseLocation, pk=request.POST.get("location"), is_active=True)
            result = create_move_request(
                request, sku, location=location, qty=qty, comment=request.POST.get("comment", "")
            )
            messages.success(
                request,
                f"Заявка на перемещение №{result.get('request_id')} создана; заданий: {result.get('tasks_created', 0)}.",
            )
        except (ValueError, ValidationError) as exc:
            messages.error(request, str(exc))
        return redirect("warehouse_goods:detail", pk=sku.pk)


class GoodsTransferView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        try:
            qty = _int_between(request.POST.get("qty"))
            order = create_transfer_draft(sku, qty, request.user, request.POST.get("comment", ""))
        except (ValueError, ValidationError) as exc:
            messages.error(request, str(exc))
            return redirect("warehouse_goods:detail", pk=sku.pk)
        messages.success(request, f"Создан черновик передачи {order.number}. Заполните получателя и подтвердите заявку.")
        return redirect(f"{reverse('shipping:create')}?edit=1&order={order.pk}")


class GoodsLabelPdfView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        try:
            width = _int_between(request.POST.get("width_mm"), 20, 200)
            height = _int_between(request.POST.get("height_mm"), 20, 250)
            copies = _int_between(request.POST.get("copies"), 1, 500)
            pdf = build_label_pdf(
                request.POST.get("image", ""), width_mm=width, height_mm=height, copies=copies
            )
            log_goods_action(sku, "label_pdf", request.user, copies=copies, width=width, height=height)
        except Exception as exc:
            return JsonResponse({"ok": False, "error": f"Не удалось создать PDF: {exc}"}, status=400)
        response = HttpResponse(pdf, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{sku.sku_code}-labels.pdf"'
        return response


class GoodsLabelQueueView(GoodsAccessMixin, View):
    def post(self, request, pk):
        sku = _goods_or_404(pk)
        try:
            width = _int_between(request.POST.get("width_mm"), 20, 200)
            height = _int_between(request.POST.get("height_mm"), 20, 250)
            copies = _int_between(request.POST.get("copies"), 1, 500)
            job = queue_label(
                sku,
                data_url=request.POST.get("image", ""),
                template_key=request.POST.get("template_key", "item"),
                width_mm=width,
                height_mm=height,
                copies=copies,
                user=request.user,
            )
        except Exception as exc:
            return JsonResponse({"ok": False, "error": f"Не удалось поставить в печать: {exc}"}, status=400)
        return JsonResponse({"ok": True, "job_id": job.pk, "message": f"Задание №{job.pk} отправлено в очередь."})
