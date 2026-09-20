"""Orders UI views and helper logic."""

import base64
import json
import re
from io import BytesIO
from decimal import Decimal, InvalidOperation
from datetime import datetime, time, timedelta
from pathlib import Path
from urllib.parse import unquote_plus, urlencode

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.generic import TemplateView
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill

from audit.models import OrderAuditEntry, log_order_action, log_staff_overaction, log_stock_move
from client_cabinet.portal_access import SECTION_REQUESTS
from client_cabinet.portal_members import resolve_portal_agency_for_user
from agent.models import DeviceAgent
from employees.models import Employee
from employees.access import RoleRequiredMixin, get_request_role, resolve_cabinet_url, is_staff_role
from orders.access_policy import order_detail_access_allowed
from fullbox.container_codes import (
    issue_container_code,
    issue_pallet_code,
    normalize_container_kind,
    rewrite_container_code_goods_type,
)
from fullbox.order_numbers import (
    format_order_number,
    next_public_order_number,
    normalize_order_type_for_number,
    order_number_sequence,
)
from fullbox.timeutils import business_localtime, parse_business_datetime
from marking.models import MarkingCode
from marking.utils import extract_processing_items
from processing_app.stages import processing_stage_label
from reachtruck.models import MoveTask
from sku.models import Agency, SKU, SKUBarcode
from orders.models import ReceivingOrderAttachment
from orders.services import ReceivingNomenclatureError, ReceivingWorkflowService
from orders.shipping_returns import (
    SHIPPING_RETURN_GOODS_TYPE,
    SHIPPING_RETURN_GOODS_TYPE_LABEL,
    resolve_shipping_return_order,
    search_shipping_return_orders,
    shipping_return_source_items,
)
from orders.title_truth import build_order_runtime_payload, resolve_order_title
from sklad.models import WarehouseStockSnapshot
from sklad.services.warehouse_commands import WarehouseCommandService
from sklad.services.warehouse_policy import WarehouseActionPolicy
from sklad.services.stock_operations import OperationalStockService
from sklad.services.warehouse_state import (
    WarehouseGoodsStateResolver,
    WarehouseMovementResolver,
    WarehouseStateCode,
)
from sklad.services.warehouse_write_path import WarehouseWritePathService
from sklad.services.stock_availability import StockAvailabilityService
from todo.models import Task
from stockmap.views import _OS_CELLS_PER_TIER, _OS_ROW_SECTIONS, _OS_TIERS
from reachtruck.services import create_batch_move_tasks

_MOJIBAKE_FIXUPS = {
    "Р—Р°СЏРІРєР° Р·Р°РІРµСЂС€РµРЅР°": "Заявка завершена",
    "РўРѕРІР°СЂ РїСЂРёРЅСЏС‚ Рё СЂР°Р·РјРµС‰РµРЅ РЅР° СЃРєР»Р°РґРµ": "Товар принят и размещен на складе",
    "РўРѕРІР°СЂ РІРѕР·РІСЂР°С‰РµРЅ РЅР° СЃРєР»Р°Рґ": "Товар возвращен на склад",
}
_JOURNAL_ENTRY_LIMIT = 80
_JOURNAL_RAW_SCAN_LIMIT = 5000


_IP_PREFIX_RE = re.compile(r"\bиндивидуальный предприниматель\b", re.IGNORECASE)
_TEMPLATE_DOCS_DIR = settings.BASE_DIR / "static" / "docs"
_ACT_DOCS_DIR = settings.MEDIA_ROOT / "acts"

_ACT_TEMPLATE_FILE = _TEMPLATE_DOCS_DIR / "receiving_act_template.xlsx"
_MX1_TEMPLATE_FILE = _TEMPLATE_DOCS_DIR / "mx1_template.xlsx"
_EXECUTOR_LABEL = (
    'ООО "Фуллбокс", ИНН 5001149130, КПП 503101001, 142450, Московская обл, '
    "г.о. Богородский, г Старая Купавна, ул Магистральная, д. 59, помещ. 55, "
    "тел.: +79778088386"
)
_OKUD_MX1 = "0335001"


def _download_template_file(file_name: str, download_name: str):
    file_path = _TEMPLATE_DOCS_DIR / file_name
    if not file_path.exists():
        raise Http404("Файл не найден")
    return FileResponse(open(file_path, "rb"), as_attachment=True, filename=download_name)


def download_receiving_template(request):
    return _download_template_file("receiving_template.xlsx", "Шаблон приемки.xlsx")


def download_sku_upload_template(request):
    return _download_template_file("sku_upload_template.xlsx", "Шаблон для загрузки номенклатуры.xlsx")


def _shorten_ip_name(name: str) -> str:
    if not name:
        return "-"
    normalized = _IP_PREFIX_RE.sub("ИП", name)
    return " ".join(normalized.split())


def _repair_mojibake_text(value) -> str:
    text = str(value or "")
    if not text:
        return ""
    fixed = _MOJIBAKE_FIXUPS.get(text)
    if fixed:
        return fixed

    def _score(raw: str) -> int:
        return (
            raw.count("Р")
            + raw.count("С")
            + raw.count("Ð")
            + raw.count("Ñ")
            + raw.count("Â")
            + raw.count("ð")
            + raw.count("à")
            + raw.count("î")
            + raw.count("вЂ")
            + raw.count("в„")
            + sum(3 for char in raw if 0x80 <= ord(char) <= 0x9F)
        )

    def _raw_bytes(raw: str, source_encoding: str) -> bytes | None:
        chunks: list[bytes] = []
        for char in raw:
            code = ord(char)
            if source_encoding == "latin-1":
                if code > 0xFF:
                    return None
                chunks.append(bytes([code]))
                continue
            if source_encoding == "cp1251" and 0x80 <= code <= 0x9F:
                chunks.append(bytes([code]))
                continue
            try:
                chunks.append(char.encode(source_encoding))
            except UnicodeEncodeError:
                return None
        return b"".join(chunks)

    def _candidate(raw: str, source_encoding: str, target_encoding: str) -> str | None:
        raw_bytes = _raw_bytes(raw, source_encoding)
        if raw_bytes is None:
            return None
        try:
            return raw_bytes.decode(target_encoding)
        except UnicodeDecodeError:
            return None

    current = text
    for _ in range(4):
        candidates = [current]
        for source_encoding, target_encoding in (
            ("cp1251", "utf-8"),
            ("latin-1", "utf-8"),
            ("latin-1", "cp1251"),
        ):
            repaired = _candidate(current, source_encoding, target_encoding)
            if repaired:
                candidates.append(repaired)
        best = min(candidates, key=_score)
        if _score(best) >= _score(current):
            break
        current = best
    return current


def _short_name(full_name: str) -> str:
    if not full_name:
        return "-"
    parts = [part for part in full_name.split() if part]
    if not parts:
        return "-"
    surname = parts[0]
    initials = "".join(f"{part[0].upper()}." for part in parts[1:3] if part)
    return f"{surname} {initials}".strip()


def _order_type_label(order_type: str) -> str:
    if not order_type:
        return "-"
    labels = {"receiving": "PR", "packing": "ЗУ", "processing": "OBR", "shipping": "OTG", "other": "OTH"}
    return labels.get(order_type, order_type)


def _display_order_number(order_type: str, order_id: str) -> str:
    return format_order_number(order_type, order_id)


def _first_active_employee_by_roles(*roles: str) -> Employee | None:
    for role in [role for role in roles if role]:
        employee = (
            Employee.objects.filter(role=role, is_active=True)
            .order_by("full_name")
            .first()
        )
        if employee:
            return employee
    return None


def _facsimile_url(employee: Employee | None) -> str:
    if not employee or not getattr(employee, "facsimile", None):
        return ""
    file_handle = None
    try:
        file_handle = employee.facsimile.open("rb")
        content = file_handle.read()
    except (FileNotFoundError, OSError, ValueError):
        return ""
    finally:
        if file_handle is not None:
            file_handle.close()
    if not content:
        return ""
    encoded = base64.b64encode(content).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _barcode_value_for_sku(sku, size: str | None) -> str:
    if not sku:
        return "-"
    barcodes = list(getattr(sku, "barcodes", []).all())
    size_value = (size or "").strip()
    if size_value:
        for barcode in barcodes:
            if (barcode.size or "").strip() == size_value:
                return barcode.value
    sku_code_barcode = str(getattr(sku, "code", "") or "").strip()
    if sku_code_barcode:
        return sku_code_barcode
    if not barcodes:
        return "-"
    primary = next((barcode for barcode in barcodes if barcode.is_primary), None)
    return primary.value if primary else barcodes[0].value


def _has_receiving_items(payload: dict) -> bool:
    items = payload.get("items") or []
    for item in items:
        for key in ("sku_code", "name", "qty", "size"):
            if str(item.get(key) or "").strip():
                return True
    return False


def _is_sent_to_manager(payload: dict) -> bool:
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    if status_value in {"sent_unconfirmed", "send", "submitted"}:
        return True
    return "подтверждени" in status_label


def _order_title_label(
    order_type: str,
    order_id: str,
    payload: dict | None = None,
    *,
    entries=None,
    created_at=None,
) -> str:
    return resolve_order_title(
        order_type,
        order_id,
        payload=payload,
        entries=entries,
        created_at=created_at,
        variant="base",
    )


def _journal_status_label(entry):
    payload = entry.payload or {}
    status = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
    status_label = _repair_mojibake_text(
        payload.get("status_label")
        or payload.get("processing_stage_label")
        or payload.get("stage_label")
        or ""
    ).strip()
    normalized_label = status_label.lower()
    if status == "draft" or "черновик" in normalized_label:
        return "Черновик"
    if status in {"sent_unconfirmed", "send", "submitted"} or "подтверждени" in normalized_label:
        return "Ждет подтверждения"
    if status_label:
        return status_label
    if status in {"done", "completed", "closed", "finished"}:
        return "Выполнена"
    if status in {"warehouse", "on_warehouse", "storekeeper"}:
        return "В ожидании поставки товара"
    return _repair_mojibake_text(payload.get("status") or "-") or "-"


def _journal_next_step_label(entry) -> str:
    payload = entry.payload or {}
    for key in ("next_step_label", "next_step", "next_action_label"):
        value = _repair_mojibake_text(payload.get(key) or "").strip()
        if value:
            return value
    return ""


def _next_step_label_from_entry(entry, audience: str = "default") -> str:
    payload = entry.payload or {}
    order_type = str(getattr(entry, "order_type", "") or "").strip().lower()
    order_id = str(getattr(entry, "order_id", "") or "")
    agency = getattr(entry, "agency", None)
    if order_type == "receiving":
        return WarehouseGoodsStateResolver.resolve_for_receiving_order(
            order_id=order_id,
            agency=agency,
            payload=payload,
        ).next_step_for(audience)
    if order_type == "processing":
        return WarehouseGoodsStateResolver.resolve_for_processing_order(
            order_id=order_id,
            agency=agency,
            payload=payload,
        ).next_step_for(audience)
    return ""


def _next_step_tone_from_label(label: str | None) -> str:
    normalized = str(label or "").strip().lower()
    if normalized == "выдай задание ричтракеру":
        return "danger"
    if normalized == "ожидаем ричтракер":
        return "pending"
    return ""


def _is_draft_entry(entry) -> bool:
    payload = entry.payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").lower()
    return status_value == "draft" or "черновик" in status_label


def _receiving_status_bits(entries) -> tuple[str, str]:
    if not entries:
        return "", ""
    status_entry = _current_status_entry(entries) or entries[-1]
    payload = status_entry.payload or {}
    status_value = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").strip().lower()
    return status_value, status_label


def _is_awaiting_manager_receiving(entries) -> bool:
    status_value, status_label = _receiving_status_bits(entries)
    return status_value in {"sent_unconfirmed", "send", "submitted"} or "подтверждени" in status_label


def _is_receiving_at_or_after_warehouse(entries) -> bool:
    if not entries:
        return False
    status_entry = _current_status_entry(entries) or entries[-1]
    if _warehouse_status_from_entry(status_entry):
        return True
    status_value, status_label = _receiving_status_bits(entries)
    if status_value in {"warehouse", "on_warehouse", "storekeeper"}:
        return True
    if "товар принят" in status_label:
        return True
    return False


def _can_client_edit_draft(entries, client_agency) -> bool:
    if not client_agency or not entries:
        return False
    if not any(entry.agency_id == client_agency.id for entry in entries if entry.agency_id):
        return False
    status_entry = _current_status_entry(entries)
    return _is_draft_entry(status_entry or entries[-1])


def _can_client_edit_receiving(entries, client_agency) -> bool:
    """Client may directly edit only a receiving draft."""
    if not client_agency or not entries:
        return False
    if not any(entry.agency_id == client_agency.id for entry in entries if entry.agency_id):
        return False
    status_entry = _current_status_entry(entries) or entries[-1]
    return _is_draft_entry(status_entry)


def _can_manager_edit_receiving(entries) -> bool:
    """Manager may correct receiving until it is handed to warehouse."""
    if not entries:
        return False
    if _act_storekeeper_signed(entries) or _act_manager_signed(entries):
        return False
    if _is_receiving_at_or_after_warehouse(entries):
        return False
    return True


def _is_receiving_draft_entries(entries) -> bool:
    if not entries:
        return False
    status_entry = _current_status_entry(entries)
    return _is_draft_entry(status_entry or entries[-1])


def _is_status_entry(entry) -> bool:
    payload = entry.payload or {}
    if entry.action == "status":
        return True
    return bool(
        payload.get("status")
        or payload.get("status_label")
        or payload.get("submit_action")
    )


def _manager_label() -> str:
    manager = (
        Employee.objects.filter(role="manager", is_active=True)
        .order_by("full_name")
        .first()
    )
    return _short_name(manager.full_name) if manager else "-"


def _storekeeper_label() -> str:
    storekeeper = (
        Employee.objects.filter(role="storekeeper", is_active=True)
        .order_by("full_name")
        .first()
    )
    return _short_name(storekeeper.full_name) if storekeeper else "-"


def _processing_head_label() -> str:
    head = (
        Employee.objects.filter(role="processing_head", is_active=True)
        .order_by("full_name")
        .first()
    )
    return _short_name(head.full_name) if head else "-"


def _storekeeper_full_name() -> str:
    storekeeper = (
        Employee.objects.filter(role="storekeeper", is_active=True)
        .order_by("full_name")
        .first()
    )
    return storekeeper.full_name.strip() if storekeeper and storekeeper.full_name else ""


def _processing_head_full_name() -> str:
    head = (
        Employee.objects.filter(role="processing_head", is_active=True)
        .order_by("full_name")
        .first()
    )
    return head.full_name.strip() if head and head.full_name else ""


def _manager_full_name() -> str:
    manager = (
        Employee.objects.filter(role="manager", is_active=True)
        .order_by("full_name")
        .first()
    )
    return manager.full_name.strip() if manager and manager.full_name else ""


def _journal_action_label(entry, manager_label: str, storekeeper_label: str) -> str:
    payload = entry.payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").lower()
    order_type = str(getattr(entry, "order_type", "") or "").strip().lower()
    if status_value == "draft" or "черновик" in status_label:
        return "У клиента"
    if status_value in {"done", "completed", "closed", "finished"} or "выполн" in status_label or "заверш" in status_label:
        return "Выполнено"
    if order_type == "processing" or "обработ" in status_label:
        return "У руководителя обработки"
    if status_value in {"warehouse", "storekeeper"} or "склад" in status_label or "ожидании поставки" in status_label:
        if storekeeper_label and storekeeper_label != "-":
            return f"У кладовщика {storekeeper_label}"
        return "У кладовщика"
    if manager_label and manager_label != "-":
        return f"В работе у менеджера {manager_label}"
    return "В работе у менеджера"


def _status_label_from_entry(entry, audience: str = "default"):
    payload = entry.payload or {}
    order_type = str(getattr(entry, "order_type", "") or "").strip().lower()
    act = (payload.get("act") or "").lower()
    act_state = (payload.get("act_state") or "").lower()
    if act == "placement":
        if order_type == "receiving":
            return WarehouseGoodsStateResolver.resolve_for_receiving_order(
                order_id=str(getattr(entry, "order_id", "") or ""),
                agency=getattr(entry, "agency", None),
                payload=payload,
            ).label_for(audience)
        if act_state == "closed":
            return "Товар принят и размещен на складе"
        return "Размещение на складе"
    if order_type == "processing":
        return processing_stage_label(payload, fallback=payload.get("status_label") or payload.get("status") or "-")
    if order_type == "receiving":
        return WarehouseGoodsStateResolver.resolve_for_receiving_order(
            order_id=str(getattr(entry, "order_id", "") or ""),
            agency=getattr(entry, "agency", None),
            payload=payload,
        ).label_for(audience)
    client_response = (payload.get("act_client_response") or "").lower()
    if client_response == "confirmed":
        return "Акт приемки подтвержден клиентом"
    if client_response == "dispute":
        return "Клиент заявил разногласия по акту приемки"
    if payload.get("act_sent"):
        return "Акт отправлен клиенту" if not payload.get("act_viewed") else "Выполнена"
    if payload.get("act_storekeeper_signed") and not payload.get("act_manager_signed"):
        return "Принято складом, акт приемки отправлен менеджеру"
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").lower()
    if "взята в работу" in status_label:
        return "Взята в работу"
    if status_value in {"sent_unconfirmed", "send", "submitted"} or "подтверждени" in status_label:
        return "Ждет подтверждения"
    if "товар принят" in status_label:
        return _repair_mojibake_text(payload.get("status_label") or payload.get("status") or "-")
    if status_value in {"warehouse", "on_warehouse"} or "ожидании поставки" in status_label or "на складе" in status_label:
        return "В ожидании поставки товара"
    return _repair_mojibake_text(payload.get("status_label") or payload.get("status") or "-")


def _current_status_entry(entries):
    latest_draft = None
    for entry in reversed(entries):
        if not _is_status_entry(entry):
            continue
        # A stale browser tab may autosave a draft after the order was sent.
        # That technical draft must not roll the manager-facing status back.
        if _is_draft_entry(entry):
            latest_draft = latest_draft or entry
            continue
        return entry
    return latest_draft or (entries[-1] if entries else None)


def _latest_status_payload_value(entries, key: str, default=None):
    for entry in reversed(entries):
        if not _is_status_entry(entry):
            continue
        payload = entry.payload or {}
        value = payload.get(key)
        if value not in (None, ""):
            return value
    return default


def _detail_history_status_label(entry, *, audience: str = "default", current_entry=None, current_label: str | None = None):
    if current_entry is not None and getattr(entry, "id", None) == getattr(current_entry, "id", None) and current_label:
        return current_label
    if getattr(entry, "order_type", "") != "receiving":
        return _status_label_from_entry(entry, audience=audience)
    payload = entry.payload or {}
    act = str(payload.get("act") or "").strip().lower()
    status_value = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").strip()
    if act == "placement":
        return _repair_mojibake_text(status_label or payload.get("act_label") or "Акт размещения")
    if act == "receiving":
        return _repair_mojibake_text(status_label or payload.get("act_label") or "Акт приемки")
    normalized_label = status_label.lower()
    if status_value == "draft" or "черновик" in normalized_label:
        return "Черновик"
    if status_value in {"sent_unconfirmed", "send", "submitted"} or "подтверждени" in normalized_label:
        return "Ждет подтверждения"
    if status_value in {"warehouse", "on_warehouse"} or "ожидании поставки" in normalized_label:
        return "В ожидании поставки товара"
    if status_label:
        return status_label
    return _repair_mojibake_text(payload.get("status") or "-") or "-"


def _detail_history_actor_label(entry, *, client_view=False, client_label: str | None = None):
    if getattr(entry, "order_type", "") != "receiving" or getattr(entry, "action", "") != "status":
        return _history_actor_label(entry, client_view=client_view, client_label=client_label)
    payload = entry.payload or {}
    storekeeper_name = str(
        payload.get("storekeeper_name")
        or payload.get("act_storekeeper_name")
        or ""
    ).strip()
    if storekeeper_name:
        return f"Кладовщик {storekeeper_name}"
    manager_name = str(payload.get("manager_name") or payload.get("act_manager_name") or "").strip()
    if manager_name:
        return f"Менеджер {manager_name}"
    if entry.user or entry.agency:
        return _actor_label(entry.user, entry.agency, client_view=client_view)
    status_value = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
    status_label = _repair_mojibake_text(payload.get("status_label") or "").strip().lower()
    if status_value == "draft" or "черновик" in status_label:
        return "Клиент"
    if status_value in {"warehouse", "on_warehouse"} or any(token in status_label for token in ("склад", "прием", "приём", "ожидании поставки")):
        return "Кладовщик"
    return "Менеджер"


def _current_responsible_label(entry):
    if not entry:
        return "-"
    payload = entry.payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    agency = entry.agency
    if entry.order_type == "processing":
        processing_state = WarehouseGoodsStateResolver.resolve_for_processing_order(
            order_id=str(getattr(entry, "order_id", "") or ""),
            agency=agency,
            payload=payload,
        )
        if processing_state.code in {
            WarehouseStateCode.RESERVED_FOR_PROCESSING,
            WarehouseStateCode.MOVING_TO_PROCESSING,
            WarehouseStateCode.IN_PROCESSING_ZONE,
            WarehouseStateCode.PROCESSING_IN_PROGRESS,
        }:
            head_full = _processing_head_full_name()
            return (
                f"Руководитель обработки {head_full}"
                if head_full
                else "Руководитель обработки"
            )
        if processing_state.code == WarehouseStateCode.PLACED_AFTER_PROCESSING:
            return "Выполнено"
    if entry.order_type == "receiving":
        receiving_state = WarehouseGoodsStateResolver.resolve_for_receiving_order(
            order_id=str(getattr(entry, "order_id", "") or ""),
            agency=agency,
            payload=payload,
        )
        if receiving_state.code in {
            WarehouseStateCode.RECEIVED_UNPLACED,
            WarehouseStateCode.PLACED_IN_RECEIVING,
            WarehouseStateCode.STORED,
        }:
            stored_name = (payload.get("storekeeper_name") or "").strip()
            storekeeper = (
                _signed_employee_from_payload(payload, "storekeeper_employee_id")
                or _signed_employee_from_payload(payload, "act_storekeeper_employee_id")
            )
            if storekeeper and storekeeper.full_name:
                return f"Кладовщик {storekeeper.full_name.strip()}"
            if stored_name:
                return f"Кладовщик {stored_name}"
            storekeeper_full = _storekeeper_full_name()
            return f"Кладовщик {storekeeper_full}" if storekeeper_full else "Кладовщик"
    if status_value == "draft" or "черновик" in status_label:
        base = (agency.fio_agn or agency.agn_name) if agency else None
        base = _shorten_ip_name(base) if base else None
        return f"Клиент {base}" if base else "Клиент"
    if status_value in {"done", "completed", "closed", "finished"} or "выполн" in status_label:
        return "Выполнено"
    if status_value == "processing_head" or (
        "передан" in status_label and "обработ" in status_label
    ):
        head_full = _processing_head_full_name()
        return (
            f"Руководитель обработки {head_full}"
            if head_full
            else "Руководитель обработки"
        )
    if entry.order_type == "processing" and (
        status_value == "processing_in_work" or "взята" in status_label
    ):
        head_full = _processing_head_full_name()
        return (
            f"Руководитель обработки {head_full}"
            if head_full
            else "Руководитель обработки"
        )
    if status_value in {"warehouse", "on_warehouse"} or any(token in status_label for token in ("склад", "прием", "приём", "ожидании поставки")):
        stored_name = (payload.get("storekeeper_name") or "").strip()
        storekeeper = (
            _signed_employee_from_payload(payload, "storekeeper_employee_id")
            or _signed_employee_from_payload(payload, "act_storekeeper_employee_id")
        )
        if storekeeper and storekeeper.full_name:
            return f"Кладовщик {storekeeper.full_name.strip()}"
        if stored_name:
            return f"Кладовщик {stored_name}"
        storekeeper_full = _storekeeper_full_name()
        return f"Кладовщик {storekeeper_full}" if storekeeper_full else "Кладовщик"
    manager = _manager_full_name()
    return f"Менеджер {manager}" if manager else "Менеджер"


def _employee_actor_label(user):
    if not user or not getattr(user, "pk", None):
        return None
    employee = (
        Employee.objects.filter(user_id=user.pk, is_active=True)
        .only("full_name", "role")
        .first()
    )
    if not employee:
        return None
    name = (
        employee.full_name
        or user.get_full_name()
        or user.get_username()
        or str(user)
    ).strip()
    role_label = str(employee.get_role_display() or "").strip()
    return f"{role_label} {name}".strip()


def _actor_label(user=None, agency=None, client_view=False):
    if user:
        employee_label = _employee_actor_label(user)
        if employee_label:
            return employee_label
        base = user.get_full_name() or user.get_username() or str(user)
        base = _shorten_ip_name(base)
        if getattr(user, "is_staff", False):
            return base
        return f"Клиент {base}"
    if client_view and agency:
        base = agency.fio_agn or agency.agn_name or "клиент"
        base = _shorten_ip_name(base)
        return f"Клиент {base}"
    if agency:
        base = agency.fio_agn or agency.agn_name or "клиент"
        base = _shorten_ip_name(base)
        return f"Клиент {base}"
    return "-"


def _history_actor_label(entry, client_view=False, client_label: str | None = None):
    if not entry:
        return "-"
    if entry.action == "status":
        payload = entry.payload or {}
        if entry.order_type == "processing":
            processing_stage = str(payload.get("processing_stage") or "").strip()
            status_value = str(payload.get("status") or payload.get("submit_action") or "").strip().lower()
            status_label = _repair_mojibake_text(payload.get("status_label") or "").strip().lower()
            description = _repair_mojibake_text(getattr(entry, "description", "") or "").strip().lower()
            approval_text = f"{status_label} {description}"
            is_manager_approval = (
                processing_stage == "manager_approved"
                or (
                    status_value == "processing_head"
                    and "утвержд" in approval_text
                    and "менедж" in approval_text
                )
            )
            if is_manager_approval:
                approved_by_name = str(payload.get("approved_by_name") or "").strip()
                if approved_by_name:
                    return f"Менеджер {approved_by_name}"
                if entry.user:
                    return _actor_label(entry.user, entry.agency, client_view=client_view)
            processing_state = WarehouseGoodsStateResolver.resolve_for_processing_order(
                order_id=str(getattr(entry, "order_id", "") or ""),
                agency=getattr(entry, "agency", None),
                payload=payload,
            )
            if processing_state.code in {
                WarehouseStateCode.RESERVED_FOR_PROCESSING,
                WarehouseStateCode.MOVING_TO_PROCESSING,
                WarehouseStateCode.IN_PROCESSING_ZONE,
                WarehouseStateCode.PROCESSING_IN_PROGRESS,
            }:
                head_full = _processing_head_full_name()
                return (
                    f"Руководитель обработки {head_full}"
                    if head_full
                    else "Руководитель обработки"
                )
        if entry.order_type == "receiving":
            receiving_state = WarehouseGoodsStateResolver.resolve_for_receiving_order(
                order_id=str(getattr(entry, "order_id", "") or ""),
                agency=getattr(entry, "agency", None),
                payload=payload,
            )
            if receiving_state.code in {
                WarehouseStateCode.RECEIVED_UNPLACED,
                WarehouseStateCode.PLACED_IN_RECEIVING,
                WarehouseStateCode.STORED,
            }:
                stored_name = (payload.get("storekeeper_name") or "").strip()
                storekeeper = (
                    _signed_employee_from_payload(payload, "storekeeper_employee_id")
                    or _signed_employee_from_payload(payload, "act_storekeeper_employee_id")
                )
                if storekeeper and storekeeper.full_name:
                    return f"Кладовщик {storekeeper.full_name.strip()}"
                if stored_name:
                    return f"Кладовщик {stored_name}"
                storekeeper_full = _storekeeper_full_name()
                return f"Кладовщик {storekeeper_full}" if storekeeper_full else "Кладовщик"
    if entry.action == "create":
        if entry.agency:
            base = entry.agency.fio_agn or entry.agency.agn_name or "клиент"
            base = _shorten_ip_name(base)
            return f"Клиент {base}"
        if client_label and client_label != "-":
            return f"Клиент {client_label}"
        if client_view:
            return "Клиент"
    if entry.action == "update":
        employee_label = _employee_actor_label(entry.user)
        if employee_label:
            return employee_label
        if entry.user:
            base = entry.user.get_full_name() or entry.user.get_username() or str(entry.user)
            base = base.strip()
            if base:
                lowered = base.lower()
                if lowered in {"manager", "менеджер"}:
                    manager_full = _manager_full_name()
                    if manager_full:
                        return f"Менеджер {manager_full}"
                if len(base.split()) >= 2:
                    return f"Менеджер {base}"
                short = _short_name(base)
                if short and short.lower() not in {"manager", "менеджер"}:
                    return f"Менеджер {short}"
        manager_full = _manager_full_name()
        return f"Менеджер {manager_full}" if manager_full else "Менеджер"
    return _actor_label(entry.user, entry.agency, client_view=client_view)


def _history_action_label(entry):
    payload = entry.payload or {}
    act = (payload.get("act") or "").lower()
    if entry.action == "status" and act == "receiving":
        return payload.get("act_label") or "Акт приемки"
    if entry.action == "status" and act == "placement":
        return payload.get("act_label") or "Акт размещения"
    return entry.get_action_display()


def _parse_qty_value(raw: str | None) -> int | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        return int(text)
    except (TypeError, ValueError):
        return None


def _build_receiving_summary(
    *,
    receiving_mode: str,
    receiving_started: bool = True,
    status_label: str,
    display_items: list | None,
    placement_pallets: list | None,
    placement_loose_boxes: list | None,
    act_payload: dict | None = None,
    placement_payload: dict | None = None,
    order_payload: dict | None = None,
    finished_at=None,
) -> dict:
    """Сводка для блока «Результат приемки» на карточке заявки."""
    mode = str(receiving_mode or "").strip().lower()
    is_cz = mode == "cz"
    if not receiving_started:
        mode_badge = "Не выбран"
        mode_label = "Режим будет выбран при начале приемки"
    else:
        mode_badge = "ЧЗ" if is_cz else "Обычная"
        mode_label = "Приемка с ЧЗ" if is_cz else "Обычная приемка"

    items = [row for row in (display_items or []) if isinstance(row, dict)]
    planned_total = 0
    accepted_total = 0
    has_any_actual = False
    mismatch_count = 0
    for row in items:
        planned = _parse_qty_value(row.get("qty"))
        if planned is None:
            planned = 0
        planned_total += planned
        actual = _parse_qty_value(row.get("actual_qty"))
        if actual is None:
            continue
        has_any_actual = True
        accepted_total += actual
        if actual != planned:
            mismatch_count += 1

    item_actual_total = accepted_total if has_any_actual else "—"

    act_payload = act_payload if isinstance(act_payload, dict) else {}
    placement_payload = placement_payload if isinstance(placement_payload, dict) else {}
    order_payload = order_payload if isinstance(order_payload, dict) else {}

    pallets = list(placement_pallets or [])
    loose_boxes = list(placement_loose_boxes or [])
    if not pallets and placement_payload.get("act_pallets"):
        pallets = list(placement_payload.get("act_pallets") or [])
    # Плоский act_boxes без pallet_code НЕ подмешиваем в loose: те же короба
    # уже лежат во вложенных boxes палет — иначе счётчик удваивается.

    pallet_count = len(pallets)
    if not pallet_count:
        pallet_count = len([
            row for row in (placement_payload.get("act_pallets") or []) if isinstance(row, dict)
        ]) or len([
            row for row in (act_payload.get("act_pallets") or []) if isinstance(row, dict)
        ])

    def _box_identity(box) -> str:
        if isinstance(box, str):
            return str(box).strip()
        if isinstance(box, dict):
            code = str(box.get("code") or box.get("barcode") or "").strip()
            if code:
                return code
        return ""

    # Факт коробов: уникальные коды. Иначе act_boxes без pallet_code
    # дублируют вложенные boxes палет (915 → 1830).
    box_identities: set[str] = set()
    anonymous_boxes = 0

    def _register_box(box) -> None:
        nonlocal anonymous_boxes
        identity = _box_identity(box)
        if identity:
            box_identities.add(identity)
        elif isinstance(box, dict) or isinstance(box, str):
            anonymous_boxes += 1

    for pallet in pallets:
        boxes = getattr(pallet, "boxes", None) if not isinstance(pallet, dict) else pallet.get("boxes")
        if boxes is None and isinstance(pallet, dict):
            boxes = []
        for box in boxes or []:
            _register_box(box)

    # Плоский act_boxes — только если ещё не собрали короба с палет,
    # либо как дополнение уникальными кодами (не второй раз те же).
    flat_boxes = []
    for source in (placement_payload, act_payload):
        raw_boxes = source.get("act_boxes") or []
        if raw_boxes:
            flat_boxes = [box for box in raw_boxes if isinstance(box, dict) or isinstance(box, str)]
            break

    if loose_boxes:
        for box in loose_boxes:
            _register_box(box)
    elif flat_boxes:
        if box_identities or anonymous_boxes:
            for box in flat_boxes:
                identity = _box_identity(box)
                if identity and identity not in box_identities:
                    box_identities.add(identity)
        else:
            # Нет структуры палет — считаем плоский список факта.
            for box in flat_boxes:
                _register_box(box)

    box_count = len(box_identities) + anonymous_boxes
    if receiving_started and not box_count:
        expected = _parse_qty_value(
            act_payload.get("expected_boxes")
            or order_payload.get("expected_boxes")
        )
        if expected:
            box_count = expected

    if receiving_started and not has_any_actual:
        # Приёмка без товарных позиций: считаем места/короба как принятое.
        fallback_qty = _parse_qty_value(
            act_payload.get("accepted_total")
            or act_payload.get("qty_total")
            or act_payload.get("expected_boxes")
            or order_payload.get("expected_boxes")
        )
        if fallback_qty is None and box_count:
            fallback_qty = box_count
        if fallback_qty is not None:
            accepted_total = fallback_qty
            has_any_actual = True

    chz_units = 0
    for source in (act_payload, placement_payload):
        units = source.get("act_units") or []
        if units:
            chz_units = len([unit for unit in units if isinstance(unit, dict)])
            break
    if not receiving_started:
        chz_label = "—"
    elif is_cz:
        chz_label = str(chz_units) if chz_units else ("Да" if has_any_actual else "—")
    else:
        chz_label = "Нет"

    if mismatch_count:
        mismatch_badge = f"Есть · {mismatch_count}"
        mismatch_tone = "warn"
    elif has_any_actual and planned_total and accepted_total == planned_total:
        mismatch_badge = "Нет"
        mismatch_tone = "ok"
    elif has_any_actual and planned_total == 0:
        mismatch_badge = "—"
        mismatch_tone = ""
    else:
        mismatch_badge = "—"
        mismatch_tone = ""

    finished_label = "—"
    if finished_at:
        finished_label = _format_datetime_value(finished_at) or "—"
    else:
        for key in (
            "act_manager_signed_at",
            "act_storekeeper_signed_at",
            "finished_at",
            "completed_at",
            "closed_at",
        ):
            raw = act_payload.get(key) or placement_payload.get(key)
            if raw:
                finished_label = _format_datetime_value(raw) or str(raw)
                break

    status_short = "Не начата" if not receiving_started else (str(status_label or "").strip() or "—")
    if len(status_short) > 72:
        status_short = status_short[:69].rstrip() + "…"

    return {
        "mode_badge": mode_badge,
        "mode_label": mode_label,
        "status": status_short,
        "planned_total": planned_total if items else "—",
        "actual_total": item_actual_total if items else "—",
        "accepted_total": accepted_total if has_any_actual else "—",
        "pallet_count": pallet_count or "—",
        "box_count": box_count or "—",
        "chz_label": chz_label,
        "mismatch_badge": mismatch_badge,
        "mismatch_tone": mismatch_tone,
        "finished_at": finished_label,
    }


def _processing_card_id(card: dict) -> str:
    if not isinstance(card, dict):
        return ""
    value = card.get("id") or card.get("article") or card.get("sku") or ""
    return str(value).strip()


def _processing_card_sets(payload: dict) -> tuple[set[str], set[str]]:
    payload = payload or {}
    processed = set()
    placed = set()
    for card in payload.get("cards") or []:
        card_id = _processing_card_id(card)
        if not card_id:
            continue
        if card.get("processed_at") or card.get("processed_done") or card.get("processed"):
            processed.add(card_id)
        if card.get("placed_at") or card.get("placed_done"):
            placed.add(card_id)
    for value in payload.get("processed_cards") or []:
        text = str(value).strip()
        if text:
            processed.add(text)
    for value in payload.get("placed_cards") or []:
        text = str(value).strip()
        if text:
            placed.add(text)
    return processed, placed


def _payload_has_value(payload: dict, *keys: str) -> bool:
    payload = payload or {}
    for key in keys:
        value = payload.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text and text not in {"-", "0", "0.0"}:
            return True
    return False


def _processing_receiving_items(
    payload: dict,
    agency_id: int | None,
    only_cards: set[str] | None = None,
) -> list[dict]:
    totals: dict[tuple[str, str], int] = {}
    names: dict[tuple[str, str], str] = {}
    allowed_articles = set()
    allowed_card_ids = {str(value or "").strip().lower() for value in (only_cards or set()) if str(value or "").strip()}

    def result_article_from_row(row: dict, card: dict | None = None) -> str:
        card = card if isinstance(card, dict) else {}
        row = row if isinstance(row, dict) else {}
        return str(
            row.get("result_article")
            or card.get("result_article")
            or row.get("article")
            or row.get("sku")
            or card.get("article")
            or card.get("sku")
            or ""
        ).strip()

    def result_name_from_row(row: dict, card: dict | None = None) -> str:
        card = card if isinstance(card, dict) else {}
        row = row if isinstance(row, dict) else {}
        return str(
            row.get("result_product_name")
            or card.get("result_product_name")
            or row.get("product_name")
            or row.get("name")
            or card.get("product_name")
            or card.get("name")
            or ""
        ).strip()

    if only_cards:
        for card in payload.get("cards") or []:
            card_id = _processing_card_id(card)
            if not card_id or card_id not in only_cards:
                continue
            article = str(card.get("article") or card.get("sku") or "").strip()
            if article:
                allowed_articles.add(article.lower())
            for row in card.get("rows") or []:
                if not isinstance(row, dict):
                    continue
                result_article = result_article_from_row(row, card)
                if result_article:
                    allowed_articles.add(result_article.lower())

    def is_allowed_row(row: dict, article_value: str) -> bool:
        if not only_cards:
            return True
        row_card_id = str((row or {}).get("card_id") or "").strip().lower()
        if row_card_id and row_card_id in allowed_card_ids:
            return True
        if not allowed_articles:
            return False
        return article_value.lower() in allowed_articles

    results = payload.get("processing_results") or []
    if isinstance(results, list) and results:
        for row in results:
            if not isinstance(row, dict):
                continue
            sku = str(row.get("article") or row.get("sku") or row.get("sku_code") or "").strip()
            if not sku or not is_allowed_row(row, sku):
                continue
            size = str(row.get("size") or "").strip()
            processed_qty = _parse_qty_value(row.get("processed")) or 0
            shipped = _parse_qty_value(row.get("shipped_qty")) or 0
            qty = max(processed_qty - shipped, 0)
            if qty <= 0:
                continue
            key = (sku, size)
            totals[key] = totals.get(key, 0) + qty
            name = result_name_from_row(row)
            if name and key not in names:
                names[key] = name
    if not totals:
        if only_cards:
            for card in payload.get("cards") or []:
                card_id = _processing_card_id(card)
                if not card_id or card_id not in only_cards:
                    continue
                for row in card.get("rows") or []:
                    if not isinstance(row, dict):
                        continue
                    sku = result_article_from_row(row, card)
                    if not sku:
                        continue
                    size = str(row.get("size") or "").strip()
                    qty = _parse_qty_value(row.get("qty")) or 0
                    if qty <= 0:
                        continue
                    key = (sku, size)
                    totals[key] = totals.get(key, 0) + qty
                    name = result_name_from_row(row, card)
                    if name and key not in names:
                        names[key] = name
        else:
            for row in extract_processing_items(payload):
                if not isinstance(row, dict):
                    continue
                sku = str(row.get("result_article") or row.get("sku_code") or row.get("article") or row.get("sku") or "").strip()
                if not sku:
                    continue
                size = str(row.get("size") or "").strip()
                qty = _parse_qty_value(row.get("qty")) or 0
                if qty <= 0:
                    continue
                key = (sku, size)
                totals[key] = totals.get(key, 0) + qty
                name = str(row.get("result_product_name") or row.get("name") or row.get("product_name") or "").strip()
                if name and key not in names:
                    names[key] = name
    if not totals:
        return []
    sku_codes = {sku for (sku, _) in totals.keys() if sku}
    sku_map = {}
    if sku_codes and agency_id:
        for sku in SKU.objects.filter(agency_id=agency_id, sku_code__in=sku_codes, deleted=False):
            sku_map[sku.sku_code] = sku
    items = []
    for (sku, size), qty in totals.items():
        sku_obj = sku_map.get(sku)
        name = names.get((sku, size)) or ((sku_obj.name or "").strip() if sku_obj else "")
        items.append(
            {
                "sku_code": sku,
                "name": name or "-",
                "size": size,
                "actual_qty": qty,
            }
        )
    return items


def _parse_int_value(raw) -> int:
    try:
        return int(str(raw).strip())
    except (TypeError, ValueError):
        return 0


def _normalize_zone_code(raw: str) -> str:
    text = (raw or "").strip()
    if not text:
        return ""
    if re.search(r"^pr$", text, re.IGNORECASE) or re.search(
        r"зона приемки|поле приемки", text, re.IGNORECASE
    ):
        return "PR"
    if re.search(r"^otg?$", text, re.IGNORECASE) or re.search(
        r"зона отгрузки|отгрузк", text, re.IGNORECASE
    ):
        return "OTG"
    if re.search(r"^mr$", text, re.IGNORECASE) or re.search(
        r"между ряд", text, re.IGNORECASE
    ):
        return "MR"
    if re.search(r"^os$", text, re.IGNORECASE) or re.search(
        r"основн|стеллаж|ряд|полк|секци|ярус|ячейк", text, re.IGNORECASE
    ):
        return "OS"
    return text.upper()


def _location_parts(location_value, pallet=None):
    pallet = pallet or {}
    zone = ""
    row = 0
    section = 0
    tier = 0
    cell = 0
    if isinstance(location_value, dict):
        zone = _normalize_zone_code(location_value.get("zone") or "")
        row = _parse_int_value(location_value.get("row") or pallet.get("row"))
        section = _parse_int_value(location_value.get("section"))
        tier = _parse_int_value(location_value.get("tier"))
        cell = _parse_int_value(location_value.get("cell"))
    elif isinstance(location_value, str):
        zone = _normalize_zone_code(location_value)
    if not zone:
        zone = _normalize_zone_code(pallet.get("zone") or "")
    if zone == "OS" and not (row and section and tier and cell):
        row = row or _parse_int_value(pallet.get("row"))
        section = section or _parse_int_value(pallet.get("section"))
        tier = tier or _parse_int_value(pallet.get("tier"))
        cell = cell or _parse_int_value(pallet.get("cell"))
    if zone == "MR" and not row:
        row = _parse_int_value(pallet.get("row"))
    return {
        "zone": zone,
        "row": row,
        "section": section,
        "tier": tier,
        "cell": cell,
    }


def _item_key(sku: str | None, name: str | None, size: str | None) -> str:
    sku_part = (sku or "").strip().lower()
    name_part = (name or "").strip().lower()
    size_part = (size or "").strip().lower()
    return f"{sku_part}|{name_part}|{size_part}"


def _receiving_sku_map(agency_id: int | None, items: list[dict]) -> dict[str, SKU]:
    if not agency_id or not items:
        return {}
    sku_codes = {
        str(item.get("sku_code") or item.get("sku") or "").strip()
        for item in items
        if str(item.get("sku_code") or item.get("sku") or "").strip()
    }
    if not sku_codes:
        return {}
    result = {}
    for sku in SKU.objects.filter(agency_id=agency_id, deleted=False, sku_code__in=sku_codes):
        key = str(sku.sku_code or "").strip().lower()
        if key and key not in result:
            result[key] = sku
    return result


def _receiving_marked_items_map(agency_id: int | None, items: list[dict]) -> dict[str, dict]:
    sku_map = _receiving_sku_map(agency_id, items)
    result = {}
    fallback_result = {}
    for item in items or []:
        key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
        sku_code_key = str(item.get("sku_code") or item.get("sku") or "").strip().lower()
        sku = sku_map.get(sku_code_key)
        candidate = {
            "sku_code": str(item.get("sku_code") or item.get("sku") or "").strip(),
            "name": str(item.get("name") or "").strip(),
            "size": str(item.get("size") or "").strip(),
            "barcode": str(item.get("barcode") or "").strip() or _barcode_value_for_sku(sku, item.get("size")),
            "qty": _parse_qty_value(item.get("qty") or item.get("actual_qty")) or 0,
            "sku_id": sku.id if sku else None,
        }
        if sku and sku.honest_sign:
            result[key] = candidate
        else:
            fallback_result[key] = candidate
    return result or fallback_result


def _collect_receiving_marking_units(
    order_id: str,
    valid_box_codes: set[str],
    box_to_pallet: dict[str, str],
    marked_items_map: dict[str, dict],
) -> tuple[list[dict], dict[str, int], bool]:
    units = []
    totals = {}
    has_orphan_units = False
    marked_by_pair = {
        (
            str(item.get("sku_code") or "").strip().lower(),
            str(item.get("size") or "").strip().lower(),
        ): (item_key, item)
        for item_key, item in marked_items_map.items()
    }
    queryset = (
        MarkingCode.objects.filter(order_type="receiving", order_id=order_id, used_at__isnull=False)
        .order_by("used_at", "created_at", "id")
    )
    for code in queryset:
        matched = marked_by_pair.get(
            (
                str(code.sku_code or "").strip().lower(),
                str(code.size or "").strip().lower(),
            )
        )
        if not matched:
            has_orphan_units = True
            continue
        matched_key, item = matched
        box_code = str(code.box_barcode or "").strip()
        if not box_code or box_code not in valid_box_codes:
            has_orphan_units = True
            continue
        totals[matched_key] = int(totals.get(matched_key, 0)) + 1
        units.append(
            {
                "sku_code": item.get("sku_code") or "",
                "name": item.get("name") or "",
                "size": item.get("size") or "",
                "barcode": code.barcode or item.get("barcode") or "",
                "marking_code": code.code,
                "box_code": box_code,
                "pallet_code": box_to_pallet.get(box_code, ""),
                "qty": 1,
            }
        )
    return units, totals, has_orphan_units




def _warehouse_status_from_entry(entry) -> bool:
    if not entry:
        return False
    payload = entry.payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    return status_value in {"warehouse", "on_warehouse"} or "склад" in status_label or "ожидании поставки" in status_label


def _act_entry_from_entries(entries, act_type: str = "receiving"):
    for entry in reversed(entries or []):
        if (entry.payload or {}).get("act") == act_type:
            return entry
    return None


def _find_act_entry(entries, act_type: str, label_hint: str):
    return ReceivingWorkflowService._find_document_act_entry(entries, act_type, label_hint)


def _is_done_status(entry) -> bool:
    if not entry:
        return False
    payload = entry.payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    if status_value in {"done", "completed", "closed", "finished"}:
        return True
    return "выполн" in status_label


def _act_storekeeper_signed_from_payload(payload: dict) -> bool:
    return bool((payload or {}).get("act_storekeeper_signed"))


def _act_manager_signed_from_payload(payload: dict) -> bool:
    return bool((payload or {}).get("act_manager_signed"))


def _act_storekeeper_signed(entries) -> bool:
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    if not act_entry:
        return False
    return _act_storekeeper_signed_from_payload(act_entry.payload or {})


def _act_manager_signed(entries) -> bool:
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    if not act_entry:
        return False
    return _act_manager_signed_from_payload(act_entry.payload or {})


def _signed_employee_from_payload(payload: dict, employee_key: str) -> Employee | None:
    raw_id = (payload or {}).get(employee_key)
    if raw_id in (None, ""):
        return None
    try:
        employee_id = int(raw_id)
    except (TypeError, ValueError):
        return None
    return Employee.objects.filter(id=employee_id, is_active=True).first()


def _assigned_client_manager(agency) -> Employee | None:
    try:
        from client_cabinet.other_requests import resolve_client_manager

        manager = resolve_client_manager(agency)
        if manager:
            return manager
    except Exception:
        pass
    return _first_active_employee_by_roles("manager", "head_manager")


def _create_manager_sign_task(order_id, agency, request, observer=None):
    manager = _assigned_client_manager(agency)
    if not manager:
        return
    route = f"/orders/receiving/{order_id}/act/print/"
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}" if agency else "Клиент: -"
    due_date = timezone.localtime()
    for task in (
        Task.objects.filter(
            route=f"/orders/receiving/{order_id}/",
            assigned_to__role__in=["manager", "head_manager"],
        )
        .exclude(status="done")
        .select_related("assigned_to")
    ):
        task.status = "done"
        task.updated_at = due_date
        task.save(update_fields=["status", "updated_at"])
    existing = (
        Task.objects.filter(
            route=route,
            assigned_to__role__in=["manager", "head_manager"],
        )
        .exclude(status="done")
        .first()
    )
    if existing:
        existing.description = description
        existing.assigned_to = manager
        existing.observer = observer
        existing.due_date = due_date
        if existing.status == "done":
            existing.status = "in_progress"
        existing.save()
        return
    Task.objects.create(
        title=f"Подписать акт приемки по заявке №{order_id}",
        description=description,
        route=route,
        assigned_to=manager,
        observer=observer,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=due_date,
        status="in_progress",
    )


def _create_manager_client_response_task(order_id, agency, request, response: str):
    manager = _assigned_client_manager(agency)
    if not manager:
        return
    response_label = "подтвердил акт приемки" if response == "confirmed" else "заявил разногласия по акту приемки"
    title = f"Клиент {response_label} по заявке №{order_id}"
    route = f"/orders/receiving/{order_id}/act/print/"
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}" if agency else "Клиент: -"
    existing = (
        Task.objects.filter(
            route=route,
            assigned_to__role__in=["manager", "head_manager"],
            title=title,
        )
        .exclude(status="done")
        .first()
    )
    if existing:
        existing.description = description
        existing.assigned_to = manager
        existing.due_date = timezone.localtime()
        if existing.status == "done":
            existing.status = "in_progress"
        existing.save()
        return
    Task.objects.create(
        title=title,
        description=description,
        route=route,
        assigned_to=manager,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=timezone.localtime(),
        status="in_progress",
    )
def _latest_payload_from_entries(entries):
    return build_order_runtime_payload("receiving", entries=entries)




def _flow_state_from_entries(entries):
    for entry in reversed(entries or []):
        payload = entry.payload or {}
        flow_state = payload.get("flow_state")
        if isinstance(flow_state, dict):
            return flow_state
        boxes = payload.get("flow_boxes")
        pallets = payload.get("flow_pallets")
        if boxes or pallets:
            return {
                "boxes": boxes or [],
                "pallets": pallets or [],
                "activeBox": payload.get("flow_active_box") or "",
                "activePallet": payload.get("flow_active_pallet") or "",
            }
    return {}


def _receiving_flow_draft_payload(flow_state):
    state = flow_state if isinstance(flow_state, dict) else {}
    boxes = []
    nonempty_box_codes = set()
    for raw_box in state.get("boxes") or []:
        if not isinstance(raw_box, dict):
            continue
        items = [
            dict(item)
            for item in (raw_box.get("items") or [])
            if isinstance(item, dict) and (_parse_qty_value(item.get("qty")) or 0) > 0
        ]
        if not items:
            continue
        box = dict(raw_box)
        box["items"] = items
        boxes.append(box)
        code = str(box.get("code") or "").strip()
        if code:
            nonempty_box_codes.add(code)

    pallets = []
    for raw_pallet in state.get("pallets") or []:
        if not isinstance(raw_pallet, dict):
            continue
        direct_items = [
            dict(item)
            for item in (raw_pallet.get("items") or [])
            if isinstance(item, dict) and (_parse_qty_value(item.get("qty")) or 0) > 0
        ]
        box_refs = []
        for raw_box in raw_pallet.get("boxes") or []:
            code = str(raw_box.get("code") if isinstance(raw_box, dict) else raw_box or "").strip()
            if code and code in nonempty_box_codes:
                box_refs.append(raw_box)
        if not direct_items and not box_refs:
            continue
        pallet = dict(raw_pallet)
        pallet["items"] = direct_items
        pallet["boxes"] = box_refs
        pallets.append(pallet)
    return {"act_boxes": boxes, "act_pallets": pallets}


def _flow_closed_from_entries(entries):
    for entry in reversed(entries or []):
        payload = entry.payload or {}
        if payload.get("flow_closed"):
            return True
        if payload.get("flow_reopened"):
            return False
    return False


_RECEIVING_WAREHOUSE_ROLES = {"storekeeper", "processing_head"}
_RECEIVING_MOVE_CREATE_ROLES = {"storekeeper", "processing_head", "manager", "head_manager", "director", "admin"}
_RECEIVING_MOVE_BLOCKING_STATUSES = {"created", "in_progress", "done"}


def _stock_move_location_label(location: dict | None) -> str:
    parts = _location_parts(location or {})
    zone = _normalize_zone_code(parts.get("zone") or "") or "PR"
    row = _parse_int_value(parts.get("row"))
    section = _parse_int_value(parts.get("section"))
    tier = _parse_int_value(parts.get("tier"))
    cell = _parse_int_value(parts.get("cell"))
    if zone == "PR":
        return "PR · Зона приемки"
    if zone == "OTG":
        return "OTG · Зона отгрузки"
    if zone == "MR":
        return f"MR · Между рядами · Ряд {row}" if row else "MR · Между рядами"
    if zone == "OS":
        if row and section and tier and cell:
            return f"OS · Ряд {row} · Секция {section} · Ярус {tier} · Ячейка {cell}"
        if row:
            return f"OS · Ряд {row}"
        return "OS · Основной склад"
    return zone


def _normalize_receiving_move_location(raw_location, fallback_payload=None) -> dict:
    return ReceivingWorkflowService._normalize_receiving_move_location(raw_location, fallback_payload)


def _suggest_receiving_destinations(
    pallets,
    *,
    exclude_order_type: str,
    exclude_order_id: str,
    agency_id: int | None = None,
) -> dict[str, dict]:
    return ReceivingWorkflowService.suggest_receiving_destinations(
        pallets,
        exclude_order_type=exclude_order_type,
        exclude_order_id=exclude_order_id,
        agency_id=agency_id,
        row_sections=_OS_ROW_SECTIONS,
        tiers=_OS_TIERS,
        cells_per_tier=_OS_CELLS_PER_TIER,
    )


def _receiving_warehouse_move_rows(order_id: str, placement_pallets, *, agency_id: int | None = None) -> list[dict]:
    return ReceivingWorkflowService.build_receiving_warehouse_move_rows(
        order_id,
        placement_pallets,
        agency_id=agency_id,
        row_sections=_OS_ROW_SECTIONS,
        tiers=_OS_TIERS,
        cells_per_tier=_OS_CELLS_PER_TIER,
    )


def _parse_receiving_warehouse_destinations(raw) -> tuple[dict[str, dict] | None, str]:
    return ReceivingWorkflowService.parse_receiving_warehouse_destinations(raw)


def _next_stock_move_number() -> str:
    max_number = 0
    for value in (
        OrderAuditEntry.objects.filter(order_type="stock_move")
        .values_list("order_id", flat=True)
        .distinct()
    ):
        text = str(value or "").strip()
        if not text.isdigit():
            continue
        number = int(text)
        if number > max_number:
            max_number = number
    next_number = max_number + 1
    while OrderAuditEntry.objects.filter(order_type="stock_move", order_id=str(next_number)).exists():
        next_number += 1
    return str(next_number)


def _latest_receiving_moves_by_pallet(order_id: str) -> dict[str, dict]:
    return ReceivingWorkflowService._latest_receiving_moves_by_pallet(order_id)


def _receiving_warehouse_move_progress(order_id: str, placement_pallets) -> dict:
    return ReceivingWorkflowService.build_receiving_warehouse_move_progress(order_id, placement_pallets)


def _create_receiving_warehouse_moves(
    order_id: str,
    entries: list[OrderAuditEntry],
    request,
    destinations_by_pallet: dict[str, dict] | None = None,
) -> tuple[int, int, int, int]:
    result = ReceivingWorkflowService.create_receiving_warehouse_moves(
        order_id=str(order_id or ""),
        entries=entries,
        role=get_request_role(request) or "storekeeper",
        user=request.user if request.user.is_authenticated else None,
        destinations_by_pallet=destinations_by_pallet,
    )
    return (
        int(result.created_count or 0),
        int(result.skipped_existing_count or 0),
        int(result.skipped_missing_destination_count or 0),
        int(result.total_count or 0),
    )


def _client_agency_from_request(request):
    if not request.user.is_authenticated:
        return None
    role = get_request_role(request)
    if is_staff_role(role):
        return None
    agency = resolve_portal_agency_for_user(request.user, required_section=SECTION_REQUESTS)
    if agency is None:
        return None
    client_id = request.GET.get("client") or request.GET.get("agency")
    if client_id and str(agency.pk) != str(client_id):
        return None
    return agency


_RECEIVING_ATTACHMENT_ALLOWED_STAFF_ROLES = {
    "manager",
    "head_manager",
    "processing_head",
    "storekeeper",
    "admin",
    "accountant",
    "project_manager",
}


def _active_receiving_attachments(order_id: str, agency: Agency | None):
    order_key = str(order_id or "").strip()
    if not order_key or agency is None:
        return ReceivingOrderAttachment.objects.none()
    cutoff = timezone.now() - timedelta(days=ReceivingOrderAttachment.RETENTION_DAYS)
    return (
        ReceivingOrderAttachment.objects.select_related("uploaded_by")
        .filter(order_id=order_key, agency=agency, uploaded_at__gte=cutoff)
        .order_by("-uploaded_at")
    )


def _receiving_attachment_names(order_id: str, agency: Agency | None) -> list[str]:
    return [attachment.filename for attachment in _active_receiving_attachments(order_id, agency)]


def _save_receiving_attachments(order_id: str, agency: Agency | None, request) -> list[str]:
    order_key = str(order_id or "").strip()
    if not order_key or agency is None:
        return []
    uploaded_names: list[str] = []
    for uploaded_file in request.FILES.getlist("documents"):
        if not uploaded_file or not getattr(uploaded_file, "name", ""):
            continue
        attachment = ReceivingOrderAttachment.objects.create(
            order_id=order_key,
            agency=agency,
            uploaded_by=request.user if request.user.is_authenticated else None,
            file=uploaded_file,
        )
        uploaded_names.append(attachment.filename)
    return uploaded_names


def _can_access_receiving_attachment(request, attachment: ReceivingOrderAttachment) -> bool:
    client_agency = _client_agency_from_request(request)
    if client_agency is not None:
        return int(client_agency.id or 0) == int(attachment.agency_id or 0)
    role = get_request_role(request)
    return role in _RECEIVING_ATTACHMENT_ALLOWED_STAFF_ROLES


@login_required
def receiving_attachment_download(request, order_id: str, attachment_id: int):
    attachment = get_object_or_404(
        ReceivingOrderAttachment.objects.select_related("uploaded_by", "agency"),
        pk=attachment_id,
        order_id=str(order_id or "").strip(),
    )
    if not _can_access_receiving_attachment(request, attachment):
        return HttpResponseForbidden("Доступ запрещен")
    if attachment.is_expired or not attachment.file:
        raise Http404("Файл недоступен")
    return FileResponse(
        attachment.file.open("rb"),
        as_attachment=True,
        filename=attachment.filename,
    )


@login_required
def receiving_shipping_return_order_search(request, order_id: str):
    if get_request_role(request) != "storekeeper":
        return JsonResponse({"ok": False, "error": "forbidden"}, status=403)
    latest = (
        OrderAuditEntry.objects.filter(
            order_id=str(order_id or "").strip(),
            order_type="receiving",
        )
        .select_related("agency")
        .order_by("-created_at", "-id")
        .first()
    )
    if latest is None or latest.agency is None:
        return JsonResponse({"ok": False, "error": "order_not_found"}, status=404)
    results = []
    for source_order in search_shipping_return_orders(
        agency=latest.agency,
        query=request.GET.get("q") or "",
    ):
        items = shipping_return_source_items(
            source_order,
            exclude_receiving_order_id=str(order_id or ""),
        )
        available_qty = sum(int(item.get("qty") or 0) for item in items)
        if available_qty <= 0:
            continue
        shipped_at = source_order.shipped_at or source_order.updated_at
        results.append(
            {
                "number": source_order.number,
                "display": format_order_number("shipping", source_order.number),
                "status": source_order.status,
                "shipped_at": (
                    timezone.localtime(shipped_at).strftime("%d.%m.%Y")
                    if shipped_at
                    else ""
                ),
                "available_qty": available_qty,
                "item_count": len(items),
                "box_count": int(source_order.expected_boxes or 0),
            }
        )
    return JsonResponse({"ok": True, "results": results})


_PHONE_DIGITS_RE = re.compile(r"\D+")


def _normalize_driver_phone(value: str) -> str:
    digits = _PHONE_DIGITS_RE.sub("", value or "")
    if not digits:
        return ""
    if len(digits) == 10 and digits.startswith("9"):
        digits = f"7{digits}"
    elif len(digits) == 11 and digits.startswith("8"):
        digits = f"7{digits[1:]}"
    elif len(digits) == 11 and digits.startswith("7"):
        digits = f"7{digits[1:]}"
    else:
        return ""
    return f"+{digits}"


def _is_valid_driver_phone(value: str) -> bool:
    return bool(_normalize_driver_phone(value))


def _min_receiving_eta(now):
    current = business_localtime(now)
    base = current.replace(second=0, microsecond=0)
    remainder = base.minute % 5
    if remainder or current.second or current.microsecond:
        base += timedelta(minutes=(5 - remainder) if remainder else 5)
    return base


def _default_receiving_eta(now):
    return _min_receiving_eta(business_localtime(now) + timedelta(hours=1))


def _format_payload_value(value):
    if value is None or value == "":
        return "-"
    text = _repair_mojibake_text(value).strip()
    return text or "-"


def _repair_mojibake_text(value) -> str:
    text = str(value or "")
    if not text:
        return ""
    for broken, fixed in _MOJIBAKE_FIXUPS.items():
        text = text.replace(broken, fixed)

    repair_steps = (
        ("latin-1", "utf-8"),
        ("latin-1", "cp1251"),
        ("cp1251", "utf-8"),
        ("cp1252", "utf-8"),
    )
    mojibake_chars = "ÐÑÂÃðàáâãäå¸êëìíîïñòóôõö÷øùúûüýþÿ"
    mojibake_markers = (
        "Р ",
        "РЎ",
        "РІвЂ",
        "РІР‚",
        "Рљ",
        "Рџ",
        "Рґ",
        "Рµ",
        "Рѕ",
        "Р°",
        "СЃ",
        "С‚",
        "вЂ",
        "в„",
        "Ð",
        "Ñ",
        "Â",
    )

    def _raw_bytes(raw: str, source_encoding: str) -> bytes | None:
        chunks: list[bytes] = []
        for char in raw:
            code = ord(char)
            if source_encoding == "latin-1":
                if code > 0xFF:
                    return None
                chunks.append(bytes([code]))
                continue
            if source_encoding == "cp1251" and 0x80 <= code <= 0x9F:
                chunks.append(bytes([code]))
                continue
            try:
                chunks.append(char.encode(source_encoding))
            except UnicodeEncodeError:
                return None
        return b"".join(chunks)

    def _candidate(raw: str, source_encoding: str, target_encoding: str) -> str | None:
        raw_bytes = _raw_bytes(raw, source_encoding)
        if raw_bytes is None:
            return None
        try:
            return raw_bytes.decode(target_encoding)
        except UnicodeDecodeError:
            return None

    def _quality(raw: str) -> tuple[int, int, int, int]:
        mojibake = raw.count("\ufffd") * 3
        mojibake += sum(raw.count(marker) * 4 for marker in mojibake_markers)
        mojibake += sum(3 for char in raw if 0x80 <= ord(char) <= 0x9F)
        mojibake += sum(2 for char in raw if char in mojibake_chars)
        cyrillic = sum(1 for ch in raw if "\u0400" <= ch <= "\u04FF")
        readable = sum(1 for ch in raw if ch.isalnum() or ch.isspace() or ch in "-_:;.,/()[]{}?#+\"'")
        return (mojibake, -cyrillic, -readable, len(raw))

    queue = [text]
    seen: set[str] = set()
    variants: list[str] = []
    while queue and len(seen) < 80:
        current = queue.pop(0)
        if current in seen:
            continue
        seen.add(current)
        variants.append(current)
        for encoding, decoding in repair_steps:
            repaired = _candidate(current, encoding, decoding)
            if repaired and repaired not in seen and _quality(repaired) <= _quality(current):
                queue.append(repaired)
    result = min(variants, key=_quality)
    for broken, fixed in _MOJIBAKE_FIXUPS.items():
        result = result.replace(broken, fixed)
    return result


def _format_party_label(name: str | None, address: str | None, phone: str | None) -> str:
    parts = []
    if name:
        parts.append(name.strip())
    if address:
        parts.append(address.strip())
    if phone:
        parts.append(f"тел.: {phone.strip()}")
    return ", ".join([part for part in parts if part]) or "-"


def _format_payload_list(value):
    if value is None or value == "":
        return "-"
    if isinstance(value, (list, tuple, set)):
        items = [_repair_mojibake_text(item).strip() for item in value if _repair_mojibake_text(item).strip()]
        return ", ".join(items) if items else "-"
    text = _repair_mojibake_text(value).strip()
    return text or "-"


def _resolve_choice_value(value, other):
    other_text = (other or "").strip()
    if other_text:
        return other_text
    return (value or "").strip()


def _place_type_label(value: str) -> str:
    text = _format_payload_value(value)
    if text == "-":
        return text
    labels = {
        "pallet": "Паллет",
        "box": "Короб",
        "bag": "Мешок",
    }
    return labels.get(text, text)


def _format_datetime_value(value):
    text = _format_payload_value(value)
    if text == "-":
        return text
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return text
    return parsed.strftime("%d.%m.%Y, %H:%M")


_ISO_DATETIME_RE = re.compile(r"\b\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2})?\b")


def _format_message_text(text: str) -> str:
    if not text:
        return ""
    text = _repair_mojibake_text(text)

    def _replace(match):
        raw = match.group(0)
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError:
            return raw
        return parsed.strftime("%d.%m.%Y, %H:%M")

    return _ISO_DATETIME_RE.sub(_replace, text)


def _is_manager_staff_role(role: str | None) -> bool:
    return role in {"manager", "head_manager", "director", "admin", "developer"}


def _order_cancel_payload(payload: dict | None, *, reason: str, actor: str = "manager") -> dict:
    data = dict(payload or {})
    data["status"] = "cancelled"
    data["submit_action"] = "cancelled"
    data["status_label"] = "Отменена"
    data["cancel_reason"] = str(reason or "").strip()
    data["cancelled_by"] = actor
    data["cancelled_at"] = timezone.localtime().isoformat()
    return data


def _close_order_tasks_for_cancel(order_type: str, order_id: str) -> None:
    route = f"/orders/{order_type}/{order_id}/"
    Task.objects.filter(route=route).exclude(status="done").update(status="done", updated_at=timezone.localtime())


def _normalize_header(value) -> str:
    text = str(value or "").strip()
    if not text or text.lower() == "nan":
        return ""
    return re.sub(r"\s+", " ", text).lower()


def _clean_cell(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    return text


def _header_index(header_map: dict, *names: str):
    for name in names:
        idx = header_map.get(name)
        if idx is not None:
            return idx
    return None


_RECEIVING_SKU_HEADERS = (
    "артикул",
    "артикул заказчика",
    "sku",
    "sku товара",
    "артикул (sku)",
    "артикул sku",
)
_RECEIVING_QTY_HEADERS = (
    "количество",
    "кол-во",
    "колво",
    "количество товара",
    "кол-во товара",
    "кол-во, шт",
    "количество, шт",
)
_RECEIVING_BARCODE_HEADERS = (
    "штрихкод",
    "баркод",
    "шк",
    "шк товара",
    "штрихкод товара",
)


def _load_template_rows(uploaded_file):
    import pandas as pd

    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass
    df = pd.read_excel(uploaded_file, header=None, dtype=str)
    if df.empty:
        return {}, []
    header_map = {}
    header_row = df.iloc[0].tolist()
    for idx, value in enumerate(header_row):
        name = _normalize_header(value)
        if not name:
            continue
        header_map.setdefault(name, idx)
    rows = df.iloc[1:].fillna("")
    return header_map, rows


def _template_type_for_header(header_map: dict) -> str | None:
    if _header_index(header_map, *_RECEIVING_QTY_HEADERS) is not None and (
        _header_index(header_map, *_RECEIVING_BARCODE_HEADERS) is not None
        or _header_index(header_map, *_RECEIVING_SKU_HEADERS) is not None
    ):
        return "receiving"
    if "артикул заказчика" in header_map:
        return "sku"
    return None


def _apply_sku_template(rows, header_map: dict, agency: Agency):
    sku_code_idx = _header_index(header_map, "артикул заказчика", "артикул")
    barcode_idx = _header_index(header_map, "баркод", "штрихкод", "шк товара", "шк")
    name_idx = _header_index(header_map, "предмет", "наименование", "товар")
    size_idx = _header_index(header_map, "размер")
    brand_idx = _header_index(header_map, "бренд")
    color_idx = _header_index(header_map, "цвет")
    composition_idx = _header_index(header_map, "состав")
    gender_idx = _header_index(header_map, "пол")
    season_idx = _header_index(header_map, "сезон")
    made_in_idx = _header_index(header_map, "страна пр-ва", "страна производства")
    img_idx = _header_index(header_map, "ссылка на товар", "ссылка")

    if sku_code_idx is None:
        return ["В шаблоне номенклатуры не найдена колонка «Артикул Заказчика»."]

    errors = []
    for _, row in rows.iterrows():
        values = [ _clean_cell(value) for value in row.tolist() ]
        if not any(values):
            continue
        sku_code = _clean_cell(row.iloc[sku_code_idx])
        if not sku_code:
            continue
        name = _clean_cell(row.iloc[name_idx]) if name_idx is not None else ""
        size = _clean_cell(row.iloc[size_idx]) if size_idx is not None else ""
        barcode = _clean_cell(row.iloc[barcode_idx]) if barcode_idx is not None else ""
        brand = _clean_cell(row.iloc[brand_idx]) if brand_idx is not None else ""
        color = _clean_cell(row.iloc[color_idx]) if color_idx is not None else ""
        composition = _clean_cell(row.iloc[composition_idx]) if composition_idx is not None else ""
        gender = _clean_cell(row.iloc[gender_idx]) if gender_idx is not None else ""
        season = _clean_cell(row.iloc[season_idx]) if season_idx is not None else ""
        made_in = _clean_cell(row.iloc[made_in_idx]) if made_in_idx is not None else ""
        img = _clean_cell(row.iloc[img_idx]) if img_idx is not None else ""

        sku = SKU.objects.filter(agency=agency, sku_code=sku_code, deleted=False).first()
        if not sku:
            sku = SKU.objects.create(
                agency=agency,
                sku_code=sku_code,
                name=name or sku_code,
                brand=brand or None,
                size=size or None,
                color=color or None,
                composition=composition or None,
                gender=gender or None,
                season=season or None,
                made_in=made_in or None,
                img=img or None,
                source="manual",
            )
        else:
            updated = False
            if name and not sku.name:
                sku.name = name
                updated = True
            if brand and not sku.brand:
                sku.brand = brand
                updated = True
            if size and not sku.size:
                sku.size = size
                updated = True
            if color and not sku.color:
                sku.color = color
                updated = True
            if composition and not sku.composition:
                sku.composition = composition
                updated = True
            if gender and not sku.gender:
                sku.gender = gender
                updated = True
            if season and not sku.season:
                sku.season = season
                updated = True
            if made_in and not sku.made_in:
                sku.made_in = made_in
                updated = True
            if img and not sku.img:
                sku.img = img
                updated = True
            if updated:
                sku.save(update_fields=[
                    "name",
                    "brand",
                    "size",
                    "color",
                    "composition",
                    "gender",
                    "season",
                    "made_in",
                    "img",
                    "updated_at",
                ])

        if barcode:
            existing_barcode = (
                SKUBarcode.objects.select_related("sku")
                .filter(agency=agency, value=barcode)
                .first()
            )
            if existing_barcode and existing_barcode.sku_id != sku.id:
                errors.append(f"ШК {barcode} уже привязан к SKU {existing_barcode.sku.sku_code}.")
                continue
            if not existing_barcode:
                SKUBarcode.objects.create(
                    sku=sku,
                    value=barcode,
                    size=size or None,
                    is_primary=not sku.barcodes.exists(),
                )

    return errors


def _parse_receiving_template(rows, header_map: dict, agency: Agency):
    barcode_idx = _header_index(header_map, *_RECEIVING_BARCODE_HEADERS)
    qty_idx = _header_index(header_map, *_RECEIVING_QTY_HEADERS)
    sku_code_idx = _header_index(header_map, *_RECEIVING_SKU_HEADERS)

    if qty_idx is None or (barcode_idx is None and sku_code_idx is None):
        return [], ["В шаблоне приемки нужны колонки «Артикул/SKU» или «Штрихкод» и «Количество»."]

    sku_map = {
        sku.sku_code: sku
        for sku in SKU.objects.filter(agency=agency, deleted=False).prefetch_related("barcodes")
    }
    barcode_map = {
        barcode.value: barcode
        for barcode in SKUBarcode.objects.select_related("sku").filter(sku__agency=agency, sku__deleted=False)
    }

    items = []
    missing = []
    for _, row in rows.iterrows():
        values = [_clean_cell(value) for value in row.tolist()]
        if not any(values):
            continue
        sku_code = _clean_cell(row.iloc[sku_code_idx]) if sku_code_idx is not None else ""
        barcode = _clean_cell(row.iloc[barcode_idx]) if barcode_idx is not None else ""
        qty_raw = _clean_cell(row.iloc[qty_idx]) if qty_idx is not None else ""
        qty_value = _parse_qty_value(qty_raw)
        if not sku_code and not barcode and qty_value is None:
            continue
        if qty_value is None or qty_value <= 0:
            continue

        sku = None
        size_value = ""
        if sku_code:
            sku = sku_map.get(sku_code)
        if not sku and barcode:
            barcode_entry = barcode_map.get(barcode)
            if barcode_entry and barcode_entry.sku and not barcode_entry.sku.deleted:
                sku = barcode_entry.sku
                size_value = (barcode_entry.size or "").strip()
        if not sku:
            missing.append(sku_code or barcode)
            continue
        if not size_value:
            size_value = (sku.size or "").strip()

        items.append(
            {
                "sku_id": str(sku.id),
                "sku_code": sku.sku_code,
                "name": sku.name,
                "size": size_value,
                "qty": str(qty_value),
                "comment": "",
            }
        )

    if missing:
        sample = ", ".join(missing[:5])
        suffix = "..." if len(missing) > 5 else ""
        return [], [f"В шаблоне приемки нет позиций в номенклатуре: {sample}{suffix}"]

    return items, []


def _merge_items(items: list[dict]) -> list[dict]:
    merged = {}
    for item in items:
        sku_code = (item.get("sku_code") or "").strip()
        size = (item.get("size") or "").strip()
        if sku_code:
            key = f"{sku_code.lower()}|{size.lower()}"
        else:
            key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
        qty_value = _parse_qty_value(item.get("qty"))
        if key not in merged:
            merged[key] = dict(item)
            if qty_value is not None:
                merged[key]["qty"] = str(qty_value)
            continue
        existing_qty = _parse_qty_value(merged[key].get("qty")) or 0
        merged_qty = existing_qty + (qty_value or 0)
        merged[key]["qty"] = str(merged_qty)
        if not merged[key].get("sku_id") and item.get("sku_id"):
            merged[key]["sku_id"] = item.get("sku_id")
        for field in ("name", "brand", "color", "size", "barcode", "comment"):
            if not merged[key].get(field) and item.get(field):
                merged[key][field] = item.get(field)
    return list(merged.values())


def _process_template_uploads(template_files, agency: Agency):
    allowed_ext = {".xlsx", ".xls"}
    sku_file = None
    receiving_file = None

    for uploaded_file in template_files:
        if not uploaded_file or not getattr(uploaded_file, "name", ""):
            continue
        ext = "." + uploaded_file.name.split(".")[-1].lower()
        if ext not in allowed_ext:
            raise ValueError(f"Файл «{uploaded_file.name}» не является Excel-шаблоном.")
        header_map, rows = _load_template_rows(uploaded_file)
        template_type = _template_type_for_header(header_map)
        if not template_type:
            raise ValueError(f"Файл «{uploaded_file.name}» не похож на шаблон приемки или номенклатуры.")
        if template_type == "sku":
            if sku_file:
                raise ValueError("Можно загрузить только один шаблон номенклатуры.")
            sku_file = (rows, header_map, uploaded_file.name)
        elif template_type == "receiving":
            if receiving_file:
                raise ValueError("Можно загрузить только один шаблон приемки.")
            receiving_file = (rows, header_map, uploaded_file.name)

    template_names = []
    for entry in (sku_file, receiving_file):
        if entry:
            template_names.append(entry[2])

    if sku_file:
        rows, header_map, _ = sku_file
        errors = _apply_sku_template(rows, header_map, agency)
        if errors:
            raise ValueError("; ".join(errors))

    items = []
    if receiving_file:
        rows, header_map, _ = receiving_file
        items, errors = _parse_receiving_template(rows, header_map, agency)
        if errors:
            raise ValueError("; ".join(errors))

    return items, template_names


def _safe_doc_name(value: str) -> str:
    text = re.sub(r"[^A-Za-z0-9_-]+", "_", str(value or "")).strip("_")
    return text or "act"


def _format_doc_date(value: str | None) -> str:
    if not value:
        return business_localtime().strftime("%d.%m.%Y")
    try:
        parsed = parse_business_datetime(value)
    except ValueError:
        return str(value)
    return parsed.strftime("%d.%m.%Y")


def _agency_label(agency: Agency | None) -> str:
    if not agency:
        return "-"
    parts = []
    name = (agency.agn_name or agency.fio_agn or "").strip()
    if name:
        parts.append(name)
    inn = (agency.inn or "").strip()
    if inn:
        parts.append(f"ИНН {inn}")
    return ", ".join(parts) if parts else f"Клиент {agency.id}"


def _placement_box_map(placement_payload: dict | None) -> dict:
    items = (placement_payload or {}).get("act_items") or []
    box_map = {}
    for item in items:
        key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
        if not key:
            continue
        box_map[key] = _parse_qty_value(item.get("box_qty")) or 0
    return box_map


def _placement_box_count_map(placement_payload: dict | None) -> dict:
    boxes = (placement_payload or {}).get("act_boxes") or []
    box_counts: dict[str, int] = {}
    for box in boxes:
        if not isinstance(box, dict):
            continue
        keys = set()
        for item in (box.get("items") or []):
            if not isinstance(item, dict):
                continue
            sku_code = item.get("sku") or item.get("sku_code")
            key = _item_key(sku_code, item.get("name"), item.get("size"))
            if key:
                keys.add(key)
        for key in keys:
            box_counts[key] = box_counts.get(key, 0) + 1
    return box_counts


def _placement_pallet_count_map(placement_payload: dict | None) -> dict:
    pallets = (placement_payload or {}).get("act_pallets") or []
    pallet_counts: dict[str, int] = {}
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        keys = set()
        for item in (pallet.get("items") or []):
            if not isinstance(item, dict):
                continue
            sku_code = item.get("sku") or item.get("sku_code")
            key = _item_key(sku_code, item.get("name"), item.get("size"))
            if key:
                keys.add(key)
        for key in keys:
            pallet_counts[key] = pallet_counts.get(key, 0) + 1
    return pallet_counts


def _placement_payload_item_rows(placement_payload: dict | None) -> list[dict]:
    payload = placement_payload if isinstance(placement_payload, dict) else {}
    totals: dict[str, dict] = {}

    def add_item(raw_item):
        if not isinstance(raw_item, dict):
            return
        sku_code = str(raw_item.get("sku_code") or raw_item.get("sku") or "").strip()
        name = str(raw_item.get("name") or "").strip()
        size = str(raw_item.get("size") or "").strip()
        key = _item_key(sku_code, name, size)
        if not key:
            return
        qty = _parse_qty_value(raw_item.get("qty") or raw_item.get("actual_qty")) or 0
        if qty <= 0:
            return
        entry = totals.setdefault(
            key,
            {
                "sku_code": sku_code,
                "name": name,
                "size": size,
                "qty": None,
                "actual_qty": 0,
                "comment": raw_item.get("comment"),
                "barcode": str(raw_item.get("barcode") or "").strip(),
            },
        )
        entry["actual_qty"] += qty
        if not entry.get("barcode"):
            entry["barcode"] = str(raw_item.get("barcode") or "").strip()
        if not entry.get("comment"):
            entry["comment"] = raw_item.get("comment")

    for box in payload.get("act_boxes") or []:
        if not isinstance(box, dict):
            continue
        for item in box.get("items") or []:
            add_item(item)
    for pallet in payload.get("act_pallets") or []:
        if not isinstance(pallet, dict):
            continue
        for item in pallet.get("items") or []:
            add_item(item)

    return sorted(
        totals.values(),
        key=lambda item: (
            str(item.get("sku_code") or "").strip().lower(),
            str(item.get("name") or "").strip().lower(),
            str(item.get("size") or "").strip().lower(),
        ),
    )


def _placement_payload_container_rows(placement_payload: dict | None) -> tuple[list[dict], list[dict]]:
    payload = placement_payload if isinstance(placement_payload, dict) else {}
    raw_boxes = payload.get("act_boxes") or []
    raw_pallets = payload.get("act_pallets") or []

    def item_qty_total(items) -> int:
        total = 0
        for raw_item in items or []:
            if not isinstance(raw_item, dict):
                continue
            total += _parse_qty_value(raw_item.get("qty") or raw_item.get("actual_qty")) or 0
        return total

    def normalize_items(items) -> list[dict]:
        prepared: list[dict] = []
        for raw_item in items or []:
            if not isinstance(raw_item, dict):
                continue
            prepared.append(
                {
                    "sku_code": str(raw_item.get("sku_code") or raw_item.get("sku") or "").strip(),
                    "name": str(raw_item.get("name") or "").strip(),
                    "size": str(raw_item.get("size") or "").strip(),
                    "barcode": str(raw_item.get("barcode") or "").strip(),
                    "qty": _parse_qty_value(raw_item.get("qty") or raw_item.get("actual_qty")) or 0,
                }
            )
        return prepared

    def normalize_box(raw_box) -> dict | None:
        if not isinstance(raw_box, dict):
            return None
        code = str(raw_box.get("code") or "").strip()
        if not code:
            return None
        items = normalize_items(raw_box.get("items") or [])
        return {
            "code": code,
            "items": items,
            "items_count": len(items),
            "qty_total": sum(int(item.get("qty") or 0) for item in items),
        }

    all_boxes: list[dict] = []
    boxes_by_code: dict[str, dict] = {}
    for raw_box in raw_boxes:
        box_row = normalize_box(raw_box)
        if not box_row:
            continue
        all_boxes.append(box_row)
        boxes_by_code[box_row["code"]] = box_row

    pallet_rows: list[dict] = []
    used_box_codes: set[str] = set()
    for raw_pallet in raw_pallets:
        if not isinstance(raw_pallet, dict):
            continue
        pallet_code = str(raw_pallet.get("code") or "").strip()
        if not pallet_code:
            continue
        pallet_boxes: list[dict] = []
        for raw_box in raw_pallet.get("boxes") or []:
            if isinstance(raw_box, dict):
                box_row = normalize_box(raw_box)
            else:
                box_code = str(raw_box or "").strip()
                box_row = dict(boxes_by_code.get(box_code) or {"code": box_code, "items_count": 0, "qty_total": 0})
            if not box_row or not str(box_row.get("code") or "").strip():
                continue
            used_box_codes.add(box_row["code"])
            pallet_boxes.append(box_row)
        pallet_rows.append(
            {
                "code": pallet_code,
                "box_count": len(pallet_boxes),
                "boxes": pallet_boxes,
                "direct_items": normalize_items(raw_pallet.get("items") or []),
                "direct_qty_total": item_qty_total(raw_pallet.get("items") or []),
            }
        )

    loose_boxes = [box for box in all_boxes if box["code"] not in used_box_codes]
    return pallet_rows, loose_boxes


def _receiving_act_payload_from_placement(entries, placement_entry) -> dict:
    placement_payload = (placement_entry.payload or {}) if placement_entry else {}
    fallback_items = _placement_payload_item_rows(placement_payload)
    if not fallback_items:
        return {}
    status_entry = _current_status_entry(entries)
    if status_entry:
        payload = dict(status_entry.payload or {})
    elif entries:
        payload = dict(entries[-1].payload or {})
    else:
        payload = {}
    payload.update(
        {
            "act": "receiving",
            "act_state": "closed",
            "act_items": fallback_items,
        }
    )
    return payload


def _receiving_act_entry_or_placement(entries, act_entry):
    if act_entry and ((act_entry.payload or {}).get("act_items") or (act_entry.payload or {}).get("act_units")):
        return act_entry
    placement_entry = _find_act_entry(entries, "placement", "")
    payload = _receiving_act_payload_from_placement(entries, placement_entry)
    if not payload:
        return None
    source_entry = placement_entry or (entries[-1] if entries else None)
    if not source_entry:
        return None
    fallback_entry = type("FallbackReceivingActEntry", (), {})()
    fallback_entry.payload = payload
    fallback_entry.agency = getattr(source_entry, "agency", None)
    fallback_entry.agency_id = getattr(source_entry, "agency_id", None)
    fallback_entry.created_at = getattr(source_entry, "created_at", timezone.localtime())
    return fallback_entry


def _has_receiving_act_print_source(entries, act_entry=None, placement_entry=None) -> bool:
    if act_entry and ((act_entry.payload or {}).get("act_items") or (act_entry.payload or {}).get("act_units")):
        return True
    placement_entry = placement_entry or _find_act_entry(entries, "placement", "")
    return bool(_receiving_act_payload_from_placement(entries, placement_entry).get("act_items"))


def _build_receiving_special_box_rows(act_payload: dict | None, placement_payload: dict | None) -> list[dict]:
    act_source = act_payload if isinstance(act_payload, dict) else {}
    placement_source = placement_payload if isinstance(placement_payload, dict) else {}
    default_goods_type = ""
    for payload_source in (act_source, placement_source):
        if not isinstance(payload_source, dict):
            continue
        default_goods_type = ReceivingWorkflowService._normalize_flow_goods_type(
            payload_source.get("goods_type"),
            default_goods_type,
        )
        if default_goods_type:
            break
    return ReceivingWorkflowService.build_special_box_rows(
        boxes=placement_source.get("act_boxes") or [],
        pallets=placement_source.get("act_pallets") or [],
        default_goods_type=default_goods_type,
    )


def _build_receiving_special_item_rows(act_payload: dict | None, placement_payload: dict | None) -> list[dict]:
    act_source = act_payload if isinstance(act_payload, dict) else {}
    placement_source = placement_payload if isinstance(placement_payload, dict) else {}
    default_goods_type = ""
    for payload_source in (act_source, placement_source):
        if not isinstance(payload_source, dict):
            continue
        default_goods_type = ReceivingWorkflowService._normalize_flow_goods_type(
            payload_source.get("goods_type"),
            default_goods_type,
        )
        if default_goods_type:
            break
    return ReceivingWorkflowService.build_special_item_rows(
        boxes=placement_source.get("act_boxes") or [],
        default_goods_type=default_goods_type,
    )


def _act_items_with_barcodes(
    agency_id: int | None,
    act_items: list[dict],
    placement_box_map: dict | None = None,
) -> list[dict]:
    sku_codes = set()
    for item in act_items:
        sku_code = str(item.get("sku_code") or "").strip()
        if sku_code:
            sku_codes.add(sku_code)

    sku_by_code = {}
    if agency_id and sku_codes:
        sku_qs = SKU.objects.filter(agency_id=agency_id, sku_code__in=sku_codes, deleted=False).prefetch_related("barcodes")
        for sku in sku_qs:
            sku_by_code[sku.sku_code] = sku

    rows = []
    for item in act_items:
        sku_code = str(item.get("sku_code") or "").strip()
        size = str(item.get("size") or "").strip()
        name = str(item.get("name") or "").strip()
        barcode = str(item.get("barcode") or "").strip()
        if not barcode and sku_code:
            sku = sku_by_code.get(sku_code)
            barcode = _barcode_value_for_sku(sku, size)
        else:
            sku = sku_by_code.get(sku_code)
        brand = str(item.get("brand") or (sku.brand if sku else "") or "").strip()
        color = str(item.get("color") or (sku.color if sku else "") or "").strip()
        planned = _parse_qty_value(item.get("planned_qty") or item.get("qty")) or 0
        actual = _parse_qty_value(item.get("actual_qty") or item.get("qty")) or 0
        key = _item_key(sku_code, name, size)
        box_qty = None
        if placement_box_map and key in placement_box_map:
            box_qty = placement_box_map.get(key)
        rows.append(
            {
                "sku_code": sku_code,
                "name": name,
                "size": size,
                "barcode": barcode,
                "brand": brand,
                "color": color,
                "planned": planned,
                "actual": actual,
                "box_qty": box_qty,
            }
        )
    return rows


def _merge_row_cells(ws, row: int, left: str, right: str):
    cell_range = f"{left}{row}:{right}{row}"
    if cell_range in ws.merged_cells:
        return
    ws.merge_cells(cell_range)


def _set_cell_value(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def _find_row_with_text(ws, needles: tuple[str, ...], start_row: int) -> int | None:
    lowered = [needle.lower() for needle in needles]
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            text = cell.value.strip().lower()
            if not text:
                continue
            if any(needle in text for needle in lowered):
                return cell.row
    return None


def _fill_receiving_act_sheet(
    ws,
    order_id: str,
    agency: Agency | None,
    act_items: list[dict],
    eta_raw: str | None,
    expected_boxes: int | None = None,
    special_boxes: list[dict] | None = None,
):
    doc_date = _format_doc_date(eta_raw)
    ws["B7"] = f"Приемка № {order_id} от {doc_date}"
    ws["B9"] = f"Поставщик: {_agency_label(agency)}"

    start_row = 14
    footer_row = _find_row_with_text(ws, ("итого",), start_row) or ws.max_row + 1
    current_rows = max(0, footer_row - start_row)
    item_count = len(act_items)
    if item_count < current_rows:
        ws.delete_rows(start_row + item_count, current_rows - item_count)
    elif item_count > current_rows:
        ws.insert_rows(footer_row, item_count - current_rows)
    footer_row = start_row + item_count
    for row in range(start_row, footer_row):
        for col in range(2, 13):
            _set_cell_value(ws, row, col, None)
    for idx, item in enumerate(act_items):
        row = start_row + idx
        _merge_row_cells(ws, row, "C", "D")
        _merge_row_cells(ws, row, "E", "I")
        ws[f"B{row}"] = idx + 1
        ws[f"C{row}"] = item.get("barcode") or "-"
        name_parts = [item.get("name") or "-"]
        if item.get("sku_code"):
            name_parts.append(f"арт.: {item['sku_code']}")
        if item.get("size"):
            name_parts.append(f"р-р {item['size']}")
        ws[f"E{row}"] = ", ".join(name_parts)
        ws[f"J{row}"] = item.get("planned", 0)
        ws[f"K{row}"] = item.get("actual", 0)
        box_qty = item.get("box_qty")
        if box_qty is not None:
            ws[f"L{row}"] = box_qty
    total_planned = sum(item.get("planned", 0) or 0 for item in act_items)
    total_actual = sum(item.get("actual", 0) or 0 for item in act_items)
    total_boxes = sum((item.get("box_qty") or 0) for item in act_items if item.get("box_qty") is not None)
    if not total_boxes and expected_boxes:
        total_boxes = expected_boxes
    ws[f"J{footer_row}"] = total_planned
    ws[f"K{footer_row}"] = total_actual
    if total_boxes:
        ws[f"L{footer_row}"] = total_boxes


def _fill_mx1_sheet_one(ws, order_id: str, agency: Agency | None, act_items: list[dict]):
    ws["H18"] = order_id
    ws["C12"] = _agency_label(agency)

    start_row = 29
    capacity = 20
    items = act_items[:capacity]
    footer_row = _find_row_with_text(ws, ("итого", "èòîãî"), start_row) or (start_row + capacity)
    current_rows = max(0, footer_row - start_row)
    item_count = len(items)
    if item_count < current_rows:
        ws.delete_rows(start_row + item_count, current_rows - item_count)
    elif item_count > current_rows:
        ws.insert_rows(footer_row, item_count - current_rows)
    footer_row = start_row + item_count
    for row in range(start_row, footer_row):
        for col in range(2, 13):
            _set_cell_value(ws, row, col, None)
    for idx, item in enumerate(items):
        row = start_row + idx
        _merge_row_cells(ws, row, "B", "C")
        _merge_row_cells(ws, row, "G", "H")
        _merge_row_cells(ws, row, "I", "K")
        _merge_row_cells(ws, row, "L", "O")
        ws[f"B{row}"] = idx + 1
        name_parts = [item.get("name") or "-"]
        if item.get("sku_code"):
            name_parts.append(f"арт.: {item['sku_code']}")
        ws[f"D{row}"] = ", ".join(name_parts)
        characteristic = item.get("size") or ""
        if characteristic:
            ws[f"F{row}"] = f"размер {characteristic}"
        ws[f"G{row}"] = "шт"
        ws[f"I{row}"] = "796"
        ws[f"L{row}"] = item.get("actual", 0)
    ws[f"L{footer_row}"] = sum(item.get("actual", 0) or 0 for item in items)


def _fill_mx1_sheet_two(ws, act_items: list[dict], offset: int):
    start_row = 7
    capacity = 23
    items = act_items[offset:offset + capacity]
    footer_row = _find_row_with_text(ws, ("условия хранения",), start_row) or (start_row + capacity)
    current_rows = max(0, footer_row - start_row)
    item_count = len(items)
    if item_count < current_rows:
        ws.delete_rows(start_row + item_count, current_rows - item_count)
    elif item_count > current_rows:
        ws.insert_rows(footer_row, item_count - current_rows)
    footer_row = start_row + item_count
    for row in range(start_row, footer_row):
        for col in range(2, 15):
            _set_cell_value(ws, row, col, None)
    for idx, item in enumerate(items):
        row = start_row + idx
        _merge_row_cells(ws, row, "C", "F")
        _merge_row_cells(ws, row, "H", "I")
        _merge_row_cells(ws, row, "J", "L")
        ws[f"B{row}"] = offset + idx + 1
        name_parts = [item.get("name") or "-"]
        if item.get("sku_code"):
            name_parts.append(f"арт.: {item['sku_code']}")
        ws[f"C{row}"] = ", ".join(name_parts)
        characteristic = item.get("size") or ""
        if characteristic:
            ws[f"H{row}"] = f"размер {characteristic}"
        ws[f"J{row}"] = "шт"
        ws[f"M{row}"] = "796"
        ws[f"N{row}"] = item.get("actual", 0)


def _ensure_act_documents(
    order_id: str,
    agency: Agency | None,
    act_payload: dict,
    placement_payload: dict | None = None,
):
    if not _ACT_TEMPLATE_FILE.exists() or not _MX1_TEMPLATE_FILE.exists():
        raise Http404("Шаблоны актов не найдены.")
    placement_box_map = _placement_box_map(placement_payload)
    act_items = _act_items_with_barcodes(
        agency.id if agency else None,
        act_payload.get("act_items") or [],
        placement_box_map,
    )
    special_item_rows = _build_receiving_special_item_rows(act_payload, placement_payload)
    act_items = ReceivingWorkflowService.split_receiving_act_rows_by_goods_type(
        act_items,
        special_rows=special_item_rows,
        planned_key="planned",
        actual_key="actual",
        box_key="box_qty",
    )
    expected_boxes = _parse_qty_value(act_payload.get("expected_boxes"))
    safe_id = _safe_doc_name(order_id)
    target_dir = Path(_ACT_DOCS_DIR) / safe_id
    target_dir.mkdir(parents=True, exist_ok=True)
    act_path = target_dir / f"act_receiving_{safe_id}.xlsx"
    mx1_path = target_dir / f"mx1_{safe_id}.xlsx"
    wb = load_workbook(_ACT_TEMPLATE_FILE)
    ws = wb.active
    _fill_receiving_act_sheet(
        ws,
        order_id,
        agency,
        act_items,
        act_payload.get("eta_at"),
        expected_boxes,
    )
    wb.save(act_path)
    wb = load_workbook(_MX1_TEMPLATE_FILE)
    sheet_one = wb["МХ-1 (1стр)"]
    sheet_two = wb["МХ-1(2стр)"]
    _fill_mx1_sheet_one(sheet_one, order_id, agency, act_items)
    _fill_mx1_sheet_two(sheet_two, act_items, 20)
    wb.save(mx1_path)
    return act_path, mx1_path


def _load_order_entries(order_id: str, order_type: str = "receiving"):
    return list(
        OrderAuditEntry.objects.filter(order_id=order_id, order_type=order_type).order_by("created_at")
    )


def _act_access_allowed(request, entries):
    if not request.user.is_authenticated:
        return False
    client_agency = _client_agency_from_request(request)
    if client_agency:
        return any(entry.agency_id == client_agency.id for entry in entries if entry.agency_id)
    role = get_request_role(request)
    return request.user.is_staff or role in {"storekeeper", "manager", "head_manager", "director", "admin"}


def _placement_closed(entries) -> bool:
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    if not placement_entry:
        return False
    state = ((placement_entry.payload or {}).get("act_state") or "closed").lower()
    return state == "closed"


def download_receiving_act_doc(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    if not _placement_closed(entries):
        return redirect(f"/orders/receiving/{order_id}/act/?error=placement_required")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    act_path, _ = _ensure_act_documents(
        order_id,
        act_entry.agency,
        act_entry.payload or {},
        placement_entry.payload if placement_entry else None,
    )
    return FileResponse(open(act_path, "rb"), as_attachment=True, filename=act_path.name)


def download_receiving_act_mx1(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    if not _placement_closed(entries):
        return redirect(f"/orders/receiving/{order_id}/act/?error=placement_required")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    _, mx1_path = _ensure_act_documents(
        order_id,
        act_entry.agency,
        act_entry.payload or {},
        placement_entry.payload if placement_entry else None,
    )
    return FileResponse(open(mx1_path, "rb"), as_attachment=True, filename=mx1_path.name)


def _act_items_from_marking_units_for_print(act_units: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, str, str, str], dict] = {}
    for unit in act_units or []:
        if not isinstance(unit, dict):
            continue
        key = (
            str(unit.get("sku_code") or "").strip(),
            str(unit.get("name") or "").strip(),
            str(unit.get("size") or "").strip(),
            str(unit.get("barcode") or "").strip(),
        )
        row = grouped.setdefault(
            key,
            {
                "sku_code": key[0],
                "name": key[1],
                "size": key[2],
                "barcode": key[3],
                "planned_qty": 0,
                "actual_qty": 0,
                "comment": "",
            },
        )
        row["planned_qty"] = int(row.get("planned_qty") or 0) + 1
        row["actual_qty"] = int(row.get("actual_qty") or 0) + 1
    return list(grouped.values())


def _employee_label_for_user(user) -> str:
    if not user:
        return ""
    employee = Employee.objects.filter(user=user).first()
    if employee and employee.full_name:
        return employee.full_name
    return (getattr(user, "get_full_name", lambda: "")() or getattr(user, "username", "") or "").strip()


def _format_excel_datetime(value) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        current = value
        if timezone.is_naive(current):
            current = timezone.make_aware(current, timezone.get_current_timezone())
        return timezone.localtime(current).strftime("%d.%m.%Y %H:%M")
    formatted = _format_datetime_value(value)
    return "" if formatted == "-" else formatted


_EXCEL_FORBIDDEN_CHARS_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1C\x1E-\x1F]")


def _excel_safe_value(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return value
    text = value.replace("\x1d", "<GS>")
    return _EXCEL_FORBIDDEN_CHARS_RE.sub("", text)


def _receiving_marking_export_units(order_id: str, act_payload: dict) -> list[dict]:
    act_units = [unit for unit in (act_payload.get("act_units") or []) if isinstance(unit, dict)]
    from receiving_cz.models import ReceivingCzUnit

    live_units = (
        ReceivingCzUnit.objects.filter(order_id=order_id)
        .select_related("accepted_by")
        .order_by("accepted_at", "id")
    )
    live_units_by_code = {unit.marking_code: unit for unit in live_units}

    rows = []
    if act_units:
        for unit in act_units:
            marking_code = str(unit.get("marking_code") or "").strip()
            live_unit = live_units_by_code.get(marking_code)
            rows.append(
                {
                    "sku_code": unit.get("sku_code") or getattr(live_unit, "sku_code", "") or "",
                    "name": unit.get("name") or getattr(live_unit, "name", "") or "",
                    "size": unit.get("size") or getattr(live_unit, "size", "") or "",
                    "barcode": unit.get("barcode") or getattr(live_unit, "barcode", "") or "",
                    "marking_code": marking_code,
                    "box_code": unit.get("box_code") or getattr(live_unit, "box_code", "") or "",
                    "pallet_code": unit.get("pallet_code") or getattr(live_unit, "pallet_code", "") or "",
                    "accepted_at": _format_excel_datetime(getattr(live_unit, "accepted_at", "")),
                    "accepted_by": _employee_label_for_user(getattr(live_unit, "accepted_by", None)),
                }
            )
        return rows

    return [
        {
            "sku_code": unit.sku_code,
            "name": unit.name,
            "size": unit.size,
            "barcode": unit.barcode,
            "marking_code": unit.marking_code,
            "box_code": unit.box_code,
            "pallet_code": unit.pallet_code,
            "accepted_at": _format_excel_datetime(unit.accepted_at),
            "accepted_by": _employee_label_for_user(unit.accepted_by),
        }
        for unit in live_units_by_code.values()
    ]


def _build_receiving_marking_workbook(order_id: str, units: list[dict]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реестр ЧЗ"
    headers = [
        "№",
        "ШК",
        "Артикул",
        "Наименование",
        "Размер",
        "ЧЗ",
        "Короб",
        "Палета",
        "Дата/время приемки",
        "Кто принял",
    ]
    sheet.append(headers)
    for index, unit in enumerate(units, start=1):
        sheet.append(
            [
                index,
                unit.get("barcode") or "",
                unit.get("sku_code") or "",
                unit.get("name") or "",
                unit.get("size") or "",
                _excel_safe_value(unit.get("marking_code") or ""),
                unit.get("box_code") or "",
                unit.get("pallet_code") or "",
                unit.get("accepted_at") or "",
                unit.get("accepted_by") or "",
            ]
        )
    widths = [8, 18, 22, 38, 14, 64, 24, 24, 22, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    summary = workbook.create_sheet("Итоги")
    summary.append(["Заявка", order_id])
    summary.append([])
    summary.append(["Артикул", "Наименование", "Размер", "Принято ЧЗ", "Короба"])
    totals: dict[tuple[str, str, str], dict] = {}
    for unit in units:
        key = (
            str(unit.get("sku_code") or "").strip(),
            str(unit.get("name") or "").strip(),
            str(unit.get("size") or "").strip(),
        )
        row = totals.setdefault(key, {"count": 0, "boxes": set()})
        row["count"] += 1
        if unit.get("box_code"):
            row["boxes"].add(str(unit.get("box_code")))
    for key, row in totals.items():
        summary.append([key[0], key[1], key[2], row["count"], ", ".join(sorted(row["boxes"]))])
    for index, width in enumerate([22, 38, 14, 14, 42], start=1):
        summary.column_dimensions[summary.cell(row=3, column=index).column_letter].width = width
    return workbook


def download_receiving_act_marking_xlsx(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    units = _receiving_marking_export_units(order_id, act_entry.payload or {})
    workbook = _build_receiving_marking_workbook(order_id, units)
    output = BytesIO()
    workbook.save(output)
    filename = f"receiving_{_safe_doc_name(order_id)}_chz.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _receiving_plain_export_rows(placement_payload: dict | None, act_payload: dict | None) -> list[dict]:
    rows: list[dict] = []
    pallet_rows, loose_boxes = _placement_payload_container_rows(placement_payload)

    def add_item(pallet_code: str, box_code: str, item: dict):
        rows.append(
            {
                "pallet_code": pallet_code,
                "box_code": box_code,
                "sku_code": item.get("sku_code") or "",
                "name": item.get("name") or "",
                "size": item.get("size") or "",
                "barcode": item.get("barcode") or "",
                "qty": _parse_qty_value(item.get("qty") or item.get("actual_qty")) or 0,
                "marking_code": "",
            }
        )

    for pallet in pallet_rows:
        pallet_code = str(pallet.get("code") or "").strip()
        for box in pallet.get("boxes") or []:
            box_code = str(box.get("code") or "").strip()
            box_items = box.get("items") or []
            if box_items:
                for item in box_items:
                    add_item(pallet_code, box_code, item)
            else:
                rows.append(
                    {
                        "pallet_code": pallet_code,
                        "box_code": box_code,
                        "sku_code": "",
                        "name": "",
                        "size": "",
                        "barcode": "",
                        "qty": _parse_qty_value(box.get("qty_total")) or 0,
                        "marking_code": "",
                    }
                )
        for item in pallet.get("direct_items") or []:
            add_item(pallet_code, "", item)

    for box in loose_boxes:
        box_code = str(box.get("code") or "").strip()
        box_items = box.get("items") or []
        if box_items:
            for item in box_items:
                add_item("", box_code, item)
        else:
            rows.append(
                {
                    "pallet_code": "",
                    "box_code": box_code,
                    "sku_code": "",
                    "name": "",
                    "size": "",
                    "barcode": "",
                    "qty": _parse_qty_value(box.get("qty_total")) or 0,
                    "marking_code": "",
                }
            )

    if rows:
        return rows

    act_items_source = (act_payload or {}).get("act_items") or []
    for item in act_items_source:
        if not isinstance(item, dict):
            continue
        rows.append(
            {
                "pallet_code": "",
                "box_code": "",
                "sku_code": item.get("sku_code") or "",
                "name": item.get("name") or "",
                "size": item.get("size") or "",
                "barcode": item.get("barcode") or "",
                "qty": _parse_qty_value(item.get("actual_qty") or item.get("qty")) or 0,
                "marking_code": "",
            }
        )
    return rows


def _build_receiving_export_workbook(order_id: str, rows: list[dict]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Приемка"
    headers = ["Паллета", "Короб", "SKU", "Наименование", "Размер", "ШК", "Кол-во", "ЧЗ"]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row.get("pallet_code") or "",
                row.get("box_code") or "",
                row.get("sku_code") or "",
                row.get("name") or "",
                row.get("size") or "",
                row.get("barcode") or "",
                row.get("qty") or 0,
                _excel_safe_value(row.get("marking_code") or ""),
            ]
        )
    for index, width in enumerate([24, 24, 22, 42, 16, 22, 12, 64], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    summary = workbook.create_sheet("Итоги")
    summary.append(["Заявка", order_id])
    summary.append(["Строк", len(rows)])
    summary.append(["Всего принято", sum(_parse_qty_value(row.get("qty")) or 0 for row in rows)])
    return workbook


def _receiving_flow_current_export_rows(boxes_data: list, pallets_data: list) -> list[dict]:
    boxes_by_code = {}
    used_box_codes = set()
    rows: list[dict] = []

    if not isinstance(boxes_data, list):
        boxes_data = []
    if not isinstance(pallets_data, list):
        pallets_data = []

    for box in boxes_data:
        if not isinstance(box, dict):
            continue
        box_code = str(box.get("code") or "").strip()
        if box_code:
            boxes_by_code[box_code] = box

    def add_item(pallet_code: str, box_code: str, item: dict):
        if not isinstance(item, dict):
            return
        qty = _parse_qty_value(item.get("qty"))
        if qty is None or qty <= 0:
            return
        rows.append(
            {
                "pallet_code": pallet_code,
                "box_code": box_code,
                "sku_code": item.get("sku_code") or item.get("sku") or "",
                "name": item.get("name") or "",
                "size": item.get("size") or "",
                "barcode": item.get("barcode") or "",
                "goods_type": item.get("goods_type") or "",
                "qty": qty,
            }
        )

    for pallet in pallets_data:
        if not isinstance(pallet, dict):
            continue
        pallet_code = str(pallet.get("code") or "").strip()
        for box_code_raw in pallet.get("boxes") or []:
            box_code = str(box_code_raw or "").strip()
            if not box_code:
                continue
            used_box_codes.add(box_code)
            box = boxes_by_code.get(box_code) or {}
            for item in box.get("items") or []:
                add_item(pallet_code, box_code, item)
        for item in pallet.get("items") or []:
            add_item(pallet_code, "", item)

    for box_code, box in boxes_by_code.items():
        if box_code in used_box_codes:
            continue
        for item in box.get("items") or []:
            add_item("", box_code, item)

    return rows


def _build_receiving_flow_current_workbook(order_id: str, rows: list[dict]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Принято"
    headers = ["№", "Паллета", "Короб", "SKU", "Наименование", "Размер", "ШК", "Тип товара", "Кол-во"]
    sheet.append(headers)
    for index, row in enumerate(rows, start=1):
        sheet.append(
            [
                index,
                row.get("pallet_code") or "-",
                row.get("box_code") or "-",
                row.get("sku_code") or "-",
                row.get("name") or "-",
                row.get("size") or "-",
                row.get("barcode") or "-",
                row.get("goods_type") or "-",
                row.get("qty") or 0,
            ]
        )

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        cell.alignment = Alignment(horizontal="center")

    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width, 10), 45)

    summary = workbook.create_sheet("Итоги")
    summary.append(["Заявка", order_id])
    summary.append(["Строк", len(rows)])
    summary.append(["Всего принято", sum(_parse_qty_value(row.get("qty")) or 0 for row in rows)])
    summary.append(["Паллет", len({row.get("pallet_code") for row in rows if row.get("pallet_code")})])
    summary.append(["Коробов", len({row.get("box_code") for row in rows if row.get("box_code")})])
    for row in summary.iter_rows():
        row[0].font = Font(bold=True)
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 28
    return workbook


def export_receiving_flow_current_xlsx(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Метод не поддерживается."}, status=405)

    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")

    try:
        boxes_data = json.loads(request.POST.get("boxes_json") or "[]")
        pallets_data = json.loads(request.POST.get("pallets_json") or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({"ok": False, "error": "Не удалось прочитать текущее состояние приемки."}, status=400)

    rows = _receiving_flow_current_export_rows(boxes_data, pallets_data)
    workbook = _build_receiving_flow_current_workbook(order_id, rows)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"receiving_{_safe_doc_name(order_id)}_current.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_receiving_result_xlsx(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    act_payload = act_entry.payload or {}
    placement_payload = placement_entry.payload if placement_entry else None
    marking_units = _receiving_marking_export_units(order_id, act_payload)
    if marking_units:
        rows = [
            {
                "pallet_code": unit.get("pallet_code") or "",
                "box_code": unit.get("box_code") or "",
                "sku_code": unit.get("sku_code") or "",
                "name": unit.get("name") or "",
                "size": unit.get("size") or "",
                "barcode": unit.get("barcode") or "",
                "qty": 1,
                "marking_code": unit.get("marking_code") or "",
            }
            for unit in marking_units
        ]
    else:
        rows = _receiving_plain_export_rows(placement_payload, act_payload)

    workbook = _build_receiving_export_workbook(order_id, rows)
    output = BytesIO()
    workbook.save(output)
    filename = f"receiving_{_safe_doc_name(order_id)}_result.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def print_receiving_act(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    if not _placement_closed(entries) and not _has_receiving_act_print_source(entries, act_entry, placement_entry):
        return redirect(f"/orders/receiving/{order_id}/act/?error=placement_required")
    act_payload = act_entry.payload or {}
    placement_payload = placement_entry.payload if placement_entry else None
    agency = act_entry.agency or (entries[-1].agency if entries else None)
    doc_date = _format_doc_date(
        act_payload.get("received_at") or act_payload.get("eta_at") or act_entry.created_at.isoformat()
    )
    arrival_at = _format_datetime_value(act_payload.get("eta_at"))
    vehicle_number = _format_payload_value(act_payload.get("vehicle_number"))
    driver_phone = _format_payload_value(act_payload.get("driver_phone"))
    agency_name = "-"
    agency_inn = "-"
    agency_kpp = "-"
    agency_address = "-"
    if agency:
        agency_name = (agency.agn_name or agency.fio_agn or str(agency)).strip() or agency_name
        agency_inn = (agency.inn or "").strip() or agency_inn
        agency_kpp = (agency.kpp or "").strip() or agency_kpp
        agency_address = (agency.adres or agency.fakt_adres or "").strip() or agency_address

    act_items_raw = act_payload.get("act_items") or []
    act_units_raw = act_payload.get("act_units") or []
    box_count_map = _placement_box_count_map(placement_payload) if placement_payload else None
    boxes_payload = (placement_payload or {}).get("act_boxes") or []
    pallets_payload = (placement_payload or {}).get("act_pallets") or []
    box_codes = set()
    for box in boxes_payload:
        if not isinstance(box, dict):
            continue
        code = (box.get("code") or "").strip()
        if code:
            box_codes.add(code)
    if box_codes:
        total_boxes_unique = len(box_codes)
    else:
        total_boxes_unique = len([box for box in boxes_payload if isinstance(box, dict)])
    total_pallets = len([pallet for pallet in pallets_payload if isinstance(pallet, dict)])
    display_items = []
    total_planned = 0
    total_actual = 0
    total_boxes = 0
    show_marking_column = False
    act_items_source = act_items_raw or _act_items_from_marking_units_for_print(act_units_raw)
    act_items = _act_items_with_barcodes(agency.id if agency else None, act_items_source, box_count_map)
    special_item_rows = _build_receiving_special_item_rows(act_payload, placement_payload)
    act_items = ReceivingWorkflowService.split_receiving_act_rows_by_goods_type(
        act_items,
        special_rows=special_item_rows,
        planned_key="planned",
        actual_key="actual",
        box_key="box_qty",
    )
    for item in act_items:
        planned = item.get("planned", 0) or 0
        actual = item.get("actual", 0) or 0
        total_planned += planned
        total_actual += actual
        box_qty = item.get("box_qty")
        if box_qty is not None:
            total_boxes += box_qty
        mismatch = actual - planned
        mismatch_display = "" if mismatch == 0 else mismatch
        name_parts = [item.get("name") or "-"]
        if item.get("size"):
            name_parts.append(f"р-р {item['size']}")
        display_items.append(
            {
                "barcode": item.get("barcode") or "-",
                "sku_code": item.get("sku_code") or "-",
                "name": ", ".join(name_parts),
                "planned": planned,
                "actual": actual,
                "box_qty": "-" if box_qty is None else box_qty,
                "mismatch": mismatch_display,
                "marking_code": "-",
            }
        )
    total_boxes_label = total_boxes_unique if total_boxes_unique else "-"
    total_pallets_label = total_pallets if total_pallets else "-"
    total_mismatch = total_actual - total_planned

    storekeeper_signed = _act_storekeeper_signed_from_payload(act_payload)
    manager_signed = _act_manager_signed_from_payload(act_payload)
    storekeeper_employee = (
        _signed_employee_from_payload(act_payload, "act_storekeeper_employee_id")
        if storekeeper_signed
        else None
    )
    manager_employee = (
        _signed_employee_from_payload(act_payload, "act_manager_employee_id")
        if manager_signed
        else None
    )
    role = get_request_role(request)
    client_agency = _client_agency_from_request(request)
    client_view = bool(client_agency)
    can_storekeeper_sign = role == "storekeeper" and not storekeeper_signed and not client_view
    can_manager_sign = (
        role in {"manager", "head_manager", "director", "admin"}
        and storekeeper_signed
        and not manager_signed
        and not client_view
    )

    client_agency = _client_agency_from_request(request)
    client_view = bool(client_agency)
    default_return_url = f"/orders/receiving/{order_id}/act/"
    if client_agency:
        default_return_url = f"/client/dashboard/lk/?client={client_agency.id}"
    raw_return_url = request.GET.get("return") or request.META.get("HTTP_REFERER")
    return_url = _safe_return_url(request, raw_return_url, request.path) or default_return_url
    status_entry = _current_status_entry(entries)
    status_payload = status_entry.payload or {} if status_entry else {}
    client_response = (
        status_payload.get("act_client_response")
        or _latest_status_payload_value(entries, "act_client_response")
        or request.GET.get("response")
        or ""
    ).lower()
    if (
        not client_view
        and role in {"manager", "head_manager", "director", "admin"}
        and client_response == "confirmed"
    ):
        Task.objects.filter(
            route=f"/orders/receiving/{order_id}/act/print/",
            assigned_to__role__in=["manager", "head_manager"],
            title__startswith="Клиент ",
        ).delete()
        Task.objects.filter(
            route=f"/orders/receiving/{order_id}/",
            assigned_to__role__in=["manager", "head_manager"],
        ).exclude(status="done").update(
            status="done",
            updated_at=timezone.localtime(),
        )
    client_response_at = _format_datetime_value(
        status_payload.get("act_client_response_at")
        or _latest_status_payload_value(entries, "act_client_response_at")
    )
    sent_at_raw = status_payload.get("act_sent_at") or act_entry.created_at.isoformat()
    client_deadline = ""
    try:
        sent_at = datetime.fromisoformat(str(sent_at_raw))
        if timezone.is_naive(sent_at):
            sent_at = timezone.make_aware(sent_at, timezone.get_current_timezone())
        sent_at = timezone.localtime(sent_at)
        client_deadline = (sent_at + timedelta(hours=24)).strftime("%d.%m.%Y %H:%M")
    except (TypeError, ValueError):
        client_deadline = ""
    marking_export_units = _receiving_marking_export_units(order_id, act_payload)
    marking_export_url = ""
    if marking_export_units:
        marking_export_url = f"/orders/receiving/{order_id}/act/chz.xlsx"
        if client_agency:
            marking_export_url = f"{marking_export_url}?client={client_agency.id}"

    ctx = {
        "order_id": order_id,
        "act_label": (act_payload.get("act_label") or "Акт приемки").strip(),
        "doc_date": doc_date,
        "arrival_at": arrival_at,
        "vehicle_number": vehicle_number,
        "driver_phone": driver_phone,
        "agency_name": agency_name,
        "agency_inn": agency_inn,
        "agency_kpp": agency_kpp,
        "agency_address": agency_address,
        "agency_label": _agency_label(agency),
        "items": display_items,
        "total_planned": total_planned,
        "total_actual": total_actual,
        "total_boxes": total_boxes_label,
        "total_pallets": total_pallets_label,
        "total_mismatch": "" if total_mismatch == 0 else total_mismatch,
        "show_marking_column": show_marking_column,
        "print_date": timezone.localtime().strftime("%d.%m.%Y %H:%M"),
        "executor_label": _EXECUTOR_LABEL,
        "manager_facsimile_url": _facsimile_url(manager_employee),
        "storekeeper_facsimile_url": _facsimile_url(storekeeper_employee),
        "storekeeper_signed": storekeeper_signed,
        "manager_signed": manager_signed,
        "can_storekeeper_sign": can_storekeeper_sign,
        "can_manager_sign": can_manager_sign,
        "sign_status": request.GET.get("signed"),
        "sign_error": request.GET.get("error"),
        "return_url": return_url,
        "client_view": client_view,
        "client_param": client_agency.id if client_agency else "",
        "client_response": client_response,
        "client_response_at": client_response_at,
        "client_deadline": client_deadline,
        "marking_export_url": marking_export_url,
    }
    return render(request, "orders/receiving_act_print.html", ctx)


def sign_receiving_act_storekeeper(request, order_id: str):
    if request.method != "POST":
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role != "storekeeper":
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    if not _placement_closed(entries):
        return redirect(f"/orders/receiving/{order_id}/act/print/?error=placement_required")
    act_payload = dict(act_entry.payload or {})
    if _act_storekeeper_signed_from_payload(act_payload):
        return redirect(f"/orders/receiving/{order_id}/act/print/")
    employee = Employee.objects.filter(user=request.user, role="storekeeper", is_active=True).first()
    if not employee:
        employee = _first_active_employee_by_roles("storekeeper")
    if not employee or not getattr(employee, "facsimile", None):
        return redirect(f"/orders/receiving/{order_id}/act/print/?error=storekeeper_facsimile")
    act_payload["act_storekeeper_signed"] = True
    act_payload["act_storekeeper_signed_at"] = timezone.localtime().isoformat()
    act_payload["act_storekeeper_employee_id"] = employee.id
    if request.user.is_authenticated:
        act_payload["act_storekeeper_user_id"] = request.user.id
    if not act_payload.get("status"):
        act_payload["status"] = "warehouse"
    act_payload["status_label"] = "Принято складом, акт приемки отправлен менеджеру"
    log_order_action(
        "status",
        order_id=order_id,
        order_type="receiving",
        user=request.user if request.user.is_authenticated else None,
        agency=act_entry.agency or (entries[-1].agency if entries else None),
        description="Акт приемки подписан кладовщиком и отправлен менеджеру",
        payload=act_payload,
    )
    _create_manager_sign_task(
        order_id,
        act_entry.agency or (entries[-1].agency if entries else None),
        request,
        observer=employee,
    )
    return redirect(f"/orders/receiving/{order_id}/act/print/?signed=storekeeper")


def sign_receiving_act_manager(request, order_id: str):
    if request.method != "POST":
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in {"manager", "head_manager", "director", "admin"}:
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    if not _placement_closed(entries):
        return redirect(f"/orders/receiving/{order_id}/act/print/?error=placement_required")
    act_payload = dict(act_entry.payload or {})
    if not _act_storekeeper_signed_from_payload(act_payload):
        return redirect(f"/orders/receiving/{order_id}/act/print/?error=storekeeper_required")
    if _act_manager_signed_from_payload(act_payload):
        return redirect(f"/orders/receiving/{order_id}/act/print/?signed=manager")
    employee = Employee.objects.filter(user=request.user, is_active=True).first()
    if not employee:
        employee = _first_active_employee_by_roles("manager", "head_manager")
    if not employee or not getattr(employee, "facsimile", None):
        return redirect(f"/orders/receiving/{order_id}/act/print/?error=manager_facsimile")
    act_payload["act_manager_signed"] = True
    act_payload["act_manager_signed_at"] = timezone.localtime().isoformat()
    act_payload["act_manager_employee_id"] = employee.id
    if request.user.is_authenticated:
        act_payload["act_manager_user_id"] = request.user.id
    log_order_action(
        "status",
        order_id=order_id,
        order_type="receiving",
        user=request.user if request.user.is_authenticated else None,
        agency=act_entry.agency or (entries[-1].agency if entries else None),
        description="Акт приемки подписан менеджером",
        payload=act_payload,
    )
    entries = _load_order_entries(order_id)
    result = ReceivingWorkflowService.send_receiving_act_to_client(
        order_id=str(order_id or ""),
        entries=entries,
        user=request.user if request.user.is_authenticated else None,
    )
    if not result.applied:
        return redirect(f"/orders/receiving/{order_id}/act/print/?signed=manager&error=send")
    Task.objects.filter(
        route=f"/orders/receiving/{order_id}/act/print/",
        assigned_to__role__in=["manager", "head_manager"],
    ).delete()
    return redirect(f"/orders/receiving/{order_id}/act/print/?signed=manager")


def _client_print_url(
    order_id: str,
    client_agency,
    response: str | None = None,
    return_url: str | None = None,
) -> str:
    params = {}
    if client_agency:
        params["client"] = client_agency.id
    if response:
        params["response"] = response
    if return_url:
        params["return"] = return_url
    suffix = f"?{urlencode(params)}" if params else ""
    return f"/orders/receiving/{order_id}/act/print/{suffix}"


def _safe_return_url(request, raw_url: str | None, current_path: str | None = None) -> str:
    if not raw_url:
        return ""
    if current_path:
        current_abs = request.build_absolute_uri(current_path)
        if raw_url.startswith(current_abs) or raw_url.startswith(current_path):
            return ""
    if url_has_allowed_host_and_scheme(
        raw_url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return raw_url
    return ""


def _client_act_access_allowed(request, entries):
    client_agency = _client_agency_from_request(request)
    if not client_agency:
        return False
    if not any(entry.agency_id == client_agency.id for entry in entries if entry.agency_id):
        return False
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    return bool(act_entry)


def confirm_receiving_act_client(request, order_id: str):
    if request.method != "POST":
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _client_act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    client_agency = _client_agency_from_request(request)
    raw_return_url = request.POST.get("return") or request.GET.get("return")
    return_url = _safe_return_url(request, raw_return_url, request.path)
    status_entry = _current_status_entry(entries)
    payload = dict(status_entry.payload or {}) if status_entry else {}
    existing_response = (
        payload.get("act_client_response")
        or _latest_status_payload_value(entries, "act_client_response")
        or ""
    ).lower()
    if existing_response in {"confirmed", "dispute"}:
        return redirect(_client_print_url(order_id, client_agency, existing_response, return_url))
    payload["act_client_response"] = "confirmed"
    payload["act_client_response_at"] = timezone.localtime().isoformat()
    log_order_action(
        "status",
        order_id=order_id,
        order_type="receiving",
        user=request.user if request.user.is_authenticated else None,
        agency=entries[-1].agency if entries else None,
        description="Акт приемки подтвержден клиентом",
        payload=payload,
    )
    _create_manager_client_response_task(
        order_id,
        entries[-1].agency if entries else None,
        request,
        "confirmed",
    )
    return redirect(_client_print_url(order_id, client_agency, "confirmed", return_url))


def dispute_receiving_act_client(request, order_id: str):
    if request.method != "POST":
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _client_act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    client_agency = _client_agency_from_request(request)
    raw_return_url = request.POST.get("return") or request.GET.get("return")
    return_url = _safe_return_url(request, raw_return_url, request.path)
    status_entry = _current_status_entry(entries)
    payload = dict(status_entry.payload or {}) if status_entry else {}
    existing_response = (
        payload.get("act_client_response")
        or _latest_status_payload_value(entries, "act_client_response")
        or ""
    ).lower()
    if existing_response in {"confirmed", "dispute"}:
        return redirect(_client_print_url(order_id, client_agency, existing_response, return_url))
    payload["act_client_response"] = "dispute"
    payload["act_client_response_at"] = timezone.localtime().isoformat()
    log_order_action(
        "status",
        order_id=order_id,
        order_type="receiving",
        user=request.user if request.user.is_authenticated else None,
        agency=entries[-1].agency if entries else None,
        description="Клиент заявил разногласия по акту приемки",
        payload=payload,
    )
    _create_manager_client_response_task(
        order_id,
        entries[-1].agency if entries else None,
        request,
        "dispute",
    )
    return redirect(_client_print_url(order_id, client_agency, "dispute", return_url))


def print_receiving_act_mx1(request, order_id: str):
    entries = _load_order_entries(order_id)
    if not entries:
        raise Http404("Заявка не найдена")
    if not _act_access_allowed(request, entries):
        return HttpResponseForbidden("Доступ запрещен")
    act_entry = _find_act_entry(entries, "receiving", "акт приемки")
    act_entry = _receiving_act_entry_or_placement(entries, act_entry)
    if not act_entry:
        raise Http404("Акт приемки не найден")
    placement_entry = _find_act_entry(entries, "placement", "акт размещения")
    if not _placement_closed(entries) and not _has_receiving_act_print_source(entries, act_entry, placement_entry):
        return redirect(f"/orders/receiving/{order_id}/act/?error=placement_required")
    act_payload = act_entry.payload or {}
    placement_payload = placement_entry.payload if placement_entry else None
    agency = act_entry.agency or (entries[-1].agency if entries else None)
    doc_date = _format_doc_date(
        act_payload.get("received_at") or act_payload.get("eta_at") or act_entry.created_at.isoformat()
    )
    agency_name = "-"
    agency_address = "-"
    agency_phone = "-"
    if agency:
        agency_name = (agency.agn_name or agency.fio_agn or str(agency)).strip() or agency_name
        agency_address = (agency.adres or agency.fakt_adres or "").strip() or agency_address
        agency_phone = (agency.phone or "").strip() or agency_phone
    keeper_label = _EXECUTOR_LABEL
    depositor_label = _format_party_label(agency_name, agency_address, agency_phone)

    act_items_raw = act_payload.get("act_items") or []
    box_count_map = _placement_box_count_map(placement_payload) if placement_payload else {}
    pallet_count_map = _placement_pallet_count_map(placement_payload) if placement_payload else {}
    act_items = _act_items_with_barcodes(agency.id if agency else None, act_items_raw, box_count_map)
    display_items = []
    total_actual = 0
    for item in act_items:
        actual = item.get("actual", 0) or 0
        total_actual += actual
        key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
        pallet_qty = pallet_count_map.get(key)
        name_parts = [item.get("name") or "-"]
        if item.get("sku_code"):
            name_parts.append(f"арт.: {item['sku_code']}")
        if item.get("size"):
            name_parts.append(f"р-р {item['size']}")
        display_items.append(
            {
                "sku_code": item.get("sku_code") or "-",
                "barcode": item.get("barcode") or "-",
                "name": ", ".join(name_parts),
                "size": item.get("size") or "-",
                "actual": actual,
                "box_qty": "-" if item.get("box_qty") is None else item.get("box_qty"),
                "pallet_qty": "-" if pallet_qty is None else pallet_qty,
            }
        )

    boxes_payload = (placement_payload or {}).get("act_boxes") or []
    pallets_payload = (placement_payload or {}).get("act_pallets") or []
    box_codes = set()
    for box in boxes_payload:
        if not isinstance(box, dict):
            continue
        code = (box.get("code") or "").strip()
        if code:
            box_codes.add(code)
    if box_codes:
        total_boxes = len(box_codes)
    else:
        total_boxes = len([box for box in boxes_payload if isinstance(box, dict)])
    total_pallets = len([pallet for pallet in pallets_payload if isinstance(pallet, dict)])

    manager_employee = _first_active_employee_by_roles("manager", "head_manager")
    storekeeper_employee = _first_active_employee_by_roles("storekeeper")

    items_per_page = 18
    pages = []
    for start in range(0, max(len(display_items), 1), items_per_page):
        pages.append(
            {
                "start": start,
                "items": display_items[start:start + items_per_page],
            }
        )

    ctx = {
        "order_id": order_id,
        "doc_date": doc_date,
        "agency_label": _agency_label(agency),
        "executor_label": _EXECUTOR_LABEL,
        "okud_code": _OKUD_MX1,
        "keeper_label": keeper_label,
        "depositor_label": depositor_label,
        "pages": pages,
        "total_actual": total_actual,
        "total_boxes": total_boxes if total_boxes else "-",
        "total_pallets": total_pallets if total_pallets else "-",
        "manager_facsimile_url": _facsimile_url(manager_employee),
        "storekeeper_facsimile_url": _facsimile_url(storekeeper_employee),
    }
    return render(request, "orders/mx1_print.html", ctx)


class OrdersHomeView(RoleRequiredMixin, TemplateView):
    template_name = 'orders/index.html'
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            order_id = kwargs.get("order_id")
            if order_id:
                has_access = OrderAuditEntry.objects.filter(
                    order_id=order_id,
                    order_type=self.order_type,
                    agency=client_agency,
                ).exists()
                if not has_access:
                    return HttpResponseForbidden("Доступ запрещен")
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    @staticmethod
    def _next_order_number(order_type: str = "receiving") -> str:
        normalized_type = normalize_order_type_for_number(order_type)
        order_types = [normalized_type]
        if normalized_type == "processing":
            order_types.append("packing")
        order_ids = (
            OrderAuditEntry.objects.filter(order_type__in=order_types)
            .values_list("order_id", flat=True)
            .distinct()
        )
        return next_public_order_number(order_type, order_ids)

    def get(self, request, *args, **kwargs):
        active_tab = kwargs.get("tab", "journal")
        client_agency = getattr(request, "_client_agency", None) or _client_agency_from_request(request)
        edit_order_id = str(request.GET.get("edit") or "").strip()
        if active_tab == "receiving" and client_agency and edit_order_id:
            from client_cabinet.request_ownership import (
                PORTAL_REQUEST_OWNER_MESSAGE,
                portal_user_can_edit_request,
            )

            if not portal_user_can_edit_request(
                user=request.user,
                agency=client_agency,
                order_type="receiving",
                order_id=edit_order_id,
                request=request,
            ):
                return HttpResponseForbidden(PORTAL_REQUEST_OWNER_MESSAGE)
        if (
            active_tab == "receiving"
            and client_agency
            and not str(request.GET.get("edit") or "").strip()
            and not kwargs.get("error")
            and request.GET.get("ok") != "1"
        ):
            from client_cabinet.client_drafts import client_draft_limit_reached, find_client_draft

            draft = find_client_draft(
                agency=client_agency,
                order_type="receiving",
                user=request.user,
            )
            if (
                draft
                and draft.get("continue_url")
                and client_draft_limit_reached(
                    agency=client_agency,
                    order_type="receiving",
                    user=request.user,
                )
            ):
                return redirect(draft["continue_url"])
        status = (request.GET.get("status") or "").lower()
        ok = request.GET.get("ok") == "1"
        context_kwargs = dict(kwargs)
        submitted = context_kwargs.pop("submitted", None) or (ok and status != "draft")
        draft_saved = ok and status == "draft"
        error = context_kwargs.pop("error", None)
        ctx = self.get_context_data(
            submitted=submitted,
            draft_saved=draft_saved,
            error=error,
            **context_kwargs,
        )
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        active_tab = kwargs.get("tab", "journal")
        if active_tab == "packing":
            return self._submit_packing(request)
        if active_tab != "receiving":
            return redirect("/orders/")

        submit_action = str(request.POST.get("submit_action") or "").strip()
        autosave = str(request.POST.get("draft_autosave") or "").strip() == "1"
        if autosave:
            submit_action = "draft"

        def fail(message: str):
            if autosave:
                return JsonResponse({"ok": False, "error": message}, status=400)
            return self.get(request, tab=active_tab, error=message)

        client_agency = getattr(request, "_client_agency", None) or _client_agency_from_request(request)
        if client_agency:
            agency = client_agency
        else:
            agency_id = request.POST.get("agency_id")
            agency = Agency.objects.filter(pk=agency_id).first()
        if not agency:
            return fail("Выберите клиента.")

        edit_order_id = str(request.POST.get("edit_order_id") or "").strip()
        if client_agency and edit_order_id:
            from client_cabinet.request_ownership import (
                PORTAL_REQUEST_OWNER_MESSAGE,
                portal_user_can_edit_request,
            )

            if not portal_user_can_edit_request(
                user=request.user,
                agency=client_agency,
                order_type="receiving",
                order_id=edit_order_id,
                request=request,
            ):
                if autosave:
                    return JsonResponse(
                        {"ok": False, "error": PORTAL_REQUEST_OWNER_MESSAGE},
                        status=403,
                    )
                return HttpResponseForbidden(PORTAL_REQUEST_OWNER_MESSAGE)

        template_items = []
        template_files = []
        template_files.extend(request.FILES.getlist("template_files"))
        sku_template_file = request.FILES.get("template_sku_file")
        receiving_template_file = request.FILES.get("template_receiving_file")
        if sku_template_file:
            template_files.append(sku_template_file)
        if receiving_template_file:
            template_files.append(receiving_template_file)
        template_names = [f.name for f in template_files if getattr(f, "name", "")]
        if template_files:
            try:
                with transaction.atomic():
                    template_items, template_names = _process_template_uploads(template_files, agency)
            except ValueError as exc:
                return fail(str(exc))

        eta_raw = (request.POST.get("eta_at") or "").strip()
        min_eta = _min_receiving_eta(business_localtime())
        eta_value = min_eta
        if eta_raw:
            try:
                eta_value = parse_business_datetime(eta_raw)
            except ValueError:
                if not autosave:
                    return fail("Некорректная дата/время прибытия.")
                eta_value = min_eta
        elif not autosave:
            return fail("Укажите плановую дату и время прибытия.")
        if not autosave and eta_value < min_eta:
            return fail("Плановая дата/время не может быть в прошлом.")

        expected_boxes_raw = (request.POST.get("expected_boxes") or "").strip()
        place_type = (request.POST.get("place_type") or "").strip()
        vehicle_number = (request.POST.get("vehicle_number") or "").strip()
        driver_phone = _normalize_driver_phone(request.POST.get("driver_phone") or "")
        receiving_chz_service = ReceivingWorkflowService.normalize_receiving_chz_service(
            request.POST.get("receiving_chz_service")
        )
        receiving_route = ReceivingWorkflowService.normalize_receiving_route(
            request.POST.get("receiving_route")
        )
        try:
            expected_boxes_value = int(expected_boxes_raw)
        except (TypeError, ValueError):
            expected_boxes_value = 0
        if not autosave:
            if expected_boxes_value <= 0:
                return fail("Укажите количество мест.")
            if not place_type:
                return fail("Выберите тип мест.")
            if not vehicle_number:
                return fail("Укажите номер авто.")
            if not driver_phone or not _is_valid_driver_phone(driver_phone):
                return fail("Введите корректный телефон водителя.")
        if submit_action != "draft" and not receiving_chz_service:
            return fail("Укажите, требуется ли сканирование Честного знака.")
        if submit_action != "draft" and not receiving_route:
            return fail("Укажите маршрут после приемки: FBS или FBO.")

        sku_codes = request.POST.getlist("sku_code[]")
        sku_ids = request.POST.getlist("sku_id[]")
        names = request.POST.getlist("item_name[]")
        brands = request.POST.getlist("brand[]")
        colors = request.POST.getlist("color[]")
        sizes = request.POST.getlist("size[]")
        barcodes = request.POST.getlist("barcode[]")
        qtys = request.POST.getlist("qty[]")
        position_comments = request.POST.getlist("position_comment[]")
        items = []
        row_count = max(
            len(sku_codes),
            len(qtys),
            len(names),
            len(brands),
            len(colors),
            len(position_comments),
            len(sku_ids),
            len(sizes),
            len(barcodes),
        )
        for idx in range(row_count):
            sku_code = sku_codes[idx] if idx < len(sku_codes) else ""
            qty = qtys[idx] if idx < len(qtys) else ""
            if not sku_code and not qty:
                continue
            items.append(
                {
                    "sku_id": sku_ids[idx] if idx < len(sku_ids) else "",
                    "sku_code": sku_code,
                    "name": names[idx] if idx < len(names) else "",
                    "brand": brands[idx] if idx < len(brands) else "",
                    "color": colors[idx] if idx < len(colors) else "",
                    "size": sizes[idx] if idx < len(sizes) else "",
                    "barcode": barcodes[idx] if idx < len(barcodes) else "",
                    "qty": qty,
                    "comment": position_comments[idx] if idx < len(position_comments) else "",
                }
            )
        if template_items:
            items = _merge_items(items + template_items)
        edit_order_id = str(request.POST.get("edit_order_id") or "").strip()
        try:
            ReceivingWorkflowService._attach_receiving_nomenclature_metadata(
                agency=agency,
                items=items,
                goods_type="",
                order_id=str(edit_order_id or ""),
                order_type="receiving",
            )
        except ReceivingNomenclatureError as exc:
            # Incomplete rows while typing should not block draft autosave.
            if not autosave:
                return fail(str(exc))

        role = get_request_role(request)
        previous_entries = []
        can_client_edit = False
        if edit_order_id:
            previous_entries = list(
                OrderAuditEntry.objects.filter(
                    order_id=edit_order_id,
                    order_type="receiving",
                ).order_by("created_at")
            )
            can_client_edit = bool(client_agency and _can_client_edit_receiving(previous_entries, client_agency))
            if _is_manager_staff_role(role):
                if not _can_manager_edit_receiving(previous_entries):
                    return fail("Заявка уже передана на склад — исправление недоступно")
            elif not can_client_edit:
                return fail("Недостаточно прав для исправления заявки")
        old_payload = {}
        if edit_order_id:
            old_payload = _latest_payload_from_entries(previous_entries)
        if not receiving_chz_service and autosave:
            receiving_chz_service = ReceivingWorkflowService.normalize_receiving_chz_service(
                old_payload.get("receiving_chz_service")
            )
        if not receiving_route and autosave:
            receiving_route = ReceivingWorkflowService.normalize_receiving_route(
                old_payload.get("receiving_route"),
                default="fbo",
            )
        status_value = "draft" if submit_action == "draft" else "sent_unconfirmed"
        status_label = "Черновик" if status_value == "draft" else "Ждет подтверждения"
        if edit_order_id and old_payload and submit_action != "send" and not autosave:
            preserved_status = old_payload.get("status") or old_payload.get("submit_action")
            preserved_label = old_payload.get("status_label")
            if preserved_status:
                status_value = preserved_status
            if preserved_label:
                status_label = preserved_label
        if autosave:
            status_value = "draft"
            status_label = "Черновик"
            submit_action = "draft"
        if not str(edit_order_id or "").strip() and submit_action == "draft":
            from client_cabinet.client_drafts import resolve_client_draft_order_id

            reused_draft_id = resolve_client_draft_order_id(
                agency=agency,
                order_type="receiving",
                allow_new=True,
                user=request.user,
            )
            if reused_draft_id:
                edit_order_id = reused_draft_id
                previous_entries = list(
                    OrderAuditEntry.objects.filter(
                        order_id=edit_order_id,
                        order_type="receiving",
                    ).order_by("created_at")
                )
                old_payload = _latest_payload_from_entries(previous_entries)
        promote_draft_order_id = ""
        if (
            edit_order_id
            and submit_action == "send"
            and _is_receiving_draft_entries(previous_entries)
            and order_number_sequence("receiving", edit_order_id) is None
        ):
            promote_draft_order_id = str(edit_order_id).strip()
        order_id = (
            self._next_order_number(order_type="receiving")
            if promote_draft_order_id
            else str(edit_order_id or "").strip() or self._next_order_number(order_type="receiving")
        )
        try:
            _save_receiving_attachments(order_id, agency, request)
        except OSError:
            return fail("Не удалось сохранить вложения заявки.")
        if promote_draft_order_id:
            ReceivingOrderAttachment.objects.filter(order_id=promote_draft_order_id, agency=agency).update(
                order_id=order_id
            )
        document_names = _receiving_attachment_names(order_id, agency)
        if not document_names:
            legacy_docs = old_payload.get("documents") if isinstance(old_payload, dict) else []
            if isinstance(legacy_docs, list):
                document_names = [str(name).strip() for name in legacy_docs if str(name).strip()]
        payload = {
            "eta_at": eta_value.isoformat(),
            "expected_boxes": expected_boxes_raw,
            "place_type": place_type,
            "vehicle_number": vehicle_number,
            "driver_phone": driver_phone,
            "receiving_chz_service": receiving_chz_service,
            "receiving_route": receiving_route,
            "comment": request.POST.get("comment"),
            "submit_action": submit_action,
            "status": status_value,
            "status_label": status_label,
            "items": items,
            "documents": document_names,
            "template_files": template_names,
        }
        if edit_order_id and not promote_draft_order_id:
            try:
                submission_result = ReceivingWorkflowService.submit_receiving_order(
                    order_id=str(edit_order_id or ""),
                    agency=agency,
                    payload=payload,
                    submit_action=submit_action or "",
                    user=request.user if request.user.is_authenticated else None,
                    submitted_at=timezone.localtime(),
                    existing_order_id=str(edit_order_id or ""),
                    old_payload=old_payload,
                    dispatch_review=bool(submit_action == "send" and self.request.GET.get("client")),
                    is_autosave=autosave,
                )
            except ValueError as exc:
                return fail(str(exc))
            if submit_action == "draft":
                from client_cabinet.client_drafts import supersede_extra_client_drafts

                supersede_extra_client_drafts(
                    agency=agency,
                    order_type="receiving",
                    keep_order_id=str(edit_order_id),
                    user=request.user,
                )
            if autosave:
                return JsonResponse(
                    {
                        "ok": True,
                        "order_id": str(edit_order_id),
                        "draft_order_id": str(edit_order_id),
                        "status": submission_result.status_value,
                        "status_label": submission_result.status_label,
                    }
                )
            if submit_action == "send" and self.request.GET.get("client"):
                return redirect(f"/client/dashboard/lk/?client={agency.id}")
            return redirect(f"/orders/receiving/{edit_order_id}/")

        try:
            ReceivingWorkflowService.submit_receiving_order(
                order_id=str(order_id or ""),
                agency=agency,
                payload=payload,
                submit_action=submit_action or "",
                user=request.user if request.user.is_authenticated else None,
                submitted_at=timezone.localtime(),
                dispatch_review=bool(submit_action != "draft" and self.request.GET.get("client")),
            )
        except ValueError as exc:
            return fail(str(exc))
        if promote_draft_order_id:
            OrderAuditEntry.objects.filter(
                order_id=promote_draft_order_id,
                order_type="receiving",
                agency=agency,
            ).delete()
        if submit_action == "draft":
            from client_cabinet.client_drafts import supersede_extra_client_drafts

            supersede_extra_client_drafts(
                agency=agency,
                order_type="receiving",
                keep_order_id=str(order_id),
                user=request.user,
            )
        if autosave:
            return JsonResponse(
                {
                    "ok": True,
                    "order_id": str(order_id),
                    "draft_order_id": str(order_id),
                    "status": status_value,
                    "status_label": status_label,
                }
            )
        if submit_action != "draft" and self.request.GET.get("client"):
            return redirect(f"/client/dashboard/lk/?client={agency.id}")
        if submit_action == "draft" and self.request.GET.get("client"):
            return redirect(
                f"/orders/receiving/?client={agency.id}&edit={order_id}&ok=1&status={status_value}"
            )
        return redirect(
            f"/orders/receiving/?client={agency.id}&ok=1&status={status_value}&order={order_id}"
        )

    def _submit_packing(self, request):
        client_agency = getattr(request, "_client_agency", None) or _client_agency_from_request(request)
        if client_agency:
            agency = client_agency
        else:
            agency_id = request.POST.get("agency_id")
            agency = Agency.objects.filter(pk=agency_id).first()
        if not agency:
            return self.get(request, error="Выберите клиента.")

        email = (request.POST.get("email") or "").strip()
        fio = (request.POST.get("fio") or "").strip()
        plan_date = (request.POST.get("plan_date") or "").strip()
        total_qty_raw = (request.POST.get("total_qty") or "").strip()
        if not email:
            return self.get(request, error="Укажите email.")
        if not fio:
            return self.get(request, error="Укажите ФИО.")
        if not plan_date:
            return self.get(request, error="Укажите плановую дату выполнения.")
        try:
            total_qty_value = int(total_qty_raw)
        except (TypeError, ValueError):
            total_qty_value = 0
        if total_qty_value <= 0:
            return self.get(request, error="Укажите общее количество к обработке.")

        marketplaces = request.POST.getlist("mp[]")
        mp_other = (request.POST.get("mp_other") or "").strip()
        if not marketplaces and not mp_other:
            return self.get(request, error="Выберите маркетплейс или укажите другое.")

        payload = {
            "email": email,
            "fio": fio,
            "org": request.POST.get("org"),
            "plan_date": plan_date,
            "marketplaces": marketplaces,
            "mp_other": request.POST.get("mp_other"),
            "subject": request.POST.get("subject"),
            "total_qty": total_qty_raw,
            "box_mode": request.POST.get("box_mode"),
            "box_mode_other": request.POST.get("box_mode_other"),
            "tasks": request.POST.getlist("tasks[]"),
            "tasks_other": request.POST.get("tasks_other"),
            "marking": request.POST.get("marking"),
            "marking_other": request.POST.get("marking_other"),
            "ship_as": request.POST.get("ship_as"),
            "ship_other": request.POST.get("ship_other"),
            "has_distribution": request.POST.get("has_distribution"),
            "comments": request.POST.get("comments"),
            "files_report": [f.name for f in request.FILES.getlist("files_report")],
            "files_distribution": [f.name for f in request.FILES.getlist("files_distribution")],
            "files_cz": [f.name for f in request.FILES.getlist("files_cz")],
        }
        status_value = "sent_unconfirmed"
        status_label = "Ждет подтверждения"
        payload["status"] = status_value
        payload["status_label"] = status_label
        payload["submit_action"] = "submitted"
        order_id = self._next_order_number(order_type="packing")
        log_order_action(
            "create",
            order_id=order_id,
            order_type="packing",
            user=request.user if request.user.is_authenticated else None,
            agency=agency,
            description=f"Заявка на упаковку №{order_id}",
            payload=payload,
        )
        return redirect(
            f"/orders/packing/?client={agency.id}&ok=1&status={status_value}&order={order_id}"
        )


    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        active_tab = kwargs.get("tab", "journal")
        ctx["active_tab"] = active_tab
        ctx["submitted"] = kwargs.get("submitted", False)
        ctx["draft_saved"] = kwargs.get("draft_saved", False)
        ctx["error"] = kwargs.get("error")
        ctx["status_label"] = ctx.get("status_label", "Подготовка заявки")
        ctx["order_number"] = ctx.get("order_number", "")
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))

        client_id = self.request.GET.get("client")
        agency_id = self.request.GET.get("agency")
        agency_key = client_id or agency_id
        client_agency = getattr(self.request, "_client_agency", None) or _client_agency_from_request(self.request)
        agency = client_agency or (Agency.objects.filter(pk=agency_key).first() if agency_key else None)
        client_view = bool(client_agency)
        ctx["agency"] = agency
        ctx["client_view"] = client_view

        if active_tab == "journal":
            raw_entries = (
                OrderAuditEntry.objects.select_related("user", "agency")
                .exclude(order_type="stock_move")
                .order_by("-created_at")
            )
            if agency:
                raw_entries = raw_entries.filter(agency=agency)
            role = get_request_role(self.request)
            order_type_filter = (self.request.GET.get("order_type") or "").strip().lower()
            if role == "processing_head":
                order_type_filter = "processing"
            if order_type_filter in {"receiving", "packing", "processing", "shipping", "other"}:
                raw_entries = raw_entries.filter(order_type=order_type_filter)
            raw_entries = list(raw_entries[:_JOURNAL_RAW_SCAN_LIMIT])
            manager_label = _manager_label()
            storekeeper_label = _storekeeper_label()
            receiving_totals = {}
            for entry in raw_entries:
                if entry.order_type != "receiving":
                    continue
                payload = entry.payload or {}
                if payload.get("act") != "receiving":
                    continue
                if entry.order_id in receiving_totals:
                    continue
                total = 0
                for item in payload.get("act_items") or []:
                    qty = _parse_qty_value(item.get("actual_qty"))
                    if qty is None:
                        qty = _parse_qty_value(item.get("qty")) or 0
                    total += qty
                receiving_totals[entry.order_id] = total
            latest_by_order = {}
            status_by_order = {}
            for entry in raw_entries:
                if entry.order_id not in latest_by_order:
                    latest_by_order[entry.order_id] = entry
                if entry.order_id not in status_by_order and _is_status_entry(entry):
                    status_by_order[entry.order_id] = entry
            latest_entries = [
                status_by_order.get(order_id, latest_entry)
                for order_id, latest_entry in latest_by_order.items()
            ]
            latest_entries.sort(key=lambda item: item.created_at, reverse=True)
            if not client_view:
                latest_entries = [entry for entry in latest_entries if not _is_draft_entry(entry)]
            latest_entries = latest_entries[:_JOURNAL_ENTRY_LIMIT]
            status_audience = "client" if client_view else ("storekeeper" if role in _RECEIVING_WAREHOUSE_ROLES else "default")
            shipping_numbers = {
                str(entry.order_id or "").strip()
                for entry in latest_entries
                if entry.order_type == "shipping" and str(entry.order_id or "").strip()
            }
            shipping_order_ids_by_number = {}
            if shipping_numbers:
                from shipping.models import ShippingOrder

                shipping_order_ids_by_number = dict(
                    ShippingOrder.objects.filter(number__in=shipping_numbers).values_list("number", "id")
                )
            entries_payload = []
            for entry in latest_entries:
                if client_view and entry.order_type == "receiving" and _is_draft_entry(entry) and entry.agency_id:
                    detail_url = f"/orders/receiving/?client={entry.agency_id}&edit={entry.order_id}"
                elif entry.order_type == "shipping":
                    shipping_order_id = shipping_order_ids_by_number.get(str(entry.order_id or "").strip())
                    detail_url = f"/shipping/{shipping_order_id}/" if shipping_order_id else "/shipping/"
                    if client_view and entry.agency:
                        detail_url += f"?client={entry.agency.id}"
                else:
                    detail_url = f"/orders/{entry.order_type}/{entry.order_id}/"
                    if client_view and entry.agency:
                        detail_url += f"?client={entry.agency.id}"
                entries_payload.append(
                        {
                            "created_at": entry.created_at,
                            "order_type": entry.order_type,
                            "type_label": _order_type_label(entry.order_type),
                            "order_id": _display_order_number(entry.order_type, entry.order_id),
                            "action_label": _journal_action_label(entry, manager_label, storekeeper_label),
                            "status_label": _journal_status_label(entry),
                            "next_step_label": _journal_next_step_label(entry),
                            "agency": entry.agency,
                            "actual_qty": receiving_totals.get(entry.order_id),
                            "client_label": _shorten_ip_name(
                                (entry.agency.agn_name or str(entry.agency)) if entry.agency else "-"
                            ),
                            "detail_url": detail_url,
                    }
                )
            ctx["entries"] = entries_payload
        if active_tab == "receiving":
            ctx["agencies"] = Agency.objects.order_by("agn_name")
            ctx["current_time"] = business_localtime()
            ctx["current_time_input"] = ctx["current_time"].strftime("%Y-%m-%dT%H:%M")
            ctx["default_receiving_eta_input"] = _default_receiving_eta(
                ctx["current_time"]
            ).strftime("%Y-%m-%dT%H:%M")
            ctx["min_past_hours"] = 0
            status = self.request.GET.get("status")
            status_label = "Подготовка заявки"
            if status == "draft":
                status_label = "Черновик"
            elif status == "sent_unconfirmed":
                status_label = "Ждет подтверждения"
            elif status in {"warehouse", "on_warehouse"}:
                status_label = "В ожидании поставки товара"
            ctx["status_label"] = status_label
            ctx["order_number"] = self.request.GET.get("order", "")

            edit_order_id = self.request.GET.get("edit")
            if edit_order_id:
                role = get_request_role(self.request)
                entries = list(
                    OrderAuditEntry.objects.filter(
                        order_id=edit_order_id,
                        order_type="receiving",
                    )
                    .select_related("agency")
                    .order_by("created_at")
                )
                is_receiving_draft = _is_receiving_draft_entries(entries)
                can_client_edit = client_view and _can_client_edit_receiving(entries, client_agency)
                can_manager_edit = (not client_view) and _is_manager_staff_role(role) and _can_manager_edit_receiving(entries)
                can_submit_draft = bool(
                    (can_client_edit and is_receiving_draft)
                    or (role == "manager" and is_receiving_draft)
                )
                if _act_storekeeper_signed(entries):
                    ctx["error"] = "Заявка подписана кладовщиком, редактирование запрещено"
                elif client_view and not can_client_edit:
                    ctx["error"] = "Недостаточно прав для исправления заявки"
                elif (not client_view) and not can_manager_edit and not can_client_edit:
                    ctx["error"] = "Недостаточно прав для исправления заявки"
                elif entries and (can_client_edit or can_manager_edit):
                    edit_payload = _latest_payload_from_entries(entries)
                    if not agency:
                        agency = entries[-1].agency
                        ctx["agency"] = agency
                    ctx["edit_mode"] = True
                    ctx["edit_order_id"] = edit_order_id
                    ctx["edit_payload"] = edit_payload
                    ctx["edit_is_draft"] = can_submit_draft
                    ctx["status_label"] = "Черновик" if can_submit_draft else "Исправление заявки"
                    ctx["order_number"] = edit_order_id

            sku_options = []
            if agency:
                skus = (
                    SKU.objects.filter(agency=agency, deleted=False)
                    .prefetch_related("barcodes")
                    .order_by("sku_code")
                )
                for sku in skus:
                    barcodes = [barcode.value for barcode in sku.barcodes.all()]
                    size_map = {}
                    for barcode in sku.barcodes.all():
                        size_value = (barcode.size or "").strip()
                        if not size_value:
                            continue
                        size_map.setdefault(size_value, []).append(barcode.value)
                    sku_options.append(
                        {
                            "id": sku.id,
                            "code": sku.sku_code,
                            "name": sku.name,
                            "brand": sku.brand or "",
                            "color": sku.color or "",
                            "barcodes_joined": "|".join(barcodes),
                            "sizes_json": json.dumps(size_map, ensure_ascii=True),
                        }
                    )
            ctx["sku_options"] = sku_options
            if agency:
                from fbs.services.client_profiles import is_client_fbs_enabled

                ctx["client_fbs_enabled"] = is_client_fbs_enabled(agency)
            else:
                ctx["client_fbs_enabled"] = False
        if active_tab == "packing":
            status = self.request.GET.get("status")
            status_label = "Подготовка заявки"
            if status == "draft":
                status_label = "Черновик"
            elif status in {"sent_unconfirmed", "send", "submitted"}:
                status_label = "Ждет подтверждения"
            ctx["status_label"] = status_label
            ctx["order_number"] = self.request.GET.get("order", "")

        if active_tab == "other":
            from client_cabinet.other_requests import list_other_entries
            from client_cabinet.web_ui import _order_bucket, _order_status_label

            entries = list_other_entries(agency=agency if agency else None, limit=80)
            role = get_request_role(self.request)
            if not client_view and role not in {None, "admin", "director", "head_manager", "manager", "storekeeper"}:
                entries = []
            other_rows = []
            for entry in entries:
                payload = entry.payload or {}
                status_label = _order_status_label(entry)
                other_rows.append(
                    {
                        "order_id": entry.order_id,
                        "display_number": format_order_number("other", entry.order_id),
                        "created_at": entry.created_at,
                        "agency_name": (entry.agency.agn_name if entry.agency else "-"),
                        "category_label": payload.get("category_label") or payload.get("title") or "Прочая",
                        "description": (payload.get("description") or "")[:240],
                        "status_label": status_label,
                        "bucket": _order_bucket(entry),
                        "detail_url": f"/orders/other/{entry.order_id}/"
                        + (f"?client={entry.agency_id}" if client_view and entry.agency_id else ""),
                    }
                )
            ctx["other_rows"] = other_rows
            ctx["status_label"] = "Прочие заявки"

        return ctx


class OrdersPackingView(OrdersHomeView):
    template_name = "orders/packing.html"




class OrdersDetailView(RoleRequiredMixin, TemplateView):
    template_name = "orders/detail.html"
    order_type = "receiving"
    allowed_roles = ("manager", "storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        if request.user.is_authenticated and not order_detail_access_allowed(
            request, order_id=kwargs.get("order_id"), order_type=self.order_type,
        ):
            return HttpResponseForbidden("Доступ запрещен")
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    @staticmethod
    def _payload_from_entries(entries):
        return build_order_runtime_payload("receiving", entries=entries)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        action = request.POST.get("action")
        if self.order_type == "receiving" and action == "close_receiving_task_invalid":
            role = get_request_role(request)
            try:
                result = ReceivingWorkflowService.close_invalid_receiving_task(
                    order_id=str(order_id),
                    role=role or "",
                    user=request.user if request.user.is_authenticated else None,
                    reason_code=request.POST.get("invalid_reason") or "",
                    note=request.POST.get("invalid_note") or "",
                )
                if result.status == "review_required":
                    messages.warning(
                        request,
                        "Задача кладовщика закрыта. Заявка передана начальнику склада без изменения остатков, зон, резервов и коробов.",
                    )
                elif result.status == "review_already_pending":
                    messages.info(request, "Заявка уже передана начальнику склада.")
                else:
                    messages.success(
                        request,
                        "Заявка отмечена удалённой и перемещена в «Готовые».",
                    )
            except (PermissionError, ValueError) as exc:
                messages.error(request, str(exc))
            return redirect(f"/orders/receiving/{order_id}/")
        if self.order_type == "receiving" and action in {
            "request_warehouse_cancel",
            "approve_warehouse_cancel",
            "reject_warehouse_cancel",
        }:
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            role = get_request_role(request)
            reason = (request.POST.get("cancel_reason") or "").strip()
            try:
                if action == "request_warehouse_cancel":
                    if not _is_manager_staff_role(role):
                        return HttpResponseForbidden("Доступ запрещен")
                    ReceivingWorkflowService.request_warehouse_cancel_confirmation(
                        order_id=str(order_id),
                        entries=entries,
                        user=request.user if request.user.is_authenticated else None,
                        reason=reason,
                    )
                    messages.success(request, "Запрос на отмену отправлен складу.")
                elif action == "approve_warehouse_cancel":
                    pending = ReceivingWorkflowService.warehouse_cancel_request_payload(entries)
                    ReceivingWorkflowService.cancel_receiving_order(
                        order_id=str(order_id),
                        entries=entries,
                        role=role or "",
                        user=request.user if request.user.is_authenticated else None,
                        reason=str(pending.get("cancel_reason") or reason).strip(),
                        warehouse_approved=True,
                    )
                    messages.success(request, "Отмена подтверждена складом.")
                else:
                    ReceivingWorkflowService.reject_warehouse_cancel_confirmation(
                        order_id=str(order_id),
                        entries=entries,
                        role=role or "",
                        user=request.user if request.user.is_authenticated else None,
                    )
                    messages.success(request, "Склад отклонил отмену. Заявка остаётся в работе.")
            except (PermissionError, ValueError) as exc:
                messages.error(request, str(exc))
            client_param = request.GET.get("client")
            suffix = f"?client={client_param}" if client_param else ""
            return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}")
        if action == "cancel_order":
            reason = (request.POST.get("cancel_reason") or "").strip()
            client_param = request.GET.get("client")
            suffix = f"?client={client_param}" if client_param else ""
            role = get_request_role(request)
            if not _is_manager_staff_role(role):
                return HttpResponseForbidden("Доступ запрещен")
            if not reason:
                return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}#cancel-reason-required")
            latest = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("-created_at")
                .first()
            )
            if not latest:
                return redirect("/orders/")
            if self.order_type == "receiving":
                entries = list(
                    OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                    .select_related("agency")
                    .order_by("created_at")
                )
                try:
                    ReceivingWorkflowService.cancel_receiving_order(
                        order_id=str(order_id),
                        entries=entries,
                        role=role or "",
                        user=request.user if request.user.is_authenticated else None,
                        reason=reason,
                    )
                except PermissionError:
                    messages.error(
                        request,
                        "Склад уже начал работу. Отправьте запрос на отмену для подтверждения складом.",
                    )
                except ValueError as exc:
                    messages.error(request, str(exc))
            else:
                payload = _order_cancel_payload(latest.payload or {}, reason=reason, actor="manager")
                log_order_action(
                    "status",
                    order_id=order_id,
                    order_type=self.order_type,
                    user=request.user if request.user.is_authenticated else None,
                    agency=latest.agency,
                    description=f"Заявка отменена менеджером. Причина: {reason}",
                    payload=payload,
                )
                _close_order_tasks_for_cancel(self.order_type, str(order_id))
            return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}")
        if action == "send_to_warehouse":
            if self.order_type != "receiving" or not ReceivingWorkflowService.can_confirm_receiving_to_warehouse(request.user):
                return HttpResponseForbidden("Отправить приемку на склад может только менеджер")
            latest = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("-created_at")
                .first()
            )
            if not latest:
                return redirect("/orders/")
            try:
                ReceivingWorkflowService.confirm_receiving_to_warehouse(
                    order_id=str(order_id or ""),
                    agency=latest.agency if latest else None,
                    status_payload=latest.payload or {},
                    user=request.user if request.user.is_authenticated else None,
                    submitted_at=timezone.localtime(),
                )
            except PermissionError:
                return HttpResponseForbidden("Отправить приемку на склад может только менеджер")
            client_param = request.GET.get("client")
            suffix = f"?client={client_param}" if client_param else ""
            return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}")
        if action == "send_act_to_client":
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            role = get_request_role(request)
            if role not in {"manager", "head_manager", "director", "admin"}:
                return redirect("/orders/")
            ReceivingWorkflowService.send_receiving_act_to_client(
                order_id=str(order_id or ""),
                entries=entries,
                user=request.user if request.user.is_authenticated else None,
            )
            client_param = request.GET.get("client")
            suffix = f"?client={client_param}" if client_param else ""
            return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}")
        if action == "take_over_receiving_work":
            if self.order_type != "receiving":
                return redirect(f"/orders/{self.order_type}/{order_id}/")
            role = get_request_role(request) or ""
            result = ReceivingWorkflowService.take_over_receiving_work(
                order_id=str(order_id or ""),
                role=role,
                user=request.user if request.user.is_authenticated else None,
            )
            if result.status in {"taken_over", "already_owned"}:
                if result.status == "taken_over":
                    previous_name = str(result.meta.get("previous_storekeeper_name") or "другого кладовщика")
                    messages.success(
                        request,
                        f"Заявка передана вам от {previous_name}. Продолжаем существующую приемку.",
                    )
                payload = result.payload or {}
                receiving_mode = str(payload.get("receiving_mode") or "standard").strip().lower()
                goods_type = str(payload.get("goods_type") or "").strip().lower()
                from receiving_cz.models import ReceivingCzUnit

                has_cz_units = ReceivingCzUnit.objects.filter(order_id=str(order_id or "")).exists()
                if (receiving_mode == "cz" or has_cz_units) and goods_type != "no":
                    return redirect(f"/orders/receiving/{order_id}/cz-flow/")
                return redirect(f"/orders/receiving/{order_id}/flow/")
            error_messages = {
                "forbidden": "Забрать заявку может только кладовщик.",
                "storekeeper_not_found": "Активный профиль кладовщика не найден.",
                "warehouse_review_pending": "Заявка ожидает проверки начальником склада.",
                "flow_closed": "Приемка уже закрыта, продолжить ее нельзя.",
                "not_in_progress": "У заявки больше нет текущего владельца. Обновите страницу.",
                "missing": "Заявка на приемку не найдена.",
            }
            messages.error(request, error_messages.get(result.status, "Не удалось забрать заявку в работу."))
            return redirect(f"/orders/receiving/{order_id}/")
        if action == "create_receiving_act":
            goods_type = (request.POST.get("goods_type") or "").strip().lower()
            receiving_mode = (request.POST.get("receiving_mode") or "").strip().lower()
            shipping_return_order_number = (
                request.POST.get("shipping_return_order_number") or ""
            ).strip()
            role = get_request_role(request)
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("created_at")
            )
            existing_flow_state = _flow_state_from_entries(entries)
            existing_flow_has_data = ReceivingWorkflowService.receiving_flow_has_items(
                existing_flow_state
            )
            if existing_flow_has_data and not _flow_closed_from_entries(entries):
                existing_status_entry = _current_status_entry(entries)
                existing_status_payload = (
                    existing_status_entry.payload
                    if existing_status_entry and isinstance(existing_status_entry.payload, dict)
                    else {}
                )
                existing_goods_type = str(
                    existing_status_payload.get("goods_type") or ""
                ).strip().lower()
                existing_receiving_mode = ReceivingWorkflowService._normalize_receiving_mode_for_goods_type(
                    existing_status_payload.get("receiving_mode") or "standard",
                    existing_goods_type,
                )
                if existing_receiving_mode == "cz" and existing_goods_type != "no":
                    return redirect(f"/orders/{self.order_type}/{order_id}/cz-flow/")
                return redirect(f"/orders/{self.order_type}/{order_id}/flow/")
            if role == "storekeeper":
                if goods_type not in ReceivingWorkflowService.RECEIVING_GOODS_TYPE_LABELS:
                    messages.error(request, "Выберите тип товара перед началом приемки.")
                    return redirect(f"/orders/{self.order_type}/{order_id}/")
                request_payload = ReceivingWorkflowService._latest_payload(entries)
                mode_constraint_error = ReceivingWorkflowService.receiving_mode_constraint_error(
                    payload=request_payload,
                    goods_type=goods_type,
                    receiving_mode=receiving_mode,
                )
                if mode_constraint_error:
                    messages.error(request, mode_constraint_error)
                    return redirect(f"/orders/{self.order_type}/{order_id}/")
                if goods_type == SHIPPING_RETURN_GOODS_TYPE:
                    agency = entries[-1].agency if entries else None
                    source_order = resolve_shipping_return_order(
                        agency=agency,
                        number=shipping_return_order_number,
                    )
                    if source_order is None:
                        messages.error(
                            request,
                            "Выберите отгруженную заявку этого клиента для возврата.",
                        )
                        return redirect(f"/orders/{self.order_type}/{order_id}/")
                    if not shipping_return_source_items(
                        source_order,
                        exclude_receiving_order_id=str(order_id or ""),
                    ):
                        messages.error(
                            request,
                            "По этой отгрузке нет товара, доступного для возврата.",
                        )
                        return redirect(f"/orders/{self.order_type}/{order_id}/")
                start_result = ReceivingWorkflowService.start_receiving_work(
                    order_id=str(order_id or ""),
                    entries=entries,
                    role=role,
                    user=request.user if request.user.is_authenticated else None,
                )
                if start_result.status == "assigned_to_other_storekeeper":
                    owner_name = str(start_result.payload.get("storekeeper_name") or "другого кладовщика")
                    messages.error(request, f"Заявка уже в работе у {owner_name}.")
                    return redirect(f"/orders/{self.order_type}/{order_id}/")
                entries = list(
                    OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                    .select_related("agency")
                    .order_by("created_at")
                )
            result = ReceivingWorkflowService.configure_receiving_act(
                order_id=str(order_id or ""),
                entries=entries,
                role=role or "",
                goods_type=goods_type,
                receiving_mode=receiving_mode,
                shipping_return_order_number=shipping_return_order_number,
                user=request.user if request.user.is_authenticated else None,
            )
            if not result.applied:
                messages.error(
                    request,
                    result.description or "Не удалось настроить приемку.",
                )
                return redirect(f"/orders/{self.order_type}/{order_id}/")
            next_mode = str((result.payload or {}).get("receiving_mode") or receiving_mode).strip().lower()
            next_goods_type = str((result.payload or {}).get("goods_type") or goods_type).strip().lower()
            if next_mode == "cz" and next_goods_type != "no":
                return redirect(f"/orders/{self.order_type}/{order_id}/cz-flow/")
            return redirect(f"/orders/{self.order_type}/{order_id}/flow/")
        comment = (request.POST.get("comment") or "").strip()
        if comment:
            latest = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("-created_at")
                .first()
            )
            log_order_action(
                "comment",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency if latest else None,
                description=comment,
                payload={"comment": comment},
            )
        client_param = request.GET.get("client")
        suffix = f"?client={client_param}" if client_param else ""
        return redirect(f"/orders/{self.order_type}/{order_id}/{suffix}")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        client_agency = getattr(self.request, "_client_agency", None) or _client_agency_from_request(self.request)
        client_view = bool(client_agency)
        entries = (
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )
        entries_list = list(entries)
        latest = entries_list[-1] if entries_list else None
        status_entry = _current_status_entry(entries_list)
        payload = self._payload_from_entries(entries_list)
        ctx["order_id"] = order_id
        ctx["order_type"] = self.order_type
        ctx["order_title"] = _order_title_label(
            self.order_type,
            order_id,
            payload,
            entries=entries_list,
            created_at=entries_list[0].created_at if entries_list else None,
        )
        ctx["agency"] = latest.agency if latest else client_agency
        current_agency = ctx["agency"]
        ctx["client_view"] = client_view
        client_param = self.request.GET.get("client") or self.request.GET.get("agency")
        client_id = None
        if client_param:
            try:
                client_id = int(client_param)
            except (TypeError, ValueError):
                client_id = None
        if not client_id and client_view:
            if client_agency and getattr(client_agency, "id", None):
                client_id = client_agency.id
            elif latest and latest.agency_id:
                client_id = latest.agency_id
        if client_id:
            ctx["client_cabinet_url"] = f"/client/dashboard/lk/?client={client_id}"
            ctx["client_orders_url"] = f"/orders/?client={client_id}"
        status_text = _status_label_from_entry(status_entry) if status_entry else "-"
        next_step_label = ""
        next_step_tone = ""
        processing_result = None
        receiving_result = None
        if self.order_type == "processing":
            processing_result = WarehouseGoodsStateResolver.resolve_for_processing_order(
                order_id=str(order_id or ""),
                agency=latest.agency if latest else client_agency,
                payload=payload,
            )
            status_text = processing_result.label_for("default")
            next_step_label = processing_result.next_step_for("default")
        elif self.order_type == "receiving":
            status_audience = "client" if client_view else ("storekeeper" if get_request_role(self.request) == "storekeeper" else "default")
            receiving_state_payload = dict(payload)
            if status_entry:
                receiving_state_payload.update(dict(status_entry.payload or {}))
            receiving_result = WarehouseGoodsStateResolver.resolve_for_receiving_order(
                order_id=str(order_id or ""),
                agency=latest.agency if latest else client_agency,
                payload=receiving_state_payload,
            )
            status_text = receiving_result.label_for(status_audience)
            next_step_label = receiving_result.next_step_for(status_audience)
            next_step_tone = _next_step_tone_from_label(next_step_label)
        ctx["status_label"] = status_text
        ctx["cancel_reason"] = (payload.get("cancel_reason") or "").strip()
        ctx["is_cancelled"] = str(payload.get("status") or payload.get("submit_action") or "").strip().lower() in {"cancelled", "canceled"}
        ctx["is_deleted_as_erroneous"] = bool(payload.get("deleted_as_erroneous"))
        ctx["next_step_label"] = next_step_label
        ctx["next_step_tone"] = next_step_tone
        ctx["responsible"] = _current_responsible_label(status_entry)
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        flow_state = _flow_state_from_entries(entries_list)
        flow_closed = _flow_closed_from_entries(entries_list)
        flow_has_data = ReceivingWorkflowService.receiving_flow_has_items(flow_state)
        status_payload = status_entry.payload or {} if status_entry else {}
        flow_has_draft = bool(
            isinstance(flow_state, dict)
            and (
                flow_state.get("activeBox")
                or flow_state.get("activePallet")
                or flow_state.get("boxes")
                or flow_state.get("pallets")
            )
        )
        receiving_claimed = bool(
            status_payload.get("receiving_tsd_claimed")
            or status_payload.get("storekeeper_employee_id")
            or str(status_payload.get("storekeeper_name") or "").strip()
        )
        receiving_started = bool(
            flow_closed or flow_has_data or flow_has_draft or receiving_claimed
        )
        ctx["receiving_started"] = receiving_started
        goods_type = (status_payload.get("goods_type") or "").strip().lower()
        goods_type_labels = {
            "op": "Оптовый",
            "gv": "Готовый",
            "br": "Брак",
            "vz": "Возврат",
            "rh": "Расходный",
            "no": "Не обработанный",
            SHIPPING_RETURN_GOODS_TYPE: SHIPPING_RETURN_GOODS_TYPE_LABEL,
        }
        receiving_mode = ReceivingWorkflowService._normalize_receiving_mode_for_goods_type(
            status_payload.get("receiving_mode") or "standard",
            goods_type,
        )
        receiving_chz_service = ReceivingWorkflowService.normalize_receiving_chz_service(
            status_payload.get("receiving_chz_service")
            or payload.get("receiving_chz_service")
        )
        ctx["receiving_chz_service"] = receiving_chz_service
        ctx["receiving_chz_service_label"] = ReceivingWorkflowService.receiving_chz_service_label(
            receiving_chz_service
        )
        receiving_route = ReceivingWorkflowService.receiving_route_from_entries(entries_list)
        ctx["receiving_route"] = receiving_route
        ctx["receiving_route_label"] = ReceivingWorkflowService.receiving_route_label(
            receiving_route
        )
        ctx["receiving_mode_required_by_client"] = ReceivingWorkflowService.receiving_mode_required_by_client(
            {"receiving_chz_service": receiving_chz_service}
        )
        ctx["goods_type"] = goods_type
        ctx["goods_type_label"] = goods_type_labels.get(goods_type, "")
        ctx["shipping_return_order_number"] = str(
            status_payload.get("shipping_return_order_number") or ""
        ).strip()
        ctx["shipping_return_order_display"] = (
            format_order_number("shipping", ctx["shipping_return_order_number"])
            if ctx["shipping_return_order_number"]
            else ""
        )
        items = payload.get("items") or []
        marked_items_map = _receiving_marked_items_map(latest.agency_id if latest else None, items)
        if receiving_mode == "cz" and not marked_items_map and items:
            receiving_mode = "standard"
        if self.order_type == "receiving" and not receiving_started:
            receiving_mode = ""
        ctx["has_marked_receiving_items"] = bool(marked_items_map)
        ctx["receiving_mode"] = receiving_mode
        can_send_to_warehouse = False
        role = get_request_role(self.request)
        signed_by_storekeeper = _act_storekeeper_signed(entries_list)
        signed_by_manager = _act_manager_signed(entries_list)
        can_edit_order = False
        if client_view:
            can_edit_order = _can_client_edit_receiving(entries_list, client_agency)
        elif _is_manager_staff_role(role):
            can_edit_order = _can_manager_edit_receiving(entries_list)
        if "товар принят" in (status_text or "").lower():
            can_edit_order = False
        if signed_by_storekeeper or signed_by_manager:
            can_edit_order = False
        ctx["can_edit_order"] = can_edit_order
        if can_edit_order:
            if client_view and client_agency:
                ctx["receiving_edit_url"] = (
                    f"/orders/receiving/?client={client_agency.id}&edit={order_id}"
                )
            elif current_agency and getattr(current_agency, "id", None):
                ctx["receiving_edit_url"] = f"/orders/receiving/?edit={order_id}&agency={current_agency.id}"
            else:
                ctx["receiving_edit_url"] = f"/orders/receiving/?edit={order_id}"
        else:
            ctx["receiving_edit_url"] = ""
        ctx["can_cancel_order"] = bool((not client_view) and _is_manager_staff_role(role) and not ctx["is_cancelled"] and not _is_done_status(status_entry))
        warehouse_cancel_request = (
            ReceivingWorkflowService.warehouse_cancel_request_payload(entries_list)
            if self.order_type == "receiving"
            else {}
        )
        requires_warehouse_cancel = bool(
            self.order_type == "receiving"
            and ReceivingWorkflowService.requires_warehouse_cancel_confirmation(entries_list)
        )
        receiving_cancel_is_final = bool(
            self.order_type == "receiving"
            and ReceivingWorkflowService.receiving_cancel_is_final(entries_list)
        )
        ctx["warehouse_cancel_request_pending"] = bool(warehouse_cancel_request)
        ctx["warehouse_cancel_reason"] = str(
            warehouse_cancel_request.get("cancel_reason") or ""
        ).strip()
        ctx["can_request_warehouse_cancel"] = bool(
            self.order_type == "receiving"
            and not client_view
            and _is_manager_staff_role(role)
            and not ctx["is_cancelled"]
            and not _is_done_status(status_entry)
            and requires_warehouse_cancel
            and not warehouse_cancel_request
        )
        ctx["can_review_warehouse_cancel"] = bool(
            self.order_type == "receiving"
            and not client_view
            and role in _RECEIVING_WAREHOUSE_ROLES
            and warehouse_cancel_request
        )
        warehouse_invalid_review = (
            ReceivingWorkflowService.warehouse_invalid_review_payload(entries_list)
            if self.order_type == "receiving"
            else {}
        )
        current_storekeeper = (
            ReceivingWorkflowService._resolve_storekeeper(
                self.request.user if self.request.user.is_authenticated else None
            )
            if role == "storekeeper" and not client_view
            else None
        )
        receiving_work_owner = (
            ReceivingWorkflowService._receiving_work_owner(entries_list)
            if self.order_type == "receiving"
            else {}
        )
        try:
            receiving_work_owner_id = int(receiving_work_owner.get("storekeeper_employee_id") or 0)
        except (TypeError, ValueError):
            receiving_work_owner_id = 0
        receiving_work_owner_name = str(receiving_work_owner.get("storekeeper_name") or "").strip()
        receiving_work_owned_by_current_user = bool(
            current_storekeeper
            and receiving_work_owner_id
            and current_storekeeper.id == receiving_work_owner_id
        )
        ctx["receiving_work_owner_name"] = receiving_work_owner_name
        ctx["receiving_work_owned_by_current_user"] = receiving_work_owned_by_current_user
        receiving_access_allowed = False
        receiving_access = None
        if self.order_type == "receiving":
            receiving_access = ReceivingWorkflowService.receiving_work_access(
                entries=entries_list,
                user=self.request.user if self.request.user.is_authenticated else None,
            )
            receiving_access_allowed = receiving_access.status == "allowed"
        ctx["warehouse_invalid_review_pending"] = bool(warehouse_invalid_review)
        ctx["warehouse_invalid_reason"] = str(
            warehouse_invalid_review.get("warehouse_invalid_reason_label") or ""
        ).strip()
        ctx["warehouse_invalid_note"] = str(
            warehouse_invalid_review.get("warehouse_invalid_note") or ""
        ).strip()
        ctx["can_close_receiving_task_invalid"] = bool(
            self.order_type == "receiving"
            and role == "storekeeper"
            and not client_view
            and not ctx["is_cancelled"]
            and not _is_done_status(status_entry)
            and not warehouse_invalid_review
            and bool(current_storekeeper)
        )
        if requires_warehouse_cancel or receiving_cancel_is_final or warehouse_cancel_request:
            ctx["can_cancel_order"] = False
        if warehouse_cancel_request:
            ctx["status_label"] = "Отмена ожидает подтверждения склада"
        can_create_receiving_act = False
        if status_entry and not client_view and role == "manager":
            payload = status_entry.payload or {}
            status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
            status_label = (payload.get("status_label") or "").lower()
            can_send_to_warehouse = status_value in {"sent_unconfirmed", "send", "submitted"}
            if "склад" in status_label or "ожидании поставки" in status_label or status_value in {"warehouse", "on_warehouse"}:
                can_send_to_warehouse = False
            if status_value == "draft":
                can_send_to_warehouse = False
        if status_entry and not client_view and role in _RECEIVING_WAREHOUSE_ROLES:
            has_receiving_act = any(
                (entry.payload or {}).get("act") == "receiving" for entry in entries_list
            )
            can_create_receiving_act = WarehouseActionPolicy.can_create_receiving_act(
                receiving_result,
                has_receiving_act=has_receiving_act,
                role=role,
                allow_legacy_warehouse=bool(
                    receiving_result and receiving_result.code == WarehouseStateCode.UNKNOWN
                    and _warehouse_status_from_entry(status_entry)
                ),
            ).allowed
        ctx["can_send_to_warehouse"] = can_send_to_warehouse
        ctx["can_create_receiving_act"] = can_create_receiving_act
        if self.order_type == "receiving":
            try:
                from receiving_distribution.services.progress import build_distribution_panel

                ctx["receiving_distribution"] = build_distribution_panel(str(order_id or ""))
            except Exception:
                ctx["receiving_distribution"] = None
        if self.order_type == "receiving" and flow_has_data and not flow_closed:
            status_text = "В процессе приемки"
            ctx["status_label"] = status_text
        elif self.order_type == "receiving" and not receiving_started:
            status_text = "Не начата"
            ctx["status_label"] = status_text
        draft_flow_payload = (
            _receiving_flow_draft_payload(flow_state)
            if self.order_type == "receiving" and flow_has_data and not flow_closed
            else {}
        )
        ctx["can_continue_flow"] = bool(
            role in _RECEIVING_WAREHOUSE_ROLES
            and receiving_started
            and not flow_closed
            and (role != "storekeeper" or receiving_access_allowed)
        )
        has_receiving_act = any((entry.payload or {}).get("act") == "receiving" for entry in entries_list)
        can_open_flow = WarehouseActionPolicy.can_open_receiving_flow(
            receiving_result,
            role=role,
            client_view=client_view,
            flow_closed=flow_closed,
            can_create_receiving_act=can_create_receiving_act,
            has_receiving_act=has_receiving_act,
            flow_has_data=flow_has_data,
        ).allowed
        if (
            not can_open_flow
            and role in _RECEIVING_WAREHOUSE_ROLES
            and not client_view
            and ReceivingWorkflowService.can_resume_stored_receiving_draft(entries_list)
        ):
            can_open_flow = True
        ctx["can_open_flow"] = can_open_flow
        ctx["can_take_over_receiving_work"] = bool(
            self.order_type == "receiving"
            and role == "storekeeper"
            and not client_view
            and current_storekeeper
            and receiving_access
            and receiving_access.status == "assigned_to_other_storekeeper"
            and receiving_work_owner_name
            and not flow_closed
            and not ctx["is_cancelled"]
            and not _is_done_status(status_entry)
            and not warehouse_invalid_review
            and can_open_flow
        )
        can_reopen_receiving_flow = False
        if (
            self.order_type == "receiving"
            and role == "storekeeper"
            and not client_view
            and flow_closed
            and flow_has_data
            and not signed_by_storekeeper
            and not signed_by_manager
        ):
            reopen_policy = WarehouseActionPolicy.can_start_receiving_flow(
                receiving_result,
                flow_closed=False,
                allow_legacy_warehouse=False,
            )
            receiving_access = ReceivingWorkflowService.receiving_work_access(
                entries=entries_list,
                user=self.request.user if self.request.user.is_authenticated else None,
            )
            can_reopen_receiving_flow = bool(
                reopen_policy.allowed and receiving_access.status == "allowed"
            )
        has_live_cz_units = False
        if can_reopen_receiving_flow:
            from receiving_cz.models import ReceivingCzUnit

            has_live_cz_units = ReceivingCzUnit.objects.filter(
                order_id=str(order_id or "")
            ).exists()
        ctx["can_reopen_receiving_flow"] = can_reopen_receiving_flow
        ctx["can_view_receiving_flow"] = can_reopen_receiving_flow
        ctx["view_receiving_flow_url"] = (
            f"/orders/receiving/{order_id}/cz-flow/"
            if has_live_cz_units or receiving_mode == "cz"
            else f"/orders/receiving/{order_id}/flow/"
        )
        ctx["view_receiving_flow_label"] = "Внести изменения в приемку"
        ctx["flow_closed"] = flow_closed
        ctx["receiving_flow_url"] = f"/orders/receiving/{order_id}/flow/"
        ctx["storage_placement_url"] = f"/orders/receiving/{order_id}/flow/?open_warehouse_move=1"
        ctx["reachtruck_request_url"] = ""
        ctx["can_manage_storage_placement"] = False
        ctx["storage_placement_action_label"] = ""
        if self.order_type == "receiving" and receiving_result and not client_view:
            receiving_move = WarehouseMovementResolver.active_for_receiving_order(
                agency=latest.agency if latest else client_agency,
                order_id=str(order_id or ""),
            )
            has_reachtruck_request = bool(
                receiving_move.task_ids
                or receiving_move.active_task_count
                or receiving_move.done_task_count
                or receiving_move.blocked_task_count
            )
            if has_reachtruck_request:
                ctx["reachtruck_request_url"] = (
                    f"/reachtruck/?mobile_category=movement&mobile_request=receiving:{order_id}"
                )
        client_label = "-"
        if latest and latest.agency:
            name = latest.agency.short_name or latest.agency.agn_name or latest.agency.fio_agn or str(latest.agency)
            client_label = _shorten_ip_name(name)
        client_prefix = ""
        if latest and latest.agency:
            client_prefix = (latest.agency.pref or "").strip()
        ctx["client_label"] = client_label
        ctx["meta"] = {
            "eta_at": _format_datetime_value(payload.get("eta_at")),
            "expected_boxes": payload.get("expected_boxes"),
            "place_type": _place_type_label(payload.get("place_type")),
            "vehicle_number": payload.get("vehicle_number"),
            "driver_phone": payload.get("driver_phone"),
            "receiving_chz_service": ReceivingWorkflowService.receiving_chz_service_label(
                receiving_chz_service
            ),
            "receiving_route": ReceivingWorkflowService.receiving_route_label(
                receiving_route
            ),
            "comment": payload.get("comment"),
            "documents": payload.get("documents") or [],
        }
        act_entry_for_items = _find_act_entry(entries_list, "receiving", "акт приемки")
        placement_entry_for_items = _find_act_entry(entries_list, "placement", "акт размещения")
        act_items = ((act_entry_for_items.payload or {}).get("act_items") or []) if act_entry_for_items else []
        if not act_items and placement_entry_for_items:
            act_items = _placement_payload_item_rows(placement_entry_for_items.payload or {})
        if not act_items and draft_flow_payload:
            act_items = _placement_payload_item_rows(draft_flow_payload)
        display_items = []
        if items:
            actual_map = {}
            act_item_by_key = {}
            for act_item in act_items:
                key = _item_key(act_item.get("sku_code"), act_item.get("name"), act_item.get("size"))
                qty = _parse_qty_value(act_item.get("actual_qty"))
                if qty is None:
                    qty = _parse_qty_value(act_item.get("qty")) or 0
                actual_map[key] = (actual_map.get(key) or 0) + qty
                if key not in act_item_by_key:
                    act_item_by_key[key] = act_item
            planned_keys = set()
            for item in items:
                key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
                planned_keys.add(key)
                display_items.append(
                    {
                        "sku_code": item.get("sku_code"),
                        "name": item.get("name"),
                        "size": item.get("size"),
                        "qty": item.get("qty"),
                        "comment": item.get("comment"),
                        "actual_qty": actual_map.get(key),
                    }
                )
            for key, act_item in act_item_by_key.items():
                if key in planned_keys:
                    continue
                planned_qty = act_item.get("planned_qty")
                if planned_qty in (None, ""):
                    planned_qty = act_item.get("qty")
                actual_qty = actual_map.get(key)
                if actual_qty is None:
                    actual_qty = _parse_qty_value(act_item.get("actual_qty"))
                    if actual_qty is None:
                        actual_qty = _parse_qty_value(act_item.get("qty"))
                display_items.append(
                    {
                        "sku_code": act_item.get("sku_code"),
                        "name": act_item.get("name"),
                        "size": act_item.get("size"),
                        "qty": planned_qty if planned_qty not in (None, "") else None,
                        "comment": act_item.get("comment"),
                        "actual_qty": actual_qty,
                    }
                )
        elif act_items:
            for item in act_items:
                planned_qty = item.get("planned_qty")
                if planned_qty in (None, ""):
                    planned_qty = item.get("qty")
                actual_qty = _parse_qty_value(item.get("actual_qty"))
                display_items.append(
                    {
                        "sku_code": item.get("sku_code"),
                        "name": item.get("name"),
                        "size": item.get("size"),
                        "qty": planned_qty,
                        "comment": item.get("comment"),
                        "actual_qty": actual_qty,
                    }
                )
        special_item_rows = _build_receiving_special_item_rows(
            (act_entry_for_items.payload or {}) if act_entry_for_items else None,
            (placement_entry_for_items.payload or {}) if placement_entry_for_items else draft_flow_payload,
        )
        display_items = ReceivingWorkflowService.split_receiving_act_rows_by_goods_type(
            display_items,
            special_rows=special_item_rows,
            planned_key="qty",
            actual_key="actual_qty",
        )
        ctx["items"] = display_items
        placement_container_payload = (
            (placement_entry_for_items.payload or {})
            if placement_entry_for_items
            else draft_flow_payload
        )
        placement_pallets, placement_loose_boxes = _placement_payload_container_rows(placement_container_payload)
        ctx["placement_pallets"] = placement_pallets
        ctx["placement_loose_boxes"] = placement_loose_boxes
        if warehouse_invalid_review:
            status_text = "Передана на разбор начальнику склада"
            ctx["status_label"] = status_text
            ctx["next_step_label"] = "Начальник склада проверяет складские следы заявки"
            ctx["next_step_tone"] = "warning"
        if self.order_type == "receiving":
            receiving_agency = latest.agency if latest else client_agency
            ctx["receiving_attachments"] = list(_active_receiving_attachments(str(order_id or ""), receiving_agency))
            ctx["attachment_retention_days"] = ReceivingOrderAttachment.RETENTION_DAYS
            act_payload_for_summary = (act_entry_for_items.payload or {}) if act_entry_for_items else {}
            if not act_payload_for_summary and placement_entry_for_items:
                act_payload_for_summary = _receiving_act_payload_from_placement(
                    entries_list,
                    placement_entry_for_items,
                ) or {}
            needs_manager_act_sign = bool(
                not client_view
                and role in {"manager", "head_manager", "director", "admin"}
                and (
                    _act_storekeeper_signed(entries_list)
                    or _act_storekeeper_signed_from_payload(act_payload_for_summary)
                )
                and not (
                    _act_manager_signed(entries_list)
                    or _act_manager_signed_from_payload(act_payload_for_summary)
                )
            )
            finished_source = None
            if act_entry_for_items and getattr(act_entry_for_items, "created_at", None):
                finished_source = act_entry_for_items.created_at
            ctx["receiving_summary"] = _build_receiving_summary(
                receiving_mode=receiving_mode,
                receiving_started=receiving_started,
                status_label=status_text,
                display_items=display_items,
                placement_pallets=placement_pallets,
                placement_loose_boxes=placement_loose_boxes,
                act_payload=act_payload_for_summary,
                placement_payload=placement_container_payload,
                order_payload=payload,
                finished_at=finished_source,
            )
            ctx["receiving_items_count"] = len(display_items)
            ctx["needs_manager_act_sign"] = needs_manager_act_sign
            ctx["has_receiving_chz_export"] = bool(
                receiving_mode == "cz"
                and (
                    act_payload_for_summary.get("act_units")
                    or any(_parse_qty_value(row.get("actual_qty")) for row in display_items)
                )
            )
        else:
            ctx["receiving_summary"] = None
            ctx["receiving_items_count"] = 0
            ctx["needs_manager_act_sign"] = False
            ctx["has_receiving_chz_export"] = False
        participants = []
        seen = set()
        for entry in entries_list:
            label = _actor_label(entry.user, entry.agency, client_view=client_view)
            if label in seen:
                continue
            seen.add(label)
            participants.append(label)
        ctx["participants"] = participants
        comments = [
            {
                "created_at": entry.created_at,
                "actor_label": _actor_label(entry.user, entry.agency, client_view=client_view),
                "text": entry.description or (entry.payload or {}).get("comment") or "-",
            }
            for entry in entries_list
            if entry.action == "comment"
        ]
        comments.reverse()
        ctx["comments"] = comments
        history_audience = "client" if client_view else ("storekeeper" if role in _RECEIVING_WAREHOUSE_ROLES else "default")
        history_rows = [
            {
                "created_at": entry.created_at,
                "action_label": _history_action_label(entry),
                "status_label": _detail_history_status_label(
                    entry,
                    audience=history_audience,
                    current_entry=status_entry,
                    current_label=status_text,
                ),
                "description": _format_message_text(entry.description),
                "actor_label": _detail_history_actor_label(
                    entry,
                    client_view=client_view,
                    client_label=client_label,
                ),
            }
            for entry in reversed(entries_list)
            if entry.action != "comment"
        ]
        ctx["history"] = history_rows
        ctx["receiving_history_count"] = len(history_rows) if self.order_type == "receiving" else 0
        act_entry = _find_act_entry(entries_list, "receiving", "акт приемки")
        placement_entry = _find_act_entry(entries_list, "placement", "акт размещения")
        ctx["act_entry"] = act_entry
        ctx["placement_act_entry"] = placement_entry
        ctx["has_receiving_act"] = bool(act_entry)
        ctx["has_placement_act"] = bool(placement_entry)
        ctx["can_print_receiving_act"] = bool(
            self.order_type == "receiving"
            and _receiving_act_entry_or_placement(entries_list, act_entry)
        )
        # Кладовщик после закрытия приёмки: CTA «Подписать и отправить акт» (подпись на печатной форме).
        ctx["can_storekeeper_sign_act"] = bool(
            self.order_type == "receiving"
            and not client_view
            and role == "storekeeper"
            and ctx["can_print_receiving_act"]
            and flow_closed
            and _placement_closed(entries_list)
            and not signed_by_storekeeper
        )
        if act_entry:
            ctx["act_label"] = (act_entry.payload or {}).get("act_label") or "Акт приемки"
        else:
            ctx["act_label"] = ""
        if placement_entry:
            ctx["placement_act_label"] = (placement_entry.payload or {}).get("act_label") or "Акт размещения"
        else:
            ctx["placement_act_label"] = ""
        role = get_request_role(self.request)
        ctx["can_create_placement_act"] = bool(
            role in _RECEIVING_WAREHOUSE_ROLES and act_entry and not placement_entry
        )
        ctx["can_send_act_to_client"] = bool(
            role in {"manager", "head_manager", "director", "admin"}
            and act_entry
            and placement_entry
            and not _is_done_status(status_entry)
            and _act_manager_signed_from_payload(act_entry.payload or {})
            and not client_view
        )
        if warehouse_invalid_review:
            for key in (
                "can_close_receiving_task_invalid",
                "can_edit_order",
                "can_cancel_order",
                "can_request_warehouse_cancel",
                "can_review_warehouse_cancel",
                "can_send_to_warehouse",
                "can_create_receiving_act",
                "can_open_flow",
                "can_continue_flow",
                "can_reopen_receiving_flow",
                "can_view_receiving_flow",
                "can_manage_storage_placement",
                "can_storekeeper_sign_act",
                "can_create_placement_act",
                "can_send_act_to_client",
            ):
                ctx[key] = False
        return ctx


class PackingDetailView(OrdersDetailView):
    order_type = "packing"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries_list = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )
        payload = self._payload_from_entries(entries_list)
        marketplaces = payload.get("marketplaces") or []
        if isinstance(marketplaces, str):
            marketplaces = [marketplaces] if marketplaces else []
        mp_other = (payload.get("mp_other") or "").strip()
        if mp_other:
            marketplaces = list(marketplaces) + [mp_other]
        tasks = payload.get("tasks") or []
        if isinstance(tasks, str):
            tasks = [tasks] if tasks else []
        tasks_other = (payload.get("tasks_other") or "").strip()
        if tasks_other:
            tasks = list(tasks) + [tasks_other]
        box_mode = _resolve_choice_value(payload.get("box_mode"), payload.get("box_mode_other"))
        marking = _resolve_choice_value(payload.get("marking"), payload.get("marking_other"))
        ship_as = _resolve_choice_value(payload.get("ship_as"), payload.get("ship_other"))
        ctx["packing_fields"] = [
            {"label": "Email", "value": _format_payload_value(payload.get("email"))},
            {"label": "ФИО", "value": _format_payload_value(payload.get("fio"))},
            {"label": "Организация", "value": _format_payload_value(payload.get("org"))},
            {"label": "Плановая дата", "value": _format_payload_value(payload.get("plan_date"))},
            {"label": "Маркетплейсы", "value": _format_payload_list(marketplaces)},
            {"label": "Предмет", "value": _format_payload_value(payload.get("subject"))},
            {"label": "Общее количество", "value": _format_payload_value(payload.get("total_qty"))},
            {"label": "Кратность коробов", "value": _format_payload_value(box_mode)},
            {"label": "Работы", "value": _format_payload_list(tasks)},
            {"label": "Маркировка", "value": _format_payload_value(marking)},
            {"label": "Отгрузка", "value": _format_payload_value(ship_as)},
            {"label": "Распределение", "value": _format_payload_value(payload.get("has_distribution"))},
            {"label": "Комментарий", "value": _format_payload_value(payload.get("comments"))},
            {"label": "Файлы отчета", "value": _format_payload_list(payload.get("files_report"))},
            {"label": "Файл распределения", "value": _format_payload_list(payload.get("files_distribution"))},
            {"label": "Файлы ЧЗ", "value": _format_payload_list(payload.get("files_cz"))},
        ]
        ctx["packing_meta"] = {
            "plan_date": _format_payload_value(payload.get("plan_date")),
            "total_qty": _format_payload_value(payload.get("total_qty")),
            "marketplaces": _format_payload_list(marketplaces),
            "box_mode": _format_payload_value(box_mode),
            "marking": _format_payload_value(marking),
            "ship_as": _format_payload_value(ship_as),
            "has_distribution": _format_payload_value(payload.get("has_distribution")),
            "comment": _format_payload_value(payload.get("comments")),
        }
        ctx["items"] = []
        ctx["can_edit_order"] = False
        ctx["can_send_to_warehouse"] = False
        ctx["can_create_receiving_act"] = False
        ctx["has_receiving_act"] = False
        ctx["has_placement_act"] = False
        ctx["act_label"] = ""
        ctx["placement_act_label"] = ""
        ctx["can_send_act_to_client"] = False
        return ctx




class ReceivingActView(RoleRequiredMixin, TemplateView):
    template_name = "orders/receiving_act.html"
    order_type = "receiving"
    allowed_roles = ("storekeeper", "processing_head", "manager", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            order_id = kwargs.get("order_id")
            entries = self._load_entries(order_id)
            if not entries:
                return HttpResponseForbidden("Доступ запрещен")
            if not any(entry.agency_id == client_agency.id for entry in entries if entry.agency_id):
                return HttpResponseForbidden("Доступ запрещен")
            if not self._act_entry(entries):
                return HttpResponseForbidden("Доступ запрещен")
            if request.method != "GET":
                return HttpResponseForbidden("Доступ запрещен")
            status_entry = _current_status_entry(entries)
            if status_entry:
                payload = dict(status_entry.payload or {})
                if payload.get("act_sent") and not payload.get("act_viewed"):
                    payload["status"] = payload.get("status") or "done"
                    payload["status_label"] = payload.get("status_label") or "Выполнена"
                    payload["act_viewed"] = True
                    payload["act_viewed_at"] = timezone.localtime().isoformat()
                    log_order_action(
                        "status",
                        order_id=order_id,
                        order_type=self.order_type,
                        user=request.user if request.user.is_authenticated else None,
                        agency=entries[-1].agency if entries else None,
                        description="Акт приемки просмотрен клиентом",
                        payload=payload,
                    )
            response = request.GET.get("response")
            raw_return_url = request.GET.get("return") or request.META.get("HTTP_REFERER")
            return_url = _safe_return_url(request, raw_return_url, request.path)
            return redirect(_client_print_url(order_id, client_agency, response, return_url))
        return super().dispatch(request, *args, **kwargs)

    def _load_entries(self, order_id):
        return list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )

    def _act_entry(self, entries):
        return _act_entry_from_entries(entries, "receiving")

    def _can_create(self, entries):
        return ReceivingWorkflowService.can_create_receiving_act(entries, role="storekeeper")

    def get(self, request, *args, **kwargs):
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        ctx = self.get_context_data(ok=ok, error=error, **kwargs)
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        role = get_request_role(request)
        if role not in _RECEIVING_WAREHOUSE_ROLES:
            return HttpResponseForbidden("Доступ запрещен")
        if role == "storekeeper":
            access = ReceivingWorkflowService.receiving_work_access(
                entries=entries,
                user=request.user if request.user.is_authenticated else None,
            )
            if access.status != "allowed":
                owner_name = str(access.payload.get("storekeeper_name") or "другого кладовщика")
                return HttpResponseForbidden(f"Заявка уже в работе у {owner_name}.")
        if not self._can_create(entries):
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        payload = _latest_payload_from_entries(entries)
        items = payload.get("items") or []
        has_planned_items = bool(items)
        actual_raw = request.POST.getlist("actual_qty[]")
        eta_raw = (request.POST.get("eta_at") or "").strip()
        vehicle_number = (request.POST.get("vehicle_number") or "").strip()
        if not eta_raw:
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        try:
            eta_value = parse_business_datetime(eta_raw)
        except ValueError:
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        if not vehicle_number:
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        act_items = []
        has_mismatch = False
        for idx, item in enumerate(items):
            raw_value = actual_raw[idx] if idx < len(actual_raw) else ""
            actual_value = _parse_qty_value(raw_value)
            if actual_value is None:
                return redirect(f"/orders/receiving/{order_id}/act/?error=1")
            planned_value = _parse_qty_value(item.get("qty")) or 0
            if planned_value != actual_value:
                has_mismatch = True
            act_items.append(
                {
                    "sku_code": item.get("sku_code"),
                    "barcode": item.get("barcode"),
                    "name": item.get("name"),
                    "brand": item.get("brand"),
                    "color": item.get("color"),
                    "size": item.get("size"),
                    "planned_qty": item.get("qty"),
                    "actual_qty": actual_value,
                    "comment": item.get("comment"),
                }
            )
        extra_sku_codes = request.POST.getlist("extra_sku_code[]")
        extra_barcodes = request.POST.getlist("extra_barcode[]")
        extra_names = request.POST.getlist("extra_name[]")
        extra_brands = request.POST.getlist("extra_brand[]")
        extra_colors = request.POST.getlist("extra_color[]")
        extra_sizes = request.POST.getlist("extra_size[]")
        extra_planned = request.POST.getlist("extra_planned_qty[]")
        extra_actual = request.POST.getlist("extra_actual_qty[]")
        extra_count = max(
            len(extra_names),
            len(extra_brands),
            len(extra_colors),
            len(extra_sizes),
            len(extra_planned),
            len(extra_actual),
            len(extra_sku_codes),
            len(extra_barcodes),
        )
        for idx in range(extra_count):
            sku_code = (extra_sku_codes[idx] if idx < len(extra_sku_codes) else "").strip()
            barcode = (extra_barcodes[idx] if idx < len(extra_barcodes) else "").strip()
            name = (extra_names[idx] if idx < len(extra_names) else "").strip()
            brand = (extra_brands[idx] if idx < len(extra_brands) else "").strip()
            color = (extra_colors[idx] if idx < len(extra_colors) else "").strip()
            size = (extra_sizes[idx] if idx < len(extra_sizes) else "").strip()
            planned_raw = extra_planned[idx] if idx < len(extra_planned) else ""
            actual_raw_value = extra_actual[idx] if idx < len(extra_actual) else ""
            if not any((sku_code, barcode, name, brand, color, size, planned_raw, actual_raw_value)):
                continue
            if not (name or sku_code):
                return redirect(f"/orders/receiving/{order_id}/act/?error=1")
            planned_value = _parse_qty_value(planned_raw)
            if planned_value is None:
                planned_value = 0
            actual_value = _parse_qty_value(actual_raw_value)
            if actual_value is None:
                return redirect(f"/orders/receiving/{order_id}/act/?error=1")
            if has_planned_items and planned_value != actual_value:
                has_mismatch = True
            act_items.append(
                {
                    "sku_code": sku_code,
                    "barcode": barcode,
                    "name": name,
                    "brand": brand,
                    "color": color,
                    "size": size,
                    "planned_qty": planned_value,
                    "actual_qty": actual_value,
                    "comment": "",
                    "extra": True,
                }
            )
        if not act_items:
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        if role == "storekeeper":
            start_result = ReceivingWorkflowService.start_receiving_work(
                order_id=str(order_id or ""),
                entries=entries,
                role=role,
                user=request.user if request.user.is_authenticated else None,
            )
            if start_result.status == "assigned_to_other_storekeeper":
                owner_name = str(start_result.payload.get("storekeeper_name") or "другого кладовщика")
                return HttpResponseForbidden(f"Заявка уже в работе у {owner_name}.")
            entries = self._load_entries(order_id)
        status_entry = _current_status_entry(entries)
        latest = entries[-1]
        act_payload = dict((status_entry.payload or {}) if status_entry else {})
        goods_type = str(act_payload.get("goods_type") or "").strip().lower()
        try:
            ReceivingWorkflowService._attach_receiving_nomenclature_metadata(
                agency=latest.agency,
                items=act_items,
                goods_type=goods_type,
                order_id=str(order_id or ""),
                order_type="receiving",
            )
        except ReceivingNomenclatureError:
            return redirect(f"/orders/receiving/{order_id}/act/?error=1")
        if not act_payload.get("status"):
            act_payload["status"] = "warehouse"
        act_payload["status_label"] = (
            "Товар принят (с расхождениями)" if has_mismatch else "Товар принят"
        )
        act_payload["eta_at"] = eta_value.isoformat()
        act_payload["vehicle_number"] = vehicle_number
        act_payload["act"] = "receiving"
        act_payload["act_label"] = (
            "Акт приемки с расхождениями" if has_mismatch else "Акт приемки"
        )
        act_payload["act_mismatch"] = has_mismatch
        act_payload["act_items"] = act_items
        act_payload["flow_closed"] = True
        act_payload["flow_closed_at"] = timezone.localtime().isoformat()
        log_order_action(
            "status",
            order_id=order_id,
            order_type=self.order_type,
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency,
            description="Создан акт приемки",
            payload=act_payload,
        )
        return redirect(f"/orders/receiving/{order_id}/placement/?ok=1")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        client_agency = getattr(self.request, "_client_agency", None)
        client_view = bool(client_agency)
        role = get_request_role(self.request)
        cabinet_url = resolve_cabinet_url(get_request_role(self.request))
        ctx.update(
            ReceivingWorkflowService.build_receiving_act_page_context(
                order_id=str(order_id or ""),
                entries=entries,
                role=role,
                client_view=client_view,
                client_agency=client_agency,
                cabinet_url=cabinet_url,
            )
        )
        ctx["ok"] = kwargs.get("ok", False)
        ctx["error"] = kwargs.get("error")
        return ctx


def _receiving_owner_lock_response(request, *, order_id: str, entries, json_response: bool = False):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head"}:
        return None
    access = ReceivingWorkflowService.receiving_work_access(
        entries=entries,
        user=request.user if request.user.is_authenticated else None,
    )
    if access.status == "allowed":
        return None
    owner_name = str(access.payload.get("storekeeper_name") or "другого кладовщика")
    message = f"Заявка уже в работе у {owner_name}."
    if json_response:
        return JsonResponse(
            {
                "ok": False,
                "error": "assigned_to_other_storekeeper",
                "storekeeper_name": owner_name,
            },
            status=409,
        )
    return HttpResponseForbidden(message)


class ReceivingFlowView(RoleRequiredMixin, TemplateView):
    template_name = "orders/receiving_flow.html"
    order_type = "receiving"
    allowed_roles = ("storekeeper", "processing_head", "manager", "head_manager", "director", "admin")

    def _load_entries(self, order_id):
        return list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )

    def _receiving_act_entry(self, entries):
        return _act_entry_from_entries(entries, "receiving")

    def _placement_act_entry(self, entries):
        return _act_entry_from_entries(entries, "placement")

    def _receiving_state_result(self, entries):
        return ReceivingWorkflowService._receiving_state_result(entries)

    def _can_start(self, entries):
        return ReceivingWorkflowService.can_start_receiving_flow(entries)

    def _start_from_post(self, request, order_id: str | None, entries):
        role = get_request_role(request)
        if role != "storekeeper":
            return HttpResponseForbidden("Доступ запрещен")
        if not order_id or not entries or _flow_closed_from_entries(entries):
            return redirect(f"/orders/receiving/{order_id}/flow/?error=1")
        start_result = ReceivingWorkflowService.start_receiving_work(
            order_id=str(order_id or ""),
            entries=entries,
            role=role,
            user=request.user if request.user.is_authenticated else None,
        )
        if start_result.status == "assigned_to_other_storekeeper":
            owner_name = str(start_result.payload.get("storekeeper_name") or "другого кладовщика")
            messages.error(request, f"Заявка уже в работе у {owner_name}.")
            return redirect(f"/orders/receiving/{order_id}/")
        if start_result.status in {"started", "already_in_progress"}:
            return redirect(f"/orders/receiving/{order_id}/flow/")
        return redirect(f"/orders/receiving/{order_id}/flow/?error=1")


    def _find_flow_state(self, entries):
        return ReceivingWorkflowService.find_receiving_flow_state(entries)

    def _normalize_flow_state(self, boxes_data, pallets_data, active_box, active_pallet):
        entries = self._load_entries(self.kwargs.get("order_id"))
        status_entry = ReceivingWorkflowService._current_status_entry(entries)
        return ReceivingWorkflowService.normalize_receiving_flow_state(
            boxes_data=boxes_data,
            pallets_data=pallets_data,
            active_box=active_box,
            active_pallet=active_pallet,
            default_goods_type=str((status_entry.payload or {}).get("goods_type") or "").strip() if status_entry else "",
        )

    def _set_goods_type(self, request, order_id, entries):
        role = get_request_role(request)
        if role not in _RECEIVING_WAREHOUSE_ROLES:
            return HttpResponseForbidden("Доступ запрещен")
        if _flow_closed_from_entries(entries):
            return redirect(f"/orders/receiving/{order_id}/flow/?goods_type_error=1")
        status_entry = ReceivingWorkflowService._current_status_entry(entries)
        status_payload = dict(status_entry.payload or {}) if status_entry else {}
        goods_type = (request.POST.get("goods_type") or "").strip().lower()
        receiving_mode = (
            (request.POST.get("receiving_mode") or "").strip().lower()
            or str(status_payload.get("receiving_mode") or "standard").strip().lower()
        )
        result = ReceivingWorkflowService.configure_receiving_act(
            order_id=str(order_id or ""),
            entries=entries,
            role=role or "",
            goods_type=goods_type,
            receiving_mode=receiving_mode,
            user=request.user if request.user.is_authenticated else None,
        )
        if not result.applied:
            return redirect(f"/orders/receiving/{order_id}/flow/?goods_type_error=1")
        return redirect(f"/orders/receiving/{order_id}/flow/?goods_type_updated=1")

    def _save_flow_draft(self, request, order_id, entries):
        boxes_raw = request.POST.get("boxes_json") or "[]"
        pallets_raw = request.POST.get("pallets_json") or "[]"
        result = ReceivingWorkflowService.save_receiving_flow_draft(
            order_id=str(order_id or ""),
            entries=entries,
            role=get_request_role(request) or "",
            boxes_raw=boxes_raw,
            pallets_raw=pallets_raw,
            active_box=request.POST.get("active_box") or "",
            active_pallet=request.POST.get("active_pallet") or "",
            received_at=request.POST.get("received_at") or "",
            draft_version=request.POST.get("flow_client_version") or "",
            draft_base_version=request.POST.get("flow_base_version") or "",
            user=request.user if request.user.is_authenticated else None,
        )
        if result.status == "forbidden":
            return JsonResponse({"ok": False, "error": "forbidden"}, status=403)
        if result.status == "closed":
            return JsonResponse({"ok": False, "error": "closed"}, status=400)
        if result.status == "not_allowed":
            return JsonResponse({"ok": False, "error": "not_allowed"}, status=400)
        if result.status == "invalid_json":
            return JsonResponse({"ok": False, "error": "invalid_json"}, status=400)
        if result.status == "pallet_composition_locked":
            pallet_codes = result.meta.get("pallet_codes") or []
            return JsonResponse(
                {
                    "ok": False,
                    "error": "pallet_composition_locked",
                    "pallet_codes": pallet_codes,
                    "message": (
                        "Состав паллеты нельзя менять после передачи на размещение. "
                        "Для изменения выполните складскую корректировку."
                    ),
                },
                status=409,
            )
        if result.status == "goods_type_locked":
            container_codes = result.meta.get("container_codes") or []
            return JsonResponse(
                {
                    "ok": False,
                    "error": "goods_type_locked",
                    "container_codes": container_codes,
                    "message": (
                        "Тип приемки зафиксирован при начале работы и не может быть изменен."
                    ),
                },
                status=409,
            )
        if result.status == "mixed_pallet_goods_types":
            pallet_codes = result.meta.get("pallet_codes") or []
            return JsonResponse(
                {
                    "ok": False,
                    "error": "mixed_pallet_goods_types",
                    "pallet_codes": pallet_codes,
                    "message": (
                        "Готовый и необработанный товар нельзя объединять на одной палете. "
                        "Разделите короба по палетам одного типа."
                    ),
                },
                status=409,
            )
        if result.meta.get("stale_ignored"):
            return JsonResponse(
                {
                    "ok": False,
                    "error": "stale_draft",
                    "flow_client_version": result.meta.get("flow_client_version", 0),
                },
                status=409,
            )
        return JsonResponse(
            {
                "ok": True,
                "flow_client_version": result.payload.get("flow_client_version", 0),
            }
        )

    def _reopen_flow(self, request, order_id, entries):
        result = ReceivingWorkflowService.reopen_receiving_flow(
            order_id=str(order_id or ""),
            entries=entries,
            role=get_request_role(request) or "",
            user=request.user if request.user.is_authenticated else None,
        )
        if result.status == "already_open":
            return redirect(f"/orders/receiving/{order_id}/flow/")
        if result.status == "denied":
            if result.reason == "role_forbidden":
                return HttpResponseForbidden("Доступ запрещен")
            return redirect(f"/orders/receiving/{order_id}/flow/?error=1")
        return redirect(f"/orders/receiving/{order_id}/flow/?ok=1")

    def _send_to_warehouse(self, request, order_id, entries):
        role = get_request_role(request)
        if role not in _RECEIVING_MOVE_CREATE_ROLES:
            return HttpResponseForbidden("Доступ запрещен")
        order_key = str(order_id or "")
        draft_token = request.POST.get("warehouse_draft_token") or ""
        destinations_raw = request.POST.get("warehouse_destinations_json")
        auto_pallet_codes: list[str] = []
        try:
            parsed_pallet_codes = json.loads(request.POST.get("warehouse_pallet_codes_json") or "[]")
        except (TypeError, ValueError, json.JSONDecodeError):
            parsed_pallet_codes = []
        if isinstance(parsed_pallet_codes, list):
            auto_pallet_codes = [
                str(raw_code or "").strip()
                for raw_code in parsed_pallet_codes
                if str(raw_code or "").strip()
            ]
        if auto_pallet_codes and not str(destinations_raw or "").strip():
            auto_destinations = ReceivingWorkflowService.collect_receiving_warehouse_move_destinations(
                order_id=order_key,
                entries=entries,
                pallet_codes=auto_pallet_codes,
                agency_id=getattr(entries[-1], "agency_id", None) if entries else None,
                row_sections=_OS_ROW_SECTIONS,
                tiers=_OS_TIERS,
                cells_per_tier=_OS_CELLS_PER_TIER,
            )
            if not auto_destinations:
                return redirect(
                    f"/orders/receiving/{order_id}/flow/?warehouse_move=missing_destination"
                    f"&warehouse_created=0&warehouse_skipped=0&warehouse_missing={len(auto_pallet_codes)}"
                )
            result = ReceivingWorkflowService.create_receiving_warehouse_moves(
                order_id=order_key,
                entries=entries,
                role=role or "",
                user=request.user if request.user.is_authenticated else None,
                destinations_by_pallet=auto_destinations,
                draft_token=draft_token,
            )
        else:
            result = ReceivingWorkflowService.send_receiving_to_storage(
                order_id=order_key,
                entries=entries,
                role=role or "",
                user=request.user if request.user.is_authenticated else None,
                destinations_raw=destinations_raw,
                draft_token=draft_token,
            )
        if result.status == "not_ready":
            return redirect(f"/orders/receiving/{order_id}/flow/?warehouse_move=not_ready")
        if result.status == "invalid_destination":
            error_text = quote_plus(result.error_message)
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=invalid_destination"
                f"&warehouse_error={error_text}"
            )
        created = int(result.created_count or 0)
        skipped_existing = int(result.skipped_existing_count or 0)
        skipped_missing = int(result.skipped_missing_destination_count or 0)
        total = int(result.total_count or 0)
        if result.status == "ok":
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=ok"
                f"&warehouse_created={created}"
                f"&warehouse_skipped={skipped_existing}"
                f"&warehouse_missing={skipped_missing}"
            )
        if result.status == "missing_destination":
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=missing_destination"
                f"&warehouse_created=0&warehouse_skipped=0&warehouse_missing={skipped_missing}"
            )
        if result.status == "exists":
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=exists"
                f"&warehouse_created=0&warehouse_skipped={skipped_existing}&warehouse_missing=0"
            )
        if result.status == "blocked":
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=blocked"
                f"&warehouse_created=0"
                f"&warehouse_skipped={skipped_existing}"
                f"&warehouse_missing={skipped_missing}"
            )
        return redirect(f"/orders/receiving/{order_id}/flow/?warehouse_move=none")

    def _cancel_warehouse_move(self, request, order_id, entries):
        role = get_request_role(request)
        if role not in _RECEIVING_MOVE_CREATE_ROLES:
            return HttpResponseForbidden("Доступ запрещен")
        move_order_id = str(request.POST.get("warehouse_move_order_id") or "").strip()
        if not move_order_id:
            return redirect(f"/orders/receiving/{order_id}/flow/?warehouse_move=none")
        latest_moves = ReceivingWorkflowService._latest_receiving_moves_by_pallet(str(order_id or ""))
        allowed_order_ids = {
            str((payload or {}).get("order_id") or "").strip()
            for payload in latest_moves.values()
            if str((payload or {}).get("order_id") or "").strip()
        }
        if move_order_id not in allowed_order_ids:
            return redirect(f"/orders/receiving/{order_id}/flow/?warehouse_move=none")
        from reachtruck.services.ui_flows import cancel_move_before_take

        ok, message = cancel_move_before_take(
            legacy_order_id=move_order_id,
            actor_role=role or "",
        )
        if not ok:
            error_text = quote_plus(message)
            return redirect(
                f"/orders/receiving/{order_id}/flow/?warehouse_move=cancel_error"
                f"&warehouse_error={error_text}"
            )
        return redirect(f"/orders/receiving/{order_id}/flow/?warehouse_move=cancelled")

    def get(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        if get_request_role(request) == "storekeeper":
            access = ReceivingWorkflowService.receiving_work_access(
                entries=entries,
                user=request.user if request.user.is_authenticated else None,
            )
            if access.status == "assigned_to_other_storekeeper":
                owner_name = str(access.payload.get("storekeeper_name") or "другого кладовщика")
                messages.error(request, f"Заявка уже в работе у {owner_name}.")
                return redirect(f"/orders/receiving/{order_id}/")
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        warehouse_error = str(request.GET.get("warehouse_error") or "").strip()
        ctx = self.get_context_data(
            ok=ok,
            error=error,
            warehouse_error=warehouse_error,
            **kwargs,
        )
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        owner_lock = _receiving_owner_lock_response(
            request,
            order_id=str(order_id or ""),
            entries=entries,
            json_response=request.POST.get("flow_action") == "draft",
        )
        if owner_lock:
            return owner_lock
        if request.POST.get("flow_action") == "start":
            return self._start_from_post(request, order_id, entries)
        if request.POST.get("flow_action") == "draft":
            return self._save_flow_draft(request, order_id, entries)
        if request.POST.get("flow_action") == "reopen":
            return self._reopen_flow(request, order_id, entries)
        if request.POST.get("flow_action") == "set_goods_type":
            return self._set_goods_type(request, order_id, entries)
        if request.POST.get("flow_action") == "send_to_warehouse":
            return redirect(f"/orders/receiving/{order_id}/flow/")
        if request.POST.get("flow_action") == "cancel_warehouse_move":
            return self._cancel_warehouse_move(request, order_id, entries)
        role = get_request_role(request)
        if role not in _RECEIVING_WAREHOUSE_ROLES:
            return HttpResponseForbidden("Доступ запрещен")
        if _flow_closed_from_entries(entries):
            return redirect(f"/orders/receiving/{order_id}/flow/")
        if not self._can_start(entries):
            return redirect(f"/orders/receiving/{order_id}/flow/?error=1")
        received_at_raw = (request.POST.get("received_at") or "").strip()
        try:
            received_at = parse_business_datetime(received_at_raw) if received_at_raw else business_localtime()
        except ValueError:
            return redirect(f"/orders/receiving/{order_id}/flow/?error=1")
        use_saved_draft = (request.POST.get("flow_use_saved_draft") or "").strip() == "1"
        if use_saved_draft:
            flow_state = ReceivingWorkflowService.find_receiving_flow_state(entries)
            boxes_raw = json.dumps(flow_state.get("boxes") or [], ensure_ascii=False)
            pallets_raw = json.dumps(flow_state.get("pallets") or [], ensure_ascii=False)
        else:
            boxes_raw = request.POST.get("boxes_json") or "[]"
            pallets_raw = request.POST.get("pallets_json") or "[]"
        preparation = ReceivingWorkflowService.prepare_receiving_flow_completion(
            order_id=str(order_id or ""),
            entries=entries,
            boxes_raw=boxes_raw,
            pallets_raw=pallets_raw,
        )
        if preparation.status != "ok":
            ReceivingWorkflowService.save_receiving_flow_draft(
                order_id=str(order_id or ""),
                entries=entries,
                role=role or "",
                boxes_raw=boxes_raw,
                pallets_raw=pallets_raw,
                active_box=request.POST.get("active_box") or "",
                active_pallet=request.POST.get("active_pallet") or "",
                user=request.user if request.user.is_authenticated else None,
            )
            return_error_messages = {
                "shipping_return_source_not_found": "Не найдена выбранная отгрузка для возврата.",
                "shipping_return_item_not_in_source": "В приемке есть товар, которого не было в выбранной отгрузке.",
                "shipping_return_qty_exceeded": "Количество возврата превышает доступное количество по выбранной отгрузке.",
                "receiving_mixed_pallet_goods_types": (
                    "Нельзя завершить приемку: готовый и необработанный товар находятся "
                    "на одной палете. Разделите короба по палетам одного типа."
                ),
            }
            if preparation.reason in return_error_messages:
                query = urlencode(
                    {
                        "error": "warehouse_locked",
                        "warehouse_error": return_error_messages[preparation.reason],
                    }
                )
                return redirect(f"/orders/receiving/{order_id}/flow/?{query}")
            return redirect(f"/orders/receiving/{order_id}/flow/?error=1")

        latest = entries[-1]
        try:
            completion_result = ReceivingWorkflowService.complete_receiving_flow(
                order_id=str(order_id or ""),
                agency=latest.agency,
                status_payload=preparation.status_payload,
                has_mismatch=preparation.has_mismatch,
                receiving_mode=preparation.receiving_mode,
                act_items=preparation.act_items,
                placement_items=preparation.placement_items,
                boxes=preparation.boxes,
                pallets=preparation.pallets,
                flow_state=preparation.flow_state,
                act_units=preparation.act_units,
                eta_at=preparation.normalized_eta,
                received_at=received_at.isoformat(),
                vehicle_number=preparation.vehicle_number,
                has_closed_placement_act=preparation.has_closed_placement_act,
                flow_client_version=request.POST.get("flow_client_version", ""),
                receiving_location_code=request.POST.get("receiving_location_code", ""),
                user=request.user if request.user.is_authenticated else None,
                submitted_at=timezone.localtime(),
            )
        except ValidationError as exc:
            message = "; ".join(str(item) for item in getattr(exc, "messages", []) if str(item)) or str(exc)
            query = urlencode(
                {
                    "error": "services_required",
                    "services_error": message,
                }
            )
            return redirect(f"/orders/receiving/{order_id}/flow/?{query}")
        except ValueError as exc:
            query = urlencode(
                {
                    "error": "warehouse_locked",
                    "warehouse_error": str(exc),
                }
            )
            return redirect(f"/orders/receiving/{order_id}/flow/?{query}")
        if not completion_result.already_completed:
            try:
                from receiving_distribution.services.lifecycle import sync_accepted_from_flow_boxes
                from receiving_distribution.services.reserve_linked import sync_linked_shipping_reserves

                sync_accepted_from_flow_boxes(
                    receiving_order_id=str(order_id or ""),
                    boxes=boxes_raw,
                    pallets=pallets_raw,
                    user=request.user if request.user.is_authenticated else None,
                    source="orders.receiving_flow.complete",
                )
                # Резерв связанных OTG только после фактического complete (остатки уже на складе).
                sync_linked_shipping_reserves(
                    receiving_order_id=str(order_id or ""),
                    user=request.user if request.user.is_authenticated else None,
                    source="orders.receiving_flow.complete",
                )
            except Exception:
                pass
        return redirect(f"/orders/receiving/{order_id}/flow/?ok=1")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        role = get_request_role(self.request)
        entries = self._load_entries(order_id)
        ctx.update(
            ReceivingWorkflowService.build_receiving_flow_page_context(
                order_id=str(order_id or ""),
                entries=entries,
                role=role,
                user=self.request.user,
                session_key=self.request.session.session_key or "",
                query_params={
                    "warehouse_move": self.request.GET.get("warehouse_move"),
                    "warehouse_created": self.request.GET.get("warehouse_created"),
                    "warehouse_skipped": self.request.GET.get("warehouse_skipped"),
                    "warehouse_missing": self.request.GET.get("warehouse_missing"),
                    "warehouse_error": unquote_plus(str(self.request.GET.get("warehouse_error") or "").strip()),
                },
                row_sections=_OS_ROW_SECTIONS,
                tiers=_OS_TIERS,
                cells_per_tier=_OS_CELLS_PER_TIER,
            )
        )
        flow_client_version = 0
        for entry in reversed(entries):
            entry_payload = entry.payload or {}
            if isinstance(entry_payload.get("flow_state"), dict):
                flow_client_version = ReceivingWorkflowService._parse_flow_client_version(
                    entry_payload.get("flow_client_version")
                )
                break
        ctx["flow_client_version"] = flow_client_version
        client_agency = next((entry.agency for entry in reversed(entries) if entry.agency_id), None)
        if role in {"manager", "head_manager", "director", "admin"} and client_agency is not None:
            ctx["cabinet_url"] = f"/client/dashboard/lk/?client={client_agency.id}"
        else:
            ctx["cabinet_url"] = resolve_cabinet_url(role)
        ctx["os_config"] = {
            "row_sections": _OS_ROW_SECTIONS,
            "tiers": _OS_TIERS,
            "cells_per_tier": _OS_CELLS_PER_TIER,
            "mr_rows": [1, 2, 3, 4],
        }
        ctx["ok"] = kwargs.get("ok", False)
        ctx["error"] = kwargs.get("error")
        ctx["services_error"] = str(self.request.GET.get("services_error") or "").strip()
        ctx["can_capture_receiving_services"] = bool(
            role in _RECEIVING_WAREHOUSE_ROLES and not ctx.get("flow_locked", False)
        )
        ctx["goods_type_updated"] = self.request.GET.get("goods_type_updated") == "1"
        ctx["goods_type_error"] = self.request.GET.get("goods_type_error") == "1"
        ctx["container_code_url"] = f"/orders/receiving/{order_id}/flow/container-code/"
        ctx["box_code_rebind_url"] = f"/orders/receiving/{order_id}/flow/rebind-box-code/"
        receiving_agency = entries[-1].agency if entries else None
        ctx["receiving_attachments"] = list(_active_receiving_attachments(str(order_id or ""), receiving_agency))
        ctx["attachment_retention_days"] = ReceivingOrderAttachment.RETENTION_DAYS
        return ctx


def _receiving_flow_box_action_meta(action: str) -> tuple[str, str]:
    mapping = {
        "edit": ("update", "Редактирование короба"),
        "delete": ("delete", "Удаление короба"),
        "delete_batch": ("delete", "Удаление группы коробов"),
        "move_batch": ("update", "Перемещение группы коробов"),
        "print_batch": ("update", "Печать этикеток группы коробов"),
    }
    return mapping.get(action, ("", ""))


def receiving_flow_box_action(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in _RECEIVING_WAREHOUSE_ROLES:
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id, "receiving")
    owner_lock = _receiving_owner_lock_response(
        request,
        order_id=str(order_id or ""),
        entries=entries,
        json_response=True,
    )
    if owner_lock:
        return owner_lock
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    action = (payload.get("action") or "").lower()
    action_kind, action_label = _receiving_flow_box_action_meta(action)
    if not action_label:
        return JsonResponse({"ok": False, "error": "invalid_action"}, status=400)
    result = ReceivingWorkflowService.log_receiving_flow_box_action(
        order_id=str(order_id or ""),
        action=action,
        payload=payload,
        user=request.user if request.user.is_authenticated else None,
    )
    if result.status == "missing_box":
        return JsonResponse({"ok": False, "error": "missing_box"}, status=400)
    if result.status == "missing_boxes":
        return JsonResponse({"ok": False, "error": "missing_boxes"}, status=400)
    return JsonResponse({"ok": True})


def receiving_flow_pallet_placement_permission(request, order_id: str):
    if request.method not in {"GET", "POST"}:
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role != "storekeeper":
        return HttpResponseForbidden("Разрешить размещение может только кладовщик")
    entries = _load_order_entries(order_id, "receiving")
    owner_lock = _receiving_owner_lock_response(
        request,
        order_id=str(order_id or ""),
        entries=entries,
        json_response=True,
    )
    if owner_lock:
        return owner_lock
    if request.method == "GET":
        from fbs.services.receiving_destinations import receiving_fbs_destination_options
        from fbs.services.receiving_movements import receiving_fbs_route_rows
        agency_id = next((entry.agency_id for entry in reversed(entries) if entry.agency_id), None)
        if not agency_id:
            return JsonResponse({"ok": False, "error": "missing_order"}, status=404)
        return JsonResponse({
            "ok": True,
            "options": receiving_fbs_destination_options(agency_id=agency_id),
            "routes": receiving_fbs_route_rows(agency_id=agency_id, order_id=order_id),
        })
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    pallet_code = str(payload.get("pallet_code") or "").strip()
    result = ReceivingWorkflowService.allow_receiving_flow_pallet_placement(
        order_id=str(order_id or ""),
        pallet_code=pallet_code,
        receiving_location_code=str(payload.get("receiving_location_code") or "").strip(),
        fbs_cell_id=payload.get("fbs_cell_id"),
        user=request.user,
        flow_client_version=payload.get("flow_client_version"),
    )
    status_key = str(result.get("status") or "")
    if status_key in {"allowed", "already_allowed", "already_placed"}:
        return JsonResponse(
            {
                "ok": True,
                "status": status_key,
                "pallet_code": result.get("pallet_code") or pallet_code,
                "takeout_status": result.get("takeout_status") or {},
                "fbs_route": result.get("fbs_route") or {},
            }
        )
    if status_key in {"nomenclature_invalid", "stale_draft", "mixed_pallet_goods_types"}:
        return JsonResponse({"ok": False, "error": status_key, "message": result.get("message")}, status=409)
    if status_key in {
        "receiving_location_required",
        "invalid_receiving_location",
        "receiving_location_occupied",
        "fbs_destination_required",
        "fbs_destination_invalid",
    }:
        return JsonResponse(
            {"ok": False, "error": status_key, "message": result.get("message")},
            status=409,
        )
    if status_key == "pallet_open":
        return JsonResponse(
            {"ok": False, "error": "pallet_open", "message": "Сначала закройте паллету."},
            status=409,
        )
    if status_key == "missing_order":
        return JsonResponse({"ok": False, "error": "missing_order"}, status=404)
    return JsonResponse({"ok": False, "error": "missing_pallet"}, status=404)


def receiving_flow_container_code(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Р”РѕСЃС‚СѓРї Р·Р°РїСЂРµС‰РµРЅ")
    role = get_request_role(request)
    if role not in _RECEIVING_WAREHOUSE_ROLES:
        return HttpResponseForbidden("Р”РѕСЃС‚СѓРї Р·Р°РїСЂРµС‰РµРЅ")
    entries = _load_order_entries(order_id, "receiving")
    owner_lock = _receiving_owner_lock_response(
        request,
        order_id=str(order_id or ""),
        entries=entries,
        json_response=True,
    )
    if owner_lock:
        return owner_lock
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    latest_with_agency = next((entry for entry in reversed(entries) if entry.agency_id), None)
    if latest_with_agency is None or not latest_with_agency.agency_id:
        return JsonResponse({"ok": False, "error": "order_agency_not_found"}, status=404)
    goods_type = str(payload.get("goods_type") or "").strip().lower()
    if not goods_type:
        for entry in reversed(entries):
            entry_payload = entry.payload if isinstance(entry.payload, dict) else {}
            goods_type = str(entry_payload.get("goods_type") or "").strip().lower()
            if goods_type:
                break
    container_kind = normalize_container_kind(payload.get("kind"))
    if container_kind == "pallet":
        issued = issue_pallet_code(
            agency=latest_with_agency.agency,
            goods_type=goods_type,
            order_type="receiving",
            order_id=order_id,
        )
    else:
        issued = issue_container_code(
            agency=latest_with_agency.agency,
            goods_type=goods_type,
            received_at=str(payload.get("received_at") or "").strip() or None,
        )
    return JsonResponse({"ok": True, **issued})


def receiving_flow_rebind_box_code(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in _RECEIVING_WAREHOUSE_ROLES:
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id, "receiving")
    owner_lock = _receiving_owner_lock_response(
        request,
        order_id=str(order_id or ""),
        entries=entries,
        json_response=True,
    )
    if owner_lock:
        return owner_lock
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    previous_code = str(payload.get("previous_code") or "").strip()
    next_code = str(payload.get("next_code") or "").strip()
    goods_type = str(payload.get("goods_type") or "").strip()
    if not previous_code or not next_code:
        return JsonResponse({"ok": False, "error": "missing_code"}, status=400)
    if goods_type:
        expected_code = rewrite_container_code_goods_type(previous_code, goods_type)
        if expected_code != next_code:
            return JsonResponse({"ok": False, "error": "invalid_target_code"}, status=400)
    updated_count = MarkingCode.objects.filter(
        order_type="receiving",
        order_id=str(order_id or ""),
        box_barcode=previous_code,
    ).update(box_barcode=next_code)
    return JsonResponse({"ok": True, "updated": updated_count})


def receiving_flow_set_item_weight(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in _RECEIVING_WAREHOUSE_ROLES:
        return HttpResponseForbidden("Доступ запрещен")
    entries = _load_order_entries(order_id, "receiving")
    owner_lock = _receiving_owner_lock_response(
        request,
        order_id=str(order_id or ""),
        entries=entries,
        json_response=True,
    )
    if owner_lock:
        return owner_lock
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    sku_code = str(payload.get("sku_code") or "").strip()
    if not sku_code:
        return JsonResponse({"ok": False, "error": "missing_sku_code"}, status=400)
    raw_weight = str(payload.get("weight_kg") or "").strip().replace(",", ".")
    if not raw_weight:
        return JsonResponse({"ok": False, "error": "missing_weight"}, status=400)
    try:
        weight_kg = Decimal(raw_weight)
    except (InvalidOperation, ValueError):
        return JsonResponse({"ok": False, "error": "invalid_weight"}, status=400)
    if weight_kg <= 0:
        return JsonResponse({"ok": False, "error": "invalid_weight"}, status=400)
    weight_kg = weight_kg.quantize(Decimal("0.001"))

    latest = entries[-1] if entries else None
    agency_id = latest.agency_id if latest else None
    if not agency_id:
        return JsonResponse({"ok": False, "error": "agency_not_found"}, status=404)
    sku = SKU.objects.filter(agency_id=agency_id, sku_code=sku_code, deleted=False).first()
    if not sku:
        return JsonResponse({"ok": False, "error": "sku_not_found"}, status=404)

    sku.weight_kg = weight_kg
    sku.weight_gross_kg = weight_kg
    sku.save(update_fields=["weight_kg", "weight_gross_kg", "updated_at"])
    return JsonResponse(
        {
            "ok": True,
            "sku_code": sku.sku_code,
            "weight_kg": f"{weight_kg:.3f}",
        }
    )


class PlacementActView(RoleRequiredMixin, TemplateView):
    template_name = "orders/placement_act.html"
    order_type = "receiving"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin", "manager")

    def _load_entries(self, order_id):
        return list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )

    def _receiving_act_entry(self, entries):
        return _act_entry_from_entries(entries, "receiving")

    def _placement_act_entry(self, entries):
        return _act_entry_from_entries(entries, "placement")

    def _receiving_state_result(self, entries):
        if not entries:
            return None
        status_entry = _current_status_entry(entries)
        status_payload = status_entry.payload or {} if status_entry else {}
        latest = entries[-1]
        return WarehouseGoodsStateResolver.resolve_for_receiving_order(
            order_id=str(latest.order_id or ""),
            agency=latest.agency,
            payload=status_payload,
        )

    def _can_create(self, entries):
        if not entries:
            return False
        if self.order_type == "receiving":
            return WarehouseActionPolicy.can_create_receiving_placement(
                self._receiving_state_result(entries),
                has_receiving_act=bool(self._receiving_act_entry(entries)),
            ).allowed
        return bool(self._receiving_act_entry(entries))

    def get(self, request, *args, **kwargs):
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        ctx = self.get_context_data(ok=ok, error=error, **kwargs)
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        if request.POST.get("flow_action") == "draft":
            return self._save_flow_draft(request, order_id, entries)
        action = (request.POST.get("action") or "close").lower()
        role = get_request_role(request)
        if role not in _RECEIVING_WAREHOUSE_ROLES:
            return redirect(f"/orders/receiving/{order_id}/placement/")
        if action == "open" and _act_storekeeper_signed(entries):
            return redirect(f"/orders/receiving/{order_id}/placement/?error=signed")
        if not self._can_create(entries):
            return redirect(f"/orders/receiving/{order_id}/placement/?error=1")
        placement_act = self._placement_act_entry(entries)
        if action == "open":
            if not placement_act:
                return redirect(f"/orders/receiving/{order_id}/placement/")
            payload = placement_act.payload or {}
            current_state = (payload.get("act_state") or "closed").lower()
            if self.order_type == "receiving":
                result = ReceivingWorkflowService.open_receiving_placement(
                    order_id=str(order_id or ""),
                    entries=entries,
                    role=role or "",
                    user=request.user if request.user.is_authenticated else None,
                )
                if result.status == "already_open":
                    return redirect(f"/orders/receiving/{order_id}/placement/")
                if result.status == "denied":
                    return redirect(f"/orders/receiving/{order_id}/placement/?error=1")
            elif current_state == "open":
                return redirect(f"/orders/receiving/{order_id}/placement/")
            if self.order_type == "receiving":
                return redirect(f"/orders/receiving/{order_id}/placement/")
            else:
                act_payload = dict(payload)
                if not act_payload.get("status"):
                    act_payload["status"] = "warehouse"
                act_payload["status_label"] = "Размещение на складе"
                act_payload["act"] = "placement"
                act_payload["act_label"] = act_payload.get("act_label") or "Акт размещения"
                act_payload["act_state"] = "open"
            latest = entries[-1]
            log_order_action(
                "status",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency,
                description="Открыт акт размещения",
                payload=act_payload,
            )
            if self.order_type != "receiving":
                OperationalStockService.clear_order_placement(
                    latest.agency,
                    self.order_type,
                    order_id,
                )
            return redirect(f"/orders/receiving/{order_id}/placement/")
        receiving_act = self._receiving_act_entry(entries)
        if not receiving_act:
            return redirect(f"/orders/receiving/{order_id}/placement/?error=1")
        preparation = ReceivingWorkflowService.prepare_receiving_placement_close(
            order_id=str(order_id or ""),
            entries=entries,
            boxes_raw=request.POST.get("boxes_json") or "[]",
            pallets_raw=request.POST.get("pallets_json") or "[]",
            order_type=self.order_type,
        )
        if preparation.status != "ok":
            return redirect(f"/orders/receiving/{order_id}/placement/?error=1")
        status_entry = _current_status_entry(entries)
        latest = entries[-1]
        ReceivingWorkflowService.close_receiving_placement(
            order_id=str(order_id or ""),
            agency=latest.agency,
            status_payload=(status_entry.payload or {}) if status_entry else {},
            placement_items=preparation.placement_items,
            boxes=preparation.boxes,
            pallets=preparation.pallets,
            has_closed_act=preparation.has_closed_act,
            user=request.user if request.user.is_authenticated else None,
            submitted_at=timezone.localtime(),
        )
        return redirect(f"/orders/receiving/{order_id}/placement/?ok=1")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        role = get_request_role(self.request)
        ctx.update(ReceivingWorkflowService.build_receiving_placement_page_context(
            order_id=str(order_id or ""),
            entries=entries,
            role=role,
            order_type=self.order_type,
        ))
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        ctx["os_config"] = {
            "row_sections": _OS_ROW_SECTIONS,
            "tiers": _OS_TIERS,
            "cells_per_tier": _OS_CELLS_PER_TIER,
            "mr_rows": [1, 2, 3, 4],
        }
        ctx["ok"] = kwargs.get("ok", False)
        ctx["error"] = kwargs.get("error")
        return ctx


class ProcessingPlacementActView(PlacementActView):
    order_type = "processing"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin", "manager")

    def _load_entries(self, order_id):
        return list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )

    def _ensure_receiving_act(self, request, order_id, entries):
        if not entries:
            return entries
        latest = entries[-1]
        status_entry = _current_status_entry(entries)
        base_payload = dict((status_entry.payload or latest.payload or {}))
        processed_cards, placed_cards = _processing_card_sets(base_payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        items = _processing_receiving_items(
            base_payload,
            latest.agency_id,
            ready_cards if ready_cards else None,
        )
        if not items:
            return entries
        existing_act = self._receiving_act_entry(entries)
        if existing_act:
            existing_items = (existing_act.payload or {}).get("act_items") or []

            def _normalize_items(source):
                return sorted(
                    [
                        (
                            str(item.get("sku_code") or "").strip(),
                            str(item.get("size") or "").strip(),
                            int(item.get("actual_qty") or 0),
                        )
                        for item in source
                        if isinstance(item, dict)
                    ]
                )

            if _normalize_items(existing_items) == _normalize_items(items):
                return entries
        act_payload = dict(base_payload)
        act_payload.update(
            {
                "act": "receiving",
                "act_label": "Акт приемки (обработка)",
                "act_state": "closed",
                "act_items": items,
            }
        )
        log_order_action(
            "update",
            order_id=order_id,
            order_type=self.order_type,
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency,
            description="Подготовлен состав для размещения после обработки",
            payload=act_payload,
        )
        return self._load_entries(order_id)

    def _can_create(self, entries):
        if not entries:
            return False
        receiving_act = self._receiving_act_entry(entries)
        items = (receiving_act.payload or {}).get("act_items") if receiving_act else []
        latest = entries[-1]
        status_entry = _current_status_entry(entries)
        status_payload = status_entry.payload or latest.payload or {}
        processing_result = WarehouseGoodsStateResolver.resolve_for_processing_order(
            order_id=str(latest.order_id or ""),
            agency=latest.agency,
            payload=status_payload,
        )
        return WarehouseActionPolicy.can_create_processing_placement(
            processing_result,
            has_items=bool(items),
        ).allowed

    def get(self, request, *args, **kwargs):
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        ctx = self.get_context_data(ok=ok, error=error, **kwargs)
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        entries = self._ensure_receiving_act(request, order_id, entries)
        action = (request.POST.get("action") or "close").lower()
        role = get_request_role(request)
        base_url = f"/orders/processing/{order_id}/placement/"
        if role != "storekeeper":
            return redirect(base_url)
        if not self._can_create(entries):
            return redirect(f"{base_url}?error=1")
        status_entry = _current_status_entry(entries)
        latest = entries[-1]
        base_payload = dict((status_entry.payload or latest.payload or {}))
        processed_cards, placed_cards = _processing_card_sets(base_payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        placement_act = self._placement_act_entry(entries)
        if action == "open":
            if not placement_act:
                return redirect(base_url)
            payload = placement_act.payload or {}
            current_state = (payload.get("act_state") or "closed").lower()
            if current_state == "open":
                return redirect(base_url)
            act_payload = dict(base_payload)
            act_payload.update(payload)
            act_payload["act"] = "placement"
            act_payload["act_label"] = act_payload.get("act_label") or "Акт размещения"
            act_payload["act_state"] = "open"
            log_order_action(
                "update",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency,
                description="Открыт акт размещения после обработки",
                payload=act_payload,
            )
            OperationalStockService.clear_order_placement(
                latest.agency,
                self.order_type,
                order_id,
            )
            return redirect(base_url)
        if not ready_cards:
            return redirect(f"{base_url}?error=1")
        act_items = _processing_receiving_items(
            base_payload,
            latest.agency_id,
            ready_cards,
        )
        if not act_items:
            return redirect(f"{base_url}?error=1")
        placement_items = []
        boxes_raw = request.POST.get("boxes_json") or "[]"
        pallets_raw = request.POST.get("pallets_json") or "[]"
        try:
            boxes_data = json.loads(boxes_raw)
            pallets_data = json.loads(pallets_raw)
        except (TypeError, ValueError):
            return redirect(f"{base_url}?error=1")
        totals = {}

        def add_total(item, qty, container_type):
            key = _item_key(item.get("sku"), item.get("name"), item.get("size"))
            if not key:
                return
            entry = totals.setdefault(key, {"box": 0, "pallet": 0, "total": 0})
            entry[container_type] += qty
            entry["total"] += qty

        for box in boxes_data if isinstance(boxes_data, list) else []:
            items = box.get("items") if isinstance(box, dict) else None
            for item in items or []:
                qty = _parse_qty_value(item.get("qty"))
                if qty is None:
                    continue
                add_total(item, qty, "box")
        for pallet in pallets_data if isinstance(pallets_data, list) else []:
            items = pallet.get("items") if isinstance(pallet, dict) else None
            for item in items or []:
                qty = _parse_qty_value(item.get("qty"))
                if qty is None:
                    continue
                add_total(item, qty, "pallet")
        placement_items = []
        for item in act_items:
            key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
            entry = totals.get(key, {"box": 0, "pallet": 0, "total": 0})
            actual_qty = _parse_qty_value(item.get("actual_qty")) or 0
            if entry["total"] > actual_qty:
                return redirect(f"{base_url}?error=1")
            if entry["total"] != actual_qty:
                return redirect(f"{base_url}?error=1")
            placement_items.append(
                {
                    "sku_code": item.get("sku_code"),
                    "name": item.get("name"),
                    "size": item.get("size"),
                    "actual_qty": actual_qty,
                    "box_qty": entry["box"],
                    "pallet_qty": entry["pallet"],
                }
            )
        placed_cards_list = base_payload.get("placed_cards") or []
        if not isinstance(placed_cards_list, list):
            placed_cards_list = []
        for card_id in ready_cards:
            if card_id and card_id not in placed_cards_list:
                placed_cards_list.append(card_id)
        base_payload["placed_cards"] = placed_cards_list
        now = timezone.localtime().isoformat()
        for card in base_payload.get("cards") or []:
            card_id = _processing_card_id(card)
            if card_id and card_id in ready_cards:
                card["placed_at"] = card.get("placed_at") or now
                card["placed_done"] = True
        act_payload = dict(base_payload)
        act_payload["act"] = "placement"
        act_payload["act_label"] = "Акт размещения"
        act_payload["act_state"] = "closed"
        act_payload["act_items"] = placement_items
        act_payload["act_boxes"] = boxes_data if isinstance(boxes_data, list) else []
        act_payload["act_pallets"] = pallets_data if isinstance(pallets_data, list) else []
        log_order_action(
            "update",
            order_id=order_id,
            order_type=self.order_type,
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency,
            description="Создан акт размещения после обработки",
            payload=act_payload,
        )
        OperationalStockService.replace_order_placement(
            latest.agency,
            self.order_type,
            order_id,
            act_payload,
            performed_by=request.user if request.user.is_authenticated else None,
        )
        return redirect(f"{base_url}?ok=1")

    def get_context_data(self, **kwargs):
        ctx = TemplateView.get_context_data(self, **kwargs)
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        entries = self._ensure_receiving_act(self.request, order_id, entries)
        latest = entries[-1] if entries else None
        status_entry = _current_status_entry(entries)
        status_payload = status_entry.payload or {} if status_entry else {}
        base_payload = dict(status_payload or (latest.payload or {}) if latest else {})
        processed_cards, placed_cards = _processing_card_sets(base_payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        receiving_act = self._receiving_act_entry(entries)
        placement_act = self._placement_act_entry(entries)
        receiving_items = (receiving_act.payload or {}).get("act_items") if receiving_act else []
        placement_items = (placement_act.payload or {}).get("act_items") if placement_act else []
        display_items = []
        source_items = placement_items or receiving_items
        for item in source_items:
            display_items.append(
                {
                    "sku_code": item.get("sku_code") or "",
                    "barcode": item.get("barcode") or "",
                    "name": item.get("name") or "-",
                    "size": item.get("size") or "-",
                    "actual_qty": item.get("actual_qty") or 0,
                    "box_qty": item.get("box_qty") or 0,
                    "pallet_qty": item.get("pallet_qty") or 0,
                }
            )
        catalog_items = []
        remaining_items = []
        for item in receiving_items:
            catalog_items.append(
                {
                    "sku_code": item.get("sku_code") or "",
                    "name": item.get("name") or "",
                    "size": item.get("size") or "",
                }
            )
            remaining_items.append(
                {
                    "sku_code": item.get("sku_code") or "",
                    "barcode": item.get("barcode") or "",
                    "name": item.get("name") or "",
                    "size": item.get("size") or "",
                    "actual_qty": item.get("actual_qty") or 0,
                }
            )
        barcode_map = {}
        sku_by_code = {}
        if latest and latest.agency_id and receiving_items:
            sku_codes = {
                (item.get("sku_code") or "").strip()
                for item in receiving_items
                if (item.get("sku_code") or "").strip()
            }
            if sku_codes:
                for sku in (
                    SKU.objects.filter(
                        agency_id=latest.agency_id,
                        sku_code__in=sku_codes,
                        deleted=False,
                    )
                    .prefetch_related("barcodes")
                ):
                    sku_by_code[sku.sku_code] = sku
                    for barcode in sku.barcodes.all():
                        value = (barcode.value or "").strip()
                        if not value:
                            continue
                        barcode_map.setdefault(
                            value,
                            {
                                "sku": sku.sku_code,
                                "name": sku.name,
                                "size": (barcode.size or sku.size or "").strip(),
                            },
                        )
                    sku_code_barcode = (sku.code or "").strip()
                    if sku_code_barcode:
                        barcode_map.setdefault(
                            sku_code_barcode,
                            {
                                "sku": sku.sku_code,
                                "name": sku.name,
                                "size": (sku.size or "").strip(),
                            },
                        )
        for row in display_items:
            if str(row.get("barcode") or "").strip():
                continue
            sku = sku_by_code.get(str(row.get("sku_code") or "").strip())
            barcode_value = _barcode_value_for_sku(sku, row.get("size"))
            if barcode_value and barcode_value != "-":
                row["barcode"] = barcode_value
        for row in remaining_items:
            if str(row.get("barcode") or "").strip():
                continue
            sku = sku_by_code.get(str(row.get("sku_code") or "").strip())
            barcode_value = _barcode_value_for_sku(sku, row.get("size"))
            if barcode_value and barcode_value != "-":
                row["barcode"] = barcode_value
        client_label = "-"
        if latest and latest.agency:
            name = latest.agency.short_name or latest.agency.agn_name or latest.agency.fio_agn or str(latest.agency)
            client_label = _shorten_ip_name(name)
        role = get_request_role(self.request)
        can_submit = role in _RECEIVING_WAREHOUSE_ROLES
        act_state = "open"
        if placement_act:
            act_state = (placement_act.payload or {}).get("act_state") or "closed"
        can_open_act = can_submit and act_state == "closed" and bool(receiving_items)
        boxes_data = (placement_act.payload or {}).get("act_boxes") if placement_act else []
        pallets_data = (placement_act.payload or {}).get("act_pallets") if placement_act else []
        occupied_cells = StockAvailabilityService.occupied_os_cells(
            exclude_order_type=self.order_type,
            exclude_order_id=order_id,
            include_agency=True,
        )
        os_config = {
            "row_sections": _OS_ROW_SECTIONS,
            "tiers": _OS_TIERS,
            "cells_per_tier": _OS_CELLS_PER_TIER,
            "mr_rows": [1, 2, 3, 4],
        }
        cz_only = _payload_has_value(
            base_payload,
            "marking_5840_each_qty",
            "marking_5840_each_needed",
        )
        placement_payload = placement_act.payload or {} if placement_act else {}
        latest_payload = latest.payload or {} if latest else {}
        goods_type = ""
        goods_type_label = ""
        for payload_source in (placement_payload, latest_payload, status_payload):
            if not isinstance(payload_source, dict):
                continue
            goods_type = str(payload_source.get("goods_type") or "").strip().lower() or goods_type
            goods_type_label = str(payload_source.get("goods_type_label") or "").strip() or goods_type_label
            if goods_type and goods_type_label:
                break
        if not goods_type and not goods_type_label:
            goods_type = "gv"
        if not goods_type_label and goods_type:
            goods_type_label = ReceivingWorkflowService.RECEIVING_GOODS_TYPE_LABELS.get(goods_type, goods_type)
        ctx.update(
            {
                "order_id": order_id,
                "client_label": client_label,
                "status_label": _status_label_from_entry(status_entry) if status_entry else "-",
                "cabinet_url": resolve_cabinet_url(get_request_role(self.request)),
                "goods_type": goods_type,
                "goods_type_label": goods_type_label,
                "order_detail_url": f"/orders/processing/{order_id}/",
                "can_submit": can_submit,
                "can_open_act": can_open_act,
                "signed_by_storekeeper": False,
                "act_exists": bool(placement_act),
                "act_state": act_state,
                "items": display_items,
                "catalog_items": catalog_items,
                "remaining_items": remaining_items,
                "barcode_map": barcode_map,
                "boxes_data": boxes_data,
                "pallets_data": pallets_data,
                "os_config": os_config,
                "occupied_cells": occupied_cells,
                "current_agency_id": latest.agency_id if latest else 0,
                "cz_only": cz_only,
                "container_code_url": f"/orders/receiving/{order_id}/flow/container-code/",
                "marking_scan_url": f"/marking/processing/{order_id}/scan/",
                "ready_cards_count": len(ready_cards),
                "ok": kwargs.get("ok", False),
                "error": kwargs.get("error"),
            }
        )
        return ctx






def _other_decimal(value):
    text = str(value or "").strip().replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, TypeError, ValueError):
        return None


def _other_service_rows_from_post(post) -> list[dict[str, str]]:
    names = post.getlist("service_name")
    qtys = post.getlist("service_qty")
    units = post.getlist("service_unit")
    comments = post.getlist("service_comment")
    rows: list[dict[str, str]] = []
    for idx, name in enumerate(names):
        title = str(name or "").strip()
        qty = str(qtys[idx] if idx < len(qtys) else "").strip()
        unit = str(units[idx] if idx < len(units) else "").strip() or "шт."
        comment = str(comments[idx] if idx < len(comments) else "").strip()
        if not title and not qty and not comment:
            continue
        rows.append(
            {
                "name": title or "Услуга без названия",
                "qty": qty,
                "unit": unit,
                "comment": comment,
            }
        )
    return rows


def _other_services_to_text(rows: list[dict[str, str]]) -> str:
    if not rows:
        return ""
    lines = ["Выполненные услуги:"]
    for row in rows:
        qty = row.get("qty") or "—"
        unit = row.get("unit") or ""
        comment = row.get("comment") or "—"
        lines.append(f"- {row.get('name')}: {qty} {unit}; комментарий: {comment}")
    return "\n".join(lines)


def _other_result_from_post(post) -> dict:
    what_done = (post.get("what_done") or post.get("result_comment") or "").strip()
    qty = (post.get("processed_qty") or "").strip()
    boxes = (post.get("processed_boxes") or "").strip()
    pallets = (post.get("processed_pallets") or "").strip()
    discrepancy = (post.get("discrepancy") or "").strip()
    executor_comment = (post.get("executor_comment") or "").strip()
    lines = []
    if what_done:
        lines.append(f"Что было сделано: {what_done}")
    if qty:
        lines.append(f"Обработано единиц: {qty}")
    if boxes:
        lines.append(f"Обработано коробов: {boxes}")
    if pallets:
        lines.append(f"Обработано паллет: {pallets}")
    if discrepancy:
        lines.append(f"Расхождения: {discrepancy}")
    if executor_comment:
        lines.append(f"Комментарий исполнителя: {executor_comment}")
    return {
        "result_outcome": post.get("result_outcome") or "done_full",
        "result_comment": "\n".join(lines),
        "result_qty": _other_decimal(qty),
        "result_unit": (post.get("result_unit") or "шт").strip() or "шт",
        "what_done": what_done,
        "service_rows": _other_service_rows_from_post(post),
    }


class OrdersOtherDetailView(RoleRequiredMixin, TemplateView):
    template_name = "orders/other_detail.html"
    order_type = "other"
    allowed_roles = ("manager", "storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            order_id = kwargs.get("order_id")
            if order_id:
                has_access = OrderAuditEntry.objects.filter(
                    order_id=order_id,
                    order_type=self.order_type,
                    agency=client_agency,
                ).exists()
                if not has_access:
                    return HttpResponseForbidden("Доступ запрещен")
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        from client_cabinet.other_requests import (
            accept_other_by_warehouse,
            cancel_other_request,
            latest_other_entry,
            save_other_attachments,
            send_other_to_warehouse,
            take_other_in_work,
        )
        from client_cabinet.models import OtherRequestAttachment, OtherRequestComment
        from client_cabinet.other_request_workflow import (
            OtherRequestError,
            approve_and_close,
            cancel_request,
            complete_by_executor,
            ensure_other_request_from_audit,
            reject_warehouse_cancel_confirmation,
            request_manager_clarification,
            request_warehouse_cancel_confirmation,
            save_executor_result,
            start_by_executor,
            warehouse_cancel_request_payload,
        )
        from todo.models import Task

        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/other/")
        action = (request.POST.get("action") or "").strip().lower()
        role = get_request_role(request)
        client_agency = getattr(request, "_client_agency", None) or _client_agency_from_request(request)
        client_param = request.GET.get("client")
        suffix = f"?client={client_param}" if client_param else ""
        redirect_url = f"/orders/other/{order_id}/{suffix}"

        if action == "send_to_warehouse" and role in {"manager", "head_manager", "director", "admin"}:
            stub = Task(route=f"/orders/other/{order_id}/")
            send_other_to_warehouse(stub, request)
            return redirect(redirect_url)
        if action == "accept_warehouse" and role == "storekeeper":
            if accept_other_by_warehouse(order_id=order_id, user=request.user):
                messages.success(request, "Заявка принята складом")
            else:
                messages.error(request, "Заявку нельзя принять в текущем статусе")
            return redirect(redirect_url)
        if action == "start_work" and role == "storekeeper":
            req = ensure_other_request_from_audit(order_id)
            if req is not None:
                try:
                    start_by_executor(req, user=request.user)
                    messages.success(request, "Заявка взята в работу")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "take_in_work" and role == "storekeeper":
            if take_other_in_work(order_id=order_id, user=request.user):
                messages.success(request, "Заявка взята в работу")
            else:
                messages.error(request, "Сначала примите заявку складом")
            return redirect(redirect_url)
        if action in {"save_result", "request_clarification"} and role == "storekeeper":
            req = ensure_other_request_from_audit(order_id)
            if req is not None:
                result = _other_result_from_post(request.POST)
                services_text = _other_services_to_text(result["service_rows"])
                try:
                    save_executor_result(
                        req,
                        user=request.user,
                        result_outcome=result["result_outcome"],
                        result_comment=result["result_comment"],
                        result_qty=result["result_qty"],
                        result_unit=result["result_unit"],
                        services_text=services_text,
                    )
                    if action == "request_clarification":
                        reason = (request.POST.get("clarification_text") or request.POST.get("executor_comment") or "").strip()
                        request_manager_clarification(req, user=request.user, reason=reason)
                        OtherRequestComment.objects.create(
                            request=req,
                            author=request.user,
                            visibility=OtherRequestComment.VISIBILITY_INTERNAL,
                            text=f"Запрос уточнения: {reason}",
                        )
                        messages.success(request, "Запрос уточнения отправлен менеджеру")
                    else:
                        messages.success(request, "Результат сохранён")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "complete" and role in {"storekeeper", "manager", "head_manager", "director", "admin"}:
            if role == "storekeeper":
                req = ensure_other_request_from_audit(order_id)
                if req is not None:
                    result = _other_result_from_post(request.POST)
                    services_text = _other_services_to_text(result["service_rows"])
                    missing = []
                    if not result["what_done"]:
                        missing.append("что было сделано")
                    if missing:
                        messages.error(request, "Нельзя передать заявку менеджеру. Не заполнено: " + ", ".join(missing) + ".")
                        return redirect(redirect_url)
                    try:
                        save_executor_result(
                            req,
                            user=request.user,
                            result_outcome=result["result_outcome"],
                            result_comment=result["result_comment"],
                            result_qty=result["result_qty"],
                            result_unit=result["result_unit"],
                            services_text=services_text,
                        )
                        complete_by_executor(
                            req,
                            user=request.user,
                            result_outcome=result["result_outcome"],
                            result_comment=result["result_comment"],
                            result_qty=result["result_qty"],
                            result_unit=result["result_unit"],
                        )
                        messages.success(request, "Результат передан менеджеру на проверку")
                    except OtherRequestError as exc:
                        messages.error(request, exc.message)
                return redirect(redirect_url)
            req = ensure_other_request_from_audit(order_id)
            if req is not None:
                try:
                    approve_and_close(req, user=request.user)
                    messages.success(request, "Заявка закрыта, факты услуг переданы в биллинг")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "add_comment" and role in {"storekeeper", "manager", "head_manager", "director", "admin"}:
            req = ensure_other_request_from_audit(order_id)
            text = (request.POST.get("comment") or "").strip()
            if req is not None and text:
                visibility = (request.POST.get("visibility") or OtherRequestComment.VISIBILITY_INTERNAL).strip()
                OtherRequestComment.objects.create(
                    request=req,
                    author=request.user,
                    visibility=visibility if visibility in {"client", "internal"} else "internal",
                    text=text,
                )
                messages.success(request, "Комментарий добавлен")
            return redirect(redirect_url)
        if action == "upload_attachment" and role in {"storekeeper", "manager", "head_manager", "director", "admin"}:
            req = ensure_other_request_from_audit(order_id)
            files = request.FILES.getlist("files")
            purpose = (request.POST.get("purpose") or OtherRequestAttachment.PURPOSE_RESULT).strip()
            if req is not None and files:
                saved = save_other_attachments(
                    order_id=order_id,
                    agency=req.agency,
                    user=request.user,
                    files=files,
                    purpose=purpose,
                    request_obj=req,
                )
                messages.success(request, f"Загружено файлов: {len(saved)}")
            else:
                messages.error(request, "Выберите файлы для загрузки")
            return redirect(redirect_url)
        if action == "request_warehouse_cancel":
            req = ensure_other_request_from_audit(order_id)
            if req is not None and role in {"manager", "head_manager", "director", "admin"}:
                reason = (request.POST.get("reason") or request.POST.get("cancel_reason") or "").strip()
                try:
                    request_warehouse_cancel_confirmation(req, user=request.user, reason=reason)
                    messages.success(request, "Запрос на отмену отправлен складу")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "approve_warehouse_cancel" and role in {"storekeeper", "processing_head"}:
            req = ensure_other_request_from_audit(order_id)
            if req is not None:
                pending = warehouse_cancel_request_payload(req)
                try:
                    cancel_request(
                        req,
                        user=request.user,
                        reason=str(pending.get("cancel_reason") or "").strip(),
                        warehouse_approved=True,
                    )
                    messages.success(request, "Отмена подтверждена складом")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "reject_warehouse_cancel" and role in {"storekeeper", "processing_head"}:
            req = ensure_other_request_from_audit(order_id)
            if req is not None:
                try:
                    reject_warehouse_cancel_confirmation(req, user=request.user)
                    messages.success(request, "Склад отклонил отмену. Заявка остаётся в работе")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        if action == "cancel":
            if role in {"manager", "head_manager", "director", "admin"} or client_agency:
                reason = (request.POST.get("reason") or request.POST.get("cancel_reason") or "").strip()
                if not reason and not client_agency:
                    messages.error(request, "Укажите причину отмены для клиента.")
                    return redirect(f"{redirect_url}#cancel-reason-required")
                try:
                    cancel_other_request(order_id=order_id, user=request.user, reason=reason)
                    messages.success(request, "Заявка отменена")
                except OtherRequestError as exc:
                    messages.error(request, exc.message)
            return redirect(redirect_url)
        return redirect(redirect_url)

    def get_context_data(self, **kwargs):
        from client_cabinet.models import OtherRequest, OtherRequestAttachment
        from client_cabinet.other_request_workflow import (
            ensure_other_request_from_audit,
            warehouse_cancel_request_payload,
        )
        from client_cabinet.other_requests import STATUS_LABELS, latest_other_entry, list_other_attachments
        from client_cabinet.web_ui import _order_bucket, _order_status_label

        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        latest = latest_other_entry(order_id) if order_id else None
        if not latest:
            raise Http404("Заявка не найдена")
        payload = dict(latest.payload or {})
        status_value = str(payload.get("status") or "").lower()
        workflow_status = str(payload.get("workflow_status") or "").lower()
        item = ensure_other_request_from_audit(order_id)
        if item is not None:
            workflow_status = item.status
            status_label = item.status_label
        else:
            status_label = _order_status_label(latest)
        warehouse_cancel_request = (
            warehouse_cancel_request_payload(item) if item is not None else {}
        )
        if warehouse_cancel_request:
            status_label = "Отмена ожидает подтверждения склада"
        bucket = _order_bucket(latest)
        role = get_request_role(self.request)
        client_agency = getattr(self.request, "_client_agency", None) or _client_agency_from_request(self.request)
        client_view = bool(client_agency)
        history = list(
            OrderAuditEntry.objects.filter(order_type=self.order_type, order_id=order_id)
            .select_related("user")
            .order_by("-created_at")[:40]
        )
        created_entry = (
            OrderAuditEntry.objects.filter(order_type=self.order_type, order_id=order_id)
            .order_by("created_at")
            .first()
        )
        is_terminal = status_value in {"completed", "done", "cancelled"} or workflow_status in {
            "closed",
            "cancelled",
            "rejected",
        }
        warehouse_started = status_value in {"warehouse_accepted", "in_work"} or workflow_status in {
            "warehouse_accepted",
            "in_progress",
            "paused",
            "awaiting_manager_check",
            "done_by_department",
            "rework",
        }
        can_storekeeper_accept = (
            (not client_view)
            and (not is_terminal)
            and role == "storekeeper"
            and workflow_status == "awaiting_department"
        )
        can_storekeeper_take = (not client_view) and (not is_terminal) and role == "storekeeper" and (
            workflow_status in {"warehouse_accepted", "rework"}
        )
        can_storekeeper_complete = (
            (not client_view)
            and (not is_terminal)
            and role == "storekeeper"
            and workflow_status == "in_progress"
        )
        can_manager_confirm = (
            (not client_view)
            and (not is_terminal)
            and role in {"manager", "head_manager", "director", "admin"}
            and workflow_status in {"awaiting_manager_check", "done_by_department"}
        )
        attachments = list_other_attachments(order_id=order_id, agency=latest.agency) if latest.agency else []
        client_files = [a for a in attachments if a["purpose"] == OtherRequestAttachment.PURPOSE_CLIENT]
        internal_files = [a for a in attachments if a["purpose"] == OtherRequestAttachment.PURPOSE_INTERNAL]
        result_files = [a for a in attachments if a["purpose"] == OtherRequestAttachment.PURPOSE_RESULT]
        comments = list(item.comments.select_related("author").order_by("created_at")) if item is not None else []
        links = list(item.links.all()) if item is not None else []
        other_service_facts = []
        if item is not None and item.agency_id:
            from billing.warehouse_services import facts_payload, list_facts

            try:
                other_service_facts = facts_payload(
                    list_facts(
                        client=item.agency,
                        order_type="other",
                        order_id=item.public_number,
                    )
                )
            except ValidationError:
                other_service_facts = []

        ctx.update(
            {
                "order_id": order_id,
                "order_type": self.order_type,
                "display_number": format_order_number(self.order_type, order_id),
                "agency": latest.agency,
                "item": item,
                "payload": payload,
                "status_value": status_value,
                "workflow_status": workflow_status,
                "status_label": status_label,
                "bucket": bucket,
                "client_view": client_view,
                "role": role,
                "history": history,
                "created_at": created_entry.created_at if created_entry else latest.created_at,
                "is_terminal": is_terminal,
                "cabinet_url": resolve_cabinet_url(role),
                "can_send_to_warehouse": (not client_view)
                and (not is_terminal)
                and role in {"manager", "head_manager", "director", "admin"}
                and status_value in {"submitted", "draft"},
                "can_accept_warehouse": can_storekeeper_accept,
                "can_take_in_work": (not client_view)
                and can_storekeeper_take,
                "can_complete": can_storekeeper_complete or can_manager_confirm,
                "can_save_result": can_storekeeper_complete,
                "can_request_clarification": can_storekeeper_complete,
                "complete_label": "Подтвердить выполнение" if can_storekeeper_complete else "Подтвердить и закрыть",
                "can_cancel": (not is_terminal)
                and (not warehouse_started)
                and (
                    client_view
                    or role in {"manager", "head_manager", "director", "admin"}
                ),
                "warehouse_cancel_request_pending": bool(warehouse_cancel_request),
                "warehouse_cancel_reason": str(
                    warehouse_cancel_request.get("cancel_reason") or ""
                ).strip(),
                "can_request_warehouse_cancel": bool(
                    not client_view
                    and not is_terminal
                    and warehouse_started
                    and role in {"manager", "head_manager", "director", "admin"}
                    and not warehouse_cancel_request
                ),
                "can_review_warehouse_cancel": bool(
                    not client_view
                    and role in {"storekeeper", "processing_head"}
                    and warehouse_cancel_request
                ),
                "status_choices": STATUS_LABELS,
                "attachments": attachments,
                "client_files": client_files,
                "internal_files": internal_files,
                "result_files": result_files,
                "comments": comments,
                "links": links,
                "other_service_facts": other_service_facts,
                "wsfm_order_type": "other",
                "wsfm_order_id": item.public_number if item is not None else order_id,
                "wsfm_client_id": item.agency_id if item is not None else latest.agency_id,
                "wsfm_source": "other_manual",
                "wsfm_title": "Фактически оказанные услуги",
                "wsfm_cancel_label": "Вернуться к заявке",
                "wsfm_items_label": "Обработано товара",
            }
        )
        return ctx
